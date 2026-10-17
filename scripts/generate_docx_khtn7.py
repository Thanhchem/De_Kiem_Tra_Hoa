# -*- coding: utf-8 -*-
"""
Script sinh file Word (.docx) chuẩn FORM ĐỀ KIỂM TRA cho Đề Cương Ôn Tập Giữa Kì I KHTN 7
Năm học 2026 - 2027 (Thầy TRẦN VĂN THẠNH)
Xuất bản 3 phiên bản:
1. Bản Học sinh: San_Pham/De_Cuong_KHTN_7_Giua_HK1_HocSinh.docx
2. Bản Giáo viên: San_Pham/De_Cuong_KHTN_7_Giua_HK1_GiaoVien.docx
3. Bản Dạy trên màn hình: San_Pham/De_Cuong_KHTN_7_Giua_HK1_DayTrenManHinh.docx
"""

import os
import sys
import re
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.path.append(os.path.dirname(__file__))
from khtn7_data import (
    VAT_LY_LY_THUYET, VAT_LY_VAN_DUNG,
    HOA_HOC_LY_THUYET, HOA_HOC_VAN_DUNG,
    SINH_HOC_TRAC_NGHIEM, SINH_HOC_TU_LUAN
)

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
os.makedirs(OUT_DIR, exist_ok=True)

# Bảng màu chuẩn quy định
COLOR_HF_RED = RGBColor(211, 47, 47)        # Coral Red (#D32F2F) cho Header & Footer
COLOR_HF_GREEN_HEX = "8BC390"               # Sage Green (#8BC390) cho viền Header & Footer
COLOR_PRIMARY = RGBColor(46, 125, 50)       # Xanh lá đậm (#2E7D32) cho Tiêu đề Phần (I, II, III)
COLOR_SUBSEC = RGBColor(0, 105, 92)         # Xanh lam ngọc (#00695C) cho mục A, B
COLOR_CORRECT = RGBColor(198, 40, 40)       # Đỏ (#C62828) cho đáp án đúng Bản Giáo Viên
COLOR_SOLUTION = RGBColor(21, 101, 192)     # Xanh lam (#1565C0) cho tiêu đề Lời giải
COLOR_SCREEN_TITLE = RGBColor(106, 27, 154) # Tím đậm (#6A1B9A) cho tiêu đề bản Dạy Màn Hình
COLOR_GRAY = RGBColor(97, 97, 97)
COLOR_BORDER_GREY_HEX = "CCCCCC"

def add_p_border_bottom(p, color="8BC390", sz="6"):
    """Thêm đường kẻ viền dưới cho đoạn văn bản."""
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_p_border_top(p, color="8BC390", sz="6"):
    """Thêm đường kẻ viền trên cho đoạn văn bản."""
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def set_run_style(run, font_name="Times New Roman", font_size=10.5, bold=False, italic=False, subscript=False, superscript=False, color=None):
    """Cài đặt style chi tiết cho từng run text."""
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    """Chèn trường số trang tự động PAGE vào chân trang."""
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def render_rich_text(p, text, base_font="Times New Roman", base_size=10.5, base_bold=False, base_italic=False, base_color=None):
    """
    Sinh văn bản định dạng phong phú từ ký hiệu markdown:
    - *bold*
    - _italic_
    - ~subscript~ (CO~2~, H~2~O)
    - ^superscript^ (10^-24^)
    """
    tokens = re.split(r'(~[^~]+~|\^[^\^]+\^|\*[^\*]+\*|_[^_]+_)', text)
    for token in tokens:
        if not token:
            continue
        is_bold = base_bold
        is_italic = base_italic
        is_sub = False
        is_sup = False
        content = token

        if token.startswith('*') and token.endswith('*'):
            is_bold = True
            content = token[1:-1]
        elif token.startswith('_') and token.endswith('_'):
            is_italic = True
            content = token[1:-1]
        elif token.startswith('~') and token.endswith('~'):
            is_sub = True
            content = token[1:-1]
        elif token.startswith('^') and token.endswith('^'):
            is_sup = True
            content = token[1:-1]

        run = p.add_run(content)
        run_color = COLOR_CORRECT if (is_bold and not base_bold and not base_color) else base_color
        set_run_style(run, font_name=base_font, font_size=base_size, bold=is_bold, italic=is_italic,
                      subscript=is_sub, superscript=is_sup, color=run_color)

def setup_header_footer(doc):
    """Cài đặt Header & Footer chuẩn theo FORM_DE_KIEM_TRA."""
    normal_style = doc.styles['Normal']
    normal_style.paragraph_format.tab_stops.clear_all()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    section.start_type = WD_SECTION_START.NEW_PAGE
    section.different_first_page_header_footer = False
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    # 1. Header (Bảng 1x2, rộng 18.0 cm)
    header = section.header
    for p in header.paragraphs:
        p.text = ""
    tbl_h = header.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl_h.cell(0, 0)
    c_right = tbl_h.cell(0, 1)
    c_left.width = Cm(8.0)
    c_right.width = Cm(10.0)

    p_hl = c_left.paragraphs[0]
    p_hl.paragraph_format.space_before = Pt(0)
    p_hl.paragraph_format.space_after = Pt(2)
    p_hl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_hl = p_hl.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)

    p_hr = c_right.paragraphs[0]
    p_hr.paragraph_format.space_before = Pt(0)
    p_hr.paragraph_format.space_after = Pt(2)
    p_hr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p_hr.add_run("Thầy ")
    set_run_style(r1, font_size=10, color=COLOR_HF_RED)
    r2 = p_hr.add_run("TRẦN VĂN THẠNH")
    set_run_style(r2, font_size=10, bold=True, color=COLOR_HF_RED)
    r3 = p_hr.add_run(" - 0777.470.803")
    set_run_style(r3, font_size=10, color=COLOR_HF_RED)

    add_p_border_bottom(p_hl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_bottom(p_hr, color=COLOR_HF_GREEN_HEX, sz="6")

    # 2. Footer (Bảng 1x2, rộng 18.0 cm)
    footer = section.footer
    for p in footer.paragraphs:
        p.text = ""
    tbl_f = footer.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_fl = tbl_f.cell(0, 0)
    c_fr = tbl_f.cell(0, 1)
    c_fl.width = Cm(14.5)
    c_fr.width = Cm(3.5)

    p_fl = c_fl.paragraphs[0]
    p_fl.paragraph_format.space_before = Pt(2)
    p_fl.paragraph_format.space_after = Pt(0)
    p_fl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_fl = p_fl.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)

    p_fr = c_fr.paragraphs[0]
    p_fr.paragraph_format.space_before = Pt(2)
    p_fr.paragraph_format.space_after = Pt(0)
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_field(p_fr, font_size=10, bold=True, color=COLOR_HF_RED)

    add_p_border_top(p_fl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_top(p_fr, color=COLOR_HF_GREEN_HEX, sz="6")

def build_title_block(doc, is_teacher=False, is_screen=False, base_size=10.5):
    """Xây dựng Title Block không dùng bảng, dùng Tab Stop phải ở 18.0 cm."""
    headers_data = [
        ("SỞ GIÁO DỤC VÀ ĐÀO TẠO THỪA THIÊN HUẾ", "ĐỀ CƯƠNG ÔN TẬP KIỂM TRA GIỮA KÌ I", 10, True, 11, True),
        ("TRƯỜNG THCS NGUYỄN THỊ MINH KHAI", "MÔN: KHOA HỌC TỰ NHIÊN 7", 10.5, True, 11.5, True),
        ("(Đề cương ôn tập toàn diện)", "NĂM HỌC 2026 - 2027", 9.5, False, 10, True)
    ]
    for l_text, r_text, l_size, l_bold, r_size, r_bold in headers_data:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)

        rl = p.add_run(l_text)
        set_run_style(rl, font_size=l_size, bold=l_bold, italic=(not l_bold))

        p.add_run('\t')
        rr = p.add_run(r_text)
        set_run_style(rr, font_size=r_size, bold=r_bold, italic=(not r_bold))

    if is_teacher:
        p_gv = doc.add_paragraph()
        p_gv.paragraph_format.space_before = Pt(3)
        p_gv.paragraph_format.space_after = Pt(3)
        p_gv.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rgv = p_gv.add_run("(BẢN GIÁO VIÊN -- CÓ ĐÁP ÁN & LỜI GIẢI CHI TIẾT)")
        set_run_style(rgv, font_size=11.5, bold=True, color=COLOR_CORRECT)
    elif is_screen:
        p_sc = doc.add_paragraph()
        p_sc.paragraph_format.space_before = Pt(3)
        p_sc.paragraph_format.space_after = Pt(3)
        p_sc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rsc = p_sc.add_run("(BẢN TRÌNH CHIẾU / GIẢNG DẠY MÀN HÌNH -- CÓ CHỪA CHỖ TRỐNG)")
        set_run_style(rsc, font_size=11.5, bold=True, color=COLOR_SCREEN_TITLE)

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(4)
    p_info.paragraph_format.space_after = Pt(3)
    render_rich_text(p_info, "*Họ, tên học sinh:* ............................................................ *Lớp:* ............. *Trường:* ....................................................", base_size=base_size)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(4)
    add_p_border_bottom(p_div, color=COLOR_BORDER_GREY_HEX, sz="4")

def add_major_section_header(doc, title, base_size=12):
    """Tiêu đề phần lớn: I. KHTN (VẬT LÝ), II. KHTN (HÓA HỌC), III. KHTN (SINH HỌC)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    render_rich_text(p, f"*{title}*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)

def add_subsection_header(doc, title, base_size=11):
    """Tiêu đề mục nhỏ: A. Lý thuyết, B. Vận dụng."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, f"*{title}*", base_size=base_size, base_bold=True, base_color=COLOR_SUBSEC)

def render_question(doc, q_text, base_size=10.5, line_spacing=1.15):
    """Trình bày nội dung câu hỏi."""
    lines = q_text.strip().split("\n")
    for idx, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = line_spacing
        p.paragraph_format.space_before = Pt(3 if idx == 0 else 1)
        p.paragraph_format.space_after = Pt(1)
        
        # In đậm tiền tố Câu X. hoặc a., b., c., d.
        m_cau = re.match(r'^(Câu\s+\d+\.)\s*(.*)$', line.strip())
        m_sub = re.match(r'^([a-d]\.)\s*(.*)$', line.strip())
        m_dash = re.match(r'^(-)\s*(.*)$', line.strip())
        
        if m_cau:
            lbl = m_cau.group(1)
            rest = m_cau.group(2)
            r_lbl = p.add_run(lbl + " ")
            set_run_style(r_lbl, font_size=base_size, bold=True)
            render_rich_text(p, rest, base_size=base_size, base_bold=False)
        elif m_sub:
            p.paragraph_format.left_indent = Cm(0.5)
            lbl = m_sub.group(1)
            rest = m_sub.group(2)
            r_lbl = p.add_run(lbl + " ")
            set_run_style(r_lbl, font_size=base_size, bold=True)
            render_rich_text(p, rest, base_size=base_size, base_bold=False)
        elif m_dash:
            p.paragraph_format.left_indent = Cm(0.5)
            r_lbl = p.add_run("• ")
            set_run_style(r_lbl, font_size=base_size, bold=True)
            render_rich_text(p, m_dash.group(2), base_size=base_size, base_bold=False)
        else:
            if idx > 0:
                p.paragraph_format.left_indent = Cm(0.5)
            render_rich_text(p, line, base_size=base_size, base_bold=False)

def add_solution(doc, solution_text, solution_image=None, base_size=10):
    """Trình bày Lời giải chi tiết theo chuẩn không dùng bảng."""
    p_head = doc.add_paragraph()
    p_head.paragraph_format.left_indent = Cm(0.4)
    p_head.paragraph_format.space_before = Pt(4)
    p_head.paragraph_format.space_after = Pt(2)
    render_rich_text(p_head, "► *Lời giải chi tiết:*", base_size=base_size + 0.5, base_bold=True, base_color=COLOR_SOLUTION)

    lines = solution_text.strip().split("\n")
    for line in lines:
        if line.startswith("*Lời giải chi tiết:*") or line.startswith("►"):
            continue
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        render_rich_text(p, line, base_size=base_size)

    if solution_image and os.path.exists(solution_image):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        run_img = p_img.add_run()
        run_img.add_picture(solution_image, width=Cm(6.8))

def add_writing_lines(doc, num_lines=4, dot_count=100, is_screen=False):
    """Thêm dòng chấm cho học sinh làm bài hoặc giáo viên ghi chú."""
    color = RGBColor(160, 160, 160) if is_screen else RGBColor(190, 190, 190)
    for _ in range(num_lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("." * dot_count)
        set_run_style(r, font_size=9.5, color=color)

def add_screen_drawing_box(doc, caption="KHUNG GHI CHÉP / VẼ HÌNH TRÊN MÀN HÌNH", height_cm=6.0):
    """Tạo khung viền nét đứt thoáng đãng để giáo viên vẽ / ghi chú trực tiếp trên màn hình."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Cm(17.5)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        r'<w:tcBorders %s>'
        r'  <w:top w:val="dashed" w:sz="6" w:space="0" w:color="9E9E9E"/>'
        r'  <w:left w:val="dashed" w:sz="6" w:space="0" w:color="9E9E9E"/>'
        r'  <w:bottom w:val="dashed" w:sz="6" w:space="0" w:color="9E9E9E"/>'
        r'  <w:right w:val="dashed" w:sz="6" w:space="0" w:color="9E9E9E"/>'
        r'</w:tcBorders>' % nsdecls('w')
    )
    tcPr.append(tcBorders)
    
    shd = parse_xml(r'<w:shd %s w:fill="FAFAFA"/>' % nsdecls('w'))
    tcPr.append(shd)
    
    # Đoạn văn bản chú thích bên trong khung
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"[{caption}]")
    set_run_style(r, font_size=9, italic=True, color=RGBColor(158, 158, 158))
    
    # Thêm khoảng trống
    for _ in range(int(height_cm * 1.5)):
        p_sp = cell.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(2)
        p_sp.paragraph_format.space_after = Pt(2)
        r_sp = p_sp.add_run("")
    
    # Khoảng cách sau bảng
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def add_styled_data_table(doc, headers, rows):
    """Tạo bảng số liệu khoa học (như Câu 6 Vật lý) thẩm mỹ, viền thanh mảnh, nền tiêu đề xám nhạt."""
    num_cols = len(headers)
    tbl = doc.add_table(rows=1 + len(rows), cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    col_w_0 = Cm(3.6)
    col_w_rest = Cm((17.5 - 3.6) / (num_cols - 1))

    # Header Row
    hdr_cells = tbl.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].width = col_w_0 if i == 0 else col_w_rest
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(h_text)
        set_run_style(r, font_size=10, bold=True)
        
        tcPr = hdr_cells[i]._tc.get_or_add_tcPr()
        shd = parse_xml(r'<w:shd %s w:fill="E8ECEF"/>' % nsdecls('w'))
        tcPr.append(shd)
        tcBorders = parse_xml(
            r'<w:tcBorders %s>'
            r'  <w:top w:val="single" w:sz="6" w:space="0" w:color="8BC390"/>'
            r'  <w:left w:val="single" w:sz="4" w:space="0" w:color="B0BEC5"/>'
            r'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="8BC390"/>'
            r'  <w:right w:val="single" w:sz="4" w:space="0" w:color="B0BEC5"/>'
            r'</w:tcBorders>' % nsdecls('w')
        )
        tcPr.append(tcBorders)

    # Data Rows
    for r_idx, row_data in enumerate(rows):
        row_cells = tbl.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].width = col_w_0 if c_idx == 0 else col_w_rest
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(val)
            set_run_style(r, font_size=10, bold=(c_idx == 0))
            
            tcPr = row_cells[c_idx]._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                r'<w:tcBorders %s>'
                r'  <w:top w:val="single" w:sz="4" w:space="0" w:color="CFD8DC"/>'
                r'  <w:left w:val="single" w:sz="4" w:space="0" w:color="B0BEC5"/>'
                r'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="8BC390"/>'
                r'  <w:right w:val="single" w:sz="4" w:space="0" w:color="B0BEC5"/>'
                r'</w:tcBorders>' % nsdecls('w')
            )
            tcPr.append(tcBorders)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(3)

def build_document(doc_type="student"):
    """
    doc_type:
    - 'student': Bản Học sinh
    - 'teacher': Bản Giáo viên (có đầy đủ đáp án, lời giải chi tiết, hình minh họa)
    - 'screen': Bản Dạy trên màn hình (cỡ chữ thoáng, chừa chỗ trống rộng rãi và khung vẽ)
    """
    doc = docx.Document()
    setup_header_footer(doc)
    
    is_teacher = (doc_type == "teacher")
    is_screen = (doc_type == "screen")
    
    base_font_size = 11.5 if is_screen else 10.5
    line_spacing = 1.25 if is_screen else 1.15
    
    # Title Block
    build_title_block(doc, is_teacher=is_teacher, is_screen=is_screen, base_size=base_font_size)

    # =========================================================================
    # PHẦN I. KHTN (VẬT LÝ)
    # =========================================================================
    add_major_section_header(doc, "I. KHTN (VẬT LÝ)", base_size=12.5 if is_screen else 11.5)
    
    # A. Lý thuyết
    add_subsection_header(doc, "A. Lý thuyết:", base_size=11.5 if is_screen else 11)
    for item in VAT_LY_LY_THUYET:
        render_question(doc, item["question"], base_size=base_font_size, line_spacing=line_spacing)
        if is_teacher:
            add_solution(doc, item["solution"], base_size=10.0)
        elif is_screen:
            add_writing_lines(doc, num_lines=item.get("blank_lines_screen", 8), is_screen=True)
        else:
            add_writing_lines(doc, num_lines=item.get("blank_lines_student", 4))

    # B. Vận dụng
    add_subsection_header(doc, "B. Vận dụng:", base_size=11.5 if is_screen else 11)
    for item in VAT_LY_VAN_DUNG:
        render_question(doc, item["question"], base_size=base_font_size, line_spacing=line_spacing)
        
        # Hình ảnh câu hỏi nếu có (như Câu 5)
        if "image_path" in item and os.path.exists(item["image_path"]):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(3)
            p_img.paragraph_format.space_after = Pt(2)
            run_img = p_img.add_run()
            run_img.add_picture(item["image_path"], width=Cm(7.2))
            if "image_caption" in item:
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(0)
                p_cap.paragraph_format.space_after = Pt(3)
                r_cap = p_cap.add_run(f"_{item['image_caption']}_")
                set_run_style(r_cap, font_size=9.5, italic=True, color=COLOR_GRAY)

        # Bảng dữ liệu nếu có (như Câu 6)
        if "table_data" in item:
            t_data = item["table_data"]
            add_styled_data_table(doc, t_data["headers"], t_data["rows"])

        # Lời giải hoặc Chỗ trống làm bài
        if is_teacher:
            add_solution(doc, item["solution"], solution_image=item.get("solution_image"), base_size=10.0)
        elif is_screen:
            if item.get("screen_box", False):
                add_screen_drawing_box(doc, caption="KHUNG VẼ ĐỒ THỊ TRỰC TIẾP TRÊN MÀN HÌNH", height_cm=6.5)
            add_writing_lines(doc, num_lines=item.get("blank_lines_screen", 8), is_screen=True)
        else:
            add_writing_lines(doc, num_lines=item.get("blank_lines_student", 5))

    # =========================================================================
    # PHẦN II. KHTN (HÓA HỌC)
    # =========================================================================
    add_major_section_header(doc, "II. KHTN (HÓA HỌC)", base_size=12.5 if is_screen else 11.5)
    
    # A. Lý thuyết
    add_subsection_header(doc, "A. Lý thuyết:", base_size=11.5 if is_screen else 11)
    for item in HOA_HOC_LY_THUYET:
        render_question(doc, item["question"], base_size=base_font_size, line_spacing=line_spacing)
        if is_teacher:
            add_solution(doc, item["solution"], base_size=10.0)
        elif is_screen:
            add_writing_lines(doc, num_lines=item.get("blank_lines_screen", 10), is_screen=True)
        else:
            add_writing_lines(doc, num_lines=item.get("blank_lines_student", 5))

    # B. Một số bài tập vận dụng
    add_subsection_header(doc, "B. Một số bài tập vận dụng:", base_size=11.5 if is_screen else 11)
    for item in HOA_HOC_VAN_DUNG:
        render_question(doc, item["question"], base_size=base_font_size, line_spacing=line_spacing)
        if is_teacher:
            add_solution(doc, item["solution"], solution_image=item.get("solution_image"), base_size=10.0)
        elif is_screen:
            if item.get("screen_box", False):
                caption = "KHUNG VẼ MÔ HÌNH NGUYÊN TỬ TRÊN MÀN HÌNH"
                add_screen_drawing_box(doc, caption=caption, height_cm=6.0)
            add_writing_lines(doc, num_lines=item.get("blank_lines_screen", 8), is_screen=True)
        else:
            add_writing_lines(doc, num_lines=item.get("blank_lines_student", 5))

    # =========================================================================
    # PHẦN III. KHTN (SINH HỌC)
    # =========================================================================
    add_major_section_header(doc, "III. KHTN (SINH HỌC)", base_size=12.5 if is_screen else 11.5)
    
    # A. Trắc nghiệm / Trọng tâm lý thuyết
    add_subsection_header(doc, "A. Trắc nghiệm & Câu hỏi ôn tập trọng tâm:", base_size=11.5 if is_screen else 11)
    for item in SINH_HOC_TRAC_NGHIEM:
        render_question(doc, item["question"], base_size=base_font_size, line_spacing=line_spacing)
        if is_teacher:
            add_solution(doc, item["solution"], base_size=10.0)
        elif is_screen:
            add_writing_lines(doc, num_lines=item.get("blank_lines_screen", 8), is_screen=True)
        else:
            add_writing_lines(doc, num_lines=item.get("blank_lines_student", 4))

    # B. Tự luận
    add_subsection_header(doc, "B. Tự luận:", base_size=11.5 if is_screen else 11)
    for item in SINH_HOC_TU_LUAN:
        render_question(doc, item["question"], base_size=base_font_size, line_spacing=line_spacing)
        if is_teacher:
            add_solution(doc, item["solution"], base_size=10.0)
        elif is_screen:
            add_writing_lines(doc, num_lines=item.get("blank_lines_screen", 12), is_screen=True)
        else:
            add_writing_lines(doc, num_lines=item.get("blank_lines_student", 8))

    return doc

def main():
    files_to_gen = [
        ("student", os.path.join(OUT_DIR, "De_Cuong_KHTN_7_Giua_HK1_HocSinh.docx")),
        ("teacher", os.path.join(OUT_DIR, "De_Cuong_KHTN_7_Giua_HK1_GiaoVien.docx")),
        ("screen",  os.path.join(OUT_DIR, "De_Cuong_KHTN_7_Giua_HK1_DayTrenManHinh.docx"))
    ]

    for doc_type, out_path in files_to_gen:
        print(f"Generating [{doc_type}]: {out_path} ...")
        doc = build_document(doc_type=doc_type)
        doc.save(out_path)
        print(f"-> Successfully saved: {out_path} ({os.path.getsize(out_path)} bytes)")

    print("\nAll 3 Word documents generated successfully!")

if __name__ == "__main__":
    main()
