# -*- coding: utf-8 -*-
"""
Script tạo file Word (.docx) chuẩn FORM ĐỀ KIỂM TRA cho Đề Cương Ôn Tập KHTN 8
Trường THCS Nguyễn Thị Minh Khai - Năm học 2026 - 2027
Tạo 2 phiên bản:
1. Bản Học sinh: San_Pham/De_Cuong_KHTN_8_Giua_HK1_Minh_Khai.docx
2. Bản Giáo viên: San_Pham/De_Cuong_KHTN_8_Giua_HK1_Minh_Khai_GiaoVien.docx
"""

import os
import sys
import re
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.path.append(os.path.dirname(__file__))
sys.stdout.reconfigure(encoding='utf-8')
from khtn8_data_part1 import ALL_PART1
from khtn8_data_part2_3_4 import PART2_INFO, PART3_INFO, PART4_ESSAY

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
os.makedirs(OUT_DIR, exist_ok=True)

# Bảng màu chuẩn quy định
COLOR_HF_RED = RGBColor(211, 47, 47)        # Coral Red (#D32F2F) cho Header & Footer
COLOR_HF_GREEN_HEX = "8BC390"               # Sage Green (#8BC390) cho viền Header & Footer
COLOR_PRIMARY = RGBColor(46, 125, 50)       # Xanh lá đậm (#2E7D32) cho Tiêu đề Phần
COLOR_CORRECT = RGBColor(198, 40, 40)       # Đỏ (#C62828) cho đáp án đúng Bản Giáo Viên
COLOR_SOLUTION = RGBColor(21, 101, 192)     # Xanh lam (#1565C0) cho tiêu đề Lời giải
COLOR_GRAY = RGBColor(97, 97, 97)
COLOR_BORDER_GREY_HEX = "CCCCCC"

def add_p_border_bottom(p, color="8BC390", sz="6"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_p_border_top(p, color="8BC390", sz="6"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def set_run_style(run, font_name="Times New Roman", font_size=10.5, bold=False, italic=False, subscript=False, superscript=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def render_rich_text(p, text, base_font="Times New Roman", base_size=10.5, base_bold=False, base_italic=False, base_color=None):
    tokens = re.split(r'(~[^~]+~|\^[^\^]+\^|\*[^\*]+\*|_[^_]+_)', text)
    for token in tokens:
        if not token:
            continue
        is_bold = base_bold
        is_italic = base_italic
        is_sub = False
        is_sup = False
        content = token

        if token.startswith('*') and token.endswith('*'):
            is_bold = True
            content = token[1:-1]
        elif token.startswith('_') and token.endswith('_'):
            is_italic = True
            content = token[1:-1]
        elif token.startswith('~') and token.endswith('~'):
            is_sub = True
            content = token[1:-1]
        elif token.startswith('^') and token.endswith('^'):
            is_sup = True
            content = token[1:-1]

        run = p.add_run(content)
        run_color = COLOR_CORRECT if (is_bold and not base_bold and not base_color) else base_color
        set_run_style(run, font_name=base_font, font_size=base_size, bold=is_bold, italic=is_italic,
                      subscript=is_sub, superscript=is_sup, color=run_color)

def setup_header_footer(doc):
    normal_style = doc.styles['Normal']
    normal_style.paragraph_format.tab_stops.clear_all()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    section.start_type = WD_SECTION_START.NEW_PAGE
    section.different_first_page_header_footer = False
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    # 1. Header (Bảng 1x2, rộng 18.0 cm)
    header = section.header
    for p in header.paragraphs:
        p.text = ""
    tbl_h = header.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl_h.cell(0, 0)
    c_right = tbl_h.cell(0, 1)
    c_left.width = Cm(8.0)
    c_right.width = Cm(10.0)

    p_hl = c_left.paragraphs[0]
    p_hl.paragraph_format.space_before = Pt(0)
    p_hl.paragraph_format.space_after = Pt(2)
    p_hl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_hl = p_hl.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)

    p_hr = c_right.paragraphs[0]
    p_hr.paragraph_format.space_before = Pt(0)
    p_hr.paragraph_format.space_after = Pt(2)
    p_hr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p_hr.add_run("Thầy ")
    set_run_style(r1, font_size=10, color=COLOR_HF_RED)
    r2 = p_hr.add_run("TRẦN VĂN THẠNH")
    set_run_style(r2, font_size=10, bold=True, color=COLOR_HF_RED)
    r3 = p_hr.add_run(" - 0777.470.803")
    set_run_style(r3, font_size=10, color=COLOR_HF_RED)

    add_p_border_bottom(p_hl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_bottom(p_hr, color=COLOR_HF_GREEN_HEX, sz="6")

    # 2. Footer (Bảng 1x2, rộng 18.0 cm)
    footer = section.footer
    for p in footer.paragraphs:
        p.text = ""
    tbl_f = footer.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_fl = tbl_f.cell(0, 0)
    c_fr = tbl_f.cell(0, 1)
    c_fl.width = Cm(14.5)
    c_fr.width = Cm(3.5)

    p_fl = c_fl.paragraphs[0]
    p_fl.paragraph_format.space_before = Pt(2)
    p_fl.paragraph_format.space_after = Pt(0)
    p_fl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_fl = p_fl.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)

    p_fr = c_fr.paragraphs[0]
    p_fr.paragraph_format.space_before = Pt(2)
    p_fr.paragraph_format.space_after = Pt(0)
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_field(p_fr, font_size=10, bold=True, color=COLOR_HF_RED)

    add_p_border_top(p_fl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_top(p_fr, color=COLOR_HF_GREEN_HEX, sz="6")

def build_title_block(doc, is_teacher=False):
    headers_data = [
        ("UBND PHƯỜNG THUỶ XUÂN", "ĐỀ CƯƠNG ÔN TẬP KIỂM TRA GIỮA HỌC KÌ I", 10, True, 11, True),
        ("TRƯỜNG THCS NGUYỄN THỊ MINH KHAI", "MÔN: KHTN 8 - NĂM HỌC 2026 - 2027", 10.5, True, 11, True),
        ("(Đề cương có 12 trang)", "Thời gian ôn tập: Học kì I", 9.5, False, 10, True)
    ]
    for l_text, r_text, l_size, l_bold, r_size, r_bold in headers_data:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)

        rl = p.add_run(l_text)
        set_run_style(rl, font_size=l_size, bold=l_bold, italic=(not l_bold))

        p.add_run('\t')
        rr = p.add_run(r_text)
        set_run_style(rr, font_size=r_size, bold=r_bold, italic=(not r_bold))

    if is_teacher:
        p_gv = doc.add_paragraph()
        p_gv.paragraph_format.space_before = Pt(3)
        p_gv.paragraph_format.space_after = Pt(3)
        p_gv.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rgv = p_gv.add_run("(BẢN GIÁO VIÊN -- CÓ ĐÁP ÁN & LỜI GIẢI CHI TIẾT)")
        set_run_style(rgv, font_size=11, bold=True, color=COLOR_CORRECT)

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(4)
    p_info.paragraph_format.space_after = Pt(3)
    render_rich_text(p_info, "*Họ, tên học sinh:* ............................................................ *Lớp:* ............. *Số báo danh:* .............", base_size=10.5)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(4)
    add_p_border_bottom(p_div, color=COLOR_BORDER_GREY_HEX, sz="4")

def add_section_header(doc, title, subtitle=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, f"*{title}*", base_size=11.5, base_bold=True, base_color=COLOR_PRIMARY)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(4)
        render_rich_text(p2, f"_{subtitle}_", base_size=10.5, base_italic=True)

def render_question_prompt(p, text, base_size=10.5):
    m = re.match(r'^(Câu\s+\d+\.)\s*(.*)$', text.strip(), re.DOTALL)
    if m:
        lbl = m.group(1)
        rest = m.group(2)
        r_lbl = p.add_run(lbl + " ")
        set_run_style(r_lbl, font_size=base_size, bold=True)
        render_rich_text(p, rest, base_size=base_size, base_bold=False)
    else:
        render_rich_text(p, text, base_size=base_size, base_bold=False)

def render_single_option(p, opt_str, is_teacher=False, base_size=10.5):
    clean_str = opt_str.strip()
    is_correct = False
    if clean_str.startswith('*') and clean_str.endswith('*'):
        is_correct = True
        clean_str = clean_str[1:-1].strip()

    m = re.match(r'^([A-D]\.\s*)(.*)$', clean_str)
    if m:
        lbl = m.group(1)
        body = m.group(2)
    else:
        lbl = ""
        body = clean_str

    if is_teacher and is_correct:
        if lbl:
            r_lbl = p.add_run(lbl)
            set_run_style(r_lbl, font_size=base_size, bold=True, color=COLOR_CORRECT)
        render_rich_text(p, body, base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
    else:
        if lbl:
            r_lbl = p.add_run(lbl)
            set_run_style(r_lbl, font_size=base_size, bold=True, color=None)
        render_rich_text(p, body, base_size=base_size, base_bold=False, base_color=None)

def render_options_no_table(doc, opts, is_teacher=False):
    clean_opts = [re.sub(r'[*_~^]', '', opt) for opt in opts]
    max_len = max(len(opt) for opt in clean_opts) if clean_opts else 0

    if max_len <= 18 and len(opts) == 4:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(4.5), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(13.5), WD_TAB_ALIGNMENT.LEFT)

        render_single_option(p, opts[0], is_teacher)
        p.add_run('\t')
        render_single_option(p, opts[1], is_teacher)
        p.add_run('\t')
        render_single_option(p, opts[2], is_teacher)
        p.add_run('\t')
        render_single_option(p, opts[3], is_teacher)

    elif max_len <= 38 and len(opts) == 4:
        p1 = doc.add_paragraph()
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.tab_stops.clear_all()
        p1.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
        render_single_option(p1, opts[0], is_teacher)
        p1.add_run('\t')
        render_single_option(p1, opts[1], is_teacher)

        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(1)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.tab_stops.clear_all()
        p2.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
        render_single_option(p2, opts[2], is_teacher)
        p2.add_run('\t')
        render_single_option(p2, opts[3], is_teacher)

    else:
        for opt in opts:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.left_indent = Cm(0.5)
            render_single_option(p, opt, is_teacher)

def render_solution_block(doc, solution_text, font_size=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(0.6)

    r_sol = p.add_run("► Lời giải chi tiết:\n")
    set_run_style(r_sol, font_size=10.5, bold=True, color=COLOR_SOLUTION)

    lines = solution_text.strip().split('\n')
    for i, line in enumerate(lines):
        render_rich_text(p, line, base_size=font_size, base_color=COLOR_GRAY)
        if i < len(lines) - 1:
            p.add_run('\n')

def add_writing_lines(doc, num_lines=4):
    for _ in range(num_lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run("..................................................................................................................................................................................")
        set_run_style(r, font_size=9, color=COLOR_GRAY)

def generate_exam_docx(is_teacher=False, output_filename="test.docx"):
    doc = docx.Document()
    setup_header_footer(doc)
    build_title_block(doc, is_teacher=is_teacher)

    # Tiêu đề A
    p_a = doc.add_paragraph()
    p_a.paragraph_format.space_before = Pt(8)
    p_a.paragraph_format.space_after = Pt(2)
    render_rich_text(p_a, "*A. PHẦN TRẮC NGHIỆM KHÁCH QUAN*", base_size=12, base_bold=True, base_color=COLOR_PRIMARY)

    # Phần I
    add_section_header(doc, "Phần I. Câu trắc nghiệm nhiều lựa chọn", "Mỗi câu hỏi, học sinh chọn một phương án đúng nhất.")

    for q in ALL_PART1:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, q["question"], base_size=10.5)

        # Options
        render_options_no_table(doc, q["options"], is_teacher=is_teacher)

        # Solution for teacher
        if is_teacher and "solution" in q:
            render_solution_block(doc, f"Đáp án đúng: {q['correct']}. {q['solution']}")

    # Phần II
    add_section_header(doc, PART2_INFO["title"], PART2_INFO["subtitle"])

    # Yêu cầu cần đạt
    for r in PART2_INFO["req"]:
        p_req = doc.add_paragraph()
        p_req.paragraph_format.space_before = Pt(1)
        p_req.paragraph_format.space_after = Pt(1)
        p_req.paragraph_format.left_indent = Cm(0.4)
        render_rich_text(p_req, f"_{r}_", base_size=10, base_italic=True, base_color=COLOR_GRAY)

    for q in PART2_INFO["questions"]:
        p_ctx = doc.add_paragraph()
        p_ctx.paragraph_format.space_before = Pt(5)
        p_ctx.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_ctx, q["context"], base_size=10.5)

        for item_text, is_true, expl in q["items"]:
            p_it = doc.add_paragraph()
            p_it.paragraph_format.space_before = Pt(1)
            p_it.paragraph_format.space_after = Pt(1)
            p_it.paragraph_format.left_indent = Cm(0.6)

            render_rich_text(p_it, item_text, base_size=10.5)

            if is_teacher:
                status_str = " [ĐÚNG]" if is_true else " [SAI]"
                r_st = p_it.add_run(status_str)
                set_run_style(r_st, font_size=10.5, bold=True, color=COLOR_CORRECT)
                if expl:
                    p_exp = doc.add_paragraph()
                    p_exp.paragraph_format.space_before = Pt(0)
                    p_exp.paragraph_format.space_after = Pt(2)
                    p_exp.paragraph_format.left_indent = Cm(1.0)
                    render_rich_text(p_exp, f"→ _{expl}_", base_size=10, base_italic=True, base_color=COLOR_GRAY)
            else:
                r_box = p_it.add_run("   (Đúng / Sai)")
                set_run_style(r_box, font_size=9.5, italic=True, color=COLOR_GRAY)

    # Phần III
    add_section_header(doc, PART3_INFO["title"], PART3_INFO["subtitle"])

    for q in PART3_INFO["questions"]:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, q["question"], base_size=10.5)

        # Nếu có bảng ở câu 8
        if q.get("has_table") and "table_data" in q:
            # Chèn ảnh bình tam giác bóng bay
            img_flask = r"c:\Antigravity_Thanh\San_Pham\images_khtn8\khtn8_flask_balloon.png"
            if os.path.exists(img_flask):
                p_img = doc.add_paragraph()
                p_img.paragraph_format.space_before = Pt(2)
                p_img.paragraph_format.space_after = Pt(3)
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.add_run().add_picture(img_flask, width=Cm(12.5))

            t_data = q["table_data"]
            tbl = doc.add_table(rows=len(t_data), cols=2)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit = False

            # Format table
            for row_idx, row in enumerate(tbl.rows):
                col0_val, col1_val = t_data[row_idx]
                c0, c1 = row.cells[0], row.cells[1]
                c0.width = Cm(13.0)
                c1.width = Cm(5.0)

                p0 = c0.paragraphs[0]
                p0.paragraph_format.space_before = Pt(2)
                p0.paragraph_format.space_after = Pt(2)
                is_hdr = (row_idx == 0)
                render_rich_text(p0, f"*{col0_val}*" if is_hdr else col0_val, base_size=10, base_bold=is_hdr)

                p1 = c1.paragraphs[0]
                p1.paragraph_format.space_before = Pt(2)
                p1.paragraph_format.space_after = Pt(2)
                render_rich_text(p1, f"*{col1_val}*" if is_hdr else col1_val, base_size=10, base_bold=is_hdr)

                # Add borders
                for c in [c0, c1]:
                    tcPr = c._element.get_or_add_tcPr()
                    borders = parse_xml(
                        f'<w:tcBdr {nsdecls("w")}>'
                        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
                        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
                        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
                        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
                        f'</w:tcBdr>'
                    )
                    tcPr.append(borders)

        # Ô ghi đáp án
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(2)
        p_ans.paragraph_format.space_after = Pt(2)
        p_ans.paragraph_format.left_indent = Cm(0.5)

        r_ans_lbl = p_ans.add_run("☞ Đáp án: ")
        set_run_style(r_ans_lbl, font_size=10.5, bold=True)

        if is_teacher:
            r_ans_val = p_ans.add_run(f"[  {q['ans']}  ]")
            set_run_style(r_ans_val, font_size=10.5, bold=True, color=COLOR_CORRECT)
            render_solution_block(doc, q["solution"])
        else:
            r_ans_box = p_ans.add_run("[ . . . . ]")
            set_run_style(r_ans_box, font_size=10.5, bold=True, color=COLOR_GRAY)

    # Phần B: TỰ LUẬN
    p_b = doc.add_paragraph()
    p_b.paragraph_format.space_before = Pt(8)
    p_b.paragraph_format.space_after = Pt(4)
    render_rich_text(p_b, "*B. TỰ LUẬN*", base_size=12, base_bold=True, base_color=COLOR_PRIMARY)

    for q in PART4_ESSAY:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(5)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, q["question"], base_size=10.5)

        # Nếu có ảnh ở câu 11
        if q.get("has_image") and "image_path" in q:
            img_p = os.path.join(r"c:\Antigravity_Thanh", q["image_path"])
            if os.path.exists(img_p):
                p_img = doc.add_paragraph()
                p_img.paragraph_format.space_before = Pt(3)
                p_img.paragraph_format.space_after = Pt(3)
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.add_run().add_picture(img_p, width=Cm(14.0))

        if is_teacher:
            render_solution_block(doc, q["solution"], font_size=10)
        else:
            # Bản học sinh: chèn dòng kẻ chấm làm bài
            add_writing_lines(doc, num_lines=5)

    # Chúc các em thi tốt
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(12)
    p_end.paragraph_format.space_after = Pt(6)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("CHÚC CÁC EM THI TỐT!!!")
    set_run_style(r_end, font_size=12, bold=True, color=COLOR_PRIMARY)

    safe_save(doc, output_filename)

def safe_save(doc, target_path):
    try:
        doc.save(target_path)
        print(f"Saved: {target_path}")
    except PermissionError:
        base, ext = os.path.splitext(target_path)
        fallback_path = f"{base}_CapNhat{ext}"
        doc.save(fallback_path)
        print(f"File đang mở trong Word! Đã lưu bản cập nhật vào: {fallback_path}")

if __name__ == "__main__":
    fn_hs = os.path.join(OUT_DIR, "De_Cuong_KHTN_8_Giua_HK1_Minh_Khai.docx")
    fn_gv = os.path.join(OUT_DIR, "De_Cuong_KHTN_8_Giua_HK1_Minh_Khai_GiaoVien.docx")

    print("Dang tao Ban Hoc Sinh...")
    generate_exam_docx(is_teacher=False, output_filename=fn_hs)

    print("Dang tao Ban Giao Vien...")
    generate_exam_docx(is_teacher=True, output_filename=fn_gv)

    print("Hoan tat 100% ca 2 ban Word!")
