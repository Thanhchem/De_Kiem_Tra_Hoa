# -*- coding: utf-8 -*-
"""
Script tạo file Word (.docx) chuẩn "FORM ĐỀ KIỂM TRA" cho:
ĐỀ KIỂM TRA SỐ 7 - CHƯƠNG 3: HỢP CHẤT CHỨA NITROGEN (HÓA HỌC 12)
Gồm 2 phiên bản:
1. De_Kiem_Tra_Hoa_12_De07.docx (Bản Học Sinh)
2. De_Kiem_Tra_Hoa_12_De07_GiaoVien.docx (Bản Giáo Viên có lời giải chi tiết)
Tuân thủ 100% quy chuẩn: Header & Footer, Tab Stops (NO TABLE trong đề/lời giải), Sub/Superscript công thức hóa học.
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
IMG_DIR = r"c:\Antigravity_Thanh\San_Pham\images_crop_de07"

COLOR_HF_RED = RGBColor(211, 47, 47)      # Coral Red (#D32F2F)
COLOR_HF_GREEN_HEX = "8BC390"             # Sage Green (#8BC390)
COLOR_PRIMARY = RGBColor(46, 125, 50)     # Dark Green (#2E7D32)
COLOR_CORRECT = RGBColor(198, 40, 40)     # Red for correct answers (#C62828)
COLOR_SOLUTION = RGBColor(21, 101, 192)   # Blue for solutions (#1565C0)
COLOR_GRAY = RGBColor(97, 97, 97)

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def set_run_style(run, font_name="Times New Roman", font_size=12, bold=False, italic=False, subscript=False, superscript=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def add_chem_run(p, text, base_font="Times New Roman", base_size=12, bold=False, italic=False, color=None):
    tokens = re.split(r'(\^\{[^}]+\}|\^[0-9\+\-]+|~\{[^}]+\}|~[0-9]+)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('^'):
            content = tok[2:-1] if tok.startswith('^{') else tok[1:]
            r = p.add_run(content)
            set_run_style(r, font_name=base_font, font_size=base_size, bold=bold, italic=italic, superscript=True, color=color)
        elif tok.startswith('~'):
            content = tok[2:-1] if tok.startswith('~{') else tok[1:]
            r = p.add_run(content)
            set_run_style(r, font_name=base_font, font_size=base_size, bold=bold, italic=italic, subscript=True, color=color)
        else:
            r = p.add_run(tok)
            set_run_style(r, font_name=base_font, font_size=base_size, bold=bold, italic=italic, color=color)

def setup_page(doc):
    for sec in doc.sections:
        sec.top_margin = Cm(1.5)
        sec.bottom_margin = Cm(1.5)
        sec.left_margin = Cm(1.5)
        sec.right_margin = Cm(1.5)
        sec.header_distance = Cm(0.6)
        sec.footer_distance = Cm(0.6)
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)

def set_cell_border_bottom(cell, color="8BC390", sz="6"):
    tcPr = cell._tc.get_or_add_tcPr()
    xml = f'<w:tcBorders {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/></w:tcBorders>'
    tcPr.append(parse_xml(xml))

def set_cell_border_top(cell, color="8BC390", sz="6"):
    tcPr = cell._tc.get_or_add_tcPr()
    xml = f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/></w:tcBorders>'
    tcPr.append(parse_xml(xml))

def set_cell_margins_zero(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    xml = f'<w:tcMar {nsdecls("w")}><w:top w:w="0" w:type="dxa"/><w:bottom w:w="0" w:type="dxa"/><w:left w:w="0" w:type="dxa"/><w:right w:w="0" w:type="dxa"/></w:tcMar>'
    tcPr.append(parse_xml(xml))

def setup_header_footer(doc):
    sec = doc.sections[0]
    
    # HEADER
    hdr = sec.header
    htable = hdr.add_table(rows=1, cols=2, width=Cm(18.0))
    htable.autofit = False
    htable.columns[0].width = Cm(9.0)
    htable.columns[1].width = Cm(9.0)
    
    cell_l, cell_r = htable.rows[0].cells[0], htable.rows[0].cells[1]
    set_cell_margins_zero(cell_l)
    set_cell_margins_zero(cell_r)
    set_cell_border_bottom(cell_l, color=COLOR_HF_GREEN_HEX, sz="6")
    set_cell_border_bottom(cell_r, color=COLOR_HF_GREEN_HEX, sz="6")
    
    p_hl = cell_l.paragraphs[0]
    p_hl.paragraph_format.space_before = Pt(0)
    p_hl.paragraph_format.space_after = Pt(2)
    p_hl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_hl = p_hl.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)
    
    p_hr = cell_r.paragraphs[0]
    p_hr.paragraph_format.space_before = Pt(0)
    p_hr.paragraph_format.space_after = Pt(2)
    p_hr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_hr1 = p_hr.add_run("Thầy ")
    set_run_style(r_hr1, font_size=10, color=COLOR_HF_RED)
    r_hr2 = p_hr.add_run("TRẦN VĂN THẠNH")
    set_run_style(r_hr2, font_size=10, bold=True, color=COLOR_HF_RED)
    r_hr3 = p_hr.add_run(" - 0777.470.803")
    set_run_style(r_hr3, font_size=10, color=COLOR_HF_RED)
    
    # FOOTER
    ftr = sec.footer
    ftable = ftr.add_table(rows=1, cols=2, width=Cm(18.0))
    ftable.autofit = False
    ftable.columns[0].width = Cm(14.0)
    ftable.columns[1].width = Cm(4.0)
    
    cell_fl, cell_fr = ftable.rows[0].cells[0], ftable.rows[0].cells[1]
    set_cell_margins_zero(cell_fl)
    set_cell_margins_zero(cell_fr)
    set_cell_border_top(cell_fl, color=COLOR_HF_GREEN_HEX, sz="6")
    set_cell_border_top(cell_fr, color=COLOR_HF_GREEN_HEX, sz="6")
    
    p_fl = cell_fl.paragraphs[0]
    p_fl.paragraph_format.space_before = Pt(2)
    p_fl.paragraph_format.space_after = Pt(0)
    p_fl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_fl = p_fl.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)
    
    p_fr = cell_fr.paragraphs[0]
    p_fr.paragraph_format.space_before = Pt(2)
    p_fr.paragraph_format.space_after = Pt(0)
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_field(p_fr, font_size=10, bold=True, color=COLOR_HF_RED)

def add_header_block(doc, is_teacher=False):
    lines = [
        ("SỞ GIÁO DỤC VÀ ĐÀO TẠO THỪA THIÊN HUẾ", "ĐỀ KIỂM TRA ĐỊNH KỲ CHƯƠNG 3 (HỢP CHẤT CHỨA NITROGEN)", True, True, 10.5, 11),
        ("ĐỀ CHÍNH THỨC", "MÔN: HOÁ HỌC – LỚP 12", True, True, 11, 11),
        ("(Đề thi có 04 trang)", "Thời gian làm bài: 50 phút (không kể thời gian phát đề)", False, True, 10, 10.5),
        ("", "Mã đề thi: 107", False, True, 10, 11)
    ]
    
    for left_text, right_text, l_bold, r_bold, l_size, r_size in lines:
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.05
        
        if left_text:
            r_l = p.add_run(left_text)
            set_run_style(r_l, font_size=l_size, bold=l_bold, italic=("Đề thi có" in left_text))
        
        p.add_run("\t")
        
        if right_text:
            r_r = p.add_run(right_text)
            set_run_style(r_r, font_size=r_size, bold=r_bold, italic=("Thời gian" in right_text), color=COLOR_PRIMARY if "CHƯƠNG 3" in right_text else None)

    if is_teacher:
        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_t.paragraph_format.space_before = Pt(2)
        p_t.paragraph_format.space_after = Pt(2)
        r_t = p_t.add_run("(BẢN GIÁO VIÊN — CÓ LỜI GIẢI CHI TIẾT)")
        set_run_style(r_t, font_size=11, bold=True, color=COLOR_CORRECT)

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(4)
    p_info.paragraph_format.space_after = Pt(2)
    r_info = p_info.add_run("Họ, tên thí sinh: .......................................................................... Lớp: ................ Số báo danh: ...............")
    set_run_style(r_info, font_size=10.5, italic=True)
    
    p_mass = doc.add_paragraph()
    p_mass.paragraph_format.space_before = Pt(0)
    p_mass.paragraph_format.space_after = Pt(6)
    r_mass = p_mass.add_run("(Cho biết nguyên tử khối: H = 1; C = 12; N = 14; O = 16; Na = 23; S = 32; Cl = 35,5; Fe = 56; Cu = 64).")
    set_run_style(r_mass, font_size=10, italic=True)

def add_part_heading(doc, part_title, part_desc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.15
    r_title = p.add_run(part_title)
    set_run_style(r_title, font_size=11.5, bold=True, color=COLOR_PRIMARY)
    
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_before = Pt(0)
    p_desc.paragraph_format.space_after = Pt(4)
    r_desc = p_desc.add_run(part_desc)
    set_run_style(r_desc, font_size=10.5, italic=True)

def add_mcq_options(doc, opt_a, opt_b, opt_c, opt_d, correct_idx=None, is_teacher=False):
    opts = [('A', opt_a), ('B', opt_b), ('C', opt_c), ('D', opt_d)]
    max_len = max(len(text) for _, text in opts)
    
    if max_len < 22:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.tab_stops.add_tab_stop(Cm(4.5), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(13.5), WD_TAB_ALIGNMENT.LEFT)
        
        for idx, (label, text) in enumerate(opts):
            if idx > 0:
                p.add_run("\t")
            is_corr = (is_teacher and idx == correct_idx)
            r_lbl = p.add_run(f"{label}. ")
            set_run_style(r_lbl, font_size=11, bold=True, color=COLOR_CORRECT if is_corr else None)
            add_chem_run(p, text, base_size=11, bold=is_corr, color=COLOR_CORRECT if is_corr else None)
    elif max_len < 46:
        for row in [(0, 1), (2, 3)]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
            
            for col_idx, opt_num in enumerate(row):
                if col_idx > 0:
                    p.add_run("\t")
                label, text = opts[opt_num]
                is_corr = (is_teacher and opt_num == correct_idx)
                r_lbl = p.add_run(f"{label}. ")
                set_run_style(r_lbl, font_size=11, bold=True, color=COLOR_CORRECT if is_corr else None)
                add_chem_run(p, text, base_size=11, bold=is_corr, color=COLOR_CORRECT if is_corr else None)
    else:
        for idx, (label, text) in enumerate(opts):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Cm(0.5)
            
            is_corr = (is_teacher and idx == correct_idx)
            r_lbl = p.add_run(f"{label}. ")
            set_run_style(r_lbl, font_size=11, bold=True, color=COLOR_CORRECT if is_corr else None)
            add_chem_run(p, text, base_size=11, bold=is_corr, color=COLOR_CORRECT if is_corr else None)

def add_solution(doc, sol_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.6)
    
    r_hdr = p.add_run("Lời giải chi tiết: ")
    set_run_style(r_hdr, font_size=10.5, bold=True, color=COLOR_SOLUTION)
    
    lines = sol_text.strip().split("\n")
    for i, line in enumerate(lines):
        if i > 0:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Cm(0.6)
        
        parts = re.split(r'(\*\*[^*]+\*\*)', line)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                r_b = p.add_run(part[2:-2])
                set_run_style(r_b, font_size=10.5, bold=True, color=COLOR_CORRECT if "Chọn" in part or "Đáp số" in part or "Đáp án" in part else None)
            else:
                add_chem_run(p, part, base_size=10.5, italic=False)

def build_docx_de07(is_teacher=False):
    doc = docx.Document()
    setup_page(doc)
    setup_header_footer(doc)
    add_header_block(doc, is_teacher=is_teacher)
    
    # ========================== PHẦN I ==========================
    add_part_heading(
        doc,
        "PHẦN I (4,5 điểm – 18 câu): Câu trắc nghiệm nhiều phương án lựa chọn.",
        "Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án."
    )
    
    part1_data = [
        (
            1,
            "Chất nào sau đây là amine thơm?",
            "Aniline.", "Cyclohexylamine.", "Alanine.", "Trimethylamine.",
            0,
            "Aniline (C~6~H~5~NH~2~) có nhóm amino (-NH~2~) liên kết trực tiếp với vòng benzene nên là một amine thơm.\n**Chọn A.**"
        ),
        (
            2,
            "Aniline có công thức phân tử là:",
            "C~3~H~7~O~2~N", "C~2~H~5~O~2~N", "C~7~H~9~N", "C~6~H~7~N",
            3,
            "Aniline có công thức cấu tạo C~6~H~5~NH~2~ => Công thức phân tử là C~6~H~7~N.\n**Chọn D.**"
        ),
        (
            3,
            "Phân tử amine có cấu tạo tương tự phân tử nào sau đây nhất?",
            "H~2~O", "NH~3~", "HF", "N~2~",
            1,
            "Amine có thể coi là dẫn xuất của ammonia (NH~3~) khi thay thế một hay nhiều nguyên tử hydrogen bằng gốc hydrocarbon. Trong phân tử amine, nguyên tử nitrogen còn 1 cặp electron chưa liên kết tương tự NH~3~.\n**Chọn B.**"
        ),
        (
            4,
            "Số nhóm amino (-NH~2~) có trong một phân tử Lysine là",
            "1", "3", "4", "2",
            3,
            "Lysine (acid 2,6-diaminohexanoic) có công thức: H~2~N-[CH~2~]~4~-CH(NH~2~)-COOH. Phân tử có 2 nhóm amino (-NH~2~) và 1 nhóm carboxyl (-COOH).\n**Chọn D.**"
        ),
        (
            5,
            "(SBT – CTST). Ở điều kiện thường, trạng thái tồn tại của amino acid là",
            "thể lỏng.", "thể khí.", "thể rắn.", "thể rắn và lỏng.",
            2,
            "Ở điều kiện thường, do tồn tại dưới dạng ion lưỡng cực có lực hút tĩnh điện liên phân tử rất mạnh, các amino acid đều là những chất rắn kết tinh, không màu, nhiệt độ nóng chảy cao.\n**Chọn C.**"
        ),
        (
            6,
            "Khi đun nóng, các ε-amino acid hoặc ω-amino acid có thể phản ứng với nhau để tạo thành polymer, đồng thời tách ra các phân tử nước gọi là phản ứng",
            "ester hóa.", "trùng ngưng.", "trùng hợp.", "hydrogen hóa.",
            1,
            "Phản ứng giữa các phân tử amino acid tạo polymer kèm theo sự giải phóng phân tử nhỏ (H~2~O) là phản ứng trùng ngưng (polycondensation).\n**Chọn B.**"
        ),
        (
            7,
            "Peptide: H~2~NCH~2~CONHCH(CH~3~)CONHCH~2~COOH được biểu diễn dạng kí hiệu như sau",
            "Gly – Gly – Ala.", "Gly – Ala – Gly.", "Ala – Gly – Gly.", "Ala – Gly – Ala.",
            1,
            "Các mắt xích lần lượt là:\n- H~2~N-CH~2~-CO- (Gly)\n- -NH-CH(CH~3~)-CO- (Ala)\n- -NH-CH~2~-COOH (Gly)\n=> Kí hiệu là Gly-Ala-Gly.\n**Chọn B.**"
        ),
        (
            8,
            "Trong phân tử tetrapeptide Gly-Val-Glu-Ala, amino acid đầu C là",
            "Gly.", "Val.", "Glu.", "Ala.",
            3,
            "Trong peptide viết theo quy ước: amino acid đầu N ở bên trái (Gly) và amino acid đầu C ở bên phải (Ala).\n**Chọn D.**"
        ),
        (
            9,
            "Insulin, một loại hormone thuộc loại protein, được sản sinh bởi tuyến tuỵ, có chức năng điều hoà quá trình chuyển hoá glucose trong cơ thể. Insulin thúc đẩy sự hấp thu glucose của các tế bào và dự trữ glucose dư thừa trong gan và cơ. Insulin là",
            "protein đơn giản.", "protein phức tạp.", "dipeptide.", "monomer.",
            0,
            "Insulin là một protein đơn giản gồm 51 amino acid được tạo thành từ 2 chuỗi polypeptide liên kết với nhau bằng cầu disulfide, chỉ được cấu tạo từ các gốc α-amino acid.\n**Chọn A.**"
        ),
        (
            10,
            "(SBT – KNTT). Trong cơ thể, enzyme có chức năng nào sau đây?",
            "Cấu trúc tế bào.", "Chất điện giải.", "Chất dự trữ năng lượng.", "Xúc tác sinh học.",
            3,
            "Hầu hết các enzyme là protein có chức năng làm chất xúc tác sinh học, làm tăng tốc độ các phản ứng hóa sinh trong cơ thể với tính đặc hiệu rất cao.\n**Chọn D.**"
        ),
        (
            11,
            "(SBT – CTST). Amine nào sau đây phản ứng được với nitrous acid tạo thành muối diazonium bền ở nhiệt độ thấp?",
            "Methanamine.", "Methanediamine.", "Benzenamine.", "Phenylmethanamine.",
            2,
            "Amine bậc 1 thơm như benzenamine (aniline, C~6~H~5~NH~2~) phản ứng với nitrous acid (HNO~2~) ở nhiệt độ thấp (0 - 5 °C) tạo muối diazonium bền vững: C~6~H~5~NH~2~ + HNO~2~ + HCl -> [C~6~H~5~N~2~]^+^Cl^-^ + 2H~2~O.\n**Chọn C.**"
        ),
        (
            12,
            "(SBT – CTST). Dạng ion chủ yếu nào của amino acid có trong môi trường acid mạnh (pH thấp)?",
            "cau12_img", "", "", "",
            3,
            "Trong môi trường acid mạnh (pH thấp, dư H^+^), nhóm carboxyl -COO^-^ nhận proton trở thành -COOH, và nhóm amino nhận proton thành -NH~3~^+^. Do đó amino acid tồn tại chủ yếu ở dạng cation: H~3~N^+^-CH(CH~3~)-COOH.\n**Chọn D.**"
        ),
        (
            13,
            "(SBT – CĐ). Một tripeptide X được cấu thành từ 2 phân tử Ala và 1 phân tử Gly. Công thức cấu tạo của X không thể là",
            "Ala-Ala-Gly.", "Ala-Gly-Ala.", "Gly-Ala-Ala.", "Gly-Ala-Gly.",
            3,
            "Tripeptide cấu tạo từ 2 Ala và 1 Gly chỉ có thể có 3 đồng phân là: Ala-Ala-Gly, Ala-Gly-Ala, Gly-Ala-Ala. Công thức Gly-Ala-Gly gồm 2 Gly và 1 Ala nên không phù hợp.\n**Chọn D.**"
        ),
        (
            14,
            "(SBT – CĐ). Cho peptide X có công thức cấu tạo như hình bên dưới:\nKhi thuỷ phân hoàn toàn X trong môi trường NaOH thu được sản phẩm hữu cơ có công thức là",
            "H~2~NCH(CH~3~)COOH và H~2~NCH~2~COOH.",
            "H~2~NCH(CH~3~)COONa và H~2~NCH~2~COONa.",
            "H~2~NCH(CH~3~)COONa và H~2~NCH~2~COOH.",
            "H~2~NCH(CH~3~)COOH và H~2~NCH~2~COONa.",
            1,
            "Peptide X là Ala-Gly (H~2~N-CH(CH~3~)-CO-NH-CH~2~-COOH). Trong môi trường kiềm NaOH dư, cả liên kết peptide và nhóm -COOH đều phản ứng tạo muối sodium của các amino acid tương ứng: H~2~NCH(CH~3~)COONa và H~2~NCH~2~COONa.\n**Chọn B.**"
        ),
        (
            15,
            "(SBT – CĐ). Leucine có công thức cấu tạo HOOCCH(NH~2~)CH~2~CH(CH~3~)~2~, là amino acid có khả năng điều hòa sự tổng hợp protein của cơ. Tên theo danh pháp thay thế của leucine là",
            "2-aminoisohexanoic acid.", "2-amino-4-methylpentanoic acid.", "4-amino-2-methylpentanoic acid.", "2-amino-isohexanoic acid.",
            1,
            "Đánh số mạch carbon chính dài nhất từ carbon nhóm carboxyl (-COOH là số 1):\nC1: -COOH, C2: -CH(NH~2~)-, C3: -CH~2~-, C4: -CH(CH~3~)-, C5: -CH~3~.\n=> Tên thay thế: 2-amino-4-methylpentanoic acid.\n**Chọn B.**"
        ),
        (
            16,
            "(SBT – CTST). Dung dịch của chất nào sau đây có môi trường base?",
            "cau16_img", "", "", "",
            0,
            "Chất A là lysine ở dạng ion lưỡng cực: ^+^H~3~N-[CH~2~]~4~-CH(NH~2~)-COO^-^, có 2 nhóm amino (-NH~2~) và 1 nhóm carboxyl (-COOH) nên dung dịch có môi trường base (pH > 7, làm quỳ tím hóa xanh).\n**Chọn A.**"
        ),
        (
            17,
            "(SBT – CTST). Nguyên nhân chủ yếu dẫn đến tính tan của amino acid trong nước là do",
            "phân tử phân cực mạnh nên dễ tan trong nước.",
            "cấu tạo lưỡng cực của phân tử, có tính kị nước nên ít tan trong nước.",
            "năng lượng liên kết của phân tử lớn, khó phá vỡ nên ít tan trong nước.",
            "hình thành liên kết hydrogen với nước nên dễ tan trong nước.",
            0,
            "Trong tinh thể và trong dung dịch, amino acid tồn tại chủ yếu dưới dạng ion lưỡng cực, phân tử có sự phân li điện tích mạnh (cực mạnh). Các ion này tương tác tĩnh điện mạnh với các phân tử nước phân cực (solvat hóa), giúp chúng dễ dàng hòa tan trong nước.\n**Chọn A.**"
        ),
        (
            18,
            "Cho các phát biểu sau:\n(a) Phenylamine có lực base yếu hơn ammonia.\n(b) Phenylamine tan ít trong nước nhưng tan tốt trong dung dịch HCl.\n(c) Khi nấu canh cua, hiện tượng riêu cua nổi lên trên là do sự đông tụ protein.\n(d) Trong tơ tằm có các gốc α-amino acid.\n(e) Dùng giấm ăn hoặc chanh khử được mùi tanh trong cá do amine gây ra.\nSố phát biểu đúng là",
            "5.", "2.", "3.", "4.",
            0,
            "(a) Đúng, vòng benzene hút electron làm giảm mật độ điện tích trên nguyên tử N => lực base phenylamine < NH~3~.\n(b) Đúng, phenylamine tác dụng với HCl tạo muối tan C~6~H~5~NH~3~Cl.\n(c) Đúng, protein trong thịt cua bị đông tụ bởi nhiệt độ khi đun sôi, nổi lên thành riêu cua.\n(d) Đúng, tơ tằm cấu tạo từ protein thiên nhiên (fibroin), chứa các gốc α-amino acid.\n(e) Đúng, acid acetic trong giấm và acid citric trong chanh phản ứng trung hòa các amine gây tanh (như trimethylamine), chuyển thành muối không bay hơi.\n=> Cả 5 phát biểu đều đúng.\n**Chọn A.**"
        )
    ]
    
    for q in part1_data:
        num, text, oa, ob, oc, od, corr, sol = q
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        
        r_num = p.add_run(f"Câu {num}. ")
        set_run_style(r_num, font_size=11, bold=True)
        add_chem_run(p, text, base_size=11)
        
        if num == 12:
            # Câu 12 options image
            p_img = doc.add_paragraph()
            p_img.paragraph_format.space_before = Pt(2)
            p_img.paragraph_format.space_after = Pt(3)
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.add_run().add_picture(os.path.join(IMG_DIR, "cau12_options.png"), width=Cm(12.5))
        elif num == 14:
            # Câu 14 dipeptide structure
            p_img = doc.add_paragraph()
            p_img.paragraph_format.space_before = Pt(2)
            p_img.paragraph_format.space_after = Pt(3)
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.add_run().add_picture(os.path.join(IMG_DIR, "cau14_dipeptide.png"), width=Cm(5.0))
            add_mcq_options(doc, oa, ob, oc, od, correct_idx=corr, is_teacher=is_teacher)
        elif num == 16:
            # Câu 16 options image
            p_img = doc.add_paragraph()
            p_img.paragraph_format.space_before = Pt(2)
            p_img.paragraph_format.space_after = Pt(3)
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.add_run().add_picture(os.path.join(IMG_DIR, "cau16_options.png"), width=Cm(13.0))
        else:
            add_mcq_options(doc, oa, ob, oc, od, correct_idx=corr, is_teacher=is_teacher)
            
        if is_teacher and sol:
            add_solution(doc, sol)
            
    # ========================== PHẦN II ==========================
    add_part_heading(
        doc,
        "PHẦN II (4,0 điểm – 4 câu): Câu trắc nghiệm đúng sai.",
        "Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai."
    )
    
    part2_data = [
        (
            1,
            "(SBT – KNTT). Các phát biểu về điều chế và ứng dụng của amine:",
            [
                ("a", "Một số amine có thể được điều chế bằng cách alkyl hoá ammonia.", True,
                 "Đúng, alkyl hoá ammonia bằng dẫn xuất halogen (RX + NH~3~ -> RNH~2~ + HX) là một phương pháp điều chế amine."),
                ("b", "Một số amine có thể được điều chế bằng cách khử hợp chất nitro.", True,
                 "Đúng, khử hợp chất nitro (ví dụ khử nitrobenzene bằng Fe/HCl hoặc H~2~/Ni) thu được amine thơm."),
                ("c", "Amine được sử dụng để tổng hợp một số loại dược phẩm.", True,
                 "Đúng, nhiều amine và dẫn xuất amine là nguyên liệu tổng hợp thuốc kháng sinh, gây tê, giảm đau..."),
                ("d", "Amine được sử dụng để tổng hợp một số loại polymer.", True,
                 "Đúng, ví dụ hexamethylenediamine phản ứng với adipic acid để tổng hợp nylon-6,6.")
            ],
            None
        ),
        (
            2,
            "(SBT – CTST). Cho một số hợp chất chứa nguyên tố nitrogen như hình dưới đây:\nVề tính chất hoá học của các chất, hãy cho biết phát biểu sau đúng hay sai:",
            [
                ("a", "Đều có phản ứng với dung dịch HCl, sinh ra muối.", True,
                 "Đúng, cả 3 chất đều chứa nhóm amino (-NH~2~) có tính base nên đều phản ứng với acid HCl tạo muối."),
                ("b", "Có 1 chất tham gia phản ứng thuỷ phân.", True,
                 "Đúng, chỉ có chất thứ 3 là dipeptide (Gly-Gly) chứa liên kết peptide (-CO-NH-) nên tham gia phản ứng thủy phân."),
                ("c", "Có thể trùng ngưng tạo thành peptide.", False,
                 "Sai, ethylamine không có nhóm -COOH nên không trùng ngưng; chất thứ hai là β-amino acid không tạo peptide theo định nghĩa peptide (gồm các gốc α-amino acid)."),
                ("d", "Chỉ tạo phức với Cu(OH)~2~ tạo thành dung dịch có màu xanh.", False,
                 "Sai, ethylamine và amino acid tạo phức chất màu xanh lam thẫm đặc trưng, còn dipeptide không cho phản ứng màu biuret màu tím nhưng tạo phức khác.")
            ],
            "p2_cau2_compounds.png"
        ),
        (
            3,
            "Collagen được coi là một loại protein quan trọng đối với làn da, trong đó peptide lại có vai trò lớn thúc đẩy sản xuất collagen. Điều này giải thích cho việc peptide có khả năng chống lại sự lão hóa, giúp làm mờ các vết nhăn làm làn da luôn căng mịn.",
            [
                ("a", "Trong một phân tử tetrapeptide mạch hở có 4 liên kết peptide.", False,
                 "Sai, phân tử tetrapeptide cấu tạo từ 4 gốc amino acid liên kết với nhau bằng 3 liên kết peptide (n - 1 liên kết)."),
                ("b", "Peptide là loại hợp chất chứa gốc α-amino acid.", True,
                 "Đúng, theo định nghĩa, phân tử peptide được cấu tạo từ các đơn vị α-amino acid liên kết bằng liên kết peptide."),
                ("c", "Trong môi trường base, dipeptide mạch hở tác dụng được với Cu(OH)~2~ cho hợp chất màu tím.", False,
                 "Sai, phản ứng màu biuret chỉ xảy ra với các peptide có từ 2 liên kết peptide trở lên (từ tripeptide trở lên). Dipeptide không cho phản ứng này."),
                ("d", "Các hợp chất peptide kém bền trong môi trường base nhưng bền trong môi trường acid.", False,
                 "Sai, liên kết peptide bị thủy phân trong cả môi trường acid lẫn môi trường base, do đó kém bền trong cả hai môi trường.")
            ],
            None
        ),
        (
            4,
            "Thành phần tạo nên vị ngọt đặc trưng của nước mắm (được sản xuất từ cá) và nước tương (được sản xuất từ đậu nành) là các amino acid tạo thành từ sự thuỷ phân hoàn toàn của protein có trong cá hoặc đậu nành.",
            [
                ("a", "Độ đạm tương ứng với hàm lượng oxygen có trong nước mắm, nước tương.", False,
                 "Sai, độ đạm của nước mắm là tổng hàm lượng nitrogen (N) có trong 1 lít nước mắm (g N/L), không phải oxygen."),
                ("b", "Độ đạm có tỉ lệ thuận với hàm lượng amino acid có trong nước mắm, nước tương.", True,
                 "Đúng, hàm lượng amino acid càng cao thì hàm lượng nitrogen hòa tan càng lớn, độ đạm càng cao."),
                ("c", "Vì hàm lượng nitrogen quyết định đến độ đạm trong khi nitrogen có trong thành phần của amino acid nên độ đạm có tỉ lệ nghịch với hàm lượng amino acid.", False,
                 "Sai, nitrogen có trong amino acid nên tỉ lệ thuận, không phải tỉ lệ nghịch."),
                ("d", "Độ đạm trong nước mắm là quá trình thủy phân cá, các chất protein được thủy phân tạo thành các amino acid nước mắm sẽ có mùi thơm đặc trưng riêng.", True,
                 "Đúng, quá trình ủ chượp cá với muối phân giải protein thành các amino acid tạo nên độ đạm và hương vị đặc trưng của nước mắm.")
            ],
            None
        )
    ]
    
    for q in part2_data:
        num, text, items, img_file = q
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        
        r_num = p.add_run(f"Câu {num}. ")
        set_run_style(r_num, font_size=11, bold=True)
        add_chem_run(p, text, base_size=11)
        
        if img_file:
            p_img = doc.add_paragraph()
            p_img.paragraph_format.space_before = Pt(2)
            p_img.paragraph_format.space_after = Pt(3)
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.add_run().add_picture(os.path.join(IMG_DIR, img_file), width=Cm(12.5))
            
        for label, item_text, is_true, sol_item in items:
            p_item = doc.add_paragraph()
            p_item.paragraph_format.space_before = Pt(1)
            p_item.paragraph_format.space_after = Pt(2)
            p_item.paragraph_format.line_spacing = 1.15
            p_item.paragraph_format.left_indent = Cm(0.6)
            
            r_l = p_item.add_run(f"{label}) ")
            set_run_style(r_l, font_size=11, bold=True)
            add_chem_run(p_item, item_text, base_size=11)
            
            if is_teacher:
                status_str = " (ĐÚNG)" if is_true else " (SAI)"
                r_status = p_item.add_run(status_str)
                set_run_style(r_status, font_size=11, bold=True, color=COLOR_CORRECT)
                
        if is_teacher:
            sol_full = "\n".join([f"- **Ý {lbl})**: {sol_it}" for lbl, _, _, sol_it in items])
            ans_summary = ", ".join([f"{lbl} - {'Đ' if t else 'S'}" for lbl, _, t, _ in items])
            sol_full += f"\n**Đáp án: {ans_summary}.**"
            add_solution(doc, sol_full)

    # ========================== PHẦN III ==========================
    add_part_heading(
        doc,
        "PHẦN III (1,5 điểm – 6 câu): Câu trắc nghiệm yêu cầu trả lời ngắn.",
        "Thí sinh trả lời từ câu 1 đến câu 6. Mỗi câu hỏi thí sinh điền câu trả lời vào ô tương ứng."
    )
    
    part3_data = [
        (
            1,
            "Cho dãy các chất: H~2~NCH~2~COOH, HCOONH~4~, (CH~3~NH~3~)~2~CO~3~, C~6~H~5~NH~2~, C~2~H~5~NH~2~, CH~3~COOH, H~2~NCH~2~CONHCH(CH~3~)COOH. Số lượng chất trong dãy phản ứng được với cả 2 dung dịch NaOH và dung dịch HCl là",
            "4",
            "Các chất phản ứng được với cả dung dịch NaOH và HCl gồm:\n"
            "1. H~2~NCH~2~COOH (glycine - amino acid lưỡng tính).\n"
            "2. HCOONH~4~ (muối ammonium của acid yếu HCOOH).\n"
            "3. (CH~3~NH~3~)~2~CO~3~ (muối alkylammonium của acid yếu H~2~CO~3~).\n"
            "4. H~2~NCH~2~CONHCH(CH~3~)COOH (dipeptide Gly-Ala, lưỡng tính).\n"
            "- C~6~H~5~NH~2~ và C~2~H~5~NH~2~ chỉ tác dụng với HCl (tính base).\n"
            "- CH~3~COOH chỉ tác dụng với NaOH (tính acid).\n"
            "=> Có 4 chất thỏa mãn.\n"
            "**Đáp số: 4.**"
        ),
        (
            2,
            "Cho các chất sau: CH~3~COONH~4~, CH~3~COONH~3~CH~3~, C~2~H~5~NH~2~, H~2~NCH~2~COOC~2~H~5~. Số chất trong dãy vừa tác dụng được với NaOH và vừa tác dụng được với HCl trong dung dịch là:",
            "3",
            "Các chất vừa tác dụng với NaOH vừa tác dụng với HCl gồm:\n"
            "1. CH~3~COONH~4~ (muối ammonium của acid yếu).\n"
            "2. CH~3~COONH~3~CH~3~ (muối methylammonium của acid yếu).\n"
            "3. H~2~NCH~2~COOC~2~H~5~ (ester của glycine: nhóm -NH~2~ tác dụng HCl, nhóm ester -COOC~2~H~5~ thủy phân trong NaOH).\n"
            "- C~2~H~5~NH~2~ chỉ tác dụng với HCl.\n"
            "=> Có 3 chất thỏa mãn.\n"
            "**Đáp số: 3.**"
        ),
        (
            3,
            "Nếu thuỷ phân không hoàn toàn pentapeptide Gly-Ala-Gly-Ala-Gly thì thu được tối đa bao nhiêu dipeptide khác nhau?",
            "2",
            "Pentapeptide: Gly-Ala-Gly-Ala-Gly.\n"
            "Khi thủy phân không hoàn toàn tạo các đoạn 2 mắt xích liên tiếp gồm:\n"
            "- Đoạn 1: Gly-Ala\n"
            "- Đoạn 2: Ala-Gly\n"
            "- Đoạn 3: Gly-Ala\n"
            "- Đoạn 4: Ala-Gly\n"
            "=> Thu được tối đa 2 dipeptide khác nhau là Gly-Ala và Ala-Gly.\n"
            "**Đáp số: 2.**"
        ),
        (
            4,
            "Tiến hành thí nghiệm phản ứng màu biuret của lòng trắng trứng (protein) theo các bước sau đây:\n"
            "Bước 1: Cho vào ống nghiệm 1 giọt dung dịch CuSO~4~ 2% + 1 mL dung dịch NaOH 30%.\n"
            "Bước 2: Lắc nhẹ, gạn lớp dung dịch để giữ kết tủa.\n"
            "Bước 3: Thêm 4 mL dung dịch lòng trắng trứng vào ống nghiệm, lắc đều.\n"
            "Cho các nhận định sau:\n"
            "(1) Cần lấy dư dung dịch NaOH để đảm bảo môi trường cho phản ứng tạo phức.\n"
            "(2) Có thể thay thế dung dịch lòng trắng trứng bằng dung dịch Gly-Ala.\n"
            "(3) Sau bước 3, kết tủa bị hoà tan và dung dịch có màu tím đặc trưng.\n"
            "(4) Sau bước 1, trong ống nghiệm xuất hiện kết tủa màu xanh lam.\n"
            "Có bao nhiêu nhận định đúng trong các nhận định trên?",
            "3",
            "(1) Đúng, phản ứng màu biuret tạo phức chelate giữa Cu^2+^ và các liên kết peptide chỉ xảy ra trong môi trường kiềm.\n"
            "(2) Sai, Gly-Ala là dipeptide chỉ có 1 liên kết peptide, không cho phản ứng màu biuret (cần từ 2 liên kết peptide trở lên).\n"
            "(3) Đúng, sau bước 3 lòng trắng trứng hòa tan kết tủa Cu(OH)~2~ tạo phức chất màu tím đặc trưng.\n"
            "(4) Đúng, sau bước 1 xuất hiện kết tủa màu xanh lam Cu(OH)~2~.\n"
            "=> Có 3 nhận định đúng: (1), (3), (4).\n"
            "**Đáp số: 3.**"
        ),
        (
            5,
            "Xác định phân tử khối gần đúng của một hemoglobin (huyết cầu tố) chứa 0,4% Fe về khối lượng (mỗi phân tử hemoglobin chỉ chứa 1 nguyên tử iron).",
            "14000",
            "Mỗi phân tử hemoglobin chứa 1 nguyên tử Fe (M~Fe~ = 56 amu).\n"
            "Khối lượng phân tử gần đúng của hemoglobin là:\n"
            "M = (56 * 100) / 0,4 = 14 000 amu.\n"
            "**Đáp số: 14000.**"
        ),
        (
            6,
            "Một khẩu phần thịt gà nướng 100g chứa khoảng 25g protein. Một người trong ngày đã ăn 150g thịt gà, 200g cá hồi (chứa 22g protein/100g) và 50g phô mai (chứa 7g protein/100g). Hãy tính tổng lượng protein (theo đơn vị gram) người đó đã hấp thụ từ các thực phẩm này.",
            "85",
            "Lượng protein từ từng loại thực phẩm:\n"
            "- Thịt gà: 150 * (25 / 100) = 37,5 g.\n"
            "- Cá hồi: 200 * (22 / 100) = 44 g.\n"
            "- Phô mai: 50 * (7 / 100) = 3,5 g.\n"
            "=> Tổng lượng protein hấp thụ: 37,5 + 44 + 3,5 = 85 g.\n"
            "**Đáp số: 85.**"
        )
    ]
    
    for q in part3_data:
        num, text, ans, sol = q
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        
        r_num = p.add_run(f"Câu {num}. ")
        set_run_style(r_num, font_size=11, bold=True)
        add_chem_run(p, text, base_size=11)
        
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(1)
        p_ans.paragraph_format.space_after = Pt(3)
        p_ans.paragraph_format.left_indent = Cm(0.6)
        
        r_ans_lbl = p_ans.add_run("Đáp số: ")
        set_run_style(r_ans_lbl, font_size=11, italic=True)
        
        if is_teacher:
            r_ans_val = p_ans.add_run(ans)
            set_run_style(r_ans_val, font_size=11, bold=True, color=COLOR_CORRECT)
        else:
            r_dots = p_ans.add_run(". . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . .")
            set_run_style(r_dots, font_size=11, color=COLOR_GRAY)
            
        if is_teacher and sol:
            add_solution(doc, sol)

    # Summary table for teacher version
    if is_teacher:
        p_tbl_hdr = doc.add_paragraph()
        p_tbl_hdr.paragraph_format.space_before = Pt(12)
        p_tbl_hdr.paragraph_format.space_after = Pt(4)
        r_thdr = p_tbl_hdr.add_run("BẢNG TỔNG HỢP ĐÁP ÁN ĐỀ SỐ 07")
        set_run_style(r_thdr, font_size=11.5, bold=True, color=COLOR_PRIMARY)
        
        p_p1 = doc.add_paragraph()
        p_p1.paragraph_format.space_before = Pt(2)
        p_p1.paragraph_format.space_after = Pt(2)
        r_p1 = p_p1.add_run("Phần I: ")
        set_run_style(r_p1, font_size=10.5, bold=True)
        p1_ans_str = "1.A  2.D  3.B  4.D  5.C  6.B  7.B  8.D  9.A  10.D  11.C  12.D  13.D  14.B  15.B  16.A  17.A  18.A"
        r_p1_val = p_p1.add_run(p1_ans_str)
        set_run_style(r_p1_val, font_size=10.5, bold=True, color=COLOR_CORRECT)
        
        p_p2 = doc.add_paragraph()
        p_p2.paragraph_format.space_before = Pt(2)
        p_p2.paragraph_format.space_after = Pt(2)
        r_p2 = p_p2.add_run("Phần II: ")
        set_run_style(r_p2, font_size=10.5, bold=True)
        p2_ans_str = "Câu 1: a-Đ, b-Đ, c-Đ, d-Đ  |  Câu 2: a-Đ, b-Đ, c-S, d-S  |  Câu 3: a-S, b-Đ, c-S, d-S  |  Câu 4: a-S, b-Đ, c-S, d-Đ"
        r_p2_val = p_p2.add_run(p2_ans_str)
        set_run_style(r_p2_val, font_size=10.5, bold=True, color=COLOR_CORRECT)
        
        p_p3 = doc.add_paragraph()
        p_p3.paragraph_format.space_before = Pt(2)
        p_p3.paragraph_format.space_after = Pt(4)
        r_p3 = p_p3.add_run("Phần III: ")
        set_run_style(r_p3, font_size=10.5, bold=True)
        p3_ans_str = "Câu 1: 4  |  Câu 2: 3  |  Câu 3: 2  |  Câu 4: 3  |  Câu 5: 14000  |  Câu 6: 85"
        r_p3_val = p_p3.add_run(p3_ans_str)
        set_run_style(r_p3_val, font_size=10.5, bold=True, color=COLOR_CORRECT)

    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(10)
    p_end.paragraph_format.space_after = Pt(6)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("---------- HẾT ----------")
    set_run_style(r_end, font_size=11, bold=True)

    out_name = "De_Kiem_Tra_Hoa_12_De07_GiaoVien.docx" if is_teacher else "De_Kiem_Tra_Hoa_12_De07.docx"
    out_path = os.path.join(OUT_DIR, out_name)
    doc.save(out_path)
    print(f"Saved: {out_path}")

if __name__ == "__main__":
    build_docx_de07(is_teacher=False)
    build_docx_de07(is_teacher=True)
