# -*- coding: utf-8 -*-
import sys
import docx

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

for name in ['HocSinh', 'Azota']:
    f = rf"c:\Antigravity_Thanh\San_Pham\De_Kiem_Tra_GK1_Hoa_12_De20_{name}.docx"
    doc = docx.Document(f)
    print(f"=== {name} ===")
    for i in range(min(15, len(doc.paragraphs))):
        p = doc.paragraphs[i]
        if "cảnh báo các con ong khác" in p.text:
            print(f"[{i}] {p.text}")
            p_next = doc.paragraphs[i+1]
            imgs = p_next._element.xpath('.//*[local-name()="blip" or local-name()="imagedata"]')
            print(f"[{i+1}] (imgs={len(imgs)}) text: '{p_next.text}'")
            break
