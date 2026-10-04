import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

CROP_DIR = r"c:\Antigravity_Thanh\San_Pham\images_crop"
OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"

COLOR_HF_RED = RGBColor(211, 47, 47)      # Coral Red (#D32F2F)
COLOR_HF_GREEN_HEX = "8BC390"             # Sage Green for horizontal lines (#8BC390)

COLOR_PRIMARY = RGBColor(46, 125, 50)     # Dark Green / Teal for headers in body
COLOR_CORRECT = RGBColor(198, 40, 40)     # Red for correct answers
COLOR_SOLUTION = RGBColor(21, 101, 192)   # Blue for solutions
COLOR_GRAY = RGBColor(97, 97, 97)

def add_p_border_bottom(p, color="8BC390", sz="8"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_p_border_top(p, color="8BC390", sz="8"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def set_run_style(run, font_name="Times New Roman", font_size=12, bold=False, italic=False, subscript=False, superscript=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def render_rich_text(p, text, base_font="Times New Roman", base_size=12, base_bold=False, base_italic=False, base_color=None):
    tokens = re.split(r'(~[^~]+~|\^[^\^]+\^|\*[^\*]+\*)', text)
    for token in tokens:
        if not token:
            continue
        is_bold = base_bold
        is_italic = base_italic
        is_sub = False
        is_sup = False
        content = token

        if token.startswith('*') and token.endswith('*'):
            is_bold = True
            content = token[1:-1]
        elif token.startswith('~') and token.endswith('~'):
            is_sub = True
            content = token[1:-1]
        elif token.startswith('^') and token.endswith('^'):
            is_sup = True
            content = token[1:-1]

        run = p.add_run(content)
        set_run_style(run, font_name=base_font, font_size=base_size, bold=is_bold, italic=is_italic,
                      subscript=is_sub, superscript=is_sup, color=base_color)

def setup_header_footer(doc):
    normal_style = doc.styles['Normal']
    normal_style.paragraph_format.tab_stops.clear_all()
    for pos in [0.5, 1.5, 5.0, 9.5, 14.0]:
        normal_style.paragraph_format.tab_stops.add_tab_stop(Cm(pos), WD_TAB_ALIGNMENT.LEFT)

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    section.start_type = WD_SECTION_START.NEW_PAGE
    section.different_first_page_header_footer = False
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    sectPr = section._sectPr
    vAlign = OxmlElement('w:vAlign')
    vAlign.set(qn('w:val'), 'top')
    sectPr.append(vAlign)

    header = section.header
    hp = header.paragraphs[0]
    hp.text = ""
    hp.paragraph_format.tab_stops.clear_all()
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
    hp.paragraph_format.space_before = Pt(0)
    hp.paragraph_format.space_after = Pt(2)

    r_hl = hp.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)

    hp.add_run('\t')

    r_hr1 = hp.add_run("Thầy ")
    set_run_style(r_hr1, font_size=10, italic=True, color=COLOR_HF_RED)
    r_hr2 = hp.add_run("TRẦN VĂN THẠNH")
    set_run_style(r_hr2, font_size=10, bold=True, italic=True, color=COLOR_HF_RED)
    r_hr3 = hp.add_run(" - 0777.470.803")
    set_run_style(r_hr3, font_size=10, italic=True, color=COLOR_HF_RED)

    add_p_border_bottom(hp, color=COLOR_HF_GREEN_HEX, sz="8")

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.paragraph_format.tab_stops.clear_all()
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
    fp.paragraph_format.space_before = Pt(4)
    fp.paragraph_format.space_after = Pt(0)

    r_fl = fp.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế | CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)

    fp.add_run('\t')
    add_page_number_field(fp, font_size=10, bold=True, color=COLOR_HF_RED)

    add_p_border_top(fp, color=COLOR_HF_GREEN_HEX, sz="8")

def add_exam_header_block(doc, is_teacher=False):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = tbl._tbl.tblPr
    tblBorders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/><w:insideH w:val="none"/><w:insideV w:val="none"/></w:tblBorders>')
    tbl_pr.append(tblBorders)

    c0 = tbl.cell(0, 0)
    c0.width = Cm(6.5)
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_after = Pt(2)
    p0.paragraph_format.line_spacing = 1.15
    r = p0.add_run("SỞ GD&ĐT THỪA THIÊN HUẾ\nTRƯỜNG THPT CAO THẮNG\n")
    set_run_style(r, font_size=10.5, bold=True)
    r_code = p0.add_run("Mã đề thi: 103")
    set_run_style(r_code, font_size=10.5, bold=True, color=COLOR_PRIMARY)

    c1 = tbl.cell(0, 1)
    c1.width = Cm(11.5)
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_after = Pt(2)
    p1.paragraph_format.line_spacing = 1.15
    r = p1.add_run("ĐỀ KIỂM TRA CUỐI KÌ II, NĂM HỌC 2023-2024\nMÔN: HÓA HỌC 10\n")
    set_run_style(r, font_size=10.5, bold=True)
    if is_teacher:
        r_sub = p1.add_run("(HƯỚNG DẪN CHẤM VÀ LỜI GIẢI CHI TIẾT)")
        set_run_style(r_sub, font_size=10, bold=True, color=COLOR_CORRECT)
    else:
        r_sub = p1.add_run("Thời gian làm bài: 45 phút (không kể thời gian giao đề)")
        set_run_style(r_sub, font_size=10, italic=True)

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(4)
    p_info.paragraph_format.space_after = Pt(2)
    p_info.paragraph_format.line_spacing = 1.15
    if not is_teacher:
        r = p_info.add_run("Họ và tên: ................................................................ Số báo danh: ........................")
        set_run_style(r, font_size=11)

    p_ntk = doc.add_paragraph()
    p_ntk.paragraph_format.space_before = Pt(0)
    p_ntk.paragraph_format.space_after = Pt(4)
    r = p_ntk.add_run("Cho biết Nguyên tử khối của: Na = 23; K = 39; Al = 27; Br = 80; Cl = 35,5; I = 127; O = 16; H = 1.")
    set_run_style(r, font_size=10.5, italic=True)

def add_section_title(doc, title_text, score_text=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"■ {title_text}")
    set_run_style(r, font_size=12, bold=True, color=COLOR_PRIMARY)
    if score_text:
        r_s = p.add_run(f" {score_text}")
        set_run_style(r_s, font_size=11, bold=True, italic=True)

def add_instruction_line(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_style(r, font_size=11, italic=True)

def add_question_prompt(doc, q_num, prompt_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_label = p.add_run(f"Câu {q_num}. ")
    set_run_style(r_label, font_size=12, bold=True)
    render_rich_text(p, prompt_text, base_size=12)
    return p

def add_choices_tabbed(doc, choices, correct_idx=None, is_teacher=False, layout_mode="4cols"):
    labels = ["A. ", "B. ", "C. ", "D. "]

    def set_standard_tabs(p):
        p.paragraph_format.tab_stops.clear_all()
        for pos in [0.5, 1.5, 5.0, 9.5, 14.0]:
            p.paragraph_format.tab_stops.add_tab_stop(Cm(pos), WD_TAB_ALIGNMENT.LEFT)

    if layout_mode == "4cols":
        p = doc.add_paragraph()
        set_standard_tabs(p)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15

        p.add_run('\t')
        is_corr0 = (is_teacher and correct_idx == 0)
        r0 = p.add_run(labels[0])
        set_run_style(r0, font_size=12, bold=True, color=COLOR_CORRECT if is_corr0 else None)
        render_rich_text(p, choices[0] + "  ", base_size=12, base_bold=is_corr0, base_color=COLOR_CORRECT if is_corr0 else None)

        tabs_to_b = '\t\t' if len(choices[0]) < 5 else '\t'
        p.add_run(tabs_to_b)
        is_corr1 = (is_teacher and correct_idx == 1)
        r1 = p.add_run(labels[1])
        set_run_style(r1, font_size=12, bold=True, color=COLOR_CORRECT if is_corr1 else None)
        render_rich_text(p, choices[1] + "  ", base_size=12, base_bold=is_corr1, base_color=COLOR_CORRECT if is_corr1 else None)

        p.add_run('\t')
        is_corr2 = (is_teacher and correct_idx == 2)
        r2 = p.add_run(labels[2])
        set_run_style(r2, font_size=12, bold=True, color=COLOR_CORRECT if is_corr2 else None)
        render_rich_text(p, choices[2] + "  ", base_size=12, base_bold=is_corr2, base_color=COLOR_CORRECT if is_corr2 else None)

        p.add_run('\t')
        is_corr3 = (is_teacher and correct_idx == 3)
        r3 = p.add_run(labels[3])
        set_run_style(r3, font_size=12, bold=True, color=COLOR_CORRECT if is_corr3 else None)
        render_rich_text(p, choices[3], base_size=12, base_bold=is_corr3, base_color=COLOR_CORRECT if is_corr3 else None)

    elif layout_mode == "2cols":
        p1 = doc.add_paragraph()
        set_standard_tabs(p1)
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.line_spacing = 1.15

        p1.add_run('\t')
        is_corr0 = (is_teacher and correct_idx == 0)
        r0 = p1.add_run(labels[0])
        set_run_style(r0, font_size=12, bold=True, color=COLOR_CORRECT if is_corr0 else None)
        render_rich_text(p1, choices[0], base_size=12, base_bold=is_corr0, base_color=COLOR_CORRECT if is_corr0 else None)

        tabs_to_b = '\t\t' if len(choices[0]) < 12 else '\t'
        p1.add_run(tabs_to_b)
        is_corr1 = (is_teacher and correct_idx == 1)
        r1 = p1.add_run(labels[1])
        set_run_style(r1, font_size=12, bold=True, color=COLOR_CORRECT if is_corr1 else None)
        render_rich_text(p1, choices[1], base_size=12, base_bold=is_corr1, base_color=COLOR_CORRECT if is_corr1 else None)

        p2 = doc.add_paragraph()
        set_standard_tabs(p2)
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.line_spacing = 1.15

        p2.add_run('\t')
        is_corr2 = (is_teacher and correct_idx == 2)
        r2 = p2.add_run(labels[2])
        set_run_style(r2, font_size=12, bold=True, color=COLOR_CORRECT if is_corr2 else None)
        render_rich_text(p2, choices[2], base_size=12, base_bold=is_corr2, base_color=COLOR_CORRECT if is_corr2 else None)

        tabs_to_d = '\t\t' if len(choices[2]) < 12 else '\t'
        p2.add_run(tabs_to_d)
        is_corr3 = (is_teacher and correct_idx == 3)
        r3 = p2.add_run(labels[3])
        set_run_style(r3, font_size=12, bold=True, color=COLOR_CORRECT if is_corr3 else None)
        render_rich_text(p2, choices[3], base_size=12, base_bold=is_corr3, base_color=COLOR_CORRECT if is_corr3 else None)

    else: # 1col
        for i in range(4):
            p = doc.add_paragraph()
            set_standard_tabs(p)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1 if i < 3 else 2)
            p.paragraph_format.line_spacing = 1.15

            p.add_run('\t')
            is_corr = (is_teacher and i == correct_idx)
            r_lbl = p.add_run(labels[i])
            set_run_style(r_lbl, font_size=12, bold=True, color=COLOR_CORRECT if is_corr else None)
            render_rich_text(p, choices[i], base_size=12, base_bold=is_corr, base_color=COLOR_CORRECT if is_corr else None)

def add_solution_block(doc, solution_text, choose_text=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.5)

    r_lbl = p.add_run("Lời giải: ")
    set_run_style(r_lbl, font_size=11.5, bold=True, color=COLOR_SOLUTION)

    render_rich_text(p, solution_text, base_size=11.5)

    if choose_text:
        r_c = p.add_run(f" ➔ {choose_text}")
        set_run_style(r_c, font_size=11.5, bold=True, color=COLOR_CORRECT)

def add_image_centered(doc, image_name, width_cm=6.0):
    img_path = os.path.join(CROP_DIR, image_name)
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        run.add_picture(img_path, width=Cm(width_cm))

def build_exam_document_hoa10(is_teacher=False):
    doc = docx.Document()
    setup_header_footer(doc)
    add_exam_header_block(doc, is_teacher=is_teacher)

    # ==========================
    # PHẦN A. TRẮC NGHIỆM
    # ==========================
    add_section_title(doc, "A. PHẦN TRẮC NGHIỆM (28 CÂU, 07 ĐIỂM).")

    # Câu 1
    add_question_prompt(doc, 1, "Enthalpy tạo thành của một chất là")
    add_choices_tabbed(doc, [
        "nhiệt kèm theo phản ứng trong điều kiện chuẩn.",
        "nhiệt kèm theo phản ứng tạo thành 1 mol chất đó từ các đơn chất bền nhất.",
        "nhiệt kèm theo phản ứng tạo thành 1 mol chất đó từ các hợp chất bền nhất.",
        "nhiệt kèm theo phản ứng trong quá trình đẳng áp."
    ], correct_idx=1, is_teacher=is_teacher, layout_mode="1col")
    if is_teacher:
        add_solution_block(doc, "Enthalpy tạo thành (nhiệt tạo thành) của một chất là lượng nhiệt kèm theo phản ứng tạo thành 1 mol chất đó từ các đơn chất ở dạng bền nhất.", "Chọn B.")

    # Câu 2
    add_question_prompt(doc, 2, "Tiến hành các thí nghiệm sau:\n"
                                "   (a) Cho một mẩu đá vôi (CaCO~3~) vào dung dịch HCl.\n"
                                "   (b) Cho KBr tác dụng với dung dịch H~2~SO~4~ đặc nóng.\n"
                                "   (c) Cho dung dịch AgNO~3~ vào dung dịch HI.\n"
                                "   (d) Cho dung dịch AgNO~3~ vào dung dịch HF.\n"
                                "Sau khi kết thúc thí nghiệm, số phản ứng xảy ra thuộc loại phản ứng oxi hóa – khử là")
    add_choices_tabbed(doc, ["2", "4", "3", "1"], correct_idx=3, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "(a) CaCO~3~ + 2HCl -> CaCl~2~ + CO~2~ ↑ + H~2~O (phản ứng trao đổi, không phải oxi hóa - khử).\n"
                                "(b) 2KBr + 2H~2~SO~4~ (đ, t^o^) -> K~2~SO~4~ + Br~2~ + SO~2~ ↑ + 2H~2~O (Br^-1 tăng lên Br^0, S^+6 giảm xuống S^+4 -> phản ứng oxi hóa - khử).\n"
                                "(c) AgNO~3~ + HI -> AgI ↓ + HNO~3~ (phản ứng trao đổi, không phải oxi hóa - khử).\n"
                                "(d) AgNO~3~ + HF -> không phản ứng (AgF tan tốt).\n"
                                "Vậy chỉ có 1 phản ứng oxi hóa - khử (thí nghiệm b).", "Chọn D.")

    # Câu 3
    add_question_prompt(doc, 3, "Cho 6 gam Zn hạt vào một cốc đựng dung dịch H~2~SO~4~ 4M (dư) ở nhiệt độ thường. Nếu giữ nguyên các điều kiện khác, chỉ biến đổi một trong các điều kiện sau đây:\n"
                                "   (1) Thay 6 gam Zn hạt bằng 6 gam Zn bột.\n"
                                "   (2) Thay dung dịch H~2~SO~4~ 4 M bằng dung dịch H~2~SO~4~ 2 M.\n"
                                "   (3) Thực hiện phản ứng ở nhiệt độ cao hơn (khoảng 50 °C).\n"
                                "   (4) Dùng thể tích dung dịch H~2~SO~4~ 4 M gấp đôi ban đầu.\n"
                                "Những biến đổi nào làm tăng tốc độ phản ứng?")
    add_choices_tabbed(doc, ["(1), (3).", "(1), (3), (4).", "(2), (3), (4).", "(1), (2), (3), (4)."], correct_idx=0, is_teacher=is_teacher, layout_mode="2cols")
    if is_teacher:
        add_solution_block(doc, "(1) Dùng Zn bột làm tăng diện tích tiếp xúc -> tăng tốc độ phản ứng.\n"
                                "(2) Giảm nồng độ axit từ 4M xuống 2M -> giảm tốc độ phản ứng.\n"
                                "(3) Tăng nhiệt độ lên 50 °C -> tăng tốc độ phản ứng.\n"
                                "(4) Tăng thể tích dung dịch (nồng độ không đổi, axit đã dư sẵn) -> không làm đổi tốc độ phản ứng.\n"
                                "Các biến đổi làm tăng tốc độ phản ứng là (1) và (3).", "Chọn A.")

    # Câu 4
    add_question_prompt(doc, 4, "Dẫn 2,479 lít (đkc) khí chlorine vào 200 ml dung dịch sodium hydroxide 1,5M ở nhiệt độ thường. Biết phản ứng xảy ra hoàn toàn, khối lượng chất tan có trong dung dịch sau phản ứng là")
    add_choices_tabbed(doc, ["7,45 gam.", "13,30 gam.", "17,30 gam.", "5,85 gam."], correct_idx=2, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "n(Cl~2~) = 2,479 / 24,79 = 0,1 mol; n(NaOH ban đầu) = 0,2 × 1,5 = 0,3 mol.\n"
                                "Phương trình: Cl~2~ + 2NaOH -> NaCl + NaClO + H~2~O.\n"
                                "Tỉ lệ: n(Cl~2~) = 0,1 < n(NaOH)/2 = 0,15 -> NaOH dư 0,3 - 0,2 = 0,1 mol.\n"
                                "Sau phản ứng:\n"
                                "- m(NaCl) = 0,1 × 58,5 = 5,85 gam.\n"
                                "- m(NaClO) = 0,1 × 74,5 = 7,45 gam.\n"
                                "- m(NaOH dư) = 0,1 × 40 = 4,00 gam.\n"
                                "Tổng khối lượng chất tan = 5,85 + 7,45 + 4,00 = 17,30 gam.", "Chọn C.")

    # Câu 5
    add_question_prompt(doc, 5, "Có phản ứng: CH~4~ (g) + 2O~2~ (g) -(t^o^)-> CO~2~ (g) + 2H~2~O (g)\n"
                                "Cho biết giá trị năng lượng liên kết của một số liên kết:")
    # Bảng năng lượng liên kết
    tbl_e = doc.add_table(rows=2, cols=5)
    tbl_e.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_e_pr = tbl_e._tbl.tblPr
    tblBorders_e = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/><w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/><w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/><w:insideV w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/></w:tblBorders>')
    tbl_e_pr.append(tblBorders_e)
    headers_e = ["Liên kết", "C – H", "O – H", "C = O", "O = O"]
    values_e = ["Eb (kJ/mol)", "413", "467", "745", "498"]
    widths_e = [Cm(3.5), Cm(2.5), Cm(2.5), Cm(2.5), Cm(2.5)]
    for j in range(5):
        cell_h = tbl_e.cell(0, j)
        cell_h.width = widths_e[j]
        p_h = cell_h.paragraphs[0]
        p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_h.paragraph_format.space_before = Pt(2)
        p_h.paragraph_format.space_after = Pt(2)
        r_h = p_h.add_run(headers_e[j])
        set_run_style(r_h, font_size=11, bold=True)

        cell_v = tbl_e.cell(1, j)
        cell_v.width = widths_e[j]
        p_v = cell_v.paragraphs[0]
        p_v.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_v.paragraph_format.space_before = Pt(2)
        p_v.paragraph_format.space_after = Pt(2)
        r_v = p_v.add_run(values_e[j])
        set_run_style(r_v, font_size=11)

    p_c5 = doc.add_paragraph("Biến thiên enthalpy chuẩn của phản ứng trên là")
    p_c5.paragraph_format.space_before = Pt(2)
    p_c5.paragraph_format.space_after = Pt(2)
    add_choices_tabbed(doc, ["- 710 kJ.", "+668 kJ.", "-978 kJ.", "-434 kJ."], correct_idx=0, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "Δr H^o~298~ = Σ Eb(chất đầu) - Σ Eb(sản phẩm)\n"
                                "Σ Eb(chất đầu) = 4 · Eb(C-H) + 2 · Eb(O=O) = 4 × 413 + 2 × 498 = 1652 + 996 = 2648 kJ.\n"
                                "Σ Eb(sản phẩm) = 2 · Eb(C=O) + 4 · Eb(O-H) = 2 × 745 + 4 × 467 = 1490 + 1868 = 3358 kJ.\n"
                                "Δr H^o~298~ = 2648 - 3358 = -710 kJ.", "Chọn A.")

    # Câu 6
    add_question_prompt(doc, 6, "Phản ứng nào dưới đây chứng minh tính khử của các ion halide?")
    add_choices_tabbed(doc, [
        "2HCl + CuO -> CuCl~2~ + H~2~O",
        "BaCl~2~ + H~2~SO~4~ -> BaSO~4~ + 2HCl",
        "2HBr + H~2~SO~4~ -> Br~2~ + SO~2~ + 2H~2~O",
        "HI + NaOH -> NaI + H~2~O"
    ], correct_idx=2, is_teacher=is_teacher, layout_mode="2cols")
    if is_teacher:
        add_solution_block(doc, "Trong phản ứng: 2HBr + H~2~SO~4~ -> Br~2~ + SO~2~ + 2H~2~O, ion Br^- có số oxi hóa tăng từ -1 lên 0 trong Br~2~, thể hiện tính khử.", "Chọn C.")

    # Câu 7
    add_question_prompt(doc, 7, "Hydrogen halide có nhiệt độ sôi cao nhất là")
    add_choices_tabbed(doc, ["HCl.", "HI.", "HBr.", "HF."], correct_idx=3, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "HF có nhiệt độ sôi cao nhất do các phân tử HF tạo được liên kết hydrogen liên phân tử mạnh với nhau.", "Chọn D.")

    # Câu 8
    add_question_prompt(doc, 8, "Nhận định nào sau đây đúng?")
    add_choices_tabbed(doc, [
        "Bất cứ phản ứng nào cũng cần tăng áp suất để tăng tốc độ phản ứng.",
        "Chất xúc tác là chất làm tăng tốc độ phản ứng và biến mất sau khi phản ứng kết thúc.",
        "Chất xúc tác làm tăng tốc độ phản ứng nhưng vẫn được bảo toàn về chất và lượng khi kết thúc phản ứng.",
        "Bất cứ phản ứng nào cũng cần chất xúc tác để tăng tốc độ phản ứng."
    ], correct_idx=2, is_teacher=is_teacher, layout_mode="1col")
    if is_teacher:
        add_solution_block(doc, "Chất xúc tác làm tăng tốc độ phản ứng nhưng không bị tiêu hao (vẫn được bảo toàn về lượng và chất) sau khi kết thúc phản ứng.", "Chọn C.")

    # Câu 9
    add_question_prompt(doc, 9, "Dựa vào phương trình nhiệt hóa học của phản ứng:\n"
                                "CO~2~(g) -> CO(g) + 1/2 O~2~(g)    Δr H^o~298~ = +280 kJ\n"
                                "Giá trị Δr H^o~298~ của phản ứng: 2CO~2~(g) -> 2CO(g) + O~2~(g) là")
    add_choices_tabbed(doc, ["+560 kJ", "-420 kJ", "+140 kJ", "-1120 kJ"], correct_idx=0, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "Phản ứng thứ hai có hệ số gấp đôi phản ứng thứ nhất -> Δr H^o~298~ = 2 × (+280 kJ) = +560 kJ.", "Chọn A.")

    # Câu 10
    add_question_prompt(doc, 10, "Trong tự nhiên, các halogen")
    add_choices_tabbed(doc, [
        "tồn tại ở cả dạng đơn chất và hợp chất.",
        "chỉ tồn tại ở dạng muối halide.",
        "chỉ tồn tại ở dạng đơn chất.",
        "chỉ tồn tại ở dạng hợp chất."
    ], correct_idx=3, is_teacher=is_teacher, layout_mode="2cols")
    if is_teacher:
        add_solution_block(doc, "Do hoạt tính hóa học mạnh, trong tự nhiên các halogen không tồn tại ở dạng đơn chất mà chỉ tồn tại ở dạng hợp chất.", "Chọn D.")

    # Câu 11
    add_question_prompt(doc, 11, "Phát biểu *không* chính xác là")
    add_choices_tabbed(doc, [
        "Trong tất cả các hợp chất, fluorine chỉ có số oxi hóa -1.",
        "Trong tất cả các hợp chất, các halogen chỉ có số oxi hóa là -1.",
        "Trong hợp chất với hydrogen và kim loại, các halogen luôn thể hiện số oxi hóa là -1.",
        "Tính oxi hóa của halogen giảm dần từ fluorine đến iodine."
    ], correct_idx=1, is_teacher=is_teacher, layout_mode="1col")
    if is_teacher:
        add_solution_block(doc, "Ngoài fluorine chỉ có số oxi hóa -1, các halogen khác (Cl, Br, I) còn có các số oxi hóa dương (+1, +3, +5, +7) trong các hợp chất chứa oxygen hoặc fluorine. Do đó phát biểu B sai.", "Chọn B.")

    # Câu 12
    add_question_prompt(doc, 12, "Ý nghĩa của hệ số nhiệt độ Van't Hoff (γ) là")
    add_choices_tabbed(doc, [
        "giá trị của γ càng lớn thì ảnh hưởng của nhiệt độ đến tốc độ phản ứng càng mạnh.",
        "giá trị của γ càng lớn thì tốc độ phản ứng càng giảm.",
        "giá trị của γ càng lớn thì ảnh hưởng của nhiệt độ đến tốc độ phản ứng càng yếu.",
        "giá trị của γ càng lớn thì tốc độ phản ứng càng tăng."
    ], correct_idx=0, is_teacher=is_teacher, layout_mode="1col")
    if is_teacher:
        add_solution_block(doc, "Hệ số nhiệt độ Van't Hoff γ cho biết tốc độ phản ứng tăng lên bao nhiêu lần khi nhiệt độ tăng 10 °C. Do đó, γ càng lớn thì ảnh hưởng của nhiệt độ đến tốc độ phản ứng càng mạnh.", "Chọn A.")

    # Câu 13
    add_question_prompt(doc, 13, "Dung dịch dùng để nhận biết các ion halide là")
    add_choices_tabbed(doc, ["NaOH.", "quỳ tím.", "HCl.", "AgNO~3~."], correct_idx=3, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "Dung dịch AgNO~3~ tạo kết tủa đặc trưng với các ion halide: AgCl trắng, AgBr vàng nhạt, AgI vàng đậm.", "Chọn D.")

    # Câu 14
    add_question_prompt(doc, 14, "Trong các phản ứng hoá học, để chuyển thành anion, nguyên tử của các nguyên tố halogen đã nhận hay nhường bao nhiêu electron?")
    add_choices_tabbed(doc, [
        "Nhường đi 1 electron.",
        "Nhận thêm 1 electron.",
        "Nhận thêm 2 electron.",
        "Nhường đi 7 electron."
    ], correct_idx=1, is_teacher=is_teacher, layout_mode="2cols")
    if is_teacher:
        add_solution_block(doc, "Nguyên tử halogen có 7 electron lớp ngoài cùng (ns^2^np^5^), dễ nhận thêm 1 electron để đạt cấu hình electron bền vững của khí hiếm: X + 1e -> X^-.", "Chọn B.")

    # Câu 15
    add_question_prompt(doc, 15, "Chuẩn bị 3 ống nghiệm, các dung dịch có cùng nồng độ 0,1M.\n"
                                "Tiến hành thí nghiệm theo các bước:\n"
                                "Bước 1: Cho vào mỗi ống nghiệm khoảng 2 mL dung dịch các chất như hình vẽ:\n"
                                "   - Ống (1): dung dịch NaCl\n"
                                "   - Ống (2): dung dịch NaBr\n"
                                "   - Ống (3): dung dịch NaI")
    add_image_centered(doc, "fig_cau15_ongnghiem.png", width_cm=7.0)
    p15_b = doc.add_paragraph("Bước 2: Cho vào mỗi ống nghiệm khoảng 1mL nước chlorine.\n"
                              "Bước 3: Thêm tiếp vào mỗi ống nghiệm vài giọt dung dịch AgNO~3~.\n"
                              "Cho các phát biểu sau:\n"
                              "   (1) Cả ba ống nghiệm đều xuất hiện kết tủa trắng ở bước 2.\n"
                              "   (2) Sau bước 3, chỉ có ống nghiệm (1) thu được kết tủa trắng.\n"
                              "   (3) Từ hiện tượng quan sát được có thể kết luận tính oxi hoá của chlorine mạnh hơn của bromine và iodine.\n"
                              "   (4) Dung dịch ở ống (1) không thể phản ứng với nước bromine.\n"
                              "Số phát biểu đúng là")
    p15_b.paragraph_format.space_before = Pt(2)
    p15_b.paragraph_format.space_after = Pt(2)
    add_choices_tabbed(doc, ["2.", "3.", "1.", "4."], correct_idx=0, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "(1) Sai, ở bước 2 không có kết tủa nào xuất hiện.\n"
                                "(2) Sai, ở bước 3 sau khi Cl~2~ đẩy Br^-, I^- tạo NaCl thì cả 3 ống nghiệm đều chứa Cl^- và tạo kết tủa trắng AgCl (hoặc AgCl lẫn AgBr/AgI nếu còn dư).\n"
                                "(3) Đúng, vì chlorine đẩy được bromine và iodine ra khỏi muối halide tương ứng.\n"
                                "(4) Đúng, vì bromine có tính oxi hoá yếu hơn chlorine nên không phản ứng với dung dịch NaCl.\n"
                                "Vậy có 2 phát biểu đúng là (3) và (4).", "Chọn A.")

    # Câu 16
    add_question_prompt(doc, 16, "Cho các biến đổi sau:\n"
                                "   (1) Đun nóng chất tham gia.\n"
                                "   (2) Thêm chất xúc tác phù hợp.\n"
                                "   (3) Pha loãng dung dịch.\n"
                                "   (4) Giảm nhiệt độ.\n"
                                "   (5) Tăng nhiệt độ.\n"
                                "   (6) Giảm diện tích bề mặt tiếp xúc.\n"
                                "   (7) Tăng nồng độ chất phản ứng.\n"
                                "   (8) Chia nhỏ chất phản ứng thành nhiều mảnh nhỏ.\n"
                                "Có bao nhiêu biến đổi làm tăng tốc độ phản ứng?")
    add_choices_tabbed(doc, ["5.", "7.", "6.", "4."], correct_idx=0, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "Các biến đổi làm tăng tốc độ phản ứng là: (1) Đun nóng; (2) Thêm chất xúc tác; (5) Tăng nhiệt độ; (7) Tăng nồng độ; (8) Chia nhỏ chất phản ứng. Tổng cộng có 5 biến đổi.", "Chọn A.")

    # Câu 17
    add_question_prompt(doc, 17, "Cấu hình electron lớp ngoài cùng của nguyên tử các nguyên tố halogen là")
    add_choices_tabbed(doc, ["ns^2^np^6^.", "ns^2^np^4^.", "ns^2^np^5^.", "ns^2^np^3^."], correct_idx=2, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "Các nguyên tố halogen thuộc nhóm VIIA có 7 electron ở lớp ngoài cùng với cấu hình dạng ns^2^np^5^.", "Chọn C.")

    # Câu 18
    add_question_prompt(doc, 18, "Trong phản ứng: 3Cu + 8HNO~3~ -> 3Cu(NO~3~)~2~ + 2NO + 4H~2~O. Số phân tử nitric acid (HNO~3~) đóng vai trò chất oxi hóa là")
    add_choices_tabbed(doc, ["4", "2", "8", "6"], correct_idx=1, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "Trong 8 phân tử HNO~3~ tham gia phản ứng, có 2 phân tử bị khử tạo thành NO (đóng vai trò chất oxi hóa) và 6 phân tử đóng vai trò tạo môi trường muối Cu(NO~3~)~2~.", "Chọn B.")

    # Câu 19
    add_question_prompt(doc, 19, "Loại liên kết yếu được hình thành giữa nguyên tử H (đã liên kết với một nguyên tử có độ âm điện lớn, thường là F, O, N) với một nguyên tử khác (có độ âm điện lớn, thường là F, O, N) còn cặp electron hóa trị chưa tham gia liên kết là")
    add_choices_tabbed(doc, [
        "liên kết cộng hóa trị có cực",
        "liên kết hydrogen",
        "liên kết ion",
        "liên kết cộng hóa trị không cực"
    ], correct_idx=1, is_teacher=is_teacher, layout_mode="2cols")
    if is_teacher:
        add_solution_block(doc, "Đây là định nghĩa chuẩn của liên kết hydrogen.", "Chọn B.")

    # Câu 20
    add_question_prompt(doc, 20, "Tốc độ của một phản ứng có dạng: v = k · C~A~^x^ · C~B~^y^ (A, B là 2 chất khác nhau). Nếu tăng nồng độ chất A lên 2 lần, nồng độ chất B không đổi thì tốc độ phản ứng tăng 8 lần. Giá trị của x là")
    add_choices_tabbed(doc, ["6.", "3.", "8.", "4."], correct_idx=1, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "v' / v = (2C~A~)^x^ / C~A~^x^ = 2^x^ = 8 -> 2^x^ = 2^3^ -> x = 3.", "Chọn B.")

    # Câu 21
    add_question_prompt(doc, 21, "Yếu tố áp suất ảnh hưởng đến tốc độ phản ứng nào dưới đây?")
    add_choices_tabbed(doc, [
        "2CO(g) + O~2~(g) -> 2CO~2~(g).",
        "Fe(s) + S(s) -> FeS(s).",
        "NaOH(aq) + HCl(aq) -> NaCl(aq) + H~2~O(l).",
        "NH~4~Cl(s) -> NH~3~(g) + HCl(g)."
    ], correct_idx=0, is_teacher=is_teacher, layout_mode="2cols")
    if is_teacher:
        add_solution_block(doc, "Áp suất chỉ ảnh hưởng đến tốc độ phản ứng khi có chất tham gia (chất phản ứng) ở thể khí. Phản ứng A có CO(g) và O~2~(g) tham gia nên áp suất ảnh hưởng đến tốc độ.", "Chọn A.")

    # Câu 22
    add_question_prompt(doc, 22, "Trong phản ứng Zn + CuSO~4~ -> ZnSO~4~ + Cu. Chất bị oxi hóa là")
    add_choices_tabbed(doc, ["CuSO~4~", "Zn", "ZnSO~4~", "Cu"], correct_idx=1, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "Zn có số oxi hóa tăng từ 0 lên +2 -> Zn là chất khử -> Zn bị oxi hóa.", "Chọn B.")

    # Câu 23
    add_question_prompt(doc, 23, "Đặc điểm chung của đơn chất halogen là")
    add_choices_tabbed(doc, [
        "Vừa có tính oxi hóa vừa có tính khử.",
        "Có tính oxi hóa mạnh.",
        "Tác dụng mạnh với nước.",
        "Ở điều kiện thường là chất khí."
    ], correct_idx=1, is_teacher=is_teacher, layout_mode="2cols")
    if is_teacher:
        add_solution_block(doc, "Các đơn chất halogen đều có tính oxi hóa mạnh do có 7e lớp ngoài cùng, dễ nhận thêm 1 electron để đạt cấu hình bền vững.", "Chọn B.")

    # Câu 24
    add_question_prompt(doc, 24, "Cho phương trình hóa học của phản ứng: 2A + B -> C. Biểu thức tính tốc độ tức thời của phản ứng là")
    add_choices_tabbed(doc, [
        "v = k · C~A~ · C~B~",
        "v = k · C~A~^2^ · C~B~",
        "v = 2k · C~A~^2^ · C~B~",
        "v = k · C~A~ · C~B~^2^"
    ], correct_idx=1, is_teacher=is_teacher, layout_mode="2cols")
    if is_teacher:
        add_solution_block(doc, "Theo định luật tác dụng khối lượng: v = k · C~A~^2^ · C~B~.", "Chọn B.")

    # Câu 25
    add_question_prompt(doc, 25, "Cho các phát biểu sau:\n"
                                "   (1) Nước Javel có khả năng tẩy màu và sát khuẩn.\n"
                                "   (2) Cho Cl~2~ vào dung dịch NaOH đun nóng ta thu được nước Javel.\n"
                                "   (3) Hydrofluoric acid là acid yếu.\n"
                                "   (4) Tính khử của các ion halide tăng dần theo thứ tự: F^-, Cl^-, Br^-, I^-.\n"
                                "   (5) Khí chlorine ẩm và nước chlorine đều có tính tẩy màu.\n"
                                "Trong các phát biểu trên, số phát biểu đúng là")
    add_choices_tabbed(doc, ["1", "3.", "4.", "2."], correct_idx=2, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "(1) Đúng, do NaClO có tính oxi hóa mạnh.\n"
                                "(2) Sai, Cl~2~ phản ứng với NaOH đun nóng tạo muối chlorate NaClO~3~ chứ không tạo nước Javel.\n"
                                "(3) Đúng, HF là axit yếu.\n"
                                "(4) Đúng, tính khử tăng từ F^- đến I^-.\n"
                                "(5) Đúng, do tạo HClO có tính tẩy màu.\n"
                                "Có 4 phát biểu đúng là (1), (3), (4), (5).", "Chọn C.")

    # Câu 26
    add_question_prompt(doc, 26, "Cho phương trình nhiệt hoá học:\n"
                                "2Al (s) + Fe~2~O~3~ (s) -> Al~2~O~3~ (s) + 2Fe (s)    Δr H^o~298~ = -851,5 kJ.\n"
                                "Phản ứng này là phản ứng")
    add_choices_tabbed(doc, ["tỏa nhiệt.", "trao đổi.", "trung hòa.", "thu nhiệt."], correct_idx=0, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "Vì Δr H^o~298~ = -851,5 kJ < 0 nên đây là phản ứng tỏa nhiệt.", "Chọn A.")

    # Câu 27
    add_question_prompt(doc, 27, "Trong phản ứng hoá học, tốc độ phản ứng")
    add_choices_tabbed(doc, [
        "tỉ lệ nghịch với nhiệt độ của phản ứng.",
        "tăng khi nhiệt độ của phản ứng tăng.",
        "giảm khi nhiệt độ của phản ứng tăng.",
        "không đổi khi nhiệt độ của phản ứng tăng."
    ], correct_idx=1, is_teacher=is_teacher, layout_mode="2cols")
    if is_teacher:
        add_solution_block(doc, "Khi tăng nhiệt độ, chuyển động của các phân tử tăng lên làm tăng số va chạm hiệu quả, do đó tốc độ phản ứng tăng.", "Chọn B.")

    # Câu 28
    add_question_prompt(doc, 28, "“Bảo quản thức ăn trong tủ lạnh để thức ăn lâu bị ôi thiu”. Yếu tố nào đã ảnh hưởng đến tốc độ phản ứng trong tình huống thực tiễn trên?")
    add_choices_tabbed(doc, ["Diện tích bề mặt.", "Nhiệt độ.", "Chất xúc tác.", "Nồng độ."], correct_idx=1, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "Nhiệt độ thấp trong tủ lạnh làm chậm tốc độ các phản ứng phân hủy thức ăn do vi sinh vật gây ra.", "Chọn B.")

    # ==========================
    # PHẦN B. TỰ LUẬN
    # ==========================
    add_section_title(doc, "B. PHẦN TỰ LUẬN (03 CÂU, 03 ĐIỂM).")

    # Câu 29
    add_question_prompt(doc, 29, "Hoàn thành phương trình hóa học của các phản ứng sau:\n"
                                "   a. Br~2~ + .... -> KBr + I~2~\n"
                                "   b. KMnO~4~ + ...... -> ..... + ...... Cl~2~ + H~2~O\n"
                                "   c. Fe + Cl~2~ -(t^o^)-> ....\n"
                                "   d. Cl~2~ + NaOH -(t^o^)-> NaCl + ...... + H~2~O")
    if is_teacher:
        add_solution_block(doc, "a. Br~2~ + 2KI -> 2KBr + I~2~\n"
                                "b. 2KMnO~4~ + 16HCl -> 2KCl + 2MnCl~2~ + 5Cl~2~ ↑ + 8H~2~O\n"
                                "c. 2Fe + 3Cl~2~ -(t^o^)-> 2FeCl~3~\n"
                                "d. 3Cl~2~ + 6NaOH -(t^o^)-> 5NaCl + NaClO~3~ + 3H~2~O")

    # Câu 30
    add_question_prompt(doc, 30, "Để hoà tan hết một mẫu Al trong dung dịch HCl ở 25°C cần 36 phút. Cũng mẫu Al đó tan hết trong dung dịch acid nói trên ở 45°C trong 4 phút. Hỏi để hoà tan hết mẫu Al đó trong dung dịch acid nói trên ở 60°C thì cần thời gian bao nhiêu giây?")
    if is_teacher:
        add_solution_block(doc, "Thời gian phản ứng tỉ lệ nghịch với tốc độ phản ứng:\n"
                                "v~45~ / v~25~ = t~25~ / t~45~ = 36 / 4 = 9.\n"
                                "Theo quy tắc Van't Hoff:\n"
                                "v~45~ / v~25~ = γ^[(45 - 25)/10]^ = γ^2^ = 9 => γ = 3.\n"
                                "Ở 60°C so với 45°C:\n"
                                "v~60~ / v~45~ = γ^[(60 - 45)/10]^ = 3^1,5^ = 3√3 ≈ 5,196.\n"
                                "Thời gian hòa tan hết mẫu Al ở 60°C là:\n"
                                "t~60~ = t~45~ / (3√3) = 4 phút / (3√3) = 240 giây / (3√3) = 80 / √3 ≈ 46,19 giây.\n"
                                "(Hoặc tính theo 25°C: t~60~ = (36 × 60) / 3^3,5^ = 2160 / (27√3) = 80 / √3 ≈ 46,19 giây).", "Đáp số: 46,19 giây (khoảng 46,2 giây).")

    # Câu 31
    add_question_prompt(doc, 31, "Theo tính toán của các nhà khoa học, để phòng bệnh bướu cổ và một số bệnh khác, mỗi người cần bổ sung 1,5·10^-4^ g nguyên tố iodine mỗi ngày. Nếu lượng iodine đó chỉ được bổ sung từ muối iodised (còn gọi là muối iot, chứa 25g KI trong một tấn muối) thì mỗi người cần bao nhiêu gam muối mỗi ngày?")
    if is_teacher:
        add_solution_block(doc, "Khối lượng mol của KI: M~KI~ = 39 + 127 = 166 g/mol.\n"
                                "Khối lượng nguyên tố iodine có trong 25 g KI là:\n"
                                "m~I~ = 25 × (127 / 166) ≈ 19,1265 g.\n"
                                "Trong 1 tấn muối (1 000 000 g = 10^6^ g) chứa 19,1265 g iodine, vậy hàm lượng iodine trong muối là:\n"
                                "%m~I~ = 19,1265 / 10^6^ = 1,91265 · 10^-5^ (g I / g muối).\n"
                                "Khối lượng muối iodised cần dùng mỗi ngày để bổ sung 1,5·10^-4^ g iodine là:\n"
                                "m~muối~ = (1,5 · 10^-4^) / (1,91265 · 10^-5^) = (1,5 · 10^-4^ × 166 × 10^6^) / (25 × 127) = 24900 / 3175 ≈ 7,84 g.", "Đáp số: Khoảng 7,84 gam muối mỗi ngày.")

    p_end = doc.add_paragraph("---------HẾT---------")
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_end.paragraph_format.space_before = Pt(8)
    p_end.paragraph_format.space_after = Pt(4)
    set_run_style(p_end.runs[0], font_size=11, bold=True)

    filename = "De_Kiem_Tra_Hoa_10_Ma103_GiaoVien.docx" if is_teacher else "De_Kiem_Tra_Hoa_10_Ma103.docx"
    out_path = os.path.join(OUT_DIR, filename)
    doc.save(out_path)
    print(f"Successfully generated: {out_path}")

if __name__ == "__main__":
    build_exam_document_hoa10(is_teacher=False)
    build_exam_document_hoa10(is_teacher=True)
