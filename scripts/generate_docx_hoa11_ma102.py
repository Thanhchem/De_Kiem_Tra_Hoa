# -*- coding: utf-8 -*-
"""
Script tạo file Word (.docx) chuẩn "Form Đề Kiểm Tra" cho:
ĐỀ KIỂM TRA GIỮA HỌC KỲ I - MÔN HOÁ HỌC 11 - TRƯỜNG THPT HOÁ CHÂU - MÃ ĐỀ 102
Áp dụng 100% quy chuẩn:
- Không dùng bảng trong nội dung
- Header & Footer chuẩn
- Tab stops chuẩn
- Subscript/superscript triệt để
- Xuất bản 2 file: Bản Học sinh và Bản Giáo viên
"""

import os
import sys
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Thêm đường dẫn scripts vào sys.path để import form_de_kiem_tra
sys.path.insert(0, os.path.dirname(__file__))
from form_de_kiem_tra import (
    setup_header_footer,
    build_title_block,
    add_section_header,
    render_question_prompt,
    render_options_no_table,
    render_statement,
    add_solution_no_table,
    add_student_writing_lines,
    render_rich_text,
    COLOR_CORRECT,
    COLOR_PRIMARY
)

def create_exam_doc(is_teacher=False):
    doc = docx.Document()
    setup_header_footer(doc)

    build_title_block(
        doc,
        school_line="SỞ GIÁO DỤC VÀ ĐÀO TẠO THÀNH PHỐ HUẾ\nTRƯỜNG THPT HOÁ CHÂU",
        exam_title="KIỂM TRA GIỮA HỌC KỲ I - NĂM HỌC 2025 - 2026",
        subject_line="MÔN: HOÁ HỌC - LỚP 11",
        num_pages="03",
        duration="45 phút",
        exam_code="102",
        atomic_masses="H = 1; C = 12; N = 14; O = 16; Na = 23; Al = 27; S = 32; Cl = 35,5; K = 39; Ca = 40; Fe = 56; Cu = 64; Ba = 137",
        is_teacher=is_teacher
    )

    # ==================== PHẦN I ====================
    add_section_header(
        doc,
        title="A. PHẦN TRẮC NGHIỆM",
        subtitle=None
    )
    add_section_header(
        doc,
        title="PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn (3,0 điểm)",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 12. Mỗi câu hỏi thí sinh chỉ chọn một phương án."
    )

    p1_questions = [
        {
            "q": "Câu 1. Cho phương trình: CH~3~COOH + H~2~O ⇌ CH~3~COO^-^ + H~3~O^+^. Trong phản ứng thuận, theo thuyết Brønsted – Lowry chất nào là acid?",
            "opts": [
                "*A. CH~3~COOH.*",
                "B. CH~3~COO^-^.",
                "C. H~3~O^+^.",
                "D. H~2~O."
            ],
            "sol": "► *Lời giải chi tiết:*\nTheo thuyết Brønsted – Lowry, acid là chất cho proton (H^+^). Trong phản ứng thuận, CH~3~COOH nhường proton H^+^ cho H~2~O để tạo thành CH~3~COO^-^, do đó CH~3~COOH là acid. (H~2~O nhận proton H^+^ nên đóng vai trò là base).\n*Chọn A.*"
        },
        {
            "q": "Câu 2. Vai trò của NH~3~ trong phản ứng: 4NH~3~ + 5O~2~ → 4NO + 6H~2~O (xúc tác Pt, t^o^) là",
            "opts": [
                "*A. chất khử.*",
                "B. chất oxi hóa.",
                "C. acid.",
                "D. base."
            ],
            "sol": "► *Lời giải chi tiết:*\nSố oxi hóa của N trong NH~3~ là -3, sau phản ứng trong NO là +2 (tăng từ -3 lên +2). Do số oxi hóa tăng nên NH~3~ nhường electron, đóng vai trò là chất khử.\n*Chọn A.*"
        },
        {
            "q": "Câu 3. Để xác định nồng độ của một dung dịch NaOH, người ta đã tiến hành chuẩn độ bằng dung dịch HCl 0,1 M. Để chuẩn độ 10 mL dung dịch NaOH này cần 12 mL dung dịch HCl. Nồng độ của dung dịch NaOH trên là",
            "opts": [
                "A. 1,2 M.",
                "B. 0,012 M.",
                "C. 0,1 M.",
                "*D. 0,12 M.*"
            ],
            "sol": "► *Lời giải chi tiết:*\nPhương trình phản ứng chuẩn độ:\nNaOH + HCl → NaCl + H~2~O\nTại điểm tương đương: n~NaOH~ = n~HCl~ ⇔ C~NaOH~ × V~NaOH~ = C~HCl~ × V~HCl~\n⇒ C~NaOH~ = (0,1 × 12) / 10 = 0,12 M.\n*Chọn D.*"
        },
        {
            "q": "Câu 4. Phản ứng thuận nghịch là phản ứng",
            "opts": [
                "A. xảy ra giữa hai chất khí.",
                "*B. xảy ra theo hai chiều ngược nhau trong cùng điều kiện.*",
                "C. có phương trình hoá học được biểu diễn bằng mũi tên một chiều.",
                "D. chỉ xảy ra theo một chiều nhất định."
            ],
            "sol": "► *Lời giải chi tiết:*\nPhản ứng thuận nghịch là phản ứng xảy ra theo hai chiều ngược nhau trong cùng điều kiện xác định.\n*Chọn B.*"
        },
        {
            "q": "Câu 5. HNO~3~ tinh khiết là chất lỏng không màu, nhưng dung dịch HNO~3~ để lâu thường chuyển sang màu vàng là do",
            "opts": [
                "A. khi để lâu HNO~3~ bị khử bởi các chất của môi trường.",
                "B. HNO~3~ tan nhiều trong nước.",
                "C. dung dịch HNO~3~ có tính oxi hóa mạnh.",
                "*D. dung dịch HNO~3~ có hoà tan một lượng nhỏ NO~2~.*"
            ],
            "sol": "► *Lời giải chi tiết:*\nDưới tác dụng của ánh sáng hoặc nhiệt độ, HNO~3~ bị phân hủy một phần: 4HNO~3~ → 4NO~2~ + O~2~ + 2H~2~O. Khí NO~2~ sinh ra có màu nâu đỏ tan lại trong dung dịch acid làm cho dung dịch chuyển sang màu vàng.\n*Chọn D.*"
        },
        {
            "q": "Câu 6. Cho cân bằng (trong bình kín) sau: CO(g) + H~2~O(g) ⇌ CO~2~(g) + H~2~(g)   Δ~r~H~298~^o^ < 0. Trong các yếu tố: (1) tăng nhiệt độ; (2) thêm một lượng hơi nước; (3) thêm một lượng H~2~; (4) tăng áp suất chung của hệ; (5) dùng chất xúc tác. Dãy gồm các yếu tố đều làm thay đổi cân bằng của hệ là",
            "opts": [
                "A. (2), (3), (4).",
                "*B. (1), (2), (3).*",
                "C. (1), (4), (5).",
                "D. (1), (2), (4)."
            ],
            "sol": "► *Lời giải chi tiết:*\n- (1) Tăng nhiệt độ: do phản ứng tỏa nhiệt (ΔH < 0), cân bằng chuyển dịch theo chiều nghịch.\n- (2) Thêm hơi nước (tăng nồng độ chất phản ứng): cân bằng chuyển dịch theo chiều thuận.\n- (3) Thêm H~2~ (tăng nồng độ sản phẩm): cân bằng chuyển dịch theo chiều nghịch.\n- (4) Áp suất: do tổng số mol khí 2 vế bằng nhau (1 + 1 = 1 + 1 = 2) nên thay đổi áp suất không làm chuyển dịch cân bằng.\n- (5) Chất xúc tác: không làm chuyển dịch cân bằng.\nVậy các yếu tố làm thay đổi cân bằng là (1), (2), (3).\n*Chọn B.*"
        },
        {
            "q": "Câu 7. Phát biểu nào dưới đây *không đúng*?",
            "opts": [
                "A. H~2~SO~4~ đặc là chất hút nước mạnh.",
                "*B. Khi pha loãng sulfuric acid chỉ được cho từ từ nước vào acid.*",
                "C. H~2~SO~4~ loãng có đầy đủ tính chất chung của acid.",
                "D. Khi tiếp xúc với H~2~SO~4~ đặc dễ gây bỏng nặng."
            ],
            "sol": "► *Lời giải chi tiết:*\nQuy tắc an toàn khi pha loãng acid sulfuric đặc là rót từ từ acid vào nước và khuấy nhẹ bằng đũa thủy tinh, tuyệt đối không được làm ngược lại (rót nước vào acid) vì acid đặc háo nước tỏa nhiệt rất mạnh làm nước sôi đột ngột và bắn các giọt acid nóng ra ngoài gây bỏng nguy hiểm.\n*Chọn B.*"
        },
        {
            "q": "Câu 8. Cho các phát biểu sau:\n(a) Sulfur là chất rắn, màu trắng, không tan trong nước.\n(b) Sulfur và sulfur dioxide vừa có tính oxi hóa, vừa có tính khử.\n(c) Khi phản ứng với hydrogen, sulfur thể hiện tính oxi hóa.\n(d) Nước thải sinh hoạt là một trong các nguồn phát thải khí SO~2~.\n(e) Sulfur dioxide được sử dụng để tẩy trắng vải sợi, bột giấy, sản xuất sulfuric acid và diệt nấm mốc.\nSố phát biểu *đúng* là",
            "opts": [
                "A. 4.",
                "B. 1.",
                "C. 2.",
                "*D. 3.*"
            ],
            "sol": "► *Lời giải chi tiết:*\n- (a) Sai vì sulfur là chất rắn màu vàng, không phải màu trắng.\n- (b) Đúng vì S(0) và S(+4) trong SO~2~ có số oxi hóa trung gian nên vừa có tính oxi hóa vừa có tính khử.\n- (c) Đúng vì H~2~ + S → H~2~S, số oxi hóa của S giảm từ 0 xuống -2 nên S là chất oxi hóa.\n- (d) Sai vì nguồn phát thải SO~2~ chủ yếu từ đốt than đá, dầu mỏ, luyện kim màu và núi lửa; nước thải sinh hoạt chủ yếu sinh ra H~2~S, NH~3~, CH~4~.\n- (e) Đúng vì SO~2~ được ứng dụng rộng rãi để tẩy trắng bột giấy, sản xuất H~2~SO~4~ và diệt nấm mốc.\nVậy có 3 phát biểu đúng: (b), (c), (e).\n*Chọn D.*"
        },
        {
            "q": "Câu 9. Để tạo độ xốp cho một số loại bánh, có thể dùng chất nào sau đây?",
            "opts": [
                "*A. NH~4~HCO~3~.*",
                "B. CaCO~3~.",
                "C. NaCl.",
                "D. (NH~4~)~3~PO~4~."
            ],
            "sol": "► *Lời giải chi tiết:*\nNH~4~HCO~3~ (bột nở) khi nướng bánh bị nhiệt phân sinh ra các khí NH~3~, CO~2~ và hơi nước làm bánh nở xốp: NH~4~HCO~3~ → NH~3~↑ + CO~2~↑ + H~2~O↑.\n*Chọn A.*"
        },
        {
            "q": "Câu 10. Ở nhiệt độ thường, nitrogen khá trơ về mặt hoạt động là do",
            "opts": [
                "*A. Phân tử nitrogen có liên kết ba khá bền.*",
                "B. Nitrogen có độ âm điện lớn nhất trong nhóm.",
                "C. Phân tử nitrogen không phân cực.",
                "D. Nitrogen có bán kính nguyên tử nhỏ."
            ],
            "sol": "► *Lời giải chi tiết:*\nPhân tử nitrogen có liên kết ba N≡N rất bền vững với năng lượng liên kết rất lớn (E~b~ = 945 kJ/mol), nên ở nhiệt độ thường nitrogen trơ về mặt hóa học.\n*Chọn A.*"
        },
        {
            "q": "Câu 11. Tính chất nào sau đây *không phải* tính chất vật lí của sulfur?",
            "opts": [
                "A. Thể rắn ở điều kiện thường.",
                "B. Không tan trong nước.",
                "C. Màu vàng ở điều kiện thường.",
                "*D. Không tan trong benzene.*"
            ],
            "sol": "► *Lời giải chi tiết:*\nSulfur không tan trong nước nhưng tan nhiều trong các dung môi hữu cơ không phân cực như benzene (C~6~H~6~), toluene, CS~2~... Do đó 'Không tan trong benzene' không phải tính chất vật lí của sulfur.\n*Chọn D.*"
        },
        {
            "q": "Câu 12. Nitric acid thường được sử dụng để tẩy rửa ống, bề mặt kim loại do có",
            "opts": [
                "A. tính khử.",
                "*B. tính oxi hoá mạnh.*",
                "C. tính acid yếu.",
                "D. tính khử và tính acid mạnh."
            ],
            "sol": "► *Lời giải chi tiết:*\nNitric acid có tính oxi hóa rất mạnh và tính acid mạnh, có khả năng hòa tan các oxit kim loại và hầu hết các kim loại bám bẩn trên bề mặt đường ống để làm sạch kim loại.\n*Chọn B.*"
        }
    ]

    for item in p1_questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, item["q"])

        render_options_no_table(doc, item["opts"], is_teacher=is_teacher)

        if is_teacher:
            add_solution_no_table(doc, item["sol"])

    # ==================== PHẦN II ====================
    add_section_header(
        doc,
        title="PHẦN II. Câu trắc nghiệm đúng sai (2,0 điểm)",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 2. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai."
    )

    p2_questions = [
        {
            "lead": "Câu 1. Cho phản ứng tổng hợp ammonia: N~2~(g) + 3H~2~(g) ⇌ 2NH~3~(g)   Δ~r~H~298~^o^ = -92 kJ (xúc tác Fe, t^o^, p).",
            "stmts": [
                ("a) Khi tăng nhiệt độ của hệ thì cân bằng chuyển dịch theo chiều nghịch.", "Đúng"),
                ("b) Phản ứng trên là phản ứng thu nhiệt theo chiều thuận.", "Sai"),
                ("c) Để phản ứng tổng hợp NH~3~ đạt hiệu suất cao thì người ta thực hiện ở nhiệt độ rất thấp và áp suất cao.", "Sai"),
                ("d) Trong phản ứng trên N~2~ đóng vai trò là chất khử.", "Sai")
            ],
            "sol": "► *Lời giải chi tiết:*\na) Đúng vì phản ứng thuận có ΔH < 0 (tỏa nhiệt), khi tăng nhiệt độ cân bằng chuyển dịch theo chiều thu nhiệt (chiều nghịch).\nb) Sai vì ΔH = -92 kJ < 0 nên phản ứng thuận là phản ứng tỏa nhiệt.\nc) Sai vì ở nhiệt độ rất thấp tốc độ phản ứng xảy ra vô cùng chậm và xúc tác mất hoạt tính, nên thực tế công nghiệp phải thực hiện ở nhiệt độ thích hợp (400 - 450 °C) và áp suất cao.\nd) Sai vì số oxi hóa của N giảm từ 0 (N~2~) xuống -3 (NH~3~) nên N~2~ đóng vai trò là chất oxi hóa."
        },
        {
            "lead": "Câu 2. Chọn phát biểu đúng – sai:",
            "stmts": [
                ("a) Dung dịch sulfuric acid đặc hòa tan được tất cả các kim loại.", "Sai"),
                ("b) Hòa tan 18,4 gam hỗn hợp X gồm Fe và Cu trong dung dịch H~2~SO~4~ (đặc, nóng, dư), thu được 8,6765 Lít SO~2~ (đkc) và dung dịch Y. Phần trăm khối lượng của Fe trong hỗn hợp X là 30,43%.", "Đúng"),
                ("c) Dung dịch HNO~3~ đặc, nguội phản ứng được với Al và Fe.", "Sai"),
                ("d) Sulfuric acid tan tốt trong nước, quá trình hòa tan tỏa nhiệt mạnh.", "Đúng")
            ],
            "sol": "► *Lời giải chi tiết:*\na) Sai vì H~2~SO~4~ đặc không hòa tan được Au và Pt.\nb) Đúng. n~SO2~ = 8,6765 / 24,79 = 0,35 mol.\nĐặt n~Fe~ = x mol, n~Cu~ = y mol.\nTa có hệ phương trình:\n- Khối lượng: 56x + 64y = 18,4\n- Bảo toàn electron: 3x + 2y = 2 × n~SO2~ = 0,7\nGiải hệ được: x = 0,1 mol và y = 0,2 mol.\n⇒ m~Fe~ = 0,1 × 56 = 5,6 gam ⇒ %m~Fe~ = (5,6 / 18,4) × 100% ≈ 30,43%.\nc) Sai vì Al và Fe bị thụ động hóa trong HNO~3~ đặc, nguội.\nd) Đúng vì H~2~SO~4~ tan vô hạn trong nước và tỏa một lượng nhiệt rất lớn do tạo hydrate."
        }
    ]

    for item in p2_questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, item["lead"])

        for stmt, ans in item["stmts"]:
            p_s = doc.add_paragraph()
            p_s.paragraph_format.left_indent = Cm(0.5)
            p_s.paragraph_format.space_before = Pt(1)
            p_s.paragraph_format.space_after = Pt(1)
            render_statement(p_s, stmt, ans=ans, is_teacher=is_teacher)

        if is_teacher:
            add_solution_no_table(doc, item["sol"])

    # ==================== PHẦN III ====================
    add_section_header(
        doc,
        title="PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn (2,0 điểm)",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 8."
    )

    p3_questions = [
        {
            "q": "Câu 1. Cho các phản ứng:\n(a) S + O~2~ → SO~2~ (t^o^)\n(b) S + 3F~2~ → SF~6~\n(c) Hg + S → HgS\n(d) H~2~ + S → H~2~S\nSố phản ứng trong đó sulfur đơn chất đóng vai trò chất khử là bao nhiêu?",
            "ans": "2",
            "sol": "► *Lời giải chi tiết:*\nTrong (a): S^0^ → S^+4^ + 4e ⇒ S là chất khử.\nTrong (b): S^0^ → S^+6^ + 6e ⇒ S là chất khử.\nTrong (c): S^0^ + 2e → S^-2^ ⇒ S là chất oxi hóa.\nTrong (d): S^0^ + 2e → S^-2^ ⇒ S là chất oxi hóa.\nVậy có 2 phản ứng (a) và (b) trong đó S đóng vai trò là chất khử.\n► *Đáp án: 2*"
        },
        {
            "q": "Câu 2. Cho các chất: (1) KNO~3~, (2) HCl, (3) HClO, (4) H~2~S, (5) Cu(OH)~2~, (6) Ba(OH)~2~, (7) NaHCO~3~. Có bao nhiêu chất điện li mạnh trong các chất trên?",
            "ans": "4",
            "sol": "► *Lời giải chi tiết:*\nChất điện li mạnh gồm muối tan, acid mạnh, base mạnh:\n- (1) KNO~3~ (muối tan),\n- (2) HCl (acid mạnh),\n- (6) Ba(OH)~2~ (base mạnh),\n- (7) NaHCO~3~ (muối tan).\nCác chất còn lại: HClO, H~2~S là acid yếu; Cu(OH)~2~ là base không tan (chất điện li yếu).\nVậy có 4 chất điện li mạnh.\n► *Đáp án: 4*"
        },
        {
            "q": "Câu 3. Cho các phản ứng sau:\n(1) NaOH + HCl → NaCl + H~2~O\n(2) H~2~ + I~2~ ⇌ 2HI\n(3) CaCO~3~ ⇌ CaO + CO~2~ (trong bình kín)\n(4) 2KClO~3~ → 2KCl + 3O~2~ (t^o^)\nCó bao nhiêu phản ứng thuận nghịch trong các phản ứng trên?",
            "ans": "2",
            "sol": "► *Lời giải chi tiết:*\nCác phản ứng thuận nghịch được biểu diễn bằng mũi tên ⇌: (2) và (3).\nPhản ứng (1) và (4) là phản ứng một chiều.\nVậy có 2 phản ứng thuận nghịch.\n► *Đáp án: 2*"
        },
        {
            "q": "Câu 4. Cho các nhận định sau về cấu tạo phân tử nitric acid:\n(1) Liên kết O – H phân cực về oxygen.\n(2) Nguyên tử N có số oxi hóa là +5.\n(3) Nguyên tử N có hóa trị bằng 4.\n(4) Phân tử HNO~3~ chứa liên kết cho – nhận N → O.\nCó bao nhiêu phát biểu đúng?",
            "ans": "4",
            "sol": "► *Lời giải chi tiết:*\nCả 4 phát biểu đều đúng. Độ âm điện O > H nên liên kết O - H phân cực về O; số oxi hóa của N là +5; N ở chu kì 2 có 4 orbital hóa trị nên hóa trị tối đa bằng 4 (IV); liên kết giữa N và một O là liên kết phối trí (cho – nhận N → O).\n► *Đáp án: 4*"
        },
        {
            "q": "Câu 5. Cho phản ứng: aFe + bHNO~3~ → cFe(NO~3~)~3~ + dNO↑ + eH~2~O. Hệ số tỉ lượng a, b, c, d, e là những số nguyên dương có tỉ lệ tối giản. Tổng (a + b) bằng bao nhiêu?",
            "ans": "5",
            "sol": "► *Lời giải chi tiết:*\nCân bằng phản ứng oxi hóa - khử:\n1Fe + 4HNO~3~ → 1Fe(NO~3~)~3~ + 1NO↑ + 2H~2~O\nCác hệ số: a = 1, b = 4, c = 1, d = 1, e = 2.\nTổng (a + b) = 1 + 4 = 5.\n► *Đáp án: 5*"
        },
        {
            "q": "Câu 6. Cho sơ đồ chuyển hóa nitrogen trong khí quyển thành phân đạm:\nN~2~ → NO (1) → NO~2~ (2) → HNO~3~ (3) → NO~3~^-^ (4).\nCó bao nhiêu phản ứng thuộc loại oxi hóa – khử trong sơ đồ trên?",
            "ans": "3",
            "sol": "► *Lời giải chi tiết:*\n- (1) N~2~(0) + O~2~(0) → 2NO(+2): oxi hóa - khử.\n- (2) 2NO(+2) + O~2~(0) → 2NO~2~(+4): oxi hóa - khử.\n- (3) 4NO~2~(+4) + O~2~(0) + 2H~2~O → 4HNO~3~(+5): oxi hóa - khử.\n- (4) HNO~3~(+5) → H^+^ + NO~3~^-^(+5): phản ứng điện li hoặc trao đổi tạo muối, số oxi hóa của N không thay đổi (+5).\nVậy có 3 phản ứng oxi hóa - khử.\n► *Đáp án: 3*"
        },
        {
            "q": "Câu 7. Cho m gam Fe tác dụng hết với dung dịch HNO~3~ loãng (dư), sau phản ứng thu được 4,958 lít (đkc) khí NO là sản phẩm khử duy nhất. Tính giá trị của m?",
            "ans": "11,2",
            "sol": "► *Lời giải chi tiết:*\nSố mol NO = 4,958 / 24,79 = 0,2 mol.\nPhương trình: Fe + 4HNO~3~ → Fe(NO~3~)~3~ + NO↑ + 2H~2~O\nTheo phương trình hoặc bảo toàn electron: n~Fe~ = n~NO~ = 0,2 mol.\nKhối lượng sắt: m = 0,2 × 56 = 11,2 gam.\n► *Đáp án: 11,2*"
        },
        {
            "q": "Câu 8. Có các thí nghiệm sau:\n(a) Nhúng thanh sắt (iron) vào dung dịch H~2~SO~4~ loãng, nguội.\n(b) Sục khí SO~2~ vào nước bromine.\n(c) Cho dung dịch BaCl~2~ tác dụng với dung dịch H~2~SO~4~ loãng.\n(d) Nhúng lá nhôm (aluminium) vào dung dịch H~2~SO~4~ đặc, nguội.\nSố thí nghiệm xảy ra phản ứng hoá học là bao nhiêu?",
            "ans": "3",
            "sol": "► *Lời giải chi tiết:*\n- (a) Fe + H~2~SO~4~ → FeSO~4~ + H~2~↑ (có phản ứng).\n- (b) SO~2~ + Br~2~ + 2H~2~O → 2HBr + H~2~SO~4~ (có phản ứng làm mất màu nước bromine).\n- (c) BaCl~2~ + H~2~SO~4~ → BaSO~4~↓ + 2HCl (có phản ứng tạo kết tủa trắng).\n- (d) Al bị thụ động hóa trong H~2~SO~4~ đặc, nguội (không xảy ra phản ứng).\nVậy có 3 thí nghiệm xảy ra phản ứng hóa học: (a), (b), (c).\n► *Đáp án: 3*"
        }
    ]

    for item in p3_questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, item["q"])

        if is_teacher:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Cm(0.5)
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(1)
            render_rich_text(p_ans, f"► *Đáp án:* *{item['ans']}*", base_size=10.5, base_bold=True, base_color=COLOR_CORRECT)
            add_solution_no_table(doc, item["sol"])
        else:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Cm(0.5)
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(3)
            render_rich_text(p_ans, "*Trả lời:* ............................................................", base_size=10.5)

    # ==================== PHẦN IV ====================
    add_section_header(
        doc,
        title="B. PHẦN TỰ LUẬN (3,0 điểm)",
        subtitle=None
    )

    essay_questions = [
        {
            "num": "Câu 1: (0,5 điểm)",
            "prompt": "a) Ở các vùng quê, người dân thường dùng phèn chua (K~2~SO~4~.Al~2~(SO~4~)~3~.24H~2~O) để làm trong nước nhờ ứng dụng của phản ứng thuỷ phân ion Al^3+^. Giải thích?\nb) Chất hay ion nào là acid, là base trong phản ứng thuỷ phân Al^3+^?",
            "sol": "► *Lời giải chi tiết:*\na) Khi hòa tan phèn chua vào nước, ion Al^3+^ bị thủy phân theo cân bằng:\nAl^3+^ + 3H~2~O ⇌ Al(OH)~3~↓ + 3H^+^\n(hoặc [Al(H~2~O)~6~]^3+^ + H~2~O ⇌ [Al(OH)(H~2~O)~5~]^2+^ + H~3~O^+^).\nKết tủa Al(OH)~3~ dạng keo xốp có bề mặt tiếp xúc và khả năng hấp phụ rất lớn, lơ lửng trong nước sẽ kéo dính các hạt bụi bẩn, tạp chất cơ học lắng nhanh xuống đáy, làm nước trở nên trong sạch.\nb) Theo thuyết Brønsted – Lowry trong cân bằng thủy phân trên:\n- Ion Al^3+^ (dạng [Al(H~2~O)~6~]^3+^) nhường proton H^+^ nên đóng vai trò là *acid*.\n- Phân tử H~2~O nhận proton H^+^ nên đóng vai trò là *base*.",
            "lines": 6
        },
        {
            "num": "Câu 2: (1,5 điểm)",
            "prompt": "a) Viết sơ đồ sản xuất nitric acid trong công nghiệp.\nb) Tính số kg dung dịch HNO~3~ 60% điều chế được từ 340 kg ammonia, biết rằng hiệu suất của toàn bộ quá trình là 90%.",
            "sol": "► *Lời giải chi tiết:*\na) Sơ đồ sản xuất nitric acid trong công nghiệp (phương pháp Ostwald):\nNH~3~ ⎯⎯⎯⎯⎯⎯⎯⎯⎯→ NO ⎯⎯⎯→ NO~2~ ⎯⎯⎯⎯⎯⎯⎯⎯⎯→ HNO~3~\n(+O~2~, Pt, 850-900 °C)       (+O~2~)        (+O~2~, +H~2~O)\nCác phương trình phản ứng hóa học xảy ra:\n(1) 4NH~3~ + 5O~2~ → 4NO + 6H~2~O (xúc tác Pt, 850 - 900 °C)\n(2) 2NO + O~2~ → 2NO~2~\n(3) 4NO~2~ + O~2~ + 2H~2~O → 4HNO~3~\n\nb) Tính toán:\n- Số mol ammonia: n~NH3~ = 340 kg / 17 kg/kmol = 20 kmol.\n- Theo sơ đồ lí thuyết: NH~3~ → HNO~3~\n⇒ n~HNO3(LT)~ = n~NH3~ = 20 kmol.\n- Do hiệu suất của cả quá trình là H = 90%:\n⇒ n~HNO3(TT)~ = 20 × (90 / 100) = 18 kmol.\n- Khối lượng chất tan HNO~3~ nguyên chất:\nm~HNO3~ = 18 kmol × 63 kg/kmol = 1134 kg.\n- Khối lượng dung dịch HNO~3~ 60% thu được:\nm~dd HNO3 60%~ = (1134 × 100) / 60 = 1890 kg.\n*Đáp số:* 1890 kg dung dịch HNO~3~ 60%.",
            "lines": 10
        },
        {
            "num": "Câu 3: (1,0 điểm)",
            "prompt": "a) Bình đựng dung dịch H~2~SO~4~ đặc để trong không khí ẩm lâu ngày thì khối lượng bình có thay đổi không? Vì sao?\nb) Sulfur và quặng pyrite sắt là các nguyên liệu chính trong công nghiệp sản xuất sulfuric acid. Tại một nhà máy, cứ đốt cháy 1 tấn quặng pyrite sắt (chứa 84% khối lượng FeS~2~) bằng không khí, thu được tối đa V m^3^ khí SO~2~ (đkc). Tính giá trị của V?",
            "sol": "► *Lời giải chi tiết:*\na) Khối lượng của bình đựng dung dịch H~2~SO~4~ đặc để trong không khí ẩm lâu ngày sẽ *tăng lên*.\n*Giải thích:* Do acid sulfuric đặc có tính háo nước rất mạnh, nó hút hơi nước có trong không khí ẩm vào bình làm lượng nước tăng lên, dẫn đến khối lượng bình tăng lên.\n\nb) Tính toán:\n- Khối lượng FeS~2~ nguyên chất trong 1 tấn quặng pyrite:\nm~FeS2~ = 1 tấn × 84% = 0,84 tấn = 840 kg = 840 000 gam.\n- Khối lượng mol của FeS~2~: M = 56 + 32 × 2 = 120 g/mol.\n- Số mol FeS~2~: n~FeS2~ = 840 000 / 120 = 7 000 mol = 7 kmol.\n- Phương trình hóa học phản ứng đốt cháy quặng pyrite:\n4FeS~2~ + 11O~2~ → 2Fe~2~O~3~ + 8SO~2~ (t^o^)\n- Theo phản ứng (hoặc bảo toàn nguyên tố S):\nn~SO2~ = 2 × n~FeS2~ = 2 × 7 = 14 kmol = 14 000 mol.\n- Thể tích khí SO~2~ thu được tối đa ở điều kiện chuẩn (đkc, 25 °C, 1 bar, 24,79 L/mol):\nV = 14 000 × 0,02479 m^3^ = 347,06 m^3^ (tương đương 347 060 Lít).\n*(Ghi chú: Nếu tính theo điều kiện tiêu chuẩn cũ 22,4 L/mol thì V = 14 000 × 0,0224 = 313,6 m^3^).* \n*Đáp số:* V = 347,06 m^3^ (đkc).",
            "lines": 10
        }
    ]

    for item in essay_questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*{item['num']}*", base_size=10.5, base_bold=True)

        for line in item["prompt"].split("\n"):
            p_line = doc.add_paragraph()
            p_line.paragraph_format.left_indent = Cm(0.5)
            p_line.paragraph_format.space_before = Pt(1)
            p_line.paragraph_format.space_after = Pt(1)
            render_rich_text(p_line, line, base_size=10.5)

        if is_teacher:
            add_solution_no_table(doc, item["sol"])
        else:
            add_student_writing_lines(doc, num_lines=item["lines"])

    # Dòng kết thúc
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(10)
    p_end.paragraph_format.space_after = Pt(4)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_end, "---------- *HẾT* ----------", base_size=10.5, base_bold=True)

    p_rule = doc.add_paragraph()
    p_rule.paragraph_format.space_before = Pt(0)
    p_rule.paragraph_format.space_after = Pt(0)
    p_rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_rule, "_(Học sinh không được sử dụng tài liệu. Giám thị không giải thích gì thêm)_", base_size=9.5, base_italic=True)

    return doc

def main():
    out_dir = r"c:\Antigravity_Thanh\San_Pham"
    os.makedirs(out_dir, exist_ok=True)

    # 1. Tạo Bản Học sinh
    doc_student = create_exam_doc(is_teacher=False)
    path_student = os.path.join(out_dir, "De_Kiem_Tra_Hoa_11_Ma102.docx")
    doc_student.save(path_student)
    print(f"Đã tạo Bản Học sinh: {path_student}")

    # 2. Tạo Bản Giáo viên
    doc_teacher = create_exam_doc(is_teacher=True)
    path_teacher = os.path.join(out_dir, "De_Kiem_Tra_Hoa_11_Ma102_GiaoVien.docx")
    doc_teacher.save(path_teacher)
    print(f"Đã tạo Bản Giáo viên: {path_teacher}")

if __name__ == "__main__":
    main()
