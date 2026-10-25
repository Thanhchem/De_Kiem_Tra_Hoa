# -*- coding: utf-8 -*-
"""
Xây dựng Ngân hàng câu hỏi Hóa học 10 – Chương 1: Cấu tạo nguyên tử (kèm Mở đầu)
Phân loại chi tiết theo:
- Chủ đề / Bài học:
  1. Mở đầu: Đối tượng và phương pháp nghiên cứu Hóa học
  2. Bài 1: Thành phần của nguyên tử
  3. Bài 2: Nguyên tố hóa học
  4. Bài 3: Cấu trúc lớp vỏ electron của nguyên tử
- Dạng thức câu hỏi:
  + Phần I (P1): Trắc nghiệm nhiều phương án lựa chọn (4 phương án A, B, C, D)
  + Phần II (P2): Trắc nghiệm Đúng / Sai (4 ý a, b, c, d có phân loại mức độ từng ý)
  + Phần III (P3): Trắc nghiệm Trả lời ngắn & Bài toán vận dụng
- Mức độ nhận thức: Nhận biết [NB] – Thông hiểu [TH] – Vận dụng [VD]
- Xuất cơ sở dữ liệu JSON chuẩn + Bộ 2 file Word (.docx) Học sinh & Giáo viên theo đúng "FORM ĐỀ KIỂM TRA".
"""

import os
import io
import re
import json
import struct
import sys
from difflib import SequenceMatcher
from PIL import Image
import docx
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

CVL_DOCX = r"C:\Users\Admin\.gemini\antigravity\brain\f9681822-b9cc-4e9c-bdb3-77f39eb3578e\scratch\cvl_hoa10_c1.docx"
TTB_DOCX = r"C:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\2026 - Phan dang de thi thu truong, so. TTB1-20260711T201135Z-2-001\2026 - Phan dang de thi thu truong, so. TTB1\X. PD 2026 - Hoa 10. CD10 - SGK va CDHT. TTB. ok - DA chi tiet.docx"

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
BANK_DIR = r"c:\Antigravity_Thanh\San_Pham\Ngan_Hang_Hoa_10"
IMG_DIR = r"c:\Antigravity_Thanh\San_Pham\images_hoa10"
os.makedirs(OUT_DIR, exist_ok=True)
os.makedirs(BANK_DIR, exist_ok=True)
os.makedirs(IMG_DIR, exist_ok=True)

CHAPTER_INFO = {
    "chapter_number": 1,
    "code": "C1",
    "name": "CHƯƠNG 1: CẤU TẠO NGUYÊN TỬ",
    "short_name": "Cấu tạo nguyên tử",
    "file_prefix": "Ngan_Hang_Hoa_10_Chuong_1_Cau_Tao_Nguyen_Tu",
    "topics": [
        "Mở đầu: Đối tượng và phương pháp nghiên cứu Hóa học",
        "Bài 1: Thành phần của nguyên tử",
        "Bài 2: Nguyên tố hóa học",
        "Bài 3: Cấu trúc lớp vỏ electron của nguyên tử",
    ]
}

LEVEL_NAMES = {
    "NB": "Biết",
    "TH": "Hiểu",
    "VD": "Vận dụng"
}

COLOR_HF_RED = RGBColor(211, 47, 47)      # #D32F2F Coral Red
COLOR_PRIMARY = RGBColor(46, 125, 50)     # #2E7D32 Dark Green
COLOR_SOLUTION = RGBColor(21, 101, 192)   # #1565C0 Blue
COLOR_CORRECT = RGBColor(198, 40, 40)     # #C62828 Red
COLOR_TAG = RGBColor(0, 105, 92)          # #00695C Teal

SYMBOL_MAP = {
    'a': 'α', 'b': 'β', 'g': 'γ', 'd': 'δ', 'p': 'π', 'r': 'ρ',
    '-': '-', '+': '+', '=': '=', '»': '≈', '£': '≤', '³': '≥',
    '´': '×', '®': '→', 'Û': '⇔', 'Þ': '⇒', '±': '±', '°': '°',
    'S': 'Σ', 'D': 'Δ'
}
SUB_MAP = str.maketrans('0123456789+-=()nxyzab', '₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎ₙₓᵧ₂ₐ₆')
SUP_MAP = str.maketrans('0123456789+-=()nxyzab', '⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿˣʸᶻᵃᵇ')


def iter_block_items(doc):
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    for child in doc.element.body.iterchildren():
        if child.tag == qn('w:p'):
            yield ('p', Paragraph(child, doc))
        elif child.tag == qn('w:tbl'):
            yield ('tbl', Table(child, doc))


def get_cau_num_ids(doc):
    num_part = doc.part.numbering_part._element
    abs_map = {}
    for an in num_part.findall(qn('w:abstractNum')):
        aid = an.get(qn('w:abstractNumId'))
        lvls = [lt.get(qn('w:val')) or '' for lt in an.findall('.//' + qn('w:lvlText'))]
        abs_map[aid] = lvls[0] if lvls else ''
    cau_nids = set()
    for n in num_part.findall(qn('w:num')):
        nid = n.get(qn('w:numId'))
        aid_el = n.find(qn('w:abstractNumId'))
        aid = aid_el.get(qn('w:val')) if aid_el is not None else ''
        if abs_map.get(aid, '').startswith('Câu'):
            cau_nids.add(nid)
    return cau_nids


def decode_wmf_eq(blob):
    pos = 22 if blob[:4] == b'\xd7\xcd\xc6\x9a' else 0
    pos += 18
    tokens = []
    hlines = []
    cur_font = ''
    cur_h = 0
    cur_x, cur_y = 0, 0
    while pos + 6 <= len(blob):
        size_words, func = struct.unpack_from('<IH', blob, pos)
        if size_words < 3:
            break
        rec = blob[pos:pos + size_words*2]
        if func == 0x0214:
            cur_y, cur_x = struct.unpack_from('<hh', rec, 6)
        elif func == 0x0213:
            ly, lx = struct.unpack_from('<hh', rec, 6)
            if abs(ly - cur_y) <= 5 and abs(lx - cur_x) >= 150:
                hlines.append((min(cur_x, lx), max(cur_x, lx), cur_y))
            cur_x, cur_y = lx, ly
        elif func == 0x02FB:
            cur_h = abs(struct.unpack_from('<h', rec, 6)[0])
            fn_raw = rec[24:].split(b'\x00')[0]
            cur_font = fn_raw.decode('latin1', errors='ignore')
        elif func == 0x0A32:
            y, x, slen, opts = struct.unpack_from('<hhhh', rec, 6)
            if x == 0 and y == 0:
                x, y = cur_x, cur_y
            off = 14 + (8 if (opts & 0x0006) else 0)
            s = rec[off:off+slen].decode('cp1252', errors='ignore')
            dx_off = off + slen + (slen % 2)
            if dx_off + slen*2 <= len(rec):
                dxs = list(struct.unpack_from(f'<{slen}h', rec, dx_off))
            else:
                dxs = [180] * slen
            cx = x
            buf_x = cx
            buf_chars = []
            for idx_c, ch in enumerate(s):
                if 'Symbol' in cur_font:
                    ch = SYMBOL_MAP.get(ch, ch)
                if not buf_chars:
                    buf_x = cx
                buf_chars.append(ch)
                adv = dxs[idx_c] if idx_c < len(dxs) else 180
                cx += adv
                if adv > 310 and idx_c + 1 < len(s):
                    tokens.append((buf_x, y, cur_h, ''.join(buf_chars)))
                    buf_chars = []
            if buf_chars:
                tokens.append((buf_x, y, cur_h, ''.join(buf_chars)))
        pos += size_words * 2
    if not tokens:
        return ''
    max_h = max(t[2] for t in tokens) if tokens else 384

    def render_token_group(t_list):
        if not t_list:
            return ''
        m_ys = [t[1] for t in t_list if t[2] >= max_h * 0.85]
        b_y = sum(m_ys)/len(m_ys) if m_ys else t_list[0][1]
        t_list = sorted(t_list, key=lambda t: (t[0] // 60, t[1], t[0]))
        res = []
        for x, y, h, txt in t_list:
            if h < max_h * 0.8:
                if y < b_y - 30:
                    res.append(txt.translate(SUP_MAP))
                elif y > b_y + 30:
                    res.append(txt.translate(SUB_MAP))
                else:
                    res.append(txt)
            else:
                res.append(txt)
        return ''.join(res)

    if not hlines:
        return render_token_group(tokens)

    frac_lines = []
    for x1, x2, ly in hlines:
        above = [t for t in tokens if x1 - 80 <= t[0] <= x2 + 80 and t[1] < ly]
        below = [t for t in tokens if x1 - 80 <= t[0] <= x2 + 80 and t[1] > ly]
        if above and below:
            frac_lines.append((x1, x2, ly))
    if not frac_lines:
        return render_token_group(tokens)

    fx1, fx2, fly = max(frac_lines, key=lambda l: l[1] - l[0])
    left_toks = [t for t in tokens if t[0] < fx1 - 60]
    right_toks = [t for t in tokens if t[0] > fx2 + 40]
    num_toks = [t for t in tokens if fx1 - 60 <= t[0] <= fx2 + 40 and t[1] <= fly]
    den_toks = [t for t in tokens if fx1 - 60 <= t[0] <= fx2 + 40 and t[1] > fly]
    return f"{render_token_group(left_toks)}({render_token_group(num_toks)})/({render_token_group(den_toks)}){render_token_group(right_toks)}"


def clean_extracted_text(s):
    if not s:
        return ""
    s = s.replace('\uf0ad\uf0af', '↑↓').replace('\xad\xaf', '↑↓')
    s = s.replace('\uf0ad', '↑').replace('\xad', '↑')
    s = s.replace('\uf0af', '↓').replace('\xaf', '↓')
    # Fix multi-line MathType system of equations and cube roots
    s = s.replace(
        'ìíî2Z2Z+-NN==1652⇒ìíîZN==1718',
        '{ 2Z + N = 52 ; 2Z - N = 16 } ⇒ { Z = 17 ; N = 18 }'
    )
    s = re.sub(
        r'ìïíïî2Z2Z4ZXXY.*?ZZXY==2616',
        '{ (2Z_X + N_X) + 2(2Z_Y + N_Y) = 178 ; (2Z_X + 4Z_Y) - (N_X + 2N_Y) = 54 ; 4Z_Y - 2Z_X = 12 } ⇒ { Z_X + 2Z_Y = 58 ; N_X + 2N_Y = 62 ; 2Z_Y - Z_X = 6 } ⇒ { Z_X = 26 ; Z_Y = 16 }. ',
        s
    )
    s = s.replace(
        'ìïíïîAx+C=y(12x+13y)/(=100100)=12,0098⇒ìíîxy==99,02%0,98%',
        '{ A_C = (12x + 13y)/100 = 12,0098 ; x + y = 100 } ⇒ { x = 99,02% ; y = 0,98% }'
    )
    s = s.replace(
        'ìïíïînn⁶³₆₅CuCu==0,1.72,7%0,1.27,3%==0,0727mol0,0273mol',
        'n(⁶³Cu) = 0,1 × 72,7% = 0,0727 mol ; n(⁶⁵Cu) = 0,1 × 27,3% = 0,0273 mol'
    )
    s = s.replace(
        'ìïíïîa + b = 100___ACu=(35a + 37b)/(100) = 35,5⇔ìíîa =75b =25',
        '{ a + b = 100 ; A_Cl = (35a + 37b)/100 = 35,5 } ⇔ { a = 75 ; b = 25 }'
    )
    s = s.replace('A__Cl=35,5', 'A_Cl = 35,5')
    s = s.replace('Vðr=(4)/(34,1867)³=10×(m⁻)⁴⁵³', 'V = (4/3)πr³ = 4,1867×10⁻⁴⁵ (m³)')
    s = s.replace('Vðr=(4)/(34,1867)³=10×(m⁻)¹²³', 'V = (4/3)πr³ = 4,1867×10⁻¹² (m³)')
    s = s.replace('r=3×4πV¹=(3×3,17×10⁻²³)/(4π)', 'r = ∛((3×V₁)/(4π)) = ∛((3×3,17×10⁻²³)/(4π))')
    s = s.replace('r=(3×8,85×10⁻²⁴)/(4π)', 'r = ∛((3×8,85×10⁻²⁴)/(4π))')
    s = s.replace('r=(3×1,58×10⁻²³)/(4π)', 'r = ∛((3×1,58×10⁻²³)/(4π))')
    s = s.replace('r=(3×8,72×10⁻²⁴)/(4π)', 'r = ∛((3×8,72×10⁻²⁴)/(4π))')
    s = s.replace('(1,47×10⁻⁸³)', '(1,47×10⁻⁸)³')
    s = s.replace('(1,34×10⁻⁸³)', '(1,34×10⁻⁸)³')
    s = s.replace('(1,44×10⁻⁸³)', '(1,44×10⁻⁸)³')
    s = s.replace('%V=VVthucmol×100%', '%V = (V_thực / V_mol)×100%')
    s = s.replace('Vmol=MD=', 'V_mol = M/D = ')
    s = s.replace('65.100(-x)', '65(100-x)')
    s = s.replace('\xa0', ' ')
    s = re.sub(r'^[\uf0b7\uf0a7\uf0d8\u2022\-\+]\s*', '', s).strip()
    return s


def get_para_text_full(p, doc_part):
    parts = []
    for child in p._p.iterchildren():
        tag = child.tag.split('}')[-1]
        if tag == 'r':
            r_fonts = child.xpath('./*[local-name()="rPr"]/*[local-name()="rFonts"]')
            is_symbol_run = False
            if r_fonts:
                fa = r_fonts[0].get(qn('w:ascii')) or ''
                fh = r_fonts[0].get(qn('w:hAnsi')) or ''
                if 'Symbol' in fa or 'Symbol' in fh:
                    is_symbol_run = True
            va = child.xpath('./*[local-name()="rPr"]/*[local-name()="vertAlign"]')
            va_val = va[0].get(qn('w:val')) if va else None

            for rchild in child.iterchildren():
                rtag = rchild.tag.split('}')[-1]
                if rtag == 't' and rchild.text:
                    txt = rchild.text
                    if is_symbol_run:
                        sym_map = {'a': 'α', 'b': 'β', 'g': 'γ', '®': '→', 'Û': '⇔', 'Þ': '⇒', '»': '≈', '´': '×', '-': '-'}
                        txt = ''.join(sym_map.get(c, c) for c in txt)
                    if va_val == 'subscript' and txt.strip():
                        parts.append(f"~{txt.strip()}~")
                    elif va_val == 'superscript' and txt.strip():
                        parts.append(f"^{txt.strip()}^")
                    else:
                        parts.append(txt)
                elif rtag == 'tab':
                    parts.append('\t')
                elif rtag in ('object', 'pict'):
                    imd = rchild.xpath('.//*[local-name()="imagedata"]')
                    if imd:
                        rid = imd[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                        if rid and rid in doc_part.related_parts:
                            part = doc_part.related_parts[rid]
                            if part.partname.endswith('.wmf'):
                                eq_str = decode_wmf_eq(part.blob)
                                if eq_str:
                                    parts.append(eq_str)
                            elif part.partname.endswith('image147.png') or len(part.blob) < 350:
                                parts.append('≈')
        elif tag in ('oMath', 'oMathPara'):
            for mnode in child.iter():
                if mnode.tag.endswith('}t') and mnode.text:
                    parts.append(mnode.text)
    return clean_extracted_text("".join(parts))


def extract_real_images_from_el(el, doc_part):
    blobs = []
    for blip in el.xpath('.//*[local-name()="blip"]'):
        rid = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
        if rid and rid in doc_part.related_parts:
            part = doc_part.related_parts[rid]
            if not part.partname.endswith('.wmf') and len(part.blob) >= 500:
                blobs.append(part.blob)
    for pict in el.xpath('.//*[local-name()="pict"]'):
        for imd in pict.xpath('.//*[local-name()="imagedata"]'):
            rid = imd.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
            if rid and rid in doc_part.related_parts:
                part = doc_part.related_parts[rid]
                if not part.partname.endswith('.wmf') and len(part.blob) >= 500:
                    blobs.append(part.blob)
    return blobs


def norm_for_compare(s):
    if not s:
        return ""
    sub_map = str.maketrans("₀₁₂₃₄₅₆₇₈₉⁰¹²³⁴⁵⁶⁷⁸⁹−–—", "01234567890123456789---")
    s = s.translate(sub_map)
    s = re.sub(r'[\~\^]', '', s)
    s = re.sub(r'^Câu\s+\d+[\.\:]?\s*', '', s, flags=re.IGNORECASE)
    s = re.sub(r'^\[[^\]]+\]\s*', '', s)
    s = s.lower()
    s = re.sub(r'[\s\.,;:\-\(\)\[\]\'\"“”\?\!]+', ' ', s).strip()
    return s


def classify_p2_item_level_hoa10(stmt_text):
    t = stmt_text.lower()
    if any(k in t for k in [
        'nguyên tử khối trung bình', 'tổng số hạt', 'phần trăm khối lượng',
        'bất đẳng thức', 'phóng to', 'sân vận động', 'thay thế', 'thay vì',
        'tiêm doping', 'mẫu nước tiểu', '2,479', 'avogadro', 'hấp thụ dư',
        'khối lượng thực tế', 'khối lượng tuyệt đối', 'nhẩm nhanh', 'phương trình'
    ]) or re.search(r'\d+[,\.]\d+', t):
        return "VD"
    if any(k in t for k in [
        'xu hướng', 'nhường', 'nhận thêm', 'độc thân', 'cấu hình electron',
        'bão hòa', 'tính chất hóa học giống', 'lệch về', 'va chạm',
        'trung hòa về điện', 'bằng một nửa', 'tương ứng với', 'quy tắc hund'
    ]):
        return "TH"
    return "NB"


def classify_question_level_hoa10(part, sec, topic, stem_text, opt_text, sol_text, item_levels=None):
    full = (stem_text + " " + opt_text).lower()
    if sec in ("DANG_1", "DANG_2", "DANG_3"):
        if sec == "DANG_3":
            return "VD"
        if any(k in full for k in [
            "biện luận", "ion", "phân tử", "quặng pyrite", "hợp chất", "doping",
            "tổng số khối", "nhiều hơn", "tỉ lệ", "thể tích của 10 gam"
        ]):
            return "VD"
        return "TH"

    if part == "P2":
        vals = list((item_levels or {}).values())
        if vals.count("VD") >= 2:
            return "VD"
        if vals.count("NB") >= 3 and vals.count("VD") == 0:
            return "NB"
        return "TH"

    if part == "P3":
        if any(k in full for k in [
            "dịch truyền", "đám mây dông", "sấm sét", "ghép số thứ tự của các mệnh đề sai",
            "suy ngược ra đặc tính", "chất a:", "ngoại lệ về cấu hình"
        ]):
            return "VD"
        return "TH"

    # P1
    if any(k in full for k in [
        "số phát biểu đúng", "số phát biểu sai", "số phát biểu không đúng",
        "số hiện tượng vật lí là", "số hiện tượng hóa học là",
        "số cấu hình electron của nguyên tố kim loại", "điện tích lớp vỏ nguyên tử là",
        "hơn kém nhau là 2", "hơn kém nhau là 3", "(a)", "(1)"
    ]):
        if "số phát biểu" in full or "hơn kém nhau" in full or "điện tích lớp vỏ" in full or "số cấu hình" in full:
            return "VD"
        return "TH"

    if any(k in full for k in [
        "không đúng", "phát biểu nào sau đây đúng", "nhận định nào sau đây",
        "cùng một nguyên tố", "đồng vị của", "loại phân tử", "số hợp chất",
        "độc thân", "orbital chứa electron", "phân lớp bão hoà", "ion ", "anion ", "cation ",
        "fe (z = 26)", "zn (z = 30)", "cobalt", "3d", "rutherford – bohr"
    ]):
        return "TH"

    return "NB"


def infer_topic_for_ttb(stem_text, opt_text):
    full = (stem_text + " " + opt_text).lower()
    if any(k in full for k in [
        "cấu hình", "1s", "2s", "2p", "3s", "3p", "lớp ngoài cùng",
        "orbital", "độc thân", "phân mức năng lượng"
    ]):
        return "Bài 3: Cấu trúc lớp vỏ electron của nguyên tử"
    if any(k in full for k in [
        "đồng vị", "kí hiệu nguyên tử", "số khối", "số hiệu nguyên tử"
    ]):
        return "Bài 2: Nguyên tố hóa học"
    return "Bài 1: Thành phần của nguyên tử"


def extract_underlined_p1_answer(paragraphs):
    for p in paragraphs:
        full_t = p.text
        if not re.search(r'[A-D]\.', full_t):
            continue
        for r in p.runs:
            if r.underline and r.text.strip():
                m = re.search(r'([A-D])', r.text.strip())
                if m:
                    return m.group(1)
        for r_el in p._element.xpath('.//w:r'):
            u_el = r_el.xpath('./w:rPr/w:u')
            if u_el:
                val = u_el[0].get(qn('w:val'))
                if val is None or val != 'none':
                    rt = "".join(r_el.xpath('.//w:t/text()')).strip()
                    m = re.search(r'^([A-D])\b', rt)
                    if m:
                        return m.group(1)
    return None


def resolve_bai_toan_answer(sec, stem_text, sol_text):
    st = stem_text.lower()
    if "argon" in st or "agon" in st:
        return "a) 39,99; b) 6,1975 L"
    if "magnesium có ba đồng vị" in st:
        return "24Mg: 79%; 25Mg: 10%"
    if "copper (đồng) được sử dụng làm dây dẫn" in st and "6,3546" in st:
        return "a) 72,7%; b) 0,0727 mol & 0,0273 mol; c) 28,7%"
    if "hclox" in st or "26,12%" in st:
        return "HClO4"
    if "sao neutron" in st:
        return "a) 3,97.10^17 kg/m3; b) 1660,5 tấn"
    if "calcium là một" in st and "1,55" in st:
        return "1,96 Å"
    if "iron là một nguyên tố" in st and "7,87" in st:
        return "1,29.10^-8 cm"
    if "magnesium là kim loại nhẹ" in st and "1,74" in st:
        return "1,53.10^-8 cm"
    if "copper là kim loại dẫn điện" in st and "8,96" in st:
        return "1,28.10^-8 cm"
    if "titanium là kim loại" in st and "147 pm" in st:
        return "4,43 g/cm3"
    if "zinc là kim loại" in st and "134 pm" in st:
        return "7,38 g/cm3"
    if "gold (au)" in st and "19,36" in st:
        return "74%"
    # Fallback regex on sol_text
    for ln in reversed(sol_text.splitlines()):
        m = re.search(r'(?:Đáp\s+(?:án|số)|Kết\s+luận)\s*[:\.]?\s*([^\n]+)', ln, flags=re.IGNORECASE)
        if m:
            raw_a = m.group(1).strip().rstrip('.')
            m_num = re.match(r'^([+-]?\d+(?:[\,\.]\d+)?)\b', raw_a)
            if m_num and len(raw_a) < 12:
                return m_num.group(1).replace('.', ',')
            return raw_a
    return ""


def parse_cvl_hoa10_c1():
    doc = docx.Document(CVL_DOCX)
    cau_nids = get_cau_num_ids(doc)
    blocks = list(iter_block_items(doc))

    start_idx = 0
    end_idx = len(blocks)
    for i, (bt, b) in enumerate(blocks):
        if bt == 'p':
            t = get_para_text_full(b, doc.part)
            if t == 'HỆ THỐNG BÀI TẬP' and start_idx == 0:
                start_idx = i
            elif t == 'MỤC LỤC' and i > start_idx:
                end_idx = i
                break

    cur_topic = "Mở đầu: Đối tượng và phương pháp nghiên cứu Hóa học"
    cur_sec = ""
    cur_part = "P1"
    questions = []
    cur_q = None

    def flush_q():
        nonlocal cur_q
        if cur_q:
            if cur_q["sec"] == "I":
                if len(cur_q["p2_items"]) == 4:
                    cur_q["part"] = "P2"
                    questions.append(cur_q)
            else:
                questions.append(cur_q)
            cur_q = None

    for i in range(start_idx + 1, end_idx):
        bt, b = blocks[i]
        if bt == 'p':
            t = get_para_text_full(b, doc.part)
            imgs = extract_real_images_from_el(b._p, doc.part)
            if not t and not imgs:
                continue

            if t == "MỞ ĐẦU":
                flush_q()
                cur_topic = "Mở đầu: Đối tượng và phương pháp nghiên cứu Hóa học"
                continue
            elif t.startswith("Bài 1:"):
                flush_q()
                cur_topic = "Bài 1: Thành phần của nguyên tử"
                continue
            elif t.startswith("Bài 2:"):
                flush_q()
                cur_topic = "Bài 2: Nguyên tố hóa học"
                continue
            elif t.startswith("Bài 3:"):
                flush_q()
                cur_topic = "Bài 3: Cấu trúc lớp vỏ electron của nguyên tử"
                continue
            elif t == "CÁC DẠNG BÀI TOÁN":
                flush_q()
                cur_sec = "BAI_TOAN"
                cur_part = "P3"
                continue
            elif t.startswith("DẠNG 1:"):
                flush_q()
                cur_topic = "Bài 1: Thành phần của nguyên tử"
                cur_sec = "DANG_1"
                cur_part = "P3"
                continue
            elif t.startswith("DẠNG 2:"):
                flush_q()
                cur_topic = "Bài 2: Nguyên tố hóa học"
                cur_sec = "DANG_2"
                cur_part = "P3"
                continue
            elif t.startswith("DẠNG 3:"):
                flush_q()
                cur_topic = "Bài 1: Thành phần của nguyên tử"
                cur_sec = "DANG_3"
                cur_part = "P3"
                continue

            if t.startswith("I. BÀI TẬP ÔN TẬP LÝ THUYẾT"):
                flush_q()
                cur_sec = "I"
                cur_part = "P2"
                continue
            elif t.startswith("II. BÀI TẬP TRẮC NGHIỆM"):
                flush_q()
                cur_sec = "II"
                cur_part = "P1"
                continue
            elif t.startswith("III. BÀI TẬP TRẮC NGHIỆM ĐÚNG SAI"):
                flush_q()
                cur_sec = "III_DS"
                cur_part = "P2"
                continue
            elif t.startswith("III. BÀI TẬP TRẢ LỜI NGẮN") or t.startswith("IV. BÀI TẬP TRẢ LỜI NGẮN"):
                flush_q()
                cur_sec = "TLN"
                cur_part = "P3"
                continue
            elif re.match(r'^[123]\.\s+[A-ZĐÁÀẢÃẠĂẮẰẲẴẶÂẤẦẨẪẬÉÈẺẼẸÊẾỀỂỄỆÍÌỈĨỊÓÒỎÕỌÔỐỒỔỖỘƠỚỜỞỠỢÚÙỦŨỤƯỨỪỬỮỰÝỲỶỸỴ]', t):
                flush_q()
                continue

            nids = b._p.xpath('.//w:numPr/w:numId/@w:val')
            is_cau = bool(nids and nids[0] in cau_nids)
            m_text_cau = re.match(r'^Câu\s+\d+[\.\:]\s*(.*)', t)

            if is_cau or m_text_cau:
                flush_q()
                stem_first = m_text_cau.group(1).strip() if m_text_cau else t
                cur_q = {
                    "topic": cur_topic,
                    "sec": cur_sec,
                    "part": cur_part,
                    "source_tag": "[THPT Châu Văn Liêm – 2026]",
                    "stem_lines": [stem_first] if stem_first else [],
                    "opt_lines": [],
                    "p2_items": {},
                    "sol_lines": [],
                    "underlined_opt": None,
                    "in_sol": False,
                    "img_blobs": list(imgs),
                    "tables": [],
                }
            elif cur_q is not None:
                if imgs and not cur_q["in_sol"]:
                    cur_q["img_blobs"].extend(imgs)
                if not t:
                    continue
                if t.startswith(("📝", "Hướng dẫn giải", "Đáp án và giải thích", "Đáp án:", "Đáp án đúng:")):
                    cur_q["in_sol"] = True
                    cur_q["sol_lines"].append(t)
                elif cur_q["in_sol"]:
                    cur_q["sol_lines"].append(t)
                else:
                    if cur_part == "P1" and (re.match(r'^[A-D]\.', t) or (len(imgs) == 4 and 'A.' in t)):
                        if len(imgs) == 4 and 'A.' in t and 'B.' in t:
                            cur_q["opt_lines"].append("A. [↑↑]\tB. [↑][↑][↑]\tC. [↑↓][↑↓][   ]\tD. [↑][↓][↑]")
                        else:
                            cur_q["opt_lines"].append(t)
                        u_ans = extract_underlined_p1_answer([b])
                        if u_ans:
                            cur_q["underlined_opt"] = u_ans
                    elif cur_part == "P2" and re.match(r'^[a-d]\)', t):
                        k = t[0].lower()
                        cur_q["p2_items"][k] = t[2:].strip()
                    else:
                        cur_q["stem_lines"].append(t)

        elif bt == 'tbl':
            if len(b.rows) == 1 and len(b.columns) == 1:
                cell = b.cell(0, 0)
                has_q = False
                for cp in cell.paragraphs:
                    nids = cp._p.xpath('.//w:numPr/w:numId/@w:val')
                    ct = get_para_text_full(cp, doc.part)
                    if (nids and nids[0] in cau_nids) or re.match(r'^Câu\s+\d+[\.\:]', ct):
                        has_q = True
                        break
                if has_q:
                    for cp in cell.paragraphs:
                        nids = cp._p.xpath('.//w:numPr/w:numId/@w:val')
                        ct = get_para_text_full(cp, doc.part)
                        if not ct:
                            continue
                        if (nids and nids[0] in cau_nids) or re.match(r'^Câu\s+\d+[\.\:]', ct):
                            flush_q()
                            ct_clean = re.sub(r'^Câu\s+\d+[\.\:]\s*', '', ct).strip()
                            cur_q = {
                                "topic": cur_topic,
                                "sec": cur_sec,
                                "part": cur_part,
                                "source_tag": "[THPT Châu Văn Liêm – 2026]",
                                "stem_lines": [ct_clean],
                                "opt_lines": [],
                                "p2_items": {},
                                "sol_lines": [],
                                "underlined_opt": None,
                                "in_sol": False,
                                "img_blobs": [],
                                "tables": [],
                            }
                        elif cur_q is not None:
                            if ct.startswith(("📝", "Hướng dẫn giải", "Đáp án")):
                                cur_q["in_sol"] = True
                                cur_q["sol_lines"].append(ct)
                            elif cur_q["in_sol"]:
                                cur_q["sol_lines"].append(ct)
                            else:
                                cur_q["stem_lines"].append(ct)
                continue

            if cur_q is not None and not cur_q["in_sol"]:
                if len(b.rows) == 1 and len(b.columns) == 2:
                    for c in b.rows[0].cells:
                        for cp in c.paragraphs:
                            ct = get_para_text_full(cp, doc.part)
                            if ct:
                                cur_q["stem_lines"].append(ct)
                else:
                    t_data = []
                    for r in b.rows:
                        row_vals = [get_para_text_full(c.paragraphs[0], doc.part) if c.paragraphs else "" for c in r.cells]
                        if any(row_vals):
                            t_data.append(row_vals)
                    if t_data:
                        cur_q["tables"].append(t_data)

    flush_q()
    return questions


def parse_ttb_hoa10_c1():
    if not os.path.exists(TTB_DOCX):
        return []
    doc = docx.Document(TTB_DOCX)
    ttb_qs = []
    cur_part = "P1"
    cur_lvl = "NB"
    cur_q = None

    def flush_ttb():
        nonlocal cur_q
        if cur_q:
            ttb_qs.append(cur_q)
            cur_q = None

    for i in range(2, 109):
        p = doc.paragraphs[i]
        t = get_para_text_full(p, doc.part)
        imgs = extract_real_images_from_el(p._p, doc.part)
        if not t and not imgs:
            continue
        if "Trắc nghiệm nhiều phương án" in t:
            flush_ttb()
            cur_part = "P1"
            continue
        elif "Trắc nghiệm đúng" in t:
            flush_ttb()
            cur_part = "P2"
            continue
        elif "Trắc nghiệm trả lời ngắn" in t:
            flush_ttb()
            cur_part = "P3"
            continue
        elif "MỨC ĐỘ BIẾT" in t:
            flush_ttb()
            cur_lvl = "NB"
            continue
        elif "MỨC ĐỘ HIỂU" in t:
            flush_ttb()
            cur_lvl = "TH"
            continue
        elif "MỨC ĐỘ VẬN DỤNG" in t:
            flush_ttb()
            cur_lvl = "VD"
            continue

        m_q = re.match(r'^Câu\s+(\d+)[\.\:]\s*(.*)', t)
        if m_q:
            flush_ttb()
            rest = m_q.group(2).strip()
            src_m = re.match(r'^(\[[^\]]+\])\s*(.*)', rest)
            src_tag = src_m.group(1) if src_m else "[TTB – 2026]"
            stem_1 = src_m.group(2).strip() if src_m else rest
            cur_q = {
                "topic": "",
                "sec": "TTB",
                "part": cur_part,
                "level_override": cur_lvl,
                "source_tag": src_tag,
                "stem_lines": [stem_1] if stem_1 else [],
                "opt_lines": [],
                "p2_items": {},
                "sol_lines": [],
                "underlined_opt": None,
                "in_sol": False,
                "img_blobs": list(imgs),
                "tables": [],
            }
        elif cur_q is not None:
            if imgs and not cur_q["in_sol"]:
                cur_q["img_blobs"].extend(imgs)
            if not t:
                continue
            if t.startswith("Hướng dẫn giải"):
                cur_q["in_sol"] = True
                continue
            if cur_q["in_sol"]:
                cur_q["sol_lines"].append(t)
            else:
                if cur_part == "P1" and re.match(r'^[A-D]\.', t):
                    cur_q["opt_lines"].append(t)
                    u_ans = extract_underlined_p1_answer([p])
                    if u_ans:
                        cur_q["underlined_opt"] = u_ans
                else:
                    cur_q["stem_lines"].append(t)

    flush_ttb()
    for q in ttb_qs:
        q["topic"] = infer_topic_for_ttb(" ".join(q["stem_lines"]), " ".join(q["opt_lines"]))
    return ttb_qs


def finalize_question_record(raw_q):
    sol_text = "\n".join(raw_q["sol_lines"]).strip()
    sol_text = re.sub(r'^📝\s*Hướng dẫn giải:?\s*', '', sol_text).strip()
    sol_text = re.sub(r'^Hướng dẫn giải:?\s*', '', sol_text).strip()
    sol_text = re.sub(r'^Đáp án và giải thích:?\s*', '', sol_text).strip()

    ans = ""
    if raw_q["part"] == "P1":
        ans = raw_q["underlined_opt"] or ""
        if not ans:
            m = re.search(r'Đáp\s+(?:án|số)\s*[:\.]?\s*([A-D])\b', sol_text, flags=re.IGNORECASE)
            if m:
                ans = m.group(1).upper()
        # Bổ sung thủ công: TTB Câu 3 (Al: p = 13, n = 14 → D)
        if not ans and "Aluminium" in " ".join(raw_q["stem_lines"]) and "n = 27-13=14" in sol_text.replace(" ", " "):
            ans = "D"
    elif raw_q["part"] == "P2":
        p2_ans = {}
        m_inline = re.findall(r'([a-d])\s*[\-\:]\s*(ĐÚNG|SAI|Đúng|Sai)', sol_text)
        for k_m, v_m in m_inline:
            p2_ans[k_m.lower()] = "Đ" if v_m.upper() == "ĐÚNG" else "S"
        for ln in raw_q["sol_lines"]:
            m = re.match(r'^([a-d])\)\s*(Đúng|Sai)', ln.strip(), flags=re.IGNORECASE)
            if m:
                p2_ans[m.group(1).lower()] = "Đ" if m.group(2).lower() == "đúng" else "S"
        if len(p2_ans) == 4:
            ans = f"a:{p2_ans['a']} - b:{p2_ans['b']} - c:{p2_ans['c']} - d:{p2_ans['d']}"
    elif raw_q["part"] == "P3":
        if raw_q["sec"] in ("DANG_1", "DANG_2", "DANG_3"):
            ans = resolve_bai_toan_answer(raw_q["sec"], " ".join(raw_q["stem_lines"]), sol_text)
        else:
            for ln in raw_q["sol_lines"]:
                m = re.search(r'Đáp\s+(?:án|số)\s*[:\.]?\s*([^\n\.]+)', ln, flags=re.IGNORECASE)
                if m:
                    raw_a = m.group(1).strip()
                    m_num = re.match(r'^([+-]?\d+(?:[\,\.]\d+)?)', raw_a)
                    if m_num:
                        ans = m_num.group(1).replace('.', ',')
                    else:
                        ans = raw_a
                    break

    item_levels = {}
    if raw_q["part"] == "P2":
        for k in ["a", "b", "c", "d"]:
            item_levels[k] = classify_p2_item_level_hoa10(raw_q["p2_items"].get(k, ""))

    lvl = raw_q.get("level_override") or classify_question_level_hoa10(
        raw_q["part"],
        raw_q["sec"],
        raw_q["topic"],
        " ".join(raw_q["stem_lines"]),
        " ".join(raw_q["opt_lines"]),
        sol_text,
        item_levels
    )

    return {
        "part": raw_q["part"],
        "topic": raw_q["topic"],
        "level": lvl,
        "item_levels": item_levels,
        "source_tag": raw_q["source_tag"],
        "stem_lines": raw_q["stem_lines"],
        "opt_lines": raw_q["opt_lines"],
        "p2_items": raw_q["p2_items"],
        "answer": ans,
        "solution": sol_text,
        "img_blobs": raw_q["img_blobs"],
        "tables": raw_q["tables"],
    }


# ==============================================================================
# WORD FORMATTING HELPERS (100% CHUẨN "FORM ĐỀ KIỂM TRA")
# ==============================================================================

UNICODE_SUP_TO_ASCII = str.maketrans('⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿˣʸᶻᵃᵇ', '0123456789+-=()nxyzab')
UNICODE_SUB_TO_ASCII = str.maketrans('₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎ₙₓᵧ₂ₐ₆', '0123456789+-=()nxyzab')


def normalize_chem_markers_hoa10(text):
    if not text:
        return ""
    s = text
    # Convert Unicode superscripts and subscripts to ^...^ and ~...~
    s = re.sub(
        r'([⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿˣʸᶻᵃᵇ]+)',
        lambda m: f"^{m.group(1).translate(UNICODE_SUP_TO_ASCII)}^",
        s
    )
    s = re.sub(
        r'([₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎ₙₓᵧₐ]+)',
        lambda m: f"~{m.group(1).translate(UNICODE_SUB_TO_ASCII)}~",
        s
    )
    # Merge adjacent ^a^^b^ or ~a~~b~
    s = re.sub(r'\^\^', '', s)
    s = re.sub(r'\~\~', '', s)

    # Electron configuration subshells: 1s2, 2s2, 2p6, 3s2, 3p6, 3d10, 4s2, 4p5, 4f14, etc.
    s = re.sub(r'([1-7][spdf])(\d{1,2})(?=[1-7\s\.,;:\)\]\-]|$)', r'\1^\2^', s)
    # Standalone s2, p6, d10, f14, p4, d5, d9
    s = re.sub(r'\b([spdf])(\d{1,2})\b', r'\1^\2^', s)
    # Powers of 10: 10-19, 10-27, 10-31, 10^23, 102 pm, 10-2 pm, 104 pm
    s = re.sub(r'\b10([+\-–−]\d+)\b', r'10^\1^', s)
    s = re.sub(r'\b10(2|4)\s*(pm\b)', r'10^\1^ \2', s)
    # Units: cm3, m3, dm3, oC
    s = re.sub(r'\b(cm|dm|m)(2|3)\b', r'\1^\2^', s)
    s = re.sub(r'(\d+)\s*oC\b', r'\1^o^C', s)
    # Common ions: X2+, X2-, M2+, M3+, R+, R2+, O2-, Fe2+, Fe3+, Co2+, Co3+, Cr3+, Al3+, Cl-, F-, Na+, K+, Ca2+, Mg2+
    s = re.sub(r'\b(Fe|Al|Cr|Co|M|X|R)\s*([23][+\-–])', r'\1^\2^', s)
    s = re.sub(r'\b(Ca|Mg|Cu|Zn|Ba|O|S)\s*(2[+\-–])', r'\1^\2^', s)
    s = re.sub(r'\b(Na|K|Ag|H|Cl|F|Br|I|R|X|M)([+\-])(?=\s|[\.,;:\)]|$)', r'\1^\2^', s)
    # Chemical formulas with subscripts (exclude 1s..7f)
    def chem_sub_repl(m):
        prefix, num = m.group(1), m.group(2)
        if re.match(r'^[1-7]?[spdf]$', prefix):
            return f"{prefix}^{num}^"
        return f"{prefix}~{num}~"

    for _ in range(3):
        s = re.sub(r'([A-Za-z\)\]])(\d+)(?=[A-Za-z\(\[\)\]\s\.,;:\-\+\=⇌→\~]|$)', chem_sub_repl, s)

    return s


def add_chem_runs_hoa10(p, text, bold=False, italic=False, underline=False, color=None, size=11.5):
    if not text:
        return
    s = normalize_chem_markers_hoa10(text)
    tokens = re.split(r'(\~[^\~]+\~|\^[^\^]+\^)', s)
    for tok in tokens:
        if not tok:
            continue
        r = p.add_run()
        r.font.name = "Times New Roman"
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        r.underline = underline
        if color:
            r.font.color.rgb = color
        if tok.startswith('~') and tok.endswith('~') and len(tok) > 2:
            r.text = tok[1:-1]
            r.font.subscript = True
        elif tok.startswith('^') and tok.endswith('^') and len(tok) > 2:
            r.text = tok[1:-1]
            r.font.superscript = True
        else:
            r.text = tok


def parse_abcd_from_opt_lines(opt_lines):
    full = "\t".join(opt_lines)
    parts = re.split(r'(?:^|\t|\s{2,})([A-D])\.\s*', full)
    opts = {"A": "", "B": "", "C": "", "D": ""}
    if len(parts) >= 9:
        for i in range(1, len(parts) - 1, 2):
            k = parts[i].strip()
            v = parts[i + 1].strip()
            if k in opts:
                opts[k] = v
    else:
        for line in opt_lines:
            m = re.findall(r'([A-D])\.\s*([^\t]+?)(?=(?:\s+[A-D]\.\s)|$)', line)
            for k, v in m:
                opts[k] = v.strip()
    return opts


def add_options_paragraph_no_table(doc, opts, correct_ans, is_teacher):
    max_len = max((len(v) for v in opts.values()), default=0)
    if max_len <= 16:
        layout = 4
    elif max_len <= 38:
        layout = 2
    else:
        layout = 1

    def render_opt(p, key):
        is_corr = is_teacher and (key == correct_ans)
        r_lbl = p.add_run(f"{key}. ")
        r_lbl.bold = True
        r_lbl.underline = is_corr
        r_lbl.font.name = "Times New Roman"
        r_lbl.font.size = Pt(11.5)
        if is_corr:
            r_lbl.font.color.rgb = COLOR_CORRECT
        add_chem_runs_hoa10(
            p,
            opts.get(key, ""),
            bold=is_corr,
            color=(COLOR_CORRECT if is_corr else None),
            size=11.5
        )

    if layout == 4:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        ts = p.paragraph_format.tab_stops
        ts.add_tab_stop(Cm(4.8), WD_TAB_ALIGNMENT.LEFT)
        ts.add_tab_stop(Cm(9.3), WD_TAB_ALIGNMENT.LEFT)
        ts.add_tab_stop(Cm(13.8), WD_TAB_ALIGNMENT.LEFT)
        for idx, k in enumerate(["A", "B", "C", "D"]):
            if idx > 0:
                p.add_run("\t")
            render_opt(p, k)
    elif layout == 2:
        for pair in [("A", "B"), ("C", "D")]:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(9.3), WD_TAB_ALIGNMENT.LEFT)
            render_opt(p, pair[0])
            p.add_run("\t")
            render_opt(p, pair[1])
    else:
        for k in ["A", "B", "C", "D"]:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1.5)
            render_opt(p, k)


def add_data_table_to_doc(doc, tbl_data):
    if not tbl_data:
        return
    rows = len(tbl_data)
    cols = max(len(r) for r in tbl_data)
    tbl = doc.add_table(rows=rows, cols=cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl._tbl.tblPr.append(parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="8BC390"/>'
        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="8BC390"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="8BC390"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="8BC390"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="8BC390"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="8BC390"/>'
        f'</w:tblBorders>'
    ))
    for r_i, row in enumerate(tbl_data):
        for c_i, val in enumerate(row):
            if c_i < cols:
                cell = tbl.cell(r_i, c_i)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_i == 0 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(1.5)
                p.paragraph_format.space_after = Pt(1.5)
                add_chem_runs_hoa10(p, val, bold=(r_i == 0), size=10.5)


def add_solution_block_hoa10(doc, q):
    p_hd = doc.add_paragraph()
    p_hd.paragraph_format.left_indent = Cm(0.6)
    p_hd.paragraph_format.space_before = Pt(2)
    p_hd.paragraph_format.space_after = Pt(1)

    r_hd = p_hd.add_run("Hướng dẫn giải: ")
    r_hd.bold = True
    r_hd.font.name = "Times New Roman"
    r_hd.font.size = Pt(11)
    r_hd.font.color.rgb = COLOR_SOLUTION

    if q["part"] == "P1" and q["answer"]:
        r_ans = p_hd.add_run(f"Chọn {q['answer']}.")
        r_ans.bold = True
        r_ans.font.name = "Times New Roman"
        r_ans.font.size = Pt(11)
        r_ans.font.color.rgb = COLOR_CORRECT
    elif q["part"] == "P2" and q["answer"]:
        r_ans = p_hd.add_run(f"Đáp án: {q['answer']}")
        r_ans.bold = True
        r_ans.font.name = "Times New Roman"
        r_ans.font.size = Pt(11)
        r_ans.font.color.rgb = COLOR_CORRECT
    elif q["part"] == "P3" and q["answer"]:
        r_ans = p_hd.add_run("Đáp án: ")
        r_ans.bold = True
        r_ans.font.name = "Times New Roman"
        r_ans.font.size = Pt(11)
        r_ans.font.color.rgb = COLOR_CORRECT
        add_chem_runs_hoa10(p_hd, str(q["answer"]), bold=True, color=COLOR_CORRECT, size=11)

    sol = (q.get("solution") or "").strip()
    if sol:
        for line in sol.splitlines():
            line_s = line.strip()
            if not line_s:
                continue
            if re.match(r'^Đáp\s+(?:án|số)\s*[:\.]?\s*[A-D0-9\,\.\s]+$', line_s, flags=re.IGNORECASE):
                continue
            p_ln = doc.add_paragraph()
            p_ln.paragraph_format.left_indent = Cm(0.6)
            p_ln.paragraph_format.space_before = Pt(0.5)
            p_ln.paragraph_format.space_after = Pt(1.5)
            add_chem_runs_hoa10(p_ln, line_s, size=11)


def setup_page_and_hf(doc):
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11.5)

    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)
        section.header_distance = Cm(0.6)
        section.footer_distance = Cm(0.6)
        section.different_first_page_header_footer = False

        header = section.header
        header.is_linked_to_previous = False
        for p in list(header.paragraphs):
            p._element.getparent().remove(p._element)
        htbl = header.add_table(rows=1, cols=2, width=Cm(18.0))
        htbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        htbl._tbl.tblPr.append(parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            f'<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="4" w:color="8BC390"/>'
            f'<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            f'<w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            f'<w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            f'</w:tblBorders>'
        ))
        c0, c1 = htbl.cell(0, 0), htbl.cell(0, 1)
        c0.width, c1.width = Cm(9.0), Cm(9.0)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run("Tài liệu lưu hành nội bộ")
        r0.italic = True
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(10)
        r0.font.color.rgb = COLOR_HF_RED

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_after = Pt(2)
        r1a = p1.add_run("Thầy ")
        r1b = p1.add_run("TRẦN VĂN THẠNH")
        r1b.bold = True
        r1c = p1.add_run(" - 0777.470.803")
        for r in (r1a, r1b, r1c):
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.font.color.rgb = COLOR_HF_RED

        footer = section.footer
        footer.is_linked_to_previous = False
        for p in list(footer.paragraphs):
            p._element.getparent().remove(p._element)
        ftbl = footer.add_table(rows=1, cols=2, width=Cm(18.0))
        ftbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        ftbl._tbl.tblPr.append(parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="4" w:color="8BC390"/>'
            f'<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            f'<w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            f'<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            f'<w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            f'<w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            f'</w:tblBorders>'
        ))
        fc0, fc1 = ftbl.cell(0, 0), ftbl.cell(0, 1)
        fc0.width, fc1.width = Cm(15.0), Cm(3.0)
        fp0 = fc0.paragraphs[0]
        fp0.paragraph_format.space_before = Pt(2)
        fr0 = fp0.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
        fr0.italic = True
        fr0.font.name = "Times New Roman"
        fr0.font.size = Pt(9.5)
        fr0.font.color.rgb = COLOR_HF_RED

        fp1 = fc1.paragraphs[0]
        fp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp1.paragraph_format.space_before = Pt(2)
        fr1 = fp1.add_run("Trang ")
        fr1.bold = True
        fr1.font.name = "Times New Roman"
        fr1.font.size = Pt(10)
        fr1.font.color.rgb = COLOR_HF_RED
        fp1._p.append(parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w')))


def build_word_bank_hoa10_c1(questions, is_teacher):
    doc = docx.Document()
    setup_page_and_hf(doc)

    # 1. Title Block
    p_t1 = doc.add_paragraph()
    p_t1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t1.paragraph_format.space_before = Pt(0)
    p_t1.paragraph_format.space_after = Pt(2)
    r_t1 = p_t1.add_run(f"NGÂN HÀNG CÂU HỎI HÓA HỌC 10 – {CHAPTER_INFO['name']}")
    r_t1.bold = True
    r_t1.font.name = "Times New Roman"
    r_t1.font.size = Pt(14)
    r_t1.font.color.rgb = COLOR_PRIMARY

    p_t2 = doc.add_paragraph()
    p_t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t2.paragraph_format.space_before = Pt(0)
    p_t2.paragraph_format.space_after = Pt(2)
    r_t2 = p_t2.add_run("PHÂN LOẠI THEO CHỦ ĐỀ, DẠNG THỨC & MỨC ĐỘ NHẬN THỨC (BIẾT – HIỂU – VẬN DỤNG)")
    r_t2.bold = True
    r_t2.font.name = "Times New Roman"
    r_t2.font.size = Pt(10.5)
    r_t2.font.color.rgb = COLOR_SOLUTION

    p_t3 = doc.add_paragraph()
    p_t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t3.paragraph_format.space_before = Pt(0)
    p_t3.paragraph_format.space_after = Pt(6)
    ver_note = (
        "(BẢN GIÁO VIÊN — CÓ ĐÁP ÁN, LỜI GIẢI CHI TIẾT & MÃ ĐỊNH DANH MA TRẬN)"
        if is_teacher else
        "(BẢN HỌC SINH — ĐỀ ÔN LUYỆN PHÂN LOẠI THEO MỨC ĐỘ NHẬN THỨC)"
    )
    r_t3 = p_t3.add_run(ver_note)
    r_t3.bold = True
    r_t3.italic = True
    r_t3.font.name = "Times New Roman"
    r_t3.font.size = Pt(11)
    r_t3.font.color.rgb = COLOR_HF_RED

    # 2. Matrix Summary Table
    tbl = doc.add_table(rows=1, cols=7)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl._tbl.tblPr.append(parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="8BC390"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="8BC390"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="8BC390"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="8BC390"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="8BC390"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="8BC390"/>'
        f'</w:tblBorders>'
    ))
    headers = ["Chủ đề / Bài học", "Dạng thức", "Biết (NB)", "Hiểu (TH)", "Vận dụng (VD)", "Tổng số câu", "Phạm vi câu"]
    for i, h in enumerate(headers):
        p = tbl.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.font.color.rgb = COLOR_PRIMARY

    for part_code, part_label in [
        ("P1", "Phần I (TN 4 lựa chọn)"),
        ("P2", "Phần II (Đúng / Sai)"),
        ("P3", "Phần III (Trả lời ngắn & BT)")
    ]:
        for topic in CHAPTER_INFO["topics"]:
            sub_qs = [q for q in questions if q["part"] == part_code and q["topic"] == topic]
            if not sub_qs:
                continue
            row_cells = tbl.add_row().cells
            nb = sum(1 for q in sub_qs if q["level"] == "NB")
            th = sum(1 for q in sub_qs if q["level"] == "TH")
            vd = sum(1 for q in sub_qs if q["level"] == "VD")
            q_nums = [q["q_num"] for q in sub_qs]
            q_range = f"Câu {min(q_nums)} – {max(q_nums)}" if len(q_nums) > 1 else f"Câu {q_nums[0]}"
            if part_code == "P2":
                it_nb = sum(1 for q in sub_qs for lv in q.get("item_levels", {}).values() if lv == "NB")
                it_th = sum(1 for q in sub_qs for lv in q.get("item_levels", {}).values() if lv == "TH")
                it_vd = sum(1 for q in sub_qs for lv in q.get("item_levels", {}).values() if lv == "VD")
                nb_str = f"{nb} câu ({it_nb} ý)"
                th_str = f"{th} câu ({it_th} ý)"
                vd_str = f"{vd} câu ({it_vd} ý)"
            else:
                nb_str, th_str, vd_str = str(nb), str(th), str(vd)

            vals = [topic, part_label, nb_str, th_str, vd_str, f"{len(sub_qs)} câu", q_range]
            for i, val in enumerate(vals):
                p = row_cells[i].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i in (0, 1) else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(1.5)
                p.paragraph_format.space_after = Pt(1.5)
                r = p.add_run(val)
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.5)
                if i == 5:
                    r.bold = True

    tot_cells = tbl.add_row().cells
    tot_nb = sum(1 for q in questions if q["level"] == "NB")
    tot_th = sum(1 for q in questions if q["level"] == "TH")
    tot_vd = sum(1 for q in questions if q["level"] == "VD")
    tot_vals = [
        "TỔNG CỘNG TOÀN CHƯƠNG 1",
        "3 Phần (I, II, III)",
        f"{tot_nb} câu",
        f"{tot_th} câu",
        f"{tot_vd} câu",
        f"{len(questions)} câu",
        f"Câu 1 – {len(questions)}"
    ]
    for i, val in enumerate(tot_vals):
        p = tot_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 1 else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(val)
        r.bold = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.font.color.rgb = COLOR_HF_RED

    # 3. Render Questions by Part & Topic
    part_headers = {
        "P1": "PHẦN I. CÂU TRẮC NGHIỆM NHIỀU PHƯƠNG ÁN LỰA CHỌN",
        "P2": "PHẦN II. CÂU TRẮC NGHIỆM ĐÚNG / SAI (4 Ý a, b, c, d)",
        "P3": "PHẦN III. CÂU TRẮC NGHIỆM TRẢ LỜI NGẮN & BÀI TOÁN VẬN DỤNG"
    }
    cur_part = None
    cur_topic = None

    for q in questions:
        if q["part"] != cur_part:
            cur_part = q["part"]
            cur_topic = None
            p_pt = doc.add_paragraph()
            p_pt.paragraph_format.space_before = Pt(10)
            p_pt.paragraph_format.space_after = Pt(4)
            r_pt = p_pt.add_run(part_headers[cur_part])
            r_pt.bold = True
            r_pt.font.name = "Times New Roman"
            r_pt.font.size = Pt(12.5)
            r_pt.font.color.rgb = COLOR_HF_RED

        if q["topic"] != cur_topic:
            cur_topic = q["topic"]
            p_tp = doc.add_paragraph()
            p_tp.paragraph_format.space_before = Pt(6)
            p_tp.paragraph_format.space_after = Pt(3)
            r_tp = p_tp.add_run(f"♦ CHỦ ĐỀ: {cur_topic.upper()}")
            r_tp.bold = True
            r_tp.font.name = "Times New Roman"
            r_tp.font.size = Pt(11.5)
            r_tp.font.color.rgb = COLOR_PRIMARY

        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)

        r_qn = p_q.add_run(f"Câu {q['q_num']}. ")
        r_qn.bold = True
        r_qn.font.name = "Times New Roman"
        r_qn.font.size = Pt(11.5)
        r_qn.font.color.rgb = COLOR_PRIMARY

        lv_name = LEVEL_NAMES[q["level"]]
        ans_str = f" | ĐA: {q['answer']}" if (is_teacher and q["answer"]) else ""
        r_tag = p_q.add_run(f"[{q['id']} | {lv_name}{ans_str}] ")
        r_tag.bold = True
        r_tag.font.name = "Times New Roman"
        r_tag.font.size = Pt(10.5)
        r_tag.font.color.rgb = COLOR_TAG

        if q.get("source_tag") and "Châu Văn Liêm" not in q["source_tag"]:
            r_src = p_q.add_run(f"{q['source_tag']} ")
            r_src.bold = True
            r_src.font.name = "Times New Roman"
            r_src.font.size = Pt(11)
            r_src.font.color.rgb = COLOR_SOLUTION

        if q["stem_lines"]:
            add_chem_runs_hoa10(p_q, q["stem_lines"][0], size=11.5)
            for extra_ln in q["stem_lines"][1:]:
                p_ex = doc.add_paragraph()
                p_ex.paragraph_format.left_indent = Cm(0.4)
                p_ex.paragraph_format.space_before = Pt(1)
                p_ex.paragraph_format.space_after = Pt(1)
                add_chem_runs_hoa10(p_ex, extra_ln, size=11.5)

        for img_p in q.get("image_paths", []):
            if os.path.exists(img_p):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(2)
                p_img.paragraph_format.space_after = Pt(2)
                w_cm = Cm(3.5) if len(q["image_paths"]) >= 4 else Cm(6.8)
                p_img.add_run().add_picture(img_p, width=w_cm)

        for tbl_d in q.get("tables", []):
            add_data_table_to_doc(doc, tbl_d)

        if q["part"] == "P1":
            opts = parse_abcd_from_opt_lines(q["opt_lines"])
            add_options_paragraph_no_table(doc, opts, q["answer"], is_teacher)
            if is_teacher:
                add_solution_block_hoa10(doc, q)

        elif q["part"] == "P2":
            ans_map = {}
            for part_ans in q["answer"].split("-"):
                if ":" in part_ans:
                    k_a, v_a = part_ans.strip().split(":", 1)
                    ans_map[k_a.strip()] = v_a.strip()
            for key in ["a", "b", "c", "d"]:
                stmt_txt = q.get("p2_items", {}).get(key, "")
                p_it = doc.add_paragraph()
                p_it.paragraph_format.left_indent = Cm(0.4)
                p_it.paragraph_format.space_before = Pt(1)
                p_it.paragraph_format.space_after = Pt(1.5)
                r_k = p_it.add_run(f"{key}) ")
                r_k.bold = True
                r_k.font.name = "Times New Roman"
                r_k.font.size = Pt(11.5)
                il_lv = q.get("item_levels", {}).get(key, "TH")
                r_mt = p_it.add_run(f"[{LEVEL_NAMES.get(il_lv, 'Hiểu')}] ")
                r_mt.bold = True
                r_mt.font.name = "Times New Roman"
                r_mt.font.size = Pt(10.5)
                r_mt.font.color.rgb = COLOR_SOLUTION
                add_chem_runs_hoa10(p_it, stmt_txt, size=11.5)
                if is_teacher and key in ans_map:
                    ds_val = "Đúng" if ans_map[key] == "Đ" else "Sai"
                    r_ds = p_it.add_run(f"   → [{ds_val}]")
                    r_ds.bold = True
                    r_ds.font.name = "Times New Roman"
                    r_ds.font.size = Pt(11)
                    r_ds.font.color.rgb = COLOR_CORRECT
            if is_teacher:
                add_solution_block_hoa10(doc, q)

        elif q["part"] == "P3":
            if not is_teacher:
                p_ans_line = doc.add_paragraph()
                p_ans_line.paragraph_format.left_indent = Cm(0.4)
                p_ans_line.paragraph_format.space_before = Pt(1)
                p_ans_line.paragraph_format.space_after = Pt(2)
                r_al = p_ans_line.add_run("Đáp số / Trả lời: .......................................................................................")
                r_al.italic = True
                r_al.font.name = "Times New Roman"
                r_al.font.size = Pt(11)
            else:
                add_solution_block_hoa10(doc, q)

    suffix = "GiaoVien" if is_teacher else "HocSinh"
    fname = f"{CHAPTER_INFO['file_prefix']}_{suffix}.docx"
    out_p1 = os.path.join(OUT_DIR, fname)
    out_p2 = os.path.join(BANK_DIR, fname)
    doc.save(out_p1)
    doc.save(out_p2)
    print(f"  + Saved {suffix} Word file: {out_p2}")


if __name__ == "__main__":
    raw_cvl = parse_cvl_hoa10_c1()
    raw_ttb = parse_ttb_hoa10_c1()
    print(f"Parsed CVL: {len(raw_cvl)} questions | Parsed TTB: {len(raw_ttb)} questions")

    all_raw = raw_cvl + raw_ttb
    deduped = []
    seen_norms = []
    dup_count = 0

    for rq in all_raw:
        fq = finalize_question_record(rq)
        st_norm = norm_for_compare(" ".join(fq["stem_lines"]))
        opt_norm = norm_for_compare(" ".join(fq["opt_lines"]))
        full_sig = f"{fq['part']}||{st_norm}||{opt_norm}"
        is_dup = False
        for prev_part, prev_st, prev_sig in seen_norms:
            if fq["part"] == prev_part:
                if full_sig == prev_sig:
                    is_dup = True
                    break
                if rq.get("sec") == "TTB" and (st_norm == prev_st or SequenceMatcher(None, st_norm, prev_st).ratio() >= 0.85):
                    is_dup = True
                    break
        if is_dup:
            dup_count += 1
        else:
            seen_norms.append((fq["part"], st_norm, full_sig))
            deduped.append(fq)

    print(f"Filtered out {dup_count} duplicate questions -> {len(deduped)} unique questions!")

    # Sort by Part (P1 -> P2 -> P3), then Topic order, then Level (NB -> TH -> VD)
    part_order = {"P1": 1, "P2": 2, "P3": 3}
    topic_order = {t: i for i, t in enumerate(CHAPTER_INFO["topics"])}
    level_order = {"NB": 1, "TH": 2, "VD": 3}

    deduped.sort(key=lambda q: (
        part_order.get(q["part"], 9),
        topic_order.get(q["topic"], 9),
        level_order.get(q["level"], 9)
    ))

    final_questions = []
    for q_num, q in enumerate(deduped, 1):
        qid = f"H10-C1-{q['part']}-{q['level']}-{q_num:03d}"
        img_paths = []
        for k, blob in enumerate(q.get("img_blobs", []), 1):
            fpath = os.path.join(IMG_DIR, f"h10_c1_q{q_num:03d}_img{k}.png")
            try:
                im = Image.open(io.BytesIO(blob))
                im.save(fpath, format="PNG")
                img_paths.append(fpath)
            except Exception:
                pass
        q_rec = {
            "q_num": q_num,
            "id": qid,
            "part": q["part"],
            "level": q["level"],
            "topic": q["topic"],
            "source_tag": q["source_tag"],
            "stem_lines": q["stem_lines"],
            "opt_lines": q["opt_lines"],
            "p2_items": q["p2_items"],
            "item_levels": q["item_levels"],
            "answer": q["answer"],
            "solution": q["solution"],
            "image_paths": img_paths,
            "tables": q["tables"],
        }
        final_questions.append(q_rec)

    bank_data = {
        "C1": {
            "grade": 10,
            "chapter_number": 1,
            "chapter_name": CHAPTER_INFO["name"],
            "topics": CHAPTER_INFO["topics"],
            "total_questions": len(final_questions),
            "questions": final_questions,
        }
    }

    json_c1_path = os.path.join(BANK_DIR, "Ngan_Hang_Hoa_10_Chuong_1_Cau_Tao_Nguyen_Tu.json")
    json_full_path = os.path.join(BANK_DIR, "Ngan_Hang_Hoa_10_Tron_Bo.json")
    with open(json_c1_path, "w", encoding="utf-8") as f:
        json.dump(bank_data, f, ensure_ascii=False, indent=2)
    with open(json_full_path, "w", encoding="utf-8") as f:
        json.dump(bank_data, f, ensure_ascii=False, indent=2)
    print(f"Saved JSON Question Bank ({len(final_questions)} questions): {json_c1_path}")

    build_word_bank_hoa10_c1(final_questions, is_teacher=False)
    build_word_bank_hoa10_c1(final_questions, is_teacher=True)

    print("\n=== THỐNG KÊ CHI TIẾT NGÂN HÀNG HÓA 10 - CHƯƠNG 1 ===")
    for pt in ["P1", "P2", "P3"]:
        sub = [q for q in final_questions if q["part"] == pt]
        nb = sum(1 for q in sub if q["level"] == "NB")
        th = sum(1 for q in sub if q["level"] == "TH")
        vd = sum(1 for q in sub if q["level"] == "VD")
        print(f"  + {pt}: {len(sub)} câu (Biết: {nb}, Hiểu: {th}, Vận dụng: {vd})")
