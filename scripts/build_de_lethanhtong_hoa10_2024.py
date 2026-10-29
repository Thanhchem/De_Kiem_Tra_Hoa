# -*- coding: utf-8 -*-
"""
Script tạo Đề Kiểm Tra Giữa Học Kì 1 Hóa Học 10 - THPT Lê Thánh Tông (2024 - 2025)
Chuẩn 100% FORM ĐỀ KIỂM TRA (Thầy TRẦN VĂN THẠNH)
Sinh cả 2 phiên bản:
1. Bản Học sinh: De_Kiem_Tra_GK1_Hoa_10_LeThanhTong_2025_HocSinh.docx
2. Bản Giáo viên: De_Kiem_Tra_GK1_Hoa_10_LeThanhTong_2025_GiaoVien.docx
"""

import os
import sys

if sys.stdout.encoding.lower() != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

sys.path.append(os.path.abspath(r"C:\Antigravity_Thanh\scripts"))
from form_de_kiem_tra import (
    setup_header_footer,
    build_title_block,
    add_section_header,
    render_question_prompt,
    render_options_no_table,
    render_statement,
    add_solution_no_table,
    add_student_writing_lines,
    render_rich_text,
    set_run_style,
    COLOR_PRIMARY,
    COLOR_CORRECT,
    COLOR_SOLUTION
)

IMG_DIR = r"C:\Antigravity_Thanh\San_Pham\images_gk1_hoa10_lethanhtong"

def add_centered_image(doc, img_path, width_cm=6.0):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(img_path, width=Cm(width_cm))

def build_exam_lethanhtong_10(is_teacher=False, output_path=None):
    doc = docx.Document()
    setup_header_footer(doc)

    # 1. Khối tiêu đề
    build_title_block(
        doc,
        school_line="SỞ GIÁO DỤC VÀ ĐÀO TẠO - THPT LÊ THÁNH TÔNG",
        exam_title="KIỂM TRA GIỮA HỌC KÌ 1 - NĂM HỌC 2024 – 2025",
        subject_line="MÔN: HÓA HỌC 10",
        num_pages="04",
        duration="45 phút",
        exam_code="103",
        atomic_masses="H = 1; C = 12; N = 14; O = 16; Na = 23; Mg = 24; Al = 27; Cl = 35,5; K = 39; Ca = 40; Ag = 108",
        is_teacher=is_teacher
    )

    # =========================================================================
    # PHẦN I
    # =========================================================================
    add_section_header(
        doc,
        "PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án."
    )

    # Câu 1
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 1. Trong nguyên tử, hạt mang điện là")
    opts1 = [
        "A. proton và neuton.",
        "*B. proton và electron.*" if is_teacher else "B. proton và electron.",
        "C. electron.",
        "D. electron và neutron."
    ]
    render_options_no_table(doc, opts1, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Trong nguyên tử: hạt proton mang điện tích dương (+1), hạt electron mang điện tích âm (-1), còn hạt neutron không mang điện.\n➔ Đáp án đúng: B.")

    # Câu 2
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 2. Số electron tối đa trong phân lớp d là")
    opts2 = [
        "A. 14.",
        "*B. 10.*" if is_teacher else "B. 10.",
        "C. 2.",
        "D. 6."
    ]
    render_options_no_table(doc, opts2, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Số electron tối đa trong các phân lớp: s chứa tối đa 2e; p chứa tối đa 6e; d chứa tối đa 10e; f chứa tối đa 14e.\n➔ Đáp án đúng: B.")

    # Câu 3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 3. Cho các nguyên tố Na, Mg, K. Thứ tự tính kim loại giảm dần của các nguyên tố đó là")
    opts3 = [
        "*A. K > Na > Mg.*" if is_teacher else "A. K > Na > Mg.",
        "B. K > Mg > Na.",
        "C. Mg > Na > K.",
        "D. Na > Mg > K."
    ]
    render_options_no_table(doc, opts3, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "• Trong cùng chu kì 3: Na (Z=11) và Mg (Z=12). Theo chiều tăng điện tích hạt nhân, tính kim loại giảm dần: Na > Mg.\n• Trong cùng nhóm IA: Na (Z=11) và K (Z=19). Theo chiều tăng điện tích hạt nhân, tính kim loại tăng dần: K > Na.\n➔ Thứ tự tính kim loại giảm dần: K > Na > Mg.\n➔ Đáp án đúng: A.")

    # Câu 4
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 4. Chọn phương án sai trong các phương án sau:")
    opts4 = [
        "A. Trong một ô orbital có tối đa 2 electron có chiều tự quay ngược nhau.",
        "B. Các electron được sắp xếp vào các ô orbital sao cho số electron độc thân là cực đại.",
        "C. Trong nguyên tử các electron được sắp xếp theo mức năng lượng từ thấp đến cao.",
        "*D. Trong một ô orbital có tối đa 2 electron có chiều tự quay cùng chiều.*" if is_teacher else "D. Trong một ô orbital có tối đa 2 electron có chiều tự quay cùng chiều."
    ]
    render_options_no_table(doc, opts4, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Theo nguyên lí Pauli: Trong một orbital nguyên tử chỉ có tối đa 2 electron và chúng phải có chiều tự quay (spin) NGƯỢC NHAU. Do đó nhận định D là sai.\n➔ Đáp án đúng: D.")

    # Câu 5
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 5. Electron thuộc lớp nào sau đây liên kết chặt chẽ nhất với hạt nhân?")
    opts5 = [
        "A. Lớp L.",
        "B. Lớp M.",
        "*C. Lớp K.*" if is_teacher else "C. Lớp K.",
        "D. Lớp N."
    ]
    render_options_no_table(doc, opts5, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Lớp K (n = 1) là lớp trong cùng, ở gần hạt nhân nhất, có mức năng lượng thấp nhất nên liên kết chặt chẽ nhất với hạt nhân.\n➔ Đáp án đúng: C.")

    # Câu 6
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 6. Orbital s chứa tối đa bao nhiêu electron?")
    opts6 = [
        "*A. 2 electron.*" if is_teacher else "A. 2 electron.",
        "B. 3 electron.",
        "C. 1 electron.",
        "D. 4 electron."
    ]
    render_options_no_table(doc, opts6, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Mỗi một orbital nguyên tử (AO) chỉ chứa tối đa 2 electron có spin ngược nhau.\n➔ Đáp án đúng: A.")

    # Câu 7
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 7. Cấu hình electron của nguyên tử Aluminium (Z = 13) theo ô orbital có số electron độc thân là")
    opts7 = [
        "*A. 1.*" if is_teacher else "A. 1.",
        "B. 3.",
        "C. 4.",
        "D. 2."
    ]
    render_options_no_table(doc, opts7, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Al (Z = 13): 1s^2 2s^2 2p^6 3s^2 3p^1.\nPhân lớp ngoài cùng 3p^1 có 1 electron điền vào 1 ô orbital ([↑][ ][ ]) => Có 1 electron độc thân.\n➔ Đáp án đúng: A.")

    # Câu 8
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 8. Nguyên tử của nguyên tố Nitrogen có 7 electron. Kết luận nào sau đây là đúng?")
    add_centered_image(doc, os.path.join(IMG_DIR, "cau8_nitrogen.png"), width_cm=3.0)
    opts8 = [
        "A. Lớp electron ngoài cùng của Nitrogen có 3 electron.",
        "B. Nguyên tử Nitrogen có 5 electron độc thân.",
        "C. Lớp L (lớp thứ 2) của Nitrogen có 3 electron.",
        "*D. Lớp electron ngoài cùng của Nitrogen có 5 electron.*" if is_teacher else "D. Lớp electron ngoài cùng của Nitrogen có 5 electron."
    ]
    render_options_no_table(doc, opts8, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Cấu hình electron của Nitrogen (Z = 7): 1s^2 2s^2 2p^3.\n• Lớp 1 (lớp K): 2 electron.\n• Lớp 2 (lớp L, lớp ngoài cùng): 2 + 3 = 5 electron.\n• Số electron độc thân ở phân lớp 2p là 3.\n➔ Đáp án đúng: D.")

    # Câu 9
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 9. Cho sơ đồ cấu tạo nguyên tử sau. Nguyên tử đó có số khối là")
    add_centered_image(doc, os.path.join(IMG_DIR, "cau9_so_do_nguyen_tu.png"), width_cm=5.0)
    opts9 = [
        "A. 12.",
        "B. 11.",
        "*C. 23.*" if is_teacher else "C. 23.",
        "D. 32."
    ]
    render_options_no_table(doc, opts9, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Dựa vào sơ đồ cấu tạo nguyên tử:\n• Trong hạt nhân có 11 proton và 12 neutron.\n• Số khối A = p + n = 11 + 12 = 23 (đây là nguyên tử Na).\n➔ Đáp án đúng: C.")

    # Câu 10
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 10. Kí hiệu nguyên tử Magnesium được cho tại hình sau. Số hạt proton, neutron và electron trong nguyên tử Magnesium lần lượt là")
    add_centered_image(doc, os.path.join(IMG_DIR, "cau10_magnesium.png"), width_cm=3.2)
    opts10 = [
        "A. 24, 12, 12.",
        "*B. 12, 12, 12.*" if is_teacher else "B. 12, 12, 12.",
        "C. 12, 24, 12.",
        "D. 12, 12, 24."
    ]
    render_options_no_table(doc, opts10, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Kí hiệu ^24_12Mg cho biết:\n• Số hiệu nguyên tử Z = p = e = 12.\n• Số neutron n = A - Z = 24 - 12 = 12.\nVậy số hạt proton, neutron, electron lần lượt là 12, 12, 12.\n➔ Đáp án đúng: B.")

    # Câu 11
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 11. Trong một chu kì, bán kính nguyên tử các nguyên tố:")
    opts11 = [
        "A. Tăng theo chiều tăng dần của điện tích hạt nhân.",
        "B. Tăng theo chiều tăng của tính phi kim.",
        "C. Giảm theo chiều tăng của tính kim loại.",
        "*D. Giảm theo chiều tăng dần của điện tích hạt nhân.*" if is_teacher else "D. Giảm theo chiều tăng dần của điện tích hạt nhân."
    ]
    render_options_no_table(doc, opts11, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Trong một chu kì, khi đi từ trái sang phải theo chiều tăng dần của điện tích hạt nhân: số lớp electron không đổi nhưng điện tích hạt nhân tăng làm tăng lực hút giữa hạt nhân với electron ngoài cùng, khiến bán kính nguyên tử giảm dần.\n➔ Đáp án đúng: D.")

    # Câu 12
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 12. Trong bảng tuần hoàn, các nguyên tố hoá học được sắp xếp theo ba nguyên tắc, nguyên tắc nào sau đây là đúng?")
    opts12 = [
        "A. Các nguyên tố có cùng số lớp electron trong nguyên tử được xếp thành một cột.",
        "*B. Các nguyên tố được xếp theo chiều điện tích hạt nhân tăng dần.*" if is_teacher else "B. Các nguyên tố được xếp theo chiều điện tích hạt nhân tăng dần.",
        "C. Các nguyên tố có cùng số electron hoá trị được xếp cùng một hàng.",
        "D. Các nguyên tố được xếp theo chiều nguyên tử khối tăng dần."
    ]
    render_options_no_table(doc, opts12, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Ba nguyên tắc sắp xếp các nguyên tố trong bảng tuần hoàn hiện đại:\n1. Các nguyên tố được sắp xếp theo chiều tăng dần của điện tích hạt nhân.\n2. Các nguyên tố có cùng số lớp electron được xếp vào cùng một hàng (chu kì).\n3. Các nguyên tố có cùng số electron hoá trị được xếp vào cùng một cột (nhóm).\n➔ Đáp án đúng: B.")

    # Câu 13
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 13. Hình ảnh sau đây là của orbital nào?")
    add_centered_image(doc, os.path.join(IMG_DIR, "cau13_orbital_px.png"), width_cm=4.5)
    opts13 = [
        "A. Orbital s",
        "B. Orbital p~z~",
        "*C. Orbital p~x~*" if is_teacher else "C. Orbital p~x~",
        "D. Orbital p~y~"
    ]
    render_options_no_table(doc, opts13, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Hình ảnh mô tả orbital p có dạng hình số 8 nổi nằm định hướng dọc theo trục toạ độ x, đó chính là orbital p~x~.\n➔ Đáp án đúng: C.")

    # Câu 14
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 14. Nhóm IA trong bảng tuần hoàn có tên gọi là")
    opts14 = [
        "A. nhóm halogen.",
        "B. nhóm kim loại kiềm thổ.",
        "*C. nhóm kim loại kiềm.*" if is_teacher else "C. nhóm kim loại kiềm.",
        "D. nhóm khí hiếm."
    ]
    render_options_no_table(doc, opts14, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "• Nhóm IA (trừ hydrogen): nhóm kim loại kiềm.\n• Nhóm IIA: nhóm kim loại kiềm thổ.\n• Nhóm VIIA: nhóm halogen.\n• Nhóm VIIIA: nhóm khí hiếm.\n➔ Đáp án đúng: C.")

    # Câu 15
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 15. Trong những cặp sau đây, cặp nào là đồng vị của nhau?")
    opts15 = [
        "A. O~2~ và O~3~.",
        "*B. ^16_8O và ^17_8O.*" if is_teacher else "B. ^16_8O và ^17_8O.",
        "C. ^40_19K và ^40_18Ar.",
        "D. ^40_19K và ^40_20Ca."
    ]
    render_options_no_table(doc, opts15, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Đồng vị là các nguyên tử của cùng một nguyên tố hoá học, có cùng số proton (cùng Z) nhưng khác nhau về số neutron (khác số khối A). Cặp ^16_8O và ^17_8O đều có Z = 8.\n➔ Đáp án đúng: B.")

    # Câu 16
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 16. Nguyên tử nguyên tố X có cấu hình electron lớp ngoài cùng là 3p^4. Vị trí của X trong bảng tuần hoàn là")
    opts16 = [
        "A. nhóm IVA, chu kì 3.",
        "B. nhóm VI B, chu kì 4.",
        "C. nhóm IV B, chu kì 4.",
        "*D. nhóm VI A, chu kì 3.*" if is_teacher else "D. nhóm VI A, chu kì 3."
    ]
    render_options_no_table(doc, opts16, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Lớp ngoài cùng có phân lớp 3p^4 thì phân lớp 3s trước đó đã bão hòa (3s^2). Cấu hình e lớp ngoài cùng là 3s^2 3p^4:\n• Có 3 lớp electron => Chu kì 3.\n• Có 6 electron hóa trị ở lớp ngoài cùng, là nguyên tố p => Nhóm VIA.\n➔ Đáp án đúng: D.")

    # Câu 17
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 17. Xét các nguyên tố trong cùng một chu kì, đại lượng nào sau đây không biến đổi tuần hoàn?")
    opts17 = [
        "*A. Số lớp electron.*" if is_teacher else "A. Số lớp electron.",
        "B. Tính kim loại, tính phi kim.",
        "C. Độ âm điện.",
        "D. Bán kính nguyên tử."
    ]
    render_options_no_table(doc, opts17, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Trong cùng một chu kì, số lớp electron của các nguyên tử luôn không đổi (bằng số thứ tự chu kì), do đó không biến đổi tuần hoàn.\n➔ Đáp án đúng: A.")

    # Câu 18
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 18. Nguyên tố X có Z = 11 thuộc loại nguyên tố nào sau đây?")
    opts18 = [
        "A. Nguyên tố p.",
        "B. Nguyên tố f.",
        "*C. Nguyên tố s.*" if is_teacher else "C. Nguyên tố s.",
        "D. Nguyên tố d."
    ]
    render_options_no_table(doc, opts18, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Cấu hình electron của X (Z = 11): 1s^2 2s^2 2p^6 3s^1. Electron cuối cùng được điền vào phân lớp s, do đó X là nguyên tố s.\n➔ Đáp án đúng: C.")

    # =========================================================================
    # PHẦN II
    # =========================================================================
    add_section_header(
        doc,
        "PHẦN II. Câu trắc nghiệm đúng sai.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai."
    )

    # Câu 1 P2
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 1. Trong lịch sử các thuyết về mô hình nguyên tử có mô hình hành tinh nguyên tử (mô hình Rutherford - Bohr) và mô hình hiện đại của nguyên tử.")
    add_centered_image(doc, os.path.join(IMG_DIR, "p2_cau1_mo_hinh.png"), width_cm=11.0)
    stmts1 = [
        ("a) Khái niệm về xác suất tìm thấy electron xuất phát từ mô hình hành tinh nguyên tử.", "Sai"),
        ("b) Với nguyên tử hydrogen, mô hình (1) là mô hình hiện đại, mô hình (2) là mô hình hành tinh nguyên tử.", "Sai"),
        ("c) Trong mô hình (2) các electron chuyển động rất nhanh xung quanh hạt nhân, không theo một quỹ đạo xác định.", "Đúng"),
        ("d) Theo mô hình (1), electron chuyển động trên quỹ đạo xác định xung quanh hạt nhân.", "Đúng")
    ]
    for stmt, ans in stmts1:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.left_indent = Cm(0.5)
        p_st.paragraph_format.space_before = Pt(1)
        p_st.paragraph_format.space_after = Pt(1)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        sol_text = (
            "• Ý a SAI: Khái niệm về xác suất tìm thấy electron (đám mây electron) xuất phát từ mô hình hiện đại (cơ học lượng tử).\n"
            "• Ý b SAI: Mô hình (1) là mô hình hành tinh nguyên tử (Rutherford - Bohr); mô hình (2) là mô hình hiện đại.\n"
            "• Ý c ĐÚNG: Theo mô hình hiện đại, electron chuyển động rất nhanh quanh hạt nhân trong khu vực không gian không theo quỹ đạo xác định.\n"
            "• Ý d ĐÚNG: Theo mô hình hành tinh nguyên tử (Rutherford - Bohr), electron quay quanh hạt nhân theo những quỹ đạo xác định."
        )
        add_solution_no_table(doc, sol_text)

    # Câu 2 P2
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 2. Chlorine là chất có nhiều ứng dụng như khử trùng nước sinh hoạt, sản xuất các chất tẩy rửa, nguyên tố này có Z = 17 và có số khối bằng 35.")
    stmts2 = [
        ("a) Trong hạt nhân nguyên tử chlorine có 17 hạt không mang điện.", "Sai"),
        ("b) Nguyên tử chlorine có 17 electron.", "Đúng"),
        ("c) Chlorine có 9 electron trên phân lớp p và là nguyên tố p.", "Sai"),
        ("d) Nguyên tử chlorine có 2 đồng vị bền, Hydrogen có 3 đồng vị bền. Từ các đồng vị này Có thể tạo ra 8 loại phân tử HCl.", "Sai")
    ]
    for stmt, ans in stmts2:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.left_indent = Cm(0.5)
        p_st.paragraph_format.space_before = Pt(1)
        p_st.paragraph_format.space_after = Pt(1)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        sol_text = (
            "• Ý a SAI: Hạt không mang điện trong hạt nhân là neutron: n = A - Z = 35 - 17 = 18 hạt (17 là số proton mang điện dương).\n"
            "• Ý b ĐÚNG: Nguyên tử trung hòa về điện nên e = p = Z = 17.\n"
            "• Ý c SAI: Cấu hình electron của Cl là 1s^2 2s^2 2p^6 3s^2 3p^5. Tổng số electron trên các phân lớp p là 6 + 5 = 11 hạt (chứ không phải 9).\n"
            "• Ý d SAI: Số phân tử HCl tối đa = (3 đồng vị H) × (2 đồng vị Cl) = 6 loại phân tử (chứ không phải 8)."
        )
        add_solution_no_table(doc, sol_text)

    # Câu 3 P2
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 3. Trong nguyên tử nguyên tố X có tổng số electron trên phân lớp s là 6, lớp ngoài cùng của nguyên tử X có 2 electron.")
    stmts3 = [
        ("a) X có 1 electron độc thân và là nguyên tố p.", "Sai"),
        ("b) Nguyên tố X là kim loại, dễ nhường electron.", "Đúng"),
        ("c) Tổng số electron trên lớp L của X là 8.", "Đúng"),
        ("d) X có 2 lớp electron.", "Sai")
    ]
    for stmt, ans in stmts3:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.left_indent = Cm(0.5)
        p_st.paragraph_format.space_before = Pt(1)
        p_st.paragraph_format.space_after = Pt(1)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        sol_text = (
            "Phân tích: Tổng số e trên phân lớp s là 6 => Gồm 1s^2, 2s^2, 3s^2.\n"
            "Lớp ngoài cùng có 2 electron => Lớp ngoài cùng là 3s^2.\n"
            "Cấu hình e đầy đủ của X là 1s^2 2s^2 2p^6 3s^2 (Z = 12: Magnesium - Mg).\n"
            "• Ý a SAI: X có cấu hình [Ne]3s^2 nên có 0 electron độc thân và là nguyên tố s.\n"
            "• Ý b ĐÚNG: X có 2 electron lớp ngoài cùng nên là kim loại kiềm thổ, dễ nhường 2e để đạt cấu hình bền.\n"
            "• Ý c ĐÚNG: Lớp L (n = 2) có cấu hình 2s^2 2p^6 gồm 2 + 6 = 8 electron.\n"
            "• Ý d SAI: X có 3 lớp electron (lớp 1, 2, 3) chứ không phải 2 lớp."
        )
        add_solution_no_table(doc, sol_text)

    # Câu 4 P2
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 4. Tìm hiểu các nguyên tố hóa học Natri (sodium, _11Na) và Potasium (_19K) trong bảng tuần hoàn.")
    stmts4 = [
        ("a) Theo xu hướng biến đổi tính kim loại, K có tính kim loại mạnh hơn Na.", "Đúng"),
        ("b) Đều thuộc chu kì 3 trong bảng tuần hoàn", "Sai"),
        ("c) Tính base của sodium hydroxide yếu hơn tính base của potasium hydroxide.", "Đúng"),
        ("d) Na và K đều có tính chất hóa học cơ bản giống nhau.", "Đúng")
    ]
    for stmt, ans in stmts4:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.left_indent = Cm(0.5)
        p_st.paragraph_format.space_before = Pt(1)
        p_st.paragraph_format.space_after = Pt(1)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        sol_text = (
            "• Ý a ĐÚNG: Na và K cùng thuộc nhóm IA, đi từ trên xuống dưới theo chiều Z tăng thì tính kim loại tăng dần: K > Na.\n"
            "• Ý b SAI: Na (Z=11) có 3 lớp e thuộc chu kì 3; K (Z=19) có 4 lớp e thuộc chu kì 4.\n"
            "• Ý c ĐÚNG: Theo quy luật nhóm IA, tính base của hydroxide tăng dần từ trên xuống: NaOH yếu hơn KOH.\n"
            "• Ý d ĐÚNG: Cả Na và K đều có 1 electron hóa trị ở lớp ngoài cùng (ns^1) nên có tính chất hóa học tương tự nhau."
        )
        add_solution_no_table(doc, sol_text)

    # =========================================================================
    # PHẦN III
    # =========================================================================
    add_section_header(
        doc,
        "PHẦN III. Trắc nghiệm trả lời ngắn và tự luận.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 6."
    )

    # Câu 1 P3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 1. Nguyên tử của nguyên tố Nitrogen ^14_7N có bao nhiêu neutron?")
    if is_teacher:
        add_solution_no_table(doc, "Số neutron của nguyên tử ^14_7N là: N = A - Z = 14 - 7 = 7.\n➔ Đáp số: 7.")
    else:
        add_student_writing_lines(doc, num_lines=2)

    # Câu 2 P3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 2. Lớp L có số orbital bằng bao nhiêu?")
    if is_teacher:
        add_solution_no_table(doc, "Lớp L ứng với số lượng tử chính n = 2.\nLớp L có 2 phân lớp: phân lớp 2s (1 orbital) và phân lớp 2p (3 orbital).\nTổng số orbital trong lớp L là: 1 + 3 = 4 (hoặc n^2 = 2^2 = 4 orbital).\n➔ Đáp số: 4.")
    else:
        add_student_writing_lines(doc, num_lines=2)

    # Câu 3 P3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 3. Nguyên tử của nguyên tố Y có tổng số hạt trong nguyên tử là 34. Số hạt mang điện nhiều hơn số hạt không mang điện là 10. Số khối của nguyên tử Y là bao nhiêu?")
    if is_teacher:
        sol_text = (
            "Gọi p, e, n lần lượt là số hạt proton, electron và neutron của nguyên tử Y (e = p).\n"
            "Theo bài ra, ta có hệ phương trình:\n"
            "  { 2p + n = 34  (1)\n"
            "  { 2p - n = 10  (2)\n"
            "Cộng (1) và (2): 4p = 44 => p = 11 => e = 11.\n"
            "Thay p = 11 vào (1): n = 34 - 2 × 11 = 12.\n"
            "Số khối của nguyên tử Y: A = p + n = 11 + 12 = 23 (nguyên tố Na).\n"
            "➔ Đáp số: 23."
        )
        add_solution_no_table(doc, sol_text)
    else:
        add_student_writing_lines(doc, num_lines=4)

    # Câu 4 P3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 4. Nguyên tử của các nguyên tố X (Z = 6); Y (Z = 20).")
    p_a = doc.add_paragraph()
    render_question_prompt(p_a, "a. Viết cấu hình electron nguyên tử của X và Y.")
    p_b = doc.add_paragraph()
    render_question_prompt(p_b, "b. Cho biết X, Y có tính chất của kim loại, phi kim hay khí hiếm? Giải thích?")
    if is_teacher:
        sol_text = (
            "a) Cấu hình electron nguyên tử:\n"
            "• Nguyên tố X (Z = 6): 1s^2 2s^2 2p^2 (hoặc [He]2s^2 2p^2).\n"
            "• Nguyên tố Y (Z = 20): 1s^2 2s^2 2p^6 3s^2 3p^6 4s^2 (hoặc [Ar]4s^2).\n\n"
            "b) Dự đoán tính chất của X và Y:\n"
            "• X có 4 electron ở lớp ngoài cùng (2s^2 2p^2) và thuộc chu kì nhỏ (chu kì 2) nên X là phi kim (nguyên tố Carbon - C).\n"
            "• Y có 2 electron ở lớp ngoài cùng (4s^2) nên Y là kim loại (nguyên tố Calcium - Ca)."
        )
        add_solution_no_table(doc, sol_text)
    else:
        add_student_writing_lines(doc, num_lines=6)

    # Câu 5 P3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 5. Trong phân tử MX~2~ có tổng số hạt cơ bản bằng 164 hạt, trong đó số hạt mang điện nhiều hơn số hạt không mang điện là 52 hạt. Số khối của nguyên tử M lớn hơn số khối của nguyên tử X là 5. Tổng số hạt trong nguyên tử M lớn hơn trong nguyên tử X là 8 hạt. Xác định số hiệu nguyên tử của M và X và cho biết công thức MX~2~.")
    if is_teacher:
        sol_text = (
            "Gọi p_M, n_M lần lượt là số proton, neutron của nguyên tử M (e_M = p_M).\n"
            "Gọi p_X, n_X lần lượt là số proton, neutron của nguyên tử X (e_X = p_X).\n"
            "Theo bài ra:\n"
            "1) Tổng số hạt cơ bản trong phân tử MX~2~ là 164:\n"
            "   (2p_M + n_M) + 2(2p_X + n_X) = 164\n"
            "   <=> (2p_M + 4p_X) + (n_M + 2n_X) = 164   (1)\n"
            "2) Số hạt mang điện nhiều hơn số hạt không mang điện là 52:\n"
            "   (2p_M + 4p_X) - (n_M + 2n_X) = 52    (2)\n"
            "Từ (1) và (2) suy ra:\n"
            "   2p_M + 4p_X = (164 + 52) / 2 = 108 <=> p_M + 2p_X = 54  (3)\n"
            "   n_M + 2n_X = (164 - 52) / 2 = 56   (4)\n\n"
            "3) Số khối của M lớn hơn số khối của X là 5:\n"
            "   A_M - A_X = (p_M + n_M) - (p_X + n_X) = 5  (5)\n"
            "4) Tổng số hạt của M lớn hơn của X là 8:\n"
            "   (2p_M + n_M) - (2p_X + n_X) = 8            (6)\n"
            "Lấy (6) trừ (5) ta được:\n"
            "   p_M - p_X = 8 - 5 = 3                      (7)\n\n"
            "Kết hợp (3) và (7) ta có hệ:\n"
            "   { p_M + 2p_X = 54\n"
            "   { p_M - p_X = 3\n"
            "Trừ hai phương trình: 3p_X = 51 => p_X = 17 (Chlorine - Cl).\n"
            "Thay vào (7): p_M = 3 + 17 = 20 (Calcium - Ca).\n\n"
            "➔ Kết luận:\n"
            "• Số hiệu nguyên tử của M: Z_M = 20 (Calcium - Ca).\n"
            "• Số hiệu nguyên tử của X: Z_X = 17 (Chlorine - Cl).\n"
            "• Công thức phân tử MX~2~ là: CaCl~2~ (Calcium chloride)."
        )
        add_solution_no_table(doc, sol_text)
    else:
        add_student_writing_lines(doc, num_lines=8)

    # Câu 6 P3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 6. Cho 5,85 gam muối NaX tác dụng với dung dịch AgNO~3~ dư, thu được 14,35 gam kết tủa trắng. Nguyên tố X có hai đồng vị ^35X (x~1~%) và ^37X (x~2~%). Xác định x~1~ và x~2~?")
    if is_teacher:
        sol_text = (
            "Phương trình phản ứng hoá học:\n"
            "  NaX + AgNO~3~ ➔ AgX↓ + NaNO~3~\n"
            "Theo phương trình: n_NaX = n_AgX\n"
            "<=> 5,85 / (23 + Ā_X) = 14,35 / (108 + Ā_X)\n"
            "<=> 5,85 × (108 + Ā_X) = 14,35 × (23 + Ā_X)\n"
            "<=> 631,8 + 5,85·Ā_X = 330,05 + 14,35·Ā_X\n"
            "<=> 8,5·Ā_X = 301,75 => Ā_X = 35,5 (X là nguyên tố Chlorine - Cl).\n\n"
            "Nguyên tố X có 2 đồng vị ^35X (x~1~%) và ^37X (x~2~%):\n"
            "Ta có: x~1~ + x~2~ = 100 <=> x~2~ = 100 - x~1~.\n"
            "Nguyên tử khối trung bình:\n"
            "  Ā_X = [35·x~1~ + 37(100 - x~1~)] / 100 = 35,5\n"
            "  <=> 35·x~1~ + 3700 - 37·x~1~ = 3550\n"
            "  <=> 2·x~1~ = 150 => x~1~ = 75%.\n"
            "  => x~2~ = 100% - 75% = 25%.\n\n"
            "➔ Kết luận: x~1~ = 75%; x~2~ = 25%."
        )
        add_solution_no_table(doc, sol_text)
    else:
        add_student_writing_lines(doc, num_lines=8)

    # Lưu file
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Đã tạo thành công: {output_path}")

if __name__ == "__main__":
    out_hs = r"C:\Antigravity_Thanh\San_Pham\De_Kiem_Tra_GK1_Hoa_10_LeThanhTong_2025_HocSinh.docx"
    out_gv = r"C:\Antigravity_Thanh\San_Pham\De_Kiem_Tra_GK1_Hoa_10_LeThanhTong_2025_GiaoVien.docx"
    build_exam_lethanhtong_10(is_teacher=False, output_path=out_hs)
    build_exam_lethanhtong_10(is_teacher=True, output_path=out_gv)
