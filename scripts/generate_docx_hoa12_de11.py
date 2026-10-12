# -*- coding: utf-8 -*-
"""
Script tạo file Word (.docx) chuẩn "FORM ĐỀ KIỂM TRA" cho:
ĐỀ KIỂM TRA SỐ 11 - CHƯƠNG 4: POLYMER (HÓA HỌC 12)
Gồm 2 phiên bản:
1. De_Kiem_Tra_Hoa_12_De11.docx (Bản Học Sinh)
2. De_Kiem_Tra_Hoa_12_De11_GiaoVien.docx (Bản Giáo Viên có lời giải chi tiết)
Tuân thủ 100% quy chuẩn: Header & Footer, Tab Stops (NO TABLE trong đề/lời giải), Sub/Superscript công thức hóa học.
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
IMG_DIR = r"c:\Antigravity_Thanh\San_Pham\images_crop_de11"

COLOR_HF_RED = RGBColor(211, 47, 47)      # Coral Red (#D32F2F)
COLOR_HF_GREEN_HEX = "8BC390"             # Sage Green (#8BC390)
COLOR_PRIMARY = RGBColor(46, 125, 50)     # Dark Green (#2E7D32)
COLOR_CORRECT = RGBColor(198, 40, 40)     # Red for correct answers (#C62828)
COLOR_SOLUTION = RGBColor(21, 101, 192)   # Blue for solutions (#1565C0)
COLOR_GRAY = RGBColor(97, 97, 97)

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def set_run_style(run, font_name="Times New Roman", font_size=12, bold=False, italic=False, subscript=False, superscript=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def add_chem_run(p, text, base_font="Times New Roman", base_size=12, bold=False, italic=False, color=None):
    tokens = re.split(r'(\^\{[^}]+\}|\^[0-9\+\-]+|~\{[^}]+\}|~[0-9]+)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('^'):
            content = tok[2:-1] if tok.startswith('^{') else tok[1:]
            r = p.add_run(content)
            set_run_style(r, font_name=base_font, font_size=base_size, bold=bold, italic=italic, superscript=True, color=color)
        elif tok.startswith('~'):
            content = tok[2:-1] if tok.startswith('~{') else tok[1:]
            r = p.add_run(content)
            set_run_style(r, font_name=base_font, font_size=base_size, bold=bold, italic=italic, subscript=True, color=color)
        else:
            r = p.add_run(tok)
            set_run_style(r, font_name=base_font, font_size=base_size, bold=bold, italic=italic, color=color)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(r'<w:tcBorders %s/>' % nsdecls('w'))
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = parse_xml(r'<%s %s w:val="%s" w:sz="%s" w:space="0" w:color="%s"/>' %
                                (tag, nsdecls('w'), edge_data.get('val', 'single'),
                                 edge_data.get('sz', '4'), edge_data.get('color', 'auto')))
            tcBorders.append(element)
        else:
            tag = 'w:{}'.format(edge)
            element = parse_xml(r'<%s %s w:val="none"/>' % (tag, nsdecls('w')))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def set_cell_margins(cell, top=50, bottom=50, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(r'<w:tcMar %s><w:top w:w="%d" w:type="dxa"/><w:bottom w:w="%d" w:type="dxa"/><w:left w:w="%d" w:type="dxa"/><w:right w:w="%d" w:type="dxa"/></w:tcMar>' % (nsdecls('w'), top, bottom, left, right))
    tcPr.append(tcMar)

def setup_page_and_header_footer(doc):
    for sec in doc.sections:
        sec.top_margin = Cm(1.5)
        sec.bottom_margin = Cm(1.5)
        sec.left_margin = Cm(1.5)
        sec.right_margin = Cm(1.5)
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.header_distance = Cm(0.6)
        sec.footer_distance = Cm(0.6)
        
        # --- HEADER ---
        hdr = sec.header
        hdr_p = hdr.paragraphs[0]
        hdr_p.text = ""
        hdr_p.paragraph_format.space_before = Pt(0)
        hdr_p.paragraph_format.space_after = Pt(0)
        
        tbl_hdr = hdr.add_table(rows=1, cols=2, width=Cm(18.0))
        tbl_hdr.autofit = False
        tbl_hdr.columns[0].width = Cm(9.0)
        tbl_hdr.columns[1].width = Cm(9.0)
        
        c0 = tbl_hdr.cell(0, 0)
        c1 = tbl_hdr.cell(0, 1)
        c0.width = Cm(9.0)
        c1.width = Cm(9.0)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run("Tài liệu lưu hành nội bộ")
        set_run_style(r0, font_name="Times New Roman", font_size=10, italic=True, color=COLOR_HF_RED)
        
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(2)
        r1a = p1.add_run("Thầy ")
        set_run_style(r1a, font_name="Times New Roman", font_size=10, color=COLOR_HF_RED)
        r1b = p1.add_run("TRẦN VĂN THẠNH")
        set_run_style(r1b, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED)
        r1c = p1.add_run(" - 0777.470.803")
        set_run_style(r1c, font_name="Times New Roman", font_size=10, color=COLOR_HF_RED)
        
        set_cell_border(c0, bottom={'val': 'single', 'sz': '6', 'color': COLOR_HF_GREEN_HEX})
        set_cell_border(c1, bottom={'val': 'single', 'sz': '6', 'color': COLOR_HF_GREEN_HEX})
        set_cell_margins(c0, top=0, bottom=40, left=0, right=0)
        set_cell_margins(c1, top=0, bottom=40, left=0, right=0)
        
        # --- FOOTER ---
        ftr = sec.footer
        ftr_p = ftr.paragraphs[0]
        ftr_p.text = ""
        ftr_p.paragraph_format.space_before = Pt(0)
        ftr_p.paragraph_format.space_after = Pt(0)
        
        tbl_ftr = ftr.add_table(rows=1, cols=2, width=Cm(18.0))
        tbl_ftr.autofit = False
        tbl_ftr.columns[0].width = Cm(13.5)
        tbl_ftr.columns[1].width = Cm(4.5)
        
        fc0 = tbl_ftr.cell(0, 0)
        fc1 = tbl_ftr.cell(0, 1)
        fc0.width = Cm(13.5)
        fc1.width = Cm(4.5)
        
        fp0 = fc0.paragraphs[0]
        fp0.paragraph_format.space_before = Pt(2)
        fp0.paragraph_format.space_after = Pt(0)
        fr0 = fp0.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
        set_run_style(fr0, font_name="Times New Roman", font_size=9.5, italic=True, color=COLOR_HF_RED)
        
        fp1 = fc1.paragraphs[0]
        fp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp1.paragraph_format.space_before = Pt(2)
        fp1.paragraph_format.space_after = Pt(0)
        add_page_number_field(fp1, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED)
        
        set_cell_border(fc0, top={'val': 'single', 'sz': '6', 'color': COLOR_HF_GREEN_HEX})
        set_cell_border(fc1, top={'val': 'single', 'sz': '6', 'color': COLOR_HF_GREEN_HEX})
        set_cell_margins(fc0, top=40, bottom=0, left=0, right=0)
        set_cell_margins(fc1, top=40, bottom=0, left=0, right=0)

def add_exam_title_block(doc, is_teacher=False):
    lines_info = [
        ("SỞ GD&ĐT THỪA THIÊN HUẾ", "ĐỀ KIỂM TRA CHƯƠNG 4: POLYMER", True, True),
        ("TRƯỜNG THPT HAI BÀ TRƯNG", "MÔN: HÓA HỌC 12 - NĂM HỌC 2026 - 2027", False, True),
        ("ĐỀ CHÍNH THỨC", "Thời gian làm bài: 50 phút (không kể thời gian phát đề)", True, False, False, True),
        ("(Đề thi gồm 04 trang)", "MÃ ĐỀ THI: 111", False, True, True, False)
    ]
    
    for item in lines_info:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
        
        r_l = p.add_run(item[0])
        set_run_style(r_l, font_name="Times New Roman", font_size=11, bold=item[2], italic=item[4] if len(item)>4 else False)
        
        p.add_run("\t")
        
        r_r = p.add_run(item[1])
        set_run_style(r_r, font_name="Times New Roman", font_size=11, bold=item[3], italic=item[5] if len(item)>5 else False)

    if is_teacher:
        p_note = doc.add_paragraph()
        p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_note.paragraph_format.space_before = Pt(2)
        p_note.paragraph_format.space_after = Pt(2)
        r_note = p_note.add_run("(BẢN GIÁO VIÊN -- CÓ LỜI GIẢI CHI TIẾT)")
        set_run_style(r_note, font_name="Times New Roman", font_size=11.5, bold=True, color=COLOR_HF_RED)

    # Student info line
    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(3)
    p_info.paragraph_format.space_after = Pt(1)
    r_info = p_info.add_run("Họ, tên thí sinh: .......................................................................... Lớp: ............. SBD: .............................")
    set_run_style(r_info, font_name="Times New Roman", font_size=11, italic=True)

    # Atomic mass note
    p_mass = doc.add_paragraph()
    p_mass.paragraph_format.space_before = Pt(0)
    p_mass.paragraph_format.space_after = Pt(4)
    r_mass = p_mass.add_run("(Cho biết nguyên tử khối: H = 1; C = 12; N = 14; O = 16; Na = 23; Al = 27; S = 32; Cl = 35,5; K = 39; Ca = 40; Fe = 56; Cu = 64; Zn = 65; Br = 80; Ag = 108; Ba = 137).")
    set_run_style(r_mass, font_name="Times New Roman", font_size=9.5, italic=True, color=COLOR_GRAY)

def add_part_heading(doc, part_title, part_instruction):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    r_title = p.add_run(part_title)
    set_run_style(r_title, font_name="Times New Roman", font_size=11.5, bold=True, color=COLOR_PRIMARY)
    
    if part_instruction:
        r_inst = p.add_run(" " + part_instruction)
        set_run_style(r_inst, font_name="Times New Roman", font_size=11, italic=True)

def add_solution_paragraph(doc, solution_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.6)
    
    r_label = p.add_run("Lời giải chi tiết: ")
    set_run_style(r_label, font_name="Times New Roman", font_size=10.5, bold=True, color=COLOR_SOLUTION)
    
    add_chem_run(p, solution_text, base_font="Times New Roman", base_size=10.5, italic=True)

# ----------------- DỮ LIỆU ĐỀ KIỂM TRA SỐ 11 -----------------

MCQ_QUESTIONS = [
    {
        "num": 1,
        "text": "(SBT – KNTT). Chất nào dưới đây thuộc loại polymer?",
        "options": ["A. Glucose.", "B. Fructose.", "C. Saccharose.", "D. Cellulose."],
        "correct": "D",
        "sol": "Chọn D. Cellulose là polysaccharide có công thức (C~6~H~10~O~5~)~n~ thuộc loại polymer thiên nhiên. Glucose, fructose là monosaccharide; saccharose là disaccharide."
    },
    {
        "num": 2,
        "text": "(SBT – KNTT). Tính chất vật lí chung của polymer là",
        "options": [
            "A. chất lỏng, không màu, không tan trong nước.",
            "B. chất lỏng, không màu, tan tốt trong nước.",
            "C. chất rắn, không bay hơi, dễ tan trong nước.",
            "D. chất rắn, không bay hơi, không tan trong nước."
        ],
        "correct": "D",
        "sol": "Chọn D. Hầu hết các polymer ở điều kiện thường là chất rắn, không bay hơi, không tan trong nước và không có nhiệt độ nóng chảy xác định."
    },
    {
        "num": 3,
        "text": "Polymer nào sau đây bị thủy phân trong dung dịch NaOH?",
        "options": ["A. Poly(vinyl acetate).", "B. Polyethylene.", "C. Polypropylene.", "D. Polybuta-1,3-diene."],
        "correct": "A",
        "sol": "Chọn A. Poly(vinyl acetate) (PVAc) chứa nhóm chức ester trong mạch nhánh: (-CH~2~-CH(OCOCH~3~)-)~n~ nên bị thủy phân trong môi trường kiềm tạo poly(vinyl alcohol) và muối acetate. Các polymer còn lại chỉ có liên kết C-C bền vững."
    },
    {
        "num": 4,
        "text": "Polymer nào sau đây có thể bị thuỷ phân hoàn toàn trong môi trường acid hoặc môi trường base thu được amino acid?",
        "options": ["A. Tinh bột.", "B. Tơ capron.", "C. Polypropylene.", "D. Poly(vinyl chloride)."],
        "correct": "B",
        "sol": "Chọn B. Tơ capron (nylon-6) là một polyamide (-NH-[CH~2~]~5~-CO-)~n~, khi thủy phân hoàn toàn trong môi trường acid hoặc base sẽ tạo ra amino acid tương ứng (hoặc muối của nó) là 6-aminohexanoic acid (ε-aminocaproic acid)."
    },
    {
        "num": 5,
        "text": "(SBT – KNTT). Trên hộp xốp cách nhiệt, hộp đựng thức ăn mang về, cốc, chén đĩa dùng một lần,... thường được in kí hiệu như hình bên. Polymer dùng làm các đồ dùng đó được tổng hợp từ monomer nào sau đây?",
        "image": "cau5_ps_symbol.png",
        "options": [
            "A. CH~2~=CH~2~.",
            "B. CH~2~=CH-CH~3~.",
            "C. CH~2~=CH-C~6~H~5~.",
            "D. CH~2~=CHCl."
        ],
        "correct": "C",
        "sol": "Chọn C. Kí hiệu số 6 với chữ PS đại diện cho polystyrene. Polystyrene được điều chế bằng phản ứng trùng hợp styrene (vinylbenzene) CH~2~=CH-C~6~H~5~."
    },
    {
        "num": 6,
        "text": "(SBT – KNTT). Phân tử polymer nào sau đây chỉ chứa hai loại nguyên tố?",
        "options": ["A. Poly(methyl methacrylate).", "B. Poly(vinyl chloride).", "C. Polyacrylonitrile.", "D. Polypropylene."],
        "correct": "D",
        "sol": "Chọn D. Polypropylene (PP) có công thức (-CH~2~-CH(CH~3~)-)~n~ chỉ chứa 2 nguyên tố C và H. Poly(methyl methacrylate) chứa C, H, O; PVC chứa C, H, Cl; polyacrylonitrile chứa C, H, N."
    },
    {
        "num": 7,
        "text": "(SBT – KNTT). Tơ tằm, sợi bông, len thuộc loại tơ nào sau đây?",
        "options": ["A. Tơ tự nhiên.", "B. Tơ tổng hợp.", "C. Tơ bán tổng hợp.", "D. Tơ nhân tạo."],
        "correct": "A",
        "sol": "Chọn A. Tơ tằm, sợi bông, len là các loại tơ có sẵn trong thiên nhiên nên thuộc loại tơ tự nhiên (tơ thiên nhiên)."
    },
    {
        "num": 8,
        "text": "(SBT – KNTT). Tơ nào sau đây thuộc loại tơ bán tổng hợp?",
        "options": ["A. Tơ nylon-6,6.", "B. Tơ cellulose acetate.", "C. Tơ capron.", "D. Tơ tằm."],
        "correct": "B",
        "sol": "Chọn B. Tơ cellulose acetate và tơ visco là tơ bán tổng hợp (tơ nhân tạo) được chế biến bằng con đường hóa học từ polymer thiên nhiên (cellulose). Tơ nylon-6,6 và tơ capron là tơ tổng hợp; tơ tằm là tơ tự nhiên."
    },
    {
        "num": 9,
        "text": "(SBT – KNTT). Polymer nào sau đây không thuộc loại cao su?",
        "options": ["A. Poly(methyl methacrylate).", "B. Polychloroprene.", "C. Polyisoprene.", "D. Polybuta-1,3-diene."],
        "correct": "A",
        "sol": "Chọn A. Poly(methyl methacrylate) (thủy tinh hữu cơ plexiglass) thuộc loại chất dẻo. Polychloroprene, polyisoprene, polybuta-1,3-diene đều là các loại cao su."
    },
    {
        "num": 10,
        "text": "Cho các polymer: polyethylene, tơ nitron, tơ capron, nylon-6,6, tinh bột, protein, cao su isoprene và cao su buna-N. Số polymer có chứa nitrogen trong phân tử là",
        "options": ["A. 5.", "B. 6.", "C. 7.", "D. 4."],
        "correct": "A",
        "sol": "Chọn A. Các polymer chứa nitrogen gồm 5 chất: tơ nitron (-CH~2~-CH(CN)-)~n~, tơ capron (-NH-[CH~2~]~5~-CO-)~n~, nylon-6,6 (-NH-[CH~2~]~6~-NH-CO-[CH~2~]~4~-CO-)~n~, protein (chuỗi polypeptide chứa N), cao su buna-N (chứa mắt xích acrylonitrile có N)."
    },
    {
        "num": 11,
        "text": "Cho các polymer sau: cellulose, amylose, amylopectin, poly(ethylene terephthalate), polybuta-1,3-diene, polyisoprene. Số polymer thiên nhiên là",
        "options": ["A. 5.", "B. 2.", "C. 3.", "D. 4."],
        "correct": "D",
        "sol": "Chọn D. Các polymer thiên nhiên gồm 4 chất: cellulose, amylose, amylopectin, polyisoprene (thành phần chính của cao su thiên nhiên). PET và polybuta-1,3-diene là polymer tổng hợp."
    },
    {
        "num": 12,
        "text": "Có một số hợp chất sau: (1) ethylene, (2) vinyl chloride, (3) adipic acid, (4) phenol, (5) acrylonitrile, (6) buta-1,3-diene. Những chất nào có thể tham gia phản ứng trùng hợp?",
        "options": [
            "A. (1), (2), (5), (6).",
            "B. (1), (2), (3), (4).",
            "C. (1), (4), (5), (6).",
            "D. (2), (3), (4), (5)."
        ],
        "correct": "A",
        "sol": "Chọn A. Các chất tham gia phản ứng trùng hợp chứa liên kết đôi C=C: (1) ethylene, (2) vinyl chloride, (5) acrylonitrile, (6) buta-1,3-diene. Adipic acid và phenol chỉ tham gia phản ứng trùng ngưng."
    },
    {
        "num": 13,
        "text": "Dãy nào sau đây đều thuộc loại polymer tổng hợp?",
        "options": [
            "A. Tơ capron, tơ nitron, cao su buna.",
            "B. Polystyrene, tơ tằm, tơ nylon-6,6.",
            "C. Tơ cellulose acetate, cao su buna-S, tơ nylon-6.",
            "D. Tơ viscose, tơ olon, tơ nylon-7."
        ],
        "correct": "A",
        "sol": "Chọn A. Tơ capron, tơ nitron, cao su buna đều là polymer tổng hợp. B có tơ tằm (thiên nhiên); C có tơ cellulose acetate (bán tổng hợp); D có tơ viscose (bán tổng hợp)."
    },
    {
        "num": 14,
        "text": "Cho các polymer sau: polybuta-1,3-diene, poly(methyl methacrylate), polyacrylonitrile, nylon-6,6. Số polymer được điều chế bằng phản ứng trùng ngưng là",
        "options": ["A. 3.", "B. 2.", "C. 4.", "D. 1."],
        "correct": "D",
        "sol": "Chọn D. Chỉ có 1 polymer điều chế bằng phản ứng trùng ngưng là nylon-6,6 (trùng ngưng hexamethylenediamine và adipic acid). Ba chất còn lại được điều chế bằng phản ứng trùng hợp."
    },
    {
        "num": 15,
        "text": "Phát biểu nào sau đây sai?",
        "options": [
            "A. Trùng ngưng buta-1,3-diene thu được polymer dùng để sản xuất chất dẻo.",
            "B. Trùng hợp ethylene thu được polymer dùng để sản xuất chất dẻo.",
            "C. Trùng ngưng ε-aminocaproic acid thu được polycaproamide.",
            "D. Trùng hợp vinyl cyanide thu được polymer dùng để sản xuất tơ nitron (tơ olon)."
        ],
        "correct": "A",
        "sol": "Chọn A. Phát biểu A sai vì buta-1,3-diene tham gia phản ứng trùng hợp (không phải trùng ngưng) và thu được polybutadiene dùng sản xuất cao su buna (không phải chất dẻo)."
    },
    {
        "num": 16,
        "text": "Không nên ủi (là) quá nóng quần áo bằng nylon, len, tơ tằm vì",
        "options": [
            "A. len, tơ tằm, tơ nylon có các nhóm liên kết trong phân tử kém bền với nhiệt.",
            "B. len, tơ tằm, tơ nylon mềm mại.",
            "C. len, tơ tằm, tơ nylon là những sợi thấm nước.",
            "D. len, tơ tằm, tơ nylon không thể là phẳng."
        ],
        "correct": "A",
        "sol": "Chọn A. Len, tơ tằm (protein) và tơ nylon (polyamide) đều chứa các liên kết peptide/amide -CO-NH- kém bền với nhiệt, khi ủi quá nóng sợi tơ dễ bị co rúm, giòn gãy và hỏng vải."
    },
    {
        "num": 17,
        "text": "(SBT – CTST). Khối lượng của một đoạn mạch nylon-6,6 là 27 346 amu và của một đoạn mạch capron (nylon-6) là 17 176 amu. Số lượng mắt xích trong đoạn mạch nylon-6,6 và capron nêu trên lần lượt là",
        "options": ["A. 113 và 152.", "B. 121 và 114.", "C. 121 và 152.", "D. 113 và 114."],
        "correct": "C",
        "sol": "Chọn C.\n• Mắt xích nylon-6,6 có công thức [-NH-(CH~2~)~6~-NH-CO-(CH~2~)~4~-CO-] với M = 226 amu.\n  => Số mắt xích = 27346 / 226 = 121.\n• Mắt xích capron có công thức [-NH-(CH~2~)~5~-CO-] với M = 113 amu.\n  => Số mắt xích = 17176 / 113 = 152."
    },
    {
        "num": 18,
        "text": "Tơ là những vật liệu polymer hình sợi dài và mảnh với độ bền nhất định. Cho các phát biểu sau về tơ:\n(a) Mạch phân nhánh, xếp song song với nhau.\n(b) Mềm, dai, không độc, tương đối bền với các dung môi thông thường và có khả năng nhuộm màu.\n(c) Nguyên liệu quan trọng trong ngành luyện kim.\n(d) Tơ tự nhiên: Là tơ có sẵn trong thiên nhiên như bông, sợi lanh, len lông cừu, tơ tằm,...\n(e) Tơ tổng hợp: Là tơ được chế tạo từ polymer tổng hợp như polyamide (nylon, capron,...).\n(f) Tơ bán tổng hợp: Là tơ xuất phát từ nguồn thiên nhiên nhưng được chế biến thêm bằng phương pháp hoá học (như tơ visco, tơ cellulose acetate,...).\nSố phát biểu đúng là",
        "options": ["A. 6.", "B. 3.", "C. 5.", "D. 4."],
        "correct": "D",
        "sol": "Chọn D. Các phát biểu đúng là (b), (d), (e), (f) (gồm 4 phát biểu). Phát biểu (a) sai vì tơ có cấu tạo mạch không phân nhánh; (c) sai vì tơ là nguyên liệu quan trọng của ngành dệt may, không phải luyện kim."
    }
]

TF_QUESTIONS = [
    {
        "num": 1,
        "intro": "(SBT – KNTT). Polymer X có thể chịu được nhiệt độ lên tới 160 oC nên được dùng làm ống dẫn nước nóng, hộp đựng thực phẩm có thể sử dụng trong lò vi sóng,... Các vật dụng làm từ X thường được in kí hiệu như hình bên.",
        "image": "p2_cau1_pp_symbol.png",
        "items": [
            ("a", "X được tổng hợp từ phản ứng trùng hợp pent-1-ene.", False, "Sai. X là polypropylene (PP), được điều chế bằng phản ứng trùng hợp propene (propylene) CH~2~=CH-CH~3~."),
            ("b", "Hộp nhựa làm từ X có thể đựng nước sôi mà không bị biến dạng.", True, "Đúng. PP có khả năng chịu nhiệt lên tới 160 oC, cao hơn nhiệt độ sôi của nước (100 oC)."),
            ("c", "X thuộc loại polymer nhiệt dẻo.", True, "Đúng. PP là chất nhiệt dẻo, khi đun nóng chảy thành chất lỏng nhớt và đông đặc lại khi để nguội."),
            ("d", "Nhựa làm từ X thuộc loại nhựa có thể tái chế.", True, "Đúng. Kí hiệu số 5 PP thuộc hệ thống mã phân loại nhựa tái chế quốc tế.")
        ]
    },
    {
        "num": 2,
        "intro": "(SBT - CD). Tơ visco được sản xuất phổ biến từ đầu thế kỉ XX, là một loại chất liệu được làm từ bột gỗ của cây như tre, đậu nành, mía,..., qua quá trình xử lí hoá học được sợi visco. Sợi visco thấm hút mồ hôi và thoáng khí, mềm mại nên rất phổ biến trong việc sản xuất quần áo, đặc biệt là trang phục mùa hè vì nó giúp người mặc cảm thấy mát mẻ và thoải mái. Những phát biểu nào sau đây là đúng?",
        "items": [
            ("a", "Tơ visco thuộc loại tơ thiên nhiên.", False, "Sai. Tơ visco thuộc loại tơ bán tổng hợp (tơ nhân tạo)."),
            ("b", "Tơ visco thuộc loại tơ bán tổng hợp.", True, "Đúng. Sợi tơ visco có nguồn gốc thiên nhiên nhưng được chế biến bằng hóa chất."),
            ("c", "Sợi visco có thành phần chính là cellulose đã được xử lí hoá chất.", True, "Đúng. Bột gỗ cellulose được xử lý với dung dịch NaOH và CS~2~ thành dung dịch keo nhớt visco rồi ép qua lỗ nhỏ vào bể acid để tái tạo sợi cellulose."),
            ("d", "Tơ là vật liệu polymer hình sợi, dài, mảnh, có độ bền nhất định, mạch không nhánh.", True, "Đúng. Đây là định nghĩa và đặc điểm cấu trúc chuẩn mực của vật liệu tơ sợi.")
        ]
    },
    {
        "num": 3,
        "intro": "(SBT - CD). Keo dán dùng để kết dính các vật liệu và được sử dụng rộng rãi trong đời sống, sản xuất. Những phát biểu nào sau đây là đúng?",
        "items": [
            ("a", "Nhựa vá săm là dung dịch keo của cao su trong dung môi hữu cơ dùng để vá săm xe.", True, "Đúng. Nhựa vá săm xe là dung dịch keo của cao su (thường là cao su thiên nhiên hoặc tổng hợp) trong dung môi hữu cơ dễ bay hơi."),
            ("b", "Keo dán epoxy gồm hai thành phần là hợp chất có chứa hai nhóm epoxy ở hai đầu và chất đóng rắn.", True, "Đúng. Keo epoxy (keo AB) gồm nhựa epoxy và chất đóng rắn (thường là các polyamine)."),
            ("c", "Bản chất của keo dán epoxy là tạo ra polymer có cấu trúc mạng không gian bền chắc, giúp gắn kết tốt hai vật liệu lại với nhau.", True, "Đúng. Khi trộn 2 thành phần, phản ứng mở vòng epoxy tạo liên kết ngang hình thành mạng lưới không gian bền chắc kết dính bề mặt."),
            ("d", "Khi sử dụng keo dán poly(urea-formaldehyde) cần bổ sung chất đóng rắn để tạo polymer có mạch phân nhánh.", False, "Sai. Chất đóng rắn tạo mạng lưới polymer không gian (mạng lưới 3 chiều), không phải mạch phân nhánh.")
        ]
    },
    {
        "num": 4,
        "intro": "(SBT – KNTT). Poly(butylene adipate terephthalate) (PBAT) là một polymer có khả năng phân huỷ sinh học, có tên thương mại là Ecoflex. PBAT có đặc tính tương tự như polyethylene mật độ thấp (LDPE) nên nó được sử dụng làm túi nylon, bao bì thực phẩm phân huỷ sinh học. PBAT được điều chế từ ba monomer sau đây:",
        "image": "p2_cau4_pbat_monomers.png",
        "items": [
            ("a", "PBAT thuộc loại polyester.", True, "Đúng. PBAT được tổng hợp từ dicarboxylic acid và diol, mạch chứa các liên kết ester -COO-."),
            ("b", "Phản ứng tổng hợp PBAT thuộc loại phản ứng trùng hợp.", False, "Sai. Đây là phản ứng trùng ngưng giữa dicarboxylic acid và diol, có giải phóng các phân tử nhỏ H~2~O."),
            ("c", "Một mắt xích PBAT gồm 3 nhóm ester.", False, "Sai. Một mắt xích PBAT chứa 4 nhóm ester do đồng trùng ngưng giữa 2 phân tử diacid và 2 phân tử diol butane-1,4-diol."),
            ("d", "Túi nylon làm từ PBAT thân thiện môi trường hơn so với LDPE.", True, "Đúng. PBAT có khả năng phân huỷ sinh học hoàn toàn trong điều kiện ủ tự nhiên thành CO~2~, H~2~O và mùn hữu cơ, thân thiện hơn nhiều so với LDPE khó phân hủy.")
        ]
    }
]

SA_QUESTIONS = [
    {
        "num": 1,
        "text": "(SBT – CTST). Có bao nhiêu polymer thiên nhiên trong các vật liệu sau: bông, tơ tằm, len, cellulose, tơ nitron?",
        "ans": "4",
        "sol": "Có 4 polymer thiên nhiên là: bông, tơ tằm, len, cellulose.\n(Tơ nitron là polymer tổng hợp điều chế từ vinyl cyanide)."
    },
    {
        "num": 2,
        "text": "Cho các polymer: (1) polyethylene, (2) poly(methyl methacrylate), (3) polybuta-1,3-diene, (4) polystyrene, (5) poly(vinyl acetate) và (6) tơ nylon-6,6. Trong các polymer trên, có bao nhiêu polymer có thể bị thuỷ phân trong dung dịch acid và dung dịch kiềm?",
        "ans": "3",
        "sol": "Có 3 polymer bị thủy phân trong dung dịch acid và kiềm:\n• (2) poly(methyl methacrylate) và (5) poly(vinyl acetate): chứa nhóm chức ester.\n• (6) tơ nylon-6,6: chứa nhóm chức amide (-CO-NH-).\nCác chất (1), (3), (4) chỉ chứa liên kết C-C bền vững nên không bị thủy phân."
    },
    {
        "num": 3,
        "text": "Một phân tử polyethylene có khối lượng phân tử bằng 56000 amu. Hệ số polymer hóa của phân tử polyethylene này là bao nhiêu?",
        "ans": "2000",
        "sol": "Polyethylene có dạng (-CH~2~-CH~2~-)~n~ với phân tử khối một mắt xích M = 28 amu.\nHệ số polymer hóa n = 56000 / 28 = 2000."
    },
    {
        "num": 4,
        "text": "Một mắt xích của polymer X gồm C, H, N. Hệ số polymer hóa của polymer này là 500 và có phân tử khối là 56500 amu. X chỉ có 1 nguyên tử N trong mỗi mắt xích. Mắt xích của polymer X có bao nhiêu nguyên tử carbon?",
        "ans": "6",
        "sol": "Khối lượng của một mắt xích là:\nM_mắt xích = 56500 / 500 = 113 amu.\nVì M = 113 amu và chứa 1 nguyên tử N, đây là mắt xích của polycaproamide (nylon-6): [-NH-(CH~2~)~5~-CO-] (C~6~H~11~NO).\nSố nguyên tử carbon trong một mắt xích là 6."
    },
    {
        "num": 5,
        "text": "Capron là một tơ sợi tổng hợp, được điều chế từ monomer là caprolactam. Một loại tơ capron có khối lượng phân tử là 14 916 amu. Số đơn vị mắt xích có trong phân tử loại tơ sợi này là bao nhiêu?",
        "image": "p3_cau5_caprolactam.png",
        "ans": "132",
        "sol": "Khi trùng hợp caprolactam mở vòng thu được tơ capron có mắt xích là [-NH-(CH~2~)~5~-CO-] với phân tử khối M = 113 amu.\nSố đơn vị mắt xích n = 14916 / 113 = 132."
    },
    {
        "num": 6,
        "text": "Polymer A trong suốt, được dùng làm hộp đựng thực phẩm, đồ chơi trẻ em, vỏ đĩa CD, DVD,... Trong công nghiệp sản xuất chất dẻo, polymer A được điều chế theo sơ đồ dưới đây:\nTừ 100 kg benzene và 32 m^3 ethylene (ở 25 oC, 1 bar), với hiệu suất mỗi quá trình (1), (2), (3) lần lượt là 60%, 55% và 60%, hãy tính khối lượng polymer A thu được theo đơn vị kg (kết quả làm tròn đến chữ số hàng phần mười).",
        "image": "p3_cau6_scheme.png",
        "ans": "26,4",
        "sol": "• Tính số mol các chất tham gia:\nn_benzene = 100 / 78 ≈ 1,282 kmol.\nn_ethylene = 32 / 24,79 ≈ 1,291 kmol.\nDo n_benzene < n_ethylene nên hiệu suất tính theo benzene.\n• Polymer A là polystyrene (-CH(C~6~H~5~)-CH~2~-)~n~ với phân tử khối mắt xích M = 104 kg/kmol.\n• Hiệu suất cả quá trình: H_tổng = 60% * 55% * 60% = 19,8%.\n• Khối lượng polymer A thu được:\nm_A = 1,282 * 104 * 19,8% ≈ 26,4 kg."
    }
]

def generate_exam_docx(is_teacher=False):
    filename = "De_Kiem_Tra_Hoa_12_De11_GiaoVien.docx" if is_teacher else "De_Kiem_Tra_Hoa_12_De11.docx"
    filepath = os.path.join(OUT_DIR, filename)
    
    doc = docx.Document()
    setup_page_and_header_footer(doc)
    add_exam_title_block(doc, is_teacher=is_teacher)
    
    # --- PHẦN I ---
    add_part_heading(doc, "PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.", "Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án.")
    
    for q in MCQ_QUESTIONS:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(1)
        p_q.paragraph_format.line_spacing = 1.15
        
        r_qnum = p_q.add_run(f"Câu {q['num']}. ")
        set_run_style(r_qnum, font_name="Times New Roman", font_size=11, bold=True)
        
        add_chem_run(p_q, q["text"], base_font="Times New Roman", base_size=11)
        
        # Nếu có hình ảnh kèm theo (như Câu 5 có biểu tượng PS)
        if q.get("image"):
            img_path = os.path.join(IMG_DIR, q["image"])
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(1)
                p_img.paragraph_format.space_after = Pt(2)
                run_img = p_img.add_run()
                run_img.add_picture(img_path, width=Cm(2.2))

        # Format options using Tabs (NO TABLE)
        opts = q["options"]
        max_opt_len = max(len(opt) for opt in opts)
        
        if max_opt_len < 22:
            # 1 dòng 4 lựa chọn (Tab stops 4.5, 9.0, 13.5 cm)
            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.space_before = Pt(1)
            p_opt.paragraph_format.space_after = Pt(2)
            p_opt.paragraph_format.line_spacing = 1.15
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(4.5), WD_TAB_ALIGNMENT.LEFT)
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(13.5), WD_TAB_ALIGNMENT.LEFT)
            
            for i, opt in enumerate(opts):
                is_correct = (opt.startswith(q["correct"] + ".")) and is_teacher
                bold = is_correct or opt[:2] in ["A.", "B.", "C.", "D."]
                color = COLOR_CORRECT if is_correct else None
                
                prefix = opt[:3]
                body = opt[3:]
                r_pre = p_opt.add_run(prefix)
                set_run_style(r_pre, font_name="Times New Roman", font_size=11, bold=True, color=color)
                add_chem_run(p_opt, body, base_font="Times New Roman", base_size=11, bold=is_correct, color=color)
                
                if i < 3:
                    p_opt.add_run("\t")
        elif max_opt_len < 46:
            # 2 dòng 2 lựa chọn (Tab stop 9.0 cm)
            for row in [(0, 1), (2, 3)]:
                p_opt = doc.add_paragraph()
                p_opt.paragraph_format.space_before = Pt(1)
                p_opt.paragraph_format.space_after = Pt(1.5)
                p_opt.paragraph_format.line_spacing = 1.15
                p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
                
                for idx_in_row, opt_idx in enumerate(row):
                    opt = opts[opt_idx]
                    is_correct = (opt.startswith(q["correct"] + ".")) and is_teacher
                    color = COLOR_CORRECT if is_correct else None
                    
                    prefix = opt[:3]
                    body = opt[3:]
                    r_pre = p_opt.add_run(prefix)
                    set_run_style(r_pre, font_name="Times New Roman", font_size=11, bold=True, color=color)
                    add_chem_run(p_opt, body, base_font="Times New Roman", base_size=11, bold=is_correct, color=color)
                    
                    if idx_in_row == 0:
                        p_opt.add_run("\t")
        else:
            # 4 dòng riêng biệt
            for opt in opts:
                p_opt = doc.add_paragraph()
                p_opt.paragraph_format.space_before = Pt(1)
                p_opt.paragraph_format.space_after = Pt(1.5)
                p_opt.paragraph_format.line_spacing = 1.15
                
                is_correct = (opt.startswith(q["correct"] + ".")) and is_teacher
                color = COLOR_CORRECT if is_correct else None
                
                prefix = opt[:3]
                body = opt[3:]
                r_pre = p_opt.add_run(prefix)
                set_run_style(r_pre, font_name="Times New Roman", font_size=11, bold=True, color=color)
                add_chem_run(p_opt, body, base_font="Times New Roman", base_size=11, bold=is_correct, color=color)

        if is_teacher and q.get("sol"):
            add_solution_paragraph(doc, q["sol"])

    # --- PHẦN II ---
    add_part_heading(doc, "PHẦN II. Câu trắc nghiệm đúng sai.", "Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.")
    
    for q in TF_QUESTIONS:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        
        r_qnum = p_q.add_run(f"Câu {q['num']}. ")
        set_run_style(r_qnum, font_name="Times New Roman", font_size=11, bold=True)
        
        add_chem_run(p_q, q["intro"], base_font="Times New Roman", base_size=11)
        
        if q.get("image"):
            img_path = os.path.join(IMG_DIR, q["image"])
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(2)
                p_img.paragraph_format.space_after = Pt(2)
                run_img = p_img.add_run()
                # Width depends on image
                w_cm = Cm(2.2) if "pp_symbol" in q["image"] else Cm(10.5)
                run_img.add_picture(img_path, width=w_cm)

        for item_key, item_text, is_true, item_sol in q["items"]:
            p_item = doc.add_paragraph()
            p_item.paragraph_format.space_before = Pt(1)
            p_item.paragraph_format.space_after = Pt(1.5)
            p_item.paragraph_format.line_spacing = 1.15
            p_item.paragraph_format.left_indent = Cm(0.5)
            
            r_k = p_item.add_run(f"{item_key}) ")
            set_run_style(r_k, font_name="Times New Roman", font_size=11, bold=True)
            
            add_chem_run(p_item, item_text, base_font="Times New Roman", base_size=11)
            
            if is_teacher:
                status_str = " [ĐÚNG]" if is_true else " [SAI]"
                color_st = COLOR_CORRECT if is_true else COLOR_PRIMARY
                r_st = p_item.add_run(status_str)
                set_run_style(r_st, font_name="Times New Roman", font_size=11, bold=True, color=color_st)
                
                p_sol = doc.add_paragraph()
                p_sol.paragraph_format.space_before = Pt(0)
                p_sol.paragraph_format.space_after = Pt(2)
                p_sol.paragraph_format.left_indent = Cm(1.0)
                r_sol_lbl = p_sol.add_run("-> ")
                set_run_style(r_sol_lbl, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_SOLUTION)
                add_chem_run(p_sol, item_sol, base_font="Times New Roman", base_size=10, italic=True)

    # --- PHẦN III ---
    add_part_heading(doc, "PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn.", "Thí sinh trả lời từ câu 1 đến câu 6.")
    
    for q in SA_QUESTIONS:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        
        r_qnum = p_q.add_run(f"Câu {q['num']}. ")
        set_run_style(r_qnum, font_name="Times New Roman", font_size=11, bold=True)
        
        add_chem_run(p_q, q["text"], base_font="Times New Roman", base_size=11)
        
        if q.get("image"):
            img_path = os.path.join(IMG_DIR, q["image"])
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(2)
                p_img.paragraph_format.space_after = Pt(2)
                run_img = p_img.add_run()
                w_cm = Cm(2.0) if "caprolactam" in q["image"] else Cm(12.0)
                run_img.add_picture(img_path, width=w_cm)

        if not is_teacher:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(3)
            p_ans.paragraph_format.left_indent = Cm(0.5)
            r_ans_lbl = p_ans.add_run("Trả lời: .................................................................................................................................")
            set_run_style(r_ans_lbl, font_name="Times New Roman", font_size=11, italic=True)
        else:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(2)
            p_ans.paragraph_format.left_indent = Cm(0.5)
            
            r_ans_lbl = p_ans.add_run("Đáp số: ")
            set_run_style(r_ans_lbl, font_name="Times New Roman", font_size=11, bold=True, color=COLOR_CORRECT)
            r_ans_val = p_ans.add_run(q["ans"])
            set_run_style(r_ans_val, font_name="Times New Roman", font_size=11, bold=True, color=COLOR_CORRECT)
            
            add_solution_paragraph(doc, q["sol"])

    # Bảng tổng hợp đáp án ở cuối đề Giáo Viên
    if is_teacher:
        p_table_title = doc.add_paragraph()
        p_table_title.paragraph_format.space_before = Pt(10)
        p_table_title.paragraph_format.space_after = Pt(3)
        r_tt = p_table_title.add_run("BẢNG TỔNG HỢP ĐÁP ÁN ĐỀ SỐ 11")
        set_run_style(r_tt, font_name="Times New Roman", font_size=11.5, bold=True, color=COLOR_PRIMARY)
        
        # Phần I Table
        p_p1 = doc.add_paragraph()
        p_p1.paragraph_format.space_before = Pt(3)
        p_p1.paragraph_format.space_after = Pt(1)
        r_p1 = p_p1.add_run("Phần I: Trắc nghiệm nhiều phương án lựa chọn")
        set_run_style(r_p1, font_name="Times New Roman", font_size=10.5, bold=True)
        
        tbl_ans1 = doc.add_table(rows=2, cols=18)
        tbl_ans1.autofit = False
        for c_idx in range(18):
            cell_q = tbl_ans1.cell(0, c_idx)
            cell_a = tbl_ans1.cell(1, c_idx)
            cell_q.width = Cm(1.0)
            cell_a.width = Cm(1.0)
            
            pq = cell_q.paragraphs[0]
            pq.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pq.paragraph_format.space_before = Pt(1)
            pq.paragraph_format.space_after = Pt(1)
            rq = pq.add_run(str(c_idx + 1))
            set_run_style(rq, font_name="Times New Roman", font_size=9, bold=True)
            
            pa = cell_a.paragraphs[0]
            pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pa.paragraph_format.space_before = Pt(1)
            pa.paragraph_format.space_after = Pt(1)
            ra = pa.add_run(MCQ_QUESTIONS[c_idx]["correct"])
            set_run_style(ra, font_name="Times New Roman", font_size=9.5, bold=True, color=COLOR_CORRECT)
            
            for cell in [cell_q, cell_a]:
                set_cell_border(cell, top={'val':'single','sz':'4','color':'CCCCCC'},
                                      bottom={'val':'single','sz':'4','color':'CCCCCC'},
                                      left={'val':'single','sz':'4','color':'CCCCCC'},
                                      right={'val':'single','sz':'4','color':'CCCCCC'})
                set_cell_margins(cell, top=20, bottom=20, left=10, right=10)

        # Phần II Table
        p_p2 = doc.add_paragraph()
        p_p2.paragraph_format.space_before = Pt(5)
        p_p2.paragraph_format.space_after = Pt(1)
        r_p2 = p_p2.add_run("Phần II: Đúng/Sai")
        set_run_style(r_p2, font_name="Times New Roman", font_size=10.5, bold=True)
        
        tbl_ans2 = doc.add_table(rows=5, cols=5)
        tbl_ans2.autofit = False
        headers_p2 = ["Câu", "a", "b", "c", "d"]
        for j, h in enumerate(headers_p2):
            cell = tbl_ans2.cell(0, j)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(h)
            set_run_style(r, font_name="Times New Roman", font_size=9.5, bold=True)
            set_cell_border(cell, top={'val':'single','sz':'4','color':'CCCCCC'},
                                  bottom={'val':'single','sz':'4','color':'CCCCCC'},
                                  left={'val':'single','sz':'4','color':'CCCCCC'},
                                  right={'val':'single','sz':'4','color':'CCCCCC'})
            set_cell_margins(cell, top=20, bottom=20, left=20, right=20)
            
        for i_q, q_tf in enumerate(TF_QUESTIONS):
            cell_q = tbl_ans2.cell(i_q + 1, 0)
            p = cell_q.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(f"Câu {q_tf['num']}")
            set_run_style(r, font_name="Times New Roman", font_size=9.5, bold=True)
            set_cell_border(cell_q, top={'val':'single','sz':'4','color':'CCCCCC'},
                                  bottom={'val':'single','sz':'4','color':'CCCCCC'},
                                  left={'val':'single','sz':'4','color':'CCCCCC'},
                                  right={'val':'single','sz':'4','color':'CCCCCC'})
            set_cell_margins(cell_q, top=20, bottom=20, left=20, right=20)
            
            for j_item, item in enumerate(q_tf["items"]):
                cell_item = tbl_ans2.cell(i_q + 1, j_item + 1)
                p = cell_item.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                st_txt = "Đ" if item[2] else "S"
                st_color = COLOR_CORRECT if item[2] else COLOR_PRIMARY
                r = p.add_run(st_txt)
                set_run_style(r, font_name="Times New Roman", font_size=9.5, bold=True, color=st_color)
                set_cell_border(cell_item, top={'val':'single','sz':'4','color':'CCCCCC'},
                                      bottom={'val':'single','sz':'4','color':'CCCCCC'},
                                      left={'val':'single','sz':'4','color':'CCCCCC'},
                                      right={'val':'single','sz':'4','color':'CCCCCC'})
                set_cell_margins(cell_item, top=20, bottom=20, left=20, right=20)

        # Phần III Table
        p_p3 = doc.add_paragraph()
        p_p3.paragraph_format.space_before = Pt(5)
        p_p3.paragraph_format.space_after = Pt(1)
        r_p3 = p_p3.add_run("Phần III: Trả lời ngắn")
        set_run_style(r_p3, font_name="Times New Roman", font_size=10.5, bold=True)
        
        tbl_ans3 = doc.add_table(rows=2, cols=6)
        tbl_ans3.autofit = False
        for c_idx in range(6):
            cell_q = tbl_ans3.cell(0, c_idx)
            cell_a = tbl_ans3.cell(1, c_idx)
            cell_q.width = Cm(2.0)
            cell_a.width = Cm(2.0)
            
            pq = cell_q.paragraphs[0]
            pq.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pq.paragraph_format.space_before = Pt(1)
            pq.paragraph_format.space_after = Pt(1)
            rq = pq.add_run(f"Câu {c_idx + 1}")
            set_run_style(rq, font_name="Times New Roman", font_size=9.5, bold=True)
            
            pa = cell_a.paragraphs[0]
            pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pa.paragraph_format.space_before = Pt(1)
            pa.paragraph_format.space_after = Pt(1)
            ra = pa.add_run(SA_QUESTIONS[c_idx]["ans"])
            set_run_style(ra, font_name="Times New Roman", font_size=9.5, bold=True, color=COLOR_CORRECT)
            
            for cell in [cell_q, cell_a]:
                set_cell_border(cell, top={'val':'single','sz':'4','color':'CCCCCC'},
                                      bottom={'val':'single','sz':'4','color':'CCCCCC'},
                                      left={'val':'single','sz':'4','color':'CCCCCC'},
                                      right={'val':'single','sz':'4','color':'CCCCCC'})
                set_cell_margins(cell, top=20, bottom=20, left=20, right=20)

    doc.save(filepath)
    print(f"Generated: {filepath}")

if __name__ == "__main__":
    generate_exam_docx(is_teacher=False)
    generate_exam_docx(is_teacher=True)
