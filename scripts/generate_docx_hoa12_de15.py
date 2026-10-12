# -*- coding: utf-8 -*-
"""
Script tạo 2 file Word (.docx) chuẩn Form Đề Kiểm Tra cho:
ĐỀ SỐ 15 - CHƯƠNG 5: PIN ĐIỆN – ĐIỆN PHÂN (HÓA HỌC 12)
Mã đề: 115
Bao gồm:
1. Bản Học sinh: De_Kiem_Tra_Hoa_12_De15.docx
2. Bản Giáo viên: De_Kiem_Tra_Hoa_12_De15_GiaoVien.docx (Có đáp án đỏ, lời giải chi tiết xanh, bảng đáp án)
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"

COLOR_CORAL_RED = RGBColor(211, 47, 47)      # #D32F2F
COLOR_SAGE_GREEN = "8BC390"                 # #8BC390
COLOR_DEEP_BLUE = RGBColor(21, 101, 192)    # #1565C0
COLOR_DARK_GREEN = RGBColor(46, 125, 50)    # #2E7D32
COLOR_BLACK = RGBColor(0, 0, 0)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.59)     # 1.5 cm
        s.bottom_margin = Inches(0.59)  # 1.5 cm
        s.left_margin = Inches(0.59)    # 1.5 cm
        s.right_margin = Inches(0.59)   # 1.5 cm
        s.header_distance = Inches(0.24) # 0.6 cm
        s.footer_distance = Inches(0.24) # 0.6 cm

        # --- HEADER ---
        header = s.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.text = ""
        hp.paragraph_format.space_before = Pt(0)
        hp.paragraph_format.space_after = Pt(0)

        tbl_h = header.add_table(rows=1, cols=2, width=Inches(7.08)) # 18.0 cm
        tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
        c0, c1 = tbl_h.rows[0].cells
        c0.width = Inches(3.54)
        c1.width = Inches(3.54)
        set_cell_margins(c0, 60, 60, 0, 0)
        set_cell_margins(c1, 60, 60, 0, 0)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run("Tài liệu lưu hành nội bộ")
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(10)
        r0.font.italic = True
        r0.font.color.rgb = COLOR_CORAL_RED

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1_a = p1.add_run("Thầy ")
        r1_a.font.name = "Times New Roman"
        r1_a.font.size = Pt(10)
        r1_a.font.color.rgb = COLOR_CORAL_RED
        r1_b = p1.add_run("TRẦN VĂN THẠNH")
        r1_b.font.name = "Times New Roman"
        r1_b.font.size = Pt(10)
        r1_b.font.bold = True
        r1_b.font.color.rgb = COLOR_CORAL_RED
        r1_c = p1.add_run(" - 0777.470.803")
        r1_c.font.name = "Times New Roman"
        r1_c.font.size = Pt(10)
        r1_c.font.color.rgb = COLOR_CORAL_RED

        tblPr_h = tbl_h._tbl.tblPr
        tblBorders_h = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{COLOR_SAGE_GREEN}"/>'
            f'</w:tblBorders>'
        )
        tblPr_h.append(tblBorders_h)

        # --- FOOTER ---
        footer = s.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.text = ""
        fp.paragraph_format.space_before = Pt(0)
        fp.paragraph_format.space_after = Pt(0)

        tbl_f = footer.add_table(rows=1, cols=2, width=Inches(7.08))
        tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
        fc0, fc1 = tbl_f.rows[0].cells
        fc0.width = Inches(5.31)
        fc1.width = Inches(1.77)
        set_cell_margins(fc0, 60, 60, 0, 0)
        set_cell_margins(fc1, 60, 60, 0, 0)

        fp0 = fc0.paragraphs[0]
        fp0.paragraph_format.space_before = Pt(0)
        fp0.paragraph_format.space_after = Pt(0)
        fr0 = fp0.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
        fr0.font.name = "Times New Roman"
        fr0.font.size = Pt(9.5)
        fr0.font.italic = True
        fr0.font.color.rgb = COLOR_CORAL_RED

        fp1 = fc1.paragraphs[0]
        fp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp1.paragraph_format.space_before = Pt(0)
        fp1.paragraph_format.space_after = Pt(0)
        fr1 = fp1.add_run("Trang ")
        fr1.font.name = "Times New Roman"
        fr1.font.size = Pt(10)
        fr1.font.bold = True
        fr1.font.color.rgb = COLOR_CORAL_RED

        fldSimple = OxmlElement('w:fldSimple')
        fldSimple.set(qn('w:instr'), 'PAGE')
        fp1._p.append(fldSimple)

        tblPr_f = tbl_f._tbl.tblPr
        tblBorders_f = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{COLOR_SAGE_GREEN}"/>'
            f'</w:tblBorders>'
        )
        tblPr_f.append(tblBorders_f)

def add_chem_run(paragraph, text, bold=False, italic=False, color=None, font_size=11):
    tokens = re.split(r'(\~[0-9\+\-mnx]+\~|\^[0-9\+\-mnx]+\^)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('~') and tok.endswith('~'):
            content = tok[1:-1]
            r = paragraph.add_run(content)
            r.font.subscript = True
        elif tok.startswith('^') and tok.endswith('^'):
            content = tok[1:-1]
            r = paragraph.add_run(content)
            r.font.superscript = True
        else:
            r = paragraph.add_run(tok)
            
        r.font.name = "Times New Roman"
        r.font.size = Pt(font_size)
        r.font.bold = bold
        r.font.italic = italic
        if color:
            r.font.color.rgb = color

def create_title_block(doc, is_teacher=False):
    lines = [
        ("SỞ GIÁO DỤC VÀ ĐÀO TẠO THỪA THIÊN HUẾ", "ĐỀ KIỂM TRA ĐỊNH KỲ CHƯƠNG 5 (PIN ĐIỆN – ĐIỆN PHÂN)", True, True),
        ("TRƯỜNG THPT HAI BÀ TRƯNG", "MÔN: HOÁ HỌC – LỚP 12", True, True),
        ("ĐỀ CHÍNH THỨC", "Thời gian làm bài: 50 phút (không kể thời gian phát đề)", True, False),
        ("(Đề thi có 04 trang)", "Mã đề thi: 115", False, True)
    ]
    for left_text, right_text, left_bold, right_bold in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.08), WD_TAB_ALIGNMENT.RIGHT)

        r_left = p.add_run(left_text)
        r_left.font.name = "Times New Roman"
        r_left.font.size = Pt(10.5)
        r_left.font.bold = left_bold
        if not left_bold:
            r_left.font.italic = True

        p.add_run("\t")

        r_right = p.add_run(right_text)
        r_right.font.name = "Times New Roman"
        r_right.font.size = Pt(10.5)
        r_right.font.bold = right_bold
        if not right_bold:
            r_right.font.italic = True
        if "ĐỀ KIỂM TRA" in right_text:
            r_right.font.color.rgb = COLOR_DARK_GREEN

    if is_teacher:
        p_gv = doc.add_paragraph()
        p_gv.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_gv.paragraph_format.space_before = Pt(4)
        p_gv.paragraph_format.space_after = Pt(4)
        r_gv = p_gv.add_run("(BẢN GIÁO VIÊN -- CÓ LỜI GIẢI CHI TIẾT)")
        r_gv.font.name = "Times New Roman"
        r_gv.font.size = Pt(12)
        r_gv.font.bold = True
        r_gv.font.color.rgb = COLOR_CORAL_RED

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(3)
    p_info.paragraph_format.space_after = Pt(2)
    r_info = p_info.add_run("Họ, tên thí sinh: ....................................................................................... Lớp: ........................ SBD: ....................")
    r_info.font.name = "Times New Roman"
    r_info.font.size = Pt(10)
    r_info.font.italic = True

    p_ntk = doc.add_paragraph()
    p_ntk.paragraph_format.space_before = Pt(0)
    p_ntk.paragraph_format.space_after = Pt(5)
    r_ntk = p_ntk.add_run("(Cho biết nguyên tử khối: H = 1; C = 12; N = 14; O = 16; Na = 23; Mg = 24; Al = 27; S = 32; Cl = 35,5; K = 39; Ca = 40; Mn = 55; Fe = 56; Cu = 64; Zn = 65; Ag = 108; Cd = 112; Au = 197; Pb = 207; F = 96500 C/mol).")
    r_ntk.font.name = "Times New Roman"
    r_ntk.font.size = Pt(9.5)
    r_ntk.font.italic = True

def add_section_header(doc, title, subtitle):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(title)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.bold = True

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(4)
    p_sub.paragraph_format.line_spacing = 1.15
    r_sub = p_sub.add_run(subtitle)
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(10.5)
    r_sub.font.italic = True

def add_loigiai(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28) # 0.7 cm
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15

    r_lbl = p.add_run("Lời giải chi tiết:\n")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(10)
    r_lbl.font.bold = True
    r_lbl.font.color.rgb = COLOR_DEEP_BLUE

    for line in text.strip().split("\n"):
        p_line = doc.add_paragraph()
        p_line.paragraph_format.left_indent = Inches(0.28)
        p_line.paragraph_format.space_before = Pt(0)
        p_line.paragraph_format.space_after = Pt(2)
        p_line.paragraph_format.line_spacing = 1.15
        
        is_bold = line.startswith("Chọn ") or line.startswith("Đáp số:") or line.startswith("=>")
        add_chem_run(p_line, line, bold=is_bold, color=COLOR_DEEP_BLUE, font_size=10)

def add_mcq_options(doc, options, correct_key=None, is_teacher=False, layout="1_row"):
    """
    Tạo các phương án A, B, C, D bằng Paragraph thuần tuý kết hợp Tab stops, TUYỆT ĐỐI KHÔNG DÙNG BẢNG.
    """
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15

    if layout == "1_row":
        p.paragraph_format.tab_stops.add_tab_stop(Inches(1.77), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(3.54), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(5.31), WD_TAB_ALIGNMENT.LEFT)
        for i, (k, txt) in enumerate(options):
            if i > 0:
                p.add_run("\t")
            is_c = (k == correct_key) and is_teacher
            col = COLOR_CORAL_RED if is_c else None
            add_chem_run(p, f"{k}. ", bold=True, color=col)
            add_chem_run(p, txt, bold=is_c, color=col)

    elif layout == "2_rows":
        p.paragraph_format.tab_stops.add_tab_stop(Inches(3.54), WD_TAB_ALIGNMENT.LEFT)
        for i in range(2):
            k, txt = options[i]
            if i > 0:
                p.add_run("\t")
            is_c = (k == correct_key) and is_teacher
            col = COLOR_CORAL_RED if is_c else None
            add_chem_run(p, f"{k}. ", bold=True, color=col)
            add_chem_run(p, txt, bold=is_c, color=col)

        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(3)
        p2.paragraph_format.line_spacing = 1.15
        p2.paragraph_format.tab_stops.add_tab_stop(Inches(3.54), WD_TAB_ALIGNMENT.LEFT)
        for i in range(2, 4):
            k, txt = options[i]
            if i > 2:
                p2.add_run("\t")
            is_c = (k == correct_key) and is_teacher
            col = COLOR_CORAL_RED if is_c else None
            add_chem_run(p2, f"{k}. ", bold=True, color=col)
            add_chem_run(p2, txt, bold=is_c, color=col)

    else: # 4_rows
        for i, (k, txt) in enumerate(options):
            if i > 0:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
            is_c = (k == correct_key) and is_teacher
            col = COLOR_CORAL_RED if is_c else None
            add_chem_run(p, f"{k}. ", bold=True, color=col)
            add_chem_run(p, txt, bold=is_c, color=col)

def generate_exam_docx(is_teacher=False):
    doc = docx.Document()
    add_header_footer(doc)
    create_title_block(doc, is_teacher=is_teacher)

    # ==============================================================
    # PHẦN I
    # ==============================================================
    add_section_header(
        doc,
        "PHẦN I (4,5 điểm -- 18 câu): Câu trắc nghiệm nhiều phương án lựa chọn.",
        "Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án."
    )

    p1_data = [
        (
            1,
            "(SBT – KNTT). Trong pin điện hoá Zn - Cu, ở anode (cực âm) xảy ra quá trình",
            [
                ("A", "oxi hoá Zn thành ion Zn^2+^."),
                ("B", "khử ion Cu^2+^ thành Cu."),
                ("C", "khử Cu thành ion Cu^2+^."),
                ("D", "oxi hoá ion Zn^2+^ thành Zn.")
            ],
            "A",
            "2_rows",
            "Ở anode (cực âm) của pin Zn - Cu xảy ra quá trình oxi hoá kẽm: Zn → Zn^2+^ + 2e.\n"
            "Chọn A."
        ),
        (
            2,
            "Pin điện là gì?",
            [
                ("A", "Một thiết bị lưu trữ và cung cấp điện năng."),
                ("B", "Một thiết bị dùng để đo nhiệt độ."),
                ("C", "Một thiết bị dùng để đo áp suất."),
                ("D", "Một thiết bị dùng để lưu trữ năng lượng cơ học.")
            ],
            "A",
            "4_rows",
            "Pin điện là thiết bị biến đổi năng lượng hoá học của phản ứng oxi hoá - khử thành điện năng, dùng để lưu trữ và cung cấp điện năng.\n"
            "Chọn A."
        ),
        (
            3,
            "Pin Galvani còn được gọi là gì?",
            [
                ("A", "Pin nhiên liệu."),
                ("B", "Pin điện hóa."),
                ("C", "Pin nhiệt điện."),
                ("D", "Pin quang điện.")
            ],
            "B",
            "1_row",
            "Pin Galvani là tên gọi truyền thống của pin điện hoá (phản ứng hoá học tự diễn biến sinh ra dòng điện).\n"
            "Chọn B."
        ),
        (
            4,
            "Điện cực dương trong pin gọi là gì?",
            [
                ("A", "Anode."),
                ("B", "Cathode."),
                ("C", "Điện cực."),
                ("D", "Cực âm.")
            ],
            "B",
            "1_row",
            "Trong pin điện hoá, cực dương là nơi xảy ra quá trình khử (nhận electron), được gọi là cathode.\n"
            "Chọn B."
        ),
        (
            5,
            "(SBT – KNTT). Trong quá trình mạ bạc cho một chiếc vòng bằng thép thì ở anode xảy ra quá trình",
            [
                ("A", "Ag → Ag^+^ + 1e"),
                ("B", "Fe → Fe^2+^ + 2e"),
                ("C", "2H~2~O → 4H^+^ + O~2~ + 4e"),
                ("D", "C → C^4+^ + 4e")
            ],
            "A",
            "1_row",
            "Quá trình mạ bạc dùng phương pháp dương cực tan: Anode làm bằng bạc (Ag), cathode là chiếc vòng cần mạ.\n"
            "Ở anode xảy ra quá trình oxi hoá bạc tan vào dung dịch: Ag → Ag^+^ + 1e.\n"
            "Chọn A."
        ),
        (
            6,
            "(SBT – KNTT). Khi điện phân dung dịch NaCl bằng dòng điện một chiều (với điện cực trơ, có màng ngăn xốp) thì ở cathode xảy ra quá trình",
            [
                ("A", "oxi hoá H~2~O thành H^+^ và O~2~."),
                ("B", "khử Cl^-^ thành Cl~2~."),
                ("C", "oxi hoá Cl^-^ thành Cl~2~."),
                ("D", "khử H~2~O thành H~2~ và OH^-^.")
            ],
            "D",
            "2_rows",
            "Tại cathode (cực âm), Na^+^ có tính oxi hoá yếu nên không bị khử trong nước; H~2~O bị khử trước:\n"
            "2H~2~O + 2e → H~2~↑ + 2OH^-^.\n"
            "Chọn D."
        ),
        (
            7,
            "(SBT – KNTT). Khi điện phân dung dịch CuSO~4~ bằng dòng điện một chiều (với điện cực anode bằng Cu) thì ở anode xảy ra quá trình",
            [
                ("A", "oxi hoá H~2~O thành H^+^ và O~2~."),
                ("B", "khử Cu^2+^ thành Cu."),
                ("C", "oxi hoá Cu thành Cu^2+^."),
                ("D", "khử H~2~O thành H~2~ và OH^-^.")
            ],
            "C",
            "2_rows",
            "Hiện tượng dương cực tan: Anode làm bằng Cu nên bị oxi hoá tan vào dung dịch: Cu → Cu^2+^ + 2e.\n"
            "Chọn C."
        ),
        (
            8,
            "So sánh hai kim loại A và B có thế điện cực chuẩn lần lượt là -0,25 V và +0,34 V, kim loại nào có khả năng khử mạnh hơn?",
            [
                ("A", "Kim loại A."),
                ("B", "Kim loại B."),
                ("C", "Cả hai đều có khả năng khử như nhau."),
                ("D", "Không thể xác định.")
            ],
            "A",
            "2_rows",
            "Kim loại có thế điện cực chuẩn âm hơn (nhỏ hơn) thì có tính khử mạnh hơn.\n"
            "Vì E°(A) = -0,25 V < E°(B) = +0,34 V nên kim loại A có khả năng khử mạnh hơn B.\n"
            "Chọn A."
        ),
        (
            9,
            "(SBT – CTST). Cho pin điện hoá Zn-Cu. Quá trình xảy ra ở cực dương của pin là",
            [
                ("A", "Zn^2+^ + 2e → Zn"),
                ("B", "Cu^2+^ + 2e → Cu"),
                ("C", "Zn → Zn^2+^ + 2e"),
                ("D", "Cu → Cu^2+^ + 2e")
            ],
            "B",
            "1_row",
            "Trong pin Zn - Cu, cực dương là cathode (thanh Cu). Ở cực dương xảy ra sự khử ion Cu^2+^: Cu^2+^ + 2e → Cu.\n"
            "Chọn B."
        ),
        (
            10,
            "(SBT – CTST). Cho pin điện hoá Al-Pb. Quá trình xảy ra ở cực âm của pin là",
            [
                ("A", "Al^3+^ + 3e → Al"),
                ("B", "Pb^2+^ + 2e → Pb"),
                ("C", "Al → Al^3+^ + 3e"),
                ("D", "Pb → Pb^2+^ + 2e")
            ],
            "C",
            "1_row",
            "Vì E°(Al^3+^/Al) (-1,66 V) < E°(Pb^2+^/Pb) (-0,13 V) nên Al đóng vai trò cực âm (anode).\n"
            "Tại cực âm Al bị oxi hoá: Al → Al^3+^ + 3e.\n"
            "Chọn C."
        ),
        (
            11,
            "(SBT – CTST). Phát biểu nào sau đây không đúng về pin Galvani?",
            [
                ("A", "Ở điện cực dương xảy ra quá trình oxi hoá."),
                ("B", "Cathode là kim loại yếu hơn, đóng vai trò điện cực dương của pin."),
                ("C", "Anode là kim loại mạnh hơn, đóng vai trò điện cực âm của pin."),
                ("D", "Phản ứng hóa học diễn ra trong pin kèm theo sự giải phóng điện năng.")
            ],
            "A",
            "4_rows",
            "Ở điện cực dương (cathode) của pin luôn xảy ra quá trình khử (nhận electron), không phải oxi hoá. Phát biểu A là sai.\n"
            "Chọn A."
        ),
        (
            12,
            "(SBT – CTST). Khi điện phân dung dịch nào sau đây, tại anode xảy ra quá trình oxi hóa nước?",
            [
                ("A", "Dung dịch ZnCl~2~."),
                ("B", "Dung dịch CuCl~2~."),
                ("C", "Dung dịch AgNO~3~."),
                ("D", "Dung dịch MgCl~2~.")
            ],
            "C",
            "2_rows",
            "Các muối ZnCl~2~, CuCl~2~, MgCl~2~ chứa ion Cl^-^ dễ bị oxi hoá hơn nước thành Cl~2~ ở anode.\n"
            "Dung dịch AgNO~3~ chứa NO~3~^-^ (gốc acid có oxi khó bị oxi hoá), do đó nước bị oxi hoá ở anode giải phóng O~2~ và H^+^.\n"
            "Chọn C."
        ),
        (
            13,
            "(SBT – CTST). Điện phân dung dịch CuCl~2~ với điện cực trơ. Nồng độ của CuCl~2~ trong dung dịch",
            [
                ("A", "giảm dần."),
                ("B", "tăng dần."),
                ("C", "không thay đổi."),
                ("D", "tăng dần rồi giảm dần.")
            ],
            "A",
            "2_rows",
            "Phản ứng điện phân: CuCl~2~ → Cu + Cl~2~↑. CuCl~2~ bị phân huỷ liên tục làm nồng độ CuCl~2~ giảm dần.\n"
            "Chọn A."
        ),
        (
            14,
            "(SBT – CTST). Điện phân dung dịch CuSO~4~ với hai điện cực trơ. Sau một thời gian, màu xanh của dung dịch nhạt dần do:",
            [
                ("A", "khí H~2~ sinh ra đã khử màu của dung dịch."),
                ("B", "dung dịch bị pha loãng dần nên màu xanh nhạt dần."),
                ("C", "ion Cu^2+^ bị khử dần thành đồng kim loại."),
                ("D", "ion Cu^2+^ được tạo thêm.")
            ],
            "C",
            "4_rows",
            "Màu xanh của dung dịch do ion Cu^2+^ quy định. Khi điện phân, ion Cu^2+^ bị khử thành Cu kim loại ở cathode làm nồng độ Cu^2+^ giảm dần, khiến màu xanh nhạt dần.\n"
            "Chọn C."
        ),
        (
            15,
            "(SBT – CTST). Điện phân điện cực trơ, có màng ngăn một dung dịch chứa các ion: Fe^2+^, Fe^3+^, Cu^2+^, Cl^-^. Thứ tự điện phân xảy ra ở cathode là",
            [
                ("A", "Fe^2+^, Fe^3+^, Cu^2+^."),
                ("B", "Fe^2+^, Cu^2+^, Fe^3+^."),
                ("C", "Fe^3+^, Cu^2+^, Fe^2+^."),
                ("D", "Fe^3+^, Fe^2+^, Cu^2+^.")
            ],
            "C",
            "2_rows",
            "Cation có tính oxi hoá mạnh hơn sẽ bị khử trước tại cathode. Thứ tự tính oxi hoá: Fe^3+^ > Cu^2+^ > Fe^2+^.\n"
            "Do đó thứ tự bị khử ở cathode là: Fe^3+^ → Cu^2+^ → Fe^2+^.\n"
            "Chọn C."
        ),
        (
            16,
            "(SBT – CTST). Cho pin điện hoá Mn-Cd có E°pin(Mn-Cd) = 0,79 V và E°(Cd^2+^/Cd) = -0,40 V. Thế điện cực chuẩn của cặp oxi hoá khử Mn^2+^/Mn có giá trị là",
            [
                ("A", "0,27 V."),
                ("B", "-1,19 V."),
                ("C", "0,26 V."),
                ("D", "-0,25 V.")
            ],
            "B",
            "1_row",
            "Công thức: E°pin = E°(Cd^2+^/Cd) - E°(Mn^2+^/Mn).\n"
            "=> E°(Mn^2+^/Mn) = E°(Cd^2+^/Cd) - E°pin = -0,40 - 0,79 = -1,19 V.\n"
            "Chọn B."
        ),
        (
            17,
            "(SBT – CTST). Quá trình điện phân dung dịch NaCl (điện cực trơ, có màng ngăn) và điện phân NaCl nóng chảy (điện cực trơ) có điểm giống nhau là",
            [
                ("A", "tại cathode xảy ra sự khử ion Na^+^."),
                ("B", "tại cathode xảy ra sự khử phân tử H~2~O."),
                ("C", "tại anode xảy ra sự oxi hóa ion Cl^-^."),
                ("D", "tại anode xảy ra sự oxi hóa phân tử H~2~O.")
            ],
            "C",
            "2_rows",
            "Ở anode (cực dương) của cả hai quá trình đều xảy ra sự oxi hoá ion Cl^-^ thành khí Cl~2~: 2Cl^-^ → Cl~2~ + 2e.\n"
            "Chọn C."
        ),
        (
            18,
            "Cho sơ đồ phản ứng:\n"
            "(1) NaCl + H~2~O → X + Y↑ (ở cathode) + Z↑ (ở anode)\n"
            "(2) X + CO~2~ dư → T\n"
            "Phát biểu nào sau đây đúng?",
            [
                ("A", "Chất khí Z có thể khử được CaO thành Ca ở nhiệt độ cao."),
                ("B", "Dung dịch X có tính tẩy màu, sát trùng, thường gọi là Javel."),
                ("C", "Chất T được dùng làm thuốc giảm đau dạ dày."),
                ("D", "Chất khí Y không có màu, mùi, vị và Y có thể duy trì sự cháy.")
            ],
            "C",
            "4_rows",
            "Phương trình (1): 2NaCl + 2H~2~O → 2NaOH (X) + H~2~↑ (Y, cathode) + Cl~2~↑ (Z, anode).\n"
            "Phương trình (2): NaOH + CO~2~ dư → NaHCO~3~ (T).\n"
            "Chất T là NaHCO~3~ (baking soda) dùng làm thuốc chữa đau dạ dày do thừa axit. Nhận định C đúng.\n"
            "Chọn C."
        )
    ]

    for q_num, content, opts, ans_key, layout, sol in p1_data:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        add_chem_run(p_q, f"Câu {q_num}. ", bold=True)
        add_chem_run(p_q, content)

        add_mcq_options(doc, opts, correct_key=ans_key, is_teacher=is_teacher, layout=layout)
        if is_teacher:
            add_loigiai(doc, sol)

    # ==============================================================
    # PHẦN II
    # ==============================================================
    add_section_header(
        doc,
        "PHẦN II (4,0 điểm -- 4 câu): Câu trắc nghiệm đúng sai.",
        "Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai."
    )

    p2_data = [
        (
            1,
            "(SBT – CĐ). Chọn tính đúng sai cho mỗi phát biểu sau về acquy:",
            [
                ("a", "Một ưu điểm của acquy là tái sử dụng được nhiều lần.", True),
                ("b", "Phản ứng xảy ra trong acquy cũng giống như phản ứng xảy ra trong pin Galvani nhưng có thể đảo ngược.", True),
                ("c", "Acquy không gây ô nhiễm môi trường.", False),
                ("d", "Acquy là nguồn điện hoá học có thể hoạt động liên tục.", False)
            ],
            "a) Đúng. Acquy là pin thứ cấp có thể sạc lại để dùng nhiều lần.\n"
            "b) Đúng. Khi phóng điện phản ứng tự phát sinh điện; khi sạc phản ứng bị cưỡng bức theo chiều ngược lại.\n"
            "c) Sai. Acquy chứa kim loại nặng độc hại (chì, axit H~2~SO~4~, lithi...) gây ô nhiễm nghiêm trọng nếu thải bỏ bừa bãi.\n"
            "d) Sai. Acquy chỉ hoạt động trong một thời gian rồi cạn điện, phải nạp lại mới dùng tiếp được."
        ),
        (
            2,
            "(SBT – CTST). Em hãy cho biết những thí nghiệm sau đây có thể tạo thành một pin điện hóa là đúng hay sai:",
            [
                ("a", "Một điện cực nhôm và một điện cực đồng cùng nhúng vào dung dịch nước muối.", True),
                ("b", "Một điện cực nhôm và một điện cực đồng cùng nhúng vào nước cất.", False),
                ("c", "Hai điện cực làm bằng đồng cùng nhúng vào dung dịch nước vôi trong.", False),
                ("d", "Hai điện cực làm bằng nhựa nhúng vào dầu hoả.", False)
            ],
            "Điều kiện tạo thành pin điện hoá: Hai điện cực có bản chất khác nhau cùng tiếp xúc với dung dịch chất điện li và nối kín mạch ngoài.\n"
            "a) Đúng. Al và Cu khác bản chất, nước muối NaCl là chất điện li.\n"
            "b) Sai. Nước cất không dẫn điện (không phải dung dịch điện li).\n"
            "c) Sai. Cả 2 cực đều bằng Cu (cùng bản chất) thì hiệu số điện thế bằng 0.\n"
            "d) Sai. Nhựa không dẫn điện, dầu hoả không phải chất điện li."
        ),
        (
            3,
            "(SBT – KNTT). Trong quá trình điện phân dung dịch CuSO~4~ với anode bằng đồng. Chọn tính đúng sai cho mỗi phát biểu sau:",
            [
                ("a", "Ở anode xảy ra quá trình oxi hóa nước.", False),
                ("b", "Khối lượng anode không thay đổi.", False),
                ("c", "Nồng độ CuSO~4~ trong dung dịch giảm dần.", False),
                ("d", "Khối lượng cathode tăng.", True)
            ],
            "Hiện tượng dương cực tan:\n"
            "- Ở anode (cực Cu): Cu bị oxi hoá tan ra (Cu → Cu^2+^ + 2e) => khối lượng anode giảm, nước không bị oxi hoá.\n"
            "- Ở cathode: Cu^2+^ bị khử bám vào cathode (Cu^2+^ + 2e → Cu) => khối lượng cathode tăng.\n"
            "- Lượng Cu^2+^ tan ra ở anode bằng lượng Cu^2+^ mất đi ở cathode => nồng độ CuSO~4~ không đổi.\n"
            "a) Sai; b) Sai; c) Sai; d) Đúng."
        ),
        (
            4,
            "(SBT – KNTT). Điện phân dung dịch NaCl bão hoà (với điện cực trơ, màng ngăn xốp) đến khi nồng độ NaCl giảm đi một nửa thì dừng điện phân. Chọn tính đúng sai cho mỗi phát biểu sau:",
            [
                ("a", "Dung dịch sau điện phân làm phenolphthalein chuyển màu hồng.", True),
                ("b", "Ở cathode chỉ xảy ra quá trình khử ion Na^+^.", False),
                ("c", "Số mol khí Cl~2~ thoát ra ở anode bằng số mol H~2~ thoát ra ở cathode.", True),
                ("d", "Thứ tự điện phân ở anode là H~2~O, Cl^-^.", False)
            ],
            "Phương trình: 2NaCl + 2H~2~O → 2NaOH + H~2~↑ (cathode) + Cl~2~↑ (anode).\n"
            "a) Đúng. Sinh ra dung dịch bazơ NaOH làm phenolphthalein đổi màu hồng.\n"
            "b) Sai. Ở cathode chỉ xảy ra sự khử H~2~O, ion Na^+^ không bị khử.\n"
            "c) Đúng. Tỉ lệ mol khí H~2~ : Cl~2~ = 1 : 1.\n"
            "d) Sai. Ở anode ion Cl^-^ bị oxi hoá trước H~2~O."
        )
    ]

    for q_num, content, sub_items, sol in p2_data:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        add_chem_run(p_q, f"Câu {q_num}. ", bold=True)
        add_chem_run(p_q, content)

        for sub_lbl, sub_txt, is_true in sub_items:
            p_sub = doc.add_paragraph()
            p_sub.paragraph_format.left_indent = Inches(0.2)
            p_sub.paragraph_format.space_before = Pt(0)
            p_sub.paragraph_format.space_after = Pt(1.5)
            p_sub.paragraph_format.line_spacing = 1.15

            add_chem_run(p_sub, f"{sub_lbl}) ", bold=True)
            add_chem_run(p_sub, sub_txt)

            if is_teacher:
                status_txt = " [ĐÚNG]" if is_true else " [SAI]"
                add_chem_run(p_sub, status_txt, bold=True, color=COLOR_CORAL_RED)

        if is_teacher:
            add_loigiai(doc, sol)

    # ==============================================================
    # PHẦN III
    # ==============================================================
    add_section_header(
        doc,
        "PHẦN III (1,5 điểm -- 6 câu): Câu trắc nghiệm yêu cầu trả lời ngắn.",
        "Thí sinh trả lời từ câu 1 đến câu 6."
    )

    p3_data = [
        (
            1,
            "Cho giá trị thế điện cực chuẩn: Al^3+^/Al (-1,676 V); Zn^2+^/Zn (-0,763 V); Fe^2+^/Fe (-0,440 V); Cu^2+^/Cu (+0,340 V); Ag^+^/Ag (+0,799 V); Au^3+^/Au (+1,520 V). Hãy cho biết có bao nhiêu kim loại trong bảng có tính khử mạnh hơn Cu?",
            "3",
            "Kim loại có tính khử mạnh hơn Cu khi thế điện cực chuẩn E° < E°(Cu^2+^/Cu) = +0,340 V.\n"
            "Các kim loại thoả mãn là: Al (-1,676 V), Zn (-0,763 V), Fe (-0,440 V). Tổng cộng có 3 kim loại.\n"
            "Đáp số: 3."
        ),
        (
            2,
            "Cho pin điện hóa Zn^2+^/Zn và Cu^2+^/Cu. Cho E°(Cu^2+^/Cu) = +0,34 V và E°(Zn^2+^/Zn) = -0,76 V. Tính sức điện động chuẩn của pin (theo đơn vị V).",
            "1,1",
            "Sức điện động chuẩn của pin:\n"
            "E°pin = E°(Cu^2+^/Cu) - E°(Zn^2+^/Zn) = 0,34 - (-0,76) = 1,10 V.\n"
            "Đáp số: 1,1."
        ),
        (
            3,
            "Cho các kim loại sau: Na, Ca, Fe, Al, Zn, Cu, Ag. Số kim loại được điều chế trong công nghiệp bằng phương pháp điện phân hợp chất nóng chảy của chúng là bao nhiêu?",
            "3",
            "Các kim loại hoạt động hoá học mạnh (từ Al trở về trước) được điều chế bằng phương pháp điện phân nóng chảy hợp chất:\n"
            "- Na (điện phân nóng chảy NaCl).\n"
            "- Ca (điện phân nóng chảy CaCl~2~).\n"
            "- Al (điện phân nóng chảy Al~2~O~3~ trong cryolite).\n"
            "Tổng cộng có 3 kim loại (Na, Ca, Al).\n"
            "Đáp số: 3."
        ),
        (
            4,
            "Cho các chất rắn sau đây: NaCl, CaCO~3~, CuSO~4~, NaOH, CaCl~2~, FeCl~2~, CuCl~2~, MgCl~2~, Al~2~O~3~. Số chất có thể dùng để điều chế kim loại bằng phương pháp điện phân nóng chảy là bao nhiêu?",
            "5",
            "Các chất điều chế được kim loại khi điện phân nóng chảy:\n"
            "1. 2NaCl → 2Na + Cl~2~ (thu được Na).\n"
            "2. 2NaOH → 2Na + 1/2 O~2~ + H~2~O (thu được Na).\n"
            "3. CaCl~2~ → Ca + Cl~2~ (thu được Ca).\n"
            "4. MgCl~2~ → Mg + Cl~2~ (thu được Mg).\n"
            "5. 2Al~2~O~3~ → 4Al + 3O~2~ (thu được Al).\n"
            "Tổng cộng có 5 chất.\n"
            "Đáp số: 5."
        ),
        (
            5,
            "Một viên pin lithium-ion của điện thoại có điện áp danh định 3,7 V và dung lượng 4000 mAh. Điện thoại tiêu thụ trung bình dòng điện 500 mA khi sử dụng liên tục. Hỏi điện thoại có thể hoạt động liên tục trong bao lâu trước khi cần sạc lại (theo đơn vị giờ)?",
            "8",
            "Thời gian hoạt động liên tục của điện thoại:\n"
            "t = Dung lượng pin (mAh) / Dòng tiêu thụ (mA) = 4000 / 500 = 8 giờ.\n"
            "Đáp số: 8."
        ),
        (
            6,
            "Một nhà máy muốn mạ một lớp đồng lên một tấm kim loại có diện tích 1 m^2^ để chống ăn mòn. Họ sử dụng dung dịch CuSO~4~ và dòng điện có cường độ 500 A trong thời gian 2 giờ. Biết khối lượng mol của Cu là 64 g/mol. Hãy tính khối lượng đồng (kg) bám vào tấm kim loại. (Làm tròn kết quả đến hàng phần trăm).",
            "1,19",
            "Điện lượng cung cấp:\n"
            "Q = I * t = 500 * (2 * 3600) = 3 600 000 C.\n"
            "Số mol electron trao đổi:\n"
            "n_e = Q / F = 3 600 000 / 96500 ≈ 37,3057 mol.\n"
            "Quá trình khử ở cathode: Cu^2+^ + 2e → Cu.\n"
            "Số mol đồng sinh ra: n_Cu = n_e / 2 = 37,3057 / 2 ≈ 18,6528 mol.\n"
            "Khối lượng đồng bám vào:\n"
            "m = 18,6528 * 64 ≈ 1193,78 g ≈ 1,19 kg.\n"
            "Đáp số: 1,19."
        )
    ]

    for q_num, content, ans, sol in p3_data:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        add_chem_run(p_q, f"Câu {q_num}. ", bold=True)
        add_chem_run(p_q, content)

        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(1)
        p_ans.paragraph_format.space_after = Pt(3)
        p_ans.paragraph_format.line_spacing = 1.15
        if is_teacher:
            add_chem_run(p_ans, "Đáp số: ", bold=True)
            add_chem_run(p_ans, ans, bold=True, color=COLOR_CORAL_RED)
            add_loigiai(doc, sol)
        else:
            add_chem_run(p_ans, "Trả lời: ...........................................................................................................................................................", italic=True)

    # Summary table for teacher version
    if is_teacher:
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(12)
        p_title.paragraph_format.space_after = Pt(4)
        r_t = p_title.add_run("BẢNG TỔNG HỢP ĐÁP ÁN ĐỀ SỐ 15")
        r_t.font.name = "Times New Roman"
        r_t.font.size = Pt(13)
        r_t.font.bold = True
        r_t.font.color.rgb = COLOR_DARK_GREEN

        p_p1 = doc.add_paragraph()
        p_p1.paragraph_format.space_before = Pt(4)
        p_p1.paragraph_format.space_after = Pt(2)
        r_p1 = p_p1.add_run("PHẦN I: TRẮC NGHIỆM NHIỀU PHƯƠNG ÁN LỰA CHỌN")
        r_p1.font.name = "Times New Roman"
        r_p1.font.size = Pt(11)
        r_p1.font.bold = True
        r_p1.font.color.rgb = COLOR_DARK_GREEN

        t1 = doc.add_table(rows=2, cols=18)
        t1.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j in range(18):
            c_h = t1.cell(0, j)
            c_h.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_h = c_h.paragraphs[0].add_run(str(j + 1))
            r_h.font.name = "Times New Roman"
            r_h.font.size = Pt(9.5)
            r_h.font.bold = True

            c_v = t1.cell(1, j)
            c_v.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_v = c_v.paragraphs[0].add_run(p1_data[j][3])
            r_v.font.name = "Times New Roman"
            r_v.font.size = Pt(9.5)
            r_v.font.bold = True
            r_v.font.color.rgb = COLOR_CORAL_RED

        p_p2 = doc.add_paragraph()
        p_p2.paragraph_format.space_before = Pt(6)
        p_p2.paragraph_format.space_after = Pt(2)
        r_p2 = p_p2.add_run("PHẦN II: TRẮC NGHIỆM ĐÚNG SAI")
        r_p2.font.name = "Times New Roman"
        r_p2.font.size = Pt(11)
        r_p2.font.bold = True
        r_p2.font.color.rgb = COLOR_DARK_GREEN

        t2 = doc.add_table(rows=5, cols=5)
        t2.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers_p2 = ["Câu", "Lệnh hỏi a", "Lệnh hỏi b", "Lệnh hỏi c", "Lệnh hỏi d"]
        for j, h in enumerate(headers_p2):
            c = t2.cell(0, j)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = c.paragraphs[0].add_run(h)
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.font.bold = True

        for i in range(4):
            c_num = t2.cell(i + 1, 0)
            c_num.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = c_num.paragraphs[0].add_run(str(i + 1))
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.font.bold = True
            for j in range(4):
                c_val = t2.cell(i + 1, j + 1)
                c_val.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                is_t = p2_data[i][2][j][2]
                val_str = "Đ" if is_t else "S"
                r_v = c_val.paragraphs[0].add_run(val_str)
                r_v.font.name = "Times New Roman"
                r_v.font.size = Pt(10)
                r_v.font.bold = True
                if is_t:
                    r_v.font.color.rgb = COLOR_CORAL_RED

        p_p3 = doc.add_paragraph()
        p_p3.paragraph_format.space_before = Pt(6)
        p_p3.paragraph_format.space_after = Pt(2)
        r_p3 = p_p3.add_run("PHẦN III: TRẢ LỜI NGẮN")
        r_p3.font.name = "Times New Roman"
        r_p3.font.size = Pt(11)
        r_p3.font.bold = True
        r_p3.font.color.rgb = COLOR_DARK_GREEN

        t3 = doc.add_table(rows=2, cols=6)
        t3.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j in range(6):
            c_h = t3.cell(0, j)
            c_h.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = c_h.paragraphs[0].add_run(f"Câu {j + 1}")
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.font.bold = True

            c_v = t3.cell(1, j)
            c_v.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_v = c_v.paragraphs[0].add_run(p3_data[j][2])
            r_v.font.name = "Times New Roman"
            r_v.font.size = Pt(10)
            r_v.font.bold = True
            r_v.font.color.rgb = COLOR_CORAL_RED

    file_name = "De_Kiem_Tra_Hoa_12_De15_GiaoVien.docx" if is_teacher else "De_Kiem_Tra_Hoa_12_De15.docx"
    out_path = os.path.join(OUT_DIR, file_name)
    doc.save(out_path)
    print(f"Saved: {out_path}")

if __name__ == "__main__":
    generate_exam_docx(is_teacher=False)
    generate_exam_docx(is_teacher=True)
