# -*- coding: utf-8 -*-
"""
Script tạo trọn bộ file Word (.docx) ĐỀ CƯƠNG ÔN TẬP GIỮA KÌ I - KHTN 9 (ĐẦY ĐỦ 3 PHÂN MÔN: HÓA HỌC, VẬT LÍ, SINH HỌC)
Trường THCS Kim Long - Thầy TRẦN VĂN THẠNH - 0777.470.803

Tạo 3 phiên bản Trọn Bộ:
1. Bản Học sinh: San_Pham/De_Cuong_KHTN_9_Giua_HK1_Kim_Long_TronBo_HocSinh.docx
2. Bản Giáo viên: San_Pham/De_Cuong_KHTN_9_Giua_HK1_Kim_Long_TronBo_GiaoVien.docx
3. Bản Dạy trên màn hình: San_Pham/De_Cuong_KHTN_9_Giua_HK1_Kim_Long_TronBo_DayTrenManHinh.docx
"""

import os
import sys
import re
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.path.append(os.path.dirname(__file__))
sys.stdout.reconfigure(encoding='utf-8')

from khtn9_kimlong_data import (
    INFO as INFO_KL, BAI_1_QUESTIONS,
    HOA_PART1_QUESTIONS, HOA_PART2_QUESTIONS,
    HOA_PART3_QUESTIONS, HOA_ESSAY_QUESTIONS
)

from khtn9_vatli_sinhhoc_data import (
    PHYSICS_CHOICE_QUESTIONS, PHYSICS_TF_QUESTIONS,
    PHYSICS_SHORT_QUESTIONS, PHYSICS_ESSAY_QUESTIONS,
    BIOLOGY_THEORY, BIOLOGY_CHOICE_QUESTIONS, BIOLOGY_TF_QUESTIONS,
    BIOLOGY_SHORT_QUESTIONS, BIOLOGY_ESSAY_QUESTIONS
)

ROOT_DIR = r"c:\Antigravity_Thanh"
OUT_DIR = os.path.join(ROOT_DIR, "San_Pham")
ONEDRIVE_DIRS = [
    r"C:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn9\KHTN9-DEKIEMTRADINHKI",
    r"C:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn9"
]

COLOR_HF_RED = RGBColor(211, 47, 47)
COLOR_HF_GREEN_HEX = "8BC390"
COLOR_PRIMARY = RGBColor(46, 125, 50)
COLOR_CORRECT = RGBColor(198, 40, 40)
COLOR_SOLUTION = RGBColor(21, 101, 192)
COLOR_SCREEN_TITLE = RGBColor(106, 27, 154)
COLOR_GRAY = RGBColor(97, 97, 97)
COLOR_BORDER_GREY_HEX = "CCCCCC"

def add_p_border_bottom(p, color="8BC390", sz="6"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_p_border_top(p, color="8BC390", sz="6"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex="F1F8E9"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="CCCCCC", sz="4"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def set_run_style(run, font_name="Times New Roman", font_size=10.5, bold=False, italic=False, subscript=False, superscript=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def render_rich_text(p, text, base_font="Times New Roman", base_size=10.5, base_bold=False, base_italic=False, base_color=None):
    tokens = re.split(r'(~[^~]+~|\^[^\^]+\^|\*[^\*]+\*|_[^_]+_)', text)
    for token in tokens:
        if not token:
            continue
        is_bold = base_bold
        is_italic = base_italic
        is_sub = False
        is_sup = False
        content = token
        if token.startswith('*') and token.endswith('*'):
            is_bold = True
            content = token[1:-1]
        elif token.startswith('_') and token.endswith('_'):
            is_italic = True
            content = token[1:-1]
        elif token.startswith('~') and token.endswith('~'):
            is_sub = True
            content = token[1:-1]
        elif token.startswith('^') and token.endswith('^'):
            is_sup = True
            content = token[1:-1]

        run = p.add_run(content)
        run_color = COLOR_CORRECT if (is_bold and not base_bold and not base_color) else base_color
        set_run_style(run, font_name=base_font, font_size=base_size, bold=is_bold, italic=is_italic,
                      subscript=is_sub, superscript=is_sup, color=run_color)

def setup_header_footer(doc):
    normal_style = doc.styles['Normal']
    normal_style.paragraph_format.tab_stops.clear_all()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    section.start_type = WD_SECTION_START.NEW_PAGE
    section.different_first_page_header_footer = False
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    header = section.header
    for p in header.paragraphs:
        p.text = ""
    tbl_h = header.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left, c_right = tbl_h.cell(0, 0), tbl_h.cell(0, 1)
    c_left.width, c_right.width = Cm(8.0), Cm(10.0)

    p_hl = c_left.paragraphs[0]
    p_hl.paragraph_format.space_before, p_hl.paragraph_format.space_after = Pt(0), Pt(2)
    p_hl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_hl = p_hl.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)

    p_hr = c_right.paragraphs[0]
    p_hr.paragraph_format.space_before, p_hr.paragraph_format.space_after = Pt(0), Pt(2)
    p_hr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p_hr.add_run("Thầy ")
    set_run_style(r1, font_size=10, color=COLOR_HF_RED)
    r2 = p_hr.add_run("TRẦN VĂN THẠNH")
    set_run_style(r2, font_size=10, bold=True, color=COLOR_HF_RED)
    r3 = p_hr.add_run(" - 0777.470.803")
    set_run_style(r3, font_size=10, color=COLOR_HF_RED)

    add_p_border_bottom(p_hl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_bottom(p_hr, color=COLOR_HF_GREEN_HEX, sz="6")

    footer = section.footer
    for p in footer.paragraphs:
        p.text = ""
    tbl_f = footer.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_fl, c_fr = tbl_f.cell(0, 0), tbl_f.cell(0, 1)
    c_fl.width, c_fr.width = Cm(14.5), Cm(3.5)

    p_fl = c_fl.paragraphs[0]
    p_fl.paragraph_format.space_before, p_fl.paragraph_format.space_after = Pt(2), Pt(0)
    p_fl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_fl = p_fl.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)

    p_fr = c_fr.paragraphs[0]
    p_fr.paragraph_format.space_before, p_fr.paragraph_format.space_after = Pt(2), Pt(0)
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_field(p_fr, font_size=10, bold=True, color=COLOR_HF_RED)

    add_p_border_top(p_fl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_top(p_fr, color=COLOR_HF_GREEN_HEX, sz="6")

def build_title_block(doc, is_teacher=False, is_screen=False, base_size=10.5):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_l, cell_r = tbl.cell(0, 0), tbl.cell(0, 1)
    cell_l.width, cell_r.width = Cm(8.0), Cm(10.0)

    p_tl1 = cell_l.paragraphs[0]
    p_tl1.paragraph_format.space_before, p_tl1.paragraph_format.space_after = Pt(0), Pt(1)
    p_tl1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tl1, "*TRƯỜNG THCS KIM LONG*", base_size=10.5, base_bold=True)

    p_tl2 = cell_l.add_paragraph()
    p_tl2.paragraph_format.space_before, p_tl2.paragraph_format.space_after = Pt(0), Pt(1)
    p_tl2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tl2, "*TỔ KHOA HỌC TỰ NHIÊN*", base_size=10.5, base_bold=True)

    p_tl3 = cell_l.add_paragraph()
    p_tl3.paragraph_format.space_before, p_tl3.paragraph_format.space_after = Pt(0), Pt(0)
    p_tl3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tl3, "_(Tài liệu lưu hành nội bộ)_", base_size=9.5, base_italic=True)

    p_tr1 = cell_r.paragraphs[0]
    p_tr1.paragraph_format.space_before, p_tr1.paragraph_format.space_after = Pt(0), Pt(1)
    p_tr1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tr1, "*ĐỀ CƯƠNG ÔN TẬP GIỮA KÌ I - KHTN 9*", base_size=11, base_bold=True)

    p_tr2 = cell_r.add_paragraph()
    p_tr2.paragraph_format.space_before, p_tr2.paragraph_format.space_after = Pt(0), Pt(1)
    p_tr2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tr2, "*TRỌN BỘ 3 PHÂN MÔN: HÓA HỌC - VẬT LÍ - SINH HỌC*", base_size=11, base_bold=True, base_color=COLOR_PRIMARY)

    p_tr3 = cell_r.add_paragraph()
    p_tr3.paragraph_format.space_before, p_tr3.paragraph_format.space_after = Pt(0), Pt(0)
    p_tr3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tr3, "_Chương trình GDPT 2018 - Năm học 2026 - 2027_", base_size=9.5, base_italic=True)

    if is_teacher:
        p_gv = doc.add_paragraph()
        p_gv.paragraph_format.space_before, p_gv.paragraph_format.space_after = Pt(4), Pt(2)
        p_gv.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rgv = p_gv.add_run("(BẢN GIÁO VIÊN -- CÓ ĐÁP ÁN & LỜI GIẢI CHI TIẾT TRỌN BỘ)")
        set_run_style(rgv, font_size=11.5, bold=True, color=COLOR_CORRECT)
    elif is_screen:
        p_sc = doc.add_paragraph()
        p_sc.paragraph_format.space_before, p_sc.paragraph_format.space_after = Pt(4), Pt(2)
        p_sc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rsc = p_sc.add_run("(BẢN TRÌNH CHIẾU / GIẢNG DẠY MÀN HÌNH TRỌN BỘ -- CÓ CHỪA CHỖ TRỐNG)")
        set_run_style(rsc, font_size=11.5, bold=True, color=COLOR_SCREEN_TITLE)
    else:
        p_hs = doc.add_paragraph()
        p_hs.paragraph_format.space_before, p_hs.paragraph_format.space_after = Pt(4), Pt(2)
        p_hs.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rhs = p_hs.add_run("(BẢN HỌC SINH TỰ LUYỆN TRỌN BỘ -- CÓ DÒNG LÀM BÀI TỰ LUẬN)")
        set_run_style(rhs, font_size=11, bold=True, color=COLOR_PRIMARY)

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before, p_info.paragraph_format.space_after = Pt(4), Pt(3)
    p_info.paragraph_format.tab_stops.clear_all()
    p_info.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
    render_rich_text(p_info, "*Họ và tên học sinh:* .................................................... *Lớp:* ......... *Trường:* .......................................", base_size=base_size)
    p_info.add_run("\t")
    r_md = p_info.add_run("Mã đề 901")
    set_run_style(r_md, font_size=10.5, bold=True)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before, p_div.paragraph_format.space_after = Pt(0), Pt(4)
    add_p_border_bottom(p_div, color=COLOR_BORDER_GREY_HEX, sz="4")

def add_major_section_header(doc, title, subtitle=None, base_size=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(10), Pt(3)
    p.paragraph_format.keep_with_next = True
    render_rich_text(p, f"*{title}*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)
    if subtitle:
        render_rich_text(p, f" - _{subtitle}_", base_size=base_size * 0.92, base_italic=True)
    add_p_border_bottom(p, color="2E7D32", sz="8")

def add_section_title(doc, title, subtitle=None, base_size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(7), Pt(2)
    p.paragraph_format.keep_with_next = True
    render_rich_text(p, f"*{title}*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)
    if subtitle:
        render_rich_text(p, f" {subtitle}", base_size=base_size * 0.95, base_italic=True)

def add_solution_block(doc, solution_text, base_size=10):
    p_sol = doc.add_paragraph()
    p_sol.paragraph_format.left_indent = Cm(0.6)
    p_sol.paragraph_format.space_before, p_sol.paragraph_format.space_after = Pt(2), Pt(2)
    p_sol.paragraph_format.line_spacing = 1.15

    r_lbl = p_sol.add_run("► Lời giải chi tiết:\n")
    set_run_style(r_lbl, font_size=base_size, bold=True, color=COLOR_SOLUTION)

    lines = solution_text.split('\n')
    for idx, line in enumerate(lines):
        line_str = line.strip()
        if not line_str:
            continue
        if idx > 0:
            p_sol = doc.add_paragraph()
            p_sol.paragraph_format.left_indent = Cm(0.6)
            p_sol.paragraph_format.space_before, p_sol.paragraph_format.space_after = Pt(1), Pt(1)
            p_sol.paragraph_format.line_spacing = 1.15
        render_rich_text(p_sol, line_str, base_size=base_size)

def add_screen_teaching_space(doc, num_lines=3):
    p_box = doc.add_paragraph()
    p_box.paragraph_format.left_indent = Cm(0.5)
    p_box.paragraph_format.space_before, p_box.paragraph_format.space_after = Pt(3), Pt(5)
    pPr = p_box._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="dashed" w:sz="4" w:space="4" w:color="B0BEC5"/><w:bottom w:val="dashed" w:sz="4" w:space="4" w:color="B0BEC5"/><w:left w:val="single" w:sz="12" w:space="6" w:color="7E57C2"/></w:pBdr>'
    pPr.append(parse_xml(xml))
    r_hint = p_box.add_run("✎ Ghi chú & Lời giải trực tiếp trên màn hình:\n")
    set_run_style(r_hint, font_size=9.5, bold=True, color=COLOR_SCREEN_TITLE)
    for _ in range(num_lines):
        p_box.add_run("\n")

def add_student_writing_lines(doc, num_lines=4):
    for _ in range(num_lines):
        p_line = doc.add_paragraph()
        p_line.paragraph_format.left_indent = Cm(0.5)
        p_line.paragraph_format.space_before, p_line.paragraph_format.space_after = Pt(1), Pt(2)
        r = p_line.add_run("...................................................................................................................................................................")
        set_run_style(r, font_size=9, color=RGBColor(189, 189, 189))

def render_mc_questions(doc, questions, is_teacher=False, is_screen=False, base_size=10.5):
    for q in questions:
        qid = q["id"]
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before, p_q.paragraph_format.space_after = Pt(3), Pt(1)
        p_q.paragraph_format.line_spacing = 1.15
        p_q.paragraph_format.keep_with_next = True
        render_rich_text(p_q, f"*Câu {qid}.* {q['content']}", base_size=base_size)

        opts = q["options"]
        lens = [len(opt) for opt in opts]
        max_len = max(lens)

        if max_len <= 18:
            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.space_before, p_opt.paragraph_format.space_after = Pt(1), Pt(2)
            p_opt.paragraph_format.tab_stops.clear_all()
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(4.8), WD_TAB_ALIGNMENT.LEFT)
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(13.6), WD_TAB_ALIGNMENT.LEFT)
            
            p_opt.add_run("\t")
            for idx, opt_text in enumerate(opts):
                opt_lbl = opt_text[:1]
                is_c = (is_teacher and opt_lbl == q["correct"])
                if is_c:
                    render_rich_text(p_opt, f"*{opt_text}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
                else:
                    render_rich_text(p_opt, opt_text, base_size=base_size)
                if idx < 3:
                    p_opt.add_run("\t")
        elif max_len <= 38:
            for row_idx in range(2):
                p_opt = doc.add_paragraph()
                p_opt.paragraph_format.space_before, p_opt.paragraph_format.space_after = Pt(1), Pt(1)
                p_opt.paragraph_format.tab_stops.clear_all()
                p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
                p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)
                
                p_opt.add_run("\t")
                opt_text1 = opts[row_idx * 2]
                opt_lbl1 = opt_text1[:1]
                is_c1 = (is_teacher and opt_lbl1 == q["correct"])
                if is_c1:
                    render_rich_text(p_opt, f"*{opt_text1}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
                else:
                    render_rich_text(p_opt, opt_text1, base_size=base_size)
                
                p_opt.add_run("\t")
                opt_text2 = opts[row_idx * 2 + 1]
                opt_lbl2 = opt_text2[:1]
                is_c2 = (is_teacher and opt_lbl2 == q["correct"])
                if is_c2:
                    render_rich_text(p_opt, f"*{opt_text2}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
                else:
                    render_rich_text(p_opt, opt_text2, base_size=base_size)
        else:
            for opt_text in opts:
                p_opt = doc.add_paragraph()
                p_opt.paragraph_format.left_indent = Cm(0.5)
                p_opt.paragraph_format.space_before, p_opt.paragraph_format.space_after = Pt(1), Pt(1)
                opt_lbl = opt_text[:1]
                is_c = (is_teacher and opt_lbl == q["correct"])
                if is_c:
                    render_rich_text(p_opt, f"*{opt_text}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
                else:
                    render_rich_text(p_opt, opt_text, base_size=base_size)

        if is_teacher:
            add_solution_block(doc, f"Đáp án đúng: *{q['correct']}*\n" + q["solution"], base_size=base_size - 0.5)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=3)

def render_tf_questions(doc, questions, is_teacher=False, is_screen=False, base_size=10.5):
    for q in questions:
        qid = q["id"]
        q_title = q.get("content") or q.get("title", "")
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before, p_q.paragraph_format.space_after = Pt(4), Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        p_q.paragraph_format.keep_with_next = True
        render_rich_text(p_q, f"*Câu {qid}.* {q_title}", base_size=base_size)

        for st in q["statements"]:
            sub_id = st.get("id") or st.get("label", "")
            st_text = st["text"]
            correct_val = "Đúng" if st["correct"] else "Sai"

            p_st = doc.add_paragraph()
            p_st.paragraph_format.left_indent = Cm(0.5)
            p_st.paragraph_format.space_before, p_st.paragraph_format.space_after = Pt(1), Pt(1)
            p_st.paragraph_format.line_spacing = 1.15

            if is_teacher:
                r_tag = p_st.add_run(f"{sub_id}) ")
                set_run_style(r_tag, font_size=base_size, bold=True)
                ans_color = COLOR_CORRECT if st["correct"] else RGBColor(0, 105, 92)
                r_ans = p_st.add_run(f"[{correct_val.upper()}] ")
                set_run_style(r_ans, font_size=base_size, bold=True, color=ans_color)
                render_rich_text(p_st, st_text, base_size=base_size)
            else:
                r_tag = p_st.add_run(f"{sub_id}) ")
                set_run_style(r_tag, font_size=base_size, bold=True)
                render_rich_text(p_st, st_text, base_size=base_size)

        if is_teacher:
            exp_text = q.get("explanation")
            if not exp_text and "statements" in q:
                exp_text = "\n".join([f"- Ý {st.get('label') or st.get('id')}: {st.get('explain', '')}" for st in q["statements"] if st.get('explain')])
            if exp_text:
                add_solution_block(doc, exp_text, base_size=base_size - 0.5)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=3)

def render_short_questions(doc, questions, is_teacher=False, is_screen=False, base_size=10.5):
    for q in questions:
        qid = q["id"]
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before, p_q.paragraph_format.space_after = Pt(3), Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        p_q.paragraph_format.keep_with_next = True
        render_rich_text(p_q, f"*Câu {qid}.* {q['content']}", base_size=base_size)

        if is_teacher:
            add_solution_block(doc, f"Đáp án: *{q['correct']}*\n" + q["solution" if "solution" in q else "explanation"], base_size=base_size - 0.5)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=3)
        else:
            p_ans_box = doc.add_paragraph()
            p_ans_box.paragraph_format.left_indent = Cm(0.5)
            p_ans_box.paragraph_format.space_before, p_ans_box.paragraph_format.space_after = Pt(1), Pt(2)
            r_box = p_ans_box.add_run("Trả lời: .............................................................................................................................................")
            set_run_style(r_box, font_size=10, italic=True, color=COLOR_GRAY)

def render_essay_questions(doc, questions, is_teacher=False, is_screen=False, base_size=10.5):
    for q in questions:
        qid = q["id"]
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before, p_q.paragraph_format.space_after = Pt(4), Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        p_q.paragraph_format.keep_with_next = True
        render_rich_text(p_q, f"*Câu {qid}.* {q['content']}", base_size=base_size)

        if q.get("has_image") and os.path.exists(q.get("image_path", "")):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(3)
            p_img.paragraph_format.space_after = Pt(1)
            run_img = p_img.add_run()
            run_img.add_picture(q["image_path"], width=Cm(6.5))
            if q.get("image_caption"):
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(0)
                p_cap.paragraph_format.space_after = Pt(3)
                r_cap = p_cap.add_run(q["image_caption"])
                set_run_style(r_cap, font_size=9, italic=True, color=COLOR_GRAY)

        if is_teacher:
            add_solution_block(doc, q["solution"], base_size=base_size - 0.5)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=5)
        else:
            add_student_writing_lines(doc, num_lines=5)

def render_theory_section(doc, theory_list, base_size=10.5):
    for sec in theory_list:
        p_th = doc.add_paragraph()
        p_th.paragraph_format.space_before, p_th.paragraph_format.space_after = Pt(5), Pt(2)
        p_th.paragraph_format.keep_with_next = True
        render_rich_text(p_th, f"*{sec['title']}*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)

        for item in sec["items"]:
            p_item = doc.add_paragraph()
            p_item.paragraph_format.left_indent = Cm(0.4)
            p_item.paragraph_format.space_before, p_item.paragraph_format.space_after = Pt(1), Pt(2)
            p_item.paragraph_format.line_spacing = 1.15
            render_rich_text(p_item, f"- {item}", base_size=base_size)

def build_tronbo_docx(output_path, mode="student"):
    is_teacher = (mode == "teacher")
    is_screen = (mode == "screen")
    base_size = 11.5 if is_screen else 10.5

    doc = docx.Document()
    setup_header_footer(doc)
    build_title_block(doc, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    # -------------------------------------------------------------------------
    # MỤC 1: BÀI 1: DỤNG CỤ, HÓA CHẤT & THUYẾT TRÌNH KHOA HỌC
    # -------------------------------------------------------------------------
    add_major_section_header(
        doc,
        "BÀI 1: NHẬN BIẾT DỤNG CỤ, HÓA CHẤT - THUYẾT TRÌNH VẤN ĐỀ KHOA HỌC",
        "Kĩ năng thực hành & Phương pháp nghiên cứu KHTN 9",
        base_size=base_size + 0.8
    )
    render_mc_questions(doc, BAI_1_QUESTIONS, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    # -------------------------------------------------------------------------
    # MỤC 2: PHÂN MÔN KHTN 1 (HÓA HỌC)
    # -------------------------------------------------------------------------
    add_major_section_header(
        doc,
        "PHÂN MÔN KHTN 1: HÓA HỌC (HỢP CHẤT HỮU CƠ, HYDROCARBON, NHIÊN LIỆU & ETHYLIC ALCOHOL)",
        "Hóa học hữu cơ lớp 9",
        base_size=base_size + 0.8
    )

    add_section_title(doc, "Phần 1: Trắc nghiệm nhiều phương án lựa chọn (Hóa học - 18 câu)", base_size=base_size + 0.3)
    render_mc_questions(doc, HOA_PART1_QUESTIONS, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    add_section_title(doc, "Phần 2: Trắc nghiệm đúng sai (Hóa học - 4 câu)", base_size=base_size + 0.3)
    render_tf_questions(doc, HOA_PART2_QUESTIONS, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    add_section_title(doc, "Phần 3: Trắc nghiệm trả lời ngắn (Hóa học - 4 câu)", base_size=base_size + 0.3)
    render_short_questions(doc, HOA_PART3_QUESTIONS, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    add_section_title(doc, "Phần Tự Luận: Bài tập tự luận Hóa học (4 câu)", base_size=base_size + 0.3)
    render_essay_questions(doc, HOA_ESSAY_QUESTIONS, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    # -------------------------------------------------------------------------
    # MỤC 3: PHÂN MÔN KHTN 2 (VẬT LÍ)
    # -------------------------------------------------------------------------
    add_major_section_header(
        doc,
        "PHÂN MÔN KHTN 2: VẬT LÍ (CƠ NĂNG - CÔNG VÀ CÔNG SUẤT - NĂNG LƯỢNG)",
        "Động năng, thế năng, cơ năng, năng lượng tái tạo",
        base_size=base_size + 0.8
    )

    add_section_title(doc, "Phần 1: Trắc nghiệm nhiều phương án lựa chọn (Vật lí - 13 câu)", base_size=base_size + 0.3)
    render_mc_questions(doc, PHYSICS_CHOICE_QUESTIONS, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    add_section_title(doc, "Phần 2: Trắc nghiệm đúng sai (Vật lí - 4 câu)", base_size=base_size + 0.3)
    render_tf_questions(doc, PHYSICS_TF_QUESTIONS, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    add_section_title(doc, "Phần 3: Trắc nghiệm trả lời ngắn (Vật lí - 5 câu)", base_size=base_size + 0.3)
    render_short_questions(doc, PHYSICS_SHORT_QUESTIONS, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    add_section_title(doc, "Phần Tự Luận: Bài tập tự luận Vật lí (12 câu)", base_size=base_size + 0.3)
    render_essay_questions(doc, PHYSICS_ESSAY_QUESTIONS, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    # -------------------------------------------------------------------------
    # MỤC 4: PHÂN MÔN KHTN 3 (SINH HỌC)
    # -------------------------------------------------------------------------
    add_major_section_header(
        doc,
        "PHÂN MÔN KHTN 3: SINH HỌC (DI TRUYỀN HỌC & GENE)",
        "Di truyền và biến dị - Menđen - DNA và RNA",
        base_size=base_size + 0.8
    )

    add_section_title(doc, "I. Tóm tắt nội dung ôn tập lý thuyết trọng tâm Sinh học", base_size=base_size + 0.3)
    render_theory_section(doc, BIOLOGY_THEORY, base_size=base_size)

    add_section_title(doc, "II. Trắc nghiệm Sinh học - Phần 1: Câu hỏi nhiều lựa chọn (19 câu)", base_size=base_size + 0.3)
    render_mc_questions(doc, BIOLOGY_CHOICE_QUESTIONS, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    add_section_title(doc, "II. Trắc nghiệm Sinh học - Phần 2: Trắc nghiệm đúng sai (5 câu)", base_size=base_size + 0.3)
    render_tf_questions(doc, BIOLOGY_TF_QUESTIONS, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    add_section_title(doc, "II. Trắc nghiệm Sinh học - Phần 3: Câu hỏi trả lời ngắn (15 câu)", base_size=base_size + 0.3)
    render_short_questions(doc, BIOLOGY_SHORT_QUESTIONS, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    add_section_title(doc, "III. Tự luận Sinh học: Sơ đồ lai & Cấu trúc phân tử DNA (4 câu)", base_size=base_size + 0.3)
    render_essay_questions(doc, BIOLOGY_ESSAY_QUESTIONS, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    # Kết thúc
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before, p_end.paragraph_format.space_after = Pt(12), Pt(4)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("--- HẾT TRỌN BỘ ĐỀ CƯƠNG KHTN 9 ---")
    set_run_style(r_end, font_size=11, bold=True)

    doc.save(output_path)
    print(f"Created Tron Bo: {output_path}")

def main():
    files = [
        ("De_Cuong_KHTN_9_Giua_HK1_Kim_Long_TronBo_HocSinh.docx", "student"),
        ("De_Cuong_KHTN_9_Giua_HK1_Kim_Long_TronBo_GiaoVien.docx", "teacher"),
        ("De_Cuong_KHTN_9_Giua_HK1_Kim_Long_TronBo_DayTrenManHinh.docx", "screen"),
    ]
    for filename, mode in files:
        out_file = os.path.join(OUT_DIR, filename)
        build_tronbo_docx(out_file, mode=mode)
        for onedrive_folder in ONEDRIVE_DIRS:
            if os.path.exists(onedrive_folder):
                try:
                    dest = os.path.join(onedrive_folder, filename)
                    shutil.copy2(out_file, dest)
                    print(f"Copied Tron Bo to OneDrive successfully: {dest}")
                except Exception as e:
                    print(f"Copy error: {e}")

if __name__ == "__main__":
    main()
