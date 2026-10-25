# -*- coding: utf-8 -*-
"""
CÔNG CỤ RÚT VÀ SOẠN ĐỀ KIỂM TRA HÓA HỌC 11 THEO MA TRẬN TỰ ĐỘNG
Hỗ trợ rút câu hỏi từ Ngân hàng 1.996 câu (Ngan_Hang_Hoa_11_Tron_Bo.json)
Xuất đề chuẩn quy cách "FORM ĐỀ KIỂM TRA" (Bản Học sinh & Bản Giáo viên).
"""

import os
import json
import random
import re
import sys
import docx
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

sys.path.insert(0, r'c:\Antigravity_Thanh\scripts')
from build_ngan_hang_hoa10_c1 import (
    setup_page_and_hf, add_chem_runs_hoa10, parse_abcd_from_opt_lines,
    add_options_paragraph_no_table, add_data_table_to_doc, add_solution_block_hoa10,
    LEVEL_NAMES, COLOR_HF_RED, COLOR_PRIMARY, COLOR_SOLUTION, COLOR_CORRECT, COLOR_TAG
)

BANK_JSON = r"c:\Antigravity_Thanh\San_Pham\Ngan_Hang_Hoa_11\Ngan_Hang_Hoa_11_Tron_Bo.json"
OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"


def load_bank_11():
    with open(BANK_JSON, "r", encoding="utf-8") as f:
        return json.load(f)


def print_bank_summary():
    bank = load_bank_11()
    print("=== TỔNG QUAN NGÂN HÀNG CÂU HỎI HÓA HỌC 11 ===")
    total = 0
    for ch_code, ch_data in bank.items():
        qs = ch_data["questions"]
        total += len(qs)
        p1 = sum(1 for q in qs if q["part"] == "P1")
        p2 = sum(1 for q in qs if q["part"] == "P2")
        p3 = sum(1 for q in qs if q["part"] == "P3")
        print(f"* [{ch_code}] {ch_data['chapter_name']} ({len(qs)} câu) -> P1: {p1}, P2: {p2}, P3: {p3}")
    print(f"=> TỔNG CỘNG: {total} CÂU HỎI ĐÃ SẴN SÀNG ĐỂ RÚT ĐỀ!\n")


def select_questions_from_bank(matrix_spec, seed=None):
    """
    matrix_spec: list of dicts:
    [
        {"chapter": "C1", "part": "P1", "level": "NB", "count": 5},
        ...
    ]
    """
    bank = load_bank_11()
    rng = random.Random(seed)
    selected = {"P1": [], "P2": [], "P3": []}
    used_ids = set()

    for item in matrix_spec:
        ch = item["chapter"]
        pt = item["part"]
        lv = item.get("level")
        cnt = item["count"]
        tp = item.get("topic")

        pool = [
            q for q in bank[ch]["questions"]
            if q["part"] == pt
            and q["id"] not in used_ids
            and (lv is None or q["level"] == lv)
            and (tp is None or q["topic"] == tp)
        ]

        if len(pool) < cnt:
            # Fallback level if needed
            pool = [
                q for q in bank[ch]["questions"]
                if q["part"] == pt and q["id"] not in used_ids
            ]

        chosen = rng.sample(pool, min(cnt, len(pool)))
        for q in chosen:
            used_ids.add(q["id"])
            selected[pt].append(q)

    return selected


def export_exam_word(exam_title, made, selected_qs, out_prefix, time_minutes=50):
    for is_teacher in [False, True]:
        doc = docx.Document()
        setup_page_and_hf(doc)

        # Header Title Block
        p_t1 = doc.add_paragraph()
        p_t1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_t1.paragraph_format.space_before = Pt(0)
        p_t1.paragraph_format.space_after = Pt(2)
        r_t1 = p_t1.add_run(exam_title.upper())
        r_t1.bold = True
        r_t1.font.name = "Times New Roman"
        r_t1.font.size = Pt(13)
        r_t1.font.color.rgb = COLOR_PRIMARY

        p_t2 = doc.add_paragraph()
        p_t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_t2.paragraph_format.space_before = Pt(0)
        p_t2.paragraph_format.space_after = Pt(2)
        r_t2 = p_t2.add_run(f"MÔN: HÓA HỌC 11 – THỜI GIAN LÀM BÀI: {time_minutes} PHÚT – MÃ ĐỀ: {made}")
        r_t2.bold = True
        r_t2.font.name = "Times New Roman"
        r_t2.font.size = Pt(10.5)

        p_t3 = doc.add_paragraph()
        p_t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_t3.paragraph_format.space_before = Pt(0)
        p_t3.paragraph_format.space_after = Pt(4)
        note_txt = (
            f"(BẢN GIÁO VIÊN — CÓ ĐÁP ÁN VÀ LỜI GIẢI CHI TIẾT – MÃ ĐỀ {made})"
            if is_teacher else
            "Họ và tên thí sinh: .............................................................. Lớp: ............. SBD: ............."
        )
        r_t3 = p_t3.add_run(note_txt)
        r_t3.bold = True
        r_t3.font.name = "Times New Roman"
        r_t3.font.size = Pt(10.5)
        if is_teacher:
            r_t3.font.color.rgb = COLOR_HF_RED
            r_t3.italic = True

        p_ntk = doc.add_paragraph()
        p_ntk.paragraph_format.space_before = Pt(0)
        p_ntk.paragraph_format.space_after = Pt(6)
        r_ntk = p_ntk.add_run("(Cho biết nguyên tử khối: H = 1; C = 12; N = 14; O = 16; Na = 23; Mg = 24; Al = 27; S = 32; Cl = 35,5; K = 39; Ca = 40; Fe = 56; Cu = 64; Zn = 65; Br = 80; Ag = 108; Ba = 137)")
        r_ntk.italic = True
        r_ntk.font.name = "Times New Roman"
        r_ntk.font.size = Pt(9.5)

        q_counter = 1

        # Phần I
        if selected_qs["P1"]:
            p_p1 = doc.add_paragraph()
            p_p1.paragraph_format.space_before = Pt(8)
            p_p1.paragraph_format.space_after = Pt(2)
            r_p1 = p_p1.add_run("PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn. Thí sinh trả lời từ câu 1 đến câu " + str(len(selected_qs["P1"])) + ".")
            r_p1.bold = True
            r_p1.font.name = "Times New Roman"
            r_p1.font.size = Pt(11)
            r_p1.font.color.rgb = COLOR_HF_RED

            p_p1_sub = doc.add_paragraph()
            p_p1_sub.paragraph_format.space_before = Pt(0)
            p_p1_sub.paragraph_format.space_after = Pt(4)
            r_p1_sub = p_p1_sub.add_run("Mỗi câu hỏi thí sinh chỉ chọn một phương án.")
            r_p1_sub.italic = True
            r_p1_sub.font.name = "Times New Roman"
            r_p1_sub.font.size = Pt(10)

            for q in selected_qs["P1"]:
                p_q = doc.add_paragraph()
                p_q.paragraph_format.space_before = Pt(3)
                p_q.paragraph_format.space_after = Pt(2)
                r_num = p_q.add_run(f"Câu {q_counter}. ")
                r_num.bold = True
                r_num.font.name = "Times New Roman"
                r_num.font.size = Pt(11)
                r_num.font.color.rgb = COLOR_PRIMARY

                if is_teacher and q.get("answer"):
                    r_tg = p_q.add_run(f"[{q['id']} | ĐA: {q['answer']}] ")
                    r_tg.bold = True
                    r_tg.font.name = "Times New Roman"
                    r_tg.font.size = Pt(10)
                    r_tg.font.color.rgb = COLOR_TAG

                add_chem_runs_hoa10(p_q, "\n".join(q["stem_lines"]).strip(), size=11)

                for imp in q.get("image_paths", []):
                    if os.path.exists(imp):
                        p_im = doc.add_paragraph()
                        p_im.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        r_im = p_im.add_run()
                        r_im.add_picture(imp, width=Cm(11.0))

                for tbl_d in q.get("tables", []):
                    add_data_table_to_doc(doc, tbl_d)

                opts = parse_abcd_from_opt_lines(q["opt_lines"])
                add_options_paragraph_no_table(doc, opts, q["answer"], is_teacher)

                if is_teacher:
                    add_solution_block_hoa10(doc, q)
                q_counter += 1

        # Phần II
        if selected_qs["P2"]:
            p2_start = q_counter
            p2_end = q_counter + len(selected_qs["P2"]) - 1
            p_p2 = doc.add_paragraph()
            p_p2.paragraph_format.space_before = Pt(8)
            p_p2.paragraph_format.space_after = Pt(2)
            r_p2 = p_p2.add_run(f"PHẦN II. Câu trắc nghiệm đúng sai. Thí sinh trả lời từ câu {p2_start} đến câu {p2_end}.")
            r_p2.bold = True
            r_p2.font.name = "Times New Roman"
            r_p2.font.size = Pt(11)
            r_p2.font.color.rgb = COLOR_HF_RED

            p_p2_sub = doc.add_paragraph()
            p_p2_sub.paragraph_format.space_before = Pt(0)
            p_p2_sub.paragraph_format.space_after = Pt(4)
            r_p2_sub = p_p2_sub.add_run("Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.")
            r_p2_sub.italic = True
            r_p2_sub.font.name = "Times New Roman"
            r_p2_sub.font.size = Pt(10)

            for q in selected_qs["P2"]:
                p_q = doc.add_paragraph()
                p_q.paragraph_format.space_before = Pt(3)
                p_q.paragraph_format.space_after = Pt(2)
                r_num = p_q.add_run(f"Câu {q_counter}. ")
                r_num.bold = True
                r_num.font.name = "Times New Roman"
                r_num.font.size = Pt(11)
                r_num.font.color.rgb = COLOR_PRIMARY

                if is_teacher and q.get("answer"):
                    r_tg = p_q.add_run(f"[{q['id']} | {q['answer']}] ")
                    r_tg.bold = True
                    r_tg.font.name = "Times New Roman"
                    r_tg.font.size = Pt(10)
                    r_tg.font.color.rgb = COLOR_TAG

                add_chem_runs_hoa10(p_q, "\n".join(q["stem_lines"]).strip(), size=11)

                for imp in q.get("image_paths", []):
                    if os.path.exists(imp):
                        p_im = doc.add_paragraph()
                        p_im.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        r_im = p_im.add_run()
                        r_im.add_picture(imp, width=Cm(11.0))

                for tbl_d in q.get("tables", []):
                    add_data_table_to_doc(doc, tbl_d)

                p2_ans_map = {}
                if q.get("answer"):
                    for m in re.finditer(r'([a-d])\s*[:\.]\s*([ĐS])', q["answer"]):
                        p2_ans_map[m.group(1).lower()] = m.group(2)

                for item_k in ["a", "b", "c", "d"]:
                    item_txt = q["p2_items"].get(item_k, "")
                    if not item_txt:
                        continue
                    p_it = doc.add_paragraph()
                    p_it.paragraph_format.left_indent = Cm(0.6)
                    p_it.paragraph_format.space_before = Pt(1.5)
                    p_it.paragraph_format.space_after = Pt(1.5)
                    corr_val = p2_ans_map.get(item_k, "")
                    if is_teacher and corr_val:
                        r_it = p_it.add_run(f"{item_k}) [{corr_val}] ")
                        r_it.bold = True
                        r_it.font.name = "Times New Roman"
                        r_it.font.size = Pt(10.5)
                        r_it.font.color.rgb = COLOR_CORRECT if corr_val == "Đ" else COLOR_PRIMARY
                    else:
                        r_it = p_it.add_run(f"{item_k}) ")
                        r_it.bold = True
                        r_it.font.name = "Times New Roman"
                        r_it.font.size = Pt(10.5)
                    add_chem_runs_hoa10(p_it, item_txt, size=10.5)

                if is_teacher:
                    add_solution_block_hoa10(doc, q)
                q_counter += 1

        # Phần III
        if selected_qs["P3"]:
            p3_start = q_counter
            p3_end = q_counter + len(selected_qs["P3"]) - 1
            p_p3 = doc.add_paragraph()
            p_p3.paragraph_format.space_before = Pt(8)
            p_p3.paragraph_format.space_after = Pt(2)
            r_p3 = p_p3.add_run(f"PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn. Thí sinh trả lời từ câu {p3_start} đến câu {p3_end}.")
            r_p3.bold = True
            r_p3.font.name = "Times New Roman"
            r_p3.font.size = Pt(11)
            r_p3.font.color.rgb = COLOR_HF_RED

            p_p3_sub = doc.add_paragraph()
            p_p3_sub.paragraph_format.space_before = Pt(0)
            p_p3_sub.paragraph_format.space_after = Pt(4)
            r_p3_sub = p_p3_sub.add_run("Thí sinh điền kết quả hoặc số liệu tương ứng vào bài làm.")
            r_p3_sub.italic = True
            r_p3_sub.font.name = "Times New Roman"
            r_p3_sub.font.size = Pt(10)

            for q in selected_qs["P3"]:
                p_q = doc.add_paragraph()
                p_q.paragraph_format.space_before = Pt(3)
                p_q.paragraph_format.space_after = Pt(2)
                r_num = p_q.add_run(f"Câu {q_counter}. ")
                r_num.bold = True
                r_num.font.name = "Times New Roman"
                r_num.font.size = Pt(11)
                r_num.font.color.rgb = COLOR_PRIMARY

                if is_teacher and q.get("answer"):
                    r_tg = p_q.add_run(f"[{q['id']} | ĐA: {q['answer']}] ")
                    r_tg.bold = True
                    r_tg.font.name = "Times New Roman"
                    r_tg.font.size = Pt(10)
                    r_tg.font.color.rgb = COLOR_TAG

                add_chem_runs_hoa10(p_q, "\n".join(q["stem_lines"]).strip(), size=11)

                for imp in q.get("image_paths", []):
                    if os.path.exists(imp):
                        p_im = doc.add_paragraph()
                        p_im.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        r_im = p_im.add_run()
                        r_im.add_picture(imp, width=Cm(11.0))

                for tbl_d in q.get("tables", []):
                    add_data_table_to_doc(doc, tbl_d)

                if not is_teacher:
                    p_ans = doc.add_paragraph()
                    p_ans.paragraph_format.left_indent = Cm(0.6)
                    p_ans.paragraph_format.space_before = Pt(1)
                    p_ans.paragraph_format.space_after = Pt(2)
                    r_ans = p_ans.add_run("Đáp số / Trả lời: .......................................................................................")
                    r_ans.italic = True
                    r_ans.font.name = "Times New Roman"
                    r_ans.font.size = Pt(10.5)
                else:
                    add_solution_block_hoa10(doc, q)
                q_counter += 1

        suffix = "GiaoVien" if is_teacher else "HocSinh"
        fname = f"{out_prefix}_{suffix}.docx"
        out_path = os.path.join(OUT_DIR, fname)
        doc.save(out_path)
        print(f"  + Xuất file {suffix}: {out_path}")


if __name__ == "__main__":
    print_bank_summary()
