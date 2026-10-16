# -*- coding: utf-8 -*-
"""
Script sinh tự động 2 file Word (.docx) cho ĐỀ SỐ 18 (Chương 6: HỢP CHẤT CARBONYL - CARBOXYLIC ACID):
1. Bản Học sinh (De_Kiem_Tra_Hoa_11_De18.docx)
2. Bản Giáo viên (De_Kiem_Tra_Hoa_11_De18_GiaoVien.docx)

Tuân thủ nghiêm ngặt Quy chuẩn Form Đề Kiểm Tra (GEMINI.md):
- Tuyệt đối KHÔNG dùng bảng (Table) trong nội dung câu hỏi/đáp án/lời giải.
- Dùng Paragraph + Tab stops chuẩn cho các phương án trắc nghiệm.
- Công thức hóa học xử lý triệt để chỉ số trên/dưới (~ cho sub, ^ cho sup).
- Header & Footer chuẩn với viền màu Sage Green (#8BC390) và chữ Coral Red (#D32F2F).
- Bản Giáo viên có ghi chú đỏ và lời giải thụt lề màu xanh #1565C0.
"""

import os
import sys

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH

# Đảm bảo import được module form_de_kiem_tra
sys.path.insert(0, os.path.dirname(__file__))
from form_de_kiem_tra import (
    setup_header_footer,
    build_title_block,
    add_section_header,
    render_question_prompt,
    render_options_no_table,
    render_statement,
    add_solution_no_table,
    add_student_writing_lines,
    render_rich_text,
    set_run_style,
    COLOR_CORRECT,
    COLOR_SOLUTION,
    COLOR_PRIMARY
)

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
IMG_DIR = os.path.join(OUT_DIR, "images_crop_hoa11_de18")

def add_p(doc, space_before=3, space_after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    return p

def add_centered_picture(doc, img_path, width_cm=8.0):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(img_path, width=Cm(width_cm))

def build_exam_document(is_teacher=False):
    doc = docx.Document()
    setup_header_footer(doc)

    build_title_block(
        doc,
        school_line="SỞ GIÁO DỤC VÀ ĐÀO TẠO THỪA THIÊN HUẾ",
        exam_title="ĐỀ KIỂM TRA ĐỊNH KỲ CHƯƠNG 6 (HỢP CHẤT CARBONYL - CARBOXYLIC ACID)",
        subject_line="MÔN: HOÁ HỌC - LỚP 11",
        num_pages="04",
        duration="50 phút",
        exam_code="118",
        atomic_masses="H = 1; C = 12; N = 14; O = 16; Na = 23; Mg = 24; Al = 27; P = 31; S = 32; Cl = 35,5; K = 39; Ca = 40; Cr = 52; Mn = 55; Fe = 56; Cu = 64; Zn = 65; Br = 80; Ag = 108; Ba = 137.",
        is_teacher=is_teacher
    )

    # =========================================================================
    # PHẦN I
    # =========================================================================
    add_section_header(
        doc,
        title="PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án."
    )

    # Câu 1
    p = add_p(doc)
    render_question_prompt(p, "Câu 1. Gọi tên một hợp chất có công thức cấu tạo như sau: (CH~3~)~3~C-CHO")
    opts_c1 = [
        "A. 2-ethyl-2-methylethanal.",
        "B. 2,2-dimethylpentanal.",
        "C. Tert-butyletanal.",
        "*D. 2,2-dimethylpropanal.*"
    ]
    render_options_no_table(doc, opts_c1, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Mạch chính dài nhất chứa nhóm chức -CHO gồm 3 nguyên tử carbon (đánh số 1 từ carbon nhóm -CHO), tại vị trí carbon số 2 có 2 nhóm thế methyl (-CH~3~).\n"
            "Tên thay thế: 2,2-dimethylpropanal.\n"
            "=> Chọn D."
        )

    # Câu 2
    p = add_p(doc)
    render_question_prompt(p, "Câu 2. Công thức cấu tạo thu gọn của acetone là")
    opts_c2 = [
        "A. OHC-CHO.",
        "*B. CH~3~COCH~3~.*",
        "C. CH~3~CHO.",
        "D. CH~2~=CHCHO."
    ]
    render_options_no_table(doc, opts_c2, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Acetone (dimethyl ketone hay propan-2-one) có công thức cấu tạo thu gọn là CH~3~COCH~3~.\n"
            "=> Chọn B."
        )

    # Câu 3
    p = add_p(doc)
    render_question_prompt(p, "Câu 3. Tên gọi của aldehyde có khối lượng phân tử nhỏ nhất là")
    opts_c3 = [
        "A. aldehyde acetic.",
        "*B. methanal.*",
        "C. acetone.",
        "D. methanol."
    ]
    render_options_no_table(doc, opts_c3, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Aldehyde đơn giản nhất và có phân tử khối nhỏ nhất là HCHO (methanal hay formaldehyde, M = 30).\n"
            "=> Chọn B."
        )

    # Câu 4
    p = add_p(doc)
    render_question_prompt(p, "Câu 4. Phản ứng giữa CH~3~COCH~3~ với HCN thuộc loại phản ứng")
    opts_c4 = [
        "A. phản ứng thuận nghịch.",
        "*B. phản ứng cộng.*",
        "C. phản ứng thế.",
        "D. phản ứng phân hủy."
    ]
    render_options_no_table(doc, opts_c4, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Liên kết đôi >C=O trong phân tử acetone tham gia phản ứng cộng nucleophile với HCN tạo thành cyanohydrin: CH~3~COCH~3~ + HCN -> (CH~3~)~2~C(OH)CN. Đây là phản ứng cộng.\n"
            "=> Chọn B."
        )

    # Câu 5
    p = add_p(doc)
    render_question_prompt(p, "Câu 5. Điều kiện của phản ứng giữa aldehyde hoặc ketone với iodine trong dung dịch NaOH tạo iodoform là")
    opts_c5 = [
        "*A. aldehyde hoặc ketone có nhóm methyl cạnh nhóm carbonyl (CH~3~CO-).*",
        "B. aldehyde hoặc ketone có nhóm methyl cạnh nhóm carbonyl (C~2~H~3~CO-).",
        "C. aldehyde hoặc ketone có nhóm methyl cạnh nhóm carbonyl (C~2~H~5~CO-).",
        "D. aldehyde hoặc ketone có nhóm methyl cạnh nhóm carbonyl (C~3~H~7~CO-)."
    ]
    render_options_no_table(doc, opts_c5, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Phản ứng tạo iodoform (kết tủa vàng CHI~3~) đặc trưng cho các hợp chất chứa nhóm methyl ketone (CH~3~-C=O) hoặc acetaldehyde (CH~3~-CHO) và các alcohol có nhóm CH~3~-CH(OH)-.\n"
            "=> Chọn A."
        )

    # Câu 6
    p = add_p(doc)
    render_question_prompt(p, "Câu 6 (SBT - CTST). Cho dãy các chất sau: HCHO, CH~3~COOH, CH~3~CHO, HCOOH, C~2~H~5~OH. Số chất trong dãy có thể tham gia phản ứng tráng bạc là")
    opts_c6 = [
        "*A. 3.*",
        "B. 6.",
        "C. 4.",
        "D. 5."
    ]
    render_options_no_table(doc, opts_c6, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Các chất có phản ứng tráng bạc phải có nhóm formyl (-CHO). Trong dãy gồm 3 chất thỏa mãn: HCHO, CH~3~CHO và HCOOH.\n"
            "=> Chọn A."
        )

    # Câu 7
    p = add_p(doc)
    render_question_prompt(p, "Câu 7 (SBT - CTST). Có 3 dung dịch: CH~3~CHO, CH~3~COOH, HCOOH đựng trong 3 lọ mất nhãn. Hóa chất có thể dùng để phân biệt ba dung dịch trên là")
    opts_c7 = [
        "A. quỳ tím, CuO.",
        "B. quỳ tím, Na.",
        "*C. quỳ tím, dung dịch AgNO~3~ trong NH~3~ dư.*",
        "D. dung dịch AgNO~3~ trong NH~3~ dư, CuO."
    ]
    render_options_no_table(doc, opts_c7, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- Dùng quỳ tím: CH~3~CHO không đổi màu quỳ; CH~3~COOH và HCOOH làm quỳ tím hóa đỏ.\n"
            "- Cho dung dịch AgNO~3~/NH~3~ (đun nóng) vào 2 acid: HCOOH có nhóm -CHO nên tham gia phản ứng tráng bạc tạo kết tủa Ag sáng bóng; CH~3~COOH không có hiện tượng.\n"
            "=> Chọn C."
        )

    # Câu 8
    p = add_p(doc)
    render_question_prompt(p, "Câu 8 (SBT - CD). Phản ứng nào sau đây thể hiện tính oxi hóa của propanal?")
    opts_c8 = [
        "A. C~2~H~5~CHO + 2[Ag(NH~3~)~2~]OH -> C~2~H~5~COONH~4~ + 3NH~3~ + 2Ag v + H~2~O.",
        "B. C~2~H~5~CHO + Br~2~ + H~2~O -> C~2~H~5~COOH + 2HBr.",
        "C. C~2~H~5~CHO + 2Cu(OH)~2~ + NaOH -> C~2~H~5~COONa + Cu~2~O v + 3H~2~O.",
        "*D. C~2~H~5~CHO + 2[H] --(LiAlH~4~)--> CH~3~CH~2~CH~2~OH.*"
    ]
    render_options_no_table(doc, opts_c8, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Trong phản ứng với chất khử LiAlH~4~, propanal nhận H (số oxi hóa của carbon carbonyl giảm từ +1 xuống -1 trong alcohol bậc 1) nên propanal đóng vai trò là chất oxi hóa.\n"
            "Ở các phản ứng A, B, C, propanal bị oxi hóa thành carboxylic acid/muối carboxylate.\n"
            "=> Chọn D."
        )

    # Câu 9
    p = add_p(doc)
    render_question_prompt(p, "Câu 9 (SBT - CD). Khi cho ethanal phản ứng với Cu(OH)~2~ trong môi trường kiềm ở nhiệt độ thích hợp, hiện tượng nào sau đây sẽ xảy ra?")
    opts_c9 = [
        "A. Cu(OH)~2~ bị tan ra, tạo dung dịch màu xanh.",
        "B. Có mùi chua của giấm, do phản ứng sinh ra acetic acid.",
        "*C. Tạo kết tủa đỏ gạch do phản ứng sinh ra Cu~2~O.*",
        "D. Sinh ra CuO màu đen."
    ]
    render_options_no_table(doc, opts_c9, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Ethanal bị oxi hóa bởi Cu(OH)~2~ trong môi trường kiềm đun nóng tạo kết tủa Cu~2~O màu đỏ gạch:\n"
            "CH~3~CHO + 2Cu(OH)~2~ + NaOH --t^o^--> CH~3~COONa + Cu~2~O v (đỏ gạch) + 3H~2~O.\n"
            "=> Chọn C."
        )

    # Câu 10
    p = add_p(doc)
    render_question_prompt(p, "Câu 10. Có thể dùng thuốc thử nào trong các chất dưới đây để nhận biết các chất: ethyl alcohol, glycerol, acetic aldehyde đựng trong ba lọ mất nhãn?")
    opts_c10 = [
        "*A. Cu(OH)~2~ trong môi trường kiềm (hoặc Copper(II) hydroxide, Na).*",
        "B. Quỳ tím.",
        "C. Kim loại sodium.",
        "D. Dung dịch AgNO~3~/NH~3~."
    ]
    render_options_no_table(doc, opts_c10, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- Cho Cu(OH)~2~ ở nhiệt độ thường: glycerol hòa tan Cu(OH)~2~ tạo dung dịch màu xanh lam thẫm (phức đồng); ethyl alcohol và acetic aldehyde không phản ứng.\n"
            "- Đun nóng 2 mẫu còn lại với Cu(OH)~2~ trong kiềm: acetic aldehyde tạo kết tủa Cu~2~O đỏ gạch; ethyl alcohol không đổi.\n"
            "=> Chọn A."
        )

    # Câu 11
    p = add_p(doc)
    render_question_prompt(p, "Câu 11. Có bao nhiêu aldehyde hai chức có công thức đơn giản nhất là C~2~H~3~O?")
    opts_c11 = [
        "*A. 2.*",
        "B. 4.",
        "C. 1.",
        "D. 3."
    ]
    render_options_no_table(doc, opts_c11, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Aldehyde hai chức có 2 nhóm -CHO nên phân tử chứa 2 nguyên tử O.\n"
            "Công thức đơn giản nhất là C~2~H~3~O => Công thức phân tử có dạng (C~2~H~3~O)~n~. Với n = 2 => C~4~H~6~O~2~.\n"
            "Các đồng phân aldehyde 2 chức mạch hở gồm:\n"
            "(1) OHC-CH~2~-CH~2~-CHO (butanedial)\n"
            "(2) OHC-CH(CH~3~)-CHO (2-methylpropanedial)\n"
            "=> Có 2 đồng phân thỏa mãn. Chọn A."
        )

    # Câu 12
    p = add_p(doc)
    render_question_prompt(p, "Câu 12. Có bao nhiêu phản ứng xảy ra khi cho các chất C~6~H~5~OH, NaHCO~3~, NaOH, HCl tác dụng với nhau từng đôi một?")
    opts_c12 = [
        "A. 3.",
        "*B. 4.*",
        "C. 5.",
        "D. 6."
    ]
    render_options_no_table(doc, opts_c12, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Các phản ứng xảy ra khi trộn từng đôi một:\n"
            "(1) C~6~H~5~OH + NaOH -> C~6~H~5~ONa + H~2~O\n"
            "(2) NaHCO~3~ + NaOH -> Na~2~CO~3~ + H~2~O\n"
            "(3) NaHCO~3~ + HCl -> NaCl + CO~2~ ^ + H~2~O\n"
            "(4) NaOH + HCl -> NaCl + H~2~O\n"
            "(Lưu ý: Phenol tính acid yếu hơn carbonic acid H~2~CO~3~ nên phenol không phản ứng với NaHCO~3~).\n"
            "=> Có 4 phản ứng. Chọn B."
        )

    # Câu 13
    p = add_p(doc)
    render_question_prompt(p, "Câu 13. Cho các hợp chất có công thức sau: CH~3~OH; CH~3~CHO; HCOOH; C~2~H~5~NH~2~; C~6~H~6~; CH~3~COCH~3~; C~2~H~6~; C~2~H~5~Cl. Trong các hợp chất trên, có bao nhiêu chất chứa nhóm carbonyl trong phân tử?")
    opts_c13 = [
        "*A. 3.*",
        "B. 4.",
        "C. 2.",
        "D. 5."
    ]
    render_options_no_table(doc, opts_c13, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Các hợp chất chứa nhóm carbonyl (>C=O) trong dãy gồm:\n"
            "- CH~3~CHO (aldehyde, chứa -CHO)\n"
            "- CH~3~COCH~3~ (ketone, chứa >C=O)\n"
            "- HCOOH (carboxylic acid, chứa -COOH trong đó có nhóm C=O)\n"
            "=> Có 3 hợp chất chứa nhóm C=O. Chọn A."
        )

    # Câu 14
    p = add_p(doc)
    render_question_prompt(p, "Câu 14. Quế có vị cay, mùi thơm nồng, được sử dụng phổ biến làm gia vị, vị thuốc trong Đông y. Hợp chất hữu cơ X tạo mùi đặc trưng của quế, có công thức phân tử là C~9~H~8~O. Trong phân tử X chứa vòng benzene. X tham gia phản ứng tráng bạc và có đồng phân hình học dạng trans. Xác định công thức cấu tạo của X.")
    
    # Chèn ảnh grid các phương án A, B, C, D
    c14_img = os.path.join(IMG_DIR, "c14_grid_teacher.png" if is_teacher else "c14_grid_student.png")
    add_centered_picture(doc, c14_img, width_cm=14.5)

    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- X có phản ứng tráng bạc => X có nhóm chức aldehyde (-CHO).\n"
            "- X có đồng phân hình học dạng trans => liên kết đôi C=C có các nhóm thế khác nhau ở mỗi carbon nối đôi.\n"
            "- Công thức phân tử C~9~H~8~O có vòng benzene (C~6~H~5~-) và mạch nhánh chứa -CH=CH-CHO.\n"
            "=> X là (E)-3-phenylpropenal (trans-cinnamaldehyde), công thức cấu tạo tương ứng với phương án A.\n"
            "=> Chọn A."
        )

    # Câu 15
    p = add_p(doc)
    render_question_prompt(p, "Câu 15. Một thí nghiệm được mô tả như hình sau đây:")
    add_centered_picture(doc, os.path.join(IMG_DIR, "c15_experiment.png"), width_cm=8.5)
    p2 = add_p(doc)
    render_rich_text(p2, "Chất lỏng thu được ở ống nghiệm B có mùi táo, có tên gọi là:")
    opts_c15 = [
        "A. ethyl formate.",
        "*B. methyl propionate.*",
        "C. ethyl propionate.",
        "D. propyl formate."
    ]
    render_options_no_table(doc, opts_c15, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Ống nghiệm A chứa hỗn hợp CH~3~OH (methanol), CH~3~CH~2~COOH (propionic acid) và xúc tác H~2~SO~4~ đặc được đun nóng.\n"
            "Phản ứng ester hóa: CH~3~CH~2~COOH + CH~3~OH <=(H~2~SO~4~ dac, t^o^)=> CH~3~CH~2~COOCH~3~ + H~2~O.\n"
            "Ester tạo thành là methyl propionate có nhiệt độ sôi thấp, bay hơi sang ống nghiệm B ngâm nước đá và ngưng tụ lại thành chất lỏng có mùi táo thơm đặc trưng.\n"
            "=> Chọn B."
        )

    # Câu 16
    p = add_p(doc)
    render_question_prompt(p, "Câu 16. Cho giá trị K~a~ của một số carboxylic acid ở 25 ^o^C dưới đây:")
    add_centered_picture(doc, os.path.join(IMG_DIR, "c16_table.png"), width_cm=14.0)
    p2 = add_p(doc)
    render_rich_text(p2, "Tính acid của carboxylic acid nào là mạnh nhất?")
    opts_c16 = [
        "*A. HCOOH.*",
        "B. CH~3~COOH.",
        "C. CH~3~CH~2~COOH.",
        "D. C~6~H~5~COOH."
    ]
    render_options_no_table(doc, opts_c16, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Giá trị hằng số phân li acid K~a~ càng lớn thì tính acid càng mạnh.\n"
            "Trong 4 acid: HCOOH (17,72.10^-5^) > C~6~H~5~COOH (6,25.10^-5^) > CH~3~COOH (1,75.10^-5^) > CH~3~CH~2~COOH (1,35.10^-5^).\n"
            "Do đó HCOOH có tính acid mạnh nhất.\n"
            "=> Chọn A."
        )

    # Câu 17
    p = add_p(doc)
    render_question_prompt(p, "Câu 17 (SBT - CD). Để trung hòa 40 mL giấm ăn cần 25 mL dung dịch NaOH 1 M. Biết khối lượng riêng của giấm ăn xấp xỉ là 1 g/mL. Mẫu giấm ăn này có nồng độ phần trăm là")
    opts_c17 = [
        "A. 3,5%.",
        "*B. 3,75%.*",
        "C. 4%.",
        "D. 5%."
    ]
    render_options_no_table(doc, opts_c17, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Số mol NaOH: n~NaOH~ = 0,025 . 1 = 0,025 mol.\n"
            "Phương trình phản ứng: CH~3~COOH + NaOH -> CH~3~COONa + H~2~O\n"
            "=> n~CH3COOH~ = n~NaOH~ = 0,025 mol.\n"
            "Khối lượng CH~3~COOH: m~CH3COOH~ = 0,025 . 60 = 1,5 gam.\n"
            "Khối lượng dung dịch giấm: m~dd~ = V . D = 40 . 1 = 40 gam.\n"
            "Nồng độ phần trăm của acetic acid trong giấm: C% = (1,5 / 40) . 100% = 3,75%.\n"
            "=> Chọn B."
        )

    # Câu 18
    p = add_p(doc)
    render_question_prompt(p, "Câu 18. Đun 12 gam acetic acid với một lượng dư ethanol (H~2~SO~4~ đặc làm xúc tác). Đến khi dừng thí nghiệm thu được 8,36 gam ester. Tính hiệu suất phản ứng ester hóa.")
    opts_c18 = [
        "A. 65,15%.",
        "B. 54,54%.",
        "C. 34,84%.",
        "*D. 47,50%.*"
    ]
    render_options_no_table(doc, opts_c18, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "n~CH3COOH~ = 12 / 60 = 0,2 mol.\n"
            "Ester tạo thành là ethyl acetate CH~3~COOC~2~H~5~ (M = 88).\n"
            "n~ester thu được~ = 8,36 / 88 = 0,095 mol.\n"
            "Phương trình phản ứng: CH~3~COOH + C~2~H~5~OH <=(H~2~SO~4~ dac, t^o^)=> CH~3~COOC~2~H~5~ + H~2~O\n"
            "Theo lí thuyết: n~ester lt~ = n~CH3COOH~ = 0,2 mol.\n"
            "Hiệu suất phản ứng: H = (0,095 / 0,2) . 100% = 47,50%.\n"
            "=> Chọn D."
        )

    # =========================================================================
    # PHẦN II
    # =========================================================================
    add_section_header(
        doc,
        title="PHẦN II. Câu trắc nghiệm đúng sai.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai."
    )

    # Câu 1
    p = add_p(doc)
    render_question_prompt(p, "Câu 1. Muscone là hợp chất tạo nên mùi thơm đặc trưng của xạ hương, có công thức cấu tạo như sau:")
    add_centered_picture(doc, os.path.join(IMG_DIR, "muscone.png"), width_cm=7.0)

    stmts_c1 = [
        ("a) Muscone thuộc hợp chất ketone.", "ĐÚNG"),
        ("b) Muscone có công thức phân tử là C~16~H~30~O.", "ĐÚNG"),
        ("c) Muscone có thể tham gia phản ứng tráng bạc với thuốc thử Tollens.", "SAI"),
        ("d) Muscone có nhiệt độ sôi cao hơn so với ethanal.", "ĐÚNG")
    ]
    for stmt, ans in stmts_c1:
        p_st = add_p(doc, space_before=1, space_after=1)
        p_st.paragraph_format.left_indent = Cm(0.5)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- a) Đúng: Phân tử muscone có chứa nhóm carbonyl (>C=O) liên kết trực tiếp với hai nguyên tử carbon trong vòng kín nên thuộc loại ketone.\n"
            "- b) Đúng: Vòng gồm 15 nguyên tử C + 1 nhóm -CH~3~ nhánh = 16 nguyên tử C; phân tử có 1 vòng no và 1 liên kết đôi C=O (độ bất bão hòa k = 2), công thức phân tử là C~16~H~30~O (M = 238).\n"
            "- c) Sai: Ketone không tham gia phản ứng tráng bạc với thuốc thử Tollens.\n"
            "- d) Đúng: Phân tử khối của muscone (M = 238) lớn hơn rất nhiều so với ethanal (M = 44) nên lực phân tán van der Waals giữa các phân tử muscone lớn hơn nhiều, dẫn đến nhiệt độ sôi của muscone (~328 °C) cao hơn rất nhiều so với ethanal (20,8 °C)."
        )

    # Câu 2
    p = add_p(doc)
    render_question_prompt(p, "Câu 2. Trong nọc độc của một số côn trùng như kiến, ong, muỗi... có chứa hợp chất acid X gây bỏng rát và ngứa.")
    stmts_c2 = [
        ("a) X là acetic acid, có công thức cấu tạo là CH~3~COOH.", "SAI"),
        ("b) X là formic acid, có công thức phân tử là HCOOH.", "ĐÚNG"),
        ("c) Khi bị côn trùng đốt thì nên bôi vôi vào vết côn trùng đốt nó sẽ trung hòa lượng acid X làm cho vết đốt đỡ ngứa rát: 2HCOOH + Ca(OH)~2~ -> (HCOO)~2~Ca + 2H~2~O.", "ĐÚNG"),
        ("d) Cho dung dịch acid X vào ống nghiệm chứa dung dịch AgNO~3~/NH~3~, sau đó đun nóng nhẹ thì thấy thành ống nghiệm sáng bóng như gương.", "ĐÚNG")
    ]
    for stmt, ans in stmts_c2:
        p_st = add_p(doc, space_before=1, space_after=1)
        p_st.paragraph_format.left_indent = Cm(0.5)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- a) Sai: Acid trong nọc côn trùng là formic acid (methanoic acid), không phải acetic acid.\n"
            "- b) Đúng: Hợp chất acid X là formic acid có công thức phân tử là HCOOH.\n"
            "- c) Đúng: Dung dịch nước vôi Ca(OH)~2~ có tính kiềm giúp trung hòa acid HCOOH làm giảm cảm giác đau buốt và ngứa rát.\n"
            "- d) Đúng: Phân tử HCOOH có chứa nhóm aldehyde (-CHO) nên có khả năng phản ứng với thuốc thử Tollens sinh ra bạc kim loại bám vào thành ống nghiệm."
        )

    # Câu 3
    p = add_p(doc)
    render_question_prompt(p, "Câu 3. Tiến hành phản ứng tráng bạc aldehyde acetic với thuốc thử Tollens theo các bước sau đây:\n"
                              "  Bước 1: Rửa sạch ống nghiệm bằng cách nhỏ vài giọt dung dịch kiềm NaOH, đun nóng nhẹ, tráng đều, sau đó đổ đi và tráng lại bằng nước cất.\n"
                              "  Bước 2: Nhỏ vào ống nghiệm 2 giọt dung dịch AgNO~3~ và 1 giọt dung dịch NH~3~, xuất hiện kết tủa nâu xám của silver oxide, nhỏ tiếp dung dịch NH~3~ đến khi kết tủa tan hết.\n"
                              "  Bước 3: Nhỏ tiếp vào ống nghiệm 2 giọt dung dịch CH~3~CHO.\n"
                              "  Bước 4: Lắc đều ống nghiệm, đun cách thủy (trong cốc nước nóng) vài phút ở 60 – 70 ^o^C.")
    stmts_c3 = [
        ("a) Sau bước 4 quan sát thấy thành ống nghiệm sáng bóng như gương.", "ĐÚNG"),
        ("b) Ở bước 3 nếu thay aldehyde acetic bằng acetone thì thu được hiện tượng tương tự.", "SAI"),
        ("c) Ở bước 4 xảy ra phản ứng oxi hóa – khử trong đó aldehyde acetic là chất bị khử.", "SAI"),
        ("d) Sản phẩm hữu cơ thu được trong dung dịch sau bước 4 có công thức phân tử là C~2~H~7~NO~2~.", "ĐÚNG")
    ]
    for stmt, ans in stmts_c3:
        p_st = add_p(doc, space_before=1, space_after=1)
        p_st.paragraph_format.left_indent = Cm(0.5)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- a) Đúng: Phản ứng tạo kim loại Ag bám vào thành ống nghiệm tạo gương bạc sáng bóng.\n"
            "- b) Sai: Acetone là ketone không tham gia phản ứng tráng bạc.\n"
            "- c) Sai: CH~3~CHO đóng vai trò là chất khử (bị oxi hóa thành ion ammonium acetate CH~3~COONH~4~).\n"
            "- d) Đúng: Sản phẩm hữu cơ tạo thành là muối ammonium acetate (CH~3~COONH~4~), có CTPT là C~2~H~7~NO~2~."
        )

    # Câu 4
    p = add_p(doc)
    render_question_prompt(p, "Câu 4. Hãy chọn các nhận định về hiện tượng và phương trình hóa học xảy ra đúng trong các trường hợp sau:")
    stmts_c4 = [
        ("a) Cho acetic aldehyde tác dụng với nước bromine thấy nước bromine bị mất màu do xảy ra phản ứng: CH~3~CHO + Br~2~ + H~2~O -> CH~3~COOH + 2HBr.", "ĐÚNG"),
        ("b) Cho formic aldehyde tác dụng với dung dịch AgNO~3~ trong NH~3~ dư thấy xuất hiện kết tủa trắng sáng do xảy ra phản ứng: HCHO + 4[Ag(NH~3~)~2~]OH --t^o^--> (NH~4~)~2~CO~3~ + 4Ag v + 6NH~3~ + 2H~2~O.", "ĐÚNG"),
        ("c) Cho acetone tác dụng với I~2~ trong dung dịch NaOH dư không có hiện tượng gì xảy ra.", "SAI"),
        ("d) Cho acetic aldehyde tác dụng với Cu(OH)~2~ trong NaOH dư thấy xuất hiện kết tủa đỏ gạch do xảy ra phản ứng: C~2~H~5~CHO + 2Cu(OH)~2~ + NaOH --t^o^--> C~2~H~5~COONa + Cu~2~O v + 3H~2~O.", "SAI")
    ]
    for stmt, ans in stmts_c4:
        p_st = add_p(doc, space_before=1, space_after=1)
        p_st.paragraph_format.left_indent = Cm(0.5)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- a) Đúng: Acetic aldehyde bị nước bromine oxi hóa thành acetic acid làm mất màu nước bromine.\n"
            "- b) Đúng: Formic aldehyde (HCHO) phản ứng tráng bạc tỉ lệ 1 : 4 Ag tạo kết tủa bạc kim loại sáng bóng.\n"
            "- c) Sai: Acetone chứa nhóm CH~3~CO- phản ứng với I~2~/NaOH tạo kết tủa màu vàng iodoform (CHI~3~).\n"
            "- d) Sai: Acetic aldehyde là CH~3~CHO, phương trình trên viết sai công thức của chất phản ứng (viết thành propionaldehyde C~2~H~5~CHO)."
        )

    # =========================================================================
    # PHẦN III
    # =========================================================================
    add_section_header(
        doc,
        title="PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 6."
    )

    # Câu 1
    p = add_p(doc)
    render_question_prompt(p, "Câu 1. Cho các hóa chất sau: dung dịch HCN, dung dịch Br~2~, dung dịch NaOH, dung dịch [Ag(NH~3~)~2~]OH, LiAlH~4~, Cu(OH)~2~/NaOH, I~2~/NaOH. Có bao nhiêu chất phản ứng được với acetone ở điều kiện nhiệt độ thích hợp?")
    if is_teacher:
        p_ans = add_p(doc)
        render_rich_text(p_ans, "Đáp án: *3*")
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Các chất phản ứng được với acetone (CH~3~COCH~3~) gồm 3 chất:\n"
            "(1) Dung dịch HCN (phản ứng cộng nucleophile):\n"
            "    CH~3~COCH~3~ + HCN -> (CH~3~)~2~C(OH)CN\n"
            "(2) LiAlH~4~ (phản ứng khử carbonyl thành alcohol bậc 2):\n"
            "    CH~3~COCH~3~ + 2[H] --(LiAlH~4~)--> CH~3~CH(OH)CH~3~\n"
            "(3) Dung dịch I~2~/NaOH (phản ứng tạo iodoform kết tủa vàng):\n"
            "    CH~3~COCH~3~ + 3I~2~ + 4NaOH -> CH~3~COONa + CHI~3~ v + 3NaI + 3H~2~O\n"
            "Acetone không phản ứng với dung dịch Br~2~, dung dịch NaOH, [Ag(NH~3~)~2~]OH, Cu(OH)~2~/NaOH đun nóng."
        )

    # Câu 2
    p = add_p(doc)
    render_question_prompt(p, "Câu 2. Aldehyde (X) no, đơn chức, mạch hở không nhánh. Phân tử khối của (X) được xác định thông qua kết quả phổ khối lượng với peak ion phân tử có giá trị m/z lớn nhất trong hình dưới đây. Khối lượng mol của X là bao nhiêu g/mol?")
    add_centered_picture(doc, os.path.join(IMG_DIR, "ms_spectrum.png"), width_cm=8.5)
    if is_teacher:
        p_ans = add_p(doc)
        render_rich_text(p_ans, "Đáp án: *72*")
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Trên phổ khối lượng (MS), peak ion phân tử [M^+^] có giá trị m/z lớn nhất là 72.\n"
            "=> Phân tử khối của X là 72 amu, tương ứng với khối lượng mol M~X~ = 72 g/mol.\n"
            "(X là aldehyde no, đơn chức, mạch hở C~n~H~2n~O: 14n + 16 = 72 => n = 4, công thức cấu tạo là CH~3~CH~2~CH~2~CHO)."
        )

    # Câu 3
    p = add_p(doc)
    render_question_prompt(p, "Câu 3. Cho các chất sau: Cu, Na, Mg, KOH, Na~2~CO~3~, NaCl, HCOOH, C~2~H~5~OH. Có bao nhiêu chất phản ứng được với CH~3~COOH?")
    if is_teacher:
        p_ans = add_p(doc)
        render_rich_text(p_ans, "Đáp án: *5*")
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Có 5 chất phản ứng được với CH~3~COOH là: Na, Mg, KOH, Na~2~CO~3~, C~2~H~5~OH.\n"
            "(1) 2CH~3~COOH + 2Na -> 2CH~3~COONa + H~2~ ^\n"
            "(2) 2CH~3~COOH + Mg -> (CH~3~COO)~2~Mg + H~2~ ^\n"
            "(3) CH~3~COOH + KOH -> CH~3~COOK + H~2~O\n"
            "(4) 2CH~3~COOH + Na~2~CO~3~ -> 2CH~3~COONa + CO~2~ ^ + H~2~O\n"
            "(5) CH~3~COOH + C~2~H~5~OH <=(H~2~SO~4~ dac, t^o^)=> CH~3~COOC~2~H~5~ + H~2~O\n"
            "(Cu đứng sau H nên không phản ứng; NaCl là muối của acid mạnh không phản ứng; HCOOH không phản ứng)."
        )

    # Câu 4
    p = add_p(doc)
    render_question_prompt(p, "Câu 4. Hợp chất carboxylic acid X có công thức đơn giản nhất là CH~2~O. Phổ MS cho thấy X có các tín hiệu sau:")
    add_centered_picture(doc, os.path.join(IMG_DIR, "c4_table_ms.png"), width_cm=8.0)
    p2 = add_p(doc)
    render_rich_text(p2, "Biết mảnh [M^+^] có giá trị m/z lớn nhất. Tổng số nguyên tử trong một phân tử của X là bao nhiêu?")
    if is_teacher:
        p_ans = add_p(doc)
        render_rich_text(p_ans, "Đáp án: *8*")
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- Giá trị m/z lớn nhất trên phổ MS là 60 => phân tử khối của X là 60 amu.\n"
            "- Công thức đơn giản nhất là CH~2~O (M = 30) => CTPT có dạng (CH~2~O)~n~.\n"
            "  Ta có: 30n = 60 => n = 2 => CTPT của X là C~2~H~4~O~2~ (acetic acid CH~3~COOH).\n"
            "- Tổng số nguyên tử trong phân tử C~2~H~4~O~2~ là: 2 (C) + 4 (H) + 2 (O) = 8 nguyên tử."
        )

    # Câu 5
    p = add_p(doc)
    render_question_prompt(p, "Câu 5. Nước ép chanh chứa khoảng 5% (khoảng 0,3 mol/lít) citric acid, điều này giúp chanh có vị chua và độ pH từ 2 – 3. Citric acid có công thức cấu tạo là:")
    add_centered_picture(doc, os.path.join(IMG_DIR, "citric_acid.png"), width_cm=7.5)
    p2 = add_p(doc)
    render_rich_text(p2, "Phân tử khối của citric acid là bao nhiêu amu?")
    if is_teacher:
        p_ans = add_p(doc)
        render_rich_text(p_ans, "Đáp án: *192*")
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Từ công thức cấu tạo: HOOC-CH~2~-C(OH)(COOH)-CH~2~-COOH, ta xác định công thức phân tử của citric acid là C~6~H~8~O~7~.\n"
            "Phân tử khối: M = 12 . 6 + 1 . 8 + 16 . 7 = 72 + 8 + 112 = 192 amu."
        )

    # Câu 6
    p = add_p(doc)
    render_question_prompt(p, "Câu 6. Thực hiện phản ứng oxi hóa 4,958 L C~2~H~4~ (đkc) bằng O~2~ (xúc tác PdCl~2~, CuCl~2~) thu được chất X đơn chức. Cho toàn bộ lượng chất X tác dụng với hydrogen cyanide (HCN) dư, thu được 7,1 gam CH~3~CH(CN)OH (cyanohydrin). Hiệu suất của quá trình tạo CH~3~CH(CN)OH từ C~2~H~4~ là x%. Giá trị của x là bao nhiêu?")
    if is_teacher:
        p_ans = add_p(doc)
        render_rich_text(p_ans, "Đáp án: *50*")
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- Số mol C~2~H~4~: n~C2H4~ = 4,958 / 24,79 = 0,2 mol.\n"
            "- Khối lượng mol của CH~3~CH(CN)OH: M = 15 + 13 + 26 + 17 = 71 g/mol.\n"
            "- Số mol CH~3~CH(CN)OH thu được: n~cyanohydrin~ = 7,1 / 71 = 0,1 mol.\n"
            "- Sơ đồ phản ứng:\n"
            "    2C~2~H~4~ + O~2~ --(PdCl~2~, CuCl~2~)--> 2CH~3~CHO\n"
            "    CH~3~CHO + HCN -> CH~3~CH(CN)OH\n"
            "  Theo sơ đồ: 1 mol C~2~H~4~ -> 1 mol CH~3~CHO -> 1 mol cyanohydrin.\n"
            "=> n~cyanohydrin lí thuyết~ = n~C2H4~ = 0,2 mol.\n"
            "- Hiệu suất quá trình: x% = (0,1 / 0,2) . 100% = 50% => x = 50."
        )

    # =========================================================================
    # PHẦN IV
    # =========================================================================
    add_section_header(
        doc,
        title="PHẦN IV. Tự luận.",
        subtitle="Thí sinh trình bày bài giải chi tiết vào giấy làm bài."
    )

    # Câu 1
    p = add_p(doc)
    render_question_prompt(p, "Câu 1 (1,0 điểm). Hoàn thành dãy chuyển hóa sau bằng các phương trình hóa học:")
    p_seq = add_p(doc)
    render_rich_text(p_seq, "    Ethane --(1)--> Ethyl chloride --(2)--> Ethanol --(3)--> Ethanal --(4)--> Acetic acid")
    
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "(1) CH~3~-CH~3~ + Cl~2~ --as (1:1)--> CH~3~-CH~2~Cl + HCl\n"
            "(2) CH~3~-CH~2~Cl + NaOH --t^o^--> CH~3~-CH~2~OH + NaCl\n"
            "(3) CH~3~-CH~2~OH + CuO --t^o^--> CH~3~-CHO + Cu v + H~2~O\n"
            "(4) CH~3~-CHO + Br~2~ + H~2~O -> CH~3~-COOH + 2HBr\n"
            "    (hoặc: 2CH~3~CHO + O~2~ --xt, t^o^--> 2CH~3~COOH)"
        )
    else:
        add_student_writing_lines(doc, num_lines=8)

    # Câu 2
    p = add_p(doc)
    render_question_prompt(p, "Câu 2 (1,0 điểm). Để xác định được hợp chất hữu cơ X, người ta thực hiện các phân tích sau:\n"
                              "– Bằng phương pháp phân tích nguyên tố, xác định được X chứa %C = 40,0%; %H = 6,67%, còn lại là oxygen.\n"
                              "– Bằng phương pháp đo phổ MS, kết quả cho thấy trên phổ xuất hiện peak ion phân tử [M^+^] có giá trị m/z bằng 60.\n"
                              "– Bằng phương pháp đo phổ IR, kết quả cho thấy trên phổ xuất hiện peak rộng trong vùng từ 2 500 đến 3 200 cm^-1^, và một peak hấp thụ mạnh ở 1 745 cm^-1^.\n"
                              "Xác định công thức phân tử và công thức cấu tạo thu gọn của hợp chất X.")
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "1. Xác định công thức phân tử của X:\n"
            "- Phần trăm khối lượng của oxygen: %O = 100% - 40,0% - 6,67% = 53,33%.\n"
            "- Phổ MS có peak ion phân tử [M^+^] có m/z = 60 => Khối lượng mol phân tử M~X~ = 60 g/mol.\n"
            "- Đặt công thức phân tử của X là C~x~H~y~O~z~. Ta có:\n"
            "    x = (40,0 . 60) / (12 . 100) = 2\n"
            "    y = (6,67 . 60) / (1 . 100) = 4\n"
            "    z = (53,33 . 60) / (16 . 100) = 2\n"
            "=> Công thức phân tử của X là C~2~H~4~O~2~.\n\n"
            "2. Xác định công thức cấu tạo thu gọn của X:\n"
            "- Phổ IR xuất hiện:\n"
            "  + Peak rộng trong vùng 2 500 – 3 200 cm^-1^: đặc trưng cho liên kết O–H của nhóm carboxylic acid (-COOH).\n"
            "  + Peak hấp thụ mạnh ở 1 745 cm^-1^: đặc trưng cho liên kết C=O trong nhóm carboxyl (-COOH).\n"
            "- Kết hợp CTPT C~2~H~4~O~2~ chứa nhóm chức -COOH => X là carboxylic acid no, đơn chức có 2 nguyên tử C.\n"
            "=> Công thức cấu tạo thu gọn của X là CH~3~COOH (acetic acid)."
        )
    else:
        add_student_writing_lines(doc, num_lines=12)

    # Dòng kết thúc
    p_end = add_p(doc, space_before=10, space_after=10)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_end = p_end.add_run("---------- HẾT ----------")
    run_end.bold = True

    # KIỂM TRA ĐIỀU KIỆN 0 BẢNG TRONG THÂN BÀI
    # Bảng duy nhất được phép tồn tại là Header Block ở đầu đề thi (1 bảng)
    assert len(doc.tables) <= 1, f"LỖI: Phát hiện {len(doc.tables)} bảng trong tài liệu! Quy chuẩn Form Đề Kiểm Tra yêu cầu tuyệt đối 0 table trong nội dung câu hỏi/đáp án."

    return doc

def main():
    os.makedirs(OUT_DIR, exist_ok=True)

    # 1. Sinh bản Học sinh
    print("Đang tạo bản Học sinh (De_Kiem_Tra_Hoa_11_De18.docx)...")
    doc_hs = build_exam_document(is_teacher=False)
    hs_path = os.path.join(OUT_DIR, "De_Kiem_Tra_Hoa_11_De18.docx")
    doc_hs.save(hs_path)
    print(f" -> Đã lưu: {hs_path}")

    # 2. Sinh bản Giáo viên
    print("Đang tạo bản Giáo viên (De_Kiem_Tra_Hoa_11_De18_GiaoVien.docx)...")
    doc_gv = build_exam_document(is_teacher=True)
    gv_path = os.path.join(OUT_DIR, "De_Kiem_Tra_Hoa_11_De18_GiaoVien.docx")
    doc_gv.save(gv_path)
    print(f" -> Đã lưu: {gv_path}")

if __name__ == "__main__":
    main()
