# -*- coding: utf-8 -*-
r"""
Script biên soạn trọn bộ tài liệu chuẩn FORM ĐỀ KIỂM TRA cho:
ĐỀ CƯƠNG ÔN TẬP GIỮA HỌC KÌ I – MÔN KHTN 7
TRƯỜNG THCS THỐNG NHẤT – NĂM HỌC 2026 - 2027
1. De_Cuong_GK1_KHTN7_ThongNhat_HocSinh.docx
2. De_Cuong_GK1_KHTN7_ThongNhat_GiaoVien.docx
3. De_Cuong_GK1_KHTN7_ThongNhat_DayTrenManHinh.pdf (Slide 16:9 Widescreen KHÔNG CÓ ĐÁP ÁN)
Và bản Word: De_Cuong_GK1_KHTN7_ThongNhat_DayTrenManHinh.docx
Thầy TRẦN VĂN THẠNH - 0777.470.803
"""

import os
import sys
import subprocess
import shutil

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

import docx
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# Import các hàm chuẩn từ form_de_kiem_tra
sys.path.insert(0, r"c:\Antigravity_Thanh\scripts")
from form_de_kiem_tra import (
    setup_header_footer,
    build_title_block,
    add_section_header,
    render_question_prompt,
    render_options_no_table,
    render_statement,
    add_solution_no_table,
    add_student_writing_lines,
    render_rich_text,
    set_run_style,
    COLOR_PRIMARY,
    COLOR_CORRECT,
    COLOR_SOLUTION,
    COLOR_HF_RED,
    COLOR_HF_GREEN_HEX
)

# Import toàn bộ dữ liệu 3 phân môn
from khtn7_thongnhat_data_full import (
    VATLI_THEORY, VATLI_MCQ, VATLI_TF, VATLI_SHORT,
    HOAHOC_THEORY, HOAHOC_MCQ, HOAHOC_TF, HOAHOC_SHORT, HOAHOC_ESSAY,
    SINHHOC_THEORY, SINHHOC_MCQ, SINHHOC_TF_C21, SINHHOC_ESSAY
)

OUTPUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
SCRATCH_DIR = r"c:\Antigravity_Thanh\scratch"
IMAGES_DIR = r"c:\Antigravity_Thanh\San_Pham\images_khtn7_thongnhat"
os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(SCRATCH_DIR, exist_ok=True)
os.makedirs(IMAGES_DIR, exist_ok=True)

TARGET_ONEDRIVE_DIRS = [
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn7\KHTN7-IN-DEKIEMTRADINHKI\New folder (3)",
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn7\KHTN7-IN-DEKIEMTRADINHKI",
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn7"
]

# ==============================================================================
# 1. HÀM TẠO FILE WORD
# ==============================================================================

def generate_word_document(mode="student"):
    is_teacher = (mode == "teacher")
    is_screen = (mode == "screen")

    if mode == "teacher":
        filename = "De_Cuong_GK1_KHTN7_ThongNhat_GiaoVien.docx"
    elif mode == "screen":
        filename = "De_Cuong_GK1_KHTN7_ThongNhat_DayTrenManHinh.docx"
    else:
        filename = "De_Cuong_GK1_KHTN7_ThongNhat_HocSinh.docx"

    filepath = os.path.join(OUTPUT_DIR, filename)

    doc = docx.Document()
    setup_header_footer(doc)

    build_title_block(
        doc,
        school_line="TRƯỜNG THCS THỐNG NHẤT – TỔ KHOA HỌC TỰ NHIÊN",
        exam_title="ĐỀ CƯƠNG ÔN TẬP GIỮA HỌC KÌ I",
        subject_line="MÔN: KHOA HỌC TỰ NHIÊN 7 (TRỌN BỘ 3 PHÂN MÔN)",
        num_pages="12",
        duration="90 phút",
        exam_code="KHTN7-TN",
        atomic_masses=None,
        is_teacher=is_teacher
    )

    base_size = 11.5 if is_screen else 10.5

    # --------------------------------------------------------------------------
    # CHỦ ĐỀ 1: NĂNG LƯỢNG VÀ SỰ BIẾN ĐỔI (VẬT LÍ)
    # --------------------------------------------------------------------------
    add_section_header(doc, "CHỦ ĐỀ: NĂNG LƯỢNG VÀ SỰ BIẾN ĐỔI (VẬT LÍ)", "A. PHẦN LÝ THUYẾT:")
    for title, points in VATLI_THEORY:
        p_th = doc.add_paragraph()
        p_th.paragraph_format.left_indent = Cm(0.3)
        p_th.paragraph_format.space_before = Pt(3)
        p_th.paragraph_format.space_after = Pt(1)
        render_rich_text(p_th, f"*{title}*", base_size=base_size, base_color=COLOR_PRIMARY)
        for pt in points:
            p_pt = doc.add_paragraph()
            p_pt.paragraph_format.left_indent = Cm(0.6)
            p_pt.paragraph_format.space_before = Pt(1)
            p_pt.paragraph_format.space_after = Pt(1)
            render_rich_text(p_pt, f"• {pt}", base_size=base_size)

    # I. Trắc nghiệm nhiều lựa chọn Vật lí
    add_section_header(doc, "B. PHẦN BÀI TẬP THAM KHẢO (VẬT LÍ)\nI. TRẮC NGHIỆM NHIỀU LỰA CHỌN (12 CÂU)", "Chọn phương án trả lời đúng nhất trong mỗi câu sau:")
    for q in VATLI_MCQ:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, f"Câu {q['id']}. {q['content']}", base_size=base_size)
        render_options_no_table(doc, q["options"], is_teacher=is_teacher)
        if is_teacher and q.get("solution"):
            add_solution_no_table(doc, f"Đáp án đúng: *{q['answer']}*.\n{q['solution']}")
        elif is_screen:
            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.left_indent = Cm(0.6)
            p_sp.paragraph_format.space_before = Pt(2)
            p_sp.paragraph_format.space_after = Pt(2)
            render_rich_text(p_sp, "✎ Ghi chú / Giải thích: ........................................................................................................", base_size=base_size-1)

    # II. Trắc nghiệm Đúng - Sai Vật lí
    add_section_header(doc, "II. TRẮC NGHIỆM ĐÚNG – SAI (11 CÂU)", "Trong mỗi ý ở mỗi câu, học sinh chọn Đúng hoặc Sai:")
    for q in VATLI_TF:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, q["prompt"], base_size=base_size)
        for item in q["items"]:
            p_it = doc.add_paragraph()
            p_it.paragraph_format.left_indent = Cm(0.5)
            p_it.paragraph_format.space_before = Pt(1)
            p_it.paragraph_format.space_after = Pt(1)
            ans_str = "Đúng" if item["ans"] else "Sai"
            if is_teacher:
                render_statement(p_it, f"{item['id']}) {item['text']}", ans=ans_str, is_teacher=True, base_size=base_size)
                p_rea = doc.add_paragraph()
                p_rea.paragraph_format.left_indent = Cm(1.0)
                p_rea.paragraph_format.space_before = Pt(0)
                p_rea.paragraph_format.space_after = Pt(2)
                render_rich_text(p_rea, f"↳ _Giải thích:_ {item['reason']}", base_size=base_size-1, base_italic=True, base_color=COLOR_SOLUTION)
            else:
                render_rich_text(p_it, f"{item['id']}) {item['text']}   [   ] Đúng     [   ] Sai", base_size=base_size)

    # III. Trả lời ngắn Vật lí
    add_section_header(doc, "III. TRẮC NGHIỆM TRẢ LỜI NGẮN (7 CÂU)", "Học sinh tính toán và viết câu trả lời ngắn gọn vào chỗ trống:")
    for q in VATLI_SHORT:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, f"Câu {q['id']}. {q['content']}", base_size=base_size)
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.left_indent = Cm(0.8)
        p_ans.paragraph_format.space_before = Pt(1)
        p_ans.paragraph_format.space_after = Pt(2)
        if is_teacher:
            render_rich_text(p_ans, f"➔ *Đáp án:* {q['answer']}", base_size=base_size, base_color=COLOR_CORRECT)
            add_solution_no_table(doc, q["solution"])
        elif is_screen:
            render_rich_text(p_ans, "➔ Trả lời: ...................................................................................................................", base_size=base_size)
        else:
            render_rich_text(p_ans, "Trả lời: .....................................................................................................................", base_size=base_size)

    # --------------------------------------------------------------------------
    # CHỦ ĐỀ 2: CHẤT VÀ SỰ BIẾN ĐỔI CHẤT (HÓA HỌC)
    # --------------------------------------------------------------------------
    add_section_header(doc, "CHỦ ĐỀ: CHẤT VÀ SỰ BIẾN ĐỔI CHẤT (HÓA HỌC)", "A. PHẦN LÝ THUYẾT:")
    for title, points in HOAHOC_THEORY:
        p_th = doc.add_paragraph()
        p_th.paragraph_format.left_indent = Cm(0.3)
        p_th.paragraph_format.space_before = Pt(3)
        p_th.paragraph_format.space_after = Pt(1)
        render_rich_text(p_th, f"*{title}*", base_size=base_size, base_color=COLOR_PRIMARY)
        for pt in points:
            p_pt = doc.add_paragraph()
            p_pt.paragraph_format.left_indent = Cm(0.6)
            p_pt.paragraph_format.space_before = Pt(1)
            p_pt.paragraph_format.space_after = Pt(1)
            render_rich_text(p_pt, f"• {pt}", base_size=base_size)

    # I. Trắc nghiệm khách quan Hóa học (Nhiều lựa chọn)
    add_section_header(doc, "B. PHẦN BÀI TẬP (HÓA HỌC)\nI. TRẮC NGHIỆM KHÁCH QUAN\n1. Câu hỏi trắc nghiệm nhiều lựa chọn (10 câu)", "Chọn phương án trả lời đúng nhất:")
    for q in HOAHOC_MCQ:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, f"Câu {q['id']}. {q['content']}", base_size=base_size)
        render_options_no_table(doc, q["options"], is_teacher=is_teacher)
        if is_teacher and q.get("solution"):
            add_solution_no_table(doc, f"Đáp án đúng: *{q['answer']}*.\n{q['solution']}")
        elif is_screen:
            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.left_indent = Cm(0.6)
            p_sp.paragraph_format.space_before = Pt(2)
            p_sp.paragraph_format.space_after = Pt(2)
            render_rich_text(p_sp, "✎ Ghi chú / Giải thích: ........................................................................................................", base_size=base_size-1)

    # 2. Trắc nghiệm Đúng/Sai Hóa học
    add_section_header(doc, "2. Trắc nghiệm đúng/sai (Trong mỗi ý a, b, c, d ở mỗi câu, chọn đúng hoặc sai):", "")
    for q in HOAHOC_TF:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, q["prompt"], base_size=base_size)

        if q.get("image"):
            img_path = os.path.join(IMAGES_DIR, q["image"])
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(3)
                p_img.paragraph_format.space_after = Pt(3)
                p_img.add_run().add_picture(img_path, width=Inches(4.2))

        for item in q["items"]:
            p_it = doc.add_paragraph()
            p_it.paragraph_format.left_indent = Cm(0.5)
            p_it.paragraph_format.space_before = Pt(1)
            p_it.paragraph_format.space_after = Pt(1)
            ans_str = "Đúng" if item["ans"] else "Sai"
            if is_teacher:
                render_statement(p_it, f"{item['id']}) {item['text']}", ans=ans_str, is_teacher=True, base_size=base_size)
                p_rea = doc.add_paragraph()
                p_rea.paragraph_format.left_indent = Cm(1.0)
                p_rea.paragraph_format.space_before = Pt(0)
                p_rea.paragraph_format.space_after = Pt(2)
                render_rich_text(p_rea, f"↳ _Giải thích:_ {item['reason']}", base_size=base_size-1, base_italic=True, base_color=COLOR_SOLUTION)
            else:
                render_rich_text(p_it, f"{item['id']}) {item['text']}   [   ] Đúng     [   ] Sai", base_size=base_size)

    # 3. Câu hỏi trả lời ngắn Hóa học
    add_section_header(doc, "3. Câu hỏi trả lời ngắn (7 câu)", "Học sinh suy luận và viết câu trả lời ngắn gọn vào chỗ trống:")
    for q in HOAHOC_SHORT:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, f"Câu {q['id']}. {q['content']}", base_size=base_size)

        if q.get("image"):
            img_path = os.path.join(IMAGES_DIR, q["image"])
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(3)
                p_img.paragraph_format.space_after = Pt(3)
                p_img.add_run().add_picture(img_path, width=Inches(2.5))

        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.left_indent = Cm(0.8)
        p_ans.paragraph_format.space_before = Pt(1)
        p_ans.paragraph_format.space_after = Pt(2)
        if is_teacher:
            render_rich_text(p_ans, f"➔ *Đáp án:* {q['answer']}", base_size=base_size, base_color=COLOR_CORRECT)
            add_solution_no_table(doc, q["solution"])
        elif is_screen:
            render_rich_text(p_ans, "➔ Trả lời: ...................................................................................................................", base_size=base_size)
        else:
            render_rich_text(p_ans, "Trả lời: .....................................................................................................................", base_size=base_size)

    # II. Tự luận Hóa học
    add_section_header(doc, "II. PHẦN TỰ LUẬN (HÓA HỌC - 7 CÂU)", "Học sinh trình bày chi tiết câu trả lời vào bài làm:")
    for q in HOAHOC_ESSAY:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, f"Câu {q['id']}. {q['content']}", base_size=base_size)
        if is_teacher:
            add_solution_no_table(doc, q["solution"])
        elif is_screen:
            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.left_indent = Cm(0.6)
            render_rich_text(p_sp, "✎ Bài giải chi tiết: ........................................................................................................", base_size=base_size)
        else:
            add_student_writing_lines(doc, num_lines=5)

    # --------------------------------------------------------------------------
    # CHỦ ĐỀ 3: VẬT SỐNG (SINH HỌC)
    # --------------------------------------------------------------------------
    add_section_header(doc, "CHỦ ĐỀ: VẬT SỐNG (SINH HỌC)", "I. LÝ THUYẾT TRỌNG TÂM:")
    for title, points in SINHHOC_THEORY:
        p_th = doc.add_paragraph()
        p_th.paragraph_format.left_indent = Cm(0.3)
        p_th.paragraph_format.space_before = Pt(3)
        p_th.paragraph_format.space_after = Pt(1)
        render_rich_text(p_th, f"*{title}*", base_size=base_size, base_color=COLOR_PRIMARY)
        for pt in points:
            p_pt = doc.add_paragraph()
            p_pt.paragraph_format.left_indent = Cm(0.6)
            p_pt.paragraph_format.space_before = Pt(1)
            p_pt.paragraph_format.space_after = Pt(1)
            render_rich_text(p_pt, f"• {pt}", base_size=base_size)

    # II. Trắc nghiệm Sinh học
    add_section_header(doc, "II. BÀI TẬP THAM KHẢO (SINH HỌC)\nA. TRẮC NGHIỆM (20 CÂU)", "Chọn phương án trả lời đúng nhất:")
    for q in SINHHOC_MCQ:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, f"Câu {q['id']}. {q['content']}", base_size=base_size)
        render_options_no_table(doc, q["options"], is_teacher=is_teacher)
        if is_teacher and q.get("solution"):
            add_solution_no_table(doc, f"Đáp án đúng: *{q['answer']}*.\n{q['solution']}")
        elif is_screen:
            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.left_indent = Cm(0.6)
            p_sp.paragraph_format.space_before = Pt(2)
            p_sp.paragraph_format.space_after = Pt(2)
            render_rich_text(p_sp, "✎ Ghi chú / Giải thích: ........................................................................................................", base_size=base_size-1)

    # Trắc nghiệm Đúng/Sai Sinh học Câu 21
    add_section_header(doc, "BẢNG ĐÁNH GIÁ ĐÚNG / SAI QUÁ TRÌNH QUANG HỢP (CÂU 21)", "")
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(4)
    p_q.paragraph_format.space_after = Pt(2)
    render_question_prompt(p_q, SINHHOC_TF_C21["prompt"], base_size=base_size)
    for item in SINHHOC_TF_C21["items"]:
        p_it = doc.add_paragraph()
        p_it.paragraph_format.left_indent = Cm(0.5)
        p_it.paragraph_format.space_before = Pt(1)
        p_it.paragraph_format.space_after = Pt(1)
        ans_str = "Đúng" if item["ans"] else "Sai"
        if is_teacher:
            render_statement(p_it, f"{item['id']}) {item['text']}", ans=ans_str, is_teacher=True, base_size=base_size)
            p_rea = doc.add_paragraph()
            p_rea.paragraph_format.left_indent = Cm(1.0)
            p_rea.paragraph_format.space_before = Pt(0)
            p_rea.paragraph_format.space_after = Pt(2)
            render_rich_text(p_rea, f"↳ _Giải thích:_ {item['reason']}", base_size=base_size-1, base_italic=True, base_color=COLOR_SOLUTION)
        else:
            render_rich_text(p_it, f"{item['id']}) {item['text']}   [   ] Đúng     [   ] Sai", base_size=base_size)

    # B. Tự luận Sinh học
    add_section_header(doc, "B. PHẦN TỰ LUẬN (SINH HỌC - 7 CÂU)", "Học sinh trình bày chi tiết câu trả lời vào bài làm:")
    for q in SINHHOC_ESSAY:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, f"Câu {q['id']}. {q['content']}", base_size=base_size)
        if is_teacher:
            add_solution_no_table(doc, q["solution"])
        elif is_screen:
            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.left_indent = Cm(0.6)
            render_rich_text(p_sp, "✎ Bài giải chi tiết: ........................................................................................................", base_size=base_size)
        else:
            add_student_writing_lines(doc, num_lines=5)

    # Dòng kết thúc
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(14)
    p_end.paragraph_format.space_after = Pt(6)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("--- HẾT ĐỀ CƯƠNG ÔN TẬP GIỮA KÌ I KHTN 7 ---")
    set_run_style(r_end, font_size=11, bold=True, italic=True, color=COLOR_PRIMARY)

    doc.save(filepath)
    print(f"[OK] Đã tạo thành công file Word: {filepath} ({os.path.getsize(filepath):,} bytes)")
    return filepath


# ==============================================================================
# 2. HÀM TẠO FILE LATEX & PDF TRÌNH CHIẾU TIVI 16:9 (KHÔNG CÓ ĐÁP ÁN)
# ==============================================================================

def generate_tivi_presentation_pdf():
    tex_filename = "De_Cuong_GK1_KHTN7_ThongNhat_DayTrenManHinh.tex"
    pdf_filename = "De_Cuong_GK1_KHTN7_ThongNhat_DayTrenManHinh.pdf"
    tex_path = os.path.join(SCRATCH_DIR, tex_filename)
    pdf_dest = os.path.join(OUTPUT_DIR, pdf_filename)

    tex_lines = []
    # Preamble chuẩn 16:9 widescreen
    tex_lines.append(r"""\documentclass{article}
\usepackage[utf8]{vietnam}
\usepackage{times}
\usepackage{amsmath,amssymb}
\usepackage{tabularx}
\usepackage{array}
\usepackage{booktabs}
\usepackage{graphicx}
\usepackage{xcolor}
\usepackage[
  paperwidth=29.7cm,
  paperheight=16.706cm,
  top=0.9cm,
  bottom=0.9cm,
  left=1.2cm,
  right=1.2cm,
  footskip=0pt,
  headheight=0pt,
  headsep=0pt
]{geometry}

\pagestyle{empty}
\setlength{\parindent}{0pt}
\setlength{\parskip}{0pt}

\definecolor{brandRed}{RGB}{211, 47, 47}
\definecolor{brandGreen}{RGB}{46, 125, 50}
\definecolor{brandBlue}{RGB}{21, 101, 192}
\definecolor{textColor}{RGB}{20, 20, 20}

\color{textColor}

\begin{document}
""")

    # Slide 0: Tiêu đề
    tex_lines.append(r"""
% SLIDE 0: TIÊU ĐỀ
\begin{center}
\vspace*{0.2cm}
{\fontsize{24}{30}\selectfont \textbf{\color{brandRed}TRƯỜNG THCS THỐNG NHẤT -- NĂM HỌC 2026 -- 2027}}\\[10pt]
{\fontsize{34}{42}\selectfont \textbf{\color{brandBlue}ĐỀ CƯƠNG ÔN TẬP KIỂM TRA GIỮA HỌC KÌ I}}\\[8pt]
{\fontsize{28}{36}\selectfont \textbf{\color{brandGreen}MÔN: KHOA HỌC TỰ NHIÊN 7 (TRỌN BỘ 3 PHÂN MÔN)}}\\[8pt]
{\fontsize{22}{28}\selectfont \textit{(Vật lí: Năng lượng \& biến đổi -- Hóa học: Chất \& biến đổi chất -- Sinh học: Vật sống)}}\\[16pt]
\rule{16cm}{1.5pt}\\[20pt]
{\fontsize{26}{34}\selectfont \textbf{Giáo viên hướng dẫn: Thầy TRẦN VĂN THẠNH -- 0777.470.803}}\\[8pt]
{\fontsize{22}{28}\selectfont \textit{Cơ sở 1: 6/15 Nguyễn Hoàng, Kim Long | Cơ sở 2: 24 Đặng Thái Thân, TP Huế}}\\[16pt]
{\fontsize{24}{30}\selectfont \textbf{TỔNG HỢP TOÀN BỘ CÂU HỎI ÔN TẬP VÀ HƯỚNG DẪN DẠY MÀN HÌNH}}
\end{center}
\newpage
""")

    # --------------------------------------------------------------------------
    # CHỦ ĐỀ 1: VẬT LÍ (SLIDES)
    # --------------------------------------------------------------------------
    # Slide 1: Lý thuyết Vật lí
    tex_lines.append(r"""
% SLIDE 1: LÝ THUYẾT VẬT LÍ
{\fontsize{25}{32}\selectfont \textbf{\color{brandBlue}CHỦ ĐỀ 1: NĂNG LƯỢNG VÀ SỰ BIẾN ĐỔI (VẬT LÍ) -- LÝ THUYẾT}}
\vspace{8pt}

\fontsize{21}{28}\selectfont
\textbf{1. Cổng quang điện và đồng hồ đo thời gian hiện số:}
\begin{itemize}
  \setlength{\itemsep}{2pt}
  \item Cổng quang điện đóng vai trò như công tắc tự động điều khiển đóng/mở đồng hồ hiện số.
  \item Gồm đầu phát và đầu thu tia hồng ngoại. Khi có vật chắn ngang, tín hiệu điện truyền qua dây cáp tới đồng hồ.
\end{itemize}

\vspace{6pt}
\textbf{2. Tốc độ chuyển động:}
\begin{itemize}
  \setlength{\itemsep}{2pt}
  \item Đại lượng đặc trưng cho mức độ nhanh chậm: $v = \dfrac{s}{t}$ ($s$: quãng đường; $t$: thời gian).
  \item Đơn vị hợp pháp trong hệ SI: $\text{m/s}$. Đổi đơn vị: $1\text{ m/s} = 3,6\text{ km/h}$.
\end{itemize}
\newpage
""")

    # Slide 2: Vật lí MCQ 1 & 2
    tex_lines.append(r"""
% SLIDE 2: VẬT LÍ MCQ 1 & 2
\fontsize{26}{35}\selectfont
\textbf{Câu 1.} Trong bộ thiết bị thí nghiệm, cổng quang điện đóng vai trò chính là gì?
\vspace{4pt}

\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{A.} Dụng cụ đo trực tiếp khối lượng của vật. & \textbf{B.} Công tắc điều khiển mở/đóng đồng hồ hiện số.\\[6pt]
\textbf{C.} Nguồn điện cung cấp cho hệ thống. & \textbf{D.} Dụng cụ đo trực tiếp kích thước vật rắn.
\end{tabularx}

\vspace{18pt}

\textbf{Câu 2.} Cổng quang điện bao gồm những bộ phận chính nào?
\vspace{4pt}

\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{A.} Đầu phát tia hồng ngoại D1, đầu thu D2 và dây nối. & \textbf{B.} Màn hình tinh thể lỏng và nút điều chỉnh.\\[6pt]
\textbf{C.} Hộp pin mặt trời và công tắc nguồn điện. & \textbf{D.} Cảm ứng từ và chốt cắm tín hiệu số.
\end{tabularx}
\newpage
""")

    # Slide 3: Vật lí MCQ 3 & 4
    tex_lines.append(r"""
% SLIDE 3: VẬT LÍ MCQ 3 & 4
\fontsize{25}{34}\selectfont
\textbf{Câu 3.} Dây cáp nối giữa cổng quang điện và đồng hồ hiện số có tác dụng nào sau đây?
\vspace{4pt}

\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{A.} Chỉ dùng truyền tín hiệu ánh sáng. & \textbf{B.} Chỉ cấp dòng điện một chiều cho cổng quang.\\[6pt]
\textbf{C.} Cung cấp điện và gửi tín hiệu điện tới đồng hồ. & \textbf{D.} Giữ cố định cổng quang trên giá đỡ quang học.
\end{tabularx}

\vspace{16pt}

\textbf{Câu 4.} Trên mặt trước của đồng hồ đo thời gian hiện số, nút MODE dùng để làm gì?
\vspace{4pt}

\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{A.} Đưa số hiển thị về trạng thái 0.000. & \textbf{B.} Chọn chế độ làm việc cho đồng hồ.\\[6pt]
\textbf{C.} Chọn giới hạn đo và độ chia nhỏ nhất. & \textbf{D.} Bật hoặc tắt nguồn điện chính.
\end{tabularx}
\newpage
""")

    # Slide 4: Vật lí MCQ 5 & 6
    tex_lines.append(r"""
% SLIDE 4: VẬT LÍ MCQ 5 & 6
\fontsize{25}{34}\selectfont
\textbf{Câu 5.} Nút RESET trên mặt trước của đồng hồ đo thời gian hiện số có tác dụng gì?
\vspace{4pt}

\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{A.} Thay đổi chế độ đo lường thời gian. & \textbf{B.} Khởi động lại hệ thống nguồn điện.\\[6pt]
\textbf{C.} Cho đồng hồ quay về chỉ số 0.000 ban đầu. & \textbf{D.} Lưu kết quả đo cuối cùng vào bộ nhớ.
\end{tabularx}

\vspace{16pt}

\textbf{Câu 6.} Những bộ phận nào sau đây nằm ở mặt sau của đồng hồ đo thời gian hiện số?
\vspace{4pt}

\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{A.} Nút thang đo, nút mode và màn hình số. & \textbf{B.} Công tắc điện, 3 ổ cắm cổng quang A, B, C và ổ điện.\\[6pt]
\textbf{C.} Chỉ có duy nhất ổ cắm cổng quang A. & \textbf{D.} Đèn báo tín hiệu hồng ngoại.
\end{tabularx}
\newpage
""")

    # Slide 5: Vật lí MCQ 7 & 8
    tex_lines.append(r"""
% SLIDE 5: VẬT LÍ MCQ 7 & 8
\fontsize{24}{33}\selectfont
\textbf{Câu 7.} Tại sao dùng cổng quang điện kết hợp đồng hồ hiện số lại chính xác hơn đồng hồ bấm giây?
\vspace{3pt}

\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{A.} Tự động ghi nhận tín hiệu, loại bỏ độ trễ do phản xạ tay người. & \textbf{B.} Sử dụng nguồn điện lưới quốc gia có điện áp lớn hơn.\\[4pt]
\textbf{C.} Cổng quang phát tia laser đẩy vật chuyển động nhanh hơn. & \textbf{D.} Thiết bị tự động chia vận tốc không cần tính toán.
\end{tabularx}

\vspace{14pt}

\textbf{Câu 8.} Khi ước lượng thời gian chuyển động của vật lớn hơn $10\text{ s}$, cần chọn thang đo nào của đồng hồ hiện số?
\vspace{3pt}

\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{A.} $9,999\text{ s} - 0,0001\text{ s}$ & \textbf{B.} $99\text{ s} - 1\text{ s}$\\[4pt]
\textbf{C.} $10\text{ s} - 9\text{ s}$ & \textbf{D.} $99,99\text{ s} - 0,01\text{ s}$
\end{tabularx}
\newpage
""")

    # Slide 6: Vật lí MCQ 9, 10, 11, 12
    tex_lines.append(r"""
% SLIDE 6: VẬT LÍ MCQ 9, 10, 11, 12
\fontsize{23}{31}\selectfont
\textbf{Câu 9.} Dụng cụ phù hợp nhất để đo thời gian chạy của vận động viên $100\text{ m}$ là:
\begin{tabularx}{\linewidth}{@{}X X X X@{}}
\textbf{A.} Cổng quang. & \textbf{B.} Đồng hồ bấm giây. & \textbf{C.} Đồng hồ cát. & \textbf{D.} Đồng hồ treo tường.
\end{tabularx}

\vspace{8pt}
\textbf{Câu 10.} Đại lượng nào sau đây cho biết mức độ nhanh hay chậm của chuyển động?
\begin{tabularx}{\linewidth}{@{}X X X X@{}}
\textbf{A.} Quãng đường. & \textbf{B.} Khối lượng. & \textbf{C.} Tốc độ. & \textbf{D.} Thời gian.
\end{tabularx}

\vspace{8pt}
\textbf{Câu 11.} Công thức nào sau đây dùng để tính tốc độ chuyển động?
\begin{tabularx}{\linewidth}{@{}X X X X@{}}
\textbf{A.} $v = s / t$. & \textbf{B.} $v = s \cdot t$. & \textbf{C.} $s = v / t$. & \textbf{D.} $t = v \cdot t$.
\end{tabularx}

\vspace{8pt}
\textbf{Câu 12.} Trong các đơn vị sau, đơn vị nào là đơn vị đo tốc độ?
\begin{tabularx}{\linewidth}{@{}X X X X@{}}
\textbf{A.} $\text{km/h}$. & \textbf{B.} $\text{m}\cdot\text{s}$. & \textbf{C.} $\text{km}\cdot\text{h}$. & \textbf{D.} $\text{s/m}$.
\end{tabularx}
\newpage
""")

    # Slide 7: Vật lí Trả lời ngắn 1, 2, 3, 4
    tex_lines.append(r"""
% SLIDE 7: VẬT LÍ TRẢ LỜI NGẮN 1, 2, 3, 4
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}III. TRẮC NGHIỆM TRẢ LỜI NGẮN (VẬT LÍ -- CÂU 1 ĐẾN 4)}}
\vspace{10pt}

\fontsize{22}{30}\selectfont
\textbf{Câu 1.} Đoàn tàu hỏa đi từ A đến B cách nhau $30\text{ km}$ trong $30\text{ phút}$. Tốc độ của đoàn tàu là: $\dots\dots\dots\dots\dots\text{km/h}$.

\vspace{14pt}
\textbf{Câu 2.} Người đi bộ với tốc độ $4\text{ km/h}$ trong $30\text{ phút}$. Quãng đường người đó đi được là: $\dots\dots\dots\dots\dots\text{km}$.

\vspace{14pt}
\textbf{Câu 3.} Người đi xe đạp với tốc độ $18\text{ km/h}$ trong $20\text{ phút}$. Quãng đường đi được là: $\dots\dots\dots\dots\dots\text{km}$.

\vspace{14pt}
\textbf{Câu 4.} Máy bay bay với tốc độ $800\text{ km/h}$ trên quãng đường $1400\text{ km}$. Thời gian bay là: $\dots\dots\dots\dots\dots\text{phút}$.
\newpage
""")

    # Slide 8: Vật lí Trả lời ngắn 5, 6, 7
    tex_lines.append(r"""
% SLIDE 8: VẬT LÍ TRẢ LỜI NGẮN 5, 6, 7
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}III. TRẮC NGHIỆM TRẢ LỜI NGẮN (VẬT LÍ -- CÂU 5, 6, 7)}}
\vspace{8pt}

\fontsize{21}{29}\selectfont
\textbf{Câu 5.} Ô tô đi đoạn đầu với tốc độ $15\text{ m/s}$ trong $30\text{ phút}$, sau đó đi tiếp với tốc độ $60\text{ km/h}$ trong $45\text{ phút}$.\\
Tổng quãng đường ô tô đi được trong $75\text{ phút}$ là: $\dots\dots\dots\dots\dots\text{km}$.

\vspace{14pt}
\textbf{Câu 6.} Ô tô đi đoạn đầu với tốc độ $12\text{ m/s}$ trong $20\text{ phút}$, sau đó đi tiếp với tốc độ $54\text{ km/h}$ trong $40\text{ phút}$.\\
Tốc độ trung bình của ô tô trong suốt thời gian $60\text{ phút}$ là: $\dots\dots\dots\dots\dots\text{m/s}$.

\vspace{14pt}
\textbf{Câu 7.} Người đi bộ trong $10\text{ phút}$ với tốc độ $1,5\text{ m/s}$, sau đó chạy trong $5\text{ phút}$ với tốc độ $3\text{ m/s}$.\\
Tốc độ trung bình trong suốt thời gian chuyển động là: $\dots\dots\dots\dots\dots\text{km/h}$.
\newpage
""")

    # --------------------------------------------------------------------------
    # CHỦ ĐỀ 2: HÓA HỌC (SLIDES)
    # --------------------------------------------------------------------------
    # Slide 9: Lý thuyết Hóa học
    tex_lines.append(r"""
% SLIDE 9: LÝ THUYẾT HÓA HỌC
{\fontsize{25}{32}\selectfont \textbf{\color{brandBlue}CHỦ ĐỀ 2: CHẤT VÀ SỰ BIẾN ĐỔI CHẤT (HÓA HỌC) -- LÝ THUYẾT}}
\vspace{8pt}

\fontsize{21}{28}\selectfont
\textbf{1. Cấu tạo nguyên tử:}
\begin{itemize}
  \setlength{\itemsep}{2pt}
  \item Hạt nhân ở tâm: chứa hạt proton (mang điện tích $+1$) và neutron (không mang điện).
  \item Lớp vỏ: gồm các hạt electron (mang điện tích $-1$) chuyển động rất nhanh quanh hạt nhân theo từng lớp.
  \item Khối lượng nguyên tử: đo bằng $\text{amu}$, tập trung hầu hết ở hạt nhân ($m \approx p + n$).
\end{itemize}

\vspace{6pt}
\textbf{2. Nguyên tố hóa học:}
\begin{itemize}
  \setlength{\itemsep}{2pt}
  \item Tập hợp những nguyên tử cùng loại có cùng số proton trong hạt nhân.
  \item Kí hiệu hóa học gồm 1 hoặc 2 chữ cái: chữ đầu viết hoa, chữ sau viết thường.
\end{itemize}
\newpage
""")

    # Slide 10: Hóa học Đúng/Sai Câu 2 (Mô hình Boron, Sodium, Sulfur)
    tex_lines.append(r"""
% SLIDE 10: HÓA HỌC ĐÚNG/SAI CÂU 2
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}HÓA HỌC -- TRẮC NGHIỆM ĐÚNG SAI (MÔ HÌNH NGUYÊN TỬ)}}
\vspace{4pt}

\begin{minipage}[c]{0.52\linewidth}
\fontsize{20}{28}\selectfont
\textbf{Câu 2.} Cho mô hình nguyên tử như hình bên:
\begin{itemize}
  \setlength{\itemsep}{4pt}
  \item[a)] Nguyên tử boron có 5 proton trong hạt nhân.\hfill $[\quad]$
  \item[b)] Nguyên tử sodium có 3 lớp electron.\hfill $[\quad]$
  \item[c)] Nguyên tử sulfur có 4 electron lớp ngoài cùng.\hfill $[\quad]$
  \item[d)] Cả ba nguyên tử đều có số proton bằng electron.\hfill $[\quad]$
\end{itemize}
\end{minipage}
\hfill
\begin{minipage}[c]{0.46\linewidth}
\centering
\includegraphics[width=0.98\linewidth]{San_Pham/images_khtn7_thongnhat/mo_hinh_boron_sodium_sulfur.png}
\end{minipage}
\newpage
""")

    # Slide 11: Hóa học Trả lời ngắn Câu 5 (Mô hình Silicon)
    tex_lines.append(r"""
% SLIDE 11: HÓA HỌC TRẢ LỜI NGẮN CÂU 5 (SILICON)
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}HÓA HỌC -- TRẢ LỜI NGẮN: MÔ HÌNH NGUYÊN TỬ SILICON}}
\vspace{6pt}

\begin{minipage}[c]{0.60\linewidth}
\fontsize{21}{29}\selectfont
\textbf{Câu 5.} Quan sát mô hình nguyên tử silicon hình bên:
\begin{enumerate}
  \setlength{\itemsep}{6pt}
  \item[a)] Các electron phân bố thành $\dots\dots\dots$ lớp theo thứ tự từ trong ra ngoài là 2/8/4.
  \item[b)] Tổng số electron của nguyên tử silicon là $\dots\dots\dots$
  \item[c)] Silicon có số hiệu nguyên tử $Z = 14$. Tổng số hạt mang điện trong hạt nhân là $\dots\dots\dots$
\end{enumerate}

\vspace{12pt}
\textbf{Trả lời:} $\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots$
\end{minipage}
\hfill
\begin{minipage}[c]{0.38\linewidth}
\centering
\includegraphics[width=0.88\linewidth]{San_Pham/images_khtn7_thongnhat/mo_hinh_silicon.png}
\end{minipage}
\newpage
""")

    # --------------------------------------------------------------------------
    # CHỦ ĐỀ 3: SINH HỌC (SLIDES)
    # --------------------------------------------------------------------------
    # Slide 12: Lý thuyết Sinh học
    tex_lines.append(r"""
% SLIDE 12: LÝ THUYẾT SINH HỌC
{\fontsize{25}{32}\selectfont \textbf{\color{brandBlue}CHỦ ĐỀ 3: VẬT SỐNG (SINH HỌC) -- LÝ THUYẾT TRỌNG TÂM}}
\vspace{8pt}

\fontsize{21}{28}\selectfont
\textbf{1. Trao đổi chất và chuyển hóa năng lượng:}
\begin{itemize}
  \setlength{\itemsep}{2pt}
  \item Trao đổi chất: Lấy chất từ môi trường ngoài, biến đổi tạo chất cần thiết, thải chất cặn bã ra ngoài.
  \item Chuyển hóa năng lượng: Quang năng $\rightarrow$ Hóa năng $\rightarrow$ Nhiệt năng, cơ năng...
\end{itemize}

\vspace{6pt}
\textbf{2. Quang hợp & Hô hấp tế bào:}
\begin{itemize}
  \setlength{\itemsep}{2pt}
  \item Quang hợp: $\text{Nước} + \text{Carbon dioxide} \xrightarrow[\text{Diệp lục}]{\text{Ánh sáng}} \text{Glucose} + \text{Oxygen}$.
  \item Hô hấp tế bào: Phân giải chất hữu cơ giải phóng năng lượng ATP cung cấp cho hoạt động sống.
\end{itemize}
\newpage
""")

    # Slide 13: Sinh học Đúng/Sai Câu 21 (Phần 1: Các ý a -> e)
    tex_lines.append(r"""
% SLIDE 13: SINH HỌC ĐÚNG/SAI C21 (PHẦN 1)
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}SINH HỌC -- BẢNG ĐÚNG SAI VỀ QUANG HỢP (PHẦN 1)}}
\vspace{6pt}

\fontsize{20}{28}\selectfont
\textbf{Câu 21.} Xác định mỗi phát biểu sau về quang hợp là Đúng hay Sai:

\vspace{6pt}
\textbf{a)} Các yếu tố chủ yếu ảnh hưởng: ánh sáng, nước, oxygen và nhiệt độ.\hfill $[\quad]$\\[6pt]
\textbf{b)} Nhiệt độ môi trường quá cao hoặc quá thấp làm giảm sút quang hợp.\hfill $[\quad]$\\[6pt]
\textbf{c)} Năng lượng ánh sáng (quang năng) chuyển hóa thành hóa năng trong chất hữu cơ.\hfill $[\quad]$\\[6pt]
\textbf{d)} Vượt quá điểm bão hòa ánh sáng, quang hợp tiếp tục tăng mạnh không ngừng.\hfill $[\quad]$\\[6pt]
\textbf{e)} Cường độ ánh sáng càng tăng vô hạn thì tốc độ quang hợp tăng theo liên tục.\hfill $[\quad]$
\newpage
""")

    # Slide 14: Sinh học Đúng/Sai Câu 21 (Phần 2: Các ý f -> k)
    tex_lines.append(r"""
% SLIDE 14: SINH HỌC ĐÚNG/SAI C21 (PHẦN 2)
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}SINH HỌC -- BẢNG ĐÚNG SAI VỀ QUANG HỢP (PHẦN 2)}}
\vspace{6pt}

\fontsize{20}{28}\selectfont
\textbf{Câu 21.} Xác định mỗi phát biểu sau về quang hợp là Đúng hay Sai:

\vspace{6pt}
\textbf{f)} Nồng độ $\text{CO}_2$ trong không khí quá thấp làm giảm hiệu suất quang hợp.\hfill $[\quad]$\\[6pt]
\textbf{g)} Nước vừa là nguyên liệu vừa tham gia điều tiết đóng mở khí khổng.\hfill $[\quad]$\\[6pt]
\textbf{h)} Mỗi loài thực vật chỉ quang hợp tốt nhất ở nhiệt độ tối ưu nhất định.\hfill $[\quad]$\\[6pt]
\textbf{j)} Sản phẩm quang hợp chỉ phục vụ thực vật, không có vai trò với động vật.\hfill $[\quad]$\\[6pt]
\textbf{k)} Lá cây là cơ quan chủ yếu quang hợp nhờ chứa tế bào có nhiều lục lạp.\hfill $[\quad]$
\newpage
""")

    # Slide 15: Sinh học Đúng/Sai Câu 21 (Phần 3: Các ý l -> p)
    tex_lines.append(r"""
% SLIDE 15: SINH HỌC ĐÚNG/SAI C21 (PHẦN 3)
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}SINH HỌC -- BẢNG ĐÚNG SAI VỀ QUANG HỢP (PHẦN 3)}}
\vspace{6pt}

\fontsize{20}{28}\selectfont
\textbf{Câu 21.} Xác định mỗi phát biểu sau về quang hợp là Đúng hay Sai:

\vspace{8pt}
\textbf{l)} Khí khổng trên biểu bì lá giúp thu nhận khí oxygen làm nguyên liệu quang hợp.\hfill $[\quad]$\\[8pt]
\textbf{m)} Gân lá chứa mạch gỗ và mạch rây vận chuyển nước và chất hữu cơ.\hfill $[\quad]$\\[8pt]
\textbf{n)} Thiếu $\text{CO}_2$, quang hợp bị hạn chế dù đủ ánh sáng và nước.\hfill $[\quad]$\\[8pt]
\textbf{p)} Nhiệt độ môi trường càng tăng thì quang hợp luôn tăng, kể cả vượt ngưỡng thích hợp.\hfill $[\quad]$
\newpage
""")

    # Slide 16: Sinh học Tự luận Câu 1 & 2
    tex_lines.append(r"""
% SLIDE 16: SINH HỌC TỰ LUẬN CÂU 1 & 2
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}B. PHẦN TỰ LUẬN (SINH HỌC -- CÂU 1 \& CÂU 2)}}
\vspace{8pt}

\fontsize{21}{29}\selectfont
\textbf{Câu 1.} Vì sao khi bảo quản rau xanh, củ, quả tươi trong gia đình, chúng ta thường bọc trong túi nilon có đục lỗ nhỏ hoặc màng bọc rồi để vào ngăn mát tủ lạnh, chứ không để ở nhiệt độ phòng hoặc để trong ngăn đá?

\vspace{14pt}
\textbf{Câu 2.} Sau những trận mưa lớn gây ngập úng ruộng lúa, luống rau, người nông dân phải bơm tháo nước ra khỏi ruộng thật nhanh. Giải thích cơ sở khoa học dưới góc độ hô hấp tế bào?

\vspace*{2cm}
\begin{center}
{\fontsize{20}{26}\selectfont \textit{\color{gray}(Khoảng trống trình chiếu tương tác / Giáo viên và học sinh chữa bài trực tiếp trên bảng)}}
\end{center}
\newpage
""")

    # Slide 17: Sinh học Tự luận Câu 5 & 6
    tex_lines.append(r"""
% SLIDE 17: SINH HỌC TỰ LUẬN CÂU 5 & 6
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}B. PHẦN TỰ LUẬN (SINH HỌC -- CÂU 5 \& CÂU 6)}}
\vspace{8pt}

\fontsize{21}{29}\selectfont
\textbf{Câu 5.} Tại sao trước khi đem cất giữ (bảo quản) các loại hạt (thóc, ngô, đậu,...), người nông dân thường phải phơi khô? Giải thích cơ sở khoa học dựa trên hiểu biết về hô hấp tế bào?

\vspace{14pt}
\textbf{Câu 6.} Việc xây dựng các công viên cây xanh trong các khu đô thị, khu công nghiệp có ý nghĩa như thế nào đối với môi trường và đời sống con người?

\vspace*{2cm}
\begin{center}
{\fontsize{20}{26}\selectfont \textit{\color{gray}(Khoảng trống trình chiếu tương tác / Giáo viên và học sinh chữa bài trực tiếp trên bảng)}}
\end{center}

\end{document}
""")

    with open(tex_path, "w", encoding="utf-8") as f:
        f.write("".join(tex_lines))
    print(f"[OK] Đã ghi file LaTeX: {tex_path}")

    # Biên dịch pdflatex 2 lần để ổn định layout
    cmd = ["pdflatex", "-interaction=nonstopmode", "-output-directory", SCRATCH_DIR, tex_path]
    res1 = subprocess.run(cmd, capture_output=True, text=True)
    res2 = subprocess.run(cmd, capture_output=True, text=True)

    pdf_scratch = os.path.join(SCRATCH_DIR, pdf_filename)
    if os.path.exists(pdf_scratch):
        shutil.copy2(pdf_scratch, pdf_dest)
        print(f"[OK] Đã biên dịch và copy PDF thành công: {pdf_dest} ({os.path.getsize(pdf_dest):,} bytes)")
        return pdf_dest
    else:
        print("[ERROR] Không tìm thấy file PDF được tạo. Log:")
        print(res2.stdout[-1500:])
        return None


# ==============================================================================
# HÀM MAIN
# ==============================================================================

def main():
    print("=" * 75)
    print("BẮT ĐẦU TẠO 3 FILE CHO ĐỀ CƯƠNG ÔN TẬP GIỮA KÌ I KHTN 7 (THCS THỐNG NHẤT)")
    print("1. FILE HỌC SINH (.docx)")
    print("2. FILE GIÁO VIÊN (.docx có lời giải chi tiết)")
    print("3. FILE DẠY TRÊN MÀN HÌNH (Slide PDF 16:9 widescreen cho Tivi 65 inch)")
    print("=" * 75)

    f_hs = generate_word_document(mode="student")
    f_gv = generate_word_document(mode="teacher")
    f_screen_docx = generate_word_document(mode="screen")
    f_screen_pdf = generate_tivi_presentation_pdf()

    all_files = [f_hs, f_gv, f_screen_docx, f_screen_pdf]

    # Đồng bộ sang các thư mục OneDrive
    print("\n[+] Đồng bộ các file sang các thư mục OneDrive:")
    for od in TARGET_ONEDRIVE_DIRS:
        if os.path.exists(od):
            for src_path in all_files:
                if src_path and os.path.exists(src_path):
                    dest = os.path.join(od, os.path.basename(src_path))
                    try:
                        shutil.copy2(src_path, dest)
                        print(f"    -> Đã copy {os.path.basename(src_path)} vào: {od}")
                    except Exception as e:
                        print(f"    [!] Cảnh báo file đang mở hoặc bận tại {od}: {e}")

    print("\n" + "=" * 75)
    print("HOÀN THÀNH TẤT CẢ 3 FILE THEO ĐÚNG CHUẨN FORM ĐỀ KIỂM TRA!")
    print("=" * 75)

if __name__ == "__main__":
    main()
