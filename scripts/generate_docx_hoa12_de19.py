# -*- coding: utf-8 -*-
"""
Script tạo 2 file Word (.docx) chuẩn Form Đề Kiểm Tra cho:
ĐỀ SỐ 19 - CHƯƠNG 7: KIM LOẠI NHÓM IA & NHÓM IIA (HÓA HỌC 12)
Mã đề thi: 119
Bao gồm:
1. Bản Học sinh: De_Kiem_Tra_Hoa_12_De19.docx
2. Bản Giáo viên: De_Kiem_Tra_Hoa_12_De19_GiaoVien.docx (Có đáp án đỏ, lời giải chi tiết xanh, bảng đáp án)
"""

import os
import sys
sys.stdout.reconfigure(encoding='utf-8')
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"

COLOR_CORAL_RED = RGBColor(211, 47, 47)      # #D32F2F
COLOR_SAGE_GREEN = "8BC390"                 # #8BC390
COLOR_DEEP_BLUE = RGBColor(21, 101, 192)    # #1565C0
COLOR_DARK_GREEN = RGBColor(46, 125, 50)    # #2E7D32
COLOR_BLACK = RGBColor(0, 0, 0)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.59)     # 1.5 cm
        s.bottom_margin = Inches(0.59)  # 1.5 cm
        s.left_margin = Inches(0.59)    # 1.5 cm
        s.right_margin = Inches(0.59)   # 1.5 cm
        s.header_distance = Inches(0.24) # 0.6 cm
        s.footer_distance = Inches(0.24) # 0.6 cm

        # --- HEADER ---
        header = s.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.text = ""
        hp.paragraph_format.space_before = Pt(0)
        hp.paragraph_format.space_after = Pt(0)

        tbl_h = header.add_table(rows=1, cols=2, width=Inches(7.08)) # 18.0 cm
        tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
        c0, c1 = tbl_h.rows[0].cells
        c0.width = Inches(3.54)
        c1.width = Inches(3.54)
        set_cell_margins(c0, 60, 60, 0, 0)
        set_cell_margins(c1, 60, 60, 0, 0)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run("Tài liệu lưu hành nội bộ")
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(10)
        r0.font.italic = True
        r0.font.color.rgb = COLOR_CORAL_RED

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1_a = p1.add_run("Thầy ")
        r1_a.font.name = "Times New Roman"
        r1_a.font.size = Pt(10)
        r1_a.font.color.rgb = COLOR_CORAL_RED
        r1_b = p1.add_run("TRẦN VĂN THẠNH")
        r1_b.font.name = "Times New Roman"
        r1_b.font.size = Pt(10)
        r1_b.font.bold = True
        r1_b.font.color.rgb = COLOR_CORAL_RED
        r1_c = p1.add_run(" - 0777.470.803")
        r1_c.font.name = "Times New Roman"
        r1_c.font.size = Pt(10)
        r1_c.font.color.rgb = COLOR_CORAL_RED

        tblPr_h = tbl_h._tbl.tblPr
        tblBorders_h = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{COLOR_SAGE_GREEN}"/>'
            f'</w:tblBorders>'
        )
        tblPr_h.append(tblBorders_h)

        # --- FOOTER ---
        footer = s.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.text = ""
        fp.paragraph_format.space_before = Pt(0)
        fp.paragraph_format.space_after = Pt(0)

        tbl_f = footer.add_table(rows=1, cols=2, width=Inches(7.08))
        tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
        fc0, fc1 = tbl_f.rows[0].cells
        fc0.width = Inches(5.31)
        fc1.width = Inches(1.77)
        set_cell_margins(fc0, 60, 60, 0, 0)
        set_cell_margins(fc1, 60, 60, 0, 0)

        fp0 = fc0.paragraphs[0]
        fp0.paragraph_format.space_before = Pt(0)
        fp0.paragraph_format.space_after = Pt(0)
        fr0 = fp0.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
        fr0.font.name = "Times New Roman"
        fr0.font.size = Pt(9.5)
        fr0.font.italic = True
        fr0.font.color.rgb = COLOR_CORAL_RED

        fp1 = fc1.paragraphs[0]
        fp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp1.paragraph_format.space_before = Pt(0)
        fp1.paragraph_format.space_after = Pt(0)
        fr1_a = fp1.add_run("Trang ")
        fr1_a.font.name = "Times New Roman"
        fr1_a.font.size = Pt(10)
        fr1_a.font.bold = True
        fr1_a.font.color.rgb = COLOR_CORAL_RED

        fldSimple = OxmlElement('w:fldSimple')
        fldSimple.set(qn('w:instr'), 'PAGE')
        r_page = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        rFont = OxmlElement('w:rFonts')
        rFont.set(qn('w:ascii'), 'Times New Roman')
        rFont.set(qn('w:hAnsi'), 'Times New Roman')
        b = OxmlElement('w:b')
        c_el = OxmlElement('w:color')
        c_el.set(qn('w:val'), 'D32F2F')
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '20')
        rPr.append(rFont)
        rPr.append(b)
        rPr.append(c_el)
        rPr.append(sz)
        r_page.append(rPr)
        t_el = OxmlElement('w:t')
        t_el.text = "1"
        r_page.append(t_el)
        fldSimple.append(r_page)
        fp1._p.append(fldSimple)

        tblPr_f = tbl_f._tbl.tblPr
        tblBorders_f = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{COLOR_SAGE_GREEN}"/>'
            f'</w:tblBorders>'
        )
        tblPr_f.append(tblBorders_f)

def add_chem_run(paragraph, text, bold=False, italic=False, color=None, font_size=11):
    parts = re.split(r'(~.*?~|\^.*?\^)', text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run()
        run.font.name = "Times New Roman"
        run.font.size = Pt(font_size)
        run.bold = bold
        run.italic = italic
        if color:
            run.font.color.rgb = color

        if part.startswith('~') and part.endswith('~'):
            run.text = part[1:-1]
            run.font.subscript = True
        elif part.startswith('^') and part.endswith('^'):
            run.text = part[1:-1]
            run.font.superscript = True
        else:
            run.text = part

def create_title_block(doc, is_teacher=False):
    lines_info = [
        ("SỞ GIÁO DỤC VÀ ĐÀO TẠO THỪA THIÊN HUẾ", "ĐỀ KIỂM TRA ĐỊNH KỲ CHƯƠNG 7 (KIM LOẠI NHÓM IA & IIA)", True, True, COLOR_DARK_GREEN),
        ("TRƯỜNG THPT HAI BÀ TRƯNG", "MÔN: HOÁ HỌC -- LỚP 12", True, True, None),
        ("ĐỀ CHÍNH THỨC", "Thời gian làm bài: 50 phút (không kể thời gian phát đề)", True, False, None),
        ("(Đề thi có 04 trang)", "Mã đề thi: 119", False, True, None)
    ]

    for left_txt, right_txt, left_bold, right_bold, right_color in lines_info:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.08), WD_TAB_ALIGNMENT.RIGHT)

        r_l = p.add_run(left_txt)
        r_l.font.name = "Times New Roman"
        r_l.font.size = Pt(11)
        r_l.bold = left_bold
        if not left_bold:
            r_l.italic = True

        p.add_run("\t")

        r_r = p.add_run(right_txt)
        r_r.font.name = "Times New Roman"
        r_r.font.size = Pt(11)
        r_r.bold = right_bold
        if right_color:
            r_r.font.color.rgb = right_color
        if not right_bold:
            r_r.italic = True

    if is_teacher:
        p_note = doc.add_paragraph()
        p_note.paragraph_format.space_before = Pt(4)
        p_note.paragraph_format.space_after = Pt(4)
        p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_n = p_note.add_run("(BẢN GIÁO VIÊN -- CÓ LỜI GIẢI CHI TIẾT)")
        r_n.font.name = "Times New Roman"
        r_n.font.size = Pt(12)
        r_n.font.bold = True
        r_n.font.color.rgb = COLOR_CORAL_RED

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(3)
    p_info.paragraph_format.space_after = Pt(2)
    r_info = p_info.add_run("Họ, tên thí sinh: ..................................................................... Lớp: ........................ SBD: ............................")
    r_info.font.name = "Times New Roman"
    r_info.font.size = Pt(10.5)
    r_info.font.italic = True

    p_mass = doc.add_paragraph()
    p_mass.paragraph_format.space_before = Pt(0)
    p_mass.paragraph_format.space_after = Pt(5)
    r_mass = p_mass.add_run("(Cho biết nguyên tử khối: H = 1; C = 12; N = 14; O = 16; Na = 23; Mg = 24; Al = 27; P = 31; S = 32; Cl = 35,5; K = 39; Ca = 40; Fe = 56; Cu = 64; Zn = 65; Br = 80; Rb = 85,5; Ag = 108; Cs = 133; Ba = 137).")
    r_mass.font.name = "Times New Roman"
    r_mass.font.size = Pt(9.5)
    r_mass.font.italic = True

def add_part_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = COLOR_DARK_GREEN

def add_explanation(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.line_spacing = 1.15

    r_lbl = p.add_run("Lời giải chi tiết: ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(10)
    r_lbl.font.bold = True
    r_lbl.font.italic = True
    r_lbl.font.color.rgb = COLOR_DEEP_BLUE

    add_chem_run(p, text, font_size=10, italic=False, color=COLOR_DEEP_BLUE)

def render_mcq(doc, q_num, text, options, correct_key, explanation, is_teacher=False, img_path=None, img_width=Inches(3.5)):
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(3)
    p_q.paragraph_format.space_after = Pt(2)
    p_q.paragraph_format.line_spacing = 1.15

    r_lbl = p_q.add_run(f"Câu {q_num}: ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(11)
    r_lbl.bold = True

    add_chem_run(p_q, text, font_size=11)

    if img_path and os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(2)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture(img_path, width=img_width)

    max_len = max(len(opt[1]) for opt in options)
    if max_len <= 15:
        p_opts = doc.add_paragraph()
        p_opts.paragraph_format.space_before = Pt(0)
        p_opts.paragraph_format.space_after = Pt(2)
        p_opts.paragraph_format.line_spacing = 1.15
        p_opts.paragraph_format.tab_stops.add_tab_stop(Inches(1.77), WD_TAB_ALIGNMENT.LEFT)
        p_opts.paragraph_format.tab_stops.add_tab_stop(Inches(3.54), WD_TAB_ALIGNMENT.LEFT)
        p_opts.paragraph_format.tab_stops.add_tab_stop(Inches(5.31), WD_TAB_ALIGNMENT.LEFT)

        for i, (k, txt) in enumerate(options):
            is_c = (k == correct_key) and is_teacher
            r_k = p_opts.add_run(f"{k}. ")
            r_k.font.name = "Times New Roman"
            r_k.font.size = Pt(11)
            r_k.font.bold = is_c or True
            if is_c:
                r_k.font.color.rgb = COLOR_CORAL_RED
            add_chem_run(p_opts, txt, bold=is_c, color=COLOR_CORAL_RED if is_c else None, font_size=11)
            if i < 3:
                p_opts.add_run("\t")

    elif max_len <= 35:
        for row_opts in [options[:2], options[2:]]:
            p_opts = doc.add_paragraph()
            p_opts.paragraph_format.space_before = Pt(0)
            p_opts.paragraph_format.space_after = Pt(2)
            p_opts.paragraph_format.line_spacing = 1.15
            p_opts.paragraph_format.tab_stops.add_tab_stop(Inches(3.54), WD_TAB_ALIGNMENT.LEFT)

            for i, (k, txt) in enumerate(row_opts):
                is_c = (k == correct_key) and is_teacher
                r_k = p_opts.add_run(f"{k}. ")
                r_k.font.name = "Times New Roman"
                r_k.font.size = Pt(11)
                r_k.font.bold = is_c or True
                if is_c:
                    r_k.font.color.rgb = COLOR_CORAL_RED
                add_chem_run(p_opts, txt, bold=is_c, color=COLOR_CORAL_RED if is_c else None, font_size=11)
                if i == 0:
                    p_opts.add_run("\t")
    else:
        for k, txt in options:
            p_opts = doc.add_paragraph()
            p_opts.paragraph_format.space_before = Pt(0)
            p_opts.paragraph_format.space_after = Pt(2)
            p_opts.paragraph_format.line_spacing = 1.15
            is_c = (k == correct_key) and is_teacher
            r_k = p_opts.add_run(f"{k}. ")
            r_k.font.name = "Times New Roman"
            r_k.font.size = Pt(11)
            r_k.font.bold = is_c or True
            if is_c:
                r_k.font.color.rgb = COLOR_CORAL_RED
            add_chem_run(p_opts, txt, bold=is_c, color=COLOR_CORAL_RED if is_c else None, font_size=11)

    if is_teacher:
        add_explanation(doc, explanation)

def render_tf(doc, q_num, text, sub_items, explanation, is_teacher=False, img_path=None, img_width=Inches(3.5)):
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(3)
    p_q.paragraph_format.space_after = Pt(2)
    p_q.paragraph_format.line_spacing = 1.15

    r_lbl = p_q.add_run(f"Câu {q_num}: ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(11)
    r_lbl.bold = True

    add_chem_run(p_q, text, font_size=11)

    if img_path and os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(2)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture(img_path, width=img_width)

    for sub_k, sub_txt, is_true in sub_items:
        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(1.5)
        p_sub.paragraph_format.left_indent = Inches(0.2)
        p_sub.paragraph_format.line_spacing = 1.15

        r_sk = p_sub.add_run(f"{sub_k}) ")
        r_sk.font.name = "Times New Roman"
        r_sk.font.size = Pt(11)
        r_sk.font.bold = True

        add_chem_run(p_sub, sub_txt, font_size=11)

        if is_teacher:
            r_ans = p_sub.add_run(f" --> {'ĐÚNG' if is_true else 'SAI'}")
            r_ans.font.name = "Times New Roman"
            r_ans.font.size = Pt(10.5)
            r_ans.font.bold = True
            r_ans.font.color.rgb = COLOR_CORAL_RED

    if is_teacher:
        add_explanation(doc, explanation)

def render_short(doc, q_num, text, answer, explanation, is_teacher=False):
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(3)
    p_q.paragraph_format.space_after = Pt(2)
    p_q.paragraph_format.line_spacing = 1.15

    r_lbl = p_q.add_run(f"Câu {q_num}: ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(11)
    r_lbl.bold = True

    add_chem_run(p_q, text, font_size=11)

    if is_teacher:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(0)
        p_ans.paragraph_format.space_after = Pt(1)
        p_ans.paragraph_format.left_indent = Inches(0.25)
        r_ak = p_ans.add_run("Đáp số: ")
        r_ak.font.name = "Times New Roman"
        r_ak.font.size = Pt(10.5)
        r_ak.font.bold = True
        r_ak.font.color.rgb = COLOR_CORAL_RED
        r_av = p_ans.add_run(str(answer))
        r_av.font.name = "Times New Roman"
        r_av.font.size = Pt(10.5)
        r_av.font.bold = True
        r_av.font.color.rgb = COLOR_CORAL_RED

        add_explanation(doc, explanation)
    else:
        p_dot = doc.add_paragraph()
        p_dot.paragraph_format.space_before = Pt(0)
        p_dot.paragraph_format.space_after = Pt(2)
        r_d = p_dot.add_run("Trả lời: .....................................................................................................................................................")
        r_d.font.name = "Times New Roman"
        r_d.font.size = Pt(10)
        r_d.font.italic = True

def add_answer_table(doc):
    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_before = Pt(8)
    p_t.paragraph_format.space_after = Pt(4)
    r_t = p_t.add_run("BẢNG ĐÁP ÁN THAM KHẢO")
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(11)
    r_t.font.bold = True
    r_t.font.color.rgb = COLOR_CORAL_RED

    p_p1 = doc.add_paragraph()
    p_p1.paragraph_format.space_before = Pt(0)
    p_p1.paragraph_format.space_after = Pt(2)
    r1 = p_p1.add_run("PHẦN I: ")
    r1.font.bold = True
    r1.font.color.rgb = COLOR_DARK_GREEN
    p_p1.add_run("1.A   2.C   3.A   4.B   5.A   6.C   7.B   8.A   9.B   10.B   11.A   12.C   13.A   14.A   15.A   16.A   17.D   18.B")

    p_p2 = doc.add_paragraph()
    p_p2.paragraph_format.space_before = Pt(2)
    p_p2.paragraph_format.space_after = Pt(2)
    r2 = p_p2.add_run("PHẦN II: ")
    r2.font.bold = True
    r2.font.color.rgb = COLOR_DARK_GREEN
    p_p2.add_run("Câu 1: a-S, b-Đ, c-S, d-Đ   |   Câu 2: a-Đ, b-S, c-S, d-S\nCâu 3: a-S, b-S, c-Đ, d-Đ   |   Câu 4: a-Đ, b-Đ, c-S, d-Đ")

    p_p3 = doc.add_paragraph()
    p_p3.paragraph_format.space_before = Pt(2)
    p_p3.paragraph_format.space_after = Pt(2)
    r3 = p_p3.add_run("PHẦN III: ")
    r3.font.bold = True
    r3.font.color.rgb = COLOR_DARK_GREEN
    p_p3.add_run("Câu 1: 4   |   Câu 2: 138   |   Câu 3: 3   |   Câu 4: 3   |   Câu 5: 27,4   |   Câu 6: 6")

# DỮ LIỆU ĐỀ 19
P1_DATA = [
    (1, "Các đại dương là những kho muối vô tận với nhiều khoáng chất có giá trị dinh dưỡng cao. Trong nước biển, hai nguyên tố kim loại có nhiều nhất là",
     [("A", "sodium và magnesium."), ("B", "đồng (copper) và kẽm (zinc)."), ("C", "nhôm (aluminium) và sắt (iron)."), ("D", "vàng (gold) và bạc (silver).")],
     "A", "Trong nước biển, hai cation kim loại dồi dào nhất là Na^+^ và Mg^2+^. Do đó, sodium và magnesium là hai nguyên tố kim loại có nhiều nhất trong nước biển."),

    (2, "Cấu hình electron lớp ngoài cùng của nguyên tử kim loại kiềm là",
     [("A", "ns^2^np^2^."), ("B", "ns^2^np^1^."), ("C", "ns^1^."), ("D", "ns^2^.")],
     "C", "Các nguyên tử kim loại kiềm thuộc nhóm IA, có 1 electron ở lớp ngoài cùng với cấu hình electron chung là ns^1^."),

    (3, "Kim loại nào sau đây là kim loại kiềm?",
     [("A", "Na."), ("B", "Ca."), ("C", "Al."), ("D", "Fe.")],
     "A", "Sodium (Na) thuộc nhóm IA nên là kim loại kiềm. Ca thuộc nhóm IIA (kiềm thổ), Al nhóm IIIA, Fe nhóm VIIIB."),

    (4, "Cho một mẩu Na vào dung dịch MgSO~4~ dư, thu được kết tủa X và chất khí Y. Hai chất X, Y lần lượt là",
     [("A", "Mg và H~2~."), ("B", "Mg(OH)~2~ và H~2~."), ("C", "Mg và O~2~."), ("D", "Na~2~SO~4~ và H~2~.")],
     "B", "Na phản ứng mãnh liệt với nước trong dung dịch: 2Na + 2H~2~O -> 2NaOH + H~2~^ (khí Y là H~2~). Sau đó, NaOH sinh ra phản ứng trao đổi với MgSO~4~ dư: 2NaOH + MgSO~4~ -> Mg(OH)~2~v (kết tủa X) + Na~2~SO~4~."),

    (5, "Trong cốc nước chứa nhiều các ion sau: Ca^2+^, Mg^2+^, Cl^-^, SO~4~^2-^. Nước trong cốc trên thuộc loại",
     [("A", "có tính cứng vĩnh cửu."), ("B", "không có tính cứng."), ("C", "có tính cứng tạm thời."), ("D", "có tính cứng toàn phần.")],
     "A", "Nước cứng vĩnh cửu là nước cứng chứa các ion Ca^2+^, Mg^2+^ cùng với các anion Cl^-^, SO~4~^2-^ (không chứa ion HCO~3~^-^)."),

    (6, "Nhận định nào không đúng về cấu tạo và tính chất của các kim loại nhóm IIA?",
     [("A", "Khối lượng riêng tương đối nhỏ, là những kim loại nhẹ hơn nhôm (trừ Ba)."),
      ("B", "Độ cứng cao hơn kim loại nhóm IA, nhưng mềm hơn kim loại nhôm."),
      ("C", "Mạng tinh thể của chúng đều có kiểu lập phương tâm khối."),
      ("D", "Nhiệt độ nóng chảy và nhiệt độ sôi tương đối thấp (trừ Be).")],
     "C", "Các kim loại nhóm IIA không cùng một kiểu mạng tinh thể: Be và Mg có mạng lục phương; Ca và Sr có mạng lập phương tâm diện; Ba có mạng lập phương tâm khối."),

    (7, "Phương trình hóa học của phản ứng nào sau đây không đúng?",
     [("A", "Ca + 2H~2~O -> Ca(OH)~2~ + H~2~."),
      ("B", "BeO + H~2~O -> Be(OH)~2~."),
      ("C", "Mg + H~2~SO~4~ -> MgSO~4~ + H~2~."),
      ("D", "Be + 2NaOH -> Na~2~BeO~2~ + H~2~.")],
     "B", "Beryllium (Be) và oxit của nó (BeO) có lớp màng bảo vệ bền vững và tính lưỡng tính đặc biệt, hoàn toàn không tác dụng với nước ở bất kì điều kiện nhiệt độ nào."),

    (8, "Cách nào sau đây không thu được NaOH sau phản ứng?",
     [("A", "Điện phân dung dịch NaCl không có màng ngăn xốp."),
      ("B", "Cho kim loại Na tác dụng với nước."),
      ("C", "Điện phân dung dịch NaCl có màng ngăn xốp, điện cực trơ."),
      ("D", "Cho dung dịch Ca(OH)~2~ tác dụng với dung dịch Na~2~CO~3~.")],
     "A", "Khi điện phân dung dịch NaCl không có màng ngăn xốp, khí Cl~2~ sinh ra ở anode sẽ tác dụng ngay với dung dịch NaOH sinh ra ở cathode tạo thành nước Javel (NaCl + NaClO + H~2~O), do đó không thu được NaOH."),

    (9, "Hai chất nào sau đây được dùng để làm mềm nước có tính cứng vĩnh cửu?",
     [("A", "NaCl và Ca(OH)~2~."), ("B", "Na~2~CO~3~ và Na~3~PO~4~."), ("C", "Na~2~CO~3~ và Ca(OH)~2~."), ("D", "Na~2~CO~3~ và HCl.")],
     "B", "Để làm mềm nước có tính cứng vĩnh cửu hoặc toàn phần, người ta dùng dung dịch chứa ion CO~3~^2-^ hoặc PO~4~^3-^ (như Na~2~CO~3~, Na~3~PO~4~) để kết tủa hoàn toàn Ca^2+^ và Mg^2+^ dưới dạng CaCO~3~, MgCO~3~ hoặc Ca~3~(PO~4~)~2~, Mg~3~(PO~4~)~2~."),

    (10, "Hợp chất Y của calcium là thành phần chính của vỏ các loại ốc, sò,... Ngoài ra Y được sử dụng rộng rãi trong sản xuất vôi, xi măng, thuỷ tinh, chất phụ gia trong công nghiệp thực phẩm. Hợp chất Y là",
     [("A", "CaO."), ("B", "CaCO~3~."), ("C", "Ca(OH)~2~."), ("D", "Ca~3~(PO~4~)~2~.")],
     "B", "Thành phần chính của vỏ các loài ốc, sò, san hô, đá vôi, đá phấn, cẩm thạch là calcium carbonate (CaCO~3~)."),

    (11, "Kim loại không phản ứng với nước ở nhiệt độ thường là",
     [("A", "Be."), ("B", "Ca."), ("C", "Li."), ("D", "K.")],
     "A", "Beryllium (Be) trơ và không tác dụng với nước ngay cả ở nhiệt độ cao đỏ do có màng oxit bền bảo vệ. Ca, Li, K đều tác dụng mãnh liệt với nước ở nhiệt độ thường."),

    (12, "Xét phản ứng phân huỷ muối carbonate của kim loại nhóm IIA: MCO~3~(s) -> MO(s) + CO~2~(g) (Δ~r~H°~298~). Từ MgCO~3~ đến BaCO~3~, biến thiên enthalpy chuẩn của phản ứng biến đổi như thế nào?",
     [("A", "Không đổi."), ("B", "Giảm dần."), ("C", "Tăng dần."), ("D", "Không có quy luật.")],
     "C", "Độ bền nhiệt của muối carbonate kim loại nhóm IIA tăng dần từ MgCO~3~ đến BaCO~3~ (nhiệt độ bắt đầu phân hủy tăng dần: MgCO~3~ < CaCO~3~ < SrCO~3~ < BaCO~3~). Do đó, phản ứng thu nhiệt ngày càng nhiều hơn, biến thiên enthalpy chuẩn Δ~r~H°~298~ tăng dần."),

    (13, "Ở nhiệt độ thường, độ tan của các hydroxide tăng dần trong dãy từ Mg(OH)~2~ đến Ba(OH)~2~. Từ thông tin này có thể dự đoán được khả năng phản ứng với nước của các kim loại từ Mg đến Ba biến đổi như thế nào?",
     [("A", "Tăng dần."), ("B", "Không đổi."), ("C", "Không có quy luật."), ("D", "Giảm dần.")],
     "A", "Khi kim loại nhóm IIA phản ứng với nước, hydroxide tạo thành nếu dễ tan sẽ nhanh chóng khuếch tán vào dung dịch, giải phóng bề mặt kim loại tiếp xúc trực tiếp với nước. Do độ tan hydroxide tăng dần từ Mg(OH)~2~ (kết tủa) đến Ba(OH)~2~ (tan tốt), khả năng phản ứng với nước của các kim loại tăng dần từ Mg đến Ba."),

    (14, "Nước chứa nhiều các ion nào sau đây có tính cứng toàn phần?",
     [("A", "Mg^2+^, Ca^2+^, HCO~3~^-^, SO~4~^2-^, Cl^-^."),
      ("B", "Na^+^, K^+^, SO~4~^2-^, Cl^-^."),
      ("C", "Mg^2+^, Ca^2+^, HCO~3~^-^."),
      ("D", "Mg^2+^, Ca^2+^, SO~4~^2-^, Cl^-^.")],
     "A", "Nước cứng toàn phần là loại nước cứng bao gồm cả tính cứng tạm thời (chứa HCO~3~^-^) và tính cứng vĩnh cửu (chứa Cl^-^, SO~4~^2-^) cùng với các cation Ca^2+^, Mg^2+^."),

    (15, "Phân tích một mẫu nước tự nhiên thấy chứa nhiều các ion: Na^+^, Ca^2+^, HCO~3~^-^, Cl^-^ và SO~4~^2-^. Chất nào sau đây có thể làm mềm mẫu nước trên?",
     [("A", "Na~2~CO~3~."), ("B", "Ca(OH)~2~."), ("C", "NaOH."), ("D", "HCl.")],
     "A", "Mẫu nước trên chứa cả ion Ca^2+^, HCO~3~^-^ và Cl^-^, SO~4~^2-^ nên là nước cứng toàn phần. Để làm mềm triệt để nước cứng toàn phần, chất tối ưu và kinh tế nhất là dùng dung dịch soda Na~2~CO~3~: Ca^2+^ + CO~3~^2-^ -> CaCO~3~v."),

    (16, "Một loại nước cứng khi đun sôi thì trở thành nước mềm. Trong loại nước này có hoà tan những hợp chất nào sau đây?",
     [("A", "Ca(HCO~3~)~2~ và Mg(HCO~3~)~2~."),
      ("B", "Ca(HCO~3~)~2~ và MgSO~4~."),
      ("C", "CaSO~4~ và MgCl~2~."),
      ("D", "MgCl~2~ và CaCl~2~.")],
     "A", "Nước mất tính cứng hoàn toàn khi đun sôi là nước cứng tạm thời. Trong nước chứa các muối hydrogencarbonate Ca(HCO~3~)~2~ và Mg(HCO~3~)~2~. Khi đun sôi: Ca(HCO~3~)~2~ -> CaCO~3~v + CO~2~^ + H~2~O và Mg(HCO~3~)~2~ -> MgCO~3~v (hoặc Mg(OH)~2~v) + CO~2~^ + H~2~O."),

    (17, "Tiến hành thí nghiệm theo các bước sau:\nBước 1: Chuẩn bị hai ống nghiệm, ống (1) chứa 2 mL dung dịch CaCl~2~ 1 M, ống (2) chứa 2 mL dung dịch BaCl~2~ 1 M.\nBước 2: Nhỏ đồng thời vào mỗi ống nghiệm 3 giọt dung dịch CuSO~4~ 1 M, thấy ống (1) xuất hiện kết tủa chậm hơn và ít hơn so với ống (2).\nNhận định nào sau đây đúng khi so sánh CaSO~4~ với BaSO~4~?",
     [("A", "Khó nhiệt phân hơn."), ("B", "Khó thuỷ phân hơn."), ("C", "Dễ kết tủa hơn."), ("D", "Dễ tan hơn.")],
     "D", "Ống (1) tạo CaSO~4~, ống (2) tạo BaSO~4~. Vì ống (1) xuất hiện kết tủa chậm hơn và lượng kết tủa ít hơn hẳn so với ống (2) nên độ tan của CaSO~4~ lớn hơn BaSO~4~ nhiều (CaSO~4~ là chất ít tan, trong khi BaSO~4~ là chất thực tế không tan). Như vậy, CaSO~4~ dễ tan hơn BaSO~4~."),

    (18, "Cho sơ đồ chuyển hoá sau:\nCaO -(+X)-> Y -(+Z)-> CaCO~3~ -(+Z+X)-> T -(+E)-> CaSO~4~\nBiết: X, Y, Z, T, E là các hợp chất khác nhau; mỗi mũi tên ứng với một phương trình hoá học. Các chất Z, E thoả mãn sơ đồ trên lần lượt là",
     [("A", "Na~2~CO~3~, H~2~SO~4~."), ("B", "CO~2~, KHSO~4~."), ("C", "NaHCO~3~, Na~2~SO~4~."), ("D", "CO~2~, BaSO~4~.")],
     "B", "X là H~2~O: CaO + H~2~O -> Ca(OH)~2~ (Y).\nZ là CO~2~: Ca(OH)~2~ + CO~2~ -> CaCO~3~v + H~2~O.\n+Z+X: CaCO~3~ + CO~2~ + H~2~O -> Ca(HCO~3~)~2~ (T).\nE là KHSO~4~: Ca(HCO~3~)~2~ + 2KHSO~4~ -> CaSO~4~v + K~2~SO~4~ + 2CO~2~^ + 2H~2~O.")
]

P2_DATA = [
    (1, "Quặng sylvinite là một khoáng chất phổ biến có thành phần chính là NaCl.KCl. Sự phụ thuộc của độ tan các muối vào nhiệt độ được biểu diễn ở đồ thị sau:",
     [("a", "Độ tan của KCl giảm chậm khi giảm nhiệt độ từ 100 °C về 0 °C.", False),
      ("b", "Tách được KCl khỏi dung dịch NaCl bằng phương pháp kết tinh phân đoạn.", True),
      ("c", "Độ tan của NaCl tăng nhanh khi nhiệt độ từ 0 °C đến 100 °C.", False),
      ("d", "Độ tan của KCl giảm nhanh hơn của NaCl khi giảm nhiệt độ từ 100 °C về 0 °C.", True)],
     "San_Pham/images_crop_de19/p2_cau1_do_thi_do_tan.png", Inches(4.5),
     "a) Sai vì đường biểu diễn độ tan của KCl rất dốc, khi hạ nhiệt độ từ 100 °C xuống 0 °C độ tan của KCl giảm rất nhanh (từ ~56 g xuống ~28 g / 100 g nước).\n"
     "b) Đúng vì độ dốc của hai đường tan chênh lệch rất lớn: khi làm lạnh dung dịch bão hòa ở nhiệt độ cao, KCl kết tinh tách ra với lượng rất lớn trong khi NaCl vẫn hầu như tan trong dung dịch.\n"
     "c) Sai vì đường độ tan của NaCl gần như nằm ngang (biến thiên rất ít, từ 35,7 g ở 0 °C lên 39,8 g ở 100 °C / 100 g nước).\n"
     "d) Đúng vì khi nhiệt độ giảm từ 100 °C về 0 °C, độ dốc của KCl lớn hơn nhiều so với NaCl."),

    (2, "Mỗi phát biểu sau đây về kim loại kiềm là đúng hay sai?",
     [("a", "Thứ tự giảm dần tính kim loại của các kim loại kiềm là: Cs, Rb, K, Na, Li.", True),
      ("b", "Phương pháp chung để điều chế kim loại kiềm là điện phân dung dịch.", False),
      ("c", "Để bảo quản kim loại Na cần ngâm Na trong cồn tinh khiết.", False),
      ("d", "Na~2~O tan trong nước tạo dung dịch trong suốt và thoát ra khí H~2~.", False)],
     None, None,
     "a) Đúng. Trong nhóm IA đi từ trên xuống dưới theo chiều tăng của Z (Li đến Cs), bán kính nguyên tử tăng, năng lượng ion hóa giảm nên tính kim loại tăng dần: Li < Na < K < Rb < Cs. Vậy thứ tự giảm dần là Cs > Rb > K > Na > Li.\n"
     "b) Sai. Do ion kim loại kiềm có tính oxi hóa rất yếu (thế điện cực chuẩn rất âm), không thể khử được chúng trong môi trường nước nên phương pháp duy nhất là điện phân nóng chảy muối halide của chúng.\n"
     "c) Sai. Na phản ứng mãnh liệt với cồn tinh khiết (ethanol) giải phóng khí H~2~: 2Na + 2C~2~H~5~OH -> 2C~2~H~5~ONa + H~2~^. Để bảo quản Na, bắt buộc phải ngâm trong dầu hỏa khan.\n"
     "d) Sai. Na~2~O tác dụng với nước là phản ứng hóa hợp không có sự thay đổi số oxi hóa: Na~2~O + H~2~O -> 2NaOH, dung dịch trong suốt và không giải phóng khí."),

    (3, "Cho các mẫu nước được đánh dấu X, Y, Z, T có chứa các ion theo bảng sau:\n"
        "- Mẫu X: Ca^2+^, Mg^2+^ (nhiều), Cl^-^, SO~4~^2-^\n"
        "- Mẫu Y: Ca^2+^, Mg^2+^ (nhiều), HCO~3~^-^\n"
        "- Mẫu Z: Na^+^, K^+^, HCO~3~^-^\n"
        "- Mẫu T: Ca^2+^, Mg^2+^ (nhiều), Cl^-^, SO~4~^2-^, HCO~3~^-^",
     [("a", "Nước X và Y đều là nước có tính cứng vĩnh cửu.", False),
      ("b", "Có thể loại bỏ tính cứng của nước T bằng cách đun nóng.", False),
      ("c", "Có thể loại bỏ tính cứng của nước Y bằng dung dịch Na~2~CO~3~ dư.", True),
      ("d", "Có thể phân biệt X, Y, Z và T bằng cách đun nóng nhẹ và dùng dung dịch NaOH.", True)],
     None, None,
     "a) Sai. Nước X là nước cứng vĩnh cửu, còn nước Y là nước cứng tạm thời.\n"
     "b) Sai. Nước T là nước cứng toàn phần. Khi đun nóng, chỉ các ion Ca^2+^, Mg^2+^ gắn với HCO~3~^-^ kết tủa, các ion Ca^2+^, Mg^2+^ liên kết với Cl^-^, SO~4~^2-^ vẫn còn tồn tại trong dung dịch nên nước vẫn còn tính cứng vĩnh cửu.\n"
     "c) Đúng. Dung dịch Na~2~CO~3~ cung cấp CO~3~^2-^ sẽ làm kết tủa hoàn toàn Ca^2+^ và Mg^2+^ trong nước Y.\n"
     "d) Đúng. Đun nóng nhẹ: Y và T xuất hiện kết tủa trắng (do HCO~3~^-^ bị nhiệt phân), X và Z không có kết tủa. Khi cho NaOH vào các mẫu: với nhóm đun nóng có kết tủa (Y và T), mẫu Y chứa tỉ lệ ion tạo kết tủa hết còn T có dư ion vĩnh cửu; với nhóm không có hiện tượng khi đun nóng (X và Z), mẫu X cho kết tủa với NaOH (Mg^2+^ + 2OH^-^ -> Mg(OH)~2~v), mẫu Z hoàn toàn không kết tủa."),

    (4, "Hợp chất của kim loại calcium có nhiều ứng dụng quan trọng trong đời sống và sản xuất.",
     [("a", "Hợp chất của calcium được dùng để đúc tượng, bó bột khi gãy xương là thạch cao nung (CaSO~4~.H~2~O).", True),
      ("b", "Vôi sống (CaO) được dùng phổ biến để khử chua cho đất trồng trọt.", True),
      ("c", "Phương trình hoá học của phản ứng tạo thạch nhũ trong các hang động đá vôi là: CaCO~3~ + CO~2~ + H~2~O -> Ca(HCO~3~)~2~.", False),
      ("d", "Thành phần hoá học chính của vỏ các loài ốc, sò, hến là calcium carbonate (CaCO~3~).", True)],
     None, None,
     "a) Đúng. Thạch cao sống CaSO~4~.2H~2~O nung ở ~150-180 °C tạo thạch cao nung CaSO~4~.H~2~O hoặc CaSO~4~.0,5H~2~O. Khi nhào với nước nó đông cứng nhanh và nở thể tích nên dùng nặn tượng, đúc khuôn, bó bột.\n"
     "b) Đúng. Vôi sống CaO hoặc vôi tôi Ca(OH)~2~ tác dụng trung hòa các axit hữu cơ và vô cơ trong đất phèn, giúp khử chua.\n"
     "c) Sai. Phản ứng tạo thạch nhũ là quá trình phân hủy muối hydrogencarbonate nhỏ giọt từ trần hang: Ca(HCO~3~)~2~ -> CaCO~3~v + CO~2~^ + H~2~O. Phản ứng CaCO~3~ + CO~2~ + H~2~O -> Ca(HCO~3~)~2~ là quá trình hòa tan đá vôi (xâm thực đá vôi).\n"
     "d) Đúng. Vỏ sò, ốc, hến cấu tạo chủ yếu từ CaCO~3~ chiếm trên 95% khối lượng.")
]

P3_DATA = [
    (1, "Cho các kim loại sau: Li, Na, Ca, Mg, Fe, Cu, K, Zn, Ag, Rb. Có bao nhiêu kim loại kiềm trong dãy trên?",
     "4",
     "Các kim loại kiềm thuộc phân nhóm chính nhóm I (nhóm IA), gồm: Li, Na, K, Rb (tổng cộng 4 kim loại kiềm).\nCác kim loại còn lại: Ca, Mg thuộc nhóm IIA (kiềm thổ); Fe, Cu, Zn, Ag là kim loại chuyển tiếp."),

    (2, "Cho sơ đồ các phản ứng sau theo đúng tỉ lệ mol:\n"
        "(1) 2X~1~ + 2X~2~ -> 2X~3~ + H~2~\n"
        "(2) X~3~ + CO~2~ -> X~4~\n"
        "(3) X~3~ + X~4~ -> X~5~ + X~2~\n"
        "(4) 2X~6~ + 3X~5~ + 3X~2~ -> 2Fe(OH)~3~ + 3CO~2~ + 6KCl\n"
        "Khối lượng mol của chất X~5~ là bao nhiêu g/mol?",
     "138",
     "Từ phản ứng (4), sinh ra KCl nên cation kim loại trong các chất là potassium (K).\n"
     "- Phản ứng (4): 2FeCl~3~ + 3K~2~CO~3~ + 3H~2~O -> 2Fe(OH)~3~v + 3CO~2~^ + 6KCl. Vậy X~6~ là FeCl~3~, X~5~ là K~2~CO~3~, X~2~ là H~2~O.\n"
     "- Phản ứng (1): 2K + 2H~2~O -> 2KOH + H~2~^ => X~1~ là K, X~3~ là KOH.\n"
     "- Phản ứng (2): KOH + CO~2~ -> KHCO~3~ => X~4~ là KHCO~3~.\n"
     "- Phản ứng (3): KOH + KHCO~3~ -> K~2~CO~3~ + H~2~O.\n"
     "Vậy X~5~ là K~2~CO~3~. Khối lượng mol M(X~5~) = 39 * 2 + 12 + 16 * 3 = 138 g/mol."),

    (3, "Cho các phát biểu sau:\n"
        "(1) Soda là chất bột màu trắng, tan trong nước tạo môi trường trung tính.\n"
        "(2) Soda dùng làm mềm nước cứng.\n"
        "(3) Soda bền với nhiệt hơn so với baking soda.\n"
        "(4) Chất béo bị thuỷ phân trong dung dịch soda tạo thành xà phòng.\n"
        "(5) Có thể dùng baking soda thay cho soda trong việc tẩy rửa lớp dầu, mỡ bám vào bồn cầu.\n"
        "Có bao nhiêu phát biểu đúng trong các phát biểu trên?",
     "3",
     "Các phát biểu đúng là (2), (3), (4):\n"
     "- (1) Sai vì ion CO~3~^2-^ bị thuỷ phân nấc 1 tạo ion OH^-^ làm dung dịch có môi trường kiềm (pH > 7).\n"
     "- (2) Đúng vì ion CO~3~^2-^ tạo kết tủa với Ca^2+^ và Mg^2+^ làm mềm cả nước cứng tạm thời và vĩnh cửu.\n"
     "- (3) Đúng vì Na~2~CO~3~ (soda) nóng chảy mà không bị nhiệt phân, trong khi NaHCO~3~ (baking soda) bị phân huỷ ngay ở nhiệt độ khoảng 70-100 °C.\n"
     "- (4) Đúng vì dung dịch soda có tính kiềm, thủy phân este chất béo (este của glycerol) tạo xà phòng (muối natri của axit béo).\n"
     "- (5) Sai vì baking soda (NaHCO~3~) tạo môi trường kiềm rất yếu, khả năng thuỷ phân và nhũ hoá dầu mỡ kém hơn nhiều so với soda."),

    (4, "Trong số các chất sau: Ca(OH)~2~, Na~2~CO~3~, Na~3~PO~4~, KCl, có bao nhiêu chất có thể dùng để làm mềm nước có tính cứng tạm thời?",
     "3",
     "Có 3 chất là Ca(OH)~2~, Na~2~CO~3~ và Na~3~PO~4~:\n"
     "- Ca(OH)~2~ (dùng lượng vừa đủ): Ca^2+^ + HCO~3~^-^ + OH^-^ -> CaCO~3~v + H~2~O.\n"
     "- Na~2~CO~3~: Ca^2+^ + CO~3~^2-^ -> CaCO~3~v và Mg^2+^ + CO~3~^2-^ -> MgCO~3~v.\n"
     "- Na~3~PO~4~: 3Ca^2+^ + 2PO~4~^3-^ -> Ca~3~(PO~4~)~2~v.\n"
     "KCl không phản ứng với ion Ca^2+^, Mg^2+^ hay HCO~3~^-^ nên không làm mềm nước cứng."),

    (5, "Trong công nghiệp, người ta sản xuất xút (NaOH) từ muối ăn bằng phương pháp điện phân dung dịch có màng ngăn xốp. Tính khối lượng (tấn) NaCl cần dùng để sản xuất 15 tấn NaOH với hiệu suất của toàn bộ quá trình là 80% (kết quả làm tròn đến hàng phần mười).",
     "27,4",
     "Phương trình hoá học điện phân dung dịch có màng ngăn:\n"
     "2NaCl + 2H~2~O -(đpdd, màng ngăn)-> 2NaOH + Cl~2~^ + H~2~^\n"
     "Khối lượng mol: M(NaCl) = 58,5 g/mol; M(NaOH) = 40 g/mol.\n"
     "Theo lý thuyết, để thu được 15 tấn NaOH thì khối lượng NaCl cần là:\n"
     "m(NaCl)~lt~ = (15 * 58,5) / 40 = 21,9375 tấn.\n"
     "Vì hiệu suất phản ứng H = 80%, khối lượng NaCl thực tế cần dùng là:\n"
     "m(NaCl)~tt~ = 21,9375 / 0,80 = 27,421875 tấn ≈ 27,4 tấn."),

    (6, "Carnallite là muối khoáng kép ngậm nước có công thức dạng KCl.MgCl~2~.nH~2~O. Khi nung nóng 5,55 gam carnallite, thu được 3,39 gam muối khan. Mặt khác, nếu cho 5,55 gam carnallite tác dụng hoàn toàn với dung dịch NaOH dư, thu được kết tủa. Lọc và nung kết tủa đến khối lượng không đổi thì khối lượng chất rắn giảm 0,36 gam so với khối lượng kết tủa ban đầu. Xác định giá trị của n.",
     "6",
     "- Khối lượng nước kết tinh bay ra khi nung muối khoáng:\n"
     "m(H~2~O) = 5,55 - 3,39 = 2,16 g => n(H~2~O) = 2,16 / 18 = 0,12 mol.\n"
     "- Khi cho mẫu carnallite tác dụng với NaOH dư:\n"
     "MgCl~2~ + 2NaOH -> Mg(OH)~2~v + 2NaCl\n"
     "Nung kết tủa Mg(OH)~2~ đến khối lượng không đổi:\n"
     "Mg(OH)~2~ -(t°)-> MgO + H~2~O^\n"
     "Độ giảm khối lượng của kết tủa chính là khối lượng nước thoát ra từ phản ứng nhiệt phân Mg(OH)~2~:\n"
     "m(H~2~O sinh ra) = 0,36 g => n(H~2~O) = 0,36 / 18 = 0,02 mol.\n"
     "=> n(Mg(OH)~2~) = 0,02 mol => n(MgCl~2~) = 0,02 mol.\n"
     "- Trong công thức KCl.MgCl~2~.nH~2~O, tỉ lệ n = n(H~2~O) / n(MgCl~2~) = 0,12 / 0,02 = 6.")
]

def build_exam_docx(is_teacher=False):
    doc = docx.Document()
    add_header_footer(doc)
    create_title_block(doc, is_teacher=is_teacher)

    # PHẦN I
    add_part_heading(doc, "PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn. Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án.")
    for q_num, text, options, c_key, expl in P1_DATA:
        render_mcq(doc, q_num, text, options, c_key, expl, is_teacher=is_teacher)

    # PHẦN II
    add_part_heading(doc, "PHẦN II. Câu trắc nghiệm đúng sai. Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.")
    for q_num, text, sub_items, img_path, img_w, expl in P2_DATA:
        render_tf(doc, q_num, text, sub_items, expl, is_teacher=is_teacher, img_path=img_path, img_width=img_w)

    # PHẦN III
    add_part_heading(doc, "PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn. Thí sinh trả lời từ câu 1 đến câu 6.")
    for q_num, text, ans, expl in P3_DATA:
        render_short(doc, q_num, text, ans, expl, is_teacher=is_teacher)

    if is_teacher:
        add_answer_table(doc)

    out_name = "De_Kiem_Tra_Hoa_12_De19_GiaoVien.docx" if is_teacher else "De_Kiem_Tra_Hoa_12_De19.docx"
    out_path = os.path.join(OUT_DIR, out_name)
    doc.save(out_path)
    print(f"Đã lưu: {out_path}")

def main():
    print("--- Bắt đầu tạo 2 file Word cho ĐỀ SỐ 19 (Mã 119) ---")
    build_exam_docx(is_teacher=False)
    build_exam_docx(is_teacher=True)
    print("--- Hoàn tất sinh file Word cho ĐỀ SỐ 19! ---")

if __name__ == "__main__":
    main()
