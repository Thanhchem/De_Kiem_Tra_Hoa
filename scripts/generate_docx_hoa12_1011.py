import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

CROP_DIR = r"c:\Antigravity_Thanh\San_Pham\images_crop"
OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"

COLOR_HF_RED = RGBColor(211, 47, 47)      # Coral Red (#D32F2F)
COLOR_HF_GREEN_HEX = "8BC390"             # Sage Green for horizontal lines (#8BC390)

COLOR_PRIMARY = RGBColor(46, 125, 50)     # Dark Green / Teal for headers in body
COLOR_CORRECT = RGBColor(198, 40, 40)     # Red for correct answers
COLOR_SOLUTION = RGBColor(21, 101, 192)   # Blue for solutions
COLOR_GRAY = RGBColor(97, 97, 97)

def add_p_border_bottom(p, color="8BC390", sz="8"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_p_border_top(p, color="8BC390", sz="8"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def set_run_style(run, font_name="Times New Roman", font_size=12, bold=False, italic=False, subscript=False, superscript=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def render_rich_text(p, text, base_font="Times New Roman", base_size=12, base_bold=False, base_italic=False, base_color=None):
    tokens = re.split(r'(~[^~]+~|\^[^\^]+\^|\*[^\*]+\*)', text)
    for token in tokens:
        if not token:
            continue
        is_bold = base_bold
        is_italic = base_italic
        is_sub = False
        is_sup = False
        content = token

        if token.startswith('*') and token.endswith('*'):
            is_bold = True
            content = token[1:-1]
        elif token.startswith('~') and token.endswith('~'):
            is_sub = True
            content = token[1:-1]
        elif token.startswith('^') and token.endswith('^'):
            is_sup = True
            content = token[1:-1]

        run = p.add_run(content)
        set_run_style(run, font_name=base_font, font_size=base_size, bold=is_bold, italic=is_italic,
                      subscript=is_sub, superscript=is_sup, color=base_color)

def setup_header_footer(doc):
    normal_style = doc.styles['Normal']
    normal_style.paragraph_format.tab_stops.clear_all()
    for pos in [0.5, 1.5, 5.0, 9.5, 14.0]:
        normal_style.paragraph_format.tab_stops.add_tab_stop(Cm(pos), WD_TAB_ALIGNMENT.LEFT)

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    section.start_type = WD_SECTION_START.NEW_PAGE
    section.different_first_page_header_footer = False
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    header = section.header
    for p in header.paragraphs:
        p.text = ""
    tbl_h = header.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl_h.cell(0, 0)
    c_right = tbl_h.cell(0, 1)
    c_left.width = Cm(8.0)
    c_right.width = Cm(10.0)

    p_hl = c_left.paragraphs[0]
    p_hl.paragraph_format.space_before = Pt(0)
    p_hl.paragraph_format.space_after = Pt(2)
    p_hl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_hl = p_hl.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)

    p_hr = c_right.paragraphs[0]
    p_hr.paragraph_format.space_before = Pt(0)
    p_hr.paragraph_format.space_after = Pt(2)
    p_hr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_hr1 = p_hr.add_run("Thầy ")
    set_run_style(r_hr1, font_size=10, italic=True, color=COLOR_HF_RED)
    r_hr2 = p_hr.add_run("TRẦN VĂN THẠNH")
    set_run_style(r_hr2, font_size=10, bold=True, italic=True, color=COLOR_HF_RED)
    r_hr3 = p_hr.add_run(" - 0777.470.803")
    set_run_style(r_hr3, font_size=10, italic=True, color=COLOR_HF_RED)

    add_p_border_bottom(p_hl, color=COLOR_HF_GREEN_HEX, sz="8")
    add_p_border_bottom(p_hr, color=COLOR_HF_GREEN_HEX, sz="8")

    footer = section.footer
    for p in footer.paragraphs:
        p.text = ""
    tbl_f = footer.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    cf_left = tbl_f.cell(0, 0)
    cf_right = tbl_f.cell(0, 1)
    cf_left.width = Cm(14.5)
    cf_right.width = Cm(3.5)

    p_fl = cf_left.paragraphs[0]
    p_fl.paragraph_format.space_before = Pt(2)
    p_fl.paragraph_format.space_after = Pt(0)
    p_fl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_fl = p_fl.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)

    p_fr = cf_right.paragraphs[0]
    p_fr.paragraph_format.space_before = Pt(2)
    p_fr.paragraph_format.space_after = Pt(0)
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_field(p_fr, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED)

    add_p_border_top(p_fl, color=COLOR_HF_GREEN_HEX, sz="8")
    add_p_border_top(p_fr, color=COLOR_HF_GREEN_HEX, sz="8")

def format_cell_borders(cell, top="none", bottom="none", left="none", right="none", color="auto", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    xml = f'<w:tcBorders {nsdecls("w")}>'
    for side, val in [("top", top), ("left", left), ("bottom", bottom), ("right", right)]:
        if val != "none":
            xml += f'<w:{side} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        else:
            xml += f'<w:{side} w:val="none"/>'
    xml += '</w:tcBorders>'
    tcPr.append(parse_xml(xml))

def build_title_block(doc):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    c0 = tbl.cell(0, 0)
    c1 = tbl.cell(0, 1)
    c0.width = Cm(7.5)
    c1.width = Cm(10.5)

    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(1)
    p0.paragraph_format.line_spacing = 1.15
    render_rich_text(p0, "*THPT NGUYỄN ĐÌNH CHIỂU*\n*ĐỀ CHÍNH THỨC*\n*MÃ ĐỀ: 1011*", base_size=10.5)

    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(1)
    p1.paragraph_format.line_spacing = 1.15
    render_rich_text(p1, "*ĐỀ KIỂM TRA GIỮA KÌ I NĂM HỌC 2025 - 2026*\n*Môn thi: HÓA HỌC 12*\n_Thời gian làm bài: 45 phút (không tính thời gian phát đề)_", base_size=10.5)

    for c in [c0, c1]:
        format_cell_borders(c, "none", "none", "none", "none")

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(4)
    p_info.paragraph_format.space_after = Pt(4)
    r_info = p_info.add_run("Họ và tên học sinh: ....................................................................................   Lớp: 12A1")
    set_run_style(r_info, font_size=11, bold=True)
    add_p_border_bottom(p_info, color="CCCCCC", sz="4")

def add_section_header(doc, title_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"■ {title_text}")
    set_run_style(r, font_size=11.5, bold=True, color=COLOR_PRIMARY)

def add_solution_box(doc, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    c = tbl.cell(0, 0)
    c.width = Cm(18.0)
    
    tcPr = c._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F4F9F4"/>')
    tcPr.append(shd)
    format_cell_borders(c, left="single", color=COLOR_HF_GREEN_HEX, sz="18")
    
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.2)
    p.paragraph_format.line_spacing = 1.15
    
    r_lbl = p.add_run("Lời giải chi tiết:\n")
    set_run_style(r_lbl, font_size=10.5, bold=True, italic=True, color=COLOR_SOLUTION)
    render_rich_text(p, text, base_size=10.5, base_italic=False, base_color=RGBColor(50, 50, 50))

def add_mcq_choices_grid(doc, choices, correct_idx=-1, is_teacher=False):
    max_len = max(len(c) for c in choices)
    if max_len <= 18:
        tbl = doc.add_table(rows=1, cols=4)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        widths = [Cm(4.5), Cm(4.5), Cm(4.5), Cm(4.5)]
        for i, (ch, w) in enumerate(zip(choices, widths)):
            cell = tbl.cell(0, i)
            cell.width = w
            format_cell_borders(cell, "none", "none", "none", "none")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            prefix = f"{chr(65+i)}. "
            is_corr = (i == correct_idx) and is_teacher
            r_pre = p.add_run(prefix)
            set_run_style(r_pre, font_size=11, bold=is_corr, color=COLOR_CORRECT if is_corr else None)
            render_rich_text(p, ch, base_size=11, base_bold=is_corr, base_color=COLOR_CORRECT if is_corr else None)
    elif max_len <= 45:
        tbl = doc.add_table(rows=2, cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        widths = [Cm(9.0), Cm(9.0)]
        for i in range(4):
            r_idx = i // 2
            c_idx = i % 2
            cell = tbl.cell(r_idx, c_idx)
            cell.width = widths[c_idx]
            format_cell_borders(cell, "none", "none", "none", "none")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            prefix = f"{chr(65+i)}. "
            is_corr = (i == correct_idx) and is_teacher
            r_pre = p.add_run(prefix)
            set_run_style(r_pre, font_size=11, bold=is_corr, color=COLOR_CORRECT if is_corr else None)
            render_rich_text(p, choices[i], base_size=11, base_bold=is_corr, base_color=COLOR_CORRECT if is_corr else None)
    else:
        for i, ch in enumerate(choices):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Cm(0.5)
            prefix = f"{chr(65+i)}. "
            is_corr = (i == correct_idx) and is_teacher
            r_pre = p.add_run(prefix)
            set_run_style(r_pre, font_size=11, bold=is_corr, color=COLOR_CORRECT if is_corr else None)
            render_rich_text(p, ch, base_size=11, base_bold=is_corr, base_color=COLOR_CORRECT if is_corr else None)

def add_image_centered(doc, img_name, width_cm=8.0):
    img_path = os.path.join(CROP_DIR, img_name)
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(img_path, width=Cm(width_cm))

def build_exam_document(is_teacher=False):
    doc = docx.Document()
    setup_header_footer(doc)
    build_title_block(doc)

    # ==================== PHẦN TRẮC NGHIỆM ====================
    add_section_header(doc, "A. PHẦN TRẮC NGHIỆM (7 điểm)")
    
    # Phần I
    p_p1 = doc.add_paragraph()
    p_p1.paragraph_format.space_before = Pt(4)
    p_p1.paragraph_format.space_after = Pt(4)
    render_rich_text(p_p1, "*PHẦN I (3 điểm – 12 câu): Câu trắc nghiệm nhiều phương án lựa chọn.* Thí sinh trả lời từ câu 1 đến câu 12. Mỗi câu hỏi thí sinh chỉ chọn 1 phương án.", base_size=11, base_italic=True)

    # Q1
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 1.* Cellulose trinitrate được điều chế từ cellulose và nitric acid đặc có xúc tác sulfuric acid đặc, nóng. Để có 44,55 kg cellulose trinitrate, cần dùng dung dịch chứa m kg nitric acid (hiệu suất phản ứng đạt 90%). Giá trị của m là")
    q1_ch = ["28,350 kg.", "21,234 kg.", "25,515 kg.", "31,500 kg."]
    add_mcq_choices_grid(doc, q1_ch, correct_idx=3, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Phương trình: [C~6~H~7~O~2~(OH)~3~]~n~ + 3n HNO~3~ —(H~2~SO~4~ đặc, t°)→ [C~6~H~7~O~2~(ONO~2~)~3~]~n~ + 3n H~2~O\nKhối lượng mol: M~trinitrate~ = 297 g/mol; 3·M~HNO3~ = 3 × 63 = 189 g/mol.\nLượng HNO~3~ theo lý thuyết: m~LT~ = 44,55 × (189 / 297) = 28,35 kg.\nDo hiệu suất đạt 90%, lượng HNO~3~ thực tế cần là:\nm = 28,35 / 90% = 28,35 / 0,9 = *31,500 kg*.\n*Chọn D.*")

    # Q2
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 2.* Methyl salicylate (chất X) là sản phẩm tự nhiên của rất nhiều loại cây, thường được kết hợp với các loại tinh dầu khác dùng làm thuốc bôi ngoài da, thuốc xoa bóp, cao dán giảm đau, chống viêm. X có công thức cấu tạo như sau:")
    add_image_centered(doc, "fig_cau2_methyl_salicylate.png", width_cm=3.8)
    p_sub2 = doc.add_paragraph()
    p_sub2.paragraph_format.space_before = Pt(1)
    p_sub2.paragraph_format.space_after = Pt(2)
    render_rich_text(p_sub2, "Cho các phát biểu sau:\n(a) Công thức phân tử của X là C~8~H~8~O~3~.\n(b) Phân tử X chứa 31,58% oxygen về khối lượng.\n(c) a mol X phản ứng tối đa với a mol Na, sinh ra a mol H~2~.\n(d) a mol X phản ứng tối đa với 2a mol NaOH.\n(e) X là hợp chất hữu cơ tạp chức, chứa đồng thời chức ester và chức alcohol.\nSố phát biểu *sai* là")
    q2_ch = ["1", "2", "4", "3"]
    add_mcq_choices_grid(doc, q2_ch, correct_idx=1, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "- (a) Đúng: CTPT của HOC~6~H~4~COOCH~3~ là C~8~H~8~O~3~ (M = 152 g/mol).\n- (b) Đúng: %m~O~ = (48 / 152) × 100% ≈ 31,58%.\n- (c) Sai: X chỉ có 1 nhóm -OH phenol nên 1 mol X tác dụng Na chỉ sinh ra 0,5 mol H~2~ (a mol X sinh ra 0,5a mol H~2~).\n- (d) Đúng: 1 mol X phản ứng với 1 NaOH (ở chức ester) và 1 NaOH (ở nhóm phenol) → tỉ lệ 1 : 2.\n- (e) Sai: Nhóm -OH liên kết trực tiếp với vòng benzene là nhóm phenol, không phải alcohol.\nVậy có 2 phát biểu sai là (c) và (e).\n*Chọn B.*")

    # Q3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 3.* Cho các chất sau: glucose, fructose, maltose, saccharose, cellulose và tinh bột. Số chất tạo phức màu xanh lam với Cu(OH)~2~ trong môi trường kiềm là")
    q3_ch = ["3", "2", "4", "1"]
    add_mcq_choices_grid(doc, q3_ch, correct_idx=2, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Các polyalcohol có nhiều nhóm -OH kề nhau hòa tan Cu(OH)~2~ ở nhiệt độ thường tạo phức màu xanh lam gồm 4 chất: glucose, fructose, maltose, saccharose. Cellulose và tinh bột không phản ứng.\n*Chọn C.*")

    # Q4
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 4.* Cho các chất: ethanol, acetic acid, methyl fomate, propionic acid. Chất nào có nhiệt độ sôi thấp nhất?")
    q4_ch = ["methyl fomate.", "acetic acid.", "ethanol.", "propionic acid."]
    add_mcq_choices_grid(doc, q4_ch, correct_idx=0, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Carboxylic acid và alcohol đều có liên kết hydrogen liên phân tử bền. Ester (methyl formate) không có liên kết hydrogen giữa các phân tử nên có nhiệt độ sôi thấp nhất (khoảng 31,5°C).\n*Chọn A.*")

    # Q5
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 5.* Cho các chất: CH~3~[CH~2~]~14~COONa, CH~3~[CH~2~]~10~CH~2~OSO~3~Na, CH~3~[CH~2~]~7~CH=CH[CH~2~]~7~COONa, CH~3~[CH~2~]~16~COOK, CH~3~COONa, (C~15~H~31~COO)~3~C~3~H~5~. Số chất có thể là thành phần chính của xà phòng là")
    q5_ch = ["3", "4", "1", "2"]
    add_mcq_choices_grid(doc, q5_ch, correct_idx=0, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Thành phần chính của xà phòng là muối sodium hoặc potassium của acid béo. Có 3 chất thỏa mãn:\n1. CH~3~[CH~2~]~14~COONa (sodium palmitate)\n2. CH~3~[CH~2~]~7~CH=CH[CH~2~]~7~COONa (sodium oleate)\n3. CH~3~[CH~2~]~16~COOK (potassium stearate).\n(CH~3~[CH~2~]~10~CH~2~OSO~3~Na là chất giặt rửa tổng hợp; CH~3~COONa là muối mạch ngắn; chất cuối là chất béo).\n*Chọn A.*")

    # Q6
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 6.* Để rửa sạch chai lọ đựng dung dịch aniline, nên dùng cách nào sau đây?")
    q6_ch = [
        "Rửa bằng dung dịch NaOH sau đó rửa lại bằng nước.",
        "Rửa bằng dung dịch HCl sau đó rửa lại bằng nước.",
        "Rửa bằng xà phòng.",
        "Rửa bằng nước."
    ]
    add_mcq_choices_grid(doc, q6_ch, correct_idx=1, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Aniline ít tan trong nước và bám vào thành chai lọ. Dùng dung dịch HCl để chuyển aniline thành muối C~6~H~5~NH~3~Cl tan tốt trong nước, sau đó tráng lại bằng nước sạch.\n*Chọn B.*")

    # Q7
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 7.* Cho sơ đồ chuyển hóa sau:\nX(l) + 2NaOH(aq) —→ CH~2~(COONa)~2~(aq) + CH~3~OH(aq) + C~2~H~5~OH(aq).\nNhận xét nào sau đây *không đúng* về chất X trong phản ứng trên?")
    p_img7 = doc.add_paragraph()
    p_img7.paragraph_format.space_before = Pt(1)
    p_img7.paragraph_format.space_after = Pt(1)
    render_rich_text(p_img7, "A. Công thức cấu tạo của (X) được biểu diễn như sau:")
    add_image_centered(doc, "fig_cau7_malonate.png", width_cm=7.5)
    q7_ch = [
        "Công thức cấu tạo của (X) như hình trên.",
        "X là ester no có hai nhóm chức có công thức phân tử C~6~H~10~O~4~.",
        "Tên của X là ethyl methyl malonate.",
        "X có nhiệt độ sôi cao vì có liên kết hydrogen giữa các phân tử."
    ]
    add_mcq_choices_grid(doc, q7_ch, correct_idx=3, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Ester X không có nguyên tử H linh động gắn với O nên không tạo được liên kết hydrogen liên phân tử, do đó nhiệt độ sôi của X thấp hơn nhiều so với alcohol hay acid tương ứng.\n*Chọn D.*")

    # Q8
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 8.* Tiến hành thí nghiệm điều chế ethyl acetate theo các bước sau:\n- Bước 1: Cho 1 mL C~2~H~5~OH, 1 mL CH~3~COOH và vài giọt dung dịch H~2~SO~4~ đặc vào ống nghiệm.\n- Bước 2: Lắc đều ống nghiệm, đun cách thủy (trong nồi nước nóng) khoảng 5 - 6 phút ở 65 - 70°C.\n- Bước 3: Làm lạnh, sau đó rót 2 mL dung dịch NaCl bão hòa vào ống nghiệm.\nPhát biểu nào sau đây là *sai*?")
    q8_ch = [
        "H~2~SO~4~ đặc có vai trò vừa làm chất xúc tác vừa làm tăng hiệu suất tạo sản phẩm.",
        "Sau bước 3, chất lỏng trong ống nghiệm tách thành hai lớp.",
        "Mục đích chính của việc thêm dung dịch NaCl bão hòa là để tránh phân hủy sản phẩm.",
        "Sau bước 2, trong ống nghiệm vẫn còn C~2~H~5~OH và CH~3~COOH."
    ]
    add_mcq_choices_grid(doc, q8_ch, correct_idx=2, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Mục đích chính của việc thêm dung dịch NaCl bão hòa là để làm tăng khối lượng riêng của lớp dung dịch nước và làm giảm độ tan của ethyl acetate, giúp ethyl acetate nhẹ hơn tách lớp rõ rệt nổi lên trên.\n*Chọn C.*")

    # Q9
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 9.* Cho sơ đồ chuyển hóa:\nX (C~10~H~16~O~7~N~2~) —(NaOH dư)—→ Y —(HCl dư)—→ Z\nBiết X là dipeptide của một α-amino acid T có cấu tạo không phân nhánh; mỗi mũi tên ứng với một phương trình hóa học của phản ứng giữa hai chất tương ứng. Phát biểu nào sau đây *đúng*?")
    q9_ch = [
        "Phần trăm khối lượng của nguyên tố chlorine trong phân tử chất Z chiếm 19,452%.",
        "Ở điều kiện thường, chất T dễ tan trong nước và có nhiệt độ nóng chảy cao.",
        "Chất Y dùng làm gia vị thức ăn (gọi là mì chính hay bột ngọt).",
        "X tác dụng tối đa với dung dịch NaOH theo tỉ lệ 1 : 3."
    ]
    add_mcq_choices_grid(doc, q9_ch, correct_idx=1, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Từ X: 2 T = C~10~H~16~O~7~N~2~ + H~2~O = C~10~H~18~O~8~N~2~ → T là C~5~H~9~O~4~N (glutamic acid).\nGlutamic acid (T) là amino acid dạng ion lưỡng cực nên ở điều kiện thường là chất rắn kết tinh, dễ tan trong nước và có nhiệt độ nóng chảy cao.\n*Chọn B.*")

    # Q10
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 10.* Carbohydrate (X) có công thức cấu tạo dưới đây:")
    add_image_centered(doc, "fig_cau10_maltose.png", width_cm=8.5)
    p_sub10 = doc.add_paragraph()
    p_sub10.paragraph_format.space_before = Pt(1)
    p_sub10.paragraph_format.space_after = Pt(2)
    render_rich_text(p_sub10, "Nhận định nào sau đây là *đúng* khi nói về (X)?")
    q10_ch = [
        "(X) có thể là saccharose.",
        "(X) còn được gọi là đường mạch nha được sản xuất từ ngũ cốc.",
        "(X) không có tính khử.",
        "(X) được cấu tạo từ 1 đơn vị α-glucose và 1 đơn vị β-fructose qua liên kết α-1,4-glycoside."
    ]
    add_mcq_choices_grid(doc, q10_ch, correct_idx=1, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Cấu tạo gồm 2 gốc α-glucose liên kết với nhau qua liên kết α-1,4-glycoside, ở đầu C1 còn nhóm -OH hemiacetal. Đây là maltose (đường mạch nha, được sản xuất từ tinh bột ngũ cốc).\n*Chọn B.*")

    # Q11
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 11.* Glutamic acid có vai trò quan trọng trong quá trình xây dựng cấu trúc tế bào của con người. Ngoài ra, muối monosodium glutamate còn được dùng chế biến gia vị thức ăn (bột ngọt hay mì chính). Glutamic acid có cấu trúc như hình vẽ bên dưới và có điểm đẳng điện pI = 3,2 (pI là giá trị pH mà khi đó amino acid có nồng độ ion lưỡng cực là cực đại. Khi pH < pI thì amino acid đó tồn tại chủ yếu ở dạng cation, còn khi pH > pI thì amino acid đó tồn tại chủ yếu ở dạng anion)")
    add_image_centered(doc, "fig_cau11_glutamic_acid.png", width_cm=6.5)
    p_sub11 = doc.add_paragraph()
    p_sub11.paragraph_format.space_before = Pt(1)
    p_sub11.paragraph_format.space_after = Pt(2)
    render_rich_text(p_sub11, "Cho các phát biểu sau:\n(a) Glutamic acid thuộc loại hợp chất hữu cơ tạp chức, trong phân tử chứa hai loại nhóm chức.\n(b) Tên thay thế của glutamic acid là 2-aminopentane-1,5-dioic acid.\n(c) Trong dung dịch pH = 3,2, glutamic acid tồn tại chủ yếu ở dạng HOOC–CH~2~–CH~2~–CH(NH~2~)–COO^-^.\n(d) Trong dung dịch pH = 6, có thể tách hỗn hợp gồm glutamic acid và lysine (pI = 9,7) bằng phương pháp điện di.\nSố phát biểu *đúng* là")
    q11_ch = ["1", "4", "2", "3"]
    add_mcq_choices_grid(doc, q11_ch, correct_idx=3, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "- (a) Đúng: Chứa nhóm -COOH và -NH~2~ (tạp chức).\n- (b) Đúng: Mạch 5C chứa 2 nhóm -COOH ở C1, C5 và -NH~2~ ở C2.\n- (c) Sai: Ở pH = pI = 3,2 dạng chủ yếu là ion lưỡng cực HOOC-CH~2~-CH~2~-CH(NH~3~^+^)-COO^-^.\n- (d) Đúng: Ở pH = 6, glutamic acid (pH > pI) tồn tại dạng anion di chuyển về cực dương; lysine (pH < pI = 9,7) tồn tại dạng cation di chuyển về cực âm → tách được bằng điện di.\nVậy có 3 phát biểu đúng: (a), (b), (d).\n*Chọn D.*")

    # Q12
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 12.* Cho các phát biểu sau:\n(a) Chất béo được gọi chung là triglyceride.\n(b) Chất béo nhẹ hơn nước, không tan trong nước nhưng tan nhiều trong dung môi hữu cơ.\n(c) Phản ứng thủy phân chất béo trong môi trường acid là phản ứng thuận nghịch.\n(d) Ở nhiệt độ thường, chất béo chứa nhiều gốc acid béo không no thường ở thể rắn.\nSố phát biểu *đúng* là")
    q12_ch = ["1.", "3.", "2.", "4."]
    add_mcq_choices_grid(doc, q12_ch, correct_idx=1, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "- (a), (b), (c) đều đúng.\n- (d) Sai: Chất béo chứa nhiều gốc acid béo không no ở thể lỏng (dầu thực vật).\nVậy có 3 phát biểu đúng: (a), (b), (c).\n*Chọn B.*")

    # ==================== PHẦN II: ĐÚNG / SAI ====================
    add_section_header(doc, "PHẦN II (3 điểm – 3 câu): Câu trắc nghiệm đúng sai")
    p_p2_intro = doc.add_paragraph()
    p_p2_intro.paragraph_format.space_before = Pt(1)
    p_p2_intro.paragraph_format.space_after = Pt(3)
    render_rich_text(p_p2_intro, "Thí sinh trả lời từ câu 1 đến câu 3. Trong mỗi ý a), b), c), d) ở mỗi câu thí sinh chọn đúng hoặc sai.", base_size=11, base_italic=True)

    # Câu 1 Phần II
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 1.* Triglyceride đóng vai trò là nguồn cung cấp năng lượng và chuyên chở các chất béo trong quá trình trao đổi chất. Cho triglyceride X có công thức cấu tạo như hình sau:")
    add_image_centered(doc, "fig_cau1_p2_triglyceride.png", width_cm=14.0)

    tf1_items = [
        ("a", "X được cấu tạo từ các acid béo là palmitic acid, oleic acid và linoleic acid.", "Sai", "Gốc (3) có 3 liên kết đôi C=C nên là linolenic acid (chứ không phải linoleic acid có 2 liên kết đôi C=C)."),
        ("b", "Thủy phân hoàn toàn 427 kg triglyceride X bằng lượng NaOH dư thu được 441 kg muối của acid béo.", "Đúng", "M~X~ = 854 g/mol; Khối lượng 3 muối sodium = 278 + 304 + 300 = 882 g/mol. Với 427 kg X (bằng nửa mol): m~muối~ = 427 × (882 / 854) = 441 kg."),
        ("c", "Trong các acid béo cấu tạo nên X có một acid béo thuộc loại acid béo omega -8.", "Sai", "Oleic acid là acid béo omega-9 (18 - 9 = 9); linolenic acid là acid béo omega-3 (18 - 15 = 3). Không có acid béo omega-8."),
        ("d", "Tổng số nguyên tử trong 1 phân tử X là 158.", "Sai", "CTPT của X là C~55~H~98~O~6~ có tổng số nguyên tử = 55 + 98 + 6 = 159 nguyên tử.")
    ]
    for tag, text_item, ans_tf, exp_tf in tf1_items:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.space_before = Pt(1)
        p_item.paragraph_format.space_after = Pt(1)
        p_item.paragraph_format.left_indent = Cm(0.5)
        prefix = f"{tag}) "
        if is_teacher:
            prefix += f"[{ans_tf}] "
            r_pre = p_item.add_run(prefix)
            set_run_style(r_pre, font_size=11, bold=True, color=COLOR_CORRECT if ans_tf == "Đúng" else COLOR_PRIMARY)
        else:
            r_pre = p_item.add_run(prefix)
            set_run_style(r_pre, font_size=11, bold=True)
        render_rich_text(p_item, text_item, base_size=11)
        if is_teacher:
            p_exp = doc.add_paragraph()
            p_exp.paragraph_format.space_before = Pt(0)
            p_exp.paragraph_format.space_after = Pt(2)
            p_exp.paragraph_format.left_indent = Cm(1.0)
            render_rich_text(p_exp, f"→ *{ans_tf}*: {exp_tf}", base_size=10, base_color=COLOR_GRAY, base_italic=True)

    # Câu 2 Phần II
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 2.* Các peptide có phản ứng thủy phân trong môi trường acid và môi trường kiềm, ngoài ra các peptide có từ 2 liên kết peptide trở lên phản ứng với Cu(OH)~2~ trong môi trường kiềm tạo thành phức chất màu tím đặc trưng, gọi là phản ứng màu biuret.")

    tf2_items = [
        ("a", "Dung dịch của các polypeptide hoà tan Cu(OH)~2~ cho dung dịch có màu xanh tím.", "Đúng", "Phản ứng màu biuret tạo phức màu tím (xanh tím) đặc trưng."),
        ("b", "Thủy phân hoàn toàn 0,1 mol Glu–Ala–Lys cần vừa đủ 300 mL dung dịch KOH 1M.", "Sai", "Glu-Ala-Lys có 2 liên kết peptide + 1 nhóm -COOH đầu C + 1 nhóm -COOH ở nhánh Glu → phản ứng tối đa với 4 KOH. Cần 0,4 mol KOH (400 mL)."),
        ("c", "Thủy phân hoàn toàn 2,17 gam tripeptide mạch hở X (được tạo nên từ hai α-amino acid có công thức dạng H~2~NC~x~H~y~COOH) bằng dung dịch NaOH dư, thu được 3,19 gam muối. Mặt khác, thủy phân hoàn toàn 3,255 gam X bằng dung dịch HCl dư, thu được 5,4375 gam muối.", "Đúng", "Từ TN1: n~X~ = (3,19 - 2,17) / (3×40 - 18) = 0,01 mol → M~X~ = 217 g/mol. Ở TN2: 3,255 g ứng với 0,015 mol X → m~muối~ = 3,255 + 0,015×2×18 + 0,015×3×36,5 = 5,4375 gam."),
        ("d", "Gly–Ala không có phản ứng màu biuret với Cu(OH)~2~.", "Đúng", "Dipeptide chỉ có 1 liên kết peptide nên không có phản ứng màu biuret.")
    ]
    for tag, text_item, ans_tf, exp_tf in tf2_items:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.space_before = Pt(1)
        p_item.paragraph_format.space_after = Pt(1)
        p_item.paragraph_format.left_indent = Cm(0.5)
        prefix = f"{tag}) "
        if is_teacher:
            prefix += f"[{ans_tf}] "
            r_pre = p_item.add_run(prefix)
            set_run_style(r_pre, font_size=11, bold=True, color=COLOR_CORRECT if ans_tf == "Đúng" else COLOR_PRIMARY)
        else:
            r_pre = p_item.add_run(prefix)
            set_run_style(r_pre, font_size=11, bold=True)
        render_rich_text(p_item, text_item, base_size=11)
        if is_teacher:
            p_exp = doc.add_paragraph()
            p_exp.paragraph_format.space_before = Pt(0)
            p_exp.paragraph_format.space_after = Pt(2)
            p_exp.paragraph_format.left_indent = Cm(1.0)
            render_rich_text(p_exp, f"→ *{ans_tf}*: {exp_tf}", base_size=10, base_color=COLOR_GRAY, base_italic=True)

    # Câu 3 Phần II
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 3.* Khi nấu rượu hoặc làm bánh mì từ bột bánh mì có chứa men (một loại nấm) và đường. Khi bột được để ở nơi ấm áp, các tế bào nấm men sẽ ăn đường để lấy năng lượng. Enzyme trong nấm men xúc tác cho phản ứng được gọi là quá trình lên men sinh ra ethanol và khí carbon dioxide:\n(C~6~H~10~O~5~)~n~ —(+H~2~O, enzyme)—→ n C~6~H~12~O~6~ —(enzyme)—→ 2n C~2~H~5~OH + 2n CO~2~")

    tf3_items = [
        ("a", "Để điều chế 5 lít ethyl alcohol 46° cần 5,4 kg bột bánh mì trên (chứa 75% tinh bột, còn lại là tạp chất trơ). Biết hiệu suất của cả quá trình là 80% và khối lượng riêng của ethyl alcohol nguyên chất là 0,8 g/mL.", "Đúng", "m~C2H5OH~ = 5000 × 0,46 × 0,8 = 1840 g (40 mol). m~tinh bột(LT)~ = 20 × 162 = 3240 g = 3,24 kg. Với H=80%: m~tinh bột(TT)~ = 3,24 / 0,8 = 4,05 kg. Lượng bột bánh mì = 4,05 / 0,75 = 5,4 kg."),
        ("b", "Nếu đun nóng hỗn hợp lên men để tách C~2~H~5~OH, đó là phương pháp chưng cất dựa trên sự khác nhau về độ tan.", "Sai", "Phương pháp chưng cất dựa trên sự khác nhau về nhiệt độ sôi của các chất."),
        ("c", "Lên men rượu luôn xảy ra tối ưu ở khoảng 30–35°C, nếu nhiệt độ thấp nấm men sẽ chết.", "Sai", "Ở nhiệt độ thấp nấm men chỉ bị giảm hoạt tính/ức chế chứ không chết."),
        ("d", "Nếu cung cấp nhiều oxy cho môi trường lên men, hiệu suất tạo ethanol sẽ tăng.", "Sai", "Quá trình lên men rượu là kị khí; cung cấp oxy sẽ chuyển sang hô hấp hiếu khí tạo CO~2~ và H~2~O hoặc oxy hóa ethanol thành acetic acid, làm giảm hiệu suất tạo ethanol.")
    ]
    for tag, text_item, ans_tf, exp_tf in tf3_items:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.space_before = Pt(1)
        p_item.paragraph_format.space_after = Pt(1)
        p_item.paragraph_format.left_indent = Cm(0.5)
        prefix = f"{tag}) "
        if is_teacher:
            prefix += f"[{ans_tf}] "
            r_pre = p_item.add_run(prefix)
            set_run_style(r_pre, font_size=11, bold=True, color=COLOR_CORRECT if ans_tf == "Đúng" else COLOR_PRIMARY)
        else:
            r_pre = p_item.add_run(prefix)
            set_run_style(r_pre, font_size=11, bold=True)
        render_rich_text(p_item, text_item, base_size=11)
        if is_teacher:
            p_exp = doc.add_paragraph()
            p_exp.paragraph_format.space_before = Pt(0)
            p_exp.paragraph_format.space_after = Pt(2)
            p_exp.paragraph_format.left_indent = Cm(1.0)
            render_rich_text(p_exp, f"→ *{ans_tf}*: {exp_tf}", base_size=10, base_color=COLOR_GRAY, base_italic=True)

    # ==================== PHẦN III: TRẢ LỜI NGẮN ====================
    add_section_header(doc, "PHẦN III (1 điểm – 4 câu): Câu trắc nghiệm yêu cầu trả lời ngắn")
    p_p3_intro = doc.add_paragraph()
    p_p3_intro.paragraph_format.space_before = Pt(1)
    p_p3_intro.paragraph_format.space_after = Pt(3)
    render_rich_text(p_p3_intro, "Thí sinh trả lời từ câu 1 đến câu 4.", base_size=11, base_italic=True)

    # Câu 1 Phần III
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 1.* Một loại chất béo có chứa 75% tristearine về khối lượng. Xà phòng hóa hoàn toàn 17,8 kg chất béo này trong dung dịch KOH, đun nóng thu được x bánh xà phòng. Biết rằng trong mỗi bánh xà phòng có chứa 63 gam potassium stearate. Giá trị của x là bao nhiêu?")
    if is_teacher:
        add_solution_box(doc, "1. Khối lượng tristearin nguyên chất: m = 17,8 × 75% = 13,35 kg = 13350 gam.\n2. n~tristearin~ = 13350 / 890 = 15 mol.\n3. Phản ứng: (C~17~H~35~COO)~3~C~3~H~5~ + 3KOH → 3 C~17~H~35~COOK + C~3~H~5~(OH)~3~\n→ n~potassium stearate~ = 3 × 15 = 45 mol.\n4. Khối lượng potassium stearate: m = 45 × 322 = 14490 gam.\n5. Số bánh xà phòng: x = 14490 / 63 = *230* bánh.\n*Đáp số: 230*")

    # Câu 2 Phần III
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 2.* Aspirin là một hợp chất được sử dụng làm giảm đau, hạ sốt được điều chế theo phản ứng sau:\n(CH~3~CO)~2~O + HOC~6~H~4~COOH ⇌ CH~3~COOC~6~H~4~COOH + CH~3~COOH\nĐể sản xuất 500 viên thuốc aspirin cần tối thiểu 20,7 gam salicylic acid. Biết rằng mỗi viên thuốc có chứa y mg aspirin và hiệu suất phản ứng đạt 75%. Giá trị của y là bao nhiêu?")
    if is_teacher:
        add_solution_box(doc, "1. n~salicylic acid~ = 20,7 / 138 = 0,15 mol.\n2. n~aspirin (TT)~ = 0,15 × 75% = 0,1125 mol.\n3. Tổng khối lượng aspirin tạo thành: m = 0,1125 × 180 = 20,25 gam = 20250 mg.\n4. Hàm lượng aspirin trong mỗi viên thuốc: y = 20250 / 500 = *40,5 mg*.\n*Đáp số: 40,5*")

    # Câu 3 Phần III
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 3.* Cho các nhận định sau:\n(a) Protein dạng hình cầu không tan trong nước, còn dạng hình sợi tan trong nước.\n(b) Một trong những tính chất hóa học đặc trưng của protein là phản ứng thủy phân.\n(c) Phản ứng của protein với nitrous acid cho sản phẩm màu vàng.\n(d) Sự đông tụ không làm thay đổi cấu tạo ban đầu của protein bị biến đổi.\n(e) Trong cơ thể, enzyme không đóng vai trò là chất xúc tác sinh học.\nSố phát biểu đúng là")
    if is_teacher:
        add_solution_box(doc, "Chỉ có duy nhất phát biểu (b) đúng. (a sai vì protein hình cầu tan tạo keo, hình sợi không tan; c sai vì phản ứng màu xanthoproteic với HNO~3~ đặc; d sai vì làm biến đổi cấu trúc không gian; e sai vì enzyme là chất xúc tác sinh học).\n*Đáp số: 1*")

    # Câu 4 Phần III
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 4.* Để tráng một số lượng gương soi có diện tích bề mặt 35 dm^2^ với độ dày 0,08 μm người ta đun nóng dung dịch chứa 30,6 gam glucose với một lượng dung dịch AgNO~3~ trong amoniac. Biết khối lượng riêng của silver là 10,49 kg/L, hiệu suất phản ứng tráng gương là 80% (tính theo glucose). Có tối đa bao nhiêu chiếc gương soi được sản xuất ra? *(Kết quả làm tròn đến hàng đơn vị)*.")
    if is_teacher:
        add_solution_box(doc, "1. n~glucose~ = 30,6 / 180 = 0,17 mol.\n2. Phản ứng: C~6~H~12~O~6~ → 2Ag\n→ n~Ag~ = 0,17 × 2 × 80% = 0,272 mol → m~Ag~ = 0,272 × 108 = 29,376 gam.\n3. D = 10,49 kg/L = 10,49 g/cm^3^ → V~tổng~ = 29,376 / 10,49 ≈ 2,80038 cm^3^.\n4. Thể tích Ag cần cho 1 chiếc gương: V~1~ = 3500 cm^2^ × (0,08 × 10^-4^ cm) = 0,028 cm^3^.\n5. Số lượng gương sản xuất được: N = 2,80038 / 0,028 ≈ *100* chiếc.\n*Đáp số: 100*")

    # ==================== PHẦN B: TỰ LUẬN ====================
    add_section_header(doc, "B. PHẦN TỰ LUẬN (3,0 điểm)")

    # Câu 1 Tự luận
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 1 (1,5 điểm).* Viết công thức cấu tạo, gọi tên thay thế và gốc chức của amin bậc 2, bậc 3 có công thức phân tử C~4~H~11~N?")
    if is_teacher:
        add_solution_box(doc, "Các amin bậc 2 và bậc 3 của C~4~H~11~N gồm:\n\n*1. Các amin bậc 2:*\n- CH~3~–NH–CH~2~CH~2~CH~3~: Tên thay thế: N-methylpropan-1-amine; Tên gốc-chức: methylpropylamine.\n- CH~3~–NH–CH(CH~3~)~2~: Tên thay thế: N-methylpropan-2-amine; Tên gốc-chức: isopropylmethylamine.\n- CH~3~CH~2~–NH–CH~2~CH~3~: Tên thay thế: N-ethylethanamine; Tên gốc-chức: diethylamine.\n\n*2. Amin bậc 3:*\n- (CH~3~)~2~N–CH~2~CH~3~: Tên thay thế: N,N-dimethylethanamine; Tên gốc-chức: ethyldimethylamine.")

    # Câu 2 Tự luận
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 2 (1,5 điểm).*\na. So sánh điểm giống và khác nhau giữa xà phòng và chất giặt rửa tổng hợp?\nb. Trình bày cơ chế giặt rửa của chúng?")
    if is_teacher:
        add_solution_box(doc, "a. So sánh xà phòng và chất giặt rửa tổng hợp:\n- *Giống nhau*: Đều có phân tử gồm 2 phần: đầu ưa nước (phân cực/ion) và đuôi kị nước (mạch hydrocarbon dài); đều làm giảm sức căng bề mặt của nước và có tác dụng tẩy rửa dầu mỡ.\n- *Khác nhau*:\n  + Xà phòng: Muối sodium/potassium của acid béo thiên nhiên, sản xuất từ mỡ động vật, dầu thực vật; bị mất tác dụng trong nước cứng do tạo kết tủa với Ca^2+^, Mg^2+^; dễ bị phân hủy sinh học.\n  + Chất giặt rửa tổng hợp: Muối sodium của alkylsulfate hoặc alkylbenzenesulfonate tổng hợp từ dầu mỏ; giặt tốt trong cả nước cứng và môi trường acid; một số loại khó phân hủy sinh học gây ô nhiễm.\n\nb. Cơ chế giặt rửa:\n- Khi giặt, phần đuôi kị nước cắm vào vết dầu mỡ, phần đầu ưa nước hướng ra ngoài tiếp xúc với nước.\n- Dưới tác dụng cơ học (vò, khuấy), vết bẩn bị chia nhỏ thành các hạt nhũ tương li ti lơ lửng trong nước và bị cuốn trôi theo dòng nước xả.")

    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_end.paragraph_format.space_before = Pt(10)
    p_end.paragraph_format.space_after = Pt(4)
    r_end = p_end.add_run("--- HẾT ---")
    set_run_style(r_end, font_size=11, bold=True)

    filename = "De_Kiem_Tra_Hoa_12_Ma1011_GiaoVien.docx" if is_teacher else "De_Kiem_Tra_Hoa_12_Ma1011.docx"
    save_path = os.path.join(OUT_DIR, filename)
    doc.save(save_path)
    print(f"Saved: {save_path}")

if __name__ == "__main__":
    build_exam_document(is_teacher=False)
    build_exam_document(is_teacher=True)
