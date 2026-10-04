import docx
from docx.shared import Pt, RGBColor
import re

def render_rich_text(p, text, base_font="Times New Roman", base_size=12, base_bold=False, base_italic=False, base_color=None):
    """
    Renders text with ~subscript~, ^superscript^, **bold**, *italic* tags into docx paragraph runs.
    """
    pattern = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|~[^~]+~|\^[^^]+\^)')
    tokens = pattern.split(text)

    for token in tokens:
        if not token:
            continue
        
        is_bold = base_bold
        is_italic = base_italic
        is_sub = False
        is_sup = False
        content = token

        if token.startswith('**') and token.endswith('**'):
            is_bold = True
            content = token[2:-2]
        elif token.startswith('*') and token.endswith('*'):
            is_italic = True
            content = token[1:-1]
        elif token.startswith('~') and token.endswith('~'):
            is_sub = True
            content = token[1:-1]
        elif token.startswith('^') and token.endswith('^'):
            is_sup = True
            content = token[1:-1]

        run = p.add_run(content)
        run.font.name = base_font
        run.font.size = Pt(base_size)
        run.font.bold = is_bold
        run.font.italic = is_italic
        run.font.subscript = is_sub
        run.font.superscript = is_sup
        if base_color:
            run.font.color.rgb = base_color

doc = docx.Document()
p = doc.add_paragraph()
render_rich_text(p, "Cho x mol phenol (C~6~H~5~OH) tác dụng với Na dư, thấy thoát ra 0,1 mol khí H~2~. Dung tích 1 m^3^ và 120 µg·m^-3^.")

doc.save(r"c:\Antigravity_Thanh\test_rich.docx")
print("Saved test_rich.docx successfully!")
