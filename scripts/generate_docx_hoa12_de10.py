# -*- coding: utf-8 -*-
"""
Script tạo file Word (.docx) chuẩn "FORM ĐỀ KIỂM TRA" cho:
ĐỀ KIỂM TRA SỐ 10 - CHƯƠNG 4: POLYMER (HÓA HỌC 12)
Gồm 2 phiên bản:
1. De_Kiem_Tra_Hoa_12_De10.docx (Bản Học Sinh)
2. De_Kiem_Tra_Hoa_12_De10_GiaoVien.docx (Bản Giáo Viên có lời giải chi tiết)
Tuân thủ 100% quy chuẩn: Header & Footer, Tab Stops (NO TABLE trong đề/lời giải), Sub/Superscript công thức hóa học.
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
IMG_DIR = r"c:\Antigravity_Thanh\San_Pham\images_crop_de10"

COLOR_HF_RED = RGBColor(211, 47, 47)      # Coral Red (#D32F2F)
COLOR_HF_GREEN_HEX = "8BC390"             # Sage Green (#8BC390)
COLOR_PRIMARY = RGBColor(46, 125, 50)     # Dark Green (#2E7D32)
COLOR_CORRECT = RGBColor(198, 40, 40)     # Red for correct answers (#C62828)
COLOR_SOLUTION = RGBColor(21, 101, 192)   # Blue for solutions (#1565C0)
COLOR_GRAY = RGBColor(97, 97, 97)

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def set_run_style(run, font_name="Times New Roman", font_size=12, bold=False, italic=False, subscript=False, superscript=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def add_chem_run(p, text, base_font="Times New Roman", base_size=12, bold=False, italic=False, color=None):
    tokens = re.split(r'(\^\{[^}]+\}|\^[0-9\+\-]+|~\{[^}]+\}|~[0-9]+)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('^'):
            content = tok[2:-1] if tok.startswith('^{') else tok[1:]
            r = p.add_run(content)
            set_run_style(r, font_name=base_font, font_size=base_size, bold=bold, italic=italic, superscript=True, color=color)
        elif tok.startswith('~'):
            content = tok[2:-1] if tok.startswith('~{') else tok[1:]
            r = p.add_run(content)
            set_run_style(r, font_name=base_font, font_size=base_size, bold=bold, italic=italic, subscript=True, color=color)
        else:
            r = p.add_run(tok)
            set_run_style(r, font_name=base_font, font_size=base_size, bold=bold, italic=italic, color=color)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(r'<w:tcBorders %s/>' % nsdecls('w'))
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = parse_xml(r'<%s %s w:val="%s" w:sz="%s" w:space="0" w:color="%s"/>' %
                                (tag, nsdecls('w'), edge_data.get('val', 'single'),
                                 edge_data.get('sz', '4'), edge_data.get('color', 'auto')))
            tcBorders.append(element)
        else:
            tag = 'w:{}'.format(edge)
            element = parse_xml(r'<%s %s w:val="none"/>' % (tag, nsdecls('w')))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def set_cell_margins(cell, top=50, bottom=50, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(r'<w:tcMar %s><w:top w:w="%d" w:type="dxa"/><w:bottom w:w="%d" w:type="dxa"/><w:left w:w="%d" w:type="dxa"/><w:right w:w="%d" w:type="dxa"/></w:tcMar>' % (nsdecls('w'), top, bottom, left, right))
    tcPr.append(tcMar)

def setup_page_and_header_footer(doc):
    for sec in doc.sections:
        sec.top_margin = Cm(1.5)
        sec.bottom_margin = Cm(1.5)
        sec.left_margin = Cm(1.5)
        sec.right_margin = Cm(1.5)
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.header_distance = Cm(0.6)
        sec.footer_distance = Cm(0.6)
        
        # --- HEADER ---
        hdr = sec.header
        hdr_p = hdr.paragraphs[0]
        hdr_p.text = ""
        hdr_p.paragraph_format.space_before = Pt(0)
        hdr_p.paragraph_format.space_after = Pt(0)
        
        tbl_hdr = hdr.add_table(rows=1, cols=2, width=Cm(18.0))
        tbl_hdr.autofit = False
        tbl_hdr.columns[0].width = Cm(9.0)
        tbl_hdr.columns[1].width = Cm(9.0)
        
        c0 = tbl_hdr.cell(0, 0)
        c1 = tbl_hdr.cell(0, 1)
        c0.width = Cm(9.0)
        c1.width = Cm(9.0)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run("Tài liệu lưu hành nội bộ")
        set_run_style(r0, font_name="Times New Roman", font_size=10, italic=True, color=COLOR_HF_RED)
        
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(2)
        r1a = p1.add_run("Thầy ")
        set_run_style(r1a, font_name="Times New Roman", font_size=10, color=COLOR_HF_RED)
        r1b = p1.add_run("TRẦN VĂN THẠNH")
        set_run_style(r1b, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED)
        r1c = p1.add_run(" - 0777.470.803")
        set_run_style(r1c, font_name="Times New Roman", font_size=10, color=COLOR_HF_RED)
        
        set_cell_border(c0, bottom={'val': 'single', 'sz': '6', 'color': COLOR_HF_GREEN_HEX})
        set_cell_border(c1, bottom={'val': 'single', 'sz': '6', 'color': COLOR_HF_GREEN_HEX})
        set_cell_margins(c0, top=0, bottom=40, left=0, right=0)
        set_cell_margins(c1, top=0, bottom=40, left=0, right=0)
        
        # --- FOOTER ---
        ftr = sec.footer
        ftr_p = ftr.paragraphs[0]
        ftr_p.text = ""
        ftr_p.paragraph_format.space_before = Pt(0)
        ftr_p.paragraph_format.space_after = Pt(0)
        
        tbl_ftr = ftr.add_table(rows=1, cols=2, width=Cm(18.0))
        tbl_ftr.autofit = False
        tbl_ftr.columns[0].width = Cm(13.5)
        tbl_ftr.columns[1].width = Cm(4.5)
        
        fc0 = tbl_ftr.cell(0, 0)
        fc1 = tbl_ftr.cell(0, 1)
        fc0.width = Cm(13.5)
        fc1.width = Cm(4.5)
        
        fp0 = fc0.paragraphs[0]
        fp0.paragraph_format.space_before = Pt(2)
        fp0.paragraph_format.space_after = Pt(0)
        fr0 = fp0.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
        set_run_style(fr0, font_name="Times New Roman", font_size=9.5, italic=True, color=COLOR_HF_RED)
        
        fp1 = fc1.paragraphs[0]
        fp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp1.paragraph_format.space_before = Pt(2)
        fp1.paragraph_format.space_after = Pt(0)
        add_page_number_field(fp1, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED)
        
        set_cell_border(fc0, top={'val': 'single', 'sz': '6', 'color': COLOR_HF_GREEN_HEX})
        set_cell_border(fc1, top={'val': 'single', 'sz': '6', 'color': COLOR_HF_GREEN_HEX})
        set_cell_margins(fc0, top=40, bottom=0, left=0, right=0)
        set_cell_margins(fc1, top=40, bottom=0, left=0, right=0)

def add_exam_title_block(doc, is_teacher=False):
    lines_info = [
        ("SỞ GD&ĐT THỪA THIÊN HUẾ", "ĐỀ KIỂM TRA CHƯƠNG 4: POLYMER", True, True),
        ("TRƯỜNG THPT HAI BÀ TRƯNG", "MÔN: HÓA HỌC 12 - NĂM HỌC 2026 - 2027", False, True),
        ("ĐỀ CHÍNH THỨC", "Thời gian làm bài: 50 phút (không kể thời gian phát đề)", True, False, False, True),
        ("(Đề thi gồm 04 trang)", "MÃ ĐỀ THI: 110", False, True, True, False)
    ]
    
    for item in lines_info:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
        
        r_l = p.add_run(item[0])
        set_run_style(r_l, font_name="Times New Roman", font_size=11, bold=item[2], italic=item[4] if len(item)>4 else False)
        
        p.add_run("\t")
        
        r_r = p.add_run(item[1])
        set_run_style(r_r, font_name="Times New Roman", font_size=11, bold=item[3], italic=item[5] if len(item)>5 else False)

    if is_teacher:
        p_note = doc.add_paragraph()
        p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_note.paragraph_format.space_before = Pt(2)
        p_note.paragraph_format.space_after = Pt(2)
        r_note = p_note.add_run("(BẢN GIÁO VIÊN -- CÓ LỜI GIẢI CHI TIẾT)")
        set_run_style(r_note, font_name="Times New Roman", font_size=11.5, bold=True, color=COLOR_HF_RED)

    # Student info line
    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(3)
    p_info.paragraph_format.space_after = Pt(1)
    r_info = p_info.add_run("Họ, tên thí sinh: .......................................................................... Lớp: ............. SBD: .............................")
    set_run_style(r_info, font_name="Times New Roman", font_size=11, italic=True)

    # Atomic mass note
    p_mass = doc.add_paragraph()
    p_mass.paragraph_format.space_before = Pt(0)
    p_mass.paragraph_format.space_after = Pt(4)
    r_mass = p_mass.add_run("(Cho biết nguyên tử khối: H = 1; C = 12; N = 14; O = 16; Na = 23; Al = 27; S = 32; Cl = 35,5; K = 39; Ca = 40; Fe = 56; Cu = 64; Zn = 65; Br = 80; Ag = 108; Ba = 137).")
    set_run_style(r_mass, font_name="Times New Roman", font_size=9.5, italic=True, color=COLOR_GRAY)

def add_part_heading(doc, part_title, part_instruction):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    r_title = p.add_run(part_title)
    set_run_style(r_title, font_name="Times New Roman", font_size=11.5, bold=True, color=COLOR_PRIMARY)
    
    if part_instruction:
        r_inst = p.add_run(" " + part_instruction)
        set_run_style(r_inst, font_name="Times New Roman", font_size=11, italic=True)

def add_solution_paragraph(doc, solution_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.6)
    
    r_label = p.add_run("Lời giải chi tiết: ")
    set_run_style(r_label, font_name="Times New Roman", font_size=10.5, bold=True, color=COLOR_SOLUTION)
    
    add_chem_run(p, solution_text, base_font="Times New Roman", base_size=10.5, italic=True)

# ----------------- DỮ LIỆU ĐỀ KIỂM TRA SỐ 10 -----------------

MCQ_QUESTIONS = [
    {
        "num": 1,
        "text": "Polymer nào sau đây thuộc loại polymer thiên nhiên?",
        "options": ["A. Tơ visco.", "B. Poly(vinyl chloride).", "C. Tinh bột.", "D. Polyethylene."],
        "correct": "C",
        "sol": "Chọn C. Tinh bột và tơ tằm, cellulose... là các polymer thiên nhiên có sẵn trong thực vật hoặc động vật. Tơ visco là bán tổng hợp; PVC và polyethylene là polymer tổng hợp."
    },
    {
        "num": 2,
        "text": "Polymer nào sau đây có công thức (-CH~2~-CH(CN)-)~n~?",
        "options": ["A. Poly(methyl methacrylate).", "B. Polyethylene.", "C. Polyacrylonitrile.", "D. Poly(vinyl chloride)."],
        "correct": "C",
        "sol": "Chọn C. (-CH~2~-CH(CN)-)~n~ là polyacrylonitrile (hay polyaacrilonitrin, dùng chế tạo tơ nitron/olon). PE là (-CH~2~-CH~2~-)~n~; PVC là (-CH~2~-CHCl-)~n~."
    },
    {
        "num": 3,
        "text": "Poly(vinyl chloride) được điều chế trực tiếp từ monomer nào sau đây?",
        "options": ["A. CH~2~=CH-CN.", "B. CH~2~=CH~2~.", "C. CH~2~=CH-CH=CH~2~.", "D. CH~2~=CH-Cl."],
        "correct": "D",
        "sol": "Chọn D. Poly(vinyl chloride) (PVC) được điều chế trực tiếp bằng phản ứng trùng hợp vinyl chloride: n CH~2~=CH-Cl -> (-CH~2~-CHCl-)~n~."
    },
    {
        "num": 4,
        "text": "Polyacrylonitrile được điều chế trực tiếp từ monomer nào sau đây?",
        "options": ["A. CH~2~=CH-Cl.", "B. CH~2~=CH~2~.", "C. CH~2~=CH-CN.", "D. CH~2~=CH-CH=CH~2~."],
        "correct": "C",
        "sol": "Chọn C. Polyacrylonitrile được điều chế bằng phản ứng trùng hợp vinyl cyanide (acrylonitrile): n CH~2~=CH-CN -> (-CH~2~-CH(CN)-)~n~."
    },
    {
        "num": 5,
        "text": "Polypropylene là chất dẻo được sử dụng phổ biến thứ hai sau polyethylene. Trùng hợp chất nào sau đây thu được polypropylene?",
        "options": ["A. CH~2~=CH-Cl.", "B. CH~2~=CH~2~.", "C. CH~2~=CH-C~6~H~5~.", "D. CH~2~=CH-CH~3~."],
        "correct": "D",
        "sol": "Chọn D. Trùng hợp propylene (propene) CH~2~=CH-CH~3~ thu được polypropylene (PP): n CH~2~=CH-CH~3~ -> (-CH~2~-CH(CH~3~)-)~n~."
    },
    {
        "num": 6,
        "text": "Trùng hợp monomer CH~2~=CH-Cl thu được chất dẻo nào sau đây?",
        "options": ["A. PE.", "B. PP.", "C. PVC.", "D. PS."],
        "correct": "C",
        "sol": "Chọn C. Trùng hợp vinyl chloride CH~2~=CH-Cl thu được poly(vinyl chloride) (PVC)."
    },
    {
        "num": 7,
        "text": "Tơ sợi nào sau đây thuộc loại tơ tự nhiên?",
        "options": ["A. Sợi bông.", "B. Nitron.", "C. Nylon-6,6.", "D. Cellulose acetate."],
        "correct": "A",
        "sol": "Chọn A. Sợi bông (chứa gần 98% cellulose) có sẵn trong tự nhiên thuộc loại tơ thiên nhiên. Nitron và nylon-6,6 là tơ tổng hợp; cellulose acetate là tơ bán tổng hợp (nhân tạo)."
    },
    {
        "num": 8,
        "text": "Trùng hợp chất nào sau đây thu được polyacrylonitrile dùng để sản xuất tơ nitron?",
        "options": ["A. CH~2~=CH-Cl.", "B. CH~2~=CH-CN.", "C. CH~2~=CH~2~.", "D. CH~2~=CH-CH~3~."],
        "correct": "B",
        "sol": "Chọn B. Monomer trùng hợp tạo tơ nitron là acrylonitrile (vinyl cyanide) CH~2~=CH-CN."
    },
    {
        "num": 9,
        "text": "Trùng hợp chất nào sau đây thu được cao su isoprene?",
        "options": ["A. CH~2~=CH-CH=CH~2~.", "B. CH~2~=CCl-CH=CH~2~.", "C. CH~2~=C(CH~3~)-CH=CH~2~.", "D. CH~2~=C(CH~3~)-CCl=CH~2~."],
        "correct": "C",
        "sol": "Chọn C. Isoprene (2-methylbuta-1,3-diene) có công thức cấu tạo là CH~2~=C(CH~3~)-CH=CH~2~. Khi trùng hợp thu được polyisoprene (cao su isoprene)."
    },
    {
        "num": 10,
        "text": "Dãy polymer đều được điều chế từ phản ứng trùng ngưng là",
        "options": [
            "A. Teflon, polyethylene, PVC.",
            "B. Cao su buna, nylon-7, tơ acetate.",
            "C. Nylon-6, poly(vinyl alcohol), thủy tinh plexiglass.",
            "D. Nhựa PPF, nylon-7, tơ lapsan."
        ],
        "correct": "D",
        "sol": "Chọn D. Nhựa PPF (trùng ngưng phenol và formaldehyde), nylon-7 (trùng ngưng 7-aminoheptanoic acid), tơ lapsan (trùng ngưng terephthalic acid và ethylene glycol). A, B, C đều chứa polymer điều chế từ trùng hợp hoặc chế hóa polymer khác."
    },
    {
        "num": 11,
        "text": "Trong các polymer: polyethylene, poly(methyl methacrylate), tinh bột, tơ tằm, tơ nylon-6, số polymer thiên nhiên là",
        "options": ["A. 4.", "B. 2.", "C. 5.", "D. 3."],
        "correct": "B",
        "sol": "Chọn B. Các polymer thiên nhiên là: tinh bột, tơ tằm (gồm 2 polymer). Polyethylene, poly(methyl methacrylate) và nylon-6 là các polymer tổng hợp."
    },
    {
        "num": 12,
        "text": "Cho các polymer sau: (-CH~2~-CH~2~-)~n~; (-CH~2~-CH=CH-CH~2~-)~n~; (-NH-CH~2~-CO-)~n~. Công thức của các monomer để khi trùng hợp hoặc trùng ngưng tạo ra các polymer trên lần lượt là",
        "options": [
            "A. CH~2~=CHCl, CH~3~-CH=CH-CH~3~, CH~3~-CH(NH~2~)-COOH.",
            "B. CH~2~=CH~2~, CH~2~=CH-CH=CH~2~, H~2~N-CH~2~-COOH.",
            "C. CH~2~=CH~2~, CH~3~-CH=C=CH~2~, NH~2~-CH~2~-COOH.",
            "D. CH~2~=CH~2~, CH~3~-CH=CH-CH~3~, H~2~N-CH~2~-CH~2~-COOH."
        ],
        "correct": "B",
        "sol": "Chọn B. (-CH~2~-CH~2~-)~n~ tạo từ CH~2~=CH~2~; (-CH~2~-CH=CH-CH~2~-)~n~ tạo từ buta-1,3-diene CH~2~=CH-CH=CH~2~; (-NH-CH~2~-CO-)~n~ tạo từ glycine H~2~N-CH~2~-COOH."
    },
    {
        "num": 13,
        "text": "Dãy nào sau đây chỉ gồm các polymer tổng hợp?",
        "options": [
            "A. Polypropylene, cellulose, nylon-7, nylon-6,6.",
            "B. Polypropylene, polybutadiene, nylon-7, nylon-6,6.",
            "C. Polypropylene, tinh bột, nylon-7, cao su thiên nhiên.",
            "D. Tinh bột, cellulose, cao su thiên nhiên, polybutadiene."
        ],
        "correct": "B",
        "sol": "Chọn B. Polypropylene, polybutadiene, nylon-7, nylon-6,6 đều là các polymer được tổng hợp từ các hợp chất hóa học. Cellulose, tinh bột, cao su thiên nhiên là polymer thiên nhiên."
    },
    {
        "num": 14,
        "text": "Cho sơ đồ chuyển hoá: Glucose -> (enzyme) X -> (ZnO, 450 oC) Y -> (xt, to, p) Cao su Buna. Hai chất X, Y lần lượt là",
        "options": [
            "A. CH~3~CH~2~OH và CH~3~CHO.",
            "B. CH~3~CH~2~OH và CH~2~=CH~2~.",
            "C. CH~2~CH~2~OH và CH~3~-CH=CH-CH~3~.",
            "D. CH~3~CH~2~OH và CH~2~=CH-CH=CH~2~."
        ],
        "correct": "D",
        "sol": "Chọn D. Glucose lên men rượu tạo ethanol (X: CH~3~CH~2~OH). Sau đó ethanol được chuyển hóa theo phản ứng Lebedev: 2 C~2~H~5~OH -> (ZnO, 450 oC) CH~2~=CH-CH=CH~2~ (Y) + 2 H~2~O + H~2~. Trùng hợp buta-1,3-diene thu được cao su Buna."
    },
    {
        "num": 15,
        "text": "Nhóm các polymer được dùng làm tơ là",
        "options": [
            "A. Polyacrylonitrile, poly(methyl methacrylate).",
            "B. Poly(vinyl chloride), polybuta-1,3-diene.",
            "C. Polyacrylonitrile, poly(hexamethylene adipamide).",
            "D. Poly(hexamethylene adipamide), poly(vinyl chloride)."
        ],
        "correct": "C",
        "sol": "Chọn C. Polyacrylonitrile được dùng làm tơ nitron (olon), poly(hexamethylene adipamide) được dùng làm tơ nylon-6,6. Poly(methyl methacrylate) và PVC dùng làm chất dẻo, polybuta-1,3-diene dùng làm cao su."
    },
    {
        "num": 16,
        "text": "Nhận xét về tính chất vật lí chung của polymer nào dưới đây không đúng?",
        "options": [
            "A. Hầu hết là những chất rắn, không bay hơi, không có nhiệt độ nóng chảy xác định.",
            "B. Khi nóng chảy, đa số polymer cho chất lỏng nhớt, để nguội sẽ rắn lại gọi là chất nhiệt dẻo.",
            "C. Một số polymer không nóng chảy khi đun mà bị phân hủy, gọi là chất nhiệt rắn.",
            "D. Polymer không tan trong nước và trong bất kỳ dung môi nào."
        ],
        "correct": "D",
        "sol": "Chọn D. Phát biểu D sai vì một số polymer có thể tan trong các dung môi hữu cơ thích hợp, ví dụ cao su tan trong benzene, xăng; poly(vinyl alcohol) tan được trong nước nóng..."
    },
    {
        "num": 17,
        "text": "Cho các phát biểu sau:\n(1) Tơ viscose, tơ acetate đều thuộc loại tơ tổng hợp.\n(2) Polyethylene và poly(vinyl chloride) là sản phẩm của phản ứng trùng hợp.\n(3) Tơ nylon-6,6 được điều chế từ hexamethylenediamine và stearic acid.\n(4) Cao su thiên nhiên không tan trong nước cũng như trong xăng, benzene.\n(5) Tơ nitron (olon) được tổng hợp từ vinyl cyanide (acrylonitrile).\nSố phát biểu đúng là",
        "options": ["A. 4.", "B. 3.", "C. 5.", "D. 2."],
        "correct": "D",
        "sol": "Chọn D. Các phát biểu đúng là (2) và (5). (1) sai vì tơ viscose, tơ acetate là tơ bán tổng hợp (nhân tạo). (3) sai vì nylon-6,6 điều chế từ hexamethylenediamine và adipic acid. (4) sai vì cao su thiên nhiên tan được trong xăng, benzene. Vậy có 2 phát biểu đúng."
    },
    {
        "num": 18,
        "text": "Một polymer X được xác định có phân tử khối là 39062,5 amu với hệ số trùng hợp để tạo nên polymer này là 625. Polymer X là",
        "options": [
            "A. Poly(vinyl chloride) (PVC).",
            "B. Polypropylene (PP).",
            "C. Polyethylene (PE).",
            "D. Polystyrene (PS)."
        ],
        "correct": "A",
        "sol": "Chọn A. Khối lượng một mắt xích: M_mắt xích = 39062,5 / 625 = 62,5 amu. Ta có M_PVC = 14 + 13 + 35,5 = 62,5 amu. Vậy polymer X là poly(vinyl chloride) (PVC)."
    }
]

TF_QUESTIONS = [
    {
        "num": 1,
        "intro": "Hộp xốp đựng thực phẩm chế biến sẵn thường làm bằng polystyrene (PS).",
        "items": [
            ("a", "Polystyrene được điều chế bằng phản ứng trùng ngưng.", False, "Sai. Polystyrene được điều chế bằng phản ứng trùng hợp styrene (C~6~H~5~-CH=CH~2~)."),
            ("b", "Có thể sử dụng hộp nhựa PS để đựng thức ăn nóng hoặc cho vào lò vi sóng.", False, "Sai. Hộp xốp PS kém bền nhiệt, ở nhiệt độ cao có thể bị biến dạng và giải phóng độc tố monomer styrene gây hại cho sức khỏe, không dùng đựng đồ quá nóng (>70-80 oC) hay quay lò vi sóng."),
            ("c", "Styrene dùng để điều chế ra PS.", True, "Đúng. Phương trình: n C~6~H~5~-CH=CH~2~ -> (-CH(C~6~H~5~)-CH~2~-)~n~."),
            ("d", "Khi đồng trùng hợp buta-1,3-diene với styrene ta thu được cao su buna-S.", True, "Đúng. Phản ứng đồng trùng hợp: n CH~2~=CH-CH=CH~2~ + m C~6~H~5~-CH=CH~2~ -> Cao su buna-S.")
        ]
    },
    {
        "num": 2,
        "intro": "Tơ sợi là một nguyên liệu quan trọng trong ngành dệt may và sản xuất vật liệu. Tơ sợi được sản xuất từ các nguồn nguyên liệu tự nhiên hoặc tổng hợp. Các loại tơ sợi phổ biến bao gồm tơ sợi tự nhiên như tơ tằm, lông cừu, sợi cotton và tơ sợi tổng hợp như nylon,... Những phát biểu nào sau đây là đúng?",
        "items": [
            ("a", "Tơ nylon-6,6 và tơ capron thuộc loại tơ polyamide.", True, "Đúng. Chúng đều chứa các liên kết amide -CO-NH- trong mạch phân tử polymer."),
            ("b", "Tơ nylon, tơ tằm, tơ visco đều bền với nhiệt độ.", False, "Sai. Tơ nylon (polyamide) và tơ tằm (protein) đều kém bền với nhiệt độ và kém bền trong môi trường axit/kiềm."),
            ("c", "Quần áo được dệt bằng sợi len lông cừu, tơ tằm không nên giặt với xà phòng có độ kiềm cao.", True, "Đúng. Len lông cừu và tơ tằm là protein, chứa nhiều liên kết peptide dễ bị thủy phân trong môi trường kiềm mạnh làm hỏng vải."),
            ("d", "Tơ capron và tơ olon đều có thành phần chứa nhóm -CO-NH-.", False, "Sai. Tơ capron chứa nhóm -CO-NH-, nhưng tơ olon (nitron) là (-CH~2~-CH(CN)-)~n~ chứa nhóm cyano -CN, không chứa nhóm -CO-NH-.")
        ]
    },
    {
        "num": 3,
        "intro": "Poly(ethylene terephthalate) (viết tắt là PET) là một polymer được điều chế từ terephthalic acid và ethylene glycol. PET được sử dụng để sản xuất tơ, chai đựng nước uống, hộp đựng thực phẩm. Để thuận lợi cho việc nhận biết, sử dụng và tái chế thì các đồ nhựa làm từ vật liệu chứa PET thường được in kí hiệu như hình bên.",
        "image": "p2_cau3_pet.png",
        "items": [
            ("a", "PET thuộc loại polyester.", True, "Đúng. PET được tạo thành từ phản ứng giữa dicarboxylic acid và diol, mạch polymer chứa các nhóm chức ester -COO-."),
            ("b", "Tơ được chế tạo từ PET thuộc loại tơ tổng hợp.", True, "Đúng. Tơ lapsan (dacron) sản xuất từ PET thuộc loại tơ tổng hợp."),
            ("c", "Trong một mắt xích PET, phần trăm khối lượng carbon là 62,5%.", True, "Đúng. Một mắt xích PET có công thức C~10~H~8~O~4~ (M = 192 amu). Phần trăm khối lượng carbon: %C = (12 * 10 / 192) * 100% = 62,5%."),
            ("d", "Phản ứng tổng hợp PET từ terephthalic acid và ethylene glycol thuộc loại phản ứng trùng hợp.", False, "Sai. Đây là phản ứng trùng ngưng (tách ra các phân tử nhỏ H~2~O cùng với mạch polymer).")
        ]
    },
    {
        "num": 4,
        "intro": "Công ty The Goodyear Tire & Rubber là một trong những công ty lốp xe lớn nhất thế giới khởi lập năm 1898. Năm 1971, lốp Goodyear trở thành bánh xe đầu tiên lăn trên Mặt Trăng... Tên công ty được đặt theo tên của nhà tiên phong Charles Goodyear, người khám phá ra phương pháp kết hợp giữa nguyên tố S (lưu huỳnh) với cao su để tạo ra một loại cao su mới.",
        "items": [
            ("a", "Loại cao su này có tên là cao su lưu hóa.", True, "Đúng. Cao su khi tác dụng với lưu huỳnh ở điều kiện thích hợp được gọi là cao su lưu hóa."),
            ("b", "Cao su lưu hóa có cấu trúc mạch phân nhánh.", False, "Sai. Cao su lưu hóa có cấu trúc mạng không gian 3 chiều do các cầu nối disulfide liên kết các mạch polymer lại với nhau."),
            ("c", "Cao su lưu hóa làm tăng cao tính bền hơn với nhiệt và các tác nhân khác, độ đàn hồi cao, chống thấm khí, chống thấm ẩm, tăng khả năng chịu được sự ma sát, va chạm tốt hơn cao su thiên nhiên.", True, "Đúng. Quá trình lưu hóa cải thiện vượt bậc cơ tính và độ bền nhiệt của cao su."),
            ("d", "Cao su lưu hóa làm tăng tính bền cơ học nhờ các cầu nối trisulfide (-S-S-S-) giữa các phân tử polyisoprene.", False, "Sai. Cầu nối tạo mạng lưới không gian chủ yếu là cầu nối disulfide (-S-S-) hoặc monosulfide (-S-), không phải trisulfide.")
        ]
    }
]

SA_QUESTIONS = [
    {
        "num": 1,
        "text": "Có bao nhiêu loại polymer có thể tạo thành từ adipic acid, 1,2-ethylene diamine và ethylene glycol bằng phản ứng trùng ngưng?",
        "ans": "2",
        "sol": "Có 2 loại polymer có thể tạo thành khi các cặp chất phản ứng với nhau:\n• Adipic acid phản ứng với 1,2-ethylene diamine tạo thành polyamide.\n• Adipic acid phản ứng với ethylene glycol tạo thành polyester.\n(1,2-ethylene diamine và ethylene glycol không trùng ngưng với nhau). Vậy có 2 loại polymer."
    },
    {
        "num": 2,
        "text": "Cho dãy các polymer sau: polyethylene, cellulose, nylon-6,6, amylose, nylon-6, tơ nitron, polybuta-1,3-diene, tơ visco, poly(phenol formaldehyde), poly(vinyl acetate). Có bao nhiêu polymer tổng hợp trong dãy trên?",
        "ans": "7",
        "sol": "Các polymer tổng hợp trong dãy gồm 7 chất: polyethylene, nylon-6,6, nylon-6, tơ nitron, polybuta-1,3-diene, poly(phenol formaldehyde), poly(vinyl acetate).\n(Cellulose và amylose là polymer thiên nhiên; tơ visco là polymer bán tổng hợp)."
    },
    {
        "num": 3,
        "text": "Khối lượng phân tử của một loại thủy tinh hữu cơ plexiglass là 25000 amu. Số mắt xích trong phân tử thủy tinh hữu cơ đó là bao nhiêu?",
        "ans": "250",
        "sol": "Thủy tinh hữu cơ plexiglass là poly(methyl methacrylate) có mắt xích là [-CH~2~-C(CH~3~)(COOCH~3~)-] với phân tử khối M = 100 amu.\nSố mắt xích = 25000 / 100 = 250."
    },
    {
        "num": 4,
        "text": "Khối lượng phân tử của một đoạn polyethylene (PE) là 28000 amu. Số mắt xích của đoạn polymer này là bao nhiêu?",
        "ans": "1000",
        "sol": "Polyethylene có dạng (-CH~2~-CH~2~-)~n~ với phân tử khối một mắt xích M = 28 amu.\nSố mắt xích (hệ số polymer hóa) n = 28000 / 28 = 1000."
    },
    {
        "num": 5,
        "text": "Cellulose triacetate được điều chế từ cellulose và acetic anhydride. Tính khối lượng cellulose triacetate (theo đơn vị tấn) tạo thành từ 1 tấn cellulose nếu hiệu suất của quá trình điều chế là 67%. (Làm tròn kết quả đến chữ số thập phân thứ hai).",
        "ans": "1,19",
        "sol": "Phương trình phản ứng theo mắt xích:\n[C~6~H~7~O~2~(OH)~3~]~n~ + 3n (CH~3~CO)~2~O -> [C~6~H~7~O~2~(OOC-CH~3~)~3~]~n~ + 3n CH~3~COOH\nCứ 162 gam cellulose tạo 288 gam cellulose triacetate.\nKhối lượng lý thuyết = 1 * (288 / 162) = 1,7778 tấn.\nKhối lượng thực tế thu được (H = 67%): m = 1,7778 * 67% = 1,1911 tấn ≈ 1,19 tấn."
    },
    {
        "num": 6,
        "text": "Một loại cao su Buna-S có phần trăm khối lượng carbon là 90,225%; m gam cao su này cộng tối đa với 9,6 gam bromine. Giá trị của m (theo đơn vị gam) là bao nhiêu? (Làm tròn kết quả đến chữ số thập phân thứ hai).",
        "ans": "5,32",
        "sol": "• Cao su Buna-S có dạng (C~4~H~6~)~a~.(C~8~H~8~)~b~.\n%C = [12(4a + 8b) / (54a + 104b)] * 100% = 90,225%\n=> 48a + 96b = 0,90225(54a + 104b) => 2,166b ≈ 0,7215a => a / b ≈ 3 / 1.\nVậy cao su có dạng (C~4~H~6~)~3b~.(C~8~H~8~)~b~.\n• Mỗi mắt xích buta-1,3-diene còn 1 nối đôi C=C phản ứng với 1 phân tử Br~2~: n_C4H6 = n_Br2 = 9,6 / 160 = 0,06 mol.\nVì a : b = 3 : 1 nên n_C8H8 = 0,06 / 3 = 0,02 mol.\n=> m = m_C4H6 + m_C8H8 = 0,06 * 54 + 0,02 * 104 = 3,24 + 2,08 = 5,32 gam."
    }
]

def generate_exam_docx(is_teacher=False):
    filename = "De_Kiem_Tra_Hoa_12_De10_GiaoVien.docx" if is_teacher else "De_Kiem_Tra_Hoa_12_De10.docx"
    filepath = os.path.join(OUT_DIR, filename)
    
    doc = docx.Document()
    setup_page_and_header_footer(doc)
    add_exam_title_block(doc, is_teacher=is_teacher)
    
    # --- PHẦN I ---
    add_part_heading(doc, "PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.", "Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án.")
    
    for q in MCQ_QUESTIONS:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(1)
        p_q.paragraph_format.line_spacing = 1.15
        
        r_qnum = p_q.add_run(f"Câu {q['num']}. ")
        set_run_style(r_qnum, font_name="Times New Roman", font_size=11, bold=True)
        
        add_chem_run(p_q, q["text"], base_font="Times New Roman", base_size=11)
        
        # Format options using Tabs (NO TABLE)
        opts = q["options"]
        max_opt_len = max(len(opt) for opt in opts)
        
        if max_opt_len < 22:
            # 1 dòng 4 lựa chọn (Tab stops 4.5, 9.0, 13.5 cm)
            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.space_before = Pt(1)
            p_opt.paragraph_format.space_after = Pt(2)
            p_opt.paragraph_format.line_spacing = 1.15
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(4.5), WD_TAB_ALIGNMENT.LEFT)
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(13.5), WD_TAB_ALIGNMENT.LEFT)
            
            for i, opt in enumerate(opts):
                is_correct = (opt.startswith(q["correct"] + ".")) and is_teacher
                bold = is_correct or opt[:2] in ["A.", "B.", "C.", "D."]
                color = COLOR_CORRECT if is_correct else None
                
                # Letter prefix
                prefix = opt[:3]
                body = opt[3:]
                r_pre = p_opt.add_run(prefix)
                set_run_style(r_pre, font_name="Times New Roman", font_size=11, bold=True, color=color)
                add_chem_run(p_opt, body, base_font="Times New Roman", base_size=11, bold=is_correct, color=color)
                
                if i < 3:
                    p_opt.add_run("\t")
        elif max_opt_len < 46:
            # 2 dòng 2 lựa chọn (Tab stop 9.0 cm)
            for row in [(0, 1), (2, 3)]:
                p_opt = doc.add_paragraph()
                p_opt.paragraph_format.space_before = Pt(1)
                p_opt.paragraph_format.space_after = Pt(1.5)
                p_opt.paragraph_format.line_spacing = 1.15
                p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
                
                for idx_in_row, opt_idx in enumerate(row):
                    opt = opts[opt_idx]
                    is_correct = (opt.startswith(q["correct"] + ".")) and is_teacher
                    color = COLOR_CORRECT if is_correct else None
                    
                    prefix = opt[:3]
                    body = opt[3:]
                    r_pre = p_opt.add_run(prefix)
                    set_run_style(r_pre, font_name="Times New Roman", font_size=11, bold=True, color=color)
                    add_chem_run(p_opt, body, base_font="Times New Roman", base_size=11, bold=is_correct, color=color)
                    
                    if idx_in_row == 0:
                        p_opt.add_run("\t")
        else:
            # 4 dòng riêng biệt
            for opt in opts:
                p_opt = doc.add_paragraph()
                p_opt.paragraph_format.space_before = Pt(1)
                p_opt.paragraph_format.space_after = Pt(1.5)
                p_opt.paragraph_format.line_spacing = 1.15
                
                is_correct = (opt.startswith(q["correct"] + ".")) and is_teacher
                color = COLOR_CORRECT if is_correct else None
                
                prefix = opt[:3]
                body = opt[3:]
                r_pre = p_opt.add_run(prefix)
                set_run_style(r_pre, font_name="Times New Roman", font_size=11, bold=True, color=color)
                add_chem_run(p_opt, body, base_font="Times New Roman", base_size=11, bold=is_correct, color=color)

        if is_teacher and q.get("sol"):
            add_solution_paragraph(doc, q["sol"])

    # --- PHẦN II ---
    add_part_heading(doc, "PHẦN II. Câu trắc nghiệm đúng sai.", "Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.")
    
    for q in TF_QUESTIONS:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        
        r_qnum = p_q.add_run(f"Câu {q['num']}. ")
        set_run_style(r_qnum, font_name="Times New Roman", font_size=11, bold=True)
        
        add_chem_run(p_q, q["intro"], base_font="Times New Roman", base_size=11)
        
        # Nếu có hình ảnh kèm theo (như Câu 3 có biểu tượng PET)
        if q.get("image"):
            img_path = os.path.join(IMG_DIR, q["image"])
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(2)
                p_img.paragraph_format.space_after = Pt(2)
                run_img = p_img.add_run()
                run_img.add_picture(img_path, width=Cm(3.2))

        for item_key, item_text, is_true, item_sol in q["items"]:
            p_item = doc.add_paragraph()
            p_item.paragraph_format.space_before = Pt(1)
            p_item.paragraph_format.space_after = Pt(1.5)
            p_item.paragraph_format.line_spacing = 1.15
            p_item.paragraph_format.left_indent = Cm(0.5)
            
            r_k = p_item.add_run(f"{item_key}) ")
            set_run_style(r_k, font_name="Times New Roman", font_size=11, bold=True)
            
            add_chem_run(p_item, item_text, base_font="Times New Roman", base_size=11)
            
            if is_teacher:
                status_str = " [ĐÚNG]" if is_true else " [SAI]"
                color_st = COLOR_CORRECT if is_true else COLOR_PRIMARY
                r_st = p_item.add_run(status_str)
                set_run_style(r_st, font_name="Times New Roman", font_size=11, bold=True, color=color_st)
                
                # Explanation
                p_sol = doc.add_paragraph()
                p_sol.paragraph_format.space_before = Pt(0)
                p_sol.paragraph_format.space_after = Pt(2)
                p_sol.paragraph_format.left_indent = Cm(1.0)
                r_sol_lbl = p_sol.add_run("-> ")
                set_run_style(r_sol_lbl, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_SOLUTION)
                add_chem_run(p_sol, item_sol, base_font="Times New Roman", base_size=10, italic=True)

    # --- PHẦN III ---
    add_part_heading(doc, "PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn.", "Thí sinh trả lời từ câu 1 đến câu 6.")
    
    for q in SA_QUESTIONS:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        
        r_qnum = p_q.add_run(f"Câu {q['num']}. ")
        set_run_style(r_qnum, font_name="Times New Roman", font_size=11, bold=True)
        
        add_chem_run(p_q, q["text"], base_font="Times New Roman", base_size=11)
        
        if not is_teacher:
            # Student version dotted lines
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(3)
            p_ans.paragraph_format.left_indent = Cm(0.5)
            r_ans_lbl = p_ans.add_run("Trả lời: .................................................................................................................................")
            set_run_style(r_ans_lbl, font_name="Times New Roman", font_size=11, italic=True)
        else:
            # Teacher version
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(2)
            p_ans.paragraph_format.left_indent = Cm(0.5)
            
            r_ans_lbl = p_ans.add_run("Đáp số: ")
            set_run_style(r_ans_lbl, font_name="Times New Roman", font_size=11, bold=True, color=COLOR_CORRECT)
            r_ans_val = p_ans.add_run(q["ans"])
            set_run_style(r_ans_val, font_name="Times New Roman", font_size=11, bold=True, color=COLOR_CORRECT)
            
            add_solution_paragraph(doc, q["sol"])

    # Bảng tổng hợp đáp án ở cuối đề Giáo Viên
    if is_teacher:
        p_table_title = doc.add_paragraph()
        p_table_title.paragraph_format.space_before = Pt(10)
        p_table_title.paragraph_format.space_after = Pt(3)
        r_tt = p_table_title.add_run("BẢNG TỔNG HỢP ĐÁP ÁN ĐỀ SỐ 10")
        set_run_style(r_tt, font_name="Times New Roman", font_size=11.5, bold=True, color=COLOR_PRIMARY)
        
        # Phần I Table
        p_p1 = doc.add_paragraph()
        p_p1.paragraph_format.space_before = Pt(3)
        p_p1.paragraph_format.space_after = Pt(1)
        r_p1 = p_p1.add_run("Phần I: Trắc nghiệm nhiều phương án lựa chọn")
        set_run_style(r_p1, font_name="Times New Roman", font_size=10.5, bold=True)
        
        tbl_ans1 = doc.add_table(rows=2, cols=18)
        tbl_ans1.autofit = False
        for c_idx in range(18):
            cell_q = tbl_ans1.cell(0, c_idx)
            cell_a = tbl_ans1.cell(1, c_idx)
            cell_q.width = Cm(1.0)
            cell_a.width = Cm(1.0)
            
            pq = cell_q.paragraphs[0]
            pq.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pq.paragraph_format.space_before = Pt(1)
            pq.paragraph_format.space_after = Pt(1)
            rq = pq.add_run(str(c_idx + 1))
            set_run_style(rq, font_name="Times New Roman", font_size=9, bold=True)
            
            pa = cell_a.paragraphs[0]
            pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pa.paragraph_format.space_before = Pt(1)
            pa.paragraph_format.space_after = Pt(1)
            ra = pa.add_run(MCQ_QUESTIONS[c_idx]["correct"])
            set_run_style(ra, font_name="Times New Roman", font_size=9.5, bold=True, color=COLOR_CORRECT)
            
            for cell in [cell_q, cell_a]:
                set_cell_border(cell, top={'val':'single','sz':'4','color':'CCCCCC'},
                                      bottom={'val':'single','sz':'4','color':'CCCCCC'},
                                      left={'val':'single','sz':'4','color':'CCCCCC'},
                                      right={'val':'single','sz':'4','color':'CCCCCC'})
                set_cell_margins(cell, top=20, bottom=20, left=10, right=10)

        # Phần II Table
        p_p2 = doc.add_paragraph()
        p_p2.paragraph_format.space_before = Pt(5)
        p_p2.paragraph_format.space_after = Pt(1)
        r_p2 = p_p2.add_run("Phần II: Đúng/Sai")
        set_run_style(r_p2, font_name="Times New Roman", font_size=10.5, bold=True)
        
        tbl_ans2 = doc.add_table(rows=5, cols=5)
        tbl_ans2.autofit = False
        headers_p2 = ["Câu", "a", "b", "c", "d"]
        for j, h in enumerate(headers_p2):
            cell = tbl_ans2.cell(0, j)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(h)
            set_run_style(r, font_name="Times New Roman", font_size=9.5, bold=True)
            set_cell_border(cell, top={'val':'single','sz':'4','color':'CCCCCC'},
                                  bottom={'val':'single','sz':'4','color':'CCCCCC'},
                                  left={'val':'single','sz':'4','color':'CCCCCC'},
                                  right={'val':'single','sz':'4','color':'CCCCCC'})
            set_cell_margins(cell, top=20, bottom=20, left=20, right=20)
            
        for i_q, q_tf in enumerate(TF_QUESTIONS):
            cell_q = tbl_ans2.cell(i_q + 1, 0)
            p = cell_q.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(f"Câu {q_tf['num']}")
            set_run_style(r, font_name="Times New Roman", font_size=9.5, bold=True)
            set_cell_border(cell_q, top={'val':'single','sz':'4','color':'CCCCCC'},
                                  bottom={'val':'single','sz':'4','color':'CCCCCC'},
                                  left={'val':'single','sz':'4','color':'CCCCCC'},
                                  right={'val':'single','sz':'4','color':'CCCCCC'})
            set_cell_margins(cell_q, top=20, bottom=20, left=20, right=20)
            
            for j_item, item in enumerate(q_tf["items"]):
                cell_item = tbl_ans2.cell(i_q + 1, j_item + 1)
                p = cell_item.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                st_txt = "Đ" if item[2] else "S"
                st_color = COLOR_CORRECT if item[2] else COLOR_PRIMARY
                r = p.add_run(st_txt)
                set_run_style(r, font_name="Times New Roman", font_size=9.5, bold=True, color=st_color)
                set_cell_border(cell_item, top={'val':'single','sz':'4','color':'CCCCCC'},
                                      bottom={'val':'single','sz':'4','color':'CCCCCC'},
                                      left={'val':'single','sz':'4','color':'CCCCCC'},
                                      right={'val':'single','sz':'4','color':'CCCCCC'})
                set_cell_margins(cell_item, top=20, bottom=20, left=20, right=20)

        # Phần III Table
        p_p3 = doc.add_paragraph()
        p_p3.paragraph_format.space_before = Pt(5)
        p_p3.paragraph_format.space_after = Pt(1)
        r_p3 = p_p3.add_run("Phần III: Trả lời ngắn")
        set_run_style(r_p3, font_name="Times New Roman", font_size=10.5, bold=True)
        
        tbl_ans3 = doc.add_table(rows=2, cols=6)
        tbl_ans3.autofit = False
        for c_idx in range(6):
            cell_q = tbl_ans3.cell(0, c_idx)
            cell_a = tbl_ans3.cell(1, c_idx)
            cell_q.width = Cm(2.0)
            cell_a.width = Cm(2.0)
            
            pq = cell_q.paragraphs[0]
            pq.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pq.paragraph_format.space_before = Pt(1)
            pq.paragraph_format.space_after = Pt(1)
            rq = pq.add_run(f"Câu {c_idx + 1}")
            set_run_style(rq, font_name="Times New Roman", font_size=9.5, bold=True)
            
            pa = cell_a.paragraphs[0]
            pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pa.paragraph_format.space_before = Pt(1)
            pa.paragraph_format.space_after = Pt(1)
            ra = pa.add_run(SA_QUESTIONS[c_idx]["ans"])
            set_run_style(ra, font_name="Times New Roman", font_size=9.5, bold=True, color=COLOR_CORRECT)
            
            for cell in [cell_q, cell_a]:
                set_cell_border(cell, top={'val':'single','sz':'4','color':'CCCCCC'},
                                      bottom={'val':'single','sz':'4','color':'CCCCCC'},
                                      left={'val':'single','sz':'4','color':'CCCCCC'},
                                      right={'val':'single','sz':'4','color':'CCCCCC'})
                set_cell_margins(cell, top=20, bottom=20, left=20, right=20)

    doc.save(filepath)
    print(f"Generated: {filepath}")

if __name__ == "__main__":
    generate_exam_docx(is_teacher=False)
    generate_exam_docx(is_teacher=True)
