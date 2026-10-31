import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_slide():
    prs = Presentation()
    # 16:9 Widescreen: 13.333 x 7.5 inches
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    slide_layout = prs.slide_layouts[6] # blank
    slide = prs.slides.add_slide(slide_layout)
    
    # Background: very soft cool off-white
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = RGBColor(248, 250, 252)
    bg_shape.line.fill.background()
    
    # ---------------- Top Header Bar ----------------
    # Top decorative line
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.12))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = RGBColor(21, 101, 192)
    top_bar.line.fill.background()
    
    # Category Tag / Badge
    tag_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(0.35), Inches(3.2), Inches(0.4))
    tag_box.fill.solid()
    tag_box.fill.fore_color.rgb = RGBColor(232, 245, 233)
    tag_box.line.color.rgb = RGBColor(76, 175, 80)
    tag_tf = tag_box.text_frame
    tag_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tag_p = tag_tf.paragraphs[0]
    tag_p.text = "🌱 KHOA HỌC TỰ NHIÊN • SINH HỌC"
    tag_p.font.size = Pt(11)
    tag_p.font.bold = True
    tag_p.font.color.rgb = RGBColor(46, 125, 50)
    tag_p.alignment = PP_ALIGN.CENTER
    
    # Main Title
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.8), Inches(12.133), Inches(0.7))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "VÌ SAO CƠ THỂ ĐA BÀO CẦN NHIỀU TẾ BÀO CHUYÊN HOÁ?"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = RGBColor(17, 24, 39)
    
    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.35), Inches(12.133), Inches(0.45))
    stf = sub_box.text_frame
    stf.word_wrap = True
    sp = stf.paragraphs[0]
    sp.text = "So sánh đặc điểm trao đổi chất & quy luật tỉ lệ S/V giữa Cơ thể đơn bào và Cơ thể đa bào"
    sp.font.size = Pt(13)
    sp.font.italic = True
    sp.font.color.rgb = RGBColor(75, 85, 99)
    
    # ---------------- 2 Cards Layout ----------------
    # Card 1: CƠ THỂ ĐƠN BÀO (Left)
    c1_left = Inches(0.6)
    c1_top = Inches(1.9)
    card_w = Inches(5.9)
    card_h = Inches(4.35)
    
    # Outer Card Shape
    c1_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left, c1_top, card_w, card_h)
    c1_card.fill.solid()
    c1_card.fill.fore_color.rgb = RGBColor(255, 255, 255)
    c1_card.line.color.rgb = RGBColor(167, 243, 208)
    c1_card.line.width = Pt(1.5)
    
    # Card 1 Header banner
    c1_hdr = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left, c1_top, card_w, Inches(0.65))
    c1_hdr.fill.solid()
    c1_hdr.fill.fore_color.rgb = RGBColor(16, 149, 114)
    c1_hdr.line.fill.background()
    c1_htf = c1_hdr.text_frame
    c1_htf.vertical_anchor = MSO_ANCHOR.MIDDLE
    c1_hp = c1_htf.paragraphs[0]
    c1_hp.text = "CƠ THỂ ĐƠN BÀO (Ví dụ: Trùng biến hình)"
    c1_hp.font.size = Pt(14)
    c1_hp.font.bold = True
    c1_hp.font.color.rgb = RGBColor(255, 255, 255)
    c1_hp.alignment = PP_ALIGN.CENTER
    
    # Card 1 Image & Details
    amoeba_img_path = 'San_Pham/amoeba.png'
    if os.path.exists(amoeba_img_path):
        slide.shapes.add_picture(amoeba_img_path, Inches(4.5), Inches(2.7), width=Inches(1.8))
        
    c1_text_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.65), Inches(3.6), Inches(3.45))
    ctf1 = c1_text_box.text_frame
    ctf1.word_wrap = True
    
    items1 = [
        ("Kích thước cơ thể:", " Rất nhỏ (kích thước hiển vi)."),
        ("Tỉ lệ S/V (Diện tích / Thể tích):", " RẤT LỚN (diện tích tiếp xúc môi trường lớn so với thể tích)."),
        ("Cơ chế trao đổi chất:", " Dinh dưỡng & chất thải dễ dàng khuếch tán TRỰC TIẾP qua màng sinh chất."),
        ("Tốc độ trao đổi chất:", " Nhanh chóng, đáp ứng đầy đủ và kịp thời mọi nhu cầu sống."),
        ("Mức độ chuyên hoá:", " KHÔNG CÓ tế bào chuyên hoá (1 tế bào đảm đương toàn bộ hoạt động sống).")
    ]
    
    for i, (label, desc) in enumerate(items1):
        cp = ctf1.paragraphs[0] if i == 0 else ctf1.add_paragraph()
        cp.space_after = Pt(7)
        r1 = cp.add_run()
        r1.text = "• " + label
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = RGBColor(6, 95, 70)
        
        r2 = cp.add_run()
        r2.text = desc
        r2.font.bold = False
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = RGBColor(31, 41, 55)
        
    # Card 2: CƠ THỂ ĐA BÀO (Right)
    c2_left = Inches(6.833)
    c2_top = Inches(1.9)
    
    # Outer Card Shape
    c2_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left, c2_top, card_w, card_h)
    c2_card.fill.solid()
    c2_card.fill.fore_color.rgb = RGBColor(255, 255, 255)
    c2_card.line.color.rgb = RGBColor(191, 219, 254)
    c2_card.line.width = Pt(1.5)
    
    # Card 2 Header banner
    c2_hdr = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left, c2_top, card_w, Inches(0.65))
    c2_hdr.fill.solid()
    c2_hdr.fill.fore_color.rgb = RGBColor(30, 64, 175)
    c2_hdr.line.fill.background()
    c2_htf = c2_hdr.text_frame
    c2_htf.vertical_anchor = MSO_ANCHOR.MIDDLE
    c2_hp = c2_htf.paragraphs[0]
    c2_hp.text = "CƠ THỂ ĐA BÀO (Ví dụ: Con voi, con gà, con người...)"
    c2_hp.font.size = Pt(14)
    c2_hp.font.bold = True
    c2_hp.font.color.rgb = RGBColor(255, 255, 255)
    c2_hp.alignment = PP_ALIGN.CENTER
    
    # Card 2 Image & Details
    elephant_img_path = 'San_Pham/elephant.png'
    if os.path.exists(elephant_img_path):
        slide.shapes.add_picture(elephant_img_path, Inches(10.9), Inches(2.7), width=Inches(1.65))
        
    c2_text_box = slide.shapes.add_textbox(Inches(7.033), Inches(2.65), Inches(3.7), Inches(3.45))
    ctf2 = c2_text_box.text_frame
    ctf2.word_wrap = True
    
    items2 = [
        ("Kích thước cơ thể:", " Thường có kích thước rất lớn."),
        ("Tỉ lệ S/V (Diện tích / Thể tích):", " RẤT NHỎ (diện tích bề mặt rất nhỏ so với thể tích khổng lồ)."),
        ("Cơ chế trao đổi chất:", " Vận chuyển trực tiếp qua màng ngoài KHÔNG ĐỦ cho các tế bào sâu bên trong."),
        ("Tốc độ trao đổi chất:", " Khuếch tán thụ động quá chậm, không cung cấp kịp dưỡng chất & oxi."),
        ("Mức độ chuyên hoá:", " BẮT BUỘC có các tế bào chuyên hoá (tuần hoàn, hô hấp, tiêu hoá, bài tiết...) để vận chuyển chất.")
    ]
    
    for i, (label, desc) in enumerate(items2):
        cp = ctf2.paragraphs[0] if i == 0 else ctf2.add_paragraph()
        cp.space_after = Pt(7)
        r1 = cp.add_run()
        r1.text = "• " + label
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = RGBColor(30, 58, 138)
        
        r2 = cp.add_run()
        r2.text = desc
        r2.font.bold = False
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = RGBColor(31, 41, 55)

    # ---------------- Bottom Key Takeaway Box ----------------
    bot_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(6.38), Inches(12.133), Inches(0.95))
    bot_box.fill.solid()
    bot_box.fill.fore_color.rgb = RGBColor(254, 243, 199)
    bot_box.line.color.rgb = RGBColor(245, 158, 11)
    bot_box.line.width = Pt(1.5)
    
    btf = bot_box.text_frame
    btf.word_wrap = True
    btf.vertical_anchor = MSO_ANCHOR.MIDDLE
    bp = btf.paragraphs[0]
    
    br_icon = bp.add_run()
    br_icon.text = "💡 BẢN CHẤT SINH HỌC CỐT LÕI: "
    br_icon.font.bold = True
    br_icon.font.size = Pt(11.5)
    br_icon.font.color.rgb = RGBColor(180, 83, 9)
    
    br_text = bp.add_run()
    br_text.text = "Kích thước cơ thể càng lớn ➔ Tỉ lệ S/V càng giảm ➔ Bề mặt tiếp xúc không đủ để trao đổi chất trực tiếp cho các tế bào bên trong ➔ Cơ thể đa bào bắt buộc phải tiến hóa hệ thống CÁC TẾ BÀO VÀ HỆ CƠ QUAN CHUYÊN HOÁ để đảm bảo nuôi dưỡng toàn bộ cơ thể."
    br_text.font.size = Pt(11)
    br_text.font.bold = False
    br_text.font.color.rgb = RGBColor(120, 53, 15)

    out_path = 'San_Pham/Slide_So_Sanh_Don_Bao_Da_Bao.pptx'
    prs.save(out_path)
    print(f"Successfully generated PowerPoint slide: {out_path}")

if __name__ == '__main__':
    create_slide()
