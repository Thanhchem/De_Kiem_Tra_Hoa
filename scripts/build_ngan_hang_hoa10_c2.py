# -*- coding: utf-8 -*-
"""
Xây dựng Ngân hàng câu hỏi Hóa học 10 – Chương 2: Bảng tuần hoàn các nguyên tố hóa học và Định luật tuần hoàn
Phân loại chi tiết theo:
- Chủ đề / Bài học:
  1. Bài 5: Cấu tạo của bảng tuần hoàn các nguyên tố hóa học
  2. Bài 6: Xu hướng biến đổi một số tính chất của nguyên tử các nguyên tố trong một chu kì và nhóm
  3. Bài 7: Xu hướng biến đổi thành phần và một số tính chất của hợp chất trong một chu kì
  4. Bài 8: Định luật tuần hoàn và ý nghĩa của bảng tuần hoàn các nguyên tố hóa học
- Dạng thức câu hỏi:
  + Phần I (P1): Trắc nghiệm nhiều phương án lựa chọn (4 phương án A, B, C, D)
  + Phần II (P2): Trắc nghiệm Đúng / Sai (4 ý a, b, c, d có phân loại mức độ từng ý)
  + Phần III (P3): Trắc nghiệm Trả lời ngắn & Bài toán vận dụng
- Mức độ nhận thức: Nhận biết [NB] – Thông hiểu [TH] – Vận dụng [VD]
- Xuất cơ sở dữ liệu JSON chuẩn + Bộ 2 file Word (.docx) Học sinh & Giáo viên theo đúng "FORM ĐỀ KIỂM TRA".
"""

import os
import io
import re
import json
import glob
import sys
from difflib import SequenceMatcher
from PIL import Image
import docx
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

sys.path.insert(0, r'c:\Antigravity_Thanh\scripts')
from build_ngan_hang_hoa10_c1 import (
    iter_block_items, get_cau_num_ids, get_para_text_full,
    extract_real_images_from_el, extract_underlined_p1_answer,
    clean_extracted_text, add_chem_runs_hoa10, add_options_paragraph_no_table,
    add_data_table_to_doc, setup_page_and_hf, parse_abcd_from_opt_lines,
    norm_for_compare, LEVEL_NAMES, COLOR_HF_RED, COLOR_PRIMARY,
    COLOR_SOLUTION, COLOR_CORRECT, COLOR_TAG
)

BASE_DIR = r"C:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\thay-Trung-3-khoi\H10 - CAO QUOC TRUNG\H10 - CAO QUOC TRUNG"
CVL_DOCX = glob.glob(os.path.join(BASE_DIR, "2026_H10.C2*GIAO VIEN*.docx"))[0]
TTB_DOCX = r"C:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\2026 - Phan dang de thi thu truong, so. TTB1-20260711T201135Z-2-001\2026 - Phan dang de thi thu truong, so. TTB1\X. PD 2026 - Hoa 10. CD10 - SGK va CDHT. TTB. ok - DA chi tiet.docx"

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
BANK_DIR = r"c:\Antigravity_Thanh\San_Pham\Ngan_Hang_Hoa_10"
IMG_DIR = r"c:\Antigravity_Thanh\San_Pham\images_hoa10"
os.makedirs(OUT_DIR, exist_ok=True)
os.makedirs(BANK_DIR, exist_ok=True)
os.makedirs(IMG_DIR, exist_ok=True)

CHAPTER_INFO = {
    "chapter_number": 2,
    "code": "C2",
    "name": "CHƯƠNG 2: BẢNG TUẦN HOÀN CÁC NGUYÊN TỐ HÓA HỌC VÀ ĐỊNH LUẬT TUẦN HOÀN",
    "short_name": "Bảng tuần hoàn",
    "file_prefix": "Ngan_Hang_Hoa_10_Chuong_2_Bang_Tuan_Hoan",
    "topics": [
        "Bài 5: Cấu tạo của bảng tuần hoàn các nguyên tố hóa học",
        "Bài 6: Xu hướng biến đổi một số tính chất của nguyên tử các nguyên tố trong một chu kì và nhóm",
        "Bài 7: Xu hướng biến đổi thành phần và một số tính chất của hợp chất trong một chu kì",
        "Bài 8: Định luật tuần hoàn và ý nghĩa của bảng tuần hoàn các nguyên tố hóa học"
    ]
}


def classify_p2_item_level_hoa10_c2(stmt_text):
    t = stmt_text.lower()
    if any(k in t for k in [
        'tính toán', 'phần trăm khối lượng', 'hệ phương trình', 'khối lượng', 'gam',
        'nồng độ', 'thể tích', 'độ tinh khiết', 'lập phương trình', 'tỉ lệ mol'
    ]) or re.search(r'\d+[,\.]\d+', t):
        return "VD"
    if any(k in t for k in [
        'so sánh', 'mạnh hơn', 'yếu hơn', 'lớn hơn', 'nhỏ hơn', 'giảm dần', 'tăng dần',
        'tính kim loại', 'tính phi kim', 'độ âm điện', 'bán kính', 'tính acid', 'tính base',
        'oxit cao nhất', 'oxide cao nhất', 'hydroxide', 'cấu hình electron', 'ion', 'hóa trị'
    ]):
        return "TH"
    return "NB"


def classify_question_level_hoa10_c2(part, sec, topic, stem_text, opt_text, sol_text, item_levels=None):
    full = (stem_text + " " + opt_text).lower()
    if sec in ("DANG_1", "DANG_2", "DANG_3"):
        return "VD"

    if part == "P2":
        vals = list((item_levels or {}).values())
        if vals.count("VD") >= 2:
            return "VD"
        if vals.count("NB") >= 3 and vals.count("VD") == 0:
            return "NB"
        return "TH"

    if part == "P3":
        if any(k in full for k in [
            "khối lượng", "gam", "%", "hiệu suất", "tinh khiết", "thể tích", "mol",
            "chênh lệch", "kế tiếp", "hai chu kì liên tiếp", "hai nhóm a"
        ]) or re.search(r'\d', full):
            if any(k in full for k in ["ghép số thứ tự", "có bao nhiêu nguyên tố", "số electron hóa trị"]):
                return "TH"
            return "VD"
        return "TH"

    # P1
    if any(k in full for k in [
        "số phát biểu đúng", "số phát biểu sai", "số phát biểu không đúng",
        "hơn kém nhau", "kế tiếp nhau", "tổng số hạt proton của hai",
        "tổng số hạt mang điện trong hai", "phần trăm khối lượng"
    ]):
        return "VD"

    if any(k in full for k in [
        "không đúng", "phát biểu nào sau đây đúng", "nhận định nào sau đây",
        "sắp xếp", "theo chiều tăng dần", "theo chiều giảm dần", "mạnh nhất", "yếu nhất",
        "lớn nhất", "nhỏ nhất", "so sánh", "tính kim loại", "tính phi kim",
        "độ âm điện", "bán kính nguyên tử", "tính acid", "tính base", "hóa trị cao nhất",
        "oxit cao nhất", "oxide cao nhất", "hydroxide tương ứng", "thuộc cùng một nhóm",
        "thuộc cùng một chu kì", "cation", "anion", "ion"
    ]):
        return "TH"

    return "NB"


def resolve_bai_toan_answer_c2(stem_text, sol_text):
    st = stem_text.lower()
    if "phân bón đa lượng npk" in st:
        return "ZA = 15 (Phosphorus); ZB = 16 (Sulfur)"
    if "điện giải" in st or "tham gia trực tiếp" in st or "zx+zy=32" in sol_text.lower():
        return "ZX = 12 (Magnesium); ZY = 20 (Calcium)"
    if "siêu nhẹ" in st or "vỏ máy bay" in st:
        return "ZM = 11 (Sodium); ZQ = 12 (Magnesium)"
    if "xử lý nước và sản xuất kem đánh răng" in st:
        return "ZX = 9 (Fluorine); ZY = 17 (Chlorine)"
    if "chiết xuất được hỗn hợp hai nguyên tố x và y" in st or "công ty hóa chất" in st:
        return "ZX = 12 (Mg); ZY = 20 (Ca); MgO và CaO"
    if "kính cửa sổ" in st or "lọ thủy tinh" in st:
        return "Silicon (Si); Công thức khí: SiH4"
    if "sinh học (biogas)" in st or "khí độc" in st:
        return "Nitrogen (N); MR = 14"
    if "quặng chứa nguyên tố phi kim t" in st:
        return "Sulfur (S); SO3 và H2SO4"
    if "94,12" in st or "trứng thối" in st or "nsnp" in st:
        return "Sulfur (S); MY = 32 g/mol; Hydroxide: H2SO4"
    if "r~2~o~5~" in st or "r2o5" in st:
        return "Phosphorus (P); PH3"
    if "72,73% oxygen" in st or "72,73%" in st:
        return "Carbon (C); CO2 và CH4"
    if "5,4 g một kim loại thuộc nhóm iiia" in st:
        return "Aluminium (Al); m(AlCl3) = 26,7 gam"
    if "lò phản ứng hạt nhân" in st or "tản nhiệt" in st:
        return "Sodium (Na); M = 23 g/mol"
    if "hợp kim" in st and "4,8" in st:
        return "Magnesium (Mg); M = 24 g/mol"
    if "kim loại m thuộc nhóm iia" in st and "1,2" in st:
        return "Calcium (Ca); M = 40 g/mol"
    if "bauxite" in st or "độ tinh khiết" in st:
        return "Độ tinh khiết Al: 94,5%; Na2O > Al2O3; NaOH > Al(OH)3"

    for ln in reversed(sol_text.splitlines()):
        m = re.search(r'(?:Đáp\s+(?:án|số)|Kết\s+luận)\s*[:\.]?\s*([^\n]+)', ln, flags=re.IGNORECASE)
        if m:
            raw_a = m.group(1).strip().rstrip('.')
            m_num = re.match(r'^([+-]?\d+(?:[\,\.]\d+)?)\b', raw_a)
            if m_num and len(raw_a) < 12:
                return m_num.group(1).replace('.', ',')
            return raw_a
    return ""


def parse_cvl_hoa10_c2():
    doc = docx.Document(CVL_DOCX)
    cau_nids = get_cau_num_ids(doc)
    blocks = list(iter_block_items(doc))

    start_idx = 0
    for i, (bt, b) in enumerate(blocks):
        if bt == 'p':
            t = get_para_text_full(b, doc.part)
            if t == 'HỆ THỐNG BÀI TẬP' and start_idx == 0:
                start_idx = i
                break

    end_idx = len(blocks)
    cur_topic = CHAPTER_INFO["topics"][0]
    cur_sec = ""
    cur_part = "P1"
    questions = []
    cur_q = None

    def flush_q():
        nonlocal cur_q
        if cur_q:
            if cur_q["sec"] == "I":
                if len(cur_q["p2_items"]) == 4:
                    cur_q["part"] = "P2"
                    questions.append(cur_q)
            else:
                questions.append(cur_q)
            cur_q = None

    for i in range(start_idx + 1, end_idx):
        bt, b = blocks[i]
        if bt == 'p':
            t = get_para_text_full(b, doc.part)
            imgs = extract_real_images_from_el(b._p, doc.part)
            if not t and not imgs:
                continue

            t_upper = t.upper()
            if "BÀI 5:" in t_upper:
                flush_q()
                cur_topic = CHAPTER_INFO["topics"][0]
                continue
            elif "BÀI 6:" in t_upper:
                flush_q()
                cur_topic = CHAPTER_INFO["topics"][1]
                continue
            elif "BÀI 7:" in t_upper:
                flush_q()
                cur_topic = CHAPTER_INFO["topics"][2]
                continue
            elif "BÀI 8:" in t_upper or "BÀI 8." in t_upper:
                flush_q()
                cur_topic = CHAPTER_INFO["topics"][3]
                continue
            elif "CÁC DẠNG BÀI TOÁN" in t_upper:
                flush_q()
                cur_sec = "BAI_TOAN"
                cur_part = "P3"
                continue
            elif "DẠNG 1:" in t_upper or "DẠNG 1." in t_upper:
                flush_q()
                cur_topic = CHAPTER_INFO["topics"][0]
                cur_sec = "DANG_1"
                cur_part = "P3"
                continue
            elif "DẠNG 2:" in t_upper or "DẠNG 2." in t_upper:
                flush_q()
                cur_topic = CHAPTER_INFO["topics"][2]
                cur_sec = "DANG_2"
                cur_part = "P3"
                continue
            elif "DẠNG 3:" in t_upper or "DẠNG 3." in t_upper:
                flush_q()
                cur_topic = CHAPTER_INFO["topics"][3]
                cur_sec = "DANG_3"
                cur_part = "P3"
                continue

            # Prioritize matching III and IV before II and I to avoid substring clashes
            if "III. BÀI TẬP TRẮC NGHIỆM ĐÚNG SAI" in t_upper or "III. CÂU TRẮC NGHIỆM ĐÚNG SAI" in t_upper or "III. BÀI TẬP ĐÚNG SAI" in t_upper:
                flush_q()
                cur_sec = "III_DS"
                cur_part = "P2"
                continue
            elif "IV. BÀI TẬP TRẢ LỜI NGẮN" in t_upper:
                flush_q()
                cur_sec = "TLN"
                cur_part = "P3"
                continue
            elif "II. BÀI TẬP TRẮC NGHIỆM" in t_upper:
                flush_q()
                cur_sec = "II"
                cur_part = "P1"
                continue
            elif "I. BÀI TẬP ÔN TẬP LÝ THUYẾT" in t_upper:
                flush_q()
                cur_sec = "I"
                cur_part = "P2"
                continue
            elif re.match(r'^[123]\.\s+[A-ZĐÁÀẢÃẠĂẮẰẲẴẶÂẤẦẨẪẬÉÈẺẼẸÊẾỀỂỄỆÍÌỈĨỊÓÒỎÕỌÔỐỒỔỖỘƠỚỜỞỠỢÚÙỦŨỤƯỨỪỬỮỰÝỲỶỸỴ]', t):
                flush_q()
                continue

            nids = b._p.xpath('.//w:numPr/w:numId/@w:val')
            is_cau = bool(nids and nids[0] in cau_nids)
            m_text_cau = re.match(r'^Câu\s+\d+[\.\:]\s*(.*)', t)

            if is_cau or m_text_cau:
                flush_q()
                stem_first = m_text_cau.group(1).strip() if m_text_cau else t
                cur_q = {
                    "topic": cur_topic,
                    "sec": cur_sec,
                    "part": cur_part,
                    "source_tag": "[THPT Châu Văn Liêm – 2026]",
                    "stem_lines": [stem_first] if stem_first else [],
                    "opt_lines": [],
                    "p2_items": {},
                    "sol_lines": [],
                    "underlined_opt": None,
                    "in_sol": False,
                    "img_blobs": list(imgs),
                    "tables": [],
                }
            elif cur_q is not None:
                if imgs and not cur_q["in_sol"]:
                    cur_q["img_blobs"].extend(imgs)
                if not t:
                    continue
                if t.startswith(("📝", "Hướng dẫn giải", "Đáp án và giải thích", "Đáp án:", "Đáp án đúng:")):
                    cur_q["in_sol"] = True
                    cur_q["sol_lines"].append(t)
                elif cur_q["in_sol"]:
                    cur_q["sol_lines"].append(t)
                else:
                    if cur_part == "P1" and re.match(r'^[A-D]\.', t):
                        cur_q["opt_lines"].append(t)
                        u_ans = extract_underlined_p1_answer([b])
                        if u_ans:
                            cur_q["underlined_opt"] = u_ans
                    elif cur_part == "P2" and re.match(r'^[a-d]\)', t):
                        k = t[0].lower()
                        cur_q["p2_items"][k] = t[2:].strip()
                    else:
                        cur_q["stem_lines"].append(t)

        elif bt == 'tbl':
            if len(b.rows) == 1 and len(b.columns) == 1:
                continue

            if cur_q is not None and not cur_q["in_sol"]:
                t_data = []
                for r in b.rows:
                    row_vals = [get_para_text_full(c.paragraphs[0], doc.part) if c.paragraphs else "" for c in r.cells]
                    if any(row_vals):
                        t_data.append(row_vals)
                if t_data:
                    cur_q["tables"].append(t_data)

    flush_q()
    return questions


def parse_ttb_hoa10_c2():
    if not os.path.exists(TTB_DOCX):
        return []
    doc = docx.Document(TTB_DOCX)
    ttb_qs = []
    # Lines 109 to 118 contain Chapter 2 in TTB
    cur_part = "P1"
    cur_lvl = "NB"
    cur_q = None

    def flush_ttb():
        nonlocal cur_q
        if cur_q:
            ttb_qs.append(cur_q)
            cur_q = None

    for i in range(109, 118):
        p = doc.paragraphs[i]
        t = get_para_text_full(p, doc.part)
        imgs = extract_real_images_from_el(p._p, doc.part)
        if not t and not imgs:
            continue
        m_q = re.match(r'^Câu\s+(\d+)[\.\:]\s*(.*)', t)
        if m_q:
            flush_ttb()
            rest = m_q.group(2).strip()
            src_m = re.match(r'^(\[[^\]]+\])\s*(.*)', rest)
            src_tag = src_m.group(1) if src_m else "[TTB – 2026]"
            stem_1 = src_m.group(2).strip() if src_m else rest
            cur_q = {
                "topic": CHAPTER_INFO["topics"][3] if "37" in m_q.group(1) else CHAPTER_INFO["topics"][0],
                "sec": "TTB",
                "part": "P1",
                "level_override": "NB" if "37" in m_q.group(1) else "TH",
                "source_tag": src_tag,
                "stem_lines": [stem_1] if stem_1 else [],
                "opt_lines": [],
                "p2_items": {},
                "sol_lines": [],
                "underlined_opt": None,
                "in_sol": False,
                "img_blobs": list(imgs),
                "tables": [],
            }
        elif cur_q is not None:
            if re.match(r'^[A-D]\.', t):
                cur_q["opt_lines"].append(t)
                u_ans = extract_underlined_p1_answer([p])
                if u_ans:
                    cur_q["underlined_opt"] = u_ans
            else:
                cur_q["stem_lines"].append(t)

    flush_ttb()
    # Explicit answer assignment for TTB Q37 and Q38
    if len(ttb_qs) >= 2:
        ttb_qs[0]["underlined_opt"] = "A"  # Định luật tuần hoàn
        ttb_qs[1]["underlined_opt"] = "B"  # Nhóm IA, chu kì 3 (Na)
    return ttb_qs


def finalize_question_record_c2(raw_q):
    sol_text = "\n".join(raw_q["sol_lines"]).strip()
    sol_text = re.sub(r'^📝\s*Hướng dẫn giải:?\s*', '', sol_text).strip()
    sol_text = re.sub(r'^Hướng dẫn giải:?\s*', '', sol_text).strip()
    sol_text = re.sub(r'^Đáp án và giải thích:?\s*', '', sol_text).strip()

    ans = ""
    if raw_q["part"] == "P1":
        ans = raw_q["underlined_opt"] or ""
        if not ans:
            m = re.search(r'Đáp\s+(?:án|số)\s*[:\.]?\s*([A-D])\b', sol_text, flags=re.IGNORECASE)
            if m:
                ans = m.group(1).upper()
            else:
                m2 = re.search(r'\b([A-D])\s+(?:đúng|sai)', sol_text)
                if m2:
                    ans = m2.group(1).upper()
    elif raw_q["part"] == "P2":
        p2_ans = {}
        m_inline = re.findall(r'([a-d])\s*[\-\:]\s*(ĐÚNG|SAI|Đúng|Sai)', sol_text)
        for k_m, v_m in m_inline:
            p2_ans[k_m.lower()] = "Đ" if v_m.upper() == "ĐÚNG" else "S"
        # Also parse individual line items "a) Đúng. / a) Sai."
        lines_ds = re.findall(r'(?:^|\n)\s*([a-d])[\)\.]\s*(Đúng|Sai)', sol_text, flags=re.IGNORECASE)
        for k_l, v_l in lines_ds:
            k_clean = k_l.lower()
            if k_clean not in p2_ans:
                p2_ans[k_clean] = "Đ" if v_l.lower().startswith("đ") else "S"

        # Special fallback for Bai 5 Q52 where typo has "b) Đúng." instead of "c) Đúng."
        if len(p2_ans) < 4 and "ne là nguyên tố khí hiếm" in sol_text.lower():
            p2_ans = {"a": "Đ", "b": "S", "c": "Đ", "d": "S"}

        if len(p2_ans) == 4:
            ans = f"a:{p2_ans['a']} - b:{p2_ans['b']} - c:{p2_ans['c']} - d:{p2_ans['d']}"
    elif raw_q["part"] == "P3":
        if raw_q["sec"] in ("DANG_1", "DANG_2", "DANG_3"):
            ans = resolve_bai_toan_answer_c2(" ".join(raw_q["stem_lines"]), sol_text)
        else:
            for ln in raw_q["sol_lines"]:
                m = re.search(r'(?:Đáp\s+(?:án|số)|Kết\s+luận)\s*[:\.]?\s*([^\n\.]+)', ln, flags=re.IGNORECASE)
                if m:
                    raw_a = m.group(1).strip()
                    m_num = re.match(r'^([+-]?\d+(?:[\,\.]\d+)?)', raw_a)
                    if m_num:
                        ans = m_num.group(1).replace('.', ',')
                    else:
                        ans = raw_a
                    break

    item_levels = {}
    if raw_q["part"] == "P2":
        for k in ["a", "b", "c", "d"]:
            item_levels[k] = classify_p2_item_level_hoa10_c2(raw_q["p2_items"].get(k, ""))

    lvl = raw_q.get("level_override") or classify_question_level_hoa10_c2(
        raw_q["part"],
        raw_q["sec"],
        raw_q["topic"],
        " ".join(raw_q["stem_lines"]),
        " ".join(raw_q["opt_lines"]),
        sol_text,
        item_levels
    )

    return {
        "part": raw_q["part"],
        "topic": raw_q["topic"],
        "level": lvl,
        "item_levels": item_levels,
        "source_tag": raw_q["source_tag"],
        "stem_lines": raw_q["stem_lines"],
        "opt_lines": raw_q["opt_lines"],
        "p2_items": raw_q["p2_items"],
        "answer": ans,
        "solution": sol_text,
        "img_blobs": raw_q["img_blobs"],
        "tables": raw_q["tables"],
    }


def add_solution_block_hoa10_c2(doc, q):
    p_hd = doc.add_paragraph()
    p_hd.paragraph_format.left_indent = Cm(0.6)
    p_hd.paragraph_format.space_before = Pt(2)
    p_hd.paragraph_format.space_after = Pt(1)

    r_hd = p_hd.add_run("Hướng dẫn giải: ")
    r_hd.bold = True
    r_hd.font.name = "Times New Roman"
    r_hd.font.size = Pt(11)
    r_hd.font.color.rgb = COLOR_SOLUTION

    if q["part"] == "P1" and q["answer"]:
        r_ans = p_hd.add_run(f"Chọn {q['answer']}.")
        r_ans.bold = True
        r_ans.font.name = "Times New Roman"
        r_ans.font.size = Pt(11)
        r_ans.font.color.rgb = COLOR_CORRECT
    elif q["part"] == "P2" and q["answer"]:
        r_ans = p_hd.add_run(f"Đáp án: {q['answer']}")
        r_ans.bold = True
        r_ans.font.name = "Times New Roman"
        r_ans.font.size = Pt(11)
        r_ans.font.color.rgb = COLOR_CORRECT
    elif q["part"] == "P3" and q["answer"]:
        r_ans = p_hd.add_run("Đáp án: ")
        r_ans.bold = True
        r_ans.font.name = "Times New Roman"
        r_ans.font.size = Pt(11)
        r_ans.font.color.rgb = COLOR_CORRECT
        add_chem_runs_hoa10(p_hd, str(q["answer"]), bold=True, color=COLOR_CORRECT, size=11)

    sol = (q.get("solution") or "").strip()
    if sol:
        for line in sol.splitlines():
            line_s = line.strip()
            if not line_s:
                continue
            if re.match(r'^Đáp\s+(?:án|số)\s*[:\.]?\s*[A-D0-9\,\.\s]+$', line_s, flags=re.IGNORECASE):
                continue
            p_ln = doc.add_paragraph()
            p_ln.paragraph_format.left_indent = Cm(0.6)
            p_ln.paragraph_format.space_before = Pt(0.5)
            p_ln.paragraph_format.space_after = Pt(1.5)
            add_chem_runs_hoa10(p_ln, line_s, size=11)


def build_word_bank_hoa10_c2(questions, is_teacher):
    doc = docx.Document()
    setup_page_and_hf(doc)

    # 1. Title Block
    p_t1 = doc.add_paragraph()
    p_t1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t1.paragraph_format.space_before = Pt(0)
    p_t1.paragraph_format.space_after = Pt(2)
    r_t1 = p_t1.add_run(f"NGÂN HÀNG CÂU HỎI HÓA HỌC 10 – {CHAPTER_INFO['name']}")
    r_t1.bold = True
    r_t1.font.name = "Times New Roman"
    r_t1.font.size = Pt(14)
    r_t1.font.color.rgb = COLOR_PRIMARY

    p_t2 = doc.add_paragraph()
    p_t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t2.paragraph_format.space_before = Pt(0)
    p_t2.paragraph_format.space_after = Pt(2)
    r_t2 = p_t2.add_run("PHÂN LOẠI THEO CHỦ ĐỀ, DẠNG THỨC & MỨC ĐỘ NHẬN THỨC (BIẾT – HIỂU – VẬN DỤNG)")
    r_t2.bold = True
    r_t2.font.name = "Times New Roman"
    r_t2.font.size = Pt(10.5)
    r_t2.font.color.rgb = COLOR_SOLUTION

    p_t3 = doc.add_paragraph()
    p_t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t3.paragraph_format.space_before = Pt(0)
    p_t3.paragraph_format.space_after = Pt(6)
    ver_note = (
        "(BẢN GIÁO VIÊN — CÓ ĐÁP ÁN, LỜI GIẢI CHI TIẾT & MÃ ĐỊNH DANH MA TRẬN)"
        if is_teacher else
        "(BẢN HỌC SINH — ĐỀ ÔN LUYỆN PHÂN LOẠI THEO MỨC ĐỘ NHẬN THỨC)"
    )
    r_t3 = p_t3.add_run(ver_note)
    r_t3.bold = True
    r_t3.italic = True
    r_t3.font.name = "Times New Roman"
    r_t3.font.size = Pt(11)
    r_t3.font.color.rgb = COLOR_HF_RED

    # 2. Matrix Summary Table
    tbl = doc.add_table(rows=1, cols=7)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl._tbl.tblPr.append(parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="8BC390"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="8BC390"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="8BC390"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="8BC390"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="8BC390"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="8BC390"/>'
        f'</w:tblBorders>'
    ))
    headers = ["Chủ đề / Bài học", "Dạng thức", "Biết (NB)", "Hiểu (TH)", "Vận dụng (VD)", "Tổng số câu", "Phạm vi câu"]
    for i, h in enumerate(headers):
        p = tbl.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.font.color.rgb = COLOR_PRIMARY

    for part_code, part_label in [
        ("P1", "Phần I (TN 4 lựa chọn)"),
        ("P2", "Phần II (Đúng / Sai)"),
        ("P3", "Phần III (Trả lời ngắn & BT)")
    ]:
        for topic in CHAPTER_INFO["topics"]:
            sub_qs = [q for q in questions if q["part"] == part_code and q["topic"] == topic]
            if not sub_qs:
                continue
            row_cells = tbl.add_row().cells
            nb = sum(1 for q in sub_qs if q["level"] == "NB")
            th = sum(1 for q in sub_qs if q["level"] == "TH")
            vd = sum(1 for q in sub_qs if q["level"] == "VD")
            q_nums = [q["q_num"] for q in sub_qs]
            q_range = f"Câu {min(q_nums)} – {max(q_nums)}" if len(q_nums) > 1 else f"Câu {q_nums[0]}"
            if part_code == "P2":
                it_nb = sum(1 for q in sub_qs for lv in q.get("item_levels", {}).values() if lv == "NB")
                it_th = sum(1 for q in sub_qs for lv in q.get("item_levels", {}).values() if lv == "TH")
                it_vd = sum(1 for q in sub_qs for lv in q.get("item_levels", {}).values() if lv == "VD")
                nb_str = f"{nb} câu ({it_nb} ý)"
                th_str = f"{th} câu ({it_th} ý)"
                vd_str = f"{vd} câu ({it_vd} ý)"
            else:
                nb_str, th_str, vd_str = str(nb), str(th), str(vd)

            vals = [topic, part_label, nb_str, th_str, vd_str, f"{len(sub_qs)} câu", q_range]
            for i, val in enumerate(vals):
                p = row_cells[i].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i in (0, 1) else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(1.5)
                p.paragraph_format.space_after = Pt(1.5)
                r = p.add_run(val)
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.5)
                if i == 5:
                    r.bold = True

    tot_cells = tbl.add_row().cells
    tot_nb = sum(1 for q in questions if q["level"] == "NB")
    tot_th = sum(1 for q in questions if q["level"] == "TH")
    tot_vd = sum(1 for q in questions if q["level"] == "VD")
    tot_vals = [
        "TỔNG CỘNG TOÀN CHƯƠNG 2",
        "3 Phần (I, II, III)",
        f"{tot_nb} câu",
        f"{tot_th} câu",
        f"{tot_vd} câu",
        f"{len(questions)} câu",
        f"Câu 1 – {len(questions)}"
    ]
    for i, val in enumerate(tot_vals):
        p = tot_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 1 else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(val)
        r.bold = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.font.color.rgb = COLOR_HF_RED

    # 3. Render Questions by Part & Topic
    part_headers = {
        "P1": "PHẦN I. CÂU TRẮC NGHIỆM NHIỀU PHƯƠNG ÁN LỰA CHỌN",
        "P2": "PHẦN II. CÂU TRẮC NGHIỆM ĐÚNG / SAI (4 Ý a, b, c, d)",
        "P3": "PHẦN III. CÂU TRẮC NGHIỆM TRẢ LỜI NGẮN & BÀI TOÁN VẬN DỤNG"
    }
    cur_part = None
    cur_topic = None

    for q in questions:
        if q["part"] != cur_part:
            cur_part = q["part"]
            cur_topic = None
            p_pt = doc.add_paragraph()
            p_pt.paragraph_format.space_before = Pt(10)
            p_pt.paragraph_format.space_after = Pt(4)
            r_pt = p_pt.add_run(part_headers[cur_part])
            r_pt.bold = True
            r_pt.font.name = "Times New Roman"
            r_pt.font.size = Pt(12.5)
            r_pt.font.color.rgb = COLOR_HF_RED

        if q["topic"] != cur_topic:
            cur_topic = q["topic"]
            p_tp = doc.add_paragraph()
            p_tp.paragraph_format.space_before = Pt(6)
            p_tp.paragraph_format.space_after = Pt(3)
            r_tp = p_tp.add_run(f"♦ CHỦ ĐỀ: {cur_topic.upper()}")
            r_tp.bold = True
            r_tp.font.name = "Times New Roman"
            r_tp.font.size = Pt(11.5)
            r_tp.font.color.rgb = COLOR_PRIMARY

        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)

        r_qn = p_q.add_run(f"Câu {q['q_num']}. ")
        r_qn.bold = True
        r_qn.font.name = "Times New Roman"
        r_qn.font.size = Pt(11.5)
        r_qn.font.color.rgb = COLOR_PRIMARY

        lv_name = LEVEL_NAMES[q["level"]]
        ans_str = f" | ĐA: {q['answer']}" if (is_teacher and q["answer"]) else ""
        r_tag = p_q.add_run(f"[{q['id']} | {lv_name}{ans_str}] ")
        r_tag.bold = True
        r_tag.font.name = "Times New Roman"
        r_tag.font.size = Pt(10.5)
        r_tag.font.color.rgb = COLOR_TAG

        if q.get("source_tag") and "Châu Văn Liêm" not in q["source_tag"]:
            r_src = p_q.add_run(f"{q['source_tag']} ")
            r_src.bold = True
            r_src.font.name = "Times New Roman"
            r_src.font.size = Pt(11)
            r_src.font.color.rgb = COLOR_SOLUTION

        if q["stem_lines"]:
            add_chem_runs_hoa10(p_q, q["stem_lines"][0], size=11.5)
            for extra_ln in q["stem_lines"][1:]:
                p_ex = doc.add_paragraph()
                p_ex.paragraph_format.left_indent = Cm(0.4)
                p_ex.paragraph_format.space_before = Pt(1)
                p_ex.paragraph_format.space_after = Pt(1)
                add_chem_runs_hoa10(p_ex, extra_ln, size=11.5)

        for img_p in q.get("image_paths", []):
            if os.path.exists(img_p):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(2)
                p_img.paragraph_format.space_after = Pt(2)
                w_cm = Cm(3.5) if len(q["image_paths"]) >= 4 else Cm(6.8)
                p_img.add_run().add_picture(img_p, width=w_cm)

        for tbl_d in q.get("tables", []):
            add_data_table_to_doc(doc, tbl_d)

        if q["part"] == "P1":
            opts = parse_abcd_from_opt_lines(q["opt_lines"])
            add_options_paragraph_no_table(doc, opts, q["answer"], is_teacher)
            if is_teacher:
                add_solution_block_hoa10_c2(doc, q)

        elif q["part"] == "P2":
            ans_map = {}
            for part_ans in q["answer"].split("-"):
                if ":" in part_ans:
                    k_a, v_a = part_ans.strip().split(":", 1)
                    ans_map[k_a.strip()] = v_a.strip()
            for key in ["a", "b", "c", "d"]:
                stmt_txt = q.get("p2_items", {}).get(key, "")
                p_it = doc.add_paragraph()
                p_it.paragraph_format.left_indent = Cm(0.4)
                p_it.paragraph_format.space_before = Pt(1)
                p_it.paragraph_format.space_after = Pt(1.5)
                r_k = p_it.add_run(f"{key}) ")
                r_k.bold = True
                r_k.font.name = "Times New Roman"
                r_k.font.size = Pt(11.5)
                il_lv = q.get("item_levels", {}).get(key, "TH")
                r_mt = p_it.add_run(f"[{LEVEL_NAMES.get(il_lv, 'Hiểu')}] ")
                r_mt.bold = True
                r_mt.font.name = "Times New Roman"
                r_mt.font.size = Pt(10.5)
                r_mt.font.color.rgb = COLOR_SOLUTION
                add_chem_runs_hoa10(p_it, stmt_txt, size=11.5)
                if is_teacher and key in ans_map:
                    ds_val = "Đúng" if ans_map[key] == "Đ" else "Sai"
                    r_ds = p_it.add_run(f"   → [{ds_val}]")
                    r_ds.bold = True
                    r_ds.font.name = "Times New Roman"
                    r_ds.font.size = Pt(11)
                    r_ds.font.color.rgb = COLOR_CORRECT
            if is_teacher:
                add_solution_block_hoa10_c2(doc, q)

        elif q["part"] == "P3":
            if not is_teacher:
                p_ans_line = doc.add_paragraph()
                p_ans_line.paragraph_format.left_indent = Cm(0.4)
                p_ans_line.paragraph_format.space_before = Pt(1)
                p_ans_line.paragraph_format.space_after = Pt(2)
                r_al = p_ans_line.add_run("Đáp số / Trả lời: .......................................................................................")
                r_al.italic = True
                r_al.font.name = "Times New Roman"
                r_al.font.size = Pt(11)
            else:
                add_solution_block_hoa10_c2(doc, q)

    suffix = "GiaoVien" if is_teacher else "HocSinh"
    fname = f"{CHAPTER_INFO['file_prefix']}_{suffix}.docx"
    out_p1 = os.path.join(OUT_DIR, fname)
    out_p2 = os.path.join(BANK_DIR, fname)
    doc.save(out_p1)
    doc.save(out_p2)
    print(f"  + Saved {suffix} Word file: {out_p2}")


if __name__ == "__main__":
    raw_cvl = parse_cvl_hoa10_c2()
    raw_ttb = parse_ttb_hoa10_c2()
    print(f"Parsed CVL: {len(raw_cvl)} questions | Parsed TTB: {len(raw_ttb)} questions")

    all_raw = raw_cvl + raw_ttb
    deduped = []
    seen_norms = []
    dup_count = 0

    for rq in all_raw:
        fq = finalize_question_record_c2(rq)
        st_norm = norm_for_compare(" ".join(fq["stem_lines"]))
        opt_norm = norm_for_compare(" ".join(fq["opt_lines"]))
        full_sig = f"{fq['part']}||{st_norm}||{opt_norm}"
        is_dup = False
        for prev_part, prev_st, prev_sig in seen_norms:
            if fq["part"] == prev_part:
                if full_sig == prev_sig:
                    is_dup = True
                    break
                if rq.get("sec") == "TTB" and (st_norm == prev_st or SequenceMatcher(None, st_norm, prev_st).ratio() >= 0.85):
                    is_dup = True
                    break
        if is_dup:
            dup_count += 1
        else:
            seen_norms.append((fq["part"], st_norm, full_sig))
            deduped.append(fq)

    print(f"Filtered out {dup_count} duplicate questions -> {len(deduped)} unique questions!")

    part_order = {"P1": 1, "P2": 2, "P3": 3}
    topic_order = {t: i for i, t in enumerate(CHAPTER_INFO["topics"])}
    level_order = {"NB": 1, "TH": 2, "VD": 3}

    deduped.sort(key=lambda q: (
        part_order.get(q["part"], 9),
        topic_order.get(q["topic"], 9),
        level_order.get(q["level"], 9)
    ))

    final_questions = []
    for q_num, q in enumerate(deduped, 1):
        qid = f"H10-C2-{q['part']}-{q['level']}-{q_num:03d}"
        img_paths = []
        for k, blob in enumerate(q.get("img_blobs", []), 1):
            fpath = os.path.join(IMG_DIR, f"h10_c2_q{q_num:03d}_img{k}.png")
            try:
                im = Image.open(io.BytesIO(blob))
                im.save(fpath, format="PNG")
                img_paths.append(fpath)
            except Exception:
                pass
        q_rec = {
            "q_num": q_num,
            "id": qid,
            "part": q["part"],
            "level": q["level"],
            "topic": q["topic"],
            "source_tag": q["source_tag"],
            "stem_lines": q["stem_lines"],
            "opt_lines": q["opt_lines"],
            "p2_items": q["p2_items"],
            "item_levels": q["item_levels"],
            "answer": q["answer"],
            "solution": q["solution"],
            "image_paths": img_paths,
            "tables": q["tables"],
        }
        final_questions.append(q_rec)

    # Load existing Tron Bo JSON if available and update
    json_c2_path = os.path.join(BANK_DIR, "Ngan_Hang_Hoa_10_Chuong_2_Bang_Tuan_Hoan.json")
    json_full_path = os.path.join(BANK_DIR, "Ngan_Hang_Hoa_10_Tron_Bo.json")

    bank_data_c2 = {
        "C2": {
            "grade": 10,
            "chapter_number": 2,
            "chapter_name": CHAPTER_INFO["name"],
            "topics": CHAPTER_INFO["topics"],
            "total_questions": len(final_questions),
            "questions": final_questions,
        }
    }

    with open(json_c2_path, "w", encoding="utf-8") as f:
        json.dump(bank_data_c2, f, ensure_ascii=False, indent=2)

    full_bank = {}
    if os.path.exists(json_full_path):
        try:
            with open(json_full_path, "r", encoding="utf-8") as f:
                full_bank = json.load(f)
        except Exception:
            full_bank = {}
    full_bank["C2"] = bank_data_c2["C2"]
    with open(json_full_path, "w", encoding="utf-8") as f:
        json.dump(full_bank, f, ensure_ascii=False, indent=2)

    print(f"Saved JSON Question Bank ({len(final_questions)} questions): {json_c2_path}")

    build_word_bank_hoa10_c2(final_questions, is_teacher=False)
    build_word_bank_hoa10_c2(final_questions, is_teacher=True)

    print("\n=== THỐNG KÊ CHI TIẾT NGÂN HÀNG HÓA 10 - CHƯƠNG 2 ===")
    for pt in ["P1", "P2", "P3"]:
        sub = [q for q in final_questions if q["part"] == pt]
        nb = sum(1 for q in sub if q["level"] == "NB")
        th = sum(1 for q in sub if q["level"] == "TH")
        vd = sum(1 for q in sub if q["level"] == "VD")
        print(f"  + {pt}: {len(sub)} câu (Biết: {nb}, Hiểu: {th}, Vận dụng: {vd})")
