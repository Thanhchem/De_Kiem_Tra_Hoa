# -*- coding: utf-8 -*-
"""
Script tạo trọn bộ 10 Đề ôn tập Chương 1, 2, 3 Hóa học 12 dạng Word (.docx)
theo đúng chuẩn FORM ĐỀ KIỂM TRA (Thầy TRẦN VĂN THẠNH).
Xuất ra:
- 10 file lẻ Bản Học sinh
- 10 file lẻ Bản Giáo viên
- 01 file gộp lớn Bản Học sinh (trọn bộ 10 đề)
- 01 file gộp lớn Bản Giáo viên (trọn bộ 10 đề)
"""

import os
import sys
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

# Import module form chuẩn
sys.path.append(os.path.abspath(r"C:\Antigravity_Thanh\scripts"))
from form_de_kiem_tra import (
    setup_header_footer,
    build_title_block,
    add_section_header,
    render_question_prompt,
    render_options_no_table,
    render_statement,
    add_solution_no_table,
    add_student_writing_lines,
    render_rich_text,
    set_run_style,
    COLOR_PRIMARY,
    COLOR_CORRECT,
    COLOR_SOLUTION
)

from parse_and_build_10_de import parse_all_exams
from solutions_10_de import SOLUTIONS

OUT_DIR = r"C:\Antigravity_Thanh\San_Pham"
SUB_DIR = os.path.join(OUT_DIR, "10_De_On_Tap_Chuong_123_Hoa12")
os.makedirs(SUB_DIR, exist_ok=True)

def render_single_exam_into_doc(doc, exam, is_teacher=False, is_first=True):
    de_num = exam["de_num"]
    sol_dict = SOLUTIONS.get(de_num, {})
    sol_p1 = sol_dict.get("p1", {})
    sol_p2 = sol_dict.get("p2", {})
    sol_p3 = sol_dict.get("p3", {})

    # Nếu không phải đề đầu tiên trong file gộp, thêm section mới
    if not is_first:
        section = doc.add_section(WD_SECTION_START.NEW_PAGE)
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)
        section.header_distance = Cm(0.6)
        section.footer_distance = Cm(0.6)
        # Giữ header & footer kế thừa

    # 1. Khối tiêu đề
    exam_code_str = f"10{de_num:02d}" if de_num < 10 else f"1{de_num:02d}"
    build_title_block(
        doc,
        school_line="TỔNG HỢP ÔN TẬP ĐỊNH KỲ HOÁ HỌC 12",
        exam_title=f"ĐỀ ÔN TẬP CHƯƠNG 1, 2, 3 - ĐỀ SỐ {de_num:02d}",
        subject_line="MÔN: HOÁ HỌC - LỚP 12",
        num_pages="04",
        duration="50 phút",
        exam_code=exam_code_str,
        atomic_masses="H = 1; C = 12; N = 14; O = 16; Na = 23; Mg = 24; S = 32; Cl = 35,5; K = 39; Ca = 40; Fe = 56; Cu = 64; Br = 80; Ag = 108",
        is_teacher=is_teacher
    )

    # 2. PHẦN I
    add_section_header(
        doc,
        "PHẦN I. Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án.",
        subtitle="Mỗi câu trả lời đúng thí sinh được 0,25 điểm"
    )

    for q in exam["part1"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1)
        render_question_prompt(p, f"Câu {q['idx']}. {q['prompt']}")

        # Chuẩn bị options
        formatted_opts = []
        for opt in q["opts"]:
            opt_letter = opt[:1] # A, B, C, D
            if is_teacher and opt_letter == q["correct"]:
                # Đánh dấu đáp án đúng
                formatted_opts.append(f"*{opt}*")
            else:
                formatted_opts.append(opt)

        render_options_no_table(doc, formatted_opts, is_teacher=is_teacher)

        if is_teacher:
            sol_text = sol_p1.get(q["idx"], f"➔ Đáp án đúng: {q['correct']}.")
            add_solution_no_table(doc, sol_text)

    # 3. PHẦN II
    add_section_header(
        doc,
        "PHẦN II. Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.",
        subtitle="Điểm tối đa cho 01 câu hỏi là 1,0 điểm (Đúng 1 ý: 0,1đ; 2 ý: 0,25đ; 3 ý: 0,5đ; 4 ý: 1,0đ)"
    )

    for q in exam["part2"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        render_question_prompt(p, f"Câu {q['idx']}. {q['intro']}")

        for stmt_text, ans_val in q["items"]:
            if not stmt_text:
                continue
            p_st = doc.add_paragraph()
            p_st.paragraph_format.left_indent = Cm(0.5)
            p_st.paragraph_format.space_before = Pt(1)
            p_st.paragraph_format.space_after = Pt(1)
            render_statement(p_st, stmt_text, ans=ans_val, is_teacher=is_teacher)

        if is_teacher:
            sol_text = sol_p2.get(q["idx"], "")
            if sol_text:
                add_solution_no_table(doc, sol_text)

    # 4. PHẦN III
    add_section_header(
        doc,
        "PHẦN III. Thí sinh trả lời từ câu 1 đến câu 6.",
        subtitle="Mỗi câu trả lời đúng thí sinh được 0,25 điểm"
    )

    for q in exam["part3"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        render_question_prompt(p, f"Câu {q['idx']}. {q['prompt']}")

        if not is_teacher:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.space_before = Pt(2)
            p_ans.paragraph_format.space_after = Pt(2)
            render_rich_text(p_ans, "*Trả lời:* ...................................................................................................................................................")
        else:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.space_before = Pt(2)
            p_ans.paragraph_format.space_after = Pt(2)
            render_rich_text(p_ans, f"*Trả lời:* *{q['ans']}*", base_color=COLOR_CORRECT)
            sol_text = sol_p3.get(q["idx"], "")
            if sol_text:
                add_solution_no_table(doc, sol_text)

    # Dòng kết thúc đề
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(8)
    p_end.paragraph_format.space_after = Pt(8)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run(f"-------------------------- HẾT ĐỀ SỐ {de_num:02d} --------------------------")
    set_run_style(r_end, font_size=10.5, bold=True)

def build_all():
    print("=== ĐANG PHÂN TÍCH VÀ CHUẨN HÓA DỮ LIỆU 10 ĐỀ ===")
    exams = parse_all_exams()
    print(f"Đã nạp {len(exams)} đề thi.")

    # 1. Tạo 10 file lẻ Học sinh và Giáo viên
    for exam in exams:
        de_num = exam["de_num"]
        print(f"-> Đang tạo Đề {de_num:02d}...")

        # File Học sinh
        doc_hs = docx.Document()
        setup_header_footer(doc_hs)
        render_single_exam_into_doc(doc_hs, exam, is_teacher=False, is_first=True)
        hs_path = os.path.join(SUB_DIR, f"De_On_Tap_Hoa_12_De{de_num:02d}_HocSinh.docx")
        doc_hs.save(hs_path)

        # File Giáo viên
        doc_gv = docx.Document()
        setup_header_footer(doc_gv)
        render_single_exam_into_doc(doc_gv, exam, is_teacher=True, is_first=True)
        gv_path = os.path.join(SUB_DIR, f"De_On_Tap_Hoa_12_De{de_num:02d}_GiaoVien.docx")
        doc_gv.save(gv_path)

    print("=== ĐÃ TẠO XONG 20 FILE LẺ (10 HỌC SINH + 10 GIÁO VIÊN) ===")

    # 2. Tạo 2 file gộp lớn trọn bộ 10 Đề
    print("-> Đang tạo file gộp lớn 10 đề Bản Học sinh...")
    doc_full_hs = docx.Document()
    setup_header_footer(doc_full_hs)
    for i, exam in enumerate(exams):
        render_single_exam_into_doc(doc_full_hs, exam, is_teacher=False, is_first=(i == 0))
    full_hs_path = os.path.join(OUT_DIR, "De_On_Tap_Hoa_12_Chuong_123_Tron_Bo_10_De_HocSinh.docx")
    doc_full_hs.save(full_hs_path)
    print(f"Đã lưu: {full_hs_path}")

    print("-> Đang tạo file gộp lớn 10 đề Bản Giáo viên...")
    doc_full_gv = docx.Document()
    setup_header_footer(doc_full_gv)
    for i, exam in enumerate(exams):
        render_single_exam_into_doc(doc_full_gv, exam, is_teacher=True, is_first=(i == 0))
    full_gv_path = os.path.join(OUT_DIR, "De_On_Tap_Hoa_12_Chuong_123_Tron_Bo_10_De_GiaoVien.docx")
    doc_full_gv.save(full_gv_path)
    print(f"Đã lưu: {full_gv_path}")

    print("=== HOÀN TẤT TẤT CẢ FILE WORD! ===")

if __name__ == "__main__":
    build_all()
