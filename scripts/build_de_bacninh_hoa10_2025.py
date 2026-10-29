# -*- coding: utf-8 -*-
"""
Script tạo Đề Kiểm Tra Giữa Học Kì 1 Hóa Học 10 - Bắc Ninh 2025-2026
Chuẩn 100% FORM ĐỀ KIỂM TRA (Thầy TRẦN VĂN THẠNH)
Sinh cả 2 phiên bản:
1. Bản Học sinh: De_Kiem_Tra_GK1_Hoa_10_BacNinh_2025_HocSinh.docx
2. Bản Giáo viên: De_Kiem_Tra_GK1_Hoa_10_BacNinh_2025_GiaoVien.docx
"""

import os
import sys

if sys.stdout.encoding.lower() != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

sys.path.append(os.path.abspath(r"C:\Antigravity_Thanh\scripts"))
from form_de_kiem_tra import (
    setup_header_footer,
    build_title_block,
    add_section_header,
    render_question_prompt,
    render_options_no_table,
    render_statement,
    add_solution_no_table,
    add_student_writing_lines,
    render_rich_text,
    set_run_style,
    COLOR_PRIMARY,
    COLOR_CORRECT,
    COLOR_SOLUTION
)

IMG_DIR = r"C:\Antigravity_Thanh\San_Pham\images_gk1_hoa10_bacninh"

def add_centered_image(doc, img_path, width_cm=8.0):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(img_path, width=Cm(width_cm))

def build_exam_bacninh_10(is_teacher=False, output_path=None):
    doc = docx.Document()
    setup_header_footer(doc)

    # 1. Khối tiêu đề
    build_title_block(
        doc,
        school_line="SỞ GD&ĐT BẮC NINH",
        exam_title="KIỂM TRA GIỮA HỌC KÌ 1 - NĂM HỌC 2025 - 2026",
        subject_line="MÔN: HÓA HỌC 10",
        num_pages="03",
        duration="45 phút",
        exam_code="102",
        atomic_masses="H = 1; C = 12; N = 14; O = 16; Na = 23; Mg = 24; Al = 27; P = 31; S = 32; Cl = 35,5; K = 39; Ca = 40; Fe = 56; Cu = 63,54",
        is_teacher=is_teacher
    )

    # =========================================================================
    # PHẦN I
    # =========================================================================
    add_section_header(
        doc,
        "PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án."
    )

    # Câu 1
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 1. Số neutron trong nguyên tử ^56_26Fe là")
    opts1 = [
        "*A. 30.*" if is_teacher else "A. 30.",
        "B. 56.",
        "C. 26.",
        "D. 13."
    ]
    render_options_no_table(doc, opts1, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Số neutron: N = A - Z = 56 - 26 = 30.\n➔ Đáp án đúng: A.")

    # Câu 2
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 2. Số electron tối đa trong phân lớp p là")
    opts2 = [
        "A. 2.",
        "B. 10.",
        "C. 14.",
        "*D. 6.*" if is_teacher else "D. 6."
    ]
    render_options_no_table(doc, opts2, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Số electron tối đa trong các phân lớp: s tối đa 2; p tối đa 6; d tối đa 10; f tối đa 14.\n➔ Đáp án đúng: D.")

    # Câu 3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 3. Sự biến đổi nào sau đây không phải là một hiện tượng hóa học?")
    opts3 = [
        "A. Than củi cháy, tạo thành khí carbonic và hơi nước.",
        "*B. Nước đá hóa lỏng khi để ở nhiệt độ thường.*" if is_teacher else "B. Nước đá hóa lỏng khi để ở nhiệt độ thường.",
        "C. Lá nhôm tan trong dung dịch acid HCl.",
        "D. Khí hydrogen cháy trong oxygen, tạo thành nước."
    ]
    render_options_no_table(doc, opts3, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Nước đá hóa lỏng là hiện tượng vật lí (sự chuyển thể từ thể rắn sang thể lỏng), không tạo thành chất mới.\n➔ Đáp án đúng: B.")

    # Câu 4
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 4. Trong nguyên tử, loại hạt cơ bản mang điện âm là")
    opts4 = [
        "A. neutron.",
        "*B. electron.*" if is_teacher else "B. electron.",
        "C. proton, electron.",
        "D. proton."
    ]
    render_options_no_table(doc, opts4, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Trong nguyên tử: hạt proton mang điện tích dương (+1), hạt electron mang điện tích âm (-1), hạt neutron không mang điện.\n➔ Đáp án đúng: B.")

    # Câu 5
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 5. Số nguyên tố hoá học có trong chu kì 1 của bảng tuần hoàn là")
    opts5 = [
        "*A. 2 nguyên tố.*" if is_teacher else "A. 2 nguyên tố.",
        "B. 18 nguyên tố.",
        "C. 8 nguyên tố.",
        "D. 20 nguyên tố."
    ]
    render_options_no_table(doc, opts5, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Chu kì 1 chỉ gồm 2 nguyên tố là H (Z = 1) và He (Z = 2).\n➔ Đáp án đúng: A.")

    # Câu 6
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 6. Nguyên tố X thuộc nhóm VIA của bảng tuần hoàn. Số electron ở lớp ngoài cùng của nguyên tử X là")
    opts6 = [
        "*A. 6.*" if is_teacher else "A. 6.",
        "B. 4.",
        "C. 5.",
        "D. 2."
    ]
    render_options_no_table(doc, opts6, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Đối với các nguyên tố nhóm A, số electron lớp ngoài cùng = số thứ tự nhóm = 6.\n➔ Đáp án đúng: A.")

    # Câu 7
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 7. “Các ...(1)... của cùng một nguyên tố hóa học có số ...(2)... khác nhau là đồng vị của nhau”. Nội dung phù hợp trong các ô trống (1), (2) lần lượt là")
    opts7 = [
        "A. nguyên tử, proton.",
        "B. nguyên tố, proton.",
        "C. nguyên tố, neutron.",
        "*D. nguyên tử, neutron.*" if is_teacher else "D. nguyên tử, neutron."
    ]
    render_options_no_table(doc, opts7, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Đồng vị là các nguyên tử của cùng một nguyên tố hóa học có cùng số proton nhưng có số neutron khác nhau.\n➔ Đáp án đúng: D.")

    # Câu 8
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 8. Cấu hình electron ở trạng thái cơ bản của nguyên tử nguyên tố X có tổng số electron trong các phân lớp p là 10. Nguyên tố X là")
    opts8 = [
        "A. P (Z=15).",
        "B. Ar (Z=18).",
        "*C. S (Z=16).*" if is_teacher else "C. S (Z=16).",
        "D. Cl (Z=17)."
    ]
    render_options_no_table(doc, opts8, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Tổng số electron ở các phân lớp p là 10 => Cấu hình các phân lớp p gồm: 2p^6 3p^4.\nCấu hình electron đầy đủ của X: 1s^2 2s^2 2p^6 3s^2 3p^4 => Z = 16 (Sulfur - S).\n➔ Đáp án đúng: C.")

    # Câu 9
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 9. Sự phân bố electron theo ô orbital nào dưới đây là đúng quy tắc?")
    add_centered_image(doc, os.path.join(IMG_DIR, "cau9_options.png"), width_cm=14.0)
    opts9 = [
        "A.",
        "*B.*" if is_teacher else "B.",
        "C.",
        "D."
    ]
    render_options_no_table(doc, opts9, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Theo quy tắc Hund, trong cùng một phân lớp, các electron có xu hướng phân bố đều vào các orbital sao cho số electron độc thân là tối đa và các electron này phải có chiều spin song song cùng chiều (mũi tên cùng hướng lên: [↑][↑][↑]).\n- Hình A vi phạm quy tắc Hund (ghép đôi sớm khi còn orbital trống).\n- Hình C, D vi phạm nguyên lí Pauli (2 electron cùng spin trong 1 ô orbital).\n➔ Đáp án đúng: B.")

    # Câu 10
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 10. Nguyên tố X có cấu hình electron nguyên tử: 1s^2 2s^2 2p^6 3s^1. Trong bảng tuần hoàn, X thuộc chu kì")
    opts10 = [
        "*A. 3.*" if is_teacher else "A. 3.",
        "B. 1.",
        "C. 2.",
        "D. 4."
    ]
    render_options_no_table(doc, opts10, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Số thứ tự chu kì = số lớp electron của nguyên tử. Cấu hình có 3 lớp electron (lớp 1, 2, 3) nên X thuộc chu kì 3.\n➔ Đáp án đúng: A.")

    # Câu 11
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 11. Trong bảng tuần hoàn các nguyên tố hoá học do Mendeleev đề xuất, các nguyên tố được sắp xếp theo chiều tăng dần của ..(1).. Trong bảng tuần hoàn các nguyên tố hoá học hiện đại, các nguyên tố được sắp xếp theo chiều tăng dần của ..(2).. Nội dung phù hợp trong các ô trống (1), (2) lần lượt là")
    opts11 = [
        "A. số electron hoá trị, khối lượng nguyên tử.",
        "B. số hiệu nguyên tử, khối lượng nguyên tử.",
        "*C. khối lượng nguyên tử, số hiệu nguyên tử.*" if is_teacher else "C. khối lượng nguyên tử, số hiệu nguyên tử.",
        "D. số electron hoá trị, số hiệu nguyên tử."
    ]
    render_options_no_table(doc, opts11, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Bảng tuần hoàn Mendeleev sắp xếp theo chiều tăng dần của khối lượng nguyên tử (nguyên tử khối). Bảng tuần hoàn hiện đại sắp xếp theo chiều tăng dần của điện tích hạt nhân (số hiệu nguyên tử Z).\n➔ Đáp án đúng: C.")

    # Câu 12
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 12. Trong tự nhiên, nguyên tố carbon (cacbon, C) có hai đồng vị bền là ^12C và ^13C chiếm tương ứng khoảng 98,89% và 1,11% về số nguyên tử. Nguyên tử khối trung bình của carbon là")
    opts12 = [
        "A. 12,022.",
        "*B. 12,011.*" if is_teacher else "B. 12,011.",
        "C. 12,055.",
        "D. 12,500."
    ]
    render_options_no_table(doc, opts12, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Nguyên tử khối trung bình của C:\nĀ = (12 × 98,89 + 13 × 1,11) / 100 = (1186,68 + 14,43) / 100 = 1201,11 / 100 = 12,0111 ≈ 12,011.\n➔ Đáp án đúng: B.")

    # Câu 13
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 13. Fluorine và hợp chất của nó được sử dụng làm chất chống sâu răng, chất cách điện, chất làm lạnh, vật liệu chống dính,... Nguyên tử fluorine chứa 9 hạt electron và 10 hạt neutron. Tổng số hạt proton, electron và neutron trong nguyên tử fluorine là")
    opts13 = [
        "A. 19.",
        "*B. 28.*" if is_teacher else "B. 28.",
        "C. 30.",
        "D. 32."
    ]
    render_options_no_table(doc, opts13, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Trong nguyên tử trung hòa về điện: p = e = 9. Số neutron n = 10.\nTổng số hạt = p + e + n = 9 + 9 + 10 = 28.\n➔ Đáp án đúng: B.")

    # Câu 14
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 14. Chromium được sử dụng nhiều trong luyện kim để chế tạo hợp kim chống ăn mòn và đánh bóng bề mặt. Nguyên tử chromium có cấu hình electron viết gọn là [Ar]3d^5 4s^1. Vị trí chromium trong bảng tuần hoàn là")
    opts14 = [
        "*A. ô số 24, chu kì 4, nhóm VIB.*" if is_teacher else "A. ô số 24, chu kì 4, nhóm VIB.",
        "B. ô số 24, chu kì 4, nhóm VIA.",
        "C. ô số 17, chu kì 4, nhóm IA.",
        "D. ô số 27, chu kì 4, nhóm IB."
    ]
    render_options_no_table(doc, opts14, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "- Ô nguyên tố: Z = 18 + 5 + 1 = 24 => Ô số 24.\n- Chu kì: Có 4 lớp electron => Chu kì 4.\n- Nhóm: Electron hóa trị là 3d^5 4s^1, tổng số electron hóa trị = 5 + 1 = 6, nguyên tố d nên thuộc phân nhóm phụ B => Nhóm VIB.\n➔ Đáp án đúng: A.")

    # Câu 15
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 15. Các nguyên tố trong cùng một chu kì có đặc điểm gì chung?")
    opts15 = [
        "*A. Số lớp electron bằng nhau.*" if is_teacher else "A. Số lớp electron bằng nhau.",
        "B. Số electron lớp ngoài cùng bằng nhau.",
        "C. Số electron hoá trị bằng nhau.",
        "D. Cùng số AO."
    ]
    render_options_no_table(doc, opts15, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Chu kì là dãy các nguyên tố mà nguyên tử của chúng có cùng số lớp electron, được sắp xếp theo chiều tăng dần của điện tích hạt nhân.\n➔ Đáp án đúng: A.")

    # Câu 16
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 16. Lớp M có số phân lớp electron là")
    opts16 = [
        "A. 2.",
        "B. 1.",
        "*C. 3.*" if is_teacher else "C. 3.",
        "D. 4."
    ]
    render_options_no_table(doc, opts16, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Lớp M ứng với số lượng tử chính n = 3, có 3 phân lớp electron: 3s, 3p, 3d.\n➔ Đáp án đúng: C.")

    # Câu 17
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 17. Nội dung nào dưới đây không phải đối tượng nghiên cứu của hóa học?")
    opts17 = [
        "A. Ứng dụng của chất.",
        "B. Tính chất và sự biến đổi của chất.",
        "C. Thành phần, cấu trúc của chất.",
        "*D. Sự lớn lên và sinh sản của tế bào.*" if is_teacher else "D. Sự lớn lên và sinh sản của tế bào."
    ]
    render_options_no_table(doc, opts17, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Đối tượng nghiên cứu của hóa học là chất và sự biến đổi của chất. Sự lớn lên và sinh sản của tế bào thuộc phạm vi nghiên cứu của Sinh học.\n➔ Đáp án đúng: D.")

    # Câu 18
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 18. Nguyên tố Y thuộc chu kì 4, nhóm IA của bảng tuần hoàn. Phát biểu nào sau đây về Y là đúng?")
    opts18 = [
        "A. Y có độ âm điện lớn nhất và bán kính nguyên tử lớn nhất trong chu kì 4.",
        "B. Y có độ âm điện lớn nhất và bán kính nguyên tử nhỏ nhất trong chu kì 4.",
        "*C. Y có độ âm điện nhỏ nhất và bán kính nguyên tử lớn nhất trong chu kì 4.*" if is_teacher else "C. Y có độ âm điện nhỏ nhất và bán kính nguyên tử lớn nhất trong chu kì 4.",
        "D. Y có độ âm điện nhỏ nhất và bán kính nguyên tử nhỏ nhất trong chu kì 4."
    ]
    render_options_no_table(doc, opts18, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Trong một chu kì theo chiều Z tăng từ trái sang phải: bán kính nguyên tử giảm dần, độ âm điện tăng dần.\nDo Y ở nhóm IA (đầu chu kì 4) nên Y có bán kính nguyên tử lớn nhất và độ âm điện nhỏ nhất trong chu kì 4.\n➔ Đáp án đúng: C.")

    # =========================================================================
    # PHẦN II
    # =========================================================================
    add_section_header(
        doc,
        "PHẦN II. Câu trắc nghiệm đúng sai.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai."
    )

    # Câu 1 P2
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 1. Khí carbon monoxide (CO) là một khí độc sinh ra khi đốt cháy các nhiên liệu hóa thạch, CO là một trong những nguyên nhân gây tử vong trong nhiều vụ cháy. Giả sử carbon có 2 đồng vị ^12_6C; ^13_6C và oxygen có 3 đồng vị là ^16_8O; ^17_8O; ^18_8O.")
    stmts1 = [
        ("a) Trong nguyên tử ^13_6C có 7 hạt neutron.", "Đúng"),
        ("b) Trong nguyên tử ^16_8O có tổng số hạt mang điện là 8 hạt.", "Sai"),
        ("c) Nguyên tử oxygen có 6 electron ở phân lớp ngoài cùng.", "Sai"),
        ("d) Từ các đồng vị trên có thể tạo thành tối đa 5 loại phân tử CO.", "Sai")
    ]
    for stmt, ans in stmts1:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.left_indent = Cm(0.5)
        p_st.paragraph_format.space_before = Pt(1)
        p_st.paragraph_format.space_after = Pt(1)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        sol_text = (
            "• Ý a ĐÚNG: N = A - Z = 13 - 6 = 7 hạt neutron.\n"
            "• Ý b SAI: Tổng số hạt mang điện trong ^16_8O là (p + e) = 8 + 8 = 16 hạt (8 là số hạt proton hoặc electron riêng rẽ).\n"
            "• Ý c SAI: Cấu hình e của oxygen (Z=8) là 1s^2 2s^2 2p^4. Lớp ngoài cùng (lớp 2) có 6 electron, nhưng phân lớp ngoài cùng là 2p chỉ có 4 electron.\n"
            "• Ý d SAI: Số phân tử CO tối đa tạo thành = (số đồng vị C) × (số đồng vị O) = 2 × 3 = 6 phân tử: "
            "(^12C^16O, ^12C^17O, ^12C^18O, ^13C^16O, ^13C^17O, ^13C^18O)."
        )
        add_solution_no_table(doc, sol_text)

    # Câu 2 P2
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 2. Nguyên tố X có cấu hình electron nguyên tử là 1s^2 2s^2 2p^6 3s^2 3p^5.")
    stmts2 = [
        ("a) Số hiệu nguyên tử của X bằng 17.", "Đúng"),
        ("b) Tổng số electron trong các phân lớp p của X là 11.", "Đúng"),
        ("c) Ở trạng thái cơ bản, nguyên tử X có 2 electron độc thân.", "Sai"),
        ("d) Nguyên tố X thuộc ô thứ 17, chu kì 3, nhóm VA.", "Sai")
    ]
    for stmt, ans in stmts2:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.left_indent = Cm(0.5)
        p_st.paragraph_format.space_before = Pt(1)
        p_st.paragraph_format.space_after = Pt(1)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        sol_text = (
            "• Ý a ĐÚNG: Z = 2 + 2 + 6 + 2 + 5 = 17.\n"
            "• Ý b ĐÚNG: Các phân lớp p là 2p^6 và 3p^5, tổng số electron = 6 + 5 = 11.\n"
            "• Ý c SAI: Phân lớp 3p^5 có 3 orbital: 2 orbital chứa cặp e ghép đôi ([↑↓][↑↓]), 1 orbital chứa 1 e độc thân ([↑]) => Chỉ có 1 electron độc thân.\n"
            "• Ý d SAI: Lớp ngoài cùng có 7 electron (3s^2 3p^5), nguyên tố p => X thuộc nhóm VIIA (chứ không phải VA)."
        )
        add_solution_no_table(doc, sol_text)

    # Câu 3 P2
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 3. Hình dưới mô tả orbital (a) và orbital (b) chứa electron trong nguyên tử sodium (Na) ở trạng thái cơ bản. Mức năng lượng của orbital (a) cao hơn orbital (b).")
    add_centered_image(doc, os.path.join(IMG_DIR, "cau3_orbitals.png"), width_cm=9.5)
    stmts3 = [
        ("a. Electron trong các orbital (a) và (b) thuộc cùng lớp electron.", "Sai"),
        ("b. Số electron trong 1 orbital (b) gấp ba số electron trong orbital (a).", "Sai"),
        ("c. Electron trên orbital (a) nằm gần hạt nhân hơn electron trên orbital (b).", "Sai"),
        ("d. Orbital (a) và (b) Khác nhau về định hướng trong không gian.", "Đúng")
    ]
    for stmt, ans in stmts3:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.left_indent = Cm(0.5)
        p_st.paragraph_format.space_before = Pt(1)
        p_st.paragraph_format.space_after = Pt(1)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        sol_text = (
            "Phân tích: Nguyên tử Na (Z=11) có cấu hình: 1s^2 2s^2 2p^6 3s^1.\n"
            "Orbital (a) có dạng hình cầu (orbital s), orbital (b) có dạng hình số 8 nổi (orbital p).\n"
            "Vì năng lượng của (a) cao hơn (b) nên (a) là orbital 3s, còn (b) là orbital 2p.\n"
            "• Ý a SAI: Orbital (a) là 3s thuộc lớp thứ 3 (lớp M), orbital (b) là 2p thuộc lớp thứ 2 (lớp L) => Khác lớp electron.\n"
            "• Ý b SAI: Trong nguyên tử Na: orbital 3s có 1e; mỗi 1 orbital trong phân lớp 2p có 2e (gấp đôi chứ không phải gấp ba).\n"
            "• Ý c SAI: Orbital 3s có năng lượng cao hơn và ở lớp thứ 3 nên nằm xa hạt nhân hơn orbital 2p ở lớp thứ 2.\n"
            "• Ý d ĐÚNG: Orbital (a) đối xứng cầu không định hướng trong không gian; orbital (b) hình số 8 nổi định hướng theo các trục x, y, z trong không gian."
        )
        add_solution_no_table(doc, sol_text)

    # Câu 4 P2
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 4. Cho các nguyên tố X, Y, T, R cùng một chu kỳ và thuộc nhóm A trong bảng tuần hoàn hóa học. Bán kính nguyên tử như hình vẽ sau:")
    add_centered_image(doc, os.path.join(IMG_DIR, "cau4_radii.png"), width_cm=10.0)
    stmts4 = [
        ("a) Nguyên tử có giá trị độ âm điện lớn nhất là T.", "Đúng"),
        ("b) Nguyên tố có tính kim loại mạnh nhất là Y.", "Đúng"),
        ("c) Nguyên tử của các nguyên tố này đều có cùng số lớp electron.", "Đúng"),
        ("d) Hydroxide của X có tính lưỡng tính thì Hydroxide của T có tính base.", "Sai")
    ]
    for stmt, ans in stmts4:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.left_indent = Cm(0.5)
        p_st.paragraph_format.space_before = Pt(1)
        p_st.paragraph_format.space_after = Pt(1)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        sol_text = (
            "Phân tích: Dựa vào hình vẽ, bán kính giảm dần: r_Y > r_R > r_X > r_T.\n"
            "Trong cùng 1 chu kì, theo chiều Z tăng từ trái sang phải: bán kính nguyên tử giảm dần => Thứ tự từ trái sang phải trong chu kì là: Y, R, X, T.\n"
            "• Ý a ĐÚNG: Đi từ trái sang phải trong chu kì độ âm điện tăng dần => T đứng sau cùng bên phải nên có độ âm điện lớn nhất.\n"
            "• Ý b ĐÚNG: Tính kim loại giảm dần từ trái sang phải => Y đứng đầu bên trái nên có tính kim loại mạnh nhất.\n"
            "• Ý c ĐÚNG: Các nguyên tố trong cùng một chu kì luôn có cùng số lớp electron.\n"
            "• Ý d SAI: Từ trái sang phải, tính base của hydroxide giảm dần, tính acid tăng dần. Vì T đứng sau X nên nếu hydroxide của X lưỡng tính thì hydroxide của T phải có tính acid, không thể có tính base."
        )
        add_solution_no_table(doc, sol_text)

    # =========================================================================
    # PHẦN III
    # =========================================================================
    add_section_header(
        doc,
        "PHẦN III. Trắc nghiệm trả lời ngắn và tự luận.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 6."
    )

    # Câu 1 P3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 1. Mỗi orbital nguyên tử chứa tối đa bao nhiêu electron?")
    if is_teacher:
        add_solution_no_table(doc, "Theo nguyên lí Pauli, trong mỗi orbital nguyên tử chỉ chứa tối đa 2 electron và 2 electron này có chiều tự quay (spin) ngược nhau.\n➔ Đáp số: 2.")
    else:
        add_student_writing_lines(doc, num_lines=2)

    # Câu 2 P3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 2. Cho các nguyên tố có số hiệu nguyên tử lần lượt là: 17, 13, 16, 19, 20. Trong các nguyên tố trên, có bao nhiêu nguyên tố là phi kim?")
    if is_teacher:
        sol_text = (
            "Xét cấu hình electron của từng nguyên tố:\n"
            "- Z = 17 (Cl): [Ne]3s^2 3p^5 => 7e ngoài cùng => Phi kim.\n"
            "- Z = 13 (Al): [Ne]3s^2 3p^1 => 3e ngoài cùng => Kim loại.\n"
            "- Z = 16 (S): [Ne]3s^2 3p^4 => 6e ngoài cùng => Phi kim.\n"
            "- Z = 19 (K): [Ar]4s^1 => 1e ngoài cùng => Kim loại.\n"
            "- Z = 20 (Ca): [Ar]4s^2 => 2e ngoài cùng => Kim loại.\n"
            "Vậy có 2 nguyên tố là phi kim (Z = 16, 17).\n"
            "➔ Đáp số: 2."
        )
        add_solution_no_table(doc, sol_text)
    else:
        add_student_writing_lines(doc, num_lines=3)

    # Câu 3 P3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 3. Cho hai nguyên tử: ^19_9F và ^24_12Mg.")
    p_a = doc.add_paragraph()
    render_question_prompt(p_a, "a. Xác định số hạt proton, electron, neutron trong mỗi nguyên tử.")
    p_b = doc.add_paragraph()
    render_question_prompt(p_b, "b. Viết cấu hình electron nguyên tử của hai nguyên tố trên.")
    if is_teacher:
        sol_text = (
            "a) Thành phần các loại hạt trong mỗi nguyên tử:\n"
            "• Nguyên tử ^19_9F:\n"
            "  - Số proton: p = Z = 9.\n"
            "  - Số electron: e = p = 9.\n"
            "  - Số neutron: n = A - Z = 19 - 9 = 10.\n"
            "• Nguyên tử ^24_12Mg:\n"
            "  - Số proton: p = Z = 12.\n"
            "  - Số electron: e = p = 12.\n"
            "  - Số neutron: n = A - Z = 24 - 12 = 12.\n\n"
            "b) Cấu hình electron nguyên tử:\n"
            "• Nguyên tử F (Z = 9): 1s^2 2s^2 2p^5 (hoặc [He]2s^2 2p^5).\n"
            "• Nguyên tử Mg (Z = 12): 1s^2 2s^2 2p^6 3s^2 (hoặc [Ne]3s^2)."
        )
        add_solution_no_table(doc, sol_text)
    else:
        add_student_writing_lines(doc, num_lines=6)

    # Câu 4 P3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 4. Nguyên tố X là nguyên tố quan trọng trong cơ thể sống. Hợp chất chứa X là thành phần quan trọng trong các phân tử sinh học DNA, RNA. Các tế bào sống cũng sử dụng hợp chất của X để vận chuyển năng lượng thông qua ATP. Nguyên tử của nguyên tố X có tổng số hạt cơ bản (proton, electron, neutron) là 46, trong đó tổng số hạt mang điện nhiều hơn tổng số hạt không mang điện là 14 hạt.")
    p_a = doc.add_paragraph()
    render_question_prompt(p_a, "a) Xác định thành phần cấu tạo (số proton, số electron, số neutron) của nguyên tử X.")
    p_b = doc.add_paragraph()
    render_question_prompt(p_b, "b) Viết cấu hình electron của X, xác định vị trí của X (ô, chu kì, nhóm) trong bảng tuần hoàn.")
    if is_teacher:
        sol_text = (
            "a) Gọi p, e, n lần lượt là số hạt proton, electron và neutron của nguyên tử X.\n"
            "Vì nguyên tử trung hòa về điện nên e = p.\n"
            "Theo bài ra, ta có hệ phương trình:\n"
            "  { 2p + n = 46  (1)\n"
            "  { 2p - n = 14  (2)\n"
            "Cộng vế với vế của (1) và (2) ta được: 4p = 60 => p = 15 => e = 15.\n"
            "Thay p = 15 vào (1): n = 46 - 2 × 15 = 16.\n"
            "Vậy thành phần cấu tạo của nguyên tử X là: 15 proton, 15 electron và 16 neutron. (X là Phosphorus - P).\n\n"
            "b) Cấu hình electron và vị trí của X trong bảng tuần hoàn:\n"
            "• Cấu hình electron của X (Z = 15): 1s^2 2s^2 2p^6 3s^2 3p^3 (viết gọn: [Ne]3s^2 3p^3).\n"
            "• Vị trí trong bảng tuần hoàn:\n"
            "  - Ô nguyên tố: Ô số 15 (do Z = 15).\n"
            "  - Chu kì: Chu kì 3 (do có 3 lớp electron).\n"
            "  - Nhóm: Nhóm VA (do có 5 electron hóa trị ở lớp ngoài cùng 3s^2 3p^3 và electron cuối cùng điền vào phân lớp p)."
        )
        add_solution_no_table(doc, sol_text)
    else:
        add_student_writing_lines(doc, num_lines=8)

    # Câu 5 P3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 5. Nguyên tử nguyên tố X ở trạng thái cơ bản có hai lớp electron, trong đó có hai electron độc thân. Vậy có bao nhiêu nguyên tố có thể thỏa mãn điều kiện trên? Biểu diễn cấu hình electron trên các ô lượng tử của nguyên tố đó.")
    if is_teacher:
        sol_text = (
            "Nguyên tử X có 2 lớp electron (lớp 1: 1s^2; lớp 2 gồm phân lớp 2s và 2p).\n"
            "Phân lớp 2s chỉ có 1 orbital, chứa tối đa 2 electron ghép đôi [↑↓].\n"
            "Phân lớp 2p có 3 ô orbital. Để có 2 electron độc thân ở trạng thái cơ bản, theo quy tắc Hund có 2 khả năng phân bố:\n"
            "• Khả năng 1: Phân lớp 2p có 2 electron (2p^2).\n"
            "  2 electron phân bố vào 2 ô orbital với chiều spin song song: [↑][↑][ ].\n"
            "  Cấu hình electron: 1s^2 2s^2 2p^2 (Z = 6: Carbon - C).\n"
            "  Biểu diễn theo ô lượng tử (AO):\n"
            "    1s: [↑↓]   2s: [↑↓]   2p: [↑][↑][ ]\n\n"
            "• Khả năng 2: Phân lớp 2p có 4 electron (2p^4).\n"
            "  1 ô ghép đôi và 2 ô còn lại mỗi ô có 1 e độc thân: [↑↓][↑][↑].\n"
            "  Cấu hình electron: 1s^2 2s^2 2p^4 (Z = 8: Oxygen - O).\n"
            "  Biểu diễn theo ô lượng tử (AO):\n"
            "    1s: [↑↓]   2s: [↑↓]   2p: [↑↓][↑][↑]\n\n"
            "➔ Kết luận: Có 2 nguyên tố thỏa mãn điều kiện là Carbon (C) và Oxygen (O)."
        )
        add_solution_no_table(doc, sol_text)
    else:
        add_student_writing_lines(doc, num_lines=7)

    # Câu 6 P3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 6. Trong tự nhiên, nguyên tố Copper (đồng) có hai đồng vị là ^63Cu và ^65Cu. Nguyên tử khối trung bình của đồng là 63,54.")
    p_a = doc.add_paragraph()
    render_question_prompt(p_a, "a. Tính % số nguyên tử Cu trong mỗi đồng vị trên.")
    p_b = doc.add_paragraph()
    render_question_prompt(p_b, "b. Xác định thành phần phần trăm theo khối lượng của mỗi đồng vị Cu trong 48,77 gam Cu(OH)2, (biết H = 1, O = 16).")
    if is_teacher:
        sol_text = (
            "a) Gọi % số nguyên tử của đồng vị ^63Cu là x (%), thì % của ^65Cu là (100 - x) (%).\n"
            "Ta có: Ā_Cu = [63x + 65(100 - x)] / 100 = 63,54\n"
            "<=> 63x + 6500 - 65x = 6354 <=> -2x = -146 <=> x = 73%.\n"
            "Vậy:\n"
            "• % số nguyên tử ^63Cu = 73%.\n"
            "• % số nguyên tử ^65Cu = 100% - 73% = 27%.\n\n"
            "b) Tính phần trăm theo khối lượng của mỗi đồng vị trong 48,77 gam Cu(OH)2:\n"
            "Khối lượng mol phân tử của Cu(OH)2:\n"
            "M_Cu(OH)2 = 63,54 + 2 × (16 + 1) = 97,54 (g/mol).\n"
            "Số mol Cu(OH)2 có trong 48,77 gam:\n"
            "n_Cu(OH)2 = 48,77 / 97,54 = 0,5 mol.\n"
            "Trong 0,5 mol Cu(OH)2 có 0,5 mol nguyên tử Cu:\n"
            "• Số mol ^63Cu = 0,5 × 73% = 0,365 mol.\n"
            "  Khối lượng ^63Cu: m_^63Cu = 0,365 × 63 = 22,995 gam.\n"
            "  % khối lượng ^63Cu trong Cu(OH)2 = (22,995 / 48,77) × 100% ≈ 47,15%.\n"
            "• Số mol ^65Cu = 0,5 × 27% = 0,135 mol.\n"
            "  Khối lượng ^65Cu: m_^65Cu = 0,135 × 65 = 8,775 gam.\n"
            "  % khối lượng ^65Cu trong Cu(OH)2 = (8,775 / 48,77) × 100% ≈ 17,99%."
        )
        add_solution_no_table(doc, sol_text)
    else:
        add_student_writing_lines(doc, num_lines=8)

    # Lưu file
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Đã tạo thành công: {output_path}")

if __name__ == "__main__":
    out_hs = r"C:\Antigravity_Thanh\San_Pham\De_Kiem_Tra_GK1_Hoa_10_BacNinh_2025_HocSinh.docx"
    out_gv = r"C:\Antigravity_Thanh\San_Pham\De_Kiem_Tra_GK1_Hoa_10_BacNinh_2025_GiaoVien.docx"
    build_exam_bacninh_10(is_teacher=False, output_path=out_hs)
    build_exam_bacninh_10(is_teacher=True, output_path=out_gv)
