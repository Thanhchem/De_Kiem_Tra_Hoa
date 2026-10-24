# -*- coding: utf-8 -*-
import os
import re
import docx
from docx.oxml.ns import qn
import sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

from build_ngan_hang_hoa12_c123 import CHAPTERS, iter_block_items

IMG_OUT = r"c:\Antigravity_Thanh\San_Pham\images_gk1_hoa12"
os.makedirs(IMG_OUT, exist_ok=True)


def extract_chapter_assets(cnum):
    """Trích xuất hình ảnh và bảng dữ liệu trong thân câu hỏi của từng câu trong CD1, CD2, CD3."""
    doc = docx.Document(CHAPTERS[cnum]["da_file"])
    assets = {}  # q_num -> {"images": [paths], "tables": [[[cell_text,...],...]]}
    curr_q = None
    in_sol = False

    def ensure_q(q):
        if q not in assets:
            assets[q] = {"images": [], "tables": []}

    def extract_images_from_p(p, q):
        # Find all blip r:embed attributes
        blips = p._element.xpath('.//*[local-name()="blip"]')
        for blip in blips:
            rId = blip.get(qn('r:embed'))
            if rId and rId in doc.part.related_parts:
                part = doc.part.related_parts[rId]
                idx = len(assets[q]["images"]) + 1
                fname = f"C{cnum}_Q{q:03d}_{idx}.png"
                fpath = os.path.join(IMG_OUT, fname)
                try:
                    import io
                    from PIL import Image
                    im = Image.open(io.BytesIO(part.blob))
                    im.save(fpath, format="PNG")
                    assets[q]["images"].append(fpath)
                except Exception:
                    pass

    for btype, block in iter_block_items(doc):
        if btype == 'p':
            txt = block.text.strip()
            m = re.match(r'^Câu\s+(\d+)[\.\:]', txt)
            if m and (not curr_q or not in_sol or int(m.group(1)) == curr_q + 1):
                curr_q = int(m.group(1))
                in_sol = False
                ensure_q(curr_q)
            elif txt.startswith("Hướng dẫn giải"):
                in_sol = True
            if curr_q and not in_sol:
                ensure_q(curr_q)
                extract_images_from_p(block, curr_q)
        elif btype == 'tbl':
            # Check if table starts a question (2-col layout table)
            started_new_q = False
            for row in block.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        m = re.match(r'^Câu\s+(\d+)[\.\:]', p.text.strip())
                        if m and (not curr_q or int(m.group(1)) == curr_q + 1):
                            curr_q = int(m.group(1))
                            in_sol = False
                            ensure_q(curr_q)
                            started_new_q = True
            if curr_q and not in_sol:
                ensure_q(curr_q)
                # Extract any images inside table cells
                has_img = False
                for row in block.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            before_cnt = len(assets[curr_q]["images"])
                            extract_images_from_p(p, curr_q)
                            if len(assets[curr_q]["images"]) > before_cnt:
                                has_img = True
                if not started_new_q and not has_img:
                    # It's a data table inside the question stem!
                    tbl_data = []
                    for row in block.rows:
                        row_cells = [ " ".join(p.text.strip() for p in c.paragraphs if p.text.strip()) for c in row.cells ]
                        tbl_data.append(row_cells)
                    if tbl_data:
                        assets[curr_q]["tables"].append(tbl_data)
    return assets


if __name__ == "__main__":
    for c in [1, 2, 3]:
        a = extract_chapter_assets(c)
        with_img = [q for q, v in a.items() if v["images"]]
        with_tbl = [q for q, v in a.items() if v["tables"]]
        print(f"CD{c}: {len(with_img)} questions with images {with_img}, {len(with_tbl)} questions with data tables {with_tbl}")
