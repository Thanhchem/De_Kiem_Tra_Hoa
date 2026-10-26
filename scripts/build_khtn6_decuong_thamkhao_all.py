# -*- coding: utf-8 -*-
"""
Script biên soạn trọn bộ 3 tài liệu chuẩn FORM ĐỀ KIỂM TRA cho:
ĐỀ CƯƠNG ÔN TẬP GIỮA HỌC KÌ I - MÔN KHTN 6 (NĂM HỌC 2026 - 2027) - BẢN THAM KHẢO
1. De_Cuong_GK1_KHTN6_ThamKhao_HocSinh.docx
2. De_Cuong_GK1_KHTN6_ThamKhao_GiaoVien.docx
3. De_Cuong_GK1_KHTN6_ThamKhao_DayTrenManHinh.pdf (Slide 16:9 Widescreen cho Tivi 65 inch)
Và bản Word: De_Cuong_GK1_KHTN6_ThamKhao_DayTrenManHinh.docx
Thầy TRẦN VĂN THẠNH - 0777.470.803
"""

import os
import sys
import subprocess
import shutil

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# Import các hàm chuẩn từ form_de_kiem_tra
sys.path.insert(0, r"c:\Antigravity_Thanh\scripts")
from form_de_kiem_tra import (
    setup_header_footer,
    build_title_block,
    add_section_header,
    render_question_prompt,
    render_options_no_table,
    render_statement,
    add_solution_no_table,
    add_student_writing_lines,
    render_rich_text,
    set_run_style,
    COLOR_PRIMARY,
    COLOR_CORRECT,
    COLOR_SOLUTION,
    COLOR_HF_RED,
    COLOR_HF_GREEN_HEX
)

OUTPUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
SCRATCH_DIR = r"c:\Antigravity_Thanh\scratch"
os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(SCRATCH_DIR, exist_ok=True)

ONEDRIVE_DIRS = [
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn6\KHTN6-DEKIEMTRADINHKI\New folder",
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn6\KHTN6-DEKIEMTRADINHKI",
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn6"
]

COLOR_PART = RGBColor(198, 40, 40)   # Đỏ đậm cho Phân môn
COLOR_SUBSECTION = RGBColor(0, 77, 64) # Xanh lục trầm cho bài tập

def add_part_heading(doc, title):
    """Tiêu đề phân môn lớn."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    set_run_style(r, font_size=12.5, bold=True, color=COLOR_PART)
    from form_de_kiem_tra import add_p_border_bottom
    add_p_border_bottom(p, color="C62828", sz="8")

# ==============================================================================
# DỮ LIỆU ĐỀ CƯƠNG THAM KHẢO KHTN 6
# ==============================================================================

VATLI_THEORY = [
    {
        "num": "1",
        "title": "Khoa học tự nhiên là gì? Các lĩnh vực chủ yếu của khoa học tự nhiên.",
        "content_student": [
            "• *Khoa học tự nhiên (KHTN):* .......................................................................................................................",
            "• *Các lĩnh vực chủ yếu của KHTN:* ............................................................................................................"
        ],
        "solution": (
            "• *Khái niệm:* Khoa học tự nhiên là ngành khoa học nghiên cứu về các sự vật, hiện tượng, quy luật tự nhiên, "
            "những ảnh hưởng của chúng đến con người và môi trường sống.\n"
            "• *Các lĩnh vực chủ yếu của KHTN:*\n"
            "  + *Vật lí học:* Nghiên cứu về vật chất, quy luật vận động, lực, năng lượng và sự biến đổi năng lượng.\n"
            "  + *Hóa học:* Nghiên cứu về chất và sự biến đổi của chất.\n"
            "  + *Sinh học:* Nghiên cứu về các vật sống, mối quan hệ giữa các sinh vật với nhau và với môi trường.\n"
            "  + *Khoa học Trái Đất:* Nghiên cứu về Trái Đất và bầu khí quyển của nó.\n"
            "  + *Thiên văn học:* Nghiên cứu về quy luật vận động và biến đổi của các vật thể trên bầu trời (vũ trụ)."
        )
    },
    {
        "num": "2",
        "title": "Vật sống là gì? Vật không sống là gì? Nêu ví dụ.",
        "content_student": [
            "• *Vật sống (vật hữu sinh):* .......................................................................................................................",
            "• *Vật không sống (vật vô sinh):* ..............................................................................................................."
        ],
        "solution": (
            "• *Vật sống (vật hữu sinh):* Là vật có các biểu hiện sống như: trao đổi chất và chuyển hóa năng lượng với môi trường, "
            "sinh trưởng (lớn lên), vận động, cảm ứng và sinh sản. Khi đến một giai đoạn nhất định sẽ già và chết.\n"
            "  *Ví dụ:* Con người, con sư tử, cây hoa mai, cây lúa, vi khuẩn, nấm rơm, chim bồ câu...\n"
            "• *Vật không sống (vật vô sinh):* Là vật không có các biểu hiện sống đặc trưng (không trao đổi chất, không lớn lên, không sinh sản).\n"
            "  *Ví dụ:* Hòn đá, hạt cát, cái bàn, bút chì, cây cầu, giọt nước, chiếc xe đạp..."
        )
    },
    {
        "num": "3",
        "title": "Nêu các quy định về an toàn trong phòng thực hành, các kí hiệu cảnh báo trong phòng thực hành.",
        "content_student": [
            "• *Các quy định về an toàn:* .......................................................................................................................",
            "• *Các kí hiệu cảnh báo thông dụng:* ..........................................................................................................."
        ],
        "solution": (
            "• *Các quy định về an toàn trong phòng thực hành:*\n"
            "  1. Mặc trang phục bảo hộ thích hợp (áo blouse, đeo kính bảo hộ, găng tay khi cần thiết); cột tóc gọn gàng.\n"
            "  2. Chỉ được làm thí nghiệm khi có sự hướng dẫn và cho phép của giáo viên.\n"
            "  3. Tuyệt đối không ăn, uống, đùa nghịch, chạy nhảy trong phòng thực hành.\n"
            "  4. Không được nếm hoặc ngửi trực tiếp hóa chất; đọc kĩ nhãn mác hóa chất trước khi dùng.\n"
            "  5. Nắm vững vị trí và cách dùng bình chữa cháy, hộp sơ cứu, vòi rửa mắt khẩn cấp.\n"
            "  6. Khi kết thúc: thu dọn dụng cụ, bỏ chất thải đúng nơi quy định, lau bàn và rửa sạch tay bằng xà phòng.\n"
            "• *Các kí hiệu cảnh báo phổ biến:*\n"
            "  + Kí hiệu chất độc (hình đầu lâu xương chéo).\n"
            "  + Kí hiệu chất ăn mòn (hình axit đổ làm thủng bề mặt kim loại hoặc da tay).\n"
            "  + Kí hiệu chất dễ cháy (hình ngọn lửa).\n"
            "  + Kí hiệu nguy hiểm về điện (hình tia sét).\n"
            "  + Kí hiệu bắt buộc mang bảo hộ (hình tròn nền xanh viền trắng)."
        )
    },
    {
        "num": "4",
        "title": "Nêu đơn vị, dụng cụ đo chiều dài trong hệ thống đo lường chính thức của Việt Nam, các bước đo.",
        "content_student": [
            "• *Đơn vị đo & Dụng cụ đo:* .......................................................................................................................",
            "• *Quy trình các bước đo:* .........................................................................................................................."
        ],
        "solution": (
            "• *Đơn vị đo:* Đơn vị đo độ dài hợp pháp của nước Việt Nam theo hệ SI là *mét* (kí hiệu: *m*).\n"
            "  Các đơn vị khác: km, dm, cm, mm (1 km = 1000 m; 1 m = 10 dm = 100 cm = 1000 mm).\n"
            "• *Dụng cụ đo:* Thước thẳng, thước cuộn, thước dây, thước kẹp...\n"
            "• *Quy trình 5 bước đo chiều dài:*\n"
            "  - *Bước 1 (Ước lượng):* Ước lượng chiều dài cần đo để chọn thước có GHĐ và ĐCNN phù hợp.\n"
            "  - *Bước 2 (Chọn thước):* Chọn thước có GHĐ lớn hơn hoặc bằng chiều dài cần đo và ĐCNN thích hợp.\n"
            "  - *Bước 3 (Đặt thước):* Đặt thước dọc theo chiều dài cần đo, sao cho một đầu của vật ngang bằng vạch số 0.\n"
            "  - *Bước 4 (Đặt mắt):* Đặt mắt nhìn vuông góc với cạnh thước ở đầu kia của vật.\n"
            "  - *Bước 5 (Đọc và ghi kết quả):* Đọc giá trị theo vạch chia gần nhất với đầu kia của vật và ghi theo ĐCNN của thước."
        )
    },
    {
        "num": "5",
        "title": "Nêu đơn vị, dụng cụ đo khối lượng trong hệ thống đo lường chính thức của Việt Nam, các bước đo.",
        "content_student": [
            "• *Đơn vị đo & Dụng cụ đo:* .......................................................................................................................",
            "• *Quy trình các bước đo:* .........................................................................................................................."
        ],
        "solution": (
            "• *Đơn vị đo:* Đơn vị đo khối lượng hợp pháp của nước ta là *kilôgam* (kí hiệu: *kg*).\n"
            "  Các đơn vị khác: tấn, tạ, yến, hg (lạng), dag, g, mg (1 tấn = 1000 kg; 1 tạ = 100 kg; 1 lạng = 100 g; 1 kg = 1000 g).\n"
            "• *Dụng cụ đo:* Cân đồng hồ, cân điện tử, cân y tế, cân đòn, cân Rô-béc-van (Roberval)...\n"
            "• *Quy trình 5 bước đo khối lượng (cân đồng hồ / cân điện tử):*\n"
            "  - *Bước 1:* Ước lượng khối lượng của vật cần cân.\n"
            "  - *Bước 2:* Chọn cân có GHĐ và ĐCNN thích hợp.\n"
            "  - *Bước 3:* Hiệu chỉnh cân về đúng vạch số 0 (hoặc nhấn nút Tare/Zero).\n"
            "  - *Bước 4:* Đặt vật cần cân lên chính giữa đĩa cân.\n"
            "  - *Bước 5:* Đọc và ghi kết quả đo theo vạch chia gần nhất với đầu kim cân."
        )
    },
    {
        "num": "6",
        "title": "Nêu đơn vị, dụng cụ đo thời gian trong hệ thống đo lường chính thức của Việt Nam, các bước đo.",
        "content_student": [
            "• *Đơn vị đo & Dụng cụ đo:* .......................................................................................................................",
            "• *Quy trình các bước đo:* .........................................................................................................................."
        ],
        "solution": (
            "• *Đơn vị đo:* Đơn vị đo thời gian hợp pháp trong hệ SI là *giây* (kí hiệu: *s*).\n"
            "  Các đơn vị khác: phút (min), giờ (h), ngày, tháng, năm, mili giây (ms). Quy đổi: 1 ngày = 24 h; 1 h = 60 min = 3600 s; 1 s = 1000 ms.\n"
            "• *Dụng cụ đo:* Đồng hồ đeo tay, đồng hồ treo tường, đồng hồ để bàn, đồng hồ bấm giây, đồng hồ cát...\n"
            "• *Quy trình 5 bước đo thời gian (dùng đồng hồ bấm giây):*\n"
            "  - *Bước 1:* Ước lượng khoảng thời gian cần đo.\n"
            "  - *Bước 2:* Chọn đồng hồ bấm giây phù hợp.\n"
            "  - *Bước 3:* Hiệu chỉnh đồng hồ về số 0 trước khi bắt đầu đo.\n"
            "  - *Bước 4:* Nhấn nút bắt đầu đo (Start) đồng thời lúc sự kiện bắt đầu; nhấn nút dừng (Stop) khi kết thúc.\n"
            "  - *Bước 5:* Đọc và ghi kết quả đo theo chỉ số hiển thị."
        )
    }
]

VATLI_EXERCISES = [
    {
        "num": "1",
        "title": "Xác định ĐCNN của các kết quả đo sau:",
        "sub_items": [
            ("a) 2,5 cm", "ĐCNN = 0,1 cm (hoặc 1 mm) hoặc 0,5 cm. (Vì kết quả đo được ghi đến hàng phần mười của cm)."),
            ("b) 3 dm", "ĐCNN = 1 dm (hoặc 0,5 dm, 0,1 dm rơi đúng vạch 3 dm). ĐCNN lớn nhất là 1 dm."),
            ("c) 12,2 mm", "ĐCNN = 0,1 mm hoặc 0,2 mm. (Kết quả ghi đến hàng phần mười của mm, thường dùng thước kẹp)."),
            ("d) 15,1 cm", "ĐCNN = 0,1 cm (hoặc 1 mm). (Kết quả ghi chính xác đến 1 chữ số thập phân của cm).")
        ],
        "solution": (
            "• *Nguyên tắc:* Kết quả đo luôn được ghi theo Độ chia nhỏ nhất (ĐCNN) của dụng cụ đo.\n"
            "  a) *2,5 cm:* ĐCNN là *0,1 cm* (tức 1 mm) hoặc *0,5 cm*.\n"
            "  b) *3 dm:* ĐCNN là *1 dm* (hoặc 0,1 dm; 0,5 dm mà giá trị đo trùng vạch 3 dm).\n"
            "  c) *12,2 mm:* ĐCNN là *0,1 mm* hoặc *0,2 mm*.\n"
            "  d) *15,1 cm:* ĐCNN là *0,1 cm* (tức 1 mm)."
        )
    },
    {
        "num": "2",
        "title": "Để đo diện tích của một vườn cỏ có kích thước 25 × 30 (m). Nếu trong tay em có hai chiếc thước: một thước gấp có giới hạn đo (GHĐ) 2 m và một thước cuộn có GHĐ 20 m. Em sẽ dùng thước nào để cho kết quả đo chính xác hơn? Vì sao?",
        "solution": (
            "• *Trả lời:* Em sẽ chọn dùng *thước cuộn có GHĐ 20 m*.\n"
            "• *Giải thích:*\n"
            "  - Vườn cỏ có kích thước rất lớn (chiều dài 30 m, chiều rộng 25 m).\n"
            "  - Nếu dùng thước gấp có GHĐ 2 m, ta phải đặt thước đo liên tiếp rất nhiều lần (cạnh 25 m cần đo khoảng 13 lần, cạnh 30 m cần đo 15 lần). "
            "Mỗi lần dịch chuyển và đánh dấu thước đều tạo ra sai số, việc đo nhiều lần làm tích lũy sai số rất lớn và tốn nhiều công sức.\n"
            "  - Khi dùng thước cuộn có GHĐ 20 m, mỗi cạnh chỉ cần đo 2 lần (cạnh 25 m đo 1 lần 20 m và 1 lần 5 m; cạnh 30 m đo 1 lần 20 m và 1 lần 10 m), "
            "giúp giảm thiểu tối đa số lần ghép thước, từ đó cho kết quả đo nhanh chóng và có độ chính xác cao hơn rất nhiều."
        )
    },
    {
        "num": "3",
        "title": "Có 20 túi đường, ban đầu mỗi túi có khối lượng 1 kg, sau đó người ta cho thêm mỗi túi 2 lạng đường nữa. Khối lượng của 20 túi đường khi đó là bao nhiêu kg?",
        "solution": (
            "• *Tóm tắt & Đổi đơn vị:* 1 lạng = 100 g = 0,1 kg  ➔  2 lạng = 200 g = 0,2 kg.\n"
            "• *Lời giải chi tiết:*\n"
            "  - *Cách 1:*\n"
            "    Khối lượng của mỗi túi đường sau khi thêm là:\n"
            "      1 kg + 0,2 kg = 1,2 kg.\n"
            "    Tổng khối lượng của 20 túi đường là:\n"
            "      m = 20 × 1,2 kg = 24 kg.\n"
            "  - *Cách 2:*\n"
            "    Khối lượng 20 túi đường ban đầu: 20 × 1 kg = 20 kg.\n"
            "    Lượng đường được thêm vào 20 túi: 20 × 2 lạng = 40 lạng = 4 kg.\n"
            "    Tổng khối lượng khi đó: 20 kg + 4 kg = 24 kg.\n"
            "• *Đáp số:* 24 kg."
        )
    },
    {
        "num": "4",
        "title": "Để thực hiện đo thời gian từ cổng trường đến lớp học em dùng loại đồng hồ nào? Giải thích sự lựa chọn đó.",
        "solution": (
            "• *Trả lời:* Em sẽ chọn dùng *đồng hồ bấm giây* (hoặc đồng hồ đeo tay có tính năng bấm giờ thể thao / điện thoại có ứng dụng bấm giờ).\n"
            "• *Giải thích:*\n"
            "  - Thời gian đi từ cổng trường vào đến lớp học là một khoảng thời gian ngắn (thường từ vài chục giây đến vài phút).\n"
            "  - Đồng hồ bấm giây có nút khởi động (Start) và tạm dừng (Stop) linh hoạt, giúp bấm bắt đầu ngay khi bước qua cổng trường và dừng ngay khi bước vào lớp.\n"
            "  - Đồng hồ bấm giây có ĐCNN nhỏ (thường đến 0,01 s hoặc 1 s), cho độ chính xác cao; đồng thời nhỏ gọn, dễ mang theo khi di chuyển.\n"
            "  - Các loại đồng hồ treo tường hoặc để bàn cồng kềnh, không mang theo được và không có tính năng bấm khoảng thời gian chính xác."
        )
    },
    {
        "num": "5",
        "title": "Một ô tô khởi hành ở Hà Nội lúc 8 giờ 15 phút. Sau khi chạy được 1 giờ thì đến Hải Dương. Ô tô nghỉ ở Hải Dương 20 phút, sau đó chạy 50 phút nữa thì đến Hải Phòng. Xác định thời điểm mà ô tô đến Hải Dương và Hải Phòng.",
        "solution": (
            "• *Lời giải chi tiết:*\n"
            "  1. Thời điểm ô tô đến Hải Dương là:\n"
            "     8 giờ 15 phút + 1 giờ = 9 giờ 15 phút.\n"
            "  2. Thời điểm ô tô nghỉ xong và rời Hải Dương là:\n"
            "     9 giờ 15 phút + 20 phút = 9 giờ 35 phút.\n"
            "  3. Thời điểm ô tô đến Hải Phòng là:\n"
            "     9 giờ 35 phút + 50 phút = 9 giờ 85 phút.\n"
            "     Vì 85 phút = 1 giờ 25 phút nên: 9 giờ 85 phút = 10 giờ 25 phút.\n"
            "• *Đáp số:*\n"
            "  - Đến Hải Dương lúc: *9 giờ 15 phút*.\n"
            "  - Đến Hải Phòng lúc: *10 giờ 25 phút*."
        )
    }
]

SINH_DATA = [
    {
        "lesson": "Bài 18: Tế bào – Đơn vị cơ bản của sự sống",
        "questions": [
            {
                "num": "1",
                "q": "Tế bào là gì?",
                "ans": (
                    "Tế bào là đơn vị cơ sở cấu tạo nên mọi cơ thể sống và thực hiện các chức năng sống cơ bản. "
                    "Mọi cơ thể sinh vật (từ đơn bào như vi khuẩn, trùng roi đến đa bào như thực vật, động vật, con người) "
                    "đều được cấu tạo từ một hoặc nhiều tế bào. Mọi hoạt động sống (trao đổi chất, sinh trưởng, cảm ứng, sinh sản) "
                    "đều diễn ra trong tế bào."
                )
            },
            {
                "num": "2",
                "q": "Nêu đặc điểm về kích thước, hình dạng của tế bào.",
                "ans": (
                    "• *Về kích thước:*\n"
                    "  - Phần lớn các tế bào có kích thước hiển vi rất nhỏ bé (khoảng từ vài micromet μm đến vài chục micromet), "
                    "mắt thường không nhìn thấy được mà phải quan sát dưới kính hiển vi quang học (ví dụ: tế bào vi khuẩn, tế bào hồng cầu, tế bào biểu bì vảy hành...).\n"
                    "  - Một số ít tế bào có kích thước lớn có thể quan sát bằng mắt thường hoặc kính lúp (ví dụ: tép bưởi, tế bào trứng cá, tế bào lòng đỏ trứng gà/chim...).\n"
                    "• *Về hình dạng:*\n"
                    "  - Tế bào có hình dạng rất đa dạng: hình cầu (tế bào trứng), hình đĩa (hồng cầu người), hình thoi (tế bào cơ trơn), "
                    "hình sợi dài (tế bào cơ vân), hình sao có nhiều sợi nhánh (tế bào thần kinh), hình đa giác (tế bào biểu bì vảy hành), hình que (trực khuẩn)...\n"
                    "  - Hình dạng của mỗi loại tế bào luôn thích nghi và phù hợp tối ưu với chức năng mà tế bào đó đảm nhận trong cơ thể."
                )
            }
        ]
    },
    {
        "lesson": "Bài 19: Cấu tạo và chức năng các thành phần của tế bào",
        "questions": [
            {
                "num": "1",
                "q": "Nêu cấu tạo và chức năng của các thành phần chính cấu tạo nên tế bào.",
                "ans": (
                    "Mọi tế bào đều gồm có 3 thành phần chính sau:\n"
                    "1. *Màng tế bào (màng sinh chất):*\n"
                    "   - *Cấu tạo:* Lớp màng mỏng bao bọc bên ngoài khối chất tế bào. (Ở tế bào thực vật còn có thêm thành tế bào dày bằng xenlulozơ bao ngoài màng tế bào giúp cố định hình dạng).\n"
                    "   - *Chức năng:* Bao bọc, bảo vệ tế bào và kiểm soát sự trao đổi chất có chọn lọc giữa tế bào với môi trường ngoài.\n"
                    "2. *Chất tế bào (tế bào chất):*\n"
                    "   - *Cấu tạo:* Vùng nằm giữa màng tế bào và nhân (hoặc vùng nhân), là chất lỏng dạng keo chứa nhiều bào quan (ti thể, lục lạp ở thực vật, lưới nội chất, không bào...).\n"
                    "   - *Chức năng:* Là nơi diễn ra hầu hết các hoạt động trao đổi chất và chuyển hóa năng lượng của tế bào.\n"
                    "3. *Nhân tế bào (hoặc Vùng nhân ở tế bào nhân sơ):*\n"
                    "   - *Cấu tạo:* Chứa vật chất di truyền (DNA). Tế bào nhân thực có màng nhân bao bọc hoàn chỉnh; tế bào nhân sơ chưa có màng nhân bao bọc.\n"
                    "   - *Chức năng:* Là trung tâm điều khiển mọi hoạt động sống của tế bào; lưu giữ và truyền đạt thông tin di truyền."
                )
            }
        ]
    }
]

# ==============================================================================
# 1. HÀM SINH FILE WORD (HỌC SINH / GIÁO VIÊN / DẠY MÀN HÌNH)
# ==============================================================================

def generate_word_document(mode="student"):
    """
    mode: 'student', 'teacher', 'screen'
    """
    is_teacher = (mode == "teacher")
    is_screen = (mode == "screen")

    if mode == "teacher":
        filename = "De_Cuong_GK1_KHTN6_ThamKhao_GiaoVien.docx"
    elif mode == "screen":
        filename = "De_Cuong_GK1_KHTN6_ThamKhao_DayTrenManHinh.docx"
    else:
        filename = "De_Cuong_GK1_KHTN6_ThamKhao_HocSinh.docx"

    filepath = os.path.join(OUTPUT_DIR, filename)

    doc = docx.Document()
    setup_header_footer(doc)

    build_title_block(
        doc,
        school_line="TRƯỜNG TRUNG HỌC CƠ SỞ",
        exam_title="ĐỀ CƯƠNG ÔN TẬP GIỮA HỌC KÌ I - MÔN KHTN 6",
        subject_line="NĂM HỌC 2026 - 2027 (HỌC SINH DÙNG THAM KHẢO)",
        num_pages="03",
        duration="45 phút",
        exam_code="KHTN6-TK",
        atomic_masses=None,
        is_teacher=is_teacher
    )

    base_size = 11.5 if is_screen else 10.5

    # --------------------------------------------------------------------------
    # PHÂN MÔN VẬT LÍ
    # --------------------------------------------------------------------------
    add_part_heading(doc, "PHẦN I. PHÂN MÔN VẬT LÍ")

    # A. Lý thuyết
    add_section_header(doc, "A. LÝ THUYẾT TRỌNG TÂM")
    for item in VATLI_THEORY:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, f"Câu {item['num']}. {item['title']}", base_size=base_size)

        if is_teacher:
            add_solution_no_table(doc, item["solution"])
        elif is_screen:
            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.left_indent = Cm(0.6)
            p_sp.paragraph_format.space_before = Pt(2)
            p_sp.paragraph_format.space_after = Pt(2)
            render_rich_text(p_sp, "✎ Ghi chú giảng bài: ........................................................................................................", base_size=base_size-1)
        else:
            for l in item["content_student"]:
                p_l = doc.add_paragraph()
                p_l.paragraph_format.left_indent = Cm(0.5)
                p_l.paragraph_format.space_before = Pt(1)
                p_l.paragraph_format.space_after = Pt(1)
                render_rich_text(p_l, l, base_size=base_size)
            add_student_writing_lines(doc, num_lines=4)

    # B. Bài tập
    add_section_header(doc, "B. BÀI TẬP TỰ LUẬN & ỨNG DỤNG")
    for ex in VATLI_EXERCISES:
        p_ex = doc.add_paragraph()
        p_ex.paragraph_format.space_before = Pt(4)
        p_ex.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_ex, f"Bài {ex['num']}. {ex['title']}", base_size=base_size)

        if "sub_items" in ex:
            for sub_q, sub_ans in ex["sub_items"]:
                p_sub = doc.add_paragraph()
                p_sub.paragraph_format.left_indent = Cm(0.5)
                p_sub.paragraph_format.space_before = Pt(1)
                p_sub.paragraph_format.space_after = Pt(1)
                if is_teacher:
                    render_rich_text(p_sub, f"*{sub_q}:*  ➔  {sub_ans}", base_size=base_size)
                else:
                    render_rich_text(p_sub, f"{sub_q}: .................................................................................................................", base_size=base_size)

        if is_teacher:
            add_solution_no_table(doc, ex["solution"])
        elif is_screen:
            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.left_indent = Cm(0.6)
            render_rich_text(p_sp, "✎ Bài giải chi tiết: ........................................................................................................", base_size=base_size)
        else:
            add_student_writing_lines(doc, num_lines=4)

    # Ghi chú SBT
    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(6)
    p_note.paragraph_format.space_after = Pt(4)
    p_note.paragraph_format.left_indent = Cm(0.4)
    render_rich_text(p_note, "*Chú ý:* Học sinh chủ động làm các bài tập trắc nghiệm trong Sách bài tập (SBT) từ Bài 1 đến Bài 6.", base_size=base_size, base_italic=True, base_color=COLOR_SOLUTION)

    # --------------------------------------------------------------------------
    # PHÂN MÔN SINH HỌC
    # --------------------------------------------------------------------------
    add_part_heading(doc, "PHẦN II. PHÂN MÔN SINH HỌC")

    for section in SINH_DATA:
        add_section_header(doc, section["lesson"].upper())
        for q in section["questions"]:
            p_q = doc.add_paragraph()
            p_q.paragraph_format.space_before = Pt(4)
            p_q.paragraph_format.space_after = Pt(2)
            render_question_prompt(p_q, f"Câu {q['num']}. {q['q']}", base_size=base_size)

            if is_teacher:
                add_solution_no_table(doc, q["ans"])
            elif is_screen:
                p_sp = doc.add_paragraph()
                p_sp.paragraph_format.left_indent = Cm(0.6)
                render_rich_text(p_sp, "✎ Trả lời / Phân tích: ...................................................................................................", base_size=base_size)
            else:
                add_student_writing_lines(doc, num_lines=4)

    # Dòng kết thúc
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(12)
    p_end.paragraph_format.space_after = Pt(6)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("--- HẾT ---")
    set_run_style(r_end, font_size=11, bold=True, italic=True, color=COLOR_PRIMARY)

    doc.save(filepath)
    print(f"[OK] Đã tạo thành công file Word: {filepath} ({os.path.getsize(filepath):,} bytes)")
    return filepath


# ==============================================================================
# 2. HÀM SINH FILE LATEX & BIÊN DỊCH PDF TRÌNH CHIẾU TIVI 16:9
# ==============================================================================

def generate_tivi_presentation_pdf():
    tex_filename = "De_Cuong_GK1_KHTN6_ThamKhao_DayTrenManHinh.tex"
    pdf_filename = "De_Cuong_GK1_KHTN6_ThamKhao_DayTrenManHinh.pdf"
    tex_path = os.path.join(SCRATCH_DIR, tex_filename)
    pdf_dest = os.path.join(OUTPUT_DIR, pdf_filename)

    tex_lines = []
    # Preamble chuẩn 16:9 widescreen không viền đen
    tex_lines.append(r"""\documentclass{article}
\usepackage[utf8]{vietnam}
\usepackage{times}
\usepackage{amsmath,amssymb}
\usepackage{tabularx}
\usepackage{array}
\usepackage{booktabs}
\usepackage{xcolor}
\usepackage[
  paperwidth=29.7cm,
  paperheight=16.706cm,
  top=0.9cm,
  bottom=0.9cm,
  left=1.2cm,
  right=1.2cm,
  footskip=0pt,
  headheight=0pt,
  headsep=0pt
]{geometry}

\pagestyle{empty}
\setlength{\parindent}{0pt}
\setlength{\parskip}{0pt}

\definecolor{brandRed}{RGB}{211, 47, 47}
\definecolor{brandGreen}{RGB}{46, 125, 50}
\definecolor{brandBlue}{RGB}{21, 101, 192}
\definecolor{textColor}{RGB}{20, 20, 20}

\color{textColor}

\begin{document}
""")

    # --------------------------------------------------------------------------
    # SLIDE 0: TIÊU ĐỀ
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 0: TIÊU ĐỀ ĐỀ CƯƠNG
% ==============================================================================
\begin{center}
\vspace*{0.2cm}
{\fontsize{24}{30}\selectfont \textbf{\color{brandRed}TRƯỜNG TRUNG HỌC CƠ SỞ -- NĂM HỌC 2026 -- 2027}}\\[10pt]
{\fontsize{34}{42}\selectfont \textbf{\color{brandBlue}ĐỀ CƯƠNG ÔN TẬP GIỮA HỌC KÌ I}}\\[8pt]
{\fontsize{28}{36}\selectfont \textbf{\color{brandGreen}MÔN: KHOA HỌC TỰ NHIÊN 6}}\\[8pt]
{\fontsize{22}{28}\selectfont \textit{(Học sinh dùng tham khảo -- Phân môn Vật lí \& Sinh học)}}\\[16pt]
\rule{16cm}{1.5pt}\\[20pt]
{\fontsize{26}{34}\selectfont \textbf{Giáo viên hướng dẫn: Thầy TRẦN VĂN THẠNH -- 0777.470.803}}\\[8pt]
{\fontsize{22}{28}\selectfont \textit{Cơ sở 1: 6/15 Nguyễn Hoàng, Kim Long | Cơ sở 2: 24 Đặng Thái Thân, TP Huế}}\\[16pt]
{\fontsize{24}{30}\selectfont \textbf{CẤU TRÚC: LÝ THUYẾT TRỌNG TÂM \& BÀI TẬP VẬN DỤNG}}
\end{center}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 1: VẬT LÍ - CÂU 1 & CÂU 2
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 1: VẬT LÍ - CÂU 1 & CÂU 2
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandRed}I. PHÂN MÔN VẬT LÍ -- LÝ THUYẾT (CÂU 1 \& CÂU 2)}}
\vspace{4pt}

\fontsize{19}{24}\selectfont
\textbf{1. Khoa học tự nhiên (KHTN) \& Các lĩnh vực chủ yếu:}
\begin{itemize}
  \setlength{\itemsep}{1pt}
  \item \textbf{Khái niệm:} KHTN nghiên cứu về các sự vật, hiện tượng, quy luật tự nhiên và ảnh hưởng của chúng đến con người, môi trường.
  \item \textbf{5 lĩnh vực chủ yếu:}
  \begin{itemize}
    \setlength{\itemsep}{0pt}
    \item \textit{Vật lí học:} Nghiên cứu vật chất, chuyển động, lực và năng lượng.
    \item \textit{Hóa học:} Nghiên cứu về chất và sự biến đổi của chất.
    \item \textit{Sinh học:} Nghiên cứu về sinh vật sống và môi trường.
    \item \textit{Khoa học Trái Đất:} Nghiên cứu về Trái Đất và khí quyển.
    \item \textit{Thiên văn học:} Nghiên cứu các thiên thể và vũ trụ.
  \end{itemize}
\end{itemize}

\vspace{3pt}
\textbf{2. Phân biệt Vật sống và Vật không sống:}
\begin{itemize}
  \setlength{\itemsep}{1pt}
  \item \textbf{Vật sống (hữu sinh):} Có trao đổi chất, lớn lên, vận động, cảm ứng, sinh sản (\textit{VD: em bé, cây lúa, con chó, vi khuẩn}).
  \item \textbf{Vật không sống (vô sinh):} Không có các dấu hiệu sống trên (\textit{VD: hòn đá, chiếc bàn, cây cầu, giọt nước}).
\end{itemize}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 2: VẬT LÍ - CÂU 3: AN TOÀN PHÒNG THỰC HÀNH
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 2: VẬT LÍ - CÂU 3: AN TOÀN PHÒNG THỰC HÀNH
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandRed}CÂU 3: AN TOÀN \& KÍ HIỆU CẢNH BÁO TRONG PHÒNG THỰC HÀNH}}
\vspace{6pt}

\fontsize{20}{26}\selectfont
\textbf{1. Các quy tắc an toàn bắt buộc:}
\begin{itemize}
  \item Mặc áo blouse, đeo kính và găng tay bảo hộ khi cần; cột tóc gọn gàng.
  \item Chỉ làm thí nghiệm khi có sự cho phép và hướng dẫn của giáo viên.
  \item Tuyệt đối không ăn, uống, đùa nghịch, chạy nhảy trong phòng thực hành.
  \item Không nếm hoặc ngửi trực tiếp hóa chất; đọc kĩ nhãn trước khi dùng.
  \item Nắm vững vị trí bình chữa cháy, hộp sơ cứu; thu dọn và rửa tay sau khi làm xong.
\end{itemize}

\vspace{6pt}
\textbf{2. Các kí hiệu cảnh báo thông dụng:}
\begin{itemize}
  \item \textbf{Chất độc:} Hình đầu lâu xương chéo $\rightarrow$ Nguy hiểm chết người.
  \item \textbf{Chất ăn mòn:} Hình axit đổ làm thủng kim loại/da tay $\rightarrow$ Gây bỏng rát.
  \item \textbf{Chất dễ cháy:} Hình ngọn lửa $\rightarrow$ Tránh xa nguồn nhiệt.
  \item \textbf{Nguy hiểm điện:} Hình tia sét $\rightarrow$ Cẩn trọng nguồn điện cao thế.
\end{itemize}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 3: VẬT LÍ - CÂU 4: ĐO CHIỀU DÀI
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 3: VẬT LÍ - CÂU 4: ĐO CHIỀU DÀI
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandRed}CÂU 4: ĐO CHIỀU DÀI TRONG HỆ THỐNG ĐO LƯỜNG HỢP PHÁP}}
\vspace{6pt}

\fontsize{20}{26}\selectfont
\begin{itemize}
  \item \textbf{Đơn vị đo hợp pháp (SI):} Là \textbf{mét} (kí hiệu: $\text{m}$).\\
  \textit{Đơn vị khác:} $1\text{ km} = 1000\text{ m}$; $1\text{ m} = 10\text{ dm} = 100\text{ cm} = 1000\text{ mm}$.
  \item \textbf{Dụng cụ đo:} Thước thẳng, thước cuộn, thước dây, thước kẹp\dots
  \item \textbf{Khái niệm quan trọng:}
  \begin{itemize}
    \item \textbf{GHĐ (Giới hạn đo):} Là độ dài lớn nhất ghi trên thước.
    \item \textbf{ĐCNN (Độ chia nhỏ nhất):} Là độ dài giữa hai vạch chia liên tiếp trên thước.
  \end{itemize}
  \item \textbf{Quy trình 5 bước đo chiều dài:}
  \begin{enumerate}
    \item \textit{Bước 1:} Ước lượng chiều dài cần đo.
    \item \textit{Bước 2:} Chọn thước có GHĐ và ĐCNN thích hợp.
    \item \textit{Bước 3:} Đặt thước dọc theo chiều dài, 1 đầu ngang bằng vạch số 0.
    \item \textit{Bước 4:} Đặt mắt nhìn vuông góc với cạnh thước ở đầu kia của vật.
    \item \textit{Bước 5:} Đọc và ghi kết quả đo theo vạch chia gần nhất và theo ĐCNN.
  \end{enumerate}
\end{itemize}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 4: VẬT LÍ - CÂU 5 & CÂU 6: ĐO KHỐI LƯỢNG & ĐO THỜI GIAN
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 4: VẬT LÍ - CÂU 5 & CÂU 6: ĐO KHỐI LƯỢNG & ĐO THỜI GIAN
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandRed}CÂU 5 \& CÂU 6: ĐO KHỐI LƯỢNG VÀ ĐO THỜI GIAN}}
\vspace{6pt}

\fontsize{19}{25}\selectfont
\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{\color{brandBlue}1. ĐO KHỐI LƯỢNG (CÂU 5):} & \textbf{\color{brandGreen}2. ĐO THỜI GIAN (CÂU 6):} \\
• \textbf{Đơn vị hợp pháp (SI):} Kilôgam ($\text{kg}$). & • \textbf{Đơn vị cơ bản (SI):} Giây ($\text{s}$). \\
• \textit{Đơn vị khác:} Tấn, tạ, yến, lạng ($\text{hg}$), $\text{g}$, $\text{mg}$. & • \textit{Đơn vị khác:} Giờ ($\text{h}$), phút ($\text{min}$), ngày, tháng, năm\dots \\
($1\text{ tấn} = 1000\text{ kg}$; $1\text{ lạng} = 100\text{ g}$; $1\text{ kg} = 1000\text{ g}$) & ($1\text{ ngày} = 24\text{ h}$; $1\text{ h} = 60\text{ min} = 3600\text{ s}$; $1\text{ s} = 1000\text{ ms}$) \\
• \textbf{Dụng cụ đo:} Cân (đồng hồ, điện tử, y tế, Roberval). & • \textbf{Dụng cụ đo:} Đồng hồ (bấm giây, đeo tay, treo tường). \\
• \textbf{5 bước cân:}\newline
1. Ước lượng khối lượng.\newline
2. Chọn cân có GHĐ, ĐCNN phù hợp.\newline
3. Hiệu chỉnh kim cân về vạch 0 (hoặc Tare).\newline
4. Đặt vật lên đĩa cân.\newline
5. Đọc và ghi kết quả theo ĐCNN. & • \textbf{5 bước đo thời gian (đồng hồ bấm giây):}\newline
1. Ước lượng khoảng thời gian.\newline
2. Chọn đồng hồ bấm giây phù hợp.\newline
3. Hiệu chỉnh đồng hồ về số 0.\newline
4. Nhấn Start lúc bắt đầu, nhấn Stop lúc kết thúc.\newline
5. Đọc và ghi kết quả hiển thị. \\
\end{tabularx}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 5: BÀI TẬP 1 (XÁC ĐỊNH ĐCNN)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 5: BÀI TẬP 1 (XÁC ĐỊNH ĐCNN)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}BÀI TẬP 1: XÁC ĐỊNH ĐCNN TỪ KẾT QUẢ ĐO}}
\vspace{6pt}

\fontsize{21}{29}\selectfont
\textbf{Đề bài:} Xác định ĐCNN của các kết quả đo sau:

\vspace{6pt}
\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{a)} $2,5\text{ cm}$ & \textbf{b)} $3\text{ dm}$\\[6pt]
\textbf{c)} $12,2\text{ mm}$ & \textbf{d)} $15,1\text{ cm}$
\end{tabularx}

\vspace{14pt}
{\fontsize{22}{28}\selectfont \textbf{\color{brandRed}Hướng dẫn giải chi tiết:}}
\begin{itemize}
  \item \textbf{Nguyên tắc:} Kết quả đo luôn được ghi theo đơn vị và bậc thập phân của ĐCNN.
  \item \textbf{a) $2,5\text{ cm}$:} ĐCNN là \textbf{$0,1\text{ cm}$ ($1\text{ mm}$)} hoặc \textbf{$0,5\text{ cm}$}.
  \item \textbf{b) $3\text{ dm}$:} ĐCNN lớn nhất là \textbf{$1\text{ dm}$} (hoặc $0,5\text{ dm}$, $0,1\text{ dm}$ nếu trùng vạch).
  \item \textbf{c) $12,2\text{ mm}$:} ĐCNN là \textbf{$0,1\text{ mm}$} hoặc \textbf{$0,2\text{ mm}$} (thường gặp ở thước kẹp).
  \item \textbf{d) $15,1\text{ cm}$:} ĐCNN là \textbf{$0,1\text{ cm}$ ($1\text{ mm}$)}.
\end{itemize}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 6: BÀI TẬP 2 (CHỌN THƯỚC ĐO VƯỜN CỎ)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 6: BÀI TẬP 2 (CHỌN THƯỚC ĐO VƯỜN CỎ)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}BÀI TẬP 2: LỰA CHỌN DỤNG CỤ ĐO DIỆN TÍCH}}
\vspace{6pt}

\fontsize{21}{28}\selectfont
\textbf{Đề bài:} Để đo diện tích một vườn cỏ có kích thước $25 \times 30\text{ (m)}$. Trong tay em có 2 thước:
\begin{itemize}
  \item Thước gấp có $\text{GHĐ} = 2\text{ m}$.
  \item Thước cuộn có $\text{GHĐ} = 20\text{ m}$.
\end{itemize}
Em sẽ dùng thước nào để cho kết quả đo chính xác hơn? Vì sao?

\vspace{10pt}
{\fontsize{22}{28}\selectfont \textbf{\color{brandRed}Hướng dẫn giải chi tiết:}}
\begin{itemize}
  \item \textbf{Lựa chọn:} Dùng \textbf{thước cuộn có GHĐ $20\text{ m}$}.
  \item \textbf{Vì sao:}
  \begin{itemize}
    \item Nếu dùng thước gấp $2\text{ m}$, cạnh $25\text{ m}$ phải đặt đo 13 lần, cạnh $30\text{ m}$ đo 15 lần. Mỗi lần nhấc thước và đánh dấu đều gây sai số, tích lũy sai số rất lớn.
    \item Thước cuộn $20\text{ m}$ chỉ cần đo 2 lần cho mỗi cạnh (1 lần $20\text{ m}$ + 1 lần $5\text{ m}$ hoặc $10\text{ m}$), giảm tối đa số lần ghép thước nên kết quả \textbf{nhanh hơn và chính xác hơn nhiều}.
  \end{itemize}
\end{itemize}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 7: BÀI TẬP 3 (TÍNH KHỐI LƯỢNG 20 TÚI ĐƯỜNG)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 7: BÀI TẬP 3 (TÍNH KHỐI LƯỢNG 20 TÚI ĐƯỜNG)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}BÀI TẬP 3: TÍNH TOÁN KHỐI LƯỢNG}}
\vspace{4pt}

\fontsize{20}{26}\selectfont
\textbf{Đề bài:} Có 20 túi đường, ban đầu mỗi túi nặng $1\text{ kg}$, sau đó người ta cho thêm mỗi túi $2\text{ lạng}$ đường nữa. Khối lượng của 20 túi đường khi đó là bao nhiêu kg?

\vspace{6pt}
{\fontsize{21}{27}\selectfont \textbf{\color{brandRed}Hướng dẫn giải chi tiết:}}
\begin{itemize}
  \setlength{\itemsep}{2pt}
  \item \textbf{Đổi đơn vị:} $1\text{ lạng} = 100\text{ g} = 0,1\text{ kg} \implies 2\text{ lạng} = 200\text{ g} = \mathbf{0,2\text{ kg}}$.
  \item \textbf{Cách 1:} Khối lượng mỗi túi sau khi thêm: $1\text{ kg} + 0,2\text{ kg} = 1,2\text{ kg}$.\\
  $\rightarrow$ Tổng khối lượng 20 túi: $m = 20 \times 1,2\text{ kg} = \mathbf{24\text{ kg}}$.
  \item \textbf{Cách 2:} Khối lượng ban đầu: $20 \times 1\text{ kg} = 20\text{ kg}$.\\
  $\rightarrow$ Lượng thêm vào: $20 \times 2\text{ lạng} = 40\text{ lạng} = 4\text{ kg} \implies m = 20 + 4 = \mathbf{24\text{ kg}}$.
  \item \textbf{Đáp số:} $\mathbf{24\text{ kg}}$.
\end{itemize}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 8: BÀI TẬP 4 (CHỌN ĐỒNG HỒ ĐO THỜI GIAN)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 8: BÀI TẬP 4 (CHỌN ĐỒNG HỒ ĐO THỜI GIAN)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}BÀI TẬP 4: LỰA CHỌN ĐỒNG HỒ ĐO THỜI GIAN}}
\vspace{6pt}

\fontsize{21}{28}\selectfont
\textbf{Đề bài:} Để thực hiện đo thời gian từ cổng trường đến lớp học em dùng loại đồng hồ nào? Giải thích sự lựa chọn đó.

\vspace{10pt}
{\fontsize{22}{28}\selectfont \textbf{\color{brandRed}Hướng dẫn giải chi tiết:}}
\begin{itemize}
  \item \textbf{Lựa chọn:} Sử dụng \textbf{đồng hồ bấm giây} (hoặc đồng hồ đeo tay/điện thoại có chức năng bấm giờ Stop-watch).
  \item \textbf{Giải thích:}\newline
  1. Khoảng thời gian đi từ cổng trường đến lớp khá ngắn (khoảng từ 1 đến vài phút).\newline
  2. Đồng hồ bấm giây có nút khởi động (Start) và dừng (Stop) tức thời: bấm Start ngay lúc bước qua cổng và Stop ngay khi bước vào lớp.\newline
  3. Độ chia nhỏ nhất của đồng hồ bấm giây rất nhỏ ($0,01\text{ s}$ đến $1\text{ s}$) cho kết quả đo chính xác.\newline
  4. Đồng hồ bấm giây nhỏ gọn, cầm tay tiện lợi khi di chuyển (khác với đồng hồ treo tường hoặc đồng hồ quả lắc cồng kềnh cố định).
\end{itemize}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 9: BÀI TẬP 5 (XÁC ĐỊNH THỜI ĐIỂM Ô TÔ ĐẾN HẢI DƯƠNG & HẢI PHÒNG)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 9: BÀI TẬP 5 (THỜI ĐIỂM CHUYỂN ĐỘNG CỦA Ô TÔ)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}BÀI TẬP 5: XÁC ĐỊNH THỜI ĐIỂM CHUYỂN ĐỘNG}}
\vspace{6pt}

\fontsize{21}{28}\selectfont
\textbf{Đề bài:} Ô tô khởi hành từ Hà Nội lúc $8\text{ giờ } 15\text{ phút}$. Sau khi chạy được $1\text{ giờ}$ thì đến Hải Dương. Ô tô nghỉ ở Hải Dương $20\text{ phút}$, sau đó chạy $50\text{ phút}$ nữa thì đến Hải Phòng. Xác định thời điểm ô tô đến Hải Dương và Hải Phòng.

\vspace{10pt}
{\fontsize{22}{28}\selectfont \textbf{\color{brandRed}Hướng dẫn giải chi tiết:}}
\begin{enumerate}
  \item \textbf{Thời điểm ô tô đến Hải Dương:}
  \[
  8\text{ giờ } 15\text{ phút} + 1\text{ giờ} = \mathbf{9\text{ giờ } 15\text{ phút}}.
  \]
  \item \textbf{Thời điểm ô tô rời Hải Dương:}
  \[
  9\text{ giờ } 15\text{ phút} + 20\text{ phút} = 9\text{ giờ } 35\text{ phút}.
  \]
  \item \textbf{Thời điểm ô tô đến Hải Phòng:}
  \[
  9\text{ giờ } 35\text{ phút} + 50\text{ phút} = 9\text{ giờ } 85\text{ phút} = \mathbf{10\text{ giờ } 25\text{ phút}} \quad (85\text{p} = 1\text{h } 25\text{p}).
  \]
\end{enumerate}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 10: SINH HỌC - BÀI 18 (TẾ BÀO – ĐƠN VỊ CƠ BẢN CỦA SỰ SỐNG)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 10: SINH HỌC - BÀI 18: TẾ BÀO
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandRed}II. PHÂN MÔN SINH HỌC -- BÀI 18: TẾ BÀO}}
\vspace{6pt}

\fontsize{20}{26}\selectfont
\textbf{1. Tế bào là gì?}
\begin{itemize}
  \item Tế bào là \textbf{đơn vị cơ sở cấu tạo nên mọi cơ thể sống} và thực hiện chức năng sống.
  \item Mọi cơ thể (từ vi khuẩn đến con người) đều cấu tạo từ tế bào. Mọi hoạt động sống đều diễn ra trong tế bào.
\end{itemize}

\vspace{5pt}
\textbf{2. Kích thước và hình dạng của tế bào:}
\begin{itemize}
  \item \textbf{Kích thước:}
  \begin{itemize}
    \item Đa số có kích thước hiển vi ($1 - 100\,\mu\text{m}$), phải dùng kính hiển vi quang học quan sát (\textit{tế bào vi khuẩn, hồng cầu, vảy hành\dots}).
    \item Một số ít có kích thước lớn nhìn thấy bằng mắt thường/kính lúp (\textit{trứng cá, tép bưởi, trứng gà\dots}).
  \end{itemize}
  \item \textbf{Hình dạng:} Rất đa dạng: hình cầu (trứng), hình đĩa (hồng cầu), hình thoi (cơ trơn), hình sao (nơ-ron thần kinh), hình đa giác (vảy hành)\dots
  \item \textbf{Quy luật:} Hình dạng luôn thích nghi phù hợp tối ưu với \textbf{chức năng} mà tế bào đảm nhận.
\end{itemize}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 11: SINH HỌC - BÀI 19 (CẤU TẠO & CHỨC NĂNG TẾ BÀO)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 11: SINH HỌC - BÀI 19: CẤU TẠO VÀ CHỨC NĂNG TẾ BÀO
% ==============================================================================
{\fontsize{23}{29}\selectfont \textbf{\color{brandRed}BÀI 19: CẤU TẠO VÀ CHỨC NĂNG CÁC THÀNH PHẦN CỦA TẾ BÀO}}
\vspace{4pt}

\fontsize{18}{23}\selectfont
Mọi tế bào đều được cấu tạo từ 3 thành phần chính:

\vspace{3pt}
\begin{tabularx}{\linewidth}{|m{4.2cm}|m{8.0cm}|X|}
\hline
\rowcolor{gray!15}
\textbf{Thành phần chính} & \textbf{Đặc điểm cấu tạo} & \textbf{Chức năng cơ bản} \\
\hline
\textbf{1. Màng tế bào\newline (Màng sinh chất)} & Lớp màng mỏng bao ngoài. (Thực vật có thêm \textbf{thành tế bào} cellulose cứng cáp). & Bao bọc, bảo vệ tế bào; kiểm soát sự trao đổi chất chọn lọc giữa tế bào với môi trường. \\
\hline
\textbf{2. Chất tế bào\newline (Tế bào chất)} & Dạng dịch keo lỏng (bào tương) chứa các bào quan (ti thể, lục lạp ở thực vật, không bào\dots). & Nơi diễn ra hầu hết các hoạt động trao đổi chất và chuyển hóa năng lượng của tế bào. \\
\hline
\textbf{3. Nhân tế bào\newline (hoặc Vùng nhân)} & Chứa vật chất di truyền ($\text{DNA}$). Nhân thực có màng nhân bao bọc; nhân sơ chưa có. & Trung tâm điều khiển mọi hoạt động sống; lưu giữ và truyền đạt thông tin di truyền. \\
\hline
\end{tabularx}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 12: BẢNG TỔNG HỢP ĐÁP ÁN TOÀN BỘ ĐỀ CƯƠNG
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 12: BẢNG TỔNG HỢP ĐÁP ÁN ĐỀ CƯƠNG
% ==============================================================================
\begin{center}
{\fontsize{23}{29}\selectfont \textbf{\color{brandRed}BẢNG TỔNG HỢP ĐÁP SỐ VÀ CHỐT KIẾN THỨC}}
\end{center}
\vspace{2pt}

\fontsize{18}{23}\selectfont
\begin{tabularx}{\linewidth}{|m{5.2cm}|X|}
\hline
\rowcolor{gray!15}
\textbf{Nội dung câu hỏi / Bài tập} & \textbf{Đáp số \& Kết luận cốt lõi} \\
\hline
\textbf{Bài 1 (ĐCNN kết quả đo)} & a) $0,1\text{ cm}$ (hoặc $0,5\text{ cm}$) \quad b) $1\text{ dm}$ \quad c) $0,1\text{ mm}$ (hoặc $0,2\text{ mm}$) \quad d) $0,1\text{ cm}$. \\
\hline
\textbf{Bài 2 (Đo vườn cỏ $25\times 30\text{m}$)} & Chọn \textbf{thước cuộn GHĐ $20\text{ m}$} (ít lần đo $\rightarrow$ giảm tối đa sai số). \\
\hline
\textbf{Bài 3 (Khối lượng 20 túi đường)} & $\mathbf{24\text{ kg}}$ ($20\text{ túi} \times 1,2\text{ kg} = 24\text{ kg}$, với $2\text{ lạng} = 0,2\text{ kg}$). \\
\hline
\textbf{Bài 4 (Đo thời gian vào lớp)} & Chọn \textbf{đồng hồ bấm giây} (thao tác Start/Stop linh hoạt, ĐCNN nhỏ). \\
\hline
\textbf{Bài 5 (Thời điểm ô tô)} & Đến Hải Dương: \textbf{9 giờ 15 phút}; Đến Hải Phòng: \textbf{10 giờ 25 phút}. \\
\hline
\textbf{Sinh học Bài 18 (Tế bào)} & Đơn vị cấu trúc và chức năng cơ bản; hình dạng phù hợp chức năng. \\
\hline
\textbf{Sinh học Bài 19 (Cấu tạo)} & Gồm 3 phần: \textbf{Màng tế bào} (bảo vệ), \textbf{Chất tế bào} (trao đổi chất), \textbf{Nhân} (điều khiển). \\
\hline
\end{tabularx}

\end{document}
""")

    with open(tex_path, "w", encoding="utf-8") as f:
        f.write("".join(tex_lines))
    print(f"[OK] Đã ghi file LaTeX: {tex_path}")

    # Biên dịch pdflatex 2 lần để ổn định layout
    cmd = ["pdflatex", "-interaction=nonstopmode", "-output-directory", SCRATCH_DIR, tex_path]
    res1 = subprocess.run(cmd, capture_output=True, text=True)
    res2 = subprocess.run(cmd, capture_output=True, text=True)

    pdf_scratch = os.path.join(SCRATCH_DIR, pdf_filename)
    if os.path.exists(pdf_scratch):
        shutil.copy2(pdf_scratch, pdf_dest)
        print(f"[OK] Đã biên dịch và copy PDF thành công: {pdf_dest} ({os.path.getsize(pdf_dest):,} bytes)")
        return pdf_dest
    else:
        print("[ERROR] Không tìm thấy file PDF được tạo. Log:")
        print(res2.stdout[-1500:])
        return None

# ==============================================================================
# HÀM MAIN
# ==============================================================================

def main():
    print("=" * 75)
    print("BẮT ĐẦU TẠO 3 FILE CHO ĐỀ CƯƠNG ÔN TẬP GIỮA HỌC KÌ I KHTN 6 (THAM KHẢO)")
    print("1. FILE HỌC SINH (.docx)")
    print("2. FILE GIÁO VIÊN (.docx có lời giải chi tiết)")
    print("3. FILE DẠY TRÊN MÀN HÌNH (Slide PDF 16:9 widescreen cho Tivi 65 inch)")
    print("=" * 75)

    f_hs = generate_word_document(mode="student")
    f_gv = generate_word_document(mode="teacher")
    f_screen_docx = generate_word_document(mode="screen")
    f_screen_pdf = generate_tivi_presentation_pdf()

    # Đồng bộ OneDrive nếu có
    all_files = [f_hs, f_gv, f_screen_docx, f_screen_pdf]
    print("\n[+] Đồng bộ các file sang OneDrive (nếu có):")
    for od in ONEDRIVE_DIRS:
        if os.path.exists(od):
            for src_path in all_files:
                if src_path and os.path.exists(src_path):
                    shutil.copy2(src_path, os.path.join(od, os.path.basename(src_path)))
                    print(f"    -> Đã copy {os.path.basename(src_path)} vào {od}")

    print("\n" + "=" * 75)
    print("HOÀN THÀNH TẤT CẢ 3 FILE THEO ĐÚNG CHUẨN FORM ĐỀ KIỂM TRA!")
    print("=" * 75)

if __name__ == "__main__":
    main()
