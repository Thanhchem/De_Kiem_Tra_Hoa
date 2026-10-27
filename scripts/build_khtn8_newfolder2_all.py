# -*- coding: utf-8 -*-
r"""
Script biên soạn trọn bộ 3 tài liệu chuẩn FORM ĐỀ KIỂM TRA cho:
ĐỀ CƯƠNG ÔN TẬP GIỮA KÌ I - MÔN KHTN 8 (K8)
Thư mục: c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn8\KHTN8-IN-DEKIEMTRADINHKI\New folder (2)
1. De_Cuong_GK1_KHTN8_K8_HocSinh.docx
2. De_Cuong_GK1_KHTN8_K8_GiaoVien.docx
3. De_Cuong_GK1_KHTN8_K8_DayTrenManHinh.pdf (Slide 16:9 Widescreen cho Tivi 65 inch)
Và bản Word: De_Cuong_GK1_KHTN8_K8_DayTrenManHinh.docx
Thầy TRẦN VĂN THẠNH - 0777.470.803
"""

import os
import sys
import subprocess
import shutil

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# Import các hàm chuẩn từ form_de_kiem_tra
sys.path.insert(0, r"c:\Antigravity_Thanh\scripts")
from form_de_kiem_tra import (
    setup_header_footer,
    build_title_block,
    add_section_header,
    render_question_prompt,
    render_options_no_table,
    render_statement,
    add_solution_no_table,
    add_student_writing_lines,
    render_rich_text,
    set_run_style,
    COLOR_PRIMARY,
    COLOR_CORRECT,
    COLOR_SOLUTION,
    COLOR_HF_RED,
    COLOR_HF_GREEN_HEX
)

OUTPUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
SCRATCH_DIR = r"c:\Antigravity_Thanh\scratch"
os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(SCRATCH_DIR, exist_ok=True)

TARGET_ONEDRIVE_DIRS = [
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn8\KHTN8-IN-DEKIEMTRADINHKI\New folder (2)",
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn8\KHTN8-IN-DEKIEMTRADINHKI",
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn8"
]

# ==============================================================================
# DỮ LIỆU ĐỀ CƯƠNG KHTN 8 (K8)
# ==============================================================================

MCQ_DATA = [
    {
        "id": 1,
        "content": "Trong học tập môn Khoa học tự nhiên, việc quan sát trực tiếp các sự vật, hiện tượng để thu thập thông tin thuộc phương pháp nào sau đây?",
        "options": [
            "A. Phương pháp quan sát.",
            "B. Phương pháp suy luận.",
            "C. Phương pháp mô hình hóa.",
            "D. Phương pháp tính toán."
        ],
        "answer": "A",
        "solution": "Phương pháp quan sát là phương pháp thu thập thông tin trực tiếp về các sự vật, hiện tượng trong tự nhiên bằng các giác quan hoặc thông qua các thiết bị hỗ trợ quan sát (kính lúp, kính hiển vi...)."
    },
    {
        "id": 2,
        "content": "Khi tiến hành một thí nghiệm khoa học, việc đầu tiên cần làm để bảo đảm an toàn và thực hiện đúng quy trình là gì?",
        "options": [
            "A. Thực hiện thí nghiệm ngay lập tức.",
            "B. Đọc và tìm hiểu hướng dẫn thí nghiệm.",
            "C. Thu dọn dụng cụ sau thí nghiệm.",
            "D. Ghi kết quả trước khi thí nghiệm."
        ],
        "answer": "B",
        "solution": "Trước khi làm bất kì thí nghiệm nào, bước đầu tiên bắt buộc là phải đọc và tìm hiểu kĩ hướng dẫn thí nghiệm (dụng cụ, hóa chất, quy tắc an toàn và các bước tiến hành) để tránh gây tai nạn và thực hiện chuẩn xác."
    },
    {
        "id": 3,
        "content": "Nồng độ phần trăm của dung dịch cho biết số gam chất tan có trong",
        "options": [
            "A. 100 gam dung dịch.",
            "B. 100 gam dung môi.",
            "C. 1 lít dung dịch.",
            "D. 1 lít dung môi."
        ],
        "answer": "A",
        "solution": "Theo định nghĩa: Nồng độ phần trăm (kí hiệu C%) của dung dịch cho biết số gam chất tan có trong 100 gam dung dịch."
    },
    {
        "id": 4,
        "content": "Phương trình hóa học dùng để biểu diễn",
        "options": [
            "A. thành phần các chất trong hỗn hợp.",
            "B. phản ứng hóa học bằng công thức hóa học.",
            "C. sự thay đổi trạng thái của chất.",
            "D. khối lượng riêng của các chất."
        ],
        "answer": "B",
        "solution": "Phương trình hóa học là phương thức dùng công thức hóa học để biểu diễn ngắn gọn, khoa học một phản ứng hóa học."
    },
    {
        "id": 5,
        "content": "Theo định luật bảo toàn khối lượng, trong một phản ứng hóa học, tổng khối lượng các chất tham gia phản ứng",
        "options": [
            "A. lớn hơn tổng khối lượng sản phẩm.",
            "B. nhỏ hơn tổng khối lượng sản phẩm.",
            "C. bằng tổng khối lượng các sản phẩm.",
            "D. luôn bằng không."
        ],
        "answer": "C",
        "solution": "Định luật bảo toàn khối lượng phát biểu: \"Trong một phản ứng hóa học, tổng khối lượng của các chất tham gia phản ứng bằng tổng khối lượng của các sản phẩm tạo thành\"."
    }
]

TF_DATA = [
    {
        "id": 6,
        "prompt": "Câu 6. Xác định mỗi phát biểu sau là đúng hay sai:",
        "items": [
            {
                "id": "a",
                "text": "Nồng độ mol của dung dịch cho biết số mol chất tan có trong một lít dung dịch.",
                "ans": True,
                "reason": "Đúng theo định nghĩa nồng độ mol (C_M = n / V)."
            },
            {
                "id": "b",
                "text": "Đơn vị thường dùng của nồng độ mol là mol/L.",
                "ans": True,
                "reason": "Đúng, nồng độ mol có đơn vị là mol/L hoặc kí hiệu là M."
            },
            {
                "id": "c",
                "text": "Khối lượng dung dịch bằng khối lượng chất tan cộng khối lượng dung môi.",
                "ans": True,
                "reason": "Đúng theo công thức bảo toàn khối lượng dung dịch: m_dd = m_ct + m_dm."
            },
            {
                "id": "d",
                "text": "Định luật bảo toàn khối lượng chỉ đúng với phản ứng không tạo ra chất khí.",
                "ans": False,
                "reason": "Sai. Định luật bảo toàn khối lượng nghiệm đúng cho mọi phản ứng hóa học (kể cả phản ứng có tạo ra chất khí, nếu đo trong hệ kín khối lượng vẫn hoàn toàn bằng nhau)."
            }
        ]
    },
    {
        "id": 7,
        "prompt": "Câu 7. Cho các phát biểu sau về nồng độ mol:",
        "items": [
            {
                "id": "a",
                "text": "Nồng độ mol của dung dịch cho biết số mol chất tan có trong 1000 lít dung dịch.",
                "ans": False,
                "reason": "Sai. Nồng độ mol cho biết số mol chất tan có trong 1 lít dung dịch (không phải 1000 lít)."
            },
            {
                "id": "b",
                "text": "Đơn vị thường dùng của nồng độ mol là mol/lít (mol/L).",
                "ans": True,
                "reason": "Đúng, đơn vị nồng độ mol là mol/L hoặc M."
            },
            {
                "id": "c",
                "text": "Hòa tan 0,5 mol NaCl vào nước, thu được 1 lít dd NaCl thì nồng độ mol của dung dịch là 0,5 M.",
                "ans": True,
                "reason": "Đúng, áp dụng C_M = n / V = 0,5 mol / 1 L = 0,5 M."
            }
        ]
    },
    {
        "id": 8,
        "prompt": "Câu 8. Xác định mỗi phát biểu sau là đúng hay sai:",
        "items": [
            {
                "id": "a",
                "text": "Trong phản ứng hóa học, tổng khối lượng các chất tham gia phản ứng bằng tổng khối lượng các sản phẩm.",
                "ans": True,
                "reason": "Đúng, đây là nội dung cốt lõi của định luật bảo toàn khối lượng."
            },
            {
                "id": "b",
                "text": "Định luật bảo toàn khối lượng được áp dụng cho các phản ứng hóa học.",
                "ans": True,
                "reason": "Đúng, định luật bảo toàn khối lượng được áp dụng cho mọi phản ứng hóa học để cân bằng phương trình và tính toán lượng chất."
            },
            {
                "id": "c",
                "text": "Khi đốt cháy hoàn toàn 12 g magnesium trong oxygen, thu được 20 g magnesium oxide thì khối lượng oxygen đã phản ứng là 8 g.",
                "ans": True,
                "reason": "Đúng. Áp dụng ĐLBTKL: m_O2 = m_MgO - m_Mg = 20 - 12 = 8 g."
            }
        ]
    },
    {
        "id": 9,
        "prompt": "Câu 9. Xác định mỗi phát biểu sau là đúng hay sai:",
        "items": [
            {
                "id": "a",
                "text": "Khối lượng dung dịch bằng tổng khối lượng của chất tan và dung môi.",
                "ans": True,
                "reason": "Đúng, m_dd = m_ct + m_dm."
            },
            {
                "id": "b",
                "text": "Khi cho thêm chất tan vào 1 dung dịch chưa bão hòa, nồng độ phần trăm của dung dịch sẽ giảm.",
                "ans": False,
                "reason": "Sai. Khi cho thêm chất tan vào dung dịch chưa bão hòa thì lượng chất tan tăng lên, nồng độ phần trăm của dung dịch sẽ tăng lên."
            },
            {
                "id": "c",
                "text": "Khi pha thêm 100ml nước cất vào 100ml dung dịch NaOH nồng độ 1M, nồng độ mol của dung dịch NaOH thu được giảm đi một nửa, còn 0,5M.",
                "ans": True,
                "reason": "Đúng. Thể tích dung dịch tăng từ 100 mL lên 200 mL (gấp 2 lần) trong khi số mol NaOH không đổi, do đó nồng độ mol giảm một nửa: C_M' = 1 M / 2 = 0,5 M."
            },
            {
                "id": "d",
                "text": "Cho 7 g iron phản ứng vừa đủ với 4 g sulfur, thu được 11 g iron sulfide.",
                "ans": True,
                "reason": "Đúng. Theo định luật bảo toàn khối lượng: m_FeS = m_Fe + m_S = 7 + 4 = 11 g."
            }
        ]
    }
]

SHORT_DATA = [
    {
        "id": 10,
        "content": "Viết công thức tính nồng độ phần trăm của dung dịch.",
        "solution": (
            "• Công thức tính nồng độ phần trăm:\n"
            "  C% = (m_ct / m_dd) × 100%\n"
            "• Trong đó:\n"
            "  + C%: Nồng độ phần trăm của dung dịch (%).\n"
            "  + m_ct: Khối lượng chất tan (gam).\n"
            "  + m_dd: Khối lượng dung dịch (gam), với m_dd = m_ct + m_dm (m_dm là khối lượng dung môi)."
        )
    },
    {
        "id": 11,
        "content": "Hòa tan 15 g muối vào 85 g nước. Tính nồng độ phần trăm của dung dịch thu được.",
        "solution": (
            "• Khối lượng dung dịch thu được là:\n"
            "  m_dd = m_muối + m_nước = 15 + 85 = 100 (g).\n"
            "• Nồng độ phần trăm của dung dịch muối thu được là:\n"
            "  C% = (m_ct / m_dd) × 100% = (15 / 100) × 100% = 15%.\n"
            "• Đáp số: 15%."
        )
    },
    {
        "id": 12,
        "content": "Đốt cháy hoàn toàn 6 g carbon trong 16 g oxygen, thu được carbon dioxide. Tính khối lượng carbon dioxide tạo thành theo định luật bảo toàn khối lượng.",
        "solution": (
            "• Phương trình chữ: Carbon + Oxygen ➔ Carbon dioxide.\n"
            "• Theo định luật bảo toàn khối lượng:\n"
            "  m_C + m_O2 = m_CO2\n"
            "  ➔ m_CO2 = 6 + 16 = 22 (g).\n"
            "• Đáp số: 22 g."
        )
    },
    {
        "id": 13,
        "content": "Cần bao nhiêu mol NaOH để pha chế 500 mL dung dịch NaOH 2 M?",
        "solution": (
            "• Đổi thể tích: V = 500 mL = 0,5 L.\n"
            "• Số mol NaOH cần dùng là:\n"
            "  n_NaOH = C_M × V = 2 × 0,5 = 1 (mol).\n"
            "• Đáp số: 1 mol."
        )
    },
    {
        "id": 14,
        "content": "Cần thêm bao nhiêu gam nước vào 100 g dung dịch NaCl 20% để thu được dung dịch NaCl nồng độ 5%?",
        "solution": (
            "• Khối lượng chất tan NaCl có trong 100 g dung dịch NaCl 20% là:\n"
            "  m_NaCl = (100 × 20) / 100 = 20 (g).\n"
            "• Khi pha thêm nước, khối lượng chất tan NaCl không đổi (vẫn là 20 g).\n"
            "• Khối lượng dung dịch NaCl 5% mới thu được là:\n"
            "  m_dd mới = (m_NaCl × 100%) / C% = (20 × 100) / 5 = 400 (g).\n"
            "• Khối lượng nước cần thêm vào là:\n"
            "  m_nước thêm = m_dd mới - m_dd ban đầu = 400 - 100 = 300 (g).\n"
            "• Đáp số: 300 g."
        )
    }
]

ESSAY_DATA = [
    {
        "id": 1,
        "content": "Viết các CT tính toán? Phát biểu định luật bảo toàn khối lượng? Giải thích.",
        "solution": (
            "1. CÁC CÔNG THỨC TÍNH TOÁN CƠ BẢN TRONG HÓA HỌC:\n"
            "   • Khối lượng dung dịch: m_dd = m_ct + m_dm (gam).\n"
            "   • Nồng độ phần trăm: C% = (m_ct / m_dd) × 100%  ➔  m_ct = (m_dd × C%) / 100%;  m_dd = (m_ct × 100%) / C%.\n"
            "   • Nồng độ mol: C_M = n / V (mol/L hoặc M)  ➔  n = C_M × V;  V = n / C_M.\n"
            "   • Số mol theo khối lượng: n = m / M (mol)  ➔  m = n × M;  M = m / n.\n"
            "   • Số mol theo thể tích chất khí (ở điều kiện chuẩn 25 °C, 1 bar): n = V / 24,79  ➔  V = n × 24,79 (lít).\n"
            "2. PHÁT BIỂU ĐỊNH LUẬT BẢO TOÀN KHỐI LƯỢNG:\n"
            "   \"Trong một phản ứng hóa học, tổng khối lượng của các chất tham gia phản ứng bằng tổng khối lượng của các sản phẩm tạo thành.\"\n"
            "   • Công thức: Giả sử có phản ứng A + B ➔ C + D thì: m_A + m_B = m_C + m_D.\n"
            "3. GIẢI THÍCH ĐỊNH LUẬT:\n"
            "   Trong phản ứng hóa học, chỉ có liên kết giữa các nguyên tử bị biến đổi, làm cho phân tử này biến thành phân tử khác. "
            "Còn số nguyên tử của mỗi nguyên tố được bảo toàn (không thay đổi trước và sau phản ứng). "
            "Vì khối lượng của mỗi nguyên tử không đổi, nên tổng khối lượng của các chất tham gia phản ứng luôn bằng tổng khối lượng của các sản phẩm tạo thành."
        )
    },
    {
        "id": 2,
        "content": "Phương trình hóa học là gì? Nêu ý nghĩa của pthh?",
        "solution": (
            "1. KHÁI NIỆM PHƯƠNG TRÌNH HÓA HỌC:\n"
            "   • Phương trình hóa học là phương thức biểu diễn ngắn gọn phản ứng hóa học bằng công thức hóa học của các chất tham gia và chất sản phẩm.\n"
            "   • Gồm CTHH của các chất phản ứng ở vế trái, CTHH của các sản phẩm ở vế phải, nối với nhau bằng dấu mũi tên (➔), "
            "có các hệ số thích hợp đặt trước các CTHH để số nguyên tử mỗi nguyên tố ở hai vế bằng nhau.\n"
            "   • Ví dụ: 2H_2 + O_2 --t°--> 2H_2O.\n"
            "2. Ý NGHĨA CỦA PHƯƠNG TRÌNH HÓA HỌC:\n"
            "   • Cho biết các chất tham gia phản ứng và các chất sản phẩm tạo thành.\n"
            "   • Cho biết tỉ lệ về số nguyên tử, số phân tử giữa các chất trong phản ứng (chính là tỉ lệ các hệ số đứng trước mỗi chất).\n"
            "   • Cho biết tỉ lệ về số mol, tỉ lệ về khối lượng (hoặc thể tích chất khí) giữa các chất, làm cơ sở khoa học để tính toán lượng chất phản ứng hoặc tạo thành."
        )
    }
]

# ==============================================================================
# 1. HÀM TẠO FILE WORD
# ==============================================================================

def generate_word_document(mode="student"):
    is_teacher = (mode == "teacher")
    is_screen = (mode == "screen")

    if mode == "teacher":
        filename = "De_Cuong_GK1_KHTN8_K8_GiaoVien.docx"
    elif mode == "screen":
        filename = "De_Cuong_GK1_KHTN8_K8_DayTrenManHinh.docx"
    else:
        filename = "De_Cuong_GK1_KHTN8_K8_HocSinh.docx"

    filepath = os.path.join(OUTPUT_DIR, filename)

    doc = docx.Document()
    setup_header_footer(doc)

    build_title_block(
        doc,
        school_line="TRƯỜNG TRUNG HỌC CƠ SỞ",
        exam_title="ĐỀ CƯƠNG ÔN TẬP KIỂM TRA GIỮA KÌ I - MÔN KHTN 8",
        subject_line="NĂM HỌC 2026 - 2027 (K8)",
        num_pages="03",
        duration="45 phút",
        exam_code="KHTN8-GK1",
        atomic_masses="H = 1; C = 12; O = 16; Na = 23; Mg = 24; S = 32; Cl = 35,5; Fe = 56",
        is_teacher=is_teacher
    )

    base_size = 11.5 if is_screen else 10.5

    # I. TRẮC NGHIỆM
    add_section_header(doc, "I. PHẦN TRẮC NGHIỆM NHIỀU LỰA CHỌN", "Chọn một phương án trả lời đúng nhất trong mỗi câu sau:")
    for q in MCQ_DATA:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, f"Câu {q['id']}. {q['content']}", base_size=base_size)

        render_options_no_table(doc, q["options"], is_teacher=is_teacher)

        if is_teacher and q.get("solution"):
            add_solution_no_table(doc, f"Đáp án đúng: *{q['answer']}*.\n{q['solution']}")
        elif is_screen:
            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.left_indent = Cm(0.6)
            p_sp.paragraph_format.space_before = Pt(2)
            p_sp.paragraph_format.space_after = Pt(2)
            render_rich_text(p_sp, "✎ Ghi chú / Giải thích: ........................................................................................................", base_size=base_size-1)

    # TRẢ LỜI ĐÚNG - SAI
    add_section_header(doc, "TRẢ LỜI ĐÚNG - SAI", "Trong mỗi ý ở mỗi câu, học sinh chọn Đúng hoặc Sai:")
    for q in TF_DATA:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, q["prompt"], base_size=base_size)

        for item in q["items"]:
            p_it = doc.add_paragraph()
            p_it.paragraph_format.left_indent = Cm(0.5)
            p_it.paragraph_format.space_before = Pt(1)
            p_it.paragraph_format.space_after = Pt(1)
            ans_str = "Đúng" if item["ans"] else "Sai"

            if is_teacher:
                render_statement(p_it, f"{item['id']}) {item['text']}", ans=ans_str, is_teacher=True, base_size=base_size)
                p_rea = doc.add_paragraph()
                p_rea.paragraph_format.left_indent = Cm(1.0)
                p_rea.paragraph_format.space_before = Pt(0)
                p_rea.paragraph_format.space_after = Pt(2)
                render_rich_text(p_rea, f"↳ _Giải thích:_ {item['reason']}", base_size=base_size-1, base_italic=True, base_color=COLOR_SOLUTION)
            else:
                render_rich_text(p_it, f"{item['id']}) {item['text']}   [   ] Đúng     [   ] Sai", base_size=base_size)

    # TRẢ LỜI NGẮN
    add_section_header(doc, "TRẢ LỜI NGẮN", "Học sinh tính toán và viết câu trả lời ngắn gọn vào chỗ trống:")
    for q in SHORT_DATA:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, f"Câu {q['id']}. {q['content']}", base_size=base_size)

        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.left_indent = Cm(0.8)
        p_ans.paragraph_format.space_before = Pt(1)
        p_ans.paragraph_format.space_after = Pt(2)

        if is_teacher:
            add_solution_no_table(doc, q["solution"])
        elif is_screen:
            render_rich_text(p_ans, "➔ Trả lời: ...................................................................................................................", base_size=base_size)
        else:
            add_student_writing_lines(doc, num_lines=3)

    # II. PHẦN TỰ LUẬN
    add_section_header(doc, "II. PHẦN TỰ LUẬN", "Học sinh trình bày chi tiết lời giải vào bài làm:")
    for q in ESSAY_DATA:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, f"Câu {q['id']}. {q['content']}", base_size=base_size)

        if is_teacher:
            add_solution_no_table(doc, q["solution"])
        elif is_screen:
            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.left_indent = Cm(0.6)
            render_rich_text(p_sp, "✎ Bài giải chi tiết: ........................................................................................................", base_size=base_size)
        else:
            add_student_writing_lines(doc, num_lines=5)

    # Dòng kết thúc
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(12)
    p_end.paragraph_format.space_after = Pt(6)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("--- HẾT ---")
    set_run_style(r_end, font_size=11, bold=True, italic=True, color=COLOR_PRIMARY)

    doc.save(filepath)
    print(f"[OK] Đã tạo thành công file Word: {filepath} ({os.path.getsize(filepath):,} bytes)")
    return filepath


# ==============================================================================
# 2. HÀM TẠO FILE LATEX & PDF TRÌNH CHIẾU TIVI 16:9
# ==============================================================================

def generate_tivi_presentation_pdf():
    tex_filename = "De_Cuong_GK1_KHTN8_K8_DayTrenManHinh.tex"
    pdf_filename = "De_Cuong_GK1_KHTN8_K8_DayTrenManHinh.pdf"
    tex_path = os.path.join(SCRATCH_DIR, tex_filename)
    pdf_dest = os.path.join(OUTPUT_DIR, pdf_filename)

    tex_lines = []
    # Preamble chuẩn 16:9 widescreen
    tex_lines.append(r"""\documentclass{article}
\usepackage[utf8]{vietnam}
\usepackage{times}
\usepackage{amsmath,amssymb}
\usepackage{tabularx}
\usepackage{array}
\usepackage{booktabs}
\usepackage{xcolor}
\usepackage[version=4]{mhchem}
\usepackage[
  paperwidth=29.7cm,
  paperheight=16.706cm,
  top=0.9cm,
  bottom=0.9cm,
  left=1.2cm,
  right=1.2cm,
  footskip=0pt,
  headheight=0pt,
  headsep=0pt
]{geometry}

\pagestyle{empty}
\setlength{\parindent}{0pt}
\setlength{\parskip}{0pt}

\definecolor{brandRed}{RGB}{211, 47, 47}
\definecolor{brandGreen}{RGB}{46, 125, 50}
\definecolor{brandBlue}{RGB}{21, 101, 192}
\definecolor{textColor}{RGB}{20, 20, 20}

\color{textColor}

\begin{document}
""")

    # --------------------------------------------------------------------------
    # SLIDE 0: TIÊU ĐỀ
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 0: TIÊU ĐỀ ĐỀ CƯƠNG
% ==============================================================================
\begin{center}
\vspace*{0.2cm}
{\fontsize{24}{30}\selectfont \textbf{\color{brandRed}TRƯỜNG TRUNG HỌC CƠ SỞ -- NĂM HỌC 2026 -- 2027}}\\[10pt]
{\fontsize{34}{42}\selectfont \textbf{\color{brandBlue}ĐỀ CƯƠNG ÔN TẬP GIỮA KÌ I -- MÔN KHTN 8}}\\[8pt]
{\fontsize{28}{36}\selectfont \textbf{\color{brandGreen}MÔN: KHOA HỌC TỰ NHIÊN 8 (K8)}}\\[8pt]
{\fontsize{22}{28}\selectfont \textit{(Phản ứng hóa học -- Định luật bảo toàn khối lượng -- Dung dịch \& Nồng độ)}}\\[16pt]
\rule{16cm}{1.5pt}\\[20pt]
{\fontsize{26}{34}\selectfont \textbf{Giáo viên hướng dẫn: Thầy TRẦN VĂN THẠNH -- 0777.470.803}}\\[8pt]
{\fontsize{22}{28}\selectfont \textit{Cơ sở 1: 6/15 Nguyễn Hoàng, Kim Long | Cơ sở 2: 24 Đặng Thái Thân, TP Huế}}\\[16pt]
{\fontsize{24}{30}\selectfont \textbf{CẤU TRÚC: TRẮC NGHIỆM ABCD -- ĐÚNG/SAI -- TRẢ LỜI NGẮN -- TỰ LUẬN}}
\end{center}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 1: TRẮC NGHIỆM CÂU 1 & CÂU 2
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 1: TRẮC NGHIỆM CÂU 1 & CÂU 2
% ==============================================================================
\fontsize{26}{35}\selectfont
\textbf{Câu 1.} Trong học tập môn Khoa học tự nhiên, việc quan sát trực tiếp các sự vật, hiện tượng để thu thập thông tin thuộc phương pháp nào sau đây?
\vspace{4pt}

\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{A.} Phương pháp quan sát. & \textbf{B.} Phương pháp suy luận.\\[6pt]
\textbf{C.} Phương pháp mô hình hóa. & \textbf{D.} Phương pháp tính toán.
\end{tabularx}

\vspace{16pt}

\textbf{Câu 2.} Khi tiến hành một thí nghiệm khoa học, việc đầu tiên cần làm để bảo đảm an toàn và thực hiện đúng quy trình là gì?
\vspace{4pt}

\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{A.} Thực hiện thí nghiệm ngay lập tức. & \textbf{B.} Đọc và tìm hiểu hướng dẫn thí nghiệm.\\[6pt]
\textbf{C.} Thu dọn dụng cụ sau thí nghiệm. & \textbf{D.} Ghi kết quả trước khi thí nghiệm.
\end{tabularx}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 2: TRẮC NGHIỆM CÂU 3 & CÂU 4
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 2: TRẮC NGHIỆM CÂU 3 & CÂU 4
% ==============================================================================
\fontsize{26}{35}\selectfont
\textbf{Câu 3.} Nồng độ phần trăm của dung dịch cho biết số gam chất tan có trong
\vspace{4pt}

\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{A.} $100\text{ gam}$ dung dịch. & \textbf{B.} $100\text{ gam}$ dung môi.\\[6pt]
\textbf{C.} $1\text{ lít}$ dung dịch. & \textbf{D.} $1\text{ lít}$ dung môi.
\end{tabularx}

\vspace{16pt}

\textbf{Câu 4.} Phương trình hóa học dùng để biểu diễn
\vspace{4pt}

\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{A.} thành phần các chất trong hỗn hợp. & \textbf{B.} phản ứng hóa học bằng công thức hóa học.\\[6pt]
\textbf{C.} sự thay đổi trạng thái của chất. & \textbf{D.} khối lượng riêng của các chất.
\end{tabularx}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 3: TRẮC NGHIỆM CÂU 5
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 3: TRẮC NGHIỆM CÂU 5 (KHÔNG CÓ ĐÁP ÁN)
% ==============================================================================
\fontsize{27}{36}\selectfont
\textbf{Câu 5.} Theo định luật bảo toàn khối lượng, trong một phản ứng hóa học, tổng khối lượng các chất tham gia phản ứng
\vspace{8pt}

\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{A.} lớn hơn tổng khối lượng sản phẩm. & \textbf{B.} nhỏ hơn tổng khối lượng sản phẩm.\\[8pt]
\textbf{C.} bằng tổng khối lượng các sản phẩm. & \textbf{D.} luôn bằng không.
\end{tabularx}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 4: ĐÚNG - SAI CÂU 6 (KHÔNG CÓ ĐÁP ÁN)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 4: ĐÚNG - SAI CÂU 6 (KHÔNG CÓ ĐÁP ÁN)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}TRẮC NGHIỆM ĐÚNG SAI -- CÂU 6}}
\vspace{8pt}

\fontsize{22}{30}\selectfont
\textbf{Câu 6.} Xác định mỗi phát biểu sau là đúng hay sai:

\vspace{12pt}
\textbf{a)} Nồng độ mol của dung dịch cho biết số mol chất tan có trong một lít dung dịch.\hfill $[\quad]$\\[10pt]
\textbf{b)} Đơn vị thường dùng của nồng độ mol là $\text{mol/L}$.\hfill $[\quad]$\\[10pt]
\textbf{c)} Khối lượng dung dịch bằng khối lượng chất tan cộng khối lượng dung môi.\hfill $[\quad]$\\[10pt]
\textbf{d)} Định luật bảo toàn khối lượng chỉ đúng với phản ứng không tạo ra chất khí.\hfill $[\quad]$
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 5: ĐÚNG - SAI CÂU 7 (KHÔNG CÓ ĐÁP ÁN)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 5: ĐÚNG - SAI CÂU 7 (KHÔNG CÓ ĐÁP ÁN)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}TRẮC NGHIỆM ĐÚNG SAI -- CÂU 7}}
\vspace{8pt}

\fontsize{22}{30}\selectfont
\textbf{Câu 7.} Cho các phát biểu sau về nồng độ mol:

\vspace{12pt}
\textbf{a)} Nồng độ mol của dung dịch cho biết số mol chất tan có trong $1000\text{ lít}$ dung dịch.\hfill $[\quad]$\\[10pt]
\textbf{b)} Đơn vị thường dùng của nồng độ mol là mol/lít ($\text{mol/L}$).\hfill $[\quad]$\\[10pt]
\textbf{c)} Hòa tan $0,5\text{ mol } \ce{NaCl}$ vào nước, thu được $1\text{ L dd } \ce{NaCl}$ thì nồng độ mol của dung dịch là $0,5\text{ M}$.\hfill $[\quad]$
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 6: ĐÚNG - SAI CÂU 8 (KHÔNG CÓ ĐÁP ÁN)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 6: ĐÚNG - SAI CÂU 8 (KHÔNG CÓ ĐÁP ÁN)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}TRẮC NGHIỆM ĐÚNG SAI -- CÂU 8}}
\vspace{8pt}

\fontsize{22}{30}\selectfont
\textbf{Câu 8.} Xác định mỗi phát biểu sau là đúng hay sai:

\vspace{12pt}
\textbf{a)} Trong phản ứng hóa học, tổng khối lượng các chất tham gia phản ứng bằng tổng khối lượng các sản phẩm.\hfill $[\quad]$\\[10pt]
\textbf{b)} Định luật bảo toàn khối lượng được áp dụng cho các phản ứng hóa học.\hfill $[\quad]$\\[10pt]
\textbf{c)} Khi đốt cháy hoàn toàn $12\text{ g magnesium}$ trong oxygen, thu được $20\text{ g magnesium oxide}$ thì khối lượng oxygen đã phản ứng là $8\text{ g}$.\hfill $[\quad]$
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 7: ĐÚNG - SAI CÂU 9 (KHÔNG CÓ ĐÁP ÁN)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 7: ĐÚNG - SAI CÂU 9 (KHÔNG CÓ ĐÁP ÁN)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}TRẮC NGHIỆM ĐÚNG SAI -- CÂU 9}}
\vspace{6pt}

\fontsize{21}{29}\selectfont
\textbf{Câu 9.} Xác định mỗi phát biểu sau là đúng hay sai:

\vspace{10pt}
\textbf{a)} Khối lượng dung dịch bằng tổng khối lượng của chất tan và dung môi.\hfill $[\quad]$\\[8pt]
\textbf{b)} Khi cho thêm chất tan vào 1 dung dịch chưa bão hòa, nồng độ phần trăm của dung dịch sẽ giảm.\hfill $[\quad]$\\[8pt]
\textbf{c)} Khi pha thêm $100\text{ mL}$ nước cất vào $100\text{ mL dd } \ce{NaOH } 1\text{M}$, nồng độ mol của dung dịch $\ce{NaOH}$ thu được giảm đi một nửa, còn $0,5\text{M}$.\hfill $[\quad]$\\[8pt]
\textbf{d)} Cho $7\text{ g iron}$ phản ứng vừa đủ với $4\text{ g sulfur}$, thu được $11\text{ g iron sulfide}$.\hfill $[\quad]$
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 8: TRẢ LỜI NGẮN CÂU 10 & CÂU 11 (KHÔNG CÓ ĐÁP ÁN)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 8: TRẢ LỜI NGẮN CÂU 10 & CÂU 11 (KHÔNG CÓ ĐÁP ÁN)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}PHẦN TRẢ LỜI NGẮN -- CÂU 10 \& CÂU 11}}
\vspace{10pt}

\fontsize{22}{30}\selectfont
\textbf{Câu 10.} Viết công thức tính nồng độ phần trăm của dung dịch.
\vspace{6pt}

\textbf{Trả lời:} $\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots$

\vspace{18pt}
\textbf{Câu 11.} Hòa tan $15\text{ g}$ muối vào $85\text{ g}$ nước. Tính nồng độ phần trăm của dung dịch thu được.
\vspace{6pt}

\textbf{Trả lời:} $\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots$
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 9: TRẢ LỜI NGẮN CÂU 12 & CÂU 13 (KHÔNG CÓ ĐÁP ÁN)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 9: TRẢ LỜI NGẮN CÂU 12 & CÂU 13 (KHÔNG CÓ ĐÁP ÁN)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}PHẦN TRẢ LỜI NGẮN -- CÂU 12 \& CÂU 13}}
\vspace{10pt}

\fontsize{21}{29}\selectfont
\textbf{Câu 12.} Đốt cháy hoàn toàn $6\text{ g carbon}$ trong $16\text{ g oxygen}$, thu được carbon dioxide. Tính khối lượng $\ce{CO2}$ tạo thành theo ĐLBTKL.
\vspace{6pt}

\textbf{Trả lời:} $\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots$

\vspace{16pt}
\textbf{Câu 13.} Cần bao nhiêu mol $\ce{NaOH}$ để pha chế $500\text{ mL}$ dung dịch $\ce{NaOH } 2\text{M}$?
\vspace{6pt}

\textbf{Trả lời:} $\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots$
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 10: TRẢ LỜI NGẮN CÂU 14 (KHÔNG CÓ ĐÁP ÁN)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 10: TRẢ LỜI NGẮN CÂU 14 (KHÔNG CÓ ĐÁP ÁN)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}PHẦN TRẢ LỜI NGẮN -- CÂU 14: BÀI TOÁN PHA LOÃNG}}
\vspace{10pt}

\fontsize{23}{32}\selectfont
\textbf{Đề bài:} Cần thêm bao nhiêu gam nước vào $100\text{ g}$ dung dịch $\ce{NaCl } 20\%$ để thu được dung dịch $\ce{NaCl}$ nồng độ $5\%$?

\vspace{16pt}
\textbf{Trả lời:} $\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots$

\vspace*{2.5cm}
\begin{center}
{\fontsize{20}{26}\selectfont \textit{\color{gray}(Khoảng trống trình chiếu tương tác / Giáo viên và học sinh chữa bài trực tiếp trên bảng)}}
\end{center}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 11: TỰ LUẬN CÂU 1 (KHÔNG CÓ ĐÁP ÁN)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 11: TỰ LUẬN CÂU 1 (KHÔNG CÓ ĐÁP ÁN)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}PHẦN TỰ LUẬN -- CÂU 1 (1,0 ĐIỂM)}}
\vspace{10pt}

\fontsize{22}{30}\selectfont
\textbf{Đề bài:}
\begin{enumerate}
  \setlength{\itemsep}{10pt}
  \item[1)] Viết các công thức tính toán cơ bản: Nồng độ phần trăm ($C\%$), Nồng độ mol ($C_M$), Khối lượng dung dịch ($m_{\text{dd}}$) và Số mol ($n$).
  \item[2)] Phát biểu nội dung Định luật bảo toàn khối lượng. Giải thích vì sao trong một phản ứng hóa học, tổng khối lượng các chất luôn được bảo toàn?
\end{enumerate}

\vspace*{2cm}
\begin{center}
{\fontsize{20}{26}\selectfont \textit{\color{gray}(Khoảng trống trình chiếu tương tác / Giáo viên và học sinh chữa bài trực tiếp trên bảng)}}
\end{center}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 12: TỰ LUẬN CÂU 2 (KHÔNG CÓ ĐÁP ÁN)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 12: TỰ LUẬN CÂU 2 (KHÔNG CÓ ĐÁP ÁN)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}PHẦN TỰ LUẬN -- CÂU 2 (1,0 ĐIỂM)}}
\vspace{10pt}

\fontsize{23}{32}\selectfont
\textbf{Đề bài:}
\begin{enumerate}
  \setlength{\itemsep}{10pt}
  \item[1)] Nêu khái niệm Phương trình hóa học (PTHH). Cho ví dụ minh họa.
  \item[2)] Nêu ý nghĩa của Phương trình hóa học trong học tập và nghiên cứu hóa học.
\end{enumerate}

\vspace*{2.5cm}
\begin{center}
{\fontsize{20}{26}\selectfont \textit{\color{gray}(Khoảng trống trình chiếu tương tác / Giáo viên và học sinh chữa bài trực tiếp trên bảng)}}
\end{center}

\end{document}
""")

    with open(tex_path, "w", encoding="utf-8") as f:
        f.write("".join(tex_lines))
    print(f"[OK] Đã ghi file LaTeX: {tex_path}")

    # Biên dịch pdflatex 2 lần để ổn định layout
    cmd = ["pdflatex", "-interaction=nonstopmode", "-output-directory", SCRATCH_DIR, tex_path]
    res1 = subprocess.run(cmd, capture_output=True, text=True)
    res2 = subprocess.run(cmd, capture_output=True, text=True)

    pdf_scratch = os.path.join(SCRATCH_DIR, pdf_filename)
    if os.path.exists(pdf_scratch):
        shutil.copy2(pdf_scratch, pdf_dest)
        print(f"[OK] Đã biên dịch và copy PDF thành công: {pdf_dest} ({os.path.getsize(pdf_dest):,} bytes)")
        return pdf_dest
    else:
        print("[ERROR] Không tìm thấy file PDF được tạo. Log:")
        print(res2.stdout[-1500:])
        return None

# ==============================================================================
# HÀM MAIN
# ==============================================================================

def main():
    print("=" * 75)
    print("BẮT ĐẦU TẠO 3 FILE CHO ĐỀ CƯƠNG KHTN 8 (K8) TẠI NEW FOLDER (2)")
    print("1. FILE HỌC SINH (.docx)")
    print("2. FILE GIÁO VIÊN (.docx có lời giải chi tiết)")
    print("3. FILE DẠY TRÊN MÀN HÌNH (Slide PDF 16:9 widescreen cho Tivi 65 inch)")
    print("=" * 75)

    f_hs = generate_word_document(mode="student")
    f_gv = generate_word_document(mode="teacher")
    f_screen_docx = generate_word_document(mode="screen")
    f_screen_pdf = generate_tivi_presentation_pdf()

    all_files = [f_hs, f_gv, f_screen_docx, f_screen_pdf]

    # Đồng bộ sang các thư mục OneDrive
    print("\n[+] Đồng bộ các file sang các thư mục OneDrive:")
    for od in TARGET_ONEDRIVE_DIRS:
        if os.path.exists(od):
            for src_path in all_files:
                if src_path and os.path.exists(src_path):
                    dest = os.path.join(od, os.path.basename(src_path))
                    try:
                        shutil.copy2(src_path, dest)
                        print(f"    -> Đã copy {os.path.basename(src_path)} vào: {od}")
                    except Exception as e:
                        print(f"    [!] Cảnh báo file đang mở hoặc bận tại {od}: {e}")

    print("\n" + "=" * 75)
    print("HOÀN THÀNH TẤT CẢ 3 FILE THEO ĐÚNG CHUẨN FORM ĐỀ KIỂM TRA!")
    print("=" * 75)

if __name__ == "__main__":
    main()
