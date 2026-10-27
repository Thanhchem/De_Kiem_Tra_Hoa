# -*- coding: utf-8 -*-
"""
Script biên soạn trọn bộ 3 file:
1. Đề cương Học sinh (.docx)
2. Đề cương Giáo viên có lời giải chi tiết (.docx)
3. Slide Dạy trên màn hình Tivi 65 inch tỉ lệ 16:9 (.pdf và .docx) - KHÔNG CÓ ĐÁP ÁN

ĐỀ CƯƠNG ÔN TẬP KIỂM TRA GIỮA HỌC KỲ I - KHTN 7 (NĂM HỌC 2026 - 2027)
TRƯỜNG THCS CHU VĂN AN
(Gồm 3 phân môn: Sinh học, Hóa học, Vật lí)
"""

import sys
if sys.stdout:
    sys.stdout.reconfigure(encoding='utf-8')
if sys.stderr:
    sys.stderr.reconfigure(encoding='utf-8')

import os
import shutil
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

# ==============================================================================
# HẰNG SỐ MÀU SẮC & ĐỊNH DẠNG THEO GEMINI.md (FORM ĐỀ KIỂM TRA)
# ==============================================================================
COLOR_PRIMARY_HEX = "D32F2F"       # Đỏ san hô (Coral Red)
COLOR_SECONDARY_HEX = "8BC390"     # Xanh xô thơm (Sage Green)
COLOR_TEXT_HEX = "1A1A1A"          # Đen than nhẹ (Off-black)
COLOR_SOLUTION_TITLE_HEX = "1565C0"# Xanh dương đậm cho tiêu đề lời giải
COLOR_GRAY_HEX = "757575"

COLOR_PRIMARY = RGBColor(0xD3, 0x2F, 0x2F)
COLOR_SECONDARY = RGBColor(0x8B, 0xC3, 0x90)
COLOR_TEXT = RGBColor(0x1A, 0x1A, 0x1A)
COLOR_SOLUTION_TITLE = RGBColor(0x15, 0x65, 0xC0)
COLOR_GRAY = RGBColor(0x75, 0x75, 0x75)

FONT_NAME = "Times New Roman"

BASE_DIR = r"c:\Antigravity_Thanh"
SAN_PHAM_DIR = os.path.join(BASE_DIR, "San_Pham")
SCRATCH_DIR = os.path.join(BASE_DIR, "scratch")
ONEDRIVE_DIR1 = r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn7\KHTN7-IN-DEKIEMTRADINHKI"
ONEDRIVE_DIR2 = r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn7"

os.makedirs(SAN_PHAM_DIR, exist_ok=True)
os.makedirs(SCRATCH_DIR, exist_ok=True)

# ==============================================================================
# DỮ LIỆU ĐỀ CƯƠNG KHTN 7 THCS CHU VĂN AN
# ==============================================================================

# ------------------------------------------------------------------------------
# PHÂN MÔN SINH HỌC
# ------------------------------------------------------------------------------
BIO_THEORY = [
    {
        "lesson": "Bài 22. Vai trò của trao đổi chất và chuyển hóa năng lượng ở sinh vật",
        "sections": [
            {
                "title": "1. Khái niệm trao đổi chất và chuyển hóa năng lượng ở sinh vật",
                "content": (
                    "• Trao đổi chất: Là tập hợp các biến đổi hóa học trong tế bào của cơ thể sinh vật và sự trao đổi các chất giữa cơ thể với môi trường nhằm đảm bảo duy trì sự sống.\n"
                    "• Chuyển hóa năng lượng: Là sự biến đổi năng lượng từ dạng này sang dạng khác trong các hoạt động sống (ví dụ: chuyển hóa từ quang năng thành hóa năng trong quang hợp; chuyển hóa từ hóa năng thành nhiệt năng, cơ năng, điện năng,...)."
                )
            },
            {
                "title": "2. Vai trò của quá trình trao đổi chất và chuyển hóa năng lượng",
                "content": (
                    "• Cung cấp nguyên liệu: Tạo ra các chất hữu cơ cần thiết cấu tạo nên tế bào, mô, cơ quan và cơ thể sinh vật.\n"
                    "• Cung cấp năng lượng: Giải phóng năng lượng tích lũy (dưới dạng ATP) phục vụ mọi hoạt động sống như sinh trưởng, phát triển, sinh sản, cảm ứng và vận động.\n"
                    "• Đào thải chất thải: Loại bỏ các chất cặn bã, độc hại sinh ra từ quá trình trao đổi chất ra khỏi cơ thể để duy trì cân bằng nội môi."
                )
            }
        ]
    },
    {
        "lesson": "Bài 23. Quang hợp ở thực vật",
        "sections": [
            {
                "title": "1. Khái niệm quang hợp",
                "content": (
                    "• Quang hợp là quá trình lá cây sử dụng năng lượng ánh sáng mặt trời đã được diệp lục hấp thu để tổng hợp chất hữu cơ (chủ yếu là glucose, tinh bột) từ nước (H2O) và khí carbon dioxide (CO2), đồng thời giải phóng khí oxygen (O2)."
                )
            },
            {
                "title": "2. Phương trình quang hợp",
                "content": (
                    "• Phương trình chữ:\n"
                    "  Nước + Khí carbon dioxide --(Ánh sáng / Diệp lục)--> Glucose + Khí oxygen\n"
                    "• Phương trình hóa học tổng quát:\n"
                    "  6CO2 + 6H2O --(Ánh sáng / Diệp lục)--> C6H12O6 + 6O2"
                )
            },
            {
                "title": "3. Một số yếu tố ảnh hưởng đến quang hợp",
                "content": (
                    "• Ánh sáng: Cường độ ánh sáng và thành phần quang phổ ảnh hưởng trực tiếp đến hiệu quả quang hợp.\n"
                    "• Khí carbon dioxide (CO2): Nồng độ CO2 ngoài môi trường từ 0,03% đến khoảng 0,2% làm tăng cường độ quang hợp.\n"
                    "• Nước: Vừa là nguyên liệu trực tiếp của quang hợp, vừa tham gia điều tiết đóng mở khí khổng để trao đổi khí và thoát hơi nước.\n"
                    "• Nhiệt độ: Nhiệt độ tối ưu cho quang hợp thường dao động từ 25 °C đến 35 °C.\n"
                    "• Các nguyên tố khoáng: Tham gia cấu tạo phân tử diệp lục (như N, Mg) và các enzyme xúc tác."
                )
            },
            {
                "title": "4. Vận dụng những hiểu biết về quang hợp ở thực vật",
                "content": (
                    "• Trồng cây với mật độ hợp lý, tỉa cành, tạo tán để cây đón nhận tối đa ánh sáng mặt trời.\n"
                    "• Cung cấp đầy đủ nước, bón phân cân đối, đúng thời điểm cho cây trồng.\n"
                    "• Trồng cây trong nhà kính kết hợp điều khiển nhiệt độ, độ ẩm và chiếu sáng nhân tạo (đèn LED) để tăng năng suất và thu hoạch trái vụ.\n"
                    "• Tích cực trồng và bảo vệ cây xanh trong trường học, khu dân cư để điều hòa khí hậu và làm sạch môi trường."
                )
            }
        ]
    },
    {
        "lesson": "Bài 25. Hô hấp tế bào",
        "sections": [
            {
                "title": "1. Khái niệm hô hấp tế bào",
                "content": (
                    "• Hô hấp tế bào là quá trình phân giải hợp chất hữu cơ (chủ yếu là glucose) diễn ra trong ti thể, giải phóng năng lượng tích lũy trong các liên kết hóa học để tạo thành ATP cung cấp cho mọi hoạt động sống của cơ thể, đồng thời tạo ra sản phẩm phụ là khí carbon dioxide và nước."
                )
            },
            {
                "title": "2. Phương trình hô hấp tế bào",
                "content": (
                    "• Phương trình chữ:\n"
                    "  Glucose + Khí oxygen ----> Khí carbon dioxide + Nước + Năng lượng (ATP + Nhiệt)\n"
                    "• Phương trình hóa học tổng quát:\n"
                    "  C6H12O6 + 6O2 ----> 6CO2 + 6H2O + Năng lượng (ATP + Q)"
                )
            },
            {
                "title": "3. Một số yếu tố ảnh hưởng đến hô hấp tế bào",
                "content": (
                    "• Hàm lượng nước: Nước là môi trường và tham gia phản ứng thủy phân; hàm lượng nước tăng thì hô hấp tế bào tăng.\n"
                    "• Nhiệt độ: Ảnh hưởng trực tiếp đến hoạt tính enzyme hô hấp. Nhiệt độ thích hợp là 30 °C - 35 °C.\n"
                    "• Nồng độ khí oxygen: Oxygen là chất tham gia trực tiếp vào hô hấp hiếu khí. Nồng độ O2 giảm dưới 5% làm hô hấp tế bào giảm mạnh.\n"
                    "• Nồng độ khí carbon dioxide: CO2 là sản phẩm; nồng độ CO2 cao (trên 3% - 5%) sẽ ức chế quá trình hô hấp tế bào."
                )
            },
            {
                "title": "4. Vận dụng những hiểu biết về hô hấp tế bào",
                "content": (
                    "• Trong bảo quản lương thực, thực phẩm:\n"
                    "  - Bảo quản khô: Phơi hoặc sấy khô hạt lúa, ngô, đậu nhằm giảm hàm lượng nước, hạn chế hô hấp.\n"
                    "  - Bảo quản lạnh: Giữ rau, củ, quả trong tủ lạnh/kho lạnh để hạ nhiệt độ làm giảm hoạt tính enzyme hô hấp.\n"
                    "  - Bảo quản trong điều kiện nồng độ CO2 cao hoặc hút chân không: Làm giảm lượng O2 tiếp xúc, ngăn ngừa hô hấp và vi sinh vật gây thối rữa.\n"
                    "• Trong chăm sóc sức khỏe con người:\n"
                    "  - Vận động và tập luyện thể dục thể thao hợp lý để thúc đẩy hô hấp tế bào tạo năng lượng.\n"
                    "  - Giữ môi trường sống trong lành, thông thoáng, không gian phòng ngủ cần mở hé cửa để tránh tích tụ CO2."
                )
            }
        ]
    }
]

# ------------------------------------------------------------------------------
# PHÂN MÔN HÓA HỌC
# ------------------------------------------------------------------------------
CHEM_THEORY = [
    {
        "lesson": "Bài 2. Nguyên tử",
        "points": [
            "Nguyên tử là hạt vô cùng nhỏ, trung hòa về điện.",
            "Cấu tạo nguyên tử gồm 2 phần chính:\n"
            "  + Hạt nhân: Nằm ở tâm nguyên tử, chứa proton (kí hiệu là p, mang điện tích dương +1) và neutron (kí hiệu là n, không mang điện).\n"
            "  + Vỏ nguyên tử: Chứa các electron (kí hiệu là e, mang điện tích âm -1) chuyển động rất nhanh xung quanh hạt nhân.",
            "Trong nguyên tử trung hòa về điện: Số hạt proton luôn bằng số hạt electron (số p = số e).",
            "Khối lượng nguyên tử tập trung chủ yếu ở hạt nhân (vì khối lượng của electron rất nhỏ, không đáng kể so với proton và neutron).",
            "Các electron trong vỏ nguyên tử được sắp xếp thành từng lớp electron quanh hạt nhân từ trong ra ngoài (lớp 1 tối đa 2e, lớp 2 tối đa 8e, lớp 3 tối đa 8e/18e,...)."
        ]
    },
    {
        "lesson": "Bài 3. Nguyên tố hóa học",
        "points": [
            "Nguyên tố hóa học là tập hợp những nguyên tử cùng loại, có cùng số proton trong hạt nhân.",
            "Số proton trong hạt nhân là số đặc trưng của một nguyên tố hóa học. Các nguyên tử thuộc cùng một nguyên tố hóa học thì đều có tính chất hóa học tương tự nhau.",
            "Mỗi nguyên tố hóa học được biểu diễn bằng một hoặc hai chữ cái, trong đó chữ cái đầu viết in hoa, chữ cái thứ hai (nếu có) viết thường. Đó là kí hiệu hóa học (ví dụ: H, C, N, O, Na, Fe, Cu,...)."
        ]
    }
]

CHEM_SHORT_QUESTIONS = [
    {
        "id": 1,
        "content": "Nguyên tử có 10 hạt electron trung hòa về điện thì số hạt proton là bao nhiêu?",
        "answer": "10 hạt proton",
        "solution": "Do nguyên tử luôn trung hòa về điện nên số proton bằng số electron: số p = số e = *10* hạt."
    },
    {
        "id": 2,
        "content": "Kí hiệu của ba loại hạt cấu tạo nên nguyên tử là?",
        "answer": "Proton kí hiệu là p, neutron kí hiệu là n, electron kí hiệu là e",
        "solution": "Ba loại hạt cấu tạo nên nguyên tử gồm: proton kí hiệu là *p*, neutron kí hiệu là *n*, và electron kí hiệu là *e*."
    },
    {
        "id": 3,
        "content": "Một nguyên tử có 11 proton. Hỏi nguyên tử trung hòa về điện có bao nhiêu electron?",
        "answer": "11 electron",
        "solution": "Trong nguyên tử trung hòa về điện, số electron luôn bằng số proton: số e = số p = *11* electron (đây chính là nguyên tử Sodium - Na)."
    },
    {
        "id": 4,
        "content": "Nguyên tử có 8 proton và 8 electron mang điện tích tổng cộng bằng bao nhiêu?",
        "answer": "0 (trung hòa về điện)",
        "solution": "Mỗi hạt proton mang 1 đơn vị điện tích dương (+1), mỗi electron mang 1 đơn vị điện tích âm (-1). Vì nguyên tử có 8 proton mang điện tích +8 và 8 electron mang điện tích -8, nên điện tích tổng cộng là (+8) + (-8) = *0* (nguyên tử trung hòa về điện)."
    }
]

CHEM_ESSAY_QUESTIONS = [
    {
        "id": 1,
        "title": "Bài 1",
        "content": (
            "Một nguyên tử có 12 proton, 12 neutron và 12 electron.\n"
            "a) Xác định số hạt proton, neutron và electron.\n"
            "b) Tính tổng số hạt trong nguyên tử.\n"
            "c) Giải thích vì sao nguyên tử trung hòa về điện."
        ),
        "solution": (
            "a) Số hạt proton là *12*, số hạt neutron là *12*, số hạt electron là *12*.\n"
            "b) Tổng số hạt cấu tạo nên nguyên tử là: N = p + n + e = 12 + 12 + 12 = *36* (hạt).\n"
            "c) Giải thích: Hạt nhân có 12 proton mang tổng điện tích dương là +12, lớp vỏ có 12 electron mang tổng điện tích âm là -12. Do tổng điện tích dương triệt tiêu hoàn toàn với tổng điện tích âm ((+12) + (-12) = 0) nên nguyên tử *trung hòa về điện*."
        )
    },
    {
        "id": 2,
        "title": "Bài 2",
        "content": (
            "Nguyên tử X có tổng số hạt proton, neutron và electron là 23. Biết số proton là 7 và nguyên tử trung hòa về điện.\n"
            "a) Tính số electron.\n"
            "b) Tính số neutron."
        ),
        "solution": (
            "a) Vì nguyên tử X trung hòa về điện nên số hạt electron bằng số hạt proton: e = p = *7* (hạt).\n"
            "b) Tổng số hạt trong nguyên tử X là p + n + e = 23.\n"
            "   Thay số: 7 + n + 7 = 23  ==>  14 + n = 23  ==>  n = 23 - 14 = *9* (hạt).\n"
            "   Vậy nguyên tử X có *9* neutron (nguyên tử X là đồng vị của Nitrogen)."
        )
    },
    {
        "id": 3,
        "title": "Bài 3",
        "content": "Phân biệt proton, neutron và electron về vị trí, điện tích.",
        "solution": (
            "Phân biệt ba loại hạt cấu tạo nên nguyên tử:\n"
            "• Hạt proton (p): Nằm trong hạt nhân nguyên tử, mang điện tích dương (+1).\n"
            "• Hạt neutron (n): Nằm trong hạt nhân nguyên tử, không mang điện (điện tích bằng 0).\n"
            "• Hạt electron (e): Nằm ở lớp vỏ nguyên tử, chuyển động rất nhanh xung quanh hạt nhân, mang điện tích âm (-1)."
        )
    },
    {
        "id": 4,
        "title": "Bài 4",
        "content": (
            "Một nguyên tử có 17 proton và 18 neutron.\n"
            "a) Nguyên tử này có bao nhiêu electron nếu trung hòa về điện?\n"
            "b) Có bao nhiêu hạt trong hạt nhân?\n"
            "c) Tổng số hạt của nguyên tử là bao nhiêu?"
        ),
        "solution": (
            "a) Vì nguyên tử trung hòa về điện nên: số electron = số proton = *17* (hạt) (đây là nguyên tử Chlorine - Cl).\n"
            "b) Số hạt nằm trong hạt nhân gồm proton và neutron:\n"
            "   Số hạt trong hạt nhân = p + n = 17 + 18 = *35* (hạt).\n"
            "c) Tổng số hạt cấu tạo nên nguyên tử:\n"
            "   Tổng số hạt = p + n + e = 17 + 18 + 17 = *52* (hạt)."
        )
    },
    {
        "id": 5,
        "title": "Bài 5",
        "content": "Giải thích sự khác nhau giữa nguyên tử và nguyên tố hóa học. Cho ví dụ minh họa.",
        "solution": (
            "Sự khác biệt cơ bản giữa nguyên tử và nguyên tố hóa học:\n"
            "• Nguyên tử là một đối tượng cụ thể: Là hạt vô cùng nhỏ, trung hòa về điện, gồm hạt nhân và vỏ electron, có khối lượng và kích thước xác định.\n"
            "• Nguyên tố hóa học là một khái niệm tập hợp: Là tập hợp gồm tất cả các nguyên tử có cùng số proton trong hạt nhân.\n"
            "• Ví dụ minh họa:\n"
            "  - Một nguyên tử Carbon có 6 proton, 6 neutron, 6 electron là một nguyên tử cụ thể.\n"
            "  - Tập hợp tất cả các nguyên tử có 6 hạt proton trong hạt nhân (dù có thể chứa 6 neutron, 7 neutron hay 8 neutron) đều thuộc về cùng một nguyên tố hóa học là *nguyên tố Carbon* (kí hiệu là C)."
        )
    }
]

# ------------------------------------------------------------------------------
# PHÂN MÔN VẬT LÝ
# ------------------------------------------------------------------------------
PHYS_THEORY_QUESTIONS = [
    {
        "id": 1,
        "content": "Nêu ý nghĩa của tốc độ.",
        "solution": "Tốc độ là đại lượng cho biết mức độ nhanh hay chậm của chuyển động. Nếu trong cùng một khoảng thời gian, vật nào đi được quãng đường dài hơn thì chuyển động nhanh hơn (tốc độ lớn hơn); ngược lại đi được quãng đường ngắn hơn thì chuyển động chậm hơn (tốc độ nhỏ hơn)."
    },
    {
        "id": 2,
        "content": "Nêu công thức tính tốc độ, giải thích các đại lượng trong công thức.",
        "solution": (
            "• Công thức tính tốc độ: v = s / t\n"
            "• Giải thích các đại lượng:\n"
            "  + v: Tốc độ chuyển động của vật (đơn vị: m/s hoặc km/h).\n"
            "  + s: Quãng đường vật đi được (đơn vị: m hoặc km).\n"
            "  + t: Thời gian đi hết quãng đường đó (đơn vị: s hoặc h)."
        )
    },
    {
        "id": 3,
        "content": "Liệt kê một số đơn vị tốc độ thường dùng.",
        "solution": (
            "• Đơn vị tốc độ trong hệ đo lường chính thức (SI) ở nước ta là *mét trên giây (m/s)*.\n"
            "• Đơn vị tốc độ thường dùng trong đời sống và giao thông là *kilômét trên giờ (km/h)*.\n"
            "• Ngoài ra còn có các đơn vị khác như: mét trên phút (m/min), xentimét trên giây (cm/s), kilômét trên giây (km/s), dặm trên giờ (mph)..."
        )
    },
    {
        "id": 4,
        "content": "Mô tả sơ lược cách đo tốc độ bằng đồng hồ bấm giây và cổng quang điện trong dụng cụ thực hành ở nhà trường.",
        "solution": (
            "• Bố trí dụng cụ: Lắp máng dẫn hướng nghiêng vừa phải. Bố trí hai cổng quang điện (cổng 1 và cổng 2) cách nhau một khoảng xác định s trên máng. Nối hai cổng quang điện với đồng hồ đo thời gian hiện số (chọn chế độ đo thời gian vật chắn qua 2 cổng Mode A<->B).\n"
            "• Tiến hành đo:\n"
            "  - Thả cho xe có gắn tấm chắn sáng chuyển động từ đỉnh máng xuống.\n"
            "  - Khi mép trước tấm chắn đi qua cổng quang điện 1, đồng hồ tự động kích hoạt đếm thời gian.\n"
            "  - Khi mép trước tấm chắn đi qua cổng quang điện 2, đồng hồ tự động ngắt và hiển thị khoảng thời gian t.\n"
            "• Tính tốc độ: Áp dụng công thức v = s / t."
        )
    },
    {
        "id": 5,
        "content": "Mô tả sơ lược cách đo tốc độ bằng thiết bị “bắn tốc độ” trong kiểm tra tốc độ các phương tiện giao thông.",
        "solution": (
            "• Cấu tạo: Thiết bị bắn tốc độ gồm có một máy phát và thu sóng vô tuyến (radar) hoặc chùm tia laser, kết hợp với một camera kỹ thuật số và bộ vi xử lý điện tử.\n"
            "• Hoạt động:\n"
            "  - Thiết bị phát tín hiệu (sóng radar hoặc tia laser) đến phương tiện giao thông đang di chuyển và thu nhận tín hiệu phản xạ trở lại.\n"
            "  - Bộ vi xử lý tính toán khoảng cách và thời gian chuyển động cực ngắn giữa hai lần phát tín hiệu (hoặc camera chụp 2 bức ảnh liên tiếp cách nhau khoảng thời gian Delta t đã biết để đo quãng đường Delta s xe đi được), từ đó tự động xác định tốc độ của phương tiện: v = Delta s / Delta t.\n"
            "  - Nếu tốc độ vượt quá mức tối đa cho phép trên đoạn đường, thiết bị sẽ tự động chụp ảnh phương tiện kèm biển số, ngày giờ và tốc độ đo được làm bằng chứng xử phạt."
        )
    }
]

PHYS_EXERCISES = [
    {
        "id": 1,
        "title": "Câu 1",
        "content": (
            "Tìm số thích hợp để điền vào chỗ trống:\n"
            "a) 15 m/s = ........ km/h.\n"
            "b) 90 km/h = ........ m/s.\n"
            "c) 250 cm/s = ........ m/s = ........ km/h.\n"
            "d) 72 km/h = ........ m/s = ........ cm/s."
        ),
        "solution": (
            "Áp dụng quy tắc đổi đơn vị tốc độ (1 m/s = 3,6 km/h; 1 km/h = 1/3,6 m/s):\n"
            "a) 15 m/s = 15 × 3,6 = *54* km/h.\n"
            "b) 90 km/h = 90 / 3,6 = *25* m/s.\n"
            "c) 250 cm/s = 250 / 100 = *2,5* m/s = 2,5 × 3,6 = *9* km/h.\n"
            "d) 72 km/h = 72 / 3,6 = *20* m/s = 20 × 100 = *2000* cm/s."
        )
    },
    {
        "id": 2,
        "title": "Câu 2",
        "content": (
            "Một học sinh đi xe đạp từ nhà đến trường, xuất phát lúc 6 giờ 30 phút và đến trường lúc 6 giờ 50 phút. "
            "Biết quãng đường từ nhà đến trường dài 4 km và học sinh đi liên tục không nghỉ. "
            "Tốc độ trung bình của học sinh trên quãng đường này là bao nhiêu km/h?"
        ),
        "solution": (
            "• Thời gian học sinh đi từ nhà đến trường là:\n"
            "  t = 6 giờ 50 phút - 6 giờ 30 phút = 20 phút = 20 / 60 h = *1/3 h* (≈ 0,333 h).\n"
            "• Tốc độ trung bình của học sinh là:\n"
            "  v = s / t = 4 / (1/3) = 4 × 3 = *12 (km/h)*.\n"
            "  (Hoặc v = 4 / 0,333 ≈ 12 km/h).\n"
            "• Đáp số: *12 km/h*."
        )
    },
    {
        "id": 3,
        "title": "Câu 3",
        "content": (
            "Tốc độ trung bình của máu chảy trong động mạch của người là 20 cm/s. "
            "Giả sử một hồng cầu chuyển động liên tục với tốc độ này trong 2 giờ. "
            "Quãng đường hồng cầu đi được trong thời gian đó là bao nhiêu mét?"
        ),
        "solution": (
            "• Đổi đơn vị:\n"
            "  Tốc độ: v = 20 cm/s = 0,2 m/s.\n"
            "  Thời gian: t = 2 giờ = 2 × 3600 s = 7200 s.\n"
            "• Quãng đường tế bào hồng cầu đi được là:\n"
            "  s = v × t = 0,2 × 7200 = *1440 (m)* = 1,44 km.\n"
            "• Đáp số: *1440 m*."
        )
    },
    {
        "id": 4,
        "title": "Câu 4",
        "content": (
            "Tuyến đường sắt từ Hà Nội đến Đà Nẵng dài khoảng 880 km. "
            "Một tàu hỏa di chuyển trên tuyến đường này với tốc độ trung bình 55 km/h. "
            "Hỏi tàu hỏa mất bao nhiêu giờ để đi từ Hà Nội đến Đà Nẵng?"
        ),
        "solution": (
            "• Thời gian tàu hỏa di chuyển từ Hà Nội đến Đà Nẵng là:\n"
            "  t = s / v = 880 / 55 = *16 (giờ)*.\n"
            "• Đáp số: *16 giờ*."
        )
    },
    {
        "id": 5,
        "title": "Câu 5",
        "content": (
            "Bạn An đi xe đạp điện từ nhà đến trường, trên đoạn đường đầu dài 3 km An đi hết 8 phút. "
            "Sau đó, An đi tiếp một đoạn đường 2,2 km trong 5 phút nữa thì đến trường.\n"
            "a) Tính tốc độ của An ứng với từng đoạn đường theo đơn vị km/h.\n"
            "b) Trên đoạn đường nào An đi nhanh hơn?\n"
            "c) Tính tốc độ trung bình của An trên cả đoạn đường từ nhà đến trường theo đơn vị km/h."
        ),
        "solution": (
            "a) Tính tốc độ trên từng đoạn đường:\n"
            "  - Đoạn đường đầu: s1 = 3 km, t1 = 8 phút = 8 / 60 h = 2/15 h.\n"
            "    v1 = s1 / t1 = 3 / (2/15) = 3 × 15 / 2 = *22,5 (km/h)*.\n"
            "  - Đoạn đường sau: s2 = 2,2 km, t2 = 5 phút = 5 / 60 h = 1/12 h.\n"
            "    v2 = s2 / t2 = 2,2 / (1/12) = 2,2 × 12 = *26,4 (km/h)*.\n"
            "b) So sánh tốc độ: Vì v2 (26,4 km/h) > v1 (22,5 km/h) nên trên đoạn đường sau bạn An *đi nhanh hơn* đoạn đường đầu.\n"
            "c) Tốc độ trung bình trên cả đoạn đường từ nhà đến trường:\n"
            "  - Tổng quãng đường: s = s1 + s2 = 3 + 2,2 = 5,2 (km).\n"
            "  - Tổng thời gian: t = t1 + t2 = 8 + 5 = 13 phút = 13 / 60 (h).\n"
            "  - Tốc độ trung bình: v_tb = s / t = 5,2 / (13/60) = (5,2 × 60) / 13 = *24 (km/h)*.\n"
            "• Đáp số: a) v1 = 22,5 km/h, v2 = 26,4 km/h; b) Đoạn sau nhanh hơn; c) v_tb = 24 km/h."
        )
    },
    {
        "id": 6,
        "title": "Câu 6",
        "content": (
            "Trên một đoạn đường, cảnh sát giao thông sử dụng thiết bị bắn tốc độ có camera để đo tốc độ của các phương tiện. "
            "Camera ghi nhận một ô tô đi qua hai vạch mốc cách nhau 10 m trong thời gian 0,5 s.\n"
            "a) Tính tốc độ của ô tô theo đơn vị km/h.\n"
            "b) Biết tốc độ tối đa cho phép trên đoạn đường này là 60 km/h. Hỏi ô tô có vi phạm quy định về tốc độ không? Vì sao?"
        ),
        "solution": (
            "a) Tốc độ chuyển động của ô tô là:\n"
            "  - Theo đơn vị m/s: v = s / t = 10 / 0,5 = *20 (m/s)*.\n"
            "  - Đổi sang đơn vị km/h: v = 20 × 3,6 = *72 (km/h)*.\n"
            "b) So sánh và kết luận:\n"
            "  Vì tốc độ của ô tô là 72 km/h, lớn hơn tốc độ tối đa cho phép trên đoạn đường là 60 km/h (72 km/h > 60 km/h), nên ô tô *CÓ VI PHẠM* quy định về an toàn giao thông đường bộ (vượt quá tốc độ tối đa cho phép 12 km/h)."
        )
    }
]

# ==============================================================================
# HÀM BỔ TRỢ ĐỊNH DẠNG WORD (.docx) CHUẨN GEMINI.md
# ==============================================================================

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders_horizontal_only(table, border_color_hex="8BC390", border_sz="6"):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color_hex}"/>\n'
        f'  <w:bottom w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color_hex}"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:insideH w:val="none"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.59)    # 1.5 cm
        s.bottom_margin = Inches(0.59) # 1.5 cm
        s.left_margin = Inches(0.59)   # 1.5 cm
        s.right_margin = Inches(0.59)  # 1.5 cm
        s.header_distance = Inches(0.24) # 0.6 cm
        s.footer_distance = Inches(0.24) # 0.6 cm

        # Header Table
        header = s.header
        header.is_linked_to_previous = False
        tbl_hdr = header.add_table(rows=1, cols=2, width=Inches(7.08))
        tbl_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders_horizontal_only(tbl_hdr, border_color_hex=COLOR_SECONDARY_HEX, border_sz="6")

        c0 = tbl_hdr.cell(0, 0)
        c0.width = Inches(3.54)
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r0 = p0.add_run("Tài liệu lưu hành nội bộ")
        r0.font.name = FONT_NAME
        r0.font.size = Pt(10)
        r0.font.italic = True
        r0.font.color.rgb = COLOR_PRIMARY

        c1 = tbl_hdr.cell(0, 1)
        c1.width = Inches(3.54)
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r1 = p1.add_run("Thầy ")
        r1.font.name = FONT_NAME
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_PRIMARY
        r1_b = p1.add_run("TRẦN VĂN THẠNH")
        r1_b.font.name = FONT_NAME
        r1_b.font.size = Pt(10)
        r1_b.font.bold = True
        r1_b.font.color.rgb = COLOR_PRIMARY
        r1_ph = p1.add_run(" - 0777.470.803")
        r1_ph.font.name = FONT_NAME
        r1_ph.font.size = Pt(10)
        r1_ph.font.color.rgb = COLOR_PRIMARY

        # Footer Table
        footer = s.footer
        footer.is_linked_to_previous = False
        tbl_ftr = footer.add_table(rows=1, cols=2, width=Inches(7.08))
        tbl_ftr.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders_horizontal_only(tbl_ftr, border_color_hex=COLOR_SECONDARY_HEX, border_sz="6")

        fc0 = tbl_ftr.cell(0, 0)
        fc0.width = Inches(5.3)
        fp0 = fc0.paragraphs[0]
        fp0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fr0 = fp0.add_run("CS1: 6/15 Nguyễn Hoàng, Kim Long   |   CS2: 24 Đặng Thái Thân, TP Huế")
        fr0.font.name = FONT_NAME
        fr0.font.size = Pt(9.5)
        fr0.font.italic = True
        fr0.font.color.rgb = COLOR_PRIMARY

        fc1 = tbl_ftr.cell(0, 1)
        fc1.width = Inches(1.78)
        fp1 = fc1.paragraphs[0]
        fp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fr1 = fp1.add_run("Trang ")
        fr1.font.name = FONT_NAME
        fr1.font.size = Pt(10)
        fr1.font.bold = True
        fr1.font.color.rgb = COLOR_PRIMARY

        # Thêm Page field
        fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        fp1._p.append(fldSimple)

def add_header_block(doc, is_teacher=False, is_screen=False):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c0 = tbl.cell(0, 0)
    c0.width = Inches(3.2)
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.line_spacing = 1.15
    p0.paragraph_format.space_after = Pt(2)

    r_ubnd = p0.add_run("TRƯỜNG THCS CHU VĂN AN\n")
    r_ubnd.font.name = FONT_NAME
    r_ubnd.font.size = Pt(10.5)
    r_ubnd.font.bold = True
    r_ubnd.font.color.rgb = COLOR_TEXT

    r_nh = p0.add_run("NĂM HỌC 2026 - 2027\n")
    r_nh.font.name = FONT_NAME
    r_nh.font.size = Pt(10)
    r_nh.font.italic = True

    r_de = p0.add_run("ĐỀ CƯƠNG CHÍNH THỨC")
    r_de.font.name = FONT_NAME
    r_de.font.size = Pt(10)
    r_de.font.bold = True
    r_de.font.color.rgb = COLOR_PRIMARY

    c1 = tbl.cell(0, 1)
    c1.width = Inches(3.88)
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.line_spacing = 1.15
    p1.paragraph_format.space_after = Pt(2)

    r_title = p1.add_run("ĐỀ CƯƠNG ÔN TẬP KIỂM TRA GIỮA HỌC KỲ I\n")
    r_title.font.name = FONT_NAME
    r_title.font.size = Pt(11)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY

    r_sub = p1.add_run("MÔN: KHOA HỌC TỰ NHIÊN - LỚP 7\n")
    r_sub.font.name = FONT_NAME
    r_sub.font.size = Pt(11)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_PRIMARY

    r_note = p1.add_run("(Gồm 3 phân môn: Sinh học - Hóa học - Vật lí)")
    r_note.font.name = FONT_NAME
    r_note.font.size = Pt(10)
    r_note.font.italic = True

    # Teacher notice
    if is_teacher:
        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_t.paragraph_format.space_before = Pt(4)
        p_t.paragraph_format.space_after = Pt(6)
        r_tn = p_t.add_run("(BẢN GIÁO VIÊN -- CÓ LỜI GIẢI CHI TIẾT VÀ ĐÁP ÁN ĐẦY ĐỦ)")
        r_tn.font.name = FONT_NAME
        r_tn.font.size = Pt(10.5)
        r_tn.font.bold = True
        r_tn.font.color.rgb = COLOR_PRIMARY
    elif is_screen:
        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_t.paragraph_format.space_before = Pt(4)
        p_t.paragraph_format.space_after = Pt(6)
        r_tn = p_t.add_run("(BẢN DẠY TRÊN MÀN HÌNH TIVI 65 INCH -- KHÔNG CÓ ĐÁP ÁN)")
        r_tn.font.name = FONT_NAME
        r_tn.font.size = Pt(10.5)
        r_tn.font.bold = True
        r_tn.font.color.rgb = COLOR_PRIMARY
    else:
        p_stu = doc.add_paragraph()
        p_stu.paragraph_format.space_before = Pt(4)
        p_stu.paragraph_format.space_after = Pt(6)
        r_s = p_stu.add_run("Họ và tên học sinh: ........................................................................ Lớp: .............. Trường: .................................................")
        r_s.font.name = FONT_NAME
        r_s.font.size = Pt(10.5)
        r_s.font.italic = True

def add_part_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = COLOR_PRIMARY

def add_subpart_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = Pt(11.5)
    r.font.bold = True
    r.font.color.rgb = COLOR_SOLUTION_TITLE

def add_solution_paragraph(doc, text):
    lines = text.strip().split("\n")
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.space_before = Pt(2) if i == 0 else Pt(1)
        p.paragraph_format.space_after = Pt(2) if i == len(lines)-1 else Pt(1)
        p.paragraph_format.line_spacing = 1.15
        
        # Nếu là dòng đầu tiên, thêm tiền tố [Lời giải chi tiết]
        if i == 0:
            r_tag = p.add_run(">> Hướng dẫn giải chi tiết:\n")
            r_tag.font.name = FONT_NAME
            r_tag.font.size = Pt(10.5)
            r_tag.font.bold = True
            r_tag.font.color.rgb = COLOR_SOLUTION_TITLE
        
        # Xử lý in đậm các phần trong dấu *text*
        parts = line.split("*")
        for idx, part in enumerate(parts):
            if not part:
                continue
            r = p.add_run(part)
            r.font.name = FONT_NAME
            r.font.size = Pt(10.5)
            if idx % 2 == 1:
                r.font.bold = True
                r.font.color.rgb = COLOR_PRIMARY
            else:
                r.font.color.rgb = COLOR_TEXT

def add_dotted_lines(doc, num_lines=3):
    for _ in range(num_lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("......................................................................................................................................................................................")
        r.font.name = FONT_NAME
        r.font.size = Pt(9.5)
        r.font.color.rgb = COLOR_GRAY

# ==============================================================================
# HÀM TẠO FILE WORD (HỌC SINH HOẶC GIÁO VIÊN HOẶC MÀN HÌNH)
# ==============================================================================
def create_word_document(mode="student"):
    doc = docx.Document()
    add_header_footer(doc)
    is_teacher = (mode == "teacher")
    is_screen = (mode == "screen")
    add_header_block(doc, is_teacher=is_teacher, is_screen=is_screen)

    # --------------------------------------------------------------------------
    # I. PHÂN MÔN SINH HỌC
    # --------------------------------------------------------------------------
    add_part_heading(doc, "PHẦN I. PHÂN MÔN SINH HỌC (VẬT SỐNG)")
    for lesson_data in BIO_THEORY:
        add_subpart_heading(doc, lesson_data["lesson"])
        for sec in lesson_data["sections"]:
            p_sec = doc.add_paragraph()
            p_sec.paragraph_format.space_before = Pt(3)
            p_sec.paragraph_format.space_after = Pt(2)
            r_sec = p_sec.add_run(sec["title"])
            r_sec.font.name = FONT_NAME
            r_sec.font.size = Pt(11)
            r_sec.font.bold = True

            for line in sec["content"].strip().split("\n"):
                p_c = doc.add_paragraph()
                p_c.paragraph_format.left_indent = Inches(0.2)
                p_c.paragraph_format.space_before = Pt(1)
                p_c.paragraph_format.space_after = Pt(2)
                p_c.paragraph_format.line_spacing = 1.15
                r_c = p_c.add_run(line)
                r_c.font.name = FONT_NAME
                r_c.font.size = Pt(10.5)

    # --------------------------------------------------------------------------
    # II. PHÂN MÔN HÓA HỌC
    # --------------------------------------------------------------------------
    add_part_heading(doc, "PHẦN II. PHÂN MÔN HÓA HỌC (CHẤT VÀ SỰ BIẾN ĐỔI CHẤT)")
    add_subpart_heading(doc, "I. KIẾN THỨC TRỌNG TÂM")
    for chem_less in CHEM_THEORY:
        p_l = doc.add_paragraph()
        p_l.paragraph_format.space_before = Pt(3)
        p_l.paragraph_format.space_after = Pt(2)
        r_l = p_l.add_run(chem_less["lesson"])
        r_l.font.name = FONT_NAME
        r_l.font.size = Pt(11)
        r_l.font.bold = True

        for pt in chem_less["points"]:
            p_p = doc.add_paragraph()
            p_p.paragraph_format.left_indent = Inches(0.2)
            p_p.paragraph_format.space_before = Pt(1)
            p_p.paragraph_format.space_after = Pt(2)
            p_p.paragraph_format.line_spacing = 1.15
            r_p = p_p.add_run(f"- {pt}")
            r_p.font.name = FONT_NAME
            r_p.font.size = Pt(10.5)

    add_subpart_heading(doc, "II. CÂU HỎI VÀ BÀI TẬP ÔN TẬP")
    p_sq_title = doc.add_paragraph()
    p_sq_title.paragraph_format.space_before = Pt(3)
    p_sq_title.paragraph_format.space_after = Pt(2)
    r_sqt = p_sq_title.add_run("1. Câu hỏi trả lời ngắn:")
    r_sqt.font.name = FONT_NAME
    r_sqt.font.size = Pt(11)
    r_sqt.font.bold = True

    for item in CHEM_SHORT_QUESTIONS:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        r_num = p_q.add_run(f"Câu {item['id']}. ")
        r_num.font.name = FONT_NAME
        r_num.font.size = Pt(10.5)
        r_num.font.bold = True
        r_cnt = p_q.add_run(item["content"])
        r_cnt.font.name = FONT_NAME
        r_cnt.font.size = Pt(10.5)

        if is_teacher:
            add_solution_paragraph(doc, item["solution"])
        elif not is_screen:
            add_dotted_lines(doc, num_lines=2)

    p_eq_title = doc.add_paragraph()
    p_eq_title.paragraph_format.space_before = Pt(5)
    p_eq_title.paragraph_format.space_after = Pt(2)
    r_eqt = p_eq_title.add_run("2. Bài tập tự luận:")
    r_eqt.font.name = FONT_NAME
    r_eqt.font.size = Pt(11)
    r_eqt.font.bold = True

    for item in CHEM_ESSAY_QUESTIONS:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        r_num = p_q.add_run(f"{item['title']}. ")
        r_num.font.name = FONT_NAME
        r_num.font.size = Pt(10.5)
        r_num.font.bold = True
        r_cnt = p_q.add_run(item["content"])
        r_cnt.font.name = FONT_NAME
        r_cnt.font.size = Pt(10.5)

        if is_teacher:
            add_solution_paragraph(doc, item["solution"])
        elif not is_screen:
            add_dotted_lines(doc, num_lines=4)

    # --------------------------------------------------------------------------
    # III. PHÂN MÔN VẬT LÝ
    # --------------------------------------------------------------------------
    add_part_heading(doc, "PHẦN III. PHÂN MÔN VẬT LÝ (NĂNG LƯỢNG VÀ TỐC ĐỘ)")
    add_subpart_heading(doc, "A. PHẦN LÝ THUYẾT")
    for item in PHYS_THEORY_QUESTIONS:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        r_num = p_q.add_run(f"Câu {item['id']}. ")
        r_num.font.name = FONT_NAME
        r_num.font.size = Pt(10.5)
        r_num.font.bold = True
        r_cnt = p_q.add_run(item["content"])
        r_cnt.font.name = FONT_NAME
        r_cnt.font.size = Pt(10.5)

        if is_teacher:
            add_solution_paragraph(doc, item["solution"])
        elif not is_screen:
            add_dotted_lines(doc, num_lines=3)

    add_subpart_heading(doc, "B. PHẦN BÀI TẬP")
    for item in PHYS_EXERCISES:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        r_num = p_q.add_run(f"{item['title']}. ")
        r_num.font.name = FONT_NAME
        r_num.font.size = Pt(10.5)
        r_num.font.bold = True
        r_cnt = p_q.add_run(item["content"])
        r_cnt.font.name = FONT_NAME
        r_cnt.font.size = Pt(10.5)

        if is_teacher:
            add_solution_paragraph(doc, item["solution"])
        elif not is_screen:
            add_dotted_lines(doc, num_lines=4)

    # Dòng kết thúc
    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_end.paragraph_format.space_before = Pt(10)
    r_end = p_end.add_run("“Nhóm GV KHTN 7 – Chúc các em ôn tập và làm bài đạt kết quả xuất sắc!”")
    r_end.font.name = FONT_NAME
    r_end.font.size = Pt(10.5)
    r_end.font.italic = True
    r_end.font.color.rgb = COLOR_PRIMARY

    # Đặt tên file xuất
    if mode == "student":
        filename = "De_Cuong_GK1_KHTN7_ChuVanAn_HocSinh.docx"
    elif mode == "teacher":
        filename = "De_Cuong_GK1_KHTN7_ChuVanAn_GiaoVien.docx"
    else:
        filename = "De_Cuong_GK1_KHTN7_ChuVanAn_DayTrenManHinh.docx"

    out_path = os.path.join(SAN_PHAM_DIR, filename)
    doc.save(out_path)
    print(f"[OK] Đã tạo thành công file Word: {out_path} ({os.path.getsize(out_path):,} bytes)")
    return out_path

# ==============================================================================
# HÀM TẠO FILE SLIDE PDF DẠY TRÊN MÀN HÌNH 16:9 CHO TIVI 65 INCH (KHÔNG CÓ ĐÁP ÁN)
# ==============================================================================
def create_screen_teaching_pdf():
    tex_path = os.path.join(SCRATCH_DIR, "De_Cuong_GK1_KHTN7_ChuVanAn_DayTrenManHinh.tex")
    
    tex_lines = [r"""\documentclass[12pt,landscape]{article}
\usepackage[paperwidth=29.7cm,paperheight=16.706cm,top=1.0cm,bottom=1.0cm,left=1.2cm,right=1.2cm,headheight=0pt,headsep=0pt,footskip=0pt]{geometry}
\usepackage[utf8]{vietnam}
\usepackage{times}
\usepackage{amsmath,amssymb}
\usepackage{graphicx}
\usepackage{xcolor}
\usepackage{tabularx}
\usepackage{enumitem}

\definecolor{brandRed}{RGB}{211, 47, 47}
\definecolor{brandBlue}{RGB}{21, 101, 192}
\definecolor{brandGreen}{RGB}{46, 125, 50}
\definecolor{offBlack}{RGB}{26, 26, 26}
\definecolor{grayText}{RGB}{117, 117, 117}

\pagestyle{empty}
\setlength{\parindent}{0pt}
\setlength{\parskip}{0pt}

\begin{document}
"""]

    # --------------------------------------------------------------------------
    # SLIDE 1: TRANG TIÊU ĐỀ
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 1: TIÊU ĐỀ
% ==============================================================================
\begin{center}
\vspace*{0.6cm}
{\fontsize{24}{30}\selectfont \textbf{\color{offBlack}TRƯỜNG THCS CHU VĂN AN -- NĂM HỌC 2026 -- 2027}}\\[12pt]
{\fontsize{34}{44}\selectfont \textbf{\color{brandRed}ĐỀ CƯƠNG ÔN TẬP KIỂM TRA GIỮA HỌC KỲ I}}\\[10pt]
{\fontsize{28}{36}\selectfont \textbf{\color{brandBlue}MÔN: KHOA HỌC TỰ NHIÊN 7}}\\[8pt]
{\fontsize{22}{28}\selectfont \textit{\color{grayText}(Trọn bộ 3 phân môn: Sinh học -- Hóa học -- Vật lí)}}\\[1.2cm]

{\fontsize{20}{26}\selectfont \textbf{Giáo viên hướng dẫn:} Thầy \textbf{\color{brandRed}TRẦN VĂN THẠNH} -- 0777.470.803}\\[6pt]
{\fontsize{18}{24}\selectfont \textit{Cơ sở 1: 6/15 Nguyễn Hoàng, Kim Long \quad|\quad Cơ sở 2: 24 Đặng Thái Thân, TP Huế}}\\[1.0cm]

{\fontsize{18}{24}\selectfont \textbf{\color{brandGreen}SLIDE TRÌNH CHIẾU TƯƠNG TÁC TIVI 65 INCH (KHÔNG HIỂN THỊ ĐÁP ÁN)}}
\end{center}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 2: SINH HỌC - BÀI 22 (TRAO ĐỔI CHẤT & CHUYỂN HÓA NĂNG LƯỢNG)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 2: SINH HỌC - BÀI 22
% ==============================================================================
{\fontsize{24}{28}\selectfont \textbf{\color{brandBlue}I. PHÂN MÔN SINH HỌC -- BÀI 22}}
\vspace{4pt}

{\fontsize{21}{26}\selectfont \textbf{\color{brandRed}VAI TRÒ CỦA TRAO ĐỔI CHẤT VÀ CHUYỂN HÓA NĂNG LƯỢNG Ở SINH VẬT}}
\vspace{8pt}

\fontsize{19}{25}\selectfont
\textbf{1. Khái niệm trao đổi chất và chuyển hóa năng lượng:}
\begin{itemize}
  \setlength{\itemsep}{3pt}
  \item \textbf{Trao đổi chất:} Quá trình cơ thể lấy các chất từ môi trường ngoài, biến đổi chúng trong tế bào và thải các chất bã ra môi trường.
  \item \textbf{Chuyển hóa năng lượng:} Quá trình biến đổi năng lượng từ dạng này sang dạng khác (quang năng $\to$ hóa năng $\to$ nhiệt năng, cơ năng,...).
\end{itemize}

\vspace{6pt}
\textbf{2. Vai trò đối với cơ thể sinh vật:}
\begin{itemize}
  \setlength{\itemsep}{3pt}
  \item \textbf{Cung cấp nguyên liệu:} Cấu tạo nên tế bào, mô, cơ quan và cơ thể.
  \item \textbf{Cung cấp năng lượng:} Duy trì mọi hoạt động sống (sinh trưởng, phát triển, sinh sản, vận động, cảm ứng).
  \item \textbf{Đào thải chất cặn bã:} Đảm bảo môi trường trong cơ thể được cân bằng, ổn định.
\end{itemize}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 3: SINH HỌC - BÀI 23 (QUANG HỢP Ở THỰC VẬT)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 3: SINH HỌC - BÀI 23
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}I. PHÂN MÔN SINH HỌC -- BÀI 23}}
\vspace{6pt}

{\fontsize{22}{28}\selectfont \textbf{\color{brandRed}QUANG HỢP Ở THỰC VẬT}}
\vspace{12pt}

\fontsize{20}{27}\selectfont
\textbf{1. Khái niệm \& Bào quan thực hiện:} Diễn ra chủ yếu ở \textbf{lục lạp} của tế bào lá cây nhờ sắc tố \textbf{diệp lục}.\\
\textbf{2. Phương trình quang hợp:}
\begin{center}
{\fontsize{21}{28}\selectfont \textbf{Nước} $+$ \textbf{Carbon dioxide} $\xrightarrow[\text{Diệp lục}]{\text{Ánh sáng}}$ \textbf{Glucose} $+$ \textbf{Oxygen}}
\end{center}

\textbf{3. Các yếu tố môi trường ảnh hưởng đến quang hợp:}
\begin{itemize}
  \setlength{\itemsep}{4pt}
  \item Ánh sáng (cường độ và chất lượng ánh sáng).
  \item Hàm lượng khí $\text{CO}_2$ và lượng nước hấp thụ qua rễ.
  \item Nhiệt độ môi trường (tối ưu khoảng $25^\circ\text{C} - 35^\circ\text{C}$) và các nguyên tố khoáng.
\end{itemize}

\textbf{4. Vận dụng trong thực tiễn:} Trồng cây đúng mật độ, tỉa tán đón ánh sáng; tưới tiêu hợp lý; chiếu đèn LED quang hợp nhân tạo trong nhà kính để tăng năng suất cây trồng.
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 4: SINH HỌC - BÀI 25 (HÔ HẤP TẾ BÀO)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 4: SINH HỌC - BÀI 25
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}I. PHÂN MÔN SINH HỌC -- BÀI 25}}
\vspace{6pt}

{\fontsize{22}{28}\selectfont \textbf{\color{brandRed}HÔ HẤP TẾ BÀO}}
\vspace{12pt}

\fontsize{20}{27}\selectfont
\textbf{1. Khái niệm:} Quá trình phân giải chất hữu cơ (chủ yếu là glucose) diễn ra trong \textbf{ti thể}, giải phóng năng lượng ($\text{ATP} + \text{nhiệt}$) cung cấp cho mọi hoạt động sống.\\
\textbf{2. Phương trình hô hấp tế bào:}
\begin{center}
{\fontsize{21}{28}\selectfont \textbf{Glucose} $+$ \textbf{Oxygen} $\longrightarrow$ \textbf{Carbon dioxide} $+$ \textbf{Nước} $+$ \textbf{Năng lượng (ATP + Nhiệt)}}
\end{center}

\textbf{3. Các yếu tố ảnh hưởng:} Hàm lượng nước, nhiệt độ, nồng độ $\text{O}_2$ và nồng độ $\text{CO}_2$.\\
\textbf{4. Vận dụng bảo quản nông sản và đời sống:}
\begin{itemize}
  \setlength{\itemsep}{4pt}
  \item \textbf{Bảo quản khô:} Phơi khô hạt thóc, ngô, đậu để hạ lượng nước $\implies$ giảm hô hấp.
  \item \textbf{Bảo quản lạnh \& Hút chân không:} Giảm nhiệt độ, hạ nồng độ $\text{O}_2 \implies$ kéo dài thời gian bảo quản.
  \item \textbf{Tập thể dục, thông thoáng nhà ở:} Tăng cường tiếp nhận $\text{O}_2$ cho hô hấp tế bào.
\end{itemize}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 5: HÓA HỌC - LÝ THUYẾT NGUYÊN TỬ & NGUYÊN TỐ HÓA HỌC
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 5: HÓA HỌC - LÝ THUYẾT TRỌNG TÂM
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}II. PHÂN MÔN HÓA HỌC -- KIẾN THỨC TRỌNG TÂM}}
\vspace{10pt}

\fontsize{20}{27}\selectfont
\textbf{1. Cấu tạo nguyên tử (Bài 2):}
\begin{itemize}
  \setlength{\itemsep}{4pt}
  \item Nguyên tử là hạt vô cùng nhỏ, \textbf{trung hòa về điện}.
  \item \textbf{Hạt nhân (ở tâm):} Gồm hạt \textbf{proton} ($p$, mang điện tích $+1$) và \textbf{neutron} ($n$, không mang điện).
  \item \textbf{Vỏ nguyên tử:} Gồm các \textbf{electron} ($e$, mang điện tích $-1$) chuyển động quanh hạt nhân.
  \item Trong nguyên tử trung hòa về điện: \textbf{Số proton = Số electron} ($p = e$).
  \item Khối lượng nguyên tử tập trung chủ yếu ở \textbf{hạt nhân}.
\end{itemize}

\vspace{8pt}
\textbf{2. Nguyên tố hóa học (Bài 3):}
\begin{itemize}
  \setlength{\itemsep}{4pt}
  \item Là tập hợp những nguyên tử có \textbf{cùng số proton} trong hạt nhân.
  \item Số proton là đặc trưng cơ bản nhất của một nguyên tố hóa học.
  \item Được kí hiệu bằng 1 hoặc 2 chữ cái (chữ đầu in hoa, chữ sau in thường, ví dụ: H, C, O, Na,...).
\end{itemize}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 6: HÓA HỌC - CÂU HỎI TRẢ LỜI NGẮN (CÂU 1 -> CÂU 4)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 6: HÓA HỌC - CÂU HỎI TRẢ LỜI NGẮN
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}II. HÓA HỌC -- CÂU HỎI TRẢ LỜI NGẮN}}
\vspace{12pt}

\fontsize{21}{29}\selectfont
\textbf{Câu 1.} Nguyên tử có $10$ hạt electron trung hòa về điện thì số hạt proton là bao nhiêu?\\
\textbf{Trả lời:} $\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots$

\vspace{16pt}
\textbf{Câu 2.} Kí hiệu của ba loại hạt cấu tạo nên nguyên tử là gì?\\
\textbf{Trả lời:} $\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots$

\vspace{16pt}
\textbf{Câu 3.} Một nguyên tử có $11$ proton. Hỏi nguyên tử trung hòa về điện có bao nhiêu electron?\\
\textbf{Trả lời:} $\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots$

\vspace{16pt}
\textbf{Câu 4.} Nguyên tử có $8$ proton và $8$ electron mang điện tích tổng cộng bằng bao nhiêu?\\
\textbf{Trả lời:} $\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots$
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 7: HÓA HỌC - TỰ LUẬN BÀI 1 & BÀI 2
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 7: HÓA HỌC - BÀI TẬP TỰ LUẬN 1 & 2
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}II. HÓA HỌC -- BÀI TẬP TỰ LUẬN (BÀI 1 \& BÀI 2)}}
\vspace{10pt}

\fontsize{21}{28}\selectfont
\textbf{Bài 1.} Một nguyên tử có $12$ proton, $12$ neutron và $12$ electron.
\begin{enumerate}[label=\alph*)]
  \setlength{\itemsep}{2pt}
  \item Xác định số hạt proton, neutron và electron.
  \item Tính tổng số hạt trong nguyên tử.
  \item Giải thích vì sao nguyên tử trung hòa về điện.
\end{enumerate}

\vspace{10pt}
\textbf{Bài 2.} Nguyên tử X có tổng số hạt proton, neutron và electron là $23$. Biết số proton là $7$ và nguyên tử trung hòa về điện.
\begin{enumerate}[label=\alph*)]
  \setlength{\itemsep}{2pt}
  \item Tính số electron của nguyên tử X.
  \item Tính số neutron của nguyên tử X.
\end{enumerate}

\vspace*{1.5cm}
\begin{center}
{\fontsize{18}{24}\selectfont \textit{\color{grayText}(Khoảng trống chữa bài trực tiếp trên màn hình)}}
\end{center}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 8: HÓA HỌC - TỰ LUẬN BÀI 3, BÀI 4 & BÀI 5
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 8: HÓA HỌC - BÀI TẬP TỰ LUẬN 3, 4 & 5
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}II. HÓA HỌC -- BÀI TẬP TỰ LUẬN (BÀI 3, 4 \& 5)}}
\vspace{10pt}

\fontsize{20}{27}\selectfont
\textbf{Bài 3.} Phân biệt proton, neutron và electron về vị trí và điện tích trong nguyên tử.\\[10pt]
\textbf{Bài 4.} Một nguyên tử có $17$ proton và $18$ neutron.
\begin{enumerate}[label=\alph*)]
  \setlength{\itemsep}{2pt}
  \item Nguyên tử này có bao nhiêu electron nếu trung hòa về điện?
  \item Có bao nhiêu hạt trong hạt nhân?
  \item Tổng số hạt của nguyên tử là bao nhiêu?
\end{enumerate}

\vspace{8pt}
\textbf{Bài 5.} Giải thích sự khác nhau giữa nguyên tử và nguyên tố hóa học. Cho ví dụ minh họa.

\vspace*{1.2cm}
\begin{center}
{\fontsize{18}{24}\selectfont \textit{\color{grayText}(Khoảng trống tương tác và viết bài trên Tivi 65 inch)}}
\end{center}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 9: VẬT LÝ - PHẦN LÝ THUYẾT (CÂU 1, 2, 3)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 9: VẬT LÝ - LÝ THUYẾT (CÂU 1, 2, 3)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}III. PHÂN MÔN VẬT LÝ -- PHẦN LÝ THUYẾT (CÂU 1, 2, 3)}}
\vspace{12pt}

\fontsize{21}{29}\selectfont
\textbf{Câu 1.} Nêu ý nghĩa của tốc độ chuyển động.\\
\textbf{Trả lời:} $\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots$

\vspace{16pt}
\textbf{Câu 2.} Nêu công thức tính tốc độ, giải thích các đại lượng và đơn vị tương ứng trong công thức.\\
\textbf{Trả lời:} $\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots$

\vspace{16pt}
\textbf{Câu 3.} Liệt kê một số đơn vị đo tốc độ thường dùng trong đời sống và kỹ thuật.\\
\textbf{Trả lời:} $\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots$

\vspace*{1.5cm}
\begin{center}
{\fontsize{18}{24}\selectfont \textit{\color{grayText}(Giáo viên và học sinh trao đổi, ghi chú trực tiếp trên bảng)}}
\end{center}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 10: VẬT LÝ - PHẦN LÝ THUYẾT (CÂU 4 & CÂU 5 - ĐO TỐC ĐỘ)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 10: VẬT LÝ - LÝ THUYẾT ĐO TỐC ĐỘ (CÂU 4 & 5)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}III. VẬT LÝ -- PHƯƠNG PHÁP ĐO TỐC ĐỘ (CÂU 4 \& 5)}}
\vspace{12pt}

\fontsize{21}{29}\selectfont
\textbf{Câu 4.} Mô tả sơ lược cách đo tốc độ bằng \textbf{đồng hồ bấm giây} và \textbf{cổng quang điện} trong dụng cụ thực hành ở nhà trường.\\[12pt]
\textbf{Câu 5.} Mô tả sơ lược cách đo tốc độ bằng \textbf{thiết bị “bắn tốc độ”} trong kiểm tra tốc độ của các phương tiện giao thông.

\vspace*{3.0cm}
\begin{center}
{\fontsize{19}{25}\selectfont \textit{\color{grayText}(Khoảng trống phân tích sơ đồ thí nghiệm và nguyên lý thiết bị radar/laser)}}
\end{center}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 11: VẬT LÝ - BÀI TẬP CÂU 1 (ĐỔI ĐƠN VỊ TỐC ĐỘ)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 11: VẬT LÝ - BÀI TẬP CÂU 1 (ĐỔI ĐƠN VỊ)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}III. VẬT LÝ -- BÀI TẬP (CÂU 1: ĐỔI ĐƠN VỊ TỐC ĐỘ)}}
\vspace{12pt}

\fontsize{22}{30}\selectfont
\textbf{Câu 1.} Tìm số thích hợp để điền vào chỗ trống:
\begin{enumerate}[label=\alph*),itemsep=12pt]
  \item $15\text{ m/s} = \dots\dots\dots\dots\text{ km/h}$.
  \item $90\text{ km/h} = \dots\dots\dots\dots\text{ m/s}$.
  \item $250\text{ cm/s} = \dots\dots\dots\dots\text{ m/s} = \dots\dots\dots\dots\text{ km/h}$.
  \item $72\text{ km/h} = \dots\dots\dots\dots\text{ m/s} = \dots\dots\dots\dots\text{ cm/s}$.
\end{enumerate}

\vspace*{2.0cm}
\begin{center}
{\fontsize{18}{24}\selectfont \textit{\color{grayText}(Nhắc lại: $1\text{ m/s} = 3,6\text{ km/h}$ \quad|\quad $1\text{ km/h} = \frac{1}{3,6}\text{ m/s}$)}}
\end{center}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 12: VẬT LÝ - BÀI TẬP CÂU 2 & CÂU 3
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 12: VẬT LÝ - BÀI TẬP CÂU 2 & 3
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}III. VẬT LÝ -- BÀI TẬP TÍNH TỐC ĐỘ (CÂU 2 \& 3)}}
\vspace{12pt}

\fontsize{21}{29}\selectfont
\textbf{Câu 2.} Một học sinh đi xe đạp từ nhà đến trường, xuất phát lúc $6\text{ giờ } 30\text{ phút}$ và đến trường lúc $6\text{ giờ } 50\text{ phút}$. Biết quãng đường từ nhà đến trường dài $4\text{ km}$ và học sinh đi liên tục không nghỉ. Tốc độ trung bình của học sinh trên quãng đường này là bao nhiêu $\text{km/h}$?

\vspace{16pt}
\textbf{Câu 3.} Tốc độ trung bình của máu chảy trong động mạch của người là $20\text{ cm/s}$. Giả sử một hồng cầu chuyển động liên tục với tốc độ này trong $2\text{ giờ}$. Quãng đường hồng cầu đi được trong thời gian đó là bao nhiêu mét?

\vspace*{2.0cm}
\begin{center}
{\fontsize{18}{24}\selectfont \textit{\color{grayText}(Khoảng trống tóm tắt và trình bày lời giải)}}
\end{center}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 13: VẬT LÝ - BÀI TẬP CÂU 4 & CÂU 5
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 13: VẬT LÝ - BÀI TẬP CÂU 4 & 5
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}III. VẬT LÝ -- BÀI TẬP (CÂU 4 \& 5)}}
\vspace{10pt}

\fontsize{20}{27}\selectfont
\textbf{Câu 4.} Tuyến đường sắt từ Hà Nội đến Đà Nẵng dài khoảng $880\text{ km}$. Một tàu hỏa di chuyển trên tuyến đường này với tốc độ trung bình $55\text{ km/h}$. Hỏi tàu hỏa mất bao nhiêu giờ để đi từ Hà Nội đến Đà Nẵng?\\[10pt]
\textbf{Câu 5.} Bạn An đi xe đạp điện từ nhà đến trường, trên đoạn đường đầu dài $3\text{ km}$ An đi hết $8\text{ phút}$. Sau đó, An đi tiếp một đoạn đường $2,2\text{ km}$ trong $5\text{ phút}$ nữa thì đến trường.
\begin{enumerate}[label=\alph*)]
  \setlength{\itemsep}{2pt}
  \item Tính tốc độ của An ứng với từng đoạn đường theo đơn vị $\text{km/h}$.
  \item Trên đoạn đường nào An đi nhanh hơn?
  \item Tính tốc độ trung bình của An trên cả đoạn đường từ nhà đến trường theo đơn vị $\text{km/h}$.
\end{enumerate}

\vspace*{1.5cm}
\begin{center}
{\fontsize{18}{24}\selectfont \textit{\color{grayText}(Khoảng trống tính toán: $v = \frac{s}{t}$, $v_{tb} = \frac{s_1 + s_2}{t_1 + t_2}$)}}
\end{center}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 14: VẬT LÝ - BÀI TẬP CÂU 6 (BẮN TỐC ĐỘ CỦA CSGT)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 14: VẬT LÝ -- BÀI TẬP CÂU 6
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}III. VẬT LÝ -- BÀI TOÁN THỰC TẾ (CÂU 6: BẮN TỐC ĐỘ)}}
\vspace{12pt}

\fontsize{22}{30}\selectfont
\textbf{Câu 6.} Trên một đoạn đường, cảnh sát giao thông sử dụng thiết bị bắn tốc độ có camera để đo tốc độ của các phương tiện.\\
Camera ghi nhận một ô tô đi qua hai vạch mốc cách nhau $10\text{ m}$ trong thời gian $0,5\text{ s}$.

\begin{enumerate}[label=\alph*),itemsep=8pt]
  \item Tính tốc độ của ô tô theo đơn vị $\text{km/h}$.
  \item Biết tốc độ tối đa cho phép trên đoạn đường này là $60\text{ km/h}$. Hỏi ô tô có vi phạm quy định về tốc độ không? Vì sao?
\end{enumerate}

\vspace*{2.5cm}
\begin{center}
{\fontsize{19}{25}\selectfont \textit{\color{grayText}(Khoảng trống học sinh lên bảng giải và đối chiếu quy định an toàn giao thông)}}
\end{center}
\newpage
""")

    tex_lines.append(r"\end{document}")

    with open(tex_path, "w", encoding="utf-8") as f:
        f.write("\n".join(tex_lines))
    print(f"[OK] Đã ghi file LaTeX slide: {tex_path}")

    # Biên dịch pdflatex
    cmd = ["pdflatex", "-interaction=nonstopmode", "-output-directory", SCRATCH_DIR, tex_path]
    res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    pdf_temp = os.path.join(SCRATCH_DIR, "De_Cuong_GK1_KHTN7_ChuVanAn_DayTrenManHinh.pdf")
    pdf_out = os.path.join(SAN_PHAM_DIR, "De_Cuong_GK1_KHTN7_ChuVanAn_DayTrenManHinh.pdf")

    if os.path.exists(pdf_temp):
        shutil.copy2(pdf_temp, pdf_out)
        print(f"[OK] Đã biên dịch và copy PDF thành công: {pdf_out} ({os.path.getsize(pdf_out):,} bytes)")
    else:
        print(f"[ERR] Không tạo được PDF. Mã lỗi: {res.returncode}")
    return pdf_out

# ==============================================================================
# HÀM ĐỒNG BỘ ONEDRIVE
# ==============================================================================
def sync_to_onedrive():
    files_to_sync = [
        "De_Cuong_GK1_KHTN7_ChuVanAn_HocSinh.docx",
        "De_Cuong_GK1_KHTN7_ChuVanAn_GiaoVien.docx",
        "De_Cuong_GK1_KHTN7_ChuVanAn_DayTrenManHinh.docx",
        "De_Cuong_GK1_KHTN7_ChuVanAn_DayTrenManHinh.pdf",
    ]
    print("\n[+] Đồng bộ các file sang các thư mục OneDrive:")
    for dest_dir in [ONEDRIVE_DIR1, ONEDRIVE_DIR2]:
        if os.path.exists(dest_dir):
            for fname in files_to_sync:
                src_f = os.path.join(SAN_PHAM_DIR, fname)
                dst_f = os.path.join(dest_dir, fname)
                if os.path.exists(src_f):
                    try:
                        shutil.copy2(src_f, dst_f)
                        print(f"    -> Đã copy {fname} vào: {dest_dir}")
                    except Exception as e:
                        print(f"    [!] Cảnh báo file đang mở/khóa tại {dest_dir}: {e}")
        else:
            print(f"    [-] Không tìm thấy thư mục: {dest_dir}")

# ==============================================================================
# MAIN
# ==============================================================================
if __name__ == "__main__":
    print("=" * 75)
    print("BẮT ĐẦU TẠO 3 FILE CHO ĐỀ CƯƠNG ÔN TẬP GIỮA KỲ I KHTN 7 (THCS CHU VĂN AN)")
    print("1. FILE HỌC SINH (.docx)")
    print("2. FILE GIÁO VIÊN (.docx có lời giải chi tiết)")
    print("3. FILE DẠY TRÊN MÀN HÌNH (Slide PDF 16:9 widescreen cho Tivi 65 inch)")
    print("=" * 75)

    create_word_document(mode="student")
    create_word_document(mode="teacher")
    create_word_document(mode="screen")
    create_screen_teaching_pdf()
    sync_to_onedrive()

    print("\n" + "=" * 75)
    print("HOÀN THÀNH TẤT CẢ 3 FILE THEO ĐÚNG CHUẨN FORM ĐỀ KIỂM TRA!")
    print("=" * 75)
