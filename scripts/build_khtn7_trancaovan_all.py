# -*- coding: utf-8 -*-
r"""
Script biên soạn trọn bộ tài liệu chuẩn FORM ĐỀ KIỂM TRA cho:
ĐỀ CƯƠNG ÔN TẬP GIỮA HỌC KÌ I – MÔN KHTN 7
TRƯỜNG THCS TRẦN CAO VÂN – NĂM HỌC 2026 - 2027
1. De_Cuong_GK1_KHTN7_TranCaoVan_HocSinh.docx
2. De_Cuong_GK1_KHTN7_TranCaoVan_GiaoVien.docx
3. De_Cuong_GK1_KHTN7_TranCaoVan_DayTrenManHinh.pdf (Slide 16:9 Widescreen KHÔNG CÓ ĐÁP ÁN)
Và bản Word: De_Cuong_GK1_KHTN7_TranCaoVan_DayTrenManHinh.docx
Thầy TRẦN VĂN THẠNH - 0777.470.803
"""

import os
import sys
import subprocess
import shutil

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

import docx
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# Import các hàm chuẩn từ form_de_kiem_tra
sys.path.insert(0, r"c:\Antigravity_Thanh\scripts")
from form_de_kiem_tra import (
    setup_header_footer,
    build_title_block,
    add_section_header,
    render_question_prompt,
    render_options_no_table,
    render_statement,
    add_solution_no_table,
    add_student_writing_lines,
    render_rich_text,
    set_run_style,
    COLOR_PRIMARY,
    COLOR_CORRECT,
    COLOR_SOLUTION,
    COLOR_HF_RED,
    COLOR_HF_GREEN_HEX
)

OUTPUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
SCRATCH_DIR = r"c:\Antigravity_Thanh\scratch"
IMAGES_DIR = r"c:\Antigravity_Thanh\San_Pham\images_khtn7_trancaovan"
os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(SCRATCH_DIR, exist_ok=True)
os.makedirs(IMAGES_DIR, exist_ok=True)

TARGET_ONEDRIVE_DIRS = [
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn7\KHTN7-IN-DEKIEMTRADINHKI",
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn7"
]

# ==============================================================================
# DỮ LIỆU ĐỀ CƯƠNG KHTN 7 TRẦN CAO VÂN
# ==============================================================================

# --- PHẦN 1: HÓA HỌC (CHẤT VÀ SỰ BIẾN ĐỔI CHẤT) ---
HOA_THEORY = [
    ("Bài 2: Nguyên tử", [
        "Quan niệm ban đầu: Nguyên tử là hạt vô cùng nhỏ bé cấu tạo nên mọi chất.",
        "Mô hình Rutherford – Bohr: Hạt nhân ở tâm mang điện tích dương; các electron mang điện tích âm chuyển động xung quanh hạt nhân theo từng lớp.",
        "Cấu tạo nguyên tử: Gồm hạt nhân (proton mang điện +1, neutron không mang điện) và vỏ nguyên tử (electron mang điện -1). Trong nguyên tử: *Số proton = Số electron*.",
        "Khối lượng nguyên tử: Tập trung hầu hết ở hạt nhân, đo bằng đơn vị quốc tế amu (1 amu ≈ 1,6605 × 10^-24 g)."
    ]),
    ("Bài 3: Nguyên tố hóa học", [
        "Khái niệm: Nguyên tố hóa học là tập hợp các nguyên tử cùng loại có cùng số proton trong hạt nhân.",
        "Tên gọi và kí hiệu: Gồm 1 hoặc 2 chữ cái (chữ đầu viết hoa, chữ thứ hai viết thường)."
    ])
]

HOA_MCQ = [
    {
        "id": 1,
        "content": "Các hạt cấu tạo nên hạt nhân của hầu hết các nguyên tử là:",
        "options": [
            "A. electron và neutron.",
            "B. proton và neutron.",
            "C. neutron và electron.",
            "D. electron, proton và neutron."
        ],
        "answer": "B",
        "solution": "Hạt nhân gồm hạt proton mang điện tích dương và hạt neutron không mang điện."
    },
    {
        "id": 2,
        "content": "Các hạt cấu tạo nên hầu hết các nguyên tử là:",
        "options": [
            "A. electron và neutron.",
            "B. proton và neutron.",
            "C. neutron và electron.",
            "D. electron, proton và neutron."
        ],
        "answer": "D",
        "solution": "Cấu tạo nguyên tử gồm đầy đủ 3 loại hạt: electron (ở vỏ), proton và neutron (ở hạt nhân)."
    },
    {
        "id": 3,
        "content": "Một nguyên tử Carbon có khối lượng là 1,9926 × 10^-23 gam. Khối lượng nguyên tử Carbon tính theo đơn vị amu là? (biết 1 amu = 1,6605 × 10^-24 gam)",
        "options": [
            "A. 6 amu.",
            "B. 12 amu.",
            "C. 24 amu.",
            "D. 48 amu."
        ],
        "answer": "B",
        "solution": "Khối lượng theo amu: *m = (1,9926 × 10^-23) / (1,6605 × 10^-24) = 12 amu*."
    },
    {
        "id": 4,
        "content": "Cho thành phần các nguyên tử như sau: X (17p, 17e, 16n), Y (20p, 20e, 19n), Z (17p, 17e, 18n), T (19p, 19e, 20n). Có bao nhiêu nguyên tố hóa học?",
        "options": [
            "A. 1.",
            "B. 2.",
            "C. 3.",
            "D. 4."
        ],
        "answer": "C",
        "solution": "Nguyên tố hóa học được xác định bởi số proton. X và Z có cùng 17p (Chlorine); Y có 20p (Calcium); T có 19p (Potassium) => Có *3 nguyên tố hóa học*."
    },
    {
        "id": 5,
        "content": "Tên của nguyên tố hóa học có kí hiệu hóa học là Fe là:",
        "options": [
            "A. Fluorine.",
            "B. Iron.",
            "C. Phosphorus.",
            "D. Potassium."
        ],
        "answer": "B",
        "solution": "Fe là kí hiệu hóa học của nguyên tố Iron (Sắt)."
    }
]

HOA_TF = [
    {
        "id": 6,
        "prompt": "Câu 6. Cho các nhận định về nguyên tố hóa học và kí hiệu hóa học, xác định tính Đúng hay Sai của mỗi nhận định:",
        "items": [
            {"id": "a", "text": "Nguyên tố hóa học là tập hợp các nguyên tử có cùng số proton trong hạt nhân.", "ans": True, "reason": "Đúng định nghĩa nguyên tố hóa học."},
            {"id": "b", "text": "Kí hiệu hóa học của nguyên tố nhôm là Al, của đồng là Cu.", "ans": True, "reason": "Al (Aluminium), Cu (Copper)."},
            {"id": "c", "text": "Khối lượng nguyên tử tính theo đơn vị gam.", "ans": False, "reason": "Khối lượng nguyên tử tính theo đơn vị quốc tế amu."},
            {"id": "d", "text": "Các nguyên tử của cùng một nguyên tố hóa học có thể có số neutron khác nhau nhưng số proton luôn giống nhau.", "ans": True, "reason": "Đúng (các đồng vị có số n khác nhau nhưng số p luôn bằng nhau)."}
        ]
    }
]

HOA_SHORT = [
    {
        "id": 7,
        "content": "Cho biết tổng số hạt proton, neutron và electron trong nguyên tử của một nguyên tố X là 13, trong đó số hạt mang điện nhiều hơn số hạt không mang điện là 3. Số proton của nguyên tố X bằng bao nhiêu?",
        "answer": "4",
        "solution": "Hệ phương trình: 2p + n = 13 và 2p - n = 3 => 4p = 16 => *p = 4* (n = 5). Nguyên tố X là Beryllium (Be)."
    }
]

HOA_ESSAY = [
    {
        "id": 8,
        "content": "Cho biết nguyên tử Aluminium (Nhôm) có 13p trong hạt nhân.\na) Tìm số electron ở vỏ nguyên tử Aluminium.\nb) Vẽ sơ đồ phân bố electron vào các lớp của nguyên tử Aluminium.\nc) Cho biết Aluminium có bao nhiêu lớp electron và số electron lớp ngoài cùng.",
        "solution": (
            "a) Trong nguyên tử trung hòa về điện: Số electron = Số proton = *13 electron*.\n"
            "b) Sơ đồ phân bố electron:\n"
            "   - Hạt nhân: mang điện tích +13.\n"
            "   - Lớp 1 (trong cùng): có 2 electron.\n"
            "   - Lớp 2: có 8 electron.\n"
            "   - Lớp 3 (ngoài cùng): có 3 electron.\n"
            "c) Nguyên tử Aluminium có *3 lớp electron* và *3 electron ở lớp ngoài cùng*."
        )
    },
    {
        "id": 9,
        "content": "a) Hoàn thành kí hiệu hóa học của các nguyên tố sau: Carbon, Nitrogen, Calcium, Magnesium, Aluminium, Chlorine, Zinc, Mercury.\nb) Cho các cách viết sau: 2 O, 3 Fe, Cu, 4 Na. Hãy cho biết các cách viết này biểu thị điều gì?",
        "solution": (
            "a) Kí hiệu hóa học:\n"
            "   • Carbon: *C*           • Nitrogen: *N*\n"
            "   • Calcium: *Ca*         • Magnesium: *Mg*\n"
            "   • Aluminium: *Al*       • Chlorine: *Cl*\n"
            "   • Zinc: *Zn*            • Mercury: *Hg*\n\n"
            "b) Ý nghĩa các cách viết:\n"
            "   • 2 O: 2 nguyên tử oxygen.\n"
            "   • 3 Fe: 3 nguyên tử iron (sắt).\n"
            "   • Cu: 1 nguyên tử copper (đồng) hoặc đơn chất đồng.\n"
            "   • 4 Na: 4 nguyên tử sodium (natri)."
        )
    }
]

# --- PHẦN 2: VẬT LÍ (NĂNG LƯỢNG) ---
LY_THEORY = [
    ("1. Tốc độ chuyển động", [
        "Tốc độ là đại lượng cho biết mức độ nhanh hay chậm của chuyển động, xác định bằng chiều dài quãng đường đi được trong một đơn vị thời gian.",
        "Công thức tính tốc độ: *v = s / t* (s: quãng đường; t: thời gian; v: tốc độ).",
        "Đơn vị đo: m/s và km/h (1 m/s = 3,6 km/h)."
    ]),
    ("2. Đo tốc độ & Thiết bị bắn tốc độ", [
        "Đo bằng đồng hồ bấm giây (đo s và t rồi tính v) hoặc đồng hồ hiện số dùng cổng quang điện.",
        "Thiết bị bắn tốc độ gồm camera ghi hình và máy phát tia hồng ngoại/laser đo thời gian phương tiện vượt qua hai vạch mốc để tính tốc độ tức thời."
    ]),
    ("3. Đồ thị quãng đường – thời gian", [
        "Đồ thị biểu diễn mối liên hệ giữa quãng đường s và thời gian t.",
        "Đoạn đồ thị nằm ngang: vật đứng yên; đoạn đồ thị càng dốc: vật chuyển động càng nhanh."
    ])
]

LY_MCQ = [
    {
        "id": 1,
        "content": "Gọi s là quãng đường đi được, t là thời gian đi hết quãng đường đó, v là tốc độ chuyển động. Công thức nào sau đây dùng để tính tốc độ chuyển động?",
        "options": [
            "A. v = s × t.",
            "B. v = s / t.",
            "C. s = v / t.",
            "D. t = v / s."
        ],
        "answer": "B",
        "solution": "Công thức chuẩn: *v = s / t*."
    },
    {
        "id": 2,
        "content": "Một đoàn tàu hỏa đi từ ga A đến ga B cách nhau 30 km trong 45 phút. Tốc độ của đoàn tàu là:",
        "options": [
            "A. 60 km/h.",
            "B. 40 km/h.",
            "C. 50 km/h.",
            "D. 55 km/h."
        ],
        "answer": "B",
        "solution": "Đổi t = 45 phút = 45/60 h = 0,75 h. Tốc độ: *v = 30 / 0,75 = 40 km/h*."
    },
    {
        "id": 3,
        "content": "Đại lượng nào đặc trưng cho sự nhanh, chậm của chuyển động?",
        "options": [
            "A. Quãng đường.",
            "B. Thời gian.",
            "C. Tốc độ.",
            "D. Nhiệt độ."
        ],
        "answer": "C",
        "solution": "Tốc độ là đại lượng đặc trưng cho độ nhanh chậm của chuyển động."
    },
    {
        "id": 4,
        "content": "Đơn vị đo tốc độ là:",
        "options": [
            "A. m.",
            "B. s/m.",
            "C. m/s.",
            "D. m.s."
        ],
        "answer": "C",
        "solution": "Đơn vị đo tốc độ trong hệ SI là mét trên giây (m/s)."
    },
    {
        "id": 5,
        "content": "Chọn đáp án đúng về đơn vị đo tốc độ:",
        "options": [
            "A. Đơn vị đo tốc độ phụ thuộc vào đơn vị độ dài.",
            "B. Đơn vị đo tốc độ phụ thuộc vào đơn vị đo thời gian.",
            "C. Cả A và B đều đúng.",
            "D. Cả A và B đều sai."
        ],
        "answer": "C",
        "solution": "Vì v = s / t nên đơn vị tốc độ phụ thuộc đồng thời vào đơn vị đo độ dài và đơn vị đo thời gian."
    },
    {
        "id": 6,
        "content": "An đạp xe từ nhà đến trường mất 10 phút, biết tốc độ của An là 2 m/s. Tính quãng đường từ nhà An đến trường?",
        "options": [
            "A. 120 m.",
            "B. 1,2 m.",
            "C. 12 km.",
            "D. 1,2 km."
        ],
        "answer": "D",
        "solution": "Đổi t = 10 phút = 600 s. Quãng đường: *s = v × t = 2 × 600 = 1200 m = 1,2 km*."
    }
]

LY_TF = [
    {
        "id": 1,
        "prompt": "Câu 1. Trong bài học về đo tốc độ (KHTN 7), xác định tính Đúng hoặc Sai của mỗi nhận định sau:",
        "items": [
            {"id": "a", "text": "Để đo tốc độ của xe đồ chơi bằng đồng hồ bấm giây, ta cần sử dụng thước để đo quãng đường và sử dụng đồng hồ bấm giây để đo thời gian chạy xe.", "ans": True, "reason": "Đúng phương pháp đo gián tiếp tốc độ."},
            {"id": "b", "text": "Trong phương pháp sử dụng đồng hồ hiện số dùng cổng quang điện, đồng hồ chỉ bắt đầu đếm khi xe đã đi qua cả hai cổng quang điện.", "ans": False, "reason": "Đồng hồ bắt đầu đếm khi xe chắn cổng 1 và dừng đếm khi xe qua cổng 2."},
            {"id": "c", "text": "Thiết bị \"bắn tốc độ\" có công dụng chính là ghi lại hình ảnh và đo tốc độ để kiểm tra xem phương tiện giao thông có vi phạm quy định về tốc độ hay không.", "ans": True, "reason": "Đúng chức năng thiết bị bắn tốc độ nghiệp vụ giao thông."},
            {"id": "d", "text": "Thiết bị \"bắn tốc độ\" bắt buộc phải gắn cố định cổng quang điện trên mặt đường mới được xác định tốc độ của xe đang chạy.", "ans": False, "reason": "Thiết bị bắn tốc độ dùng sóng vô tuyến (radar) hoặc laser đo từ xa, không gắn cổng quang trên đường."}
        ]
    }
]

LY_ESSAY = [
    {
        "id": 1,
        "content": "Em hãy đổi các đơn vị sau:\na) 125 cm = ......... m;\nb) 5,03 km = ............ m;\nc) 250 m = .......... km;\nd) 1 giờ 20 phút = ...... giờ = ....... phút = ........ giây;\ne) 1500 giây = ...... giờ;\nf) 15 km/h = ...... m/s;\ng) 15 m/s = ........... km/h;\nh) 1500 cm/min = ...... m/s.",
        "solution": (
            "a) 125 cm = *1,25 m*\n"
            "b) 5,03 km = *5030 m*\n"
            "c) 250 m = *0,25 km*\n"
            "d) 1 giờ 20 phút = *4/3 giờ* (≈ 1,333 giờ) = *80 phút* = *4800 giây*\n"
            "e) 1500 giây = 1500 / 3600 = *5/12 giờ* (≈ 0,417 giờ)\n"
            "f) 15 km/h = 15 / 3,6 = *25/6 m/s* (≈ 4,17 m/s)\n"
            "g) 15 m/s = 15 × 3,6 = *54 km/h*\n"
            "h) 1500 cm/min = 15 m / 60 s = *0,25 m/s*"
        )
    },
    {
        "id": 2,
        "content": "Nói tốc độ đi xe đạp của bạn An là 15 km/h, con số này cho ta biết điều gì?",
        "solution": "Con số 15 km/h cho biết mức độ nhanh chậm của chuyển động, nghĩa là trong mỗi giờ bạn An đi được quãng đường dài 15 km (nếu chuyển động đều liên tục trong 1 giờ thì bạn An đi được 15 km)."
    },
    {
        "id": 3,
        "content": "Một chiếc ô tô được phát hiện bởi thiết bị bắn tốc độ, thời gian ô tô chạy từ vạch mốc 1 và vạch mốc 2 cách nhau 5 m là 0,2 s. Nếu tốc độ giới hạn là 22 m/s thì xe có vượt quá tốc độ cho phép không?",
        "solution": (
            "• Tốc độ của ô tô khi chạy qua hai vạch mốc là:\n"
            "  v = s / t = 5 / 0,2 = *25 (m/s)*.\n"
            "• So sánh với tốc độ giới hạn:\n"
            "  Vì 25 m/s > 22 m/s (vượt quá 3 m/s, tương đương vượt quá 10,8 km/h).\n"
            "• Kết luận: Chiếc ô tô này *ĐÃ VƯỢT QUÁ TỐC ĐỘ CHO PHÉP*."
        )
    },
    {
        "id": 4,
        "content": "Hùng đạp xe lên dốc dài 150 m với tốc độ 2 m/s, sau đó xuống dốc dài 120 m hết 30 s. Hỏi tốc độ trung bình của Hùng trên cả đoạn đường dốc?",
        "solution": (
            "• Thời gian Hùng đạp xe lên dốc:\n"
            "  t_1 = s_1 / v_1 = 150 / 2 = 75 (s).\n"
            "• Tổng quãng đường cả lên dốc và xuống dốc:\n"
            "  s = s_1 + s_2 = 150 + 120 = 270 (m).\n"
            "• Tổng thời gian chuyển động:\n"
            "  t = t_1 + t_2 = 75 + 30 = 105 (s).\n"
            "• Tốc độ trung bình trên cả đoạn đường dốc:\n"
            "  v_tb = s / t = 270 / 105 = 18/7 ≈ *2,57 (m/s)* (hoặc ≈ *9,26 km/h*)."
        )
    },
    {
        "id": 5,
        "content": "Cho bảng ghi quãng đường đi được theo thời gian của một xe ô tô:\n\n[BẢNG SỐ LIỆU]\nThời gian (h): 0 | 1 | 2 | 3 | 4 | 5 | 6\nQuãng đường (km): 0 | 60 | 120 | 180 | 180 | 220 | 260\n\nHãy vẽ đồ thị quãng đường – thời gian của ô tô và nhận xét đặc điểm chuyển động.",
        "image": "dothi_oto_bai5.png",
        "solution": (
            "• Các bước vẽ đồ thị:\n"
            "  1. Vẽ hệ trục tọa độ vuông góc: Trục hoành t (h), tỉ lệ 1 cm ứng với 1 h. Trục tung s (km), tỉ lệ 1 cm ứng với 60 km.\n"
            "  2. Xác định các điểm tọa độ: (0; 0), (1; 60), (2; 120), (3; 180), (4; 180), (5; 220), (6; 260).\n"
            "  3. Nối các điểm bằng các đoạn thẳng.\n"
            "• Nhận xét đặc điểm chuyển động:\n"
            "  - Từ t = 0 đến t = 3 h: Ô tô chuyển động đều với tốc độ v_1 = 60 km/h (đồ thị dốc lên).\n"
            "  - Từ t = 3 h đến t = 4 h: Đoạn đồ thị nằm ngang, s = 180 km không đổi. Ô tô *dừng lại nghỉ 1 giờ* (tốc độ bằng 0).\n"
            "  - Từ t = 4 h đến t = 6 h: Ô tô tiếp tục chuyển động đều đi thêm 80 km trong 2 h với tốc độ v_2 = (260 - 180) / 2 = 40 km/h."
        )
    }
]

# --- PHẦN 3: SINH HỌC (VẬT SỐNG) ---
SINH_THEORY = [
    ("Bài 21: Trao đổi chất và chuyển hóa năng lượng", [
        "Trao đổi chất là quá trình sinh vật lấy chất từ môi trường, biến đổi tạo chất cần thiết và thải chất thải ra ngoài.",
        "Chuyển hóa năng lượng là sự biến đổi năng lượng từ dạng này sang dạng khác (quang năng -> hóa năng -> nhiệt năng, cơ năng...).",
        "Trao đổi chất và chuyển hóa năng lượng là điều kiện cơ bản duy trì sự sống."
    ]),
    ("Bài 22 & 23: Quang hợp ở thực vật", [
        "Khái niệm: Lá cây hấp thụ CO2 và H2O nhờ năng lượng ánh sáng và diệp lục tổng hợp glucose và giải phóng O2.",
        "Phương trình chữ: Nước + Carbon dioxide --(Ánh sáng / Diệp lục)--> Glucose + Oxygen.",
        "Các yếu tố ảnh hưởng: Ánh sáng, nước, hàm lượng CO2, nhiệt độ môi trường."
    ]),
    ("Bài 24: Thực hành chứng minh quang hợp tạo tinh bột", [
        "Đặt cây vào chỗ tối 2 ngày để hết tinh bột trong lá; dùng băng giấy đen bịt một phần lá; đem phơi nắng 4-6 giờ; đun cách thủy trong cồn để tẩy diệp lục; nhỏ thuốc thử iodine.",
        "Hiện tượng: Phần lá hứng ánh sáng chuyển sang màu xanh tím (có tinh bột), phần bịt giấy đen giữ màu vàng nhạt."
    ]),
    ("Bài 25, 26, 27: Hô hấp tế bào & Thực hành", [
        "Phương trình: Glucose + Oxygen -> Carbon dioxide + Nước + Năng lượng (ATP + nhiệt).",
        "Các yếu tố ảnh hưởng: Hàm lượng nước, nhiệt độ, nồng độ O2, nồng độ CO2.",
        "Ứng dụng: Phơi khô hạt giống, bảo quản lạnh, hút chân không để ức chế hô hấp tế bào giúp bảo quản nông sản lâu dài."
    ])
]

SINH_MCQ = [
    {
        "id": 1,
        "content": "Quá trình quang hợp ở thực vật diễn ra chủ yếu trong bào quan nào sau đây?",
        "options": [
            "A. Lục lạp.",
            "B. Ti thể.",
            "C. Ribosome.",
            "D. Bộ máy Golgi."
        ],
        "answer": "A",
        "solution": "Lục lạp chứa chất diệp lục là bào quan chuyên trách thực hiện quang hợp."
    },
    {
        "id": 2,
        "content": "Cho các yếu tố sau: (1) Ánh sáng; (2) Nước; (3) Hàm lượng carbon dioxide; (4) Mật độ lá trên cây; (5) Nhiệt độ. Số yếu tố môi trường ảnh hưởng đến quang hợp ở thực vật là:",
        "options": [
            "A. 1.",
            "B. 2.",
            "C. 3.",
            "D. 4."
        ],
        "answer": "D",
        "solution": "Có 4 yếu tố môi trường chủ yếu ảnh hưởng trực tiếp đến quang hợp là: Ánh sáng (1), Nước (2), Hàm lượng CO2 (3) và Nhiệt độ (5)."
    },
    {
        "id": 3,
        "content": "Các sản phẩm được tạo ra trong quá trình quang hợp là:",
        "options": [
            "A. Oxygen.",
            "B. Glucose.",
            "C. Carbon dioxide.",
            "D. Cả A và B."
        ],
        "answer": "D",
        "solution": "Sản phẩm của quang hợp gồm đường glucose (chất hữu cơ) và khí oxygen."
    },
    {
        "id": 4,
        "content": "Nguyên liệu chất của quá trình quang hợp gồm:",
        "options": [
            "A. khí oxygen và glucose.",
            "B. glucose và nước.",
            "C. khí carbon dioxide, nước và năng lượng ánh sáng.",
            "D. khí carbon dioxide và nước."
        ],
        "answer": "D",
        "solution": "Hai nguyên liệu vật chất tham gia phản ứng quang hợp là khí carbon dioxide và nước."
    },
    {
        "id": 5,
        "content": "Phương trình chữ của hô hấp tế bào là:",
        "options": [
            "A. Glucose + Oxygen -> Carbon dioxide + Nước + Năng lượng (ATP + nhiệt).",
            "B. Glucose + Carbon dioxide -> Nước + Oxygen + Năng lượng (ATP + nhiệt).",
            "C. Oxygen + Carbon dioxide -> Glucose + Năng lượng (ATP + nhiệt).",
            "D. Glucose + Oxygen -> Nước + Carbon dioxide + Nhiệt."
        ],
        "answer": "A",
        "solution": "Hô hấp tế bào phân giải glucose nhờ oxygen tạo thành CO2, H2O và giải phóng năng lượng ATP cùng nhiệt."
    },
    {
        "id": 6,
        "content": "Khi vận động mạnh như chơi thể thao, lao động nặng,... nhịp hô hấp của cơ thể sẽ tăng lên vì:",
        "options": [
            "A. tốc độ hô hấp tế bào tăng dẫn đến nhu cầu thải oxygen và hấp thụ carbon dioxide tăng.",
            "B. tốc độ hô hấp tế bào tăng dẫn đến nhu cầu thải carbon dioxide và hấp thụ oxygen tăng.",
            "C. tốc độ hô hấp tế bào tăng dẫn đến nhu cầu thải nitrogen và hấp thụ oxygen tăng.",
            "D. tốc độ hô hấp tế bào tăng dẫn đến nhu cầu thải oxygen và hấp thụ nitrogen tăng."
        ],
        "answer": "B",
        "solution": "Cơ bắp cần nhiều năng lượng nên hô hấp tế bào tăng cao, đòi hỏi hấp thụ nhiều O2 và đào thải nhanh khí CO2."
    },
    {
        "id": 7,
        "content": "Trong quá trình hô hấp tế bào, chất hữu cơ bị phân giải thành sản phẩm cuối cùng là:",
        "options": [
            "A. carbon dioxide và nước.",
            "B. carbon dioxide và oxygen.",
            "C. oxygen và nitrogen.",
            "D. oxygen và nước."
        ],
        "answer": "A",
        "solution": "Sản phẩm phân giải cuối cùng là CO2 và H2O."
    },
    {
        "id": 8,
        "content": "Sắp xếp các bước tiến hành thí nghiệm chứng minh hô hấp tế bào hấp thụ khí oxygen và thải khí carbon dioxide:\n(1) Đậy kín bình thủy tinh để chỗ tối 1 ngày;\n(2) Chia đôi hạt vào hai bình C và D;\n(3) Ngâm hạt trong nước ấm 4-12 giờ;\n(4) Bình C: đưa nến đang cháy vào quan sát;\n(5) Bình D: dẫn khí vào cốc nước vôi trong quan sát.\nThứ tự đúng là:",
        "options": [
            "A. 1, 3, 4, 2, 5.",
            "B. 2, 1, 3, 4, 5.",
            "C. 3, 2, 1, 4, 5.",
            "D. 4, 1, 2, 5, 3."
        ],
        "answer": "C",
        "solution": "Quy trình chuẩn: (3) Ngâm kích thích nảy mầm -> (2) Cho vào bình C, D -> (1) Để tối cho hô hấp -> (4) Thử nến ở bình C và (5) Thử nước vôi trong ở bình D."
    }
]

SINH_TF = [
    {
        "id": 1,
        "prompt": "Nhận định đúng / sai trong những câu sau về quang hợp và hô hấp tế bào:",
        "items": [
            {"id": "a", "text": "Quang hợp được thực hiện ở tất cả các bộ phận của cây (rễ, thân, lá).", "ans": False, "reason": "Quang hợp chỉ diễn ra ở các bộ phận có lục lạp chứa diệp lục, chủ yếu là lá."},
            {"id": "b", "text": "Lá là cơ quan chủ yếu thực hiện chức năng quang hợp.", "ans": True, "reason": "Lá có diện tích bề mặt lớn, chứa nhiều lục lạp nhất."},
            {"id": "c", "text": "Cây rau dền đỏ, huyết dụ không quang hợp vì lá không có màu xanh.", "ans": False, "reason": "Lá vẫn chứa chất diệp lục quang hợp bình thường nhưng bị sắc tố đỏ che khuất."},
            {"id": "d", "text": "Toàn bộ năng lượng giải phóng từ hô hấp được tích lũy trong ATP.", "ans": False, "reason": "Một phần năng lượng bị chuyển hóa thành nhiệt năng tỏa ra môi trường."},
            {"id": "e", "text": "Ti thể là bào quan thực hiện hô hấp chủ yếu ở thực vật.", "ans": True, "reason": "Ti thể là nhà máy năng lượng của tế bào, nơi diễn ra hô hấp hiếu khí."}
        ]
    }
]

SINH_SHORT = [
    {
        "id": 1,
        "content": "Nguyên liệu trực tiếp cung cấp năng lượng cho quá trình quang hợp là gì?",
        "answer": "Ánh sáng mặt trời (quang năng)",
        "solution": "Năng lượng ánh sáng mặt trời được chất diệp lục hấp thu để biến đổi thành hóa năng."
    },
    {
        "id": 2,
        "content": "Nhiệt độ thuận lợi nhất cho phần lớn cây xanh quang hợp là bao nhiêu?",
        "answer": "25°C đến 35°C",
        "solution": "Khoảng nhiệt độ tối ưu cho hệ enzym quang hợp ở đa số thực vật là từ *25°C đến 35°C*."
    }
]

# ==============================================================================
# HÀM TẠO FILE WORD
# ==============================================================================

def generate_word_document(mode="student"):
    is_teacher = (mode == "teacher")
    is_screen = (mode == "screen")

    if mode == "teacher":
        filename = "De_Cuong_GK1_KHTN7_TranCaoVan_GiaoVien.docx"
    elif mode == "screen":
        filename = "De_Cuong_GK1_KHTN7_TranCaoVan_DayTrenManHinh.docx"
    else:
        filename = "De_Cuong_GK1_KHTN7_TranCaoVan_HocSinh.docx"

    filepath = os.path.join(OUTPUT_DIR, filename)

    doc = docx.Document()
    setup_header_footer(doc)

    build_title_block(
        doc,
        school_line="TRƯỜNG THCS TRẦN CAO VÂN",
        exam_title="ĐỀ CƯƠNG ÔN TẬP GIỮA HỌC KỲ I",
        subject_line="MÔN: KHOA HỌC TỰ NHIÊN 7 (TRỌN BỘ 5 TRANG)",
        num_pages="05",
        duration="60 phút",
        exam_code="KHTN7-TCV",
        atomic_masses=None,
        is_teacher=is_teacher
    )

    base_size = 11.5 if is_screen else 10.5

    # --------------------------------------------------------------------------
    # CHỦ ĐỀ 1: CHẤT VÀ SỰ BIẾN ĐỔI CHẤT (HÓA HỌC)
    # --------------------------------------------------------------------------
    add_section_header(doc, "KHTN – CHẤT VÀ SỰ BIẾN ĐỔI CHẤT (HÓA HỌC)", "I. KIẾN THỨC CẦN NHỚ:")
    for title, points in HOA_THEORY:
        p_th = doc.add_paragraph()
        p_th.paragraph_format.left_indent = Cm(0.3)
        p_th.paragraph_format.space_before = Pt(3)
        p_th.paragraph_format.space_after = Pt(1)
        render_rich_text(p_th, f"*{title}*", base_size=base_size, base_color=COLOR_PRIMARY)
        for pt in points:
            p_pt = doc.add_paragraph()
            p_pt.paragraph_format.left_indent = Cm(0.6)
            p_pt.paragraph_format.space_before = Pt(1)
            p_pt.paragraph_format.space_after = Pt(1)
            render_rich_text(p_pt, f"• {pt}", base_size=base_size)

    # II. BÀI TẬP MINH HỌA HÓA HỌC
    add_section_header(doc, "II. BÀI TẬP MINH HỌA (HÓA HỌC)\nPhần 1: Trắc nghiệm khách quan nhiều lựa chọn (5 câu)", "Chọn phương án trả lời đúng nhất:")
    for q in HOA_MCQ:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, f"Câu {q['id']}. {q['content']}", base_size=base_size)
        render_options_no_table(doc, q["options"], is_teacher=is_teacher)
        if is_teacher and q.get("solution"):
            add_solution_no_table(doc, f"Đáp án đúng: *{q['answer']}*.\n{q['solution']}")
        elif is_screen:
            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.left_indent = Cm(0.6)
            p_sp.paragraph_format.space_before = Pt(2)
            p_sp.paragraph_format.space_after = Pt(2)
            render_rich_text(p_sp, "✎ Ghi chú / Giải thích: ........................................................................................................", base_size=base_size-1)

    # Phần 2: Đúng Sai Hóa học
    add_section_header(doc, "Phần 2: Trắc nghiệm khách quan \"Đúng – Sai\"", "")
    for q in HOA_TF:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, q["prompt"], base_size=base_size)
        for item in q["items"]:
            p_it = doc.add_paragraph()
            p_it.paragraph_format.left_indent = Cm(0.5)
            p_it.paragraph_format.space_before = Pt(1)
            p_it.paragraph_format.space_after = Pt(1)
            ans_str = "Đúng" if item["ans"] else "Sai"
            if is_teacher:
                render_statement(p_it, f"{item['id']}) {item['text']}", ans=ans_str, is_teacher=True, base_size=base_size)
                p_rea = doc.add_paragraph()
                p_rea.paragraph_format.left_indent = Cm(1.0)
                p_rea.paragraph_format.space_before = Pt(0)
                p_rea.paragraph_format.space_after = Pt(2)
                render_rich_text(p_rea, f"↳ _Giải thích:_ {item['reason']}", base_size=base_size-1, base_italic=True, base_color=COLOR_SOLUTION)
            else:
                render_rich_text(p_it, f"{item['id']}) {item['text']}   [   ] Đúng     [   ] Sai", base_size=base_size)

    # Phần 3: Trả lời ngắn Hóa học
    add_section_header(doc, "Phần 3: Trắc nghiệm khách quan \"Trả lời ngắn\"", "")
    for q in HOA_SHORT:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, f"Câu {q['id']}. {q['content']}", base_size=base_size)
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.left_indent = Cm(0.8)
        p_ans.paragraph_format.space_before = Pt(1)
        p_ans.paragraph_format.space_after = Pt(2)
        if is_teacher:
            render_rich_text(p_ans, f"➔ *Đáp án:* {q['answer']}", base_size=base_size, base_color=COLOR_CORRECT)
            add_solution_no_table(doc, q["solution"])
        elif is_screen:
            render_rich_text(p_ans, "➔ Trả lời: ...................................................................................................................", base_size=base_size)
        else:
            render_rich_text(p_ans, "Trả lời: .....................................................................................................................", base_size=base_size)

    # Phần 4: Tự luận Hóa học
    add_section_header(doc, "Phần 4: Tự luận (Hóa học)", "Học sinh trình bày câu trả lời chi tiết vào bài làm:")
    for q in HOA_ESSAY:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, f"Câu {q['id']}. {q['content']}", base_size=base_size)
        if is_teacher:
            add_solution_no_table(doc, q["solution"])
        elif is_screen:
            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.left_indent = Cm(0.6)
            render_rich_text(p_sp, "✎ Bài giải chi tiết: ........................................................................................................", base_size=base_size)
        else:
            add_student_writing_lines(doc, num_lines=5)

    # --------------------------------------------------------------------------
    # CHỦ ĐỀ 2: KHTN – NĂNG LƯỢNG (VẬT LÍ)
    # --------------------------------------------------------------------------
    add_section_header(doc, "KHTN – NĂNG LƯỢNG (VẬT LÍ)", "I. KIẾN THỨC CẦN NHỚ:")
    for title, points in LY_THEORY:
        p_th = doc.add_paragraph()
        p_th.paragraph_format.left_indent = Cm(0.3)
        p_th.paragraph_format.space_before = Pt(3)
        p_th.paragraph_format.space_after = Pt(1)
        render_rich_text(p_th, f"*{title}*", base_size=base_size, base_color=COLOR_PRIMARY)
        for pt in points:
            p_pt = doc.add_paragraph()
            p_pt.paragraph_format.left_indent = Cm(0.6)
            p_pt.paragraph_format.space_before = Pt(1)
            p_pt.paragraph_format.space_after = Pt(1)
            render_rich_text(p_pt, f"• {pt}", base_size=base_size)

    # II. BÀI TẬP MINH HỌA VẬT LÍ
    add_section_header(doc, "II. BÀI TẬP MINH HỌA (VẬT LÍ)\n1. Trắc nghiệm nhiều lựa chọn (6 câu)", "Chọn phương án trả lời đúng nhất:")
    for q in LY_MCQ:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, f"Câu {q['id']}. {q['content']}", base_size=base_size)
        render_options_no_table(doc, q["options"], is_teacher=is_teacher)
        if is_teacher and q.get("solution"):
            add_solution_no_table(doc, f"Đáp án đúng: *{q['answer']}*.\n{q['solution']}")
        elif is_screen:
            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.left_indent = Cm(0.6)
            p_sp.paragraph_format.space_before = Pt(2)
            p_sp.paragraph_format.space_after = Pt(2)
            render_rich_text(p_sp, "✎ Ghi chú / Giải thích: ........................................................................................................", base_size=base_size-1)

    # Trắc nghiệm đúng sai Vật lí
    add_section_header(doc, "2. Trắc nghiệm đúng sai", "")
    for q in LY_TF:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, q["prompt"], base_size=base_size)
        for item in q["items"]:
            p_it = doc.add_paragraph()
            p_it.paragraph_format.left_indent = Cm(0.5)
            p_it.paragraph_format.space_before = Pt(1)
            p_it.paragraph_format.space_after = Pt(1)
            ans_str = "Đúng" if item["ans"] else "Sai"
            if is_teacher:
                render_statement(p_it, f"{item['id']}) {item['text']}", ans=ans_str, is_teacher=True, base_size=base_size)
                p_rea = doc.add_paragraph()
                p_rea.paragraph_format.left_indent = Cm(1.0)
                p_rea.paragraph_format.space_before = Pt(0)
                p_rea.paragraph_format.space_after = Pt(2)
                render_rich_text(p_rea, f"↳ _Giải thích:_ {item['reason']}", base_size=base_size-1, base_italic=True, base_color=COLOR_SOLUTION)
            else:
                render_rich_text(p_it, f"{item['id']}) {item['text']}   [   ] Đúng     [   ] Sai", base_size=base_size)

    # Tự luận Vật lí
    add_section_header(doc, "3. Phần tự luận (Vật lí - 5 bài)", "Học sinh trình bày lời giải chi tiết và vẽ đồ thị:")
    for q in LY_ESSAY:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, f"Bài {q['id']}. {q['content']}", base_size=base_size)

        if q.get("image"):
            img_path = os.path.join(IMAGES_DIR, q["image"])
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(3)
                p_img.paragraph_format.space_after = Pt(3)
                p_img.add_run().add_picture(img_path, width=Inches(3.8))

        if is_teacher:
            add_solution_no_table(doc, q["solution"])
        elif is_screen:
            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.left_indent = Cm(0.6)
            render_rich_text(p_sp, "✎ Bài giải chi tiết: ........................................................................................................", base_size=base_size)
        else:
            add_student_writing_lines(doc, num_lines=5)

    # --------------------------------------------------------------------------
    # CHỦ ĐỀ 3: KHTN – VẬT SỐNG (SINH HỌC)
    # --------------------------------------------------------------------------
    add_section_header(doc, "KHTN – VẬT SỐNG (SINH HỌC)", "I. KIẾN THỨC CẦN NHỚ:")
    for title, points in SINH_THEORY:
        p_th = doc.add_paragraph()
        p_th.paragraph_format.left_indent = Cm(0.3)
        p_th.paragraph_format.space_before = Pt(3)
        p_th.paragraph_format.space_after = Pt(1)
        render_rich_text(p_th, f"*{title}*", base_size=base_size, base_color=COLOR_PRIMARY)
        for pt in points:
            p_pt = doc.add_paragraph()
            p_pt.paragraph_format.left_indent = Cm(0.6)
            p_pt.paragraph_format.space_before = Pt(1)
            p_pt.paragraph_format.space_after = Pt(1)
            render_rich_text(p_pt, f"• {pt}", base_size=base_size)

    # II. BÀI TẬP MINH HỌA SINH HỌC
    add_section_header(doc, "II. BÀI TẬP MINH HỌA (SINH HỌC)\n1. Trắc nghiệm nhiều lựa chọn (8 câu)", "Chọn phương án trả lời đúng nhất:")
    for q in SINH_MCQ:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, f"Câu {q['id']}. {q['content']}", base_size=base_size)
        render_options_no_table(doc, q["options"], is_teacher=is_teacher)
        if is_teacher and q.get("solution"):
            add_solution_no_table(doc, f"Đáp án đúng: *{q['answer']}*.\n{q['solution']}")
        elif is_screen:
            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.left_indent = Cm(0.6)
            p_sp.paragraph_format.space_before = Pt(2)
            p_sp.paragraph_format.space_after = Pt(2)
            render_rich_text(p_sp, "✎ Ghi chú / Giải thích: ........................................................................................................", base_size=base_size-1)

    # Trắc nghiệm đúng sai Sinh học
    add_section_header(doc, "2. Nhận định đúng / sai trong những câu sau:", "")
    for q in SINH_TF:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, q["prompt"], base_size=base_size)
        for item in q["items"]:
            p_it = doc.add_paragraph()
            p_it.paragraph_format.left_indent = Cm(0.5)
            p_it.paragraph_format.space_before = Pt(1)
            p_it.paragraph_format.space_after = Pt(1)
            ans_str = "Đúng" if item["ans"] else "Sai"
            if is_teacher:
                render_statement(p_it, f"{item['id']}) {item['text']}", ans=ans_str, is_teacher=True, base_size=base_size)
                p_rea = doc.add_paragraph()
                p_rea.paragraph_format.left_indent = Cm(1.0)
                p_rea.paragraph_format.space_before = Pt(0)
                p_rea.paragraph_format.space_after = Pt(2)
                render_rich_text(p_rea, f"↳ _Giải thích:_ {item['reason']}", base_size=base_size-1, base_italic=True, base_color=COLOR_SOLUTION)
            else:
                render_rich_text(p_it, f"{item['id']}) {item['text']}   [   ] Đúng     [   ] Sai", base_size=base_size)

    # Trả lời câu hỏi ngắn Sinh học
    add_section_header(doc, "3. Trả lời câu hỏi ngắn:", "")
    for q in SINH_SHORT:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, f"• {q['content']}", base_size=base_size)
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.left_indent = Cm(0.8)
        p_ans.paragraph_format.space_before = Pt(1)
        p_ans.paragraph_format.space_after = Pt(2)
        if is_teacher:
            render_rich_text(p_ans, f"➔ *Đáp án:* {q['answer']}", base_size=base_size, base_color=COLOR_CORRECT)
            add_solution_no_table(doc, q["solution"])
        elif is_screen:
            render_rich_text(p_ans, "➔ Trả lời: ...................................................................................................................", base_size=base_size)
        else:
            render_rich_text(p_ans, "Trả lời: .....................................................................................................................", base_size=base_size)

    # Dòng kết thúc
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(14)
    p_end.paragraph_format.space_after = Pt(6)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("--- HẾT ĐỀ CƯƠNG ÔN TẬP GIỮA HỌC KỲ I KHTN 7 ---")
    set_run_style(r_end, font_size=11, bold=True, italic=True, color=COLOR_PRIMARY)

    doc.save(filepath)
    print(f"[OK] Đã tạo thành công file Word: {filepath} ({os.path.getsize(filepath):,} bytes)")
    return filepath


# ==============================================================================
# HÀM TẠO FILE LATEX & PDF TRÌNH CHIẾU TIVI 16:9 (KHÔNG CÓ ĐÁP ÁN)
# ==============================================================================

def generate_tivi_presentation_pdf():
    tex_filename = "De_Cuong_GK1_KHTN7_TranCaoVan_DayTrenManHinh.tex"
    pdf_filename = "De_Cuong_GK1_KHTN7_TranCaoVan_DayTrenManHinh.pdf"
    tex_path = os.path.join(SCRATCH_DIR, tex_filename)
    pdf_dest = os.path.join(OUTPUT_DIR, pdf_filename)

    tex_lines = []
    # Preamble chuẩn 16:9 widescreen
    tex_lines.append(r"""\documentclass{article}
\usepackage[utf8]{vietnam}
\usepackage{times}
\usepackage{amsmath,amssymb}
\usepackage{tabularx}
\usepackage{array}
\usepackage{booktabs}
\usepackage{graphicx}
\usepackage{xcolor}
\usepackage[
  paperwidth=29.7cm,
  paperheight=16.706cm,
  top=0.9cm,
  bottom=0.9cm,
  left=1.2cm,
  right=1.2cm,
  footskip=0pt,
  headheight=0pt,
  headsep=0pt
]{geometry}

\pagestyle{empty}
\setlength{\parindent}{0pt}
\setlength{\parskip}{0pt}

\definecolor{brandRed}{RGB}{211, 47, 47}
\definecolor{brandGreen}{RGB}{46, 125, 50}
\definecolor{brandBlue}{RGB}{21, 101, 192}
\definecolor{textColor}{RGB}{20, 20, 20}

\color{textColor}

\begin{document}
""")

    # Slide 0: Tiêu đề
    tex_lines.append(r"""
% SLIDE 0: TIÊU ĐỀ
\begin{center}
\vspace*{0.2cm}
{\fontsize{24}{30}\selectfont \textbf{\color{brandRed}TRƯỜNG THCS TRẦN CAO VÂN -- NĂM HỌC 2026 -- 2027}}\\[10pt]
{\fontsize{34}{42}\selectfont \textbf{\color{brandBlue}ĐỀ CƯƠNG ÔN TẬP KIỂM TRA GIỮA HỌC KỲ I}}\\[8pt]
{\fontsize{28}{36}\selectfont \textbf{\color{brandGreen}MÔN: KHOA HỌC TỰ NHIÊN 7}}\\[8pt]
{\fontsize{22}{28}\selectfont \textit{(Chất \& biến đổi chất -- Năng lượng -- Vật sống)}}\\[16pt]
\rule{16cm}{1.5pt}\\[20pt]
{\fontsize{26}{34}\selectfont \textbf{Giáo viên hướng dẫn: Thầy TRẦN VĂN THẠNH -- 0777.470.803}}\\[8pt]
{\fontsize{22}{28}\selectfont \textit{Cơ sở 1: 6/15 Nguyễn Hoàng, Kim Long | Cơ sở 2: 24 Đặng Thái Thân, TP Huế}}\\[16pt]
{\fontsize{24}{30}\selectfont \textbf{TỔNG HỢP TOÀN BỘ CÂU HỎI ÔN TẬP VÀ HƯỚNG DẪN DẠY MÀN HÌNH}}
\end{center}
\newpage
""")

    # Slide 1: Hóa học Lý thuyết
    tex_lines.append(r"""
% SLIDE 1: HÓA HỌC LÝ THUYẾT
{\fontsize{25}{32}\selectfont \textbf{\color{brandBlue}KHTN -- CHẤT VÀ SỰ BIẾN ĐỔI CHẤT (HÓA HỌC) -- LÝ THUYẾT}}
\vspace{8pt}

\fontsize{21}{28}\selectfont
\textbf{1. Cấu tạo nguyên tử:}
\begin{itemize}
  \setlength{\itemsep}{2pt}
  \item Hạt nhân: ở tâm, mang điện tích dương, gồm hạt proton ($+1$) và neutron (không mang điện).
  \item Vỏ nguyên tử: gồm các electron mang điện tích âm ($-1$) chuyển động rất nhanh quanh hạt nhân theo từng lớp.
  \item Trong nguyên tử: \textbf{Số proton = Số electron}. Khối lượng nguyên tử đo bằng $\text{amu}$ ($1\text{ amu} \approx 1,6605 \times 10^{-24}\text{ g}$).
\end{itemize}

\vspace{6pt}
\textbf{2. Nguyên tố hóa học:}
\begin{itemize}
  \setlength{\itemsep}{2pt}
  \item Là tập hợp các nguyên tử cùng loại có cùng số proton trong hạt nhân.
  \item Kí hiệu hóa học: chữ cái đầu viết hoa, chữ thứ hai (nếu có) viết thường.
\end{itemize}
\newpage
""")

    # Slide 2: Hóa học MCQ 1 & 2
    tex_lines.append(r"""
% SLIDE 2: HÓA HỌC MCQ 1 & 2
\fontsize{26}{35}\selectfont
\textbf{Câu 1.} Các hạt cấu tạo nên hạt nhân của hầu hết các nguyên tử là:
\vspace{4pt}

\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{A.} electron và neutron. & \textbf{B.} proton và neutron.\\[6pt]
\textbf{C.} neutron và electron. & \textbf{D.} electron, proton và neutron.
\end{tabularx}

\vspace{18pt}

\textbf{Câu 2.} Các hạt cấu tạo nên hầu hết các nguyên tử là:
\vspace{4pt}

\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{A.} electron và neutron. & \textbf{B.} proton và neutron.\\[6pt]
\textbf{C.} neutron và electron. & \textbf{D.} electron, proton và neutron.
\end{tabularx}
\newpage
""")

    # Slide 3: Hóa học MCQ 3, 4, 5
    tex_lines.append(r"""
% SLIDE 3: HÓA HỌC MCQ 3, 4, 5
\fontsize{22}{30}\selectfont
\textbf{Câu 3.} Nguyên tử Carbon có khối lượng là $1,9926 \times 10^{-23}\text{ g}$. Khối lượng tính theo $\text{amu}$ là? ($1\text{ amu} = 1,6605 \times 10^{-24}\text{ g}$)
\begin{tabularx}{\linewidth}{@{}X X X X@{}}
\textbf{A.} $6\text{ amu}$. & \textbf{B.} $12\text{ amu}$. & \textbf{C.} $24\text{ amu}$. & \textbf{D.} $48\text{ amu}$.
\end{tabularx}

\vspace{10pt}
\textbf{Câu 4.} Thành phần các nguyên tử: $\text{X}(17p, 17e, 16n)$, $\text{Y}(20p, 20e, 19n)$, $\text{Z}(17p, 17e, 18n)$, $\text{T}(19p, 19e, 20n)$. Có bao nhiêu nguyên tố hóa học?
\begin{tabularx}{\linewidth}{@{}X X X X@{}}
\textbf{A.} 1. & \textbf{B.} 2. & \textbf{C.} 3. & \textbf{D.} 4.
\end{tabularx}

\vspace{10pt}
\textbf{Câu 5.} Tên của nguyên tố hóa học có kí hiệu $\text{Fe}$ là:
\begin{tabularx}{\linewidth}{@{}X X X X@{}}
\textbf{A.} Fluorine. & \textbf{B.} Iron. & \textbf{C.} Phosphorus. & \textbf{D.} Potassium.
\end{tabularx}
\newpage
""")

    # Slide 4: Hóa học Đúng/Sai Câu 6
    tex_lines.append(r"""
% SLIDE 4: HÓA HỌC ĐÚNG SAI CÂU 6
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}HÓA HỌC -- TRẮC NGHIỆM ĐÚNG SAI (CÂU 6)}}
\vspace{8pt}

\fontsize{22}{30}\selectfont
\textbf{Câu 6.} Cho các nhận định về nguyên tố và kí hiệu hóa học, xác định Đúng hoặc Sai:

\vspace{12pt}
\textbf{a)} Nguyên tố hóa học là tập hợp các nguyên tử có cùng số proton trong hạt nhân.\hfill $[\quad]$\\[10pt]
\textbf{b)} Kí hiệu hóa học của nguyên tố nhôm là $\text{Al}$, của đồng là $\text{Cu}$.\hfill $[\quad]$\\[10pt]
\textbf{c)} Khối lượng nguyên tử tính theo đơn vị gam.\hfill $[\quad]$\\[10pt]
\textbf{d)} Các nguyên tử cùng một nguyên tố có thể có số neutron khác nhau nhưng proton luôn giống nhau.\hfill $[\quad]$
\newpage
""")

    # Slide 5: Hóa học Trả lời ngắn & Tự luận
    tex_lines.append(r"""
% SLIDE 5: HÓA HỌC TRẢ LỜI NGẮN & TỰ LUẬN
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}HÓA HỌC -- TRẢ LỜI NGẮN \& BÀI TẬP TỰ LUẬN}}
\vspace{8pt}

\fontsize{21}{29}\selectfont
\textbf{Câu 7.} Nguyên tử $\text{X}$ có tổng số hạt là 13, số hạt mang điện nhiều hơn không mang điện là 3. Số proton của $\text{X}$ là: $\dots\dots\dots\dots\dots$

\vspace{14pt}
\textbf{Câu 8.} Nguyên tử $\text{Aluminium}$ có $13p$ trong hạt nhân:
\begin{itemize}
  \setlength{\itemsep}{3pt}
  \item[a)] Số electron ở vỏ nguyên tử là: $\dots\dots\dots\dots\dots$
  \item[b)] Sơ đồ phân bố electron: $\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots$
  \item[c)] Số lớp electron: $\dots\dots\dots$; \quad Số electron lớp ngoài cùng: $\dots\dots\dots$
\end{itemize}

\vspace{14pt}
\textbf{Câu 9.} Viết kí hiệu hóa học: $\text{Carbon}$, $\text{Nitrogen}$, $\text{Calcium}$, $\text{Magnesium}$, $\text{Aluminium}$, $\text{Chlorine}$, $\text{Zinc}$, $\text{Mercury}$.
\newpage
""")

    # Slide 6: Vật lí Lý thuyết
    tex_lines.append(r"""
% SLIDE 6: VẬT LÍ LÝ THUYẾT
{\fontsize{25}{32}\selectfont \textbf{\color{brandBlue}KHTN -- NĂNG LƯỢNG (VẬT LÍ) -- LÝ THUYẾT TRỌNG TÂM}}
\vspace{8pt}

\fontsize{21}{28}\selectfont
\textbf{1. Tốc độ chuyển động:}
\begin{itemize}
  \setlength{\itemsep}{2pt}
  \item Đặc trưng cho mức độ nhanh chậm của chuyển động: $v = \dfrac{s}{t}$ ($s$: quãng đường; $t$: thời gian).
  \item Đơn vị hợp pháp: mét trên giây ($\text{m/s}$) hoặc kilômét trên giờ ($\text{km/h}$).
\end{itemize}

\vspace{6pt}
\textbf{2. Đo tốc độ & Thiết bị bắn tốc độ:}
\begin{itemize}
  \setlength{\itemsep}{2pt}
  \item Đo bằng đồng hồ bấm giây hoặc cổng quang điện kết hợp đồng hồ hiện số.
  \item Thiết bị bắn tốc độ ghi hình ảnh và đo tốc độ tức thời của phương tiện giao thông.
\end{itemize}
\newpage
""")

    # Slide 7: Vật lí MCQ 1 & 2
    tex_lines.append(r"""
% SLIDE 7: VẬT LÍ MCQ 1 & 2
\fontsize{26}{35}\selectfont
\textbf{Câu 1.} Công thức nào sau đây dùng để tính tốc độ chuyển động?
\vspace{4pt}

\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{A.} $v = s \times t$. & \textbf{B.} $v = s / t$.\\[6pt]
\textbf{C.} $s = v / t$. & \textbf{D.} $t = v / s$.
\end{tabularx}

\vspace{18pt}

\textbf{Câu 2.} Một đoàn tàu hỏa đi từ A đến B cách nhau $30\text{ km}$ trong $45\text{ phút}$. Tốc độ đoàn tàu là:
\vspace{4pt}

\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{A.} $60\text{ km/h}$. & \textbf{B.} $40\text{ km/h}$.\\[6pt]
\textbf{C.} $50\text{ km/h}$. & \textbf{D.} $55\text{ km/h}$.
\end{tabularx}
\newpage
""")

    # Slide 8: Vật lí MCQ 3, 4, 5, 6
    tex_lines.append(r"""
% SLIDE 8: VẬT LÍ MCQ 3, 4, 5, 6
\fontsize{22}{30}\selectfont
\textbf{Câu 3.} Đại lượng nào đặc trưng cho sự nhanh, chậm của chuyển động?
\begin{tabularx}{\linewidth}{@{}X X X X@{}}
\textbf{A.} Quãng đường. & \textbf{B.} Thời gian. & \textbf{C.} Tốc độ. & \textbf{D.} Nhiệt độ.
\end{tabularx}

\vspace{8pt}
\textbf{Câu 4.} Đơn vị đo tốc độ là:
\begin{tabularx}{\linewidth}{@{}X X X X@{}}
\textbf{A.} $\text{m}$. & \textbf{B.} $\text{s/m}$. & \textbf{C.} $\text{m/s}$. & \textbf{D.} $\text{m}\cdot\text{s}$.
\end{tabularx}

\vspace{8pt}
\textbf{Câu 5.} Đơn vị đo tốc độ phụ thuộc vào:
\begin{tabularx}{\linewidth}{@{}X X X X@{}}
\textbf{A.} đơn vị độ dài. & \textbf{B.} đơn vị thời gian. & \textbf{C.} cả A và B. & \textbf{D.} Cả A, B sai.
\end{tabularx}

\vspace{8pt}
\textbf{Câu 6.} An đạp xe đến trường mất $10\text{ phút}$ với tốc độ $2\text{ m/s}$. Quãng đường là:
\begin{tabularx}{\linewidth}{@{}X X X X@{}}
\textbf{A.} $120\text{ m}$. & \textbf{B.} $1,2\text{ m}$. & \textbf{C.} $12\text{ km}$. & \textbf{D.} $1,2\text{ km}$.
\end{tabularx}
\newpage
""")

    # Slide 9: Vật lí Đúng/Sai
    tex_lines.append(r"""
% SLIDE 9: VẬT LÍ ĐÚNG SAI
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}VẬT LÍ -- TRẮC NGHIỆM ĐÚNG SAI (ĐO TỐC ĐỘ)}}
\vspace{8pt}

\fontsize{21}{29}\selectfont
\textbf{Câu 1.} Xác định tính Đúng hoặc Sai của mỗi nhận định sau:

\vspace{10pt}
\textbf{a)} Đo tốc độ xe đồ chơi bằng đồng hồ bấm giây cần dùng thước đo quãng đường và đồng hồ đo thời gian.\hfill $[\quad]$\\[8pt]
\textbf{b)} Dùng đồng hồ hiện số và cổng quang, đồng hồ chỉ bắt đầu đếm khi xe đã đi qua cả hai cổng quang.\hfill $[\quad]$\\[8pt]
\textbf{c)} Thiết bị bắn tốc độ có công dụng ghi hình ảnh và đo tốc độ phương tiện kiểm tra an toàn giao thông.\hfill $[\quad]$\\[8pt]
\textbf{d)} Thiết bị bắn tốc độ bắt buộc phải gắn cố định cổng quang điện trên mặt đường.\hfill $[\quad]$
\newpage
""")

    # Slide 10: Vật lí Tự luận Bài 1 & 2 & 3
    tex_lines.append(r"""
% SLIDE 10: VẬT LÍ TỰ LUẬN 1, 2, 3
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}VẬT LÍ -- BÀI TẬP TỰ LUẬN (BÀI 1, 2, 3)}}
\vspace{6pt}

\fontsize{20}{28}\selectfont
\textbf{Bài 1.} Đổi các đơn vị sau:
\begin{tabularx}{\linewidth}{@{}X X X@{}}
a) $125\text{ cm} = \dots\dots\text{m}$ & b) $5,03\text{ km} = \dots\dots\text{m}$ & c) $250\text{ m} = \dots\dots\text{km}$\\[4pt]
d) $1\text{h}20\text{p} = \dots\dots\text{s}$ & e) $15\text{ km/h} = \dots\dots\text{m/s}$ & f) $15\text{ m/s} = \dots\dots\text{km/h}$
\end{tabularx}

\vspace{10pt}
\textbf{Bài 2.} Tốc độ đi xe đạp của bạn An là $15\text{ km/h}$, con số này cho biết điều gì?

\vspace{10pt}
\textbf{Bài 3.} Ô tô chạy qua 2 vạch mốc cách nhau $5\text{ m}$ hết $0,2\text{ s}$. Tốc độ giới hạn là $22\text{ m/s}$. Xe có vượt quá tốc độ cho phép không?
\newline \textbf{Trả lời:} $\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots$
\newpage
""")

    # Slide 11: Vật lí Tự luận Bài 4 & 5 (Đồ thị ô tô)
    tex_lines.append(r"""
% SLIDE 11: VẬT LÍ BÀI 4 & 5 (ĐỒ THỊ)
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}VẬT LÍ -- BÀI TẬP 4 \& 5 (ĐỒ THỊ QUÃNG ĐƯỜNG - THỜI GIAN)}}
\vspace{4pt}

\begin{minipage}[c]{0.54\linewidth}
\fontsize{19}{26}\selectfont
\textbf{Bài 4.} Hùng đạp xe lên dốc $150\text{ m}$ với tốc độ $2\text{ m/s}$, xuống dốc $120\text{ m}$ hết $30\text{ s}$. Tốc độ trung bình là bao nhiêu?

\vspace{8pt}
\textbf{Bài 5.} Bảng số liệu ô tô:
\begin{center}
\begin{tabular}{|l|c|c|c|c|c|c|c|}
\hline
$t\text{ (h)}$ & 0 & 1 & 2 & 3 & 4 & 5 & 6\\
\hline
$s\text{ (km)}$ & 0 & 60 & 120 & 180 & 180 & 220 & 260\\
\hline
\end{tabular}
\end{center}
Vẽ đồ thị và nhận xét giai đoạn dừng lại của xe.
\end{minipage}
\hfill
\begin{minipage}[c]{0.44\linewidth}
\centering
\includegraphics[height=11.0cm,keepaspectratio]{San_Pham/images_khtn7_trancaovan/dothi_oto_bai5.png}
\end{minipage}
\newpage
""")

    # Slide 12: Sinh học Lý thuyết
    tex_lines.append(r"""
% SLIDE 12: SINH HỌC LÝ THUYẾT
{\fontsize{25}{32}\selectfont \textbf{\color{brandBlue}KHTN -- VẬT SỐNG (SINH HỌC) -- LÝ THUYẾT TRỌNG TÂM}}
\vspace{8pt}

\fontsize{21}{28}\selectfont
\textbf{1. Quang hợp ở thực vật:}
\begin{itemize}
  \setlength{\itemsep}{2pt}
  \item Diễn ra chủ yếu ở lục lạp của lá cây.
  \item Phương trình chữ: $\text{Nước} + \text{Carbon dioxide} \xrightarrow[\text{Diệp lục}]{\text{Ánh sáng}} \text{Glucose} + \text{Oxygen}$.
\end{itemize}

\vspace{6pt}
\textbf{2. Hô hấp tế bào:}
\begin{itemize}
  \setlength{\itemsep}{2pt}
  \item Diễn ra chủ yếu ở ti thể của tế bào.
  \item Phương trình chữ: $\text{Glucose} + \text{Oxygen} \rightarrow \text{Carbon dioxide} + \text{Nước} + \text{Năng lượng (ATP + nhiệt)}$.
\end{itemize}
\newpage
""")

    # Slide 13: Sinh học MCQ 1, 2, 3, 4
    tex_lines.append(r"""
% SLIDE 13: SINH HỌC MCQ 1, 2, 3, 4
\fontsize{22}{30}\selectfont
\textbf{Câu 1.} Quá trình quang hợp ở thực vật diễn ra chủ yếu trong bào quan nào?
\begin{tabularx}{\linewidth}{@{}X X X X@{}}
\textbf{A.} Lục lạp. & \textbf{B.} Ti thể. & \textbf{C.} Ribosome. & \textbf{D.} Bộ máy Golgi.
\end{tabularx}

\vspace{8pt}
\textbf{Câu 2.} Số yếu tố môi trường ảnh hưởng đến quang hợp ở thực vật là:
\begin{tabularx}{\linewidth}{@{}X X X X@{}}
\textbf{A.} 1. & \textbf{B.} 2. & \textbf{C.} 3. & \textbf{D.} 4.
\end{tabularx}

\vspace{8pt}
\textbf{Câu 3.} Các sản phẩm được tạo ra trong quá trình quang hợp là:
\begin{tabularx}{\linewidth}{@{}X X X X@{}}
\textbf{A.} Oxygen. & \textbf{B.} Glucose. & \textbf{C.} Carbon dioxide. & \textbf{D.} Cả A và B.
\end{tabularx}

\vspace{8pt}
\textbf{Câu 4.} Nguyên liệu chất của quá trình quang hợp gồm:
\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{A.} khí oxygen và glucose. & \textbf{B.} glucose và nước.\\[4pt]
\textbf{C.} carbon dioxide, nước và ánh sáng. & \textbf{D.} khí carbon dioxide và nước.
\end{tabularx}
\newpage
""")

    # Slide 14: Sinh học MCQ 5, 6, 7
    tex_lines.append(r"""
% SLIDE 14: SINH HỌC MCQ 5, 6, 7
\fontsize{22}{30}\selectfont
\textbf{Câu 5.} Phương trình hô hấp tế bào là:
\vspace{2pt}
\begin{itemize}
  \setlength{\itemsep}{2pt}
  \item[\textbf{A.}] $\text{Glucose} + \text{Oxygen} \rightarrow \text{Carbon dioxide} + \text{Nước} + \text{Năng lượng (ATP + nhiệt)}$.
  \item[\textbf{B.}] $\text{Glucose} + \text{Carbon dioxide} \rightarrow \text{Nước} + \text{Oxygen} + \text{Năng lượng}$.
  \item[\textbf{C.}] $\text{Oxygen} + \text{Carbon dioxide} \rightarrow \text{Glucose} + \text{Năng lượng}$.
  \item[\textbf{D.}] $\text{Glucose} + \text{Oxygen} \rightarrow \text{Nước} + \text{Carbon dioxide} + \text{Nhiệt}$.
\end{itemize}

\vspace{6pt}
\textbf{Câu 6.} Khi vận động mạnh, nhịp hô hấp của cơ thể tăng lên vì:
\begin{itemize}
  \item[\textbf{A.}] Hô hấp tế bào tăng $\rightarrow$ nhu cầu thải oxygen và hấp thụ carbon dioxide tăng.
  \item[\textbf{B.}] Hô hấp tế bào tăng $\rightarrow$ nhu cầu thải carbon dioxide và hấp thụ oxygen tăng.
\end{itemize}

\vspace{6pt}
\textbf{Câu 7.} Hô hấp tế bào phân giải chất hữu cơ thành sản phẩm cuối cùng là:
\begin{tabularx}{\linewidth}{@{}X X X X@{}}
\textbf{A.} $\text{CO}_2$ và $\text{H}_2\text{O}$. & \textbf{B.} $\text{CO}_2$ và $\text{O}_2$. & \textbf{C.} $\text{O}_2$ và $\text{N}_2$. & \textbf{D.} $\text{O}_2$ và $\text{H}_2\text{O}$.
\end{tabularx}
\newpage
""")

    # Slide 15: Sinh học MCQ 8 & Đúng/Sai & Trả lời ngắn
    tex_lines.append(r"""
% SLIDE 15: SINH HỌC MCQ 8 & ĐÚNG SAI
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}SINH HỌC -- THÍ NGHIỆM HÔ HẤP \& CÂU HỎI NGẮN}}
\vspace{6pt}

\fontsize{20}{27}\selectfont
\textbf{Câu 8.} Sắp xếp tiến trình thí nghiệm hô hấp tế bào ở hạt nảy mầm:
\begin{tabularx}{\linewidth}{@{}X X X X@{}}
\textbf{A.} 1, 3, 4, 2, 5. & \textbf{B.} 2, 1, 3, 4, 5. & \textbf{C.} 3, 2, 1, 4, 5. & \textbf{D.} 4, 1, 2, 5, 3.
\end{tabularx}

\vspace{8pt}
\textbf{Nhận định Đúng / Sai:}
\begin{itemize}
  \setlength{\itemsep}{2pt}
  \item[1.] Quang hợp thực hiện ở tất cả các bộ phận của cây.\hfill $[\quad]$
  \item[2.] Lá là cơ quan chủ yếu thực hiện chức năng quang hợp.\hfill $[\quad]$
  \item[3.] Cây rau dền đỏ không quang hợp vì lá không có màu xanh.\hfill $[\quad]$
  \item[4.] Ti thể là bào quan thực hiện hô hấp chủ yếu ở thực vật.\hfill $[\quad]$
\end{itemize}

\vspace{6pt}
\textbf{Trả lời ngắn:}
\begin{itemize}
  \setlength{\itemsep}{2pt}
  \item[•] Nguồn năng lượng trực tiếp cho quang hợp: $\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots$
  \item[•] Nhiệt độ thuận lợi nhất cho phần lớn cây xanh quang hợp: $\dots\dots\dots\dots\dots$
\end{itemize}

\end{document}
""")

    with open(tex_path, "w", encoding="utf-8") as f:
        f.write("".join(tex_lines))
    print(f"[OK] Đã ghi file LaTeX: {tex_path}")

    # Biên dịch pdflatex 2 lần để ổn định layout
    cmd = ["pdflatex", "-interaction=nonstopmode", "-output-directory", SCRATCH_DIR, tex_path]
    res1 = subprocess.run(cmd, capture_output=True, text=True)
    res2 = subprocess.run(cmd, capture_output=True, text=True)

    pdf_scratch = os.path.join(SCRATCH_DIR, pdf_filename)
    if os.path.exists(pdf_scratch):
        shutil.copy2(pdf_scratch, pdf_dest)
        print(f"[OK] Đã biên dịch và copy PDF thành công: {pdf_dest} ({os.path.getsize(pdf_dest):,} bytes)")
        return pdf_dest
    else:
        print("[ERROR] Không tìm thấy file PDF được tạo. Log:")
        print(res2.stdout[-1500:])
        return None


# ==============================================================================
# HÀM MAIN
# ==============================================================================

def main():
    print("=" * 75)
    print("BẮT ĐẦU TẠO 3 FILE CHO ĐỀ CƯƠNG ÔN TẬP GIỮA KỲ I KHTN 7 (THCS TRẦN CAO VÂN)")
    print("1. FILE HỌC SINH (.docx)")
    print("2. FILE GIÁO VIÊN (.docx có lời giải chi tiết)")
    print("3. FILE DẠY TRÊN MÀN HÌNH (Slide PDF 16:9 widescreen cho Tivi 65 inch)")
    print("=" * 75)

    f_hs = generate_word_document(mode="student")
    f_gv = generate_word_document(mode="teacher")
    f_screen_docx = generate_word_document(mode="screen")
    f_screen_pdf = generate_tivi_presentation_pdf()

    all_files = [f_hs, f_gv, f_screen_docx, f_screen_pdf]

    # Đồng bộ sang các thư mục OneDrive
    print("\n[+] Đồng bộ các file sang các thư mục OneDrive:")
    for od in TARGET_ONEDRIVE_DIRS:
        if os.path.exists(od):
            for src_path in all_files:
                if src_path and os.path.exists(src_path):
                    dest = os.path.join(od, os.path.basename(src_path))
                    try:
                        shutil.copy2(src_path, dest)
                        print(f"    -> Đã copy {os.path.basename(src_path)} vào: {od}")
                    except Exception as e:
                        print(f"    [!] Cảnh báo file đang mở hoặc bận tại {od}: {e}")

    print("\n" + "=" * 75)
    print("HOÀN THÀNH TẤT CẢ 3 FILE THEO ĐÚNG CHUẨN FORM ĐỀ KIỂM TRA!")
    print("=" * 75)

if __name__ == "__main__":
    main()
