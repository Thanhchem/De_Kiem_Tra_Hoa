import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

CROP_DIR = r"c:\Antigravity_Thanh\San_Pham\images_crop_de01"
OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"

COLOR_HF_RED = RGBColor(211, 47, 47)      # Coral Red (#D32F2F)
COLOR_HF_GREEN_HEX = "8BC390"             # Sage Green (#8BC390)
COLOR_PRIMARY = RGBColor(46, 125, 50)     # Dark Green (#2E7D32)
COLOR_CORRECT = RGBColor(198, 40, 40)     # Red for correct answers (#C62828)
COLOR_SOLUTION = RGBColor(21, 101, 192)   # Blue for solutions (#1565C0)
COLOR_GRAY = RGBColor(97, 97, 97)

def add_p_border_bottom(p, color="8BC390", sz="8"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_p_border_top(p, color="8BC390", sz="8"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def set_run_style(run, font_name="Times New Roman", font_size=12, bold=False, italic=False, subscript=False, superscript=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def render_rich_text(p, text, base_font="Times New Roman", base_size=12, base_bold=False, base_italic=False, base_color=None):
    tokens = re.split(r'(~[^~]+~|\^[^\^]+\^|\*[^\*]+\*|_[^_]+_)', text)
    for token in tokens:
        if not token:
            continue
        is_bold = base_bold
        is_italic = base_italic
        is_sub = False
        is_sup = False
        content = token

        if token.startswith('*') and token.endswith('*'):
            is_bold = True
            content = token[1:-1]
        elif token.startswith('_') and token.endswith('_'):
            is_italic = True
            content = token[1:-1]
        elif token.startswith('~') and token.endswith('~'):
            is_sub = True
            content = token[1:-1]
        elif token.startswith('^') and token.endswith('^'):
            is_sup = True
            content = token[1:-1]

        run = p.add_run(content)
        run_color = COLOR_CORRECT if (is_bold and not base_bold and not base_color) else base_color
        set_run_style(run, font_name=base_font, font_size=base_size, bold=is_bold, italic=is_italic,
                      subscript=is_sub, superscript=is_sup, color=run_color)

def setup_header_footer(doc):
    normal_style = doc.styles['Normal']
    normal_style.paragraph_format.tab_stops.clear_all()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    section.start_type = WD_SECTION_START.NEW_PAGE
    section.different_first_page_header_footer = False
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    # 1. Header (bảng 1 hàng 2 cột, rộng 18.0 cm)
    header = section.header
    for p in header.paragraphs:
        p.text = ""
    tbl_h = header.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl_h.cell(0, 0)
    c_right = tbl_h.cell(0, 1)
    c_left.width = Cm(8.0)
    c_right.width = Cm(10.0)

    p_hl = c_left.paragraphs[0]
    p_hl.paragraph_format.space_before = Pt(0)
    p_hl.paragraph_format.space_after = Pt(2)
    p_hl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_hl = p_hl.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)

    p_hr = c_right.paragraphs[0]
    p_hr.paragraph_format.space_before = Pt(0)
    p_hr.paragraph_format.space_after = Pt(2)
    p_hr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p_hr.add_run("Thầy ")
    set_run_style(r1, font_size=10, color=COLOR_HF_RED)
    r2 = p_hr.add_run("TRẦN VĂN THẠNH")
    set_run_style(r2, font_size=10, bold=True, color=COLOR_HF_RED)
    r3 = p_hr.add_run(" - 0777.470.803")
    set_run_style(r3, font_size=10, color=COLOR_HF_RED)

    add_p_border_bottom(p_hl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_bottom(p_hr, color=COLOR_HF_GREEN_HEX, sz="6")

    # 2. Footer (bảng 1 hàng 2 cột, rộng 18.0 cm)
    footer = section.footer
    for p in footer.paragraphs:
        p.text = ""
    tbl_f = footer.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_fl = tbl_f.cell(0, 0)
    c_fr = tbl_f.cell(0, 1)
    c_fl.width = Cm(14.5)
    c_fr.width = Cm(3.5)

    p_fl = c_fl.paragraphs[0]
    p_fl.paragraph_format.space_before = Pt(2)
    p_fl.paragraph_format.space_after = Pt(0)
    p_fl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_fl = p_fl.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)

    p_fr = c_fr.paragraphs[0]
    p_fr.paragraph_format.space_before = Pt(2)
    p_fr.paragraph_format.space_after = Pt(0)
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_field(p_fr, font_size=10, bold=True, color=COLOR_HF_RED)

    add_p_border_top(p_fl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_top(p_fr, color=COLOR_HF_GREEN_HEX, sz="6")

def build_title_block_no_table(doc, is_teacher=False):
    """Tiêu đề đầu đề thi trình bày tab stop thanh lịch, không viền rườm rà."""
    headers_data = [
        ("SỞ GIÁO DỤC VÀ ĐÀO TẠO THỪA THIÊN HUẾ", "ĐỀ KIỂM TRA ĐỊNH KỲ CHƯƠNG 1 (ESTER - LIPID)", 10, True, 11, True),
        ("ĐỀ CHÍNH THỨC", "MÔN: HOÁ HỌC - LỚP 12", 10.5, True, 11, True),
        ("(Đề thi có 04 trang)", "Thời gian: 50 phút | MÃ ĐỀ THI: 101", 9.5, False, 10, True)
    ]
    for l_text, r_text, l_size, l_bold, r_size, r_bold in headers_data:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)

        rl = p.add_run(l_text)
        set_run_style(rl, font_size=l_size, bold=l_bold, italic=(not l_bold))

        p.add_run('\t')
        rr = p.add_run(r_text)
        set_run_style(rr, font_size=r_size, bold=r_bold, italic=(not r_bold))

    if is_teacher:
        p_gv = doc.add_paragraph()
        p_gv.paragraph_format.space_before = Pt(2)
        p_gv.paragraph_format.space_after = Pt(2)
        p_gv.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rgv = p_gv.add_run("(BẢN GIÁO VIÊN -- CÓ LỜI GIẢI CHI TIẾT)")
        set_run_style(rgv, font_size=11, bold=True, color=COLOR_CORRECT)

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(4)
    p_info.paragraph_format.space_after = Pt(2)
    render_rich_text(p_info, "*Họ, tên thí sinh:* ............................................................ *Lớp:* ............. *Số báo danh:* .............", base_size=11)

    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(0)
    p_note.paragraph_format.space_after = Pt(4)
    render_rich_text(p_note, "_(Cho biết nguyên tử khối: H = 1; C = 12; N = 14; O = 16; Na = 23; K = 39; Br = 80)._", base_size=10, base_italic=True)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(4)
    add_p_border_bottom(p_div, color="CCCCCC", sz="4")

def add_section_header(doc, title, subtitle=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, f"*{title}*", base_size=11.5, base_bold=True, base_color=COLOR_PRIMARY)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(4)
        render_rich_text(p2, f"_{subtitle}_", base_size=10.5, base_italic=True)

def render_question_prompt(p, text, base_size=10.5):
    """In đậm tiền tố 'Câu X.' ở đầu câu hỏi, phần nội dung câu hỏi sau đó giữ bình thường."""
    m = re.match(r'^(Câu\s+\d+\.)\s*(.*)$', text.strip(), re.DOTALL)
    if m:
        lbl = m.group(1)
        rest = m.group(2)
        r_lbl = p.add_run(lbl + " ")
        set_run_style(r_lbl, font_size=base_size, bold=True)
        render_rich_text(p, rest, base_size=base_size, base_bold=False)
    else:
        render_rich_text(p, text, base_size=base_size, base_bold=False)

def render_statement(p, stmt, ans=None, is_teacher=False, base_size=10.5):
    """In đậm tiền tố 'a)', 'b)', 'c)', 'd)' ở Phần II."""
    m = re.match(r'^([a-d]\))\s*(.*)$', stmt.strip())
    if m:
        lbl = m.group(1)
        body = m.group(2)
        r_lbl = p.add_run(lbl + " ")
        set_run_style(r_lbl, font_size=base_size, bold=True)
        render_rich_text(p, body, base_size=base_size, base_bold=False)
    else:
        render_rich_text(p, stmt, base_size=base_size, base_bold=False)

    if is_teacher and ans:
        p.add_run("  ➔ ")
        r_ans = p.add_run(f"({ans})")
        set_run_style(r_ans, font_size=base_size, bold=True, color=COLOR_CORRECT)

def render_single_option(p, opt_str, is_teacher=False, base_size=10.5):
    """
    In đậm nhãn A., B., C., D. cho mọi phương án.
    Nếu là bản giáo viên và phương án đúng: in đậm và tô màu đỏ cho cả nhãn và nội dung (kèm sub/sup).
    """
    clean_str = opt_str.strip()
    is_correct = False
    if clean_str.startswith('*') and clean_str.endswith('*'):
        is_correct = True
        clean_str = clean_str[1:-1].strip()

    m = re.match(r'^([A-D]\.\s*)(.*)$', clean_str)
    if m:
        lbl = m.group(1)
        body = m.group(2)
    else:
        lbl = ""
        body = clean_str

    if is_teacher and is_correct:
        if lbl:
            r_lbl = p.add_run(lbl)
            set_run_style(r_lbl, font_size=base_size, bold=True, color=COLOR_CORRECT)
        render_rich_text(p, body, base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
    else:
        if lbl:
            r_lbl = p.add_run(lbl)
            set_run_style(r_lbl, font_size=base_size, bold=True, color=None)
        render_rich_text(p, body, base_size=base_size, base_bold=False, base_color=None)

def render_options_no_table(doc, opts, is_teacher=False):
    """Trình bày các phương án A, B, C, D hoàn toàn KHÔNG DÙNG BẢNG, dùng paragraph có tab stops.
    In đậm nhãn A., B., C., D. cho mọi phương án. Phương án đúng in đậm màu đỏ trong bản giáo viên.
    """
    clean_opts = [re.sub(r'[*_~^]', '', opt) for opt in opts]
    max_len = max(len(o) for o in clean_opts)

    if max_len <= 18:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(4.5), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(13.5), WD_TAB_ALIGNMENT.LEFT)

        for idx, opt in enumerate(opts):
            if idx > 0:
                p.add_run('\t')
            render_single_option(p, opt, is_teacher=is_teacher, base_size=10.5)
    elif max_len <= 38:
        p1 = doc.add_paragraph()
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.tab_stops.clear_all()
        p1.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
        render_single_option(p1, opts[0], is_teacher=is_teacher, base_size=10.5)
        p1.add_run('\t')
        render_single_option(p1, opts[1], is_teacher=is_teacher, base_size=10.5)

        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(1)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.tab_stops.clear_all()
        p2.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
        render_single_option(p2, opts[2], is_teacher=is_teacher, base_size=10.5)
        p2.add_run('\t')
        render_single_option(p2, opts[3], is_teacher=is_teacher, base_size=10.5)
    else:
        for opt in opts:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            render_single_option(p, opt, is_teacher=is_teacher, base_size=10.5)

def add_solution_no_table(doc, solution_text):
    """Trình bày lời giải KHÔNG DÙNG BẢNG (đoạn văn thụt đầu dòng, phân biệt rõ nét)."""
    lines = solution_text.strip().split("\n")
    p_head = doc.add_paragraph()
    p_head.paragraph_format.left_indent = Cm(0.4)
    p_head.paragraph_format.space_before = Pt(3)
    p_head.paragraph_format.space_after = Pt(1)
    render_rich_text(p_head, "► *Lời giải chi tiết:*", base_size=10.5, base_bold=True, base_color=COLOR_SOLUTION)

    for line in lines:
        if line.startswith("*Lời giải chi tiết:*") or line.startswith("►"):
            continue
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        render_rich_text(p, line, base_size=10)

def build_exam_de01(is_teacher=False):
    doc = docx.Document()
    setup_header_footer(doc)
    build_title_block_no_table(doc, is_teacher=is_teacher)

    # =========================================================================
    # PHẦN I
    # =========================================================================
    add_section_header(doc, "PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.",
                       "Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án.")

    part1_questions = [
        {
            "num": 1,
            "q": "Câu 1. Chất nào sau đây thuộc loại ester?",
            "opts": [
                "*A. CH~3~COOC~2~H~5~.*" if is_teacher else "A. CH~3~COOC~2~H~5~.",
                "B. HOOCCH~3~.",
                "C. H~2~N-CH~2~-COOH.",
                "D. CH~3~CHO."
            ],
            "sol": "• CH~3~COOC~2~H~5~ (ethyl acetate) thuộc loại ester vì chứa nhóm chức -COO- liên kết với gốc hydrocarbon.\n• HOOCCH~3~ là carboxylic acid (acetic acid), H~2~N-CH~2~-COOH là amino acid (glycine), CH~3~CHO là aldehyde."
        },
        {
            "num": 2,
            "q": "Câu 2. Phản ứng điều chế xà phòng từ chất béo được gọi là phản ứng",
            "opts": [
                "A. ester hóa.",
                "*B. xà phòng hóa.*" if is_teacher else "B. xà phòng hóa.",
                "C. trung hòa.",
                "D. hydrate hóa."
            ],
            "sol": "Phản ứng thủy phân chất béo trong môi trường kiềm (dung dịch NaOH hoặc KOH) đun nóng để sản xuất xà phòng và glycerol được gọi là phản ứng xà phòng hóa."
        },
        {
            "num": 3,
            "q": "Câu 3. Dầu chuối là ester có tên isoamyl acetate, được điều chế từ",
            "opts": [
                "A. CH~3~OH, CH~3~COOH.",
                "B. (CH~3~)~2~CH-CH~2~OH, CH~3~COOH.",
                "C. C~2~H~5~COOH, C~2~H~5~OH.",
                "*D. CH~3~COOH, (CH~3~)~2~CH-CH~2~-CH~2~OH.*" if is_teacher else "D. CH~3~COOH, (CH~3~)~2~CH-CH~2~-CH~2~OH."
            ],
            "sol": "Isoamyl acetate có công thức cấu tạo CH~3~COOCH~2~CH~2~CH(CH~3~)~2~, được điều chế từ phản ứng ester hóa giữa acetic acid (CH~3~COOH) và isoamyl alcohol ((CH~3~)~2~CH-CH~2~-CH~2~OH) với xúc tác H~2~SO~4~ đặc, đun nóng."
        },
        {
            "num": 4,
            "q": "Câu 4. Một số ester được dùng trong hương liệu, mĩ phẩm, bột giặt là nhờ các ester",
            "opts": [
                "A. là chất lỏng dễ bay hơi.",
                "*B. có mùi thơm, an toàn với người.*" if is_teacher else "B. có mùi thơm, an toàn với người.",
                "C. có thể bay hơi nhanh sau khi sử dụng.",
                "D. đều có nguồn gốc từ thiên nhiên."
            ],
            "sol": "Nhiều ester có mùi thơm đặc trưng của các loài hoa quả chín, không độc và an toàn cho người nên được ứng dụng làm hương liệu trong thực phẩm, mỹ phẩm và chất giặt rửa."
        },
        {
            "num": 5,
            "q": "Câu 5. Thủy phân ester nào sau đây trong dung dịch NaOH thu được sodium formate?",
            "opts": [
                "A. CH~3~COOCH~3~.",
                "B. CH~3~COOC~2~H~5~.",
                "*C. HCOOC~2~H~5~.*" if is_teacher else "C. HCOOC~2~H~5~.",
                "D. CH~3~COOC~3~H~7~."
            ],
            "sol": "Phương trình phản ứng: HCOOC~2~H~5~ + NaOH -(t°)-> HCOONa (sodium formate) + C~2~H~5~OH."
        },
        {
            "num": 6,
            "q": "Câu 6. Số đồng phân ester ứng với công thức phân tử C~4~H~8~O~2~ là",
            "opts": [
                "A. 2.",
                "B. 3.",
                "*C. 4.*" if is_teacher else "C. 4.",
                "D. 5."
            ],
            "sol": "C~4~H~8~O~2~ có 4 đồng phân ester:\n  (1) HCOOCH~2~CH~2~CH~3~ (propyl formate)\n  (2) HCOOCH(CH~3~)~2~ (isopropyl formate)\n  (3) CH~3~COOCH~2~CH~3~ (ethyl acetate)\n  (4) C~2~H~5~COOCH~3~ (methyl propionate)."
        },
        {
            "num": 7,
            "q": "Câu 7. Công thức của tristearin là",
            "opts": [
                "A. (C~2~H~5~COO)~3~C~3~H~5~.",
                "*B. (C~17~H~35~COO)~3~C~3~H~5~.*" if is_teacher else "B. (C~17~H~35~COO)~3~C~3~H~5~.",
                "C. (CH~3~COO)~3~C~3~H~5~.",
                "D. (HCOO)~3~C~3~H~5~."
            ],
            "sol": "Tristearin là triglyceride của stearic acid (C~17~H~35~COOH) với glycerol, có công thức là (C~17~H~35~COO)~3~C~3~H~5~."
        },
        {
            "num": 8,
            "q": "Câu 8. Thực hiện phản ứng ester hoá giữa HOOC-COOH với hỗn hợp CH~3~OH và C~2~H~5~OH thu được tối đa bao nhiêu ester hai chức?",
            "opts": [
                "A. 2.",
                "*B. 3.*" if is_teacher else "B. 3.",
                "C. 1.",
                "D. 4."
            ],
            "sol": "Phản ứng giữa acid 2 chức HOOC-COOH và 2 alcohol CH~3~OH, C~2~H~5~OH tạo tối đa 3 ester hai chức:\n  (1) CH~3~OOC-COOCH~3~ (dimethyl oxalate)\n  (2) C~2~H~5~OOC-COOC~2~H~5~ (diethyl oxalate)\n  (3) CH~3~OOC-COOC~2~H~5~ (ethyl methyl oxalate)."
        },
        {
            "num": 9,
            "q": "Câu 9. Công thức phân tử của oleic acid là",
            "opts": [
                "A. C~2~H~5~COOH.",
                "B. HCOOOH.",
                "C. CH~3~COOH.",
                "*D. C~17~H~33~COOH.*" if is_teacher else "D. C~17~H~33~COOH."
            ],
            "sol": "Oleic acid là acid béo không no đơn chức mang 1 nối đôi C=C ở mạch carbon, có công thức là C~17~H~33~COOH."
        },
        {
            "num": 10,
            "q": "Câu 10. Công thức của triolein là",
            "opts": [
                "*A. (C~17~H~33~COO)~3~C~3~H~5~.*" if is_teacher else "A. (C~17~H~33~COO)~3~C~3~H~5~.",
                "B. (HCOO)~3~C~3~H~5~.",
                "C. (C~2~H~5~COO)~3~C~3~H~5~.",
                "D. (CH~3~COO)~3~C~3~H~5~."
            ],
            "sol": "Triolein là triglyceride của oleic acid (C~17~H~33~COOH) với glycerol, có công thức phân tử là (C~17~H~33~COO)~3~C~3~H~5~."
        },
        {
            "num": 11,
            "q": "Câu 11. Cho các chất sau: (1) alcohol ethylic, (2) acetic acid, (3) nước, (4) methyl formate. Thứ tự nhiệt độ sôi giảm dần là",
            "opts": [
                "A. (1) > (4) > (3) > (2).",
                "B. (1) > (2) > (3) > (4).",
                "C. (1) > (3) > (2) > (4).",
                "*D. (2) > (3) > (1) > (4).*" if is_teacher else "D. (2) > (3) > (1) > (4)."
            ],
            "sol": "Nhiệt độ sôi: Acetic acid (117,9 °C) > H~2~O (100 °C) > Alcohol ethylic (78,4 °C) > Methyl formate (31,5 °C).\nThứ tự nhiệt độ sôi giảm dần: (2) > (3) > (1) > (4)."
        },
        {
            "num": 12,
            "q": "Câu 12. Xà phòng và chất giặt rửa có đặc điểm chung nào sau đây?",
            "opts": [
                "A. Không tan trong nước.",
                "B. Là muối sodium hoặc potassium của acid béo.",
                "C. Là muối sulfonate hoặc sulfate của acid béo.",
                "*D. Thường có cấu tạo gồm hai phần là phần không phân cực (kị nước) và phần phân cực (ưa nước).*" if is_teacher else "D. Thường có cấu tạo gồm hai phần là phần không phân cực (kị nước) và phần phân cực (ưa nước)."
            ],
            "sol": "Đặc điểm chung về cấu tạo phân tử của xà phòng và chất giặt rửa là có tính lưỡng cực: phần đầu phân cực (ưa nước) và phần đuôi hydrocarbon dài không phân cực (kị nước, ưa dầu mỡ)."
        },
        {
            "num": 13,
            "q": "Câu 13. Cho các chất sau: CH~3~[CH~2~]~7~CH=CH[CH~2~]~7~COONa, CH~3~[CH~2~]~14~COOK, CH~3~[CH~2~]~10~COOK và CH~3~COONa. Trong các chất nêu trên, có bao nhiêu chất có thể là thành phần chính của xà phòng?",
            "opts": [
                "A. 1.",
                "B. 2.",
                "*C. 3.*" if is_teacher else "C. 3.",
                "D. 4."
            ],
            "sol": "Xà phòng là muối sodium hoặc potassium của acid béo (mạch C dài từ 12-24 nguyên tử C):\n  • CH~3~[CH~2~]~7~CH=CH[CH~2~]~7~COONa (sodium oleate: 18C) -> Thỏa mãn.\n  • CH~3~[CH~2~]~14~COOK (potassium palmitate: 16C) -> Thỏa mãn.\n  • CH~3~[CH~2~]~10~COOK (potassium laurate: 12C) -> Thỏa mãn.\n  • CH~3~COONa có mạch carbon quá ngắn (2C), không có hoạt tính giặt rửa của xà phòng.\n=> Có 3 chất có thể là thành phần chính của xà phòng."
        },
        {
            "num": 14,
            "q": "Câu 14. Phát biểu nào sau đây về xà phòng là đúng?",
            "opts": [
                "A. Xà phòng có thành phần chính là muối sodium hoặc potassium của carboxylic acid.",
                "B. Các phân tử xà phòng đều có đầu kị nước gắn với đuôi dài ưa nước.",
                "*C. Xà phòng mất tính giặt rửa khi sử dụng với nước cứng.*" if is_teacher else "C. Xà phòng mất tính giặt rửa khi sử dụng với nước cứng.",
                "D. Nhược điểm của xà phòng là khó bị phân huỷ hoặc phân huỷ chậm, do đó gây hại cho hệ sinh thái."
            ],
            "sol": "• C đúng: Nước cứng chứa ion Ca^2+^, Mg^2+^ kết tủa với anion acid béo làm mất khả năng tạo bọt và tính giặt rửa của xà phòng.\n• A sai: Phải là muối của acid béo (mạch C dài), không phải mọi carboxylic acid.\n• B sai: Đầu ưa nước gắn với đuôi kị nước.\n• D sai: Xà phòng dễ bị vi sinh vật phân hủy sinh học."
        },
        {
            "num": 15,
            "q": "Câu 15. Trong số các vật phẩm tiêu dùng sau: xà phòng bánh, dầu gội đầu, nước bồ kết và baking soda (NaHCO~3~), số vật phẩm có thành phần chất giặt rửa tự nhiên và tổng hợp là",
            "opts": [
                "A. 1.",
                "*B. 2.*" if is_teacher else "B. 2.",
                "C. 3.",
                "D. 4."
            ],
            "sol": "• Nước bồ kết chứa chất giặt rửa tự nhiên (saponin).\n• Dầu gội đầu chứa chất giặt rửa tổng hợp (sodium lauryl sulfate,...).\n=> Có 2 vật phẩm có thành phần chất giặt rửa tự nhiên và tổng hợp."
        },
        {
            "num": 16,
            "q": "Câu 16. Loại dầu nào sau đây không phải là chất béo?",
            "opts": [
                "A. dầu vừng.",
                "B. dầu oliu.",
                "C. dầu gan cá.",
                "*D. dầu luyn.*" if is_teacher else "D. dầu luyn."
            ],
            "sol": "Dầu luyn (dầu bôi trơn máy móc) là hỗn hợp các hydrocarbon từ dầu mỏ. Dầu vừng, dầu oliu, dầu gan cá đều là chất béo (triglyceride)."
        },
        {
            "num": 17,
            "q": "Câu 17. Cho các phản ứng sau:\n(1) Thuỷ phân ester trong môi trường acid.\n(2) Thuỷ phân ester trong dung dịch NaOH, đun nóng.\n(3) Cho ester tác dụng với dung dịch KOH, đun nóng.\n(4) Thuỷ phân dẫn xuất halogen trong dung dịch NaOH, đun nóng.\n(5) Cho carboxylic acid tác dụng với dung dịch NaOH.\nNhững phản ứng nào không được gọi là phản ứng xà phòng hoá?",
            "opts": [
                "A. (1), (2), (3), (4).",
                "*B. (1), (4), (5).*" if is_teacher else "B. (1), (4), (5).",
                "C. (1), (3), (4), (5).",
                "D. (3), (4), (5)."
            ],
            "sol": "Phản ứng xà phòng hóa là phản ứng thủy phân ester trong dung dịch kiềm (NaOH, KOH) đun nóng -> (2) và (3) là phản ứng xà phòng hóa. Các phản ứng (1), (4), (5) không phải là phản ứng xà phòng hóa."
        },
        {
            "num": 18,
            "q": "Câu 18. Đun nóng acid acetic với isoamyl alcohol (CH~3~)~2~CH-CH~2~CH~2~OH có H~2~SO~4~ đặc xúc tác thu được isoamyl acetate (dầu chuối). Tính lượng dầu chuối thu được từ 132,35 gam acid acetic đun nóng với 200 gam isoamyl alcohol (Biết hiệu suất phản ứng đạt 68%).",
            "opts": [
                "A. 97,5 gam.",
                "*B. 195,0 gam.*" if is_teacher else "B. 195,0 gam.",
                "C. 292,5 gam.",
                "D. 159,0 gam."
            ],
            "sol": "• Phương trình: CH~3~COOH + (CH~3~)~2~CH-CH~2~CH~2~OH <=(H~2~SO~4~ đ, t°)=> CH~3~COOCH~2~CH~2~CH(CH~3~)~2~ + H~2~O\n• n~acid~ = 132,35 / 60 ≈ 2,206 mol; n~alcohol~ = 200 / 88 ≈ 2,273 mol.\n• Do n~alcohol~ > n~acid~ nên hiệu suất tính theo CH~3~COOH:\n    n~ester~ = 2,206 × 68% = 1,5 mol.\n• m~ester~ = 1,5 × 130 = 195,0 gam."
        }
    ]

    for item in part1_questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(1)
        render_question_prompt(p_q, item["q"], base_size=10.5)

        render_options_no_table(doc, item["opts"], is_teacher=is_teacher)

        if is_teacher and item.get("sol"):
            add_solution_no_table(doc, item["sol"])

    # =========================================================================
    # PHẦN II
    # =========================================================================
    add_section_header(doc, "PHẦN II. Câu trắc nghiệm đúng sai.",
                       "Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.")

    # Câu 1
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(4)
    p_q.paragraph_format.space_after = Pt(2)
    render_question_prompt(p_q, "Câu 1. Cho các triglyceride X, Y với công thức cấu tạo sau:", base_size=10.5)

    img_xy = os.path.join(CROP_DIR, "triglyceride_XY.png")
    if os.path.exists(img_xy):
        p_img = doc.add_paragraph()
        p_img.paragraph_format.space_before = Pt(2)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(img_xy, width=Cm(15.5))

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(1)
    p_sub.paragraph_format.space_after = Pt(2)
    render_rich_text(p_sub, "Em hãy cho biết các phát biểu sau đây là đúng hay sai:", base_size=10.5, base_italic=True)

    c1_statements = [
        ("a) Triglyceride X có tên gọi là tripalmitin.", "Sai", "Gốc acid trong X có 17 C ở đuôi hydrocarbon (C~17~H~35~COO-), do đó X là tristearin ((C~17~H~35~COO)~3~C~3~H~5~), không phải tripalmitin."),
        ("b) X là chất béo no, Y là chất béo không no.", "Đúng", "X không chứa liên kết đôi C=C ở gốc acid nên là chất béo no; Y chứa các liên kết đôi C=C (mỗi gốc oleate có 1 nối đôi C=C) nên là chất béo không no."),
        ("c) X, Y đều tan tốt trong nước.", "Sai", "Chất béo nhẹ hơn nước và hầu như không tan trong nước do phân tử không phân cực."),
        ("d) Hydrogen hoá Y thu được X.", "Đúng", "Hydrogen hóa hoàn toàn triolein (Y) với xúc tác Ni, đun nóng sẽ cộng H~2~ vào liên kết đôi C=C thu được tristearin (X).")
    ]
    for stmt, ans, expl in c1_statements:
        p_stmt = doc.add_paragraph()
        p_stmt.paragraph_format.left_indent = Cm(0.5)
        p_stmt.paragraph_format.space_before = Pt(1)
        p_stmt.paragraph_format.space_after = Pt(1)
        render_statement(p_stmt, stmt, ans=ans, is_teacher=is_teacher, base_size=10.5)

    if is_teacher:
        sol_c1 = "\n".join([f"• {stmt[:2]} *{ans}*: {expl}" for stmt, ans, expl in c1_statements])
        add_solution_no_table(doc, sol_c1)

    # Câu 2
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(4)
    p_q.paragraph_format.space_after = Pt(2)
    render_question_prompt(p_q, "Câu 2. Nhiệt độ sôi và độ tan của một số ester, carboxylic acid và alcohol có cùng số nguyên tử carbon được cho trong bảng sau:", base_size=10.5)

    # Bảng dữ liệu Câu 2
    tbl_data = [
        ("Công thức", "Nhiệt độ sôi (°C)", "Độ tan ở 25 °C (g/100 g nước)"),
        ("HCOOCH~3~", "31,5", "23,0"),
        ("HCOOC~2~H~5~", "54,2", "12,0"),
        ("CH~3~COOH", "117,9", "Tan vô hạn"),
        ("C~2~H~5~COOH", "141,0", "Tan vô hạn"),
        ("C~2~H~5~OH", "78,4", "Tan vô hạn"),
        ("CH~3~CH~2~CH~2~OH", "97,2", "Tan vô hạn")
    ]
    tbl2 = doc.add_table(rows=len(tbl_data), cols=3)
    tbl2.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(tbl_data):
        for c_idx, val in enumerate(row):
            cell = tbl2.cell(r_idx, c_idx)
            cell.paragraphs[0].paragraph_format.space_before = Pt(1)
            cell.paragraphs[0].paragraph_format.space_after = Pt(1)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            is_hdr = (r_idx == 0)
            render_rich_text(cell.paragraphs[0], f"*{val}*" if is_hdr else val, base_size=10, base_bold=is_hdr)
    # Viền bảng xám mảnh
    for row in tbl2.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(r'''
                <w:tcBorders %s>
                    <w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
                    <w:left w:val="none"/>
                    <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
                    <w:right w:val="none"/>
                </w:tcBorders>
            ''' % nsdecls('w'))
            tcPr.append(tcBorders)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(2)
    p_sub.paragraph_format.space_after = Pt(2)
    render_rich_text(p_sub, "Em hãy cho biết các phát biểu sau đây là đúng hay sai:", base_size=10.5, base_italic=True)

    c2_statements = [
        ("a) Do không tạo được liên kết hydrogen giữa các phân tử nên ester có nhiệt độ sôi thấp hơn nhiệt độ sôi của carboxylic acid và alcohol có cùng số nguyên tử carbon.", "Đúng", "Phân tử ester không có liên kết hydrogen liên phân tử nên nhiệt độ sôi thấp hơn nhiều so với acid và alcohol tương ứng."),
        ("b) Do có khả năng tạo liên kết hydrogen yếu với nước nên ester thường ít tan trong nước hơn so với carboxylic acid và alcohol có cùng số carbon.", "Đúng", "Ester chỉ có thể nhận liên kết hydrogen từ nước, không tự cho liên kết hydrogen, nên độ tan trong nước kém hơn hẳn acid và alcohol."),
        ("c) Carboxylic acid có nhiệt độ sôi cao hơn alcohol có cùng số nguyên tử carbon.", "Đúng", "Liên kết hydrogen giữa các phân tử carboxylic acid dạng nhị hợp (dimer) bền hơn so với liên kết hydrogen của alcohol."),
        ("d) Methanol có khả năng tan vô hạn trong nước.", "Đúng", "Methanol (CH~3~OH) có gốc hydrocarbon nhỏ và tạo liên kết hydrogen rất bền với nước nên tan vô hạn trong nước ở mọi tỉ lệ.")
    ]
    for stmt, ans, expl in c2_statements:
        p_stmt = doc.add_paragraph()
        p_stmt.paragraph_format.left_indent = Cm(0.5)
        p_stmt.paragraph_format.space_before = Pt(1)
        p_stmt.paragraph_format.space_after = Pt(1)
        render_statement(p_stmt, stmt, ans=ans, is_teacher=is_teacher, base_size=10.5)

    if is_teacher:
        sol_c2 = "\n".join([f"• {stmt[:2]} *{ans}*: {expl}" for stmt, ans, expl in c2_statements])
        add_solution_no_table(doc, sol_c2)

    # Câu 3
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(4)
    p_q.paragraph_format.space_after = Pt(2)
    render_question_prompt(p_q, "Câu 3. Các phát biểu sau đây về xà phòng và chất giặt rửa là đúng hay sai?", base_size=10.5)

    c3_statements = [
        ("a) Xà phòng và chất giặt rửa thường có cấu tạo gồm hai phần: ưa nước và kị nước.", "Đúng", "Cả xà phòng và chất giặt rửa đều có cấu tạo lưỡng cực gồm phần đầu ưa nước và phần đuôi hydrocarbon kị nước."),
        ("b) Xà phòng hoá tripalmitin với dung dịch NaOH thu được sản phẩm là C~15~H~29~COONa và glycerol.", "Sai", "Tripalmitin có công thức (C~15~H~31~COO)~3~C~3~H~5~, khi xà phòng hóa thu được sodium palmitate là C~15~H~31~COONa, không phải C~15~H~29~COONa."),
        ("c) Chất giặt rửa tổng hợp thường được điều chế từ chất béo.", "Sai", "Chất giặt rửa tổng hợp chủ yếu được tổng hợp từ nguồn hydrocarbon của dầu mỏ, không phải từ chất béo."),
        ("d) Mỡ động vật, dầu thực vật là nguyên liệu để sản xuất xà phòng.", "Đúng", "Dầu thực vật và mỡ động vật là nguồn nguyên liệu tự nhiên chủ yếu để sản xuất xà phòng.")
    ]
    for stmt, ans, expl in c3_statements:
        p_stmt = doc.add_paragraph()
        p_stmt.paragraph_format.left_indent = Cm(0.5)
        p_stmt.paragraph_format.space_before = Pt(1)
        p_stmt.paragraph_format.space_after = Pt(1)
        render_statement(p_stmt, stmt, ans=ans, is_teacher=is_teacher, base_size=10.5)

    if is_teacher:
        sol_c3 = "\n".join([f"• {stmt[:2]} *{ans}*: {expl}" for stmt, ans, expl in c3_statements])
        add_solution_no_table(doc, sol_c3)

    # Câu 4
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(4)
    p_q.paragraph_format.space_after = Pt(2)
    render_question_prompt(p_q, "Câu 4. Các phát biểu sau đây là đúng hay sai?", base_size=10.5)

    c4_statements = [
        ("a) Chất giặt rửa thường là muối sodium alkylsulfate hoặc alkylbenzene sulfonate.", "Đúng", "Các chất giặt rửa tổng hợp thông dụng thường gặp là muối sodium alkylsulfate hoặc alkylbenzene sulfonate."),
        ("b) Phân tử chất giặt rửa gồm một đầu kị nước gắn với một đầu ưa nước.", "Đúng", "Phân tử chất hoạt động bề mặt gồm đầu phân cực ưa nước gắn với đuôi mạch carbon dài kị nước."),
        ("c) Khi giặt rửa bằng nước cứng nên sử dụng xà phòng.", "Sai", "Trong nước cứng (chứa Ca^2+^, Mg^2+^), xà phòng bị kết tủa mất tác dụng; do đó nên dùng chất giặt rửa tổng hợp."),
        ("d) Phản ứng thuỷ phân chất béo trong môi trường kiềm (NaOH, KOH) thuộc loại phản ứng xà phòng hoá.", "Đúng", "Phản ứng thủy phân chất béo trong dung dịch kiềm nóng được gọi là phản ứng xà phòng hóa.")
    ]
    for stmt, ans, expl in c4_statements:
        p_stmt = doc.add_paragraph()
        p_stmt.paragraph_format.left_indent = Cm(0.5)
        p_stmt.paragraph_format.space_before = Pt(1)
        p_stmt.paragraph_format.space_after = Pt(1)
        render_statement(p_stmt, stmt, ans=ans, is_teacher=is_teacher, base_size=10.5)

    if is_teacher:
        sol_c4 = "\n".join([f"• {stmt[:2]} *{ans}*: {expl}" for stmt, ans, expl in c4_statements])
        add_solution_no_table(doc, sol_c4)

    # =========================================================================
    # PHẦN III
    # =========================================================================
    add_section_header(doc, "PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn.",
                       "Thí sinh trả lời từ câu 1 đến câu 6. Mỗi câu hỏi thí sinh điền câu trả lời ngắn gọn theo yêu cầu.")

    part3_questions = [
        {
            "num": 1,
            "q": "Câu 1. Khi xà phòng hóa triglyceride X bằng dung dịch NaOH dư, đun nóng, thu được sản phẩm gồm glycerol, sodium oleate, sodium stearate và sodium palmitate. Số đồng phân cấu tạo thỏa mãn tính chất trên của X là bao nhiêu?",
            "ans": "3",
            "sol": "• Triglyceride X chứa 3 gốc acid béo khác nhau: oleate (C~17~H~33~COO-), stearate (C~17~H~35~COO-) và palmitate (C~15~H~31~COO-).\n• Số đồng phân cấu tạo chứa 3 gốc acid béo khác nhau gắn vào glycerol là: 3! / 2 = 3 đồng phân (do vị trí C1 và C3 trên khung glycerol có tính đối xứng).\n➔ Đáp số: 3."
        },
        {
            "num": 2,
            "q": "Câu 2. Cho triolein lần lượt vào mỗi ống nghiệm chứa riêng biệt: Na, Cu(OH)~2~, CH~3~OH, dung dịch Br~2~, dung dịch NaOH. Trong điều kiện thích hợp, số phản ứng xảy ra là bao nhiêu?",
            "ans": "2",
            "sol": "• Triolein ((C~17~H~33~COO)~3~C~3~H~5~) là triglyceride không no có 3 liên kết đôi C=C:\n  (1) Phản ứng cộng dung dịch Br~2~ vào liên kết đôi C=C (làm mất màu dung dịch brom).\n  (2) Phản ứng thủy phân trong dung dịch NaOH (phản ứng xà phòng hóa).\n• Triolein không phản ứng với Na, Cu(OH)~2~ và CH~3~OH.\n➔ Có 2 phản ứng xảy ra."
        },
        {
            "num": 3,
            "q": "Câu 3. Một loại chất béo có chứa 80% tristearin về khối lượng. Để sản xuất ba nghìn (3000) bánh xà phòng cần dùng tối thiểu x kg loại chất béo trên cho phản ứng với dung dịch NaOH, đun nóng. Biết hiệu suất phản ứng đạt 90%. Biết rằng trong mỗi bánh xà phòng có chứa 60 gam sodium stearate. Giá trị của x là bao nhiêu? (Làm tròn kết quả đến một chữ số thập phân).",
            "ans": "242,4",
            "img": os.path.join(CROP_DIR, "soap_bar.png"),
            "sol": "• Khối lượng sodium stearate cần tạo ra:\n    m = 3000 × 60 g = 180 000 g = 180 kg.\n• Phương trình hóa học:\n    (C~17~H~35~COO)~3~C~3~H~5~ + 3NaOH -(t°)-> 3C~17~H~35~COONa + C~3~H~5~(OH)~3~\n    890 kg ----------------------------------> 3 × 306 = 918 kg\n    m~tristearin (LT)~ ---------------------> 180 kg\n• Khối lượng tristearin theo lý thuyết: m~LT~ = (180 × 890) / 918 ≈ 174,51 kg.\n• Với hiệu suất phản ứng H = 90% và chất béo chứa 80% tristearin:\n    x = m~chất béo~ = 174,51 / (0,90 × 0,80) = (180 × 890 × 100 × 100) / (918 × 90 × 80) ≈ 242,37 kg ≈ 242,4 kg.\n➔ Đáp số: 242,4."
        },
        {
            "num": 4,
            "q": "Câu 4. Cho 0,1 mol butanoic acid tác dụng với 0,1 mol methyl alcohol có mặt H~2~SO~4~ đặc làm xúc tác. Tính khối lượng ester tạo thành theo gam (biết 67% alcohol chuyển hoá thành ester). (Làm tròn kết quả đến hai chữ số thập phân).",
            "ans": "6,83",
            "sol": "• Phương trình phản ứng:\n    CH~3~CH~2~CH~2~COOH + CH~3~OH <=(H~2~SO~4~ đ, t°)=> CH~3~CH~2~CH~2~COOCH~3~ + H~2~O\n• Khối lượng mol ester methyl butanoate: M = 102 g/mol.\n• Số mol ester tạo thành: n~ester~ = 0,1 × 67% = 0,067 mol.\n• Khối lượng ester thu được: m~ester~ = 0,067 × 102 = 6,834 g ≈ 6,83 gam.\n➔ Đáp số: 6,83."
        },
        {
            "num": 5,
            "q": "Câu 5. Một loại dầu thực vật trong đó thành phần chất béo chứa hai gốc linoleate, một gốc oleate và thành phần phần trăm khối lượng chất béo trong dầu thực vật là 88%. Tính chỉ số ester của dầu thực vật đó. (Biết chỉ số ester là số miligam KOH dùng để xà phòng hoá hết lượng triglyceride có trong 1 g chất béo).",
            "ans": "168",
            "sol": "• Công thức cấu tạo chất béo: (C~17~H~31~COO)~2~(C~17~H~33~COO)C~3~H~5~\n    M = 2 × 279 + 281 + 41 = 880 g/mol.\n• Trong 1 g dầu thực vật có:\n    m~chất béo~ = 1 × 88% = 0,88 g => n~chất béo~ = 0,88 / 880 = 0,001 mol.\n• Phản ứng xà phòng hóa cần:\n    n~KOH~ = 3 × n~chất béo~ = 0,003 mol.\n• Khối lượng KOH tương ứng:\n    m~KOH~ = 0,003 × 56 = 0,168 g = 168 mg.\n• Vậy chỉ số ester của loại dầu thực vật là 168.\n➔ Đáp số: 168."
        },
        {
            "num": 6,
            "q": "Câu 6. Số miligam KOH dùng để xà phòng hoá hết lượng triglyceride có trong 1 g chất béo được gọi là chỉ số ester hoá của loại chất béo đó. Tính chỉ số ester của một loại chất béo chứa 65% tristearin và 23% triolein (còn lại là tạp chất không phản ứng). (Kết quả làm tròn đến phần nguyên).",
            "ans": "166",
            "sol": "• Trong 1 g chất béo có:\n    m~tristearin~ = 0,65 g (M = 890 g/mol) => n~tristearin~ = 0,65 / 890 mol.\n    m~triolein~ = 0,23 g (M = 884 g/mol) => n~triolein~ = 0,23 / 884 mol.\n• Số mol KOH cần dùng:\n    n~KOH~ = 3 × (n~tristearin~ + n~triolein~) = 3 × (0,65 / 890 + 0,23 / 884) ≈ 0,0029715 mol.\n• Khối lượng KOH cần dùng:\n    m~KOH~ = 0,0029715 × 56 ≈ 0,1664 g = 166,4 mg ≈ 166 mg.\n• Vậy chỉ số ester của chất béo là 166.\n➔ Đáp số: 166."
        }
    ]

    for item in part3_questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, item["q"], base_size=10.5)

        if item.get("img") and os.path.exists(item["img"]):
            p_img = doc.add_paragraph()
            p_img.paragraph_format.space_before = Pt(2)
            p_img.paragraph_format.space_after = Pt(2)
            p_img.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_img.add_run().add_picture(item["img"], width=Cm(3.8))

        if is_teacher:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Cm(0.5)
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(1)
            render_rich_text(p_ans, f"➔ *Đáp án:* *{item['ans']}*", base_size=10.5, base_bold=True, base_color=COLOR_CORRECT)
            add_solution_no_table(doc, item["sol"])
        else:
            p_ans_blank = doc.add_paragraph()
            p_ans_blank.paragraph_format.left_indent = Cm(0.5)
            p_ans_blank.paragraph_format.space_before = Pt(1)
            p_ans_blank.paragraph_format.space_after = Pt(3)
            render_rich_text(p_ans_blank, "Thí sinh điền đáp án vào ô: .......................................................................................", base_size=10, base_italic=True)

    # Cuối bản giáo viên: Bảng tổng hợp đáp án nhanh
    if is_teacher:
        p_div = doc.add_paragraph()
        p_div.paragraph_format.space_before = Pt(8)
        p_div.paragraph_format.space_after = Pt(4)
        add_p_border_bottom(p_div, color="CCCCCC", sz="4")

        p_tb_title = doc.add_paragraph()
        p_tb_title.paragraph_format.space_before = Pt(4)
        p_tb_title.paragraph_format.space_after = Pt(3)
        render_rich_text(p_tb_title, "► *BẢNG TỔNG HỢP ĐÁP ÁN ĐỀ SỐ 1*", base_size=11, base_bold=True, base_color=COLOR_PRIMARY)

        # Bảng đáp án Phần I
        p_p1 = doc.add_paragraph()
        p_p1.paragraph_format.space_before = Pt(2)
        p_p1.paragraph_format.space_after = Pt(2)
        render_rich_text(p_p1, "*1. PHẦN I (Trắc nghiệm nhiều phương án):*", base_size=10, base_bold=True)

        p1_ans = [
            ("1.A", "2.B", "3.D", "4.B", "5.C", "6.C"),
            ("7.B", "8.B", "9.D", "10.A", "11.D", "12.D"),
            ("13.C", "14.C", "15.B", "16.D", "17.B", "18.B")
        ]
        for row in p1_ans:
            p_row = doc.add_paragraph()
            p_row.paragraph_format.left_indent = Cm(0.5)
            p_row.paragraph_format.space_before = Pt(1)
            p_row.paragraph_format.space_after = Pt(1)
            p_row.paragraph_format.tab_stops.clear_all()
            for ts in [2.5, 5.0, 7.5, 10.0, 12.5]:
                p_row.paragraph_format.tab_stops.add_tab_stop(Cm(ts), WD_TAB_ALIGNMENT.LEFT)
            for idx, item in enumerate(row):
                if idx > 0:
                    p_row.add_run('\t')
                render_rich_text(p_row, f"*{item}*", base_size=10, base_bold=True, base_color=COLOR_CORRECT)

        # Đáp án Phần II
        p_p2 = doc.add_paragraph()
        p_p2.paragraph_format.space_before = Pt(4)
        p_p2.paragraph_format.space_after = Pt(2)
        render_rich_text(p_p2, "*2. PHẦN II (Trắc nghiệm đúng sai):*", base_size=10, base_bold=True)
        p2_summary = [
            "• Câu 1: a) S  |  b) Đ  |  c) S  |  d) Đ",
            "• Câu 2: a) Đ  |  b) Đ  |  c) Đ  |  d) Đ",
            "• Câu 3: a) Đ  |  b) S  |  c) S  |  d) Đ",
            "• Câu 4: a) Đ  |  b) Đ  |  c) S  |  d) Đ"
        ]
        for s in p2_summary:
            p_s = doc.add_paragraph()
            p_s.paragraph_format.left_indent = Cm(0.5)
            p_s.paragraph_format.space_before = Pt(1)
            p_s.paragraph_format.space_after = Pt(1)
            render_rich_text(p_s, f"*{s}*", base_size=10, base_bold=True, base_color=COLOR_CORRECT)

        # Đáp án Phần III
        p_p3 = doc.add_paragraph()
        p_p3.paragraph_format.space_before = Pt(4)
        p_p3.paragraph_format.space_after = Pt(2)
        render_rich_text(p_p3, "*3. PHẦN III (Trả lời ngắn):*", base_size=10, base_bold=True)
        p3_summary = "• Câu 1: *3*    • Câu 2: *2*    • Câu 3: *242,4*    • Câu 4: *6,83*    • Câu 5: *168*    • Câu 6: *166*"
        p_s3 = doc.add_paragraph()
        p_s3.paragraph_format.left_indent = Cm(0.5)
        p_s3.paragraph_format.space_before = Pt(1)
        p_s3.paragraph_format.space_after = Pt(1)
        render_rich_text(p_s3, p3_summary, base_size=10, base_color=COLOR_CORRECT)

    # Save
    filename = "De_Kiem_Tra_Hoa_12_De01_GiaoVien.docx" if is_teacher else "De_Kiem_Tra_Hoa_12_De01.docx"
    filepath = os.path.join(OUT_DIR, filename)
    doc.save(filepath)
    print(f"Saved: {filepath} ({os.path.getsize(filepath) / 1024:.1f} KB)")

if __name__ == "__main__":
    print("Building Student Word document...")
    build_exam_de01(is_teacher=False)
    print("Building Teacher Word document...")
    build_exam_de01(is_teacher=True)
    print("Done!")
