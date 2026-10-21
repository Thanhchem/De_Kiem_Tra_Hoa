# -*- coding: utf-8 -*-
"""
Script tạo trọn bộ 3 file Word (.docx) chuẩn "FORM ĐỀ KIỂM TRA" cho:
HƯỚNG DẪN ÔN TẬP GIỮA HỌC KÌ I - NĂM HỌC 2026 - 2027
TRƯỜNG THCS TRẦN CAO VÂN - MÔN KHOA HỌC TỰ NHIÊN 9
Thầy TRẦN VĂN THẠNH - 0777.470.803

3 phiên bản:
1. Bản Học sinh: De_Cuong_KHTN_9_Giua_HK1_TranCaoVan_HocSinh.docx
2. Bản Giáo viên: De_Cuong_KHTN_9_Giua_HK1_TranCaoVan_GiaoVien.docx
3. Bản Dạy trên màn hình: De_Cuong_KHTN_9_Giua_HK1_TranCaoVan_DayTrenManHinh.docx
"""

import os
import sys
import re
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.path.append(os.path.dirname(__file__))
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

from khtn9_trancaovan_data import (
    INFO,
    IMG_CON_LAC,
    VATLI_QUESTIONS,
    HOA_THEORY,
    HOA_MCQ,
    HOA_TF,
    HOA_SHORT,
    SINHHOC_MCQ,
    SINHHOC_SHORT,
    SINHHOC_ESSAY,
    SINHHOC_TF
)

ROOT_DIR = r"c:\Antigravity_Thanh"
OUT_DIR = os.path.join(ROOT_DIR, "San_Pham")
os.makedirs(OUT_DIR, exist_ok=True)

ONEDRIVE_DIRS = [
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn9\KHTN9-DEKIEMTRADINHKI",
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn9"
]

# Bảng màu quy chuẩn theo FORM_DE_KIEM_TRA
COLOR_HF_RED = RGBColor(211, 47, 47)        # Coral Red (#D32F2F) cho Header & Footer
COLOR_HF_GREEN_HEX = "8BC390"               # Sage Green (#8BC390) cho viền Header & Footer
COLOR_PRIMARY = RGBColor(46, 125, 50)       # Xanh lá đậm (#2E7D32) cho Tiêu đề Phân môn / Phần
COLOR_CORRECT = RGBColor(198, 40, 40)       # Đỏ (#C62828) cho Đáp án đúng Bản Giáo Viên
COLOR_SOLUTION = RGBColor(21, 101, 192)     # Xanh lam (#1565C0) cho tiêu đề Lời giải
COLOR_SCREEN_TITLE = RGBColor(106, 27, 154) # Tím đậm (#6A1B9A) cho tiêu đề bản Dạy Màn Hình
COLOR_GRAY = RGBColor(110, 110, 110)
COLOR_LINE_DOT = RGBColor(185, 185, 185)    # Xám nhạt cho dòng chấm ..........
COLOR_BORDER_GREY_HEX = "CCCCCC"

def add_p_border_bottom(p, color="8BC390", sz="6"):
    """Thêm đường kẻ viền dưới cho đoạn văn bản."""
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_p_border_top(p, color="8BC390", sz="6"):
    """Thêm đường kẻ viền trên cho đoạn văn bản."""
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    """Cài đặt lề đệm cho ô bảng."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex="F1F8E9"):
    """Đặt màu nền cho ô bảng."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="CCCCCC", sz="4"):
    """Đặt viền thanh mảnh cho bảng."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def set_run_style(run, font_name="Times New Roman", font_size=10.5, bold=False, italic=False, subscript=False, superscript=False, color=None):
    """Cài đặt style chi tiết cho từng run text."""
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    """Chèn trường số trang tự động PAGE vào chân trang."""
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def render_rich_text(p, text, base_font="Times New Roman", base_size=10.5, base_bold=False, base_italic=False, base_color=None):
    """
    Sinh văn bản định dạng phong phú:
    - *bold*
    - _italic_
    - ~subscript~
    - ^superscript^
    """
    tokens = re.split(r'(~[^~]+~|\^[^\^]+\^|\*[^\*]+\*|_[^_]+_)', text)
    for token in tokens:
        if not token:
            continue

        is_bold = base_bold
        is_italic = base_italic
        is_sub = False
        is_sup = False
        content = token

        if token.startswith('*') and token.endswith('*'):
            is_bold = True
            content = token[1:-1]
        elif token.startswith('_') and token.endswith('_'):
            is_italic = True
            content = token[1:-1]
        elif token.startswith('~') and token.endswith('~'):
            is_sub = True
            content = token[1:-1]
        elif token.startswith('^') and token.endswith('^'):
            is_sup = True
            content = token[1:-1]

        run = p.add_run(content)
        set_run_style(run, font_name=base_font, font_size=base_size, bold=is_bold, italic=is_italic,
                      subscript=is_sub, superscript=is_sup, color=base_color)

def setup_header_footer(doc):
    """Cài đặt Header & Footer chuẩn theo FORM_DE_KIEM_TRA."""
    normal_style = doc.styles['Normal']
    normal_style.paragraph_format.tab_stops.clear_all()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    section.start_type = WD_SECTION_START.NEW_PAGE
    section.different_first_page_header_footer = False
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    # 1. Header (Bảng 1x2, rộng 18.0 cm)
    header = section.header
    for p in header.paragraphs:
        p.text = ""
    tbl_h = header.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl_h.cell(0, 0)
    c_right = tbl_h.cell(0, 1)
    c_left.width = Cm(9.0)
    c_right.width = Cm(9.0)

    p_hl = c_left.paragraphs[0]
    p_hl.paragraph_format.space_before = Pt(0)
    p_hl.paragraph_format.space_after = Pt(2)
    p_hl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_hl = p_hl.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)

    p_hr = c_right.paragraphs[0]
    p_hr.paragraph_format.space_before = Pt(0)
    p_hr.paragraph_format.space_after = Pt(2)
    p_hr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p_hr.add_run("Thầy ")
    set_run_style(r1, font_size=10, color=COLOR_HF_RED)
    r2 = p_hr.add_run("TRẦN VĂN THẠNH")
    set_run_style(r2, font_size=10, bold=True, color=COLOR_HF_RED)
    r3 = p_hr.add_run(" - 0777.470.803")
    set_run_style(r3, font_size=10, color=COLOR_HF_RED)

    add_p_border_bottom(p_hl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_bottom(p_hr, color=COLOR_HF_GREEN_HEX, sz="6")

    # 2. Footer (Bảng 1x2, rộng 18.0 cm)
    footer = section.footer
    for p in footer.paragraphs:
        p.text = ""
    tbl_f = footer.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_fl = tbl_f.cell(0, 0)
    c_fr = tbl_f.cell(0, 1)
    c_fl.width = Cm(14.5)
    c_fr.width = Cm(3.5)

    p_fl = c_fl.paragraphs[0]
    p_fl.paragraph_format.space_before = Pt(2)
    p_fl.paragraph_format.space_after = Pt(0)
    p_fl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_fl = p_fl.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)

    p_fr = c_fr.paragraphs[0]
    p_fr.paragraph_format.space_before = Pt(2)
    p_fr.paragraph_format.space_after = Pt(0)
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_field(p_fr, font_size=10, bold=True, color=COLOR_HF_RED)

    add_p_border_top(p_fl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_top(p_fr, color=COLOR_HF_GREEN_HEX, sz="6")

def build_title_block(doc, is_teacher=False, is_screen=False, base_size=10.5):
    """Khối tiêu đề theo FORM_DE_KIEM_TRA."""
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl.cell(0, 0)
    c_right = tbl.cell(0, 1)
    c_left.width = Cm(8.0)
    c_right.width = Cm(10.0)

    # Cột trái
    p_l1 = c_left.paragraphs[0]
    p_l1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l1.paragraph_format.space_before = Pt(0)
    p_l1.paragraph_format.space_after = Pt(1)
    render_rich_text(p_l1, f"*{INFO['truong']}*", base_size=11, base_bold=True)

    p_l2 = c_left.add_paragraph()
    p_l2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l2.paragraph_format.space_before = Pt(0)
    p_l2.paragraph_format.space_after = Pt(1)
    render_rich_text(p_l2, f"*{INFO['so_gddt']}*", base_size=10, base_bold=False)

    p_l3 = c_left.add_paragraph()
    p_l3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l3.paragraph_format.space_before = Pt(1)
    p_l3.paragraph_format.space_after = Pt(2)
    render_rich_text(p_l3, "_Tài liệu học tập chính thức_", base_size=9.5, base_italic=True, base_color=COLOR_GRAY)

    # Cột phải
    p_r1 = c_right.paragraphs[0]
    p_r1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r1.paragraph_format.space_before = Pt(0)
    p_r1.paragraph_format.space_after = Pt(1)
    render_rich_text(p_r1, f"*{INFO['tieude']}*", base_size=11, base_bold=True, base_color=COLOR_PRIMARY)

    p_r2 = c_right.add_paragraph()
    p_r2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r2.paragraph_format.space_before = Pt(0)
    p_r2.paragraph_format.space_after = Pt(1)
    render_rich_text(p_r2, f"*{INFO['mon']}*", base_size=10.5, base_bold=True)

    p_r3 = c_right.add_paragraph()
    p_r3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r3.paragraph_format.space_before = Pt(0)
    p_r3.paragraph_format.space_after = Pt(2)
    render_rich_text(p_r3, f"_{INFO['nam_hoc']} - {INFO['thoigian']}_", base_size=9.5, base_italic=True)

    # Dòng ghi chú phân biệt bản Giáo viên / Dạy màn hình
    if is_teacher:
        p_note = doc.add_paragraph()
        p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_note.paragraph_format.space_before = Pt(2)
        p_note.paragraph_format.space_after = Pt(2)
        render_rich_text(p_note, "*(BẢN GIÁO VIÊN -- CÓ ĐÁP ÁN VÀ LỜI GIẢI CHI TIẾT)*", base_size=10.5, base_bold=True, base_color=COLOR_CORRECT)
    elif is_screen:
        p_note = doc.add_paragraph()
        p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_note.paragraph_format.space_before = Pt(2)
        p_note.paragraph_format.space_after = Pt(2)
        render_rich_text(p_note, "*(BẢN DẠY TRÊN MÀN HÌNH -- CÓ KHÔNG GIAN TƯƠNG TÁC RỘNG RÃI)*", base_size=10.5, base_bold=True, base_color=COLOR_SCREEN_TITLE)

    # Dòng thông tin học sinh
    p_stu = doc.add_paragraph()
    p_stu.paragraph_format.space_before = Pt(3)
    p_stu.paragraph_format.space_after = Pt(2)
    render_rich_text(p_stu, "Họ và tên học sinh: ........................................................................ Lớp: ................. SBD: .................", base_size=10)

    # Dòng nguyên tử khối
    p_mass = doc.add_paragraph()
    p_mass.paragraph_format.space_before = Pt(1)
    p_mass.paragraph_format.space_after = Pt(4)
    render_rich_text(p_mass, f"_{INFO['nguyen_tu_khoi']}_", base_size=9.0, base_italic=True, base_color=COLOR_GRAY)

    # Đường phân cách nhẹ
    p_sep = doc.add_paragraph()
    p_sep.paragraph_format.space_before = Pt(0)
    p_sep.paragraph_format.space_after = Pt(4)
    add_p_border_bottom(p_sep, color=COLOR_BORDER_GREY_HEX, sz="4")

def add_part_header(doc, title_text, base_size=11.5):
    """Tiêu đề phân môn lớn."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    render_rich_text(p, f"*{title_text.upper()}*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)
    add_p_border_bottom(p, color=COLOR_HF_GREEN_HEX, sz="6")

def add_section_header(doc, section_title, base_size=10.5):
    """Tiêu đề mục nhỏ."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    render_rich_text(p, f"*{section_title}*", base_size=base_size, base_bold=True, base_color=RGBColor(27, 94, 32))

def add_solution_block(doc, sol_text, base_size=10):
    """
    Khối Lời giải chi tiết:
    - Thụt lề trái 0.6 cm.
    - Tiêu đề in đậm màu xanh lam (#1565C0).
    - TUYỆT ĐỐI KHÔNG DÙNG TABLE.
    """
    p_head = doc.add_paragraph()
    p_head.paragraph_format.left_indent = Cm(0.6)
    p_head.paragraph_format.space_before = Pt(2)
    p_head.paragraph_format.space_after = Pt(1)
    p_head.paragraph_format.keep_with_next = True
    render_rich_text(p_head, "► *Lời giải chi tiết:*", base_size=base_size, base_bold=True, base_color=COLOR_SOLUTION)

    lines = sol_text.strip().split("\n")
    for line in lines:
        if not line.strip():
            continue
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        render_rich_text(p, line, base_size=base_size)

def add_writing_lines(doc, num_lines=3):
    """Thêm các dòng chấm làm bài cho học sinh."""
    for _ in range(num_lines):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        render_rich_text(p, "........................................................................................................................................................",
                         base_size=9.5, base_color=COLOR_LINE_DOT)

def add_screen_interactive_space(doc, num_lines=4):
    """Thêm không gian tương tác cho bản dạy trên màn hình."""
    for _ in range(num_lines):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        render_rich_text(p, "─────────────────────────────────────────────────────────────────────────────",
                         base_size=9.0, base_color=RGBColor(210, 210, 210))

def render_mcq_options(doc, options, correct_letter, is_teacher, base_size=10.5):
    """
    Bố trí 4 phương án trắc nghiệm A-B-C-D:
    - TUYỆT ĐỐI KHÔNG DÙNG TABLE.
    - Dùng Tab Stops căn chỉnh chính xác.
    """
    opts = [opt.strip() for opt in options if opt.strip()]
    if not opts:
        return

    lens = [len(opt) for opt in opts]
    max_len = max(lens) if lens else 0

    if max_len <= 18:
        # 4 phương án trên 1 dòng
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(5.0), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(9.5), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(14.0), WD_TAB_ALIGNMENT.LEFT)

        for opt in opts:
            p.add_run("\t")
            letter = opt.strip()[0]
            if is_teacher and letter == correct_letter:
                render_rich_text(p, f"*{opt}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p, opt, base_size=base_size)

    elif max_len <= 45:
        # 2 dòng x 2 phương án
        p1 = doc.add_paragraph()
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.tab_stops.clear_all()
        p1.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p1.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)

        for opt in opts[:2]:
            p1.add_run("\t")
            letter = opt.strip()[0]
            if is_teacher and letter == correct_letter:
                render_rich_text(p1, f"*{opt}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p1, opt, base_size=base_size)

        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(1)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.tab_stops.clear_all()
        p2.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p2.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)

        for opt in opts[2:]:
            p2.add_run("\t")
            letter = opt.strip()[0]
            if is_teacher and letter == correct_letter:
                render_rich_text(p2, f"*{opt}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p2, opt, base_size=base_size)

    else:
        # 4 dòng riêng biệt
        for opt in opts:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            letter = opt.strip()[0]
            if is_teacher and letter == correct_letter:
                render_rich_text(p, f"*{opt}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p, opt, base_size=base_size)

def render_mcq_list(doc, questions, is_teacher, is_screen, base_size=10.5):
    """Sinh danh sách câu hỏi trắc nghiệm nhiều lựa chọn."""
    for q in questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(1)
        render_rich_text(p_q, f"*Câu {q['num']}.* {q['question']}", base_size=base_size)

        render_mcq_options(doc, q["options"], q.get("correct", ""), is_teacher, base_size=base_size)

        if is_teacher and "solution" in q and q["solution"]:
            add_solution_block(doc, q["solution"], base_size=10)
        elif is_screen:
            add_screen_interactive_space(doc, num_lines=2)

def render_tf_list(doc, questions, is_teacher, is_screen, base_size=10.5):
    """Sinh danh sách câu hỏi Đúng/Sai (dạng chùm ý a,b,c,d)."""
    for q in questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*Câu {q['num']}.* {q['intro']}", base_size=base_size)

        sol_lines = []
        for item in q["items"]:
            p_item = doc.add_paragraph()
            p_item.paragraph_format.left_indent = Cm(0.5)
            p_item.paragraph_format.space_before = Pt(1)
            p_item.paragraph_format.space_after = Pt(1)

            lbl = f"{item['label']})"
            txt = f"{lbl} {item['text']}"
            if is_teacher:
                ans_str = f" ({item['ans'].upper()})"
                render_rich_text(p_item, txt, base_size=base_size)
                r_ans = p_item.add_run(ans_str)
                set_run_style(r_ans, font_size=base_size, bold=True, color=COLOR_CORRECT if item['ans'] == "Đúng" else RGBColor(183, 28, 28))
                sol_lines.append(f"Ý {item['label']}: {item['ans']}. {item.get('explain', '')}")
            else:
                render_rich_text(p_item, f"{txt} [   ]", base_size=base_size)

        if is_teacher and sol_lines:
            add_solution_block(doc, "\n".join(sol_lines), base_size=10)
        elif is_screen:
            add_screen_interactive_space(doc, num_lines=3)

def render_short_list(doc, questions, is_teacher, is_screen, base_size=10.5):
    """Sinh danh sách câu hỏi trả lời ngắn."""
    for q in questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*Câu {q['num']}.* {q['question']}", base_size=base_size)

        if is_teacher:
            add_solution_block(doc, q["solution"], base_size=10)
        elif is_screen:
            add_screen_interactive_space(doc, num_lines=3)
        else:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Cm(0.6)
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(2)
            render_rich_text(p_ans, "Trả lời: ........................................................................................................................................", base_size=10)

def render_essay_list(doc, questions, is_teacher, is_screen, base_size=10.5):
    """Sinh danh sách câu hỏi tự luận."""
    for q in questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*Câu {q['num']}.* {q['question']}", base_size=base_size)

        if is_teacher:
            add_solution_block(doc, q["solution"], base_size=10)
        elif is_screen:
            add_screen_interactive_space(doc, num_lines=5)
        else:
            add_writing_lines(doc, num_lines=4)

def render_summary_table(doc):
    """Bảng tổng hợp đáp án nhanh cho Bản Giáo Viên ở cuối tài liệu."""
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(12)
    p_sec.paragraph_format.space_after = Pt(4)
    p_sec.paragraph_format.keep_with_next = True
    render_rich_text(p_sec, "► *BẢNG TỔNG HỢP ĐÁP ÁN TRẮC NGHIỆM KHÁCH QUAN*", base_size=11, base_bold=True, base_color=COLOR_PRIMARY)

    # Bảng 20 câu trắc nghiệm Hóa học (2 hàng x 10 cột)
    p_sub1 = doc.add_paragraph()
    p_sub1.paragraph_format.space_before = Pt(2)
    p_sub1.paragraph_format.space_after = Pt(2)
    render_rich_text(p_sub1, "*1. Đáp án trắc nghiệm Hóa học (20 câu):*", base_size=10, base_bold=True)

    tbl1 = doc.add_table(rows=2, cols=10)
    tbl1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl1, color="999999", sz="4")

    for i in range(10):
        c_q = tbl1.cell(0, i)
        c_q.width = Cm(1.7)
        set_cell_margins(c_q, 40, 40, 40, 40)
        set_cell_shading(c_q, "E8F5E9")
        p_q = c_q.paragraphs[0]
        p_q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        render_rich_text(p_q, f"*C{i+1}*", base_size=9.5, base_bold=True)

        c_a = tbl1.cell(1, i)
        c_a.width = Cm(1.7)
        set_cell_margins(c_a, 40, 40, 40, 40)
        p_a = c_a.paragraphs[0]
        p_a.alignment = WD_ALIGN_PARAGRAPH.CENTER
        render_rich_text(p_a, f"*{HOA_MCQ[i]['correct']}*", base_size=10, base_bold=True, base_color=COLOR_CORRECT)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(1)
    p_space.paragraph_format.space_after = Pt(1)

    tbl2 = doc.add_table(rows=2, cols=10)
    tbl2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl2, color="999999", sz="4")

    for i in range(10):
        idx = i + 10
        c_q = tbl2.cell(0, i)
        c_q.width = Cm(1.7)
        set_cell_margins(c_q, 40, 40, 40, 40)
        set_cell_shading(c_q, "E8F5E9")
        p_q = c_q.paragraphs[0]
        p_q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        render_rich_text(p_q, f"*C{idx+1}*", base_size=9.5, base_bold=True)

        c_a = tbl2.cell(1, i)
        c_a.width = Cm(1.7)
        set_cell_margins(c_a, 40, 40, 40, 40)
        p_a = c_a.paragraphs[0]
        p_a.alignment = WD_ALIGN_PARAGRAPH.CENTER
        render_rich_text(p_a, f"*{HOA_MCQ[idx]['correct']}*", base_size=10, base_bold=True, base_color=COLOR_CORRECT)

    # Đáp án Sinh học & Vật lí
    p_sub2 = doc.add_paragraph()
    p_sub2.paragraph_format.space_before = Pt(4)
    p_sub2.paragraph_format.space_after = Pt(2)
    render_rich_text(p_sub2, "*2. Đáp án trắc nghiệm Vật lí & Sinh học:*", base_size=10, base_bold=True)

    tbl3 = doc.add_table(rows=2, cols=3)
    tbl3.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl3, color="999999", sz="4")

    labels = ["Vật lí - Câu 5", "Sinh học - Câu 1", "Sinh học - Câu 2"]
    ans = ["A (500 W)", SINHHOC_MCQ[0]['correct'], SINHHOC_MCQ[1]['correct']]
    for i in range(3):
        c_q = tbl3.cell(0, i)
        c_q.width = Cm(5.8)
        set_cell_margins(c_q, 40, 40, 40, 40)
        set_cell_shading(c_q, "E3F2FD")
        p_q = c_q.paragraphs[0]
        p_q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        render_rich_text(p_q, f"*{labels[i]}*", base_size=9.5, base_bold=True)

        c_a = tbl3.cell(1, i)
        c_a.width = Cm(5.8)
        set_cell_margins(c_a, 40, 40, 40, 40)
        p_a = c_a.paragraphs[0]
        p_a.alignment = WD_ALIGN_PARAGRAPH.CENTER
        render_rich_text(p_a, f"*{ans[i]}*", base_size=10, base_bold=True, base_color=COLOR_CORRECT)

def build_full_document(mode="hocsinh"):
    """
    Xây dựng toàn bộ tài liệu cho 1 trong 3 chế độ:
    - 'hocsinh': Bản Học sinh
    - 'giaovien': Bản Giáo viên (có lời giải chi tiết và đáp án)
    - 'daytrenmanhinh': Bản Dạy Trên Màn Hình (không gian tương tác)
    """
    is_teacher = (mode == "giaovien")
    is_screen = (mode == "daytrenmanhinh")

    doc = docx.Document()
    setup_header_footer(doc)
    build_title_block(doc, is_teacher=is_teacher, is_screen=is_screen, base_size=10.5)

    # =========================================================================
    # PHÂN MÔN 1: NĂNG LƯỢNG (VẬT LÍ)
    # =========================================================================
    add_part_header(doc, "A. PHÂN MÔN: NĂNG LƯỢNG (VẬT LÍ)")

    for q in VATLI_QUESTIONS:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*Câu {q['id']}.* {q['question'] if 'question' in q else q['intro']}", base_size=10.5)

        qtype = q["type"]
        if qtype == "mcq":
            render_mcq_options(doc, q["options"], q.get("correct", ""), is_teacher, base_size=10.5)
            if is_teacher and "solution" in q:
                add_solution_block(doc, q["solution"], base_size=10)
            elif is_screen:
                add_screen_interactive_space(doc, num_lines=2)

        elif qtype == "tf":
            sol_lines = []
            for item in q["items"]:
                p_item = doc.add_paragraph()
                p_item.paragraph_format.left_indent = Cm(0.5)
                p_item.paragraph_format.space_before = Pt(1)
                p_item.paragraph_format.space_after = Pt(1)
                lbl = f"{item['label']})"
                txt = f"{lbl} {item['text']}"
                if is_teacher:
                    render_rich_text(p_item, txt, base_size=10.5)
                    ans_str = f" ({item['ans'].upper()})"
                    r_ans = p_item.add_run(ans_str)
                    set_run_style(r_ans, font_size=10.5, bold=True, color=COLOR_CORRECT if item['ans'] == "Đúng" else RGBColor(183, 28, 28))
                    sol_lines.append(f"Ý {item['label']}: {item['ans']}. {item.get('explain', '')}")
                else:
                    render_rich_text(p_item, f"{txt} [   ]", base_size=10.5)

            if is_teacher and sol_lines:
                add_solution_block(doc, "\n".join(sol_lines), base_size=10)
            elif is_screen:
                add_screen_interactive_space(doc, num_lines=3)

        elif qtype == "pendulum_essay":
            # Chèn hình vẽ con lắc đơn HD
            if os.path.exists(IMG_CON_LAC):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(2)
                p_img.paragraph_format.space_after = Pt(3)
                r_img = p_img.add_run()
                r_img.add_picture(IMG_CON_LAC, width=Cm(6.8))

            if is_teacher:
                add_solution_block(doc, q["solution"], base_size=10)
            elif is_screen:
                add_screen_interactive_space(doc, num_lines=5)
            else:
                add_writing_lines(doc, num_lines=5)

        elif qtype == "essay":
            if is_teacher:
                add_solution_block(doc, q["solution"], base_size=10)
            elif is_screen:
                add_screen_interactive_space(doc, num_lines=4)
            else:
                add_writing_lines(doc, num_lines=3)

    # =========================================================================
    # PHÂN MÔN 2: CHẤT VÀ SỰ BIẾN ĐỔI CHẤT (HÓA HỌC)
    # =========================================================================
    doc.add_page_break()
    add_part_header(doc, "B. PHÂN MÔN: CHẤT VÀ SỰ BIẾN ĐỔI CHẤT (HÓA HỌC)")

    # I. Kiến thức cần nhớ
    add_section_header(doc, "I. KIẾN THỨC CẦN NHỚ (BÀI 18, 19, 20, 21 - SGK KHTN 9 KNTT)")
    for item in HOA_THEORY:
        p_t = doc.add_paragraph()
        p_t.paragraph_format.space_before = Pt(3)
        p_t.paragraph_format.space_after = Pt(1)
        render_rich_text(p_t, f"*{item['title']}*", base_size=10.5, base_bold=True, base_color=RGBColor(0, 77, 64))

        for line in item['content'].split("\n"):
            p_l = doc.add_paragraph()
            p_l.paragraph_format.left_indent = Cm(0.4)
            p_l.paragraph_format.space_before = Pt(1)
            p_l.paragraph_format.space_after = Pt(1)
            render_rich_text(p_l, line, base_size=10)

    # II. Một số câu hỏi ôn tập
    add_section_header(doc, "II. MỘT SỐ CÂU HỎI ÔN TẬP")

    # 1. Trắc nghiệm nhiều lựa chọn
    add_section_header(doc, "1. Trắc nghiệm nhiều lựa chọn", base_size=10)
    render_mcq_list(doc, HOA_MCQ, is_teacher, is_screen, base_size=10.5)

    # 2. Trắc nghiệm Đúng / Sai
    add_section_header(doc, "2. Trắc nghiệm khách quan đúng – sai", base_size=10)
    render_tf_list(doc, HOA_TF, is_teacher, is_screen, base_size=10.5)

    # 3. Trả lời ngắn
    add_section_header(doc, "3. Trắc nghiệm trả lời ngắn", base_size=10)
    render_short_list(doc, HOA_SHORT, is_teacher, is_screen, base_size=10.5)

    # =========================================================================
    # PHÂN MÔN 3: VẬT SỐNG (SINH HỌC)
    # =========================================================================
    doc.add_page_break()
    add_part_header(doc, "C. PHÂN MÔN: VẬT SỐNG (SINH HỌC)")

    # 1. Trắc nghiệm nhiều lựa chọn
    add_section_header(doc, "1. Trắc nghiệm nhiều lựa chọn", base_size=10)
    render_mcq_list(doc, SINHHOC_MCQ, is_teacher, is_screen, base_size=10.5)

    # 2. Câu hỏi đếm nhận định đúng
    add_section_header(doc, "2. Câu hỏi trả lời ngắn", base_size=10)
    render_short_list(doc, SINHHOC_SHORT, is_teacher, is_screen, base_size=10.5)

    # 3. Tự luận & Sơ đồ lai
    add_section_header(doc, "3. Câu hỏi tự luận & Sơ đồ lai", base_size=10)
    render_essay_list(doc, SINHHOC_ESSAY, is_teacher, is_screen, base_size=10.5)

    # 4. Trắc nghiệm Đúng / Sai
    add_section_header(doc, "4. Trắc nghiệm khách quan đúng – sai", base_size=10)
    render_tf_list(doc, SINHHOC_TF, is_teacher, is_screen, base_size=10.5)

    # Bảng tổng hợp đáp án cho Bản Giáo Viên
    if is_teacher:
        render_summary_table(doc)

    # Lời chúc cuối tài liệu
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(16)
    p_end.paragraph_format.space_after = Pt(8)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_end, "*★ ★ ★ Chúc các em học sinh Trường THCS Trần Cao Vân ôn tập và thi tốt! ★ ★ ★*", base_size=11, base_bold=True, base_italic=True, base_color=COLOR_PRIMARY)

    return doc

def main():
    print("=" * 65)
    print("BẮT ĐẦU TẠO BỘ 3 FILE WORD KHTN 9 (THCS TRẦN CAO VÂN - HUẾ)")
    print("=" * 65)

    configs = [
        ("hocsinh", "De_Cuong_KHTN_9_Giua_HK1_TranCaoVan_HocSinh.docx", "Bản Học sinh"),
        ("giaovien", "De_Cuong_KHTN_9_Giua_HK1_TranCaoVan_GiaoVien.docx", "Bản Giáo viên"),
        ("daytrenmanhinh", "De_Cuong_KHTN_9_Giua_HK1_TranCaoVan_DayTrenManHinh.docx", "Bản Dạy Trên Màn Hình")
    ]

    for mode, filename, desc in configs:
        out_path = os.path.join(OUT_DIR, filename)
        print(f"--> Đang khởi tạo: {desc} ({filename})...")
        doc = build_full_document(mode=mode)
        doc.save(out_path)
        print(f"    [Xong] Kích thước file: {os.path.getsize(out_path):,} bytes")

        # Sao lưu sang OneDrive
        for od in ONEDRIVE_DIRS:
            if os.path.exists(od):
                try:
                    dest = os.path.join(od, filename)
                    shutil.copy2(out_path, dest)
                    print(f"    [OneDrive] Đã sao lưu sang: {dest}")
                except Exception as e:
                    print(f"    [Cảnh báo OneDrive] {e}")

    print("\n" + "=" * 65)
    print("HOÀN TẤT TẠO 3 FILE WORD KHTN 9 THÀNH CÔNG!")
    print("=" * 65)

if __name__ == "__main__":
    main()
