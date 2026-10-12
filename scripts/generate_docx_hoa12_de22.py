# -*- coding: utf-8 -*-
"""
Script tạo 2 file Word (.docx) chuẩn Form Đề Kiểm Tra cho:
ĐỀ SỐ 22 - CHƯƠNG 8: SƠ LƯỢC VỀ DÃY KIM LOẠI CHUYỂN TIẾP THỨ NHẤT VÀ PHỨC CHẤT (HÓA HỌC 12)
Mã đề thi: 122
Bao gồm:
1. Bản Học sinh: De_Kiem_Tra_Hoa_12_De22.docx
2. Bản Giáo viên: De_Kiem_Tra_Hoa_12_De22_GiaoVien.docx (Có đáp án đỏ, lời giải chi tiết xanh, bảng đáp án)
"""

import os
import sys
sys.stdout.reconfigure(encoding='utf-8')
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"

COLOR_CORAL_RED = RGBColor(211, 47, 47)      # #D32F2F
COLOR_SAGE_GREEN = "8BC390"                 # #8BC390
COLOR_DEEP_BLUE = RGBColor(21, 101, 192)    # #1565C0
COLOR_DARK_GREEN = RGBColor(46, 125, 50)    # #2E7D32
COLOR_BLACK = RGBColor(0, 0, 0)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.59)     # 1.5 cm
        s.bottom_margin = Inches(0.59)  # 1.5 cm
        s.left_margin = Inches(0.59)    # 1.5 cm
        s.right_margin = Inches(0.59)   # 1.5 cm
        s.header_distance = Inches(0.24) # 0.6 cm
        s.footer_distance = Inches(0.24) # 0.6 cm

        # --- HEADER ---
        header = s.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.text = ""
        hp.paragraph_format.space_before = Pt(0)
        hp.paragraph_format.space_after = Pt(0)

        tbl_h = header.add_table(rows=1, cols=2, width=Inches(7.08)) # 18.0 cm
        tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
        c0, c1 = tbl_h.rows[0].cells
        c0.width = Inches(3.54)
        c1.width = Inches(3.54)
        set_cell_margins(c0, 60, 60, 0, 0)
        set_cell_margins(c1, 60, 60, 0, 0)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run("Tài liệu lưu hành nội bộ")
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(10)
        r0.font.italic = True
        r0.font.color.rgb = COLOR_CORAL_RED

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1_a = p1.add_run("Thầy ")
        r1_a.font.name = "Times New Roman"
        r1_a.font.size = Pt(10)
        r1_a.font.color.rgb = COLOR_CORAL_RED
        r1_b = p1.add_run("TRẦN VĂN THẠNH")
        r1_b.font.name = "Times New Roman"
        r1_b.font.size = Pt(10)
        r1_b.font.bold = True
        r1_b.font.color.rgb = COLOR_CORAL_RED
        r1_c = p1.add_run(" - 0777.470.803")
        r1_c.font.name = "Times New Roman"
        r1_c.font.size = Pt(10)
        r1_c.font.color.rgb = COLOR_CORAL_RED

        tblPr_h = tbl_h._tbl.tblPr
        tblBorders_h = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{COLOR_SAGE_GREEN}"/>'
            f'</w:tblBorders>'
        )
        tblPr_h.append(tblBorders_h)

        # --- FOOTER ---
        footer = s.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.text = ""
        fp.paragraph_format.space_before = Pt(0)
        fp.paragraph_format.space_after = Pt(0)

        tbl_f = footer.add_table(rows=1, cols=2, width=Inches(7.08))
        tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
        fc0, fc1 = tbl_f.rows[0].cells
        fc0.width = Inches(5.31)
        fc1.width = Inches(1.77)
        set_cell_margins(fc0, 60, 60, 0, 0)
        set_cell_margins(fc1, 60, 60, 0, 0)

        fp0 = fc0.paragraphs[0]
        fp0.paragraph_format.space_before = Pt(0)
        fp0.paragraph_format.space_after = Pt(0)
        fr0 = fp0.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
        fr0.font.name = "Times New Roman"
        fr0.font.size = Pt(9.5)
        fr0.font.italic = True
        fr0.font.color.rgb = COLOR_CORAL_RED

        fp1 = fc1.paragraphs[0]
        fp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp1.paragraph_format.space_before = Pt(0)
        fp1.paragraph_format.space_after = Pt(0)
        fr1_a = fp1.add_run("Trang ")
        fr1_a.font.name = "Times New Roman"
        fr1_a.font.size = Pt(10)
        fr1_a.font.bold = True
        fr1_a.font.color.rgb = COLOR_CORAL_RED

        fldSimple = OxmlElement('w:fldSimple')
        fldSimple.set(qn('w:instr'), 'PAGE')
        r_page = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        rFont = OxmlElement('w:rFonts')
        rFont.set(qn('w:ascii'), 'Times New Roman')
        rFont.set(qn('w:hAnsi'), 'Times New Roman')
        b = OxmlElement('w:b')
        c_el = OxmlElement('w:color')
        c_el.set(qn('w:val'), 'D32F2F')
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '20')
        rPr.append(rFont)
        rPr.append(b)
        rPr.append(c_el)
        rPr.append(sz)
        r_page.append(rPr)
        t_el = OxmlElement('w:t')
        t_el.text = "1"
        r_page.append(t_el)
        fldSimple.append(r_page)
        fp1._p.append(fldSimple)

        tblPr_f = tbl_f._tbl.tblPr
        tblBorders_f = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{COLOR_SAGE_GREEN}"/>'
            f'</w:tblBorders>'
        )
        tblPr_f.append(tblBorders_f)

def add_chem_run(paragraph, text, bold=False, italic=False, color=None, font_size=11):
    parts = re.split(r'(~.*?~|\^.*?\^)', text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run()
        run.font.name = "Times New Roman"
        run.font.size = Pt(font_size)
        run.bold = bold
        run.italic = italic
        if color:
            run.font.color.rgb = color

        if part.startswith('~') and part.endswith('~'):
            run.text = part[1:-1]
            run.font.subscript = True
        elif part.startswith('^') and part.endswith('^'):
            run.text = part[1:-1]
            run.font.superscript = True
        else:
            run.text = part

def create_title_block(doc, is_teacher=False):
    lines_info = [
        ("SỞ GIÁO DỤC VÀ ĐÀO TẠO THỪA THIÊN HUẾ", "ĐỀ KIỂM TRA ĐỊNH KỲ CHƯƠNG 8 (KIM LOẠI CHUYỂN TIẾP & PHỨC CHẤT)", True, True, COLOR_DARK_GREEN),
        ("TRƯỜNG THPT HAI BÀ TRƯNG", "MÔN: HOÁ HỌC -- LỚP 12", True, True, None),
        ("ĐỀ CHÍNH THỨC", "Thời gian làm bài: 50 phút (không kể thời gian phát đề)", True, False, None),
        ("(Đề thi có 04 trang)", "Mã đề thi: 122", False, True, None)
    ]

    for left_txt, right_txt, left_bold, right_bold, right_color in lines_info:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.08), WD_TAB_ALIGNMENT.RIGHT)

        r_l = p.add_run(left_txt)
        r_l.font.name = "Times New Roman"
        r_l.font.size = Pt(11)
        r_l.bold = left_bold
        if not left_bold:
            r_l.italic = True

        p.add_run("\t")

        r_r = p.add_run(right_txt)
        r_r.font.name = "Times New Roman"
        r_r.font.size = Pt(11)
        r_r.bold = right_bold
        if right_color:
            r_r.font.color.rgb = right_color
        if not right_bold:
            r_r.italic = True

    if is_teacher:
        p_note = doc.add_paragraph()
        p_note.paragraph_format.space_before = Pt(4)
        p_note.paragraph_format.space_after = Pt(4)
        p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_n = p_note.add_run("(BẢN GIÁO VIÊN -- CÓ LỜI GIẢI CHI TIẾT)")
        r_n.font.name = "Times New Roman"
        r_n.font.size = Pt(12)
        r_n.font.bold = True
        r_n.font.color.rgb = COLOR_CORAL_RED

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(3)
    p_info.paragraph_format.space_after = Pt(2)
    r_info = p_info.add_run("Họ, tên thí sinh: ..................................................................... Lớp: ........................ SBD: ............................")
    r_info.font.name = "Times New Roman"
    r_info.font.size = Pt(10.5)
    r_info.font.italic = True

    p_mass = doc.add_paragraph()
    p_mass.paragraph_format.space_before = Pt(0)
    p_mass.paragraph_format.space_after = Pt(5)
    r_mass = p_mass.add_run("(Cho biết nguyên tử khối: H = 1; C = 12; N = 14; O = 16; Na = 23; Mg = 24; Al = 27; P = 31; S = 32; Cl = 35,5; K = 39; Ca = 40; Ti = 48; V = 51; Cr = 52; Mn = 55; Fe = 56; Co = 59; Ni = 58,7; Cu = 64; Zn = 65; Ag = 108; Cd = 112,4; Pt = 195; Au = 197; hằng số Faraday F = 96485 C/mol).")
    r_mass.font.name = "Times New Roman"
    r_mass.font.size = Pt(9.5)
    r_mass.font.italic = True

def add_part_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = COLOR_DARK_GREEN

def add_explanation(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.line_spacing = 1.15

    r_lbl = p.add_run("Lời giải chi tiết: ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(10)
    r_lbl.font.bold = True
    r_lbl.font.italic = True
    r_lbl.font.color.rgb = COLOR_DEEP_BLUE

    add_chem_run(p, text, font_size=10, italic=False, color=COLOR_DEEP_BLUE)

def render_mcq(doc, q_num, text, options, correct_key, explanation, is_teacher=False, img_path=None, img_width=Inches(3.5)):
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(3)
    p_q.paragraph_format.space_after = Pt(2)
    p_q.paragraph_format.line_spacing = 1.15

    r_lbl = p_q.add_run(f"Câu {q_num}: ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(11)
    r_lbl.bold = True

    add_chem_run(p_q, text, font_size=11)

    if img_path and os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(2)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture(img_path, width=img_width)

    max_len = max(len(opt[1]) for opt in options)
    if max_len <= 15:
        p_opts = doc.add_paragraph()
        p_opts.paragraph_format.space_before = Pt(0)
        p_opts.paragraph_format.space_after = Pt(2)
        p_opts.paragraph_format.line_spacing = 1.15
        p_opts.paragraph_format.tab_stops.add_tab_stop(Inches(1.77), WD_TAB_ALIGNMENT.LEFT)
        p_opts.paragraph_format.tab_stops.add_tab_stop(Inches(3.54), WD_TAB_ALIGNMENT.LEFT)
        p_opts.paragraph_format.tab_stops.add_tab_stop(Inches(5.31), WD_TAB_ALIGNMENT.LEFT)

        for i, (k, txt) in enumerate(options):
            is_c = (k == correct_key) and is_teacher
            r_k = p_opts.add_run(f"{k}. ")
            r_k.font.name = "Times New Roman"
            r_k.font.size = Pt(11)
            r_k.font.bold = is_c or True
            if is_c:
                r_k.font.color.rgb = COLOR_CORAL_RED
            add_chem_run(p_opts, txt, bold=is_c, color=COLOR_CORAL_RED if is_c else None, font_size=11)
            if i < 3:
                p_opts.add_run("\t")

    elif max_len <= 35:
        for row_opts in [options[:2], options[2:]]:
            p_opts = doc.add_paragraph()
            p_opts.paragraph_format.space_before = Pt(0)
            p_opts.paragraph_format.space_after = Pt(2)
            p_opts.paragraph_format.line_spacing = 1.15
            p_opts.paragraph_format.tab_stops.add_tab_stop(Inches(3.54), WD_TAB_ALIGNMENT.LEFT)

            for i, (k, txt) in enumerate(row_opts):
                is_c = (k == correct_key) and is_teacher
                r_k = p_opts.add_run(f"{k}. ")
                r_k.font.name = "Times New Roman"
                r_k.font.size = Pt(11)
                r_k.font.bold = is_c or True
                if is_c:
                    r_k.font.color.rgb = COLOR_CORAL_RED
                add_chem_run(p_opts, txt, bold=is_c, color=COLOR_CORAL_RED if is_c else None, font_size=11)
                if i == 0:
                    p_opts.add_run("\t")
    else:
        for k, txt in options:
            p_opts = doc.add_paragraph()
            p_opts.paragraph_format.space_before = Pt(0)
            p_opts.paragraph_format.space_after = Pt(2)
            p_opts.paragraph_format.line_spacing = 1.15
            is_c = (k == correct_key) and is_teacher
            r_k = p_opts.add_run(f"{k}. ")
            r_k.font.name = "Times New Roman"
            r_k.font.size = Pt(11)
            r_k.font.bold = is_c or True
            if is_c:
                r_k.font.color.rgb = COLOR_CORAL_RED
            add_chem_run(p_opts, txt, bold=is_c, color=COLOR_CORAL_RED if is_c else None, font_size=11)

    if is_teacher:
        add_explanation(doc, explanation)

def render_tf(doc, q_num, text, sub_items, explanation, is_teacher=False, img_path=None, img_width=Inches(3.5)):
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(3)
    p_q.paragraph_format.space_after = Pt(2)
    p_q.paragraph_format.line_spacing = 1.15

    r_lbl = p_q.add_run(f"Câu {q_num}: ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(11)
    r_lbl.bold = True

    add_chem_run(p_q, text, font_size=11)

    if img_path and os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(2)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture(img_path, width=img_width)

    for sub_k, sub_txt, is_true in sub_items:
        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(1.5)
        p_sub.paragraph_format.left_indent = Inches(0.2)
        p_sub.paragraph_format.line_spacing = 1.15

        r_sk = p_sub.add_run(f"{sub_k}) ")
        r_sk.font.name = "Times New Roman"
        r_sk.font.size = Pt(11)
        r_sk.bold = True

        add_chem_run(p_sub, sub_txt, font_size=11)

        if is_teacher:
            r_ans = p_sub.add_run(f" --> {'ĐÚNG' if is_true else 'SAI'}")
            r_ans.font.name = "Times New Roman"
            r_ans.font.size = Pt(10.5)
            r_ans.font.bold = True
            r_ans.font.color.rgb = COLOR_CORAL_RED

    if is_teacher:
        add_explanation(doc, explanation)

def render_short(doc, q_num, text, answer, explanation, is_teacher=False, img_path=None, img_width=Inches(3.5)):
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(3)
    p_q.paragraph_format.space_after = Pt(2)
    p_q.paragraph_format.line_spacing = 1.15

    r_lbl = p_q.add_run(f"Câu {q_num}: ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(11)
    r_lbl.bold = True

    add_chem_run(p_q, text, font_size=11)

    if img_path and os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(2)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture(img_path, width=img_width)

    if is_teacher:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(0)
        p_ans.paragraph_format.space_after = Pt(1)
        p_ans.paragraph_format.left_indent = Inches(0.25)
        r_ak = p_ans.add_run("Đáp số: ")
        r_ak.font.name = "Times New Roman"
        r_ak.font.size = Pt(10.5)
        r_ak.font.bold = True
        r_ak.font.color.rgb = COLOR_CORAL_RED
        r_av = p_ans.add_run(str(answer))
        r_av.font.name = "Times New Roman"
        r_av.font.size = Pt(10.5)
        r_av.font.bold = True
        r_av.font.color.rgb = COLOR_CORAL_RED

        add_explanation(doc, explanation)
    else:
        p_dot = doc.add_paragraph()
        p_dot.paragraph_format.space_before = Pt(0)
        p_dot.paragraph_format.space_after = Pt(2)
        r_d = p_dot.add_run("Trả lời: .....................................................................................................................................................")
        r_d.font.name = "Times New Roman"
        r_d.font.size = Pt(10)
        r_d.font.italic = True

def add_answer_table(doc):
    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_before = Pt(8)
    p_t.paragraph_format.space_after = Pt(4)
    r_t = p_t.add_run("BẢNG ĐÁP ÁN THAM KHẢO")
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(11)
    r_t.font.bold = True
    r_t.font.color.rgb = COLOR_CORAL_RED

    p_p1 = doc.add_paragraph()
    p_p1.paragraph_format.space_before = Pt(0)
    p_p1.paragraph_format.space_after = Pt(2)
    r1 = p_p1.add_run("PHẦN I: ")
    r1.font.bold = True
    r1.font.color.rgb = COLOR_DARK_GREEN
    p_p1.add_run("1.B   2.A   3.C   4.C   5.C   6.C   7.C   8.B   9.A   10.A   11.C   12.C   13.C   14.B   15.B   16.B   17.A   18.A")

    p_p2 = doc.add_paragraph()
    p_p2.paragraph_format.space_before = Pt(2)
    p_p2.paragraph_format.space_after = Pt(2)
    r2 = p_p2.add_run("PHẦN II: ")
    r2.font.bold = True
    r2.font.color.rgb = COLOR_DARK_GREEN
    p_p2.add_run("Câu 1: a-Đ, b-S, c-Đ, d-S   |   Câu 2: a-S, b-S, c-Đ, d-Đ\nCâu 3: a-Đ, b-Đ, c-Đ, d-Đ   |   Câu 4: a-Đ, b-Đ, c-S, d-Đ")

    p_p3 = doc.add_paragraph()
    p_p3.paragraph_format.space_before = Pt(2)
    p_p3.paragraph_format.space_after = Pt(2)
    r3 = p_p3.add_run("PHẦN III: ")
    r3.font.bold = True
    r3.font.color.rgb = COLOR_DARK_GREEN
    p_p3.add_run("Câu 1: 3   |   Câu 2: 6   |   Câu 3: 1   |   Câu 4: -1   |   Câu 5: 4   |   Câu 6: 17,7")

# DỮ LIỆU ĐỀ 22
P1_DATA = [
    (1, "Muối nào sau đây có khả năng làm mất màu thuốc tím trong môi trường sulfuric acid loãng?",
     [("A", "Na~2~SO~4~."), ("B", "FeSO~4~."), ("C", "MgSO~4~."), ("D", "Fe~2~(SO~4~)~3~.")],
     "B", "Ion Fe^2+^ trong FeSO~4~ có tính khử mạnh, bị ion MnO~4~^-^ (thuốc tím) oxi hoá thành Fe^3+^ trong môi trường acid, làm mất màu tím của thuốc tím:\n10FeSO~4~ + 2KMnO~4~ + 8H~2~SO~4~ -> 5Fe~2~(SO~4~)~3~ + K~2~SO~4~ + 2MnSO~4~ + 8H~2~O."),

    (2, "Từ cấu hình electron của nguyên tử Cu ở trạng thái cơ bản là [Ar]3d^10^4s^1^, xác định được cấu hình electron của ion Cu^2+^ là",
     [("A", "[Ar]3d^9^."), ("B", "[Ar]3d^8^4s^1^."), ("C", "[Ar]3d^10^."), ("D", "[Ar]3d^8^.")],
     "A", "Khi nguyên tử Cu nhường 2 electron để tạo ion Cu^2+^, nó nhường 1 electron ở phân lớp 4s trước và 1 electron ở phân lớp 3d: Cu -> Cu^2+^ + 2e. Cấu hình electron của Cu^2+^ là [Ar]3d^9^."),

    (3, "Trong dãy kim loại chuyển tiếp thứ nhất (từ Sc đến Zn), kim loại có tính dẫn điện tốt nhất là",
     [("A", "Fe."), ("B", "Ti."), ("C", "Cu."), ("D", "Mn.")],
     "C", "Trong dãy kim loại chuyển tiếp thứ nhất cũng như trong toàn bộ các kim loại (chỉ sau Ag), đồng (copper, Cu) là kim loại có độ dẫn điện và dẫn nhiệt tốt nhất."),

    (4, "Vanadium (V) có nhiều oxide có hoạt tính xúc tác như VO~2~, V~2~O~5~, V~2~O~3~,... Cấu hình electron của nguyên tử V là 1s^2^2s^2^2p^6^3s^2^3p^6^3d^3^4s^2^. Nhận xét nào sau đây là đúng về trạng thái oxi hoá của vanadium?",
     [("A", "Số oxi hoá lớn nhất là +3."),
      ("B", "Số oxi hoá lớn nhất là +2."),
      ("C", "Có nhiều số oxi hoá khác nhau trong các hợp chất."),
      ("D", "Số oxi hoá lớn nhất là +4.")],
     "C", "Do có các electron hoá trị ở cả phân lớp 3d và 4s (tổng cộng 5 electron hoá trị: 3d^3^4s^2^), vanadium thể hiện nhiều số oxi hoá khác nhau trong hợp chất như +2, +3, +4, +5 (trong đó số oxi hoá cao nhất là +5)."),

    (5, "Trong các chất sau đây, chất nào là phức chất?",
     [("A", "KAl(SO~4~)~2~.12H~2~O."),
      ("B", "FeSO~4~.(NH~4~)~2~SO~4~.6H~2~O."),
      ("C", "[Co(NH~3~)~3~(NO~2~)~3~]."),
      ("D", "KCl.MgCl~2~.6H~2~O.")],
     "C", "[Co(NH~3~)~3~(NO~2~)~3~] là một phức chất trung hoà điện gồm nguyên tử trung tâm Co^3+^ liên kết với 3 phối tử NH~3~ và 3 phối tử NO~2~^-^. Các chất còn lại là muối kép."),

    (6, "Trong phức chất [Zn(NH~3~)~4~]^2+^, số lượng phối tử liên kết với nguyên tử trung tâm là",
     [("A", "2."), ("B", "3."), ("C", "4."), ("D", "3 hoặc 4.")],
     "C", "Trong phức chất [Zn(NH~3~)~4~]^2+^, cation Zn^2+^ đóng vai trò là nguyên tử trung tâm liên kết phối trí với 4 phân tử NH~3~ (phối tử), do đó số lượng phối tử là 4."),

    (7, "Bản chất của liên kết hoá học giữa phối tử và nguyên tử trung tâm trong phức chất là liên kết",
     [("A", "ion."), ("B", "kim loại."), ("C", "cho -- nhận (phối trí)."), ("D", "hydrogen.")],
     "C", "Liên kết giữa phối tử và nguyên tử trung tâm được hình thành do phối tử đóng vai trò là base Lewis (cho cặp electron hoá trị chưa liên kết) vào orbital trống của nguyên tử trung tâm (acid Lewis), tạo thành liên kết cho -- nhận (liên kết phối trí)."),

    (8, "Trong phức chất [Fe(CN)~6~]^3-^, nguyên tử trung tâm có số phối trí là bao nhiêu?",
     [("A", "4."), ("B", "6."), ("C", "2."), ("D", "8.")],
     "B", "Ion Fe^3+^ liên kết với 6 phối tử một càng CN^-^, do đó số phối trí của nguyên tử trung tâm Fe^3+^ trong phức chất là 6."),

    (9, "Chọn phát biểu đúng nhất để giải thích sự hình thành liên kết trong phức chất [Ag(NH~3~)~2~]^+^:",
     [("A", "Liên kết trong phức chất được hình thành do phối tử NH~3~ cho cặp electron chưa liên kết vào orbital trống của nguyên tử trung tâm Ag^+^."),
      ("B", "Liên kết trong phức chất được hình thành do phối tử NH~3~ cho cặp electron đã liên kết vào orbital trống của nguyên tử trung tâm Ag^+^."),
      ("C", "Liên kết trong phức chất được hình thành do phối tử NH~3~ tạo liên kết ion với nguyên tử trung tâm Ag^+^."),
      ("D", "Liên kết trong phức chất được hình thành do nguyên tử trung tâm Ag^+^ cho cặp electron chưa liên kết vào orbital trống của phối tử NH~3~.")],
     "A", "Mỗi phối tử NH~3~ sử dụng một cặp electron tự do (chưa liên kết) trên nguyên tử nitrogen để xen phủ với một orbital lai hoá trống của cation Ag^+^, hình thành 2 liên kết cho -- nhận N->Ag."),

    (10, "Trong sự hình thành phức chất aqua của cation kim loại chuyển tiếp trong dung dịch, cation kim loại chuyển tiếp (M^n+^) thường nhận",
     [("A", "các cặp electron hoá trị riêng."),
      ("B", "các cặp electron hoá trị chung."),
      ("C", "các ion hoá trị riêng."),
      ("D", "các electron hoá trị riêng.")],
     "A", "Cation kim loại chuyển tiếp M^n+^ có các AO trống, đóng vai trò nhận các cặp electron hoá trị riêng (cặp electron tự do chưa liên kết) từ các phân tử nước H~2~O đóng vai trò phối tử."),

    (11, "Trong sự hình thành phức chất aqua của cation kim loại chuyển tiếp trong dung dịch nước, cation kim loại chuyển tiếp (M^n+^) thường nhận các cặp electron hoá trị riêng từ các phân tử",
     [("A", "H~2~."), ("B", "O~2~."), ("C", "H~2~O."), ("D", "kim loại.")],
     "C", "Phức chất aqua được tạo thành khi các cation kim loại chuyển tiếp liên kết phối trí với các phối tử là phân tử nước H~2~O thông qua cặp electron hoá trị riêng trên nguyên tử oxygen."),

    (12, "Số phối trí phổ biến nhất của các cation kim loại chuyển tiếp dãy thứ nhất trong các phức chất aqua dạng bát diện [M(OH~2~)~6~]^n+^ là",
     [("A", "2."), ("B", "4."), ("C", "6."), ("D", "8.")],
     "C", "Trong phức chất aqua dạng bát diện như [Fe(H~2~O)~6~]^2+^, [Co(H~2~O)~6~]^2+^, [Cr(H~2~O)~6~]^3+^, nguyên tử trung tâm liên kết với 6 phân tử H~2~O nên có số phối trí là 6."),

    (13, "Cho phát biểu sau: 'Khi tan trong nước, muối của các kim loại chuyển tiếp ...(1)... thành các ion. Sau đó, cation kim loại chuyển tiếp M^n+^ thường nhận các cặp electron hoá trị riêng từ các phân tử H~2~O để hình thành các liên kết cho -- nhận, tạo ra phức chất aqua có dạng tổng quát là ...(2)...'. Cụm từ cần điền vào (1) và (2) lần lượt là:",
     [("A", "điện li, [M(OH~2~)]~n~^+^."),
      ("B", "điện li, [M(OH~2~)~m~]^n^."),
      ("C", "phân li, [M(OH~2~)~m~]^n+^."),
      ("D", "phân li, [M(OH~2~)]~n~^+^.")],
     "C", "Khi tan trong nước, tinh thể muối phân li thành các cation và anion. Cation kim loại chuyển tiếp M^n+^ lập tức tạo liên kết phối trí với m phân tử H~2~O để tạo thành phức chất aqua có dạng tổng quát là [M(OH~2~)~m~]^n+^ (thường m = 6)."),

    (14, "Phức chất cis-[Pt(NH~3~)~2~Cl~2~] (cisplatin) là một hợp chất vô cơ có khả năng ức chế mạnh các tế bào ung thư tinh hoàn, buồng trứng, bàng quang,... Hãy cho biết nguyên tử trung tâm trong phức chất này là",
     [("A", "Cl^-^."), ("B", "Pt^2+^."), ("C", "OH^-^."), ("D", "NH~3~.")],
     "B", "Cisplatin có công thức [Pt(NH~3~)~2~Cl~2~]. Cation Pt^2+^ đóng vai trò là nguyên tử trung tâm, 2 phân tử NH~3~ và 2 ion Cl^-^ đóng vai trò là các phối tử."),

    (15, "Nhỏ vài giọt dung dịch HCl đặc vào ống nghiệm chứa dung dịch CuSO~4~ loãng thấy dung dịch từ màu xanh lam nhạt chuyển sang màu vàng chanh, chứng tỏ phức chất [CuCl~4~]^2-^ đã được tạo thành. Dấu hiệu nhận biết phản ứng tạo phức trên dựa vào",
     [("A", "sự ngưng tụ."), ("B", "sự biến đổi màu sắc."), ("C", "sự hoà tan chất rắn."), ("D", "sự xuất hiện kết tủa.")],
     "B", "Phản ứng thay thế phối tử: [Cu(H~2~O)~6~]^2+^ (xanh lam) + 4Cl^-^ <=> [CuCl~4~]^2-^ (vàng chanh) + 6H~2~O được nhận biết trực tiếp thông qua sự biến đổi màu sắc đặc trưng của dung dịch."),

    (16, "Cho phương trình hoá học: [Fe(H~2~O)~6~]^2+^(aq) + 6CN^-^(aq) -> [Fe(CN)~6~]^4-^(aq) + 6H~2~O(l). Ở phản ứng này, các phối tử H~2~O trong phức chất [Fe(H~2~O)~6~]^2+^ đã bị thay thế bởi phối tử nào sau đây?",
     [("A", "Cl^-^."), ("B", "CN^-^."), ("C", "OH^-^."), ("D", "NH~3~.")],
     "B", "6 phối tử H~2~O trong phức chất aqua đã bị thay thế hoàn toàn bởi 6 phối tử cyanide CN^-^ tạo phức chất hexacyanidoferrate(II) [Fe(CN)~6~]^4-^."),

    (17, "Nước có lượng đáng kể các cation Al^3+^ và Fe^3+^ được gọi là nước nhiễm phèn. Trong nước nhiễm phèn, mỗi cation này bị thuỷ phân tạo thành phức chất dạng [M(OH)~3~(H~2~O)~3~]. Chọn phát biểu sai:",
     [("A", "Cả 2 phức chất trên đều có dạng hình học là tứ diện đều."),
      ("B", "Phối tử OH^-^ và phối tử H~2~O tạo phức được với ion trung tâm là do đều có cặp electron tự do."),
      ("C", "Khi tạo phức, các ion Al^3+^ và Fe^3+^ đóng vai trò là acid Lewis do có orbital trống nhận cặp electron."),
      ("D", "Nước bị nhiễm phèn thường có độ pH thấp (môi trường acid) do quá trình thuỷ phân sinh ra ion H^+^ trong dung dịch.")],
     "A", "Phát biểu A sai vì phức chất [M(OH)~3~(H~2~O)~3~] có tổng cộng 6 phối tử (3 OH^-^ và 3 H~2~O) nên có cấu trúc hình học dạng bát diện (octahedral), không phải tứ diện đều."),

    (18, "Cho các phản ứng hoá học sau:\n"
         "(1) [Co(OH~2~)~6~]^2+^(aq) + 6NH~3~(aq) <=> [Co(NH~3~)~6~]^2+^(aq) + 6H~2~O(l)\n"
         "(2) 2[Co(NH~3~)~6~]^2+^(aq) + H~2~O~2~(aq) <=> 2[Co(NH~3~)~6~]^3+^(aq) + 2OH^-^(aq)\n"
         "(3) Au(s) + NO~3~^-^(aq) + 4H^+^(aq) + 4Cl^-^(aq) <=> [AuCl~4~]^-^(aq) + NO(g) + 2H~2~O(l)\n"
         "(4) [Ag(NH~3~)~2~]^+(aq) + 2CN^-^(aq) <=> [Ag(CN)~2~]^-^(aq) + 2NH~3~(aq)\n"
         "(5) [Fe(OH~2~)~6~]^3+^(aq) + 2Cl^-^(aq) <=> [Fe(OH~2~)~4~Cl~2~]^+(aq) + 2H~2~O(l)\n"
         "Những phản ứng xảy ra sự thay thế phối tử trong phức chất là",
     [("A", "(1), (4), (5)."), ("B", "(2), (3), (5)."), ("C", "(1), (2), (4)."), ("D", "(2), (4), (5).")],
     "A", "Các phản ứng thay thế phối tử:\n- (1): NH~3~ thay thế H~2~O trong phức cobalt(II).\n- (4): CN^-^ thay thế NH~3~ trong phức silver(I).\n- (5): Cl^-^ thay thế một phần H~2~O trong phức iron(III).\nPhản ứng (2) là phản ứng oxi hoá -- khử (Co^2+^ lên Co^3+^). Phản ứng (3) là phản ứng hoà tan oxi hoá -- khử vàng tạo phức.")
]

P2_DATA = [
    (1, "Mỗi phát biểu sau đây về heme B (thành phần cấu tạo nên hồng cầu) là đúng hay sai?",
     [("a", "Heme B là thành phần không thể thiếu cấu tạo nên hemoglobin trong tế bào hồng cầu.", True),
      ("b", "Heme B là phức chất vòng càng của ion Fe(III).", False),
      ("c", "Heme B có vai trò sống còn trong việc vận chuyển và cung cấp oxygen đến các tế bào và mô trong cơ thể.", True),
      ("d", "Máu của các loài động vật giáp xác (như tôm, cua) có chứa hàm lượng lớn heme B.", False)],
     None, None,
     "a) Đúng vì hemoglobin gồm phân tử protein globin liên kết với nhóm ngoại heme B.\n"
     "b) Sai vì trong heme B nguyên tử trung tâm là ion Fe(II) (Fe^2+^).\n"
     "c) Đúng vì Fe^2+^ trong heme B kết hợp thuận nghịch với O~2~ tạo oxyhemoglobin vận chuyển oxygen từ phổi đến mô.\n"
     "d) Sai vì máu của động vật giáp xác và thân mềm không chứa heme B (sắt) mà chứa hemocyanin (phức chất của đồng, Cu) tạo ra màu xanh lam đặc trưng."),

    (2, "Tiến hành hai thí nghiệm khảo sát sự tạo phức của ion copper(II):\n"
        "- Thí nghiệm 1 (ở 0 °C): Cho 1 mL dung dịch CuSO~4~ 0,5% màu xanh lam nhạt vào ống nghiệm. Thêm từ từ đến hết 2 mL dung dịch HCl đặc không màu vào ống nghiệm đó, thu được dung dịch chuyển sang màu vàng chanh do cân bằng: [Cu(OH~2~)~6~]^2+^(aq) + 4Cl^-^(aq) <=> [CuCl~4~]^2-^(aq) + 6H~2~O(l) có K~C~ = 4,18.10^5^.\n"
        "- Thí nghiệm 2 (ở 20 °C): Cho 1 mL dung dịch CuSO~4~ 0,5% màu xanh lam nhạt vào ống nghiệm. Thêm từ từ đến hết 2 mL dung dịch NaCl bão hoà không màu vào ống nghiệm đó, thu được dung dịch có màu xanh lam nhạt hơn ban đầu.",
     [("a", "Biểu thức tính hằng số cân bằng K~C~ của phản ứng ở thí nghiệm 1 là K~C~ = ([[CuCl~4~]^2-^].[H~2~O]^6^) / ([[Cu(OH~2~)~6~]^2+^].[Cl^-^]^4^).", False),
      ("b", "Trong thí nghiệm 1, phản ứng nghịch diễn ra thuận lợi hơn nhiều so với phản ứng thuận.", False),
      ("c", "Trong thí nghiệm 1, nồng độ anion Cl^-^ càng cao thì cân bằng càng chuyển dịch theo chiều thuận tạo phức màu vàng chanh.", True),
      ("d", "Trong thí nghiệm 2, không có dấu hiệu rõ ràng của phản ứng hình thành phức chất [CuCl~4~]^2-^.", True)],
     None, None,
     "a) Sai vì trong dung dịch loãng, nồng độ dung môi nước H~2~O được coi là hằng số và quy ước bằng 1, do đó không xuất hiện trong biểu thức K~C~: K~C~ = [[CuCl~4~]^2-^] / ([[Cu(OH~2~)~6~]^2+^].[Cl^-^]^4^).\n"
     "b) Sai vì hằng số cân bằng K~C~ = 4,18.10^5^ rất lớn (K~C~ >> 1) nên phản ứng thuận diễn ra rất mạnh và thuận lợi.\n"
     "c) Đúng vì theo nguyên lí chuyển dịch cân bằng Le Chatelier, tăng nồng độ chất phản ứng (Cl^-^) làm cân bằng chuyển dịch theo chiều thuận tạo phức [CuCl~4~]^2-^.\n"
     "d) Đúng vì dung dịch ở thí nghiệm 2 chỉ nhạt màu do pha loãng, hoàn toàn không chuyển sang màu vàng chanh đặc trưng của [CuCl~4~]^2-^."),

    (3, "Phức chất cisplatin có công thức cấu tạo như hình bên dưới và được sử dụng rộng rãi làm thuốc hoá trị chống ung thư:",
     [("a", "Dạng hình học của phân tử cisplatin là dạng vuông phẳng.", True),
      ("b", "Nguyên tử trung tâm của phức chất cisplatin là ion Pt^2+^.", True),
      ("c", "Trong phân tử cisplatin, các phối tử liên kết với nguyên tử trung tâm là NH~3~ và Cl^-^.", True),
      ("d", "Phân tử phức chất cisplatin có tổng cộng 4 phối tử.", True)],
     "San_Pham/images_crop_de22/p2_cau3_cisplatin.png", Inches(3.5),
     "a) Đúng vì ion Pt^2+^ cấu hình d^8^ lai hoá dsp^2^ tạo phức vuông phẳng (square planar).\n"
     "b) Đúng vì số oxi hoá của platin trong phân tử là +2 (Pt^2+^).\n"
     "c) Đúng vì cisplatin chứa 2 phối tử trung hoà NH~3~ và 2 phối tử mang điện âm Cl^-^.\n"
     "d) Đúng vì có 2 phối tử ammine và 2 phối tử chlorido, tổng cộng là 4 phối tử."),

    (4, "Metalloporphyrin có cấu trúc tương tự heme và được sử dụng trong y học để kiểm soát chứng tăng bilirubin máu ở trẻ sơ sinh (nguyên nhân gây vàng da ở trẻ sơ sinh). Cấu trúc của phức chất metalloporphyrin được cho như hình bên dưới:",
     [("a", "Ion trung tâm của phức chất metalloporphyrin là cation Fe^2+^.", True),
      ("b", "Số phối trí của nguyên tử trung tâm trong phức chất trên là 4.", True),
      ("c", "Metalloporphyrin là một loại muối kép chứ không phải là phức chất.", False),
      ("d", "Các liên kết phối trí trong metalloporphyrin được tạo thành do các nguyên tử N dùng cặp electron chưa liên kết cho vào orbital trống của ion Fe^2+^.", True)],
     "San_Pham/images_crop_de22/p2_cau4_metalloporphyrin.png", Inches(2.8),
     "a) Đúng vì ion trung tâm ở tâm vòng porphyrin là cation Fe^2+^.\n"
     "b) Đúng vì có 4 liên kết phối trí từ 4 nguyên tử N của vòng porphyrin đến ion Fe^2+^.\n"
     "c) Sai vì metalloporphyrin là một phức chất vòng càng (chelate complex) hữu cơ kim loại điển hình.\n"
     "d) Đúng vì nguyên tử N có cặp electron tự do đóng vai trò là base Lewis cho vào các orbital lai hoá trống của ion Fe^2+^ (acid Lewis).")
]

P3_DATA = [
    (1, "Cho các phức chất sau: [Co(NH~3~)~6~]^3+^, [Ti(OH~2~)~6~]^3+^, [Zn(CN)~4~]^2-^, [Cd(CN)~4~]^2-^, [Mo(CO)~6~]. Có bao nhiêu phức chất có dạng hình học bát diện?",
     "3",
     "Có 3 phức chất có 6 phối tử tạo dạng hình học bát diện là:\n"
     "- [Co(NH~3~)~6~]^3+^\n"
     "- [Ti(OH~2~)~6~]^3+^\n"
     "- [Mo(CO)~6~]\n"
     "(Hai phức chất [Zn(CN)~4~]^2-^ và [Cd(CN)~4~]^2-^ có 4 phối tử tạo cấu trúc tứ diện)."),

    (2, "Cho dung dịch NH~3~ đặc đến dư vào dung dịch CoCl~2~, thu được dung dịch có màu hồng chuyển dần sang nâu. Hiện tượng này được giải thích là do tất cả các phối tử H~2~O trong phức chất [Co(OH~2~)~6~]^2+^ đã được thay thế hoàn toàn bằng phối tử NH~3~. Có bao nhiêu liên kết phối trí Co-N có trong phân tử phức chất mới được tạo thành?",
     "6",
     "Phản ứng thay thế phối tử:\n"
     "[Co(OH~2~)~6~]^2+^ + 6NH~3~ -> [Co(NH~3~)~6~]^2+^ + 6H~2~O\n"
     "Mỗi phân tử NH~3~ liên kết với ion Co^2+^ qua 1 liên kết phối trí Co-N. Do có 6 phối tử NH~3~ nên có tất cả 6 liên kết Co-N."),

    (3, "Phức chất [Pt(NH~3~)~4~]^2+^ có thể bị thế bởi 1 phối tử Cl^-^ để tạo thành phức chất mới. Điện tích của phức chất tạo thành là bao nhiêu?",
     "1",
     "Phương trình thay thế phối tử:\n"
     "[Pt(NH~3~)~4~]^2+^ + Cl^-^ -> [Pt(NH~3~)~3~Cl]^+^ + NH~3~\n"
     "Trong phức chất [Pt(NH~3~)~3~Cl]^+^:\n"
     "Điện tích = Điện tích(Pt^2+^) + 3*Điện tích(NH~3~) + Điện tích(Cl^-^) = (+2) + 0 + (-1) = +1.\n"
     "Đáp số: 1 (hoặc +1)."),

    (4, "Cisplatin là một phức chất của Pt^2+^ được sử dụng rộng rãi trong phác đồ điều trị ung thư. Cisplatin có thể được tổng hợp theo sơ đồ sau:\n"
        "[PtCl~4~]^2-^ -(+NH~3~, -Cl^-)-> [PtCl~3~(NH~3~)]^X^ -(+NH~3~, -Cl^-)-> [PtCl~2~(NH~3~)~2~]\n"
        "Xác định giá trị của điện tích X.",
     "-1",
     "Trong ion phức trung gian [PtCl~3~(NH~3~)]^X^:\n"
     "- Nguyên tử trung tâm là Pt^2+^ (điện tích +2).\n"
     "- 3 phối tử Cl^-^ (mỗi phối tử mang điện tích -1).\n"
     "- 1 phối tử trung hoà NH~3~ (điện tích 0).\n"
     "=> Điện tích X = (+2) + 3*(-1) + 0 = -1.\n"
     "Đáp số: -1."),

    (5, "Phức chất [MA~x~B~y~] có dạng hình học vuông phẳng, trong đó M là nguyên tử trung tâm; A và B là các phối tử một càng; x và y là số lượng phối tử của A và B. Xác định giá trị của tổng x + y.",
     "4",
     "Phức chất có cấu trúc hình học vuông phẳng luôn có số phối trí bằng 4 đối với các phối tử một càng.\n"
     "Do đó: x + y = 4.\n"
     "Đáp số: 4."),

    (6, "Để tiến hành mạ đồng cho một tấm huy chương hình trụ tròn (ruột bằng sắt) có bán kính đáy r = 2,5 cm và chiều cao h = 0,3 cm sao cho lớp mạ đồng phủ đều mọi bề mặt với độ dày 0,1 cm. Người ta tiến hành điện phân dung dịch CuSO~4~ 0,5 M dư với anode bằng đồng tinh khiết ở cường độ dòng điện không đổi I = 2 A. Biết hiệu suất điện phân là 100%, khối lượng riêng của đồng D = 8,95 g/cm^3^, lấy π = 3,14; M(Cu) = 64 g/mol; hằng số Faraday F = 96485 C/mol. Tính thời gian t (theo đơn vị giờ) để hoàn thành quá trình mạ (làm tròn kết quả đến hàng phần mười).",
     "17,7",
     "- Thể tích tấm huy chương trước khi mạ:\n"
     "V~1~ = π.r^2^.h = 3,14 * (2,5)^2^ * 0,3 = 5,8875 cm^3^.\n"
     "- Kích thước sau khi mạ lớp đồng dày 0,1 cm:\n"
     "+ Bán kính đáy mới: R = 2,5 + 0,1 = 2,6 cm.\n"
     "+ Chiều cao mới (mạ cả 2 mặt đáy): H = 0,3 + 2 * 0,1 = 0,5 cm.\n"
     "V~2~ = π.R^2^.H = 3,14 * (2,6)^2^ * 0,5 = 10,6132 cm^3^.\n"
     "- Thể tích lượng đồng mạ bám vào huy chương:\n"
     "V(Cu) = V~2~ - V~1~ = 10,6132 - 5,8875 = 4,7257 cm^3^.\n"
     "- Khối lượng đồng đã mạ:\n"
     "m(Cu) = V(Cu) * D = 4,7257 * 8,95 ≈ 42,2950 gam.\n"
     "- Số mol Cu: n(Cu) = 42,2950 / 64 ≈ 0,66086 mol.\n"
     "- Ở cathode: Cu^2+^ + 2e -> Cu => n(e) = 2 * n(Cu) = 2 * 0,66086 = 1,32172 mol.\n"
     "- Thời gian điện phân:\n"
     "t(giây) = (n(e) * F) / I = (1,32172 * 96485) / 2 ≈ 63763 giây.\n"
     "t(giờ) = 63763 / 3600 ≈ 17,712 giờ ≈ 17,7 giờ.")
]

def build_exam_docx(is_teacher=False):
    doc = docx.Document()
    add_header_footer(doc)
    create_title_block(doc, is_teacher=is_teacher)

    # PHẦN I
    add_part_heading(doc, "PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn. Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án.")
    for q_num, text, options, c_key, expl in P1_DATA:
        render_mcq(doc, q_num, text, options, c_key, expl, is_teacher=is_teacher)

    # PHẦN II
    add_part_heading(doc, "PHẦN II. Câu trắc nghiệm đúng sai. Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.")
    for q_num, text, sub_items, img_path, img_w, expl in P2_DATA:
        render_tf(doc, q_num, text, sub_items, expl, is_teacher=is_teacher, img_path=img_path, img_width=img_w)

    # PHẦN III
    add_part_heading(doc, "PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn. Thí sinh trả lời từ câu 1 đến câu 6.")
    for q_num, text, ans, expl in P3_DATA:
        img_p = "San_Pham/images_crop_de22/p3_cau4_so_do_cisplatin.png" if q_num == 4 else None
        img_w = Inches(3.8) if q_num == 4 else Inches(3.5)
        render_short(doc, q_num, text, ans, expl, is_teacher=is_teacher, img_path=img_p, img_width=img_w)

    if is_teacher:
        add_answer_table(doc)

    out_name = "De_Kiem_Tra_Hoa_12_De22_GiaoVien.docx" if is_teacher else "De_Kiem_Tra_Hoa_12_De22.docx"
    out_path = os.path.join(OUT_DIR, out_name)
    doc.save(out_path)
    print(f"Đã lưu: {out_path}")

def main():
    print("--- Bắt đầu tạo 2 file Word cho ĐỀ SỐ 22 (Mã 122) ---")
    build_exam_docx(is_teacher=False)
    build_exam_docx(is_teacher=True)
    print("--- Hoàn tất sinh file Word cho ĐỀ SỐ 22! ---")

if __name__ == "__main__":
    main()
