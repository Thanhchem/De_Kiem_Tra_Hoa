# -*- coding: utf-8 -*-
"""
Script biên soạn trọn bộ 3 file:
1. Đề cương Học sinh (.docx)
2. Đề cương Giáo viên có lời giải chi tiết (.docx)
3. Slide Dạy trên màn hình Tivi 65 inch tỉ lệ 16:9 (.pdf và .docx) - KHÔNG CÓ ĐÁP ÁN

ĐỀ CƯƠNG ÔN TẬP THI GIỮA KÌ 1 KHỐI 9
MÔN: KHOA HỌC TỰ NHIÊN 9 (PHÂN MÔN VẬT LÍ / KHTN 2)
TRƯỜNG THCS NGUYỄN HOÀNG (NĂM HỌC 2026 - 2027)
"""

import sys
if sys.stdout:
    sys.stdout.reconfigure(encoding='utf-8')
if sys.stderr:
    sys.stderr.reconfigure(encoding='utf-8')

import os
import shutil
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

# ==============================================================================
# HẰNG SỐ MÀU SẮC & ĐỊNH DẠNG THEO GEMINI.md (FORM ĐỀ KIỂM TRA)
# ==============================================================================
COLOR_PRIMARY_HEX = "D32F2F"       # Đỏ san hô (Coral Red)
COLOR_SECONDARY_HEX = "8BC390"     # Xanh xô thơm (Sage Green)
COLOR_TEXT_HEX = "1A1A1A"          # Đen than nhẹ (Off-black)
COLOR_SOLUTION_TITLE_HEX = "1565C0"# Xanh dương đậm cho tiêu đề lời giải
COLOR_GRAY_HEX = "757575"

COLOR_PRIMARY = RGBColor(0xD3, 0x2F, 0x2F)
COLOR_SECONDARY = RGBColor(0x8B, 0xC3, 0x90)
COLOR_TEXT = RGBColor(0x1A, 0x1A, 0x1A)
COLOR_SOLUTION_TITLE = RGBColor(0x15, 0x65, 0xC0)
COLOR_GRAY = RGBColor(0x75, 0x75, 0x75)

FONT_NAME = "Times New Roman"

BASE_DIR = r"c:\Antigravity_Thanh"
SAN_PHAM_DIR = os.path.join(BASE_DIR, "San_Pham")
SCRATCH_DIR = os.path.join(BASE_DIR, "scratch")

ONEDRIVE_DIR1 = r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn9\KHTN9-DEKIEMTRADINHKI"
ONEDRIVE_DIR2 = r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn9"

os.makedirs(SAN_PHAM_DIR, exist_ok=True)
os.makedirs(SCRATCH_DIR, exist_ok=True)

# ==============================================================================
# DỮ LIỆU ĐỀ CƯƠNG KHTN 9 (PHÂN MÔN VẬT LÍ)
# ==============================================================================

MCQ_DATA = [
    {
        "id": 1,
        "question": "Dụng cụ nào dưới đây được sử dụng để quan sát và nghiên cứu tế bào?",
        "options": [
            "A. Tiêu bản nhiễm sắc thể người",
            "B. Lăng kính",
            "C. Thấu kính hội tụ",
            "D. Thấu kính phân kì"
        ],
        "answer": "A",
        "solution": "Trong nghiên cứu tế bào, kính hiển vi quang học kết hợp với *tiêu bản nhiễm sắc thể người* (hoặc tiêu bản tế bào) là dụng cụ chuyên dụng trực tiếp để quan sát cấu trúc tế bào và bộ NST. Do đó chọn phương án A."
    },
    {
        "id": 2,
        "question": "Thấu kính trong thí nghiệm về sự tạo ảnh được làm từ chất liệu nào?",
        "options": [
            "A. Kim loại",
            "B. Thủy tinh",
            "C. Cao su",
            "D. Gỗ"
        ],
        "answer": "B",
        "solution": "Thấu kính là một khối chất trong suốt đồng tính (thường được chế tạo bằng *thủy tinh* hoặc nhựa quang học trong suốt) giới hạn bởi hai mặt cong hoặc một mặt phẳng và một mặt cong. Do đó chọn phương án B."
    },
    {
        "id": 3,
        "question": "Khi làm báo cáo khoa học, mục tiêu nghiên cứu là gì?",
        "options": [
            "A. Xác định các phương pháp thực hiện thí nghiệm",
            "B. Tóm tắt kết quả nghiên cứu",
            "C. Xác định điều cần nghiên cứu và thử nghiệm",
            "D. Đề xuất các phương pháp mới"
        ],
        "answer": "C",
        "solution": "Mục tiêu nghiên cứu trong một đề tài/báo cáo khoa học là làm rõ *điều cần nghiên cứu và thử nghiệm* nhằm giải quyết vấn đề đặt ra. Do đó chọn phương án C."
    },
    {
        "id": 4,
        "question": "Đồng hồ đo điện đa năng có thể đo những đại lượng nào?",
        "options": [
            "A. Cường độ dòng điện",
            "B. Hiệu điện thế",
            "C. Điện trở",
            "D. Cả cường độ dòng điện, hiệu điện thế và điện trở"
        ],
        "answer": "D",
        "solution": "Đồng hồ vạn năng (đa năng hiện số) là thiết bị đo điện linh hoạt, có thể đo được cả *cường độ dòng điện (A/mA), hiệu điện thế (V/mV) và điện trở (Ω/kΩ)*. Do đó chọn phương án D."
    },
    {
        "id": 5,
        "question": "Bộ dụng cụ tìm hiểu tính chất ảnh qua thấu kính không có thiết bị nào sau đây:",
        "options": [
            "A. Thấu kính phân kì",
            "B. Màn chắn",
            "C. Gương phẳng",
            "D. Khe hình chữ F"
        ],
        "answer": "C",
        "solution": "Bộ thí nghiệm khảo sát ảnh qua thấu kính gồm nguồn sáng, vật sáng (vật hình chữ F hoặc ngọn nến), thấu kính (hội tụ/phân kì), màn hứng ảnh (màn chắn) và giá quang học. Bộ dụng cụ này *không có gương phẳng*. Do đó chọn phương án C."
    },
    {
        "id": 6,
        "question": "Một chất điểm có khối lượng m đang chuyển động với vận tốc v thì động năng của nó là",
        "options": [
            "A. W_d = (1/2).m.v^2",
            "B. W_d = m.v^2",
            "C. W_d = 2.m.v^2",
            "D. W_d = (1/2).m.v"
        ],
        "answer": "A",
        "solution": "Động năng của một vật khối lượng m chuyển động với tốc độ v được xác định bằng công thức: *W_d = (1/2).m.v^2*. Do đó chọn phương án A."
    },
    {
        "id": 7,
        "question": "Công thức tính công suất nào dưới đây là đúng?",
        "options": [
            "A. A = P / t",
            "B. A = F . v",
            "C. P = F . v",
            "D. P = A . t"
        ],
        "answer": "C",
        "solution": "Công suất P = A / t = (F . s) / t = *F . v* (với v là tốc độ chuyển động đều theo phương của lực F). Do đó chọn phương án C."
    },
    {
        "id": 8,
        "question": "Nếu chọn mặt đất làm mốc để tính thế năng thì trong các vật sau đây vật nào không có thế năng?",
        "options": [
            "A. Máy bay đang bay trên cao.",
            "B. Chiếc lá đang rơi.",
            "C. Một người đứng trên tầng ba của tòa nhà.",
            "D. Quả bóng nằm yên trên sàn nhà."
        ],
        "answer": "D",
        "solution": "Quả bóng nằm yên trên sàn nhà (ngang mặt đất) có độ cao h = 0 m so với mốc thế năng, do đó thế năng trọng trường của nó bằng 0 (*không có thế năng*). Do đó chọn phương án D."
    },
    {
        "id": 9,
        "question": "Nếu khối lượng của vật giảm đi 2 lần, còn tốc độ của vật tăng lên 4 lần thì động năng của vật",
        "options": [
            "A. tăng lên 2 lần.",
            "B. tăng lên 8 lần.",
            "C. giảm đi 2 lần.",
            "D. giảm đi 8 lần."
        ],
        "answer": "B",
        "solution": "Ta có: W_d = (1/2).m.v^2. Khi m' = m / 2 và v' = 4v thì:\nW_d' = (1/2) . (m / 2) . (4v)^2 = (1/2) . (m / 2) . 16v^2 = 8 . [(1/2).m.v^2] = *8.W_d* (tăng lên 8 lần). Do đó chọn phương án B."
    },
    {
        "id": 10,
        "question": "Công suất có các đơn vị đo là:",
        "options": [
            "A. J, BTU, cal.",
            "B. J, cal, W.",
            "C. W, HP, BTU.",
            "D. W, HP, BTU/h."
        ],
        "answer": "D",
        "solution": "Các đơn vị đo công suất gồm Oát (*W*, kW, MW), mã lực (*HP*) và *BTU/h* (thường dùng trong máy điều hòa, 1 BTU/h ≈ 0,293 W). Còn J, cal, BTU là đơn vị của năng lượng. Do đó chọn phương án D."
    },
    {
        "id": 11,
        "question": "Một chất điểm có khối lượng m đang chuyển động với vận tốc υ thì động năng của nó là:",
        "options": [
            "A. (1/2).m^2.υ^2",
            "B. (1/2).m.υ^2",
            "C. 2.m.υ^2",
            "D. m.υ^2"
        ],
        "answer": "B",
        "solution": "Động năng của chất điểm là: *W_d = (1/2).m.υ^2*. Do đó chọn phương án B."
    },
    {
        "id": 12,
        "question": "Nếu chọn mặt đất làm mốc để tính thế năng thì trong các vật sau đây vật nào không có thế năng?",
        "options": [
            "A. Máy bay đang bay.",
            "B. Xe máy đang chuyển động trên mặt đường.",
            "C. Chiếc lá đang rơi.",
            "D. Quyển sách đặt trên bàn."
        ],
        "answer": "B",
        "solution": "Xe máy đang chuyển động trên mặt đường (nằm ở mốc độ cao h = 0) nên thế năng trọng trường bằng 0 (*không có thế năng*). Do đó chọn phương án B."
    },
    {
        "id": 13,
        "question": "Chọn câu trả lời sai: Công suất có đơn vị là?",
        "options": [
            "A. Mã lực",
            "B. J/s",
            "C. W",
            "D. kWh"
        ],
        "answer": "D",
        "solution": "kWh (kilôoát giờ) là đơn vị đo năng lượng (công của dòng điện, 1 kWh = 3.600.000 J), *không phải đơn vị công suất*. Do đó chọn phương án D."
    },
    {
        "id": 14,
        "question": "Dụng cụ thí nghiệm nào sau đây dùng để phân tán nhiệt khi đốt?",
        "options": [
            "A. Bát sứ.",
            "B. Phễu.",
            "C. Bình cầu.",
            "D. Lưới tản nhiệt."
        ],
        "answer": "D",
        "solution": "*Lưới tản nhiệt* (lưới thép có tấm amiăng ở giữa) dùng để phân tán đều ngọn lửa đèn cồn, giúp đáy bình thủy tinh được gia nhiệt đồng đều và không bị vỡ. Do đó chọn phương án D."
    },
    {
        "id": 15,
        "question": "Trong các trường hợp sau, trường hợp nào vật có cả động năng và thế năng? Chọn mốc thế năng tại mặt đất.",
        "options": [
            "A. Một máy bay đang chuyển động trên đường băng của sân bay.",
            "B. Một ô tô đang đỗ trong bến xe.",
            "C. Một máy bay đang bay trên cao.",
            "D. Một ô tô đang chuyển động trên đường."
        ],
        "answer": "C",
        "solution": "Máy bay đang bay trên cao vừa có độ cao so với mặt đất (h > 0 ==> có thế năng), vừa có tốc độ chuyển động (v > 0 ==> có động năng). Do đó chọn phương án C."
    },
    {
        "id": 16,
        "question": "Cơ năng của một vật đang chuyển động trên mặt đất là:",
        "options": [
            "A. là tích của động năng và thế năng của vật.",
            "B. là hiệu của động năng và thế năng của vật.",
            "C. chỉ bao gồm động năng của vật.",
            "D. chỉ bao gồm thế năng của vật."
        ],
        "answer": "C",
        "solution": "Khi vật chuyển động ngay trên mặt đất (chọn mốc thế năng tại mặt đất thì h = 0 ==> W_t = 0), cơ năng W = W_d + W_t = W_d (*chỉ bao gồm động năng của vật*). Do đó chọn phương án C."
    },
    {
        "id": 17,
        "question": "Một lực F tác dụng vào vật làm vật dịch chuyển quãng đường s theo hướng của lực thì công của lực F được tính bằng công thức:",
        "options": [
            "A. A = F / s",
            "B. A = F . s",
            "C. A = s / F",
            "D. A = F . v"
        ],
        "answer": "B",
        "solution": "Công của lực F khi làm vật dịch chuyển một quãng đường s theo phương của lực là *A = F . s*. Do đó chọn phương án B."
    },
    {
        "id": 18,
        "question": "Trong môn Khoa học tự nhiên lớp 9, dụng cụ nào sau đây là dụng cụ thí nghiệm quang học?",
        "options": [
            "A. Điện kế.",
            "B. Đèn laser.",
            "C. Lưới tản nhiệt.",
            "D. Bát sứ."
        ],
        "answer": "B",
        "solution": "*Đèn laser* phát ra chùm tia sáng song song định hướng cao, là dụng cụ thí nghiệm quang học quen thuộc trong bài khúc xạ, phản xạ và thấu kính. Do đó chọn phương án B."
    },
    {
        "id": 19,
        "question": "Trong các trường hợp sau, trường hợp nào vật không có cả động năng và thế năng? Chọn mốc thế năng tại mặt đất.",
        "options": [
            "A. Một máy bay đang chuyển động trên đường băng của sân bay.",
            "B. Một ô tô đang đỗ trong bến xe.",
            "C. Một máy bay đang bay trên cao.",
            "D. Một ô tô đang chuyển động trên đường."
        ],
        "answer": "B",
        "solution": "Ô tô đang đỗ yên trong bến xe (v = 0 ==> W_d = 0) và nằm trên mặt đất (h = 0 ==> W_t = 0) nên *không có cả động năng và thế năng*. Do đó chọn phương án B."
    },
    {
        "id": 20,
        "question": "Cơ năng của một vật đang đứng yên cách mặt đất 10 m là:",
        "options": [
            "A. là tích của động năng và thế năng của vật.",
            "B. là hiệu của động năng và thế năng của vật.",
            "C. chỉ bao gồm động năng của vật.",
            "D. chỉ bao gồm thế năng của vật."
        ],
        "answer": "D",
        "solution": "Vật đứng yên nên v = 0 ==> W_d = 0. Vật ở độ cao h = 10 m ==> có thế năng W_t. Do đó cơ năng W = W_t (*chỉ bao gồm thế năng của vật*). Do đó chọn phương án D."
    },
    {
        "id": 21,
        "question": "Một lực F tác dụng vào vật làm vật dịch chuyển quãng đường s theo hướng của lực trong thời gian t thì công suất của lực F được tính bằng công thức:",
        "options": [
            "A. P = A / t = F . v",
            "B. P = A . t",
            "C. P = F / v",
            "D. P = t / A"
        ],
        "answer": "A",
        "solution": "Công suất đặc trưng cho tốc độ sinh công: *P = A / t = (F . s) / t = F . v*. Do đó chọn phương án A."
    }
]

SHORT_ANS_DATA = [
    {
        "id": 1,
        "content": "Một chiếc máy bay có khối lượng 400 tấn đang bay cách mặt đất ở độ cao 10 000 m và có tốc độ là 720 km/h. Động năng và cơ năng của máy bay là ……… .10^8 J và ……… .10^8 J.",
        "answer": "80 .10^8 J và 480 .10^8 J (với g = 10 m/s^2) hoặc 472 .10^8 J (với g = 9,8 m/s^2)",
        "solution": (
            "• Đổi đơn vị: m = 400 tấn = 400 000 kg; h = 10 000 m; v = 720 km/h = 720 / 3,6 = 200 m/s.\n"
            "• Động năng của máy bay:\n"
            "  W_d = (1/2) . m . v^2 = (1/2) × 400 000 × 200^2 = 8 000 000 000 J = *80 .10^8 J*.\n"
            "• Thế năng của máy bay (lấy g = 10 m/s^2):\n"
            "  W_t = m . g . h = 400 000 × 10 × 10 000 = 40 000 000 000 J = 400 .10^8 J.\n"
            "• Cơ năng của máy bay: W = W_d + W_t = 80 .10^8 + 400 .10^8 = *480 .10^8 J*.\n"
            "  (Nếu lấy g = 9,8 m/s^2 thì W_t = 392 .10^8 J ==> Cơ năng W = *472 .10^8 J*)."
        )
    },
    {
        "id": 2,
        "content": "Một chiếc xe đang chạy với tốc độ 36 km/h trong 10 s. Nếu công suất của động cơ xe là 5000 W thì công cơ học mà xe sinh ra là ……… kJ và lực kéo của động cơ xe là ……… N.",
        "answer": "50 kJ và 500 N",
        "solution": (
            "• Đổi đơn vị: v = 36 km/h = 10 m/s; t = 10 s; P = 5000 W.\n"
            "• Công cơ học mà xe sinh ra trong 10 s là:\n"
            "  A = P . t = 5000 × 10 = 50 000 J = *50 kJ*.\n"
            "• Lực kéo của động cơ xe là:\n"
            "  P = F . v  ==>  F = P / v = 5000 / 10 = *500 N*.\n"
            "  (Hoặc: quãng đường s = v . t = 10 × 10 = 100 m ==> F = A / s = 50 000 / 100 = 500 N)."
        )
    },
    {
        "id": 3,
        "content": "Một vật được ném thẳng đứng lên từ mặt đất, có tốc độ ném ban đầu là 14 m/s. Độ cao lớn nhất mà vật đạt được là ……… (m); ở vị trí ……… (m) thì động năng bằng thế năng. (Lấy g = 9,8 m/s^2).",
        "answer": "10 m và 5 m",
        "solution": (
            "• Chọn mốc thế năng tại mặt đất. Bỏ qua lực cản không khí, cơ năng bảo toàn.\n"
            "• Cơ năng ban đầu: W = W_d0 = (1/2) . m . v0^2.\n"
            "• Tại độ cao lớn nhất h_max thì v = 0 ==> W = m . g . h_max:\n"
            "  m . g . h_max = (1/2) . m . v0^2  ==>  h_max = v0^2 / (2g) = 14^2 / (2 × 9,8) = 196 / 19,6 = *10 (m)*.\n"
            "• Ở vị trí động năng bằng thế năng (W_d = W_t):\n"
            "  Cơ năng W = W_t + W_d = 2.W_t  ==>  m . g . h_max = 2 . m . g . h  ==>  h = h_max / 2 = 10 / 2 = *5 (m)*."
        )
    },
    {
        "id": 4,
        "content": "Một chiếc xe đang chạy với tốc độ 10 m/s trong 1/6 phút. Nếu công suất của động cơ xe là 1000 W thì công cơ học mà xe sinh ra là ……… kJ và lực kéo của động cơ xe là ……… N.",
        "answer": "10 kJ và 100 N",
        "solution": (
            "• Đổi thời gian: t = 1/6 phút = (1/6) × 60 s = 10 s.\n"
            "• Công cơ học xe sinh ra:\n"
            "  A = P . t = 1000 × 10 = 10 000 J = *10 kJ*.\n"
            "• Lực kéo của động cơ xe:\n"
            "  F = P / v = 1000 / 10 = *100 N*."
        )
    },
    {
        "id": 5,
        "content": "Một vật 10 tấn đang di chuyển theo phương ngang 10 m trong 1 s thì công suất của trọng lực có giá trị là ……… .10^5 W, công của trọng lực sinh ra trong thời gian đó là ……… .10^5 J.",
        "answer": "0 .10^5 W và 0 .10^5 J",
        "solution": (
            "• Trọng lực P luôn có phương thẳng đứng, hướng xuống dưới.\n"
            "• Vật di chuyển theo phương nằm ngang nên phương của lực vuông góc với phương dịch chuyển (góc α = 90°).\n"
            "• Công của trọng lực là: A = P . s . cos(90°) = *0 J* = *0 .10^5 J*.\n"
            "• Công suất của trọng lực: P = A / t = 0 / 1 = *0 W* = *0 .10^5 W*."
        )
    }
]

ESSAY_DATA = [
    {
        "id": 1,
        "title": "Câu 1",
        "content": (
            "Một thang cuốn ở một trung tâm thương mại có thể nâng cùng lúc 12 người lên độ cao 4 m trong thời gian 12 s. "
            "Xem lực nâng của thang cuốn bằng với trọng lượng của vật cần nâng và khối lượng trung bình của mỗi người là 65 kg. "
            "Tính công suất của thang cuốn. (Lấy g = 10 m/s^2)."
        ),
        "solution": (
            "• Tổng khối lượng của 12 người là:\n"
            "  m = 12 × 65 = 780 (kg).\n"
            "• Lực nâng của thang cuốn bằng trọng lượng của 12 người:\n"
            "  F = P = m . g = 780 × 10 = 7800 (N).\n"
            "• Công mà thang cuốn thực hiện để nâng 12 người lên độ cao h = 4 m là:\n"
            "  A = F . h = 7800 × 4 = 31 200 (J).\n"
            "• Công suất của thang cuốn là:\n"
            "  P = A / t = 31 200 / 12 = *2600 (W)* = *2,6 (kW)*.\n"
            "  (Nếu lấy g = 9,8 m/s^2: F = 7644 N ==> A = 30 576 J ==> P = *2548 W* = 2,548 kW).\n"
            "• Đáp số: *2600 W* (hoặc 2548 W)."
        )
    },
    {
        "id": 2,
        "title": "Câu 2",
        "content": (
            "Thả rơi một vật từ độ cao 10 m xuống mặt đất. Biết khối lượng của vật là 0,5 kg. (Lấy g = 10 m/s^2).\n"
            "a) Tính tốc độ của vật khi vật chạm đất.\n"
            "b) Tìm vị trí mà tại đó thế năng gấp 3 lần động năng."
        ),
        "solution": (
            "Chọn mốc thế năng tại mặt đất. Bỏ qua sức cản của không khí nên cơ năng của vật được bảo toàn.\n"
            "• Cơ năng ban đầu của vật tại độ cao h0 = 10 m (khi thả rơi, v0 = 0 m/s):\n"
            "  W = W_t0 + W_d0 = m . g . h0 + 0 = 0,5 × 10 × 10 = 50 (J).\n"
            "a) Khi vật chạm đất (h = 0 m ==> W_t = 0):\n"
            "  Cơ năng: W = W_d = (1/2) . m . v^2  ==>  50 = (1/2) × 0,5 × v^2 = 0,25.v^2\n"
            "  ==> v^2 = 50 / 0,25 = 200  ==>  v = căn(200) = *10.căn(2) m/s* ≈ *14,14 (m/s)*.\n"
            "  (Nếu lấy g = 9,8 m/s^2: v = căn(2.g.h0) = căn(2 × 9,8 × 10) = căn(196) = *14 (m/s)*).\n"
            "b) Vị trí thế năng gấp 3 lần động năng (W_t = 3.W_d ==> W_d = (1/3).W_t):\n"
            "  Cơ năng: W = W_t + W_d = W_t + (1/3).W_t = (4/3).W_t\n"
            "  Mà W = m . g . h0 và W_t = m . g . h1 nên:\n"
            "  m . g . h0 = (4/3) . m . g . h1  ==>  h1 = (3/4) . h0 = (3/4) × 10 = *7,5 (m)*.\n"
            "• Đáp số: a) v ≈ 14,14 m/s (hoặc 14 m/s); b) h1 = 7,5 m."
        )
    },
    {
        "id": 3,
        "title": "Câu 3",
        "content": "Một vật thả rơi từ độ cao h xuống mặt đất. Chứng minh tốc độ rơi của vật khi chạm đất không phụ thuộc vào khối lượng của vật.",
        "solution": (
            "Chứng minh:\n"
            "• Chọn mốc thế năng tại mặt đất. Bỏ qua sức cản của không khí, cơ năng của vật bảo toàn trong suốt quá trình rơi.\n"
            "• Cơ năng của vật ở độ cao ban đầu h (thả rơi không vận tốc đầu, v0 = 0):\n"
            "  W_1 = W_t1 + W_d1 = m . g . h + 0 = m . g . h.\n"
            "• Cơ năng của vật ngay khi chạm đất (độ cao h = 0):\n"
            "  W_2 = W_t2 + W_d2 = 0 + (1/2) . m . v^2 = (1/2) . m . v^2.\n"
            "• Theo định luật bảo toàn cơ năng (W_1 = W_2):\n"
            "  m . g . h = (1/2) . m . v^2.\n"
            "• Chia cả hai vế cho khối lượng m (m > 0):\n"
            "  g . h = (1/2) . v^2  ==>  v^2 = 2 . g . h  ==>  *v = căn(2.g.h)*.\n"
            "• Nhận xét: Trong công thức v = căn(2gh), tốc độ chạm đất chỉ phụ thuộc vào gia tốc trọng trường g và độ cao ban đầu h, hoàn toàn không phụ thuộc vào khối lượng m của vật (đpcm)."
        )
    },
    {
        "id": 4,
        "title": "Câu 4",
        "content": (
            "Một vật có khối lượng m = 500 g bắt đầu trượt từ đỉnh của một dốc dài 10 m, góc nghiêng giữa mặt dốc và mặt sàn nằm ngang là 30°. "
            "Tính thế năng của vật tại đỉnh dốc này. (Chọn mốc thế năng tại mặt sàn nằm ngang, lấy g = 10 m/s^2)."
        ),
        "solution": (
            "• Đổi đơn vị: m = 500 g = 0,5 kg; chiều dài dốc l = 10 m; góc nghiêng α = 30°.\n"
            "• Độ cao của đỉnh dốc so với mặt sàn nằm ngang là:\n"
            "  h = l . sin(α) = 10 × sin(30°) = 10 × 0,5 = *5 (m)*.\n"
            "• Thế năng trọng trường của vật tại đỉnh dốc là:\n"
            "  W_t = m . g . h = 0,5 × 10 × 5 = *25 (J)*.\n"
            "  (Nếu lấy g = 9,8 m/s^2 thì W_t = 0,5 × 9,8 × 5 = *24,5 J*).\n"
            "• Đáp số: *25 J*."
        )
    },
    {
        "id": 5,
        "title": "Câu 5",
        "content": (
            "Ném một vật nặng 0,5 kg lên cao từ mặt đất với tốc độ ban đầu là 20 m/s. (Lấy g = 10 m/s^2).\n"
            "a) Tính độ cao lớn nhất mà vật có thể đạt được.\n"
            "b) Tìm vị trí mà tại đó thế năng của vật bằng động năng."
        ),
        "solution": (
            "Chọn mốc thế năng tại mặt đất. Bỏ qua sức cản không khí, cơ năng bảo toàn.\n"
            "• Cơ năng ban đầu của vật tại mặt đất:\n"
            "  W = (1/2) . m . v0^2 = (1/2) × 0,5 × 20^2 = 0,25 × 400 = 100 (J).\n"
            "a) Khi vật lên tới độ cao lớn nhất h_max thì vận tốc v = 0 (W_d = 0):\n"
            "  W = W_t = m . g . h_max  ==>  100 = 0,5 × 10 × h_max = 5.h_max\n"
            "  ==>  h_max = 100 / 5 = *20 (m)*.\n"
            "b) Tại vị trí thế năng bằng động năng (W_t = W_d):\n"
            "  Cơ năng: W = W_t + W_d = 2.W_t  ==>  100 = 2 . (m . g . h) = 2 × 0,5 × 10 × h = 10.h\n"
            "  ==>  h = 100 / 10 = *10 (m)*.\n"
            "  (Hoặc h = h_max / 2 = 20 / 2 = 10 m).\n"
            "• Đáp số: a) h_max = 20 m; b) h = 10 m."
        )
    },
    {
        "id": 6,
        "title": "Câu 6",
        "content": (
            "Thả rơi một vật từ độ cao 10 m xuống mặt đất. Biết khối lượng của vật là 0,5 kg. (Lấy g = 10 m/s^2).\n"
            "a) Tính tốc độ của vật khi vật chạm đất.\n"
            "b) Tìm vị trí mà tại đó thế năng gấp 3 lần động năng."
        ),
        "solution": (
            "• Lời giải tương tự Câu 2:\n"
            "a) Tốc độ khi chạm đất: v = căn(2.g.h0) = căn(2 × 10 × 10) = *10.căn(2) m/s* ≈ *14,14 (m/s)*.\n"
            "b) Vị trí thế năng gấp 3 lần động năng: h = (3/4) . h0 = (3/4) × 10 = *7,5 (m)*.\n"
            "• Đáp số: a) v ≈ 14,14 m/s; b) h = 7,5 m."
        )
    }
]

# ==============================================================================
# HÀM BỔ TRỢ ĐỊNH DẠNG WORD (.docx) CHUẨN GEMINI.md
# ==============================================================================

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders_horizontal_only(table, border_color_hex="8BC390", border_sz="6"):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color_hex}"/>\n'
        f'  <w:bottom w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color_hex}"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:insideH w:val="none"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.59)    # 1.5 cm
        s.bottom_margin = Inches(0.59) # 1.5 cm
        s.left_margin = Inches(0.59)   # 1.5 cm
        s.right_margin = Inches(0.59)  # 1.5 cm
        s.header_distance = Inches(0.24) # 0.6 cm
        s.footer_distance = Inches(0.24) # 0.6 cm

        # Header Table
        header = s.header
        header.is_linked_to_previous = False
        tbl_hdr = header.add_table(rows=1, cols=2, width=Inches(7.08))
        tbl_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders_horizontal_only(tbl_hdr, border_color_hex=COLOR_SECONDARY_HEX, border_sz="6")

        c0 = tbl_hdr.cell(0, 0)
        c0.width = Inches(3.54)
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r0 = p0.add_run("Tài liệu lưu hành nội bộ")
        r0.font.name = FONT_NAME
        r0.font.size = Pt(10)
        r0.font.italic = True
        r0.font.color.rgb = COLOR_PRIMARY

        c1 = tbl_hdr.cell(0, 1)
        c1.width = Inches(3.54)
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r1 = p1.add_run("Thầy ")
        r1.font.name = FONT_NAME
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_PRIMARY
        r1_b = p1.add_run("TRẦN VĂN THẠNH")
        r1_b.font.name = FONT_NAME
        r1_b.font.size = Pt(10)
        r1_b.font.bold = True
        r1_b.font.color.rgb = COLOR_PRIMARY
        r1_ph = p1.add_run(" - 0777.470.803")
        r1_ph.font.name = FONT_NAME
        r1_ph.font.size = Pt(10)
        r1_ph.font.color.rgb = COLOR_PRIMARY

        # Footer Table
        footer = s.footer
        footer.is_linked_to_previous = False
        tbl_ftr = footer.add_table(rows=1, cols=2, width=Inches(7.08))
        tbl_ftr.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders_horizontal_only(tbl_ftr, border_color_hex=COLOR_SECONDARY_HEX, border_sz="6")

        fc0 = tbl_ftr.cell(0, 0)
        fc0.width = Inches(5.3)
        fp0 = fc0.paragraphs[0]
        fp0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fr0 = fp0.add_run("CS1: 6/15 Nguyễn Hoàng, Kim Long   |   CS2: 24 Đặng Thái Thân, TP Huế")
        fr0.font.name = FONT_NAME
        fr0.font.size = Pt(9.5)
        fr0.font.italic = True
        fr0.font.color.rgb = COLOR_PRIMARY

        fc1 = tbl_ftr.cell(0, 1)
        fc1.width = Inches(1.78)
        fp1 = fc1.paragraphs[0]
        fp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fr1 = fp1.add_run("Trang ")
        fr1.font.name = FONT_NAME
        fr1.font.size = Pt(10)
        fr1.font.bold = True
        fr1.font.color.rgb = COLOR_PRIMARY

        fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        fp1._p.append(fldSimple)

def add_header_block(doc, is_teacher=False, is_screen=False):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c0 = tbl.cell(0, 0)
    c0.width = Inches(3.2)
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.line_spacing = 1.15
    p0.paragraph_format.space_after = Pt(2)

    r_ubnd = p0.add_run("TRƯỜNG THCS NGUYỄN HOÀNG\n")
    r_ubnd.font.name = FONT_NAME
    r_ubnd.font.size = Pt(10.5)
    r_ubnd.font.bold = True
    r_ubnd.font.color.rgb = COLOR_TEXT

    r_nh = p0.add_run("NĂM HỌC 2026 - 2027\n")
    r_nh.font.name = FONT_NAME
    r_nh.font.size = Pt(10)
    r_nh.font.italic = True

    r_de = p0.add_run("ĐỀ CƯƠNG CHÍNH THỨC")
    r_de.font.name = FONT_NAME
    r_de.font.size = Pt(10)
    r_de.font.bold = True
    r_de.font.color.rgb = COLOR_PRIMARY

    c1 = tbl.cell(0, 1)
    c1.width = Inches(3.88)
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.line_spacing = 1.15
    p1.paragraph_format.space_after = Pt(2)

    r_title = p1.add_run("ĐỀ CƯƠNG ÔN TẬP THI GIỮA KÌ I\n")
    r_title.font.name = FONT_NAME
    r_title.font.size = Pt(11)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY

    r_sub = p1.add_run("MÔN: KHOA HỌC TỰ NHIÊN 9 (KHTN 2)\n")
    r_sub.font.name = FONT_NAME
    r_sub.font.size = Pt(11)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_PRIMARY

    r_note = p1.add_run("PHÂN MÔN VẬT LÍ (ĐỘNG NĂNG - THẾ NĂNG - CƠ NĂNG)")
    r_note.font.name = FONT_NAME
    r_note.font.size = Pt(10)
    r_note.font.bold = True
    r_note.font.color.rgb = COLOR_SOLUTION_TITLE

    # Teacher / Screen notice
    if is_teacher:
        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_t.paragraph_format.space_before = Pt(4)
        p_t.paragraph_format.space_after = Pt(6)
        r_tn = p_t.add_run("(BẢN GIÁO VIÊN -- CÓ LỜI GIẢI CHI TIẾT VÀ ĐÁP ÁN ĐẦY ĐỦ)")
        r_tn.font.name = FONT_NAME
        r_tn.font.size = Pt(10.5)
        r_tn.font.bold = True
        r_tn.font.color.rgb = COLOR_PRIMARY
    elif is_screen:
        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_t.paragraph_format.space_before = Pt(4)
        p_t.paragraph_format.space_after = Pt(6)
        r_tn = p_t.add_run("(BẢN DẠY TRÊN MÀN HÌNH TIVI 65 INCH -- KHÔNG CÓ ĐÁP ÁN)")
        r_tn.font.name = FONT_NAME
        r_tn.font.size = Pt(10.5)
        r_tn.font.bold = True
        r_tn.font.color.rgb = COLOR_PRIMARY
    else:
        p_stu = doc.add_paragraph()
        p_stu.paragraph_format.space_before = Pt(4)
        p_stu.paragraph_format.space_after = Pt(6)
        r_s = p_stu.add_run("Họ và tên học sinh: ........................................................................ Lớp: .............. Số báo danh: .......................................")
        r_s.font.name = FONT_NAME
        r_s.font.size = Pt(10.5)
        r_s.font.italic = True

def add_part_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = COLOR_PRIMARY

def add_subpart_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = COLOR_SOLUTION_TITLE

def add_solution_paragraph(doc, text):
    lines = text.strip().split("\n")
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.space_before = Pt(2) if i == 0 else Pt(1)
        p.paragraph_format.space_after = Pt(2) if i == len(lines)-1 else Pt(1)
        p.paragraph_format.line_spacing = 1.15
        
        if i == 0:
            r_tag = p.add_run(">> Hướng dẫn giải chi tiết:\n")
            r_tag.font.name = FONT_NAME
            r_tag.font.size = Pt(10.5)
            r_tag.font.bold = True
            r_tag.font.color.rgb = COLOR_SOLUTION_TITLE
        
        parts = line.split("*")
        for idx, part in enumerate(parts):
            if not part:
                continue
            r = p.add_run(part)
            r.font.name = FONT_NAME
            r.font.size = Pt(10.5)
            if idx % 2 == 1:
                r.font.bold = True
                r.font.color.rgb = COLOR_PRIMARY
            else:
                r.font.color.rgb = COLOR_TEXT

def add_dotted_lines(doc, num_lines=4):
    for _ in range(num_lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("......................................................................................................................................................................................")
        r.font.name = FONT_NAME
        r.font.size = Pt(9.5)
        r.font.color.rgb = COLOR_GRAY

def add_options_two_by_two(doc, options, correct_key=None, is_teacher=False):
    # Dòng 1: A và B (Tab stop ở 3.5 inch)
    p1 = doc.add_paragraph()
    p1.paragraph_format.left_indent = Inches(0.2)
    p1.paragraph_format.tab_stops.add_tab_stop(Inches(3.5), WD_TAB_ALIGNMENT.LEFT)
    p1.paragraph_format.space_before = Pt(1)
    p1.paragraph_format.space_after = Pt(1)
    p1.paragraph_format.line_spacing = 1.15

    r_a = p1.add_run(options[0])
    r_a.font.name = FONT_NAME
    r_a.font.size = Pt(10.5)
    if is_teacher and correct_key == "A":
        r_a.font.bold = True
        r_a.font.color.rgb = COLOR_PRIMARY

    p1.add_run("\t")

    r_b = p1.add_run(options[1])
    r_b.font.name = FONT_NAME
    r_b.font.size = Pt(10.5)
    if is_teacher and correct_key == "B":
        r_b.font.bold = True
        r_b.font.color.rgb = COLOR_PRIMARY

    # Dòng 2: C và D
    p2 = doc.add_paragraph()
    p2.paragraph_format.left_indent = Inches(0.2)
    p2.paragraph_format.tab_stops.add_tab_stop(Inches(3.5), WD_TAB_ALIGNMENT.LEFT)
    p2.paragraph_format.space_before = Pt(1)
    p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.line_spacing = 1.15

    r_c = p2.add_run(options[2])
    r_c.font.name = FONT_NAME
    r_c.font.size = Pt(10.5)
    if is_teacher and correct_key == "C":
        r_c.font.bold = True
        r_c.font.color.rgb = COLOR_PRIMARY

    p2.add_run("\t")

    r_d = p2.add_run(options[3])
    r_d.font.name = FONT_NAME
    r_d.font.size = Pt(10.5)
    if is_teacher and correct_key == "D":
        r_d.font.bold = True
        r_d.font.color.rgb = COLOR_PRIMARY

# ==============================================================================
# HÀM TẠO FILE WORD (HỌC SINH HOẶC GIÁO VIÊN HOẶC MÀN HÌNH)
# ==============================================================================
def create_word_document(mode="student"):
    doc = docx.Document()
    add_header_footer(doc)
    is_teacher = (mode == "teacher")
    is_screen = (mode == "screen")
    add_header_block(doc, is_teacher=is_teacher, is_screen=is_screen)

    # --------------------------------------------------------------------------
    # PHẦN I. CÂU TRẮC NGHIỆM PHƯƠNG ÁN NHIỀU LỰA CHỌN
    # --------------------------------------------------------------------------
    add_part_heading(doc, "PHẦN I. CÂU TRẮC NGHIỆM PHƯƠNG ÁN NHIỀU LỰA CHỌN")
    
    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(1)
    p_note.paragraph_format.space_after = Pt(4)
    r_n = p_note.add_run("(Thí sinh lựa chọn một phương án trả lời đúng nhất trong mỗi câu sau)")
    r_n.font.name = FONT_NAME
    r_n.font.size = Pt(10)
    r_n.font.italic = True

    for item in MCQ_DATA:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(1)
        p_q.paragraph_format.line_spacing = 1.15
        
        r_num = p_q.add_run(f"Câu {item['id']}. ")
        r_num.font.name = FONT_NAME
        r_num.font.size = Pt(10.5)
        r_num.font.bold = True
        
        r_cnt = p_q.add_run(item["question"])
        r_cnt.font.name = FONT_NAME
        r_cnt.font.size = Pt(10.5)

        add_options_two_by_two(doc, item["options"], correct_key=item["answer"], is_teacher=is_teacher)

        if is_teacher:
            add_solution_paragraph(doc, item["solution"])

    # --------------------------------------------------------------------------
    # PHẦN II. TRẮC NGHIỆM TRẢ LỜI NGẮN
    # --------------------------------------------------------------------------
    add_part_heading(doc, "PHẦN II. TRẮC NGHIỆM TRẢ LỜI NGẮN")
    
    p_sn_note = doc.add_paragraph()
    p_sn_note.paragraph_format.space_before = Pt(1)
    p_sn_note.paragraph_format.space_after = Pt(4)
    r_snn = p_sn_note.add_run("(Học sinh tính toán và điền đáp số thích hợp vào chỗ trống)")
    r_snn.font.name = FONT_NAME
    r_snn.font.size = Pt(10)
    r_snn.font.italic = True

    for item in SHORT_ANS_DATA:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        
        r_num = p_q.add_run(f"Câu {item['id']}. ")
        r_num.font.name = FONT_NAME
        r_num.font.size = Pt(10.5)
        r_num.font.bold = True
        
        r_cnt = p_q.add_run(item["content"])
        r_cnt.font.name = FONT_NAME
        r_cnt.font.size = Pt(10.5)

        if is_teacher:
            add_solution_paragraph(doc, item["solution"])
        elif not is_screen:
            p_ans_stu = doc.add_paragraph()
            p_ans_stu.paragraph_format.left_indent = Inches(0.2)
            p_ans_stu.paragraph_format.space_before = Pt(1)
            p_ans_stu.paragraph_format.space_after = Pt(2)
            r_as = p_ans_stu.add_run("Trả lời: ............................................................................................................................................")
            r_as.font.name = FONT_NAME
            r_as.font.size = Pt(9.5)
            r_as.font.color.rgb = COLOR_GRAY

    # --------------------------------------------------------------------------
    # PHẦN III. TỰ LUẬN
    # --------------------------------------------------------------------------
    add_part_heading(doc, "PHẦN III. PHẦN TỰ LUẬN")
    for item in ESSAY_DATA:
        p_eq = doc.add_paragraph()
        p_eq.paragraph_format.space_before = Pt(4)
        p_eq.paragraph_format.space_after = Pt(2)
        p_eq.paragraph_format.line_spacing = 1.15
        
        r_num = p_eq.add_run(f"{item['title']}. ")
        r_num.font.name = FONT_NAME
        r_num.font.size = Pt(10.5)
        r_num.font.bold = True
        
        r_cnt = p_eq.add_run(item["content"])
        r_cnt.font.name = FONT_NAME
        r_cnt.font.size = Pt(10.5)

        if is_teacher:
            add_solution_paragraph(doc, item["solution"])
        elif not is_screen:
            add_dotted_lines(doc, num_lines=6)

    # Dòng kết thúc
    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_end.paragraph_format.space_before = Pt(10)
    r_end = p_end.add_run("“HẾT -- Chúc các em ôn tập và làm bài đạt kết quả xuất sắc!”")
    r_end.font.name = FONT_NAME
    r_end.font.size = Pt(10)
    r_end.font.italic = True
    r_end.font.color.rgb = COLOR_PRIMARY

    # Đặt tên file xuất
    if mode == "student":
        filename = "De_Cuong_GK1_KHTN9_NguyenHoang_VatLi_HocSinh.docx"
    elif mode == "teacher":
        filename = "De_Cuong_GK1_KHTN9_NguyenHoang_VatLi_GiaoVien.docx"
    else:
        filename = "De_Cuong_GK1_KHTN9_NguyenHoang_VatLi_DayTrenManHinh.docx"

    out_path = os.path.join(SAN_PHAM_DIR, filename)
    doc.save(out_path)
    print(f"[OK] Đã tạo thành công file Word: {out_path} ({os.path.getsize(out_path):,} bytes)")
    return out_path

# ==============================================================================
# HÀM TẠO FILE SLIDE PDF DẠY TRÊN MÀN HÌNH 16:9 CHO TIVI 65 INCH (KHÔNG CÓ ĐÁP ÁN)
# ==============================================================================
def create_screen_teaching_pdf():
    tex_path = os.path.join(SCRATCH_DIR, "De_Cuong_GK1_KHTN9_NguyenHoang_VatLi_DayTrenManHinh.tex")
    
    tex_lines = [r"""\documentclass[12pt,landscape]{article}
\usepackage[paperwidth=29.7cm,paperheight=16.706cm,top=1.0cm,bottom=1.0cm,left=1.2cm,right=1.2cm,headheight=0pt,headsep=0pt,footskip=0pt]{geometry}
\usepackage[utf8]{vietnam}
\usepackage{times}
\usepackage{amsmath,amssymb}
\usepackage{graphicx}
\usepackage{xcolor}
\usepackage{tabularx}
\usepackage{enumitem}

\definecolor{brandRed}{RGB}{211, 47, 47}
\definecolor{brandBlue}{RGB}{21, 101, 192}
\definecolor{brandGreen}{RGB}{46, 125, 50}
\definecolor{offBlack}{RGB}{26, 26, 26}
\definecolor{grayText}{RGB}{117, 117, 117}

\pagestyle{empty}
\setlength{\parindent}{0pt}
\setlength{\parskip}{0pt}

\begin{document}
"""]

    # SLIDE 1: TIÊU ĐỀ
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 1: TIÊU ĐỀ
% ==============================================================================
\begin{center}
\vspace*{0.8cm}
{\fontsize{24}{30}\selectfont \textbf{\color{offBlack}TRƯỜNG THCS NGUYỄN HOÀNG -- NĂM HỌC 2026 -- 2027}}\\[12pt]
{\fontsize{34}{44}\selectfont \textbf{\color{brandRed}ĐỀ CƯƠNG ÔN TẬP THI GIỮA KÌ I}}\\[10pt]
{\fontsize{28}{36}\selectfont \textbf{\color{brandBlue}MÔN: KHOA HỌC TỰ NHIÊN 9 (KHTN 2)}}\\[8pt]
{\fontsize{22}{28}\selectfont \textit{\color{grayText}(Phân môn Vật lí: Động năng -- Thế năng -- Cơ năng -- Công và Công suất)}}\\[1.2cm]

{\fontsize{20}{26}\selectfont \textbf{Giáo viên hướng dẫn:} Thầy \textbf{\color{brandRed}TRẦN VĂN THẠNH} -- 0777.470.803}\\[6pt]
{\fontsize{18}{24}\selectfont \textit{Cơ sở 1: 6/15 Nguyễn Hoàng, Kim Long \quad|\quad Cơ sở 2: 24 Đặng Thái Thân, TP Huế}}\\[1.0cm]

{\fontsize{18}{24}\selectfont \textbf{\color{brandGreen}SLIDE TRÌNH CHIẾU TƯƠNG TÁC TIVI 65 INCH (KHÔNG HIỂN THỊ ĐÁP ÁN)}}
\end{center}
\newpage
""")

    # SLIDES TRẮC NGHIỆM NHIỀU LỰA CHỌN (2 CÂU / 1 SLIDE)
    # Câu 1 & 2
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 2: CÂU 1 & CÂU 2
% ==============================================================================
\fontsize{26}{36}\selectfont
\textbf{Câu 1.} Dụng cụ nào dưới đây được sử dụng để quan sát và nghiên cứu tế bào?
\vspace{6pt}

\begin{tabularx}{\textwidth}{@{}XX@{}}
\textbf{A.} Tiêu bản nhiễm sắc thể người & \textbf{B.} Lăng kính \\[8pt]
\textbf{C.} Thấu kính hội tụ & \textbf{D.} Thấu kính phân kì
\end{tabularx}

\vspace*{\fill}

\textbf{Câu 2.} Thấu kính trong thí nghiệm về sự tạo ảnh được làm từ chất liệu nào?
\vspace{6pt}

\begin{tabularx}{\textwidth}{@{}XX@{}}
\textbf{A.} Kim loại & \textbf{B.} Thủy tinh \\[8pt]
\textbf{C.} Cao su & \textbf{D.} Gỗ
\end{tabularx}
\vspace*{0.5cm}
\newpage
""")

    # Câu 3 & 4
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 3: CÂU 3 & CÂU 4
% ==============================================================================
\fontsize{25}{35}\selectfont
\textbf{Câu 3.} Khi làm báo cáo khoa học, mục tiêu nghiên cứu là gì?
\vspace{6pt}

\begin{tabularx}{\textwidth}{@{}XX@{}}
\textbf{A.} Xác định các phương pháp thực hiện & \textbf{B.} Tóm tắt kết quả nghiên cứu \\[8pt]
\textbf{C.} Xác định điều cần nghiên cứu \& thử nghiệm & \textbf{D.} Đề xuất các phương pháp mới
\end{tabularx}

\vspace*{\fill}

\textbf{Câu 4.} Đồng hồ đo điện đa năng có thể đo những đại lượng nào?
\vspace{6pt}

\begin{tabularx}{\textwidth}{@{}XX@{}}
\textbf{A.} Cường độ dòng điện & \textbf{B.} Hiệu điện thế \\[8pt]
\textbf{C.} Điện trở & \textbf{D.} Cả cường độ dòng điện, hiệu điện thế \& điện trở
\end{tabularx}
\vspace*{0.5cm}
\newpage
""")

    # Câu 5 & 6
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 4: CÂU 5 & CÂU 6
% ==============================================================================
\fontsize{25}{35}\selectfont
\textbf{Câu 5.} Bộ dụng cụ tìm hiểu tính chất ảnh qua thấu kính \textbf{không có} thiết bị nào sau đây:
\vspace{6pt}

\begin{tabularx}{\textwidth}{@{}XX@{}}
\textbf{A.} Thấu kính phân kì & \textbf{B.} Màn chắn \\[8pt]
\textbf{C.} Gương phẳng & \textbf{D.} Khe hình chữ F
\end{tabularx}

\vspace*{\fill}

\textbf{Câu 6.} Một chất điểm có khối lượng $m$ đang chuyển động với vận tốc $v$ thì động năng của nó là
\vspace{6pt}

\begin{tabularx}{\textwidth}{@{}XX@{}}
\textbf{A.} $W_d = \frac{1}{2}mv^2$ & \textbf{B.} $W_d = mv^2$ \\[8pt]
\textbf{C.} $W_d = 2mv^2$ & \textbf{D.} $W_d = \frac{1}{2}mv$
\end{tabularx}
\vspace*{0.5cm}
\newpage
""")

    # Câu 7 & 8
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 5: CÂU 7 & CÂU 8
% ==============================================================================
\fontsize{25}{35}\selectfont
\textbf{Câu 7.} Công thức tính công suất nào dưới đây là đúng?
\vspace{6pt}

\begin{tabularx}{\textwidth}{@{}XX@{}}
\textbf{A.} $A = \mathcal{P}/t$ & \textbf{B.} $A = F \cdot v$ \\[8pt]
\textbf{C.} $\mathcal{P} = F \cdot v$ & \textbf{D.} $\mathcal{P} = A \cdot t$
\end{tabularx}

\vspace*{\fill}

\textbf{Câu 8.} Nếu chọn mặt đất làm mốc để tính thế năng thì trong các vật sau đây vật nào \textbf{không có} thế năng?
\vspace{6pt}

\begin{tabularx}{\textwidth}{@{}XX@{}}
\textbf{A.} Máy bay đang bay trên cao. & \textbf{B.} Chiếc lá đang rơi. \\[8pt]
\textbf{C.} Một người đứng trên tầng ba của tòa nhà. & \textbf{D.} Quả bóng nằm yên trên sàn nhà.
\end{tabularx}
\vspace*{0.5cm}
\newpage
""")

    # Câu 9 & 10
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 6: CÂU 9 & CÂU 10
% ==============================================================================
\fontsize{25}{35}\selectfont
\textbf{Câu 9.} Nếu khối lượng của vật giảm đi 2 lần, còn tốc độ tăng lên 4 lần thì động năng của vật
\vspace{6pt}

\begin{tabularx}{\textwidth}{@{}XX@{}}
\textbf{A.} tăng lên 2 lần. & \textbf{B.} tăng lên 8 lần. \\[8pt]
\textbf{C.} giảm đi 2 lần. & \textbf{D.} giảm đi 8 lần.
\end{tabularx}

\vspace*{\fill}

\textbf{Câu 10.} Công suất có các đơn vị đo là:
\vspace{6pt}

\begin{tabularx}{\textwidth}{@{}XX@{}}
\textbf{A.} $\text{J, BTU, cal.}$ & \textbf{B.} $\text{J, cal, W.}$ \\[8pt]
\textbf{C.} $\text{W, HP, BTU.}$ & \textbf{D.} $\text{W, HP, BTU/h.}$
\end{tabularx}
\vspace*{0.5cm}
\newpage
""")

    # Câu 11 & 12
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 7: CÂU 11 & CÂU 12
% ==============================================================================
\fontsize{25}{35}\selectfont
\textbf{Câu 11.} Một chất điểm có khối lượng $m$ đang chuyển động với vận tốc $\upsilon$ thì động năng của nó là:
\vspace{6pt}

\begin{tabularx}{\textwidth}{@{}XX@{}}
\textbf{A.} $\frac{1}{2}m^2\upsilon^2$ & \textbf{B.} $\frac{1}{2}m\upsilon^2$ \\[8pt]
\textbf{C.} $2m\upsilon^2$ & \textbf{D.} $m\upsilon^2$
\end{tabularx}

\vspace*{\fill}

\textbf{Câu 12.} Nếu chọn mặt đất làm mốc để tính thế năng thì vật nào \textbf{không có} thế năng?
\vspace{6pt}

\begin{tabularx}{\textwidth}{@{}XX@{}}
\textbf{A.} Máy bay đang bay. & \textbf{B.} Xe máy đang chuyển động trên mặt đường. \\[8pt]
\textbf{C.} Chiếc lá đang rơi. & \textbf{D.} Quyển sách đặt trên bàn.
\end{tabularx}
\vspace*{0.5cm}
\newpage
""")

    # Câu 13 & 14
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 8: CÂU 13 & CÂU 14
% ==============================================================================
\fontsize{25}{35}\selectfont
\textbf{Câu 13.} Chọn câu trả lời sai: Công suất có đơn vị là?
\vspace{6pt}

\begin{tabularx}{\textwidth}{@{}XX@{}}
\textbf{A.} Mã lực & \textbf{B.} $\text{J/s}$ \\[8pt]
\textbf{C.} $\text{W}$ & \textbf{D.} $\text{kWh}$
\end{tabularx}

\vspace*{\fill}

\textbf{Câu 14.} Dụng cụ thí nghiệm nào sau đây dùng để phân tán nhiệt khi đốt?
\vspace{6pt}

\begin{tabularx}{\textwidth}{@{}XX@{}}
\textbf{A.} Bát sứ. & \textbf{B.} Phễu. \\[8pt]
\textbf{C.} Bình cầu. & \textbf{D.} Lưới tản nhiệt.
\end{tabularx}
\vspace*{0.5cm}
\newpage
""")

    # Câu 15 & 16
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 9: CÂU 15 & CÂU 16
% ==============================================================================
\fontsize{25}{35}\selectfont
\textbf{Câu 15.} Trường hợp nào vật có cả động năng và thế năng? (Mốc thế năng tại mặt đất).
\vspace{6pt}

\begin{tabularx}{\textwidth}{@{}XX@{}}
\textbf{A.} Máy bay chuyển động trên đường băng. & \textbf{B.} Ô tô đang đỗ trong bến xe. \\[8pt]
\textbf{C.} Máy bay đang bay trên cao. & \textbf{D.} Ô tô đang chuyển động trên đường.
\end{tabularx}

\vspace*{\fill}

\textbf{Câu 16.} Cơ năng của một vật đang chuyển động trên mặt đất là:
\vspace{6pt}

\begin{tabularx}{\textwidth}{@{}XX@{}}
\textbf{A.} tích của động năng và thế năng. & \textbf{B.} hiệu của động năng và thế năng. \\[8pt]
\textbf{C.} chỉ bao gồm động năng của vật. & \textbf{D.} chỉ bao gồm thế năng của vật.
\end{tabularx}
\vspace*{0.5cm}
\newpage
""")

    # Câu 17 & 18
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 10: CÂU 17 & CÂU 18
% ==============================================================================
\fontsize{25}{35}\selectfont
\textbf{Câu 17.} Lực $F$ làm vật dịch chuyển quãng đường $s$ theo hướng của lực thì công của lực $F$ là:
\vspace{6pt}

\begin{tabularx}{\textwidth}{@{}XX@{}}
\textbf{A.} $A = F/s$ & \textbf{B.} $A = F \cdot s$ \\[8pt]
\textbf{C.} $A = s/F$ & \textbf{D.} $A = F \cdot v$
\end{tabularx}

\vspace*{\fill}

\textbf{Câu 18.} Trong môn KHTN lớp 9, dụng cụ nào sau đây là dụng cụ thí nghiệm quang học?
\vspace{6pt}

\begin{tabularx}{\textwidth}{@{}XX@{}}
\textbf{A.} Điện kế. & \textbf{B.} Đèn laser. \\[8pt]
\textbf{C.} Lưới tản nhiệt. & \textbf{D.} Bát sứ.
\end{tabularx}
\vspace*{0.5cm}
\newpage
""")

    # Câu 19, 20 & 21
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 11: CÂU 19, CÂU 20 & CÂU 21
% ==============================================================================
\fontsize{22}{30}\selectfont
\textbf{Câu 19.} Trường hợp nào vật \textbf{không có} cả động năng và thế năng? (Mốc tại mặt đất).\\
\textbf{A.} Máy bay trên đường băng. \quad \textbf{B.} Ô tô đang đỗ ở bến. \quad \textbf{C.} Máy bay trên cao. \quad \textbf{D.} Ô tô đang chạy.

\vspace{10pt}
\textbf{Câu 20.} Cơ năng của một vật đang đứng yên cách mặt đất $10\text{ m}$ là:\\
\textbf{A.} tích của động năng và thế năng. \qquad\quad \textbf{B.} hiệu của động năng và thế năng.\\
\textbf{C.} chỉ bao gồm động năng của vật. \qquad \textbf{D.} chỉ bao gồm thế năng của vật.

\vspace{10pt}
\textbf{Câu 21.} Lực $F$ làm vật dịch chuyển theo hướng lực thì công suất của lực $F$ tính bởi:\\
\textbf{A.} $\mathcal{P} = A/t = F \cdot v$ \qquad \textbf{B.} $\mathcal{P} = A \cdot t$ \qquad \textbf{C.} $\mathcal{P} = F/v$ \qquad \textbf{D.} $\mathcal{P} = t/A$
\vspace*{0.5cm}
\newpage
""")

    # PHẦN II. TRẢ LỜI NGẮN (SLIDE 12 & 13)
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 12: TRẢ LỜI NGẮN (CÂU 1, 2, 3)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}PHẦN II. TRẮC NGHIỆM TRẢ LỜI NGẮN}}
\vspace{10pt}

\fontsize{21}{28}\selectfont
\textbf{Câu 1.} Một máy bay nặng $400\text{ tấn}$ đang bay ở độ cao $10\,000\text{ m}$ với tốc độ $720\text{ km/h}$.\\
Động năng và cơ năng của máy bay là: $\dots\dots\dots\dots \cdot 10^8\text{ J}$ \quad và \quad $\dots\dots\dots\dots \cdot 10^8\text{ J}$.

\vspace{12pt}
\textbf{Câu 2.} Một chiếc xe chạy với tốc độ $36\text{ km/h}$ trong $10\text{ s}$. Nếu công suất động cơ là $5000\text{ W}$ thì công xe sinh ra là $\dots\dots\dots\dots\text{ kJ}$ và lực kéo là $\dots\dots\dots\dots\text{ N}$.

\vspace{12pt}
\textbf{Câu 3.} Ném thẳng đứng một vật từ mặt đất lên với tốc độ ban đầu $14\text{ m/s}$.\\
Độ cao cực đại là $\dots\dots\dots\dots\text{ m}$; ở vị trí $\dots\dots\dots\dots\text{ m}$ thì động năng bằng thế năng.
\newpage
""")

    tex_lines.append(r"""
% ==============================================================================
% SLIDE 13: TRẢ LỜI NGẮN (CÂU 4 & CÂU 5)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}PHẦN II. TRẮC NGHIỆM TRẢ LỜI NGẮN (TIẾP THEO)}}
\vspace{12pt}

\fontsize{22}{30}\selectfont
\textbf{Câu 4.} Một chiếc xe đang chạy với tốc độ $10\text{ m/s}$ trong $1/6\text{ phút}$. Nếu công suất của động cơ xe là $1000\text{ W}$ thì công cơ học mà xe sinh ra là $\dots\dots\dots\dots\text{ kJ}$ và lực kéo của động cơ xe là $\dots\dots\dots\dots\text{ N}$.

\vspace{18pt}
\textbf{Câu 5.} Một vật $10\text{ tấn}$ đang di chuyển theo phương ngang $10\text{ m}$ trong $1\text{ s}$ thì:\\
Công suất của trọng lực là $\dots\dots\dots\dots \cdot 10^5\text{ W}$,\\
Công của trọng lực sinh ra trong thời gian đó là $\dots\dots\dots\dots \cdot 10^5\text{ J}$.
\newpage
""")

    # PHẦN III. TỰ LUẬN
    # Slide 14: Tự luận Câu 1 & 2
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 14: TỰ LUẬN (CÂU 1 & CÂU 2)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}PHẦN III. TỰ LUẬN (CÂU 1 \& CÂU 2)}}
\vspace{8pt}

\fontsize{20}{27}\selectfont
\textbf{Câu 1.} Một thang cuốn nâng cùng lúc $12\text{ người}$ lên độ cao $4\text{ m}$ trong $12\text{ s}$. Xem lực nâng bằng trọng lượng người và khối lượng trung bình mỗi người là $65\text{ kg}$. Tính công suất của thang cuốn.

\vspace{12pt}
\textbf{Câu 2.} Thả rơi một vật từ độ cao $10\text{ m}$ xuống đất. Biết khối lượng của vật là $0,5\text{ kg}$.
\begin{enumerate}[label=\alph*),itemsep=2pt]
  \item Tính tốc độ của vật khi vật chạm đất.
  \item Tìm vị trí mà tại đó thế năng gấp 3 lần động năng.
\end{enumerate}

\vspace*{1.5cm}
\begin{center}
{\fontsize{18}{24}\selectfont \textit{\color{grayText}(Khoảng trống chữa bài: $P = \frac{A}{t}$, bảo toàn cơ năng $W = W_t + W_d$)}}
\end{center}
\newpage
""")

    # Slide 15: Tự luận Câu 3 & 4
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 15: TỰ LUẬN (CÂU 3 \& CÂU 4)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}PHẦN III. TỰ LUẬN (CÂU 3 \& CÂU 4)}}
\vspace{8pt}

\fontsize{20}{27}\selectfont
\textbf{Câu 3.} Một vật thả rơi từ độ cao $h$ xuống mặt đất. Chứng minh tốc độ rơi của vật khi chạm đất không phụ thuộc vào khối lượng của vật.

\vspace{12pt}
\textbf{Câu 4.} Một vật có khối lượng $m = 500\text{ g}$ bắt đầu trượt từ đỉnh dốc dài $10\text{ m}$, góc nghiêng giữa mặt dốc và sàn ngang là $30^\circ$. Tính thế năng của vật tại đỉnh dốc (chọn mốc thế năng tại sàn nằm ngang).

\vspace*{1.5cm}
\begin{center}
{\fontsize{18}{24}\selectfont \textit{\color{grayText}(Khoảng trống chữa bài: $mgh = \frac{1}{2}mv^2 \implies v = \sqrt{2gh}$; $h = l\sin 30^\circ$)}}
\end{center}
\newpage
""")

    # Slide 16: Tự luận Câu 5
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 16: TỰ LUẬN (CÂU 5)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}PHẦN III. TỰ LUẬN (CÂU 5)}}
\vspace{10pt}

\fontsize{22}{30}\selectfont
\textbf{Câu 5.} Ném một vật nặng $0,5\text{ kg}$ lên cao từ mặt đất với tốc độ ban đầu là $20\text{ m/s}$.
\begin{enumerate}[label=\alph*),itemsep=8pt]
  \item Tính độ cao lớn nhất mà vật có thể đạt được.
  \item Tìm vị trí mà tại đó thế năng của vật bằng động năng.
\end{enumerate}

\vspace*{2.5cm}
\begin{center}
{\fontsize{19}{25}\selectfont \textit{\color{grayText}(Khoảng trống chữa bài: Bảo toàn cơ năng từ mặt đất lên độ cao cực đại)}}
\end{center}
\newpage
""")

    tex_lines.append(r"\end{document}")

    with open(tex_path, "w", encoding="utf-8") as f:
        f.write("\n".join(tex_lines))
    print(f"[OK] Đã ghi file LaTeX slide: {tex_path}")

    cmd = ["pdflatex", "-interaction=nonstopmode", "-output-directory", SCRATCH_DIR, tex_path]
    res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    pdf_temp = os.path.join(SCRATCH_DIR, "De_Cuong_GK1_KHTN9_NguyenHoang_VatLi_DayTrenManHinh.pdf")
    pdf_out = os.path.join(SAN_PHAM_DIR, "De_Cuong_GK1_KHTN9_NguyenHoang_VatLi_DayTrenManHinh.pdf")

    if os.path.exists(pdf_temp):
        shutil.copy2(pdf_temp, pdf_out)
        print(f"[OK] Đã biên dịch và copy PDF thành công: {pdf_out} ({os.path.getsize(pdf_out):,} bytes)")
    else:
        print(f"[ERR] Không tạo được PDF. Mã lỗi: {res.returncode}")
    return pdf_out

# ==============================================================================
# HÀM ĐỒNG BỘ ONEDRIVE
# ==============================================================================
def sync_to_onedrive():
    files_to_sync = [
        "De_Cuong_GK1_KHTN9_NguyenHoang_VatLi_HocSinh.docx",
        "De_Cuong_GK1_KHTN9_NguyenHoang_VatLi_GiaoVien.docx",
        "De_Cuong_GK1_KHTN9_NguyenHoang_VatLi_DayTrenManHinh.docx",
        "De_Cuong_GK1_KHTN9_NguyenHoang_VatLi_DayTrenManHinh.pdf",
    ]
    print("\n[+] Đồng bộ các file sang các thư mục OneDrive:")
    for dest_dir in [ONEDRIVE_DIR1, ONEDRIVE_DIR2]:
        if os.path.exists(dest_dir):
            for fname in files_to_sync:
                src_f = os.path.join(SAN_PHAM_DIR, fname)
                dst_f = os.path.join(dest_dir, fname)
                if os.path.exists(src_f):
                    try:
                        shutil.copy2(src_f, dst_f)
                        print(f"    -> Đã copy {fname} vào: {dest_dir}")
                    except Exception as e:
                        print(f"    [!] Cảnh báo file đang mở/khóa tại {dest_dir}: {e}")
        else:
            print(f"    [-] Không tìm thấy thư mục: {dest_dir}")

# ==============================================================================
# MAIN
# ==============================================================================
if __name__ == "__main__":
    print("=" * 75)
    print("BẮT ĐẦU TẠO 3 FILE CHO ĐỀ CƯƠNG ÔN TẬP GIỮA KỲ I KHTN 9 (THCS NGUYỄN HOÀNG)")
    print("PHÂN MÔN VẬT LÍ (KHTN 2)")
    print("1. FILE HỌC SINH (.docx)")
    print("2. FILE GIÁO VIÊN (.docx có lời giải chi tiết)")
    print("3. FILE DẠY TRÊN MÀN HÌNH (Slide PDF 16:9 widescreen cho Tivi 65 inch)")
    print("=" * 75)

    create_word_document(mode="student")
    create_word_document(mode="teacher")
    create_word_document(mode="screen")
    create_screen_teaching_pdf()
    sync_to_onedrive()

    print("\n" + "=" * 75)
    print("HOÀN THÀNH TẤT CẢ 3 FILE THEO ĐÚNG CHUẨN FORM ĐỀ KIỂM TRA!")
    print("=" * 75)
