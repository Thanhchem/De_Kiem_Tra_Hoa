# -*- coding: utf-8 -*-
"""
Script tạo file Word (.docx) chuẩn "FORM ĐỀ KIỂM TRA" cho Đề thi thử Giữa học kì I Hóa học 11 (Đề số 04 / Mã đề 104)
Thầy TRẦN VĂN THẠNH - Năm học 2026 - 2027

Tạo 3 phiên bản:
1. Bản Học sinh: San_Pham/De_Thi_Thu_Hoa_11_Giua_HK1_De04_HocSinh.docx
2. Bản Giáo viên: San_Pham/De_Thi_Thu_Hoa_11_Giua_HK1_De04_GiaoVien.docx
3. Bản Dạy trên màn hình: San_Pham/De_Thi_Thu_Hoa_11_Giua_HK1_De04_DayTrenManHinh.docx
"""

import os
import sys
import re
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.path.append(os.path.dirname(__file__))
from hoa11_de04_data import (
    EXAM_INFO, PART_1_QUESTIONS, PART_2_QUESTIONS, PART_3_QUESTIONS
)

ROOT_DIR = r"c:\Antigravity_Thanh"
OUT_DIR = os.path.join(ROOT_DIR, "San_Pham")
IMG_DIR = os.path.join(OUT_DIR, "images_hoa11_de04")
ONEDRIVE_DIR = r"C:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\11"
os.makedirs(OUT_DIR, exist_ok=True)
os.makedirs(IMG_DIR, exist_ok=True)

# Bảng màu chuẩn quy định trong FORM_DE_KIEM_TRA
COLOR_HF_RED = RGBColor(211, 47, 47)        # Coral Red (#D32F2F) cho Header & Footer
COLOR_HF_GREEN_HEX = "8BC390"               # Sage Green (#8BC390) cho viền Header & Footer
COLOR_PRIMARY = RGBColor(46, 125, 50)       # Xanh lá đậm (#2E7D32) cho Tiêu đề Phần
COLOR_CORRECT = RGBColor(198, 40, 40)       # Đỏ (#C62828) cho đáp án đúng Bản Giáo Viên
COLOR_SOLUTION = RGBColor(21, 101, 192)     # Xanh lam (#1565C0) cho tiêu đề Lời giải
COLOR_SCREEN_TITLE = RGBColor(106, 27, 154) # Tím đậm (#6A1B9A) cho tiêu đề bản Dạy Màn Hình
COLOR_GRAY = RGBColor(97, 97, 97)
COLOR_BORDER_GREY_HEX = "CCCCCC"

def add_p_border_bottom(p, color="8BC390", sz="6"):
    """Thêm đường kẻ viền dưới cho đoạn văn bản."""
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_p_border_top(p, color="8BC390", sz="6"):
    """Thêm đường kẻ viền trên cho đoạn văn bản."""
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    """Cài đặt lề đệm cho ô bảng."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_run_style(run, font_name="Times New Roman", font_size=10.5, bold=False, italic=False, subscript=False, superscript=False, color=None):
    """Cài đặt style chi tiết cho từng run text."""
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    """Chèn trường số trang tự động PAGE vào chân trang."""
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def render_rich_text(p, text, base_font="Times New Roman", base_size=10.5, base_bold=False, base_italic=False, base_color=None):
    """
    Sinh văn bản định dạng phong phú từ ký hiệu:
    - *bold*
    - _italic_
    - ~subscript~ (CO~2~, H~2~O, NaCl)
    - ^superscript^ (1s^2^, 4s^1^, H^+^)
    """
    tokens = re.split(r'(~[^~]+~|\^[^\^]+\^|\*[^\*]+\*|_[^_]+_)', text)
    for token in tokens:
        if not token:
            continue

        is_bold = base_bold
        is_italic = base_italic
        is_sub = False
        is_sup = False
        content = token

        if token.startswith('*') and token.endswith('*'):
            is_bold = True
            content = token[1:-1]
        elif token.startswith('_') and token.endswith('_'):
            is_italic = True
            content = token[1:-1]
        elif token.startswith('~') and token.endswith('~'):
            is_sub = True
            content = token[1:-1]
        elif token.startswith('^') and token.endswith('^'):
            is_sup = True
            content = token[1:-1]

        run = p.add_run(content)
        run_color = COLOR_CORRECT if (is_bold and not base_bold and not base_color) else base_color
        set_run_style(run, font_name=base_font, font_size=base_size, bold=is_bold, italic=is_italic,
                      subscript=is_sub, superscript=is_sup, color=run_color)

def setup_header_footer(doc):
    """Cài đặt Header & Footer chuẩn theo FORM_DE_KIEM_TRA."""
    normal_style = doc.styles['Normal']
    normal_style.paragraph_format.tab_stops.clear_all()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    section.start_type = WD_SECTION_START.NEW_PAGE
    section.different_first_page_header_footer = False
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    # 1. Header (Bảng 1x2, rộng 18.0 cm)
    header = section.header
    for p in header.paragraphs:
        p.text = ""
    tbl_h = header.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl_h.cell(0, 0)
    c_right = tbl_h.cell(0, 1)
    c_left.width = Cm(8.0)
    c_right.width = Cm(10.0)

    p_hl = c_left.paragraphs[0]
    p_hl.paragraph_format.space_before = Pt(0)
    p_hl.paragraph_format.space_after = Pt(2)
    p_hl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_hl = p_hl.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)

    p_hr = c_right.paragraphs[0]
    p_hr.paragraph_format.space_before = Pt(0)
    p_hr.paragraph_format.space_after = Pt(2)
    p_hr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p_hr.add_run("Thầy ")
    set_run_style(r1, font_size=10, color=COLOR_HF_RED)
    r2 = p_hr.add_run("TRẦN VĂN THẠNH")
    set_run_style(r2, font_size=10, bold=True, color=COLOR_HF_RED)
    r3 = p_hr.add_run(" - 0777.470.803")
    set_run_style(r3, font_size=10, color=COLOR_HF_RED)

    add_p_border_bottom(p_hl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_bottom(p_hr, color=COLOR_HF_GREEN_HEX, sz="6")

    # 2. Footer (Bảng 1x2, rộng 18.0 cm)
    footer = section.footer
    for p in footer.paragraphs:
        p.text = ""
    tbl_f = footer.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_fl = tbl_f.cell(0, 0)
    c_fr = tbl_f.cell(0, 1)
    c_fl.width = Cm(14.5)
    c_fr.width = Cm(3.5)

    p_fl = c_fl.paragraphs[0]
    p_fl.paragraph_format.space_before = Pt(2)
    p_fl.paragraph_format.space_after = Pt(0)
    p_fl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_fl = p_fl.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)

    p_fr = c_fr.paragraphs[0]
    p_fr.paragraph_format.space_before = Pt(2)
    p_fr.paragraph_format.space_after = Pt(0)
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_field(p_fr, font_size=10, bold=True, color=COLOR_HF_RED)

    add_p_border_top(p_fl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_top(p_fr, color=COLOR_HF_GREEN_HEX, sz="6")

def build_title_block(doc, is_teacher=False, is_screen=False, base_size=10.5):
    """Xây dựng Header Block không viền, đẹp mắt, chuẩn mực."""
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_l, cell_r = tbl.cell(0, 0), tbl.cell(0, 1)
    cell_l.width = Cm(8.0)
    cell_r.width = Cm(10.0)

    # Cột trái
    p_tl1 = cell_l.paragraphs[0]
    p_tl1.paragraph_format.space_before = Pt(0)
    p_tl1.paragraph_format.space_after = Pt(1)
    p_tl1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tl1, f"*{EXAM_INFO['title_left_1']}*", base_size=10.5, base_bold=True)

    p_tl2 = cell_l.add_paragraph()
    p_tl2.paragraph_format.space_before = Pt(0)
    p_tl2.paragraph_format.space_after = Pt(1)
    p_tl2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tl2, f"*{EXAM_INFO['title_left_2']}*", base_size=10.5, base_bold=True)

    p_tl3 = cell_l.add_paragraph()
    p_tl3.paragraph_format.space_before = Pt(0)
    p_tl3.paragraph_format.space_after = Pt(0)
    p_tl3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tl3, f"_{EXAM_INFO['sub_left']}_", base_size=9.5, base_italic=True)

    # Cột phải
    p_tr1 = cell_r.paragraphs[0]
    p_tr1.paragraph_format.space_before = Pt(0)
    p_tr1.paragraph_format.space_after = Pt(1)
    p_tr1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tr1, f"*{EXAM_INFO['title_right_1']}*", base_size=11, base_bold=True)

    p_tr2 = cell_r.add_paragraph()
    p_tr2.paragraph_format.space_before = Pt(0)
    p_tr2.paragraph_format.space_after = Pt(1)
    p_tr2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tr2, f"*{EXAM_INFO['title_right_2']}*", base_size=11.5, base_bold=True, base_color=COLOR_PRIMARY)

    p_tr3 = cell_r.add_paragraph()
    p_tr3.paragraph_format.space_before = Pt(0)
    p_tr3.paragraph_format.space_after = Pt(0)
    p_tr3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tr3, f"_{EXAM_INFO['time_str']}_", base_size=9.5, base_italic=True)

    # Dòng tiêu đề phụ cho Giáo viên / Màn hình
    if is_teacher:
        p_gv = doc.add_paragraph()
        p_gv.paragraph_format.space_before = Pt(4)
        p_gv.paragraph_format.space_after = Pt(2)
        p_gv.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rgv = p_gv.add_run("(BẢN GIÁO VIÊN -- CÓ ĐÁP ÁN & LỜI GIẢI CHI TIẾT)")
        set_run_style(rgv, font_size=11.5, bold=True, color=COLOR_CORRECT)
    elif is_screen:
        p_sc = doc.add_paragraph()
        p_sc.paragraph_format.space_before = Pt(4)
        p_sc.paragraph_format.space_after = Pt(2)
        p_sc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rsc = p_sc.add_run("(BẢN TRÌNH CHIẾU / GIẢNG DẠY MÀN HÌNH -- CÓ CHỪA CHỖ TRỐNG)")
        set_run_style(rsc, font_size=11.5, bold=True, color=COLOR_SCREEN_TITLE)

    # Dòng thông tin học sinh & mã đề
    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(4)
    p_info.paragraph_format.space_after = Pt(3)
    p_info.paragraph_format.tab_stops.clear_all()
    p_info.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
    
    render_rich_text(p_info, "*Họ và tên thí sinh:* .................................................... *Lớp:* ......... *Số báo danh:* .........", base_size=base_size)
    p_info.add_run("\t")
    r_md = p_info.add_run(f"Mã đề {EXAM_INFO['made']}")
    set_run_style(r_md, font_size=11, bold=True)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(4)
    add_p_border_bottom(p_div, color=COLOR_BORDER_GREY_HEX, sz="4")

def add_section_title(doc, title, subtitle=None, base_size=11):
    """Tiêu đề phần: PHẦN I, PHẦN II, PHẦN III."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, f"*{title}*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)
    if subtitle:
        render_rich_text(p, f" {subtitle}", base_size=base_size * 0.95, base_italic=True)

def add_solution_block(doc, solution_text, base_size=10):
    """Trình bày Lời giải chi tiết theo đúng quy chuẩn không dùng Table."""
    p_sol = doc.add_paragraph()
    p_sol.paragraph_format.left_indent = Cm(0.6)
    p_sol.paragraph_format.space_before = Pt(2)
    p_sol.paragraph_format.space_after = Pt(3)
    p_sol.paragraph_format.line_spacing = 1.15

    r_lbl = p_sol.add_run("Lời giải chi tiết:\n")
    set_run_style(r_lbl, font_size=base_size, bold=True, color=COLOR_SOLUTION)

    lines = solution_text.split('\n')
    for idx, line in enumerate(lines):
        line_str = line.strip()
        if not line_str:
            continue
        if idx > 0:
            p_sol = doc.add_paragraph()
            p_sol.paragraph_format.left_indent = Cm(0.6)
            p_sol.paragraph_format.space_before = Pt(1)
            p_sol.paragraph_format.space_after = Pt(2)
            p_sol.paragraph_format.line_spacing = 1.15
        render_rich_text(p_sol, line_str, base_size=base_size)

def add_screen_teaching_space(doc, num_lines=3):
    """Chừa khoảng trống rộng rãi cho bản Dạy Trên Màn Hình."""
    p_box = doc.add_paragraph()
    p_box.paragraph_format.left_indent = Cm(0.5)
    p_box.paragraph_format.space_before = Pt(3)
    p_box.paragraph_format.space_after = Pt(6)
    
    # Viền nét đứt màu xám nhẹ, viền trái nét liền tím
    pPr = p_box._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="dashed" w:sz="4" w:space="4" w:color="B0BEC5"/><w:bottom w:val="dashed" w:sz="4" w:space="4" w:color="B0BEC5"/><w:left w:val="single" w:sz="12" w:space="6" w:color="7E57C2"/></w:pBdr>'
    pPr.append(parse_xml(xml))
    
    r_hint = p_box.add_run("✎ Ghi chú & Lời giải trực tiếp trên màn hình:\n")
    set_run_style(r_hint, font_size=9.5, bold=True, color=COLOR_SCREEN_TITLE)
    for _ in range(num_lines):
        p_box.add_run("\n")

def build_exam_docx(output_path, mode="student"):
    """
    Hàm chính tạo file Word đề kiểm tra
    mode: 'student' | 'teacher' | 'screen'
    """
    is_teacher = (mode == "teacher")
    is_screen = (mode == "screen")
    base_size = 11.5 if is_screen else 10.5

    doc = docx.Document()
    setup_header_footer(doc)
    build_title_block(doc, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    # -------------------------------------------------------------------------
    # PHẦN I: 18 CÂU TRẮC NGHIỆM NHIỀU LỰA CHỌN (4,5 ĐIỂM)
    # -------------------------------------------------------------------------
    add_section_title(
        doc,
        "PHẦN I (4,5 ĐIỂM). Câu trắc nghiệm nhiều phương án lựa chọn.",
        "Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án.",
        base_size=base_size + 0.5
    )

    for q in PART_1_QUESTIONS:
        qid = q["id"]
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(1)
        p_q.paragraph_format.line_spacing = 1.15
        render_rich_text(p_q, f"*Câu {qid}.* {q['content']}", base_size=base_size)

        # Xử lý các phương án A, B, C, D bằng Tab Stops - KHÔNG DÙNG TABLE
        opts = q["options"]
        lens = [len(opt) for opt in opts]
        max_len = max(lens)

        if max_len <= 18:
            # 4 phương án trên 1 dòng
            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.space_before = Pt(1)
            p_opt.paragraph_format.space_after = Pt(2)
            p_opt.paragraph_format.tab_stops.clear_all()
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(4.8), WD_TAB_ALIGNMENT.LEFT)
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(13.6), WD_TAB_ALIGNMENT.LEFT)
            
            p_opt.add_run("\t")
            for idx, opt_text in enumerate(opts):
                opt_lbl = opt_text[:1]
                is_correct_opt = (is_teacher and opt_lbl == q["correct"])
                if is_correct_opt:
                    render_rich_text(p_opt, f"*{opt_text}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
                else:
                    render_rich_text(p_opt, opt_text, base_size=base_size)
                if idx < 3:
                    p_opt.add_run("\t")
        elif max_len <= 38:
            # 2 hàng, mỗi hàng 2 phương án
            for row_idx in range(2):
                p_opt = doc.add_paragraph()
                p_opt.paragraph_format.space_before = Pt(1)
                p_opt.paragraph_format.space_after = Pt(2)
                p_opt.paragraph_format.tab_stops.clear_all()
                p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
                p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)
                
                p_opt.add_run("\t")
                opt_text1 = opts[row_idx * 2]
                opt_lbl1 = opt_text1[:1]
                is_c1 = (is_teacher and opt_lbl1 == q["correct"])
                if is_c1:
                    render_rich_text(p_opt, f"*{opt_text1}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
                else:
                    render_rich_text(p_opt, opt_text1, base_size=base_size)
                
                p_opt.add_run("\t")
                opt_text2 = opts[row_idx * 2 + 1]
                opt_lbl2 = opt_text2[:1]
                is_c2 = (is_teacher and opt_lbl2 == q["correct"])
                if is_c2:
                    render_rich_text(p_opt, f"*{opt_text2}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
                else:
                    render_rich_text(p_opt, opt_text2, base_size=base_size)
        else:
            # 4 hàng riêng biệt
            for opt_text in opts:
                p_opt = doc.add_paragraph()
                p_opt.paragraph_format.left_indent = Cm(0.5)
                p_opt.paragraph_format.space_before = Pt(1)
                p_opt.paragraph_format.space_after = Pt(1)
                opt_lbl = opt_text[:1]
                is_c = (is_teacher and opt_lbl == q["correct"])
                if is_c:
                    render_rich_text(p_opt, f"*{opt_text}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
                else:
                    render_rich_text(p_opt, opt_text, base_size=base_size)

        # Giáo viên: Lời giải chi tiết
        if is_teacher:
            add_solution_block(doc, f"Đáp án đúng: *{q['correct']}*\n" + q["explanation"], base_size=base_size - 0.5)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=3)

    # -------------------------------------------------------------------------
    # PHẦN II: 4 CÂU TRẮC NGHIỆM ĐÚNG / SAI (4,0 ĐIỂM)
    # -------------------------------------------------------------------------
    add_section_title(
        doc,
        "PHẦN II (4,0 ĐIỂM). Câu trắc nghiệm đúng sai.",
        "Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.",
        base_size=base_size + 0.5
    )

    for q in PART_2_QUESTIONS:
        qid = q["id"]
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        render_rich_text(p_q, f"*Câu {qid}.* {q['content']}", base_size=base_size)

        # Chèn hình ảnh minh họa nếu có (Câu 1: acid rain)
        if "image" in q and q["image"]:
            img_path = os.path.join(IMG_DIR, q["image"])
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.paragraph_format.space_before = Pt(2)
                p_img.paragraph_format.space_after = Pt(3)
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.add_run().add_picture(img_path, width=Cm(6.5))

        for st in q["statements"]:
            sub_id = st["id"]
            st_text = st["text"]
            correct_val = "Đúng" if st["correct"] else "Sai"

            p_st = doc.add_paragraph()
            p_st.paragraph_format.left_indent = Cm(0.5)
            p_st.paragraph_format.space_before = Pt(1)
            p_st.paragraph_format.space_after = Pt(1)
            p_st.paragraph_format.line_spacing = 1.15

            if is_teacher:
                r_tag = p_st.add_run(f"{sub_id}) ")
                set_run_style(r_tag, font_size=base_size, bold=True)
                
                # Hiển thị nhãn [ĐÚNG] hoặc [SAI]
                ans_color = COLOR_CORRECT if st["correct"] else RGBColor(0, 105, 92)
                r_ans = p_st.add_run(f"[{correct_val.upper()}] ")
                set_run_style(r_ans, font_size=base_size, bold=True, color=ans_color)

                render_rich_text(p_st, st_text, base_size=base_size)
            else:
                r_tag = p_st.add_run(f"{sub_id}) ")
                set_run_style(r_tag, font_size=base_size, bold=True)
                render_rich_text(p_st, st_text, base_size=base_size)

        # Lời giải chi tiết Phần II
        if is_teacher:
            add_solution_block(doc, q["explanation"], base_size=base_size - 0.5)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=4)

    # -------------------------------------------------------------------------
    # PHẦN III: 6 CÂU TRẮC NGHIỆM TRẢ LỜI NGẮN (1,5 ĐIỂM)
    # -------------------------------------------------------------------------
    add_section_title(
        doc,
        "PHẦN III (1,5 ĐIỂM). Câu trắc nghiệm trả lời ngắn.",
        "Thí sinh trả lời từ câu 1 đến câu 6.",
        base_size=base_size + 0.5
    )

    for q in PART_3_QUESTIONS:
        qid = q["id"]
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        render_rich_text(p_q, f"*Câu {qid}.* {q['content']}", base_size=base_size)

        if is_teacher:
            add_solution_block(doc, f"Đáp án: *{q['correct']}*\n" + q["explanation"], base_size=base_size - 0.5)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=4)
        else:
            # Học sinh: có dòng điền đáp án ngắn gọn
            p_ans_box = doc.add_paragraph()
            p_ans_box.paragraph_format.left_indent = Cm(0.5)
            p_ans_box.paragraph_format.space_before = Pt(1)
            p_ans_box.paragraph_format.space_after = Pt(3)
            r_box = p_ans_box.add_run("Trả lời: .............................................................................................................................................")
            set_run_style(r_box, font_size=10, italic=True, color=COLOR_GRAY)

    # Kết thúc đề
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(8)
    p_end.paragraph_format.space_after = Pt(4)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("--- HẾT ---")
    set_run_style(r_end, font_size=11, bold=True)

    # -------------------------------------------------------------------------
    # BẢNG TỔNG HỢP ĐÁP ÁN (DÀNH RIÊNG CHO BẢN GIÁO VIÊN)
    # -------------------------------------------------------------------------
    if is_teacher:
        p_ans_head = doc.add_paragraph()
        p_ans_head.paragraph_format.space_before = Pt(14)
        p_ans_head.paragraph_format.space_after = Pt(4)
        p_ans_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_ah = p_ans_head.add_run("BẢNG TỔNG HỢP ĐÁP ÁN NHANH (MÃ ĐỀ 104)")
        set_run_style(r_ah, font_size=12, bold=True, color=COLOR_CORRECT)

        # 1. Bảng Phần I: 18 câu (2 hàng x 9 cột hoặc 1 hàng x 18 cột)
        p_t1 = doc.add_paragraph()
        p_t1.paragraph_format.space_before = Pt(4)
        p_t1.paragraph_format.space_after = Pt(2)
        render_rich_text(p_t1, "*1. Đáp án Phần I (Trắc nghiệm nhiều lựa chọn - 18 câu):*", base_size=10.5, base_bold=True, base_color=COLOR_PRIMARY)
        
        # Tạo bảng 4 hàng x 9 cột (Hàng 0,1: Câu 1-9; Hàng 2,3: Câu 10-18)
        tbl_p1 = doc.add_table(rows=4, cols=9)
        tbl_p1.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Hàng 1 & 2: Câu 1 đến 9
        for i in range(9):
            c0 = tbl_p1.cell(0, i); c1 = tbl_p1.cell(1, i)
            c0.width = Cm(2.0); c1.width = Cm(2.0)
            set_cell_margins(c0, top=40, bottom=40, left=40, right=40)
            set_cell_margins(c1, top=40, bottom=40, left=40, right=40)
            p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r0 = p0.add_run(f"Câu {i+1}")
            set_run_style(r0, font_size=9.5, bold=True)
            p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r1 = p1.add_run(PART_1_QUESTIONS[i]["correct"])
            set_run_style(r1, font_size=10.5, bold=True, color=COLOR_CORRECT)

        # Hàng 3 & 4: Câu 10 đến 18
        for i in range(9):
            c0 = tbl_p1.cell(2, i); c1 = tbl_p1.cell(3, i)
            c0.width = Cm(2.0); c1.width = Cm(2.0)
            set_cell_margins(c0, top=40, bottom=40, left=40, right=40)
            set_cell_margins(c1, top=40, bottom=40, left=40, right=40)
            p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r0 = p0.add_run(f"Câu {i+10}")
            set_run_style(r0, font_size=9.5, bold=True)
            p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r1 = p1.add_run(PART_1_QUESTIONS[i+9]["correct"])
            set_run_style(r1, font_size=10.5, bold=True, color=COLOR_CORRECT)

        # 2. Bảng Phần II: 4 câu Đúng/Sai
        p_t2 = doc.add_paragraph()
        p_t2.paragraph_format.space_before = Pt(6)
        p_t2.paragraph_format.space_after = Pt(2)
        render_rich_text(p_t2, "*2. Đáp án Phần II (Trắc nghiệm đúng/sai - 4 câu):*", base_size=10.5, base_bold=True, base_color=COLOR_PRIMARY)
        
        tbl_p2 = doc.add_table(rows=5, cols=5)
        tbl_p2.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers_p2 = ["Câu hỏi", "Lệnh hỏi a", "Lệnh hỏi b", "Lệnh hỏi c", "Lệnh hỏi d"]
        for j, h in enumerate(headers_p2):
            cell = tbl_p2.cell(0, j)
            cell.width = Cm(3.6)
            set_cell_margins(cell, top=40, bottom=40, left=40, right=40)
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(h)
            set_run_style(r, font_size=9.5, bold=True)
        
        for q_idx, q in enumerate(PART_2_QUESTIONS):
            row_idx = q_idx + 1
            c_qid = tbl_p2.cell(row_idx, 0)
            c_qid.width = Cm(3.6)
            set_cell_margins(c_qid, top=40, bottom=40, left=40, right=40)
            p = c_qid.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(f"Câu {q['id']}")
            set_run_style(r, font_size=9.5, bold=True)
            for st_idx, st in enumerate(q["statements"]):
                c_st = tbl_p2.cell(row_idx, st_idx + 1)
                c_st.width = Cm(3.6)
                set_cell_margins(c_st, top=40, bottom=40, left=40, right=40)
                p_st = c_st.paragraphs[0]; p_st.alignment = WD_ALIGN_PARAGRAPH.CENTER
                val_str = "Đ" if st["correct"] else "S"
                r_st = p_st.add_run(val_str)
                color_st = COLOR_CORRECT if st["correct"] else RGBColor(0, 105, 92)
                set_run_style(r_st, font_size=10, bold=True, color=color_st)

        # 3. Bảng Phần III: 6 câu trả lời ngắn
        p_t3 = doc.add_paragraph()
        p_t3.paragraph_format.space_before = Pt(6)
        p_t3.paragraph_format.space_after = Pt(2)
        render_rich_text(p_t3, "*3. Đáp án Phần III (Trả lời ngắn - 6 câu):*", base_size=10.5, base_bold=True, base_color=COLOR_PRIMARY)
        
        tbl_p3 = doc.add_table(rows=2, cols=6)
        tbl_p3.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i in range(6):
            c0 = tbl_p3.cell(0, i)
            c1 = tbl_p3.cell(1, i)
            c0.width = Cm(3.0)
            c1.width = Cm(3.0)
            set_cell_margins(c0, top=40, bottom=40, left=40, right=40)
            set_cell_margins(c1, top=40, bottom=40, left=40, right=40)
            
            p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r0 = p0.add_run(f"Câu {i+1}")
            set_run_style(r0, font_size=9.5, bold=True)
            
            p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r1 = p1.add_run(PART_3_QUESTIONS[i]["correct"])
            set_run_style(r1, font_size=11, bold=True, color=COLOR_CORRECT)

    # Lưu tài liệu
    doc.save(output_path)
    print(f"Created: {output_path}")

def main():
    files = [
        ("De_Thi_Thu_Hoa_11_Giua_HK1_De04_HocSinh.docx", "student"),
        ("De_Thi_Thu_Hoa_11_Giua_HK1_De04_GiaoVien.docx", "teacher"),
        ("De_Thi_Thu_Hoa_11_Giua_HK1_De04_DayTrenManHinh.docx", "screen"),
    ]
    for filename, mode in files:
        out_file = os.path.join(OUT_DIR, filename)
        build_exam_docx(out_file, mode=mode)
        
        # Sao chép đồng thời sang thư mục OneDrive của Thầy để Thầy mở dùng ngay
        if os.path.exists(ONEDRIVE_DIR):
            try:
                dest = os.path.join(ONEDRIVE_DIR, filename)
                shutil.copy2(out_file, dest)
                print(f"Copied to OneDrive successfully: {filename}")
            except Exception as e:
                print(f"Copy to OneDrive error: {e}")

if __name__ == "__main__":
    main()
