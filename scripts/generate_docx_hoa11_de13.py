# -*- coding: utf-8 -*-
"""
Script sinh tự động 2 file Word (.docx) cho ĐỀ SỐ 13 (Chương 5: DẪN XUẤT HALOGEN - ALCOHOL - PHENOL):
1. Bản Học sinh (De_Kiem_Tra_Hoa_11_De13.docx)
2. Bản Giáo viên (De_Kiem_Tra_Hoa_11_De13_GiaoVien.docx)

Tuân thủ nghiêm ngặt Quy chuẩn Form Đề Kiểm Tra (GEMINI.md):
- Không dùng bảng (Table) trong nội dung câu hỏi/đáp án/lời giải.
- Dùng Paragraph + Tab stops chuẩn cho các phương án trắc nghiệm.
- Công thức hóa học xử lý triệt để chỉ số trên/dưới.
- Header & Footer chuẩn với viền màu Sage Green (#8BC390) và chữ Coral Red (#D32F2F).
- Bản Giáo viên có ghi chú đỏ và lời giải thụt lề màu xanh #1565C0.
"""

import os
import sys

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH

# Đảm bảo import được module form_de_kiem_tra
sys.path.insert(0, os.path.dirname(__file__))
from form_de_kiem_tra import (
    setup_header_footer,
    build_title_block,
    add_section_header,
    render_question_prompt,
    render_options_no_table,
    render_statement,
    add_solution_no_table,
    add_student_writing_lines,
    render_rich_text,
    set_run_style,
    COLOR_CORRECT,
    COLOR_SOLUTION,
    COLOR_PRIMARY
)

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
IMG_DIR = os.path.join(OUT_DIR, "images_crop_hoa11_de13")

def add_p(doc, space_before=3, space_after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    return p

def add_centered_picture(doc, img_path, width_cm=8.0):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(img_path, width=Cm(width_cm))

def build_exam_document(is_teacher=False):
    doc = docx.Document()
    setup_header_footer(doc)

    build_title_block(
        doc,
        school_line="SỞ GIÁO DỤC VÀ ĐÀO TẠO THỪA THIÊN HUẾ",
        exam_title="ĐỀ KIỂM TRA ĐỊNH KỲ CHƯƠNG 5 (DẪN XUẤT HALOGEN - ALCOHOL - PHENOL)",
        subject_line="MÔN: HOÁ HỌC - LỚP 11",
        num_pages="04",
        duration="50 phút",
        exam_code="113",
        atomic_masses="H = 1; C = 12; N = 14; O = 16; F = 19; Na = 23; Mg = 24; Al = 27; P = 31; S = 32; Cl = 35,5; K = 39; Ca = 40; Cr = 52; Mn = 55; Fe = 56; Cu = 64; Zn = 65; Br = 80; Ag = 108; I = 127. (Thể tích khí đo ở đkc 25 °C, 1 bar là 24,79 L/mol)",
        is_teacher=is_teacher
    )

    # =========================================================================
    # PHẦN I
    # =========================================================================
    add_section_header(
        doc,
        title="PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án."
    )

    # Câu 1
    p = add_p(doc)
    render_question_prompt(p, "Câu 1. (SBT-KNTT) Dẫn xuất halogen nào sau đây có đồng phân hình học?")
    opts_c1 = [
        "A. CH~2~=CHCl.",
        "B. CH~2~=CH-CH~2~Br.",
        "*C. CH~3~CH=C(F)CH~3~.*",
        "D. (CH~3~)~2~C=CHI."
    ]
    render_options_no_table(doc, opts_c1, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "CH~3~-CH=C(F)-CH~3~ có carbon liên kết đôi thứ nhất liên kết với -H và -CH~3~ (khác nhau), carbon thứ hai liên kết với -F và -CH~3~ (khác nhau) nên tồn tại đồng phân hình học cis/trans."
        )

    # Câu 2
    p = add_p(doc)
    render_question_prompt(p, "Câu 2. (SBT-KNTT) Chất nào sau đây không phải là dẫn xuất halogen của hydrocarbon?")
    opts_c2 = [
        "A. CH~3~CH~2~Cl.",
        "B. CH~2~=CHBr.",
        "*C. ClCH~2~COOH.*",
        "D. CF~3~CH~2~Cl."
    ]
    render_options_no_table(doc, opts_c2, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Dẫn xuất halogen của hydrocarbon được tạo thành khi thay thế một hay nhiều nguyên tử H trong phân tử hydrocarbon bằng nguyên tử halogen (chỉ chứa C, H, halogen). ClCH~2~COOH còn chứa nhóm carboxyl (-COOH có oxygen), là dẫn xuất halogen của carboxylic acid."
        )

    # Câu 3
    p = add_p(doc)
    render_question_prompt(p, "Câu 3. Chất nào sau đây là dẫn xuất halogen của hydrocarbon?")
    opts_c3 = [
        "A. Cl-CH~2~-COOH.",
        "*B. C~6~H~5~-CH~2~-Cl.*",
        "C. CH~3~-CH~2~-Mg-Br.",
        "D. CH~3~-CO-Cl."
    ]
    render_options_no_table(doc, opts_c3, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "C~6~H~5~-CH~2~-Cl (benzyl chloride) chỉ chứa C, H và Cl, là dẫn xuất halogen của hydrocarbon (toluene)."
        )

    # Câu 4
    p = add_p(doc)
    render_question_prompt(p, "Câu 4. Theo chiều tăng khối lượng mol trong phân tử, độ tan trong nước của các alcohol")
    opts_c4 = [
        "A. tăng dần.",
        "B. không đổi.",
        "*C. giảm dần.*",
        "D. biến đổi không theo quy luật."
    ]
    render_options_no_table(doc, opts_c4, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Khi khối lượng mol tăng, phần gốc hydrocarbon kị nước có kích thước càng lớn, làm giảm khả năng tương tác với phân tử nước, do đó độ tan trong nước của alcohol giảm dần."
        )

    # Câu 5
    p = add_p(doc)
    render_question_prompt(p, "Câu 5. Khi đun nóng ethyl alcohol với H~2~SO~4~ đặc ở 140 °C thì sẽ tạo ra sản phẩm chính là")
    opts_c5 = [
        "A. C~2~H~4~.",
        "B. CH~3~CHO.",
        "*C. C~2~H~5~OC~2~H~5~.*",
        "D. CH~3~COOH."
    ]
    render_options_no_table(doc, opts_c5, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Ở 140 °C với xúc tác H~2~SO~4~ đặc, xảy ra phản ứng tách nước liên phân tử giữa 2 phân tử ethyl alcohol tạo diethyl ether: 2C~2~H~5~OH --(H~2~SO~4~ đặc, 140 °C)--> C~2~H~5~-O-C~2~H~5~ + H~2~O."
        )

    # Câu 6
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 6. Thủy phân dẫn xuất halogen nào sau đây sẽ thu được alcohol?\n"
        "(1) CH~3~CH~2~Cl.    (2) CH~3~CH=CHCl.    (3) C~6~H~5~CH~2~Cl.    (4) C~6~H~5~Cl."
    )
    opts_c6 = [
        "*A. (1), (3).* ",
        "B. (1), (2), (3).",
        "C. (1), (2), (4).",
        "D. (1), (2), (3), (4)."
    ]
    render_options_no_table(doc, opts_c6, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- (1) CH~3~CH~2~Cl + NaOH -> CH~3~CH~2~OH (alcohol) + NaCl.\n"
            "- (2) Cl đính vào C lai hoá sp^2 mang nối đôi, thuỷ phân tạo enol không bền chuyển thành aldehyde CH~3~CH~2~CHO.\n"
            "- (3) C~6~H~5~CH~2~Cl + NaOH -> C~6~H~5~CH~2~OH (benzyl alcohol) + NaCl.\n"
            "- (4) Cl đính trực tiếp vào vòng benzene rất trơ, không bị thuỷ phân ở điều kiện thông thường.\n"
            "=> Thu được alcohol là (1) và (3)."
        )

    # Câu 7
    p = add_p(doc)
    render_question_prompt(p, "Câu 7. Tên thay thế của alcohol có công thức cấu tạo CH~3~CH~2~CH(OH)CH~3~ là")
    opts_c7 = [
        "*A. butan-2-ol.*",
        "B. butan-3-ol.",
        "C. butan-1-ol.",
        "D. 2-methylpropan-2-ol."
    ]
    render_options_no_table(doc, opts_c7, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Mạch chính dài nhất gồm 4 nguyên tử carbon (butan), đánh số từ đầu gần nhóm -OH hơn: nhóm -OH ở vị trí C số 2 -> butan-2-ol."
        )

    # Câu 8
    p = add_p(doc)
    render_question_prompt(p, "Câu 8. Hợp chất (CH~3~)~3~COH có tên thay thế là")
    opts_c8 = [
        "*A. 2-methylpropan-2-ol.*",
        "B. 1,1-dimethylethanol.",
        "C. trimethylmethanol.",
        "D. butan-2-ol."
    ]
    render_options_no_table(doc, opts_c8, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Mạch chính gồm 3 carbon (propan), nguyên tử C số 2 mang một nhóm methyl (-CH~3~) và nhóm -OH -> 2-methylpropan-2-ol."
        )

    # Câu 9
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 9. (SBT-CTST) Keo dán phenol formaldehyde (PF) có độ kết dính cao, chịu nhiệt và nước, thường dùng để ép gỗ, dán gỗ coppha trong xây dựng. Nhựa PF là sản phẩm trùng ngưng của formaldehyde (HCHO) với"
    )
    opts_c9 = [
        "A. ethanol (C~2~H~5~OH).",
        "*B. phenol (C~6~H~5~OH).* ",
        "C. toluene (C~6~H~5~CH~3~).",
        "D. benzene (C~6~H~6~)."
    ]
    render_options_no_table(doc, opts_c9, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Nhựa phenol formaldehyde (PF) được điều chế bằng phản ứng trùng ngưng giữa phenol (C~6~H~5~OH) và formaldehyde (HCHO) với xúc tác acid hoặc base."
        )

    # Câu 10
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 10. (SBT-CTST) Liều lượng phù hợp của một số hợp chất phenol như 4-hexylresorcinol được dùng trong thuốc giảm ho, butyl paraben, BHA, BHT,... được ứng dụng làm chất bảo quản trong thực phẩm, mĩ phẩm. Vì đặc điểm quan trọng của các hợp chất này là"
    )
    opts_c10 = [
        "*A. có hoạt tính sinh học.*",
        "B. chi phí sản xuất thấp.",
        "C. dễ bảo quản.",
        "D. không gây độc hại nếu dùng dư liều lượng."
    ]
    render_options_no_table(doc, opts_c10, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Các hợp chất phenol có hoạt tính sinh học cao (kháng khuẩn, diệt nấm, chống oxy hoá mạnh mẽ), do đó ở liều lượng quy chuẩn chúng được ứng dụng làm chất sát khuẩn và chất bảo quản thực phẩm, dược mĩ phẩm."
        )

    # Câu 11
    p = add_p(doc)
    render_question_prompt(p, "Câu 11. Ảnh hưởng của gốc phenyl (C~6~H~5~-) đến nhóm -OH trong phân tử phenol thể hiện qua phản ứng giữa phenol với")
    opts_c11 = [
        "A. dung dịch H~2~SO~4~ đặc.",
        "B. H~2~ (xúc tác Ni, t°).",
        "*C. dung dịch NaOH.*",
        "D. dung dịch Br~2~ trong H~2~O."
    ]
    render_options_no_table(doc, opts_c11, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Do gốc phenyl hút electron làm tăng độ phân cực của liên kết O-H, nguyên tử H linh động hơn alcohol, làm phenol thể hiện tính acid yếu và phản ứng được với dung dịch kiềm NaOH (alcohol không có phản ứng này)."
        )

    # Câu 12
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 12. Phản ứng nào sau đây thể hiện ảnh hưởng của gốc C~6~H~5~- đối với nhóm -OH?\n"
        "(1) 2C~6~H~5~OH + 2Na -> 2C~6~H~5~ONa + H~2~^\n"
        "(2) C~6~H~5~OH + NaOH -> C~6~H~5~ONa + H~2~O\n"
        "(3) C~6~H~5~OH + 3Br~2~ -> 2,4,6-tribromophenol + 3HBr"
    )
    opts_c12 = [
        "*A. (2).* ",
        "B. (2), (3).",
        "C. (1), (2).",
        "D. (1), (3)."
    ]
    render_options_no_table(doc, opts_c12, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- Phản ứng (2) tác dụng với NaOH chứng minh tính acid của nhóm -OH tăng lên do ảnh hưởng của vòng thơm.\n"
            "- Phản ứng (1) thì cả alcohol và phenol đều tác dụng với Na.\n"
            "- Phản ứng (3) thế Br vào vòng thơm thể hiện ảnh hưởng của nhóm -OH lên vòng benzene."
        )

    # Câu 13
    p = add_p(doc)
    render_question_prompt(p, "Câu 13. Khả năng phản ứng thế nguyên tử chlorine bằng nhóm -OH của các chất được xếp theo chiều tăng dần từ trái sang phải là:")
    opts_c13 = [
        "A. Allyl chloride, phenyl chloride, propyl chloride.",
        "B. Allyl chloride, propyl chloride, phenyl chloride.",
        "C. Phenyl chloride, allyl chloride, propyl chloride.",
        "*D. Phenyl chloride, propyl chloride, allyl chloride.*"
    ]
    render_options_no_table(doc, opts_c13, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- Phenyl chloride (Cl gắn vào vòng thơm) rất trơ, khó thế nhất.\n"
            "- Propyl chloride (dẫn xuất alkyl no) thế chậm hơn trong dung dịch kiềm nóng.\n"
            "- Allyl chloride (Cl gắn cạnh liên kết đôi) rất linh động, dễ thế nhất ngay trong nước nóng.\n"
            "=> Thứ tự tăng dần: Phenyl chloride < propyl chloride < allyl chloride."
        )

    # Câu 14
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 14. (SBT-KNTT) Cho các alcohol sau: CH~3~OH; C~2~H~5~OH; HO-CH~2~-CH~2~-OH; HO-CH~2~-CH(OH)-CH~2~-OH; HO-CH~2~-CH~2~-CH~2~-OH. "
        "Số alcohol không hòa tan được Cu(OH)~2~ ở điều kiện thường là"
    )
    opts_c14 = [
        "A. 1.",
        "B. 2.",
        "*C. 3.*",
        "D. 4."
    ]
    render_options_no_table(doc, opts_c14, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Chỉ các polyalcohol có từ 2 nhóm -OH kề nhau mới hoà tan được Cu(OH)~2~ tạo dung dịch màu xanh lam (ethylene glycol và glycerol). "
            "Các alcohol không hoà tan Cu(OH)~2~ gồm: CH~3~OH, C~2~H~5~OH và HO-CH~2~-CH~2~-CH~2~-OH (2 nhóm -OH không kề nhau). Số chất là 3."
        )

    # Câu 15
    p = add_p(doc)
    render_question_prompt(p, "Câu 15. (SBT-KNTT) Nhận xét nào sau đây không đúng?")
    opts_c15 = [
        "A. Oxi hóa không hoàn toàn alcohol bậc I, thu được aldehyde.",
        "*B. Oxi hóa hoàn toàn alcohol bậc I, thu được aldehyde.*",
        "C. Oxi hoá alcohol bậc II, thu được ketone.",
        "D. Alcohol bậc III không bị oxi hoá bởi các tác nhân oxi hoá thông thường."
    ]
    render_options_no_table(doc, opts_c15, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Oxi hóa hoàn toàn alcohol là phản ứng đốt cháy tạo CO~2~ và H~2~O. Muốn thu được aldehyde phải thực hiện oxi hoá không hoàn toàn bằng CuO, t°."
        )

    # Câu 16
    p = add_p(doc)
    render_question_prompt(p, "Câu 16. Nhận xét nào sau đây là đúng khi nói về alcohol?")
    opts_c16 = [
        "A. Có thể sử dụng Cu(OH)~2~ để phân biệt ethylene glycol và propane-1,2-diol đựng trong hai lọ riêng biệt.",
        "B. Công thức chung của dãy đồng đẳng alcohol no, mạch hở là C~n~H~2n~O~x~ (n ≥ 1, x ≥ 1).",
        "*C. Các alcohol tan dễ dàng trong nước là nhờ có liên kết hydrogen giữa alcohol và các phân tử nước.*",
        "D. Khi tách nước một alcohol luôn thu được sản phẩm là alkene."
    ]
    render_options_no_table(doc, opts_c16, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Các alcohol đầu dãy tan tốt hoặc vô hạn trong nước nhờ khả năng tạo liên kết hydrogen liên phân tử mạnh mẽ giữa nhóm -OH và nước."
        )

    # Câu 17
    p = add_p(doc)
    render_question_prompt(p, "Câu 17. Cho hỗn hợp gồm ethylene và propylene phản ứng với nước có xúc tác acid thu được bao nhiêu alcohol?")
    opts_c17 = [
        "A. 2.",
        "B. 1.",
        "*C. 3.*",
        "D. 4."
    ]
    render_options_no_table(doc, opts_c17, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- Ethylene (CH~2~=CH~2~) + H~2~O -> 1 alcohol duy nhất: CH~3~CH~2~OH.\n"
            "- Propylene (CH~2~=CH-CH~3~) + H~2~O -> 2 alcohol: CH~3~-CH(OH)-CH~3~ (sp chính) và CH~3~CH~2~CH~2~OH (sp phụ).\n"
            "=> Tổng cộng thu được 3 alcohol."
        )

    # Câu 18
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 18. Saccharose là một loại đường phổ biến, sản xuất chủ yếu từ cây mía. Saccharose có cấu trúc phân tử như hình dưới:"
    )
    add_centered_picture(doc, os.path.join(IMG_DIR, "saccharose_structure.png"), width_cm=7.5)
    p_sac = add_p(doc, space_before=2, space_after=3)
    render_rich_text(p_sac, "Số nhóm chức alcohol trong một phân tử saccharose là bao nhiêu?")
    opts_c18 = [
        "A. 3.",
        "B. 5.",
        "*C. 8.*",
        "D. 11."
    ]
    render_options_no_table(doc, opts_c18, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Phân tử saccharose (C~12~H~22~O~11~) được tạo bởi một gốc alpha-glucose và một gốc beta-fructose qua liên kết glycosidic, chứa đúng 8 nhóm -OH alcohol (gồm 3 nhóm -OH bậc I và 5 nhóm -OH bậc II)."
        )

    # =========================================================================
    # PHẦN II
    # =========================================================================
    add_section_header(
        doc,
        title="PHẦN II. Câu trắc nghiệm đúng sai.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai."
    )

    # Câu 1
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 1. Khái niệm và tính chất của dẫn xuất halogen.\n"
        "Trong phòng thí nghiệm, học sinh quan sát thí nghiệm thuỷ phân ethyl bromide với dung dịch NaOH."
    )
    stmts_c1_p2 = [
        ("a) Dẫn xuất halogen là hợp chất hữu cơ trong đó một hoặc nhiều nguyên tử hydrogen được thay thế bởi nguyên tử halogen.", "ĐÚNG"),
        ("b) Ethyl bromide có thể tham gia phản ứng thế với NaOH để tạo ra ethanol và NaBr.", "ĐÚNG"),
        ("c) Phản ứng giữa ethyl bromide và NaOH tạo ra sản phẩm là NaCl và nước.", "SAI"),
        ("d) Từ kết quả thí nghiệm, có thể suy ra dẫn xuất halogen có khả năng tham gia phản ứng với OH^- để tạo hợp chất chứa nhóm -OH.", "ĐÚNG")
    ]
    for stmt, ans in stmts_c1_p2:
        p_st = add_p(doc, space_before=1, space_after=1)
        p_st.paragraph_format.left_indent = Cm(0.5)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- a) ĐÚNG: Định nghĩa chuẩn về dẫn xuất halogen.\n"
            "- b) ĐÚNG, c) SAI: Phản ứng thuỷ phân: C~2~H~5~Br + NaOH --t°--> C~2~H~5~OH + NaBr, không sinh ra NaCl hay nước.\n"
            "- d) ĐÚNG: Nhóm halogen -Br bị thay thế bởi nhóm -OH.\n"
            "=> a: ĐÚNG, b: ĐÚNG, c: SAI, d: ĐÚNG."
        )

    # Câu 2
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 2. Khái niệm và bậc của alcohol.\n"
        "Trong phòng thí nghiệm, học sinh được giao nhiệm vụ nhận diện bậc của các alcohol thông qua công thức cấu tạo."
    )
    stmts_c2_p2 = [
        ("a) Alcohol là hợp chất hữu cơ chứa nhóm -OH liên kết với nguyên tử carbon.", "SAI"),
        ("b) Công thức tổng quát của alcohol no, đơn chức, mạch hở là C~n~H~2n+1~OH (n ≥ 1).", "ĐÚNG"),
        ("c) Alcohol bậc hai là alcohol trong đó nhóm -OH liên kết với nguyên tử carbon bậc hai.", "ĐÚNG"),
        ("d) Để phân biệt các bậc của alcohol, có thể sử dụng phương pháp xác định vị trí nhóm -OH trong cấu trúc phân tử.", "ĐÚNG")
    ]
    for stmt, ans in stmts_c2_p2:
        p_st = add_p(doc, space_before=1, space_after=1)
        p_st.paragraph_format.left_indent = Cm(0.5)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- a) SAI: Nhóm -OH phải liên kết trực tiếp với nguyên tử carbon no (nếu đính vào carbon thơm là phenol, đính carbon nối đôi là enol không bền).\n"
            "- b) ĐÚNG: Công thức alcohol no, đơn chức mạch hở là C~n~H~2n+1~OH (hoặc C~n~H~2n+2~O, n ≥ 1).\n"
            "- c) ĐÚNG: Bậc alcohol là bậc của nguyên tử carbon liên kết trực tiếp với nhóm -OH.\n"
            "- d) ĐÚNG: Nhận diện vị trí carbon mang nhóm -OH xác định chính xác bậc của alcohol.\n"
            "=> a: SAI, b: ĐÚNG, c: ĐÚNG, d: ĐÚNG."
        )

    # Câu 3
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 3. Tính chất vật lý của alcohol.\n"
        "Trong thực tế, ethanol được sử dụng làm dung môi trong phòng thí nghiệm nhờ tính chất tan tốt trong nước."
    )
    stmts_c3_p2 = [
        ("a) Ethanol ở điều kiện thường tồn tại ở trạng thái lỏng.", "ĐÚNG"),
        ("b) Nhiệt độ sôi của alcohol giảm khi số nguyên tử carbon trong mạch tăng.", "SAI"),
        ("c) Alcohol có khả năng tan tốt trong nước nhờ liên kết hydrogen giữa nhóm -OH và nước.", "ĐÚNG"),
        ("d) Methanol có nhiệt độ sôi cao hơn ethanol do khối lượng phân tử nhỏ hơn.", "SAI")
    ]
    for stmt, ans in stmts_c3_p2:
        p_st = add_p(doc, space_before=1, space_after=1)
        p_st.paragraph_format.left_indent = Cm(0.5)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- a) ĐÚNG: Ethanol là chất lỏng không màu, mùi thơm nhẹ.\n"
            "- b) SAI: Khi số nguyên tử carbon tăng, phân tử khối tăng dẫn đến nhiệt độ sôi tăng dần.\n"
            "- c) ĐÚNG: Liên kết hydrogen giữa alcohol và nước giúp các alcohol nhẹ tan vô hạn trong nước.\n"
            "- d) SAI: Ethanol có phân tử khối lớn hơn và tương tác van der Waals lớn hơn nên sôi ở 78,3 °C, cao hơn methanol (64,7 °C).\n"
            "=> a: ĐÚNG, b: SAI, c: ĐÚNG, d: SAI."
        )

    # Câu 4
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 4. Thí nghiệm phenol tác dụng với nước bromine.\n"
        "Thực hiện thí nghiệm cho phenol phản ứng với nước bromine cho thấy khi nhỏ nước bromine vào dung dịch phenol, xuất hiện kết tủa trắng. Cho các mệnh đề sau:"
    )
    stmts_c4_p2 = [
        ("a) Kết tủa trắng là 2,4,6-tribromophenol.", "ĐÚNG"),
        ("b) Phản ứng thế bromine xảy ra ở vị trí ortho và para so với nhóm -OH.", "ĐÚNG"),
        ("c) Phản ứng này dùng để nhận biết phenol.", "ĐÚNG"),
        ("d) Nguyên tử bromine thế vào vòng benzene làm mất màu nước bromine.", "ĐÚNG")
    ]
    for stmt, ans in stmts_c4_p2:
        p_st = add_p(doc, space_before=1, space_after=1)
        p_st.paragraph_format.left_indent = Cm(0.5)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Phenol phản ứng dễ dàng với dung dịch nước bromine ở điều kiện thường tạo kết tủa trắng 2,4,6-tribromophenol và làm mất màu dung dịch nước bromine: "
            "C~6~H~5~OH + 3Br~2~ -> C~6~H~2~Br~3~OH↓ + 3HBr. Phản ứng ưu tiên thế vào 2 vị trí ortho và 1 vị trí para do nhóm -OH đẩy electron làm hoạt hoá vòng thơm, dùng làm phản ứng đặc trưng để nhận biết phenol.\n"
            "=> Cả 4 ý a, b, c, d đều ĐÚNG."
        )

    # =========================================================================
    # PHẦN III
    # =========================================================================
    add_section_header(
        doc,
        title="PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 6."
    )

    # Câu 1
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 1. Cho công thức khung phân tử của dẫn xuất halogen X như hình dưới đây:"
    )
    add_centered_picture(doc, os.path.join(IMG_DIR, "dan_xuat_halogen_x.png"), width_cm=4.5)
    p_x = add_p(doc, space_before=2, space_after=3)
    render_rich_text(p_x, "Số nguyên tử carbon có trong một phân tử chất X là bao nhiêu?")
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Chất X có công thức khung phân tử tương ứng với 2-chloro-3-methylbutane: (CH~3~)~2~CH-CH(Cl)-CH~3~. "
            "Công thức phân tử là C~5~H~11~Cl. Số nguyên tử carbon trong phân tử là 5.\n"
            "=> Đáp số: 5."
        )

    # Câu 2
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 2. Fluoroethane (CH~3~CH~2~F) từng được sử dụng làm chất làm lạnh, được tạo ra theo phản ứng:\n"
        "CH~2~=CH~2~ + HF -> CH~3~CH~2~F    Δ~r~H°~298~ = -73 kJ\n"
        "Dựa vào năng lượng liên kết Eb (kJ/mol) của một số liên kết: C-H: 413; C=C: 614; H-F: 565; C-C: 347. "
        "Hãy xác định năng lượng liên kết C-F (kJ/mol) trong phân tử fluoroethane."
    )
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Biến thiên enthalpy của phản ứng tính theo năng lượng liên kết:\n"
            "Δ~r~H°~298~ = [Eb(C=C) + 4Eb(C-H) + Eb(H-F)] - [Eb(C-C) + 5Eb(C-H) + Eb(C-F)]\n"
            "= Eb(C=C) + Eb(H-F) - Eb(C-C) - Eb(C-H) - Eb(C-F)\n"
            "Thay số: -73 = 614 + 565 - 347 - 413 - Eb(C-F)\n"
            "-73 = 419 - Eb(C-F) => Eb(C-F) = 419 + 73 = 492 kJ/mol.\n"
            "=> Đáp số: 492."
        )

    # Câu 3
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 3. Xylitol là một hợp chất hữu cơ được sử dụng làm chất tạo ngọt tự nhiên cho thực phẩm ăn kiêng và kẹo cao su, có công thức cấu tạo như hình dưới:"
    )
    add_centered_picture(doc, os.path.join(IMG_DIR, "xylitol_structure.png"), width_cm=6.5)
    p_xyl = add_p(doc, space_before=2, space_after=3)
    render_rich_text(p_xyl, "Trong một phân tử Xylitol có bao nhiêu nhóm chức alcohol (-OH)?")
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Công thức cấu tạo của Xylitol là HO-CH~2~-[CH(OH)]~3~-CH~2~-OH. "
            "Phân tử chứa 5 nhóm -OH (thuộc loại polyalcohol có 5 nhóm chức alcohol).\n"
            "=> Đáp số: 5."
        )

    # Câu 4
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 4. Catechin là chất chống oxi hoá mạnh có trong lá chè tươi và nhiều loại hoa quả, có tác dụng phòng ngừa bệnh tim mạch và ung thư. "
        "Công thức cấu tạo của catechin được cho như hình dưới:"
    )
    add_centered_picture(doc, os.path.join(IMG_DIR, "catechin_structure.png"), width_cm=7.0)
    p_cat = add_p(doc, space_before=2, space_after=3)
    render_rich_text(p_cat, "Số nguyên tử carbon có trong một phân tử catechin là bao nhiêu?")
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Phân tử catechin có công thức phân tử là C~15~H~14~O~6~, cấu tạo gồm 2 vòng benzene (mỗi vòng 6 carbon) nối với vòng pyran chứa 3 carbon (6 + 6 + 3 = 15 nguyên tử carbon).\n"
            "=> Đáp số: 15."
        )

    # Câu 5
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 5. Một mẫu hơi thở của người tham gia giao thông có thể tích 52,5 mL được thổi vào thiết bị Breathalyzer chứa 3 mL dung dịch K~2~Cr~2~O~7~ nồng độ 0,042 mg/mL trong môi trường acid H~2~SO~4~. "
        "Toàn bộ C~2~H~5~OH trong hơi thở bị oxi hóa hoàn toàn thành CH~3~COOH theo phản ứng:\n"
        "3C~2~H~5~OH + 2K~2~Cr~2~O~7~ + 8H~2~SO~4~ -> 3CH~3~COOH + 2K~2~SO~4~ + 2Cr~2~(SO~4~)~3~ + 11H~2~O\n"
        "Biết rằng dung dịch màu da cam của Cr~2~O~7~^2-^ chuyển hoàn toàn thành màu xanh của Cr^3+^. "
        "Nồng độ alcohol trong hơi thở của người đó là bao nhiêu mg/L? (Làm tròn đến ba chữ số thập phân)."
    )
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- Khối lượng K~2~Cr~2~O~7~ phản ứng: m = 3 mL × 0,042 mg/mL = 0,126 mg = 1,26 × 10^-4 g.\n"
            "- Số mol K~2~Cr~2~O~7~: n = (1,26 × 10^-4) / 294 ≈ 4,2857 × 10^-7 mol.\n"
            "- Số mol C~2~H~5~OH trong 52,5 mL hơi thở: n(cồn) = 1,5 × n(K~2~Cr~2~O~7~) ≈ 6,4286 × 10^-7 mol.\n"
            "- Khối lượng ethanol trong 52,5 mL hơi thở: m(cồn) = 6,4286 × 10^-7 × 46 ≈ 2,957 × 10^-5 g = 0,02957 mg.\n"
            "- Nồng độ ethanol trong 1 lít (1000 mL) hơi thở:\n"
            "  C = (0,02957 mg / 52,5 mL) × 1000 mL ≈ 0,563 mg/L.\n"
            "=> Đáp số: 0,563."
        )

    # Câu 6
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 6. Từ m kg gạo (chứa 75% khối lượng tinh bột) thực hiện quá trình lên men ethyl alcohol với hiệu suất chung 80%, thu được 10 lít ethyl alcohol 46°. "
        "Biết khối lượng riêng của ethyl alcohol là 0,8 g/mL. Giá trị của m là bao nhiêu kg? (Làm tròn đến một chữ số thập phân)."
    )
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- Thể tích ethyl alcohol nguyên chất: V = 10 × 46% = 4,6 lít = 4600 mL.\n"
            "- Khối lượng ethanol thu được: m(ethanol) = 4600 mL × 0,8 g/mL = 3680 g = 3,68 kg.\n"
            "- Sơ đồ: (C~6~H~10~O~5~)n -> n C~6~H~12~O~6~ -> 2n C~2~H~5~OH.\n"
            "  162 g tinh bột  ----->  92 g ethanol.\n"
            "  Khối lượng tinh bột theo lý thuyết: m(tb, lt) = 3,68 × (162 / 92) = 6,48 kg.\n"
            "- Do hiệu suất H = 80% và tinh bột chiếm 75% khối lượng gạo:\n"
            "  m(gạo) = 6,48 / (0,80 × 0,75) = 6,48 / 0,60 = 10,8 kg.\n"
            "=> Đáp số: 10,8."
        )

    # =========================================================================
    # PHẦN IV
    # =========================================================================
    add_section_header(
        doc,
        title="PHẦN IV. Câu hỏi tự luận.",
        subtitle="Thí sinh trình bày lời giải chi tiết vào giấy làm bài."
    )

    # Câu 1
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 1. Methyl tert-butyl ether (MTBE) có công thức cấu tạo CH~3~-O-C(CH~3~)~3~, là phụ gia pha vào xăng nhằm làm tăng chỉ số octane của xăng dầu.\n"
        "a) Viết phương trình hóa học của phản ứng tạo ra MTBE từ hai alcohol tương ứng. Vì sao phương pháp điều chế MTBE từ hai alcohol tương ứng không phù hợp để tổng hợp MTBE trong công nghiệp?\n"
        "b) Trong công nghiệp, MTBE được sản xuất bằng phản ứng cộng methanol vào 2-methylpropene. Viết phương trình hóa học của phản ứng."
    )
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "a) Phản ứng điều chế MTBE từ methanol và tert-butyl alcohol:\n"
            "   CH~3~OH + (CH~3~)~3~C-OH --(H~2~SO~4~ đặc, t°)--> CH~3~-O-C(CH~3~)~3~ + H~2~O\n"
            "   Phương pháp này không phù hợp trong công nghiệp vì xảy ra sự tạo thành hỗn hợp 3 ether khác nhau (CH~3~OCH~3~, MTBE, và di-tert-butyl ether) "
            "đồng thời tert-butyl alcohol rất dễ bị tách nước tạo alkene (2-methylpropene), dẫn đến hiệu suất thấp và rất khó tách tinh chế.\n"
            "b) Phương trình tổng hợp MTBE trong công nghiệp từ methanol và 2-methylpropene (xúc tác acid):\n"
            "   CH~3~OH + CH~2~=C(CH~3~)~2~ --(H^+, t°)--> CH~3~-O-C(CH~3~)~3~"
        )
    else:
        add_student_writing_lines(doc, num_lines=8)

    # Câu 2
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 2. (SBT-KNTT) Một đèn cồn thí nghiệm chứa 100 mL cồn 90°. Tính nhiệt lượng đèn cồn tỏa ra khi đốt cháy hết lượng cồn trên, "
        "biết khối lượng riêng của ethanol là 0,789 g/mL và nhiệt sinh ra khi đốt cháy 1 mol ethanol là 1 371 kJ/mol. (Làm tròn đến một chữ số thập phân)."
    )
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- Thể tích ethanol nguyên chất trong 100 mL cồn 90°:\n"
            "  V(ethanol) = 100 mL × 90% = 90 mL.\n"
            "- Khối lượng ethanol:\n"
            "  m = 90 mL × 0,789 g/mL = 71,01 g.\n"
            "- Số mol ethanol:\n"
            "  n = 71,01 / 46 ≈ 1,5437 mol.\n"
            "- Lượng nhiệt toả ra khi đốt cháy hoàn toàn:\n"
            "  Q = 1,5437 mol × 1371 kJ/mol ≈ 2116,4 kJ.\n"
            "  (Kết luận: Đèn cồn toả ra khoảng 2 116,4 kJ nhiệt lượng)."
        )
    else:
        add_student_writing_lines(doc, num_lines=8)

    return doc

def check_and_save():
    os.makedirs(OUT_DIR, exist_ok=True)

    # 1. Bản Học sinh
    doc_hs = build_exam_document(is_teacher=False)
    tables_hs = len(doc_hs.tables)
    print(f"Bản Học sinh có {tables_hs} bảng (Chỉ được phép có 1 bảng Header block).")
    assert tables_hs <= 1, "LỖI: Bản Học sinh chứa bảng trong nội dung!"
    path_hs = os.path.join(OUT_DIR, "De_Kiem_Tra_Hoa_11_De13.docx")
    doc_hs.save(path_hs)
    print("Đã lưu Bản Học sinh:", path_hs)

    # 2. Bản Giáo viên
    doc_gv = build_exam_document(is_teacher=True)
    tables_gv = len(doc_gv.tables)
    print(f"Bản Giáo viên có {tables_gv} bảng (Chỉ được phép có 1 bảng Header block).")
    assert tables_gv <= 1, "LỖI: Bản Giáo viên chứa bảng trong nội dung!"
    path_gv = os.path.join(OUT_DIR, "De_Kiem_Tra_Hoa_11_De13_GiaoVien.docx")
    doc_gv.save(path_gv)
    print("Đã lưu Bản Giáo viên:", path_gv)

if __name__ == "__main__":
    check_and_save()
