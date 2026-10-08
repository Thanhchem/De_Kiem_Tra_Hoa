import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
IMG_DIR = r"c:\Antigravity_Thanh\San_Pham\images_crop_101"

COLOR_HF_RED = RGBColor(211, 47, 47)      # Coral Red (#D32F2F)
COLOR_HF_GREEN_HEX = "8BC390"             # Sage Green (#8BC390)

COLOR_PRIMARY = RGBColor(46, 125, 50)     # Dark Green (#2E7D32)
COLOR_CORRECT = RGBColor(198, 40, 40)     # Red (#C62828)
COLOR_SOLUTION = RGBColor(21, 101, 192)   # Blue (#1565C0)
COLOR_GRAY = RGBColor(97, 97, 97)

def add_p_border_bottom(p, color="8BC390", sz="6"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_p_border_top(p, color="8BC390", sz="6"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def set_run_style(run, font_name="Times New Roman", font_size=12, bold=False, italic=False, subscript=False, superscript=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def render_rich_text(p, text, base_font="Times New Roman", base_size=12, base_bold=False, base_italic=False, base_color=None):
    tokens = re.split(r'(~[^~]+~|\^[^\^]+\^|\*[^\*]+\*|_[^_]+_)', text)
    for token in tokens:
        if not token:
            continue
        is_bold = base_bold
        is_italic = base_italic
        is_sub = False
        is_sup = False
        content = token

        if token.startswith('*') and token.endswith('*'):
            is_bold = True
            content = token[1:-1]
        elif token.startswith('_') and token.endswith('_'):
            is_italic = True
            content = token[1:-1]
        elif token.startswith('~') and token.endswith('~'):
            is_sub = True
            content = token[1:-1]
        elif token.startswith('^') and token.endswith('^'):
            is_sup = True
            content = token[1:-1]

        run = p.add_run(content)
        run_color = COLOR_CORRECT if (is_bold and not base_bold and not base_color) else base_color
        set_run_style(run, font_name=base_font, font_size=base_size, bold=is_bold, italic=is_italic,
                      subscript=is_sub, superscript=is_sup, color=run_color)

def setup_header_footer(doc):
    normal_style = doc.styles['Normal']
    normal_style.paragraph_format.tab_stops.clear_all()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    section.start_type = WD_SECTION_START.NEW_PAGE
    section.different_first_page_header_footer = False
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    # 1. Header (bảng 1 hàng 2 cột, rộng 18.0 cm)
    header = section.header
    for p in header.paragraphs:
        p.text = ""
    tbl_h = header.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl_h.cell(0, 0)
    c_right = tbl_h.cell(0, 1)
    c_left.width = Cm(8.0)
    c_right.width = Cm(10.0)

    p_hl = c_left.paragraphs[0]
    p_hl.paragraph_format.space_before = Pt(0)
    p_hl.paragraph_format.space_after = Pt(2)
    p_hl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_hl = p_hl.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)

    p_hr = c_right.paragraphs[0]
    p_hr.paragraph_format.space_before = Pt(0)
    p_hr.paragraph_format.space_after = Pt(2)
    p_hr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p_hr.add_run("Thầy ")
    set_run_style(r1, font_size=10, color=COLOR_HF_RED)
    r2 = p_hr.add_run("TRẦN VĂN THẠNH")
    set_run_style(r2, font_size=10, bold=True, color=COLOR_HF_RED)
    r3 = p_hr.add_run(" - 0777.470.803")
    set_run_style(r3, font_size=10, color=COLOR_HF_RED)

    add_p_border_bottom(p_hl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_bottom(p_hr, color=COLOR_HF_GREEN_HEX, sz="6")

    # 2. Footer (bảng 1 hàng 2 cột, rộng 18.0 cm)
    footer = section.footer
    for p in footer.paragraphs:
        p.text = ""
    tbl_f = footer.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_fl = tbl_f.cell(0, 0)
    c_fr = tbl_f.cell(0, 1)
    c_fl.width = Cm(14.5)
    c_fr.width = Cm(3.5)

    p_fl = c_fl.paragraphs[0]
    p_fl.paragraph_format.space_before = Pt(2)
    p_fl.paragraph_format.space_after = Pt(0)
    p_fl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_fl = p_fl.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)

    p_fr = c_fr.paragraphs[0]
    p_fr.paragraph_format.space_before = Pt(2)
    p_fr.paragraph_format.space_after = Pt(0)
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_field(p_fr, font_size=10, bold=True, color=COLOR_HF_RED)

    add_p_border_top(p_fl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_top(p_fr, color=COLOR_HF_GREEN_HEX, sz="6")

def add_exam_header_block(doc, is_teacher=False):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = tbl._tbl.tblPr
    tblBorders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/><w:insideH w:val="none"/><w:insideV w:val="none"/></w:tblBorders>')
    tbl_pr.append(tblBorders)

    c0 = tbl.cell(0, 0)
    c0.width = Cm(7.5)
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_after = Pt(2)
    p0.paragraph_format.line_spacing = 1.15
    r = p0.add_run("TRƯỜNG ĐẠI HỌC KHOA HỌC\nTRƯỜNG THPT CHUYÊN KHOA HỌC HUẾ\n")
    set_run_style(r, font_size=10, bold=True)
    r_code = p0.add_run("MÃ ĐỀ 101\n")
    set_run_style(r_code, font_size=10.5, bold=True, color=COLOR_PRIMARY)
    r_sub = p0.add_run("(Đề gồm 04 trang)")
    set_run_style(r_sub, font_size=9.5, italic=True)

    c1 = tbl.cell(0, 1)
    c1.width = Cm(10.5)
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_after = Pt(2)
    p1.paragraph_format.line_spacing = 1.15
    r = p1.add_run("ĐỀ KIỂM TRA GIỮA HỌC KỲ I -- NĂM HỌC 2024-2025\n")
    set_run_style(r, font_size=10.5, bold=True)
    r_m = p1.add_run("Môn: Hóa học – Lớp 10\n")
    set_run_style(r_m, font_size=11, bold=True, color=COLOR_PRIMARY)
    if is_teacher:
        r_sub = p1.add_run("(BẢN GIÁO VIÊN -- CÓ LỜI GIẢI CHI TIẾT)")
        set_run_style(r_sub, font_size=10.5, bold=True, color=COLOR_CORRECT)
    else:
        r_sub = p1.add_run("Thời gian làm bài: 45 phút (không kể thời gian giao đề)")
        set_run_style(r_sub, font_size=10, italic=True)

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(4)
    p_info.paragraph_format.space_after = Pt(2)
    p_info.paragraph_format.line_spacing = 1.15
    render_rich_text(p_info, "*Họ, tên thí sinh:* ................................................................ *Lớp:* ........................ *Số báo danh:* ........................", base_size=11)

    p_ntk = doc.add_paragraph()
    p_ntk.paragraph_format.space_before = Pt(0)
    p_ntk.paragraph_format.space_after = Pt(4)
    r_ntk = p_ntk.add_run("(Cho biết nguyên tử khối: H = 1; C = 12; N = 14; O = 16; Na = 23; Mg = 24; Al = 27; P = 31; S = 32; Cl = 35,5; K = 39; Ca = 40; Fe = 56; Cu = 63,54; thể tích mol khí ở đkc (25 °C, 1 bar) là 24,79 L).")
    set_run_style(r_ntk, font_size=10, italic=True)

    # Dòng kẻ phân cách
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(4)
    add_p_border_bottom(p_div, color=COLOR_HF_GREEN_HEX, sz="4")

def add_section_title(doc, title_text, score_text=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"■ {title_text}")
    set_run_style(r, font_size=12, bold=True, color=COLOR_PRIMARY)
    if score_text:
        r_s = p.add_run(f" {score_text}")
        set_run_style(r_s, font_size=11, bold=True, italic=True)

def add_instruction_line(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_style(r, font_size=11, italic=True)

def add_question_prompt(doc, q_num, prompt_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_label = p.add_run(f"Câu {q_num}. ")
    set_run_style(r_label, font_size=12, bold=True)
    render_rich_text(p, prompt_text, base_size=12)
    return p

def add_choices_tabbed(doc, choices, correct_idx=None, is_teacher=False, layout_mode="4cols"):
    labels = ["A. ", "B. ", "C. ", "D. "]

    def set_standard_tabs(p, mode):
        p.paragraph_format.tab_stops.clear_all()
        if mode == "4cols":
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(4.8), WD_TAB_ALIGNMENT.LEFT)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(13.8), WD_TAB_ALIGNMENT.LEFT)
        elif mode == "2cols":
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)
        else:
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)

    if layout_mode == "4cols":
        p = doc.add_paragraph()
        set_standard_tabs(p, "4cols")
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15

        for i in range(4):
            p.add_run('\t')
            is_corr = (is_teacher and correct_idx == i)
            r_l = p.add_run(labels[i])
            set_run_style(r_l, font_size=12, bold=True, color=COLOR_CORRECT if is_corr else None)
            render_rich_text(p, choices[i], base_size=12, base_bold=is_corr, base_color=COLOR_CORRECT if is_corr else None)

    elif layout_mode == "2cols":
        p1 = doc.add_paragraph()
        set_standard_tabs(p1, "2cols")
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.line_spacing = 1.15

        p1.add_run('\t')
        is_corr0 = (is_teacher and correct_idx == 0)
        r0 = p1.add_run(labels[0])
        set_run_style(r0, font_size=12, bold=True, color=COLOR_CORRECT if is_corr0 else None)
        render_rich_text(p1, choices[0], base_size=12, base_bold=is_corr0, base_color=COLOR_CORRECT if is_corr0 else None)

        p1.add_run('\t')
        is_corr1 = (is_teacher and correct_idx == 1)
        r1 = p1.add_run(labels[1])
        set_run_style(r1, font_size=12, bold=True, color=COLOR_CORRECT if is_corr1 else None)
        render_rich_text(p1, choices[1], base_size=12, base_bold=is_corr1, base_color=COLOR_CORRECT if is_corr1 else None)

        p2 = doc.add_paragraph()
        set_standard_tabs(p2, "2cols")
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.line_spacing = 1.15

        p2.add_run('\t')
        is_corr2 = (is_teacher and correct_idx == 2)
        r2 = p2.add_run(labels[2])
        set_run_style(r2, font_size=12, bold=True, color=COLOR_CORRECT if is_corr2 else None)
        render_rich_text(p2, choices[2], base_size=12, base_bold=is_corr2, base_color=COLOR_CORRECT if is_corr2 else None)

        p2.add_run('\t')
        is_corr3 = (is_teacher and correct_idx == 3)
        r3 = p2.add_run(labels[3])
        set_run_style(r3, font_size=12, bold=True, color=COLOR_CORRECT if is_corr3 else None)
        render_rich_text(p2, choices[3], base_size=12, base_bold=is_corr3, base_color=COLOR_CORRECT if is_corr3 else None)

    else: # 1col
        for i in range(4):
            p = doc.add_paragraph()
            set_standard_tabs(p, "1col")
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.15
            p.add_run('\t')
            is_corr = (is_teacher and correct_idx == i)
            r_l = p.add_run(labels[i])
            set_run_style(r_l, font_size=12, bold=True, color=COLOR_CORRECT if is_corr else None)
            render_rich_text(p, choices[i], base_size=12, base_bold=is_corr, base_color=COLOR_CORRECT if is_corr else None)

def add_solution_block(doc, solution_text, answer_key=""):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15

    r_sol = p.add_run("Lời giải chi tiết: ")
    set_run_style(r_sol, font_size=11, bold=True, color=COLOR_SOLUTION)

    if answer_key:
        r_ans = p.add_run(f"({answer_key}) ")
        set_run_style(r_ans, font_size=11, bold=True, color=COLOR_CORRECT)

    render_rich_text(p, solution_text, base_size=11)

def add_true_false_item(doc, item_letter, text, is_true=None, explanation="", is_teacher=False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15

    r_lbl = p.add_run(f"{item_letter}) ")
    set_run_style(r_lbl, font_size=12, bold=True)
    render_rich_text(p, text, base_size=12)

    if is_teacher and is_true is not None:
        status_txt = " --> ĐÚNG" if is_true else " --> SAI"
        r_st = p.add_run(status_txt)
        set_run_style(r_st, font_size=11.5, bold=True, color=COLOR_CORRECT)

        if explanation:
            p_exp = doc.add_paragraph()
            p_exp.paragraph_format.left_indent = Cm(1.2)
            p_exp.paragraph_format.space_before = Pt(0)
            p_exp.paragraph_format.space_after = Pt(2)
            p_exp.paragraph_format.line_spacing = 1.15
            r_g = p_exp.add_run("• Giải thích: ")
            set_run_style(r_g, font_size=11, bold=True, italic=True, color=COLOR_SOLUTION)
            render_rich_text(p_exp, explanation, base_size=11, base_italic=True)

def add_student_lines(doc, num_lines=5):
    for _ in range(num_lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run("." * 105)
        set_run_style(r, font_size=10, color=RGBColor(180, 180, 180))

# ----------------- DATA DEFINITIONS -----------------

MC_QUESTIONS = [
    {
        "num": 1,
        "prompt": "Cho nguyên tử X có 11 electron ở lớp vỏ. Điện tích hạt nhân nguyên tử X là",
        "choices": ["-1,76.10^-18^ C.", "-1,826.10^-18^ C.", "+1,826.10^-18^ C.", "+1,76.10^-18^ C."],
        "correct": 3,
        "ans_letter": "Chọn D",
        "layout": "2cols",
        "solution": "Trong nguyên tử trung hòa về điện, số electron ở vỏ bằng số proton trong hạt nhân: Z = p = e = 11. Hạt nhân mang điện tích dương: q~hạt nhân~ = +Z . e~0~ = +11 . (+1,602.10^-19^ C) ≈ +1,76.10^-18^ C."
    },
    {
        "num": 2,
        "prompt": "Cấu hình electron nguyên tử của oxygen là 1s^2^2s^2^2p^4^. Vị trí của oxygen trong bảng tuần hoàn là",
        "choices": ["chu kì 2, nhóm VIA.", "chu kì 3, nhóm VIA.", "chu kì 2, nhóm IVA.", "chu kì 2, nhóm VIB."],
        "correct": 0,
        "ans_letter": "Chọn A",
        "layout": "2cols",
        "solution": "Oxygen có 2 lớp electron nên thuộc chu kì 2. Lớp ngoài cùng (lớp n = 2) có 2 + 4 = 6 electron, electron cuối cùng điền vào phân lớp 2p (nguyên tố p) nên oxygen thuộc nhóm VIA."
    },
    {
        "num": 3,
        "prompt": "Đối tượng nghiên cứu của hóa học là gì?",
        "choices": [
            "Vật chất, năng lượng và sự vận động của chúng.",
            "Thế giới sinh vật gần gũi với đời sống hằng ngày của học sinh.",
            "Chất và sự biến đổi của chất.",
            "Nghệ thuật ngôn từ."
        ],
        "correct": 2,
        "ans_letter": "Chọn C",
        "layout": "1col",
        "solution": "Hóa học là ngành khoa học tự nhiên nghiên cứu về thành phần, cấu trúc, tính chất và sự biến đổi của chất cũng như ứng dụng của chúng trong đời sống và sản xuất."
    },
    {
        "num": 4,
        "prompt": "Số hiệu nguyên tử cho biết thông tin nào sau đây?",
        "choices": ["Số proton.", "Số neutron.", "Số khối.", "Nguyên tử khối."],
        "correct": 0,
        "ans_letter": "Chọn A",
        "layout": "4cols",
        "solution": "Số hiệu nguyên tử (Z) cho biết số đơn vị điện tích hạt nhân, đồng thời cho biết số proton trong hạt nhân (và bằng số electron trong nguyên tử trung hòa)."
    },
    {
        "num": 5,
        "prompt": "Sự phân bố electron theo ô orbital nào dưới đây là đúng?",
        "is_special_img": True,
        "correct": 0,
        "ans_letter": "Chọn A",
        "solution": "Theo nguyên lý Pauli, mỗi ô orbital chỉ chứa tối đa 2 electron có spin ngược nhau; do đó phương án B (↑↑) và D (↑↑) vi phạm nguyên lý Pauli. Theo quy tắc Hund, các electron phân bố vào các orbital sao cho số electron độc thân là cực đại và có spin song song cùng chiều; phương án C ghép đôi electron khi vẫn còn orbital trống là vi phạm quy tắc Hund. Phương án A tuân thủ đúng nguyên lý Pauli và quy tắc Hund."
    },
    {
        "num": 6,
        "prompt": "Nguyên tử Chlorine (Z = 17) có số electron hóa trị là",
        "choices": ["1.", "3.", "5.", "7."],
        "correct": 3,
        "ans_letter": "Chọn D",
        "layout": "4cols",
        "solution": "Cấu hình electron của Chlorine (Z = 17): 1s^2^2s^2^2p^6^3s^2^3p^5^ (thuộc nhóm VIIA, nguyên tố p). Với nguyên tố nhóm A, số electron hóa trị bằng số electron lớp ngoài cùng. Lớp ngoài cùng (lớp n = 3) có 2 + 5 = 7 electron, nên số electron hóa trị là 7."
    },
    {
        "num": 7,
        "prompt": "Mỗi orbital nguyên tử chứa tối đa",
        "choices": ["1 electron.", "2 electron.", "3 electron.", "4 electron."],
        "correct": 1,
        "ans_letter": "Chọn B",
        "layout": "4cols",
        "solution": "Theo nguyên lý loại trừ Pauli, mỗi orbital nguyên tử chỉ có thể chứa tối đa 2 electron và hai electron này có chiều tự quay (spin) ngược nhau."
    },
    {
        "num": 8,
        "prompt": "Trong tự nhiên, copper có 2 đồng vị là ^63^Cu và ^65^Cu, trong đó đồng vị ^65^Cu chiếm 27% nguyên tử. Phần trăm khối lượng của ^63^Cu trong Cu~2~O là: (cho biết đồng vị oxygen ^16_8^O)",
        "choices": ["73%", "63%", "32,14%", "64,29%"],
        "correct": 3,
        "ans_letter": "Chọn D",
        "layout": "4cols",
        "solution": "Phần trăm số nguyên tử của đồng vị ^63^Cu là: 100% - 27% = 73%. Nguyên tử khối trung bình của Cu là: A~Cu~ = (63 . 73 + 65 . 27) / 100 = 63,54. Khối lượng mol phân tử Cu~2~O là: M~Cu2O~ = 2 . 63,54 + 16 = 143,08 g/mol. Trong 1 mol Cu~2~O có 2 mol nguyên tử Cu, khối lượng của đồng vị ^63^Cu là: m(^63^Cu) = 2 . 73% . 63 = 91,98 gam. Phần trăm khối lượng của ^63^Cu trong Cu~2~O là: %m(^63^Cu) = (91,98 / 143,08) . 100% ≈ 64,29%."
    },
    {
        "num": 9,
        "prompt": "Cho các phát biểu sau:\n(1) Nguyên tử K có điện tích hạt nhân là +3,0438.10^-18^ C.\n(2) Khối lượng hạt nhân được xem như là khối lượng nguyên tử.\n(3) 1 amu bằng 1/12 khối lượng của nguyên tử carbon - 12.\n(4) Đường kính hạt nhân gần bằng đường kính nguyên tử.\nSố phát biểu đúng là",
        "choices": ["1.", "2.", "3.", "4."],
        "correct": 2,
        "ans_letter": "Chọn C",
        "layout": "4cols",
        "solution": "(1) ĐÚNG: Nguyên tử K có Z = 19, điện tích hạt nhân q = +19 . 1,602.10^-19^ C = +3,0438.10^-18^ C.\n(2) ĐÚNG: Do electron có khối lượng rất nhỏ không đáng kể so với proton và neutron nên khối lượng hạt nhân được xem như khối lượng nguyên tử.\n(3) ĐÚNG: Theo định nghĩa, 1 amu = 1/12 khối lượng một nguyên tử carbon-12.\n(4) SAI: Đường kính nguyên tử lớn hơn đường kính hạt nhân khoảng 10 000 lần.\nVậy có 3 phát biểu đúng là (1), (2), (3)."
    },
    {
        "num": 10,
        "prompt": "Sự phân bố electron vào các lớp và phân lớp căn cứ vào",
        "choices": ["nguyên tử khối tăng dần.", "điện tích hạt nhân tăng dần.", "số khối tăng dần.", "mức năng lượng electron."],
        "correct": 3,
        "ans_letter": "Chọn D",
        "layout": "2cols",
        "solution": "Theo nguyên lý vững bền, ở trạng thái cơ bản, các electron trong nguyên tử lần lượt chiếm các mức năng lượng từ thấp đến cao (căn cứ vào mức năng lượng của electron)."
    },
    {
        "num": 11,
        "prompt": "Số nguyên tố thuộc chu kì 3 của bảng tuần hoàn là",
        "choices": ["2.", "8.", "18.", "32."],
        "correct": 1,
        "ans_letter": "Chọn B",
        "layout": "4cols",
        "solution": "Chu kì 3 bắt đầu từ nguyên tố Sodium (~11~Na) đến Argon (~18~Ar), gồm tổng cộng 8 nguyên tố."
    },
    {
        "num": 12,
        "prompt": "Các hạt cấu tạo nên hạt nhân nguyên tử là",
        "choices": ["neutron và electron.", "electron, proton và neutron.", "electron và proton.", "proton và neutron."],
        "correct": 3,
        "ans_letter": "Chọn D",
        "layout": "2cols",
        "solution": "Hạt nhân nguyên tử được cấu tạo từ các hạt proton (mang điện tích dương) và neutron (không mang điện tích). Vỏ nguyên tử được cấu tạo từ electron."
    },
    {
        "num": 13,
        "prompt": "Năm 1897, nhà vật lý người Anh Joseph John Thomson thực hiện thí nghiệm phóng điện trong ống thủy tinh gần như chân không với hiệu điện thế lớn (15 kV). Mô hình thí nghiệm như hình vẽ bên. Nếu đặt một chong chóng nhẹ trên đường đi của tia âm cực thì chong chóng sẽ quay. Hiện tượng này chứng tỏ điều gì về tia âm cực?",
        "has_image": "cau13_clean.png",
        "choices": [
            "Tia âm cực mang điện tích âm.",
            "Tia âm cực là một loại ánh sáng trắng như ánh sáng mặt trời.",
            "Tia âm cực có phương truyền thẳng.",
            "Tia âm cực là chùm hạt vật chất chuyển động với vận tốc rất lớn."
        ],
        "correct": 3,
        "ans_letter": "Chọn D",
        "layout": "1col",
        "solution": "Chong chóng quay khi bị tia âm cực bắn vào chứng tỏ tia âm cực có động lượng, tức là chùm hạt vật chất có khối lượng và chuyển động với vận tốc rất lớn va chạm cơ học làm quay chong chóng."
    },
    {
        "num": 14,
        "prompt": "Trong tự nhiên oxygen có 3 đồng vị bền: ^16_8^O, ^17_8^O, ^18_8^O, còn carbon có 2 đồng vị bền: ^12_6^C, ^13_6^C. Số lượng phân tử CO~2~ tạo ra từ các đồng vị trên là:",
        "choices": ["8.", "10.", "12.", "6."],
        "correct": 2,
        "ans_letter": "Chọn C",
        "layout": "4cols",
        "solution": "Phân tử CO~2~ có cấu trúc dạng O = C = O đối xứng. Số cặp gồm 2 nguyên tử oxygen (kể cả giống nhau và khác nhau) là: 3 . (3 + 1) / 2 = 6 cặp (3 cặp O giống nhau và 3 cặp O khác nhau). Với 2 đồng vị của carbon, mỗi cặp oxygen kết hợp với 1 nguyên tử C sẽ tạo ra 1 phân tử CO~2~. Tổng số phân tử CO~2~ tạo thành là: 6 . 2 = 12 phân tử."
    },
    {
        "num": 15,
        "prompt": "Nguyên tử X có tổng số hạt cơ bản là 40. Trong đó tổng số hạt mang điện nhiều hơn số hạt không mang điện là 12 hạt. Nguyên tử X và số hiệu nguyên tử là",
        "choices": ["Na (Z = 11).", "Mg (Z = 12).", "Al (Z = 13).", "Cl (Z = 17)."],
        "correct": 2,
        "ans_letter": "Chọn C",
        "layout": "2cols",
        "solution": "Gọi p, n, e là số proton, neutron, electron của nguyên tử X (p = e). Ta có hệ phương trình:\n2p + n = 40 và 2p - n = 12\n=> 4p = 52 => p = 13 (Al) và n = 14. Số hiệu nguyên tử Z = p = 13, X là Aluminium (Al)."
    },
    {
        "num": 16,
        "prompt": "Trong các hạt sau đây, hạt nào không mang điện tích?",
        "choices": ["Electron.", "Neutron.", "Electron và proton.", "Proton."],
        "correct": 1,
        "ans_letter": "Chọn B",
        "layout": "2cols",
        "solution": "Proton mang điện tích dương (+1), electron mang điện tích âm (-1), neutron là hạt trung hòa về điện (không mang điện tích)."
    },
    {
        "num": 17,
        "prompt": "Nguyên tố nào sau đây thuộc nhóm A?",
        "choices": ["[Ne]3s^2^3p^3^.", "[Ar]3d^1^4s^2^.", "[Ar]3d^7^4s^2^.", "[Ar]3d^5^4s^2^."],
        "correct": 0,
        "ans_letter": "Chọn A",
        "layout": "4cols",
        "solution": "Nguyên tố nhóm A là các nguyên tố s và nguyên tố p (electron cuối cùng điền vào phân lớp s hoặc p). Cấu hình [Ne]3s^2^3p^3^ có electron cuối cùng điền vào phân lớp 3p nên là nguyên tố p, thuộc nhóm VA. Các cấu hình còn lại đều là nguyên tố d thuộc nhóm B."
    },
    {
        "num": 18,
        "prompt": "Khối lượng nguyên tử gần bằng khối lượng hạt nhân vì",
        "choices": [
            "khối lượng electron gần bằng khối lượng hạt nhân.",
            "số lượng electron quá ít.",
            "tổng khối lượng electron không đáng kể.",
            "khối lượng nhân quá lớn."
        ],
        "correct": 2,
        "ans_letter": "Chọn C",
        "layout": "1col",
        "solution": "Khối lượng nguyên tử bằng tổng khối lượng của proton, neutron và electron. Vì khối lượng của mỗi electron rất nhỏ (chỉ khoảng 1/1836 khối lượng proton) nên tổng khối lượng của các electron không đáng kể so với hạt nhân. Do đó khối lượng nguyên tử xấp xỉ khối lượng hạt nhân."
    },
    {
        "num": 19,
        "prompt": "So sánh nào dưới đây về mức năng lượng của các phân lớp là không phù hợp?",
        "choices": ["3d < 4s", "3p < 3d", "1s < 2s", "4s > 3s"],
        "correct": 0,
        "ans_letter": "Chọn A",
        "layout": "4cols",
        "solution": "Theo thứ tự mức năng lượng tăng dần: 1s < 2s < 2p < 3s < 3p < 4s < 3d < 4p... Do đó mức năng lượng của phân lớp 4s thấp hơn phân lớp 3d (4s < 3d). So sánh 3d < 4s là không phù hợp (sai)."
    },
    {
        "num": 20,
        "prompt": "Sulfur dạng kem bôi được sử dụng để điều trị mụn trứng cá. Nguyên tử sulfur có phân lớp electron ngoài cùng là 3p^4^. Phát biểu nào sau đây không đúng về nguyên tử sulfur?",
        "choices": [
            "Lớp ngoài cùng có 4 electron.",
            "Nguyên tử có 16 electron.",
            "Thuộc chu kỳ 3 trong bảng tuần hoàn.",
            "Thuộc nhóm VIA trong bảng tuần hoàn."
        ],
        "correct": 0,
        "ans_letter": "Chọn A",
        "layout": "2cols",
        "solution": "Cấu hình electron đầy đủ của sulfur là 1s^2^2s^2^2p^6^3s^2^3p^4^. Lớp ngoài cùng (lớp thứ 3) có 2 + 4 = 6 electron (chứ không phải 4 electron). Phát biểu A sai."
    }
]

TF_QUESTIONS = [
    {
        "num": 21,
        "prompt": "X là nguyên tố phổ biến thứ 4 trong vỏ trái đất, X có trong hemoglobin của máu làm nhiệm vụ vận chuyển oxygen, duy trì sự sống. Nguyên tử X có 26 proton trong hạt nhân.",
        "items": [
            {
                "letter": "a",
                "text": "X có 26 neutron trong hạt nhân.",
                "is_true": False,
                "explanation": "Đề bài cho X có 26 proton (Z = 26 là nguyên tố sắt/Fe). Đồng vị bền phổ biến nhất của sắt là ^56_26^Fe có số neutron là N = 56 - 26 = 30 neutron. Hạt nhân Fe không có 26 neutron."
            },
            {
                "letter": "b",
                "text": "X có 26 electron ở vỏ nguyên tử.",
                "is_true": True,
                "explanation": "Trong nguyên tử trung hòa về điện, số electron ở vỏ bằng số proton trong hạt nhân: e = p = 26."
            },
            {
                "letter": "c",
                "text": "X có điện tích hạt nhân là + 26.",
                "is_true": True,
                "explanation": "Hạt nhân có 26 proton nên điện tích hạt nhân là +26 (theo đơn vị điện tích nguyên tố e~0~) hoặc +26 . 1,602.10^-19^ C."
            },
            {
                "letter": "d",
                "text": "Khối lượng nguyên tử X là 26 amu.",
                "is_true": False,
                "explanation": "Khối lượng nguyên tử Fe xấp xỉ số khối (A ≈ 56 amu). Giá trị 26 chỉ là số proton hoặc số đơn vị điện tích hạt nhân."
            }
        ]
    },
    {
        "num": 22,
        "prompt": "Hình dưới mô tả orbital (a) và orbital (b) chứa electron trong nguyên tử sodium (Na) ở trạng thái cơ bản. Mức năng lượng của orbital (a) cao hơn orbital (b).",
        "has_image": "cau22_clean.png",
        "items": [
            {
                "letter": "a",
                "text": "Electron trong các orbital (a) và (b) thuộc cùng lớp electron.",
                "is_true": False,
                "explanation": "Sodium (Na, Z = 11) có cấu hình: 1s^2^2s^2^2p^6^3s^1^. Orbital (a) hình cầu có năng lượng cao hơn orbital (b) hình số 8 nổi (2p), do đó (a) là orbital 3s (lớp n = 3) còn (b) là orbital 2p (lớp n = 2). Chúng thuộc hai lớp electron khác nhau."
            },
            {
                "letter": "b",
                "text": "Số electron trong 1 orbital (b) gấp ba số electron trong orbital (a).",
                "is_true": False,
                "explanation": "Trong nguyên tử Na, mỗi orbital 2p chứa tối đa 2 electron, orbital 3s chứa 1 electron (3s^1^). Số electron trong 1 orbital (b) là 2, gấp 2 lần (chứ không phải gấp 3 lần) số electron trong orbital (a)."
            },
            {
                "letter": "c",
                "text": "Electron trên orbital (a) nằm gần hạt nhân hơn electron trên orbital (b).",
                "is_true": False,
                "explanation": "Orbital 3s (a) thuộc lớp thứ 3 nằm xa hạt nhân hơn orbital 2p (b) thuộc lớp thứ 2."
            },
            {
                "letter": "d",
                "text": "Orbital (a) và (b) khác nhau về định hướng trong không gian.",
                "is_true": True,
                "explanation": "Orbital s (a) đối xứng cầu không có định hướng ưu tiên, còn orbital p (b) có định hướng xác định trong không gian theo các trục tọa độ Ox, Oy, Oz."
            }
        ]
    }
]

ESSAY_QUESTIONS = [
    {
        "num": 23,
        "score_str": "(1,0 điểm)",
        "prompt": "Trong tự nhiên, magnesium có 3 đồng vị bền là ^24^Mg, ^25^Mg và ^26^Mg. Phương pháp phổ khối lượng xác nhận đồng vị ^26^Mg chiếm tỉ lệ phần trăm số nguyên tử là 11%. Biết rằng nguyên tử khối trung bình của Mg là 24,32. Tính % số nguyên tử của đồng vị ^24^Mg, đồng vị ^25^Mg?",
        "solution": (
            "• Gọi x (%) và y (%) lần lượt là phần trăm số nguyên tử của đồng vị ^24^Mg và ^25^Mg (điều kiện: x, y > 0).\n"
            "• Tổng phần trăm số nguyên tử của 3 đồng vị bằng 100%:\n"
            "  x + y + 11 = 100  =>  x + y = 89    (1)\n"
            "• Theo công thức tính nguyên tử khối trung bình của Mg:\n"
            "  A~tb~ = (24x + 25y + 26 . 11) / 100 = 24,32\n"
            "  => 24x + 25y + 286 = 2432\n"
            "  => 24x + 25y = 2146    (2)\n"
            "• Từ (1) và (2) ta có hệ phương trình bậc nhất hai ẩn:\n"
            "  { x + y = 89\n"
            "  { 24x + 25y = 2146\n"
            "  Giải hệ phương trình ta được: x = 79 và y = 10.\n"
            "• Kết luận: Phần trăm số nguyên tử của đồng vị ^24^Mg là *79%*, đồng vị ^25^Mg là *10%*."
        ),
        "student_lines": 8
    },
    {
        "num": 24,
        "score_str": "(1,0 điểm)",
        "prompt": (
            "Cho hai nguyên tử X (Z = 15) và Y (Z = 26)\n"
            "a) Viết cấu hình electron và xác định vị trí của mỗi nguyên tử trong bảng tuần hoàn hóa học.\n"
            "b) Cho biết X và Y là nguyên tố kim loại, phi kim hay khí hiếm? Giải thích."
        ),
        "solution": (
            "a) Cấu hình electron và vị trí trong bảng tuần hoàn:\n"
            "• Với nguyên tử X (Z = 15):\n"
            "  - Cấu hình electron: 1s^2^2s^2^2p^6^3s^2^3p^3^ (hoặc viết gọn [Ne]3s^2^3p^3^).\n"
            "  - Vị trí trong bảng tuần hoàn:\n"
            "    + Ô số: 15 (vì có Z = 15, số proton = 15).\n"
            "    + Chu kì: 3 (vì có 3 lớp electron).\n"
            "    + Nhóm: VA (vì là nguyên tố p và có 5 electron lớp ngoài cùng).\n"
            "• Với nguyên tử Y (Z = 26):\n"
            "  - Cấu hình electron: 1s^2^2s^2^2p^6^3s^2^3p^6^3d^6^4s^2^ (hoặc viết gọn [Ar]3d^6^4s^2^).\n"
            "  - Vị trí trong bảng tuần hoàn:\n"
            "    + Ô số: 26 (vì có Z = 26, số proton = 26).\n"
            "    + Chu kì: 4 (vì có 4 lớp electron).\n"
            "    + Nhóm: VIIIB (vì là nguyên tố d, tổng số electron phân lớp 3d và 4s là 6 + 2 = 8).\n"
            "b) Xác định tính chất nguyên tố:\n"
            "• Nguyên tố X là *phi kim* vì có 5 electron ở lớp ngoài cùng (3s^2^3p^3^).\n"
            "• Nguyên tố Y là *kim loại* (kim loại chuyển tiếp) vì có 2 electron ở lớp ngoài cùng (4s^2^) và phân lớp 3d chưa bão hòa."
        ),
        "student_lines": 10
    },
    {
        "num": 25,
        "score_str": "(1,0 điểm)",
        "prompt": "Hòa tan hoàn toàn 20 gam hỗn hợp 2 nguyên tố A và B thuộc nhóm IIA, ở 2 chu kì liên tiếp nhau vào dung dịch HCl dư thu được 17,353 lít khí (đktc). Xác định tên 2 nguyên tố A,B và thành phần % về khối lượng của mỗi nguyên tố trong hỗn hợp.",
        "solution": (
            "• Tính số mol khí H~2~ thoát ra theo điều kiện chuẩn (đkc: 25 °C, 1 bar):\n"
            "  n~H2~ = 17,353 / 24,79 = 0,7 mol.\n"
            "  _(Ghi chú: Đề bài ghi kí hiệu đktc theo thói quen cũ nhưng dùng số liệu chuẩn V~mol~ = 24,79 L/mol theo chương trình GDPT 2018)._\n"
            "• Gọi M~tb~ là nguyên tử khối trung bình của hai kim loại A và B.\n"
            "  Phương trình phản ứng: M~tb~ + 2HCl ---> M~tb~Cl~2~ + H~2~^\n"
            "  Theo phương trình phản ứng: n~hh kim loại~ = n~H2~ = 0,7 mol.\n"
            "• Khối lượng mol trung bình của hai kim loại:\n"
            "  M~tb~ = m~hh~ / n~hh~ = 20 / 0,7 ≈ 28,57 g/mol.\n"
            "• Xác định tên hai nguyên tố:\n"
            "  Vì A và B là 2 kim loại thuộc nhóm IIA ở hai chu kì liên tiếp:\n"
            "  Dãy kim loại nhóm IIA gồm: Be (9) < Mg (24) < Ca (40) < Sr (88) < Ba (137).\n"
            "  Vì M~A~ < M~tb~ = 28,57 < M~B~, suy ra:\n"
            "  M~A~ = 24 (Mg, chu kì 3) và M~B~ = 40 (Ca, chu kì 4).\n"
            "  Vậy hai nguyên tố cần tìm là *Magnesium (Mg)* và *Calcium (Ca)*.\n"
            "• Tính thành phần % về khối lượng của mỗi nguyên tố:\n"
            "  Gọi a và b lần lượt là số mol của Mg và Ca trong 20 gam hỗn hợp (a, b > 0):\n"
            "  Ta có hệ phương trình:\n"
            "  { a + b = 0,7 (theo số mol H~2~)\n"
            "  { 24a + 40b = 20 (khối lượng hỗn hợp)\n"
            "  Giải hệ phương trình thu được: a = 0,5 mol và b = 0,2 mol.\n"
            "  - Khối lượng của Mg: m~Mg~ = 0,5 . 24 = 12 gam.\n"
            "  - Khối lượng của Ca: m~Ca~ = 0,2 . 40 = 8 gam.\n"
            "  - Phần trăm khối lượng mỗi kim loại trong hỗn hợp:\n"
            "    %m~Mg~ = (12 / 20) . 100% = *60%*.\n"
            "    %m~Ca~ = (8 / 20) . 100% = *40%*."
        ),
        "student_lines": 12
    }
]

def build_docx(is_teacher=False):
    doc = docx.Document()
    setup_header_footer(doc)
    add_exam_header_block(doc, is_teacher=is_teacher)

    # PHẦN I
    add_section_title(doc, "PHẦN I. TRẮC NGHIỆM NHIỀU LỰA CHỌN", "(5,0 điểm -- 20 câu)")
    add_instruction_line(doc, "Thí sinh trả lời từ câu 1 đến câu 20. Mỗi câu hỏi thí sinh chỉ chọn một phương án.")

    for q in MC_QUESTIONS:
        q_num = q["num"]
        prompt = q["prompt"]
        add_question_prompt(doc, q_num, prompt)

        if q.get("has_image"):
            img_path = os.path.join(IMG_DIR, q["has_image"])
            if os.path.exists(img_path):
                p_im = doc.add_paragraph()
                p_im.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_im.paragraph_format.space_before = Pt(2)
                p_im.paragraph_format.space_after = Pt(4)
                p_im.add_run().add_picture(img_path, width=Cm(11.0))

        if q.get("is_special_img"):
            # Câu 5: Special images for choices A, B, C, D
            p_c = doc.add_paragraph()
            p_c.paragraph_format.tab_stops.clear_all()
            p_c.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
            p_c.paragraph_format.tab_stops.add_tab_stop(Cm(4.8), WD_TAB_ALIGNMENT.LEFT)
            p_c.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)
            p_c.paragraph_format.tab_stops.add_tab_stop(Cm(13.8), WD_TAB_ALIGNMENT.LEFT)
            p_c.paragraph_format.space_before = Pt(0)
            p_c.paragraph_format.space_after = Pt(2)
            p_c.paragraph_format.line_spacing = 1.15

            letters = ["A. ", "B. ", "C. ", "D. "]
            img_files = ["cau5_a.png", "cau5_b.png", "cau5_c.png", "cau5_d.png"]
            for idx in range(4):
                p_c.add_run('\t')
                is_c = (is_teacher and idx == 0)
                r_l = p_c.add_run(letters[idx])
                set_run_style(r_l, font_size=12, bold=True, color=COLOR_CORRECT if is_c else None)
                im_p = os.path.join(IMG_DIR, img_files[idx])
                if os.path.exists(im_p):
                    p_c.add_run().add_picture(im_p, height=Pt(18))
        else:
            add_choices_tabbed(doc, q["choices"], correct_idx=q["correct"], is_teacher=is_teacher, layout_mode=q["layout"])

        if is_teacher:
            add_solution_block(doc, q["solution"], answer_key=q.get("ans_letter", ""))

    # PHẦN II
    add_section_title(doc, "PHẦN II. TRẮC NGHIỆM ĐÚNG SAI", "(2,0 điểm -- 2 câu)")
    add_instruction_line(doc, "Thí sinh trả lời từ câu 1 đến câu 2. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn ĐÚNG hoặc SAI.")

    for q in TF_QUESTIONS:
        q_num = q["num"]
        prompt = q["prompt"]
        add_question_prompt(doc, q_num, prompt)

        if q.get("has_image"):
            img_path = os.path.join(IMG_DIR, q["has_image"])
            if os.path.exists(img_path):
                p_im = doc.add_paragraph()
                p_im.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_im.paragraph_format.space_before = Pt(2)
                p_im.paragraph_format.space_after = Pt(4)
                p_im.add_run().add_picture(img_path, width=Cm(10.0))

        p_cho = doc.add_paragraph()
        p_cho.paragraph_format.space_before = Pt(0)
        p_cho.paragraph_format.space_after = Pt(1)
        r_cho = p_cho.add_run("Cho các phát biểu sau:")
        set_run_style(r_cho, font_size=11.5, italic=True)

        for it in q["items"]:
            add_true_false_item(doc, it["letter"], it["text"], is_true=it["is_true"],
                                explanation=it.get("explanation", ""), is_teacher=is_teacher)

    # PHẦN III
    add_section_title(doc, "PHẦN III. TỰ LUẬN", "(3,0 điểm -- 3 câu)")
    add_instruction_line(doc, "Thí sinh trình bày chi tiết các bước giải và đáp số vào bài làm.")

    for q in ESSAY_QUESTIONS:
        q_num = q["num"]
        score = q["score_str"]
        prompt = q["prompt"]

        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        r_label = p_q.add_run(f"Câu {q_num}. ")
        set_run_style(r_label, font_size=12, bold=True)
        r_sc = p_q.add_run(f"{score} ")
        set_run_style(r_sc, font_size=11, bold=True, italic=True, color=COLOR_PRIMARY)
        render_rich_text(p_q, prompt, base_size=12)

        if is_teacher:
            add_solution_block(doc, q["solution"])
        else:
            add_student_lines(doc, num_lines=q.get("student_lines", 6))

    # FOOTER END BLOCK
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(8)
    p_end.paragraph_format.space_after = Pt(2)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end_note = p_end.add_run("Ghi chú: Cán bộ coi thi không được giải thích gì thêm.\n")
    set_run_style(r_end_note, font_size=10, italic=True)
    r_end = p_end.add_run("--- HẾT ---")
    set_run_style(r_end, font_size=11, bold=True)

    # Signature block using Paragraph Tab Stops (NO TABLES in exam content)
    p_sig1 = doc.add_paragraph()
    p_sig1.paragraph_format.space_before = Pt(10)
    p_sig1.paragraph_format.space_after = Pt(2)
    p_sig1.paragraph_format.tab_stops.clear_all()
    p_sig1.paragraph_format.tab_stops.add_tab_stop(Cm(3.5), WD_TAB_ALIGNMENT.CENTER)
    p_sig1.paragraph_format.tab_stops.add_tab_stop(Cm(14.5), WD_TAB_ALIGNMENT.CENTER)

    p_sig1.add_run('\t')
    r_d1 = p_sig1.add_run("DUYỆT")
    set_run_style(r_d1, font_size=11, bold=True)
    p_sig1.add_run('\t')
    r_r1 = p_sig1.add_run("CÁN BỘ RA ĐỀ")
    set_run_style(r_r1, font_size=11, bold=True)

    p_sig2 = doc.add_paragraph()
    p_sig2.paragraph_format.space_before = Pt(0)
    p_sig2.paragraph_format.space_after = Pt(30)
    p_sig2.paragraph_format.tab_stops.clear_all()
    p_sig2.paragraph_format.tab_stops.add_tab_stop(Cm(3.5), WD_TAB_ALIGNMENT.CENTER)
    p_sig2.paragraph_format.tab_stops.add_tab_stop(Cm(14.5), WD_TAB_ALIGNMENT.CENTER)

    p_sig2.add_run('\t')
    r_d2 = p_sig2.add_run("(Ký và ghi rõ họ tên)")
    set_run_style(r_d2, font_size=10, italic=True)
    p_sig2.add_run('\t')
    r_r2 = p_sig2.add_run("(Ký và ghi rõ họ tên)")
    set_run_style(r_r2, font_size=10, italic=True)

    p_sig3 = doc.add_paragraph()
    p_sig3.paragraph_format.space_before = Pt(0)
    p_sig3.paragraph_format.space_after = Pt(4)
    p_sig3.paragraph_format.tab_stops.clear_all()
    p_sig3.paragraph_format.tab_stops.add_tab_stop(Cm(3.5), WD_TAB_ALIGNMENT.CENTER)
    p_sig3.paragraph_format.tab_stops.add_tab_stop(Cm(14.5), WD_TAB_ALIGNMENT.CENTER)

    p_sig3.add_run('\t')
    r_name1 = p_sig3.add_run("Nguyễn Hồ Ngọc Thư")
    set_run_style(r_name1, font_size=11, bold=True)
    p_sig3.add_run('\t')
    r_name2 = p_sig3.add_run("Tôn Nữ Mỹ Phương")
    set_run_style(r_name2, font_size=11, bold=True)

    filename = "De_Kiem_Tra_Hoa_10_Ma101_GiaoVien.docx" if is_teacher else "De_Kiem_Tra_Hoa_10_Ma101.docx"
    filepath = os.path.join(OUT_DIR, filename)
    doc.save(filepath)
    print(f"Generated: {filepath}")

if __name__ == "__main__":
    build_docx(is_teacher=False)
    build_docx(is_teacher=True)
