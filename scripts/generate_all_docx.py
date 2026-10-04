import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

CROP_DIR = r"c:\Antigravity_Thanh\San_Pham\images_crop"
OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"

# Colors matching user image media_1791087945720.png
COLOR_HF_RED = RGBColor(211, 47, 47)      # Coral Red (#D32F2F)
COLOR_HF_GREEN_HEX = "8BC390"             # Sage Green for horizontal lines (#8BC390 / #A4C7A6)

COLOR_PRIMARY = RGBColor(46, 125, 50)     # Dark Green / Teal for headers in body
COLOR_CORRECT = RGBColor(198, 40, 40)     # Red for correct answers
COLOR_SOLUTION = RGBColor(21, 101, 192)   # Blue for solutions
COLOR_GRAY = RGBColor(97, 97, 97)

def add_p_border_bottom(p, color="8BC390", sz="8"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_p_border_top(p, color="8BC390", sz="8"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def set_run_style(run, font_name="Times New Roman", font_size=12, bold=False, italic=False, subscript=False, superscript=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.subscript = subscript
    run.font.superscript = superscript
    if color:
        run.font.color.rgb = color

def render_rich_text(p, text, base_font="Times New Roman", base_size=12, base_bold=False, base_italic=False, base_color=None):
    """
    Renders text with ~subscript~, ^superscript^, **bold**, *italic* tags into docx paragraph runs.
    """
    pattern = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|~[^~]+~|\^[^^]+\^)')
    tokens = pattern.split(text)

    for token in tokens:
        if not token:
            continue
        
        is_bold = base_bold
        is_italic = base_italic
        is_sub = False
        is_sup = False
        content = token

        if token.startswith('**') and token.endswith('**'):
            is_bold = True
            content = token[2:-2]
        elif token.startswith('*') and token.endswith('*'):
            is_italic = True
            content = token[1:-1]
        elif token.startswith('~') and token.endswith('~'):
            is_sub = True
            content = token[1:-1]
        elif token.startswith('^') and token.endswith('^'):
            is_sup = True
            content = token[1:-1]

        run = p.add_run(content)
        set_run_style(run, font_name=base_font, font_size=base_size, bold=is_bold, italic=is_italic,
                      subscript=is_sub, superscript=is_sup, color=base_color)

def setup_header_footer(doc):
    section = doc.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.different_first_page_header_footer = True

    # 1. DEFAULT HEADER (Pages 2+) - Matching media_1791087945720.png exactly
    header = section.header
    hp = header.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
    hp.paragraph_format.space_before = Pt(0)
    hp.paragraph_format.space_after = Pt(2)

    r_hl = hp.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)

    hp.add_run('\t')

    r_hr1 = hp.add_run("Thầy ")
    set_run_style(r_hr1, font_size=10, italic=True, color=COLOR_HF_RED)
    r_hr2 = hp.add_run("TRẦN VĂN THẠNH")
    set_run_style(r_hr2, font_size=10, bold=True, italic=True, color=COLOR_HF_RED)
    r_hr3 = hp.add_run(" - 0777.470.803")
    set_run_style(r_hr3, font_size=10, italic=True, color=COLOR_HF_RED)

    add_p_border_bottom(hp, color=COLOR_HF_GREEN_HEX, sz="8")

    # 2. DEFAULT FOOTER (Pages 2+) - Matching media_1791087945720.png exactly
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
    fp.paragraph_format.space_before = Pt(4)
    fp.paragraph_format.space_after = Pt(0)

    r_fl = fp.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế | CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)

    fp.add_run('\t')
    add_page_number_field(fp, font_size=10, bold=True, color=COLOR_HF_RED)

    add_p_border_top(fp, color=COLOR_HF_GREEN_HEX, sz="8")

    # 3. FIRST PAGE HEADER (Empty)
    first_header = section.first_page_header
    fhp = first_header.paragraphs[0]
    fhp.text = ""

    # 4. FIRST PAGE FOOTER (Has footer matching media_1791087945720.png)
    first_footer = section.first_page_footer
    ffp = first_footer.paragraphs[0]
    ffp.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
    ffp.paragraph_format.space_before = Pt(4)
    ffp.paragraph_format.space_after = Pt(0)

    r_ffl = ffp.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế | CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_ffl, font_size=9.5, italic=True, color=COLOR_HF_RED)

    ffp.add_run('\t')
    add_page_number_field(ffp, font_size=10, bold=True, color=COLOR_HF_RED)

    add_p_border_top(ffp, color=COLOR_HF_GREEN_HEX, sz="8")

def add_exam_header_block(doc, is_teacher=False):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = tbl._tbl.tblPr
    tblBorders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/><w:insideH w:val="none"/><w:insideV w:val="none"/></w:tblBorders>')
    tbl_pr.append(tblBorders)

    c0 = tbl.cell(0, 0)
    c0.width = Cm(7.5)
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_after = Pt(2)
    p0.paragraph_format.line_spacing = 1.15
    r = p0.add_run("SỞ GIÁO DỤC & ĐÀO TẠO TP HUẾ\nTRƯỜNG THPT HAI BÀ TRƯNG\n")
    set_run_style(r, font_size=11, bold=True)
    r_code = p0.add_run("Mã đề thi: 209")
    set_run_style(r_code, font_size=11, bold=True, color=COLOR_PRIMARY)

    c1 = tbl.cell(0, 1)
    c1.width = Cm(10.5)
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_after = Pt(2)
    p1.paragraph_format.line_spacing = 1.15
    r = p1.add_run("ĐỀ KIỂM TRA CUỐI KÌ II - NĂM HỌC 2024-2025\nMÔN HÓA HỌC LỚP 11\n")
    set_run_style(r, font_size=11, bold=True)
    if is_teacher:
        r_sub = p1.add_run("(HƯỚNG DẪN CHẤM VÀ LỜI GIẢI CHI TIẾT)")
        set_run_style(r_sub, font_size=10.5, bold=True, color=COLOR_CORRECT)
    else:
        r_sub = p1.add_run("Thời gian làm bài: 45 phút (không kể thời gian phát đề)")
        set_run_style(r_sub, font_size=10.5, italic=True)

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(4)
    p_info.paragraph_format.space_after = Pt(2)
    p_info.paragraph_format.line_spacing = 1.15
    if not is_teacher:
        r = p_info.add_run("Họ, tên học sinh: ................................................................ Lớp: ........................")
        set_run_style(r, font_size=11)

    p_ntk = doc.add_paragraph()
    p_ntk.paragraph_format.space_before = Pt(0)
    p_ntk.paragraph_format.space_after = Pt(4)
    r = p_ntk.add_run("Cho NTK: H = 1; C = 12; O = 16; N = 14; Ag = 108; Na = 23; Br = 80; Cl = 35,5.")
    set_run_style(r, font_size=10.5, italic=True)

def add_section_title(doc, title_text, score_text=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"■ {title_text}")
    set_run_style(r, font_size=12, bold=True, color=COLOR_PRIMARY)
    if score_text:
        r_s = p.add_run(f" {score_text}")
        set_run_style(r_s, font_size=11, bold=True, italic=True)

def add_instruction_line(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_style(r, font_size=11, italic=True)

def add_question_prompt(doc, q_num, prompt_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_label = p.add_run(f"Câu {q_num}. ")
    set_run_style(r_label, font_size=12, bold=True)
    render_rich_text(p, prompt_text, base_size=12)
    return p

def add_choices_tabbed(doc, choices, correct_idx=None, is_teacher=False, layout_mode="4cols"):
    labels = ["A. ", "B. ", "C. ", "D. "]

    if layout_mode == "4cols":
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(4.8), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(13.6), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15

        for i in range(4):
            p.add_run('\t')
            is_corr = (is_teacher and i == correct_idx)
            r_lbl = p.add_run(labels[i])
            set_run_style(r_lbl, font_size=12, bold=True, color=COLOR_CORRECT if is_corr else None)
            render_rich_text(p, choices[i] + ("  " if i < 3 else ""), base_size=12,
                             base_bold=is_corr, base_color=COLOR_CORRECT if is_corr else None)

    elif layout_mode == "2cols":
        p1 = doc.add_paragraph()
        p1.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p1.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.line_spacing = 1.15

        for i in [0, 1]:
            p1.add_run('\t')
            is_corr = (is_teacher and i == correct_idx)
            r_lbl = p1.add_run(labels[i])
            set_run_style(r_lbl, font_size=12, bold=True, color=COLOR_CORRECT if is_corr else None)
            render_rich_text(p1, choices[i], base_size=12,
                             base_bold=is_corr, base_color=COLOR_CORRECT if is_corr else None)

        p2 = doc.add_paragraph()
        p2.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p2.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.line_spacing = 1.15

        for i in [2, 3]:
            p2.add_run('\t')
            is_corr = (is_teacher and i == correct_idx)
            r_lbl = p2.add_run(labels[i])
            set_run_style(r_lbl, font_size=12, bold=True, color=COLOR_CORRECT if is_corr else None)
            render_rich_text(p2, choices[i], base_size=12,
                             base_bold=is_corr, base_color=COLOR_CORRECT if is_corr else None)

    else: # 1col
        for i in range(4):
            p = doc.add_paragraph()
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1 if i < 3 else 2)
            p.paragraph_format.line_spacing = 1.15

            p.add_run('\t')
            is_corr = (is_teacher and i == correct_idx)
            r_lbl = p.add_run(labels[i])
            set_run_style(r_lbl, font_size=12, bold=True, color=COLOR_CORRECT if is_corr else None)
            render_rich_text(p, choices[i], base_size=12,
                             base_bold=is_corr, base_color=COLOR_CORRECT if is_corr else None)

def add_solution_block(doc, solution_text, choose_text=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.5)

    r_lbl = p.add_run("Lời giải: ")
    set_run_style(r_lbl, font_size=11.5, bold=True, color=COLOR_SOLUTION)

    render_rich_text(p, solution_text, base_size=11.5)

    if choose_text:
        r_c = p.add_run(f" ➔ {choose_text}")
        set_run_style(r_c, font_size=11.5, bold=True, color=COLOR_CORRECT)

def add_image_centered(doc, image_name, width_cm=6.0):
    img_path = os.path.join(CROP_DIR, image_name)
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        run.add_picture(img_path, width=Cm(width_cm))

def build_exam_document(is_teacher=False):
    doc = docx.Document()
    setup_header_footer(doc)
    add_exam_header_block(doc, is_teacher=is_teacher)

    # ==========================
    # PHẦN 1
    # ==========================
    add_section_title(doc, "A. PHẦN TRẮC NGHIỆM: [7,0 điểm]")
    add_section_title(doc, "PHẦN 1. Câu trắc nghiệm 4 phương án", "[3,0 điểm]")
    add_instruction_line(doc, "Thí sinh trả lời từ câu 1 đến câu 12. Mỗi câu hỏi thí sinh chỉ chọn một phương án (0,25 điểm/câu).")

    # Câu 1
    add_question_prompt(doc, 1, "Cho x mol phenol (C~6~H~5~OH) tác dụng với Na dư, thấy thoát ra 0,1 mol khí H~2~. Giá trị của x là")
    add_choices_tabbed(doc, ["0,1", "0,05", "0,15", "0,2"], correct_idx=3, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "Phương trình phản ứng: C~6~H~5~OH + Na -> C~6~H~5~ONa + 1/2 H~2~ ↑.\nTa có: n~phenol~ = 2 · n~H2~ = 2 × 0,1 = 0,2 mol. Do đó x = 0,2.", "Chọn D.")

    # Câu 2
    add_question_prompt(doc, 2, "Cho các chất sau: methane, ethylene, acetylene, benzene, toluene và naphthalene. Số chất ở thể lỏng trong điều kiện thường là")
    add_choices_tabbed(doc, ["1", "2", "3", "4"], correct_idx=1, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "Trong điều kiện thường:\n- Methane (CH~4~), ethylene (C~2~H~4~), acetylene (C~2~H~2~) là các chất khí.\n- Benzene (C~6~H~6~) và toluene (C~6~H~5~CH~3~) là các chất lỏng.\n- Naphthalene (C~10~H~8~) là chất rắn hình phiến màu trắng.\nVậy có 2 chất ở thể lỏng (benzene, toluene).", "Chọn B.")

    # Câu 3
    add_question_prompt(doc, 3, "Aldehyde nào sau đây là đồng đẳng của CH~3~CHO?")
    add_choices_tabbed(doc, ["CH~2~=CH-CHO", "C~6~H~5~-OH", "CH≡C-CHO", "C~2~H~5~CHO"], correct_idx=3, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "CH~3~CHO thuộc dãy đồng đẳng aldehyde no, đơn chức, mạch hở (công thức chung C~n~H~2n~O, n ≥ 1). C~2~H~5~CHO (propanal) cũng là aldehyde no, đơn chức, mạch hở thuộc cùng dãy đồng đẳng.", "Chọn D.")

    # Câu 4
    add_question_prompt(doc, 4, "Số hợp chất hữu cơ có công thức phân tử C~3~H~8~O phản ứng được với Na là")
    add_choices_tabbed(doc, ["1", "2", "3", "4"], correct_idx=1, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "Hợp chất C~3~H~8~O tác dụng được với Na phải có nhóm chức alcohol (-OH). Có 2 đồng phân alcohol: propan-1-ol (CH~3~-CH~2~-CH~2~-OH) và propan-2-ol (CH~3~-CH(OH)-CH~3~).", "Chọn B.")

    # Câu 5
    add_question_prompt(doc, 5, "Khi nhỏ từ từ dung dịch bromine vào ống nghiệm chứa dung dịch phenol, hiện tượng quan sát được trong ống nghiệm là")
    add_choices_tabbed(doc, [
        "không xảy ra hiện tượng gì",
        "dung dịch brom mất màu và xuất hiện kết tủa trắng",
        "xuất hiện kết tủa vàng",
        "dung dịch trong suốt"
    ], correct_idx=1, is_teacher=is_teacher, layout_mode="2cols")
    if is_teacher:
        add_solution_block(doc, "Phenol tác dụng với dung dịch nước bromine theo phản ứng thế vào nhân thơm tạo kết tủa trắng 2,4,6-tribromophenol: C~6~H~5~OH + 3Br~2~ -> C~6~H~2~Br~3~OH ↓ trắng + 3HBr. Do đó dung dịch bromine bị mất màu và xuất hiện kết tủa trắng.", "Chọn B.")

    # Câu 6
    add_question_prompt(doc, 6, "Tên gọi theo danh pháp thay thế của dẫn xuất halogen có công thức khung phân tử sau là:")
    add_image_centered(doc, "fig_bromobutane.png", width_cm=4.5)
    add_choices_tabbed(doc, ["1-bromobutane", "1-bromopentane", "4-bromobutane", "2-bromobutane"], correct_idx=0, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "Mạch chính gồm 4 nguyên tử carbon (butane), đánh số từ đầu gần nhóm thế bromine nhất: C~1~-C~2~-C~3~-C~4~ -> 1-bromobutane.", "Chọn A.")

    # Câu 7
    add_question_prompt(doc, 7, "Catechin là một chất chống oxi hoá mạnh, ức chế hoạt động của các gốc tự do nên có khả năng phòng chống bệnh ung thư, nhồi máu cơ tim. Trong lá chè tươi, catechin chiếm khoảng 25 - 35% tổng trọng lượng khô. Ngoài ra, catechin còn có trong táo, lê, nho,... Công thức cấu tạo của catechin cho như hình bên dưới:")
    add_image_centered(doc, "fig_catechin.png", width_cm=7.0)
    p_sub7 = doc.add_paragraph()
    p_sub7.paragraph_format.space_before = Pt(2)
    p_sub7.paragraph_format.space_after = Pt(2)
    r7 = p_sub7.add_run("Phát biểu nào sau đây là KHÔNG đúng?")
    set_run_style(r7, font_size=12, bold=True)
    add_choices_tabbed(doc, [
        "Phân tử catechin có 5 nhóm -OH phenol",
        "Catechin thuộc loại hợp chất thơm",
        "Catechin phản ứng được với dung dịch NaOH",
        "Công thức phân tử của catechin là C~15~H~14~O~6~"
    ], correct_idx=0, is_teacher=is_teacher, layout_mode="1col")
    if is_teacher:
        add_solution_block(doc, "Phân tích cấu trúc của catechin:\n- Có 4 nhóm -OH liên kết trực tiếp với vòng benzen (nhóm -OH phenol).\n- Có 1 nhóm -OH liên kết với nguyên tử carbon no của vòng chứa dị tố O (nhóm -OH alcohol).\nVậy phát biểu 'có 5 nhóm -OH phenol' là SAI.", "Chọn A.")

    # Câu 8
    add_question_prompt(doc, 8, "Trong các chất sau đây: CH~3~CH~2~OH, CH~3~CHO, CH~3~COOH, CH~3~CH~2~CH~2~CH~3~. Chất nào có nhiệt độ sôi cao nhất?")
    add_choices_tabbed(doc, ["CH~3~CHO", "CH~3~CH~2~OH", "CH~3~COOH", "CH~3~CH~2~CH~2~CH~3~"], correct_idx=2, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "Axit carboxylic (CH~3~COOH) có liên kết hydrogen liên phân tử bền hơn alcohol (CH~3~CH~2~OH) do liên kết O-H phân cực mạnh hơn và tạo dạng dimer bền. Thứ tự nhiệt độ sôi: CH~3~COOH (118 °C) > CH~3~CH~2~OH (78,3 °C) > CH~3~CHO (20,2 °C) > C~4~H~10~ (-0,5 °C).", "Chọn C.")

    # Câu 9
    add_question_prompt(doc, 9, "Để phân biệt styrene và phenylacetylene có thể dùng chất nào sau đây?")
    add_choices_tabbed(doc, ["Khí oxygen dư", "Nước bromine", "Dung dịch KMnO~4~", "Dung dịch AgNO~3~ trong NH~3~"], correct_idx=3, is_teacher=is_teacher, layout_mode="2cols")
    if is_teacher:
        add_solution_block(doc, "Phenylacetylene (C~6~H~5~-C≡CH) có liên kết ba đầu mạch nên tác dụng với dung dịch AgNO~3~/NH~3~ tạo kết tủa vàng nhạt C~6~H~5~-C≡CAg ↓. Styrene (C~6~H~5~-CH=CH~2~) không có liên kết ba đầu mạch nên không phản ứng.", "Chọn D.")

    # Câu 10
    add_question_prompt(doc, 10, "Alkyne là những hydrocarbon mạch hở, chỉ chứa các liên kết đơn và một liên kết ba C≡C trong phân tử, có công thức chung là")
    add_choices_tabbed(doc, ["C~n~H~2n-2~ (n ≥ 2)", "C~n~H~2n+2~ (n ≥ 1)", "C~n~H~2n~ (n ≥ 2)", "C~n~H~2n-6~ (n ≥ 6)"], correct_idx=0, is_teacher=is_teacher, layout_mode="2cols")
    if is_teacher:
        add_solution_block(doc, "Alkyne mạch hở, chứa 1 liên kết ba C≡C (độ bất bão hoà k = 2) nên có công thức phân tử chung là C~n~H~2n-2~ với n ≥ 2.", "Chọn A.")

    # Câu 11
    add_question_prompt(doc, 11, "Geraniol có trong tinh dầu hoa hồng (công thức cấu tạo thu gọn như hình bên dưới) được sử dụng phổ biến trong công nghiệp hương liệu, thực phẩm,... vì có mùi thơm đặc trưng:")
    add_image_centered(doc, "fig_geraniol.png", width_cm=6.5)
    p_sub11 = doc.add_paragraph()
    p_sub11.paragraph_format.space_before = Pt(2)
    p_sub11.paragraph_format.space_after = Pt(2)
    r11 = p_sub11.add_run("Geraniol thuộc loại hợp chất hữu cơ nào sau đây?")
    set_run_style(r11, font_size=12)
    add_choices_tabbed(doc, ["Hydrocarbon", "Carboxylic acid", "Alcohol", "Aldehyde"], correct_idx=2, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "Geraniol có nhóm hydroxy (-OH) liên kết trực tiếp với nguyên tử carbon no mạch hở, do đó geraniol thuộc loại alcohol (alcohol không no, mạch hở, bậc I).", "Chọn C.")

    # Câu 12
    add_question_prompt(doc, 12, "Dẫn xuất halogen nào sau đây khi tác dụng với dung dịch NaOH đun nóng không tạo thành alcohol?")
    add_choices_tabbed(doc, ["C~6~H~5~Cl", "CH~3~CH(Br)CH~3~", "C~6~H~5~CH~2~Br", "C~2~H~5~Cl"], correct_idx=0, is_teacher=is_teacher, layout_mode="4cols")
    if is_teacher:
        add_solution_block(doc, "Chlorobenzene (C~6~H~5~Cl) có nguyên tử Cl liên kết trực tiếp với vòng benzen, liên kết C-Cl rất bền do hiệu ứng liên hợp p-π, không bị thuỷ phân trong dung dịch kiềm đun nóng ở điều kiện thường (chỉ phản ứng ở 300 °C, 200 atm tạo C~6~H~5~ONa).", "Chọn A.")

    # ==========================
    # PHẦN 2
    # ==========================
    add_section_title(doc, "PHẦN 2. Câu trắc nghiệm đúng sai", "[2,0 điểm]")
    add_instruction_line(doc, "Thí sinh trả lời từ câu 1 đến câu 2. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng (Đ) hoặc sai (S). Điểm tối đa mỗi câu là 1,0 điểm.")

    # Câu 1 TF
    add_question_prompt(doc, 1, "Nhắc đến ethanol nhiều người thường nghĩ ngay đến đồ uống có cồn, tuy nhiên đây cũng là thành phần quan trọng được sử dụng trong y tế, là dung môi phổ biến cho nhiều ngành công nghiệp...")
    tf1_items = [
        ("a) ", "Ethanol còn gọi là ethyl alcohol.", True,
         "Tên thông thường của ethanol là ethyl alcohol."),
        ("b) ", "Ethanol là chất lỏng, dễ bay hơi, không mùi.", False,
         "Ethanol là chất lỏng, dễ bay hơi, nhưng có mùi thơm nhẹ đặc trưng và vị cay nồng (không phải không mùi)."),
        ("c) ", "Ethanol được tạo thành từ phản ứng thuỷ phân bromoethane bằng dung dịch NaOH có đun nóng.", True,
         "Phương trình phản ứng: C~2~H~5~Br + NaOH -(t^o^)-> C~2~H~5~OH + NaBr."),
        ("d) ", "Cho 4,6 gam ethanol tác dụng với Na dư, thể tích khí H~2~ thu được (đkc) là 1,2395 L.", True,
         "n(C~2~H~5~OH) = 4,6 / 46 = 0,1 mol -> n(H~2~) = 0,1 / 2 = 0,05 mol -> V(H~2~) ở đkc = 0,05 × 24,79 = 1,2395 L.")
    ]
    for lbl, text, is_true, expl in tf1_items:
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        p.add_run('\t')
        r_lbl = p.add_run(lbl)
        set_run_style(r_lbl, font_size=12, bold=True)
        render_rich_text(p, text, base_size=12)
        if is_teacher:
            tag = " [ĐÚNG]" if is_true else " [SAI]"
            r_tag = p.add_run(tag)
            set_run_style(r_tag, font_size=12, bold=True, color=COLOR_CORRECT)
            p_ex = doc.add_paragraph()
            p_ex.paragraph_format.left_indent = Cm(1.0)
            p_ex.paragraph_format.space_before = Pt(0)
            p_ex.paragraph_format.space_after = Pt(2)
            r_ex = p_ex.add_run("➔ Giải thích: ")
            set_run_style(r_ex, font_size=11, bold=True, italic=True, color=COLOR_SOLUTION)
            render_rich_text(p_ex, expl, base_size=11, base_italic=True, base_color=COLOR_SOLUTION)

    # Câu 2 TF
    add_question_prompt(doc, 2, "Hợp chất carbonyl đơn giản nhất là aldehyde và ketone đơn chức. Chúng có nhiều ứng dụng trong ngành công nghiệp hoá chất cũng như trong thiên nhiên.")
    tf2_items = [
        ("a) ", "Cho 0,44 gam ethanal vào dung dịch AgNO~3~ trong NH~3~ dư, đến khi phản ứng hoàn toàn thì thu được 10,8 gam Ag.", False,
         "n(CH~3~CHO) = 0,44 / 44 = 0,01 mol -> n(Ag) = 2 × 0,01 = 0,02 mol -> m(Ag) = 0,02 × 108 = 2,16 gam (không phải 10,8 gam)."),
        ("b) ", "Trong tinh dầu thảo mộc có chứa những aldehyde, dùng dung dịch AgNO~3~ trong NH~3~ để nhận biết thành phần aldehyde trong tinh dầu.", True,
         "Phản ứng tráng bạc với thuốc thử Tollens là phản ứng đặc trưng để nhận biết nhóm chức aldehyde tạo kết tủa Ag sáng bóng."),
        ("c) ", "Tên thay thế của (CH~3~)~2~CHCH~2~CHO là 2-methylbutanal.", False,
         "Mạch chính đánh số từ nhóm -CHO là C1: CH~3~-CH(CH~3~)-CH~2~-CHO -> C1 ở -CHO, C2 là CH~2~, C3 mang nhóm methyl -CH~3~. Tên thay thế đúng là 3-methylbutanal."),
        ("d) ", "Ethanal và propanone không thuộc loại hợp chất carbonyl.", False,
         "Hợp chất carbonyl có chứa nhóm >C=O trong phân tử. Ethanal là aldehyde (CH~3~CHO) và propanone là ketone (CH~3~COCH~3~), cả hai đều là hợp chất carbonyl.")
    ]
    for lbl, text, is_true, expl in tf2_items:
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        p.add_run('\t')
        r_lbl = p.add_run(lbl)
        set_run_style(r_lbl, font_size=12, bold=True)
        render_rich_text(p, text, base_size=12)
        if is_teacher:
            tag = " [ĐÚNG]" if is_true else " [SAI]"
            r_tag = p.add_run(tag)
            set_run_style(r_tag, font_size=12, bold=True, color=COLOR_CORRECT)
            p_ex = doc.add_paragraph()
            p_ex.paragraph_format.left_indent = Cm(1.0)
            p_ex.paragraph_format.space_before = Pt(0)
            p_ex.paragraph_format.space_after = Pt(2)
            r_ex = p_ex.add_run("➔ Giải thích: ")
            set_run_style(r_ex, font_size=11, bold=True, italic=True, color=COLOR_SOLUTION)
            render_rich_text(p_ex, expl, base_size=11, base_italic=True, base_color=COLOR_SOLUTION)

    # ==========================
    # PHẦN 3
    # ==========================
    add_section_title(doc, "PHẦN 3. Câu trắc nghiệm trả lời ngắn", "[2,0 điểm]")
    add_instruction_line(doc, "Thí sinh trả lời từ câu 1 đến câu 4 (0,5 điểm/câu).")

    # Câu 1 ShortAns
    add_question_prompt(doc, 1, "Ethene và acetylene là những hydrocarbon không no đơn giản nhất và có nhiều ứng dụng quan trọng. Một ứng dụng quan trọng của acetylene là làm nhiên liệu trong đèn xì oxygen - acetylene. Khi đèn hoạt động, hai khí này được trộn vào nhau để thực hiện phản ứng đốt cháy theo sơ đồ C~2~H~2~ + O~2~ -(t^o^)-> 2CO~2~ + H~2~O. Đốt cháy hoàn toàn V lít (đkc) khí acetylene thu được 7,2 gam H~2~O. Nếu cho tất cả sản phẩm cháy hấp thụ hết vào bình đựng nước vôi trong dư thì khối lượng bình tăng m gam. Tính giá trị của m?")
    if not is_teacher:
        add_image_centered(doc, "fig_shortans.png", width_cm=4.0)
    else:
        add_solution_block(doc, "n(H~2~O) = 7,2 / 18 = 0,4 mol -> n(C~2~H~2~) = 0,4 mol.\nn(CO~2~) = 2 · n(C~2~H~2~) = 0,8 mol -> m(CO~2~) = 0,8 × 44 = 35,2 gam.\nKhi hấp thụ vào dung dịch Ca(OH)~2~ dư, khối lượng bình tăng chính bằng tổng khối lượng của CO~2~ và H~2~O:\nm = m(CO~2~) + m(H~2~O) = 35,2 + 7,2 = 42,4 gam.", "Đáp án: 42,4")

    # Câu 2 ShortAns
    add_question_prompt(doc, 2, "Phenol được dùng để sản xuất phẩm nhuộm, nhựa phenol-formaldehyde, thuốc nổ (2,4,6-trinitrophenol), chất diệt cỏ 2,4-D, chất diệt nấm mốc (các đồng phân của nitrophenol),... Do có tính diệt khuẩn nên phenol được dùng làm chất khử trùng, tẩy uế. Thuốc xịt chloraseptic chứa 1,4% phenol được dùng làm thuốc chữa đau họng. Cho các phát biểu sau:\n"
                                "   a) Phenol tan vô hạn trong nước lạnh ở điều kiện thường.\n"
                                "   b) Nhiệt độ nóng chảy của phenol cao hơn ethanol.\n"
                                "   c) Phenol có khả năng tác dụng với dung dịch bromine tạo kết tủa trắng.\n"
                                "   d) Phenol dùng để sản xuất phẩm nhuộm, chất diệt nấm mốc, thuốc nổ TNT.\n"
                                "Có bao nhiêu phát biểu đúng?")
    if not is_teacher:
        add_image_centered(doc, "fig_shortans.png", width_cm=4.0)
    else:
        add_solution_block(doc, "Phân tích các phát biểu:\n- a) Sai: Phenol ít tan trong nước lạnh (8,3 g/100 g nước ở 20 °C), tan vô hạn khi trên 66 °C.\n- b) Đúng: Phenol ở thể rắn ở đk thường (nhiệt độ nóng chảy 43 °C), trong khi ethanol là chất lỏng (nóng chảy ở -114,1 °C).\n- c) Đúng: Phenol tạo kết tủa trắng 2,4,6-tribromophenol với nước bromine.\n- d) Sai: Thuốc nổ TNT là trinitrotoluene (sản xuất từ toluene), phenol dùng để sản xuất axit picric (2,4,6-trinitrophenol).\nVậy có 2 phát biểu đúng (b, c).", "Đáp án: 2")

    # Câu 3 ShortAns
    add_question_prompt(doc, 3, "Ngày nay, nhu cầu về đồ gỗ nội thất ngày càng nhiều song nguồn gỗ tự nhiên không còn dồi dào nên việc chuyển sang sử dụng gỗ công nghiệp đang là xu hướng của nhiều nước trên thế giới. Việc sử dụng gỗ công nghiệp góp phần bảo vệ rừng, bảo vệ môi trường. Quy trình sản xuất gỗ công nghiệp là nghiền các cây gỗ trồng ngắn ngày như keo, bạch đàn, cao su,..., sau đó sử dụng keo để kết dính và ép để tạo độ dày ván gỗ. Keo được sử dụng trong gỗ công nghiệp thường chứa dư lượng formaldehyde, là một hoá chất độc hại đối với sức khoẻ con người. Tại các nước phát triển như ở châu Âu và Mỹ, dư lượng formaldehyde được kiểm soát rất nghiêm ngặt. Châu Âu quy định tiêu chuẩn dư lượng formaldehyde trong gỗ công nghiệp là 120 µg·m^-3^. Cơ quan kiểm định lấy 300 gam gỗ trong một lô gỗ của một doanh nghiệp Việt Nam xuất khẩu sang châu Âu và kiểm tra bằng phương pháp sắc kí thấy chứa 0,03 µg formaldehyde. Biết khối lượng riêng của loại gỗ này là 800 kg·m^-3^. Hàm lượng formaldehyde có trong 800 kg (hay 1 m^3^) gỗ là bao nhiêu µg?")
    if not is_teacher:
        add_image_centered(doc, "fig_shortans.png", width_cm=4.0)
    else:
        add_solution_block(doc, "Trong 300 gam gỗ có chứa 0,03 µg formaldehyde.\nKhối lượng 800 kg = 800.000 gam.\nHàm lượng formaldehyde trong 800 kg gỗ (tương đương 1 m^3^ gỗ) là:\nm = 0,03 × (800.000 / 300) = 80 µg.", "Đáp án: 80")

    # Câu 4 ShortAns
    add_question_prompt(doc, 4, "Cho các chất sau: H-CHO; H-COOH; CH~3~-COOH; C~2~H~5~-OH; CH~2~=CH-COOH; C~6~H~5~-COOH; HOOC-COOH; C~6~H~6~; HOOC-CH~2~-COOH; C~2~H~5~COOH. Có bao nhiêu hợp chất carboxylic acid trong các chất trên?")
    if not is_teacher:
        add_image_centered(doc, "fig_shortans.png", width_cm=4.0)
    else:
        add_solution_block(doc, "Carboxylic acid là các hợp chất có nhóm carboxyl (-COOH). Các chất là carboxylic acid gồm:\n1) H-COOH (methanoic acid)\n2) CH~3~-COOH (ethanoic acid)\n3) CH~2~=CH-COOH (acrylic acid)\n4) C~6~H~5~-COOH (benzoic acid)\n5) HOOC-COOH (oxalic acid)\n6) HOOC-CH~2~-COOH (malonic acid)\n7) C~2~H~5~COOH (propanoic acid)\nTổng cộng có 7 chất.", "Đáp án: 7")

    # ==========================
    # PHẦN 4: TỰ LUẬN
    # ==========================
    add_section_title(doc, "B. PHẦN TỰ LUẬN: [3,0 điểm]")
    add_instruction_line(doc, "Thí sinh trình bày bài làm chi tiết cho các câu hỏi sau (1,0 điểm/câu).")

    # Tự luận 1
    p_tl1 = doc.add_paragraph()
    p_tl1.paragraph_format.space_before = Pt(4)
    p_tl1.paragraph_format.space_after = Pt(2)
    p_tl1.paragraph_format.line_spacing = 1.15
    r_tl1 = p_tl1.add_run("Câu 1 (1,0 điểm). ")
    set_run_style(r_tl1, font_size=12, bold=True)
    render_rich_text(p_tl1, "Các nhà hoá học đã tìm ra một số dẫn xuất halogen không chứa chlorine như: CBr~2~F~2~, CF~3~-CHF~2~, CH~2~=CF-CF~3~, CF~3~CH~2~CF~2~CH~3~,... đang được sử dụng trong công nghiệp nhiệt lạnh, vì sự phân huỷ các hợp chất này nhanh chóng sau khi phát tán vào không khí nên ảnh hưởng rất ít đến tầng ozone hay sự ấm lên toàn cầu thấp. Gọi tên theo danh pháp thay thế các hợp chất đó.", base_size=12)
    if is_teacher:
        add_solution_block(doc,
            "Tên theo danh pháp thay thế của các hợp chất:\n"
            "• CBr~2~F~2~: dibromodifluoromethane [0,25 điểm]\n"
            "• CF~3~-CHF~2~: 1,1,1,2,2-pentafluoroethane [0,25 điểm]\n"
            "• CH~2~=CF-CF~3~: 2,3,3,3-tetrafluoroprop-1-ene [0,25 điểm]\n"
            "• CF~3~CH~2~CF~2~CH~3~: 1,1,1,3,3-pentafluorobutane [0,25 điểm]")
    else:
        p_blank = doc.add_paragraph()
        p_blank.paragraph_format.space_before = Pt(6)
        p_blank.paragraph_format.space_after = Pt(12)
        r_bl = p_blank.add_run("Bài làm:\n.......................................................................................................................................................................................\n.......................................................................................................................................................................................")
        set_run_style(r_bl, font_size=11, italic=True)

    # Tự luận 2
    p_tl2 = doc.add_paragraph()
    p_tl2.paragraph_format.space_before = Pt(4)
    p_tl2.paragraph_format.space_after = Pt(2)
    p_tl2.paragraph_format.line_spacing = 1.15
    r_tl2 = p_tl2.add_run("Câu 2 (1,0 điểm). ")
    set_run_style(r_tl2, font_size=12, bold=True)
    render_rich_text(p_tl2, "Trong công nghiệp chế biến đường từ mía sẽ tạo ra sản phẩm phụ, gọi là rỉ đường hay rỉ mật, sử dụng rỉ đường để lên men tạo ra ethanol trong điều kiện thích hợp, hiệu suất cả quá trình là 80%. Tính khối lượng ethanol thu được từ 1,5 tấn rỉ đường mía theo 2 phương trình:\n"
                            "   C~12~H~22~O~11~ + H~2~O -> C~6~H~12~O~6~ + C~6~H~12~O~6~\n"
                            "   C~6~H~12~O~6~ -> 2C~2~H~5~OH + 2CO~2~", base_size=12)
    if is_teacher:
        add_solution_block(doc,
            "Sơ đồ chuyển hoá tổng hợp:\n"
            "   C~12~H~22~O~11~ (342 g/mol) -> 2 C~6~H~12~O~6~ -> 4 C~2~H~5~OH (4 × 46 = 184 g/mol) [0,25 điểm]\n\n"
            "Khối lượng ethanol thu được theo lý thuyết từ 1,5 tấn rỉ đường:\n"
            "   m(lt) = 1,5 × (184 / 342) ≈ 0,8070 tấn = 807,0 kg. [0,50 điểm]\n\n"
            "Do hiệu suất toàn bộ quá trình đạt 80%, khối lượng ethanol thực tế thu được là:\n"
            "   m(tt) = 0,8070 × 80% = 0,6456 tấn = 645,6 kg. [0,25 điểm]\n"
            "Đáp số: 645,6 kg (hoặc 0,6456 tấn) ethanol.")
    else:
        p_blank = doc.add_paragraph()
        p_blank.paragraph_format.space_before = Pt(6)
        p_blank.paragraph_format.space_after = Pt(12)
        r_bl = p_blank.add_run("Bài làm:\n.......................................................................................................................................................................................\n.......................................................................................................................................................................................")
        set_run_style(r_bl, font_size=11, italic=True)

    # Tự luận 3
    p_tl3 = doc.add_paragraph()
    p_tl3.paragraph_format.space_before = Pt(4)
    p_tl3.paragraph_format.space_after = Pt(2)
    p_tl3.paragraph_format.line_spacing = 1.15
    r_tl3 = p_tl3.add_run("Câu 3 (1,0 điểm). ")
    set_run_style(r_tl3, font_size=12, bold=True)
    render_rich_text(p_tl3, "Thị trường tiêu thụ phenol trên toàn thế giới khoảng 11,67 triệu tấn trong năm 2022. Phenol được sử dụng để sản xuất nhiều loại hoá chất như bisphenol A, nhựa phenol-formaldehyde, picric acid và các chất khác. Khoảng 90% lượng phenol được sản xuất từ cumene. Khối lượng cumene cần dùng để sản xuất phenol cho năm 2022 là bao nhiêu?", base_size=12)
    if is_teacher:
        add_solution_block(doc,
            "Lượng phenol được sản xuất từ cumene trong năm 2022:\n"
            "   m(phenol) = 11,67 × 90% = 10,503 triệu tấn. [0,25 điểm]\n\n"
            "Phương trình phản ứng điều chế phenol từ cumene:\n"
            "   C~6~H~5~-CH(CH~3~)~2~ + O~2~ -(H~2~SO~4~)-> C~6~H~5~OH + CH~3~COCH~3~\n"
            "Tỉ lệ mol giữa cumene và phenol là 1 : 1.\n"
            "Khối lượng mol của cumene (C~9~H~12~) là 120 g/mol; của phenol (C~6~H~5~OH) là 94 g/mol. [0,25 điểm]\n\n"
            "Khối lượng cumene cần dùng là:\n"
            "   m(cumene) = 10,503 × (120 / 94) ≈ 13,408 triệu tấn. [0,50 điểm]\n"
            "Đáp số: Khoảng 13,41 triệu tấn cumene.")
    else:
        p_blank = doc.add_paragraph()
        p_blank.paragraph_format.space_before = Pt(6)
        p_blank.paragraph_format.space_after = Pt(12)
        r_bl = p_blank.add_run("Bài làm:\n.......................................................................................................................................................................................\n.......................................................................................................................................................................................")
        set_run_style(r_bl, font_size=11, italic=True)

    # ==========================
    # BẢNG ĐÁP ÁN (DÀNH CHO GIÁO VIÊN)
    # ==========================
    if is_teacher:
        p_div = doc.add_paragraph()
        p_div.paragraph_format.space_before = Pt(12)
        p_div.paragraph_format.space_after = Pt(4)
        p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_div = p_div.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
        set_run_style(r_div, font_size=11, color=COLOR_PRIMARY)

        p_tb_title = doc.add_paragraph()
        p_tb_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_tb_title.paragraph_format.space_before = Pt(4)
        p_tb_title.paragraph_format.space_after = Pt(4)
        r_tbt = p_tb_title.add_run("BẢNG TỔNG HỢP ĐÁP ÁN ĐỀ THI MÃ 209")
        set_run_style(r_tbt, font_size=13, bold=True, color=COLOR_PRIMARY)

        # Bảng Phần 1
        p_p1 = doc.add_paragraph()
        p_p1.paragraph_format.space_before = Pt(4)
        p_p1.paragraph_format.space_after = Pt(2)
        r_p1 = p_p1.add_run("1. Đáp án Phần 1 (0,25 điểm/câu):")
        set_run_style(r_p1, font_size=11.5, bold=True)

        t1 = doc.add_table(rows=2, cols=12)
        t1.alignment = WD_TABLE_ALIGNMENT.CENTER
        ans_p1 = ["D", "B", "D", "B", "B", "A", "A", "C", "D", "A", "C", "A"]
        for c in range(12):
            cell_h = t1.cell(0, c)
            cell_h.text = f"{c+1}"
            cell_h.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_style(cell_h.paragraphs[0].runs[0], font_size=11, bold=True)
            cell_a = t1.cell(1, c)
            cell_a.text = ans_p1[c]
            cell_a.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_style(cell_a.paragraphs[0].runs[0], font_size=11.5, bold=True, color=COLOR_CORRECT)
        t1.style = 'Table Grid'

        # Bảng Phần 2
        p_p2 = doc.add_paragraph()
        p_p2.paragraph_format.space_before = Pt(6)
        p_p2.paragraph_format.space_after = Pt(2)
        r_p2 = p_p2.add_run("2. Đáp án Phần 2 (1,0 điểm/câu):")
        set_run_style(r_p2, font_size=11.5, bold=True)

        t2 = doc.add_table(rows=3, cols=5)
        t2.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers2 = ["Câu", "Ý a)", "Ý b)", "Ý c)", "Ý d)"]
        for col_idx, h in enumerate(headers2):
            cell = t2.cell(0, col_idx)
            cell.text = h
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_style(cell.paragraphs[0].runs[0], font_size=11, bold=True)

        ans_p2 = [
            ["Câu 1", "Đ", "S", "Đ", "Đ"],
            ["Câu 2", "S", "Đ", "S", "S"]
        ]
        for row_idx, row_vals in enumerate(ans_p2):
            for col_idx, val in enumerate(row_vals):
                cell = t2.cell(row_idx + 1, col_idx)
                cell.text = val
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_run_style(cell.paragraphs[0].runs[0], font_size=11.5, bold=True,
                              color=COLOR_CORRECT if col_idx > 0 else None)
        t2.style = 'Table Grid'

        # Bảng Phần 3
        p_p3 = doc.add_paragraph()
        p_p3.paragraph_format.space_before = Pt(6)
        p_p3.paragraph_format.space_after = Pt(2)
        r_p3 = p_p3.add_run("3. Đáp án Phần 3 (0,5 điểm/câu):")
        set_run_style(r_p3, font_size=11.5, bold=True)

        t3 = doc.add_table(rows=2, cols=4)
        t3.alignment = WD_TABLE_ALIGNMENT.CENTER
        ans_p3 = ["42,4", "2", "80", "7"]
        for c in range(4):
            cell_h = t3.cell(0, c)
            cell_h.text = f"Câu {c+1}"
            cell_h.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_style(cell_h.paragraphs[0].runs[0], font_size=11, bold=True)

            cell_a = t3.cell(1, c)
            cell_a.text = ans_p3[c]
            cell_a.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_style(cell_a.paragraphs[0].runs[0], font_size=11.5, bold=True, color=COLOR_CORRECT)
        t3.style = 'Table Grid'

    # Footer note
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(10)
    p_end.paragraph_format.space_after = Pt(2)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("---------- HẾT ----------")
    set_run_style(r_end, font_size=11, bold=True)

    if not is_teacher:
        p_note = doc.add_paragraph()
        p_note.paragraph_format.space_before = Pt(2)
        p_note.paragraph_format.space_after = Pt(0)
        r_n = p_note.add_run("- Thí sinh không được sử dụng tài liệu.\n- Cán bộ coi kiểm tra không giải thích gì thêm.")
        set_run_style(r_n, font_size=10.5, italic=True)

    out_name = "De_Kiem_Tra_Hoa_11_Ma209_GiaoVien.docx" if is_teacher else "De_Kiem_Tra_Hoa_11_Ma209.docx"
    out_path = os.path.join(OUT_DIR, out_name)
    doc.save(out_path)
    print(f"Successfully generated: {out_path}")

if __name__ == "__main__":
    build_exam_document(is_teacher=False)
    build_exam_document(is_teacher=True)
