# -*- coding: utf-8 -*-
"""
Script tạo Đề Kiểm Tra Giữa Kì 1 Hóa Học 12 - Trường THPT Nguyễn Hiền (TP. Hồ Chí Minh)
Năm học 2026 - 2027. Mã đề: Gốc.
Chuẩn 100% FORM ĐỀ KIỂM TRA (Thầy TRẦN VĂN THẠNH)
Sinh 2 phiên bản:
1. Bản Học sinh: De_Kiem_Tra_GK1_Hoa_12_NguyenHien_2026_HocSinh.docx
2. Bản Giáo viên: De_Kiem_Tra_GK1_Hoa_12_NguyenHien_2026_GiaoVien.docx
"""

import os
import sys

if sys.stdout.encoding.lower() != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

sys.path.append(os.path.abspath(r"C:\Antigravity_Thanh\scripts"))
from form_de_kiem_tra import (
    setup_header_footer,
    build_title_block,
    add_section_header,
    render_question_prompt,
    render_options_no_table,
    render_statement,
    add_solution_no_table,
    add_student_writing_lines,
    render_rich_text,
    set_run_style,
    COLOR_PRIMARY,
    COLOR_CORRECT,
    COLOR_SOLUTION
)

IMG_DIR = r"C:\Antigravity_Thanh\San_Pham\images_gk1_nguyenhien"
ANILINE_HD = r"C:\Antigravity_Thanh\San_Pham\So_Do_Dieu_Che_Aniline\so_do_dieu_che_aniline.png"

def add_centered_image(doc, img_path, width_cm=8.0):
    """Chèn hình ảnh căn giữa văn bản"""
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(img_path, width=Cm(width_cm))

def build_exam_nguyenhien(is_teacher=False, output_path=None):
    doc = docx.Document()
    setup_header_footer(doc)

    # 1. Khối tiêu đề
    build_title_block(
        doc,
        school_line="SỞ GIÁO DỤC VÀ ĐÀO TẠO TP. HỒ CHÍ MINH\nTRƯỜNG T.H.P.T NGUYỄN HIỀN",
        exam_title="ĐỀ KIỂM TRA GIỮA KÌ 1. NĂM HỌC 2026 – 2027",
        subject_line="MÔN: HÓA HỌC - KHỐI 12",
        num_pages="04",
        duration="45 phút",
        exam_code="Gốc",
        atomic_masses="H = 1, C = 12, N = 14, O = 16, S = 32, Cl = 35,5, Na = 23, K = 39, Ca = 40, Br = 80",
        is_teacher=is_teacher
    )

    # =========================================================================
    # PHẦN I
    # =========================================================================
    add_section_header(
        doc,
        "PHẦN I (4,5 điểm). Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án.",
        subtitle="Mỗi câu trả lời đúng thí sinh được 0,25 điểm"
    )

    # Câu 1
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 1. Methyl acetate có công thức là")
    opts1 = [
        "A. HCOOCH~2~CH~3~.",
        "*B. CH~3~COOCH~3~.*" if is_teacher else "B. CH~3~COOCH~3~.",
        "C. CH~3~COOCH~2~CH~3~.",
        "D. CH~2~=CHCOOCH~3~."
    ]
    render_options_no_table(doc, opts1, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Methyl acetate là ester được tạo từ acetic acid (CH~3~COOH) và methyl alcohol (CH~3~OH), có công thức cấu tạo là CH~3~COOCH~3~.\n➔ Đáp án đúng: B.")

    # Câu 2
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 2. Thủy phân tristearin (C~17~H~35~COO)~3~C~3~H~5~ trong dung dịch NaOH, thu được muối ứng dụng làm xà phòng có công thức là")
    opts2 = [
        "A. C~3~H~5~COONa.",
        "*B. C~17~H~35~COONa.*" if is_teacher else "B. C~17~H~35~COONa.",
        "C. C~3~H~5~(OH)~3~.",
        "D. C~17~H~33~COONa."
    ]
    render_options_no_table(doc, opts2, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Phương trình phản ứng xà phòng hóa tristearin:\n(C~17~H~35~COO)~3~C~3~H~5~ + 3NaOH ➔ 3C~17~H~35~COONa + C~3~H~5~(OH)~3~.\nMuối sodium stearate (C~17~H~35~COONa) là thành phần chính có hoạt tính giặt rửa của xà phòng.\n➔ Đáp án đúng: B.")

    # Câu 3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 3. Nhỏ dung dịch methylamine dư vào dung dịch nào sau đây thấy xuất hiện kết tủa, sau đó kết tủa tan dần và tạo dung dịch phức chất có màu xanh?")
    opts3 = [
        "A. Dung dịch FeCl~3~.",
        "B. Dung dịch H~2~SO~4~.",
        "*C. Dung dịch CuSO~4~.*" if is_teacher else "C. Dung dịch CuSO~4~.",
        "D. Dung dịch HCl."
    ]
    render_options_no_table(doc, opts3, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "• Methylamine trong nước có tính base yếu:\n  CH~3~NH~2~ + H~2~O ⇌ CH~3~NH~3~^+ + OH^-.\n• Ban đầu xuất hiện kết tủa màu xanh lam của Cu(OH)~2~:\n  Cu^2+^ + 2OH^- ➔ Cu(OH)~2~↓ (xanh lam).\n• Khi methylamine dư, do có cặp electron tự do trên nguyên tử N, các phân tử methylamine đóng vai trò phối tử tạo phức tan tetraamine copper(II) có màu xanh lam đặc trưng:\n  Cu(OH)~2~ + 4CH~3~NH~2~ ➔ [Cu(CH~3~NH~2~)~4~](OH)~2~.\n➔ Đáp án đúng: C.")

    # Câu 4
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 4. Ester nào sau đây khi đun nóng với dung dịch NaOH không thu được sản phẩm alcohol?")
    opts4 = [
        "*A. CH~3~COOCH=CH-CH~3~.*" if is_teacher else "A. CH~3~COOCH=CH-CH~3~.",
        "B. CH~2~=CHCOOCH~2~-CH~3~.",
        "C. CH~3~COOCH~2~-CH~3~.",
        "D. CH~3~COOCH~2~-CH=CH~2~."
    ]
    render_options_no_table(doc, opts4, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "• CH~3~COOCH=CH-CH~3~ có nhóm -O- liên kết trực tiếp với nguyên tử carbon mang nối đôi C=C, sản phẩm trung gian enol CH~3~-CH=CH-OH không bền tự đồng phân hóa tạo aldehyde:\n  CH~3~COOCH=CH-CH~3~ + NaOH ➔ CH~3~COONa + CH~3~CH~2~CHO (propanal), không sinh ra alcohol.\n• Các ester B, C sinh ra C~2~H~5~OH; ester D sinh ra CH~2~=CH-CH~2~OH (allyl alcohol).\n➔ Đáp án đúng: A.")

    # Câu 5
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 5. Chất nào sau đây tác dụng với bromine thu được kết tủa màu trắng?")
    opts5 = [
        "A. glucose.",
        "B. ethyl amine.",
        "C. triolein.",
        "*D. aniline.*" if is_teacher else "D. aniline."
    ]
    render_options_no_table(doc, opts5, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Aniline (C~6~H~5~NH~2~) tác dụng với nước bromine ở điều kiện thường tạo kết tủa trắng 2,4,6-tribromoaniline do ảnh hưởng hoạt hóa mạnh của nhóm -NH~2~ đến vòng benzen:\nC~6~H~5~NH~2~ + 3Br~2~ ➔ C~6~H~2~Br~3~NH~2~↓ (trắng) + 3HBr.\n➔ Đáp án đúng: D.")

    # Câu 6
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 6. Nhiệt độ sôi của amine phụ thuộc vào khối lượng phân tử và khả năng tạo liên kết hydrogen liên phân tử. Amine nào sau đây có nhiệt độ sôi cao nhất?")
    opts6 = [
        "A. (CH~3~)~3~N.",
        "B. CH~3~NHCH~2~CH~3~.",
        "*C. CH~3~CH~2~CH~2~NH~2~.*" if is_teacher else "C. CH~3~CH~2~CH~2~NH~2~.",
        "D. CH~3~CH~2~NH~2~."
    ]
    render_options_no_table(doc, opts6, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "• So sánh 3 đồng phân có cùng CTPT C~3~H~9~N (M = 59):\n  - CH~3~CH~2~CH~2~NH~2~ (bậc 1) có 2 nguyên tử H trên nguyên tử N nên tạo liên kết hydrogen liên phân tử bền vững nhất ➔ nhiệt độ sôi cao nhất (~48 °C).\n  - CH~3~NHCH~2~CH~3~ (bậc 2) chỉ có 1 liên kết N-H nên liên kết hydrogen yếu hơn (~36 °C).\n  - (CH~3~)~3~N (bậc 3) không có liên kết N-H nên không thể tạo liên kết hydrogen liên phân tử (~3 °C).\n• CH~3~CH~2~NH~2~ có phân tử khối nhỏ hơn hẳn (M = 45, sôi ở 16,6 °C).\n➔ Đáp án đúng: C.")

    # Câu 7
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 7. Một mắt xích của cellulose có cấu trúc như hình bên:\nNitrate hoá cellulose bằng tác nhân HNO~3~ (xúc tác H~2~SO~4~) thấy tạo thành cellulose trinitrate. Các nguyên tử oxygen ở vị trí nào (trong mỗi mắt xích) đã tham gia phản ứng?")
    add_centered_image(doc, os.path.join(IMG_DIR, "cau7_cellulose.png"), width_cm=5.5)
    opts7 = [
        "*A. Vị trí (1), (3), (5).*" if is_teacher else "A. Vị trí (1), (3), (5).",
        "B. Vị trí (2), (3), (5).",
        "C. Vị trí (2), (3), (4).",
        "D. Vị trí (3), (4), (5)."
    ]
    render_options_no_table(doc, opts7, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "• Mỗi mắt xích cellulose chứa 3 nhóm -OH tự do có thể tham gia phản ứng ester hóa với HNO~3~:\n  - Vị trí (1): nhóm -OH gắn vào carbon C2.\n  - Vị trí (5): nhóm -OH gắn vào carbon C3.\n  - Vị trí (3): nhóm -CH~2~OH gắn vào carbon C6.\n• Vị trí (2) và (4) là nguyên tử oxygen trong cầu nối glycoside β-1,4 và dị vòng pyranose (liên kết ether bền vững), không tham gia phản ứng.\n➔ Ba nguyên tử oxygen tham gia phản ứng ở vị trí (1), (3), (5).\n➔ Đáp án đúng: A.")

    # Câu 8
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 8. Phát biểu nào sau đây về xà phòng và chất giặt rửa không đúng?")
    opts8 = [
        "A. Chất giặt rửa tổng hợp có thể được sản xuất từ các sản phẩm của dầu mỏ.",
        "B. Xà phòng là hỗn hợp muối sodium hoặc potassium của acid béo và một số chất phụ gia.",
        "*C. Xà phòng giặt rửa tốt trong nước cứng vì muối calcium, magnesium của acid béo tan tốt trong nước.*" if is_teacher else "C. Xà phòng giặt rửa tốt trong nước cứng vì muối calcium, magnesium của acid béo tan tốt trong nước.",
        "D. Phân tử chất giặt rửa gồm phần ưa nước và phần kị nước (ưa dầu mỡ)."
    ]
    render_options_no_table(doc, opts8, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Phát biểu C sai vì trong nước cứng chứa nhiều ion Ca^2+^ và Mg^2+^, các ion này tạo kết tủa với anion acid béo ((RCOO)~2~Ca, (RCOO)~2~Mg không tan) làm kết tủa xà phòng, mất tác dụng giặt rửa và làm hỏng sợi vải. Vì vậy xà phòng không giặt rửa tốt trong nước cứng.\n➔ Đáp án đúng: C.")

    # Câu 9
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 9. Khi nghiên cứu về phản ứng điều chế ethyl acetate từ ethanol 90° và acetic acid nguyên chất, một nhóm học sinh muốn khảo sát ảnh hưởng của lượng H~2~SO~4~ đặc đến hiệu suất điều chế ester. Để thực hiện điều này, nhóm học sinh đã xây dựng kế hoạch cho ba thí nghiệm điều chế, trong đó có liệt kê một số yếu tố sau:\n(1) Nhiệt độ khi thực hiện phản ứng.\n(2) Thể tích dung dịch H~2~SO~4~ 98%.\n(3) Thể tích ethanol và acetic acid ban đầu.\n(4) Thời gian thực hiện phản ứng.\nNhững yếu tố nào cần được thiết lập giống nhau ở cả ba thí nghiệm?")
    opts9 = [
        "*A. (1), (3), (4).*" if is_teacher else "A. (1), (3), (4).",
        "B. (1), (2), (4).",
        "C. (1), (2), (3).",
        "D. (2), (3), (4)."
    ]
    render_options_no_table(doc, opts9, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Mục tiêu là khảo sát ảnh hưởng của lượng H~2~SO~4~ đặc (đây là biến số độc lập cần thay đổi giữa 3 thí nghiệm). Do đó các yếu tố còn lại (nhiệt độ, thể tích các chất tham gia ban đầu, thời gian đun phản ứng) phải được kiểm soát giữ cố định giống nhau tuyệt đối.\n➔ Các yếu tố cần thiết lập giống nhau là: (1), (3), (4).\n➔ Đáp án đúng: A.")

    # Câu 10
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 10. Trong quá trình sản xuất bia bằng phương pháp lên men sinh học, dưới tác dụng của enzyme sẽ xảy ra quá trình chuyển hoá sau: X ➔ maltose ➔ Y. Chất X, Y tương ứng là")
    opts10 = [
        "A. tinh bột và fructose.",
        "B. cellulose và glucose.",
        "C. cellulose và fructose.",
        "*D. tinh bột và glucose.*" if is_teacher else "D. tinh bột và glucose."
    ]
    render_options_no_table(doc, opts10, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Trong công nghệ ủ mạch nha sản xuất bia:\n• Tinh bột (X) trong hạt đại mạch nảy mầm dưới tác dụng của enzyme amylase bị thủy phân thành đường maltose.\n• Sau đó maltose tiếp tục được enzyme maltase thủy phân tạo thành glucose (Y).\n• Cuối cùng glucose lên men rượu bởi nấm men tạo ethanol.\n➔ X là tinh bột, Y là glucose.\n➔ Đáp án đúng: D.")

    # Câu 11
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 11. Chất nào sau đây không thuộc nhóm hợp chất carbohydrate?")
    opts11 = [
        "A. Fructose.",
        "B. Tinh bột.",
        "*C. Glucosamine.*" if is_teacher else "C. Glucosamine.",
        "D. Glucose."
    ]
    render_options_no_table(doc, opts11, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Carbohydrate là những hợp chất hữu cơ tạp chức thường có công thức chung C~n~(H~2~O)~m~ (chỉ chứa các nguyên tố C, H, O). Glucosamine có công thức phân tử C~6~H~13~NO~5~, chứa nhóm amino (-NH~2~), là một dẫn xuất amino-sugar chứ không phải là carbohydrate thuần túy.\n➔ Đáp án đúng: C.")

    # Câu 12
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 12. Sodium stearate (C~17~H~35~COONa) có khả năng làm sạch vết dầu mỡ bám trên vải là do")
    opts12 = [
        "A. sodium stearate phản ứng hóa học với dầu mỡ tạo thành chất tan trong nước.",
        "B. gốc C~17~H~35~– ưa nước, còn nhóm –COONa kị nước nên kéo vết dầu mỡ vào nước.",
        "C. sodium stearate làm tăng sức căng bề mặt của nước nên nước dễ thấm vào vải.",
        "*D. gốc C~17~H~35~– thâm nhập vào vết dầu mỡ, nhóm –COO^– ưa nước hướng ra ngoài, giúp vết bẩn phân tán thành các hạt nhỏ trong nước.*" if is_teacher else "D. gốc C~17~H~35~– thâm nhập vào vết dầu mỡ, nhóm –COO^– ưa nước hướng ra ngoài, giúp vết bẩn phân tán thành các hạt nhỏ trong nước."
    ]
    render_options_no_table(doc, opts12, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Cơ chế tẩy rửa của xà phòng: Phân tử sodium stearate gồm đuôi kị nước (gốc hydrocarbon dài C~17~H~35~–) cắm sâu hòa tan vào giọt dầu mỡ, còn đầu ưa nước (nhóm –COO^– phân cực) hướng ra môi trường nước. Lực đẩy tĩnh điện giữa các đầu phân cực giúp bẻ nhỏ vết dầu mỡ thành các hạt micelle phân tán lơ lửng trong nước và bị rửa trôi.\n➔ Đáp án đúng: D.")

    # Câu 13
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 13. Tính chất của nhóm -OH hemiacetal trong glucose được thể hiện khi phản ứng với chất nào sau đây?")
    opts13 = [
        "A. [Ag(NH~3~)~2~]OH (t°).",
        "*B. CH~3~OH (xúc tác HCl khan).*" if is_teacher else "B. CH~3~OH (xúc tác HCl khan).",
        "C. Cu(OH)~2~/OH^- (t°).",
        "D. Cu(OH)~2~."
    ]
    render_options_no_table(doc, opts13, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Nhóm -OH hemiacetal ở vị trí C1 của glucose có khả năng phản ứng chọn lọc với alcohol (như CH~3~OH khi có xúc tác khí HCl khan) tạo thành methyl glucoside (liên kết ether glycoside). Các phản ứng tráng bạc hoặc khử Cu(OH)~2~ ở t° là của nhóm -CHO mạch hở; phản ứng tạo phức màu xanh lam với Cu(OH)~2~ ở nhiệt độ thường là tính chất của polyalcohol.\n➔ Đáp án đúng: B.")

    # Câu 14
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 14. Cấu tạo dạng mạch vòng của hai disaccharide: lactose và saccharose được biểu diễn như hình bên.\nPhát biểu nào sau đây không đúng?")
    add_centered_image(doc, os.path.join(IMG_DIR, "cau14_disaccharide.png"), width_cm=9.5)
    opts14 = [
        "A. Lactose và saccharose là đồng phân cấu tạo của nhau.",
        "B. Lactose và saccharose đều có khả năng tham gia phản ứng thủy phân trong môi trường acid.",
        "*C. Trong phân tử lactose, các gốc glucose liên kết với nhau qua liên kết α -1,4- glycoside.*" if is_teacher else "C. Trong phân tử lactose, các gốc glucose liên kết với nhau qua liên kết α -1,4- glycoside.",
        "D. Ở điều kiện thường, hai dung dịch lactose và saccharose đều hòa tan được Cu(OH)~2~ trong môi trường kiềm."
    ]
    render_options_no_table(doc, opts14, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Phát biểu C sai vì lactose được cấu tạo từ 1 gốc β-galactose và 1 gốc α-glucose liên kết với nhau qua liên kết β-1,4-glycoside (không phải cấu tạo từ hai gốc glucose qua liên kết α-1,4-glycoside như maltose).\n➔ Đáp án đúng: C.")

    # Câu 15
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 15. Đặc điểm nào sau đây không đúng đối với methylamine?")
    opts15 = [
        "A. Methylamine có công thức phân tử CH~5~N.",
        "*B. Methylamine là amine bậc hai.*" if is_teacher else "B. Methylamine là amine bậc hai.",
        "C. Methylamine làm quỳ tím ẩm đổi thành màu xanh.",
        "D. Methylamine phản ứng được với dung dịch HCl."
    ]
    render_options_no_table(doc, opts15, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Methylamine có CTCT là CH~3~NH~2~, nguyên tử nitơ liên kết với 1 gốc methyl nên thuộc loại amine bậc một (không phải amine bậc hai).\n➔ Đáp án đúng: B.")

    # Câu 16
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 16. Propyl tiglate là một loại chất tạo nên mùi hương của quả táo, có công thức cấu tạo như sau:")
    add_centered_image(doc, os.path.join(IMG_DIR, "cau16_propyl_tiglate.png"), width_cm=5.0)
    p_sub16 = doc.add_paragraph()
    p_sub16.paragraph_format.space_before = Pt(1)
    p_sub16.paragraph_format.space_after = Pt(1)
    render_rich_text(p_sub16, "Phát biểu nào sau đây không đúng?")
    opts16 = [
        "*A. Thủy phân propyl tiglate trong dung dịch NaOH thu được ethanol.*" if is_teacher else "A. Thủy phân propyl tiglate trong dung dịch NaOH thu được ethanol.",
        "B. Propyl tiglate có 2 liên kết pi (π) trong phân tử.",
        "C. Propyl tiglate là ester không no, đơn chức, mạch hở.",
        "D. Công thức phân tử của propyl tiglate là C~8~H~14~O~2~."
    ]
    render_options_no_table(doc, opts16, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Thủy phân propyl tiglate (CH~3~-CH=C(CH~3~)-COO-CH~2~CH~2~CH~3~) trong NaOH:\nCH~3~-CH=C(CH~3~)-COOCH~2~CH~2~CH~3~ + NaOH ➔ CH~3~-CH=C(CH~3~)-COONa + CH~3~CH~2~CH~2~OH.\nSản phẩm alcohol thu được là propan-1-ol (propyl alcohol), không phải ethanol. ➔ A sai.\n➔ Đáp án đúng: A.")

    # Câu 17
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 17. Phát biểu nào sau đây về fructose không đúng?")
    opts17 = [
        "*A. Fructose làm mất màu nước bromine ở điều kiện thường.*" if is_teacher else "A. Fructose làm mất màu nước bromine ở điều kiện thường.",
        "B. Fructose là đồng phân của glucose.",
        "C. Fructose thuộc loại monosaccharide.",
        "D. Fructose phản ứng được với Cu(OH)~2~ trong môi trường kiềm ở điều kiện thường."
    ]
    render_options_no_table(doc, opts17, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Fructose là một polyhydroxy ketone, trong môi trường acid/trung tính của nước bromine nó không có phản ứng và không làm mất màu nước bromine. Nước bromine chính là thuốc thử dùng để phân biệt glucose (làm mất màu) và fructose (không làm mất màu).\n➔ Đáp án đúng: A.")

    # Câu 18
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    render_question_prompt(p, "Câu 18. Trong quá trình sản xuất xà phòng, sau khi đun chất béo với dung dịch NaOH, người ta thêm dung dịch NaCl bão hòa vào hỗn hợp nhằm mục đích")
    opts18 = [
        "*A. làm giảm độ tan của xà phòng, giúp xà phòng tách ra rõ rệt và nổi lên trên.*" if is_teacher else "A. làm giảm độ tan của xà phòng, giúp xà phòng tách ra rõ rệt và nổi lên trên.",
        "B. làm xúc tác cho phản ứng xà phòng hóa.",
        "C. làm tăng tốc độ phản ứng xà phòng hóa.",
        "D. trung hòa lượng NaOH còn dư."
    ]
    render_options_no_table(doc, opts18, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Dung dịch NaCl bão hòa cung cấp ion Na^+ nồng độ cao (hiệu ứng ion chung) và làm tăng khối lượng riêng của dung dịch, khiến độ tan của muối sodium stearate/palmitate giảm mạnh, xà phòng đông kết hoàn toàn và nổi lên bề mặt để thu lấy.\n➔ Đáp án đúng: A.")

    # =========================================================================
    # PHẦN II
    # =========================================================================
    add_section_header(
        doc,
        "PHẦN II (4,0 điểm). Cấu trúc trắc nghiệm đúng sai. Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.",
        subtitle="Điểm tối đa cho 01 câu hỏi là 1,0 điểm (Đúng 1 ý: 0,1đ; 2 ý: 0,25đ; 3 ý: 0,5đ; 4 ý: 1,0đ)"
    )

    # Câu 1 Phần II
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 1. Cellulose và tinh bột đều có công thức chung là (C~6~H~10~O~5~)~n~. Tinh bột gồm hai thành phần là amylose và amylopectin. Thủy phân hoàn toàn cellulose và tinh bột trong môi trường acid đun nóng, theo phương trình hóa học sau: (C~6~H~10~O~5~)~n~ + nH~2~O ➔ [H^+, t°] ➔ nC~6~H~12~O~6~  (1)")
    stmts1 = [
        ("a) Cellulose có khả năng tạo phức chất màu xanh tím với iodine.", "Sai"),
        ("b) Phản ứng thủy phân cellulose và tinh bột theo phương trình (1) tạo thành glucose.", "Đúng"),
        ("c) Vì cellulose và tinh bột đều có công thức chung là (C~6~H~10~O~5~)~n~ nên phân tử khối của cellulose và tinh bột đều bằng nhau.", "Sai"),
        ("d) Cellulose và amylopectin đều cấu tạo từ các đơn vị glucose liên kết với nhau bằng các liên kết glycoside nên chúng đều có cấu tạo mạch không phân nhánh.", "Sai")
    ]
    for stmt, ans in stmts1:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.left_indent = Cm(0.5)
        p_st.paragraph_format.space_before = Pt(1)
        p_st.paragraph_format.space_after = Pt(1)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)

    if is_teacher:
        add_solution_no_table(doc, "a) Sai: Tinh bột tạo phức màu xanh tím với iodine do cấu trúc lò xo xoắn giữ phân tử I~2~; cellulose dạng sợi kéo dài không có tính chất này.\nb) Đúng: Thủy phân hoàn toàn tinh bột hay cellulose bằng acid đều chỉ thu được glucose (C~6~H~12~O~6~).\nc) Sai: Mặc dù cùng công thức đơn giản nhất, hệ số polime hóa n của cellulose (1.000 - 14.000) lớn hơn rất nhiều so với tinh bột, do đó phân tử khối của chúng khác nhau nhiều.\nd) Sai: Amylopectin có cấu tạo mạch phân nhánh do các liên kết α-1,6-glycoside tại điểm rẽ nhánh.")

    # Câu 2 Phần II
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 2. Trong kiểm định chất lượng dầu, chỉ số acid (AV) được định nghĩa là số miligam KOH cần thiết để trung hòa toàn bộ acid béo tự do có trong 1,00 gam chất béo; chỉ số xà phòng hóa (SV) là số miligam KOH cần thiết để xà phòng hóa hoàn toàn triglyceride và đồng thời trung hòa acid béo tự do có trong 1,00 gam chất béo. Một mẫu chất béo (X) gồm triglyceride và một lượng nhỏ acid béo tự do. Triglyceride trong mẫu chất béo (X) có công thức cấu tạo như sau:")
    add_centered_image(doc, os.path.join(IMG_DIR, "p2_cau2_triglyceride.png"), width_cm=7.5)
    
    p_ctx2 = doc.add_paragraph()
    p_ctx2.paragraph_format.space_before = Pt(2)
    p_ctx2.paragraph_format.space_after = Pt(2)
    render_rich_text(p_ctx2, "Phòng thí nghiệm tiến hành hai thí nghiệm độc lập sau:\n- Thí nghiệm 1: Hòa tan 20,00 gam chất béo (X) vào hỗn hợp dung môi thích hợp, thêm vài giọt phenolphthalein rồi chuẩn độ bằng dung dịch chuẩn KOH 0,10 M. Đến điểm kết thúc chuẩn độ, thể tích dung dịch KOH đã tiêu tốn là 6,00 mL. Giả thiết trong thí nghiệm này, liên kết ester trong chất béo không tham gia phản ứng.\n- Thí nghiệm 2: Đun nóng 40,00 gam chất béo (X) với 200,00 mL dung dịch KOH 1,00 M trong cồn đến khi phản ứng xà phòng hóa xảy ra hoàn toàn. Để trung hòa lượng KOH còn dư sau phản ứng cần vừa đủ 120,00 mL dung dịch chuẩn HCl 0,50 M. Giả thiết trong thí nghiệm này, toàn bộ liên kết ester trong chất béo và nhóm carboxylic acid trong acid béo đều tham gia phản ứng.")

    stmts2 = [
        ("a) Trong thí nghiệm 1, dung môi thích hợp để hòa tan mẫu (X) hoàn toàn phải là nước cất.", "Sai"),
        ("b) Gốc hydrocarbon mạch dài trong triglyceride chỉ gồm hai loại gốc: stearate – C~17~H~35~ và oleate – C~17~H~33~.", "Đúng"),
        ("c) Từ thí nghiệm 1, chỉ số acid (AV) của mẫu chất béo (X) bằng 3,36.", "Sai"),
        ("d) Từ thí nghiệm 2, chỉ số xà phòng hoá (SV) của mẫu chất béo (X) bằng 196,00.", "Đúng")
    ]
    for stmt, ans in stmts2:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.left_indent = Cm(0.5)
        p_st.paragraph_format.space_before = Pt(1)
        p_st.paragraph_format.space_after = Pt(1)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)

    if is_teacher:
        add_solution_no_table(doc, "a) Sai: Chất béo kị nước, không tan trong nước cất. Dung môi hòa tan mẫu chất béo chuẩn độ là cồn - ether trung tính.\nb) Đúng: Dựa vào công thức cấu tạo, triglyceride chứa 1 gốc stearate C~17~H~35~COO- và 2 gốc oleate C~17~H~33~COO-.\nc) Sai: Ở TN1: n~KOH~ = 0,10 M · 0,006 L = 0,00060 mol ➔ m~KOH~ = 0,00060 · 56 · 1000 = 33,6 mg trong 20,00 g chất béo. Chỉ số acid AV = 33,6 / 20 = 1,68 mg KOH/g chất béo (khác 3,36).\nd) Đúng: Ở TN2: n~KOH bđ~ = 0,200 L · 1,00 M = 0,200 mol; n~HCl chuẩn độ dư~ = 0,120 L · 0,50 M = 0,060 mol ➔ n~KOH pư~ = 0,200 - 0,060 = 0,140 mol trong 40,00 g chất béo. Khối lượng KOH phản ứng: m~KOH~ = 0,140 · 56 · 1000 = 7840 mg. Chỉ số xà phòng hóa SV = 7840 / 40,00 = 196,00 mg KOH/g chất béo.")

    # Câu 3 Phần II
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 3. Trong dung dịch, glucose và fructose đều tồn tại cả dạng mạch hở và mạch vòng. Quá trình chuyển từ mạch vòng sang mạch hở là do trong phân tử hai chất này có nhóm -OH hemiacetal hoặc -OH hemiketal. Hình vẽ dưới đây mô tả cấu tạo của α-glucose, β-fructose và saccharose, trong đó nhóm -OH được khoanh tròn là nhóm -OH hemiacetal của α-glucose và nhóm -OH hemiketal của β-fructose:")
    add_centered_image(doc, os.path.join(IMG_DIR, "p2_cau3_hemiacetal.png"), width_cm=10.5)

    stmts3 = [
        ("a) Đơn vị α-glucose và đơn vị β-fructose trong saccharose liên kết với nhau qua liên kết α -1,2- glycoside.", "Đúng"),
        ("b) Do cả glucose và fructose đều không bị thuỷ phân nên chúng thuộc loại monosaccharide.", "Đúng"),
        ("c) Khác với glucose và fructose, saccharose không có dạng mở vòng vì trong phân tử saccharose không chứa nhóm –OH hemiacetal và –OH hemiketal.", "Đúng"),
        ("d) Dung dịch saccharose có tính chất của polyalcohol vì phân tử saccharose chứa nhiều nhóm -OH liền kề nhau.", "Đúng")
    ]
    for stmt, ans in stmts3:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.left_indent = Cm(0.5)
        p_st.paragraph_format.space_before = Pt(1)
        p_st.paragraph_format.space_after = Pt(1)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)

    if is_teacher:
        add_solution_no_table(doc, "a) Đúng: C1 của gốc α-glucose liên kết với C2 của gốc β-fructose qua cầu oxy tạo liên kết α-1,2-glycoside.\nb) Đúng: Monosaccharide là các loại carbohydrate đơn giản nhất không thể bị thủy phân.\nc) Đúng: Do cả nhóm -OH hemiacetal của glucose (C1) và hemiketal của fructose (C2) đều tham gia liên kết glycoside, phân tử saccharose không còn nhóm -OH hemiacetal/hemiketal tự do nên không thể mở vòng.\nd) Đúng: Saccharose có 8 nhóm -OH, nhiều nhóm nằm ở vị trí kề nhau nên hòa tan Cu(OH)~2~ tạo dung dịch màu xanh lam đặc trưng của polyalcohol.")

    # Câu 4 Phần II
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 4. Một nhóm học sinh tìm hiểu ảnh hưởng của nồng độ đường đến quá trình lên men lactic, với giả thuyết: “Nồng độ glucose ban đầu càng cao thì lượng lactic acid thu được càng nhiều”. Nhóm đã thực hiện thí nghiệm như sau:\nChuẩn bị: các dung dịch glucose 2%, 5%, 10%, 15%, 20% (khối lượng/thể tích, tức g/100 mL), chế phẩm vi sinh vật lên men lactic (Lactobacillus), 6 bình thủy tinh có nắp, tủ ấm điều nhiệt.\nCách làm: Lấy 200 mL mỗi dung dịch glucose cho vào các bình tương ứng (mẫu 1 ➔ mẫu 5). Thêm cùng một lượng chế phẩm vi sinh vật vào mỗi bình, đậy nắp và ủ ở 37 °C. Bình thứ 6 chứa 200 mL glucose 10% nhưng không thêm vi sinh vật (mẫu trắng). Sau 3 ngày, lọc sạch các mẫu, để ổn định ở nhiệt độ phòng thì thấy thể tích dung dịch trong các mẫu hầu như không thay đổi là 200 mL. Lần lượt lấy 5 mL các mẫu: mẫu trắng; các mẫu 1; 2; 3; 4; 5 rồi pha loãng thành 100 mL và tiến hành chuẩn độ bằng dung dịch NaOH 0,1 M với chỉ thị phenolphthalein. Kết quả thu được như bảng sau:")

    # Bảng số liệu chuẩn hóa bằng Paragraph kết hợp Tab stops
    p_t1 = doc.add_paragraph()
    p_t1.paragraph_format.left_indent = Cm(1.2)
    p_t1.paragraph_format.tab_stops.add_tab_stop(Cm(4.5), WD_TAB_ALIGNMENT.LEFT)
    p_t1.paragraph_format.tab_stops.add_tab_stop(Cm(7.0), WD_TAB_ALIGNMENT.LEFT)
    p_t1.paragraph_format.tab_stops.add_tab_stop(Cm(9.5), WD_TAB_ALIGNMENT.LEFT)
    p_t1.paragraph_format.tab_stops.add_tab_stop(Cm(12.0), WD_TAB_ALIGNMENT.LEFT)
    p_t1.paragraph_format.tab_stops.add_tab_stop(Cm(14.5), WD_TAB_ALIGNMENT.LEFT)
    render_rich_text(p_t1, "*Mẫu:* \t*Mẫu trắng* \t*Mẫu 1* \t*Mẫu 2* \t*Mẫu 3* \t*Mẫu 4* \t*Mẫu 5*", base_bold=True)

    p_t2 = doc.add_paragraph()
    p_t2.paragraph_format.left_indent = Cm(1.2)
    p_t2.paragraph_format.tab_stops.add_tab_stop(Cm(4.5), WD_TAB_ALIGNMENT.LEFT)
    p_t2.paragraph_format.tab_stops.add_tab_stop(Cm(7.0), WD_TAB_ALIGNMENT.LEFT)
    p_t2.paragraph_format.tab_stops.add_tab_stop(Cm(9.5), WD_TAB_ALIGNMENT.LEFT)
    p_t2.paragraph_format.tab_stops.add_tab_stop(Cm(12.0), WD_TAB_ALIGNMENT.LEFT)
    p_t2.paragraph_format.tab_stops.add_tab_stop(Cm(14.5), WD_TAB_ALIGNMENT.LEFT)
    render_rich_text(p_t2, "*V~dd NaOH~ (mL):* \t0,0 \t4,8 \t8,4 \t12,0 \t9,6 \t6,0")

    p_note4 = doc.add_paragraph()
    p_note4.paragraph_format.space_before = Pt(2)
    p_note4.paragraph_format.space_after = Pt(2)
    render_rich_text(p_note4, "Biết rằng trong các mẫu chỉ chứa C~6~H~12~O~6~, CH~3~CH(OH)COOH (lactic acid) và H~2~O; glucose chỉ chuyển hóa thành lactic acid.")

    stmts4 = [
        ("a) Trong thí nghiệm trên, lượng lactic acid được tạo thành nhiều nhất trong một mẫu (200 mL) là 4,32 gam.", "Đúng"),
        ("b) Nếu thay dung dịch NaOH 0,1 M bằng dung dịch NaOH 0,2 M thì thể tích NaOH dùng để chuẩn độ mẫu 3 sẽ nhỏ hơn 5,0 mL.", "Sai"),
        ("c) Từ kết quả thí nghiệm, kết luận giả thuyết của nhóm học sinh trên là đúng.", "Sai"),
        ("d) Hiệu suất chuyển hóa glucose thành lactic acid ở mẫu 1 cao hơn ở mẫu 3.", "Đúng")
    ]
    for stmt, ans in stmts4:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.left_indent = Cm(0.5)
        p_st.paragraph_format.space_before = Pt(1)
        p_st.paragraph_format.space_after = Pt(1)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)

    if is_teacher:
        add_solution_no_table(doc, "a) Đúng: Mẫu 3 có lượng lactic acid tạo thành nhiều nhất (V~NaOH~ = 12,0 mL). Lấy 5 mL chuẩn độ hết n~NaOH~ = 0,012 · 0,1 = 1,2 · 10^-3 mol. Trong cả bình 200 mL (gấp 40 lần) có n~acid~ = 1,2 · 10^-3 · 40 = 0,048 mol ➔ m~acid~ = 0,048 · 90 = 4,32 gam.\nb) Sai: Khi dùng NaOH 0,2 M: V~NaOH~ = 1,2 · 10^-3 / 0,2 = 0,006 L = 6,0 mL > 5,0 mL.\nc) Sai: Nồng độ glucose tăng lên 15% và 20% thì lượng lactic acid thu được lại giảm (V~NaOH~ giảm xuống 9,6 và 6,0 mL) do áp suất thẩm thấu cao ức chế vi khuẩn. Do đó giả thuyết ban đầu là sai.\nd) Đúng: • Mẫu 1 (2%, 200 mL): m~glu bđ~ = 4,00 gam. Lượng acid tạo ra = 4,8 · 0,1 · 10^-3 · 40 = 0,0192 mol ➔ m~glu pư~ = (0,0192 / 2) · 180 = 1,728 gam ➔ H~1~ = (1,728 / 4,00) · 100% = 43,2%.\n• Mẫu 3 (10%, 200 mL): m~glu bđ~ = 20,00 gam. m~glu pư~ = (0,048 / 2) · 180 = 4,32 gam ➔ H~3~ = (4,32 / 20,00) · 100% = 21,6%.\n➔ Hiệu suất ở mẫu 1 cao hơn mẫu 3.")

    # =========================================================================
    # PHẦN III
    # =========================================================================
    add_section_header(
        doc,
        "PHẦN III (1,5 điểm). Thí sinh trả lời từ câu 1 đến câu 6.",
        subtitle="Mỗi câu trả lời đúng thí sinh được 0,25 điểm"
    )

    # Câu 1 Phần III
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 1. Có bao nhiêu đồng phân cấu tạo của amine bậc một có công thức phân tử C~4~H~11~N?")
    if not is_teacher:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(2)
        p_ans.paragraph_format.space_after = Pt(2)
        render_rich_text(p_ans, "*Trả lời:* ...................................................................................................................................................")
    else:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(2)
        p_ans.paragraph_format.space_after = Pt(2)
        render_rich_text(p_ans, "*Trả lời:* *4*", base_color=COLOR_CORRECT)
        add_solution_no_table(doc, "Amine bậc một có dạng C~4~H~9~NH~2~. Có 4 gốc butyl (-C~4~H~9~):\n1. CH~3~CH~2~CH~2~CH~2~NH~2~ (butan-1-amine)\n2. CH~3~CH~2~CH(NH~2~)CH~3~ (butan-2-amine)\n3. (CH~3~)~2~CHCH~2~NH~2~ (2-methylpropan-1-amine)\n4. (CH~3~)~3~CNH~2~ (2-methylpropan-2-amine)\n➔ Số đồng phân cấu tạo thỏa mãn là 4.\n➔ Đáp số: 4.")

    # Câu 2 Phần III
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 2. Cho sơ đồ điều chế amine X từ benzene:")
    img_aniline = ANILINE_HD if os.path.exists(ANILINE_HD) else os.path.join(IMG_DIR, "p3_cau2_aniline.png")
    add_centered_image(doc, img_aniline, width_cm=9.0)
    p_sub_c2 = doc.add_paragraph()
    p_sub_c2.paragraph_format.space_before = Pt(1)
    p_sub_c2.paragraph_format.space_after = Pt(1)
    render_rich_text(p_sub_c2, "Phân tử khối của X là bao nhiêu?")
    if not is_teacher:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(2)
        p_ans.paragraph_format.space_after = Pt(2)
        render_rich_text(p_ans, "*Trả lời:* ...................................................................................................................................................")
    else:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(2)
        p_ans.paragraph_format.space_after = Pt(2)
        render_rich_text(p_ans, "*Trả lời:* *93*", base_color=COLOR_CORRECT)
        add_solution_no_table(doc, "Theo sơ đồ: Benzene ➔ Nitrobenzene (C~6~H~5~NO~2~) ➔ Aniline (C~6~H~5~NH~2~).\nPhân tử khối của amine X (aniline C~6~H~5~NH~2~):\nM = 12 · 6 + 1 · 5 + 14 + 1 · 2 = 93 g/mol.\n➔ Đáp số: 93.")

    # Câu 3 Phần III
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 3. Cho các phản ứng sau:\n(1) 6nCO~2~ + 5nH~2~O ➔ [chlorophyll] ➔ (C~6~H~10~O~5~)~n~ + 6nO~2~\n(2) C~6~H~12~O~6~ ➔ [enzyme] ➔ 2CH~3~-CHOH-COOH (lactic acid)\n(3) (C~6~H~10~O~5~)~n~ + nH~2~O ➔ [H^+, t°] ➔ nC~6~H~12~O~6~ (glucose)\n(4) C~6~H~12~O~6~ ➔ [enzyme] ➔ 2C~2~H~5~OH + 2CO~2~\n(5) (C~17~H~35~COO)~3~C~3~H~5~ (tristearin) + 3NaOH ➔ [t°] ➔ 3C~17~H~35~COONa + C~3~H~5~(OH)~3~\n(6) C~6~H~12~O~6~ (glucose) + 2[Ag(NH~3~)~2~]OH ➔ [t°] ➔ CH~2~OH(CHOH)~4~COONH~4~ + 2Ag + 3NH~3~ + H~2~O\nGán số thứ tự phương trình hoá học của các phản ứng theo tên gọi: xà phòng hóa, quang hợp, lên men lactic, tráng gương và sắp xếp theo trình tự thành dãy bốn số (Ví dụ: 1234, 4321, ...)")
    if not is_teacher:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(2)
        p_ans.paragraph_format.space_after = Pt(2)
        render_rich_text(p_ans, "*Trả lời:* ...................................................................................................................................................")
    else:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(2)
        p_ans.paragraph_format.space_after = Pt(2)
        render_rich_text(p_ans, "*Trả lời:* *5126*", base_color=COLOR_CORRECT)
        add_solution_no_table(doc, "Gán số thứ tự tương ứng với mỗi quá trình:\n• Phản ứng xà phòng hóa: phương trình (5).\n• Phản ứng quang hợp: phương trình (1).\n• Phản ứng lên men lactic: phương trình (2).\n• Phản ứng tráng gương: phương trình (6).\n➔ Trình tự bốn số: 5 - 1 - 2 - 6.\n➔ Dãy số: 5126.")

    # Câu 4 Phần III
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 4. Xà phòng giặt kém hiệu quả trong nước cứng vì ion Ca^2+^ kết tủa với anion của acid béo:\nCa^2+^ + 2C~17~H~35~COO^- ➔ (C~17~H~35~COO)~2~Ca↓\nMột gia đình dùng 50 lít nước giếng khoan có nồng độ ion Ca^2+^ là 59 mg/L để giặt quần áo bằng xà phòng bánh chứa 84% sodium stearate (C~17~H~35~COONa) về khối lượng. Tính khối lượng (gam) xà phòng bánh bị hao phí do tạo kết tủa với toàn bộ lượng ion Ca^2+^ trong 50 lít nước trên. (Làm tròn kết quả đến hàng phần mười.)")
    if not is_teacher:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(2)
        p_ans.paragraph_format.space_after = Pt(2)
        render_rich_text(p_ans, "*Trả lời:* ...................................................................................................................................................")
    else:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(2)
        p_ans.paragraph_format.space_after = Pt(2)
        render_rich_text(p_ans, "*Trả lời:* *53,7*", base_color=COLOR_CORRECT)
        add_solution_no_table(doc, "• Khối lượng ion Ca^2+^ trong 50 L nước giếng:\n  m~Ca2+~ = 59 mg/L · 50 L = 2950 mg = 2,95 gam.\n• Số mol ion Ca^2+^:\n  n~Ca2+~ = 2,95 / 40 = 0,07375 mol.\n• Theo PTHH: Ca^2+^ + 2C~17~H~35~COO^- ➔ (C~17~H~35~COO)~2~Ca↓\n  Số mol C~17~H~35~COONa bị kết tủa: n = 2 · 0,07375 = 0,1475 mol.\n• Khối lượng sodium stearate nguyên chất bị tiêu tốn (M = 306 g/mol):\n  m = 0,1475 mol · 306 g/mol = 45,135 gam.\n• Khối lượng xà phòng bánh (chứa 84% muối) bị hao phí:\n  m~xà phòng~ = 45,135 / 0,84 ≈ 53,732 gam.\nLàm tròn đến hàng phần mười: 53,7 gam.\n➔ Đáp số: 53,7.")

    # Câu 5 Phần III
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 5. Biodiesel được sản xuất từ dầu ăn đã qua sử dụng (dầu thải) bằng phương trình phản ứng tổng quát như sau: (RCOO)~3~C~3~H~5~ + 3CH~3~OH ➔ 3RCOOCH~3~ + C~3~H~5~(OH)~3~.\nPhân tích thành phần của 100 gam biodiesel từ dầu chiên thải có thành phần (coi chỉ gồm 3 ester) như sau:")

    # Bảng thành phần 3 ester dùng Tab stop
    p_b1 = doc.add_paragraph()
    p_b1.paragraph_format.left_indent = Cm(1.5)
    p_b1.paragraph_format.tab_stops.add_tab_stop(Cm(6.0), WD_TAB_ALIGNMENT.LEFT)
    p_b1.paragraph_format.tab_stops.add_tab_stop(Cm(10.5), WD_TAB_ALIGNMENT.LEFT)
    render_rich_text(p_b1, "*Chất:* \t*C~15~H~31~COOCH~3~* \t*C~17~H~33~COOCH~3~* \t*C~17~H~35~COOCH~3~*", base_bold=True)

    p_b2 = doc.add_paragraph()
    p_b2.paragraph_format.left_indent = Cm(1.5)
    p_b2.paragraph_format.tab_stops.add_tab_stop(Cm(6.0), WD_TAB_ALIGNMENT.LEFT)
    p_b2.paragraph_format.tab_stops.add_tab_stop(Cm(10.5), WD_TAB_ALIGNMENT.LEFT)
    render_rich_text(p_b2, "*Khối lượng (gam):* \t40 \t50 \t10")

    p_ctx5 = doc.add_paragraph()
    p_ctx5.paragraph_format.space_before = Pt(2)
    p_ctx5.paragraph_format.space_after = Pt(2)
    render_rich_text(p_ctx5, "Hệ số phát thải (EF) là khối lượng CO~2~ (gam) tạo thành trên 1 MJ nhiệt lượng tỏa ra. Cho EF của diesel truyền thống là 74,1 g CO~2~/MJ.\nBiết rằng biodiesel có khối lượng riêng 0,880 kg/L. Đốt cháy hoàn toàn 1 lít biodiesel tỏa ra 32,0 MJ (megaJoule) nhiệt lượng. Phần carbon trong gốc acid béo (R–COO–) có nguồn gốc từ thực vật; CO~2~ sinh ra từ phần này được coi là carbon trung hòa vì cây đã hấp thụ lượng CO~2~ đó khi quang hợp. Riêng carbon trong nhóm –CH~3~ của ester có nguồn gốc từ methanol (CH~3~OH) sản xuất từ khí thiên nhiên; CO~2~ sinh ra từ quá trình cháy phần này được tính là CO~2~ có nguồn gốc hóa thạch.\nPhần trăm giảm phát thải CO~2~ nguồn gốc hóa thạch khi thay diesel bằng biodiesel (cho cùng một nhiệt lượng) được tính theo công thức:\n%Δ = [(EF~diesel~ - EF~hóa thạch~) / EF~diesel~] · 100%\nTrong đó EF~hóa thạch~ là khối lượng CO~2~ hóa thạch (gam) sinh ra trên 1 MJ nhiệt lượng khi đốt cháy biodiesel.\nTính giá trị %Δ của mẫu biodiesel trên. (Chỉ xét khí thải từ quá trình đốt cháy. Không làm tròn kết quả các phép tính trung gian, chỉ làm tròn kết quả cuối cùng đến hàng phần mười.)")

    if not is_teacher:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(2)
        p_ans.paragraph_format.space_after = Pt(2)
        render_rich_text(p_ans, "*Trả lời:* ...................................................................................................................................................")
    else:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(2)
        p_ans.paragraph_format.space_after = Pt(2)
        render_rich_text(p_ans, "*Trả lời:* *94,3*", base_color=COLOR_CORRECT)
        add_solution_no_table(doc, "• Trong 100 gam biodiesel:\n  - C~15~H~31~COOCH~3~ (M = 270 g/mol): n~1~ = 40 / 270 = 4 / 27 mol.\n  - C~17~H~33~COOCH~3~ (M = 296 g/mol): n~2~ = 50 / 296 = 25 / 148 mol.\n  - C~17~H~35~COOCH~3~ (M = 298 g/mol): n~3~ = 10 / 298 = 5 / 149 mol.\n• Mỗi phân tử ester có đúng 1 nguyên tử C hóa thạch (nhóm -CH~3~ từ methanol):\n  Số mol C hóa thạch trong 100 g biodiesel: n~C~ = 4/27 + 25/148 + 5/149 mol.\n• Khối lượng biodiesel cần đốt để tạo 1 MJ nhiệt lượng:\n  m~1MJ~ = (1000 mL · 0,880 g/mL) / 32,0 MJ = 880 / 32,0 = 27,5 gam.\n• Khối lượng CO~2~ hóa thạch phát thải trên 1 MJ (EF~hóa thạch~):\n  EF~hóa thạch~ = 44 · [27,5 / 100] · (4/27 + 25/148 + 5/149) = 12,1 · (4/27 + 25/148 + 5/149) ≈ 4,24255 gam CO~2~/MJ.\n• Phần trăm giảm phát thải CO~2~ nguồn gốc hóa thạch:\n  %Δ = [(74,1 - 4,24255) / 74,1] · 100% ≈ 94,2745% ≈ 94,3%.\n➔ Đáp số: 94,3.")

    # Câu 6 Phần III
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 6. Cellulose là thành phần chính của thành tế bào thực vật, làm nên bộ khung của cây cối. Cellulose có nhiều trong bông, đay, tre, nứa, gỗ,... Từ cellulose người ta điều chế được cellulose acetate (có thành phần chủ yếu là cellulose triacetate và cellulose diacetate) bằng phản ứng acetyl hóa theo các phương trình sau:\n(1) [C~6~H~7~O~2~(OH)~3~]~n~ + 2n(CH~3~CO)~2~O ➔ [C~6~H~7~O~2~(OCOCH~3~)~2~(OH)]~n~ + 2nCH~3~COOH\n(2) [C~6~H~7~O~2~(OH)~3~]~n~ + 3n(CH~3~CO)~2~O ➔ [C~6~H~7~O~2~(OCOCH~3~)~3~]~n~ + 3nCH~3~COOH\nAcetyl hóa hoàn toàn 16,2 gam cellulose bằng acetic anhydride (CH~3~CO)~2~O dư trong điều kiện thích hợp, thu được 26,7 gam hỗn hợp X chỉ gồm cellulose diacetate và cellulose triacetate. Giả sử toàn bộ cellulose chuyển thành X, mạch polymer không bị cắt và bỏ qua ảnh hưởng của các nhóm chức ở hai đầu mạch. Cho biết mỗi nhóm -OH được acetyl hóa chuyển thành nhóm -OCOCH~3~.\nXác định phần trăm khối lượng cellulose triacetate trong X. (làm tròn đến hàng phần mười).")
    if not is_teacher:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(2)
        p_ans.paragraph_format.space_after = Pt(2)
        render_rich_text(p_ans, "*Trả lời:* ...................................................................................................................................................")
    else:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(2)
        p_ans.paragraph_format.space_after = Pt(2)
        render_rich_text(p_ans, "*Trả lời:* *53,9*", base_color=COLOR_CORRECT)
        add_solution_no_table(doc, "• Số mol mắt xích cellulose ban đầu: n = 16,2 / 162 = 0,10 mol.\n• Gọi x, y lần lượt là số mol mắt xích cellulose diacetate (M~1~ = 246 g/mol) và cellulose triacetate (M~2~ = 288 g/mol) trong X:\n  x + y = 0,10\n  246x + 288y = 26,7\n➔ x = 0,05 mol; y = 0,05 mol.\n• Khối lượng cellulose triacetate: m = 0,05 · 288 = 14,4 gam.\n• Phần trăm khối lượng trong X:\n  %m(triacetate) = (14,4 / 26,7) · 100% ≈ 53,9326% ≈ 53,9%.\n➔ Đáp số: 53,9.")

    # Dòng kết thúc
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(8)
    p_end.paragraph_format.space_after = Pt(4)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rend = p_end.add_run("--- HẾT ---")
    set_run_style(rend, font_size=11, bold=True)

    if output_path:
        doc.save(output_path)
        print(f"Đã lưu thành công: {output_path}")

if __name__ == "__main__":
    out_dir = r"C:\Antigravity_Thanh\San_Pham"
    os.makedirs(out_dir, exist_ok=True)
    
    hs_path = os.path.join(out_dir, "De_Kiem_Tra_GK1_Hoa_12_NguyenHien_2026_HocSinh.docx")
    gv_path = os.path.join(out_dir, "De_Kiem_Tra_GK1_Hoa_12_NguyenHien_2026_GiaoVien.docx")
    
    build_exam_nguyenhien(is_teacher=False, output_path=hs_path)
    build_exam_nguyenhien(is_teacher=True, output_path=gv_path)
