# -*- coding: utf-8 -*-
"""
Tự động tạo 2 file Word cho ĐỀ SỐ 6 - HÓA HỌC 11:
- De_Kiem_Tra_Hoa_11_De06.docx (Bản Học Sinh)
- De_Kiem_Tra_Hoa_11_De06_GiaoVien.docx (Bản Giáo Viên)
Chương 2: NITROGEN – SULFUR
Tuân thủ 100% Form Đề Kiểm Tra trong GEMINI.md & form_de_kiem_tra.py:
- Tuyệt đối không dùng Table trong nội dung câu hỏi/đáp án/lời giải.
- Header & Footer chuẩn: #D32F2F, viền #8BC390, số trang tự động PAGE.
- Căn lề A4 1.5 cm.
- Chỉ số trên/dưới hóa học triệt để.
"""

import os
import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Thêm đường dẫn scripts để import module dùng chung
sys.path.append(os.path.dirname(os.path.abspath(__file__)))
from form_de_kiem_tra import (
    setup_header_footer,
    build_title_block,
    add_section_header,
    render_question_prompt,
    render_options_no_table,
    render_statement,
    add_solution_no_table,
    add_student_writing_lines,
    render_rich_text,
    set_run_style,
    COLOR_CORRECT,
    COLOR_SOLUTION,
    COLOR_PRIMARY
)

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
IMG_DIR = os.path.join(OUT_DIR, "images_crop_hoa11_de06")

def add_centered_picture(doc, img_path, width_cm=8.0):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(img_path, width=Cm(width_cm))

def build_exam_document(is_teacher=False):
    doc = docx.Document()
    setup_header_footer(doc)

    build_title_block(
        doc,
        school_line="SỞ GIÁO DỤC VÀ ĐÀO TẠO THỪA THIÊN HUẾ",
        exam_title="ĐỀ KIỂM TRA ĐỊNH KỲ CHƯƠNG 2 (NITROGEN – SULFUR)",
        subject_line="MÔN: HOÁ HỌC - LỚP 11",
        num_pages="04",
        duration="45 phút",
        exam_code="106",
        atomic_masses="H = 1; C = 12; N = 14; O = 16; Na = 23; Mg = 24; Al = 27; S = 32; Cl = 35,5; K = 39; Ca = 40; Fe = 56; Cu = 64; Zn = 65; Ag = 108; Ba = 137; Au = 197",
        is_teacher=is_teacher
    )

    # ==========================================
    # PHẦN I
    # ==========================================
    add_section_header(
        doc,
        title="PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án."
    )

    # Câu 1
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 1 (SBT - CTST). Ở trạng thái tự nhiên, nitrogen")
    opts_c1 = [
        "*A. tồn tại ở dạng đơn chất và hợp chất.*",
        "B. chỉ tồn tại ở dạng đơn chất.",
        "C. chỉ tồn tại ở dạng hợp chất.",
        "D. tự do chiếm khoảng 20% thể tích không khí."
    ]
    render_options_no_table(doc, opts_c1, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Trong tự nhiên, nitrogen tồn tại ở cả dạng đơn chất (chiếm khoảng 78% thể tích khí quyển Trái Đất) và dạng hợp chất (trong khoáng vật diêm tiêu NaNO~3~, KNO~3~, trong các hợp chất hữu cơ như protein, nucleic acid, chlorophyl,...).")

    # Câu 2
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 2 (SBT - KNTT). Trong phương pháp Ostwald, ammonia bị oxi hóa bởi oxygen không khí tạo thành sản phẩm chính là")
    opts_c2 = [
        "*A. NO.*",
        "B. N~2~.",
        "C. N~2~O.",
        "D. NO~2~."
    ]
    render_options_no_table(doc, opts_c2, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Trong phương pháp Ostwald sản xuất nitric acid, giai đoạn đầu là oxi hóa ammonia bằng không khí với xúc tác platin (Pt) ở nhiệt độ 850 - 900 °C: 4NH~3~ + 5O~2~ -(Pt, t°)-> 4NO + 6H~2~O.")

    # Câu 3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 3 [KNTT - SBT]. Cho các chất khí sau: H~2~S, NO, NO~2~, SO~2~. Số khí gây ô nhiễm môi trường khi phát thải vào không khí là:")
    opts_c3 = [
        "A. 1.",
        "*B. 4.*",
        "C. 3.",
        "D. 2."
    ]
    render_options_no_table(doc, opts_c3, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Cả 4 khí trên đều là những chất gây ô nhiễm không khí nguy hiểm: H~2~S có mùi trứng thối độc tính cao; NO và NO~2~ là tác nhân gây mưa acid và sương mù quang hóa; SO~2~ là khí độc gây mưa acid và kích ứng niêm mạc hô hấp.")

    # Câu 4
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 4 (SBT - KNTT). Oxide của nitrogen được tạo thành khi nguyên tố nitrogen trong nhiên liệu hoặc sinh khối kết hợp với oxygen dư thừa trong không khí được gọi là")
    opts_c4 = [
        "A. NO~x~ tức thời.",
        "B. NO~x~ nhiệt.",
        "*C. NO~x~ nhiên liệu.*",
        "D. NO~x~ tự nhiên."
    ]
    render_options_no_table(doc, opts_c4, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Khi đốt cháy than đá, dầu mỏ hoặc sinh khối thực vật chứa hợp chất hữu cơ của nitrogen, các nguyên tử N bị oxi hóa bởi oxygen dư tạo thành oxide của nitrogen; nguồn này được định danh là NO~x~ nhiên liệu (Fuel NO~x~).")

    # Câu 5
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 5 [CD - SBT]. Phát biểu nào sau đây về đơn chất nitrogen (N~2~) là không đúng?")
    opts_c5 = [
        "*A. Dù phân tử N~2~ có tính kém hoạt động hóa học, nhưng vẫn hoạt động hóa học mạnh hơn chlorine, Cl~2~.*",
        "B. Đơn chất nitrogen không phản ứng với hydrogen, oxygen ở điều kiện thường.",
        "C. Do có nhiệt độ rất thấp nên nitrogen lỏng được sử dụng bảo quản một số loại mẫu vật.",
        "D. Trong bầu khí quyển, khi có sấm chớp, khí nitrogen tạo các nitrogen oxide, là một nguyên nhân làm cho nước mưa có tính acid."
    ]
    render_options_no_table(doc, opts_c5, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Phát biểu A sai vì phân tử N~2~ có liên kết ba N≡N rất bền vững (năng lượng liên kết 945 kJ/mol) nên ở nhiệt độ thường N~2~ rất trơ, kém hoạt động hơn rất nhiều so với khí chlorine (Cl~2~ phản ứng mãnh liệt với nhiều chất ở nhiệt độ thường).")

    # Câu 6
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 6 (SBT - KNTT). Ở điều kiện thích hợp, sulfur dioxide đóng vai trò là chất oxi hoá khi tham gia phản ứng với chất nào sau đây?")
    opts_c6 = [
        "A. NO~2~.",
        "*B. H~2~S.*",
        "C. NaOH.",
        "D. Ca(OH)~2~."
    ]
    render_options_no_table(doc, opts_c6, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Trong phản ứng: SO~2~ + 2H~2~S ➔ 3S + 2H~2~O, số oxi hóa của sulfur trong SO~2~ giảm từ +4 về 0, do đó SO~2~ đóng vai trò là chất oxi hóa (H~2~S là chất khử).")

    # Câu 7
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 7 (SBT - CTST). Điều nào sau đây đúng về tính chất hóa học của dung dịch H~2~SO~4~ đặc?")
    opts_c7 = [
        "A. Dung dịch H~2~SO~4~ đặc có tính khử mạnh.",
        "*B. Dung dịch H~2~SO~4~ đặc có tính oxi hóa mạnh.*",
        "C. Dung dịch H~2~SO~4~ đặc vừa có tính khử, vừa có tính oxi hóa.",
        "D. Dung dịch H~2~SO~4~ đặc không có tính khử, không có tính oxi hóa."
    ]
    render_options_no_table(doc, opts_c7, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Trong phân tử H~2~SO~4~, nguyên tử sulfur có số oxi hóa cao nhất là +6 nên dung dịch H~2~SO~4~ đặc có tính oxi hóa rất mạnh (oxi hóa hầu hết các kim loại trừ Au, Pt và nhiều phi kim, hợp chất). Ngoài ra nó còn có tính háo nước mạnh.")

    # Câu 8
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 8 (SBT - KNTT). Xét phản ứng trong quá trình tạo ra NO~x~ nhiệt:\nN~2~(g) + O~2~(g) ➔ 2NO(g)   Δ~r~H°~298~ = 180,6 kJ\nNhiệt tạo thành chuẩn của NO(g) là")
    opts_c8 = [
        "A. 180,6 kJ/mol.",
        "B. -180,6 kJ/mol.",
        "C. -90,3 kJ/mol.",
        "*D. 90,3 kJ/mol.*"
    ]
    render_options_no_table(doc, opts_c8, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Phản ứng trên tạo thành 2 mol NO(g) từ các đơn chất bền ở điều kiện chuẩn:\nΔ~r~H°~298~ = 2.Δ~f~H°~298~(NO, g) = 180,6 kJ ➔ Δ~f~H°~298~(NO, g) = 180,6 / 2 = 90,3 kJ/mol.")

    # Câu 9
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 9 (SBT - KNTT). Xét cân bằng hóa học: N~2~(g) + 3H~2~(g) ⇌ 2NH~3~(g)   Δ~r~H°~298~ < 0.\nHiệu suất phản ứng khi hệ đạt cân bằng ở nhiệt độ 400 °C và 500 °C lần lượt bằng x% và y%. Mối quan hệ giữa x và y là")
    opts_c9 = [
        "A. x < y.",
        "B. x = y.",
        "*C. x > y.*",
        "D. 5x = 4y."
    ]
    render_options_no_table(doc, opts_c9, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Phản ứng thuận là phản ứng tỏa nhiệt (Δ~r~H°~298~ < 0). Khi tăng nhiệt độ từ 400 °C lên 500 °C, theo nguyên lý Le Chatelier cân bằng chuyển dịch theo chiều thu nhiệt (chiều nghịch), làm giảm lượng NH~3~ tạo thành, do đó hiệu suất giảm ➔ x > y.")

    # Câu 10
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 10. Muốn cho cân bằng của phản ứng tổng hợp ammonia chuyển dịch sang phải cần đồng thời:\nN~2~ + 3H~2~ ⇌ 2NH~3~ ; Δ~r~H°~298~ < 0")
    opts_c10 = [
        "A. Tăng áp suất và tăng nhiệt độ",
        "*B. Tăng áp suất và giảm nhiệt độ*",
        "C. Giảm áp suất và giảm nhiệt độ",
        "D. Giảm áp suất và tăng nhiệt độ"
    ]
    render_options_no_table(doc, opts_c10, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "- Về áp suất: Phía chất phản ứng có 4 mol khí, phía sản phẩm có 2 mol khí. Khi tăng áp suất, cân bằng chuyển dịch theo chiều làm giảm số mol khí (chiều thuận sang phải).\n- Về nhiệt độ: Phản ứng tỏa nhiệt (Δ~r~H°~298~ < 0), khi giảm nhiệt độ cân bằng chuyển dịch theo chiều tỏa nhiệt (chiều thuận sang phải).")

    # Câu 11
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 11. Hợp chất vô cơ X có các tính chất: X tinh khiết là chất lỏng không màu, bốc khói mạnh trong không khí ẩm; dung dịch X đặc để lâu có màu vàng; tan tốt trong nước; có tính oxi hóa mạnh. Chất X là")
    opts_c11 = [
        "A. H~3~PO~4~.",
        "*B. HNO~3~.*",
        "C. H~2~SO~4~.",
        "D. HCl."
    ]
    render_options_no_table(doc, opts_c11, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Nitric acid (HNO~3~) tinh khiết là chất lỏng không màu, bốc khói mạnh trong không khí ẩm. Dưới tác dụng của ánh sáng, HNO~3~ bị phân hủy một phần sinh ra khí NO~2~ có màu nâu đỏ tan lại trong acid làm dung dịch ngả sang màu vàng: 4HNO~3~ -(as)-> 4NO~2~ + O~2~ + 2H~2~O.")

    # Câu 12
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 12 [KNTT - SBT]. Cho cân bằng hóa học sau: 2SO~2~(g) + O~2~(g) ⇌ 2SO~3~(g)   Δ~r~H°~298~ < 0.\nKhi tăng nhiệt độ thì")
    opts_c12 = [
        "A. Tổng số mol khí trong hệ giảm.",
        "B. Hiệu suất phản ứng tăng.",
        "*C. Cân bằng chuyển dịch theo chiều nghịch.*",
        "D. Nồng độ khí sản phẩm tăng."
    ]
    render_options_no_table(doc, opts_c12, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Vì phản ứng theo chiều thuận tỏa nhiệt (Δ~r~H°~298~ < 0), khi tăng nhiệt độ thì theo nguyên lý Le Chatelier cân bằng sẽ chuyển dịch theo chiều thu nhiệt (chiều nghịch).")

    # Câu 13
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 13 [KNTT - SBT]. Trong công nghiệp sản xuất sulfuric acid, hai nguồn nguyên liệu được khai thác từ mỏ để cung cấp nguyên tố lưu huỳnh là")
    opts_c13 = [
        "A. ZnS, PbS",
        "B. H~2~S, SO~2~",
        "C. CaSO~4~, BaSO~4~",
        "*D. S, FeS~2~*"
    ]
    render_options_no_table(doc, opts_c13, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Hai nguồn nguyên liệu mỏ chính cung cấp sulfur để sản xuất sulfuric acid trong công nghiệp là sulfur tự do (S) và quặng pyrite sắt (FeS~2~).")

    # Câu 14
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 14. Nếu ta dẫn khí SO~2~ từ từ đến dư vào dung dịch bromine sẽ có phản ứng hóa học xảy ra kèm hiện tượng gì sau đây?")
    opts_c14 = [
        "A. Dung dịch không đổi màu.",
        "B. Dung dịch bị vẩn đục do tạo kết tủa và kết tủa không tan.",
        "C. Dung dịch sẽ có màu đỏ nâu đậm dần lên hơn lúc ban đầu.",
        "*D. Dung dịch sẽ nhạt màu đỏ nâu rồi mất hẳn màu.*"
    ]
    render_options_no_table(doc, opts_c14, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Khí SO~2~ có tính khử mạnh tác dụng với dung dịch nước bromine màu đỏ nâu, tạo thành các acid không màu:\nSO~2~ + Br~2~ + 2H~2~O ➔ 2HBr + H~2~SO~4~, làm dung dịch mất màu hoàn toàn.")

    # Câu 15
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 15. Hậu quả chính của hiện tượng phú dưỡng đối với môi trường nước là gì?")
    opts_c15 = [
        "A. Làm tăng độ trong suốt của nước và giúp cải thiện hệ sinh thái thủy sinh.",
        "*B. Làm giảm lượng oxygen hòa tan trong nước, gây ra hiện tượng cá chết hàng loạt và phá hủy hệ sinh thái.*",
        "C. Tăng nồng độ oxygen hòa tan trong nước, giúp thủy sinh vật phát triển mạnh mẽ hơn.",
        "D. Giảm sự phát triển của vi sinh vật trong nước."
    ]
    render_options_no_table(doc, opts_c15, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Khi xảy ra hiện tượng phú dưỡng, tảo sinh sôi bùng phát trên mặt nước làm giảm ánh sáng xuống tầng dưới. Khi tảo chết, vi khuẩn phân hủy xác tảo tiêu thụ lượng lớn oxygen hòa tan (DO), khiến nguồn nước bị cạn kiệt oxy làm cá tôm chết ngạt hàng loạt.")

    # Câu 16
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 16. Loại khí nào dưới đây là nguyên nhân chính gây ra mưa acid?")
    opts_c16 = [
        "A. Khí CO~2~",
        "B. Khí O~2~",
        "*C. Khí SO~2~ và NO~x~*",
        "D. Khí CH~4~"
    ]
    render_options_no_table(doc, opts_c16, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Khí sulfur dioxide (SO~2~) và các oxide của nitrogen (NO~x~) phát thải từ nhà máy nhiệt điện, công nghiệp và xe cộ là nguyên nhân chủ yếu tạo thành H~2~SO~4~ và HNO~3~ trong các cơn mưa acid.")

    # Câu 17
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 17. Khu vực nào có nguy cơ bị ảnh hưởng nghiêm trọng bởi mưa acid?")
    opts_c17 = [
        "A. Những khu vực có nhiều cây cối và hệ sinh thái phong phú.",
        "*B. Những khu vực gần các nhà máy điện, nhà máy công nghiệp và giao thông dày đặc.*",
        "C. Những khu vực có độ cao lớn.",
        "D. Những khu vực gần biển."
    ]
    render_options_no_table(doc, opts_c17, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(doc, "Những vùng đô thị phát triển, khu công nghiệp và các nhà máy nhiệt điện đốt than là những nơi phát thải tập trung lượng lớn SO~2~ và NO~x~ vào bầu khí quyển, nên có nguy cơ hứng chịu các trận mưa acid nặng nề nhất.")

    # Câu 18
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 18. Cho các phát biểu sau:\n(1) Phương pháp tiếp xúc gồm ba công đoạn chính sản xuất: SO~2~, SO~3~, H~2~SO~4~.\n(2) Ở giai đoạn tạo ra SO~3~, người ta chọn điều kiện phản ứng ở nhiệt độ thấp.\n(3) Nhờ có chất xúc tác nên phản ứng giữa SO~2~ và O~2~ ưu tiên diễn ra theo chiều thuận.\n(4). Khí sulfur trioxide đi từ dưới lên đỉnh tháp, sulfuric acid đặc chảy từ đỉnh tháp xuống dưới.\n(5) Tương tác van der Waals giữa các phân tử nitrogen với nước rất yếu nên ít tan trong nước.\n(6) Phân tử nitrogen nặng hơn không khí và không phân cực.\n(7) Nitrogen lỏng ở nhiệt độ thấp, -196 °C và có tính trơ.\nSố phát biểu đúng là:")
    opts_c18 = [
        "A. 3.",
        "*B. 4.*",
        "C. 5.",
        "D. 6."
    ]
    render_options_no_table(doc, opts_c18, is_teacher=is_teacher)
    if is_teacher:
        sol_c18 = (
            "Có 4 phát biểu đúng gồm: (1), (4), (5), (7).\n"
            "- (1) Đúng: 3 công đoạn chính là sản xuất SO~2~ ➔ oxi hóa SO~2~ thành SO~3~ ➔ hấp thụ SO~3~ tạo H~2~SO~4~.\n"
            "- (2) Sai: Thực tế người ta tiến hành ở nhiệt độ khoảng 450 °C để đảm bảo tốc độ phản ứng đủ nhanh.\n"
            "- (3) Sai: Chất xúc tác chỉ làm tăng tốc độ đạt cân bằng chứ không làm thay đổi vị trí cân bằng.\n"
            "- (4) Đúng: Chuyển động ngược dòng giữa khí SO~3~ đi lên và acid chảy xuống giúp tăng tối đa bề mặt tiếp xúc và hiệu suất hấp thụ.\n"
            "- (5) Đúng: Phân tử N~2~ không phân cực, liên kết van der Waals với nước rất yếu nên rất ít tan.\n"
            "- (6) Sai: M(N~2~) = 28 nhẹ hơn khối lượng mol trung bình không khí (29).\n"
            "- (7) Đúng: Nitrogen lỏng có nhiệt độ sôi -196 °C và rất trơ về mặt hóa học."
        )
        add_solution_no_table(doc, sol_c18)

    # ==========================================
    # PHẦN II
    # ==========================================
    add_section_header(
        doc,
        title="PHẦN II. Câu trắc nghiệm đúng sai.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai."
    )

    # Câu 1
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 1. Hợp chất có công thức hóa học NH~4~NO~3~ được giới chức quốc gia Lebanon xác định là nguyên nhân gây ra vụ nổ thảm khốc ở thủ đô Beirut vào ngày 04/08/2020. Tia lửa hàn trong quá trình sửa chữa nhà kho có thể đã làm 2 750 tấn NH~4~NO~3~ cất trữ phát nổ, phá hủy nhiều nhà cửa, dẫn đến nhiều người thiệt mạng.\nNH~4~NO~3~ -(t°)-> N~2~O + 2H~2~O   Δ~r~H°~298~ = -35,9 kJ")
    stmts_p2_c1 = [
        ("a) Ammonium nitrate khi ở nhiệt độ cao bị phân hủy thành khí N~2~O và hơi nước, là một phản ứng thu nhiệt và năng lượng lớn.", "SAI"),
        ("b) Khi phản ứng nổ xảy ra, năng lượng được giải phóng một cách đột ngột dưới áp lực rất cao, tăng nhanh, còn được gọi là sóng nổ hoặc sóng xung kích.", "ĐÚNG"),
        ("c) Sóng xung kích gây ra thiệt hại lớn cho môi trường xung quanh. Ammonium nitrate có thể tự phân hủy qua thời gian.", "ĐÚNG"),
        ("d) Tia lửa hàn trong quá trình sửa chữa nhà kho đã khơi mào phản ứng phân hủy ammonium nitrate gây nổ.", "ĐÚNG")
    ]
    for stmt, ans in stmts_p2_c1:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.left_indent = Cm(0.5)
        p_st.paragraph_format.space_before = Pt(1)
        p_st.paragraph_format.space_after = Pt(1)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        sol_p2_c1 = (
            "a) Sai vì Δ~r~H°~298~ = -35,9 kJ < 0 nên phản ứng phân hủy ammonium nitrate là phản ứng tỏa nhiệt mạnh.\n"
            "b) Đúng: Sự giải phóng thể tích khí khổng lồ và nhiệt năng cực nhanh tạo nên sóng áp suất xung kích khủng khiếp.\n"
            "c) Đúng: Sóng xung kích tàn phá các công trình vật chất; NH~4~NO~3~ bảo quản lâu ngày không đúng cách có thể tự phân hủy sinh nhiệt tích tụ.\n"
            "d) Đúng: Nguồn nhiệt từ tia lửa hàn đã kích hoạt phản ứng phân hủy tỏa nhiệt nhanh của khối phân bón cất trữ."
        )
        add_solution_no_table(doc, sol_p2_c1)

    # Câu 2
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 2. Nguồn nào sau đây phát sinh oxide của nitrogen trong không khí?")
    stmts_p2_c2 = [
        ("a) Đốt cháy các nhiên liệu hóa thạch.", "ĐÚNG"),
        ("b) Quá trình sản xuất và sử dụng nitric acid.", "ĐÚNG"),
        ("c) Trong khí quyển khi có sấm sét.", "ĐÚNG"),
        ("d) Quá trình quang hợp của cây xanh.", "SAI")
    ]
    for stmt, ans in stmts_p2_c2:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.left_indent = Cm(0.5)
        p_st.paragraph_format.space_before = Pt(1)
        p_st.paragraph_format.space_after = Pt(1)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        sol_p2_c2 = (
            "a) Đúng: Quá trình đốt than đá, xăng dầu ở nhiệt độ cao sinh ra NO~x~ nhiệt và NO~x~ nhiên liệu.\n"
            "b) Đúng: Khí thải các phân xưởng sản xuất và sử dụng HNO~3~ luôn chứa một lượng NO~x~.\n"
            "c) Đúng: Tia sét phóng điện tạo nhiệt độ hàng nghìn độ C kết hợp N~2~ và O~2~ tạo NO.\n"
            "d) Sai: Quá trình quang hợp của cây xanh hấp thụ CO~2~ và giải phóng O~2~, không sinh ra NO~x~."
        )
        add_solution_no_table(doc, sol_p2_c2)

    # Câu 3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 3. Sulfur dioxide là một chất có nhiều ứng dụng trong công nghiệp (dùng để sản xuất sulfuric acid, tẩy trắng bột giấy công nghiệp giấy, tẩy trắng dung dịch đường trong sản xuất đường tinh luyện,...) và giúp ngăn cản sự phát triển của một số loại vi khuẩn và nấm gây hại cho thực phẩm. Ở áp suất 1 bar và nhiệt độ 25 °C, phản ứng giữa 1 mol sulfur với oxygen xảy ra theo phương trình:\nS(s) + O~2~(g) ➔ SO~2~(g) và tỏa ra một lượng nhiệt là 296,9 kJ.")
    add_centered_picture(doc, os.path.join(IMG_DIR, "cau3_p2_so2_ungdung.png"), width_cm=11.5)
    stmts_p2_c3 = [
        ("a) Biến thiên enthalpy chuẩn của phản ứng là 296,9 kJ.", "SAI"),
        ("b) Enthalpy tạo thành chuẩn của sulfur dioxide bằng -296,9 kJ·mol^-1^.", "ĐÚNG"),
        ("c) Sulfur dioxide vừa có thể là chất khử vừa có thể là chất oxi hóa, tùy thuộc vào phản ứng mà nó tham gia.", "ĐÚNG"),
        ("d) 0,5 mol sulfur tác dụng hết oxygen giải phóng 148,45 kJ năng lượng dưới dạng nhiệt.", "ĐÚNG")
    ]
    for stmt, ans in stmts_p2_c3:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.left_indent = Cm(0.5)
        p_st.paragraph_format.space_before = Pt(1)
        p_st.paragraph_format.space_after = Pt(1)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        sol_p2_c3 = (
            "a) Sai: Vì phản ứng tỏa nhiệt nên biến thiên enthalpy chuẩn phải mang dấu âm: Δ~r~H°~298~ = -296,9 kJ.\n"
            "b) Đúng: Vì 1 mol SO~2~(g) được tạo thành từ các đơn chất bền ở điều kiện chuẩn nên Δ~f~H°~298~(SO~2~, g) = -296,9 kJ/mol.\n"
            "c) Đúng: Nguyên tử S trong SO~2~ có mức oxi hóa trung gian (+4), nên vừa có thể tăng lên +6 (thể hiện tính khử) vừa có thể giảm xuống 0 hoặc -2 (thể hiện tính oxi hóa).\n"
            "d) Đúng: 0,5 mol S phản ứng giải phóng nhiệt lượng: 0,5 . 296,9 = 148,45 kJ."
        )
        add_solution_no_table(doc, sol_p2_c3)

    # Câu 4
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 4. Sulfuric acid có hoạt tính hóa học rất mạnh:")
    stmts_p2_c4 = [
        ("a) Nhôm (aluminum) bị thụ động trong dung dịch H~2~SO~4~ loãng nguội.", "SAI"),
        ("b) Hỗn hợp gồm Cu và Fe~2~O~3~ có tỉ lệ mol 1:1 tan hết trong dung dịch H~2~SO~4~ loãng.", "ĐÚNG"),
        ("c) Khi cho H~2~SO~4~ đặc vào đường kính trắng (C~12~H~22~O~11~), khí thoát ra làm mất màu dung dịch KMnO~4~.", "ĐÚNG"),
        ("d) H~2~SO~4~ đặc có thể dùng để làm khô các khí CO~2~, Cl~2~, N~2~, H~2~S.", "SAI")
    ]
    for stmt, ans in stmts_p2_c4:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.left_indent = Cm(0.5)
        p_st.paragraph_format.space_before = Pt(1)
        p_st.paragraph_format.space_after = Pt(1)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        sol_p2_c4 = (
            "a) Sai: Kim loại Al chỉ bị thụ động hóa trong H~2~SO~4~ đặc nguội (và HNO~3~ đặc nguội), trong dung dịch H~2~SO~4~ loãng nguội nhôm tan nhanh giải phóng khí H~2~.\n"
            "b) Đúng: Fe~2~O~3~ + 3H~2~SO~4~ ➔ Fe~2~(SO~4~)~3~ + 3H~2~O; sau đó Cu + Fe~2~(SO~4~)~3~ ➔ CuSO~4~ + 2FeSO~4~. Do tỉ lệ mol 1:1 nên hỗn hợp tan hết tạo dung dịch trong suốt.\n"
            "c) Đúng: H~2~SO~4~ đặc oxi hóa carbon sinh ra khí SO~2~; khí SO~2~ thoát ra có tính khử làm mất màu thuốc tím KMnO~4~.\n"
            "d) Sai: H~2~SO~4~ đặc là chất oxi hóa mạnh sẽ phản ứng với khí khử H~2~S (H~2~SO~4~ + H~2~S ➔ S + SO~2~ + H~2~O), nên không thể dùng làm khô H~2~S."
        )
        add_solution_no_table(doc, sol_p2_c4)

    # ==========================================
    # PHẦN III
    # ==========================================
    add_section_header(
        doc,
        title="PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 6."
    )

    # Câu 1
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 1. Cho các chất sau: H~2~SO~4~, CO~2~, CO, N~2~, NH~3~. Số chất tan tốt trong nước ở nhiệt độ thường là bao nhiêu?")
    if is_teacher:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.left_indent = Cm(0.5)
        render_rich_text(p_ans, "➔ *Đáp số:* *2*", base_size=10.5)
        add_solution_no_table(doc, "Có 2 chất tan tốt trong nước ở điều kiện thường là H~2~SO~4~ (tan vô hạn và tỏa nhiều nhiệt) và NH~3~ (1 lít nước ở 20 °C hòa tan được khoảng 700 lít NH~3~ nhờ tạo liên kết hydrogen mạnh với nước). Các khí CO, N~2~ rất ít tan; CO~2~ chỉ tan một lượng nhỏ.")
    else:
        add_student_writing_lines(doc, num_lines=2)

    # Câu 2
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 2. Cho các phản ứng:\n(a) S + O~2~ -(t°)-> SO~2~\n(b) S + 3F~2~ ➔ SF~6~\n(c) Hg + S ➔ HgS\n(d) H~2~ + S -(t°)-> H~2~S\nSố phản ứng trong đó sulfur đơn chất đóng vai trò chất oxi hóa là bao nhiêu?")
    if is_teacher:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.left_indent = Cm(0.5)
        render_rich_text(p_ans, "➔ *Đáp số:* *2*", base_size=10.5)
        add_solution_no_table(doc, "Sulfur đóng vai trò chất oxi hóa khi số oxi hóa giảm từ 0 xuống -2. Điều này xảy ra ở phản ứng (c) tạo HgS và phản ứng (d) tạo H~2~S. Ở phản ứng (a) và (b), S tăng số oxi hóa lên +4 và +6 nên là chất khử. Vậy có 2 phản ứng.")
    else:
        add_student_writing_lines(doc, num_lines=2)

    # Câu 3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 3. Cho các chất: HCl (aq hoặc khí), O~2~ (t°), AlCl~3~ (aq), H~2~SO~4~ (aq), NaOH (aq), Na~2~CO~3~ (aq), CuO. Có bao nhiêu chất phản ứng được với NH~3~?")
    if is_teacher:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.left_indent = Cm(0.5)
        render_rich_text(p_ans, "➔ *Đáp số:* *5*", base_size=10.5)
        add_solution_no_table(doc, "Có 5 chất phản ứng được với NH~3~ gồm:\n(1) HCl: NH~3~ + HCl ➔ NH~4~Cl\n(2) O~2~: 4NH~3~ + 3O~2~ -(t°)-> 2N~2~ + 6H~2~O\n(3) AlCl~3~: 3NH~3~ + 3H~2~O + AlCl~3~ ➔ Al(OH)~3~↓ + 3NH~4~Cl\n(4) H~2~SO~4~: 2NH~3~ + H~2~SO~4~ ➔ (NH~4~)~2~SO~4~\n(5) CuO: 2NH~3~ + 3CuO -(t°)-> 3Cu + 3H~2~O + N~2~.\nCác chất kiềm như NaOH, Na~2~CO~3~ không phản ứng với NH~3~.")
    else:
        add_student_writing_lines(doc, num_lines=2)

    # Câu 4
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 4. Cho hỗn hợp khí (X) gồm N~2~, H~2~, NH~3~ có tỉ khối so với khí hydrogen là 8. Dẫn hỗn hợp khí (X) đi qua dung dịch H~2~SO~4~ đặc, dư thì thể tích khí còn lại một nửa. Thành phần % theo thể tích của NH~3~ trong hỗn hợp (X) là x%. Xác định giá trị x.")
    if is_teacher:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.left_indent = Cm(0.5)
        render_rich_text(p_ans, "➔ *Đáp số:* *50*", base_size=10.5)
        add_solution_no_table(doc, "Khi dẫn hỗn hợp (X) qua H~2~SO~4~ đặc dư, chỉ có khí NH~3~ bị hấp thụ hoàn toàn theo phản ứng:\n2NH~3~ + H~2~SO~4~ ➔ (NH~4~)~2~SO~4~.\nDo thể tích khí còn lại một nửa nên thể tích khí NH~3~ bị hấp thụ đúng bằng một nửa thể tích hỗn hợp ban đầu:\n%V(NH~3~) = (V_NH3 / V_hh) . 100% = 50% ➔ x = 50.")
    else:
        add_student_writing_lines(doc, num_lines=2)

    # Câu 5
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 5. Tại nhiều làng nghề thủ công mĩ nghệ, sulfur dioxide được dùng là chất chống mốc cho các sản phẩm mây tre đan. Trong một ngày, một làng nghề đốt cháy 20 kg sulfur để tạo thành sulfur dioxide. Giả thiết có 20% lượng khí SO~2~ trên bay vào khí quyển và chuyển hoá hết thành H~2~SO~4~ trong nước mưa theo sơ đồ:\nSO~2~ -(+O~2~, xt)-> SO~3~ -(+H~2~O)-> H~2~SO~4~\nThể tích nước mưa là x nghìn m^3^ bị nhiễm acid nếu nồng độ H~2~SO~4~ trong nước mưa là 1,25·10^-5^ M. Xác định giá trị của x?")
    if is_teacher:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.left_indent = Cm(0.5)
        render_rich_text(p_ans, "➔ *Đáp số:* *10*", base_size=10.5)
        add_solution_no_table(doc, "Số mol sulfur đốt cháy: n_S = 20 000 / 32 = 625 mol ➔ n(SO~2~ tạo thành) = 625 mol.\nLượng SO~2~ bay vào khí quyển: n(SO~2~ thoát) = 0,20 . 625 = 125 mol.\nTheo sơ đồ chuyển hóa: n(H~2~SO~4~) = n(SO~2~) = 125 mol.\nThể tích nước mưa bị nhiễm acid:\nV = n / C_M = 125 / (1,25 . 10^-5^) = 10 000 000 L = 10 000 m^3^ = 10 nghìn m^3^ ➔ x = 10.")
    else:
        add_student_writing_lines(doc, num_lines=3)

    # Câu 6
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 6. Tại một nhà máy phân bón, ammophos được sản xuất từ ammonia và phosphoric acid, thu được NH~4~H~2~PO~4~ và (NH~4~)~2~HPO~4~ với tỉ lệ mol là 1:1. Cho khí ammonia tác dụng vừa đủ với dung dịch chứa 5,88 tấn phosphoric acid. Tính khối lượng ammophos thu được (theo đơn vị tấn)?")
    if is_teacher:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.left_indent = Cm(0.5)
        render_rich_text(p_ans, "➔ *Đáp số:* *7,41*", base_size=10.5)
        add_solution_no_table(doc, "Số mol H~3~PO~4~ = 5,88 . 10^6^ / 98 = 60 000 mol.\nGọi số mol của NH~4~H~2~PO~4~ và (NH~4~)~2~HPO~4~ lần lượt là a và a mol (tỉ lệ 1:1):\nBảo toàn nguyên tố P: a + a = 60 000 ➔ a = 30 000 mol.\nSố mol NH~3~ cần dùng = a + 2a = 3a = 90 000 mol.\nKhối lượng NH~3~ đã phản ứng: m(NH~3~) = 90 000 . 17 = 1 530 000 g = 1,53 tấn.\nTheo định luật bảo toàn khối lượng:\nm(ammophos) = m(H~3~PO~4~) + m(NH~3~) = 5,88 + 1,53 = 7,41 tấn.")
    else:
        add_student_writing_lines(doc, num_lines=3)

    # ==========================================
    # PHẦN IV
    # ==========================================
    add_section_header(
        doc,
        title="PHẦN IV. Tự luận",
        subtitle="Học sinh giải chi tiết các câu hỏi sau vào giấy làm bài."
    )

    # Câu 1
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 1 (SBT - CD). Vàng tan trong hỗn hợp gồm dung dịch nitric acid đặc và dung dịch hydrochloric acid đặc (tỉ lệ 1 : 3 về thể tích) tạo ra hợp chất tan của Au^3+^ theo phản ứng sau:\nAu + HNO~3~ + HCl ➔ HAuCl~4~ + H~2~O + NO\na) Cân bằng phương trình hoá học của phản ứng trên theo phương pháp thăng bằng electron.\nb) Cho biết acid nào đóng vai trò chất oxi hoá trong phản ứng trên. Giải thích.")
    if is_teacher:
        sol_p4_c1 = (
            "a) Cân bằng theo phương pháp thăng bằng electron:\n"
            "   Quá trình oxi hóa: Au^0^ + 4Cl^- ➔ Au^+3^Cl~4~^- + 3e  | x 1\n"
            "   Quá trình khử:     N^+5^ + 3e ➔ N^+2^ (trong NO)       | x 1\n"
            "   Phương trình hóa học hoàn chỉnh:\n"
            "   Au + HNO~3~ + 4HCl ➔ HAuCl~4~ + 2H~2~O + NO↑\n"
            "b) Trong phản ứng trên, HNO~3~ đóng vai trò là chất oxi hóa vì số oxi hóa của nguyên tử nitrogen giảm từ +5 (trong HNO~3~) xuống +2 (trong khí NO). Acid HCl đóng vai trò môi trường và chất tạo phức."
        )
        add_solution_no_table(doc, sol_p4_c1)
    else:
        add_student_writing_lines(doc, num_lines=5)

    # Câu 2
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_question_prompt(p, "Câu 2 (SBT - CTST). Đặt hai cốc (A), (B) có khối lượng bằng nhau lên 2 đĩa cân, cân ở vị trí thăng bằng. Cho 120 gam hỗn hợp potassium hydrogencarbonate và sodium hydrogencarbonate vào cốc (A); 85 gam silver nitrate vào cốc (B). Thêm từ từ 100 gam dung dịch sulfuric acid 19,6% vào cốc (A); 100 gam dung dịch hydrochloric acid 36,5% vào cốc (B). Sau thí nghiệm cân có ở vị trí thăng bằng không? Nếu cân không ở vị trí thăng bằng thì cần thêm bao nhiêu gam dung dịch hydrochloric acid 36,5% vào cốc nào để cân trở lại vị trí thăng bằng? Giả thiết khí CO~2~ không tan trong nước, bỏ qua quá trình bay hơi của nước và hydrogen chloride.")
    if is_teacher:
        sol_p4_c2 = (
            "Hướng dẫn giải chi tiết:\n"
            "- Xét cốc (A): Đặt công thức chung của hai muối hydrogencarbonate là RHCO~3~ (M = 84 đến 100).\n"
            "  Số mol tối đa của RHCO~3~ (khi chỉ có NaHCO~3~): n(RHCO~3~) = 120 / 84 ≈ 1,43 mol.\n"
            "  Số mol H~2~SO~4~: n(H~2~SO~4~) = (100 . 19,6%) / 98 = 0,2 mol.\n"
            "  Phương trình phản ứng: 2RHCO~3~ + H~2~SO~4~ ➔ R~2~SO~4~ + 2CO~2~↑ + 2H~2~O\n"
            "  Ta có: n(RHCO~3~)/2 > 0,7 > n(H~2~SO~4~) = 0,2 nên RHCO~3~ dư, H~2~SO~4~ phản ứng hết.\n"
            "  Số mol khí CO~2~ thoát ra: n(CO~2~) = 2 . n(H~2~SO~4~) = 0,4 mol.\n"
            "  Khối lượng cốc (A) sau phản ứng:\n"
            "  m_A = m(muối) + m(dd H~2~SO~4~) - m(CO~2~) = 120 + 100 - 0,4 . 44 = 202,4 gam.\n"
            "- Xét cốc (B):\n"
            "  Số mol AgNO~3~: n(AgNO~3~) = 85 / 170 = 0,5 mol.\n"
            "  Số mol HCl: n(HCl) = (100 . 36,5%) / 36,5 = 1,0 mol.\n"
            "  Phương trình phản ứng: AgNO~3~ + HCl ➔ AgCl↓ + HNO~3~\n"
            "  Vì n(AgNO~3~) = 0,5 mol < n(HCl) = 1,0 mol nên HCl dư, AgNO~3~ phản ứng hết.\n"
            "  Do kết tủa AgCl không bay hơi và không có chất khí nào thoát ra khỏi cốc (B), nên:\n"
            "  m_B = m(AgNO~3~) + m(dd HCl) = 85 + 100 = 185 gam.\n"
            "- Kết luận và tính toán cân bằng:\n"
            "  Vì m_A = 202,4 gam > m_B = 185 gam nên sau thí nghiệm cân không ở vị trí thăng bằng (cốc A nặng hơn cốc B).\n"
            "  Để cân trở lại vị trí thăng bằng, cần thêm vào cốc (B) một lượng dung dịch HCl 36,5% có khối lượng bằng:\n"
            "  m(dd HCl 36,5% cần thêm) = m_A - m_B = 202,4 - 185 = 17,4 gam."
        )
        add_solution_no_table(doc, sol_p4_c2)
    else:
        add_student_writing_lines(doc, num_lines=8)

    # Lưu file
    filename = "De_Kiem_Tra_Hoa_11_De06_GiaoVien.docx" if is_teacher else "De_Kiem_Tra_Hoa_11_De06.docx"
    filepath = os.path.join(OUT_DIR, filename)
    doc.save(filepath)
    print(f"Saved: {filepath}")

if __name__ == "__main__":
    print("Generating De_Kiem_Tra_Hoa_11_De06.docx (Học Sinh)...")
    build_exam_document(is_teacher=False)
    print("Generating De_Kiem_Tra_Hoa_11_De06_GiaoVien.docx (Giáo Viên)...")
    build_exam_document(is_teacher=True)
    print("Done!")
