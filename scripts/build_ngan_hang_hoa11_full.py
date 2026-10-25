# -*- coding: utf-8 -*-
"""
HỆ THỐNG XÂY DỰNG TRỌN BỘ NGÂN HÀNG CÂU HỎI HÓA HỌC 11 (CHƯƠNG 1 -> CHƯƠNG 6 & CHUYÊN ĐỀ HỌC TẬP)
Tự động tích hợp, phân loại, loại bỏ trùng lặp và xuất bản theo chuẩn "FORM ĐỀ KIỂM TRA":
- 7 Cơ sở dữ liệu JSON từng chương + 01 Master JSON trọn bộ: Ngan_Hang_Hoa_11_Tron_Bo.json
- 14 File Word (.docx) chuẩn gồm bản Học sinh (_HocSinh.docx) và Giáo viên (_GiaoVien.docx)
  đầy đủ Ma trận nhận thức, Header/Footer màu chuẩn (#D32F2F / #8BC390), không dùng bảng trong nội dung.
"""

import os
import io
import re
import json
import sys
from difflib import SequenceMatcher
from PIL import Image
import docx
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

sys.path.insert(0, r'c:\Antigravity_Thanh\scripts')
from build_ngan_hang_hoa10_c1 import (
    iter_block_items, get_cau_num_ids, clean_extracted_text, get_para_text_full,
    extract_real_images_from_el, norm_for_compare, extract_underlined_p1_answer,
    normalize_chem_markers_hoa10, add_chem_runs_hoa10, parse_abcd_from_opt_lines,
    add_options_paragraph_no_table, add_data_table_to_doc, add_solution_block_hoa10,
    setup_page_and_hf, LEVEL_NAMES, COLOR_HF_RED, COLOR_PRIMARY,
    COLOR_SOLUTION, COLOR_CORRECT, COLOR_TAG
)

# Thư mục gốc dữ liệu
DIR_CVL = r"C:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\thay-Trung-3-khoi\H11 - CAO QUOC TRUNG\H11 - CAO QUOC TRUNG"
DIR_ROOT_TRUNG = r"C:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\thay-Trung-3-khoi"
TTB_DOCX = r"C:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\2026 - Phan dang de thi thu truong, so. TTB1-20260711T201135Z-2-001\2026 - Phan dang de thi thu truong, so. TTB1\X. PD 2026 - Hoa 11. CD11 - SGK va CDHT. TTB. ok - DA chi tiet.docx"

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
BANK_DIR = r"c:\Antigravity_Thanh\San_Pham\Ngan_Hang_Hoa_11"
IMG_DIR = r"c:\Antigravity_Thanh\San_Pham\images_hoa11"
os.makedirs(OUT_DIR, exist_ok=True)
os.makedirs(BANK_DIR, exist_ok=True)
os.makedirs(IMG_DIR, exist_ok=True)

CHAPTER_DEFS = {
    1: {
        "code": "C1",
        "name": "CHƯƠNG 1: CÂN BẰNG HÓA HỌC",
        "short_name": "Cân bằng hóa học",
        "file_prefix": "Ngan_Hang_Hoa_11_Chuong_1_Can_Bang_Hoa_Hoc",
        "cvl_files": [os.path.join(DIR_CVL, "2026_H11.C1. CAN BANG HOA HOC (GIAO VIEN).docx")],
        "ttb_range": (2, 430),
        "topics": [
            "Bài 1: Khái niệm về cân bằng hóa học",
            "Bài 2: Cân bằng trong dung dịch nước"
        ]
    },
    2: {
        "code": "C2",
        "name": "CHƯƠNG 2: NITROGEN – SULFUR",
        "short_name": "Nitrogen – Sulfur",
        "file_prefix": "Ngan_Hang_Hoa_11_Chuong_2_Nitrogen_Sulfur",
        "cvl_files": [os.path.join(DIR_CVL, "2026_H11.C2. NITROGEN-SULFUR (GIAO VIEN).docx")],
        "ttb_range": (430, 915),
        "topics": [
            "Bài 3: Đơn chất nitrogen",
            "Bài 4: Ammonia và một số hợp chất ammonium",
            "Bài 5: Một số hợp chất của nitrogen với oxygen",
            "Bài 6: Sulfur và sulfur dioxide",
            "Bài 7: Sulfuric acid và muối sulfate"
        ]
    },
    3: {
        "code": "C3",
        "name": "CHƯƠNG 3: ĐẠI CƯƠNG HÓA HỌC HỮU CƠ",
        "short_name": "Đại cương hóa học hữu cơ",
        "file_prefix": "Ngan_Hang_Hoa_11_Chuong_3_Dai_Cuong_Hoa_Huu_Co",
        "cvl_files": [os.path.join(DIR_CVL, "2026.H11.C3. ĐẠI CƯƠNG HÓA HỌC HỮU CƠ (GIAO VIEN).docx")],
        "ttb_range": (915, 1120),
        "topics": [
            "Bài 8: Hợp chất hữu cơ và hóa học hữu cơ",
            "Bài 9: Phương pháp tách biệt và tinh chế hợp chất hữu cơ",
            "Bài 10: Công thức phân tử hợp chất hữu cơ",
            "Bài 11: Cấu tạo hóa học hợp chất hữu cơ"
        ]
    },
    4: {
        "code": "C4",
        "name": "CHƯƠNG 4: HYDROCARBON",
        "short_name": "Hydrocarbon",
        "file_prefix": "Ngan_Hang_Hoa_11_Chuong_4_Hydrocarbon",
        "cvl_files": [os.path.join(DIR_CVL, "2026. H11.C4. HYDROCARBON (GIAO VIEN).docx")],
        "ttb_range": (1120, 1308),
        "topics": [
            "Bài 12: Alkane",
            "Bài 13: Hydrocarbon không no (Alkene – Alkyne)",
            "Bài 14: Arene (Hydrocarbon thơm)"
        ]
    },
    5: {
        "code": "C5",
        "name": "CHƯƠNG 5: DẪN XUẤT HALOGEN – ALCOHOL – PHENOL",
        "short_name": "Dẫn xuất halogen – Alcohol – Phenol",
        "file_prefix": "Ngan_Hang_Hoa_11_Chuong_5_Dan_Xuat_Halogen_Alcohol_Phenol",
        "cvl_files": [os.path.join(DIR_ROOT_TRUNG, "2026.H11.C5-DẪN XUẤT HALOGEN - ALCOHOL - PHENOL.docx")],
        "ttb_range": (1308, 1575),
        "topics": [
            "Bài 15: Dẫn xuất halogen",
            "Bài 16: Alcohol",
            "Bài 17: Phenol"
        ]
    },
    6: {
        "code": "C6",
        "name": "CHƯƠNG 6: HỢP CHẤT CARBONYL – CARBOXYLIC ACID",
        "short_name": "Carbonyl – Carboxylic acid",
        "file_prefix": "Ngan_Hang_Hoa_11_Chuong_6_Hop_Chat_Carbonyl_Carboxylic_Acid",
        "cvl_files": [
            os.path.join(DIR_ROOT_TRUNG, "2026. H11.C6 - HỢP CHẤT CARBONYL (GIAO VIEN).docx"),
            os.path.join(DIR_ROOT_TRUNG, "H11_ĐỀ ÔN TẬP CHƯƠNG 6 LẦN 1.docx"),
            os.path.join(DIR_ROOT_TRUNG, "H11_ĐỀ ÔN TẬP CHƯƠNG 6 LẦN 2.docx"),
            os.path.join(DIR_ROOT_TRUNG, "H11_ĐỀ ÔN TẬP CHƯƠNG 6 LẦN 3.docx"),
        ],
        "ttb_range": (1575, 1996),
        "topics": [
            "Bài 18: Hợp chất carbonyl (Aldehyde – Ketone)",
            "Bài 19: Carboxylic acid"
        ]
    },
    7: {
        "code": "CDHT",
        "name": "CHUYÊN ĐỀ HỌC TẬP HÓA HỌC 11",
        "short_name": "Chuyên đề học tập",
        "file_prefix": "Ngan_Hang_Hoa_11_Chuyen_De_Hoc_Tap",
        "cvl_files": [],
        "ttb_range": (1996, 2310),
        "topics": [
            "Chuyên đề 11.1: Phân bón hóa học",
            "Chuyên đề 11.2: Trải nghiệm thực hành hóa học hữu cơ",
            "Chuyên đề 11.3: Dầu mỏ và chế biến dầu mỏ"
        ]
    }
}


def classify_p2_item_level_hoa11(stmt_text):
    t = stmt_text.lower()
    if any(k in t for k in [
        'tính toán', 'hằng số', 'kc', 'ph =', 'nồng độ', 'thể tích', 'mol',
        'chuẩn độ', 'mol/l', 'hiệu suất', 'khối lượng', 'gam', 'tỉ lệ mol',
        'phần trăm', 'đốt cháy', 'hệ phương trình'
    ]) or re.search(r'\d+[,\.]\d+', t):
        return "VD"
    if any(k in t for k in [
        'chuyển dịch', 'le chatelier', 'tăng nhiệt độ', 'giảm áp suất', 'tăng áp suất',
        'thuận nghịch', 'tốc độ', 'thủy phân', 'chất chỉ thị', 'màu sắc', 'ion thu gọn',
        'brønsted', 'lưỡng tính', 'so sánh', 'môi trường', 'nhiệt độ sôi', 'độ tan',
        'liên kết hydrogen', 'đồng phân', 'bậc alcohol', 'phản ứng tráng bạc',
        'iodoform', 'tollens', 'cu(oh)2'
    ]):
        return "TH"
    return "NB"


def classify_question_level_hoa11(part, sec, topic, stem_text, opt_text, sol_text, item_levels=None):
    full = (stem_text + " " + opt_text).lower()
    if sec in ("DANG_1", "DANG_2", "DANG_3", "DANG_4", "TU_LUAN"):
        return "VD"

    if part == "P1":
        if any(k in full for k in [
            'tính', 'kc =', 'tính ph', 'chuẩn độ', 'hiệu suất', 'nồng độ mol',
            'khối lượng', 'gam', 'tỉ khối', 'hấp thụ hoàn toàn', 'thu được m gam',
            'đốt cháy hoàn toàn', 'giá trị của m', 'giá trị của v'
        ]) or re.search(r'\bkc\b\s*=\s*\d+', full) or re.search(r'\d+[,\.]\d+\s*(?:g|gam|ml|l|mol)', full):
            return "VD"
        if any(k in full for k in [
            'chuyển dịch', 'khi tăng', 'khi giảm', 'yếu tố nào', 'áp suất', 'nhiệt độ',
            'màu', 'nhạt dần', 'đậm dần', 'thuận hay nghịch', 'chiều thuận', 'chiều nghịch',
            'ph của dung dịch', 'nhỏ nhất', 'lớn nhất', 'sắp xếp', 'quỳ tím', 'phenolphthalein',
            'thuyết brønsted', 'đóng vai trò', 'lưỡng tính', 'thủy phân', 'số đồng phân',
            'sản phẩm chính', 'quy tắc', 'nhiệt độ sôi cao nhất', 'nhiệt độ sôi giảm dần',
            'thuốc thử để phân biệt', 'hiện tượng quan sát được'
        ]):
            return "TH"
        return "NB"

    elif part == "P2":
        if item_levels:
            vd_cnt = sum(1 for v in item_levels.values() if v == "VD")
            th_cnt = sum(1 for v in item_levels.values() if v == "TH")
            if vd_cnt >= 2 or (vd_cnt >= 1 and th_cnt >= 2):
                return "VD"
            if th_cnt >= 2 or vd_cnt >= 1:
                return "TH"
            return "NB"
        return "TH"

    elif part == "P3":
        if any(k in full for k in ['tính', 'nồng độ', 'thể tích', 'khối lượng', 'hiệu suất', 'kc', 'ph', 'phần trăm', 'gam']):
            return "VD"
        if any(k in full for k in ['có bao nhiêu', 'số phát biểu', 'sắp xếp', 'số chất']):
            return "TH"
        return "TH"

    return "TH"


def infer_topic_by_chapter(ch_num, text):
    t = text.lower()
    ch_info = CHAPTER_DEFS[ch_num]
    topics = ch_info["topics"]

    if ch_num == 1:
        if any(k in t for k in ['ph', 'chuẩn độ', 'dung dịch', 'điện li', 'brønsted', 'acid', 'base', 'thủy phân', 'h+', 'oh-', 'quỳ', 'phenolphthalein', 'muối', 'phèn']):
            return topics[1]
        return topics[0]

    elif ch_num == 2:
        if any(k in t for k in ['h2so4', 'sulfuric', 'sulfate', 'muối sulfate', 'oleum']):
            return topics[4]
        if any(k in t for k in ['sulfur', 'so2', 'dioxide', 'quặng pyrit', 'h2s']):
            return topics[3]
        if any(k in t for k in ['hno3', 'nitric', 'no2', 'mưa acid', 'nitrogen monoxide', 'khí no']):
            return topics[2]
        if any(k in t for k in ['ammonia', 'nh3', 'ammonium', 'nh4+', 'haber']):
            return topics[1]
        return topics[0]

    elif ch_num == 3:
        if any(k in t for k in ['cấu tạo', 'đồng phân', 'nhóm chức', 'hồng ngoại', 'ir']):
            return topics[3]
        if any(k in t for k in ['công thức phân tử', 'ctpt', 'khối phổ', 'ms', 'phân tử khối']):
            return topics[2]
        if any(k in t for k in ['tách biệt', 'tinh chế', 'chưng cất', 'chiết', 'kết tinh', 'sắc kí']):
            return topics[1]
        return topics[0]

    elif ch_num == 4:
        if any(k in t for k in ['benzene', 'toluene', 'arene', 'thơm', 'xylene', 'styrene']):
            return topics[2]
        if any(k in t for k in ['alkene', 'alkyne', 'không no', 'ethene', 'ethylene', 'acetylene', 'bromine', 'markovnikov']):
            return topics[1]
        return topics[0]

    elif ch_num == 5:
        if any(k in t for k in ['phenol', 'picric']):
            return topics[2]
        if any(k in t for k in ['alcohol', 'ethanol', 'methanol', 'glycerol', 'bậc alcohol', 'nhiệt độ sôi của alcohol']):
            return topics[1]
        return topics[0]

    elif ch_num == 6:
        if any(k in t for k in ['acid', 'acetic', 'formic', 'carboxylic', 'carboxyl', 'lên men giấm', 'ch3cooh', 'hcooh']):
            return topics[1]
        return topics[0]

    elif ch_num == 7:
        if any(k in t for k in ['dầu mỏ', 'chế biến', 'cracking', 'reforming', 'octane', 'khí mỏ dầu', 'xăng']):
            return topics[2]
        if any(k in t for k in ['thực hành', 'trải nghiệm', 'thí nghiệm', 'điều chế', 'xà phòng', 'tinh dầu']):
            return topics[1]
        return topics[0]

    return topics[0]


def parse_cvl_doc(fpath, ch_num):
    if not os.path.exists(fpath):
        return []
    doc = docx.Document(fpath)
    cau_nids = get_cau_num_ids(doc)
    blocks = list(iter_block_items(doc))

    start_idx = 0
    end_idx = len(blocks)
    for i, (bt, b) in enumerate(blocks):
        if bt == 'p':
            t = get_para_text_full(b, doc.part)
            if ('HỆ THỐNG BÀI TẬP' in t or 'BÀI TẬP' in t or 'PHẦN I' in t or 'PHẦN 1' in t) and start_idx == 0:
                start_idx = i
            elif t == 'MỤC LỤC' and i > start_idx:
                end_idx = i
                break

    cur_topic = CHAPTER_DEFS[ch_num]["topics"][0]
    cur_sec = ""
    cur_part = "P1"
    questions = []
    cur_q = None

    def flush_q():
        nonlocal cur_q
        if cur_q:
            if cur_q["sec"] in ("LT", "I") and len(cur_q["p2_items"]) != 4:
                pass
            else:
                questions.append(cur_q)
            cur_q = None

    for i in range(start_idx, end_idx):
        bt, b = blocks[i]
        if bt == 'p':
            t = get_para_text_full(b, doc.part)
            imgs = extract_real_images_from_el(b._p, doc.part)
            if not t and not imgs:
                continue

            # Topic heading checks
            t_u = t.upper()
            topics_list = CHAPTER_DEFS[ch_num]["topics"]
            new_topic = None
            if len(t) < 80:
                if ch_num == 1:
                    if t.startswith(('BÀI 1:', 'BÀI 1.')):
                        new_topic = topics_list[0]
                    elif t.startswith(('BÀI 2:', 'BÀI 2.')):
                        new_topic = topics_list[1]
                elif ch_num == 2:
                    if 'PHẦN ①' in t:
                        new_topic = topics_list[0]
                    elif 'PHẦN ②' in t:
                        new_topic = topics_list[1]
                    elif 'PHẦN ③' in t or 'DẠNG 4:' in t:
                        new_topic = topics_list[2]
                    elif 'PHẦN ④' in t or 'TÍNH CHẤT SO2' in t_u:
                        new_topic = topics_list[3]
                    elif 'PHẦN ⑤' in t or 'ĐIỀU CHẾ H2SO4' in t_u:
                        new_topic = topics_list[4]
                    elif '① NITROGEN' in t_u:
                        new_topic = topics_list[1]
                    elif '② SULFUR' in t_u:
                        new_topic = topics_list[3]
                elif ch_num == 3:
                    if t.startswith(('BÀI 10', 'BÀI 10.')):
                        new_topic = topics_list[0]
                    elif t.startswith(('BÀI 11', 'BÀI 11.')):
                        new_topic = topics_list[1]
                    elif t.startswith(('BÀI 12', 'BÀI 12.')):
                        new_topic = topics_list[2]
                    elif t.startswith(('BÀI 13', 'BÀI 13.')):
                        new_topic = topics_list[3]
                elif ch_num == 4:
                    if t.startswith(('BÀI 15', 'HỆ THỐNG BÀI TẬP ALKANE')):
                        new_topic = topics_list[0]
                    elif t.startswith(('BÀI 16', 'HỆ THỐNG BÀI TẬP ALKENE', 'HỆ THỐNG BÀI TẬP HYDROCARBON KHÔNG NO')):
                        new_topic = topics_list[1]
                    elif t.startswith(('BÀI 17', 'ARENE', 'HỆ THỐNG BÀI TẬP ARENE')):
                        new_topic = topics_list[2]
                elif ch_num == 5:
                    if t.startswith(('BÀI 19', 'BÀI 19.')):
                        new_topic = topics_list[0]
                    elif t.startswith(('BÀI 20', 'BÀI 20.')):
                        new_topic = topics_list[1]
                    elif t.startswith(('BÀI 21', 'BÀI 21.')):
                        new_topic = topics_list[2]
                elif ch_num == 6:
                    if t.startswith(('BÀI 23', 'BÀI 23.')):
                        new_topic = topics_list[0]
                    elif t.startswith(('BÀI 24', 'BÀI 24.')):
                        new_topic = topics_list[1]

            if new_topic and new_topic != cur_topic:
                flush_q()
                cur_topic = new_topic

            # Section heading checks
            if 'BÀI TẬP ÔN TẬP LÝ THUYẾT' in t or 'PHẦN 1.' in t or 'PHẦN I.' in t:
                flush_q()
                cur_sec = 'LT'
                cur_part = 'P2'
                continue
            elif ('TRẮC NGHIỆM' in t and 'ĐÚNG SAI' not in t) or 'PHẦN 2.' in t or 'PHẦN II.' in t:
                flush_q()
                cur_sec = 'TN'
                cur_part = 'P1'
                continue
            elif 'ĐÚNG SAI' in t or 'PHẦN 3.' in t or 'PHẦN III.' in t:
                flush_q()
                cur_sec = 'DS'
                cur_part = 'P2'
                continue
            elif 'TRẢ LỜI NGẮN' in t or 'PHẦN 4.' in t or 'PHẦN IV.' in t:
                flush_q()
                cur_sec = 'TLN'
                cur_part = 'P3'
                continue
            elif 'TỰ LUẬN' in t or t.startswith('DẠNG '):
                flush_q()
                cur_sec = 'TU_LUAN'
                cur_part = 'P3'
                continue

            nids = b._p.xpath('.//w:numPr/w:numId/@w:val')
            is_cau = bool((nids and nids[0] in cau_nids) or re.match(r'^Câu\s+\d+[\.:]', t))
            if is_cau:
                flush_q()
                stem_first = re.sub(r'^Câu\s+\d+[\.:]\s*', '', t).strip()
                cur_q = {
                    "topic": cur_topic,
                    "sec": cur_sec,
                    "part": cur_part,
                    "source_tag": "[THPT Châu Văn Liêm – 2026]",
                    "stem_lines": [stem_first] if stem_first else [],
                    "opt_lines": [],
                    "p2_items": {},
                    "sol_lines": [],
                    "underlined_opt": None,
                    "in_sol": False,
                    "img_blobs": list(imgs),
                    "tables": [],
                }
            elif cur_q is not None:
                if imgs and not cur_q["in_sol"]:
                    cur_q["img_blobs"].extend(imgs)
                if not t:
                    continue
                if t.startswith(("📝", "Hướng dẫn giải", "Đáp án và giải thích", "Đáp án:", "Đáp án đúng:", "Lời giải")):
                    cur_q["in_sol"] = True
                    cur_q["sol_lines"].append(t)
                elif cur_q["in_sol"]:
                    cur_q["sol_lines"].append(t)
                else:
                    if cur_part == "P1" and re.match(r'^[A-D]\.', t):
                        cur_q["opt_lines"].append(t)
                        u_ans = extract_underlined_p1_answer([b])
                        if u_ans:
                            cur_q["underlined_opt"] = u_ans
                    elif cur_part == "P2" and re.match(r'^[a-d]\)', t):
                        k = t[0].lower()
                        cur_q["p2_items"][k] = t[2:].strip()
                    else:
                        cur_q["stem_lines"].append(t)

        elif bt == 'tbl':
            if len(b.rows) == 1 and len(b.columns) == 1:
                cell = b.cell(0, 0)
                has_q = False
                for cp in cell.paragraphs:
                    nids = cp._p.xpath('.//w:numPr/w:numId/@w:val')
                    ct = get_para_text_full(cp, doc.part)
                    if (nids and nids[0] in cau_nids) or re.match(r'^Câu\s+\d+[\.:]', ct):
                        has_q = True
                        break
                if has_q:
                    for cp in cell.paragraphs:
                        nids = cp._p.xpath('.//w:numPr/w:numId/@w:val')
                        ct = get_para_text_full(cp, doc.part)
                        if not ct:
                            continue
                        if (nids and nids[0] in cau_nids) or re.match(r'^Câu\s+\d+[\.:]', ct):
                            flush_q()
                            ct_clean = re.sub(r'^Câu\s+\d+[\.:]\s*', '', ct).strip()
                            cur_q = {
                                "topic": cur_topic,
                                "sec": cur_sec,
                                "part": cur_part,
                                "source_tag": "[THPT Châu Văn Liêm – 2026]",
                                "stem_lines": [ct_clean],
                                "opt_lines": [],
                                "p2_items": {},
                                "sol_lines": [],
                                "underlined_opt": None,
                                "in_sol": False,
                                "img_blobs": [],
                                "tables": [],
                            }
                        elif cur_q is not None:
                            if ct.startswith(("📝", "Hướng dẫn giải", "Đáp án", "Lời giải")):
                                cur_q["in_sol"] = True
                                cur_q["sol_lines"].append(ct)
                            elif cur_q["in_sol"]:
                                cur_q["sol_lines"].append(ct)
                            else:
                                cur_q["stem_lines"].append(ct)
                continue

            if cur_q is not None and not cur_q["in_sol"]:
                if len(b.rows) == 1 and len(b.columns) == 2:
                    for c in b.rows[0].cells:
                        for cp in c.paragraphs:
                            ct = get_para_text_full(cp, doc.part)
                            if ct:
                                cur_q["stem_lines"].append(ct)
                else:
                    t_data = []
                    for r in b.rows:
                        row_vals = [get_para_text_full(c.paragraphs[0], doc.part) if c.paragraphs else "" for c in r.cells]
                        if any(row_vals):
                            t_data.append(row_vals)
                    if t_data:
                        cur_q["tables"].append(t_data)

    flush_q()
    return questions


def parse_ttb_section(start_para, end_para, ch_num):
    if not os.path.exists(TTB_DOCX):
        return []
    doc = docx.Document(TTB_DOCX)
    ttb_qs = []
    cur_part = "P1"
    cur_lvl = "NB"
    cur_q = None

    def flush_ttb():
        nonlocal cur_q
        if cur_q:
            ttb_qs.append(cur_q)
            cur_q = None

    for i in range(start_para, min(end_para, len(doc.paragraphs))):
        p = doc.paragraphs[i]
        t = get_para_text_full(p, doc.part)
        imgs = extract_real_images_from_el(p._p, doc.part)
        if not t and not imgs:
            continue
        if "Trắc nghiệm nhiều phương án" in t:
            flush_ttb()
            cur_part = "P1"
            continue
        elif "Trắc nghiệm đúng" in t:
            flush_ttb()
            cur_part = "P2"
            continue
        elif "Trắc nghiệm trả lời ngắn" in t:
            flush_ttb()
            cur_part = "P3"
            continue
        elif "MỨC ĐỘ BIẾT" in t:
            flush_ttb()
            cur_lvl = "NB"
            continue
        elif "MỨC ĐỘ HIỂU" in t:
            flush_ttb()
            cur_lvl = "TH"
            continue
        elif "MỨC ĐỘ VẬN DỤNG" in t:
            flush_ttb()
            cur_lvl = "VD"
            continue

        m_cau = re.match(r'^Câu\s+\d+[\.\:]\s*(.*)', t)
        if m_cau:
            flush_ttb()
            content = m_cau.group(1).strip()
            src_tag = ""
            m_src = re.match(r'^(\[[^\]]+\])\s*(.*)', content)
            if m_src:
                src_tag = m_src.group(1)
                content = m_src.group(2)
            cur_q = {
                "topic": CHAPTER_DEFS[ch_num]["topics"][0],
                "sec": "TTB",
                "part": cur_part,
                "level_override": cur_lvl,
                "source_tag": src_tag or "[Đề thi thử THPT Quốc gia 2026]",
                "stem_lines": [content] if content else [],
                "opt_lines": [],
                "p2_items": {},
                "sol_lines": [],
                "underlined_opt": None,
                "in_sol": False,
                "img_blobs": list(imgs),
                "tables": []
            }
        elif cur_q is not None:
            if imgs and not cur_q["in_sol"]:
                cur_q["img_blobs"].extend(imgs)
            if not t:
                continue
            if t.startswith(("Lời giải", "Hướng dẫn giải", "Đáp án:")):
                cur_q["in_sol"] = True
                cur_q["sol_lines"].append(t)
            elif cur_q["in_sol"]:
                cur_q["sol_lines"].append(t)
            else:
                if cur_part == "P1" and re.match(r'^[A-D]\.', t):
                    cur_q["opt_lines"].append(t)
                elif cur_part == "P2" and re.match(r'^[a-d]\)', t):
                    k = t[0].lower()
                    cur_q["p2_items"][k] = t[2:].strip()
                else:
                    cur_q["stem_lines"].append(t)

    flush_ttb()

    for q in ttb_qs:
        q["topic"] = infer_topic_by_chapter(ch_num, " ".join(q["stem_lines"]) + " " + " ".join(q["opt_lines"]))

    return ttb_qs


def finalize_question_record_full(raw_q, ch_num):
    sol_text = "\n".join(raw_q["sol_lines"]).strip()
    sol_text = re.sub(r'^📝\s*Hướng dẫn giải:?\s*', '', sol_text).strip()
    sol_text = re.sub(r'^Hướng dẫn giải:?\s*', '', sol_text).strip()
    sol_text = re.sub(r'^Đáp án và giải thích:?\s*', '', sol_text).strip()

    ans = ""
    if raw_q["part"] == "P1":
        ans = raw_q["underlined_opt"] or ""
        if not ans:
            m = re.search(r'Đáp\s+(?:án|số)\s*[:\.]?\s*([A-D])\b', sol_text, flags=re.IGNORECASE)
            if m:
                ans = m.group(1).upper()
        if not ans:
            for ln in raw_q["sol_lines"]:
                m_single = re.match(r'^([A-D])[\.\:]?\s*$', ln.strip())
                if m_single:
                    ans = m_single.group(1)
                    break
    elif raw_q["part"] == "P2":
        p2_ans = {}
        m_inline = re.findall(r'([a-d])\s*[\-\:]\s*(ĐÚNG|SAI|Đúng|Sai|Đ|S)', sol_text)
        for k_m, v_m in m_inline:
            p2_ans[k_m.lower()] = "Đ" if v_m.upper() in ("ĐÚNG", "Đ") else "S"
        for ln in raw_q["sol_lines"]:
            m = re.match(r'^([a-d])\)\s*(Đúng|Sai)', ln.strip(), flags=re.IGNORECASE)
            if m:
                p2_ans[m.group(1).lower()] = "Đ" if m.group(2).lower() == "đúng" else "S"
        if len(p2_ans) == 4:
            ans = f"a:{p2_ans['a']} - b:{p2_ans['b']} - c:{p2_ans['c']} - d:{p2_ans['d']}"
        elif p2_ans:
            ans = " - ".join(f"{k}:{v}" for k, v in sorted(p2_ans.items()))
    elif raw_q["part"] == "P3":
        for ln in raw_q["sol_lines"]:
            m = re.search(r'Đáp\s+(?:án|số)\s*[:\.]?\s*([^\n\.]+)', ln, flags=re.IGNORECASE)
            if m:
                raw_a = m.group(1).strip()
                m_num = re.match(r'^([+-]?\d+(?:[\,\.]\d+)?)', raw_a)
                if m_num:
                    ans = m_num.group(1).replace('.', ',')
                else:
                    ans = raw_a
                break
        if not ans:
            m_kl = re.search(r'Kết\s+luận\s*:\s*([^\n\.]+)', sol_text)
            if m_kl:
                val = m_kl.group(1).strip()
                m_val = re.search(r'([0-9]+(?:[,\.][0-9]+)?(?:\s*%)?(?:\s*(?:M|mL|kg|gam|g))?)', val)
                ans = m_val.group(1).strip().replace('.', ',') if m_val else val

    item_levels = {}
    if raw_q["part"] == "P2":
        for k in ["a", "b", "c", "d"]:
            item_levels[k] = classify_p2_item_level_hoa11(raw_q["p2_items"].get(k, ""))

    lvl = raw_q.get("level_override") or classify_question_level_hoa11(
        raw_q["part"],
        raw_q["sec"],
        raw_q["topic"],
        " ".join(raw_q["stem_lines"]),
        " ".join(raw_q["opt_lines"]),
        sol_text,
        item_levels
    )

    return {
        "part": raw_q["part"],
        "topic": raw_q["topic"],
        "level": lvl,
        "item_levels": item_levels,
        "source_tag": raw_q["source_tag"],
        "stem_lines": raw_q["stem_lines"],
        "opt_lines": raw_q["opt_lines"],
        "p2_items": raw_q["p2_items"],
        "answer": ans,
        "solution": sol_text,
        "img_blobs": raw_q["img_blobs"],
        "tables": raw_q["tables"],
    }


def build_chapter_word(ch_info, questions, is_teacher):
    doc = docx.Document()
    setup_page_and_hf(doc)

    # 1. Title Block
    p_t1 = doc.add_paragraph()
    p_t1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t1.paragraph_format.space_before = Pt(0)
    p_t1.paragraph_format.space_after = Pt(2)
    r_t1 = p_t1.add_run(f"NGÂN HÀNG CÂU HỎI HÓA HỌC 11 – {ch_info['name']}")
    r_t1.bold = True
    r_t1.font.name = "Times New Roman"
    r_t1.font.size = Pt(14)
    r_t1.font.color.rgb = COLOR_PRIMARY

    p_t2 = doc.add_paragraph()
    p_t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t2.paragraph_format.space_before = Pt(0)
    p_t2.paragraph_format.space_after = Pt(2)
    r_t2 = p_t2.add_run("PHÂN LOẠI THEO CHỦ ĐỀ, DẠNG THỨC & MỨC ĐỘ NHẬN THỨC (BIẾT – HIỂU – VẬN DỤNG)")
    r_t2.bold = True
    r_t2.font.name = "Times New Roman"
    r_t2.font.size = Pt(10.5)
    r_t2.font.color.rgb = COLOR_SOLUTION

    p_t3 = doc.add_paragraph()
    p_t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t3.paragraph_format.space_before = Pt(0)
    p_t3.paragraph_format.space_after = Pt(6)
    ver_note = (
        "(BẢN GIÁO VIÊN — CÓ ĐÁP ÁN, LỜI GIẢI CHI TIẾT & MÃ ĐỊNH DANH MA TRẬN)"
        if is_teacher else
        "(BẢN HỌC SINH — ĐỀ ÔN LUYỆN PHÂN LOẠI THEO MỨC ĐỘ NHẬN THỨC)"
    )
    r_t3 = p_t3.add_run(ver_note)
    r_t3.bold = True
    r_t3.italic = True
    r_t3.font.name = "Times New Roman"
    r_t3.font.size = Pt(11)
    r_t3.font.color.rgb = COLOR_HF_RED

    # 2. Matrix Summary Table
    tbl = doc.add_table(rows=1, cols=7)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl._tbl.tblPr.append(parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="8BC390"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="8BC390"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="8BC390"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="8BC390"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="8BC390"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="8BC390"/>'
        f'</w:tblBorders>'
    ))
    headers = ["Chủ đề / Bài học", "Dạng thức", "Biết (NB)", "Hiểu (TH)", "Vận dụng (VD)", "Tổng số câu", "Phạm vi câu"]
    for i, h in enumerate(headers):
        p = tbl.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.font.color.rgb = COLOR_PRIMARY

    for part_code, part_label in [
        ("P1", "Phần I (TN 4 lựa chọn)"),
        ("P2", "Phần II (Đúng / Sai)"),
        ("P3", "Phần III (Trả lời ngắn & BT)")
    ]:
        for topic in ch_info["topics"]:
            sub_qs = [q for q in questions if q["part"] == part_code and q["topic"] == topic]
            if not sub_qs:
                continue
            row_cells = tbl.add_row().cells
            nb = sum(1 for q in sub_qs if q["level"] == "NB")
            th = sum(1 for q in sub_qs if q["level"] == "TH")
            vd = sum(1 for q in sub_qs if q["level"] == "VD")
            q_nums = [q["q_num"] for q in sub_qs]
            r_range = f"Câu {min(q_nums)} – {max(q_nums)}" if q_nums else ""

            vals = [topic, part_label, str(nb), str(th), str(vd), str(len(sub_qs)), r_range]
            for c_idx, val in enumerate(vals):
                cp = row_cells[c_idx].paragraphs[0]
                cp.paragraph_format.space_before = Pt(1.5)
                cp.paragraph_format.space_after = Pt(1.5)
                if c_idx >= 2:
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cr = cp.add_run(val)
                cr.font.name = "Times New Roman"
                cr.font.size = Pt(9.5)
                if c_idx == 0:
                    cr.bold = True

    # Total row
    tot_row = tbl.add_row().cells
    tot_nb = sum(1 for q in questions if q["level"] == "NB")
    tot_th = sum(1 for q in questions if q["level"] == "TH")
    tot_vd = sum(1 for q in questions if q["level"] == "VD")
    tot_vals = ["TỔNG TOÀN CHƯƠNG", "Trọn bộ 3 phần", str(tot_nb), str(tot_th), str(tot_vd), str(len(questions)), f"1 – {len(questions)}"]
    for c_idx, val in enumerate(tot_vals):
        cp = tot_row[c_idx].paragraphs[0]
        cp.paragraph_format.space_before = Pt(2)
        cp.paragraph_format.space_after = Pt(2)
        if c_idx >= 2:
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cp.add_run(val)
        cr.bold = True
        cr.font.name = "Times New Roman"
        cr.font.size = Pt(10)
        cr.font.color.rgb = COLOR_HF_RED

    # 3. Render Questions
    cur_part = None
    cur_topic = None

    part_headers = {
        "P1": "PHẦN I. CÂU HỎI TRẮC NGHIỆM NHIỀU PHƯƠNG ÁN LỰA CHỌN",
        "P2": "PHẦN II. CÂU HỎI TRẮC NGHIỆM ĐÚNG SAI",
        "P3": "PHẦN III. CÂU HỎI TRẢ LỜI NGẮN & BÀI TOÁN VẬN DỤNG"
    }

    for q in questions:
        if q["part"] != cur_part:
            cur_part = q["part"]
            cur_topic = None
            p_pt = doc.add_paragraph()
            p_pt.paragraph_format.space_before = Pt(10)
            p_pt.paragraph_format.space_after = Pt(4)
            r_pt = p_pt.add_run(part_headers[cur_part])
            r_pt.bold = True
            r_pt.font.name = "Times New Roman"
            r_pt.font.size = Pt(12.5)
            r_pt.font.color.rgb = COLOR_HF_RED

        if q["topic"] != cur_topic:
            cur_topic = q["topic"]
            p_tp = doc.add_paragraph()
            p_tp.paragraph_format.space_before = Pt(6)
            p_tp.paragraph_format.space_after = Pt(3)
            r_tp = p_tp.add_run(f"♦ CHỦ ĐỀ: {cur_topic.upper()}")
            r_tp.bold = True
            r_tp.font.name = "Times New Roman"
            r_tp.font.size = Pt(11.5)
            r_tp.font.color.rgb = COLOR_PRIMARY

        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)

        r_qn = p_q.add_run(f"Câu {q['q_num']}. ")
        r_qn.bold = True
        r_qn.font.name = "Times New Roman"
        r_qn.font.size = Pt(11.5)
        r_qn.font.color.rgb = COLOR_PRIMARY

        lv_name = LEVEL_NAMES[q["level"]]
        ans_str = f" | ĐA: {q['answer']}" if (is_teacher and q["answer"]) else ""
        r_tag = p_q.add_run(f"[{q['id']} | {lv_name}{ans_str}] ")
        r_tag.bold = True
        r_tag.font.name = "Times New Roman"
        r_tag.font.size = Pt(10.5)
        r_tag.font.color.rgb = COLOR_TAG

        if q.get("source_tag") and "Châu Văn Liêm" not in q["source_tag"]:
            r_src = p_q.add_run(f"{q['source_tag']} ")
            r_src.bold = True
            r_src.font.name = "Times New Roman"
            r_src.font.size = Pt(10.5)
            r_src.font.color.rgb = COLOR_SOLUTION

        st_text = "\n".join(q["stem_lines"]).strip()
        add_chem_runs_hoa10(p_q, st_text, size=11.5)

        for imp in q.get("image_paths", []):
            if os.path.exists(imp):
                p_im = doc.add_paragraph()
                p_im.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_im.paragraph_format.space_before = Pt(3)
                p_im.paragraph_format.space_after = Pt(3)
                r_im = p_im.add_run()
                r_im.add_picture(imp, width=Cm(12.0))

        for tbl_data in q.get("tables", []):
            add_data_table_to_doc(doc, tbl_data)

        if q["part"] == "P1":
            opts = parse_abcd_from_opt_lines(q["opt_lines"])
            add_options_paragraph_no_table(doc, opts, q["answer"], is_teacher)
            if is_teacher:
                add_solution_block_hoa10(doc, q)
        elif q["part"] == "P2":
            p2_ans_map = {}
            if q.get("answer"):
                for m in re.finditer(r'([a-d])\s*[:\.]\s*([ĐS])', q["answer"]):
                    p2_ans_map[m.group(1).lower()] = m.group(2)

            for item_k in ["a", "b", "c", "d"]:
                item_txt = q["p2_items"].get(item_k, "")
                if not item_txt:
                    continue
                p_it = doc.add_paragraph()
                p_it.paragraph_format.left_indent = Cm(0.6)
                p_it.paragraph_format.space_before = Pt(1.5)
                p_it.paragraph_format.space_after = Pt(1.5)

                item_lv = q.get("item_levels", {}).get(item_k, "TH")
                item_lv_name = LEVEL_NAMES.get(item_lv, "Thông hiểu")
                correct_val = p2_ans_map.get(item_k, "")

                if is_teacher and correct_val:
                    r_lbl = p_it.add_run(f"{item_k}) [{item_lv_name} | {correct_val}] ")
                    r_lbl.bold = True
                    r_lbl.font.name = "Times New Roman"
                    r_lbl.font.size = Pt(11)
                    r_lbl.font.color.rgb = COLOR_CORRECT if correct_val == "Đ" else COLOR_PRIMARY
                else:
                    r_lbl = p_it.add_run(f"{item_k}) [{item_lv_name}] ")
                    r_lbl.bold = True
                    r_lbl.font.name = "Times New Roman"
                    r_lbl.font.size = Pt(11)
                    r_lbl.font.color.rgb = COLOR_TAG

                add_chem_runs_hoa10(p_it, item_txt, size=11)

            if is_teacher:
                add_solution_block_hoa10(doc, q)
        elif q["part"] == "P3":
            if not is_teacher:
                p_ans_line = doc.add_paragraph()
                p_ans_line.paragraph_format.left_indent = Cm(0.6)
                p_ans_line.paragraph_format.space_before = Pt(1)
                p_ans_line.paragraph_format.space_after = Pt(2)
                r_al = p_ans_line.add_run("Đáp số / Trả lời: .......................................................................................")
                r_al.italic = True
                r_al.font.name = "Times New Roman"
                r_al.font.size = Pt(11)
            else:
                add_solution_block_hoa10(doc, q)

    suffix = "GiaoVien" if is_teacher else "HocSinh"
    fname = f"{ch_info['file_prefix']}_{suffix}.docx"
    out_p1 = os.path.join(OUT_DIR, fname)
    out_p2 = os.path.join(BANK_DIR, fname)
    doc.save(out_p1)
    doc.save(out_p2)
    print(f"  + Saved {suffix} Word file: {out_p2}")


def build_single_chapter(ch_num):
    ch_info = CHAPTER_DEFS[ch_num]
    print(f"\n=======================================================")
    print(f"BẮT ĐẦU XỬ LÝ {ch_info['name']}")
    print(f"=======================================================")

    raw_cvl = []
    for fp in ch_info["cvl_files"]:
        raw_cvl.extend(parse_cvl_doc(fp, ch_num))

    start_p, end_p = ch_info["ttb_range"]
    raw_ttb = parse_ttb_section(start_p, end_p, ch_num)
    print(f"Parsed CVL: {len(raw_cvl)} questions | Parsed TTB: {len(raw_ttb)} questions")

    all_raw = raw_cvl + raw_ttb
    deduped = []
    seen_norms = []
    dup_count = 0

    for rq in all_raw:
        fq = finalize_question_record_full(rq, ch_num)
        st_norm = norm_for_compare(" ".join(fq["stem_lines"]))
        opt_norm = norm_for_compare(" ".join(fq["opt_lines"]))
        full_sig = f"{fq['part']}||{st_norm}||{opt_norm}"
        is_dup = False
        for prev_part, prev_st, prev_sig in seen_norms:
            if fq["part"] == prev_part:
                if full_sig == prev_sig:
                    is_dup = True
                    break
                if rq.get("sec") == "TTB" and (st_norm == prev_st or SequenceMatcher(None, st_norm, prev_st).ratio() >= 0.85):
                    is_dup = True
                    break
        if is_dup:
            dup_count += 1
        else:
            seen_norms.append((fq["part"], st_norm, full_sig))
            deduped.append(fq)

    print(f"Filtered out {dup_count} duplicate questions -> {len(deduped)} unique questions!")

    part_order = {"P1": 1, "P2": 2, "P3": 3}
    topic_order = {t: i for i, t in enumerate(ch_info["topics"])}
    level_order = {"NB": 1, "TH": 2, "VD": 3}

    deduped.sort(key=lambda q: (
        part_order.get(q["part"], 9),
        topic_order.get(q["topic"], 9),
        level_order.get(q["level"], 9)
    ))

    final_questions = []
    for q_num, q in enumerate(deduped, 1):
        qid = f"H11-{ch_info['code']}-{q['part']}-{q['level']}-{q_num:03d}"
        img_paths = []
        for k, blob in enumerate(q.get("img_blobs", []), 1):
            fpath = os.path.join(IMG_DIR, f"h11_{ch_info['code'].lower()}_q{q_num:03d}_img{k}.png")
            try:
                im = Image.open(io.BytesIO(blob))
                im.save(fpath, format="PNG")
                img_paths.append(fpath)
            except Exception:
                pass
        q_rec = {
            "q_num": q_num,
            "id": qid,
            "part": q["part"],
            "level": q["level"],
            "topic": q["topic"],
            "source_tag": q["source_tag"],
            "stem_lines": q["stem_lines"],
            "opt_lines": q["opt_lines"],
            "p2_items": q["p2_items"],
            "item_levels": q["item_levels"],
            "answer": q["answer"],
            "solution": q["solution"],
            "image_paths": img_paths,
            "tables": q["tables"],
        }
        final_questions.append(q_rec)

    # Save chapter JSON
    ch_data = {
        "grade": 11,
        "chapter_number": ch_num,
        "chapter_name": ch_info["name"],
        "topics": ch_info["topics"],
        "total_questions": len(final_questions),
        "questions": final_questions,
    }

    json_path = os.path.join(BANK_DIR, f"{ch_info['file_prefix']}.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump({ch_info["code"]: ch_data}, f, ensure_ascii=False, indent=2)
    print(f"Saved Chapter JSON ({len(final_questions)} questions): {json_path}")

    # Build Word docs
    build_chapter_word(ch_info, final_questions, is_teacher=False)
    build_chapter_word(ch_info, final_questions, is_teacher=True)

    print(f"\n--- THỐNG KÊ {ch_info['code']}: {len(final_questions)} CÂU ---")
    for pt in ["P1", "P2", "P3"]:
        sub = [q for q in final_questions if q["part"] == pt]
        nb = sum(1 for q in sub if q["level"] == "NB")
        th = sum(1 for q in sub if q["level"] == "TH")
        vd = sum(1 for q in sub if q["level"] == "VD")
        print(f"  + {pt}: {len(sub)} câu (Biết: {nb}, Hiểu: {th}, Vận dụng: {vd})")

    return ch_info["code"], ch_data


def build_all_chapters():
    master_bank = {}
    for ch_num in sorted(CHAPTER_DEFS.keys()):
        code, ch_data = build_single_chapter(ch_num)
        master_bank[code] = ch_data

    # Save Master JSON
    master_json_path = os.path.join(BANK_DIR, "Ngan_Hang_Hoa_11_Tron_Bo.json")
    with open(master_json_path, "w", encoding="utf-8") as f:
        json.dump(master_bank, f, ensure_ascii=False, indent=2)
    print(f"\n=======================================================")
    print(f"HOÀN THÀNH TOÀN BỘ NGÂN HÀNG HÓA HỌC 11!")
    print(f"Saved Master JSON: {master_json_path}")
    total_all = sum(v["total_questions"] for v in master_bank.values())
    print(f"TỔNG CỘNG TOÀN KHỐI 11: {total_all} CÂU HỎI TRỌN BỘ 7 CHƯƠNG")
    print(f"=======================================================")


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1].isdigit():
        build_single_chapter(int(sys.argv[1]))
    else:
        build_all_chapters()
