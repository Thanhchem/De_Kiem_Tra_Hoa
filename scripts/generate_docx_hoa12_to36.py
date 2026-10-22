# -*- coding: utf-8 -*-
"""
Script tạo trọn bộ 3 file Word (.docx) chuẩn "FORM ĐỀ KIỂM TRA" cho:
CHUYÊN ĐỀ: PIN ĐIỆN VÀ ĐIỆN PHÂN (TỜ SỐ 36) - HÓA HỌC 12
Thầy TRẦN VĂN THẠNH - 0777.470.803

3 phiên bản:
1. Bản Học sinh: De_To36_Hoa12_PinDien_DienPhan_HocSinh.docx
2. Bản Giáo viên: De_To36_Hoa12_PinDien_DienPhan_GiaoVien.docx
3. Bản Dạy trên màn hình: De_To36_Hoa12_PinDien_DienPhan_DayTrenManHinh.docx
"""

import os
import sys
import re
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.path.append(os.path.dirname(__file__))
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

import hoa12_to36_data as d

ROOT_DIR = r"c:\Antigravity_Thanh"
OUT_DIR = os.path.join(ROOT_DIR, "San_Pham")
os.makedirs(OUT_DIR, exist_ok=True)

# Bảng màu chuẩn theo FORM_DE_KIEM_TRA
COLOR_HF_RED = RGBColor(211, 47, 47)        # Coral Red (#D32F2F) cho Header & Footer
COLOR_HF_GREEN_HEX = "8BC390"               # Sage Green (#8BC390) cho viền Header & Footer
COLOR_PRIMARY = RGBColor(46, 125, 50)       # Xanh lá đậm (#2E7D32) cho Tiêu đề Phân môn / Phần
COLOR_CORRECT = RGBColor(198, 40, 40)       # Đỏ (#C62828) cho Đáp án đúng Bản Giáo Viên
COLOR_SOLUTION = RGBColor(21, 101, 192)     # Xanh lam (#1565C0) cho tiêu đề Lời giải
COLOR_SCREEN_TITLE = RGBColor(106, 27, 154) # Tím đậm (#6A1B9A) cho tiêu đề bản Dạy Màn Hình
COLOR_GRAY = RGBColor(110, 110, 110)
COLOR_LINE_DOT = RGBColor(185, 185, 185)    # Xám nhạt cho dòng chấm ..........
COLOR_BORDER_GREY_HEX = "CCCCCC"

def add_p_border_bottom(p, color="8BC390", sz="6"):
    """Thêm đường kẻ viền dưới cho đoạn văn bản."""
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_p_border_top(p, color="8BC390", sz="6"):
    """Thêm đường kẻ viền trên cho đoạn văn bản."""
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    """Cài đặt lề đệm cho ô bảng."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex="F1F8E9"):
    """Đặt màu nền cho ô bảng."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="CCCCCC", sz="4"):
    """Đặt viền thanh mảnh cho bảng."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def set_run_style(run, font_name="Times New Roman", font_size=10.5, bold=False, italic=False, subscript=False, superscript=False, color=None):
    """Cài đặt style chi tiết cho từng run text."""
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    """Chèn trường số trang tự động PAGE vào chân trang."""
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def render_rich_text(p, text, base_font="Times New Roman", base_size=10.5, base_bold=False, base_italic=False, base_color=None):
    """
    Sinh văn bản định dạng phong phú:
    - *bold*
    - _italic_
    - ~subscript~
    - ^superscript^
    """
    tokens = re.split(r'(~[^~]+~|\^[^\^]+\^|\*[^\*]+\*|_[^_]+_)', text)
    for token in tokens:
        if not token:
            continue

        is_bold = base_bold
        is_italic = base_italic
        is_sub = False
        is_sup = False
        content = token

        if token.startswith('*') and token.endswith('*'):
            is_bold = True
            content = token[1:-1]
        elif token.startswith('_') and token.endswith('_'):
            is_italic = True
            content = token[1:-1]
        elif token.startswith('~') and token.endswith('~'):
            is_sub = True
            content = token[1:-1]
        elif token.startswith('^') and token.endswith('^'):
            is_sup = True
            content = token[1:-1]

        run = p.add_run(content)
        set_run_style(run, font_name=base_font, font_size=base_size, bold=is_bold, italic=is_italic,
                      subscript=is_sub, superscript=is_sup, color=base_color)

def setup_header_footer(doc):
    """Cài đặt Header & Footer chuẩn theo FORM_DE_KIEM_TRA."""
    normal_style = doc.styles['Normal']
    normal_style.paragraph_format.tab_stops.clear_all()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    section.start_type = WD_SECTION_START.NEW_PAGE
    section.different_first_page_header_footer = False
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    # 1. Header (Bảng 1x2, rộng 18.0 cm)
    header = section.header
    for p in header.paragraphs:
        p.text = ""
    tbl_h = header.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl_h.cell(0, 0)
    c_right = tbl_h.cell(0, 1)
    c_left.width = Cm(9.0)
    c_right.width = Cm(9.0)

    p_hl = c_left.paragraphs[0]
    p_hl.paragraph_format.space_before = Pt(0)
    p_hl.paragraph_format.space_after = Pt(2)
    p_hl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_hl = p_hl.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)

    p_hr = c_right.paragraphs[0]
    p_hr.paragraph_format.space_before = Pt(0)
    p_hr.paragraph_format.space_after = Pt(2)
    p_hr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p_hr.add_run("Thầy ")
    set_run_style(r1, font_size=10, color=COLOR_HF_RED)
    r2 = p_hr.add_run("TRẦN VĂN THẠNH")
    set_run_style(r2, font_size=10, bold=True, color=COLOR_HF_RED)
    r3 = p_hr.add_run(" - 0777.470.803")
    set_run_style(r3, font_size=10, color=COLOR_HF_RED)

    add_p_border_bottom(p_hl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_bottom(p_hr, color=COLOR_HF_GREEN_HEX, sz="6")

    # 2. Footer (Bảng 1x2, rộng 18.0 cm)
    footer = section.footer
    for p in footer.paragraphs:
        p.text = ""
    tbl_f = footer.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_fl = tbl_f.cell(0, 0)
    c_fr = tbl_f.cell(0, 1)
    c_fl.width = Cm(14.5)
    c_fr.width = Cm(3.5)

    p_fl = c_fl.paragraphs[0]
    p_fl.paragraph_format.space_before = Pt(2)
    p_fl.paragraph_format.space_after = Pt(0)
    p_fl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_fl = p_fl.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)

    p_fr = c_fr.paragraphs[0]
    p_fr.paragraph_format.space_before = Pt(2)
    p_fr.paragraph_format.space_after = Pt(0)
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_field(p_fr, font_size=10, bold=True, color=COLOR_HF_RED)

    add_p_border_top(p_fl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_top(p_fr, color=COLOR_HF_GREEN_HEX, sz="6")

def build_title_block(doc, is_teacher=False, is_screen=False):
    """Khối tiêu đề theo FORM_DE_KIEM_TRA."""
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl.cell(0, 0)
    c_right = tbl.cell(0, 1)
    c_left.width = Cm(8.0)
    c_right.width = Cm(10.0)

    # Cột trái
    p_l1 = c_left.paragraphs[0]
    p_l1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l1.paragraph_format.space_before = Pt(0)
    p_l1.paragraph_format.space_after = Pt(1)
    render_rich_text(p_l1, "*TRUNG TÂM LUYỆN THI HÓA HỌC THẦY THẠNH*", base_size=10.5, base_bold=True)

    p_l2 = c_left.add_paragraph()
    p_l2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l2.paragraph_format.space_before = Pt(0)
    p_l2.paragraph_format.space_after = Pt(1)
    render_rich_text(p_l2, "*TÀI LIỆU CHUYÊN ĐỀ LỚP 12*", base_size=10, base_bold=False)

    p_l3 = c_left.add_paragraph()
    p_l3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l3.paragraph_format.space_before = Pt(1)
    p_l3.paragraph_format.space_after = Pt(2)
    render_rich_text(p_l3, "_Biên soạn theo chương trình GDPT 2018_", base_size=9.5, base_italic=True, base_color=COLOR_GRAY)

    # Cột phải
    p_r1 = c_right.paragraphs[0]
    p_r1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r1.paragraph_format.space_before = Pt(0)
    p_r1.paragraph_format.space_after = Pt(1)
    render_rich_text(p_r1, f"*{d.INFO['title']}*", base_size=11.5, base_bold=True, base_color=COLOR_PRIMARY)

    p_r2 = c_right.add_paragraph()
    p_r2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r2.paragraph_format.space_before = Pt(0)
    p_r2.paragraph_format.space_after = Pt(1)
    render_rich_text(p_r2, f"*MÔN: {d.INFO['subject']}*", base_size=10.5, base_bold=True)

    p_r3 = c_right.add_paragraph()
    p_r3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r3.paragraph_format.space_before = Pt(0)
    p_r3.paragraph_format.space_after = Pt(2)
    render_rich_text(p_r3, "_Hệ thống lý thuyết chuyên sâu & bài tập phân loại_", base_size=9.5, base_italic=True)

    # Dòng ghi chú phân biệt bản Giáo viên / Dạy màn hình
    if is_teacher:
        p_note = doc.add_paragraph()
        p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_note.paragraph_format.space_before = Pt(3)
        p_note.paragraph_format.space_after = Pt(2)
        render_rich_text(p_note, "*(BẢN GIÁO VIÊN -- CÓ ĐÁP ÁN VÀ LỜI GIẢI CHI TIẾT)*", base_size=11, base_bold=True, base_color=COLOR_CORRECT)
    elif is_screen:
        p_note = doc.add_paragraph()
        p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_note.paragraph_format.space_before = Pt(3)
        p_note.paragraph_format.space_after = Pt(2)
        render_rich_text(p_note, "*(BẢN DẠY TRÊN MÀN HÌNH -- CÓ KHÔNG GIAN TƯƠNG TÁC RỘNG RÃI)*", base_size=11, base_bold=True, base_color=COLOR_SCREEN_TITLE)

    # Dòng thông tin học sinh
    p_stu = doc.add_paragraph()
    p_stu.paragraph_format.space_before = Pt(3)
    p_stu.paragraph_format.space_after = Pt(2)
    render_rich_text(p_stu, "Họ và tên học sinh: ........................................................................ Lớp: ................. SBD: .................", base_size=10)

    # Dòng nguyên tử khối
    p_mass = doc.add_paragraph()
    p_mass.paragraph_format.space_before = Pt(1)
    p_mass.paragraph_format.space_after = Pt(4)
    render_rich_text(p_mass, "_(Cho biết: H = 1; C = 12; N = 14; O = 16; Na = 23; Mg = 24,3; Al = 27; S = 32; Cl = 35,5; Fe = 56; Ni = 58,7; Cu = 64; Zn = 65; Ag = 108; Au = 197; F = 96 500 C/mol)_", base_size=9.0, base_italic=True, base_color=COLOR_GRAY)

    # Đường phân cách nhẹ
    p_sep = doc.add_paragraph()
    p_sep.paragraph_format.space_before = Pt(0)
    p_sep.paragraph_format.space_after = Pt(4)
    add_p_border_bottom(p_sep, color=COLOR_BORDER_GREY_HEX, sz="4")

def add_part_header(doc, title_text, base_size=11.5):
    """Tiêu đề phân môn lớn."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    render_rich_text(p, f"*{title_text.upper()}*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)
    add_p_border_bottom(p, color=COLOR_HF_GREEN_HEX, sz="6")

def add_solution_block(doc, sol_text, base_size=10):
    """Khối Lời giải chi tiết thụt lề 0.6 cm."""
    p_head = doc.add_paragraph()
    p_head.paragraph_format.left_indent = Cm(0.6)
    p_head.paragraph_format.space_before = Pt(2)
    p_head.paragraph_format.space_after = Pt(1)
    p_head.paragraph_format.keep_with_next = True
    render_rich_text(p_head, "► *Lời giải chi tiết:*", base_size=base_size, base_bold=True, base_color=COLOR_SOLUTION)

    lines = sol_text.strip().split("\n")
    for line in lines:
        if not line.strip():
            continue
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        render_rich_text(p, line, base_size=base_size)

def add_writing_lines(doc, num_lines=3):
    """Thêm các dòng chấm làm bài cho học sinh."""
    for _ in range(num_lines):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        render_rich_text(p, "........................................................................................................................................................",
                         base_size=9.5, base_color=COLOR_LINE_DOT)

def add_screen_interactive_space(doc, num_lines=4):
    """Thêm không gian tương tác cho bản dạy trên màn hình."""
    for _ in range(num_lines):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        render_rich_text(p, "─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─ ─",
                         base_size=9.0, base_color=RGBColor(210, 210, 210))

def render_mcq_options(doc, options, correct_letter, is_teacher, base_size=10.5):
    """Bố trí các phương án trắc nghiệm bằng Tab Stops theo Quy tắc 1.7."""
    opts = [opt.strip() for opt in options if opt.strip()]
    if not opts:
        return

    lens = [len(opt) for opt in opts]
    max_len = max(lens) if lens else 0

    if max_len <= 18:
        # 4 phương án trên 1 dòng
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(5.0), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(9.5), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(14.0), WD_TAB_ALIGNMENT.LEFT)

        for opt in opts:
            p.add_run("\t")
            letter = opt.strip()[0]
            if is_teacher and letter == correct_letter:
                render_rich_text(p, f"*{opt}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p, opt, base_size=base_size)

    elif max_len <= 45:
        # 2 dòng x 2 phương án
        p1 = doc.add_paragraph()
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.tab_stops.clear_all()
        p1.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p1.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)

        for opt in opts[:2]:
            p1.add_run("\t")
            letter = opt.strip()[0]
            if is_teacher and letter == correct_letter:
                render_rich_text(p1, f"*{opt}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p1, opt, base_size=base_size)

        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(1)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.tab_stops.clear_all()
        p2.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p2.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)

        for opt in opts[2:]:
            p2.add_run("\t")
            letter = opt.strip()[0]
            if is_teacher and letter == correct_letter:
                render_rich_text(p2, f"*{opt}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p2, opt, base_size=base_size)

    else:
        # Mỗi phương án 1 dòng riêng
        for opt in opts:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            letter = opt.strip()[0]
            if is_teacher and letter == correct_letter:
                render_rich_text(p, f"*{opt}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p, opt, base_size=base_size)

def add_image_if_exists(doc, img_name, max_width_cm=14.0):
    """Chèn hình ảnh nếu tồn tại trên đĩa."""
    if not img_name:
        return
    img_path = os.path.join(d.IMG_DIR, img_name)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(3)
        p_img.paragraph_format.space_after = Pt(3)
        p_img.paragraph_format.keep_with_next = True
        try:
            run_img = p_img.add_run()
            run_img.add_picture(img_path, width=Cm(min(max_width_cm, 16.0)))
        except Exception as e:
            print(f"Warning: could not add image {img_name}: {e}")

def render_theory_section(doc, theory_data):
    """Trình bày một chuyên đề lý thuyết kèm các bảng tổng kết sư phạm."""
    add_part_header(doc, theory_data["title"], base_size=11.5)

    for sec in theory_data.get("sections", []):
        heading = sec.get("heading")
        if heading:
            p_h = doc.add_paragraph()
            p_h.paragraph_format.space_before = Pt(5)
            p_h.paragraph_format.space_after = Pt(2)
            p_h.paragraph_format.keep_with_next = True
            render_rich_text(p_h, f"*{heading}*", base_size=10.5, base_bold=True, base_color=RGBColor(27, 94, 32))

        img = sec.get("image")
        if img:
            add_image_if_exists(doc, img, max_width_cm=13.0)

        for it in sec.get("items", []):
            p_it = doc.add_paragraph()
            p_it.paragraph_format.left_indent = Cm(0.4)
            p_it.paragraph_format.space_before = Pt(1)
            p_it.paragraph_format.space_after = Pt(1.5)
            render_rich_text(p_it, it, base_size=10.0)

        table_info = sec.get("table")
        if table_info:
            headers = table_info["headers"]
            rows = table_info["rows"]
            tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(tbl, color="B0BEC5", sz="4")

            # Header row
            hdr_cells = tbl.rows[0].cells
            for col_idx, h_text in enumerate(headers):
                cell = hdr_cells[col_idx]
                set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
                set_cell_shading(cell, color_hex="E8F5E9")
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                render_rich_text(p, f"*{h_text}*", base_size=9.5, base_bold=True, base_color=RGBColor(27, 94, 32))

            # Data rows
            for r_idx, row_data in enumerate(rows):
                row_cells = tbl.rows[r_idx + 1].cells
                for c_idx, val in enumerate(row_data):
                    cell = row_cells[c_idx]
                    set_cell_margins(cell, top=60, bottom=60, left=90, right=90)
                    if r_idx % 2 == 1:
                        set_cell_shading(cell, color_hex="F9FBE7")
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    if c_idx == 0:
                        render_rich_text(p, f"*{val}*", base_size=9.5, base_bold=True)
                    else:
                        render_rich_text(p, val, base_size=9.5)

            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.space_before = Pt(0)
            p_sp.paragraph_format.space_after = Pt(3)

def render_question(doc, q, is_teacher=False, is_screen=False, q_prefix="Câu"):
    """Trình bày một câu hỏi hoàn chỉnh theo đúng chuẩn và bản tương ứng."""
    q_type = q.get("type", "mcq")
    q_num = q.get("num", 1)
    content = q.get("content", "")
    ans = q.get("ans", "")
    solution = q.get("solution", "")
    image = q.get("image")

    # Đoạn dẫn câu hỏi
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(4)
    p_q.paragraph_format.space_after = Pt(2)
    p_q.paragraph_format.keep_with_next = True

    # Nhãn câu hỏi
    render_rich_text(p_q, f"*{q_prefix} {q_num}:* ", base_size=10.5, base_bold=True, base_color=RGBColor(13, 71, 161))
    render_rich_text(p_q, content, base_size=10.5)

    # Chèn hình ảnh nếu có
    if image:
        add_image_if_exists(doc, image, max_width_cm=13.0)

    # 1. Câu hỏi trắc nghiệm nhiều phương án (MCQ)
    if q_type == "mcq":
        options = q.get("options", [])
        render_mcq_options(doc, options, ans, is_teacher=is_teacher, base_size=10.5)

    # 2. Câu hỏi trắc nghiệm Đúng / Sai (TF)
    elif q_type == "tf":
        items = q.get("items", [])
        for it in items:
            p_it = doc.add_paragraph()
            p_it.paragraph_format.left_indent = Cm(0.6)
            p_it.paragraph_format.space_before = Pt(1)
            p_it.paragraph_format.space_after = Pt(1)

            it_id = it.get("id", "a")
            it_text = it.get("text", "")
            it_ans = it.get("ans", "")

            if is_teacher:
                ans_str = "[ĐÚNG]" if it_ans.upper() in ["Đ", "D", "TRUE"] else "[SAI]"
                color_ans = COLOR_PRIMARY if "ĐÚNG" in ans_str else COLOR_CORRECT
                render_rich_text(p_it, f"*{it_id})* ", base_size=10.0, base_bold=True)
                render_rich_text(p_it, f"*{ans_str}* ", base_size=10.0, base_bold=True, base_color=color_ans)
                render_rich_text(p_it, it_text, base_size=10.0)
            elif is_screen:
                render_rich_text(p_it, f"*{it_id})* [     ] ", base_size=10.0, base_bold=True)
                render_rich_text(p_it, it_text, base_size=10.0)
            else:
                render_rich_text(p_it, f"*{it_id})* [     ] ", base_size=10.0, base_bold=True)
                render_rich_text(p_it, it_text, base_size=10.0)

    # 3. Câu hỏi tự luận / Trả lời ngắn / Tính toán
    elif q_type in ["essay", "short_answer"]:
        if is_teacher:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Cm(0.6)
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(1)
            render_rich_text(p_ans, f"► *Đáp số:* *{ans}*", base_size=10.5, base_bold=True, base_color=COLOR_CORRECT)
        elif is_screen:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Cm(0.6)
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(1)
            render_rich_text(p_ans, "► *Trả lời:* .........................................................................................", base_size=10.0)
            add_screen_interactive_space(doc, num_lines=4)
        else:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Cm(0.6)
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(1)
            render_rich_text(p_ans, "► *Trả lời:* .........................................................................................", base_size=10.0)
            add_writing_lines(doc, num_lines=3)

    # Hiển thị lời giải chi tiết cho Bản Giáo Viên
    if is_teacher and solution:
        add_solution_block(doc, solution, base_size=10.0)

def add_summary_answer_table(doc):
    """Bảng tổng hợp đáp án nhanh ở cuối tài liệu cho Bản Giáo Viên."""
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.paragraph_format.keep_with_next = True
    render_rich_text(p_title, "*BẢNG TỔNG HỢP ĐÁP ÁN NHANH TOÀN BỘ CHUYÊN ĐỀ*", base_size=11, base_bold=True, base_color=COLOR_PRIMARY)
    add_p_border_bottom(p_title, color=COLOR_HF_GREEN_HEX, sz="6")

    # Bảng đáp án Bài tập tổng hợp (Câu 1 đến Câu 36)
    p_th = doc.add_paragraph()
    p_th.paragraph_format.space_before = Pt(4)
    p_th.paragraph_format.space_after = Pt(2)
    render_rich_text(p_th, "*1. ĐÁP ÁN BÀI TẬP TỔNG HỢP (CÂU 1 - CÂU 36):*", base_size=10.0, base_bold=True)

    items = d.TONG_HOP_MCQ_TF
    cols = 6
    rows = (len(items) + cols - 1) // cols
    tbl = doc.add_table(rows=rows * 2, cols=cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color="B0BEC5", sz="4")

    for idx, q in enumerate(items):
        r_blk = idx // cols
        c = idx % cols
        r_q = r_blk * 2
        r_a = r_blk * 2 + 1

        cell_q = tbl.cell(r_q, c)
        set_cell_margins(cell_q, top=50, bottom=50, left=50, right=50)
        set_cell_shading(cell_q, color_hex="E8F5E9")
        pq = cell_q.paragraphs[0]
        pq.alignment = WD_ALIGN_PARAGRAPH.CENTER
        render_rich_text(pq, f"*{q['num']}*", base_size=9.5, base_bold=True)

        cell_a = tbl.cell(r_a, c)
        set_cell_margins(cell_a, top=50, bottom=50, left=50, right=50)
        pa = cell_a.paragraphs[0]
        pa.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if q["type"] == "mcq":
            render_rich_text(pa, f"*{q.get('ans', '')}*", base_size=9.5, base_bold=True, base_color=COLOR_CORRECT)
        elif q["type"] == "tf":
            sub_ans = "".join([it.get("ans", "") for it in q.get("items", [])])
            render_rich_text(pa, f"*{sub_ans}*", base_size=9.0, base_bold=True, base_color=COLOR_SOLUTION)

def build_single_document(version="HocSinh"):
    """Tạo một file Word theo đúng phiên bản."""
    is_teacher = (version == "GiaoVien")
    is_screen = (version == "DayTrenManHinh")

    doc = docx.Document()
    setup_header_footer(doc)
    build_title_block(doc, is_teacher=is_teacher, is_screen=is_screen)

    # 1. DẠNG 1: MẠ KIM LOẠI
    render_theory_section(doc, d.THEORY_DANG1)
    p_sec1 = doc.add_paragraph()
    p_sec1.paragraph_format.space_before = Pt(6)
    p_sec1.paragraph_format.space_after = Pt(2)
    p_sec1.paragraph_format.keep_with_next = True
    render_rich_text(p_sec1, "*BÀI TẬP VẬN DỤNG DẠNG 1: MẠ KIM LOẠI*", base_size=11.0, base_bold=True, base_color=RGBColor(27, 94, 32))
    for ex in d.EXERCISES_DANG1:
        render_question(doc, ex, is_teacher=is_teacher, is_screen=is_screen, q_prefix="Câu")

    # 2. DẠNG 2: ANODE HÓA NHÔM
    render_theory_section(doc, d.THEORY_DANG2)
    p_sec2 = doc.add_paragraph()
    p_sec2.paragraph_format.space_before = Pt(6)
    p_sec2.paragraph_format.space_after = Pt(2)
    p_sec2.paragraph_format.keep_with_next = True
    render_rich_text(p_sec2, "*BÀI TẬP VẬN DỤNG DẠNG 2: QUÁ TRÌNH ANODE HÓA NHÔM*", base_size=11.0, base_bold=True, base_color=RGBColor(27, 94, 32))
    for ex in d.EXERCISES_DANG2:
        render_question(doc, ex, is_teacher=is_teacher, is_screen=is_screen, q_prefix="Câu")

    # 3. DẠNG 3: ĐIỆN PHÂN NÓNG CHẢY Al2O3
    render_theory_section(doc, d.THEORY_DANG3)
    p_sec3 = doc.add_paragraph()
    p_sec3.paragraph_format.space_before = Pt(6)
    p_sec3.paragraph_format.space_after = Pt(2)
    p_sec3.paragraph_format.keep_with_next = True
    render_rich_text(p_sec3, "*BÀI TẬP VẬN DỤNG DẠNG 3: ĐIỆN PHÂN NÓNG CHẢY Al~2~O~3~ (HALL-HÉROULT)*", base_size=11.0, base_bold=True, base_color=RGBColor(27, 94, 32))
    for ex in d.EXERCISES_DANG3:
        render_question(doc, ex, is_teacher=is_teacher, is_screen=is_screen, q_prefix="Câu")

    # 4. DẠNG 4: TINH CHẾ KIM LOẠI (ĐỒNG)
    render_theory_section(doc, d.THEORY_DANG4)
    p_sec4 = doc.add_paragraph()
    p_sec4.paragraph_format.space_before = Pt(6)
    p_sec4.paragraph_format.space_after = Pt(2)
    p_sec4.paragraph_format.keep_with_next = True
    render_rich_text(p_sec4, "*BÀI TẬP VẬN DỤNG DẠNG 4: TINH CHẾ ĐIỆN PHÂN ĐỒNG*", base_size=11.0, base_bold=True, base_color=RGBColor(27, 94, 32))
    for ex in d.EXERCISES_DANG4:
        render_question(doc, ex, is_teacher=is_teacher, is_screen=is_screen, q_prefix="Câu")

    # 5. BÀI TẬP TỔNG HỢP (CÂU 1 ĐẾN CÂU 36)
    add_part_header(doc, "BÀI TẬP TỔNG HỢP PIN ĐIỆN VÀ ĐIỆN PHÂN (CÂU 1 - CÂU 36)", base_size=11.5)
    for ex in d.TONG_HOP_MCQ_TF:
        render_question(doc, ex, is_teacher=is_teacher, is_screen=is_screen, q_prefix="Câu")

    # 6. PHẦN III. TRẢ LỜI NGẮN (CÂU 37 ĐẾN CÂU 74)
    add_part_header(doc, "PHẦN III. CÂU HỎI TRẢ LỜI NGẮN VÀ TÍNH TOÁN THỰC TIỄN (CÂU 37 - CÂU 74)", base_size=11.5)
    for ex in d.PHAN3_SHORT_ANSWER:
        render_question(doc, ex, is_teacher=is_teacher, is_screen=is_screen, q_prefix="Câu")

    # Bảng đáp án nhanh cho Bản Giáo Viên
    if is_teacher:
        add_summary_answer_table(doc)

    out_name = f"De_To36_Hoa12_PinDien_DienPhan_{version}.docx"
    out_path = os.path.join(OUT_DIR, out_name)
    doc.save(out_path)
    print(f"-> Saved: {out_path} ({os.path.getsize(out_path):,} bytes)")
    return out_path

def main():
    print("=================================================================")
    print("BẮT ĐẦU TẠO 3 FILE WORD CHUẨN FORM CHO TỜ SỐ 36 HÓA 12...")
    print("=================================================================")

    p_hs = build_single_document("HocSinh")
    p_gv = build_single_document("GiaoVien")
    p_mh = build_single_document("DayTrenManHinh")

    print("\nHOÀN TẤT TẠO 3 FILE WORD THÀNH CÔNG!")
    print(f"1. Học sinh: {p_hs}")
    print(f"2. Giáo viên: {p_gv}")
    print(f"3. Dạy màn hình: {p_mh}")

if __name__ == "__main__":
    main()
