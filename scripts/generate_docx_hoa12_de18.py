# -*- coding: utf-8 -*-
"""
Script tạo 2 file Word (.docx) chuẩn Form Đề Kiểm Tra cho:
ĐỀ SỐ 18 - CHƯƠNG 6: ĐẠI CƯƠNG KIM LOẠI (HÓA HỌC 12)
Mã đề: 118
Bao gồm:
1. Bản Học sinh: De_Kiem_Tra_Hoa_12_De18.docx
2. Bản Giáo viên: De_Kiem_Tra_Hoa_12_De18_GiaoVien.docx (Có đáp án đỏ, lời giải chi tiết xanh, bảng đáp án)
"""

import os
import sys
sys.stdout.reconfigure(encoding='utf-8')
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"

COLOR_CORAL_RED = RGBColor(211, 47, 47)      # #D32F2F
COLOR_SAGE_GREEN = "8BC390"                 # #8BC390
COLOR_DEEP_BLUE = RGBColor(21, 101, 192)    # #1565C0
COLOR_DARK_GREEN = RGBColor(46, 125, 50)    # #2E7D32
COLOR_BLACK = RGBColor(0, 0, 0)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.59)     # 1.5 cm
        s.bottom_margin = Inches(0.59)  # 1.5 cm
        s.left_margin = Inches(0.59)    # 1.5 cm
        s.right_margin = Inches(0.59)   # 1.5 cm
        s.header_distance = Inches(0.24) # 0.6 cm
        s.footer_distance = Inches(0.24) # 0.6 cm

        # --- HEADER ---
        header = s.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.text = ""
        hp.paragraph_format.space_before = Pt(0)
        hp.paragraph_format.space_after = Pt(0)

        tbl_h = header.add_table(rows=1, cols=2, width=Inches(7.08)) # 18.0 cm
        tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
        c0, c1 = tbl_h.rows[0].cells
        c0.width = Inches(3.54)
        c1.width = Inches(3.54)
        set_cell_margins(c0, 60, 60, 0, 0)
        set_cell_margins(c1, 60, 60, 0, 0)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run("Tài liệu lưu hành nội bộ")
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(10)
        r0.font.italic = True
        r0.font.color.rgb = COLOR_CORAL_RED

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1_a = p1.add_run("Thầy ")
        r1_a.font.name = "Times New Roman"
        r1_a.font.size = Pt(10)
        r1_a.font.color.rgb = COLOR_CORAL_RED
        r1_b = p1.add_run("TRẦN VĂN THẠNH")
        r1_b.font.name = "Times New Roman"
        r1_b.font.size = Pt(10)
        r1_b.font.bold = True
        r1_b.font.color.rgb = COLOR_CORAL_RED
        r1_c = p1.add_run(" - 0777.470.803")
        r1_c.font.name = "Times New Roman"
        r1_c.font.size = Pt(10)
        r1_c.font.color.rgb = COLOR_CORAL_RED

        tblPr_h = tbl_h._tbl.tblPr
        tblBorders_h = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{COLOR_SAGE_GREEN}"/>'
            f'</w:tblBorders>'
        )
        tblPr_h.append(tblBorders_h)

        # --- FOOTER ---
        footer = s.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.text = ""
        fp.paragraph_format.space_before = Pt(0)
        fp.paragraph_format.space_after = Pt(0)

        tbl_f = footer.add_table(rows=1, cols=2, width=Inches(7.08))
        tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
        fc0, fc1 = tbl_f.rows[0].cells
        fc0.width = Inches(5.31)
        fc1.width = Inches(1.77)
        set_cell_margins(fc0, 60, 60, 0, 0)
        set_cell_margins(fc1, 60, 60, 0, 0)

        fp0 = fc0.paragraphs[0]
        fp0.paragraph_format.space_before = Pt(0)
        fp0.paragraph_format.space_after = Pt(0)
        fr0 = fp0.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
        fr0.font.name = "Times New Roman"
        fr0.font.size = Pt(9.5)
        fr0.font.italic = True
        fr0.font.color.rgb = COLOR_CORAL_RED

        fp1 = fc1.paragraphs[0]
        fp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp1.paragraph_format.space_before = Pt(0)
        fp1.paragraph_format.space_after = Pt(0)
        fr1_a = fp1.add_run("Trang ")
        fr1_a.font.name = "Times New Roman"
        fr1_a.font.size = Pt(10)
        fr1_a.font.bold = True
        fr1_a.font.color.rgb = COLOR_CORAL_RED

        fldSimple = OxmlElement('w:fldSimple')
        fldSimple.set(qn('w:instr'), 'PAGE')
        r_page = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        rFont = OxmlElement('w:rFonts')
        rFont.set(qn('w:ascii'), 'Times New Roman')
        rFont.set(qn('w:hAnsi'), 'Times New Roman')
        b = OxmlElement('w:b')
        c_el = OxmlElement('w:color')
        c_el.set(qn('w:val'), 'D32F2F')
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '20')
        rPr.append(rFont)
        rPr.append(b)
        rPr.append(c_el)
        rPr.append(sz)
        r_page.append(rPr)
        t_el = OxmlElement('w:t')
        t_el.text = "1"
        r_page.append(t_el)
        fldSimple.append(r_page)
        fp1._p.append(fldSimple)

        tblPr_f = tbl_f._tbl.tblPr
        tblBorders_f = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{COLOR_SAGE_GREEN}"/>'
            f'</w:tblBorders>'
        )
        tblPr_f.append(tblBorders_f)

def add_chem_run(paragraph, text, bold=False, italic=False, color=None, font_size=11):
    parts = re.split(r'(~.*?~|\^.*?\^)', text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run()
        run.font.name = "Times New Roman"
        run.font.size = Pt(font_size)
        run.bold = bold
        run.italic = italic
        if color:
            run.font.color.rgb = color

        if part.startswith('~') and part.endswith('~'):
            run.text = part[1:-1]
            run.font.subscript = True
        elif part.startswith('^') and part.endswith('^'):
            run.text = part[1:-1]
            run.font.superscript = True
        else:
            run.text = part

def create_title_block(doc, is_teacher=False):
    lines_info = [
        ("SỞ GIÁO DỤC VÀ ĐÀO TẠO THỪA THIÊN HUẾ", "ĐỀ KIỂM TRA ĐỊNH KỲ CHƯƠNG 6 (ĐẠI CƯƠNG KIM LOẠI)", True, True, COLOR_DARK_GREEN),
        ("TRƯỜNG THPT HAI BÀ TRƯNG", "MÔN: HOÁ HỌC -- LỚP 12", True, True, None),
        ("ĐỀ CHÍNH THỨC", "Thời gian làm bài: 50 phút (không kể thời gian phát đề)", True, False, None),
        ("(Đề thi có 04 trang)", "Mã đề thi: 118", False, True, None)
    ]

    for left_txt, right_txt, left_bold, right_bold, right_color in lines_info:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.08), WD_TAB_ALIGNMENT.RIGHT)

        r_l = p.add_run(left_txt)
        r_l.font.name = "Times New Roman"
        r_l.font.size = Pt(11)
        r_l.bold = left_bold
        if not left_bold:
            r_l.italic = True

        p.add_run("\t")

        r_r = p.add_run(right_txt)
        r_r.font.name = "Times New Roman"
        r_r.font.size = Pt(11)
        r_r.bold = right_bold
        if right_color:
            r_r.font.color.rgb = right_color
        if not right_bold:
            r_r.italic = True

    if is_teacher:
        p_note = doc.add_paragraph()
        p_note.paragraph_format.space_before = Pt(4)
        p_note.paragraph_format.space_after = Pt(4)
        p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_n = p_note.add_run("(BẢN GIÁO VIÊN -- CÓ LỜI GIẢI CHI TIẾT)")
        r_n.font.name = "Times New Roman"
        r_n.font.size = Pt(12)
        r_n.font.bold = True
        r_n.font.color.rgb = COLOR_CORAL_RED

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(3)
    p_info.paragraph_format.space_after = Pt(2)
    r_info = p_info.add_run("Họ, tên thí sinh: ..................................................................... Lớp: ........................ SBD: ............................")
    r_info.font.name = "Times New Roman"
    r_info.font.size = Pt(10.5)
    r_info.font.italic = True

    p_mass = doc.add_paragraph()
    p_mass.paragraph_format.space_before = Pt(0)
    p_mass.paragraph_format.space_after = Pt(5)
    r_mass = p_mass.add_run("(Cho biết nguyên tử khối: H = 1; C = 12; N = 14; O = 16; Na = 23; Mg = 24; Al = 27; S = 32; Cl = 35,5; K = 39; Ca = 40; Fe = 56; Cu = 64; Zn = 65; Br = 80; Ag = 108; Ba = 137; Au = 197; Pb = 207; hằng số Faraday F = 96485 C/mol).")
    r_mass.font.name = "Times New Roman"
    r_mass.font.size = Pt(9.5)
    r_mass.font.italic = True

def add_part_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = COLOR_DARK_GREEN

def add_explanation(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.line_spacing = 1.15

    r_lbl = p.add_run("Lời giải chi tiết: ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(10)
    r_lbl.font.bold = True
    r_lbl.font.italic = True
    r_lbl.font.color.rgb = COLOR_DEEP_BLUE

    add_chem_run(p, text, font_size=10, italic=False, color=COLOR_DEEP_BLUE)

def render_mcq(doc, q_num, text, options, correct_key, explanation, is_teacher=False, img_path=None, img_width=Inches(3.5)):
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(3)
    p_q.paragraph_format.space_after = Pt(2)
    p_q.paragraph_format.line_spacing = 1.15

    r_lbl = p_q.add_run(f"Câu {q_num}: ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(11)
    r_lbl.font.bold = True

    add_chem_run(p_q, text, font_size=11)

    if img_path and os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(2)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture(img_path, width=img_width)

    max_len = max(len(opt[1]) for opt in options)
    if max_len <= 15:
        p_opts = doc.add_paragraph()
        p_opts.paragraph_format.space_before = Pt(0)
        p_opts.paragraph_format.space_after = Pt(2)
        p_opts.paragraph_format.line_spacing = 1.15
        p_opts.paragraph_format.tab_stops.add_tab_stop(Inches(1.77), WD_TAB_ALIGNMENT.LEFT)
        p_opts.paragraph_format.tab_stops.add_tab_stop(Inches(3.54), WD_TAB_ALIGNMENT.LEFT)
        p_opts.paragraph_format.tab_stops.add_tab_stop(Inches(5.31), WD_TAB_ALIGNMENT.LEFT)

        for i, (k, txt) in enumerate(options):
            is_c = (k == correct_key) and is_teacher
            r_k = p_opts.add_run(f"{k}. ")
            r_k.font.name = "Times New Roman"
            r_k.font.size = Pt(11)
            r_k.font.bold = is_c or True
            if is_c:
                r_k.font.color.rgb = COLOR_CORAL_RED
            add_chem_run(p_opts, txt, bold=is_c, color=COLOR_CORAL_RED if is_c else None, font_size=11)
            if i < 3:
                p_opts.add_run("\t")

    elif max_len <= 35:
        for row_opts in [options[:2], options[2:]]:
            p_opts = doc.add_paragraph()
            p_opts.paragraph_format.space_before = Pt(0)
            p_opts.paragraph_format.space_after = Pt(2)
            p_opts.paragraph_format.line_spacing = 1.15
            p_opts.paragraph_format.tab_stops.add_tab_stop(Inches(3.54), WD_TAB_ALIGNMENT.LEFT)

            for i, (k, txt) in enumerate(row_opts):
                is_c = (k == correct_key) and is_teacher
                r_k = p_opts.add_run(f"{k}. ")
                r_k.font.name = "Times New Roman"
                r_k.font.size = Pt(11)
                r_k.font.bold = is_c or True
                if is_c:
                    r_k.font.color.rgb = COLOR_CORAL_RED
                add_chem_run(p_opts, txt, bold=is_c, color=COLOR_CORAL_RED if is_c else None, font_size=11)
                if i == 0:
                    p_opts.add_run("\t")
    else:
        for k, txt in options:
            p_opts = doc.add_paragraph()
            p_opts.paragraph_format.space_before = Pt(0)
            p_opts.paragraph_format.space_after = Pt(2)
            p_opts.paragraph_format.line_spacing = 1.15
            is_c = (k == correct_key) and is_teacher
            r_k = p_opts.add_run(f"{k}. ")
            r_k.font.name = "Times New Roman"
            r_k.font.size = Pt(11)
            r_k.font.bold = is_c or True
            if is_c:
                r_k.font.color.rgb = COLOR_CORAL_RED
            add_chem_run(p_opts, txt, bold=is_c, color=COLOR_CORAL_RED if is_c else None, font_size=11)

    if is_teacher:
        add_explanation(doc, explanation)

def render_tf(doc, q_num, text, sub_items, explanation, is_teacher=False, img_path=None, img_width=Inches(3.5)):
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(3)
    p_q.paragraph_format.space_after = Pt(2)
    p_q.paragraph_format.line_spacing = 1.15

    r_lbl = p_q.add_run(f"Câu {q_num}: ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(11)
    r_lbl.font.bold = True

    add_chem_run(p_q, text, font_size=11)

    if img_path and os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(2)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture(img_path, width=img_width)

    for sub_k, sub_txt, is_true in sub_items:
        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(1.5)
        p_sub.paragraph_format.left_indent = Inches(0.2)
        p_sub.paragraph_format.line_spacing = 1.15

        r_sk = p_sub.add_run(f"{sub_k}) ")
        r_sk.font.name = "Times New Roman"
        r_sk.font.size = Pt(11)
        r_sk.font.bold = True

        add_chem_run(p_sub, sub_txt, font_size=11)

        if is_teacher:
            r_ans = p_sub.add_run(f" --> {'ĐÚNG' if is_true else 'SAI'}")
            r_ans.font.name = "Times New Roman"
            r_ans.font.size = Pt(10.5)
            r_ans.font.bold = True
            r_ans.font.color.rgb = COLOR_CORAL_RED

    if is_teacher:
        add_explanation(doc, explanation)

def render_short(doc, q_num, text, answer, explanation, is_teacher=False, extra_table=None):
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(3)
    p_q.paragraph_format.space_after = Pt(2)
    p_q.paragraph_format.line_spacing = 1.15

    r_lbl = p_q.add_run(f"Câu {q_num}: ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(11)
    r_lbl.font.bold = True

    add_chem_run(p_q, text, font_size=11)

    if extra_table:
        # extra paragraphs for table info
        for line in extra_table:
            p_t = doc.add_paragraph()
            p_t.paragraph_format.space_before = Pt(0)
            p_t.paragraph_format.space_after = Pt(1)
            p_t.paragraph_format.left_indent = Inches(0.3)
            p_t.paragraph_format.line_spacing = 1.15
            add_chem_run(p_t, line, font_size=10.5)

    if not is_teacher:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(1)
        p_ans.paragraph_format.space_after = Pt(3)
        p_ans.paragraph_format.left_indent = Inches(0.2)
        r_ans = p_ans.add_run("Đáp án: ............................................................................................................................")
        r_ans.font.name = "Times New Roman"
        r_ans.font.size = Pt(10.5)
        r_ans.font.italic = True
    else:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(1)
        p_ans.paragraph_format.space_after = Pt(2)
        p_ans.paragraph_format.left_indent = Inches(0.2)
        r_lbl_ans = p_ans.add_run("Đáp án: ")
        r_lbl_ans.font.name = "Times New Roman"
        r_lbl_ans.font.size = Pt(11)
        r_lbl_ans.font.bold = True
        r_lbl_ans.font.color.rgb = COLOR_CORAL_RED

        r_ans_val = p_ans.add_run(f"{answer}")
        r_ans_val.font.name = "Times New Roman"
        r_ans_val.font.size = Pt(11)
        r_ans_val.font.bold = True
        r_ans_val.font.color.rgb = COLOR_CORAL_RED

        add_explanation(doc, explanation)

def add_answer_key_table(doc):
    p_h = doc.add_paragraph()
    p_h.paragraph_format.space_before = Pt(8)
    p_h.paragraph_format.space_after = Pt(4)
    r_h = p_h.add_run("BẢNG ĐÁP ÁN THAM KHẢO TOÀN ĐỀ")
    r_h.font.name = "Times New Roman"
    r_h.font.size = Pt(11)
    r_h.font.bold = True
    r_h.font.color.rgb = COLOR_CORAL_RED

    p_p1 = doc.add_paragraph()
    p_p1.paragraph_format.space_before = Pt(2)
    p_p1.paragraph_format.space_after = Pt(2)
    r_p1 = p_p1.add_run("PHẦN I: ")
    r_p1.bold = True
    r_p1.font.color.rgb = COLOR_DARK_GREEN
    p_p1.add_run("1.D   2.A   3.D   4.B   5.C   6.B   7.B   8.A   9.D   10.C   11.A   12.B   13.B   14.A   15.D   16.D   17.B   18.A")

    p_p2 = doc.add_paragraph()
    p_p2.paragraph_format.space_before = Pt(2)
    p_p2.paragraph_format.space_after = Pt(2)
    r_p2 = p_p2.add_run("PHẦN II: ")
    r_p2.bold = True
    r_p2.font.color.rgb = COLOR_DARK_GREEN
    p_p2.add_run("Câu 1: a-Đ, b-Đ, c-Đ, d-S   |   Câu 2: a-S, b-Đ, c-S, d-S\nCâu 3: a-Đ, b-Đ, c-Đ, d-S   |   Câu 4: a-Đ, b-S, c-S, d-S")

    p_p3 = doc.add_paragraph()
    p_p3.paragraph_format.space_before = Pt(2)
    p_p3.paragraph_format.space_after = Pt(2)
    r_p3 = p_p3.add_run("PHẦN III: ")
    r_p3.bold = True
    r_p3.font.color.rgb = COLOR_DARK_GREEN
    p_p3.add_run("Câu 1: 5   |   Câu 2: 1 2 4   |   Câu 3: 3142   |   Câu 4: 3 5   |   Câu 5: 25   |   Câu 6: 65,0")


def generate_de18_docx(is_teacher=False):
    doc = docx.Document()
    add_header_footer(doc)
    create_title_block(doc, is_teacher=is_teacher)

    img_dir = os.path.join(OUT_DIR, "images_crop_de18")

    # ==================== PHẦN I ====================
    add_part_heading(doc, "PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn. Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án.")

    # Câu 1
    render_mcq(doc, 1,
        "Ở điều kiện thường, kim loại nào sau đây ở trạng thái lỏng?",
        [("A", "Au."), ("B", "Cu."), ("C", "Na."), ("D", "Hg.")],
        "D",
        "Ở điều kiện thường (25 °C), thuỷ ngân (Hg) là kim loại duy nhất tồn tại ở trạng thái lỏng (nhiệt độ nóng chảy là -38,83 °C).",
        is_teacher=is_teacher
    )

    # Câu 2
    render_mcq(doc, 2,
        "Kim loại có độ cứng lớn nhất là",
        [("A", "Cr."), ("B", "Al."), ("C", "Mg."), ("D", "Na.")],
        "A",
        "Chromium (Cr) là kim loại có độ cứng lớn nhất (đạt độ cứng bằng 9 trên thang Mohs, có thể rạch được thuỷ tinh).",
        is_teacher=is_teacher
    )

    # Câu 3
    render_mcq(doc, 3,
        "Kim loại có nhiệt độ nóng chảy cao nhất là",
        [("A", "Na."), ("B", "Pb."), ("C", "Hg."), ("D", "W.")],
        "D",
        "Tungsten (W - vonfam) có nhiệt độ nóng chảy cao nhất trong các kim loại (3422 °C), được dùng làm dây tóc bóng đèn sợi đốt.",
        is_teacher=is_teacher
    )

    # Câu 4
    render_mcq(doc, 4,
        "Kim loại có khối lượng riêng nhỏ nhất là",
        [("A", "Na."), ("B", "Li."), ("C", "Fe."), ("D", "Al.")],
        "B",
        "Lithium (Li) là kim loại có khối lượng riêng nhỏ nhất (D = 0,534 g/cm^3~, nhẹ hơn nước và nhẹ nhất trong tất cả các kim loại).",
        is_teacher=is_teacher
    )

    # Câu 5
    render_mcq(doc, 5,
        "Trong mạng tinh thể kim loại chứa các cation kim loại được sắp xếp theo một trật tự nhất định cùng với các electron chuyển động",
        [("A", "theo một quỹ đạo xác định."),
         ("B", "xung quanh một vị trí xác định."),
         ("C", "tự do trong toàn bộ mạng tinh thể."),
         ("D", "trong một khu vực không gian nhất định.")],
        "C",
        "Trong mạng tinh thể kim loại, các electron hoá trị tách khỏi nguyên tử trở thành các electron tự do, chuyển động hỗn loạn không ngừng trong toàn bộ không gian giữa các ion dương kim loại.",
        is_teacher=is_teacher
    )

    # Câu 6
    render_mcq(doc, 6,
        "Hình dưới đây mô tả tính chất vật lí nào của kim loại? (hình tròn to mô tả ion kim loại, hình tròn nhỏ mô tả electron tự do)",
        [("A", "Tính dẻo."), ("B", "Tính dẫn điện."), ("C", "Tính dẫn nhiệt."), ("D", "Tính cứng.")],
        "B",
        "Khi đặt một điện trường ngoài vào kim loại (cực âm bên trái, cực dương bên phải), các electron tự do mang điện tích âm chuyển động định hướng từ cực âm về cực dương tạo thành dòng điện --> Hình vẽ mô tả tính dẫn điện của kim loại.",
        is_teacher=is_teacher,
        img_path=os.path.join(img_dir, "p1_cau6_tinh_dan_dien.png"),
        img_width=Inches(3.2)
    )

    # Câu 7
    render_mcq(doc, 7,
        "Cầu chì là một phần tử hay thiết bị bảo vệ mạch điện bằng cách làm đứt mạch điện. Cầu chì được sử dụng nhằm phòng tránh các hiện tượng quá tải trên đường dây. Dây cầu chì thường được làm bằng kim loại chì (Pb), thiếc trắng (Sn) hoặc cadmium (Cd).\nỨng dụng này dựa trên tính chất nào của các kim loại trên?",
        [("A", "Có tính dẻo cao."),
         ("B", "Có nhiệt độ nóng chảy tương đối thấp."),
         ("C", "Có độ dẫn điện cao."),
         ("D", "Có độ cứng tương đối thấp.")],
        "B",
        "Chì (327 °C), thiếc (232 °C), cadmium (321 °C) có nhiệt độ nóng chảy tương đối thấp. Khi xảy ra quá tải mạch điện, dòng điện sinh nhiệt lớn làm nóng chảy dây chì trước tiên, ngắt mạch để bảo vệ hệ thống.",
        is_teacher=is_teacher,
        img_path=os.path.join(img_dir, "p1_cau7_cau_chi.png"),
        img_width=Inches(3.4)
    )

    # Câu 8
    render_mcq(doc, 8,
        "Trong nước nguyên chất khi phản ứng xảy ra sẽ có quá trình sau:\n2H~2~O + 2e --> H~2~ + 2OH^-^   E^o^_2H~2~O/H~2~+2OH^-^ = -0,42 V.\nCho E^o^_Na^+^/Na = -2,71 V; E^o^_Cu^2+^/Cu = +0,34 V; E^o^_Mg^2+^/Mg = -2,36 V; E^o^_Al^3+^/Al = -1,68 V.\nHãy cho biết kim loại nào sau đây không phản ứng được với nước ở điều kiện chuẩn.",
        [("A", "Cu."), ("B", "Na."), ("C", "Mg."), ("D", "Al.")],
        "A",
        "Để khử được nước ở điều kiện chuẩn tạo khí H~2~, thế điện cực chuẩn của cặp oxi hoá - khử của kim loại phải âm hơn thế điện cực khử của nước ở pH = 7 (E^o^ = -0,42 V). Cặp Cu^2+^/Cu có E^o^ = +0,34 V > -0,42 V nên Cu không thể khử được nước ở điều kiện chuẩn.",
        is_teacher=is_teacher
    )

    # Câu 9
    render_mcq(doc, 9,
        "Cho các kim loại: Na, Ca, Cu, Ag, K. Số kim loại có khả năng khử được nước ở điều kiện thường giải phóng H~2~ là",
        [("A", "2."), ("B", "1."), ("C", "4."), ("D", "3.")],
        "D",
        "Có 3 kim loại có khả năng phản ứng trực tiếp với nước ở điều kiện thường giải phóng khí H~2~ là Na, Ca, K. Kim loại Cu và Ag có thế điện cực chuẩn dương, không phản ứng với nước.",
        is_teacher=is_teacher
    )

    # Câu 10
    render_mcq(doc, 10,
        "Lấy một đinh sắt đã được làm sạch bề mặt rồi ngâm vào dung dịch CuSO~4~ một thời gian. Hiện tượng nào sau đây không xuất hiện trong thí nghiệm trên?",
        [("A", "Có chất rắn màu đỏ bám lên đinh sắt."),
         ("B", "Màu xanh của dung dịch nhạt dần."),
         ("C", "Có bọt khí thoát ra mạnh trên bề mặt đinh sắt."),
         ("D", "Đinh sắt bị ăn mòn một phần.")],
        "C",
        "Phản ứng xảy ra: Fe + CuSO~4~ --> FeSO~4~ + Cu. Kim loại đồng màu đỏ bám vào đinh sắt, màu xanh lam của ion Cu^2+^ nhạt dần, đinh sắt tan một phần. Trong thí nghiệm này không có khí sinh ra, do đó hiện tượng 'có bọt khí thoát ra mạnh' là không xuất hiện.",
        is_teacher=is_teacher
    )

    # Câu 11
    render_mcq(doc, 11,
        "Nguyên tắc chung để điều chế kim loại từ hợp chất của chúng là",
        [("A", "khử cation kim loại thành kim loại."),
         ("B", "oxi hoá ion kim loại thành kim loại."),
         ("C", "oxi hoá kim loại thành ion kim loại."),
         ("D", "khử kim loại thành ion kim loại.")],
        "A",
        "Trong hợp chất, kim loại tồn tại ở dạng cation M^n+^. Nguyên tắc chung để điều chế kim loại là khử cation kim loại thành kim loại tự do: M^n+^ + ne --> M.",
        is_teacher=is_teacher
    )

    # Câu 12
    render_mcq(doc, 12,
        "Gang và thép là hai hợp kim quan trọng nhất của sắt. Chúng có rất nhiều ứng dụng trong đời sống và sản xuất. Để sản xuất gang, người ta khử oxide của sắt (Fe~2~O~3~ hoặc Fe~3~O~4~) trong quặng sắt bằng khí CO (được tạo ra do đốt cháy than cốc). Sắt được tạo ra trong quá trình sản xuất gang bằng phương pháp nào?",
        [("A", "Kết tinh."), ("B", "Nhiệt luyện."), ("C", "Thuỷ luyện."), ("D", "Chiết.")],
        "B",
        "Phương pháp dùng chất khử (như CO, C, H~2~, Al) để khử cation kim loại trong oxide ở nhiệt độ cao được gọi là phương pháp nhiệt luyện.",
        is_teacher=is_teacher
    )

    # Câu 13
    render_mcq(doc, 13,
        "Kim loại kẽm (zinc, Zn) được sản xuất trong công nghiệp từ quặng sphalerite (có thành phần chính là ZnS) theo sơ đồ: ZnS --(+O~2~, t^o^)--> ZnO --(+C, t^o^)--> Zn. Phương pháp điều chế kim loại nào đã được sử dụng trong quá trình sản xuất zinc theo sơ đồ trên?",
        [("A", "Thuỷ luyện."), ("B", "Nhiệt luyện."), ("C", "Điện phân."), ("D", "Kết tinh.")],
        "B",
        "Giai đoạn tạo kim loại Zn là phản ứng dùng than C khử oxide ZnO ở nhiệt độ cao: ZnO + C --(t^o^)--> Zn + CO. Đây là phương pháp nhiệt luyện.",
        is_teacher=is_teacher
    )

    # Câu 14
    render_mcq(doc, 14,
        "Vàng (Au) tồn tại trong tự nhiên ở dạng đơn chất. Tuy nhiên, hàm lượng Au trong quặng hoặc trong đất thường rất thấp vì vậy rất khó tách Au bằng phương pháp cơ học. Trong công nghiệp, người ta tách vàng từ quặng theo sơ đồ sau:\nQuặng chứa vàng (Au) --(+O~2~ + KCN + H~2~O)--> K[Au(CN)~2~] (aq) --(+Zn dư)--> Au (s)\nPhương pháp điều chế kim loại nào đã được sử dụng trong quá trình sản xuất vàng theo sơ đồ trên?",
        [("A", "Thuỷ luyện."), ("B", "Nhiệt luyện."), ("C", "Điện phân."), ("D", "Chiết.")],
        "A",
        "Phương pháp hoà tan vàng bằng dung dịch xyanua (KCN) tạo dung dịch phức tan, sau đó dùng kim loại khử mạnh hơn (Zn) đẩy kim loại vàng ra khỏi dung dịch được gọi là phương pháp thuỷ luyện.",
        is_teacher=is_teacher
    )

    # Câu 15
    render_mcq(doc, 15,
        "Thực hiện thí nghiệm sau:\nBước 1: Cho dung dịch NaCl 5% vào ống thuỷ tinh hình chữ U như hình bên.\nBước 2: Nhúng một thanh copper và một thanh zinc đã làm sạch vào hai đầu của ống chữ U.\nBước 3: Nối hai thanh kim loại bằng dây dẫn.\nSau bước 2, kim loại nào bị ăn mòn?",
        [("A", "Đồng."),
         ("B", "Cả hai đều bị ăn mòn."),
         ("C", "Kẽm."),
         ("D", "Không kim loại nào bị ăn mòn.")],
        "D",
        "Sau bước 2 chưa nối dây dẫn nên chưa tạo mạch kín của pin điện hoá (không có ăn mòn điện hoá). Đồng thời, dung dịch NaCl là môi trường trung tính, cả Cu và Zn đều không phản ứng với nước hay NaCl ở điều kiện thường --> Không xảy ra phản ứng hoá học --> Không kim loại nào bị ăn mòn.",
        is_teacher=is_teacher,
        img_path=os.path.join(img_dir, "p1_cau15_u_tube.png"),
        img_width=Inches(1.8)
    )

    # Câu 16
    render_mcq(doc, 16,
        "Để khảo sát quá trình ăn mòn kim loại, tiến hành ba thí nghiệm như hình dưới đây. Mỗi ống nghiệm chứa 1 chiếc đinh sắt trong các môi trường khác nhau. Phát biểu nào dưới đây là đúng?",
        [("A", "Sau một thời gian, đinh sắt trong cả ba ống nghiệm đều bị gỉ."),
         ("B", "Sau một thời gian, đinh sắt trong cả ba ống nghiệm đều không có sự thay đổi."),
         ("C", "Sau một thời gian, đinh sắt trong ống nghiệm số 1 và 3 bị gỉ."),
         ("D", "Sau một thời gian, chỉ có đinh sắt trong ống nghiệm số 1 bị gỉ.")],
        "D",
        "Ăn mòn sắt trong không khí ẩm đòi hỏi đồng thời cả nước và oxygen:\n- Ống 1: Có cả nước và oxygen trong không khí --> đinh sắt bị gỉ.\n- Ống 2: Nước cất đun sôi đã đuổi hết oxygen và phủ lớp dầu ngăn cách --> không có oxygen --> đinh sắt không bị gỉ.\n- Ống 3: Có hạt canxi clorua hút ẩm, không khí khô kín --> không có nước --> đinh sắt không bị gỉ.\nVậy chỉ có đinh sắt trong ống 1 bị gỉ.",
        is_teacher=is_teacher,
        img_path=os.path.join(img_dir, "p1_cau16_an_mon_ong_nghiem.png"),
        img_width=Inches(2.5)
    )

    # Câu 17
    render_mcq(doc, 17,
        "Trong số các thí nghiệm sau đây, thí nghiệm có tốc độ ăn mòn đinh sắt chậm nhất?",
        [("A", "(a)."), ("B", "(b)."), ("C", "(c)."), ("D", "(d).")],
        "B",
        "Ống (b) ngâm chìm hoàn toàn trong dầu thực vật, cách li hoàn toàn đinh sắt với nước và không khí ẩm nên hầu như không xảy ra quá trình ăn mòn kim loại --> Tốc độ ăn mòn chậm nhất.",
        is_teacher=is_teacher,
        img_path=os.path.join(img_dir, "p1_cau17_bon_ong_nghiem.png"),
        img_width=Inches(3.8)
    )

    # Câu 18
    render_mcq(doc, 18,
        "Tiến hành thí nghiệm của ba kim loại khác nhau với dung dịch hydrochloric acid loãng theo hình vẽ sau:\nXác định các kim loại phù hợp với hình vẽ (theo thứ tự: Kim loại W, Kim loại X, Kim loại Y)?",
        [("A", "Cu, Mg, Zn."),
         ("B", "Zn, Mg, Cu."),
         ("C", "Mg, Zn, Cu."),
         ("D", "Cu, Zn, Mg.")],
        "A",
        "Quan sát mức độ thoát khí:\n- Ống W: Hoàn toàn không có bọt khí thoát ra --> Kim loại không tác dụng với HCl loãng --> W là Cu.\n- Ống X: Thoát khí rất mạnh và mãnh liệt --> Kim loại có tính khử mạnh nhất --> X là Mg.\n- Ống Y: Thoát khí với tốc độ vừa phải --> Y là Zn.\nVậy thứ tự đúng là Cu, Mg, Zn.",
        is_teacher=is_teacher,
        img_path=os.path.join(img_dir, "p1_cau18_ba_kim_loai_hcl.png"),
        img_width=Inches(3.8)
    )

    # ==================== PHẦN II ====================
    add_part_heading(doc, "PHẦN II. Câu trắc nghiệm đúng sai. Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.")

    # Câu 1
    render_tf(doc, 1,
        "Kim loại Cu có độ dẫn điện kém hơn Ag, kim loại Al có độ dẫn điện kém hơn Au nhưng dây dẫn điện chủ yếu làm bằng kim loại Cu và Al mà không phải là Ag và Au. Ứng dụng làm dây dẫn điện của Cu và Al là do chúng",
        [("a", "đều dẫn điện tốt.", True),
         ("b", "có giá thành phù hợp với thị trường tiêu dùng.", True),
         ("c", "sản xuất được với quy mô công nghiệp.", True),
         ("d", "đều có khối lượng riêng nhỏ.", False)],
        "a) Đúng vì Cu dẫn điện rất tốt (chỉ sau Ag) và Al dẫn điện tốt (sau Ag, Cu, Au).\nb) Đúng vì Cu và Al rẻ hơn rất nhiều so với Ag và Au.\nc) Đúng vì Cu và Al có trữ lượng dồi dào, luyện kim được với quy mô công nghiệp lớn.\nd) Sai vì chỉ có Al là kim loại nhẹ (D = 2,7 g/cm^3~ < 5 g/cm^3~), còn Cu là kim loại nặng (D = 8,96 g/cm^3~ > 5 g/cm^3~).",
        is_teacher=is_teacher
    )

    # Câu 2
    render_tf(doc, 2,
        "Tiến hành thí nghiệm theo các bước sau:\nBước 1. Cho vào 3 ống nghiệm, mỗi ống 2 mL dung dịch H~2~SO~4~ 1 M.\nBước 2. Cho 3 lá kim loại có kích thước như nhau gồm lá nhôm (Al) đã làm sạch lớp bề mặt vào ống nghiệm (1), lá sắt (iron, Fe) vào ống nghiệm (2) và lá đồng (Cu) vào ống nghiệm (3).\nBiết: E^o^_Al^3+^/Al = -1,676 V; E^o^_Fe^2+^/Fe = -0,440 V; E^o^_Cu^2+^/Cu = +0,340 V.",
        [("a", "Ở bước 2, ở cả ba ống nghiệm đều có khí thoát ra.", False),
         ("b", "Tốc độ thoát khí ở ống (1) nhanh hơn ống (2).", True),
         ("c", "Nếu thay H~2~SO~4~ loãng bằng H~2~SO~4~ đặc thì hiện tượng ở bước 2 sẽ không đổi.", False),
         ("d", "Ở bước 2, nếu thêm tiếp 2 mL H~2~SO~4~ 1 M vào cả 3 ống thì tốc độ thoát khí ở cả ba ống sẽ tăng.", False)],
        "a) Sai vì Cu đứng sau H (E^o^ = +0,34 V > 0 V), không phản ứng với H~2~SO~4~ loãng nên ống 3 không có khí thoát ra.\nb) Đúng vì Al có tính khử mạnh hơn Fe (E^o^ = -1,676 V < -0,440 V) nên phản ứng diễn ra nhanh hơn, tốc độ thoát khí ở ống (1) nhanh hơn ống (2).\nc) Sai vì Al và Fe bị thụ động hoá trong H~2~SO~4~ đặc nguội; còn nếu đặc nóng thì phản ứng sinh khí SO~2~ mùi hắc thay vì khí H~2~.\nd) Sai vì thêm 2 mL dung dịch H~2~SO~4~ 1 M thì nồng độ H^+^ vẫn giữ nguyên 1 M nên tốc độ thoát khí không tăng; đồng thời ống 3 Cu vẫn không phản ứng.",
        is_teacher=is_teacher
    )

    # Câu 3
    render_tf(doc, 3,
        "Nhôm phế liệu thường lẫn các tạp chất là các chất hữu cơ và vô cơ (có trong nhãn, mác do in hoặc sơn). Khi tái chế, phế liệu được cắt, băm nhỏ rồi cho vào lò nung đến khi chảy lỏng. Phần lớn các tạp chất này biến thành xỉ lỏng, nổi lên và được vớt ra khỏi lò. Phần còn lại là nhôm tái chế ở trạng thái lỏng và được đưa vào đúc khuôn.",
        [("a", "Việc tái chế nhôm giúp giảm giá thành sản phẩm.", True),
         ("b", "Việc tái chế nhôm giúp giảm chất thải ra môi trường.", True),
         ("c", "Không nên dùng nhôm tái chế để chế tạo dụng cụ nhà bếp và y tế.", True),
         ("d", "Đem cắt, băm nhỏ nhôm phế liệu để quá trình khử Al~2~O~3~ xảy ra dễ hơn.", False)],
        "a) Đúng vì tái chế nhôm tiết kiệm khoảng 95% năng lượng so với sản xuất nhôm từ quặng bauxite, giảm chi phí và giá thành sản phẩm.\nb) Đúng vì tận dụng phế liệu, giảm lượng rác thải rắn và giảm phát thải khí nhà kính ra môi trường.\nc) Đúng vì nhôm tái chế có thể còn chứa tạp chất kim loại nặng độc hại (chì, cadmium...) không đảm bảo an toàn vệ sinh thực phẩm và y tế.\nd) Sai vì tái chế nhôm là quá trình nấu chảy vật lí để loại bỏ xỉ và đúc lại kim loại nhôm, không phải phản ứng hoá học khử Al~2~O~3~.",
        is_teacher=is_teacher
    )

    # Câu 4
    render_tf(doc, 4,
        "Tiến hành thí nghiệm theo các bước sau:\nBước 1. Cho vào cốc thuỷ tinh 30 mL dung dịch H~2~SO~4~ 0,5 M.\nBước 2. Cho một lá nhôm và một lá đồng vào cốc sao cho chúng không tiếp xúc với nhau. Để yên khoảng 1 phút.\nBước 3. Dùng dây dẫn điện nối lá nhôm và lá đồng với một vôn kế (như hình trên).\nBiết: E^o^_Al^3+^/Al = -1,676 V; E^o^_2H^+^/H~2~ = 0 V; E^o^_Cu^2+^/Cu = +0,340 V.",
        [("a", "Ở bước 2, lá nhôm bị ăn mòn hoá học còn lá đồng không bị ăn mòn.", True),
         ("b", "Ở bước 3, có sự tạo thành cặp pin điện hoá, trong đó lá nhôm là cathode và lá đồng là anode.", False),
         ("c", "Ở bước 3, khí chỉ thoát ra ở bề mặt lá đồng, còn ở lá nhôm không có khí thoát ra.", False),
         ("d", "Ở bước 3, kim vôn kế chỉ 1,336 V.", False)],
        "a) Đúng vì ở bước 2 chưa nối dây dẫn, lá Al phản ứng trực tiếp với H~2~SO~4~ (ăn mòn hoá học), lá Cu không tác dụng nên không bị ăn mòn.\nb) Sai vì Al có tính khử mạnh hơn Cu nên Al là cực âm (anode), Cu là cực dương (cathode).\nc) Sai vì ngoài quá trình ăn mòn điện hoá giải phóng khí H~2~ ở Cu (cathode), trên bề mặt Al phản ứng trực tiếp giữa Al và H~2~SO~4~ vẫn diễn ra song song (ăn mòn hoá học), nên khí H~2~ vẫn thoát ra ở lá Al.\nd) Sai vì sức điện động chuẩn của pin: E^o^_pin = E^o^_cathode - E^o^_anode = 0 - (-1,676) = 1,676 V (hoặc nếu đo pin Al-Cu là 0,340 - (-1,676) = 2,016 V), không phải 1,336 V.",
        is_teacher=is_teacher,
        img_path=os.path.join(img_dir, "p2_cau4_pin_al_cu.png"),
        img_width=Inches(2.5)
    )

    # ==================== PHẦN III ====================
    add_part_heading(doc, "PHẦN III: Câu trắc nghiệm yêu cầu trả lời ngắn. Thí sinh trả lời từ câu 1 đến câu 6.")

    # Câu 1
    table_blast_furnace = [
        "Khoảng nhiệt độ (°C)      Phản ứng",
        "400 °C:                   3Fe~2~O~3~ + CO --(t^o^)--> 2Fe~3~O~4~ + CO~2~",
        "500 - 600 °C:             Fe~3~O~4~ + CO --(t^o^)--> 3FeO + CO~2~",
        "700 - 800 °C:             FeO + CO --(t^o^)--> Fe + CO~2~",
        "1000 °C:                  CaCO~3~ --(t^o^)--> CaO + CO~2~",
        "1300 °C:                  CaO + SiO~2~ --(t^o^)--> CaSiO~3~",
        "1500 °C:                  C + CO~2~ --(t^o^)--> 2CO",
        "1800 °C:                  C + O~2~ --(t^o^)--> CO~2~"
    ]
    render_short(doc, 1,
        "Quá trình sản xuất gang từ nguyên liệu là quặng hematite, than cốc, chất chảy trong lò cao xảy ra các phản ứng chính:\nTrong các phản ứng trên có bao nhiêu phản ứng oxi hoá - khử?",
        "5",
        "Các phản ứng oxi hoá - khử (có sự thay đổi số oxi hoá của các nguyên tố):\n(1) 3Fe~2~O~3~ + CO --> 2Fe~3~O~4~ + CO~2~ (thay đổi Fe, C)\n(2) Fe~3~O~4~ + CO --> 3FeO + CO~2~ (thay đổi Fe, C)\n(3) FeO + CO --> Fe + CO~2~ (thay đổi Fe, C)\n(4) C + CO~2~ --> 2CO (thay đổi C)\n(5) C + O~2~ --> CO~2~ (thay đổi C, O)\nHai phản ứng phân huỷ CaCO~3~ và tạo xỉ CaSiO~3~ không có sự thay đổi số oxi hoá.\nTổng cộng có 5 phản ứng oxi hoá - khử.",
        is_teacher=is_teacher,
        extra_table=table_blast_furnace
    )

    # Câu 2
    render_short(doc, 2,
        "Để bảo vệ vật bằng thép khỏi bị ăn mòn điện hoá, trong thực tế người ta đã thực hiện một số cách sau:\n(1) Sơn kín bề mặt cánh cửa làm bằng thép.\n(2) Tráng kẽm lên tấm thép mỏng khi sản xuất tôn.\n(3) Gắn một số tấm kẽm lên vỏ tàu làm bằng thép.\n(4) Cho dầu mỡ lên các ốc vít trên đường ray.\nLiệt kê những cách sử dụng phương pháp bảo vệ bề mặt theo số thứ tự tăng dần.",
        "1 2 4",
        "Phương pháp bảo vệ bề mặt là phủ lên bề mặt kim loại một lớp màng cách li với môi trường: sơn (1), tráng kẽm (2), tra dầu mỡ (4).\nCách (3) là phương pháp bảo vệ điện hoá (dùng kim loại hi sinh Zn).\nVậy các cách bảo vệ bề mặt là: 1 2 4.",
        is_teacher=is_teacher
    )

    # Câu 3
    table_props = [
        "Cột A (Tính chất vật lí):",
        "  (a) Tính dẻo",
        "  (b) Tính dẫn điện",
        "  (c) Tính dẫn nhiệt",
        "  (d) Tính ánh kim",
        "Cột B (Giải thích):",
        "  (1) Dưới tác dụng của điện trường, các electron tự do trong tinh thể kim loại sẽ chuyển động có hướng, tạo nên dòng điện.",
        "  (2) Các electron tự do phản xạ được những tia sáng nhìn thấy.",
        "  (3) Do các electron tự do trong mạng tinh thể kim loại chuyển động, liên kết các ion kim loại lại với nhau.",
        "  (4) Các electron ở vùng có nhiệt độ cao truyền năng lượng cho ion dương và electron khác trong toàn khối kim loại."
    ]
    render_short(doc, 3,
        "Cho các tính chất vật lí ở cột (A) và giải thích ở cột (B) như sau:\nHãy sắp xếp thứ tự giải thích của cột B với thứ tự tính chất vật lí ở cột A cho phù hợp.",
        "3142",
        "Ghép tương ứng:\n- (a) Tính dẻo <--> (3) Electron tự do liên kết các ion kim loại khi các lớp ion trượt lên nhau.\n- (b) Tính dẫn điện <--> (1) Electron tự do chuyển động có hướng dưới tác dụng của điện trường.\n- (c) Tính dẫn nhiệt <--> (4) Electron vùng nhiệt độ cao truyền năng lượng cho ion dương và electron khác.\n- (d) Tính ánh kim <--> (2) Electron tự do phản xạ hầu hết các tia sáng nhìn thấy chiếu tới.\nThứ tự ghép đôi tương ứng với (a)-(b)-(c)-(d) là: 3142.",
        is_teacher=is_teacher,
        extra_table=table_props
    )

    # Câu 4
    render_short(doc, 4,
        "Nhôm (Al) là kim loại phổ biến nhất trong vỏ Trái Đất. Nhôm được sử dụng để chế tạo hợp kim nhẹ, bền, được dùng trong kĩ thuật hàng không, trong xây dựng và làm nội thất. Nhôm có số hiệu nguyên tử là 13 và E^o^_Al^3+^/Al = -1,676 V. Cho các phát biểu sau:\n(1) Nhôm và đồ vật làm bằng nhôm dễ bị ăn mòn trong không khí ẩm.\n(2) Ở trạng thái cơ bản, nguyên tử Al có 1 electron ở lớp ngoài cùng.\n(3) Nhôm phản ứng được với nước ngay ở điều kiện chuẩn.\n(4) Nhôm có tính khử mạnh hơn Mg (biết E^o^_Mg^2+^/Mg = -2,356 V).\n(5) Trong vỏ Trái Đất, nhôm chỉ tồn tại ở dạng hợp chất.\nLiệt kê các phát biểu đúng theo số thứ tự tăng dần.",
        "3 5",
        "(1) Sai vì nhôm được bảo vệ bởi lớp màng oxide Al~2~O~3~ mỏng, mịn và rất bền, ngăn cản nhôm tiếp xúc với không khí ẩm.\n(2) Sai vì Al có cấu hình electron [Ne]3s^2~3p^1~, có 3 electron ở lớp ngoài cùng.\n(3) Đúng vì E^o^_Al^3+^/Al = -1,676 V < -0,42 V nên về mặt nhiệt động học, Al có khả năng phản ứng với nước ở điều kiện chuẩn (nhưng phản ứng dừng lại nhanh chóng do tạo màng Al(OH)~3~ bảo vệ).\n(4) Sai vì E^o^_Mg^2+^/Mg = -2,356 V < E^o^_Al^3+^/Al = -1,676 V nên Mg có tính khử mạnh hơn Al.\n(5) Đúng vì Al hoạt động hoá học mạnh nên trong tự nhiên chỉ tồn tại dưới dạng hợp chất.\nVậy các phát biểu đúng là: 3 5.",
        is_teacher=is_teacher
    )

    # Câu 5
    render_short(doc, 5,
        "Để mạ đồng một vật dụng kim loại có tổng diện tích bề mặt là 10 cm^2~, người người ta tiến hành điện phân dung dịch CuSO~4~ với cực âm là vật dụng cần mạ và cực dương là lá đồng thô. Biết cường độ dòng điện không đổi là 2 A, hiệu suất điện phân là 90%, khối lượng riêng của tinh thể Cu là 8,94 g/cm^3~ và lượng đồng tạo ra được tính theo công thức Faraday là m = A.I.t / (n.F) (với A là nguyên tử khối của Cu = 64; I là cường độ dòng điện, F là hằng số Faraday = 96485 C/mol, n là số electron mà 1 ion Cu^2+^ nhận, t là thời gian điện phân tính bằng giây). Thời gian điện phân để lớp mạ có độ dày đồng nhất 0,1 mm là bao nhiêu phút? (Làm tròn kết quả đến hàng đơn vị)",
        "25",
        "Đổi độ dày: d = 0,1 mm = 0,01 cm.\nThể tích lớp mạ đồng: V = S . d = 10 . 0,01 = 0,1 cm^3~.\nKhối lượng Cu cần kết tủa: m = V . D = 0,1 . 8,94 = 0,894 gam.\nTheo công thức Faraday có tính hiệu suất điện phân H = 90%:\nm = (A . I . t / n . F) . (H / 100%)\n=> 0,894 = (64 . 2 . t / (2 . 96485)) . 0,9\n=> t = (0,894 . 96485) / (64 . 0,9) ≈ 1497,53 s.\nĐổi ra phút: t = 1497,53 / 60 ≈ 24,96 phút ≈ 25 phút.",
        is_teacher=is_teacher
    )

    # Câu 6
    render_short(doc, 6,
        "Đồng thau là hợp kim của đồng và kẽm (các tạp chất khác không đáng kể). Đồng thau được sử dụng để trang trí, làm ổ khoá, bánh răng, vòng bi, tay nắm cửa, vỏ đạn, các hệ thống ống nước, điện và một số nhạc cụ. Để xác định hàm lượng Cu trong một loại đồng thau, người ta cho 10 gam mảnh đồng thau vào dung dịch H~2~SO~4~ (loãng, dư); sau khi các phản ứng xảy ra hoàn toàn, lọc lấy phần chất rắn đem làm khô, cân lại thu được 6,5 gam chất rắn. Hàm lượng đồng trong loại đồng thau trên bằng bao nhiêu? (Làm tròn kết quả đến hàng phần mười)",
        "65,0",
        "Khi cho đồng thau (Cu, Zn) vào dung dịch H~2~SO~4~ loãng dư, chỉ có Zn phản ứng tan vào dung dịch:\nZn + H~2~SO~4~ --> ZnSO~4~ + H~2~↑\nKim loại Cu đứng sau H trong dãy điện hoá nên không tác dụng với H~2~SO~4~ loãng.\nChất rắn không tan còn lại chính là Cu: m_Cu = 6,5 gam.\nHàm lượng đồng trong đồng thau là:\n%Cu = (6,5 / 10) . 100% = 65,0%.",
        is_teacher=is_teacher
    )

    if is_teacher:
        add_answer_key_table(doc)

    out_filename = "De_Kiem_Tra_Hoa_12_De18_GiaoVien.docx" if is_teacher else "De_Kiem_Tra_Hoa_12_De18.docx"
    out_path = os.path.join(OUT_DIR, out_filename)
    doc.save(out_path)
    print(f"Đã tạo thành công file: {out_path}")

if __name__ == "__main__":
    generate_de18_docx(is_teacher=False)
    generate_de18_docx(is_teacher=True)
