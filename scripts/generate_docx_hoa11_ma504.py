# -*- coding: utf-8 -*-
"""
Script tạo file Word (.docx) chuẩn "Form Đề Kiểm Tra" cho:
ĐỀ KIỂM TRA GIỮA HỌC KỲ I - MÔN HOÁ HỌC 11 CTST - TRƯỜNG THPT PHONG ĐIỀN - MÃ ĐỀ 504
Áp dụng 100% quy chuẩn:
- Không dùng bảng trong nội dung
- Header & Footer chuẩn
- Tab stops chuẩn
- Subscript/superscript triệt để
- Xuất bản 2 file: Bản Học sinh và Bản Giáo viên
"""

import os
import sys
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Thêm đường dẫn scripts vào sys.path để import form_de_kiem_tra
sys.path.insert(0, os.path.dirname(__file__))
from form_de_kiem_tra import (
    setup_header_footer,
    build_title_block,
    add_section_header,
    render_question_prompt,
    render_options_no_table,
    render_statement,
    add_solution_no_table,
    add_student_writing_lines,
    render_rich_text,
    COLOR_CORRECT,
    COLOR_PRIMARY
)

def create_exam_doc_504(is_teacher=False):
    doc = docx.Document()
    setup_header_footer(doc)

    build_title_block(
        doc,
        school_line="SỞ GD&ĐT THÀNH PHỐ HUẾ\nTRƯỜNG THPT PHONG ĐIỀN",
        exam_title="KIỂM TRA GIỮA HỌC KỲ I - NĂM HỌC 2025 - 2026",
        subject_line="MÔN: HOÁ HỌC 11 CTST",
        num_pages="04",
        duration="45 phút",
        exam_code="504",
        atomic_masses="H = 1; C = 12; N = 14; O = 16; Na = 23; S = 32; Cl = 35,5; Fe = 56",
        is_teacher=is_teacher
    )

    # ==================== PHẦN I ====================
    add_section_header(
        doc,
        title="PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 12. Mỗi câu hỏi thí sinh chỉ chọn một phương án."
    )

    p1_questions = [
        {
            "q": "Câu 1. Cân bằng hóa học là một cân bằng",
            "opts": [
                "A. Đứng yên",
                "B. Tĩnh",
                "C. Cố định",
                "*D. Động*"
            ],
            "sol": "► *Lời giải chi tiết:*\nCân bằng hóa học là một cân bằng động vì tại trạng thái cân bằng, phản ứng thuận và phản ứng nghịch vẫn tiếp tục diễn ra với tốc độ bằng nhau (v~thuận~ = v~nghịch~ ≠ 0).\n*Chọn D.*"
        },
        {
            "q": "Câu 2. Trong thí nghiệm đốt cháy sulfur (S) trong không khí, hiện tượng nào sau đây có thể được quan sát?",
            "opts": [
                "A. Sulfur cháy với ngọn lửa màu đỏ và sinh ra khói trắng",
                "*B. Sulfur cháy với ngọn lửa màu xanh và sinh ra khí không màu, mùi hắc*",
                "C. Sulfur cháy với ngọn lửa màu tím và sinh ra khí không mùi",
                "D. Sulfur cháy với ngọn lửa màu vàng nhạt và sinh ra khí có mùi ngọt"
            ],
            "sol": "► *Lời giải chi tiết:*\nSulfur cháy trong không khí với ngọn lửa màu xanh nhạt, sinh ra sulfur dioxide (SO~2~) là chất khí không màu, mùi hắc: S + O~2~ → SO~2~ (t^o^).\n*Chọn B.*"
        },
        {
            "q": "Câu 3. Nạp đầy khí amoniac vào bình thuỷ tinh trong suốt, đậy bình NH~3~ bằng nút cao su có ống thuỷ tinh vuốt nhọn xuyên qua. Nhúng đầu ống thuỷ tinh vào một chậu thuỷ tinh chứa nước có pha thêm dung dịch phenolphthalein. Một lát sau, nước trong chậu phun vào bình thành những tia có màu hồng. Thí nghiệm này chứng tỏ được:",
            "opts": [
                "A. NH~3~ là chất khí.",
                "B. NH~3~ tan nhiều trong nước.",
                "C. NH~3~ có mùi khai.",
                "*D. NH~3~ có tính bazơ và tan nhiều trong nước.*"
            ],
            "sol": "► *Lời giải chi tiết:*\n- Khí NH~3~ tan rất nhiều trong nước làm giảm áp suất trong bình so với áp suất khí quyển, hút nước từ chậu phun mạnh vào bình thành các tia nước.\n- Dung dịch nước có pha phenolphthalein chuyển thành màu hồng chứng tỏ dung dịch NH~3~ có tính base.\nVậy thí nghiệm chứng minh NH~3~ tan nhiều trong nước và dung dịch có tính base.\n*Chọn D.*"
        },
        {
            "q": "Câu 4. Nitrogen monoxide được tạo thành khi mưa dông kèm theo sấm sét do phản ứng giữa nitrogen và oxygen trong không khí được gọi là",
            "opts": [
                "A. NO~x~ nhiên liệu.",
                "B. NO~x~ tức thời.",
                "*C. NO~x~ tự nhiên.*",
                "D. NO~x~ nhiệt."
            ],
            "sol": "► *Lời giải chi tiết:*\nKhí NO hình thành trong cơn mưa dông do sấm sét (hiện tượng tự nhiên) giữa N~2~ và O~2~ được gọi là NO~x~ tự nhiên.\n*Chọn C.*"
        },
        {
            "q": "Câu 5. Cho các nhận định sau về cấu tạo phân tử nitric acid:\n(a) Liên kết O – H phân cực về oxygen.\n(b) Nguyên tử N có số oxi hóa là +5.\n(c) Nguyên tử N có hóa trị bằng IV.\n(d) Liên kết cho – nhận N → O kém bền.\nSố nhận định đúng là",
            "opts": [
                "*A. 4.*",
                "B. 1.",
                "C. 2.",
                "D. 3."
            ],
            "sol": "► *Lời giải chi tiết:*\n- (a) Đúng vì độ âm điện O (3,44) lớn hơn H (2,20) nên liên kết O - H phân cực mạnh về phía O.\n- (b) Đúng vì trong HNO~3~ số oxi hóa của N là +5.\n- (c) Đúng vì nguyên tử N ở chu kì 2 có 4 orbital hóa trị nên hóa trị tối đa bằng 4 (IV).\n- (d) Đúng vì liên kết phối trí N → O kém bền, khiến phân tử HNO~3~ dễ bị nhiệt và ánh sáng phân hủy sinh ra NO~2~ và O~2~.\nVậy cả 4 nhận định đều đúng.\n*Chọn A.*"
        },
        {
            "q": "Câu 6. Đâu là phát biểu của thuyết Bronsted-Lowry?",
            "opts": [
                "A. Acid là chất nhận proton, base là chất cho proton",
                "B. Cả acid và base đều là chất nhận proton",
                "*C. Acid là chất cho proton, base là chất nhận proton*",
                "D. Cả acid và base đều là chất cho proton"
            ],
            "sol": "► *Lời giải chi tiết:*\nTheo thuyết Brønsted – Lowry, acid là chất cho proton (H^+^), base là chất nhận proton (H^+^).\n*Chọn C.*"
        },
        {
            "q": "Câu 7. Xét cân bằng: N~2~(g) + 3H~2~(g) ⇌ 2NH~3~(g). Biểu thức hằng số cân bằng của phản ứng là:",
            "opts": [
                "A. K = [N~2~][H~2~] / [NH~3~]",
                "B. K = [NH~3~] / ([N~2~][H~2~])",
                "*C. K = [NH~3~]^2^ / ([N~2~][H~2~]^3^)*",
                "D. K = ([N~2~][H~2~]^3^) / [NH~3~]^2^"
            ],
            "sol": "► *Lời giải chi tiết:*\nBiểu thức hằng số cân bằng KC bằng tích nồng độ sản phẩm mũ hệ số chia tích nồng độ chất phản ứng mũ hệ số tương ứng: K = [NH~3~]^2^ / ([N~2~][H~2~]^3^).\n*Chọn C.*"
        },
        {
            "q": "Câu 8. Hợp chất nào sau đây của sulfur được sử dụng để khử trùng và tẩy trắng trong công nghiệp?",
            "opts": [
                "*A. SO~2~*",
                "B. H~2~SO~4~",
                "C. S~8~",
                "D. H~2~S"
            ],
            "sol": "► *Lời giải chi tiết:*\nSO~2~ (sulfur dioxide) được sử dụng để khử trùng nấm mốc lương thực thực phẩm và tẩy trắng bột giấy, sợi vải trong công nghiệp.\n*Chọn A.*"
        },
        {
            "q": "Câu 9. Trong tự nhiên luôn diễn ra các quá trình chuyển hóa nitrogen từ dạng này sang dạng khác theo...",
            "opts": [
                "A. Một chu trình tuần hoàn không khép kín",
                "*B. Một chu trình tuần hoàn khép kín*",
                "C. Một chu trình ngẫu nhiên",
                "D. Một chu trình không tuần hoàn"
            ],
            "sol": "► *Lời giải chi tiết:*\nTrong tự nhiên, nitrogen liên tục luân chuyển giữa khí quyển, đất, sinh vật và nước theo một chu trình tuần hoàn khép kín (chu trình nitrogen).\n*Chọn B.*"
        },
        {
            "q": "Câu 10. pH là",
            "opts": [
                "A. Chỉ số đánh giá độ acid của một dung dịch",
                "*B. Chỉ số đánh giá độ acid hay độ base của một dung dịch*",
                "C. Chỉ số đánh giá độ base của một dung dịch",
                "D. Chỉ số đánh giá các chất điện li mạnh"
            ],
            "sol": "► *Lời giải chi tiết:*\npH là chỉ số dùng để đánh giá độ acid hay độ base của một dung dịch (pH = -log[H^+^]).\n*Chọn B.*"
        },
        {
            "q": "Câu 11. Tính chất nào sau đây là tính chất vật lý của sulfur (S)?",
            "opts": [
                "A. Là chất khí không màu",
                "*B. Là chất rắn màu vàng, không tan trong nước*",
                "C. Tan tốt trong nước",
                "D. Dễ bay hơi ở nhiệt độ thường"
            ],
            "sol": "► *Lời giải chi tiết:*\nỞ điều kiện thường, sulfur là chất rắn màu vàng, giòn, không tan trong nước.\n*Chọn B.*"
        },
        {
            "q": "Câu 12. Cho các nhận định sau về tính chất hoá học của nitric acid: (1) có tính acid mạnh; (2) có tính acid yếu; (3) có tính oxi hoá mạnh; (4) có tính khử mạnh. Số nhận định đúng là",
            "opts": [
                "A. 4.",
                "B. 1.",
                "*C. 2.*",
                "D. 3."
            ],
            "sol": "► *Lời giải chi tiết:*\nHNO~3~ là một acid mạnh và có tính oxi hóa rất mạnh (do N đạt số oxi hóa cao nhất +5). Do đó có 2 nhận định đúng: (1) và (3).\n*Chọn C.*"
        }
    ]

    for item in p1_questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, item["q"])

        render_options_no_table(doc, item["opts"], is_teacher=is_teacher)

        if is_teacher:
            add_solution_no_table(doc, item["sol"])

    # ==================== PHẦN II ====================
    add_section_header(
        doc,
        title="PHẦN II. Câu trắc nghiệm đúng sai. (2 điểm)",
        subtitle="Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai. Điểm tối đa của 01 câu hỏi là 1 điểm:\n- Thí sinh chỉ lựa chọn chính xác 01 ý trong 1 câu hỏi được 0,1 điểm.\n- Thí sinh chỉ lựa chọn chính xác 02 ý trong 1 câu hỏi được 0,25 điểm.\n- Thí sinh chỉ lựa chọn chính xác 03 ý trong 1 câu hỏi được 0,5 điểm.\n- Thí sinh chỉ lựa chọn chính xác 04 ý trong 1 câu hỏi được 1 điểm."
    )

    p2_questions = [
        {
            "lead": "Câu 1. Ammonium nitrate được dùng làm phân đạm. Năm 2020, một vụ nổ tại thủ đô Beirut, Lebanon đã cướp đi sinh mạng của hàng trăm người. Nguyên nhân của vụ nổ được do sự phân huỷ 2750 tấn ammonium nitrate trên một tàu hàng bỏ hoang theo phương trình hoá học sau:\nNH~4~NO~3~(s) → N~2~O(g) + 2H~2~O(g)   Δ~r~H~298~^o^ = -36 kJ",
            "stmts": [
                ("a) Phản ứng (1) là phản ứng thu nhiệt.", "Sai"),
                ("b) Dung dịch muối Ammonium nitrate phản ứng với dung dịch base đặc, nóng thoát ra chất khí làm quỳ tím ẩm hoá đỏ.", "Sai"),
                ("c) Khi bón phân đạm ammonium nitrate nhiều sẽ gây hiện tượng chua đất.", "Đúng"),
                ("d) Nhiệt của vụ nổ trên tương đương với nhiệt toả ra 1062 tấn thuốc nổ TNT. (biết 1 kg thuốc nổ TNT phát nổ toả ra 1,165.10^6 J ).", "Đúng")
            ],
            "sol": "► *Lời giải chi tiết:*\na) Sai vì ΔH = -36 kJ < 0 nên đây là phản ứng tỏa nhiệt.\nb) Sai vì NH~4~^+^ + OH^-^ → NH~3~↑ + H~2~O (t^o^). Khí NH~3~ có tính base làm quỳ tím ẩm hóa xanh.\nc) Đúng vì ion NH~4~^+^ bị thủy phân trong đất tạo H~3~O^+^: NH~4~^+^ + H~2~O ⇌ NH~3~ + H~3~O^+^, làm giảm pH của đất gây chua đất.\nd) Đúng. Khối lượng NH~4~NO~3~: m = 2750 tấn = 2,75.10^9 gam.\nSố mol NH~4~NO~3~: n = 2,75.10^9 / 80 = 34 375 000 mol.\nNhiệt lượng tỏa ra: Q = 34 375 000 × 36 kJ = 1,2375.10^12 J.\nKhối lượng thuốc nổ TNT tương đương:\nm~TNT~ = 1,2375.10^12 / (1,165.10^6) ≈ 1 062 232 kg ≈ 1062 tấn."
        },
        {
            "lead": "Câu 2. Một học sinh làm thí nghiệm xác định pH của đất như sau: lấy một lượng đất cho vào nước rồi lọc lấy phần dung dịch. Dùng máy đo pH đo được pH = 4,69.",
            "stmts": [
                ("a) Môi trường của dung dịch là acid.", "Đúng"),
                ("b) Loại đất trên được gọi là đất chua. Để giảm độ chua cho đất có thể bón vôi.", "Đúng"),
                ("c) Dung dịch trong cốc có [OH^-^] > [H^+^] vì pOH > pH.", "Sai"),
                ("d) Nồng độ [H^+^] trong cốc lớn hơn 0,001.", "Sai")
            ],
            "sol": "► *Lời giải chi tiết:*\na) Đúng vì pH = 4,69 < 7 nên môi trường là acid.\nb) Đúng vì đất có pH < 7 là đất chua. Bón vôi (CaO, Ca(OH)~2~) giúp trung hòa acid, giảm độ chua cho đất.\nc) Sai vì pOH = 14 - 4,69 = 9,31 > pH = 4,69 ⇒ [H^+^] = 10^-4,69^ ≈ 2,04.10^-5^ M > [OH^-^] = 10^-9,31^ ≈ 4,9.10^-10^ M.\nd) Sai vì [H^+^] = 10^-4,69^ ≈ 2,04.10^-5^ M < 0,001 M (10^-3^ M)."
        }
    ]

    for item in p2_questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, item["lead"])

        for stmt, ans in item["stmts"]:
            p_s = doc.add_paragraph()
            p_s.paragraph_format.left_indent = Cm(0.5)
            p_s.paragraph_format.space_before = Pt(1)
            p_s.paragraph_format.space_after = Pt(1)
            render_statement(p_s, stmt, ans=ans, is_teacher=is_teacher)

        if is_teacher:
            add_solution_no_table(doc, item["sol"])

    # ==================== PHẦN III ====================
    add_section_header(
        doc,
        title="PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn. (2 điểm)",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 4."
    )

    p3_questions = [
        {
            "q": "Câu 1. Tổng số liên kết σ và số liên kết π trong phân tử nitrogen là bao nhiêu?",
            "ans": "3",
            "sol": "► *Lời giải chi tiết:*\nPhân tử nitrogen (N~2~) có công thức cấu tạo N≡N gồm 1 liên kết σ và 2 liên kết π.\nTổng số liên kết σ và liên kết π = 1 + 2 = 3.\n► *Đáp án: 3*"
        },
        {
            "q": "Câu 2. Để giảm thải lượng khí thải SO~2~ ra môi trường, trong các biện pháp sau:\n(1) Sử dụng nguồn nhiên liệu sinh học thân thiện với môi trường như hydrogen, ethanol.\n(2) Sử dụng nguồn năng lượng hóa thạch.\n(3) Sử dụng nguồn năng lượng tái tạo như năng lượng mặt trời, năng lượng gió, mưa, sóng biển, thủy triều, địa nhiệt.\n(4) Xử lý khí thải công nghiệp trước khi thải ra môi trường.\nCần thực hiện biện pháp nào trong số các biện pháp trên?",
            "ans": "3",
            "sol": "► *Lời giải chi tiết:*\n- (1) Sử dụng nhiên liệu sinh học không chứa sulfur giúp giảm SO~2~ (thực hiện).\n- (2) Dùng năng lượng hóa thạch phát thải nhiều SO~2~ (không thực hiện).\n- (3) Năng lượng tái tạo sạch không sinh SO~2~ (thực hiện).\n- (4) Xử lý khí thải công nghiệp trước khi xả ra môi trường giúp lọc bỏ SO~2~ (thực hiện).\nCó 3 biện pháp cần thực hiện: (1), (3), (4).\n► *Đáp án: 3*"
        },
        {
            "q": "Câu 3. Dịch vị dạ dày của con người có chứa hydrochloric acid (HCl), với pH dao động trong khoảng 2-4. Những người bị mắc bệnh viêm loét dạ dày, tá tràng thường có pH < 2. Nabica (có thành phần chính là NaHCO~3~) là một loại thuốc được dùng để trung hoà bớt lượng acid trong dạ dày của những bệnh nhân có pH trong dạ dày thấp và thường được uống trước bữa ăn. Giả sử một bệnh nhân bị đau dạ dày đã uống 0,504 gam bột NaHCO~3~. Hãy tính số mol HCl đã được trung hòa.",
            "ans": "0,006",
            "sol": "► *Lời giải chi tiết:*\nSố mol NaHCO~3~ = 0,504 / 84 = 0,006 mol.\nPhương trình phản ứng: NaHCO~3~ + HCl → NaCl + CO~2~↑ + H~2~O\nTheo phản ứng: n~HCl~ = n~NaHCO3~ = 0,006 mol.\n► *Đáp án: 0,006*"
        },
        {
            "q": "Câu 4. Cho các phát biểu sau:\n(a) Trong không khí, Nitrogen chiếm khoảng 78% về thể tích.\n(b) Trong bầu khí quyển, khi có sấm chớp, khí nitrogen tạo các nitrogen oxide, là một nguyên nhân làm cho nước mưa có tính acid.\n(c) Trong phản ứng giữa N~2~ và H~2~ thì N~2~ là chất khử.\n(d) Phần lớn Nitrogen được sử dụng để tổng hợp NH~3~, từ đó sản xuất nitric acid, phân bón.\nSố phát biểu đúng là bao nhiêu?",
            "ans": "3",
            "sol": "► *Lời giải chi tiết:*\n- (a) Đúng: Khí nitrogen chiếm khoảng 78% thể tích không khí.\n- (b) Đúng: Sấm chớp tạo NO rồi chuyển thành NO~2~, hòa tan tạo HNO~3~ làm nước mưa có tính acid.\n- (c) Sai: N~2~ + 3H~2~ ⇌ 2NH~3~, số oxi hóa N giảm từ 0 xuống -3 nên N~2~ là chất oxi hóa.\n- (d) Đúng: Phần lớn N~2~ dùng sản xuất NH~3~ để chế tạo phân đạm và acid nitric.\nVậy có 3 phát biểu đúng: (a), (b), (d).\n► *Đáp án: 3*"
        }
    ]

    for item in p3_questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, item["q"])

        if is_teacher:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Cm(0.5)
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(1)
            render_rich_text(p_ans, f"► *Đáp án:* *{item['ans']}*", base_size=10.5, base_bold=True, base_color=COLOR_CORRECT)
            add_solution_no_table(doc, item["sol"])
        else:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Cm(0.5)
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(3)
            render_rich_text(p_ans, "*Trả lời:* ............................................................", base_size=10.5)

    # ==================== PHẦN IV ====================
    add_section_header(
        doc,
        title="PHẦN IV: Tự luận (3 điểm)",
        subtitle=None
    )

    essay_questions = [
        {
            "num": "Câu 1: (0,5 điểm)",
            "prompt": "Giải thích vì sao khi mở nắp chai nước có gas, dung dịch bị sủi bọt mạnh, sau đó giảm dần.",
            "sol": "► *Lời giải chi tiết:*\n- Trong chai nước có gas đóng kín, khí CO~2~ được nén ở áp suất cao tan vào nước ngọt theo cân bằng: CO~2~(aq) ⇌ CO~2~(g).\n- Khi mở nắp chai, áp suất trên bề mặt chất lỏng giảm đột ngột về áp suất khí quyển. Theo định luật Henry, độ tan của khí CO~2~ giảm mạnh khi áp suất giảm, làm cho một lượng lớn CO~2~ hòa tan nhanh chóng thoát ra khỏi dung dịch tạo hiện tượng sủi bọt mạnh.\n- Sau một thời gian, khí CO~2~ đã thoát ra gần hết và hệ đạt tới trạng thái cân bằng mới ở áp suất khí quyển nên hiện tượng sủi bọt giảm dần rồi dừng hẳn.",
            "lines": 5
        },
        {
            "num": "Câu 2: (0,5 điểm)",
            "prompt": "Một nhà máy thải ra 1000 m^3^ khí chứa 0,05% thể tích SO~2~ mỗi ngày. Tính khối lượng SO~2~ (tính bằng kg) thải ra trong ngày (ở đktc).",
            "sol": "► *Lời giải chi tiết:*\n- Thể tích SO~2~ thải ra trong ngày:\nV~SO2~ = 1000 m^3^ × 0,05% = 0,5 m^3^ = 500 Lít.\n- Số mol SO~2~ thải ra ở đktc (22,4 L/mol):\nn~SO2~ = 500 / 22,4 ≈ 22,3214 mol.\n- Khối lượng SO~2~ thải ra mỗi ngày:\nm~SO2~ = 22,3214 mol × 64 g/mol ≈ 1428,57 gam ≈ 1,429 kg.\n*(Ghi chú: Nếu tính theo điều kiện chuẩn đkc 24,79 L/mol thì m~SO2~ = (500 / 24,79) × 64 ≈ 1290,84 gam ≈ 1,291 kg)*.\n*Đáp số:* 1,429 kg (đktc).",
            "lines": 5
        },
        {
            "num": "Câu 3: (0,5 điểm)",
            "prompt": "Giải thích vì sao khí nitrogen (N~2~) trơ về mặt hóa học ở điều kiện thường.",
            "sol": "► *Lời giải chi tiết:*\n- Phân tử nitrogen gồm 2 nguyên tử N liên kết với nhau bằng một liên kết ba N≡N rất bền vững.\n- Năng lượng liên kết ba N≡N rất lớn (E~b~ = 945 kJ/mol), đòi hỏi năng lượng rất cao để phá vỡ liên kết.\n- Vì vậy, ở nhiệt độ thường nitrogen trơ về mặt hóa học và hầu như không tham gia các phản ứng hóa học (chỉ phản ứng ở nhiệt độ cao hoặc có tia lửa điện).",
            "lines": 5
        },
        {
            "num": "Câu 4: (1,0 điểm)",
            "prompt": "Một khu ruộng rau có lớp đất mặt dày 10 cm, khối lượng riêng 1,2 g/cm^3^, có khả năng trao đổi cation - Cation Exchange Capacity (CEC) là 5 cmolc/kg và tỉ lệ bão hòa bazo là 10%. Trong mùa mưa, tổng lượng mưa là 400 mm, trung bình pH = 2,8 (do ô nhiễm khí NO~2~ và SO~2~ trong không khí). Biết: 1 mm mưa trên 1 ha = 10 m^3^ nước = 10^4^ L; [H^+^] = 10^-pH^ mol/L; 1 mol H^+^ trung hòa 1 mol điện tích bazo (molc).\nCâu hỏi:\n1. Tính số mol H^+^ do lượng mưa này mang xuống 1 ha đất (làm tròn đến hàng chục mol).\n2. Tính lượng mol bazo sẵn có trong lớp đất (10 cm) của 1 ha.\n3. So sánh hai giá trị trên và kết luận xem mưa acid có thể gây hạ pH đất hay không. Giải thích ngắn gọn vì sao pH đất giảm lại có thể gây hại cho cây trồng.",
            "sol": "► *Lời giải chi tiết:*\n1. Tính số mol H^+^ mang xuống 1 ha đất:\n- Thể tích nước mưa trên 1 ha: V = 400 mm × 10^4^ L/mm = 4.10^6^ Lít.\n- Nồng độ H^+^: [H^+^] = 10^-2,8^ mol/L.\n- Số mol H^+^: n~H+~ = 4.10^6^ × 10^-2,8^ ≈ 6339,57 mol.\nLàm tròn đến hàng chục mol: 6340 mol.\n\n2. Tính lượng mol base sẵn có trong lớp đất (10 cm) của 1 ha:\n- Diện tích 1 ha = 10 000 m^2^; thể tích đất: V = 10 000 m^2^ × 0,1 m = 1000 m^3^ = 10^9^ cm^3^.\n- Khối lượng đất: m = 10^9^ cm^3^ × 1,2 g/cm^3^ = 1,2.10^9^ gam = 1,2.10^6^ kg.\n- Tổng khả năng trao đổi cation: CEC~tổng~ = 1,2.10^6^ kg × 0,05 molc/kg = 60 000 molc.\n- Lượng base sẵn có (tỉ lệ bão hòa base 10%):\nn~base~ = 60 000 molc × 10% = 6000 molc.\n\n3. So sánh, kết luận và giải thích:\n- So sánh: Lượng H^+^ do mưa acid mang xuống (6340 mol) lớn hơn lượng base đệm sẵn có trong đất (6000 molc).\n- Kết luận: Mưa acid vượt quá khả năng đệm của đất nên chắc chắn sẽ làm hạ mạnh pH của đất, gây chua hóa đất nghiêm trọng.\n- Giải thích tác hại khi pH đất giảm:\n+ pH giảm làm hòa tan các ion kim loại nặng độc hại (Al^3+^, Mn^2+^) gây tổn thương và thui chột hệ rễ cây trồng.\n+ Làm rửa trôi các khoáng chất dinh dưỡng thiết yếu (Ca^2+^, Mg^2+^, K^+^) khiến đất bị bạc màu cằn cỗi.\n+ Ức chế hoạt động của các vi sinh vật có ích trong đất (vi khuẩn cố định đạm, phân giải lân).",
            "lines": 12
        },
        {
            "num": "Câu 5: (0,5 điểm)",
            "prompt": "Ở trạng thái bình thường, dịch vị dạ dày thường có nồng độ [H^+^] là 2.10^-4^ M. Khi tiến hành tiêu hóa, thức ăn đi vào dạ dày làm giải phóng acid HCl và dịch vị dạ dày cũng vì vậy mà có giá trị thay đổi, khi này nồng độ ion [H^+^] là 4.10^-2^ M.\nTính giá trị pH của dạ dày ở trạng thái bình thường và khi dạ dày tiêu hóa thức ăn.",
            "sol": "► *Lời giải chi tiết:*\n- Ở trạng thái bình thường:\npH~bình thường~ = -log[H^+^] = -log(2.10^-4^) = 4 - log(2) ≈ 4 - 0,301 = 3,70.\n- Khi dạ dày tiến hành tiêu hóa thức ăn:\npH~tiêu hóa~ = -log[H^+^] = -log(4.10^-2^) = 2 - log(4) ≈ 2 - 0,602 = 1,40.\n*Đáp số:* Trạng thái bình thường pH ≈ 3,70; khi tiêu hóa thức ăn pH ≈ 1,40.",
            "lines": 6,
            "image": "San_Pham/images_crop/de504_da_day.png"
        }
    ]

    for item in essay_questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*{item['num']}*", base_size=10.5, base_bold=True)

        for line in item["prompt"].split("\n"):
            p_line = doc.add_paragraph()
            p_line.paragraph_format.left_indent = Cm(0.5)
            p_line.paragraph_format.space_before = Pt(1)
            p_line.paragraph_format.space_after = Pt(1)
            render_rich_text(p_line, line, base_size=10.5)

        if "image" in item and os.path.exists(item["image"]):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(3)
            p_img.paragraph_format.space_after = Pt(3)
            p_img.add_run().add_picture(item["image"], width=Cm(6.0))

        if is_teacher:
            add_solution_no_table(doc, item["sol"])
        else:
            add_student_writing_lines(doc, num_lines=item["lines"])

    # Dòng kết thúc
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(10)
    p_end.paragraph_format.space_after = Pt(4)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_end, "---------- *HẾT* ----------", base_size=10.5, base_bold=True)

    p_rule = doc.add_paragraph()
    p_rule.paragraph_format.space_before = Pt(0)
    p_rule.paragraph_format.space_after = Pt(0)
    p_rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_rule, "_(Học sinh không sử dụng tài liệu và bảng tuần hoàn hoá học)_", base_size=9.5, base_italic=True)

    return doc

def main():
    out_dir = r"c:\Antigravity_Thanh\San_Pham"
    os.makedirs(out_dir, exist_ok=True)

    # 1. Tạo Bản Học sinh
    doc_student = create_exam_doc_504(is_teacher=False)
    path_student = os.path.join(out_dir, "De_Kiem_Tra_Hoa_11_Ma504.docx")
    doc_student.save(path_student)
    print(f"Đã tạo Bản Học sinh: {path_student}")

    # 2. Tạo Bản Giáo viên
    doc_teacher = create_exam_doc_504(is_teacher=True)
    path_teacher = os.path.join(out_dir, "De_Kiem_Tra_Hoa_11_Ma504_GiaoVien.docx")
    doc_teacher.save(path_teacher)
    print(f"Đã tạo Bản Giáo viên: {path_teacher}")

if __name__ == "__main__":
    main()
