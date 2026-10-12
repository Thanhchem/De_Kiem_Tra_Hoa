# -*- coding: utf-8 -*-
"""
Script tạo 2 file Word (.docx) chuẩn Form Đề Kiểm Tra cho:
ĐỀ SỐ 20 - CHƯƠNG 7: KIM LOẠI NHÓM IA & NHÓM IIA (HÓA HỌC 12)
Mã đề thi: 120
Bao gồm:
1. Bản Học sinh: De_Kiem_Tra_Hoa_12_De20.docx
2. Bản Giáo viên: De_Kiem_Tra_Hoa_12_De20_GiaoVien.docx (Có đáp án đỏ, lời giải chi tiết xanh, bảng đáp án)
"""

import os
import sys
sys.stdout.reconfigure(encoding='utf-8')
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"

COLOR_CORAL_RED = RGBColor(211, 47, 47)      # #D32F2F
COLOR_SAGE_GREEN = "8BC390"                 # #8BC390
COLOR_DEEP_BLUE = RGBColor(21, 101, 192)    # #1565C0
COLOR_DARK_GREEN = RGBColor(46, 125, 50)    # #2E7D32
COLOR_BLACK = RGBColor(0, 0, 0)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.59)     # 1.5 cm
        s.bottom_margin = Inches(0.59)  # 1.5 cm
        s.left_margin = Inches(0.59)    # 1.5 cm
        s.right_margin = Inches(0.59)   # 1.5 cm
        s.header_distance = Inches(0.24) # 0.6 cm
        s.footer_distance = Inches(0.24) # 0.6 cm

        # --- HEADER ---
        header = s.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.text = ""
        hp.paragraph_format.space_before = Pt(0)
        hp.paragraph_format.space_after = Pt(0)

        tbl_h = header.add_table(rows=1, cols=2, width=Inches(7.08)) # 18.0 cm
        tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
        c0, c1 = tbl_h.rows[0].cells
        c0.width = Inches(3.54)
        c1.width = Inches(3.54)
        set_cell_margins(c0, 60, 60, 0, 0)
        set_cell_margins(c1, 60, 60, 0, 0)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run("Tài liệu lưu hành nội bộ")
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(10)
        r0.font.italic = True
        r0.font.color.rgb = COLOR_CORAL_RED

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1_a = p1.add_run("Thầy ")
        r1_a.font.name = "Times New Roman"
        r1_a.font.size = Pt(10)
        r1_a.font.color.rgb = COLOR_CORAL_RED
        r1_b = p1.add_run("TRẦN VĂN THẠNH")
        r1_b.font.name = "Times New Roman"
        r1_b.font.size = Pt(10)
        r1_b.font.bold = True
        r1_b.font.color.rgb = COLOR_CORAL_RED
        r1_c = p1.add_run(" - 0777.470.803")
        r1_c.font.name = "Times New Roman"
        r1_c.font.size = Pt(10)
        r1_c.font.color.rgb = COLOR_CORAL_RED

        tblPr_h = tbl_h._tbl.tblPr
        tblBorders_h = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{COLOR_SAGE_GREEN}"/>'
            f'</w:tblBorders>'
        )
        tblPr_h.append(tblBorders_h)

        # --- FOOTER ---
        footer = s.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.text = ""
        fp.paragraph_format.space_before = Pt(0)
        fp.paragraph_format.space_after = Pt(0)

        tbl_f = footer.add_table(rows=1, cols=2, width=Inches(7.08))
        tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
        fc0, fc1 = tbl_f.rows[0].cells
        fc0.width = Inches(5.31)
        fc1.width = Inches(1.77)
        set_cell_margins(fc0, 60, 60, 0, 0)
        set_cell_margins(fc1, 60, 60, 0, 0)

        fp0 = fc0.paragraphs[0]
        fp0.paragraph_format.space_before = Pt(0)
        fp0.paragraph_format.space_after = Pt(0)
        fr0 = fp0.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
        fr0.font.name = "Times New Roman"
        fr0.font.size = Pt(9.5)
        fr0.font.italic = True
        fr0.font.color.rgb = COLOR_CORAL_RED

        fp1 = fc1.paragraphs[0]
        fp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp1.paragraph_format.space_before = Pt(0)
        fp1.paragraph_format.space_after = Pt(0)
        fr1_a = fp1.add_run("Trang ")
        fr1_a.font.name = "Times New Roman"
        fr1_a.font.size = Pt(10)
        fr1_a.font.bold = True
        fr1_a.font.color.rgb = COLOR_CORAL_RED

        fldSimple = OxmlElement('w:fldSimple')
        fldSimple.set(qn('w:instr'), 'PAGE')
        r_page = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        rFont = OxmlElement('w:rFonts')
        rFont.set(qn('w:ascii'), 'Times New Roman')
        rFont.set(qn('w:hAnsi'), 'Times New Roman')
        b = OxmlElement('w:b')
        c_el = OxmlElement('w:color')
        c_el.set(qn('w:val'), 'D32F2F')
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '20')
        rPr.append(rFont)
        rPr.append(b)
        rPr.append(c_el)
        rPr.append(sz)
        r_page.append(rPr)
        t_el = OxmlElement('w:t')
        t_el.text = "1"
        r_page.append(t_el)
        fldSimple.append(r_page)
        fp1._p.append(fldSimple)

        tblPr_f = tbl_f._tbl.tblPr
        tblBorders_f = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{COLOR_SAGE_GREEN}"/>'
            f'</w:tblBorders>'
        )
        tblPr_f.append(tblBorders_f)

def add_chem_run(paragraph, text, bold=False, italic=False, color=None, font_size=11):
    parts = re.split(r'(~.*?~|\^.*?\^)', text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run()
        run.font.name = "Times New Roman"
        run.font.size = Pt(font_size)
        run.bold = bold
        run.italic = italic
        if color:
            run.font.color.rgb = color

        if part.startswith('~') and part.endswith('~'):
            run.text = part[1:-1]
            run.font.subscript = True
        elif part.startswith('^') and part.endswith('^'):
            run.text = part[1:-1]
            run.font.superscript = True
        else:
            run.text = part

def create_title_block(doc, is_teacher=False):
    lines_info = [
        ("SỞ GIÁO DỤC VÀ ĐÀO TẠO THỪA THIÊN HUẾ", "ĐỀ KIỂM TRA ĐỊNH KỲ CHƯƠNG 7 (KIM LOẠI NHÓM IA & IIA)", True, True, COLOR_DARK_GREEN),
        ("TRƯỜNG THPT HAI BÀ TRƯNG", "MÔN: HOÁ HỌC -- LỚP 12", True, True, None),
        ("ĐỀ CHÍNH THỨC", "Thời gian làm bài: 50 phút (không kể thời gian phát đề)", True, False, None),
        ("(Đề thi có 04 trang)", "Mã đề thi: 120", False, True, None)
    ]

    for left_txt, right_txt, left_bold, right_bold, right_color in lines_info:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.08), WD_TAB_ALIGNMENT.RIGHT)

        r_l = p.add_run(left_txt)
        r_l.font.name = "Times New Roman"
        r_l.font.size = Pt(11)
        r_l.bold = left_bold
        if not left_bold:
            r_l.italic = True

        p.add_run("\t")

        r_r = p.add_run(right_txt)
        r_r.font.name = "Times New Roman"
        r_r.font.size = Pt(11)
        r_r.bold = right_bold
        if right_color:
            r_r.font.color.rgb = right_color
        if not right_bold:
            r_r.italic = True

    if is_teacher:
        p_note = doc.add_paragraph()
        p_note.paragraph_format.space_before = Pt(4)
        p_note.paragraph_format.space_after = Pt(4)
        p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_n = p_note.add_run("(BẢN GIÁO VIÊN -- CÓ LỜI GIẢI CHI TIẾT)")
        r_n.font.name = "Times New Roman"
        r_n.font.size = Pt(12)
        r_n.font.bold = True
        r_n.font.color.rgb = COLOR_CORAL_RED

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(3)
    p_info.paragraph_format.space_after = Pt(2)
    r_info = p_info.add_run("Họ, tên thí sinh: ..................................................................... Lớp: ........................ SBD: ............................")
    r_info.font.name = "Times New Roman"
    r_info.font.size = Pt(10.5)
    r_info.font.italic = True

    p_mass = doc.add_paragraph()
    p_mass.paragraph_format.space_before = Pt(0)
    p_mass.paragraph_format.space_after = Pt(5)
    r_mass = p_mass.add_run("(Cho biết nguyên tử khối: H = 1; C = 12; N = 14; O = 16; Na = 23; Mg = 24; Al = 27; P = 31; S = 32; Cl = 35,5; K = 39; Ca = 40; Fe = 56; Cu = 64; Zn = 65; Br = 80; Rb = 85,5; Ag = 108; Cs = 133; Ba = 137).")
    r_mass.font.name = "Times New Roman"
    r_mass.font.size = Pt(9.5)
    r_mass.font.italic = True

def add_part_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = COLOR_DARK_GREEN

def add_explanation(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.line_spacing = 1.15

    r_lbl = p.add_run("Lời giải chi tiết: ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(10)
    r_lbl.font.bold = True
    r_lbl.font.italic = True
    r_lbl.font.color.rgb = COLOR_DEEP_BLUE

    add_chem_run(p, text, font_size=10, italic=False, color=COLOR_DEEP_BLUE)

def render_mcq(doc, q_num, text, options, correct_key, explanation, is_teacher=False, img_path=None, img_width=Inches(3.5)):
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(3)
    p_q.paragraph_format.space_after = Pt(2)
    p_q.paragraph_format.line_spacing = 1.15

    r_lbl = p_q.add_run(f"Câu {q_num}: ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(11)
    r_lbl.bold = True

    add_chem_run(p_q, text, font_size=11)

    if img_path and os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(2)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture(img_path, width=img_width)

    max_len = max(len(opt[1]) for opt in options)
    if max_len <= 15:
        p_opts = doc.add_paragraph()
        p_opts.paragraph_format.space_before = Pt(0)
        p_opts.paragraph_format.space_after = Pt(2)
        p_opts.paragraph_format.line_spacing = 1.15
        p_opts.paragraph_format.tab_stops.add_tab_stop(Inches(1.77), WD_TAB_ALIGNMENT.LEFT)
        p_opts.paragraph_format.tab_stops.add_tab_stop(Inches(3.54), WD_TAB_ALIGNMENT.LEFT)
        p_opts.paragraph_format.tab_stops.add_tab_stop(Inches(5.31), WD_TAB_ALIGNMENT.LEFT)

        for i, (k, txt) in enumerate(options):
            is_c = (k == correct_key) and is_teacher
            r_k = p_opts.add_run(f"{k}. ")
            r_k.font.name = "Times New Roman"
            r_k.font.size = Pt(11)
            r_k.font.bold = is_c or True
            if is_c:
                r_k.font.color.rgb = COLOR_CORAL_RED
            add_chem_run(p_opts, txt, bold=is_c, color=COLOR_CORAL_RED if is_c else None, font_size=11)
            if i < 3:
                p_opts.add_run("\t")

    elif max_len <= 35:
        for row_opts in [options[:2], options[2:]]:
            p_opts = doc.add_paragraph()
            p_opts.paragraph_format.space_before = Pt(0)
            p_opts.paragraph_format.space_after = Pt(2)
            p_opts.paragraph_format.line_spacing = 1.15
            p_opts.paragraph_format.tab_stops.add_tab_stop(Inches(3.54), WD_TAB_ALIGNMENT.LEFT)

            for i, (k, txt) in enumerate(row_opts):
                is_c = (k == correct_key) and is_teacher
                r_k = p_opts.add_run(f"{k}. ")
                r_k.font.name = "Times New Roman"
                r_k.font.size = Pt(11)
                r_k.font.bold = is_c or True
                if is_c:
                    r_k.font.color.rgb = COLOR_CORAL_RED
                add_chem_run(p_opts, txt, bold=is_c, color=COLOR_CORAL_RED if is_c else None, font_size=11)
                if i == 0:
                    p_opts.add_run("\t")
    else:
        for k, txt in options:
            p_opts = doc.add_paragraph()
            p_opts.paragraph_format.space_before = Pt(0)
            p_opts.paragraph_format.space_after = Pt(2)
            p_opts.paragraph_format.line_spacing = 1.15
            is_c = (k == correct_key) and is_teacher
            r_k = p_opts.add_run(f"{k}. ")
            r_k.font.name = "Times New Roman"
            r_k.font.size = Pt(11)
            r_k.font.bold = is_c or True
            if is_c:
                r_k.font.color.rgb = COLOR_CORAL_RED
            add_chem_run(p_opts, txt, bold=is_c, color=COLOR_CORAL_RED if is_c else None, font_size=11)

    if is_teacher:
        add_explanation(doc, explanation)

def render_tf(doc, q_num, text, sub_items, explanation, is_teacher=False, img_path=None, img_width=Inches(3.5)):
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(3)
    p_q.paragraph_format.space_after = Pt(2)
    p_q.paragraph_format.line_spacing = 1.15

    r_lbl = p_q.add_run(f"Câu {q_num}: ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(11)
    r_lbl.bold = True

    add_chem_run(p_q, text, font_size=11)

    if img_path and os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(2)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture(img_path, width=img_width)

    for sub_k, sub_txt, is_true in sub_items:
        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(1.5)
        p_sub.paragraph_format.left_indent = Inches(0.2)
        p_sub.paragraph_format.line_spacing = 1.15

        r_sk = p_sub.add_run(f"{sub_k}) ")
        r_sk.font.name = "Times New Roman"
        r_sk.font.size = Pt(11)
        r_sk.bold = True

        add_chem_run(p_sub, sub_txt, font_size=11)

        if is_teacher:
            r_ans = p_sub.add_run(f" --> {'ĐÚNG' if is_true else 'SAI'}")
            r_ans.font.name = "Times New Roman"
            r_ans.font.size = Pt(10.5)
            r_ans.font.bold = True
            r_ans.font.color.rgb = COLOR_CORAL_RED

    if is_teacher:
        add_explanation(doc, explanation)

def render_short(doc, q_num, text, answer, explanation, is_teacher=False):
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(3)
    p_q.paragraph_format.space_after = Pt(2)
    p_q.paragraph_format.line_spacing = 1.15

    r_lbl = p_q.add_run(f"Câu {q_num}: ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(11)
    r_lbl.bold = True

    add_chem_run(p_q, text, font_size=11)

    if is_teacher:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(0)
        p_ans.paragraph_format.space_after = Pt(1)
        p_ans.paragraph_format.left_indent = Inches(0.25)
        r_ak = p_ans.add_run("Đáp số: ")
        r_ak.font.name = "Times New Roman"
        r_ak.font.size = Pt(10.5)
        r_ak.font.bold = True
        r_ak.font.color.rgb = COLOR_CORAL_RED
        r_av = p_ans.add_run(str(answer))
        r_av.font.name = "Times New Roman"
        r_av.font.size = Pt(10.5)
        r_av.font.bold = True
        r_av.font.color.rgb = COLOR_CORAL_RED

        add_explanation(doc, explanation)
    else:
        p_dot = doc.add_paragraph()
        p_dot.paragraph_format.space_before = Pt(0)
        p_dot.paragraph_format.space_after = Pt(2)
        r_d = p_dot.add_run("Trả lời: .....................................................................................................................................................")
        r_d.font.name = "Times New Roman"
        r_d.font.size = Pt(10)
        r_d.font.italic = True

def add_answer_table(doc):
    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_before = Pt(8)
    p_t.paragraph_format.space_after = Pt(4)
    r_t = p_t.add_run("BẢNG ĐÁP ÁN THAM KHẢO")
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(11)
    r_t.font.bold = True
    r_t.font.color.rgb = COLOR_CORAL_RED

    p_p1 = doc.add_paragraph()
    p_p1.paragraph_format.space_before = Pt(0)
    p_p1.paragraph_format.space_after = Pt(2)
    r1 = p_p1.add_run("PHẦN I: ")
    r1.font.bold = True
    r1.font.color.rgb = COLOR_DARK_GREEN
    p_p1.add_run("1.C   2.D   3.C   4.A   5.B   6.A   7.C   8.A   9.D   10.A   11.A   12.C   13.D   14.B   15.C   16.C   17.A   18.B")

    p_p2 = doc.add_paragraph()
    p_p2.paragraph_format.space_before = Pt(2)
    p_p2.paragraph_format.space_after = Pt(2)
    r2 = p_p2.add_run("PHẦN II: ")
    r2.font.bold = True
    r2.font.color.rgb = COLOR_DARK_GREEN
    p_p2.add_run("Câu 1: a-S, b-Đ, c-Đ, d-Đ   |   Câu 2: a-Đ, b-Đ, c-S, d-S\nCâu 3: a-S, b-Đ, c-S, d-Đ   |   Câu 4: a-Đ, b-Đ, c-Đ, d-S")

    p_p3 = doc.add_paragraph()
    p_p3.paragraph_format.space_before = Pt(2)
    p_p3.paragraph_format.space_after = Pt(2)
    r3 = p_p3.add_run("PHẦN III: ")
    r3.font.bold = True
    r3.font.color.rgb = COLOR_DARK_GREEN
    p_p3.add_run("Câu 1: 4   |   Câu 2: 5   |   Câu 3: 3   |   Câu 4: 4   |   Câu 5: 196   |   Câu 6: 13,48")

# DỮ LIỆU ĐỀ 20
P1_DATA = [
    (1, "Sodium chloride là gia vị quan trọng trong thức ăn của con người. Công thức của sodium chloride là",
     [("A", "NaNO~3~."), ("B", "KNO~3~."), ("C", "NaCl."), ("D", "KCl.")],
     "C", "Muối ăn có thành phần chính là sodium chloride với công thức hoá học là NaCl."),

    (2, "Tro thực vật được sử dụng như một loại phân bón cung cấp nguyên tố kali (potassium) cho cây trồng do chứa muối potassium carbonate. Công thức của potassium carbonate là",
     [("A", "KCl."), ("B", "KOH."), ("C", "NaCl."), ("D", "K~2~CO~3~.")],
     "D", "Potassium carbonate được tạo bởi cation K^+^ và anion CO~3~^2-^, có công thức hoá học là K~2~CO~3~."),

    (3, "Khi nhiệt phân hoàn toàn NaHCO~3~ thì sản phẩm của phản ứng nhiệt phân là",
     [("A", "NaOH, CO~2~, H~2~."), ("B", "Na~2~O, CO~2~, H~2~O."), ("C", "Na~2~CO~3~, CO~2~, H~2~O."), ("D", "NaOH, CO~2~, H~2~O.")],
     "C", "Khi đun nóng, muối sodium hydrogencarbonate bị nhiệt phân theo phương trình: 2NaHCO~3~ -> Na~2~CO~3~ + CO~2~ + H~2~O. Muối Na~2~CO~3~ sinh ra bền nhiệt và không bị nhiệt phân tiếp."),

    (4, "Dãy gồm các kim loại đều phản ứng với nước ở nhiệt độ thường, tạo ra dung dịch có môi trường kiềm là",
     [("A", "Na, Ba, K."), ("B", "Be, Na, Ca."), ("C", "Na, Fe, K."), ("D", "Na, Cr, K.")],
     "A", "Các kim loại kiềm (Na, K) và kim loại kiềm thổ (Ba, Ca, Sr) phản ứng mãnh liệt với nước ở nhiệt độ thường tạo dung dịch hydroxide có tính kiềm mạnh (base tan). Be trơ với nước; Fe, Cr không phản ứng với nước ở nhiệt độ thường."),

    (5, "Cho phương trình hóa học của phản ứng nhiệt phân sau: 4M(NO~3~)~x~ -(t°)-> 2M~2~O~x~ + 4xNO~2~ + xO~2~. Kim loại M là",
     [("A", "Ca."), ("B", "Mg."), ("C", "K."), ("D", "Ag.")],
     "B", "Muối nitrate của các kim loại từ Mg đến Cu khi nhiệt phân tạo thành basic oxide, NO~2~ và O~2~: 2Mg(NO~3~)~2~ -> 2MgO + 4NO~2~ + O~2~ (với x = 2, tương đương 4Mg(NO~3~)~2~ -> 2Mg~2~O~2~ + 8NO~2~ + 2O~2~). Trong khi đó muối nitrate của Ca, K nhiệt phân tạo nitrite và O~2~; muối của Ag nhiệt phân tạo kim loại Ag tự do."),

    (6, "Cho sơ đồ chuyển hoá hoá học sau:",
     [("A", "BaCO~3~."), ("B", "CaCO~3~."), ("C", "BaSO~3~."), ("D", "MgCO~3~.")],
     "A", "Nhiệt phân X tạo ra 2 oxide: BaCO~3~ -(t°)-> BaO (B~1~) + CO~2~ (A~1~).\n"
     "- Nhánh khí A~1~ (CO~2~): CO~2~ + Ca(OH)~2~ -> CaCO~3~v (A~2~) + H~2~O; CaCO~3~ -(900 °C)-> CaO + CO~2~^.\n"
     "- Nhánh oxide B~1~ (BaO): BaO + 2HCl -> BaCl~2~ (B~2~) + H~2~O; BaCl~2~ + Na~2~SO~4~ -> BaSO~4~v (B~3~) + 2NaCl.\n"
     "Chất X thoả mãn là BaCO~3~."),

    (7, "Phát biểu nào sau đây không đúng?",
     [("A", "Đun sôi nước cứng chỉ làm mềm nước cứng tạm thời."),
      ("B", "Dung dịch Na~2~CO~3~ làm mềm nước cứng tạm thời, nước cứng vĩnh cửu."),
      ("C", "Dung dịch NH~4~Cl làm mềm nước cứng."),
      ("D", "Cho dung dịch Ca(OH)~2~ vừa đủ làm mềm nước cứng tạm thời.")],
     "C", "Dung dịch NH~4~Cl không cung cấp anion có khả năng kết tủa Ca^2+^ hay Mg^2+^, do đó không thể làm mềm nước cứng."),

    (8, "Để xử lí chất thải có tính acid trong công nghiệp và đời sống, người ta thường dùng",
     [("A", "nước vôi."), ("B", "phèn chua."), ("C", "giấm ăn."), ("D", "muối ăn.")],
     "A", "Nước vôi Ca(OH)~2~ là dung dịch kiềm rẻ tiền, an toàn và sẵn có, có tác dụng trung hòa các ion H^+^ (acid) trong nước thải: H^+^ + OH^-^ -> H~2~O."),

    (9, "Hợp chất của calcium nào sau đây không gặp trong tự nhiên ở trạng thái tự do?",
     [("A", "CaCO~3~."), ("B", "CaSO~4~."), ("C", "Ca(HCO~3~)~2~."), ("D", "CaO.")],
     "D", "Vôi sống (CaO) có tính hút ẩm rất mạnh và dễ dàng tác dụng với hơi nước và khí CO~2~ trong không khí: CaO + H~2~O -> Ca(OH)~2~; Ca(OH)~2~ + CO~2~ -> CaCO~3~ + H~2~O. Do đó trong tự nhiên không tồn tại CaO tự do."),

    (10, "Trong các kim loại sau, kim loại nào có tính khử mạnh nhất?",
     [("A", "K."), ("B", "Al."), ("C", "Mg."), ("D", "Na.")],
     "A", "Trong dãy điện hoá (hoặc theo quy luật tuần hoàn), tính khử của kim loại kiềm tăng dần theo chiều bán kính tăng: Na < K. So với các kim loại nhóm IIA (Mg) và IIIA (Al) thì potassium (K) có thế điện cực chuẩn âm nhất và có tính khử mạnh nhất."),

    (11, "Các dung dịch muối ăn, phèn chua, nước vôi trong được kí hiệu ngẫu nhiên là X, Y, Z. Một số kết quả thí nghiệm được ghi ở bảng sau:\n"
         "- Mẫu X + dung dịch phenolphthalein -> Dung dịch chuyển sang màu hồng.\n"
         "- Mẫu Z + dung dịch BaCl~2~ -> Xuất hiện kết tủa trắng.\n"
         "Các dung dịch ban đầu (muối ăn, phèn chua, nước vôi trong) tương ứng lần lượt với các kí hiệu là",
     [("A", "Y, Z, X."), ("B", "Z, X, Y."), ("C", "X, Y, Z."), ("D", "Y, X, Z.")],
     "A", "- Mẫu X làm phenolphthalein chuyển sang màu hồng nên X là dung dịch bazơ: nước vôi trong (Ca(OH)~2~).\n"
     "- Mẫu Z tạo kết tủa trắng với BaCl~2~ nên Z chứa ion SO~4~^2-^: phèn chua (KAl(SO~4~)~2~.12H~2~O).\n"
     "- Còn lại Y là muối ăn (NaCl).\n"
     "Vậy theo thứ tự (muối ăn, phèn chua, nước vôi trong) tương ứng là Y, Z, X."),

    (12, "Nguyên tố kim loại nhóm IIA được tìm thấy trong tự nhiên dưới dạng nào sau đây?\n"
         "(1) Các cation M^2+^ trong nước sông, hồ, nước ngầm, nước biển.\n"
         "(2) Các khoáng vật ít tan như carbonate, sulfate, silicate.\n"
         "(3) Các hợp chất ít tan cấu tạo nên răng, xương động vật.",
     [("A", "(1) và (2)."), ("B", "(1) và (3)."), ("C", "(1), (2) và (3)."), ("D", "(2) và (3).")],
     "C", "Cả 3 nhận định đều đúng. Kim loại kiềm thổ hoạt động hoá học mạnh nên chỉ tồn tại ở dạng hợp chất: dạng ion Ca^2+^, Mg^2+^ trong nước biển, nước ngầm; dạng khoáng chất đá vôi (CaCO~3~), thạch cao (CaSO~4~), dolomite (CaCO~3~.MgCO~3~); và cấu trúc xương, men răng dưới dạng apatite/hydroxyapatite Ca~5~(PO~4~)~3~(OH)."),

    (13, "Barium phản ứng với nước dễ dàng hơn nhiều so với magnesium ở điều kiện thường là do các nguyên nhân nào sau đây?\n"
         "(1) Barium có tính khử mạnh hơn magnesium.\n"
         "(2) Độ tan của barium hydroxide trong nước cao hơn nhiều so với magnesium hydroxide.\n"
         "(3) Bọt khí hydrogen sinh ra bám trên bề mặt magnesium nhiều hơn, cản trở phản ứng tiếp diễn.",
     [("A", "(1)."), ("B", "(1), (2) và (3)."), ("C", "(1) và (3)."), ("D", "(1) và (2).")],
     "D", "- (1) Đúng vì Ba có bán kính lớn hơn, năng lượng ion hóa thấp hơn nên tính khử mạnh hơn Mg.\n"
     "- (2) Đúng vì Ba(OH)~2~ tan tốt trong nước giải phóng bề mặt Ba liên tục, trong khi Mg(OH)~2~ ít tan bám lại làm trơ bề mặt Mg.\n"
     "- (3) Sai vì yếu tố chính là lớp màng kết tủa hydroxide và tính khử nội tại, không phải do bọt khí hydrogen."),

    (14, "Nhiệt độ phân huỷ thành oxide của các muối carbonate của kim loại nhóm IIA giảm dần theo dãy nào sau đây?",
     [("A", "MgCO~3~, CaCO~3~, SrCO~3~, BaCO~3~."),
      ("B", "BaCO~3~, SrCO~3~, CaCO~3~, MgCO~3~."),
      ("C", "BaCO~3~, CaCO~3~, SrCO~3~, MgCO~3~."),
      ("D", "MgCO~3~, BaCO~3~, SrCO~3~, CaCO~3~.")],
     "B", "Độ bền nhiệt của muối carbonate nhóm IIA tăng dần từ MgCO~3~ đến BaCO~3~ (nhiệt độ phân hủy tăng dần: MgCO~3~ (~540 °C) < CaCO~3~ (~900 °C) < SrCO~3~ (~1290 °C) < BaCO~3~ (~1360 °C)). Do đó, nhiệt độ phân hủy giảm dần theo chiều ngược lại: BaCO~3~ > SrCO~3~ > CaCO~3~ > MgCO~3~."),

    (15, "Dãy các kim loại nào sau đây tác dụng nhanh với nước ở điều kiện thường?",
     [("A", "Be, Na, Ba."), ("B", "Mg, Ca, Ba."), ("C", "Li, Ca, Ba."), ("D", "Sr, Sn, Ba.")],
     "C", "Li, Ca, Ba đều tác dụng nhanh, mãnh liệt với nước ở nhiệt độ thường. Be hoàn toàn không phản ứng; Mg phản ứng chậm với nước lạnh; Sn là kim loại trơ không tan trong nước."),

    (16, "Ion Ca^2+^ (Z = 20) đóng vai trò thiết yếu trong việc phát triển xương, giúp duy trì hoạt động của cơ bắp, kích thích máu lưu thông, điều tiết một số loại hormone,... Tổng số hạt proton và electron của ion Ca^2+^ là",
     [("A", "40."), ("B", "42."), ("C", "38."), ("D", "18.")],
     "C", "Nguyên tử Ca có Z = 20 hạt proton và 20 hạt electron. Khi tạo ion Ca^2+^, nguyên tử nhường 2 electron: số electron = 20 - 2 = 18; số proton không đổi = 20. Tổng số hạt proton và electron = 20 + 18 = 38 hạt."),

    (17, "Đặc điểm nào sau đây là đặc điểm chung của nước có tính cứng vĩnh cửu và nước có tính cứng toàn phần?",
     [("A", "Đều có thể làm mềm bằng dung dịch Na~3~PO~4~."),
      ("B", "Đều không chứa anion HCO~3~^-^."),
      ("C", "Đều bị mất một phần tính cứng khi đun sôi nước."),
      ("D", "Thành phần anion hoàn toàn giống nhau.")],
     "A", "Na~3~PO~4~ (hoặc Na~2~CO~3~) có thể dùng làm mềm cho mọi loại nước cứng (tạm thời, vĩnh cửu và toàn phần) bằng cách tạo kết tủa phosphat không tan: 3Ca^2+^ + 2PO~4~^3-^ -> Ca~3~(PO~4~)~2~v; 3Mg^2+^ + 2PO~4~^3-^ -> Mg~3~(PO~4~)~2~v."),

    (18, "Cho các phát biểu sau:\n"
         "(a) Tất cả các kim loại kiềm thổ đều tan trong nước ở nhiệt độ thường.\n"
         "(b) Các kim loại kiềm có thể đẩy được kim loại yếu hơn ra khỏi dung dịch muối của chúng.\n"
         "(c) Các ion Na^+^, Mg^2+^, Al^3+^ có cùng cấu hình electron ở trạng thái cơ bản và đều có tính oxi hoá yếu.\n"
         "(d) Các kim loại kiềm K, Rb, Cs tự bốc cháy khi tiếp xúc với nước.\n"
         "(e) Kim loại magnesium có cấu trúc mạng tinh thể lục phương.\n"
         "Trong các phát biểu trên, số phát biểu đúng là",
     [("A", "2."), ("B", "3."), ("C", "4."), ("D", "5.")],
     "B", "Có 3 phát biểu đúng là (c), (d), (e):\n"
     "- (a) Sai vì Be không phản ứng với nước, Mg tan rất chậm ở nhiệt độ thường.\n"
     "- (b) Sai vì kim loại kiềm phản ứng với nước trong dung dịch trước tạo dung dịch kiềm và khí H~2~, sau đó kiềm mới phản ứng với muối.\n"
     "- (c) Đúng vì Na^+^, Mg^2+^, Al^3+^ đều có 10 electron với cấu hình 1s^2^2s^2^2p^6^ và tính oxi hoá rất yếu.\n"
     "- (d) Đúng vì phản ứng tỏa nhiệt rất lớn làm bốc cháy khí H~2~ sinh ra.\n"
     "- (e) Đúng vì Mg có mạng tinh thể lục phương chặt khít.")
]

P2_DATA = [
    (1, "Mỗi phát biểu sau đây về hợp chất của sodium là đúng hay sai?",
     [("a", "Điện phân dung dịch NaCl có màng ngăn xốp thu được NaOH, H~2~ và O~2~.", False),
      ("b", "NaHCO~3~ là hợp chất lưỡng tính.", True),
      ("c", "Na~2~CO~3~ là nguyên liệu quan trọng để sản xuất thuỷ tinh.", True),
      ("d", "Phương pháp Solvay sản xuất NaHCO~3~ từ nguyên liệu chính là NH~3~, NaCl và CO~2~.", True)],
     "San_Pham/images_crop_de20/p2_cau1_solvay.png", Inches(4.5),
     "a) Sai vì khi điện phân dung dịch NaCl có màng ngăn, ở anode thu được khí Cl~2~ (không phải O~2~), ở cathode thu được khí H~2~ và dung dịch NaOH.\n"
     "b) Đúng vì ion HCO~3~^-^ vừa có khả năng nhường proton H^+^ (tính acid khi gặp base), vừa có khả năng nhận proton H^+^ (tính base khi gặp acid).\n"
     "c) Đúng vì soda (Na~2~CO~3~) cùng với cát thạch anh (SiO~2~) và đá vôi (CaCO~3~) là thành phần chính để sản xuất thuỷ tinh silicat (Na~2~O.CaO.6SiO~2~).\n"
     "d) Đúng vì theo phương pháp Solvay, khí NH~3~ và CO~2~ được sục vào dung dịch NaCl bão hoà: NaCl + NH~3~ + CO~2~ + H~2~O -> NaHCO~3~v + NH~4~Cl."),

    (2, "Sodium chloride là một hợp chất ion điển hình.",
     [("a", "Ở trạng thái nóng chảy, sodium chloride có khả năng dẫn điện.", True),
      ("b", "Sodium chloride có nhiệt độ nóng chảy cao.", True),
      ("c", "Trong tinh thể sodium chloride ở điều kiện thường, các ion có thể di chuyển tự do.", False),
      ("d", "Khi dùng búa đập vào hạt muối thì hạt muối bị biến dạng do có tính dẻo.", False)],
     None, None,
     "a) Đúng vì khi nóng chảy, mạng tinh thể ion bị phá vỡ, các cation Na^+^ và anion Cl^-^ chuyển động tự do tạo thành hạt mang điện dẫn điện.\n"
     "b) Đúng vì liên kết ion giữa Na^+^ và Cl^-^ có lực hút tĩnh điện rất mạnh (nhiệt độ nóng chảy khoảng 801 °C).\n"
     "c) Sai vì trong tinh thể rắn, các ion dao động quanh vị trí cân bằng cố định tại các nút mạng, không thể di chuyển tự do.\n"
     "d) Sai vì tinh thể ion có tính giòn; khi bị ngoại lực tác dụng, các lớp ion cùng dấu trượt lên nhau gây lực đẩy tĩnh điện làm tinh thể vỡ vụn."),

    (3, "Mỗi phát biểu sau đây về kim loại kiềm thổ là đúng hay sai?",
     [("a", "Các oxide của kim loại kiềm thổ đều bị khí CO khử thành kim loại tự do ở nhiệt độ cao.", False),
      ("b", "Các kim loại Ca, Mg và Ba chỉ điều chế được bằng phương pháp điện phân nóng chảy.", True),
      ("c", "Các kim loại Mg và Ba đều khử được ion Ag^+^ trong dung dịch thành Ag.", False),
      ("d", "Cho kim loại Mg vào dung dịch FeCl~3~ dư thì không thu được kim loại Fe.", True)],
     None, None,
     "a) Sai vì khí CO chỉ khử được các oxide của kim loại đứng sau nhôm trong dãy hoạt động hoá học ở nhiệt độ cao; oxide kim loại kiềm thổ rất bền nhiệt không bị CO khử.\n"
     "b) Đúng vì Ca, Mg, Ba có tính khử rất mạnh, thế điện cực chuẩn rất âm nên chỉ điều chế bằng điện phân nóng chảy muối chloride khan.\n"
     "c) Sai vì Ba tác dụng với nước trước: Ba + 2H~2~O -> Ba(OH)~2~ + H~2~^, sau đó OH^-^ tạo kết tủa Ag~2~O với ion Ag^+^, không khử được Ag^+^ thành kim loại Ag.\n"
     "d) Đúng vì khi FeCl~3~ dư, phản ứng dừng lại ở nấc tạo Fe^2+^: Mg + 2FeCl~3~ -> MgCl~2~ + 2FeCl~2~; Mg hết nên không thể khử tiếp Fe^2+^ thành Fe."),

    (4, "Hợp chất của kim loại kiềm có nhiều ứng dụng quan trọng trong thực tiễn.",
     [("a", "Sodium carbonate khan (Na~2~CO~3~, còn gọi là soda) được dùng trong công nghiệp thuỷ tinh, đồ gốm, bột giặt,...", True),
      ("b", "Sodium hydrogencarbonate (NaHCO~3~) được dùng trong công nghiệp thực phẩm làm bột nở và dùng làm thuốc giảm đau dạ dày.", True),
      ("c", "Dung dịch NaCl có nồng độ 0,9% (nước muối sinh lý) được dùng để vệ sinh, súc miệng và sát khuẩn.", True),
      ("d", "Chất thường được gọi là xút ăn da là KOH.", False)],
     None, None,
     "a) Đúng vì Na~2~CO~3~ là hoá chất công nghiệp quan trọng trong luyện thuỷ tinh, xà phòng, chất tẩy rửa.\n"
     "b) Đúng vì NaHCO~3~ khi nung sinh CO~2~ làm xốp bánh; ion HCO~3~^-^ trung hòa bớt acid HCl dư trong dạ dày: HCO~3~^-^ + H^+^ -> CO~2~^ + H~2~O.\n"
     "c) Đúng vì dung dịch NaCl 0,9% đẳng trương với máu người, dùng làm nước muối sinh lý.\n"
     "d) Sai vì xút ăn da là sodium hydroxide (NaOH); KOH được gọi là potash ăn da (potút ăn da).")
]

P3_DATA = [
    (1, "Cho các phát biểu sau:\n"
        "(1) Nhúng đầu dây inox sạch vào dung dịch muối NaCl rồi đưa vào ngọn lửa đèn khí thấy ngọn lửa nhuộm màu vàng tươi.\n"
        "(2) Muối potassium carbonate tan tốt trong nước.\n"
        "(3) Trong công nghiệp, một lượng lớn soda được dùng để sản xuất thuỷ tinh.\n"
        "(4) Sodium hydrogencarbonate là chất rắn màu trắng, bền ở nhiệt độ thường.\n"
        "(5) Điện phân dung dịch NaCl có màng ngăn xốp thì ở cathode thu được kim loại Na.\n"
        "Có bao nhiêu phát biểu đúng trong các phát biểu trên?",
     "4",
     "Có 4 phát biểu đúng là (1), (2), (3), (4):\n"
     "- (1) Đúng vì ion Na^+^ nhuộm ngọn lửa thành màu vàng đặc trưng.\n"
     "- (2) Đúng vì hầu hết các muối của kim loại kiềm đều tan rất tốt trong nước.\n"
     "- (3) Đúng vì soda chiếm tỉ lệ lớn trong phối liệu sản xuất thuỷ tinh thông thường.\n"
     "- (4) Đúng vì NaHCO~3~ bền ở nhiệt độ phòng, chỉ phân hủy rõ rệt khi đun nóng trên 70 °C.\n"
     "- (5) Sai vì ở cathode xảy ra sự khử H~2~O tạo khí H~2~ và OH^-^ (2H~2~O + 2e -> H~2~^ + 2OH^-^), không thể thu được kim loại Na."),

    (2, "Baking soda bị phân huỷ khi đun nóng tạo thành soda, khí carbon dioxide và hơi nước. Tổng hệ số nguyên tối giản của các chất trong phương trình hoá học của phản ứng nhiệt phân baking soda là bao nhiêu?",
     "5",
     "Phương trình hoá học nhiệt phân baking soda:\n"
     "2NaHCO~3~ -(t°)-> Na~2~CO~3~ + CO~2~ + H~2~O\n"
     "Các hệ số tối giản lần lượt là 2, 1, 1, 1.\n"
     "Tổng hệ số = 2 + 1 + 1 + 1 = 5."),

    (3, "Cho dãy các kim loại: Fe, Na, Ba, Ca. Có bao nhiêu kim loại trong dãy trên tác dụng được với nước ở nhiệt độ thường?",
     "3",
     "Có 3 kim loại tác dụng với nước ở nhiệt độ thường là Na, Ba, Ca tạo thành dung dịch bazơ tan tương ứng và giải phóng khí H~2~. Sắt (Fe) không phản ứng với nước ở nhiệt độ thường."),

    (4, "Có bao nhiêu chất bị phân huỷ bởi nhiệt độ trong dãy các chất sau: Na~2~CO~3~, BaCO~3~, CaSO~4~, Ca(HCO~3~)~2~, Mg(NO~3~)~2~, NaHCO~3~?",
     "4",
     "Có 4 chất bị phân huỷ bởi nhiệt độ là:\n"
     "- BaCO~3~ -(t° cao)-> BaO + CO~2~^\n"
     "- Ca(HCO~3~)~2~ -(t°)-> CaCO~3~v + CO~2~^ + H~2~O\n"
     "- 2Mg(NO~3~)~2~ -(t°)-> 2MgO + 4NO~2~^ + O~2~^\n"
     "- 2NaHCO~3~ -(t°)-> Na~2~CO~3~ + CO~2~^ + H~2~O\n"
     "(Na~2~CO~3~ và CaSO~4~ rất bền nhiệt, không bị phân huỷ ở nhiệt độ phòng thí nghiệm thông thường)."),

    (5, "Ở điều kiện thường, kim loại calcium (Ca) kết tinh theo mạng tinh thể lập phương tâm diện với khối lượng riêng thực nghiệm D = 1,55 g/cm^3^. Giả thiết các nguyên tử Ca là những hình cầu chiếm 74% thể tích tinh thể, phần còn lại là các khe rỗng. Cho biết thể tích hình cầu V = (4/3)πr^3^; số Avogadro N~A~ = 6,023.10^23^ mol^-1^; số π = 3,1416; khối lượng mol của Ca là 40 g/mol. Tính bán kính nguyên tử Ca theo đơn vị picomet (pm) (làm tròn kết quả đến phần nguyên).",
     "196",
     "- Khối lượng của 1 mol Ca là 40 gam.\n"
     "- Thể tích của 1 mol tinh thể Ca:\n"
     "V(tinh thể) = m / D = 40 / 1,55 ≈ 25,80645 cm^3^.\n"
     "- Thể tích thực tế do các nguyên tử Ca chiếm giữ:\n"
     "V(nguyên tử) = 25,80645 * 0,74 ≈ 19,09677 cm^3^.\n"
     "- Thể tích của một nguyên tử Ca:\n"
     "V(1 nguyên tử) = 19,09677 / (6,023.10^23^) ≈ 3,1706.10^-23^ cm^3^.\n"
     "- Bán kính r của nguyên tử Ca:\n"
     "r^3^ = 3V / (4π) = 3 * 3,1706.10^-23^ / (4 * 3,1416) ≈ 7,5696.10^-24^ cm^3^.\n"
     "r = (7,5696.10^-24^)^(1/3) ≈ 1,9634.10^-8^ cm = 196,34 pm ≈ 196 pm."),

    (6, "Khi điện phân 500 mL dung dịch NaCl với điện cực trơ, có màng ngăn xốp đến khi khối lượng dung dịch giảm 5,475 gam thì ngừng điện phân, thu được dung dịch X (coi thể tích dung dịch thay đổi không đáng kể và nước bay hơi không đáng kể). Xác định giá trị pH của dung dịch X (làm tròn kết quả đến hai chữ số thập phân).",
     "13,48",
     "Phương trình hoá học của phản ứng điện phân dung dịch có màng ngăn:\n"
     "2NaCl + 2H~2~O -(đpdd, màng ngăn)-> 2NaOH + Cl~2~^ + H~2~^\n"
     "Gọi x là số mol NaCl đã tham gia phản ứng (mol).\n"
     "Theo phương trình: n(Cl~2~) = 0,5x mol; n(H~2~) = 0,5x mol; n(NaOH sinh ra) = x mol.\n"
     "Khối lượng dung dịch giảm chính là khối lượng hai khí thoát ra khỏi dung dịch:\n"
     "m(giảm) = m(Cl~2~) + m(H~2~) = 71 * 0,5x + 2 * 0,5x = 36,5x = 5,475 g\n"
     "=> x = 5,475 / 36,5 = 0,15 mol.\n"
     "Nồng độ mol của ion OH^-^ trong dung dịch X:\n"
     "[OH^-^] = n(NaOH) / V = 0,15 / 0,5 = 0,3 M.\n"
     "Ta có: pOH = -log[OH^-^] = -log(0,3) ≈ 0,52288\n"
     "=> pH = 14 - pOH = 14 - 0,52288 = 13,47712 ≈ 13,48.")
]

def build_exam_docx(is_teacher=False):
    doc = docx.Document()
    add_header_footer(doc)
    create_title_block(doc, is_teacher=is_teacher)

    # PHẦN I
    add_part_heading(doc, "PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn. Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án.")
    for q_num, text, options, c_key, expl in P1_DATA:
        img_p = "San_Pham/images_crop_de20/p1_cau6_so_do.png" if q_num == 6 else None
        img_w = Inches(3.8) if q_num == 6 else Inches(3.5)
        render_mcq(doc, q_num, text, options, c_key, expl, is_teacher=is_teacher, img_path=img_p, img_width=img_w)

    # PHẦN II
    add_part_heading(doc, "PHẦN II. Câu trắc nghiệm đúng sai. Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.")
    for q_num, text, sub_items, img_path, img_w, expl in P2_DATA:
        render_tf(doc, q_num, text, sub_items, expl, is_teacher=is_teacher, img_path=img_path, img_width=img_w)

    # PHẦN III
    add_part_heading(doc, "PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn. Thí sinh trả lời từ câu 1 đến câu 6.")
    for q_num, text, ans, expl in P3_DATA:
        render_short(doc, q_num, text, ans, expl, is_teacher=is_teacher)

    if is_teacher:
        add_answer_table(doc)

    out_name = "De_Kiem_Tra_Hoa_12_De20_GiaoVien.docx" if is_teacher else "De_Kiem_Tra_Hoa_12_De20.docx"
    out_path = os.path.join(OUT_DIR, out_name)
    doc.save(out_path)
    print(f"Đã lưu: {out_path}")

def main():
    print("--- Bắt đầu tạo 2 file Word cho ĐỀ SỐ 20 (Mã 120) ---")
    build_exam_docx(is_teacher=False)
    build_exam_docx(is_teacher=True)
    print("--- Hoàn tất sinh file Word cho ĐỀ SỐ 20! ---")

if __name__ == "__main__":
    main()
