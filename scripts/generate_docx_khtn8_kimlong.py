# -*- coding: utf-8 -*-
"""
Script tạo trọn bộ 3 file Word (.docx) chuẩn "FORM ĐỀ KIỂM TRA" cho:
HƯỚNG DẪN ÔN TẬP GIỮA KÌ I - MÔN: KHOA HỌC TỰ NHIÊN 8
Trường THCS Kim Long - UBND Phường Kim Long
Thầy TRẦN VĂN THẠNH - 0777.470.803

3 phiên bản:
1. Bản Học sinh: De_Cuong_KHTN_8_Giua_HK1_Kim_Long_HocSinh.docx
2. Bản Giáo viên: De_Cuong_KHTN_8_Giua_HK1_Kim_Long_GiaoVien.docx
3. Bản Dạy trên màn hình: De_Cuong_KHTN_8_Giua_HK1_Kim_Long_DayTrenManHinh.docx
"""

import os
import sys
import re
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.path.append(os.path.dirname(__file__))
sys.stdout.reconfigure(encoding='utf-8')

from khtn8_kimlong_ontap_gk1_data import (
    INFO,
    KHTN1_THEORY, KHTN1_MCQ, KHTN1_TF, KHTN1_SHORT, KHTN1_ESSAY,
    KHTN2_THEORY, KHTN2_MCQ, KHTN2_TF, KHTN2_SHORT, KHTN2_ESSAY,
    KHTN3_MCQ, KHTN3_TF, KHTN3_SHORT, KHTN3_ESSAY
)

ROOT_DIR = r"c:\Antigravity_Thanh"
OUT_DIR = os.path.join(ROOT_DIR, "San_Pham")
ONEDRIVE_DIRS = [
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn8\KHTN8-IN-DEKIEMTRADINHKI",
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn8"
]
os.makedirs(OUT_DIR, exist_ok=True)

# Bảng màu chuẩn quy định trong FORM_DE_KIEM_TRA
COLOR_HF_RED = RGBColor(211, 47, 47)        # Coral Red (#D32F2F) cho Header & Footer
COLOR_HF_GREEN_HEX = "8BC390"               # Sage Green (#8BC390) cho viền Header & Footer
COLOR_PRIMARY = RGBColor(46, 125, 50)       # Xanh lá đậm (#2E7D32) cho Tiêu đề Phân môn / Phần
COLOR_CORRECT = RGBColor(198, 40, 40)       # Đỏ (#C62828) cho Đáp án đúng Bản Giáo Viên
COLOR_SOLUTION = RGBColor(21, 101, 192)     # Xanh lam (#1565C0) cho tiêu đề Lời giải
COLOR_SCREEN_TITLE = RGBColor(106, 27, 154) # Tím đậm (#6A1B9A) cho tiêu đề bản Dạy Màn Hình
COLOR_GRAY = RGBColor(117, 117, 117)
COLOR_BORDER_GREY_HEX = "CCCCCC"

def add_p_border_bottom(p, color="8BC390", sz="6"):
    """Thêm đường kẻ viền dưới cho đoạn văn bản."""
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_p_border_top(p, color="8BC390", sz="6"):
    """Thêm đường kẻ viền trên cho đoạn văn bản."""
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    """Cài đặt lề đệm cho ô bảng."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex="F1F8E9"):
    """Đặt màu nền cho ô bảng."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="CCCCCC", sz="4"):
    """Đặt viền thanh mảnh cho bảng."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def set_run_style(run, font_name="Times New Roman", font_size=10.5, bold=False, italic=False, subscript=False, superscript=False, color=None):
    """Cài đặt style chi tiết cho từng run text."""
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    """Chèn trường số trang tự động PAGE vào chân trang."""
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def render_rich_text(p, text, base_font="Times New Roman", base_size=10.5, base_bold=False, base_italic=False, base_color=None):
    """
    Sinh văn bản định dạng phong phú từ kí hiệu:
    - *bold*
    - _italic_
    - ~subscript~ (CO~2~, H~2~O, P~2~O~5~)
    - ^superscript^ (10^23^, m^2^, m^3^)
    """
    tokens = re.split(r'(~[^~]+~|\^[^\^]+\^|\*[^\*]+\*|_[^_]+_)', text)
    for token in tokens:
        if not token:
            continue

        is_bold = base_bold
        is_italic = base_italic
        is_sub = False
        is_sup = False
        content = token

        if token.startswith('*') and token.endswith('*'):
            is_bold = True
            content = token[1:-1]
        elif token.startswith('_') and token.endswith('_'):
            is_italic = True
            content = token[1:-1]
        elif token.startswith('~') and token.endswith('~'):
            is_sub = True
            content = token[1:-1]
        elif token.startswith('^') and token.endswith('^'):
            is_sup = True
            content = token[1:-1]

        run = p.add_run(content)
        set_run_style(run, font_name=base_font, font_size=base_size, bold=is_bold, italic=is_italic,
                      subscript=is_sub, superscript=is_sup, color=base_color)

def setup_header_footer(doc):
    """Cài đặt Header & Footer chuẩn theo FORM_DE_KIEM_TRA."""
    normal_style = doc.styles['Normal']
    normal_style.paragraph_format.tab_stops.clear_all()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    section.start_type = WD_SECTION_START.NEW_PAGE
    section.different_first_page_header_footer = False
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    # 1. Header (Bảng 1x2, rộng 18.0 cm)
    header = section.header
    for p in header.paragraphs:
        p.text = ""
    tbl_h = header.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl_h.cell(0, 0)
    c_right = tbl_h.cell(0, 1)
    c_left.width = Cm(9.0)
    c_right.width = Cm(9.0)

    p_hl = c_left.paragraphs[0]
    p_hl.paragraph_format.space_before = Pt(0)
    p_hl.paragraph_format.space_after = Pt(2)
    p_hl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_hl = p_hl.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)

    p_hr = c_right.paragraphs[0]
    p_hr.paragraph_format.space_before = Pt(0)
    p_hr.paragraph_format.space_after = Pt(2)
    p_hr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p_hr.add_run("Thầy ")
    set_run_style(r1, font_size=10, color=COLOR_HF_RED)
    r2 = p_hr.add_run("TRẦN VĂN THẠNH")
    set_run_style(r2, font_size=10, bold=True, color=COLOR_HF_RED)
    r3 = p_hr.add_run(" - 0777.470.803")
    set_run_style(r3, font_size=10, color=COLOR_HF_RED)

    add_p_border_bottom(p_hl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_bottom(p_hr, color=COLOR_HF_GREEN_HEX, sz="6")

    # 2. Footer (Bảng 1x2, rộng 18.0 cm)
    footer = section.footer
    for p in footer.paragraphs:
        p.text = ""
    tbl_f = footer.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_fl = tbl_f.cell(0, 0)
    c_fr = tbl_f.cell(0, 1)
    c_fl.width = Cm(14.5)
    c_fr.width = Cm(3.5)

    p_fl = c_fl.paragraphs[0]
    p_fl.paragraph_format.space_before = Pt(2)
    p_fl.paragraph_format.space_after = Pt(0)
    p_fl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_fl = p_fl.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)

    p_fr = c_fr.paragraphs[0]
    p_fr.paragraph_format.space_before = Pt(2)
    p_fr.paragraph_format.space_after = Pt(0)
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_field(p_fr, font_size=10, bold=True, color=COLOR_HF_RED)

    add_p_border_top(p_fl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_top(p_fr, color=COLOR_HF_GREEN_HEX, sz="6")

def build_title_block(doc, is_teacher=False, is_screen=False, base_size=10.5):
    """Khối tiêu đề theo FORM_DE_KIEM_TRA."""
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl.cell(0, 0)
    c_right = tbl.cell(0, 1)
    c_left.width = Cm(8.0)
    c_right.width = Cm(10.0)

    # Cột trái
    p_l1 = c_left.paragraphs[0]
    p_l1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l1.paragraph_format.space_before = Pt(0)
    p_l1.paragraph_format.space_after = Pt(1)
    render_rich_text(p_l1, f"*{INFO['ubnd']}*", base_size=10.5, base_bold=True)

    p_l2 = c_left.add_paragraph()
    p_l2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l2.paragraph_format.space_before = Pt(0)
    p_l2.paragraph_format.space_after = Pt(1)
    render_rich_text(p_l2, f"*{INFO['truong']}*", base_size=11, base_bold=True)

    p_l3 = c_left.add_paragraph()
    p_l3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l3.paragraph_format.space_before = Pt(0)
    p_l3.paragraph_format.space_after = Pt(2)
    render_rich_text(p_l3, "---***---", base_size=9, base_color=COLOR_GRAY)

    # Cột phải
    p_r1 = c_right.paragraphs[0]
    p_r1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r1.paragraph_format.space_before = Pt(0)
    p_r1.paragraph_format.space_after = Pt(1)
    render_rich_text(p_r1, f"*{INFO['de_title']}*", base_size=12, base_bold=True, base_color=COLOR_HF_RED)

    p_r2 = c_right.add_paragraph()
    p_r2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r2.paragraph_format.space_before = Pt(0)
    p_r2.paragraph_format.space_after = Pt(1)
    render_rich_text(p_r2, f"*{INFO['mon']}*", base_size=11.5, base_bold=True, base_color=RGBColor(21, 101, 192))

    p_r3 = c_right.add_paragraph()
    p_r3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r3.paragraph_format.space_before = Pt(0)
    p_r3.paragraph_format.space_after = Pt(2)
    render_rich_text(p_r3, f"_{INFO['nam_hoc']}_", base_size=10, base_italic=True)

    if is_teacher:
        p_tc = doc.add_paragraph()
        p_tc.paragraph_format.space_before = Pt(2)
        p_tc.paragraph_format.space_after = Pt(2)
        p_tc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rtc = p_tc.add_run("(BẢN GIÁO VIÊN -- CÓ ĐÁP ÁN VÀ LỜI GIẢI CHI TIẾT)")
        set_run_style(rtc, font_size=11, bold=True, color=COLOR_CORRECT)
    elif is_screen:
        p_sc = doc.add_paragraph()
        p_sc.paragraph_format.space_before = Pt(2)
        p_sc.paragraph_format.space_after = Pt(2)
        p_sc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rsc = p_sc.add_run("(BẢN TRÌNH CHIẾU / GIẢNG DẠY MÀN HÌNH -- CÓ KHUNG CHỪA CHỖ TRỐNG TƯƠNG TÁC)")
        set_run_style(rsc, font_size=11, bold=True, color=COLOR_SCREEN_TITLE)

    # Dòng thông tin học sinh
    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(4)
    p_info.paragraph_format.space_after = Pt(3)
    render_rich_text(p_info, "*Họ, tên học sinh:* .............................................................. *Lớp:* ............. *Phòng thi:* .............", base_size=base_size)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(4)
    add_p_border_bottom(p_div, color=COLOR_BORDER_GREY_HEX, sz="4")

def add_part_header(doc, title, base_size=12):
    """Tiêu đề lớn cho Phân môn."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    render_rich_text(p, f"*{title}*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)

def add_section_header(doc, title, base_size=11):
    """Tiêu đề mục con."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, f"*{title}*", base_size=base_size, base_bold=True, base_color=RGBColor(21, 101, 192))

def add_callout_box(doc, lines, base_size=10):
    """Khung tóm tắt lý thuyết trọng tâm."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    c.width = Cm(18.0)
    set_cell_margins(c, top=60, bottom=60, left=100, right=100)
    set_cell_shading(c, "F5F9F6")
    
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    render_rich_text(p, lines[0], base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)
    
    for line in lines[1:]:
        p_line = c.add_paragraph()
        p_line.paragraph_format.space_before = Pt(1)
        p_line.paragraph_format.space_after = Pt(1)
        render_rich_text(p_line, line, base_size=base_size)

def add_solution_block(doc, solution_text, base_size=10):
    """Trình bày Lời giải chi tiết theo chuẩn không dùng Table (thụt lề 0.6 cm)."""
    p_sol = doc.add_paragraph()
    p_sol.paragraph_format.left_indent = Cm(0.6)
    p_sol.paragraph_format.space_before = Pt(2)
    p_sol.paragraph_format.space_after = Pt(3)
    p_sol.paragraph_format.line_spacing = 1.15

    r_lbl = p_sol.add_run("Lời giải chi tiết:\n")
    set_run_style(r_lbl, font_size=base_size, bold=True, color=COLOR_SOLUTION)

    lines = solution_text.split('\n')
    for idx, line in enumerate(lines):
        line_str = line.strip()
        if not line_str:
            continue
        if idx > 0:
            p_sol = doc.add_paragraph()
            p_sol.paragraph_format.left_indent = Cm(0.6)
            p_sol.paragraph_format.space_before = Pt(1)
            p_sol.paragraph_format.space_after = Pt(1)
            p_sol.paragraph_format.line_spacing = 1.15
        render_rich_text(p_sol, line_str, base_size=base_size)

def add_screen_teaching_space(doc, num_lines=3):
    """Chừa khoảng trống rộng rãi cho bản Dạy Trên Màn Hình."""
    p_box = doc.add_paragraph()
    p_box.paragraph_format.left_indent = Cm(0.5)
    p_box.paragraph_format.space_before = Pt(3)
    p_box.paragraph_format.space_after = Pt(5)
    
    # Viền nét đứt màu xám nhạt, viền trái nét liền tím
    pPr = p_box._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="dashed" w:sz="4" w:space="4" w:color="B0BEC5"/><w:bottom w:val="dashed" w:sz="4" w:space="4" w:color="B0BEC5"/><w:left w:val="single" w:sz="12" w:space="6" w:color="7E57C2"/></w:pBdr>'
    pPr.append(parse_xml(xml))
    
    r_hint = p_box.add_run("✎ Ghi chú & Lời giải trực tiếp trên màn hình:\n")
    set_run_style(r_hint, font_size=9.5, bold=True, color=COLOR_SCREEN_TITLE)
    for _ in range(num_lines):
        p_box.add_run("\n")

def add_writing_lines(doc, num_lines=3):
    """Thêm dòng chấm chấm cho học sinh làm bài tự luận."""
    for _ in range(num_lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("................................................................................................................................................................................................................")
        set_run_style(r, font_size=9, color=COLOR_GRAY)

def render_mcq_options(doc, opts, correct_letter, is_teacher, base_size=10.5):
    """Bố trí 4 phương án A, B, C, D bằng Tab Stops - Tuyệt đối không dùng Table."""
    lens = [len(opt) for opt in opts]
    max_len = max(lens) if lens else 0

    if max_len <= 18:
        # 4 phương án trên 1 dòng
        p_opt = doc.add_paragraph()
        p_opt.paragraph_format.space_before = Pt(1)
        p_opt.paragraph_format.space_after = Pt(2)
        p_opt.paragraph_format.tab_stops.clear_all()
        p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(5.0), WD_TAB_ALIGNMENT.LEFT)
        p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(9.5), WD_TAB_ALIGNMENT.LEFT)
        p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(14.0), WD_TAB_ALIGNMENT.LEFT)

        for opt in opts:
            p_opt.add_run("\t")
            letter = opt.strip()[0]
            if is_teacher and letter == correct_letter:
                render_rich_text(p_opt, f"*{opt}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p_opt, opt, base_size=base_size)

    elif max_len <= 45:
        # 2 dòng x 2 phương án
        p1 = doc.add_paragraph()
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.tab_stops.clear_all()
        p1.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p1.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)

        for opt in opts[:2]:
            p1.add_run("\t")
            letter = opt.strip()[0]
            if is_teacher and letter == correct_letter:
                render_rich_text(p1, f"*{opt}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p1, opt, base_size=base_size)

        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(1)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.tab_stops.clear_all()
        p2.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p2.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)

        for opt in opts[2:]:
            p2.add_run("\t")
            letter = opt.strip()[0]
            if is_teacher and letter == correct_letter:
                render_rich_text(p2, f"*{opt}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p2, opt, base_size=base_size)

    else:
        # 4 dòng riêng biệt
        for opt in opts:
            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.left_indent = Cm(0.5)
            p_opt.paragraph_format.space_before = Pt(1)
            p_opt.paragraph_format.space_after = Pt(1)
            letter = opt.strip()[0]
            if is_teacher and letter == correct_letter:
                render_rich_text(p_opt, f"*{opt}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p_opt, opt, base_size=base_size)

def render_mcq_list(doc, questions, is_teacher, is_screen, base_size=10.5):
    """Sinh danh sách câu hỏi trắc nghiệm nhiều lựa chọn."""
    for q in questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(1)
        render_rich_text(p_q, f"*Câu {q['num']}.* {q['question']}", base_size=base_size)

        render_mcq_options(doc, q["options"], q.get("correct", ""), is_teacher, base_size=base_size)

        if is_teacher and "solution" in q and q["solution"]:
            add_solution_block(doc, q["solution"], base_size=10)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=2)

def render_tf_list(doc, questions, is_teacher, is_screen, base_size=10.5):
    """Sinh danh sách câu hỏi Đúng/Sai (dạng chùm ý a,b,c,d)."""
    for q in questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*Câu {q['num']}.* {q['intro']}", base_size=base_size)

        sol_lines = []
        for item in q["items"]:
            p_item = doc.add_paragraph()
            p_item.paragraph_format.left_indent = Cm(0.5)
            p_item.paragraph_format.space_before = Pt(1)
            p_item.paragraph_format.space_after = Pt(1)

            lbl = f"{item['label']})"
            txt = f"{lbl} {item['text']}"
            if is_teacher:
                ans_str = f" ({item['ans'].upper()})"
                render_rich_text(p_item, txt, base_size=base_size)
                r_ans = p_item.add_run(ans_str)
                set_run_style(r_ans, font_size=base_size, bold=True, color=COLOR_CORRECT if item['ans'] == "Đúng" else RGBColor(183, 28, 28))
                sol_lines.append(f"Ý {item['label']}: {item['ans']}. {item.get('explain', '')}")
            else:
                render_rich_text(p_item, f"{txt} [   ]", base_size=base_size)

        if is_teacher and sol_lines:
            add_solution_block(doc, "\n".join(sol_lines), base_size=10)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=3)

def render_simple_tf_list(doc, questions, is_teacher, is_screen, base_size=10.5):
    """Sinh danh sách câu hỏi Đúng/Sai đơn lẻ (như trong KHTN 3)."""
    for q in questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        
        if is_teacher:
            render_rich_text(p_q, f"*Câu {q['num']}.* {q['text']} ", base_size=base_size)
            r_ans = p_q.add_run(f"({q['ans']})")
            set_run_style(r_ans, font_size=base_size, bold=True, color=COLOR_CORRECT if q['ans'] == "Đúng" else RGBColor(183, 28, 28))
            add_solution_block(doc, q.get("explain", ""), base_size=10)
        elif is_screen:
            render_rich_text(p_q, f"*Câu {q['num']}.* {q['text']} (........)", base_size=base_size)
            add_screen_teaching_space(doc, num_lines=2)
        else:
            render_rich_text(p_q, f"*Câu {q['num']}.* {q['text']} (........)", base_size=base_size)

def render_short_list(doc, questions, is_teacher, is_screen, base_size=10.5):
    """Sinh danh sách câu hỏi trả lời ngắn."""
    for q in questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*Câu {q['num']}.* {q['question']}", base_size=base_size)

        if is_teacher:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Cm(0.5)
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(2)
            r_ans_lbl = p_ans.add_run("Đáp án: ")
            set_run_style(r_ans_lbl, font_size=base_size, bold=True, color=COLOR_CORRECT)
            render_rich_text(p_ans, f"*{q['ans']}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)

            if "solution" in q and q["solution"]:
                add_solution_block(doc, q["solution"], base_size=10)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=2)
        else:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Cm(0.5)
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(2)
            render_rich_text(p_ans, "*Trả lời:* .................................................................................................................................................", base_size=base_size)

def render_essay_list(doc, questions, is_teacher, is_screen, base_size=10.5):
    """Sinh danh sách câu hỏi tự luận."""
    for q in questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*Câu {q['num']}.* {q['question']}", base_size=base_size)

        if is_teacher:
            add_solution_block(doc, q["solution"], base_size=10)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=5)
        else:
            add_writing_lines(doc, num_lines=4)

def render_summary_table(doc):
    """Bảng tổng hợp đáp án trắc nghiệm nhanh ở cuối Bản Giáo Viên."""
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(12)
    p_sec.paragraph_format.space_after = Pt(4)
    render_rich_text(p_sec, "*BẢNG TỔNG HỢP ĐÁP ÁN TRẮC NGHIỆM NHANH*", base_size=12, base_bold=True, base_color=COLOR_HF_RED)

    # 1. KHTN 1 MCQ (15 câu)
    p1 = doc.add_paragraph()
    p1.paragraph_format.space_before = Pt(4)
    p1.paragraph_format.space_after = Pt(2)
    render_rich_text(p1, "*1. Phân môn KHTN 1 (Hóa học) - Trắc nghiệm nhiều lựa chọn*", base_size=10.5, base_bold=True, base_color=COLOR_PRIMARY)
    
    tbl1 = doc.add_table(rows=2, cols=15)
    tbl1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl1, color="B0BEC5", sz="4")
    for i in range(15):
        tbl1.cell(0, i).width = Cm(1.2)
        tbl1.cell(1, i).width = Cm(1.2)
        set_cell_shading(tbl1.cell(0, i), "E8F5E9")
        p_h = tbl1.cell(0, i).paragraphs[0]
        p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        render_rich_text(p_h, f"*{i+1}*", base_size=9.5, base_bold=True)

        p_a = tbl1.cell(1, i).paragraphs[0]
        p_a.alignment = WD_ALIGN_PARAGRAPH.CENTER
        render_rich_text(p_a, f"*{KHTN1_MCQ[i]['correct']}*", base_size=10, base_bold=True, base_color=COLOR_CORRECT)

    # 2. KHTN 2 MCQ (15 câu)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(6)
    p2.paragraph_format.space_after = Pt(2)
    render_rich_text(p2, "*2. Phân môn KHTN 2 (Vật lí) - Trắc nghiệm nhiều lựa chọn*", base_size=10.5, base_bold=True, base_color=COLOR_PRIMARY)
    
    tbl2 = doc.add_table(rows=2, cols=15)
    tbl2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl2, color="B0BEC5", sz="4")
    for i in range(15):
        tbl2.cell(0, i).width = Cm(1.2)
        tbl2.cell(1, i).width = Cm(1.2)
        set_cell_shading(tbl2.cell(0, i), "E8F5E9")
        p_h = tbl2.cell(0, i).paragraphs[0]
        p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        render_rich_text(p_h, f"*{i+1}*", base_size=9.5, base_bold=True)

        p_a = tbl2.cell(1, i).paragraphs[0]
        p_a.alignment = WD_ALIGN_PARAGRAPH.CENTER
        render_rich_text(p_a, f"*{KHTN2_MCQ[i]['correct']}*", base_size=10, base_bold=True, base_color=COLOR_CORRECT)

    # 3. KHTN 3 MCQ (10 câu)
    p3 = doc.add_paragraph()
    p3.paragraph_format.space_before = Pt(6)
    p3.paragraph_format.space_after = Pt(2)
    render_rich_text(p3, "*3. Phân môn KHTN 3 (Sinh học) - Trắc nghiệm nhiều lựa chọn*", base_size=10.5, base_bold=True, base_color=COLOR_PRIMARY)
    
    tbl3 = doc.add_table(rows=2, cols=10)
    tbl3.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl3, color="B0BEC5", sz="4")
    for i in range(10):
        tbl3.cell(0, i).width = Cm(1.8)
        tbl3.cell(1, i).width = Cm(1.8)
        set_cell_shading(tbl3.cell(0, i), "E8F5E9")
        p_h = tbl3.cell(0, i).paragraphs[0]
        p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        render_rich_text(p_h, f"*{i+1}*", base_size=10, base_bold=True)

        p_a = tbl3.cell(1, i).paragraphs[0]
        p_a.alignment = WD_ALIGN_PARAGRAPH.CENTER
        render_rich_text(p_a, f"*{KHTN3_MCQ[i]['correct']}*", base_size=10, base_bold=True, base_color=COLOR_CORRECT)

def build_full_document(mode="hocsinh"):
    """
    Xây dựng tài liệu hoàn chỉnh cho 1 trong 3 chế độ:
    - 'hocsinh': Bản Học sinh
    - 'giaovien': Bản Giáo viên (có lời giải chi tiết và đáp án)
    - 'daytrenmanhinh': Bản Dạy Trên Màn Hình (có khung tương tác rộng rãi)
    """
    is_teacher = (mode == "giaovien")
    is_screen = (mode == "daytrenmanhinh")

    doc = docx.Document()
    setup_header_footer(doc)
    build_title_block(doc, is_teacher=is_teacher, is_screen=is_screen, base_size=10.5)

    # =========================================================================
    # 1. PHÂN MÔN KHTN 1 (HÓA HỌC)
    # =========================================================================
    add_part_header(doc, "PHÂN MÔN KHTN 1 (HÓA HỌC)")
    add_callout_box(doc, KHTN1_THEORY, base_size=10)

    add_section_header(doc, "B/ BÀI TẬP: Một số dạng câu hỏi ôn tập tham khảo")
    add_section_header(doc, "I/ Trắc nghiệm nhiều lựa chọn: Chọn câu trả lời đúng trong các câu sau:", base_size=10.5)
    render_mcq_list(doc, KHTN1_MCQ, is_teacher, is_screen, base_size=10.5)

    add_section_header(doc, "II/ Trắc nghiệm đúng sai: Đánh Đúng/Sai vào lựa chọn của em:", base_size=10.5)
    render_tf_list(doc, KHTN1_TF, is_teacher, is_screen, base_size=10.5)

    add_section_header(doc, "III/ Trắc nghiệm trả lời ngắn:", base_size=10.5)
    render_short_list(doc, KHTN1_SHORT, is_teacher, is_screen, base_size=10.5)

    add_section_header(doc, "IV/ Phần tự luận:", base_size=10.5)
    render_essay_list(doc, KHTN1_ESSAY, is_teacher, is_screen, base_size=10.5)

    # =========================================================================
    # 2. PHÂN MÔN KHTN 2 (VẬT LÍ)
    # =========================================================================
    doc.add_page_break()
    add_part_header(doc, "PHÂN MÔN KHTN 2 (VẬT LÍ)")
    add_callout_box(doc, KHTN2_THEORY, base_size=10)

    add_section_header(doc, "B. BÀI TẬP: MỘT SỐ DẠNG CÂU HỎI ÔN TẬP THAM KHẢO")
    add_section_header(doc, "I. TRẮC NGHIỆM NHIỀU LỰA CHỌN (Chọn đáp án đúng nhất)", base_size=10.5)
    render_mcq_list(doc, KHTN2_MCQ, is_teacher, is_screen, base_size=10.5)

    add_section_header(doc, "II. TRẮC NGHIỆM ĐÚNG - SAI (Đánh dấu Đúng hoặc Sai cho mỗi ý)", base_size=10.5)
    render_tf_list(doc, KHTN2_TF, is_teacher, is_screen, base_size=10.5)

    add_section_header(doc, "III. TRẮC NGHIỆM TRẢ LỜI NGẮN", base_size=10.5)
    render_short_list(doc, KHTN2_SHORT, is_teacher, is_screen, base_size=10.5)

    add_section_header(doc, "IV. PHẦN TỰ LUẬN", base_size=10.5)
    render_essay_list(doc, KHTN2_ESSAY, is_teacher, is_screen, base_size=10.5)

    # =========================================================================
    # 3. PHÂN MÔN KHTN 3 (SINH HỌC)
    # =========================================================================
    doc.add_page_break()
    add_part_header(doc, "PHÂN MÔN KHTN 3 (SINH HỌC)")

    add_section_header(doc, "Phần I. Câu trắc nghiệm nhiều phương án lựa chọn", base_size=10.5)
    render_mcq_list(doc, KHTN3_MCQ, is_teacher, is_screen, base_size=10.5)

    add_section_header(doc, "Phần II. Câu trắc nghiệm đúng sai", base_size=10.5)
    render_simple_tf_list(doc, KHTN3_TF, is_teacher, is_screen, base_size=10.5)

    add_section_header(doc, "Phần III. Câu trắc nghiệm trả lời ngắn", base_size=10.5)
    render_short_list(doc, KHTN3_SHORT, is_teacher, is_screen, base_size=10.5)

    add_section_header(doc, "Phần IV. Tự luận", base_size=10.5)
    render_essay_list(doc, KHTN3_ESSAY, is_teacher, is_screen, base_size=10.5)

    # Bảng tổng hợp đáp án cho Bản Giáo Viên
    if is_teacher:
        render_summary_table(doc)

    # Lời chúc cuối tài liệu
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(16)
    p_end.paragraph_format.space_after = Pt(8)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_end, "*★ ★ ★ Chúc các em ôn tập và kiểm tra đạt kết quả cao! ★ ★ ★*", base_size=11, base_bold=True, base_italic=True, base_color=COLOR_PRIMARY)

    return doc

def main():
    print("=" * 60)
    print("BẮT ĐẦU TẠO 3 FILE WORD KHTN 8 - THCS KIM LONG")
    print("=" * 60)

    configs = [
        ("hocsinh", "De_Cuong_KHTN_8_Giua_HK1_Kim_Long_HocSinh.docx", "Bản Học sinh"),
        ("giaovien", "De_Cuong_KHTN_8_Giua_HK1_Kim_Long_GiaoVien.docx", "Bản Giáo viên"),
        ("daytrenmanhinh", "De_Cuong_KHTN_8_Giua_HK1_Kim_Long_DayTrenManHinh.docx", "Bản Dạy Trên Màn Hình")
    ]

    generated_files = []

    for mode, filename, desc in configs:
        out_path = os.path.join(OUT_DIR, filename)
        print(f"--> Đang khởi tạo: {desc} ({filename})...")
        doc = build_full_document(mode=mode)
        doc.save(out_path)
        print(f"    [Xong] Kích thước file: {os.path.getsize(out_path):,} bytes")
        generated_files.append(out_path)

        # Sao chép vào OneDrive nếu có thư mục
        for od in ONEDRIVE_DIRS:
            if os.path.exists(od):
                try:
                    dest = os.path.join(od, filename)
                    shutil.copy2(out_path, dest)
                    print(f"    [OneDrive] Đã sao lưu sang: {dest}")
                except Exception as e:
                    print(f"    [Cảnh báo OneDrive] {e}")

    print("\n" + "=" * 60)
    print("HOÀN TẤT TẠO 3 FILE WORD THÀNH CÔNG 100%!")
    print("=" * 60)

if __name__ == "__main__":
    main()
