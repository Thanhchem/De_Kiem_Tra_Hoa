# -*- coding: utf-8 -*-
"""
Kiểm tra trùng lặp và bổ sung câu hỏi mới từ tài liệu:
"TRƯỜNG THPT CHÂU VĂN LIÊM - TÀI LIỆU MÔN HOÁ HỌC 12 - CHƯƠNG 2. CARBOHYDRATE (ThS. Cao Quốc Trung)"
vào Ngân hàng câu hỏi Hóa 12 Chương 2 (Carbohydrate):
1. Cập nhật JSON: San_Pham/Ngan_Hang_Hoa_12/Ngan_Hang_Hoa_12_Tron_Bo_9_Chuong.json
2. Cập nhật 2 file Word Ngân hàng Chương 2 (_HocSinh.docx & _GiaoVien.docx) kèm Bảng tổng hợp Ma trận mới.
"""

import os
import io
import re
import json
import sys
from difflib import SequenceMatcher
from PIL import Image
import docx
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from generate_5_de_gk1_hoa12 import (
    add_chem_runs, parse_abcd_from_opt_lines, add_options_paragraph_no_table,
    add_data_table_to_doc, add_solution_block, COLOR_PRIMARY, COLOR_CORRECT,
    COLOR_SOLUTION, COLOR_TAG, COLOR_HF_RED
)
from build_ngan_hang_hoa12_c123 import (
    CHAPTERS, LEVEL_NAMES, setup_page_and_hf,
    inject_question_tags_into_doc, insert_matrix_summary_at_top,
    classify_p2_item_level, iter_block_items
)
from merge_cvl_c1_to_bank import (
    get_full_para_text, extract_rids_from_el, norm_for_compare,
    get_cau_num_ids, classify_level
)

NEW_DOCX = r"C:\Users\Admin\.gemini\antigravity\brain\f4f5522d-5a23-4349-b3d7-e0da630bfcb9\.user_uploaded\media_1792856755605_36107230.docx"
BANK_JSON = r"c:\Antigravity_Thanh\San_Pham\Ngan_Hang_Hoa_12\Ngan_Hang_Hoa_12_Tron_Bo_9_Chuong.json"
OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
BANK_DIR = r"c:\Antigravity_Thanh\San_Pham\Ngan_Hang_Hoa_12"
IMG_DIR = r"c:\Antigravity_Thanh\San_Pham\images_gk1_hoa12"
os.makedirs(IMG_DIR, exist_ok=True)


def parse_cvl_c2_questions(doc):
    cau_nids = get_cau_num_ids(doc)
    blocks = list(iter_block_items(doc))

    start_idx = 0
    end_idx = len(blocks)
    for i, (btype, blk) in enumerate(blocks):
        if btype == 'p':
            t = get_full_para_text(blk)
            if t == 'BÀI TẬP TỰ LUYỆN' and start_idx == 0:
                start_idx = i
            if t == 'MỤC LỤC' and i > start_idx:
                end_idx = i
                break

    cur_topic = "Glucose – Fructose"
    cur_part = "P1"
    questions = []
    cur_q = None

    def flush_q():
        nonlocal cur_q
        if cur_q:
            questions.append(cur_q)
            cur_q = None

    for i in range(start_idx + 1, end_idx):
        btype, blk = blocks[i]
        if btype == 'p':
            t = get_full_para_text(blk)
            imgs = extract_rids_from_el(blk._p, doc.part)
            if not t and not imgs:
                continue

            t_up = t.upper()
            if "PHẦN 1:" in t_up and "GLUCOSE" in t_up:
                flush_q()
                cur_topic = "Glucose – Fructose"
                continue
            elif "PHẦN 2:" in t_up and "SACCHAROSE" in t_up:
                flush_q()
                cur_topic = "Saccharose – Maltose"
                continue
            elif "PHẦN 3:" in t_up and "TINH BỘT" in t_up:
                flush_q()
                cur_topic = "Tinh bột – Cellulose"
                continue

            if "TRẮC NGHIỆM NHIỀU PHƯƠNG ÁN" in t_up:
                flush_q()
                cur_part = "P1"
                continue
            elif "TRẮC NGHIỆM ĐÚNG SAI" in t_up:
                flush_q()
                cur_part = "P2"
                continue
            elif "TRẮC NGHIỆM YÊU CẦU TRẢ LỜI NGẮN" in t_up:
                flush_q()
                cur_part = "P3"
                continue

            nids = blk._p.xpath('.//w:numPr/w:numId/@w:val')
            is_auto_cau = bool(nids and nids[0] in cau_nids)
            m_text_cau = re.match(r'^Câu\s+\d+[\.\:]\s*(.*)', t)

            if is_auto_cau or m_text_cau:
                flush_q()
                stem_first = m_text_cau.group(1).strip() if m_text_cau else t
                cur_q = {
                    "topic": cur_topic,
                    "part": cur_part,
                    "stem_lines": [stem_first] if stem_first else [],
                    "opt_lines": [],
                    "p2_items": {},
                    "sol_lines": [],
                    "underlined_opt": None,
                    "in_sol": False,
                    "img_blobs": list(imgs),
                    "tables": [],
                }
            elif cur_q is not None:
                if imgs and not cur_q["in_sol"]:
                    cur_q["img_blobs"].extend(imgs)
                if not t:
                    continue
                if t.startswith(("📝", "Hướng dẫn giải", "Đáp án:", "Đáp án ")):
                    cur_q["in_sol"] = True
                    cur_q["sol_lines"].append(t)
                elif cur_q["in_sol"]:
                    cur_q["sol_lines"].append(t)
                else:
                    if cur_part == "P1" and re.match(r'^[A-D]\.', t):
                        cur_q["opt_lines"].append(t)
                        for r in blk.runs:
                            if r.underline:
                                m = re.search(r'([A-D])', r.text)
                                if m:
                                    cur_q["underlined_opt"] = m.group(1)
                    elif cur_part == "P2" and re.match(r'^[a-d]\)', t):
                        k = t[0].lower()
                        cur_q["p2_items"][k] = t[2:].strip()
                    else:
                        cur_q["stem_lines"].append(t)

        elif btype == 'tbl':
            tbl_cau_started = False
            for r in blk.rows:
                for c in r.cells:
                    for cp in c.paragraphs:
                        nids = cp._p.xpath('.//w:numPr/w:numId/@w:val')
                        ct = get_full_para_text(cp)
                        if (nids and nids[0] in cau_nids) or re.match(r'^Câu\s+\d+[\.\:]', ct):
                            tbl_cau_started = True
                            break
            if tbl_cau_started:
                flush_q()
                stem_lns = []
                blobs = []
                seen_cells = set()
                for r in blk.rows:
                    for c in r.cells:
                        if id(c._tc) in seen_cells:
                            continue
                        seen_cells.add(id(c._tc))
                        blobs.extend(extract_rids_from_el(c._tc, doc.part))
                        for cp in c.paragraphs:
                            ct = get_full_para_text(cp)
                            ct = re.sub(r'^Câu\s+\d+[\.\:]\s*', '', ct).strip()
                            if ct:
                                stem_lns.append(ct)
                cur_q = {
                    "topic": cur_topic,
                    "part": cur_part,
                    "stem_lines": stem_lns,
                    "opt_lines": [],
                    "p2_items": {},
                    "sol_lines": [],
                    "underlined_opt": None,
                    "in_sol": False,
                    "img_blobs": blobs,
                    "tables": [],
                }
            elif cur_q is not None and not cur_q["in_sol"]:
                blobs = extract_rids_from_el(blk._tbl, doc.part)
                if blobs:
                    cur_q["img_blobs"].extend(blobs)
                else:
                    t_data = []
                    for r in blk.rows:
                        row_vals = [get_full_para_text(c.paragraphs[0]) if c.paragraphs else "" for c in r.cells]
                        if any(row_vals):
                            t_data.append(row_vals)
                    if t_data:
                        cur_q["tables"].append(t_data)

    flush_q()

    cleaned = []
    for q in questions:
        sol_text = "\n".join(q["sol_lines"]).strip()
        sol_text = re.sub(r'^📝\s*Hướng dẫn giải:?\s*', '', sol_text).strip()
        sol_text = re.sub(r'^Hướng dẫn giải:?\s*', '', sol_text).strip()
        ans = ""
        if q["part"] == "P1":
            ans = q["underlined_opt"] or ""
            if not ans:
                m = re.search(r'Đáp án:?\s*([A-D])', sol_text)
                if m:
                    ans = m.group(1)
        elif q["part"] == "P2":
            p2_ans = {}
            for ln in q["sol_lines"]:
                m = re.match(r'^([a-d])\)\s*(Đúng|Sai)', ln.strip(), flags=re.IGNORECASE)
                if m:
                    p2_ans[m.group(1).lower()] = "Đ" if m.group(2).lower() == "đúng" else "S"
            if len(p2_ans) == 4:
                ans = f"a:{p2_ans['a']} - b:{p2_ans['b']} - c:{p2_ans['c']} - d:{p2_ans['d']}"
        elif q["part"] == "P3":
            for ln in q["sol_lines"]:
                m = re.search(r'Đáp\s+(?:án|số)\s*[:\.]?\s*([0-9]+(?:[\,\.][0-9]+)?)', ln, flags=re.IGNORECASE)
                if m:
                    ans = m.group(1).replace('.', ',')
                    break

        stem_full = " ".join(q["stem_lines"])
        lvl = classify_level(q["part"], stem_full, sol_text)
        item_levels = {}
        if q["part"] == "P2":
            for k in ["a", "b", "c", "d"]:
                item_levels[k] = classify_p2_item_level(q["p2_items"].get(k, ""))

        cleaned.append({
            "part": q["part"],
            "topic": q["topic"],
            "level": lvl,
            "item_levels": item_levels,
            "source_tag": "[THPT Châu Văn Liêm – 2026]",
            "stem_lines": q["stem_lines"],
            "opt_lines": q["opt_lines"],
            "p2_items": q["p2_items"],
            "answer": ans,
            "solution": sol_text,
            "img_blobs": q["img_blobs"],
            "tables": q["tables"],
        })
    return cleaned


def append_new_questions_to_c2_word_files(c2_all_questions, new_questions):
    info = CHAPTERS[2]
    for is_teacher in (False, True):
        src_doc_path = info["da_file"] if is_teacher else info["in_file"]
        doc = docx.Document(src_doc_path)
        setup_page_and_hf(doc)
        inject_question_tags_into_doc(doc, c2_all_questions[:127], is_teacher=is_teacher)

        for p in list(doc.paragraphs[-5:]):
            if "HẾT" in p.text:
                p._element.getparent().remove(p._element)

        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(10)
        p_sec.paragraph_format.space_after = Pt(4)
        r_sec = p_sec.add_run("PHẦN BỔ SUNG TỪ TÀI LIỆU THPT CHÂU VĂN LIÊM – 2026 (ĐÃ LỌC TRÙNG LẶP)")
        r_sec.bold = True
        r_sec.font.name = "Times New Roman"
        r_sec.font.size = Pt(12.5)
        r_sec.font.color.rgb = COLOR_HF_RED

        cur_part_rendered = None
        for q in new_questions:
            if q["part"] != cur_part_rendered:
                cur_part_rendered = q["part"]
                part_titles = {
                    "P1": "I. BỔ SUNG TRẮC NGHIỆM NHIỀU PHƯƠNG ÁN LỰA CHỌN",
                    "P2": "II. BỔ SUNG TRẮC NGHIỆM ĐÚNG / SAI",
                    "P3": "III. BỔ SUNG TRẮC NGHIỆM TRẢ LỜI NGẮN",
                }
                p_pt = doc.add_paragraph()
                p_pt.paragraph_format.space_before = Pt(6)
                p_pt.paragraph_format.space_after = Pt(3)
                r_pt = p_pt.add_run(part_titles.get(cur_part_rendered, ""))
                r_pt.bold = True
                r_pt.font.name = "Times New Roman"
                r_pt.font.size = Pt(11.5)
                r_pt.font.color.rgb = COLOR_PRIMARY

            p_q = doc.add_paragraph()
            p_q.paragraph_format.space_before = Pt(4)
            p_q.paragraph_format.space_after = Pt(2)
            r_qn = p_q.add_run(f"Câu {q['q_num']}. ")
            r_qn.bold = True
            r_qn.font.name = "Times New Roman"
            r_qn.font.size = Pt(11.5)
            r_qn.font.color.rgb = COLOR_PRIMARY

            lv_name = LEVEL_NAMES[q["level"]]
            ans_str = f" | ĐA: {q['answer']}" if (is_teacher and q["answer"]) else ""
            r_tag = p_q.add_run(f"[{q['id']} | {lv_name} – {q['topic']}{ans_str}] ")
            r_tag.bold = True
            r_tag.font.name = "Times New Roman"
            r_tag.font.size = Pt(10.5)
            r_tag.font.color.rgb = COLOR_TAG

            if q["stem_lines"]:
                add_chem_runs(p_q, q["stem_lines"][0], size=11.5)
                for extra_ln in q["stem_lines"][1:]:
                    p_ex = doc.add_paragraph()
                    p_ex.paragraph_format.space_before = Pt(1)
                    p_ex.paragraph_format.space_after = Pt(1)
                    add_chem_runs(p_ex, extra_ln, size=11.5)

            for img_p in q.get("image_paths", []):
                if os.path.exists(img_p):
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.add_run().add_picture(img_p, width=Cm(6.5))

            for tbl_d in q.get("tables", []):
                add_data_table_to_doc(doc, tbl_d)

            if q["part"] == "P1":
                opts = parse_abcd_from_opt_lines(q["opt_lines"])
                add_options_paragraph_no_table(doc, opts, q["answer"], is_teacher)
                if is_teacher:
                    add_solution_block(doc, q, opts)
            elif q["part"] == "P2":
                ans_map = {}
                for part_ans in q["answer"].split("-"):
                    if ":" in part_ans:
                        k_a, v_a = part_ans.strip().split(":", 1)
                        ans_map[k_a.strip()] = v_a.strip()
                for key in ["a", "b", "c", "d"]:
                    stmt_txt = q.get("p2_items", {}).get(key, "")
                    p_it = doc.add_paragraph()
                    p_it.paragraph_format.left_indent = Cm(0.4)
                    p_it.paragraph_format.space_before = Pt(1)
                    p_it.paragraph_format.space_after = Pt(1.5)
                    r_k = p_it.add_run(f"{key}) ")
                    r_k.bold = True
                    r_k.font.name = "Times New Roman"
                    r_k.font.size = Pt(11.5)
                    il_lv = q.get("item_levels", {}).get(key, "TH")
                    r_mt = p_it.add_run(f"[{LEVEL_NAMES.get(il_lv, 'Hiểu')}] ")
                    r_mt.bold = True
                    r_mt.font.name = "Times New Roman"
                    r_mt.font.size = Pt(10.5)
                    r_mt.font.color.rgb = COLOR_SOLUTION
                    add_chem_runs(p_it, stmt_txt, size=11.5)
                    if is_teacher and key in ans_map:
                        ds_val = "Đúng" if ans_map[key] == "Đ" else "Sai"
                        r_ds = p_it.add_run(f"   → [{ds_val}]")
                        r_ds.bold = True
                        r_ds.font.name = "Times New Roman"
                        r_ds.font.size = Pt(11)
                        r_ds.font.color.rgb = COLOR_CORRECT
                if is_teacher:
                    add_solution_block(doc, q)
            elif q["part"] == "P3":
                if is_teacher:
                    add_solution_block(doc, q)

        insert_matrix_summary_at_top(doc, info, c2_all_questions, is_teacher=is_teacher)
        suffix = "GiaoVien" if is_teacher else "HocSinh"
        p1 = os.path.join(OUT_DIR, f"{info['file_prefix']}_{suffix}.docx")
        p2 = os.path.join(BANK_DIR, f"{info['file_prefix']}_{suffix}.docx")
        doc.save(p1)
        doc.save(p2)
        print(f"  + Updated Chapter 2 {suffix} Word file: {p2}")


if __name__ == "__main__":
    doc = docx.Document(NEW_DOCX)
    cvl_questions = parse_cvl_c2_questions(doc)
    print(f"Tổng số câu hỏi chuẩn trích xuất từ phần Bài tập tự luyện Chương 2 (THPT Châu Văn Liêm): {len(cvl_questions)} câu")

    with open(BANK_JSON, "r", encoding="utf-8") as f:
        bank = json.load(f)

    existing_c2 = [q for q in bank["C2"]["questions"] if q["q_num"] <= 127]
    print(f"Số câu hỏi gốc trong Ngân hàng Chương 2 (C2): {len(existing_c2)} câu")

    existing_norms = []
    for eq in existing_c2:
        st = norm_for_compare(" ".join(eq["stem_lines"]))
        existing_norms.append((eq["q_num"], eq["id"], eq["part"], st))

    duplicates = []
    new_unique = []
    seen_new_norms = []

    for idx, nq in enumerate(cvl_questions, 1):
        nq_norm = norm_for_compare(" ".join(nq["stem_lines"]))
        if not nq_norm:
            continue
        is_dup = False
        dup_with = None
        for eq_num, eq_id, eq_part, eq_norm in existing_norms:
            if nq["part"] == eq_part:
                if nq_norm == eq_norm or (len(nq_norm) > 35 and (nq_norm in eq_norm or eq_norm in nq_norm)):
                    is_dup = True
                    dup_with = eq_id
                    break
                ratio = SequenceMatcher(None, nq_norm, eq_norm).ratio()
                if ratio >= 0.85:
                    is_dup = True
                    dup_with = f"{eq_id} ({ratio:.0%})"
                    break
        if not is_dup:
            for s_norm in seen_new_norms:
                if nq_norm == s_norm or SequenceMatcher(None, nq_norm, s_norm).ratio() >= 0.88:
                    is_dup = True
                    dup_with = "Trùng nội bộ trong file mới"
                    break
        if is_dup:
            duplicates.append((idx, nq["part"], nq["topic"], nq["stem_lines"][0][:75], dup_with))
        else:
            seen_new_norms.append(nq_norm)
            new_unique.append(nq)

    print(f"  + Số câu TRÙNG LẶP đã lọc bỏ: {len(duplicates)} câu")
    print(f"  + Số câu MỚI HOÀN TOÀN được bổ sung vào Chương 2: {len(new_unique)} câu")

    next_qnum = 128
    final_new_qs = []
    for nq in new_unique:
        q_num = next_qnum
        next_qnum += 1
        qid = f"C2-{nq['part']}-{nq['level']}-{q_num:03d}"
        img_paths = []
        for k, blob in enumerate(nq.get("img_blobs", []), 1):
            fpath = os.path.join(IMG_DIR, f"c2_q{q_num}_img{k}.png")
            try:
                im = Image.open(io.BytesIO(blob))
                im.save(fpath, format="PNG")
                img_paths.append(fpath)
            except Exception:
                pass
        q_record = {
            "q_num": q_num,
            "id": qid,
            "part": nq["part"],
            "level": nq["level"],
            "raw_topic": nq["topic"],
            "topic": nq["topic"],
            "source_tag": nq["source_tag"],
            "stem_lines": nq["stem_lines"],
            "opt_lines": nq["opt_lines"],
            "p2_items": nq["p2_items"],
            "item_levels": nq["item_levels"],
            "answer": nq["answer"],
            "solution": nq["solution"],
            "image_paths": img_paths,
            "tables": nq.get("tables", []),
        }
        final_new_qs.append(q_record)

    c2_all = existing_c2 + final_new_qs
    bank["C2"]["questions"] = c2_all
    bank["C2"]["total_questions"] = len(c2_all)

    with open(BANK_JSON, "w", encoding="utf-8") as f:
        json.dump(bank, f, ensure_ascii=False, indent=2)
    print(f"\nĐã cập nhật JSON Ngân hàng 9 chương: Chương 2 tăng từ 127 câu -> {len(c2_all)} câu!")

    append_new_questions_to_c2_word_files(c2_all, final_new_qs)

    print("\nThống kê chi tiết các câu hỏi mới vừa bổ sung vào Chương 2:")
    for pt in ["P1", "P2", "P3"]:
        sub = [q for q in final_new_qs if q["part"] == pt]
        nb = sum(1 for q in sub if q["level"] == "NB")
        th = sum(1 for q in sub if q["level"] == "TH")
        vd = sum(1 for q in sub if q["level"] == "VD")
        print(f"  - {pt}: {len(sub)} câu mới (Biết: {nb}, Hiểu: {th}, Vận dụng: {vd})")
