import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

CROP_DIR = r"c:\Antigravity_Thanh\San_Pham\images_crop_485"
OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"

COLOR_HF_RED = RGBColor(211, 47, 47)      # Coral Red (#D32F2F)
COLOR_HF_GREEN_HEX = "8BC390"             # Sage Green for horizontal lines (#8BC390)

COLOR_PRIMARY = RGBColor(46, 125, 50)     # Dark Green / Teal for headers in body
COLOR_CORRECT = RGBColor(198, 40, 40)     # Red for correct answers
COLOR_SOLUTION = RGBColor(21, 101, 192)   # Blue for solutions
COLOR_GRAY = RGBColor(97, 97, 97)

def add_p_border_bottom(p, color="8BC390", sz="8"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_p_border_top(p, color="8BC390", sz="8"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def set_run_style(run, font_name="Times New Roman", font_size=12, bold=False, italic=False, subscript=False, superscript=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def render_rich_text(p, text, base_font="Times New Roman", base_size=12, base_bold=False, base_italic=False, base_color=None):
    tokens = re.split(r'(~[^~]+~|\^[^\^]+\^|\*[^\*]+\*|_[^_]+_)', text)
    for token in tokens:
        if not token:
            continue
        is_bold = base_bold
        is_italic = base_italic
        is_sub = False
        is_sup = False
        content = token

        if token.startswith('*') and token.endswith('*'):
            is_bold = True
            content = token[1:-1]
        elif token.startswith('_') and token.endswith('_'):
            is_italic = True
            content = token[1:-1]
        elif token.startswith('~') and token.endswith('~'):
            is_sub = True
            content = token[1:-1]
        elif token.startswith('^') and token.endswith('^'):
            is_sup = True
            content = token[1:-1]

        run = p.add_run(content)
        set_run_style(run, font_name=base_font, font_size=base_size, bold=is_bold, italic=is_italic,
                      subscript=is_sub, superscript=is_sup, color=base_color)

def setup_header_footer(doc):
    normal_style = doc.styles['Normal']
    normal_style.paragraph_format.tab_stops.clear_all()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    section.start_type = WD_SECTION_START.NEW_PAGE
    section.different_first_page_header_footer = False
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    # Header with table in header layer (standard 2-column header as required by GEMINI.md)
    header = section.header
    for p in header.paragraphs:
        p.text = ""
    tbl_h = header.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl_h.cell(0, 0)
    c_right = tbl_h.cell(0, 1)
    c_left.width = Cm(8.0)
    c_right.width = Cm(10.0)

    p_hl = c_left.paragraphs[0]
    p_hl.paragraph_format.space_before = Pt(0)
    p_hl.paragraph_format.space_after = Pt(2)
    p_hl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_hl = p_hl.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)

    p_hr = c_right.paragraphs[0]
    p_hr.paragraph_format.space_before = Pt(0)
    p_hr.paragraph_format.space_after = Pt(2)
    p_hr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_hr1 = p_hr.add_run("Thầy ")
    set_run_style(r_hr1, font_size=10, italic=True, color=COLOR_HF_RED)
    r_hr2 = p_hr.add_run("TRẦN VĂN THẠNH")
    set_run_style(r_hr2, font_size=10, bold=True, italic=True, color=COLOR_HF_RED)
    r_hr3 = p_hr.add_run(" - 0777.470.803")
    set_run_style(r_hr3, font_size=10, italic=True, color=COLOR_HF_RED)

    add_p_border_bottom(p_hl, color=COLOR_HF_GREEN_HEX, sz="8")
    add_p_border_bottom(p_hr, color=COLOR_HF_GREEN_HEX, sz="8")

    # Footer
    footer = section.footer
    for p in footer.paragraphs:
        p.text = ""
    tbl_f = footer.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    cf_left = tbl_f.cell(0, 0)
    cf_right = tbl_f.cell(0, 1)
    cf_left.width = Cm(14.5)
    cf_right.width = Cm(3.5)

    p_fl = cf_left.paragraphs[0]
    p_fl.paragraph_format.space_before = Pt(2)
    p_fl.paragraph_format.space_after = Pt(0)
    p_fl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_fl = p_fl.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)

    p_fr = cf_right.paragraphs[0]
    p_fr.paragraph_format.space_before = Pt(2)
    p_fr.paragraph_format.space_after = Pt(0)
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_field(p_fr, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED)

    add_p_border_top(p_fl, color=COLOR_HF_GREEN_HEX, sz="8")
    add_p_border_top(p_fr, color=COLOR_HF_GREEN_HEX, sz="8")

def build_title_block_no_table(doc, is_teacher=False):
    """Xây dựng phần tiêu đề đầu trang KHÔNG DÙNG BẢNG (sử dụng tab stops)."""
    lines = [
        ("SỞ GD & ĐT THÀNH PHỐ HUẾ", "ĐỀ KIỂM TRA GIỮA HỌC KÌ I – NĂM HỌC 2025-2026", True, True, 10.5, 11),
        ("TRƯỜNG THPT HOÁ CHÂU", "MÔN: HOÁ HỌC – LỚP 12", True, True, 10.5, 11),
        ("ĐỀ CHÍNH THỨC", "Thời gian làm bài: 45 phút (không kể thời gian phát đề)", True, False, 10.5, 10),
        ("(Đề gồm có 03 trang)", "MÃ ĐỀ: 485", False, True, 10, 11)
    ]

    for l_text, r_text, l_bold, r_bold, l_size, r_size in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(3.8), WD_TAB_ALIGNMENT.CENTER)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(13.5), WD_TAB_ALIGNMENT.CENTER)

        p.add_run('\t')
        rl = p.add_run(l_text)
        set_run_style(rl, font_size=l_size, bold=l_bold, italic=(not l_bold and "(" in l_text))

        p.add_run('\t')
        rr = p.add_run(r_text)
        set_run_style(rr, font_size=r_size, bold=r_bold, italic=(not r_bold))

    if is_teacher:
        p_gv = doc.add_paragraph()
        p_gv.paragraph_format.space_before = Pt(1)
        p_gv.paragraph_format.space_after = Pt(2)
        p_gv.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rgv = p_gv.add_run("(BẢN GIÁO VIÊN -- CÓ LỜI GIẢI CHI TIẾT)")
        set_run_style(rgv, font_size=10.5, bold=True, color=COLOR_CORRECT)

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(4)
    p_info.paragraph_format.space_after = Pt(2)
    render_rich_text(p_info, "*Họ, tên thí sinh:* ............................................................ *Lớp:* ............. *Số báo danh:* ............. *Mã đề 485*", base_size=11)

    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(0)
    p_note.paragraph_format.space_after = Pt(4)
    render_rich_text(p_note, "_(Cho nguyên tử khối của các nguyên tử: H = 1; C = 12; O = 16; N = 14; Cl = 35,5)_", base_size=10, base_italic=True)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(4)
    add_p_border_bottom(p_div, color="CCCCCC", sz="4")

def add_section_header(doc, title, subtitle=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, f"*{title}*", base_size=11.5, base_bold=True, base_color=COLOR_PRIMARY)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(4)
        render_rich_text(p2, f"_{subtitle}_", base_size=10.5, base_italic=True)

def render_options_no_table(doc, opts):
    """Trình bày các phương án A, B, C, D hoàn toàn KHÔNG DÙNG BẢNG, dùng paragraph có tab stops."""
    clean_opts = [re.sub(r'[*_~^]', '', opt) for opt in opts]
    max_len = max(len(o) for o in clean_opts)

    if max_len <= 18:
        # Cả 4 phương án trên 1 dòng
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(4.5), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(13.5), WD_TAB_ALIGNMENT.LEFT)

        for idx, opt in enumerate(opts):
            if idx > 0:
                p.add_run('\t')
            render_rich_text(p, opt, base_size=10.5)
    elif max_len <= 38:
        # 2 phương án trên 1 dòng (2 dòng)
        p1 = doc.add_paragraph()
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.tab_stops.clear_all()
        p1.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
        render_rich_text(p1, opts[0], base_size=10.5)
        p1.add_run('\t')
        render_rich_text(p1, opts[1], base_size=10.5)

        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(1)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.tab_stops.clear_all()
        p2.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
        render_rich_text(p2, opts[2], base_size=10.5)
        p2.add_run('\t')
        render_rich_text(p2, opts[3], base_size=10.5)
    else:
        # Mỗi phương án trên 1 dòng (4 dòng)
        for opt in opts:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            render_rich_text(p, opt, base_size=10.5)

def add_solution_no_table(doc, solution_text):
    """Trình bày lời giải KHÔNG DÙNG BẢNG (đoạn văn thụt đầu dòng, phân biệt rõ nét)."""
    lines = solution_text.strip().split("\n")
    p_head = doc.add_paragraph()
    p_head.paragraph_format.left_indent = Cm(0.4)
    p_head.paragraph_format.space_before = Pt(3)
    p_head.paragraph_format.space_after = Pt(1)
    render_rich_text(p_head, "► *Lời giải chi tiết:*", base_size=10.5, base_bold=True, base_color=COLOR_SOLUTION)

    for line in lines:
        if line.startswith("*Lời giải chi tiết:*"):
            continue
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        render_rich_text(p, line, base_size=10)

def build_exam_docx(is_teacher=False):
    doc = docx.Document()
    setup_header_footer(doc)
    build_title_block_no_table(doc, is_teacher=is_teacher)

    # PHẦN I
    add_section_header(doc, "A. TRẮC NGHIỆM KHÁCH QUAN", None)
    add_section_header(doc, "PHẦN I (3 điểm - 12 câu): Câu trắc nghiệm nhiều phương án lựa chọn.",
                       "Thí sinh trả lời từ câu 1 đến câu 12. Mỗi câu hỏi thí sinh chỉ chọn một phương án.")

    part1_questions = [
        {
            "num": 1,
            "q": "Thuỷ phân 324 gam tinh bột với hiệu suất 75% khối lượng glucose thu được là",
            "opts": ["A. 360 gam.", "*B. 270 gam.*" if is_teacher else "B. 270 gam.", "C. 300 gam.", "D. 250 gam."],
            "sol": "- Phương trình phản ứng: (C~6~H~10~O~5~)~n~ + n H~2~O -(H^+, t°)-> n C~6~H~12~O~6~\n"
                   "- n~tinh bột~ = 324 / 162 = 2 mol.\n"
                   "- Theo lí thuyết: n~glucose (LT)~ = 2 mol -> m~glucose (LT)~ = 2 × 180 = 360 gam.\n"
                   "- Do hiệu suất H = 75%: m~glucose (TT)~ = 360 × 75% = *270 gam*.\n"
                   "- *Chọn B.*"
        },
        {
            "num": 2,
            "q": "Phát biểu nào sau đây không đúng khi nhận xét về maltose?",
            "opts": [
                "A. Là một disaccharide.",
                "B. Phân tử có một nhóm -OH hemiacetal.",
                "C. Cho được phản ứng thuỷ phân.",
                "*D. Không làm mất màu nước bromine.*" if is_teacher else "D. Không làm mất màu nước bromine."
            ],
            "sol": "- Maltose là disaccharide tạo bởi hai gốc α-glucose liên kết qua liên kết α-1,4-glycosidic.\n"
                   "- Trong phân tử maltose còn 1 nhóm -OH hemiacetal tự do ở gốc glucose thứ hai, nên ở dạng mạch hở maltose có nhóm chức aldehyde (-CHO).\n"
                   "- Do đó, maltose có tính khử, có khả năng *làm mất màu nước bromine*. Phát biểu D khẳng định 'Không làm mất màu nước bromine' là sai.\n"
                   "- *Chọn D.*"
        },
        {
            "num": 3,
            "q": "Khi xà phòng hoá tristearin thu được sản phẩm là",
            "opts": [
                "A. C~15~H~31~COONa và C~2~H~5~OH.",
                "B. C~15~H~31~COOH và C~3~H~5~(OH)~3~.",
                "*C. C~17~H~35~COONa và C~3~H~5~(OH)~3~.*" if is_teacher else "C. C~17~H~35~COONa và C~3~H~5~(OH)~3~.",
                "D. C~17~H~35~COOH và C~3~H~5~(OH)~3~."
            ],
            "sol": "- Tristearin là triglyceride có công thức (C~17~H~35~COO)~3~C~3~H~5~.\n"
                   "- Phản ứng xà phòng hoá trong môi trường kiềm (NaOH):\n"
                   "  (C~17~H~35~COO)~3~C~3~H~5~ + 3 NaOH -> 3 C~17~H~35~COONa + C~3~H~5~(OH)~3~\n"
                   "- Sản phẩm thu được là sodium stearate (C~17~H~35~COONa) và glycerol (C~3~H~5~(OH)~3~).\n"
                   "- *Chọn C.*"
        },
        {
            "num": 4,
            "q": "Công thức cấu tạo dạng mạch vòng của α-glucose là",
            "img": os.path.join(CROP_DIR, "fig_cau4_glucose_clean.png"),
            "img_width": 14.5,
            "opts": [
                "*A. Cấu trúc A.*" if is_teacher else "A. Cấu trúc A.",
                "B. Cấu trúc B.",
                "C. Cấu trúc C.",
                "D. Cấu trúc D."
            ],
            "sol": "- Vòng của glucose là vòng 6 cạnh pyranose.\n"
                   "- Ở dạng α-D-glucopyranose (cấu trúc A): nhóm -OH hemiacetal ở vị trí C1 hướng xuống dưới (khác phía với nhóm -CH~2~OH ở C5 hướng lên trên); C2 hướng xuống, C3 hướng lên, C4 hướng xuống.\n"
                   "- Cấu trúc B là β-D-glucopyranose (nhóm -OH ở C1 hướng lên trên).\n"
                   "- Cấu trúc C và D là các dạng vòng furanose 5 cạnh.\n"
                   "- *Chọn A.*"
        },
        {
            "num": 5,
            "q": "Đun nóng hỗn hợp các chất trong bình cầu chứa 12 mL acetic acid (D = 1,05 g/mL), 11 mL pentyl alcohol (D = 0,81 g/mL) và 4 mL dung dịch H~2~SO~4~ đặc cùng một ít đá bọt trong 20 phút. Khối lượng ester thu được sau khi tách khỏi hỗn hợp và làm sạch là 8 gam. Hiệu suất phản ứng ester hoá đạt",
            "opts": [
                "A. 29,30%.",
                "*B. 60,78%.*" if is_teacher else "B. 60,78%.",
                "C. 66,67%.",
                "D. 72,27%."
            ],
            "sol": "- Phản ứng: CH~3~COOH + C~5~H~11~OH <=(H2SO4 dac, to)=> CH~3~COOC~5~H~11~ + H~2~O\n"
                   "- Khối lượng các chất ban đầu:\n"
                   "  + m~CH3COOH~ = 12 × 1,05 = 12,6 g -> n~CH3COOH~ = 12,6 / 60 = 0,21 mol.\n"
                   "  + m~C5H11OH~ = 11 × 0,81 = 8,91 g -> n~C5H11OH~ = 8,91 / 88 = 0,10125 mol.\n"
                   "- Do 0,10125 < 0,21 nên CH~3~COOH dư, hiệu suất phản ứng tính theo pentyl alcohol (C~5~H~11~OH).\n"
                   "- Khối lượng ester (pentyl acetate, M = 130 g/mol) theo lí thuyết:\n"
                   "  m~ester (LT)~ = 0,10125 × 130 = 13,1625 gam.\n"
                   "- Hiệu suất phản ứng ester hoá:\n"
                   "  H = (8 / 13,1625) × 100% ≈ *60,78%*.\n"
                   "- *Chọn B.*"
        },
        {
            "num": 6,
            "q": "Chất nào sau đây không phải là ester?",
            "opts": [
                "A. HCOOCH~3~.",
                "*B. C~2~H~5~COC~3~H~7~.*" if is_teacher else "B. C~2~H~5~COC~3~H~7~.",
                "C. CH~3~COOC~2~H~5~.",
                "D. (C~15~H~31~COO)~3~C~3~H~5~."
            ],
            "sol": "- A, C, D đều chứa nhóm chức ester (-COO-).\n"
                   "- B (C~2~H~5~COC~3~H~7~ hay hexan-3-one) chứa nhóm chức carbonyl (>C=O) liên kết với hai gốc hydrocarbon, đây là một ketone, không phải ester.\n"
                   "- *Chọn B.*"
        },
        {
            "num": 7,
            "q": "Cho các phát biểu sau:\n"
               "(a) Khi làm đậu phụ xảy ra sự đông tụ protein.\n"
               "(b) Enzyme bị biến tính không thể thực hiện vai trò xúc tác.\n"
               "(c) Không nên vắt chanh vào sữa khi uống.\n"
               "(d) Sự thuỷ phân protein xảy ra trong quá trình làm nước mắm hay nấu nước tương.\n"
               "(e) Mỗi enzyme có một nhiệt độ tối ưu. Tại nhiệt độ tối ưu, enzyme có hoạt tính tối đa làm tốc độ phản ứng xảy ra nhanh nhất.\n"
               "Số phát biểu đúng là:",
            "opts": ["A. 2.", "B. 3.", "C. 4.", "*D. 5.*" if is_teacher else "D. 5."],
            "sol": "- (a) Đúng. Protein trong sữa đậu nành bị đông tụ khi thêm ion kim loại hoặc acid (nước chua, thạch cao).\n"
                   "- (b) Đúng. Khi bị biến tính do nhiệt độ hoặc hóa chất, cấu trúc không gian bậc cao bị phá hủy dẫn đến trung tâm hoạt động mất chức năng xúc tác.\n"
                   "- (c) Đúng. Acid citric trong chanh làm giảm pH gây đông tụ protein casein trong sữa, làm sữa vón cục khó tiêu.\n"
                   "- (d) Đúng. Protein trong cá (làm nước mắm) hoặc đậu nành (làm nước tương) bị protease thủy phân thành các amino acid và peptide bổ dưỡng.\n"
                   "- (e) Đúng. Mỗi enzyme hoạt động tối ưu ở một khoảng nhiệt độ và pH xác định.\n"
                   "- Cả 5 phát biểu đều đúng.\n"
                   "- *Chọn D.*"
        },
        {
            "num": 8,
            "q": "Tên gọi của hợp chất C~6~H~5~-CH~2~-CH(NH~2~)-COOH là",
            "opts": [
                "A. Amino phenyl propionic acid.",
                "B. phenylalanine.",
                "C. 2-amino-3-phenylpropionic acid.",
                "*D. 2-amino-3-phenylpropanoic acid*" if is_teacher else "D. 2-amino-3-phenylpropanoic acid"
            ],
            "sol": "- Chọn mạch chính dài nhất chứa nhóm -COOH gồm 3 carbon (propanoic acid).\n"
                   "- Đánh số từ carbon của nhóm carboxyl: C1 (-COOH), C2 gắn nhóm amino (-NH~2~), C3 gắn nhóm phenyl (-C~6~H~5~).\n"
                   "- Tên thay thế chuẩn IUPAC theo chương trình mới: *2-amino-3-phenylpropanoic acid*.\n"
                   "(Tên thông thường là phenylalanine; tên bán hệ thống là α-amino-β-phenylpropionic acid).\n"
                   "- *Chọn D.*"
        },
        {
            "num": 9,
            "q": "Công thức cấu tạo của glycine là",
            "opts": [
                "A. C~2~H~5~NH~2~.",
                "*B. H~2~NCH~2~COOH.*" if is_teacher else "B. H~2~NCH~2~COOH.",
                "C. CH~3~NH~2~.",
                "D. H~2~NCH(CH~3~)COOH."
            ],
            "sol": "- Glycine (kí hiệu Gly) là amino acid đơn giản nhất, có công thức H~2~N-CH~2~-COOH (aminoethanoic acid).\n"
                   "- *Chọn B.*"
        },
        {
            "num": 10,
            "q": "Cho dãy các chất: tinh bột, cellulose, glucose, fructose, saccharose. Số chất trong dãy thuộc loại monosaccharide là",
            "opts": ["A. 1.", "B. 3.", "*C. 2.*" if is_teacher else "C. 2.", "D. 4."],
            "sol": "- Monosaccharide: glucose, fructose (2 chất).\n"
                   "- Disaccharide: saccharose.\n"
                   "- Polysaccharide: tinh bột, cellulose.\n"
                   "- Vậy có 2 chất thuộc loại monosaccharide.\n"
                   "- *Chọn C.*"
        },
        {
            "num": 11,
            "q": "Chất nào sau đây là thành phần chính của chất giặt rửa tổng hợp?",
            "img_inline": os.path.join(CROP_DIR, "fig_cau11_detergent_clean.png"),
            "opts": [
                "A. (C~17~H~35~COO)~2~Ca.",
                "B. C~15~H~31~COONa.",
                "C. C~17~H~35~COOK.",
                "*D. CH~3~[CH~2~]~11~-C~6~H~4~-SO~3~Na (sodium dodecylbenzenesulfonate)*" if is_teacher else "D. CH~3~[CH~2~]~11~-C~6~H~4~-SO~3~Na (hình bên dưới)"
            ],
            "sol": "- B (C~15~H~31~COONa) và C (C~17~H~35~COOK) là xà phòng.\n"
                   "- A ((C~17~H~35~COO)~2~Ca) là muối kết tủa không tan, làm mất tác dụng của xà phòng trong nước cứng.\n"
                   "- D là sodium 4-dodecylbenzenesulfonate (chất giặt rửa tổng hợp thuộc loại alkylbenzenesulfonate).\n"
                   "- *Chọn D.*"
        },
        {
            "num": 12,
            "q": "Số amine bậc I trong số các chất: CH~3~NH~2~, CH~3~NH~3~Cl, (NH~2~)~2~CO, CH~3~NHCH~3~, CH~3~CH~2~NH~2~, NH~2~CH~2~NH~2~, (CH~3~)~3~N, C~6~H~5~NH~2~ (aniline) là:",
            "opts": ["A. 6.", "B. 7.", "C. 5.", "*D. 4.*" if is_teacher else "D. 4."],
            "sol": "- Amine bậc I có nhóm chức -NH~2~ liên kết với gốc hydrocarbon.\n"
                   "- Xét từng chất:\n"
                   "  1. CH~3~NH~2~: amine bậc I (thỏa mãn).\n"
                   "  2. CH~3~NH~3~Cl: muối ammonium, không phải amine (loại).\n"
                   "  3. (NH~2~)~2~CO: urea (diamide), không phải amine (loại).\n"
                   "  4. CH~3~NHCH~3~: amine bậc II (loại).\n"
                   "  5. CH~3~CH~2~NH~2~: amine bậc I (thỏa mãn).\n"
                   "  6. NH~2~CH~2~NH~2~: diamine bậc I (thỏa mãn).\n"
                   "  7. (CH~3~)~3~N: amine bậc III (loại).\n"
                   "  8. C~6~H~5~NH~2~: amine bậc I (thỏa mãn).\n"
                   "- Số amine bậc I là 4 chất (CH~3~NH~2~, CH~3~CH~2~NH~2~, NH~2~CH~2~NH~2~, C~6~H~5~NH~2~).\n"
                   "- *Chọn D.*"
        }
    ]

    for item in part1_questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*Câu {item['num']}:* {item['q']}", base_size=11)

        if "img" in item and os.path.exists(item["img"]):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(2)
            p_img.paragraph_format.space_after = Pt(3)
            p_img.add_run().add_picture(item["img"], width=Cm(item.get("img_width", 12.0)))

        # Render options hoàn toàn KHÔNG DÙNG BẢNG
        render_options_no_table(doc, item["opts"])

        if "img_inline" in item and os.path.exists(item["img_inline"]):
            p_inline = doc.add_paragraph()
            p_inline.paragraph_format.left_indent = Cm(0.5)
            p_inline.paragraph_format.space_before = Pt(1)
            p_inline.paragraph_format.space_after = Pt(2)
            p_inline.add_run("   Công thức D: ")
            p_inline.add_run().add_picture(item["img_inline"], width=Cm(6.5))

        if is_teacher and "sol" in item:
            add_solution_no_table(doc, item["sol"])

    # PHẦN II: Câu trắc nghiệm đúng sai
    add_section_header(doc, "PHẦN II (2 điểm - 2 câu): Câu trắc nghiệm đúng sai.",
                       "Thí sinh trả lời từ câu 1 đến câu 2. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.")

    part2_questions = [
        {
            "num": 1,
            "intro": "Thực hiện phản ứng điều chế isoamyl acetate theo trình tự sau:\n"
                     "- *Bước 1:* Cho 2 mL isoamyl alcohol, 2 mL acetic acid nguyên chất và 2 giọt sulfuric acid đặc vào ống nghiệm khô.\n"
                     "- *Bước 2:* Lắc đều, đun cách thủy hỗn hợp 5-6 phút trong nồi nước nóng.\n"
                     "- *Bước 3:* Để nguội, rồi rót hỗn hợp sản phẩm vào ống nghiệm chứa 2 mL dung dịch NaCl bão hòa.",
            "items": [
                ("a", "Tách isoamyl acetate từ hỗn hợp sau bước 3 bằng phương pháp chiết.", "Đúng"),
                ("b", "Mục đích chính của việc cho dung dịch NaCl bão hòa nhằm tránh sự phân hủy sản phẩm.", "Sai"),
                ("c", "Ở bước 2 xảy ra phản ứng ester hóa, giải phóng hơi có mùi thơm của chuối chín.", "Đúng"),
                ("d", "Phản ứng ester hóa giữa isoamyl alcohol với acetic acid là phản ứng thuận nghịch.", "Đúng")
            ],
            "sol": "- *Ý a) Đúng.* Sau bước 3, isoamyl acetate không tan trong nước, nhẹ hơn dung dịch NaCl bão hòa nên nổi lên trên tách thành 2 lớp chất lỏng phân biệt. Do đó dùng phương pháp chiết để tách riêng ester.\n"
                   "- *Ý b) Sai.* Mục đích của việc thêm dung dịch NaCl bão hòa là làm tăng khối lượng riêng của lớp dung dịch nước và làm giảm độ tan của ester trong nước (hiện tượng kết muối - salting out), giúp ester dễ dàng tách lớp nổi lên trên; không nhằm mục đích chống phân hủy.\n"
                   "- *Ý c) Đúng.* Ở bước 2, phản ứng ester hóa tạo isoamyl acetate có mùi thơm đặc trưng của chuối chín.\n"
                   "- *Ý d) Đúng.* Phản ứng ester hóa giữa acid hữu cơ và alcohol có xúc tác acid vô cơ mạnh là phản ứng thuận nghịch (hai chiều).\n"
                   "- *Đáp án:* a - Đúng, b - Sai, c - Đúng, d - Đúng."
        },
        {
            "num": 2,
            "intro": "Em hãy cho biết các phát biểu sau đúng hay sai?",
            "items": [
                ("a", "Thuỷ phân saccharose và maltose chỉ thu được một monosaccharide duy nhất.", "Sai"),
                ("b", "Glucose và Fructose đều có thể mở vòng thành dạng mạch hở.", "Đúng"),
                ("c", "Glucose và fructose là đồng phân của nhau.", "Đúng"),
                ("d", "Thuỷ phân 1 tấn tinh bột thu được 10 tấn glucose. Hiệu suất thuỷ phân tinh bột đạt 80%.", "Sai")
            ],
            "sol": "- *Ý a) Sai.* Thủy phân maltose chỉ thu được 1 monosaccharide là α-glucose; nhưng thủy phân saccharose thu được hỗn hợp gồm cả glucose và fructose.\n"
                   "- *Ý b) Đúng.* Trong dung dịch nước, các phân tử glucose và fructose luôn tồn tại cân bằng động giữa dạng mạch vòng và dạng mạch hở chứa nhóm carbonyl.\n"
                   "- *Ý c) Đúng.* Cả glucose và fructose đều có cùng công thức phân tử C~6~H~12~O~6~ nhưng khác công thức cấu tạo nên là đồng phân của nhau.\n"
                   "- *Ý d) Sai.* 1 tấn tinh bột (C~6~H~10~O~5~)~n~ theo lí thuyết (hiệu suất 100%) chỉ cho tối đa 1 × (180/162) ≈ 1,11 tấn glucose. Với hiệu suất 80% chỉ thu được 0,889 tấn glucose, không thể thu được 10 tấn.\n"
                   "- *Đáp án:* a - Sai, b - Đúng, c - Đúng, d - Sai."
        }
    ]

    for q in part2_questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*Câu {q['num']}:* {q['intro']}", base_size=11)

        # Trình bày các ý a, b, c, d hoàn toàn KHÔNG DÙNG BẢNG
        for label, text, ans in q["items"]:
            p_item = doc.add_paragraph()
            p_item.paragraph_format.left_indent = Cm(0.5)
            p_item.paragraph_format.space_before = Pt(1)
            p_item.paragraph_format.space_after = Pt(1)
            render_rich_text(p_item, f"*{label})* {text}", base_size=10.5)

            if is_teacher:
                color = COLOR_CORRECT if ans == "Đúng" else COLOR_SOLUTION
                p_item.add_run("   ➔ ")
                r_ans = p_item.add_run(f"[{ans}]")
                set_run_style(r_ans, font_size=10.5, bold=True, color=color)
            else:
                r_choice = p_item.add_run("   [ Đúng / Sai ]")
                set_run_style(r_choice, font_size=10, italic=True, color=COLOR_GRAY)

        if is_teacher and "sol" in q:
            add_solution_no_table(doc, q["sol"])

    # PHẦN III: Câu trắc nghiệm yêu cầu trả lời ngắn
    add_section_header(doc, "PHẦN III (2,0 điểm - 8 câu): Câu trắc nghiệm yêu cầu trả lời ngắn.",
                       "Thí sinh trả lời từ câu 1 đến câu 8.")

    part3_questions = [
        {
            "num": 1,
            "q": "Lên men m gam glucose với hiệu suất 90%, lượng khí CO~2~ sinh ra hấp thụ hết vào nước vôi trong, thu được 10 gam kết tủa. Khối lượng sau phản ứng giảm 3,4 gam so với khối lượng dung dịch nước vôi trong ban đầu. Giá trị của m là bao nhiêu?",
            "ans": "15",
            "sol": "- Kết tủa là CaCO~3~: n~CaCO3~ = 10 / 100 = 0,1 mol.\n"
                   "- Khối lượng dung dịch giảm 3,4 gam:\n"
                   "  Δm~dd~ = m~CO2~ - m~kết tủa~ = -3,4 gam\n"
                   "  => m~CO2~ = 10 - 3,4 = 6,6 gam => n~CO2~ = 6,6 / 44 = 0,15 mol.\n"
                   "- Phản ứng lên men: C~6~H~12~O~6~ -> 2 CO~2~ + 2 C~2~H~5~OH\n"
                   "  n~glucose (LT)~ = n~CO2~ / 2 = 0,15 / 2 = 0,075 mol.\n"
                   "- Vì hiệu suất lên men đạt 90%:\n"
                   "  n~glucose (TT)~ = 0,075 / 0,9 = 1/12 mol.\n"
                   "  => m = (1/12) × 180 = *15 gam*.\n"
                   "- *Đáp số:* *15*."
        },
        {
            "num": 2,
            "q": "Có bao nhiêu amine bậc 1 trong các amine sau?\n"
               "(1). CH~3~CH~2~CH~2~NH~2~.  (2). CH~3~CH(NH~2~)CH~3~.  (3). CH~3~NHCH~2~CH~3~.  (4). (CH~3~)~3~N.",
            "ans": "2",
            "sol": "- (1) CH~3~CH~2~CH~2~NH~2~: nhóm -NH~2~ gắn gốc propyl -> amine bậc 1.\n"
                   "- (2) CH~3~CH(NH~2~)CH~3~: nhóm -NH~2~ gắn gốc isopropyl -> amine bậc 1.\n"
                   "- (3) CH~3~NHCH~2~CH~3~: nhóm -NH- gắn 2 gốc -> amine bậc 2.\n"
                   "- (4) (CH~3~)~3~N: nguyên tử N gắn 3 gốc methyl -> amine bậc 3.\n"
                   "- Có 2 amine bậc 1 là (1) và (2).\n"
                   "- *Đáp số:* *2*."
        },
        {
            "num": 3,
            "q": "Một ester X mạch hở, có công thức phân tử C~4~H~8~O~2~. Có bao nhiêu đồng phân cấu tạo ester của X?",
            "ans": "4",
            "sol": "- Các đồng phân cấu tạo ester của C~4~H~8~O~2~ (dạng RCOOR'):\n"
                   "  1. HCOOCH~2~CH~2~CH~3~ (propyl formate)\n"
                   "  2. HCOOCH(CH~3~)~2~ (isopropyl formate)\n"
                   "  3. CH~3~COOCH~2~CH~3~ (ethyl acetate)\n"
                   "  4. CH~3~CH~2~COOCH~3~ (methyl propionate)\n"
                   "- Tổng cộng có 4 đồng phân cấu tạo ester.\n"
                   "- *Đáp số:* *4*."
        },
        {
            "num": 4,
            "q": "Có bao nhiêu đồng phân amine bậc một ứng với công thức phân tử C~4~H~11~N?",
            "ans": "4",
            "sol": "- Amine bậc một có dạng C~4~H~9~NH~2~ (nhóm -NH~2~ gắn vào các vị trí khác nhau trên 2 khung carbon C-C-C-C và C-C(C)-C):\n"
                   "  1. CH~3~-CH~2~-CH~2~-CH~2~-NH~2~ (butan-1-amine)\n"
                   "  2. CH~3~-CH~2~-CH(NH~2~)-CH~3~ (butan-2-amine)\n"
                   "  3. (CH~3~)~2~CH-CH~2~-NH~2~ (2-methylpropan-1-amine)\n"
                   "  4. (CH~3~)~3~C-NH~2~ (2-methylpropan-2-amine)\n"
                   "- Tổng cộng có 4 đồng phân amine bậc 1.\n"
                   "- *Đáp số:* *4*."
        },
        {
            "num": 5,
            "q": "Cơ thể người mã hoá được mấy loại amino acid có trong bảng sau để tổng hợp protein cho cơ thể?",
            "img": os.path.join(CROP_DIR, "fig_cau5_amino_acids_clean.png"),
            "img_width": 14.0,
            "ans": "1",
            "sol": "- Cơ thể người chỉ sử dụng các α-amino acid tiêu chuẩn (trong số 20 amino acid tiêu chuẩn) để mã hóa di truyền tổng hợp protein.\n"
                   "- Quan sát 4 chất:\n"
                   "  + Chất A: ^+^H~3~N-CH~2~-CH~2~-COO^- (β-alanine, amino acid nhóm β, không dùng tổng hợp protein).\n"
                   "  + Chất B: CH~3~-CH(NH~3~^+^)-COO^- (dạng ion lưỡng cực của α-alanine / Ala, là α-amino acid tiêu chuẩn tham gia mã hóa protein).\n"
                   "  + Chất C: ^+^H~3~N-CH~2~-CH(NH~2~)-CH~2~-COO^- (diaminobutanoic acid, không phải amino acid tiêu chuẩn).\n"
                   "  + Chất D: H~2~N-CH~2~-OH (amino alcohol, không phải amino acid).\n"
                   "- Vậy chỉ có 1 loại amino acid (chất B) được mã hóa để tổng hợp protein.\n"
                   "- *Đáp số:* *1*."
        },
        {
            "num": 6,
            "q": "Thuỷ phân hết m gam tetrapeptide Ala-Ala-Ala-Ala (mạch hở) thu được hỗn hợp gồm 28,24 gam Ala, 32 gam Ala-Ala và 27,27 gam Ala-Ala-Ala. Giá trị của m là bao nhiêu? (Làm tròn kết quả 1 lần cuối cùng đến hàng đơn vị)",
            "ans": "81",
            "sol": "- Khối lượng mol các peptide và amino acid:\n"
                   "  M~Ala~ = 89 g/mol; M~Ala2~ = 2×89 - 18 = 160 g/mol; M~Ala3~ = 3×89 - 36 = 231 g/mol; M~Ala4~ = 4×89 - 54 = 302 g/mol.\n"
                   "- Số mol các sản phẩm thu được:\n"
                   "  n~Ala~ = 28,24 / 89 ≈ 0,3173 mol.\n"
                   "  n~Ala2~ = 32 / 160 = 0,2 mol.\n"
                   "  n~Ala3~ = 27,27 / 231 ≈ 0,11805 mol.\n"
                   "- Bảo toàn số mol gốc Ala:\n"
                   "  n~Ala (tổng)~ = 1 × n~Ala~ + 2 × n~Ala2~ + 3 × n~Ala3~ = 0,3173 + 2×0,2 + 3×0,11805 = 1,07145 mol.\n"
                   "- Do mỗi phân tử tetrapeptide Ala~4~ có chứa 4 gốc Ala:\n"
                   "  n~Ala4~ = n~Ala (tổng)~ / 4 = 1,07145 / 4 ≈ 0,26786 mol.\n"
                   "- Khối lượng tetrapeptide ban đầu:\n"
                   "  m = n~Ala4~ × M~Ala4~ = 0,26786 × 302 ≈ 80,895 gam ≈ *81 gam*.\n"
                   "- *Đáp số:* *81*."
        },
        {
            "num": 7,
            "q": "Saccharose tham gia được bao nhiêu phản ứng nào sau đây?\n"
               "(a). Phản ứng với thuốc thử Tollens.\n"
               "(b). Phản ứng với nước bromine.\n"
               "(c). Phản ứng với Cu(OH)~2~ tạo dung dịch màu xanh lam.\n"
               "(d). Phản ứng thuỷ phân trong môi trường kiềm.",
            "ans": "1",
            "sol": "- (a) Không tham gia vì saccharose không có nhóm -OH hemiacetal/hemiketal tự do, không mở vòng tạo chức aldehyde nên không tráng bạc.\n"
                   "- (b) Không tham gia vì saccharose không có tính khử, không làm mất màu nước bromine.\n"
                   "- (c) Có tham gia vì saccharose có nhiều nhóm -OH kề nhau tạo phức chất xanh lam với Cu(OH)~2~ ở nhiệt độ thường.\n"
                   "- (d) Không tham gia vì carbohydrate không bị thủy phân trong môi trường kiềm (chỉ thủy phân trong môi trường acid hoặc có enzyme).\n"
                   "- Vậy saccharose chỉ tham gia 1 phản ứng duy nhất là (c).\n"
                   "- *Đáp số:* *1*."
        },
        {
            "num": 8,
            "q": "Protein tham gia được loại phản ứng dưới đây? Viết theo số thứ tự tăng dần (1234...)\n"
               "(1). Phản ứng khử thành alcohol.\n"
               "(2). Phản ứng màu với Cu(OH)~2~.\n"
               "(3). Phản ứng màu với HNO~3~.\n"
               "(4). Phản ứng thuỷ phân.",
            "ans": "234",
            "sol": "- (1) Không xảy ra phản ứng khử protein thành alcohol.\n"
                   "- (2) Phản ứng màu biuret với Cu(OH)~2~ trong kiềm tạo dung dịch màu tím đặc trưng.\n"
                   "- (3) Phản ứng màu xanthoproteic với HNO~3~ đặc tạo kết tủa màu vàng.\n"
                   "- (4) Phản ứng thủy phân liên kết peptide tạo các amino acid.\n"
                   "- Các phản ứng protein tham gia được là: (2), (3), (4).\n"
                   "- Viết theo thứ tự tăng dần: *234*.\n"
                   "- *Đáp số:* *234*."
        }
    ]

    for item in part3_questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*Câu {item['num']}:* {item['q']}", base_size=11)

        if "img" in item and os.path.exists(item["img"]):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(2)
            p_img.paragraph_format.space_after = Pt(2)
            p_img.add_run().add_picture(item["img"], width=Cm(item.get("img_width", 12.0)))

        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.left_indent = Cm(0.5)
        p_ans.paragraph_format.space_before = Pt(1)
        p_ans.paragraph_format.space_after = Pt(2)
        if is_teacher:
            render_rich_text(p_ans, f"*Đáp số:* *{item['ans']}*", base_size=10.5, base_bold=True, base_color=COLOR_CORRECT)
        else:
            render_rich_text(p_ans, "Đáp số: ............................................................", base_size=10.5, base_color=COLOR_GRAY)

        if is_teacher and "sol" in item:
            add_solution_no_table(doc, item["sol"])

    # PHẦN IV: TỰ LUẬN
    add_section_header(doc, "B. PHẦN IV - TỰ LUẬN: 3 điểm", None)

    part4_questions = [
        {
            "num": 1,
            "score": "1,5 điểm",
            "q": "a) Cho hợp chất hữu cơ X có công thức phân tử C~3~H~6~O~2~. Hãy viết các đồng phân cấu tạo ester và gọi tên chúng.\n"
               "b) Viết công thức khung phân tử của stearic acid và oleic acid, biết oleic acid là acid béo omega-9, có liên kết đôi C=C ở dạng cis. Dự đoán nhiệt độ nóng chảy của oleic acid và stearic acid. Giải thích.",
            "sol": "*a) Đồng phân cấu tạo ester của C~3~H~6~O~2~ và tên gọi:*\n"
                   "- Độ bất bão hòa k = 1 -> ester no, đơn chức, mạch hở.\n"
                   "  1. HCOOCH~2~CH~3~: ethyl formate (hoặc ethyl methanoate).\n"
                   "  2. CH~3~COOCH~3~: methyl acetate (hoặc methyl ethanoate).\n"
                   "\n"
                   "*b) Công thức khung phân tử và nhiệt độ nóng chảy:*\n"
                   "- *Stearic acid* (C~17~H~35~COOH): acid béo no 18 carbon, công thức khung là chuỗi thẳng zigzag liên tục gồm 17 liên kết C-C nối tiếp kết thúc bằng nhóm -COOH.\n"
                   "- *Oleic acid* (cis-C~17~H~33~COOH): acid béo không no 18 carbon omega-9, có 1 liên kết đôi C=C cấu hình cis ở vị trí C9=C10. Do cấu hình cis nên mạch carbon bị gập khúc (bẻ cong) tại vị trí liên kết đôi.\n"
                   "- *Dự đoán nhiệt độ nóng chảy:* Nhiệt độ nóng chảy của stearic acid cao hơn oleic acid (stearic acid ở thể rắn, t°~nc~ ≈ 69,6°C; oleic acid ở thể lỏng, t°~nc~ ≈ 13 - 14°C ở nhiệt độ thường).\n"
                   "- *Giải thích:*\n"
                   "  + Phân tử stearic acid có mạch hydrocarbon no duỗi thẳng đều đặn, cho phép các phân tử sắp xếp sít nhau một cách chặt chẽ, tạo diện tích tiếp xúc lớn và lực tương tác Van der Waals giữa các phân tử mạnh -> nhiệt độ nóng chảy cao.\n"
                   "  + Phân tử oleic acid có liên kết đôi dạng cis làm mạch carbon bị bẻ cong, các phân tử khó xếp khít chặt vào nhau, diện tích tiếp xúc nhỏ hơn và lực tương tác liên phân tử yếu hơn nhiều -> nhiệt độ nóng chảy thấp hơn."
        },
        {
            "num": 2,
            "score": "0,75 đ",
            "q": "a) Cho các lọ đựng các dung dịch sau bị mất nhãn: glucose, fructose và tinh bột. Trình bày cách nhận biết chúng.\n"
               "b) X là hỗn hợp gồm glucose, fructose và tinh bột. Phân tích thành phần nguyên tố trong X được %C = 40,82%; %H = 6,57% (theo khối lượng). Phần trăm khối lượng tinh bột trong X là bao nhiêu?",
            "sol": "*a) Phương pháp nhận biết 3 dung dịch:*\n"
                   "- Bước 1: Trích mỗi dung dịch một ít làm mẫu thử.\n"
                   "- Bước 2: Nhỏ vài giọt dung dịch iodine (I~2~) vào 3 mẫu thử:\n"
                   "  + Mẫu thử xuất hiện màu xanh tím đặc trưng là *tinh bột*.\n"
                   "  + Hai mẫu thử còn lại (glucose và fructose) không có hiện tượng đổi màu.\n"
                   "- Bước 3: Lấy hai mẫu thử chưa nhận biết (glucose và fructose), nhỏ vài giọt nước bromine (Br~2~) vào từng ống nghiệm:\n"
                   "  + Mẫu thử làm mất màu vàng nâu của nước bromine là *glucose* (do nhóm -CHO bị oxy hóa thành nhóm -COOH):\n"
                   "    CH~2~OH[CHOH]~4~CHO + Br~2~ + H~2~O -> CH~2~OH[CHOH]~4~COOH + 2 HBr\n"
                   "  + Mẫu thử không làm mất màu nước bromine là *fructose* (trong môi trường acid yếu của nước bromine, fructose không chuyển hóa thành glucose).\n"
                   "\n"
                   "*b) Tính phần trăm khối lượng tinh bột trong X:*\n"
                   "- Glucose và fructose có cùng CTPT C~6~H~12~O~6~ (M = 180 g/mol):\n"
                   "  %C~monosaccharide~ = (6 × 12 / 180) × 100% = 40,00%.\n"
                   "- Tinh bột có công thức mắt xích C~6~H~10~O~5~ (M = 162 g/mol):\n"
                   "  %C~tinh bột~ = (6 × 12 / 162) × 100% = (72 / 162) × 100% = 44,44%.\n"
                   "- Gọi phần trăm khối lượng của tinh bột trong X là a%:\n"
                   "  %C~X~ = 40,00 × (100 - a)/100 + (72/162) × 100 × (a/100) = 40,82\n"
                   "  <=> 40,00 + a × (72/162 - 0,4) = 40,82\n"
                   "  <=> a × (2/45) = 0,82\n"
                   "  <=> a = 0,82 × 45 / 2 = *18,45%*.\n"
                   "(Kiểm tra lại theo %H: %H~X~ = 6,67% × (1 - 0,1845) + 6,17% × 0,1845 = 6,57% - hoàn toàn chính xác).\n"
                   "- *Vậy phần trăm khối lượng của tinh bột trong X là 18,45%*."
        },
        {
            "num": 3,
            "score": "0,75 đ",
            "q": "Cho 0,02 mol α-amino acid X tác dụng vừa đủ với dung dịch chứa 0,04 mol NaOH. Mặt khác 0,02 mol X tác dụng vừa đủ với 0,02 mol HCl, thu được 3,67 gam muối.\n"
               "a) Nhúng Quỳ tím vào dung dịch X thì quỳ tím chuyển sang màu gì? Giải thích.\n"
               "b) Xác định công thức cấu tạo và gọi tên X, Viết kí hiệu của X.",
            "sol": "*a) Hiện tượng khi nhúng quỳ tím vào dung dịch X:*\n"
                   "- Tỉ lệ mol n~NaOH~ / n~X~ = 0,04 / 0,02 = 2 => X có 2 nhóm carboxyl (-COOH).\n"
                   "- Tỉ lệ mol n~HCl~ / n~X~ = 0,02 / 0,02 = 1 => X có 1 nhóm amino (-NH~2~).\n"
                   "- Do số nhóm -COOH nhiều hơn số nhóm -NH~2~ (2 nhóm -COOH > 1 nhóm -NH~2~) nên dung dịch X có tính acid (pH < 7).\n"
                   "- *Hiện tượng:* Nhúng quỳ tím vào dung dịch X làm *quỳ tím chuyển sang màu đỏ (hoặc hồng)*.\n"
                   "\n"
                   "*b) Xác định CTCT, tên gọi và kí hiệu của X:*\n"
                   "- Khi tác dụng với HCl:\n"
                   "  H~2~N-R-(COOH)~2~ + HCl -> ClH~3~N-R-(COOH)~2~ (muối)\n"
                   "  n~muối~ = n~X~ = 0,02 mol.\n"
                   "  Khối lượng mol của muối: M~muối~ = 3,67 / 0,02 = 183,5 g/mol.\n"
                   "- Khối lượng mol của X: M~X~ = M~muối~ - 36,5 = 183,5 - 36,5 = 147 g/mol.\n"
                   "- Ta có: M~R~ + 16 + 2 × 45 = 147 => M~R~ = 41 (gốc -C~3~H~5~-).\n"
                   "- Vì X là α-amino acid, nhóm -NH~2~ đính ở C số 2 (Cα), nên công thức cấu tạo của X là:\n"
                   "  *HOOC-CH~2~-CH~2~-CH(NH~2~)-COOH*\n"
                   "- *Tên gọi:* *glutamic acid* (hoặc acid glutamic; tên thay thế: *2-aminopentanedioic acid*).\n"
                   "- *Kí hiệu của X:* *Glu* (hoặc kí hiệu một chữ cái: *E*)."
        }
    ]

    for item in part4_questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*Câu {item['num']} ({item['score']}):* {item['q']}", base_size=11)

        if is_teacher and "sol" in item:
            add_solution_no_table(doc, item["sol"])
        elif not is_teacher:
            p_blank = doc.add_paragraph()
            p_blank.paragraph_format.left_indent = Cm(0.5)
            p_blank.paragraph_format.space_before = Pt(2)
            p_blank.paragraph_format.space_after = Pt(4)
            render_rich_text(p_blank, "Bài làm:\n........................................................................................................................................................................................\n........................................................................................................................................................................................\n........................................................................................................................................................................................", base_size=10.5, base_color=COLOR_GRAY)

    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_end.paragraph_format.space_before = Pt(12)
    p_end.paragraph_format.space_after = Pt(4)
    render_rich_text(p_end, "*--------- HẾT ---------*", base_size=11, base_bold=True)

    filename = "De_Kiem_Tra_Hoa_12_Ma485_GiaoVien.docx" if is_teacher else "De_Kiem_Tra_Hoa_12_Ma485.docx"
    filepath = os.path.join(OUT_DIR, filename)
    doc.save(filepath)
    print(f"Saved: {filepath}")

if __name__ == "__main__":
    build_exam_docx(is_teacher=False)
    build_exam_docx(is_teacher=True)
