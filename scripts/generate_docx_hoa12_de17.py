# -*- coding: utf-8 -*-
"""
Script tạo 2 file Word (.docx) chuẩn Form Đề Kiểm Tra cho:
ĐỀ SỐ 17 - CHƯƠNG 6: ĐẠI CƯƠNG KIM LOẠI (HÓA HỌC 12)
Mã đề: 117
Bao gồm:
1. Bản Học sinh: De_Kiem_Tra_Hoa_12_De17.docx
2. Bản Giáo viên: De_Kiem_Tra_Hoa_12_De17_GiaoVien.docx (Có đáp án đỏ, lời giải chi tiết xanh, bảng đáp án)
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"

COLOR_CORAL_RED = RGBColor(211, 47, 47)      # #D32F2F
COLOR_SAGE_GREEN = "8BC390"                 # #8BC390
COLOR_DEEP_BLUE = RGBColor(21, 101, 192)    # #1565C0
COLOR_DARK_GREEN = RGBColor(46, 125, 50)    # #2E7D32
COLOR_BLACK = RGBColor(0, 0, 0)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.59)     # 1.5 cm
        s.bottom_margin = Inches(0.59)  # 1.5 cm
        s.left_margin = Inches(0.59)    # 1.5 cm
        s.right_margin = Inches(0.59)   # 1.5 cm
        s.header_distance = Inches(0.24) # 0.6 cm
        s.footer_distance = Inches(0.24) # 0.6 cm

        # --- HEADER ---
        header = s.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.text = ""
        hp.paragraph_format.space_before = Pt(0)
        hp.paragraph_format.space_after = Pt(0)

        tbl_h = header.add_table(rows=1, cols=2, width=Inches(7.08)) # 18.0 cm
        tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
        c0, c1 = tbl_h.rows[0].cells
        c0.width = Inches(3.54)
        c1.width = Inches(3.54)
        set_cell_margins(c0, 60, 60, 0, 0)
        set_cell_margins(c1, 60, 60, 0, 0)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run("Tài liệu lưu hành nội bộ")
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(10)
        r0.font.italic = True
        r0.font.color.rgb = COLOR_CORAL_RED

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1_a = p1.add_run("Thầy ")
        r1_a.font.name = "Times New Roman"
        r1_a.font.size = Pt(10)
        r1_a.font.color.rgb = COLOR_CORAL_RED
        r1_b = p1.add_run("TRẦN VĂN THẠNH")
        r1_b.font.name = "Times New Roman"
        r1_b.font.size = Pt(10)
        r1_b.font.bold = True
        r1_b.font.color.rgb = COLOR_CORAL_RED
        r1_c = p1.add_run(" - 0777.470.803")
        r1_c.font.name = "Times New Roman"
        r1_c.font.size = Pt(10)
        r1_c.font.color.rgb = COLOR_CORAL_RED

        tblPr_h = tbl_h._tbl.tblPr
        tblBorders_h = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{COLOR_SAGE_GREEN}"/>'
            f'</w:tblBorders>'
        )
        tblPr_h.append(tblBorders_h)

        # --- FOOTER ---
        footer = s.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.text = ""
        fp.paragraph_format.space_before = Pt(0)
        fp.paragraph_format.space_after = Pt(0)

        tbl_f = footer.add_table(rows=1, cols=2, width=Inches(7.08))
        tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
        fc0, fc1 = tbl_f.rows[0].cells
        fc0.width = Inches(5.31)
        fc1.width = Inches(1.77)
        set_cell_margins(fc0, 60, 60, 0, 0)
        set_cell_margins(fc1, 60, 60, 0, 0)

        fp0 = fc0.paragraphs[0]
        fp0.paragraph_format.space_before = Pt(0)
        fp0.paragraph_format.space_after = Pt(0)
        fr0 = fp0.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
        fr0.font.name = "Times New Roman"
        fr0.font.size = Pt(9.5)
        fr0.font.italic = True
        fr0.font.color.rgb = COLOR_CORAL_RED

        fp1 = fc1.paragraphs[0]
        fp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp1.paragraph_format.space_before = Pt(0)
        fp1.paragraph_format.space_after = Pt(0)
        fr1_a = fp1.add_run("Trang ")
        fr1_a.font.name = "Times New Roman"
        fr1_a.font.size = Pt(10)
        fr1_a.font.bold = True
        fr1_a.font.color.rgb = COLOR_CORAL_RED

        fldSimple = OxmlElement('w:fldSimple')
        fldSimple.set(qn('w:instr'), 'PAGE')
        r_page = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        rFont = OxmlElement('w:rFonts')
        rFont.set(qn('w:ascii'), 'Times New Roman')
        rFont.set(qn('w:hAnsi'), 'Times New Roman')
        b = OxmlElement('w:b')
        c_el = OxmlElement('w:color')
        c_el.set(qn('w:val'), 'D32F2F')
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '20')
        rPr.append(rFont)
        rPr.append(b)
        rPr.append(c_el)
        rPr.append(sz)
        r_page.append(rPr)
        t_el = OxmlElement('w:t')
        t_el.text = "1"
        r_page.append(t_el)
        fldSimple.append(r_page)
        fp1._p.append(fldSimple)

        tblPr_f = tbl_f._tbl.tblPr
        tblBorders_f = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{COLOR_SAGE_GREEN}"/>'
            f'</w:tblBorders>'
        )
        tblPr_f.append(tblBorders_f)

def add_chem_run(paragraph, text, bold=False, italic=False, color=None, font_size=11):
    parts = re.split(r'(~.*?~|\^.*?\^)', text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run()
        run.font.name = "Times New Roman"
        run.font.size = Pt(font_size)
        run.bold = bold
        run.italic = italic
        if color:
            run.font.color.rgb = color

        if part.startswith('~') and part.endswith('~'):
            run.text = part[1:-1]
            run.font.subscript = True
        elif part.startswith('^') and part.endswith('^'):
            run.text = part[1:-1]
            run.font.superscript = True
        else:
            run.text = part

def create_title_block(doc, is_teacher=False):
    lines_info = [
        ("SỞ GIÁO DỤC VÀ ĐÀO TẠO THỪA THIÊN HUẾ", "ĐỀ KIỂM TRA ĐỊNH KỲ CHƯƠNG 6 (ĐẠI CƯƠNG KIM LOẠI)", True, True, COLOR_DARK_GREEN),
        ("TRƯỜNG THPT HAI BÀ TRƯNG", "MÔN: HOÁ HỌC -- LỚP 12", True, True, None),
        ("ĐỀ CHÍNH THỨC", "Thời gian làm bài: 50 phút (không kể thời gian phát đề)", True, False, None),
        ("(Đề thi có 04 trang)", "Mã đề thi: 117", False, True, None)
    ]

    for left_txt, right_txt, left_bold, right_bold, right_color in lines_info:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.08), WD_TAB_ALIGNMENT.RIGHT)

        r_l = p.add_run(left_txt)
        r_l.font.name = "Times New Roman"
        r_l.font.size = Pt(11)
        r_l.bold = left_bold
        if not left_bold:
            r_l.italic = True

        p.add_run("\t")

        r_r = p.add_run(right_txt)
        r_r.font.name = "Times New Roman"
        r_r.font.size = Pt(11)
        r_r.bold = right_bold
        if right_color:
            r_r.font.color.rgb = right_color
        if not right_bold:
            r_r.italic = True

    if is_teacher:
        p_note = doc.add_paragraph()
        p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_note.paragraph_format.space_before = Pt(3)
        p_note.paragraph_format.space_after = Pt(3)
        r_note = p_note.add_run("(BẢN GIÁO VIÊN -- CÓ LỜI GIẢI CHI TIẾT)")
        r_note.font.name = "Times New Roman"
        r_note.font.size = Pt(12)
        r_note.font.bold = True
        r_note.font.color.rgb = COLOR_CORAL_RED

    p_stu = doc.add_paragraph()
    p_stu.paragraph_format.space_before = Pt(4)
    p_stu.paragraph_format.space_after = Pt(2)
    r_stu = p_stu.add_run("Họ, tên thí sinh: .................................................................................... Lớp: ........................ SBD: ....................")
    r_stu.font.name = "Times New Roman"
    r_stu.font.size = Pt(11)
    r_stu.italic = True

    p_mass = doc.add_paragraph()
    p_mass.paragraph_format.space_before = Pt(2)
    p_mass.paragraph_format.space_after = Pt(6)
    r_mass = p_mass.add_run(
        "(Cho biết nguyên tử khối: H = 1; C = 12; N = 14; O = 16; Na = 23; Mg = 24; Al = 27; S = 32; Cl = 35,5; K = 39; Ca = 40; Fe = 56; Cu = 64; Zn = 65; Br = 80; Ag = 108; Ba = 137; Au = 197; Pb = 207; hằng số Faraday F = 96500 C/mol)."
    )
    r_mass.font.name = "Times New Roman"
    r_mass.font.size = Pt(9.5)
    r_mass.italic = True

def add_section_header(doc, title, instruction):
    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_before = Pt(6)
    p_t.paragraph_format.space_after = Pt(2)
    r_t = p_t.add_run(title)
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(11)
    r_t.font.bold = True

    p_i = doc.add_paragraph()
    p_i.paragraph_format.space_before = Pt(0)
    p_i.paragraph_format.space_after = Pt(4)
    r_i = p_i.add_run(instruction)
    r_i.font.name = "Times New Roman"
    r_i.font.size = Pt(11)
    r_i.italic = True

def add_loigiai(doc, explanation_text):
    p_lbl = doc.add_paragraph()
    p_lbl.paragraph_format.left_indent = Inches(0.3)
    p_lbl.paragraph_format.space_before = Pt(3)
    p_lbl.paragraph_format.space_after = Pt(1)
    r_lbl = p_lbl.add_run("Lời giải chi tiết:")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(10.5)
    r_lbl.font.bold = True
    r_lbl.font.italic = True
    r_lbl.font.color.rgb = COLOR_DEEP_BLUE

    for line in explanation_text.strip().split("\n"):
        p_c = doc.add_paragraph()
        p_c.paragraph_format.left_indent = Inches(0.3)
        p_c.paragraph_format.space_before = Pt(0)
        p_c.paragraph_format.space_after = Pt(1.5)
        p_c.paragraph_format.line_spacing = 1.15
        add_chem_run(p_c, line, color=COLOR_DEEP_BLUE, font_size=10.5)

def add_mcq_options(doc, options, correct_key, is_teacher=False, layout="1_row"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15

    if layout == "1_row":
        tabs = [Inches(1.77), Inches(3.54), Inches(5.31)]
        for t in tabs:
            p.paragraph_format.tab_stops.add_tab_stop(t, WD_TAB_ALIGNMENT.LEFT)
        for i, (k, txt) in enumerate(options):
            if i > 0:
                p.add_run("\t")
            is_c = (k == correct_key) and is_teacher
            col = COLOR_CORAL_RED if is_c else None
            add_chem_run(p, f"{k}. ", bold=True, color=col)
            add_chem_run(p, txt, bold=is_c, color=col)

    elif layout == "2_rows":
        p.paragraph_format.tab_stops.add_tab_stop(Inches(3.54), WD_TAB_ALIGNMENT.LEFT)
        for i in range(2):
            k, txt = options[i]
            if i > 0:
                p.add_run("\t")
            is_c = (k == correct_key) and is_teacher
            col = COLOR_CORAL_RED if is_c else None
            add_chem_run(p, f"{k}. ", bold=True, color=col)
            add_chem_run(p, txt, bold=is_c, color=col)

        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(3)
        p2.paragraph_format.line_spacing = 1.15
        p2.paragraph_format.tab_stops.add_tab_stop(Inches(3.54), WD_TAB_ALIGNMENT.LEFT)
        for i in range(2, 4):
            k, txt = options[i]
            if i > 2:
                p2.add_run("\t")
            is_c = (k == correct_key) and is_teacher
            col = COLOR_CORAL_RED if is_c else None
            add_chem_run(p2, f"{k}. ", bold=True, color=col)
            add_chem_run(p2, txt, bold=is_c, color=col)

    else: # 4_rows
        for i, (k, txt) in enumerate(options):
            if i > 0:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
            is_c = (k == correct_key) and is_teacher
            col = COLOR_CORAL_RED if is_c else None
            add_chem_run(p, f"{k}. ", bold=True, color=col)
            add_chem_run(p, txt, bold=is_c, color=col)

def generate_exam_docx(is_teacher=False):
    doc = docx.Document()
    add_header_footer(doc)
    create_title_block(doc, is_teacher=is_teacher)

    # ==============================================================
    # PHẦN I
    # ==============================================================
    add_section_header(
        doc,
        "PHẦN I (4,5 điểm -- 18 câu): Câu trắc nghiệm nhiều phương án lựa chọn.",
        "Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án."
    )

    p1_data = [
        (
            1,
            "Ở trạng thái cơ bản, cấu hình electron lớp ngoài cùng của nguyên tử Al (Z = 13) là",
            [
                ("A", "3s^1^."),
                ("B", "3s^2^."),
                ("C", "3s^2^3p^1^."),
                ("D", "3p^1^.")
            ],
            "C",
            "1_row",
            "Nguyên tử Al có Z = 13: 1s^2^2s^2^2p^6^3s^2^3p^1^. Lớp ngoài cùng (lớp n = 3) có cấu hình là 3s^2^3p^1^.\n"
            "Chọn C."
        ),
        (
            2,
            "Ở trạng thái cơ bản, nguyên tử của nguyên tố nào sau đây có số electron lớp ngoài cùng ít nhất?",
            [
                ("A", "C (Z = 6)."),
                ("B", "Li (Z = 3)."),
                ("C", "O (Z = 8)."),
                ("D", "F (Z = 9).")
            ],
            "B",
            "1_row",
            "Cấu hình electron lớp ngoài cùng:\n"
            "- C (Z = 6): 2s^2^2p^2^ (4 electron).\n"
            "- Li (Z = 3): 2s^1^ (1 electron).\n"
            "- O (Z = 8): 2s^2^2p^4^ (6 electron).\n"
            "- F (Z = 9): 2s^2^2p^5^ (7 electron).\n"
            "Vậy Li có số electron lớp ngoài cùng ít nhất (1 electron).\n"
            "Chọn B."
        ),
        (
            3,
            "Ở trạng thái cơ bản, nguyên tử của nguyên tố nào sau đây có 3 electron ở lớp ngoài cùng?",
            [
                ("A", "~11~Na."),
                ("B", "~13~Al."),
                ("C", "~20~Ca."),
                ("D", "~26~Fe.")
            ],
            "B",
            "1_row",
            "- ~11~Na: [Ne] 3s^1^ (1 e lớp ngoài cùng).\n"
            "- ~13~Al: [Ne] 3s^2^3p^1^ (3 e lớp ngoài cùng).\n"
            "- ~20~Ca: [Ar] 4s^2^ (2 e lớp ngoài cùng).\n"
            "- ~26~Fe: [Ar] 3d^6^4s^2^ (2 e lớp ngoài cùng).\n"
            "Chọn B."
        ),
        (
            4,
            "Thành phần nào sau đây không có trong mạng tinh thể kim loại?",
            [
                ("A", "Ion kim loại."),
                ("B", "Electron."),
                ("C", "Nguyên tử kim loại."),
                ("D", "Anion gốc acid.")
            ],
            "D",
            "2_rows",
            "Mạng tinh thể kim loại bao gồm cation kim loại ở các nút mạng, các electron hoá trị tự do chuyển động và một phần nguyên tử kim loại trung hoà. Mạng tinh thể kim loại không chứa anion gốc acid.\n"
            "Chọn D."
        ),
        (
            5,
            "Đốt một sợi dây kim loại X trong bình khí chlorine (Cl~2~) thấy tạo ra khói màu nâu đỏ. X là kim loại nào sau đây?",
            [
                ("A", "Mg."),
                ("B", "Al."),
                ("C", "Fe."),
                ("D", "Cu.")
            ],
            "C",
            "1_row",
            "Sắt cháy trong khí chlorine theo phản ứng: 2Fe + 3Cl~2~ → 2FeCl~3~ (tạo khói màu nâu đỏ đặc trưng của các hạt tinh thể FeCl~3~ rắn).\n"
            "Chọn C."
        ),
        (
            6,
            "Thuỷ ngân (Hg) là chất lỏng ở điều kiện thường, dễ bay hơi và hơi thuỷ ngân rất độc. Khi nhiệt kế thuỷ ngân bị vỡ, để tránh thuỷ ngân phân tán ra môi trường xung quanh, người ta gom thuỷ ngân lại rồi trộn với chất nào sau đây (ở dạng bột)?",
            [
                ("A", "Carbon (C)."),
                ("B", "Magnesium (Mg)."),
                ("C", "Iron (Fe)."),
                ("D", "Sulfur (S).")
            ],
            "D",
            "2_rows",
            "Dùng bột lưu huỳnh (S) rắc lên thuỷ ngân vì phản ứng xảy ra ngay ở nhiệt độ thường tạo muối HgS kết tủa bền, không bay hơi và không độc hại: Hg + S → HgS.\n"
            "Chọn D."
        ),
        (
            7,
            "Tính chất hoá học đặc trưng của kim loại là",
            [
                ("A", "tính oxi hoá."),
                ("B", "tính acid."),
                ("C", "tính base."),
                ("D", "tính khử.")
            ],
            "D",
            "1_row",
            "Các nguyên tử kim loại có ít electron lớp ngoài cùng, năng lượng ion hoá thấp và bán kính nguyên tử tương đối lớn nên dễ nhường electron: M → M^n+^ + ne. Do đó, tính chất hoá học đặc trưng của kim loại là tính khử.\n"
            "Chọn D."
        ),
        (
            8,
            "Trong vỏ Trái Đất, kim loại nào sau đây có thể tồn tại ở dạng đơn chất?",
            [
                ("A", "Na, Mg."),
                ("B", "Al, Fe."),
                ("C", "Cu, Zn."),
                ("D", "Ag, Au.")
            ],
            "D",
            "1_row",
            "Các kim loại rất kém hoạt động hoá học như bạc (Ag), vàng (Au), platinum (Pt) có thế điện cực chuẩn dương lớn, trơ về mặt hoá học nên có thể tồn tại trong tự nhiên ở dạng đơn chất.\n"
            "Chọn D."
        ),
        (
            9,
            "Ion kim loại nào sau đây có nhiều nhất trong nước biển?",
            [
                ("A", "Mg^2+^."),
                ("B", "K^+^."),
                ("C", "Na^+^."),
                ("D", "Ca^2+^.")
            ],
            "C",
            "1_row",
            "Nước biển chứa hàm lượng muối NaCl rất cao (chiếm khoảng 77,8% lượng muối hoà tan), do đó cation kim loại dồi dào nhất trong nước biển là Na^+^.\n"
            "Chọn C."
        ),
        (
            10,
            "Hai nguyên tố kim loại phổ biến nhất trong vỏ Trái Đất là",
            [
                ("A", "Mg, Ca."),
                ("B", "Na, Ca."),
                ("C", "Au, Ag."),
                ("D", "Al, Fe.")
            ],
            "D",
            "1_row",
            "Trong vỏ Trái Đất, nhôm (Al, chiếm khoảng 8%) là kim loại phổ biến nhất, tiếp theo là sắt (Fe, chiếm khoảng 5%).\n"
            "Chọn D."
        ),
        (
            11,
            "Nhôm (Al) là nguyên tố phổ biến thứ ba và là kim loại phổ biến nhất trong vỏ Trái Đất. Trong tự nhiên, quặng chính chứa nhôm là bauxite và quặng này là nguyên liệu chính để sản xuất nhôm trong công nghiệp. Thành phần chính của quặng bauxite là",
            [
                ("A", "Na~3~AlF~6~."),
                ("B", "Al~2~O~3~.2H~2~O."),
                ("C", "KAl(SO~4~)~2~.12H~2~O."),
                ("D", "K~2~O.Al~2~O~3~.6SiO~2~.")
            ],
            "B",
            "2_rows",
            "Quặng bauxite có thành phần chính là nhôm oxide ngậm nước Al~2~O~3~.2H~2~O (hoặc Al~2~O~3~.xH~2~O chứa tạp chất Fe~2~O~3~, SiO~2~,...).\n"
            "Chọn B."
        ),
        (
            12,
            "Hợp kim là vật liệu kim loại chứa một kim loại cơ bản với",
            [
                ("A", "một số kim loại khác hoặc phi kim."),
                ("B", "một số oxide của kim loại đó."),
                ("C", "một số oxide kim loại khác hoặc phi kim."),
                ("D", "một số phi kim và oxide của phi kim đó.")
            ],
            "A",
            "2_rows",
            "Hợp kim là vật liệu kim loại chứa một kim loại cơ bản (kim loại nền) kết hợp với một số kim loại khác hoặc phi kim nhằm cải thiện các tính chất cơ học, vật lý và hoá học.\n"
            "Chọn A."
        ),
        (
            13,
            "Phát biểu nào sau đây sai?",
            [
                ("A", "Tính chất vật lí của hợp kim khác nhiều so với của đơn chất thành phần."),
                ("B", "Tính chất hoá học của hợp kim tương tự tính chất của đơn chất thành phần."),
                ("C", "Tính chất cơ học của hợp kim phụ thuộc vào thành phần đơn chất của hợp kim."),
                ("D", "Tính chất cơ học của hợp kim không phụ thuộc vào hàm lượng các đơn chất.")
            ],
            "D",
            "4_rows",
            "Tính chất cơ học (độ cứng, độ bền, độ đàn hồi...) của hợp kim phụ thuộc chặt chẽ vào thành phần và tỉ lệ phần trăm hàm lượng của các nguyên tố trong hợp kim. Do đó phát biểu D là sai.\n"
            "Chọn D."
        ),
        (
            14,
            "Trường hợp nào sau đây đã sử dụng phương pháp điện hoá để chống ăn mòn kim loại?",
            [
                ("A", "Mạ vàng lên quai đồng hồ."),
                ("B", "Bôi dầu mỡ lên xích xe đạp."),
                ("C", "Gắn tấm kẽm lên mặt ngoài vỏ tàu biển."),
                ("D", "Sơn kín bề mặt khung cửa thép.")
            ],
            "C",
            "2_rows",
            "Gắn tấm kẽm (Zn) lên vỏ tàu biển bằng thép là phương pháp bảo vệ điện hoá (dùng kim loại hy sinh). Kẽm có tính khử mạnh hơn sắt nên đóng vai trò là anode bị ăn mòn, bảo vệ vỏ tàu sắt không bị gỉ. Các phương án A, B, D là phương pháp phủ bảo vệ bề mặt.\n"
            "Chọn C."
        ),
        (
            15,
            "Trường hợp nào sau đây không xảy ra sự ăn mòn kim loại?",
            [
                ("A", "Tàu đánh cá làm bằng thép bị hoen gỉ sau thời gian đi biển về."),
                ("B", "Trống đồng bị chuyển màu xanh khi để lâu ngày trong không khí ẩm."),
                ("C", "Vòng tay làm bằng bạc kim loại bị hoá đen khi sử dụng lâu ngày."),
                ("D", "Nấu chảy vàng để đúc khuôn khi chế tác vàng trang sức.")
            ],
            "D",
            "4_rows",
            "Nấu chảy vàng chỉ là quá trình chuyển thể vật lý từ thể rắn sang thể lỏng, không có sự biến đổi chất hay quá trình oxi hoá kim loại nên không phải là sự ăn mòn kim loại.\n"
            "Chọn D."
        ),
        (
            16,
            "Khi kim loại bị ăn mòn, luôn xảy ra quá trình nào sau đây?",
            [
                ("A", "Oxi hoá kim loại thành ion kim loại."),
                ("B", "Khử ion kim loại thành kim loại."),
                ("C", "Khử oxygen trong không khí."),
                ("D", "Khử kim loại thành ion kim loại.")
            ],
            "A",
            "2_rows",
            "Bản chất của mọi quá trình ăn mòn kim loại là sự oxi hoá kim loại thành ion kim loại: M → M^n+^ + ne.\n"
            "Chọn A."
        ),
        (
            17,
            "Trường hợp nào sau đây kim loại bị phá huỷ chủ yếu do ăn mòn hoá học?",
            [
                ("A", "Thiết bị làm bằng thép trong lò đốt lâu ngày bị phá huỷ."),
                ("B", "Thép xây dựng bị gỉ khi để lâu ngày trong không khí ẩm."),
                ("C", "Ống nước làm bằng gang bị gỉ khi chôn dưới đất lâu ngày."),
                ("D", "Vỏ tàu biển làm bằng thép bị gỉ sau một thời gian sử dụng.")
            ],
            "A",
            "4_rows",
            "Thiết bị làm bằng thép trong lò đốt tiếp xúc trực tiếp với các khí oxi hoá (O~2~, CO~2~, SO~2~...) ở nhiệt độ rất cao, các electron chuyển trực tiếp sang môi trường mà không phát sinh dòng điện nên đây là ăn mòn hoá học. Các trường hợp còn lại đều xảy ra trong dung dịch chất điện li nên là ăn mòn điện hoá học.\n"
            "Chọn A."
        ),
        (
            18,
            "Trường hợp nào sau đây kim loại bị phá huỷ chủ yếu do ăn mòn điện hoá học?",
            [
                ("A", "Thiết bị làm bằng sắt trong lò đốt lâu ngày bị phá huỷ."),
                ("B", "Thước nhôm để trong không khí bị phủ bởi lớp Al~2~O~3~ ở bề mặt."),
                ("C", "Cánh cửa làm bằng thép nhanh bị gỉ khi để ngoài trời mưa."),
                ("D", "Đốt dây sắt trong bình khí chlorine.")
            ],
            "C",
            "4_rows",
            "Thép là hợp kim của sắt và carbon. Khi gặp nước mưa (chứa chất điện li), tạo thành vô số pin vi mô Fe - C dẫn đến ăn mòn điện hoá học của sắt.\n"
            "Chọn C."
        )
    ]

    for q_num, content, opts, ans_key, layout, sol in p1_data:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        add_chem_run(p_q, f"Câu {q_num}. ", bold=True)
        add_chem_run(p_q, content)

        add_mcq_options(doc, opts, correct_key=ans_key, is_teacher=is_teacher, layout=layout)
        if is_teacher:
            add_loigiai(doc, sol)

    # ==============================================================
    # PHẦN II
    # ==============================================================
    add_section_header(
        doc,
        "PHẦN II (4,0 điểm -- 4 câu): Câu trắc nghiệm đúng sai.",
        "Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai."
    )

    p2_data = [
        (
            1,
            "Tùy thuộc vào tính chất vật lí riêng của mỗi kim loại mà chúng được sử dụng vào những mục đích khác nhau. Mỗi phát biểu sau đây là đúng hay sai?",
            [
                ("a", "Kim loại chì (Pb) và cadmium (Cd) có nhiệt nóng chảy khá thấp nên được sử dụng làm dây chảy trong cầu chì.", True),
                ("b", "Kim loại tungsten (W) có độ bền nhiệt và nhiệt độ nóng chảy rất cao nên được sử dụng làm dây tóc bóng đèn, thiết bị sưởi.", True),
                ("c", "Do có tính dẻo và độ cứng phù hợp nên nhôm (Al) thường được gia công làm vật liệu như khung cửa, khung thiết bị.", True),
                ("d", "Do kim loại magnesium (Mg) có khối lượng riêng là 1,735 g/cm^3^ nên được dùng để chế tạo các hợp kim nặng.", False)
            ],
            None,
            "a) Đúng. Chì và cadmium có nhiệt độ nóng chảy thấp nên hợp kim của chúng được ứng dụng làm dây chảy cầu chì tự ngắt khi quá nhiệt.\n"
            "b) Đúng. Tungsten có nhiệt độ nóng chảy cao nhất trong các kim loại (3422 °C) nên được dùng làm dây tóc bóng đèn sợi đốt.\n"
            "c) Đúng. Nhôm có khối lượng nhẹ, tính dẻo cao và dễ uốn tạo hình khung cửa, khung kết cấu.\n"
            "d) Sai. Magnesium có khối lượng riêng nhỏ (1,735 g/cm^3^ < 5 g/cm^3^) nên là kim loại nhẹ, hợp kim của nó là hợp kim nhẹ."
        ),
        (
            2,
            "Đồng (Cu) là kim loại có tính khử yếu, không tan trong dung dịch H~2~SO~4~ loãng, nhưng tan được trong dung dịch H~2~SO~4~ đặc, nóng hoặc trong dung dịch H~2~SO~4~ loãng khi có mặt của O~2~ ngay ở nhiệt độ thường theo các phương trình hoá học sau:\n"
            "Cu + 2H~2~SO~4~ đặc, nóng → CuSO~4~ + SO~2~ + 2H~2~O (1)\n"
            "2Cu + 2H~2~SO~4~ loãng + O~2~ → 2CuSO~4~ + 2H~2~O (2)\n"
            "Mỗi phát biểu sau đây là đúng hay sai?",
            [
                ("a", "Trong hai phản ứng trên, H~2~SO~4~ đều đóng vai trò là chất oxi hoá.", False),
                ("b", "Cùng một lượng CuSO~4~ tạo ra thì phản ứng (1) tiêu tốn H~2~SO~4~ nhiều hơn phản ứng (2).", True),
                ("c", "Do có giá trị thế điện cực chuẩn dương nên Cu tác dụng được với dung dịch H~2~SO~4~ loãng ở (2).", False),
                ("d", "Từ phản ứng (2) chứng tỏ E°(Cu^2+^/Cu) < E°(O~2~, 4H^+^ / 2H~2~O).", True)
            ],
            None,
            "a) Sai. Trong phản ứng (2), chất oxi hoá là O~2~ (số oxi hoá giảm từ 0 xuống -2), H~2~SO~4~ chỉ đóng vai trò môi trường.\n"
            "b) Đúng. Ở phản ứng (1), để tạo 1 mol CuSO~4~ cần 2 mol H~2~SO~4~; ở phản ứng (2), để tạo 1 mol CuSO~4~ chỉ cần 1 mol H~2~SO~4~.\n"
            "c) Sai. E°(Cu^2+^/Cu) = +0,34 V > 0 V nên Cu không đẩy được H^+^ của acid loãng; phản ứng (2) xảy ra là do có mặt chất oxi hoá mạnh hơn O~2~ trong môi trường acid.\n"
            "d) Đúng. Phản ứng xảy ra tự phát chứng tỏ chất oxi hoá mạnh hơn (O~2~ trong acid, E° = +1,23 V) oxi hoá được Cu (E° = +0,34 V)."
        ),
        (
            3,
            "Cho các phát biểu sau về ứng dụng và tính chất của kim loại. Mỗi phát biểu sau đây là đúng hay sai?",
            [
                ("a", "Trong công nghiệp, nhôm được sản xuất bằng cách điện phân dung dịch AlCl~3~.", False),
                ("b", "Cho mẩu nhỏ Na vào dung dịch CuSO~4~, thu được chất rắn là kim loại Cu.", False),
                ("c", "Tái chế kim loại là giải pháp giúp con người sử dụng hiệu quả hơn nguồn tài nguyên.", True),
                ("d", "Tất cả hợp kim của sắt đều dễ bị ăn mòn trong không khí ẩm.", False)
            ],
            None,
            "a) Sai. Điện phân dung dịch AlCl~3~ tại cathode nước bị khử sinh ra H~2~, không thu được Al. Nhôm được điều chế bằng điện phân nóng chảy Al~2~O~3~.\n"
            "b) Sai. Na tác dụng với nước tạo NaOH và H~2~, sau đó NaOH tạo kết tủa Cu(OH)~2~ màu xanh, không thu được Cu kim loại.\n"
            "c) Đúng. Tái chế kim loại giúp tiết kiệm tài nguyên khoáng sản, giảm tiêu tốn năng lượng và giảm ô nhiễm môi trường.\n"
            "d) Sai. Thép không gỉ (inox) chứa Cr tạo lớp màng oxide thụ động bảo vệ rất bền vững, không bị gỉ trong không khí ẩm."
        ),
        (
            4,
            "Tiến hành thí nghiệm (như hình vẽ bên): Rót dung dịch NaCl bão hoà vào cốc 1, cốc 2, cốc 3; cho dầu nhờn vào cốc 4. Cho vào cốc 1 và cốc 4 một đinh sắt sạch; cho vào cốc 2 đinh sắt sạch được quấn bởi dây kẽm; cho vào cốc 3 đinh sắt sạch được quấn bởi dây đồng. Để 4 cốc trong không khí khoảng 5 ngày. Mỗi phát biểu sau đây là đúng hay sai?",
            [
                ("a", "Ở cốc 4, đinh sắt không bị gỉ (không bị ăn mòn).", True),
                ("b", "Ở cốc 2, đinh sắt không bị gỉ, dây Zn bị ăn mòn và có khí thoát ra.", False),
                ("c", "Ở cốc 3, đinh sắt bị gỉ nhiều nhất và dây đồng không bị ăn mòn.", True),
                ("d", "Ở cốc 1, đinh sắt bị gỉ và dung dịch có màu vàng của FeCl~2~.", False)
            ],
            r"c:\Antigravity_Thanh\San_Pham\images_crop_de17\p2_cau4_an_mon_sat.png",
            "a) Đúng. Dầu nhờn cách li đinh sắt hoàn toàn khỏi không khí và nước nên không xảy ra ăn mòn.\n"
            "b) Sai. Zn bị ăn mòn bảo vệ sắt, nhưng trong môi trường NaCl trung tính, ở cathode xảy ra sự khử oxi hoà tan (O~2~ + 2H~2~O + 4e → 4OH^-^), không có khí thoát ra.\n"
            "c) Đúng. Cặp Fe - Cu tạo pin điện hoá trong đó Fe là anode bị ăn mòn nhanh hơn nhiều so với cốc 1; Cu là cathode được bảo vệ không bị ăn mòn.\n"
            "d) Sai. Gỉ sắt là Fe~2~O~3~.nH~2~O kết tủa màu nâu đỏ; dung dịch FeCl~2~ loãng có màu lục rất nhạt hoặc không màu, không phải màu vàng."
        )
    ]

    for q_num, content, sub_items, img_path, sol in p2_data:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        add_chem_run(p_q, f"Câu {q_num}. ", bold=True)
        add_chem_run(p_q, content)

        if img_path and os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(2)
            p_img.paragraph_format.space_after = Pt(2)
            p_img.add_run().add_picture(img_path, width=Inches(4.5))

        for sub_lbl, sub_txt, is_true in sub_items:
            p_sub = doc.add_paragraph()
            p_sub.paragraph_format.left_indent = Inches(0.2)
            p_sub.paragraph_format.space_before = Pt(0)
            p_sub.paragraph_format.space_after = Pt(1.5)
            p_sub.paragraph_format.line_spacing = 1.15

            add_chem_run(p_sub, f"{sub_lbl}) ", bold=True)
            add_chem_run(p_sub, sub_txt)

            if is_teacher:
                status_txt = " [ĐÚNG]" if is_true else " [SAI]"
                add_chem_run(p_sub, status_txt, bold=True, color=COLOR_CORAL_RED)

        if is_teacher:
            add_loigiai(doc, sol)

    # ==============================================================
    # PHẦN III
    # ==============================================================
    add_section_header(
        doc,
        "PHẦN III (1,5 điểm -- 6 câu): Câu trắc nghiệm yêu cầu trả lời ngắn.",
        "Thí sinh trả lời từ câu 1 đến câu 6."
    )

    p3_data = [
        (
            1,
            "Cho các nhận định sau:\n"
            "(1) Các kim loại thuộc nhóm IA và IIA bao gồm các nguyên tố s.\n"
            "(2) Nguyên tử của hầu hết các kim loại có từ 1 đến 3 electron ở lớp ngoài cùng.\n"
            "(3) Nguyên tử kim loại dễ nhường electron hơn so với các nguyên tử phi kim.\n"
            "(4) Trong mạng tinh thể kim loại, các electron chuyển động tự do theo nhiều hướng.\n"
            "Có bao nhiêu nhận định là đúng?",
            "4",
            "Cả 4 nhận định đều đúng:\n"
            "- (1) Nhóm IA có cấu hình ns^1^, nhóm IIA có ns^2^ nên đều là nguyên tố s.\n"
            "- (2) Hầu hết nguyên tử kim loại có 1, 2 hoặc 3 electron lớp ngoài cùng.\n"
            "- (3) Kim loại có năng lượng ion hoá thấp nên dễ nhường electron hơn phi kim.\n"
            "- (4) Electron tự do chuyển động hỗn loạn không định hướng trong toàn bộ tinh thể.\n"
            "Đáp số: 4."
        ),
        (
            2,
            "Sắt (Fe) có số hiệu nguyên tử là 26. Trong vỏ Trái Đất, sắt là nguyên tố kim loại phổ biến thứ 2 (sau nhôm). Ứng dụng chủ yếu của sắt là để tạo ra các hợp kim thép dùng trong xây dựng và chế tạo. Cho các phát biểu sau:\n"
            "(a) Ở trạng thái cơ bản, nguyên tử Fe có 6 electron ở lớp ngoài cùng.\n"
            "(b) Hợp kim thép carbon khi để trong không khí ẩm sẽ bị ăn mòn điện hoá học.\n"
            "(c) Trong vỏ Trái Đất, sắt tồn tại ở dạng hợp chất chủ yếu trong các quặng.\n"
            "(d) Trong hợp chất, số oxi hoá của sắt chủ yếu là +2 và +3.\n"
            "(e) Kim loại sắt (dư) tác dụng với chlorine tạo ra sản phẩm là FeCl~2~.\n"
            "Hãy cho biết có bao nhiêu phát biểu đúng trong các phát biểu trên.",
            "3",
            "Các phát biểu đúng là (b), (c), (d):\n"
            "- (a) Sai vì Fe ([Ar] 3d^6^4s^2^) có 2 electron ở lớp ngoài cùng (lớp n = 4).\n"
            "- (b) Đúng vì cặp Fe - C tiếp xúc với màng nước điện li bị ăn mòn điện hoá.\n"
            "- (c) Đúng vì sắt tồn tại trong quặng hematite, magnetite, siderite, pyrite...\n"
            "- (d) Đúng vì sắt có số oxi hoá phổ biến là +2 và +3.\n"
            "- (e) Sai vì chlorine là chất oxi hoá mạnh, luôn oxi hoá sắt lên mức cao nhất là FeCl~3~ ngay cả khi sắt dư.\n"
            "Tổng cộng có 3 phát biểu đúng.\n"
            "Đáp số: 3."
        ),
        (
            3,
            "Tiến hành các thí nghiệm:\n"
            "(1) Đốt dây nhôm trong bình khí oxygen.\n"
            "(2) Cho đinh thép sạch vào dung dịch H~2~SO~4~ loãng.\n"
            "(3) Cho lá Al vào dung dịch CuSO~4~.\n"
            "(4) Nối dây nhôm với dây đồng rồi để trong không khí ẩm.\n"
            "Liệt kê những thí nghiệm có xảy ra ăn mòn điện hoá học theo số thứ tự tăng dần (ví dụ: 1, 2, 3).",
            "2, 3, 4",
            "- (1) Đốt nhôm trong khí O~2~ ở nhiệt độ cao là ăn mòn hoá học.\n"
            "- (2) Cặp Fe - C trong thép tiếp xúc dung dịch acid H~2~SO~4~ xảy ra ăn mòn điện hoá học.\n"
            "- (3) Al đẩy Cu^2+^ tạo kim loại Cu bám vào lá Al, tạo cặp pin Al - Cu trong dung dịch chất điện li xảy ra ăn mòn điện hoá học.\n"
            "- (4) Dây Al nối dây Cu trong không khí ẩm tạo pin Al - Cu xảy ra ăn mòn điện hoá học.\n"
            "Các thí nghiệm xảy ra ăn mòn điện hoá học là: 2, 3, 4.\n"
            "Đáp số: 2, 3, 4."
        ),
        (
            4,
            "Để xảy ra ăn mòn điện hoá học thì cần đủ các yếu tố nào sau đây?\n"
            "(1) Phải có hai điện cực khác chất nhau (hai kim loại khác nhau hoặc kim loại với phi kim).\n"
            "(2) Hai điện cực phải tiếp xúc với nhau (trực tiếp hoặc gián tiếp qua dây dẫn).\n"
            "(3) Hai điện cực phải được nối qua vôn kế hoặc ampe kế.\n"
            "(4) Hai điện cực phải cùng tiếp xúc với một dung dịch chất điện li.\n"
            "Liệt kê các yếu tố cần thiết theo số thứ tự tăng dần (ví dụ: 1, 2, 4).",
            "1, 2, 4",
            "Điều kiện để xảy ra ăn mòn điện hoá học gồm 3 yếu tố đồng thời:\n"
            "(1) Các điện cực khác nhau về bản chất.\n"
            "(2) Các điện cực phải tiếp xúc trực tiếp hoặc gián tiếp qua dây dẫn.\n"
            "(4) Các điện cực cùng tiếp xúc với một dung dịch chất điện li.\n"
            "Yếu tố (3) không cần thiết vì dòng điện ăn mòn tự khép kín mạch ngoài mà không cần thiết bị đo.\n"
            "Đáp số: 1, 2, 4."
        ),
        (
            5,
            "Trong công nghiệp sản xuất gang ở nước ta hiện nay, muốn sản xuất ra 1 tấn gang cần phải sử dụng 1,7 đến 1,8 tấn quặng sắt, 0,6 đến 0,7 tấn đá vôi làm chất trợ dung, 0,6 đến 0,8 tấn than cốc. Giả thiết trong đá vôi CaCO~3~ chiếm 97% khối lượng và 90% than cốc chuyển hoá thành CO~2~. Trong điều kiện sản xuất như trên, khi sản xuất được 1 tấn gang, nhà máy đã thải ra môi trường tối thiểu bao nhiêu m^3^ khí CO~2~ ở điều kiện chuẩn? (Làm tròn kết quả đến hàng đơn vị).",
            "1260",
            "Để lượng khí CO~2~ thải ra tối thiểu, ta lấy lượng chất tiêu tốn ở mức thấp nhất:\n"
            "- Đá vôi tối thiểu: m(đá vôi) = 0,6 tấn = 600 kg.\n"
            "  => m(CaCO~3~) = 600 * 97% = 582 kg => n(CaCO~3~) = 582 / 100 = 5,82 kmol.\n"
            "  Nhiệt phân đá vôi tạo: n_1(CO~2~) = 5,82 kmol.\n"
            "- Than cốc tối thiểu: m(than cốc) = 0,6 tấn = 600 kg.\n"
            "  Lượng C chuyển thành CO~2~: n_2(CO~2~) = (600 * 90%) / 12 = 45 kmol.\n"
            "- Tổng lượng CO~2~ tối thiểu sinh ra:\n"
            "  n(CO~2~) = 5,82 + 45 = 50,82 kmol.\n"
            "- Thể tích khí CO~2~ ở điều kiện chuẩn (25 °C, 1 bar):\n"
            "  V(CO~2~) = 50,82 * 24,79 ≈ 1259,83 ≈ 1260 m^3^.\n"
            "Đáp số: 1260."
        ),
        (
            6,
            "Bằng phân tích quang phổ, người ta phát hiện trong “khí quyển” của sao Thuỷ có kim loại X. Biết X có cấu trúc mạng tinh thể lập phương tâm khối (độ đặc khít 68%) và bán kính nguyên tử là 0,230 nm. Khối lượng riêng của nguyên tố này là 0,862 g/cm^3^. (Biết Al = 27; K = 39; Zn = 65; Ba = 137). Xác định nguyên tử khối kim loại X.",
            "39",
            "- Bán kính nguyên tử: R = 0,230 nm = 2,30 * 10^-8^ cm.\n"
            "- Thể tích 1 nguyên tử hình cầu:\n"
            "  V_1 = (4/3) * π * R^3 = (4/3) * π * (2,30 * 10^-8^)^3 ≈ 5,0965 * 10^-23^ cm^3^.\n"
            "- Xét 1 mol nguyên tử X (chứa N_A = 6,022 * 10^23^ nguyên tử):\n"
            "  V_thực = V_1 * N_A = 5,0965 * 10^-23^ * 6,022 * 10^23^ ≈ 30,691 cm^3^.\n"
            "- Do độ đặc khít là 68%, thể tích của 1 mol tinh thể là:\n"
            "  V_tinh thể = V_thực / 0,68 = 30,691 / 0,68 ≈ 45,134 cm^3^.\n"
            "- Khối lượng 1 mol tinh thể X:\n"
            "  m = V_tinh thể * D = 45,134 * 0,862 ≈ 38,91 g ≈ 39 g.\n"
            "- Nguyên tử khối kim loại X là 39 (Potassium / Kali, K).\n"
            "Đáp số: 39."
        )
    ]

    for q_num, content, ans, sol in p3_data:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        add_chem_run(p_q, f"Câu {q_num}. ", bold=True)
        add_chem_run(p_q, content)

        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(1)
        p_ans.paragraph_format.space_after = Pt(3)
        p_ans.paragraph_format.line_spacing = 1.15
        if is_teacher:
            add_chem_run(p_ans, "Đáp số: ", bold=True)
            add_chem_run(p_ans, ans, bold=True, color=COLOR_CORAL_RED)
            add_loigiai(doc, sol)
        else:
            add_chem_run(p_ans, "Trả lời: ...........................................................................................................................................................", italic=True)

    # Summary table for teacher version
    if is_teacher:
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(12)
        p_title.paragraph_format.space_after = Pt(4)
        r_sum = p_title.add_run("BẢNG TỔNG HỢP ĐÁP ÁN ĐỀ SỐ 17")
        r_sum.font.name = "Times New Roman"
        r_sum.font.size = Pt(12)
        r_sum.font.bold = True
        r_sum.font.color.rgb = COLOR_DARK_GREEN

        # Table 1: Part I
        p_p1 = doc.add_paragraph()
        p_p1.paragraph_format.space_before = Pt(4)
        p_p1.paragraph_format.space_after = Pt(2)
        r_p1 = p_p1.add_run("PHẦN I: TRẮC NGHIỆM NHIỀU PHƯƠNG ÁN LỰA CHỌN")
        r_p1.font.name = "Times New Roman"
        r_p1.font.size = Pt(10.5)
        r_p1.font.bold = True
        r_p1.font.color.rgb = COLOR_DARK_GREEN

        tbl1 = doc.add_table(rows=2, cols=18)
        tbl1.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr1 = tbl1._tbl.tblPr
        tblPr1.append(parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            f'<w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            f'</w:tblBorders>'
        ))

        for j in range(18):
            c_h = tbl1.rows[0].cells[j]
            c_h.width = Inches(0.39)
            set_cell_margins(c_h, 40, 40, 20, 20)
            p = c_h.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(j + 1))
            r.font.name = "Times New Roman"
            r.font.size = Pt(9.5)
            r.font.bold = True

            c_a = tbl1.rows[1].cells[j]
            c_a.width = Inches(0.39)
            set_cell_margins(c_a, 40, 40, 20, 20)
            p2 = c_a.paragraphs[0]
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_before = Pt(0)
            p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run(p1_data[j][3])
            r2.font.name = "Times New Roman"
            r2.font.size = Pt(9.5)
            r2.font.bold = True
            r2.font.color.rgb = COLOR_CORAL_RED

        # Table 2: Part II
        p_p2 = doc.add_paragraph()
        p_p2.paragraph_format.space_before = Pt(6)
        p_p2.paragraph_format.space_after = Pt(2)
        r_p2 = p_p2.add_run("PHẦN II: TRẮC NGHIỆM ĐÚNG SAI")
        r_p2.font.name = "Times New Roman"
        r_p2.font.size = Pt(10.5)
        r_p2.font.bold = True
        r_p2.font.color.rgb = COLOR_DARK_GREEN

        tbl2 = doc.add_table(rows=5, cols=5)
        tbl2.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr2 = tbl2._tbl.tblPr
        tblPr2.append(parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            f'<w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            f'</w:tblBorders>'
        ))

        headers_p2 = ["Câu", "Lệnh hỏi a", "Lệnh hỏi b", "Lệnh hỏi c", "Lệnh hỏi d"]
        for j, h in enumerate(headers_p2):
            cell = tbl2.rows[0].cells[j]
            cell.width = Inches(1.4)
            set_cell_margins(cell, 40, 40, 40, 40)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(h)
            r.font.name = "Times New Roman"
            r.font.size = Pt(9.5)
            r.font.bold = True

        for i, (q_num, _, sub_items, _, _) in enumerate(p2_data):
            c_q = tbl2.rows[i + 1].cells[0]
            c_q.width = Inches(1.4)
            set_cell_margins(c_q, 40, 40, 40, 40)
            p = c_q.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(q_num))
            r.font.name = "Times New Roman"
            r.font.size = Pt(9.5)
            r.font.bold = True

            for j, (_, _, is_true) in enumerate(sub_items):
                cell = tbl2.rows[i + 1].cells[j + 1]
                cell.width = Inches(1.4)
                set_cell_margins(cell, 40, 40, 40, 40)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                txt = "Đ" if is_true else "S"
                r = p.add_run(txt)
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.5)
                r.font.bold = True
                if is_true:
                    r.font.color.rgb = COLOR_CORAL_RED

        # Table 3: Part III
        p_p3 = doc.add_paragraph()
        p_p3.paragraph_format.space_before = Pt(6)
        p_p3.paragraph_format.space_after = Pt(2)
        r_p3 = p_p3.add_run("PHẦN III: TRẢ LỜI NGẮN")
        r_p3.font.name = "Times New Roman"
        r_p3.font.size = Pt(10.5)
        r_p3.font.bold = True
        r_p3.font.color.rgb = COLOR_DARK_GREEN

        tbl3 = doc.add_table(rows=2, cols=6)
        tbl3.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr3 = tbl3._tbl.tblPr
        tblPr3.append(parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            f'<w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            f'</w:tblBorders>'
        ))

        for j in range(6):
            c_h = tbl3.rows[0].cells[j]
            c_h.width = Inches(1.18)
            set_cell_margins(c_h, 40, 40, 20, 20)
            p = c_h.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(f"Câu {j + 1}")
            r.font.name = "Times New Roman"
            r.font.size = Pt(9.5)
            r.font.bold = True

            c_a = tbl3.rows[1].cells[j]
            c_a.width = Inches(1.18)
            set_cell_margins(c_a, 40, 40, 20, 20)
            p2 = c_a.paragraphs[0]
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_before = Pt(0)
            p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run(p3_data[j][2])
            r2.font.name = "Times New Roman"
            r2.font.size = Pt(9.5)
            r2.font.bold = True
            r2.font.color.rgb = COLOR_CORAL_RED

    filename = "De_Kiem_Tra_Hoa_12_De17_GiaoVien.docx" if is_teacher else "De_Kiem_Tra_Hoa_12_De17.docx"
    filepath = os.path.join(OUT_DIR, filename)
    doc.save(filepath)
    print(f"Saved: {filepath}")

if __name__ == "__main__":
    generate_exam_docx(is_teacher=False)
    generate_exam_docx(is_teacher=True)
