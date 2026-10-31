# -*- coding: utf-8 -*-
import sys
import docx

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

for name in ['De_Kiem_Tra_GK1_Hoa_12_De20_GiaoVien.docx', 'De_Kiem_Tra_GK1_Hoa_12_De20_HocSinh.docx']:
    f = rf"c:\Antigravity_Thanh\San_Pham\{name}"
    doc = docx.Document(f)
    print(f"=== {name} ===")
    for i, p in enumerate(doc.paragraphs):
        if "Câu 13." in p.text:
            print(f"[{i}] {p.text}")
            print(f"[{i+1}] {doc.paragraphs[i+1].text}")
            if "GiaoVien" in name:
                print(f"[{i+2}] {doc.paragraphs[i+2].text}")
            break
