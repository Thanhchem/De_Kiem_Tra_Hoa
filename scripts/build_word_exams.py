import os
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

CROP_DIR = r"c:\Antigravity_Thanh\images_crop"
OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"

COLOR_PRIMARY = RGBColor(46, 125, 50)       # Dark Teal / Green (Thầy Thạnh style)
COLOR_CORRECT = RGBColor(198, 40, 40)       # Red for correct answer
COLOR_SOLUTION = RGBColor(21, 101, 192)     # Blue for explanation
COLOR_GRAY = RGBColor(97, 97, 97)

def add_p_border_bottom(p, color="2E7D32", sz="6"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_p_border_top(p, color="2E7D32", sz="6"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_page_number_fields(run):
    r = run._r
    # Page
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r.append(fld1); r.append(instr1); r.append(fld2); r.append(fld3)

def set_run_font(run, name='Times New Roman', size_pt=12, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color

def setup_header_footer(doc, is_teacher=False):
    section = doc.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    # Header
    header = section.header
    hp = header.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
    hp.paragraph_format.space_after = Pt(2)

    title_left = "Tài liệu lưu hành nội bộ  |  THPT Hai Bà Trưng -- Mã đề: 209" if not is_teacher else "Tài liệu lưu hành nội bộ  |  THPT Hai Bà Trưng -- Lời giải"
    r1 = hp.add_run(title_left)
    set_run_font(r1, 'Times New Roman', 9.5, italic=True, color=COLOR_PRIMARY)

    hp.add_run('\t')

    r2 = hp.add_run("Thầy TRẦN VĂN THẠNH — 0777.470.803")
    set_run_font(r2, 'Times New Roman', 9.5, bold=True, italic=True, color=COLOR_PRIMARY)
    add_p_border_bottom(hp, color="2E7D32", sz="6")

    # Footer
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
    fp.paragraph_format.space_before = Pt(4)
    fp.paragraph_format.space_after = Pt(0)

    fr1 = fp.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế  |  CS2: 24 Đặng Thái Thân, TP Huế\nLớp Hóa Thầy Thạnh — SĐT: 0777.470.803")
    set_run_font(fr1, 'Times New Roman', 8.5, italic=True, color=COLOR_GRAY)

    fp.add_run('\t')
    fr2 = fp.add_run("Trang ")
    set_run_font(fr2, 'Times New Roman', 9.0, bold=True, color=COLOR_GRAY)
    add_page_number_fields(fr2)

    add_p_border_top(fp, color="2E7D32", sz="6")

def add_exam_header_block(doc, is_teacher=False):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Remove borders
    tbl_pr = tbl._tbl.tblPr
    tblBorders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/><w:insideH w:val="none"/><w:insideV w:val="none"/></w:tblBorders>')
    tbl_pr.append(tblBorders)

    c0 = tbl.cell(0, 0)
    c0.width = Cm(7.5)
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_after = Pt(2)
    p0.paragraph_format.line_spacing = 1.15
    r = p0.add_run("SỞ GIÁO DỤC & ĐÀO TẠO TP HUẾ\nTRƯỜNG THPT HAI BÀ TRƯNG\n")
    set_run_font(r, 'Times New Roman', 11, bold=True)
    r_code = p0.add_run("Mã đề thi: 209")
    set_run_font(r_code, 'Times New Roman', 11, bold=True, color=COLOR_PRIMARY)

    c1 = tbl.cell(0, 1)
    c1.width = Cm(10.5)
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_after = Pt(2)
    p1.paragraph_format.line_spacing = 1.15
    r = p1.add_run("ĐỀ KIỂM TRA CUỐI KÌ II - NĂM HỌC 2024-2025\nMÔN HÓA HỌC LỚP 11\n")
    set_run_font(r, 'Times New Roman', 11, bold=True)
    if is_teacher:
        r_sub = p1.add_run("(HƯỚNG DẪN CHẤM VÀ LỜI GIẢI CHI TIẾT)")
        set_run_font(r_sub, 'Times New Roman', 10.5, bold=True, color=COLOR_CORRECT)
    else:
        r_sub = p1.add_run("Thời gian làm bài: 45 phút (không kể thời gian phát đề)")
        set_run_font(r_sub, 'Times New Roman', 10.5, italic=True)

    # Info line
    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(4)
    p_info.paragraph_format.space_after = Pt(2)
    p_info.paragraph_format.line_spacing = 1.15
    if not is_teacher:
        r = p_info.add_run("Họ, tên học sinh: ................................................................ Lớp: ........................")
        set_run_font(r, 'Times New Roman', 11)

    p_ntk = doc.add_paragraph()
    p_ntk.paragraph_format.space_before = Pt(0)
    p_ntk.paragraph_format.space_after = Pt(4)
    r = p_ntk.add_run("Cho NTK: H = 1; C = 12; O = 16; N = 14; Ag = 108; Na = 23; Br = 80; Cl = 35,5.")
    set_run_font(r, 'Times New Roman', 10.5, italic=True)

def add_section_title(doc, title_text, score_text=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"■ {title_text}")
    set_run_font(r, 'Times New Roman', 12, bold=True, color=COLOR_PRIMARY)
    if score_text:
        r_s = p.add_run(f" {score_text}")
        set_run_font(r_s, 'Times New Roman', 11, bold=True, italic=True)

def add_instruction_line(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, 'Times New Roman', 11, italic=True)

def add_question_prompt(doc, q_num, prompt_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_label = p.add_run(f"Câu {q_num}. ")
    set_run_font(r_label, 'Times New Roman', 12, bold=True)
    r_txt = p.add_run(prompt_text)
    set_run_font(r_txt, 'Times New Roman', 12)
    return p

def add_choices_tabbed(doc, choices, correct_idx=None, is_teacher=False, layout_mode="4cols"):
    """
    choices: list of 4 strings [A_text, B_text, C_text, D_text]
    correct_idx: 0 for A, 1 for B, 2 for C, 3 for D
    layout_mode: "4cols", "2cols", or "1col"
    NO TABLES USED! Only paragraph tab stops!
    """
    labels = ["A. ", "B. ", "C. ", "D. "]

    if layout_mode == "4cols":
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(4.8), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(13.6), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15

        for i in range(4):
            p.add_run('\t')
            is_corr = (is_teacher and i == correct_idx)
            r_lbl = p.add_run(labels[i])
            set_run_font(r_lbl, 'Times New Roman', 12, bold=True, color=COLOR_CORRECT if is_corr else None)
            r_txt = p.add_run(choices[i] + ("  " if i < 3 else ""))
            set_run_font(r_txt, 'Times New Roman', 12, bold=is_corr, color=COLOR_CORRECT if is_corr else None)

    elif layout_mode == "2cols":
        # Line 1: A and B
        p1 = doc.add_paragraph()
        p1.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p1.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.line_spacing = 1.15

        for i in [0, 1]:
            p1.add_run('\t')
            is_corr = (is_teacher and i == correct_idx)
            r_lbl = p1.add_run(labels[i])
            set_run_font(r_lbl, 'Times New Roman', 12, bold=True, color=COLOR_CORRECT if is_corr else None)
            r_txt = p1.add_run(choices[i])
            set_run_font(r_txt, 'Times New Roman', 12, bold=is_corr, color=COLOR_CORRECT if is_corr else None)

        # Line 2: C and D
        p2 = doc.add_paragraph()
        p2.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p2.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.line_spacing = 1.15

        for i in [2, 3]:
            p2.add_run('\t')
            is_corr = (is_teacher and i == correct_idx)
            r_lbl = p2.add_run(labels[i])
            set_run_font(r_lbl, 'Times New Roman', 12, bold=True, color=COLOR_CORRECT if is_corr else None)
            r_txt = p2.add_run(choices[i])
            set_run_font(r_txt, 'Times New Roman', 12, bold=is_corr, color=COLOR_CORRECT if is_corr else None)

    else: # 1col
        for i in range(4):
            p = doc.add_paragraph()
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1 if i < 3 else 2)
            p.paragraph_format.line_spacing = 1.15

            p.add_run('\t')
            is_corr = (is_teacher and i == correct_idx)
            r_lbl = p.add_run(labels[i])
            set_run_font(r_lbl, 'Times New Roman', 12, bold=True, color=COLOR_CORRECT if is_corr else None)
            r_txt = p.add_run(choices[i])
            set_run_font(r_txt, 'Times New Roman', 12, bold=is_corr, color=COLOR_CORRECT if is_corr else None)

def add_solution_block(doc, solution_text, choose_text=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.5)

    r_lbl = p.add_run("Lời giải: ")
    set_run_font(r_lbl, 'Times New Roman', 11.5, bold=True, color=COLOR_SOLUTION)

    r_txt = p.add_run(solution_text)
    set_run_font(r_txt, 'Times New Roman', 11.5)

    if choose_text:
        r_c = p.add_run(f" ➔ {choose_text}")
        set_run_font(r_c, 'Times New Roman', 11.5, bold=True, color=COLOR_CORRECT)

def add_image_centered(doc, image_name, width_cm=6.0):
    img_path = os.path.join(CROP_DIR, image_name)
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        run.add_picture(img_path, width=Cm(width_cm))

print("Helper functions defined successfully!")
