# -*- coding: utf-8 -*-
"""
Script tạo file Word (.docx) chuẩn "FORM ĐỀ KIỂM TRA" cho:
ĐỀ KIỂM TRA SỐ 13 - CHƯƠNG 5: PIN ĐIỆN – ĐIỆN PHÂN (HÓA HỌC 12)
Gồm 2 phiên bản:
1. De_Kiem_Tra_Hoa_12_De13.docx (Bản Học Sinh)
2. De_Kiem_Tra_Hoa_12_De13_GiaoVien.docx (Bản Giáo Viên có lời giải chi tiết)
Tuân thủ 100% quy chuẩn: Header & Footer, Tab Stops (NO TABLE trong đề/lời giải), Sub/Superscript công thức hóa học.
"""

import sys
sys.stdout.reconfigure(encoding='utf-8')
import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
IMG_DIR = r"c:\Antigravity_Thanh\San_Pham\images_crop_de13"

COLOR_HF_RED = RGBColor(211, 47, 47)      # Coral Red (#D32F2F)
COLOR_HF_GREEN_HEX = "8BC390"             # Sage Green (#8BC390)
COLOR_PRIMARY = RGBColor(46, 125, 50)     # Dark Green (#2E7D32)
COLOR_CORRECT = RGBColor(198, 40, 40)     # Red for correct answers (#C62828)
COLOR_SOLUTION = RGBColor(21, 101, 192)   # Blue for solutions (#1565C0)
COLOR_GRAY = RGBColor(97, 97, 97)

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def set_run_style(run, font_name="Times New Roman", font_size=12, bold=False, italic=False, subscript=False, superscript=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def add_chem_run(p, text, base_font="Times New Roman", base_size=12, bold=False, italic=False, color=None):
    tokens = re.split(r'(\^\{[^}]+\}|\^[0-9\+\-]+|~\{[^}]+\}|~[0-9]+)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('^'):
            content = tok[2:-1] if tok.startswith('^{') else tok[1:]
            r = p.add_run(content)
            set_run_style(r, font_name=base_font, font_size=base_size, bold=bold, italic=italic, superscript=True, color=color)
        elif tok.startswith('~'):
            content = tok[2:-1] if tok.startswith('~{') else tok[1:]
            r = p.add_run(content)
            set_run_style(r, font_name=base_font, font_size=base_size, bold=bold, italic=italic, subscript=True, color=color)
        else:
            r = p.add_run(tok)
            set_run_style(r, font_name=base_font, font_size=base_size, bold=bold, italic=italic, color=color)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(r'<w:tcBorders %s/>' % nsdecls('w'))
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = parse_xml(r'<%s %s w:val="%s" w:sz="%s" w:space="0" w:color="%s"/>' %
                                (tag, nsdecls('w'), edge_data.get('val', 'single'),
                                 edge_data.get('sz', '4'), edge_data.get('color', 'auto')))
            tcBorders.append(element)
        else:
            tag = 'w:{}'.format(edge)
            element = parse_xml(r'<%s %s w:val="none"/>' % (tag, nsdecls('w')))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def set_cell_margins(cell, top=50, bottom=50, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(r'<w:tcMar %s><w:top w:w="%d" w:type="dxa"/><w:bottom w:w="%d" w:type="dxa"/><w:left w:w="%d" w:type="dxa"/><w:right w:w="%d" w:type="dxa"/></w:tcMar>' % (nsdecls('w'), top, bottom, left, right))
    tcPr.append(tcMar)

def setup_page_and_header_footer(doc):
    for sec in doc.sections:
        sec.top_margin = Cm(1.5)
        sec.bottom_margin = Cm(1.5)
        sec.left_margin = Cm(1.5)
        sec.right_margin = Cm(1.5)
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.header_distance = Cm(0.6)
        sec.footer_distance = Cm(0.6)
        
        # --- HEADER ---
        hdr = sec.header
        hdr_p = hdr.paragraphs[0]
        hdr_p.text = ""
        hdr_p.paragraph_format.space_before = Pt(0)
        hdr_p.paragraph_format.space_after = Pt(0)
        
        tbl_hdr = hdr.add_table(rows=1, cols=2, width=Cm(18.0))
        tbl_hdr.autofit = False
        tbl_hdr.columns[0].width = Cm(9.0)
        tbl_hdr.columns[1].width = Cm(9.0)
        
        c0 = tbl_hdr.cell(0, 0)
        c1 = tbl_hdr.cell(0, 1)
        c0.width = Cm(9.0)
        c1.width = Cm(9.0)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run("Tài liệu lưu hành nội bộ")
        set_run_style(r0, font_name="Times New Roman", font_size=10, italic=True, color=COLOR_HF_RED)
        
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(2)
        r1a = p1.add_run("Thầy ")
        set_run_style(r1a, font_name="Times New Roman", font_size=10, color=COLOR_HF_RED)
        r1b = p1.add_run("TRẦN VĂN THẠNH")
        set_run_style(r1b, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED)
        r1c = p1.add_run(" - 0777.470.803")
        set_run_style(r1c, font_name="Times New Roman", font_size=10, color=COLOR_HF_RED)
        
        set_cell_border(c0, bottom={'val': 'single', 'sz': '6', 'color': COLOR_HF_GREEN_HEX})
        set_cell_border(c1, bottom={'val': 'single', 'sz': '6', 'color': COLOR_HF_GREEN_HEX})
        set_cell_margins(c0, top=0, bottom=40, left=0, right=0)
        set_cell_margins(c1, top=0, bottom=40, left=0, right=0)
        
        # --- FOOTER ---
        ftr = sec.footer
        ftr_p = ftr.paragraphs[0]
        ftr_p.text = ""
        ftr_p.paragraph_format.space_before = Pt(0)
        ftr_p.paragraph_format.space_after = Pt(0)
        
        tbl_ftr = ftr.add_table(rows=1, cols=2, width=Cm(18.0))
        tbl_ftr.autofit = False
        tbl_ftr.columns[0].width = Cm(13.5)
        tbl_ftr.columns[1].width = Cm(4.5)
        
        fc0 = tbl_ftr.cell(0, 0)
        fc1 = tbl_ftr.cell(0, 1)
        fc0.width = Cm(13.5)
        fc1.width = Cm(4.5)
        
        fp0 = fc0.paragraphs[0]
        fp0.paragraph_format.space_before = Pt(2)
        fp0.paragraph_format.space_after = Pt(0)
        fr0 = fp0.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
        set_run_style(fr0, font_name="Times New Roman", font_size=9.5, italic=True, color=COLOR_HF_RED)
        
        fp1 = fc1.paragraphs[0]
        fp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp1.paragraph_format.space_before = Pt(2)
        fp1.paragraph_format.space_after = Pt(0)
        add_page_number_field(fp1, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED)
        
        set_cell_border(fc0, top={'val': 'single', 'sz': '6', 'color': COLOR_HF_GREEN_HEX})
        set_cell_border(fc1, top={'val': 'single', 'sz': '6', 'color': COLOR_HF_GREEN_HEX})
        set_cell_margins(fc0, top=40, bottom=0, left=0, right=0)
        set_cell_margins(fc1, top=40, bottom=0, left=0, right=0)

def add_exam_title_block(doc, is_teacher=False):
    lines_info = [
        ("SỞ GD&ĐT THỪA THIÊN HUẾ", "ĐỀ KIỂM TRA CHƯƠNG 5: PIN ĐIỆN – ĐIỆN PHÂN", True, True),
        ("TRƯỜNG THPT HAI BÀ TRƯNG", "MÔN: HÓA HỌC 12 - NĂM HỌC 2026 - 2027", False, True),
        ("ĐỀ CHÍNH THỨC", "Thời gian làm bài: 50 phút (không kể thời gian phát đề)", True, False, False, True),
        ("(Đề thi gồm 04 trang)", "MÃ ĐỀ THI: 113", False, True, True, False)
    ]
    
    for item in lines_info:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
        
        r_l = p.add_run(item[0])
        set_run_style(r_l, font_name="Times New Roman", font_size=11, bold=item[2], italic=item[4] if len(item)>4 else False)
        
        p.add_run("\t")
        
        r_r = p.add_run(item[1])
        set_run_style(r_r, font_name="Times New Roman", font_size=11, bold=item[3], italic=item[5] if len(item)>5 else False)

    if is_teacher:
        p_note = doc.add_paragraph()
        p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_note.paragraph_format.space_before = Pt(2)
        p_note.paragraph_format.space_after = Pt(2)
        r_note = p_note.add_run("(BẢN GIÁO VIÊN -- CÓ LỜI GIẢI CHI TIẾT)")
        set_run_style(r_note, font_name="Times New Roman", font_size=11.5, bold=True, color=COLOR_HF_RED)

    # Student info line
    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(3)
    p_info.paragraph_format.space_after = Pt(1)
    r_info = p_info.add_run("Họ, tên thí sinh: .......................................................................... Lớp: ............. SBD: .............................")
    set_run_style(r_info, font_name="Times New Roman", font_size=11, italic=True)

    # Atomic mass note
    p_mass = doc.add_paragraph()
    p_mass.paragraph_format.space_before = Pt(0)
    p_mass.paragraph_format.space_after = Pt(4)
    r_mass = p_mass.add_run("(Cho biết nguyên tử khối: H = 1; C = 12; N = 14; O = 16; Na = 23; Mg = 24; Al = 27; S = 32; Cl = 35,5; K = 39; Ca = 40; Fe = 56; Ni = 59; Cu = 64; Zn = 65; Ag = 108; Sn = 119; Pb = 207; hằng số Faraday F = 96500 C/mol).")
    set_run_style(r_mass, font_name="Times New Roman", font_size=9.5, italic=True, color=COLOR_GRAY)

def add_part_heading(doc, part_title, part_instruction):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    r_title = p.add_run(part_title)
    set_run_style(r_title, font_name="Times New Roman", font_size=11.5, bold=True, color=COLOR_PRIMARY)
    
    if part_instruction:
        r_inst = p.add_run(" " + part_instruction)
        set_run_style(r_inst, font_name="Times New Roman", font_size=11, italic=True)

def add_solution_paragraph(doc, solution_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.6)
    
    r_label = p.add_run("Lời giải chi tiết: ")
    set_run_style(r_label, font_name="Times New Roman", font_size=10.5, bold=True, color=COLOR_SOLUTION)
    
    add_chem_run(p, solution_text, base_font="Times New Roman", base_size=10.5, italic=True)

# ----------------- DỮ LIỆU ĐỀ KIỂM TRA SỐ 13 -----------------

MCQ_QUESTIONS = [
    {
        "num": 1,
        "text": "(SBT – CĐ). Sức điện động chuẩn của pin Galvani được tính như thế nào?",
        "options": [
            "A. Bằng hiệu của thế điện cực chuẩn tương ứng của điện cực dương và điện cực âm.",
            "B. Bằng tổng của thế điện cực chuẩn tương ứng của điện cực dương và điện cực âm.",
            "C. Bằng tích của thế điện cực chuẩn tương ứng của điện cực dương và điện cực âm.",
            "D. Bằng thương của thế điện cực chuẩn tương ứng của điện cực dương và điện cực âm."
        ],
        "correct": "A",
        "sol": "Chọn A. Sức điện động chuẩn của pin điện hoá (pin Galvani) bằng hiệu giữa thế điện cực chuẩn của cực dương (cathode) và cực âm (anode): E^o^_pin = E^o^_(+) - E^o^_(-) = E^o^_cathode - E^o^_anode."
    },
    {
        "num": 2,
        "text": "(SBT – KNTT). Trong pin điện hoá Zn - Cu, phản ứng hoá học xảy ra giữa hai dạng nào của các cặp oxi hoá - khử tương ứng?",
        "options": [
            "A. Zn và Cu^2+^.",
            "B. Zn và Cu.",
            "C. Zn^2+^ và Cu^2+^.",
            "D. Pb và Cu^2+^."
        ],
        "correct": "A",
        "sol": "Chọn A. Trong pin điện hoá Zn - Cu, chất khử mạnh (Zn) tác dụng với chất oxi hoá mạnh (Cu^2+^) theo phản ứng oxi hóa - khử tự diễn biến: Zn + Cu^2+^ -> Zn^2+^ + Cu."
    },
    {
        "num": 3,
        "text": "(SBT – KNTT). Trong quá trình hoạt động của pin điện Zn - Cu, dòng electron di chuyển từ",
        "options": [
            "A. cực kẽm sang cực đồng.",
            "B. cực bên phải sang cực bên trái.",
            "C. cathode sang anode.",
            "D. cực dương sang cực âm."
        ],
        "correct": "A",
        "sol": "Chọn A. Ở cực âm (anode, Zn), Zn bị oxi hoá giải phóng electron: Zn -> Zn^2+^ + 2e. Các electron này di chuyển qua dây dẫn mạch ngoài sang cực dương (cathode, Cu) để khử Cu^2+^. Vậy electron chạy từ cực kẽm sang cực đồng (từ anode sang cathode)."
    },
    {
        "num": 4,
        "text": "(SBT – KNTT). Phản ứng hoá học chính xảy ra trong quá trình điện phân nóng chảy Al~2~O~3~ trong cryolite (3NaF·AlF~3~) là",
        "options": [
            "A. 2AlF~3~ -> 2Al + 3F~2~.",
            "B. 2NaF -> 2Na + F~2~.",
            "C. 2H~2~O -> 2H~2~ + O~2~.",
            "D. 2Al~2~O~3~ -> 4Al + 3O~2~."
        ],
        "correct": "D",
        "sol": "Chọn D. Trong công nghiệp, nhôm được sản xuất bằng phương pháp điện phân nóng chảy Al~2~O~3~ với sự có mặt của cryolite. Phản ứng chính là: 2Al~2~O~3~ -> 4Al + 3O~2~."
    },
    {
        "num": 5,
        "text": "(SBT – KNTT). Việc duy trì điện áp thấp (~ 5 V) trong quá trình điện phân nóng chảy Al~2~O~3~ trong cryolite (3NaF·AlF~3~) nhằm ngăn cản quá trình nào sau đây xảy ra ở cathode?",
        "options": [
            "A. Al^3+^ + 3e -> Al.",
            "B. Na^+^ + 1e -> Na.",
            "C. F~2~ + 2e -> 2F^-^.",
            "D. O~2~ + 4e -> 2O^2-^."
        ],
        "correct": "B",
        "sol": "Chọn B. Cặp Na^+^/Na có thế điện cực chuẩn âm hơn Al^3+^/Al (E^o^ = -2,71 V so với -1,66 V). Việc duy trì điện áp thấp nhằm đảm bảo chỉ có ion Al^3+^ bị khử tại cathode mà không khử ion Na^+^ thành Na tự do."
    },
    {
        "num": 6,
        "text": "(SBT – KNTT). Trong công nghiệp, phương pháp điện phân dung dịch được sử dụng để sản xuất một lượng đáng kể kim loại nào sau đây?",
        "options": ["A. Zn.", "B. Al.", "C. Fe.", "D. Mg."],
        "correct": "A",
        "sol": "Chọn A. Trong công nghiệp, kẽm (Zn) được sản xuất chủ yếu bằng phương pháp điện phân dung dịch ZnSO~4~. Al và Mg là các kim loại có tính khử mạnh được điều chế bằng điện phân nóng chảy; Fe được sản xuất bằng phương pháp nhiệt luyện trong lò cao."
    },
    {
        "num": 7,
        "text": "(SBT – CTST). Dãy kim loại sắp xếp theo chiều tăng dần tính khử từ trái sang phải là",
        "options": [
            "A. Pb, Ni, Sn, Zn.",
            "B. Pb, Sn, Ni, Zn.",
            "C. Ni, Sn, Zn, Pb.",
            "D. Ni, Zn, Pb, Sn."
        ],
        "correct": "B",
        "sol": "Chọn B. Căn cứ vào thế điện cực chuẩn: E^o^_Zn^2+^/Zn (-0,76 V) < E^o^_Ni^2+^/Ni (-0,26 V) < E^o^_Sn^2+^/Sn (-0,14 V) < E^o^_Pb^2+^/Pb (-0,13 V). Thế điện cực càng âm thì tính khử của kim loại càng mạnh, do đó chiều tăng dần tính khử là: Pb < Sn < Ni < Zn."
    },
    {
        "num": 8,
        "text": "(SBT – CTST). Dãy cation kim loại được sắp xếp theo chiều tăng dần tính oxi hóa từ trái sang phải là",
        "options": [
            "A. Cu^2+^, Fe^2+^, Mg^2+^.",
            "B. Mg^2+^, Fe^2+^, Cu^2+^.",
            "C. Mg^2+^, Cu^2+^, Fe^2+^.",
            "D. Cu^2+^, Mg^2+^, Fe^2+^."
        ],
        "correct": "B",
        "sol": "Chọn B. Thế điện cực chuẩn của các cặp: E^o^_Mg^2+^/Mg (-2,37 V) < E^o^_Fe^2+^/Fe (-0,44 V) < E^o^_Cu^2+^/Cu (+0,34 V). Thế điện cực càng dương thì tính oxi hóa của cation càng mạnh. Chiều tăng dần tính oxi hoá là: Mg^2+^ < Fe^2+^ < Cu^2+^."
    },
    {
        "num": 9,
        "text": "(SBT – KNTT). Phản ứng hoá học xảy ra trong pin điện hoá Sn - Cu: Sn + Cu^2+^ -> Sn^2+^ + Cu. Trong quá trình hoạt động của pin điện hoá, nhận định nào sau đây là đúng?",
        "options": [
            "A. Khối lượng của điện cực Sn tăng.",
            "B. Nồng độ Sn^2+^ trong dung dịch tăng.",
            "C. Khối lượng của điện cực Cu giảm.",
            "D. Nồng độ Cu^2+^ trong dung dịch tăng."
        ],
        "correct": "B",
        "sol": "Chọn B. Ở anode, Sn bị oxi hóa: Sn -> Sn^2+^ + 2e (khối lượng điện cực Sn giảm, nồng độ ion Sn^2+^ tăng lên). Ở cathode, Cu^2+^ bị khử: Cu^2+^ + 2e -> Cu (khối lượng điện cực Cu tăng, nồng độ Cu^2+^ giảm)."
    },
    {
        "num": 10,
        "text": "(SBT – KNTT). Một pin điện hoá có điện cực Zn nhúng trong dung dịch ZnSO~4~ và điện cực Cu nhúng trong dung dịch CuSO~4~. Sau một thời gian pin đó phóng điện thì",
        "options": [
            "A. khối lượng điện cực Zn giảm còn khối lượng điện cực Cu tăng.",
            "B. khối lượng điện cực Zn tăng còn khối lượng điện cực Cu giảm.",
            "C. khối lượng cả hai điện cực Zn và Cu đều tăng.",
            "D. khối lượng cả hai điện cực Zn và Cu đều giảm."
        ],
        "correct": "A",
        "sol": "Chọn A. Cực Zn là cực âm (anode) bị tan ra (Zn -> Zn^2+^ + 2e) nên khối lượng giảm; cực Cu là cực dương (cathode) có Cu bám vào (Cu^2+^ + 2e -> Cu) nên khối lượng tăng."
    },
    {
        "num": 11,
        "text": "(SBT – KNTT). Khi điện phân dung dịch gồm CuSO~4~ và HCl (sử dụng điện cực trơ, có màng ngăn xốp), chất khí đầu tiên thoát ra ở anode là",
        "options": ["A. O~2~.", "B. Cl~2~.", "C. H~2~.", "D. SO~2~."],
        "correct": "B",
        "sol": "Chọn B. Ở anode (cực dương), ion Cl^-^ bị oxi hoá trước H~2~O theo quá trình: 2Cl^-^ -> Cl~2~ + 2e. Vì vậy chất khí đầu tiên thoát ra ở anode là Cl~2~."
    },
    {
        "num": 12,
        "text": "(SBT – CTST). Cách nào sau đây không được dùng để điều chế NaOH?",
        "options": [
            "A. Điện phân dung dịch NaCl có màng ngăn xốp, điện cực trơ.",
            "B. Cho dung dịch Ca(OH)~2~ tác dụng với dung dịch Na~2~CO~3~.",
            "C. Cho Na~2~O tác dụng với nước.",
            "D. Dẫn khí NH~3~ vào dung dịch Na~2~CO~3~."
        ],
        "correct": "D",
        "sol": "Chọn D. NH~3~ là base yếu hơn NaOH, dung dịch NH~3~ không thể phản ứng với muối Na~2~CO~3~ để giải phóng base mạnh NaOH. Các cách A, B, C đều điều chế được NaOH."
    },
    {
        "num": 13,
        "text": "(SBT – CTST). Khi điện phân NaCl nóng chảy (điện cực trơ), ở cathode xảy ra",
        "options": [
            "A. sự khử ion Cl^-^.",
            "B. sự oxi hóa ion Cl^-^.",
            "C. sự oxi hóa ion Na^+^.",
            "D. sự khử ion Na^+^."
        ],
        "correct": "D",
        "sol": "Chọn D. Tại cathode (cực âm), ion Na^+^ nhận electron chuyển thành Na kim loại: Na^+^ + 1e -> Na. Quá trình nhận electron là quá trình khử (sự khử) ion Na^+^."
    },
    {
        "num": 14,
        "text": "(SBT – CTST). Khi điện phân CaCl~2~ nóng chảy (điện cực trơ), ở anode xảy ra",
        "options": [
            "A. sự khử ion Cl^-^.",
            "B. sự khử ion Ca^2+^.",
            "C. sự oxi hóa ion Ca^2+^.",
            "D. sự oxi hóa ion Cl^-^."
        ],
        "correct": "D",
        "sol": "Chọn D. Tại anode (cực dương), ion Cl^-^ nhường electron sinh ra khí Cl~2~: 2Cl^-^ -> Cl~2~ + 2e. Quá trình nhường electron là quá trình oxi hoá (sự oxi hoá) ion Cl^-^."
    },
    {
        "num": 15,
        "text": "Một học sinh thực hiện ba thí nghiệm liên quan đến ba kim loại X, Y, Z và dung dịch muối nitrate của chúng ở điều kiện chuẩn và quan sát được các hiện tượng sau:\n(1) X không phản ứng được với dung dịch Y(NO~3~)~2~ 1 M.\n(2) Y tan trong dung dịch muối ZNO~3~ 1 M và xuất hiện tinh thể Z.\n(3) Z không phản ứng với dung dịch X(NO~3~)~2~ 1 M.\nTrật tự nào sau đây thể hiện đúng mức độ khử của ba kim loại?",
        "options": ["A. X > Y > Z.", "B. Y > X > Z.", "C. X > Z > Y.", "D. Y > Z > X."],
        "correct": "B",
        "sol": "Chọn B.\n• Từ (1): X không đẩy được Y^2+^ => tính khử của X < Y.\n• Từ (2): Y đẩy được Z^+^ => tính khử của Y > Z.\n• Từ (3): Z không đẩy được X^2+^ => tính khử của Z < X.\nKết hợp lại ta được thứ tự giảm dần tính khử: Y > X > Z."
    },
    {
        "num": 16,
        "text": "(SBT – CTST). Cho E^o^_pin(Ni-Cu) = 0,59 V và E^o^_Cu^2+^/Cu = 0,34 V. Thế điện cực chuẩn của cặp oxi hóa - khử Ni^2+^/Ni (E^o^_Ni^2+^/Ni) là",
        "options": ["A. 0,34 V.", "B. 0,35 V.", "C. -0,34 V.", "D. -0,25 V."],
        "correct": "D",
        "sol": "Chọn D. Ta có công thức: E^o^_pin(Ni-Cu) = E^o^_Cu^2+^/Cu - E^o^_Ni^2+^/Ni => E^o^_Ni^2+^/Ni = 0,34 - 0,59 = -0,25 V."
    },
    {
        "num": 17,
        "text": "(SBT – KNTT). Sức điện động chuẩn của pin điện hoá H~2~ - Cu (gồm hai điện cực ứng với hai cặp oxi hoá - khử là 2H^+^/H~2~ và Cu^2+^/Cu) đo được bằng vôn kế có điện trở vô cùng lớn là 0,340 V. Từ đó, xác định được thế điện cực chuẩn của cặp Cu^2+^/Cu là",
        "options": ["A. -0,340 V.", "B. 0,000 V.", "C. 0,680 V.", "D. +0,340 V."],
        "correct": "D",
        "sol": "Chọn D. Quy ước thế điện cực chuẩn của điện cực hydro là 0,000 V (E^o^_2H^+^/H~2~ = 0,000 V). Khi đó: E^o^_pin = E^o^_Cu^2+^/Cu - E^o^_2H^+^/H~2~ => E^o^_Cu^2+^/Cu = 0,340 + 0,000 = +0,340 V."
    },
    {
        "num": 18,
        "text": "(SBT – CTST). Điện phân một dung dịch hỗn hợp gồm: HCl, CuCl~2~, NaCl với điện cực trơ, có màng ngăn. Kết luận nào sau đây không đúng?",
        "options": [
            "A. Sau khi kết thúc quá trình điện phân, pH của dung dịch tăng so với dung dịch ban đầu.",
            "B. Các chất bị điện phân lần lượt theo thứ tự: CuCl~2~, HCl, NaCl và H~2~O.",
            "C. Quá trình điện phân NaCl làm tăng pH của dung dịch.",
            "D. Quá trình điện phân HCl làm giảm pH của dung dịch."
        ],
        "correct": "D",
        "sol": "Chọn D. Quá trình điện phân HCl: 2HCl -> H~2~ + Cl~2~ tiêu thụ ion H^+^ làm nồng độ H^+^ giảm, do đó pH của dung dịch phải tăng lên. Nhận định D 'làm giảm pH' là sai."
    }
]

TF_QUESTIONS = [
    {
        "num": 1,
        "intro": "Cho thế điện cực chuẩn của các cặp oxi hóa – khử: Cu^2+^/Cu = 0,34 V; Fe^2+^/Fe = –0,44 V; Fe^3+^/Fe^2+^ = 0,77 V; Ag^+^/Ag = 0,80 V. Các phản ứng có thể xảy ra là:",
        "items": [
            ("a", "Cu^2+^ + Fe -> Cu + Fe^2+^.", True, "Đúng. Do E^o^_Cu^2+^/Cu (+0,34 V) > E^o^_Fe^2+^/Fe (-0,44 V) nên Fe khử được ion Cu^2+^ thành Cu."),
            ("b", "Fe + 2Fe^3+^ -> 3Fe^2+^.", True, "Đúng. Do E^o^_Fe^3+^/Fe^2+^ (+0,77 V) > E^o^_Fe^2+^/Fe (-0,44 V) nên Fe khử được ion Fe^3+^ về Fe^2+^."),
            ("c", "Ag^+^ + Fe^2+^ -> Ag + Fe^3+^.", True, "Đúng. Do E^o^_Ag^+^/Ag (+0,80 V) > E^o^_Fe^3+^/Fe^2+^ (+0,77 V) nên Ag^+^ oxi hóa được ion Fe^2+^ thành Fe^3+^."),
            ("d", "Cu + 2Fe^3+^ -> Cu^2+^ + 2Fe^2+^.", True, "Đúng. Do E^o^_Fe^3+^/Fe^2+^ (+0,77 V) > E^o^_Cu^2+^/Cu (+0,34 V) nên Cu khử được ion Fe^3+^ về Fe^2+^.")
        ]
    },
    {
        "num": 2,
        "intro": "(SBT – KNTT). Xét quá trình hoạt động của một pin điện hoá Cu – Ag được thiết lập ở các điều kiện như hình vẽ bên. Cho thế điện cực chuẩn của các cặp Cu^2+^/Cu và Ag^+^/Ag lần lượt là +0,340 V và +0,799 V.",
        "image": "p2_cau2_pin_cu_ag.png",
        "image_w": Cm(6.5),
        "items": [
            ("a", "Giá trị sức điện động chuẩn của pin điện hoá trên là 0,459 V.", True, "Đúng. E^o^_pin = E^o^_Ag^+^/Ag - E^o^_Cu^2+^/Cu = 0,799 - 0,340 = 0,459 V."),
            ("b", "Ở anode xảy ra quá trình oxi hoá Cu, ở cathode xảy ra quá trình khử Ag^+^.", True, "Đúng. Cực Cu có thế điện cực nhỏ hơn nên đóng vai trò anode (xảy ra oxi hóa Cu: Cu -> Cu^2+^ + 2e), cực Ag đóng vai trò cathode (xảy ra khử Ag^+^: Ag^+^ + 1e -> Ag)."),
            ("c", "Điện cực Cu tăng khối lượng, điện cực Ag giảm khối lượng.", False, "Sai. Ở anode, Cu bị tan ra nên khối lượng cực Cu giảm; ở cathode Ag sinh ra bám vào điện cực nên khối lượng cực Ag tăng."),
            ("d", "Phản ứng hoá học xảy ra trong pin: Cu + 2Ag^+^ -> Cu^2+^ + 2Ag.", True, "Đúng. Đây là phản ứng oxi hóa - khử tổng quát diễn ra tự phát trong pin điện hoá Cu - Ag.")
        ]
    },
    {
        "num": 3,
        "intro": "Thí nghiệm điện phân được thực hiện như hình vẽ sau:",
        "image": "p2_cau3_dien_phan_nacl.png",
        "image_w": Cm(7.5),
        "items": [
            ("a", "Có khí H~2~ thoát ra ở điện cực dương, khí O~2~ thoát ra ở cực âm.", False, "Sai. Ở cực âm (cathode) khử H~2~O thoát khí H~2~; ở cực dương (anode) oxi hóa Cl^-^ thoát khí Cl~2~."),
            ("b", "Thí nghiệm này ứng dụng để điều chế NaOH trong công nghiệp.", True, "Đúng. Đây là mô hình điện phân dung dịch NaCl có màng ngăn (vách ngăn) dùng để sản xuất xút NaOH, khí Cl~2~ và H~2~ trong công nghiệp."),
            ("c", "Nếu không có vách ngăn điện cực, dung dịch thu được có tính tẩy màu.", True, "Đúng. Khi không có vách ngăn, khí Cl~2~ ở anode sẽ phản ứng với NaOH ở cathode tạo thành nước Javel (chứa NaClO) có tính tẩy màu: Cl~2~ + 2NaOH -> NaCl + NaClO + H~2~O."),
            ("d", "Dung dịch sau điện phân làm quỳ tím chuyển sang màu xanh.", True, "Đúng. Dung dịch sau điện phân chứa base mạnh NaOH nên làm quỳ tím chuyển sang màu xanh.")
        ]
    },
    {
        "num": 4,
        "intro": "Điện phân dung dịch MSO~4~ (M là kim loại) với điện cực trơ, cường độ dòng điện không đổi. Sau thời gian t giây, thu được a mol khí ở anode. Nếu thời gian điện phân là 2t giây thì tổng số mol khí thu được ở cả hai điện cực là 2,5a mol. Giả sử hiệu suất điện phân là 100%, khí sinh ra không tan trong nước.",
        "items": [
            ("a", "Tại thời điểm 2t giây, có bọt khí ở cathode.", True, "Đúng. Tại 2t giây, số mol electron trao đổi gấp đôi (ne = 8a). Số mol khí ở anode = 2a mol O~2~ < 2,5a mol tổng số mol khí, chứng tỏ đã có 0,5a mol H~2~ thoát ra ở cathode."),
            ("b", "Tại thời điểm t giây, ion M^2+^ chưa bị điện phân hết.", True, "Đúng. Tổng số mol M^2+^ ban đầu là 3,5a mol. Tại t giây lượng electron là 4a mol, ion M^2+^ mới phản ứng 2a mol < 3,5a mol nên chưa điện phân hết."),
            ("c", "Dung dịch sau điện phân có pH < 7.", True, "Đúng. Ở anode sinh ra lượng ion H^+^ (4*2a = 8a mol) lớn hơn lượng OH^-^ sinh ra ở cathode (2*0,5a = 1a mol), dung dịch có môi trường acid nên pH < 7."),
            ("d", "Khi thu được 1,8a mol khí ở anode thì vẫn chưa xuất hiện bọt khí ở cathode.", False, "Sai. Khi n_O~2~ = 1,8a mol => ne = 7,2a mol. Vì lượng e cần khử hết M^2+^ là 2 * 3,5a = 7,0a mol < 7,2a mol nên đã bắt đầu điện phân H~2~O ở cathode sinh ra 0,1a mol khí H~2~.")
        ]
    }
]

SA_QUESTIONS = [
    {
        "num": 1,
        "text": "Có 3 cốc đựng các dung dịch muối có nồng độ 1 M, cho vào mỗi cốc 1 thanh Al:\n(a) Cốc I: Pb(NO~3~)~2~ (aq)\n(b) Cốc II: Ni(NO~3~)~2~ (aq)\n(c) Cốc III: Cu(NO~3~)~2~ (aq)\nCó bao nhiêu cốc có phản ứng xảy ra?",
        "ans": "3",
        "sol": "Vì E^o^_Al^3+^/Al (-1,66 V) nhỏ hơn thế điện cực của cả ba cặp Pb^2+^/Pb (-0,13 V), Ni^2+^/Ni (-0,26 V), Cu^2+^/Cu (+0,34 V) nên kim loại Al đều đẩy được cả ba ion Pb^2+^, Ni^2+^, Cu^2+^ ra khỏi dung dịch muối của chúng. Cả 3 cốc đều có phản ứng xảy ra."
    },
    {
        "num": 2,
        "text": "Cho pin điện hóa Zn - Cu gồm hai cặp oxi hóa - khử Zn^2+^/Zn và Cu^2+^/Cu. Cho thế điện cực chuẩn E^o^_Cu^2+^/Cu = 0,34 V và E^o^_Zn^2+^/Zn = -0,76 V. Tính sức điện động chuẩn của pin (theo đơn vị V).",
        "ans": "1,1",
        "sol": "Sức điện động chuẩn của pin Zn - Cu là: E^o^_pin = E^o^_Cu^2+^/Cu - E^o^_Zn^2+^/Zn = 0,34 - (-0,76) = 1,10 V."
    },
    {
        "num": 3,
        "text": "Điện phân dung dịch hỗn hợp gồm: NaCl, Mg(NO~3~)~2~, ZnCl~2~, CuSO~4~ và Fe~2~(SO~4~)~3~ với điện cực trơ cho đến khi nước bắt đầu điện phân ở cả hai điện cực thì dừng lại. Thu được bao nhiêu chất rắn (kim loại) ở cathode?",
        "ans": "3",
        "sol": "Tại cathode (cực âm), thứ tự nhận electron của các ion: Fe^3+^ -> Fe^2+^, sau đó Cu^2+^ -> Cu, Fe^2+^ -> Fe, Zn^2+^ -> Zn. Các ion Mg^2+^, Na^+^ không bị điện phân trong dung dịch. Vậy thu được 3 kim loại bám ở cathode là Cu, Fe và Zn."
    },
    {
        "num": 4,
        "text": "Cho các phản ứng sau:\n(1) Cho Cu vào dung dịch FeSO~4~.\n(2) Cho Cu vào dung dịch Fe~2~(SO~4~)~3~.\n(3) Cho K vào dung dịch CuSO~4~.\n(4) Cho dung dịch AgNO~3~ vào dung dịch Fe(NO~3~)~2~.\nSố lượng phản ứng có thể thu được đơn chất kim loại là",
        "ans": "1",
        "sol": "• (1) Không phản ứng do Cu có tính khử yếu hơn Fe.\n• (2) Cu + Fe~2~(SO~4~)~3~ -> CuSO~4~ + 2FeSO~4~ (không sinh ra kim loại).\n• (3) K phản ứng với nước trước tạo KOH, sau đó KOH phản ứng với CuSO~4~ tạo kết tủa Cu(OH)~2~ (không thu được kim loại Cu).\n• (4) AgNO~3~ + Fe(NO~3~)~2~ -> Fe(NO~3~)~3~ + Ag↓ (thu được kim loại Ag).\nVậy chỉ có 1 phản ứng thu được kim loại."
    },
    {
        "num": 5,
        "text": "Cho một pin Galvani với điện cực Zn và Cu có sức điện động chuẩn là 1,10 V. Sử dụng pin này để thắp sáng một bóng đèn nhỏ với cường độ dòng điện chạy qua là I = 0,02 A. Nếu điện cực kẽm hao mòn 0,1 mol do pin phóng điện thì thời gian tối đa mà pin thắp sáng được bóng đèn là bao nhiêu giờ?\nCho biết công thức: Q = n*F = I*t, trong đó Q là điện lượng (C), n là số mol electron đi qua dây dẫn, I là cường độ dòng điện (A), t là thời gian (giây), F là hằng số Faraday (F = 96500 C/mol). (Làm tròn kết quả đến hàng đơn vị).",
        "ans": "268",
        "sol": "• Ở cực âm (anode): Zn -> Zn^2+^ + 2e.\n  Số mol electron sinh ra khi tiêu hao 0,1 mol Zn là: n_e = 2 * 0,1 = 0,2 mol.\n• Điện lượng: Q = n_e * F = 0,2 * 96500 = 19300 C.\n• Thời gian phóng điện: t = Q / I = 19300 / 0,02 = 965000 giây.\n• Đổi sang giờ: t = 965000 / 3600 ≈ 268,06 ≈ 268 giờ."
    },
    {
        "num": 6,
        "text": "Tiến hành điện phân dung dịch CuSO~4~ với anode bằng đồng (hiện tượng dương cực tan). Để hoà tan 100 g đồng ở anode trong 8 giờ thì cần cường độ dòng điện bằng bao nhiêu ampe? (Làm tròn kết quả đến chữ số hàng phần mười).",
        "ans": "10,5",
        "sol": "• Số mol Cu bị hòa tan ở anode: n_Cu = 100 / 64 = 1,5625 mol.\n• Quá trình ở anode: Cu -> Cu^2+^ + 2e => n_e = 2 * 1,5625 = 3,125 mol.\n• Điện lượng: Q = n_e * F = 3,125 * 96500 = 301562,5 C.\n• Thời gian: t = 8 giờ = 8 * 3600 = 28800 giây.\n• Cường độ dòng điện: I = Q / t = 301562,5 / 28800 ≈ 10,47 ≈ 10,5 A."
    }
]

def generate_exam_docx(is_teacher=False):
    filename = "De_Kiem_Tra_Hoa_12_De13_GiaoVien.docx" if is_teacher else "De_Kiem_Tra_Hoa_12_De13.docx"
    filepath = os.path.join(OUT_DIR, filename)
    
    doc = docx.Document()
    setup_page_and_header_footer(doc)
    add_exam_title_block(doc, is_teacher=is_teacher)
    
    # --- PHẦN I ---
    add_part_heading(doc, "PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.", "Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án.")
    
    for q in MCQ_QUESTIONS:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(1)
        p_q.paragraph_format.line_spacing = 1.15
        
        r_qnum = p_q.add_run(f"Câu {q['num']}. ")
        set_run_style(r_qnum, font_name="Times New Roman", font_size=11, bold=True)
        
        add_chem_run(p_q, q["text"], base_font="Times New Roman", base_size=11)
        
        if q.get("image"):
            img_path = os.path.join(IMG_DIR, q["image"])
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(1)
                p_img.paragraph_format.space_after = Pt(2)
                run_img = p_img.add_run()
                w = q.get("image_w", Cm(8.0))
                run_img.add_picture(img_path, width=w)

        # Format options using Tabs (NO TABLE)
        opts = q["options"]
        max_opt_len = max(len(opt) for opt in opts)
        
        if max_opt_len < 22:
            # 1 dòng 4 lựa chọn (Tab stops 4.5, 9.0, 13.5 cm)
            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.space_before = Pt(1)
            p_opt.paragraph_format.space_after = Pt(2)
            p_opt.paragraph_format.line_spacing = 1.15
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(4.5), WD_TAB_ALIGNMENT.LEFT)
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(13.5), WD_TAB_ALIGNMENT.LEFT)
            
            for i, opt in enumerate(opts):
                is_correct = (opt.startswith(q["correct"] + ".")) and is_teacher
                color = COLOR_CORRECT if is_correct else None
                
                prefix = opt[:3]
                body = opt[3:]
                r_pre = p_opt.add_run(prefix)
                set_run_style(r_pre, font_name="Times New Roman", font_size=11, bold=True, color=color)
                add_chem_run(p_opt, body, base_font="Times New Roman", base_size=11, bold=is_correct, color=color)
                
                if i < 3:
                    p_opt.add_run("\t")
        elif max_opt_len < 46:
            # 2 dòng 2 lựa chọn (Tab stop 9.0 cm)
            for row in [(0, 1), (2, 3)]:
                p_opt = doc.add_paragraph()
                p_opt.paragraph_format.space_before = Pt(1)
                p_opt.paragraph_format.space_after = Pt(1.5)
                p_opt.paragraph_format.line_spacing = 1.15
                p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
                
                for idx_in_row, opt_idx in enumerate(row):
                    opt = opts[opt_idx]
                    is_correct = (opt.startswith(q["correct"] + ".")) and is_teacher
                    color = COLOR_CORRECT if is_correct else None
                    
                    prefix = opt[:3]
                    body = opt[3:]
                    r_pre = p_opt.add_run(prefix)
                    set_run_style(r_pre, font_name="Times New Roman", font_size=11, bold=True, color=color)
                    add_chem_run(p_opt, body, base_font="Times New Roman", base_size=11, bold=is_correct, color=color)
                    
                    if idx_in_row == 0:
                        p_opt.add_run("\t")
        else:
            # 4 dòng riêng biệt
            for opt in opts:
                p_opt = doc.add_paragraph()
                p_opt.paragraph_format.space_before = Pt(1)
                p_opt.paragraph_format.space_after = Pt(1.5)
                p_opt.paragraph_format.line_spacing = 1.15
                
                is_correct = (opt.startswith(q["correct"] + ".")) and is_teacher
                color = COLOR_CORRECT if is_correct else None
                
                prefix = opt[:3]
                body = opt[3:]
                r_pre = p_opt.add_run(prefix)
                set_run_style(r_pre, font_name="Times New Roman", font_size=11, bold=True, color=color)
                add_chem_run(p_opt, body, base_font="Times New Roman", base_size=11, bold=is_correct, color=color)

        if is_teacher and q.get("sol"):
            add_solution_paragraph(doc, q["sol"])

    # --- PHẦN II ---
    add_part_heading(doc, "PHẦN II. Câu trắc nghiệm đúng sai.", "Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.")
    
    for q in TF_QUESTIONS:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        
        r_qnum = p_q.add_run(f"Câu {q['num']}. ")
        set_run_style(r_qnum, font_name="Times New Roman", font_size=11, bold=True)
        
        add_chem_run(p_q, q["intro"], base_font="Times New Roman", base_size=11)
        
        if q.get("image"):
            img_path = os.path.join(IMG_DIR, q["image"])
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(2)
                p_img.paragraph_format.space_after = Pt(2)
                run_img = p_img.add_run()
                w = q.get("image_w", Cm(8.0))
                run_img.add_picture(img_path, width=w)

        for item_key, item_text, is_true, item_sol in q["items"]:
            p_item = doc.add_paragraph()
            p_item.paragraph_format.space_before = Pt(1)
            p_item.paragraph_format.space_after = Pt(1.5)
            p_item.paragraph_format.line_spacing = 1.15
            p_item.paragraph_format.left_indent = Cm(0.5)
            
            r_k = p_item.add_run(f"{item_key}) ")
            set_run_style(r_k, font_name="Times New Roman", font_size=11, bold=True)
            
            add_chem_run(p_item, item_text, base_font="Times New Roman", base_size=11)
            
            if is_teacher:
                status_str = " [ĐÚNG]" if is_true else " [SAI]"
                color_st = COLOR_CORRECT if is_true else COLOR_PRIMARY
                r_st = p_item.add_run(status_str)
                set_run_style(r_st, font_name="Times New Roman", font_size=11, bold=True, color=color_st)
                
                p_sol = doc.add_paragraph()
                p_sol.paragraph_format.space_before = Pt(0)
                p_sol.paragraph_format.space_after = Pt(2)
                p_sol.paragraph_format.left_indent = Cm(1.0)
                r_sol_lbl = p_sol.add_run("-> ")
                set_run_style(r_sol_lbl, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_SOLUTION)
                add_chem_run(p_sol, item_sol, base_font="Times New Roman", base_size=10, italic=True)

    # --- PHẦN III ---
    add_part_heading(doc, "PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn.", "Thí sinh trả lời từ câu 1 đến câu 6.")
    
    for q in SA_QUESTIONS:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        
        r_qnum = p_q.add_run(f"Câu {q['num']}. ")
        set_run_style(r_qnum, font_name="Times New Roman", font_size=11, bold=True)
        
        add_chem_run(p_q, q["text"], base_font="Times New Roman", base_size=11)
        
        if q.get("image"):
            img_path = os.path.join(IMG_DIR, q["image"])
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(2)
                p_img.paragraph_format.space_after = Pt(2)
                run_img = p_img.add_run()
                w = q.get("image_w", Cm(8.0))
                run_img.add_picture(img_path, width=w)

        if not is_teacher:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(3)
            p_ans.paragraph_format.left_indent = Cm(0.5)
            r_ans_lbl = p_ans.add_run("Trả lời: .................................................................................................................................")
            set_run_style(r_ans_lbl, font_name="Times New Roman", font_size=11, italic=True)
        else:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(2)
            p_ans.paragraph_format.left_indent = Cm(0.5)
            
            r_ans_lbl = p_ans.add_run("Đáp số: ")
            set_run_style(r_ans_lbl, font_name="Times New Roman", font_size=11, bold=True, color=COLOR_CORRECT)
            r_ans_val = p_ans.add_run(q["ans"])
            set_run_style(r_ans_val, font_name="Times New Roman", font_size=11, bold=True, color=COLOR_CORRECT)
            
            add_solution_paragraph(doc, q["sol"])

    # Bảng tổng hợp đáp án ở cuối đề Giáo Viên
    if is_teacher:
        p_table_title = doc.add_paragraph()
        p_table_title.paragraph_format.space_before = Pt(10)
        p_table_title.paragraph_format.space_after = Pt(3)
        r_tt = p_table_title.add_run("BẢNG TỔNG HỢP ĐÁP ÁN ĐỀ SỐ 13")
        set_run_style(r_tt, font_name="Times New Roman", font_size=11.5, bold=True, color=COLOR_PRIMARY)
        
        # Phần I Table
        p_p1 = doc.add_paragraph()
        p_p1.paragraph_format.space_before = Pt(3)
        p_p1.paragraph_format.space_after = Pt(1)
        r_p1 = p_p1.add_run("Phần I: Trắc nghiệm nhiều phương án lựa chọn")
        set_run_style(r_p1, font_name="Times New Roman", font_size=10.5, bold=True)
        
        tbl_ans1 = doc.add_table(rows=2, cols=18)
        tbl_ans1.autofit = False
        for c_idx in range(18):
            cell_q = tbl_ans1.cell(0, c_idx)
            cell_a = tbl_ans1.cell(1, c_idx)
            cell_q.width = Cm(1.0)
            cell_a.width = Cm(1.0)
            
            pq = cell_q.paragraphs[0]
            pq.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pq.paragraph_format.space_before = Pt(1)
            pq.paragraph_format.space_after = Pt(1)
            rq = pq.add_run(str(c_idx + 1))
            set_run_style(rq, font_name="Times New Roman", font_size=9, bold=True)
            
            pa = cell_a.paragraphs[0]
            pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pa.paragraph_format.space_before = Pt(1)
            pa.paragraph_format.space_after = Pt(1)
            ra = pa.add_run(MCQ_QUESTIONS[c_idx]["correct"])
            set_run_style(ra, font_name="Times New Roman", font_size=9.5, bold=True, color=COLOR_CORRECT)
            
            for cell in [cell_q, cell_a]:
                set_cell_border(cell, top={'val':'single','sz':'4','color':'CCCCCC'},
                                      bottom={'val':'single','sz':'4','color':'CCCCCC'},
                                      left={'val':'single','sz':'4','color':'CCCCCC'},
                                      right={'val':'single','sz':'4','color':'CCCCCC'})
                set_cell_margins(cell, top=20, bottom=20, left=10, right=10)

        # Phần II Table
        p_p2 = doc.add_paragraph()
        p_p2.paragraph_format.space_before = Pt(5)
        p_p2.paragraph_format.space_after = Pt(1)
        r_p2 = p_p2.add_run("Phần II: Đúng/Sai")
        set_run_style(r_p2, font_name="Times New Roman", font_size=10.5, bold=True)
        
        tbl_ans2 = doc.add_table(rows=5, cols=5)
        tbl_ans2.autofit = False
        headers_p2 = ["Câu", "a", "b", "c", "d"]
        for j, h in enumerate(headers_p2):
            cell = tbl_ans2.cell(0, j)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(h)
            set_run_style(r, font_name="Times New Roman", font_size=9.5, bold=True)
            set_cell_border(cell, top={'val':'single','sz':'4','color':'CCCCCC'},
                                  bottom={'val':'single','sz':'4','color':'CCCCCC'},
                                  left={'val':'single','sz':'4','color':'CCCCCC'},
                                  right={'val':'single','sz':'4','color':'CCCCCC'})
            set_cell_margins(cell, top=20, bottom=20, left=10, right=10)

        for i, q in enumerate(TF_QUESTIONS):
            cell_q = tbl_ans2.cell(i + 1, 0)
            cell_q.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cell_q.paragraphs[0].paragraph_format.space_before = Pt(1)
            cell_q.paragraphs[0].paragraph_format.space_after = Pt(1)
            rq = cell_q.paragraphs[0].add_run(f"Câu {q['num']}")
            set_run_style(rq, font_name="Times New Roman", font_size=9.5, bold=True)
            set_cell_border(cell_q, top={'val':'single','sz':'4','color':'CCCCCC'},
                                    bottom={'val':'single','sz':'4','color':'CCCCCC'},
                                    left={'val':'single','sz':'4','color':'CCCCCC'},
                                    right={'val':'single','sz':'4','color':'CCCCCC'})
            set_cell_margins(cell_q, top=20, bottom=20, left=10, right=10)
            
            for j, item in enumerate(q["items"]):
                cell_item = tbl_ans2.cell(i + 1, j + 1)
                p_it = cell_item.paragraphs[0]
                p_it.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_it.paragraph_format.space_before = Pt(1)
                p_it.paragraph_format.space_after = Pt(1)
                is_true = item[2]
                st_text = "Đ" if is_true else "S"
                st_col = COLOR_CORRECT if is_true else COLOR_PRIMARY
                r_st = p_it.add_run(st_text)
                set_run_style(r_st, font_name="Times New Roman", font_size=9.5, bold=True, color=st_col)
                set_cell_border(cell_item, top={'val':'single','sz':'4','color':'CCCCCC'},
                                           bottom={'val':'single','sz':'4','color':'CCCCCC'},
                                           left={'val':'single','sz':'4','color':'CCCCCC'},
                                           right={'val':'single','sz':'4','color':'CCCCCC'})
                set_cell_margins(cell_item, top=20, bottom=20, left=10, right=10)

        # Phần III Table
        p_p3 = doc.add_paragraph()
        p_p3.paragraph_format.space_before = Pt(5)
        p_p3.paragraph_format.space_after = Pt(1)
        r_p3 = p_p3.add_run("Phần III: Trả lời ngắn")
        set_run_style(r_p3, font_name="Times New Roman", font_size=10.5, bold=True)
        
        tbl_ans3 = doc.add_table(rows=2, cols=6)
        tbl_ans3.autofit = False
        for c_idx in range(6):
            cell_q = tbl_ans3.cell(0, c_idx)
            cell_a = tbl_ans3.cell(1, c_idx)
            cell_q.width = Cm(3.0)
            cell_a.width = Cm(3.0)
            
            pq = cell_q.paragraphs[0]
            pq.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pq.paragraph_format.space_before = Pt(1)
            pq.paragraph_format.space_after = Pt(1)
            rq = pq.add_run(f"Câu {c_idx + 1}")
            set_run_style(rq, font_name="Times New Roman", font_size=9.5, bold=True)
            
            pa = cell_a.paragraphs[0]
            pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pa.paragraph_format.space_before = Pt(1)
            pa.paragraph_format.space_after = Pt(1)
            ra = pa.add_run(SA_QUESTIONS[c_idx]["ans"])
            set_run_style(ra, font_name="Times New Roman", font_size=9.5, bold=True, color=COLOR_CORRECT)
            
            for cell in [cell_q, cell_a]:
                set_cell_border(cell, top={'val':'single','sz':'4','color':'CCCCCC'},
                                      bottom={'val':'single','sz':'4','color':'CCCCCC'},
                                      left={'val':'single','sz':'4','color':'CCCCCC'},
                                      right={'val':'single','sz':'4','color':'CCCCCC'})
                set_cell_margins(cell, top=20, bottom=20, left=10, right=10)

    doc.save(filepath)
    print(f"Đã tạo thành công: {filepath}")

if __name__ == "__main__":
    generate_exam_docx(is_teacher=False)
    generate_exam_docx(is_teacher=True)
