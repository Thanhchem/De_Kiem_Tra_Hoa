# -*- coding: utf-8 -*-
"""
Script tạo trọn bộ file Word (.docx) ĐỀ CƯƠNG ÔN TẬP GIỮA KÌ I KHTN 7
Trường THCS Nguyễn Thị Minh Khai - Thầy TRẦN VĂN THẠNH - 0777.470.803

Tạo 3 phiên bản chuẩn hóa theo FORM ĐỀ KIỂM TRA:
1. Bản Học sinh: San_Pham/De_Cuong_KHTN_7_Giua_HK1_MinhKhai_HocSinh.docx
2. Bản Giáo viên: San_Pham/De_Cuong_KHTN_7_Giua_HK1_MinhKhai_GiaoVien.docx
3. Bản Dạy trên màn hình: San_Pham/De_Cuong_KHTN_7_Giua_HK1_MinhKhai_DayTrenManHinh.docx
"""

import os
import sys
import re
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.path.append(os.path.dirname(__file__))
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

from khtn7_minhkhai_data import (
    INFO, MC_QUESTIONS, TF_QUESTIONS, SHORT_QUESTIONS, ESSAY_QUESTIONS
)

ROOT_DIR = r"c:\Antigravity_Thanh"
OUT_DIR = os.path.join(ROOT_DIR, "San_Pham")
ONEDRIVE_DIR = r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn7\KHTN7-IN-DEKIEMTRADINHKI\New folder"

# Bảng màu chuẩn mực
COLOR_HF_RED = RGBColor(211, 47, 47)         # Đỏ san hô (#D32F2F)
COLOR_HF_GREEN_HEX = "8BC390"                # Xanh Sage Green
COLOR_PRIMARY = RGBColor(46, 125, 50)        # Xanh lá đậm (#2E7D32)
COLOR_CORRECT = RGBColor(198, 40, 40)        # Đỏ đậm nhấn đáp án đúng (#C62828)
COLOR_WRONG = RGBColor(0, 105, 92)           # Xanh tím than cho [SAI]
COLOR_SOLUTION = RGBColor(21, 101, 192)      # Xanh lam đậm (#1565C0) cho tiêu đề lời giải
COLOR_SCREEN_TITLE = RGBColor(106, 27, 154)  # Tím thanh lịch cho bản trình chiếu
COLOR_GRAY = RGBColor(97, 97, 97)            # Xám ghi chú
COLOR_BORDER_GREY_HEX = "CCCCCC"

def add_p_border_bottom(p, color="8BC390", sz="6"):
    """Thêm đường viền đáy cho đoạn văn bản."""
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_p_border_top(p, color="8BC390", sz="6"):
    """Thêm đường viền trên cho đoạn văn bản."""
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    """Cài đặt lề đệm cho ô bảng."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex="F1F8E9"):
    """Đặt màu nền cho ô bảng."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="CCCCCC", sz="4"):
    """Đặt viền thanh mảnh cho bảng."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def set_run_style(run, font_name="Times New Roman", font_size=10.5, bold=False, italic=False, subscript=False, superscript=False, color=None):
    """Cài đặt style chi tiết cho từng run text."""
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    """Chèn trường số trang tự động PAGE vào chân trang."""
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def render_rich_text(p, text, base_font="Times New Roman", base_size=10.5, base_bold=False, base_italic=False, base_color=None):
    """
    Sinh văn bản định dạng phong phú từ kí hiệu:
    - *bold*
    - _italic_
    - ~subscript~ (CO~2~, H~2~O, H~2~SO~4~)
    - ^superscript^ (cm^3, m/s^2)
    """
    tokens = re.split(r'(~[^~]+~|\^[^\^]+\^|\*[^\*]+\*|_[^_]+_)', text)
    for token in tokens:
        if not token:
            continue

        is_bold = base_bold
        is_italic = base_italic
        is_sub = False
        is_sup = False
        content = token

        if token.startswith('*') and token.endswith('*'):
            is_bold = True
            content = token[1:-1]
        elif token.startswith('_') and token.endswith('_'):
            is_italic = True
            content = token[1:-1]
        elif token.startswith('~') and token.endswith('~'):
            is_sub = True
            content = token[1:-1]
        elif token.startswith('^') and token.endswith('^'):
            is_sup = True
            content = token[1:-1]

        run = p.add_run(content)
        run_color = COLOR_CORRECT if (is_bold and not base_bold and not base_color) else base_color
        set_run_style(run, font_name=base_font, font_size=base_size, bold=is_bold, italic=is_italic,
                      subscript=is_sub, superscript=is_sup, color=run_color)

def setup_header_footer(doc):
    """Cài đặt Header & Footer chuẩn theo FORM_DE_KIEM_TRA."""
    normal_style = doc.styles['Normal']
    normal_style.paragraph_format.tab_stops.clear_all()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    section.start_type = WD_SECTION_START.NEW_PAGE
    section.different_first_page_header_footer = False
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    # 1. Header (Bảng 1x2, rộng 18.0 cm)
    header = section.header
    for p in header.paragraphs:
        p.text = ""
    tbl_h = header.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl_h.cell(0, 0)
    c_right = tbl_h.cell(0, 1)
    c_left.width = Cm(8.0)
    c_right.width = Cm(10.0)

    p_hl = c_left.paragraphs[0]
    p_hl.paragraph_format.space_before = Pt(0)
    p_hl.paragraph_format.space_after = Pt(2)
    p_hl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_hl = p_hl.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)

    p_hr = c_right.paragraphs[0]
    p_hr.paragraph_format.space_before = Pt(0)
    p_hr.paragraph_format.space_after = Pt(2)
    p_hr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p_hr.add_run("Thầy ")
    set_run_style(r1, font_size=10, color=COLOR_HF_RED)
    r2 = p_hr.add_run("TRẦN VĂN THẠNH")
    set_run_style(r2, font_size=10, bold=True, color=COLOR_HF_RED)
    r3 = p_hr.add_run(" - 0777.470.803")
    set_run_style(r3, font_size=10, color=COLOR_HF_RED)

    add_p_border_bottom(p_hl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_bottom(p_hr, color=COLOR_HF_GREEN_HEX, sz="6")

    # 2. Footer (Bảng 1x2, rộng 18.0 cm)
    footer = section.footer
    for p in footer.paragraphs:
        p.text = ""
    tbl_f = footer.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_fl = tbl_f.cell(0, 0)
    c_fr = tbl_f.cell(0, 1)
    c_fl.width = Cm(14.5)
    c_fr.width = Cm(3.5)

    p_fl = c_fl.paragraphs[0]
    p_fl.paragraph_format.space_before = Pt(2)
    p_fl.paragraph_format.space_after = Pt(0)
    p_fl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_fl = p_fl.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)

    p_fr = c_fr.paragraphs[0]
    p_fr.paragraph_format.space_before = Pt(2)
    p_fr.paragraph_format.space_after = Pt(0)
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_field(p_fr, font_size=10, bold=True, color=COLOR_HF_RED)

    add_p_border_top(p_fl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_top(p_fr, color=COLOR_HF_GREEN_HEX, sz="6")

def build_title_block(doc, is_teacher=False, is_screen=False, base_size=10.5):
    """Xây dựng Header Block không viền, đẹp mắt, cân đối."""
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_l, cell_r = tbl.cell(0, 0), tbl.cell(0, 1)
    cell_l.width = Cm(8.0)
    cell_r.width = Cm(10.0)

    # Cột trái
    p_tl1 = cell_l.paragraphs[0]
    p_tl1.paragraph_format.space_before, p_tl1.paragraph_format.space_after = Pt(0), Pt(1)
    p_tl1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tl1, f"*{INFO['department']}*", base_size=10, base_bold=True)

    p_tl2 = cell_l.add_paragraph()
    p_tl2.paragraph_format.space_before, p_tl2.paragraph_format.space_after = Pt(0), Pt(1)
    p_tl2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tl2, f"*{INFO['school']}*", base_size=10.5, base_bold=True)

    p_tl3 = cell_l.add_paragraph()
    p_tl3.paragraph_format.space_before, p_tl3.paragraph_format.space_after = Pt(0), Pt(0)
    p_tl3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_off = p_tl3.add_run("TÀI LIỆU ÔN TẬP ĐỘC QUYỀN")
    set_run_style(r_off, font_size=9, bold=True, color=COLOR_PRIMARY)

    # Cột phải
    p_tr1 = cell_r.paragraphs[0]
    p_tr1.paragraph_format.space_before, p_tr1.paragraph_format.space_after = Pt(0), Pt(1)
    p_tr1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tr1, f"*{INFO['title']}*", base_size=11.5, base_bold=True, base_color=COLOR_HF_RED)

    p_tr2 = cell_r.add_paragraph()
    p_tr2.paragraph_format.space_before, p_tr2.paragraph_format.space_after = Pt(0), Pt(1)
    p_tr2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tr2, f"*{INFO['subject']}*", base_size=11, base_bold=True, base_color=COLOR_PRIMARY)

    p_tr3 = cell_r.add_paragraph()
    p_tr3.paragraph_format.space_before, p_tr3.paragraph_format.space_after = Pt(0), Pt(0)
    p_tr3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tr3, f"_{INFO['year']} - {INFO['time_str']}_", base_size=9.5, base_italic=True)

    if is_teacher:
        p_gv = doc.add_paragraph()
        p_gv.paragraph_format.space_before, p_gv.paragraph_format.space_after = Pt(4), Pt(2)
        p_gv.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_gv = p_gv.add_run("(BẢN GIÁO VIÊN -- CÓ ĐÁP ÁN & LỜI GIẢI CHI TIẾT)")
        set_run_style(r_gv, font_size=11.5, bold=True, color=COLOR_CORRECT)
    elif is_screen:
        p_scr = doc.add_paragraph()
        p_scr.paragraph_format.space_before, p_scr.paragraph_format.space_after = Pt(4), Pt(2)
        p_scr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_scr = p_scr.add_run("(BẢN TRÌNH CHIẾU / GIẢNG DẠY MÀN HÌNH -- CÓ CHỪA CHỖ TRỐNG)")
        set_run_style(r_scr, font_size=11.5, bold=True, color=COLOR_SCREEN_TITLE)

    # Dòng thông tin thí sinh
    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(4)
    p_info.paragraph_format.space_after = Pt(4)
    p_info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_info = p_info.add_run("Họ, tên học sinh: ................................................................ Lớp: ............. Số báo danh: .............")
    set_run_style(r_info, font_size=10, italic=True)

    # Dải phân cách mỏng
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(4)
    add_p_border_bottom(p_div, color=COLOR_BORDER_GREY_HEX, sz="4")

def add_major_section_header(doc, title, sub_title="", base_size=11.5):
    """Tiêu đề phần lớn với thanh nền nổi bật."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True

    pPr = p._element.get_or_add_pPr()
    shd_xml = f'<w:shd {nsdecls("w")} w:fill="E8F5E9"/>'
    pbdr_xml = f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="8" w:color="2E7D32"/></w:pBdr>'
    pPr.append(parse_xml(shd_xml))
    pPr.append(parse_xml(pbdr_xml))

    r_title = p.add_run(f" {title.upper()} ")
    set_run_style(r_title, font_size=base_size, bold=True, color=COLOR_PRIMARY)
    if sub_title:
        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Cm(0.5)
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(3)
        p2.paragraph_format.keep_with_next = True
        r_sub = p2.add_run(f"({sub_title})")
        set_run_style(r_sub, font_size=base_size - 1.5, italic=True, color=COLOR_GRAY)

def add_section_title(doc, title, base_size=11.0):
    """Tiêu đề mục nhỏ."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    render_rich_text(p, f"*{title}*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)

def add_solution_block(doc, solution_text, base_size=10.0):
    """Hiển thị khối lời giải chi tiết (không dùng Table, thụt lề chuẩn)."""
    p_sol = doc.add_paragraph()
    p_sol.paragraph_format.left_indent = Cm(0.6)
    p_sol.paragraph_format.space_before = Pt(2)
    p_sol.paragraph_format.space_after = Pt(1)
    p_sol.paragraph_format.line_spacing = 1.15

    r_lbl = p_sol.add_run("Lời giải chi tiết:\n")
    set_run_style(r_lbl, font_size=base_size, bold=True, color=COLOR_SOLUTION)

    lines = solution_text.split('\n')
    for idx, line in enumerate(lines):
        line_str = line.strip()
        if not line_str:
            continue
        if idx > 0:
            p_sol = doc.add_paragraph()
            p_sol.paragraph_format.left_indent = Cm(0.6)
            p_sol.paragraph_format.space_before = Pt(1)
            p_sol.paragraph_format.space_after = Pt(1)
            p_sol.paragraph_format.line_spacing = 1.15
        render_rich_text(p_sol, line_str, base_size=base_size)

def add_screen_teaching_space(doc, num_lines=4):
    """Chừa khoảng trống tương tác cho bản Dạy Trên Màn Hình."""
    p_box = doc.add_paragraph()
    p_box.paragraph_format.left_indent = Cm(0.5)
    p_box.paragraph_format.space_before = Pt(3)
    p_box.paragraph_format.space_after = Pt(6)
    
    # Viền nét đứt nhẹ, viền trái tím đậm
    pPr = p_box._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="dashed" w:sz="4" w:space="4" w:color="B0BEC5"/><w:bottom w:val="dashed" w:sz="4" w:space="4" w:color="B0BEC5"/><w:left w:val="single" w:sz="12" w:space="6" w:color="7E57C2"/></w:pBdr>'
    pPr.append(parse_xml(xml))
    
    r_hint = p_box.add_run("✍ Ghi chú & Lời giải trực tiếp trên màn hình:\n")
    set_run_style(r_hint, font_size=9.5, bold=True, color=COLOR_SCREEN_TITLE)
    for _ in range(num_lines):
        p_box.add_run("\n")

def add_student_writing_lines(doc, num_lines=5):
    """Thêm dòng chấm cho học sinh làm bài tự luận."""
    for _ in range(num_lines):
        p_line = doc.add_paragraph()
        p_line.paragraph_format.left_indent = Cm(0.5)
        p_line.paragraph_format.space_before = Pt(1)
        p_line.paragraph_format.space_after = Pt(2)
        r = p_line.add_run("...................................................................................................................................................................")
        set_run_style(r, font_size=9, color=RGBColor(189, 189, 189))

def render_mc_questions(doc, questions, is_teacher=False, is_screen=False, base_size=10.5):
    """
    Hiển thị các câu hỏi trắc nghiệm 4 lựa chọn (A, B, C, D)
    TUYỆT ĐỐI KHÔNG DÙNG TABLE: Dùng Tab Stops chuẩn Word!
    """
    for q in questions:
        qid = q["id"]
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        p_q.paragraph_format.keep_with_next = True
        render_rich_text(p_q, f"*Câu {qid}.* {q['content']}", base_size=base_size)

        opts = q["options"]
        lens = [len(opt) for opt in opts]
        max_len = max(lens)

        if max_len <= 18:
            # 4 phương án trên 1 dòng với Tab Stops
            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.space_before = Pt(1)
            p_opt.paragraph_format.space_after = Pt(2)
            p_opt.paragraph_format.tab_stops.clear_all()
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(4.8), WD_TAB_ALIGNMENT.LEFT)
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)
            p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(13.6), WD_TAB_ALIGNMENT.LEFT)
            
            p_opt.add_run("\t")
            for idx, opt_text in enumerate(opts):
                opt_lbl = opt_text[:1]
                is_correct_opt = (is_teacher and opt_lbl == q["correct"])
                if is_correct_opt:
                    render_rich_text(p_opt, f"*{opt_text}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
                else:
                    render_rich_text(p_opt, opt_text, base_size=base_size)
                if idx < 3:
                    p_opt.add_run("\t")
        elif max_len <= 38:
            # 2 hàng, mỗi hàng 2 phương án
            # Hàng 1: A và B
            p_opt1 = doc.add_paragraph()
            p_opt1.paragraph_format.space_before = Pt(1)
            p_opt1.paragraph_format.space_after = Pt(1)
            p_opt1.paragraph_format.tab_stops.clear_all()
            p_opt1.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
            p_opt1.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)
            
            p_opt1.add_run("\t")
            optA = opts[0]
            is_cor_A = (is_teacher and optA[:1] == q["correct"])
            if is_cor_A:
                render_rich_text(p_opt1, f"*{optA}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p_opt1, optA, base_size=base_size)
            
            p_opt1.add_run("\t")
            optB = opts[1]
            is_cor_B = (is_teacher and optB[:1] == q["correct"])
            if is_cor_B:
                render_rich_text(p_opt1, f"*{optB}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p_opt1, optB, base_size=base_size)

            # Hàng 2: C và D
            p_opt2 = doc.add_paragraph()
            p_opt2.paragraph_format.space_before = Pt(1)
            p_opt2.paragraph_format.space_after = Pt(2)
            p_opt2.paragraph_format.tab_stops.clear_all()
            p_opt2.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
            p_opt2.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)

            p_opt2.add_run("\t")
            optC = opts[2]
            is_cor_C = (is_teacher and optC[:1] == q["correct"])
            if is_cor_C:
                render_rich_text(p_opt2, f"*{optC}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p_opt2, optC, base_size=base_size)
            
            p_opt2.add_run("\t")
            optD = opts[3]
            is_cor_D = (is_teacher and optD[:1] == q["correct"])
            if is_cor_D:
                render_rich_text(p_opt2, f"*{optD}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p_opt2, optD, base_size=base_size)
        else:
            # 4 hàng, mỗi phương án 1 dòng riêng
            for opt_text in opts:
                p_opt = doc.add_paragraph()
                p_opt.paragraph_format.left_indent = Cm(0.5)
                p_opt.paragraph_format.space_before = Pt(1)
                p_opt.paragraph_format.space_after = Pt(1)
                p_opt.paragraph_format.line_spacing = 1.15
                opt_lbl = opt_text[:1]
                is_correct_opt = (is_teacher and opt_lbl == q["correct"])
                if is_correct_opt:
                    render_rich_text(p_opt, f"*{opt_text}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
                else:
                    render_rich_text(p_opt, opt_text, base_size=base_size)

        if is_teacher:
            add_solution_block(doc, q["explanation"], base_size=base_size - 0.5)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=3)

def render_tf_questions(doc, questions, is_teacher=False, is_screen=False, base_size=10.5):
    """Hiển thị các câu hỏi Đúng / Sai chuẩn mực."""
    for q in questions:
        qid = q["id"]
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(5)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        p_q.paragraph_format.keep_with_next = True
        render_rich_text(p_q, f"*Câu {qid}.* {q['content']}", base_size=base_size)

        for st in q["statements"]:
            sub_id = st["id"]
            st_text = st["text"]
            correct_val = "Đúng" if st["correct"] else "Sai"

            p_st = doc.add_paragraph()
            p_st.paragraph_format.left_indent = Cm(0.5)
            p_st.paragraph_format.space_before = Pt(1)
            p_st.paragraph_format.space_after = Pt(1)
            p_st.paragraph_format.line_spacing = 1.15

            if is_teacher:
                r_tag = p_st.add_run(f"{sub_id}) ")
                set_run_style(r_tag, font_size=base_size, bold=True)
                
                ans_color = COLOR_CORRECT if st["correct"] else COLOR_WRONG
                r_ans = p_st.add_run(f"[{correct_val.upper()}] ")
                set_run_style(r_ans, font_size=base_size, bold=True, color=ans_color)

                render_rich_text(p_st, st_text, base_size=base_size)
            else:
                r_tag = p_st.add_run(f"{sub_id}) ")
                set_run_style(r_tag, font_size=base_size, bold=True)
                render_rich_text(p_st, st_text, base_size=base_size)

        if is_teacher:
            add_solution_block(doc, q["explanation"], base_size=base_size - 0.5)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=4)

def render_short_questions(doc, questions, is_teacher=False, is_screen=False, base_size=10.5):
    """Hiển thị các câu hỏi trắc nghiệm trả lời ngắn."""
    for q in questions:
        qid = q["id"]
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(5)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        p_q.paragraph_format.keep_with_next = True
        render_rich_text(p_q, f"*Câu {qid}.* {q['content']}", base_size=base_size)

        # Dòng đáp án
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.left_indent = Cm(0.5)
        p_ans.paragraph_format.space_before = Pt(1)
        p_ans.paragraph_format.space_after = Pt(2)

        if is_teacher:
            r_lbl = p_ans.add_run("Đáp án: ")
            set_run_style(r_lbl, font_size=base_size, bold=True)
            r_val = p_ans.add_run(f"{q['answer']}")
            set_run_style(r_val, font_size=base_size, bold=True, color=COLOR_CORRECT)
            add_solution_block(doc, q["explanation"], base_size=base_size - 0.5)
        elif is_screen:
            r_lbl = p_ans.add_run("Đáp án: ................................................................")
            set_run_style(r_lbl, font_size=base_size, italic=True)
            add_screen_teaching_space(doc, num_lines=3)
        else:
            r_lbl = p_ans.add_run("Đáp án: ................................................................")
            set_run_style(r_lbl, font_size=base_size, italic=True)

def render_essay_questions(doc, questions, is_teacher=False, is_screen=False, base_size=10.5):
    """Hiển thị các câu hỏi phần Tự luận."""
    for q in questions:
        qid = q["id"]
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        p_q.paragraph_format.keep_with_next = True
        
        # In nội dung câu hỏi
        render_rich_text(p_q, f"*{q['content']}*", base_size=base_size)

        if is_teacher:
            add_solution_block(doc, q["solution"], base_size=base_size - 0.5)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=q.get("blank_lines_screen", 6))
        else:
            add_student_writing_lines(doc, num_lines=q.get("blank_lines_student", 6))

def add_teacher_summary_table(doc, base_size=9.5):
    """Thêm Bảng tổng hợp đáp án trắc nghiệm ở cuối bản Giáo viên."""
    add_major_section_header(doc, "BẢNG TỔNG HỢP ĐÁP ÁN TRẮC NGHIỆM", "Dành cho Giáo viên chấm thi & tra cứu nhanh")

    # 1. Bảng 90 câu Phần I (chia thành các bảng 10 cột hoặc 15 cột)
    add_section_title(doc, "1. Đáp án Phần I: Trắc nghiệm 4 lựa chọn (Câu 1 - 90)")
    
    # 9 hàng x 10 cột
    cols = 10
    tbl = doc.add_table(rows=18, cols=cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl, color="B0BEC5", sz="4")

    for r_idx in range(9):
        # Hàng số câu: 2*r_idx
        row_q = tbl.rows[2 * r_idx]
        # Hàng đáp án: 2*r_idx + 1
        row_a = tbl.rows[2 * r_idx + 1]

        for c_idx in range(cols):
            q_num = r_idx * cols + c_idx + 1
            if q_num <= len(MC_QUESTIONS):
                q_obj = MC_QUESTIONS[q_num - 1]
                ans_val = q_obj["correct"]

                cell_q = row_q.cells[c_idx]
                cell_q.width = Cm(1.8)
                set_cell_margins(cell_q, top=30, bottom=30, left=20, right=20)
                set_cell_shading(cell_q, "E8F5E9")
                p_q = cell_q.paragraphs[0]
                p_q.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_q = p_q.add_run(f"C{q_num}")
                set_run_style(r_q, font_size=8.5, bold=True, color=COLOR_PRIMARY)

                cell_a = row_a.cells[c_idx]
                cell_a.width = Cm(1.8)
                set_cell_margins(cell_a, top=30, bottom=30, left=20, right=20)
                p_a = cell_a.paragraphs[0]
                p_a.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_a = p_a.add_run(ans_val)
                set_run_style(r_a, font_size=9, bold=True, color=COLOR_CORRECT)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(4)
    p_sp.paragraph_format.space_after = Pt(2)

    # 2. Bảng đáp án Phần II (Đúng - Sai)
    add_section_title(doc, "2. Đáp án Phần II: Trắc nghiệm Đúng - Sai")
    tbl_tf = doc.add_table(rows=3, cols=5)
    tbl_tf.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_tf.autofit = False
    set_table_borders(tbl_tf, color="B0BEC5", sz="4")

    headers_tf = ["Câu", "Ý a", "Ý b", "Ý c", "Ý d"]
    for i, h in enumerate(headers_tf):
        c = tbl_tf.rows[0].cells[i]
        c.width = Cm(3.6)
        set_cell_shading(c, "E8F5E9")
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_style(r, font_size=9.5, bold=True, color=COLOR_PRIMARY)

    for row_idx, q in enumerate(TF_QUESTIONS):
        row_cells = tbl_tf.rows[row_idx + 1].cells
        c0 = row_cells[0]
        c0.width = Cm(3.6)
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(f"Câu {q['id']}")
        set_run_style(r0, font_size=9.5, bold=True)

        for s_idx, st in enumerate(q["statements"]):
            c = row_cells[s_idx + 1]
            c.width = Cm(3.6)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            val_text = "ĐÚNG" if st["correct"] else "SAI"
            val_col = COLOR_CORRECT if st["correct"] else COLOR_WRONG
            r = p.add_run(val_text)
            set_run_style(r, font_size=9.5, bold=True, color=val_col)

    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(4)
    p_sp2.paragraph_format.space_after = Pt(2)

    # 3. Bảng đáp án Phần III (Trả lời ngắn)
    add_section_title(doc, "3. Đáp án Phần III: Trắc nghiệm Trả lời ngắn (Câu 1 - 38)")
    tbl_sh = doc.add_table(rows=8, cols=10)
    tbl_sh.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_sh.autofit = False
    set_table_borders(tbl_sh, color="B0BEC5", sz="4")

    # 4 cặp hàng x 10 cột (tổng tối đa 40 ô, cho 38 câu)
    for r_idx in range(4):
        row_q = tbl_sh.rows[2 * r_idx]
        row_a = tbl_sh.rows[2 * r_idx + 1]

        for c_idx in range(10):
            q_num = r_idx * 10 + c_idx + 1
            if q_num <= len(SHORT_QUESTIONS):
                q_obj = SHORT_QUESTIONS[q_num - 1]
                ans_val = q_obj["answer"]

                cell_q = row_q.cells[c_idx]
                cell_q.width = Cm(1.8)
                set_cell_margins(cell_q, top=30, bottom=30, left=20, right=20)
                set_cell_shading(cell_q, "E8F5E9")
                p_q = cell_q.paragraphs[0]
                p_q.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_q = p_q.add_run(f"C{q_num}")
                set_run_style(r_q, font_size=8.5, bold=True, color=COLOR_PRIMARY)

                cell_a = row_a.cells[c_idx]
                cell_a.width = Cm(1.8)
                set_cell_margins(cell_a, top=30, bottom=30, left=20, right=20)
                p_a = cell_a.paragraphs[0]
                p_a.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_a = p_a.add_run(ans_val)
                set_run_style(r_a, font_size=9, bold=True, color=COLOR_CORRECT)
            else:
                cell_q = row_q.cells[c_idx]
                cell_q.width = Cm(1.8)
                cell_a = row_a.cells[c_idx]
                cell_a.width = Cm(1.8)

def build_complete_docx(mode="student"):
    """
    Xây dựng toàn bộ tài liệu Word
    mode: 'student' | 'teacher' | 'screen'
    """
    is_teacher = (mode == "teacher")
    is_screen = (mode == "screen")
    base_size = 11.5 if is_screen else 10.5

    doc = docx.Document()
    setup_header_footer(doc)
    build_title_block(doc, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    # =========================================================================
    # A. PHẦN TRẮC NGHIỆM KHÁCH QUAN
    # =========================================================================
    add_major_section_header(
        doc,
        "A. PHẦN TRẮC NGHIỆM KHÁCH QUAN",
        "Gồm Phần I: Trắc nghiệm 4 lựa chọn, Phần II: Đúng/Sai, Phần III: Trả lời ngắn",
        base_size=base_size + 1.0
    )

    # -------------------------------------------------------------------------
    # Phần I. Câu trắc nghiệm nhiều lựa chọn
    # -------------------------------------------------------------------------
    add_section_title(
        doc,
        "Phần I. Câu trắc nghiệm nhiều lựa chọn: Mỗi câu hỏi, học sinh chọn một phương án đúng nhất (Câu 1 đến Câu 90).",
        base_size=base_size
    )
    render_mc_questions(doc, MC_QUESTIONS, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    # -------------------------------------------------------------------------
    # Phần II. Câu trắc nghiệm đúng - sai
    # -------------------------------------------------------------------------
    add_section_title(
        doc,
        "Phần II. Câu trắc nghiệm đúng - sai: Mỗi câu hỏi có 4 ý a, b, c, d - học sinh chọn đúng hoặc sai cho mỗi ý.",
        base_size=base_size
    )
    render_tf_questions(doc, TF_QUESTIONS, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    # -------------------------------------------------------------------------
    # Phần III. Câu trắc nghiệm trả lời ngắn
    # -------------------------------------------------------------------------
    add_section_title(
        doc,
        "Phần III. Câu trắc nghiệm trả lời ngắn: Mỗi câu hỏi, học sinh ghi câu trả lời không quá 4 kí tự (Câu 1 đến Câu 38).",
        base_size=base_size
    )
    render_short_questions(doc, SHORT_QUESTIONS, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    # =========================================================================
    # B. PHẦN TỰ LUẬN
    # =========================================================================
    add_major_section_header(
        doc,
        "B. PHẦN TỰ LUẬN",
        "Học sinh trình bày lời giải chi tiết, rõ ràng từng bước theo quy chuẩn khoa học (Câu 1 đến Câu 18)",
        base_size=base_size + 1.0
    )
    render_essay_questions(doc, ESSAY_QUESTIONS, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    # Lời chúc ở cuối bài
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(8)
    p_end.paragraph_format.space_after = Pt(4)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_wish = p_end.add_run("CHÚC CÁC EM ÔN TẬP TỐT! KIỂM TRA ĐẠT KẾT QUẢ CAO NHÉ!\n-----Hết-----")
    set_run_style(r_wish, font_size=base_size, bold=True, italic=True, color=COLOR_PRIMARY)

    # Đối với bản Giáo viên: Thêm bảng tổng hợp đáp án ở cuối
    if is_teacher:
        doc.add_page_break()
        add_teacher_summary_table(doc, base_size=base_size - 1.0)

    return doc

def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    os.makedirs(ONEDRIVE_DIR, exist_ok=True)

    files_info = [
        ("student", "De_Cuong_KHTN_7_Giua_HK1_MinhKhai_HocSinh.docx", "Bản Học sinh (đề thi sạch, có dòng kẻ làm bài)"),
        ("teacher", "De_Cuong_KHTN_7_Giua_HK1_MinhKhai_GiaoVien.docx", "Bản Giáo viên (đáp án đỏ, lời giải chi tiết, bảng tổng hợp đáp án)"),
        ("screen", "De_Cuong_KHTN_7_Giua_HK1_MinhKhai_DayTrenManHinh.docx", "Bản Dạy trên màn hình (cỡ chữ thoáng, chừa chỗ trống ghi chú tương tác)")
    ]

    print("=" * 70)
    print("BẮT ĐẦU XUẤT BẢN TRỌN BỘ 3 FILE WORD ĐỀ CƯƠNG KHTN 7 MINH KHAI")
    print("=" * 70)

    for mode, filename, desc in files_info:
        out_path = os.path.join(OUT_DIR, filename)
        print(f"\n[+] Đang tạo: {desc}...")
        doc = build_complete_docx(mode=mode)
        doc.save(out_path)
        print(f"    --> Đã lưu thành công tại: {out_path}")

        # Đồng bộ sang OneDrive của Thầy
        onedrive_path = os.path.join(ONEDRIVE_DIR, filename)
        shutil.copy2(out_path, onedrive_path)
        print(f"    --> Đã đồng bộ sang OneDrive: {onedrive_path}")

    print("\n" + "=" * 70)
    print("HOÀN THÀNH TẤT CẢ 3 BẢN WORD CHUẨN FORM ĐỀ KIỂM TRA!")
    print("=" * 70)

if __name__ == "__main__":
    main()
