# -*- coding: utf-8 -*-
"""
Script sinh tự động 2 file Word (.docx) cho ĐỀ SỐ 17 (Chương 6: HỢP CHẤT CARBONYL - CARBOXYLIC ACID):
1. Bản Học sinh (De_Kiem_Tra_Hoa_11_De17.docx)
2. Bản Giáo viên (De_Kiem_Tra_Hoa_11_De17_GiaoVien.docx)

Tuân thủ nghiêm ngặt Quy chuẩn Form Đề Kiểm Tra (GEMINI.md):
- Tuyệt đối KHÔNG dùng bảng (Table) trong nội dung câu hỏi/đáp án/lời giải.
- Dùng Paragraph + Tab stops chuẩn cho các phương án trắc nghiệm.
- Công thức hóa học xử lý triệt để chỉ số trên/dưới (~ cho sub, ^ cho sup).
- Header & Footer chuẩn với viền màu Sage Green (#8BC390) và chữ Coral Red (#D32F2F).
- Bản Giáo viên có ghi chú đỏ và lời giải thụt lề màu xanh #1565C0.
"""

import os
import sys

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH

# Đảm bảo import được module form_de_kiem_tra
sys.path.insert(0, os.path.dirname(__file__))
from form_de_kiem_tra import (
    setup_header_footer,
    build_title_block,
    add_section_header,
    render_question_prompt,
    render_options_no_table,
    render_statement,
    add_solution_no_table,
    add_student_writing_lines,
    render_rich_text,
    set_run_style,
    COLOR_CORRECT,
    COLOR_SOLUTION,
    COLOR_PRIMARY
)

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
IMG_DIR = os.path.join(OUT_DIR, "images_crop_hoa11_de17")

def add_p(doc, space_before=3, space_after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    return p

def add_centered_picture(doc, img_path, width_cm=8.0):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(img_path, width=Cm(width_cm))

def build_exam_document(is_teacher=False):
    doc = docx.Document()
    setup_header_footer(doc)

    build_title_block(
        doc,
        school_line="SỞ GIÁO DỤC VÀ ĐÀO TẠO THỪA THIÊN HUẾ",
        exam_title="ĐỀ KIỂM TRA ĐỊNH KỲ CHƯƠNG 6 (HỢP CHẤT CARBONYL - CARBOXYLIC ACID)",
        subject_line="MÔN: HOÁ HỌC - LỚP 11",
        num_pages="04",
        duration="50 phút",
        exam_code="117",
        atomic_masses="H = 1; C = 12; N = 14; O = 16; Na = 23; Mg = 24; Al = 27; P = 31; S = 32; Cl = 35,5; K = 39; Ca = 40; Cr = 52; Mn = 55; Fe = 56; Cu = 64; Zn = 65; Br = 80; Ag = 108; Ba = 137.",
        is_teacher=is_teacher
    )

    # =========================================================================
    # PHẦN I
    # =========================================================================
    add_section_header(
        doc,
        title="PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án."
    )

    # Câu 1
    p = add_p(doc)
    render_question_prompt(p, "Câu 1. Chất nào sau đây là acetic aldehyde?")
    opts_c1 = [
        "A. HCHO.",
        "*B. CH~3~CHO.*",
        "C. C~2~H~5~OH.",
        "D. C~2~H~6~."
    ]
    render_options_no_table(doc, opts_c1, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Acetic aldehyde (ethanal, acetaldehyde) có công thức cấu tạo là CH~3~CHO.\n"
            "=> Chọn B."
        )

    # Câu 2
    p = add_p(doc)
    render_question_prompt(p, "Câu 2. Tên theo danh pháp thay thế của chất CH~3~-CH~2~-CH~2~-CHO là gì?")
    opts_c2 = [
        "A. Propan-1-al.",
        "B. propanal.",
        "*C. butanal.*",
        "D. butyraldehyde."
    ]
    render_options_no_table(doc, opts_c2, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Mạch chính gồm 4 nguyên tử carbon, no, chứa nhóm chức -CHO ở đầu mạch => tên thay thế là butanal.\n"
            "=> Chọn C."
        )

    # Câu 3
    p = add_p(doc)
    render_question_prompt(p, "Câu 3. Benzaldehyde tác dụng với Cu(OH)~2~ trong môi trường kiềm (t°C) thu được kết tủa màu đỏ gạch là")
    opts_c3 = [
        "A. Cu.",
        "*B. Cu~2~O.*",
        "C. CuO.",
        "D. Cu~2~O, Cu."
    ]
    render_options_no_table(doc, opts_c3, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Aldehyde tác dụng với Cu(OH)~2~ trong môi trường kiềm khi đun nóng bị oxi hóa thành muối carboxylate, đồng thời Cu(II) bị khử thành kết tủa đỏ gạch Cu~2~O (copper(I) oxide).\n"
            "=> Chọn B."
        )

    # Câu 4
    p = add_p(doc)
    render_question_prompt(p, "Câu 4. Phản ứng giữa CH~3~CHO với HCN thuộc loại phản ứng")
    opts_c4 = [
        "A. phản ứng thuận nghịch.",
        "*B. phản ứng cộng.*",
        "C. phản ứng thế.",
        "D. phản ứng phân hủy."
    ]
    render_options_no_table(doc, opts_c4, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Phản ứng giữa acetaldehyde và hydrogen cyanide là phản ứng cộng nucleophile vào liên kết đôi C=O của nhóm carbonyl tạo cyanohydrin:\n"
            "CH~3~CHO + HCN -> CH~3~-CH(OH)-CN.\n"
            "=> Chọn B."
        )

    # Câu 5
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 5. [KNTT - SBT] Có bốn chất lỏng có thể tích bằng nhau là ethanol, acetone, acetaldehyde, acetic acid. "
        "Tiến hành chưng cất hỗn hợp này, sau một thời gian, hàm lượng chất nào trong bình chưng cất còn lại lớn nhất?"
    )
    opts_c5 = [
        "A. Ethanol.",
        "B. Acetone.",
        "C. Acetaldehyde.",
        "*D. Acetic acid.*"
    ]
    render_options_no_table(doc, opts_c5, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Nhiệt độ sôi của 4 chất: acetaldehyde (21°C) < acetone (56°C) < ethanol (78,3°C) < acetic acid (118°C).\n"
            "Chất có nhiệt độ sôi thấp hơn sẽ bay hơi trước. Acetic acid có nhiệt độ sôi cao nhất nên bay hơi chậm nhất, sau một thời gian chưng cất hàm lượng còn lại trong bình là lớn nhất.\n"
            "=> Chọn D."
        )

    # Câu 6
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 6. [KNTT - SBT] Cho ba chất lỏng riêng biệt sau: C~2~H~5~OH, CH~3~CHO, CH~3~COOH. "
        "Cách nào sau đây phù hợp để phân biệt ba chất lỏng trên?"
    )
    opts_c6 = [
        "A. Dùng quỳ tím, sau đó dùng dung dịch NaOH.",
        "*B. Dùng quỳ tím, sau đó dùng dung dịch AgNO~3~ trong NH~3~.*",
        "C. Dùng Na sau đó dùng dung dịch AgNO~3~ trong NH~3~.",
        "D. Dùng Na sau đó dùng quỳ tím."
    ]
    render_options_no_table(doc, opts_c6, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- Dùng quỳ tím: CH~3~COOH làm quỳ tím chuyển màu đỏ; C~2~H~5~OH và CH~3~CHO không đổi màu quỳ.\n"
            "- Dùng dung dịch AgNO~3~/NH~3~ (đun nóng): CH~3~CHO cho phản ứng tráng bạc tạo kết tủa trắng sáng bám vào thành ống nghiệm; C~2~H~5~OH không phản ứng.\n"
            "=> Chọn B."
        )

    # Câu 7
    p = add_p(doc)
    render_question_prompt(p, "Câu 7. (SBT- CD) Dãy nào sau đây gồm các chất đều tác dụng với dung dịch AgNO~3~/NH~3~?")
    opts_c7 = [
        "A. Acetaldehyde, but-1-yne, ethylene.",
        "B. Acetaldehyde, acetylene, but-2-yne.",
        "*C. Formaldehyde, vinylacetylene, propyne.*",
        "D. Formaldehyde, acetylene, ethylene."
    ]
    render_options_no_table(doc, opts_c7, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- Formaldehyde (HCHO) tham gia phản ứng tráng bạc tạo Ag.\n"
            "- Vinylacetylene (CH#C-CH=CH~2~) và propyne (CH#C-CH~3~) là các alkyne có liên kết ba đầu mạch (alk-1-yne) phản ứng tạo kết tủa vàng nhạt.\n"
            "=> Chọn C."
        )

    # Câu 8
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 8. (SBT- CD) Trong các chất sau: (1) CH~3~CH~2~CHO, (2) CH~3~CH(OH)CH~3~, (3) (CH~3~)~2~CHCHO, (4) CH~2~=CHCH~2~OH. "
        "Những chất nào phản ứng với H~2~ (Ni, t°) hoặc NaBH~4~ sinh ra cùng một sản phẩm?"
    )
    opts_c8 = [
        "*A. (1) và (4).*",
        "B. (2) và (4).",
        "C. (1) và (2).",
        "D. (3) và (4)."
    ]
    render_options_no_table(doc, opts_c8, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- (1) CH~3~CH~2~CHO bị khử tạo propan-1-ol: CH~3~CH~2~CH~2~OH.\n"
            "- (4) CH~2~=CH-CH~2~OH phản ứng cộng H~2~ (Ni, t°) tạo propan-1-ol: CH~3~CH~2~CH~2~OH.\n"
            "=> Cả hai đều sinh ra propan-1-ol.\n"
            "=> Chọn A."
        )

    # Câu 9
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 9. (SBT - CTST) Cho dãy các chất sau: HCHO, CH~3~COOH, CH~3~CHO, HCOOH, C~2~H~5~OH. "
        "Số chất trong dãy có thể tham gia phản ứng tráng bạc là"
    )
    opts_c9 = [
        "*A. 3.*",
        "B. 6.",
        "C. 4.",
        "D. 5."
    ]
    render_options_no_table(doc, opts_c9, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Các chất có phản ứng tráng bạc phải chứa nhóm chức aldehyde -CHO:\n"
            "1. HCHO\n"
            "2. CH~3~CHO\n"
            "3. HCOOH (formic acid chứa nhóm -CHO liên kết với -OH)\n"
            "=> Có 3 chất. Chọn A."
        )

    # Câu 10
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 10. Hai hợp chất hữu cơ X, Y có cùng công thức phân tử C~3~H~6~O~2~. Cả X và Y đều tác dụng với Na; "
        "X tác dụng được với NaHCO~3~ còn Y có khả năng tham gia phản ứng tráng bạc. Công thức cấu tạo của X và Y lần lượt là"
    )
    opts_c10 = [
        "A. C~2~H~5~COOH và HCOOC~2~H~5~.",
        "B. HCOOC~2~H~5~ và HOCH~2~COCH~3~.",
        "C. HCOOC~2~H~5~ và HOCH~2~CH~2~CHO.",
        "*D. C~2~H~5~COOH và CH~3~CH(OH)CHO.*"
    ]
    render_options_no_table(doc, opts_c10, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- X tác dụng NaHCO~3~ => X chứa nhóm carboxylic acid -COOH => X là C~2~H~5~COOH.\n"
            "- Y tác dụng Na và tráng bạc => Y vừa chứa nhóm -OH vừa chứa nhóm -CHO => Y là CH~3~CH(OH)CHO (hoặc HO-CH~2~CH~2~CHO).\n"
            "=> Chọn D."
        )

    # Câu 11
    p = add_p(doc)
    render_question_prompt(p, "Câu 11. Dãy gồm các chất có thể điều chế trực tiếp (bằng một phản ứng) tạo ra acetic acid là:")
    opts_c11 = [
        "A. CH~3~CHO, C~2~H~5~OH, C~2~H~5~COOCH~3~.",
        "B. CH~3~CHO, C~6~H~12~O~6~ (glucose), CH~3~OH.",
        "*C. CH~3~OH, C~2~H~5~OH, CH~3~CHO.*",
        "D. C~2~H~4~(OH)~2~, CH~3~OH, CH~3~CHO."
    ]
    render_options_no_table(doc, opts_c11, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- CH~3~OH + CO --(xt, t°)--> CH~3~COOH\n"
            "- C~2~H~5~OH + O~2~ --(men giấm)--> CH~3~COOH + H~2~O\n"
            "- 2 CH~3~CHO + O~2~ --(xt, t°)--> 2 CH~3~COOH\n"
            "=> Chọn C."
        )

    # Câu 12
    p = add_p(doc)
    render_question_prompt(p, "Câu 12. Để điều chế CH~3~CHO từ Al~4~C~3~ cần ít nhất bao nhiêu phương trình phản ứng?")
    opts_c12 = [
        "A. 2.",
        "*B. 3.*",
        "C. 4.",
        "D. 5."
    ]
    render_options_no_table(doc, opts_c12, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Cần ít nhất 3 phương trình phản ứng:\n"
            "(1) Al~4~C~3~ + 12H~2~O -> 4Al(OH)~3~ + 3CH~4~\n"
            "(2) 2CH~4~ --(1500°C, làm lạnh nhanh)--> C~2~H~2~ + 3H~2~\n"
            "(3) C~2~H~2~ + H~2~O --(HgSO~4~, 80°C)--> CH~3~CHO\n"
            "=> Chọn B."
        )

    # Câu 13
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 13. Cho dãy các chất: HCHO, CH~3~COOH, HCOONa, HCOOH, C~2~H~5~OH, HCOOCH~3~. "
        "Số chất trong dãy tham gia phản ứng tráng gương là"
    )
    opts_c13 = [
        "A. 3.",
        "B. 6.",
        "*C. 4.*",
        "D. 5."
    ]
    render_options_no_table(doc, opts_c13, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Các chất tham gia phản ứng tráng gương gồm 4 chất:\n"
            "1. HCHO (formaldehyde)\n"
            "2. HCOONa (sodium formate chứa nhóm H-C=O)\n"
            "3. HCOOH (formic acid chứa nhóm H-C=O)\n"
            "4. HCOOCH~3~ (methyl formate chứa nhóm H-C=O)\n"
            "=> Chọn C."
        )

    # Câu 14
    p = add_p(doc)
    render_question_prompt(p, "Câu 14. Aldehyde bị khử bằng NaBH~4~ hoặc LiAlH~4~ sinh ra")
    opts_c14 = [
        "A. alcohol bậc hai.",
        "*B. alcohol bậc một.*",
        "C. ketone.",
        "D. alcohol bậc ba."
    ]
    render_options_no_table(doc, opts_c14, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Khử aldehyde (R-CHO) bằng NaBH~4~ hoặc LiAlH~4~ thu được alcohol bậc một (R-CH~2~OH).\n"
            "=> Chọn B."
        )

    # Câu 15
    p = add_p(doc)
    render_question_prompt(p, "Câu 15. Cho sơ đồ phản ứng: C~2~H~5~OH -> X -> CH~3~COOH (mỗi mũi tên ứng với một phản ứng). Chất X là")
    opts_c15 = [
        "A. HCHO.",
        "B. C~2~H~5~CHO.",
        "C. CH~4~.",
        "*D. CH~3~CHO.*"
    ]
    render_options_no_table(doc, opts_c15, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "C~2~H~5~OH oxi hóa không hoàn toàn tạo acetaldehyde (CH~3~CHO), sau đó CH~3~CHO tiếp tục bị oxi hóa tạo acetic acid (CH~3~COOH).\n"
            "=> Chọn D."
        )

    # Câu 16
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 16. (SBT - CTST) Trung hoà 100 mL dung dịch carboxylic acid (X) nồng độ 0,1 M cần vừa đủ 16 gam dung dịch NaOH 5%, "
        "thu được 1,48 gam muối. Công thức cấu tạo của (X) là"
    )
    opts_c16 = [
        "*A. HOOC-CH~2~-COOH.*",
        "B. HOOC-COOH.",
        "C. CH~3~-COOH.",
        "D. HOOC-CH~2~-CH~2~-COOH."
    ]
    render_options_no_table(doc, opts_c16, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- n_X = 0,1 × 0,1 = 0,01 mol.\n"
            "- n_NaOH = (16 × 5%) / 40 = 0,02 mol.\n"
            "Tỉ lệ n_NaOH : n_X = 0,02 : 0,01 = 2 => X là dicarboxylic acid dạng HOOC-R-COOH.\n"
            "Phương trình: HOOC-R-COOH + 2NaOH -> NaOOC-R-COONa + 2H~2~O\n"
            "n_muối = n_X = 0,01 mol => M_muối = 1,48 / 0,01 = 148 g/mol.\n"
            "=> R + 2 × 67 = 148 => R = 14 (-CH~2~-).\n"
            "Vậy công thức cấu tạo của X là HOOC-CH~2~-COOH (malonic acid).\n"
            "=> Chọn A."
        )

    # Câu 17
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 17. (SBT - CTST) Cho dung dịch acetic acid có nồng độ x% tác dụng vừa đủ với dung dịch NaOH 20% "
        "thu được dung dịch muối sodium acetate có nồng độ 10,25%. Giá trị của x là"
    )
    opts_c17 = [
        "*A. 10%.*",
        "B. 15%.",
        "C. 18,67%.",
        "D. 20%."
    ]
    render_options_no_table(doc, opts_c17, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Phương trình: CH~3~COOH + NaOH -> CH~3~COONa + H~2~O\n"
            "Giả sử có a mol CH~3~COOH phản ứng:\n"
            "- m_dd NaOH = (40a / 20) × 100 = 200a gam.\n"
            "- m_CH~3~COONa = 82a gam => m_dd CH~3~COONa = (82a / 10,25) × 100 = 800a gam.\n"
            "- Bảo toàn khối lượng dung dịch: m_dd CH~3~COOH = 800a - 200a = 600a gam.\n"
            "- Nồng độ phần trăm: x% = (60a / 600a) × 100% = 10%.\n"
            "=> Chọn A."
        )

    # Câu 18
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 18. (SBT- CD) Acetic acid được tổng hợp từ nguồn khí than đá theo các phản ứng sau:\n"
        "(1) CO + 2H~2~ --(t°, xt)--> CH~3~OH\n"
        "(2) CH~3~OH + CO --(t°, xt)--> CH~3~COOH\n"
        "Giả sử hiệu suất của các phản ứng (1) và (2) đều đạt 90%. Để sản xuất 1000 lít acetic acid (D = 1,05 g/mL), "
        "cần thể tích khí CO và H~2~ (ở điều kiện chuẩn) lần lượt là"
    )
    opts_c18 = [
        "A. 964,06 m^3 và 1 928,12 m^3.",
        "B. 535,6 m^3 và 1 071,17 m^3.",
        "C. 964,06 m^3 và 964,06 m^3.",
        "*D. 1 017,6 m^3 và 1 071,2 m^3.*"
    ]
    render_options_no_table(doc, opts_c18, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Khối lượng acid cần sản xuất: m = 1 000 × 1,05 = 1 050 kg = 1 050 000 gam.\n"
            "Số mol acid: n_acid = 1 050 / 60 = 17,5 kmol.\n"
            "- Từ phản ứng (2): n_CH~3~OH cần dùng = 17,5 / 0,9 = 175/9 kmol; n_CO(2) = 175/9 kmol.\n"
            "- Từ phản ứng (1): n_CO(1) = (175/9) / 0,9 = 1750/81 kmol; n_H~2~ = 2 × n_CO(1) = 3500/81 kmol.\n"
            "- Tổng mol CO cần dùng: n_CO = 1750/81 + 175/9 = 3325/81 kmol.\n"
            "- Thể tích khí ở đkc (24,79 L/mol):\n"
            "  V_CO = (3325/81) × 24,79 = 1 017,6 m^3.\n"
            "  V_H~2~ = (3500/81) × 24,79 = 1 071,2 m^3.\n"
            "=> Chọn D."
        )


    # =========================================================================
    # PHẦN II
    # =========================================================================
    add_section_header(
        doc,
        title="PHẦN II. Câu trắc nghiệm đúng sai.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai."
    )

    # Câu 1 Phần II
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 1. Cinnamaldehyde (C~6~H~5~CH=CHCHO) là hợp chất tạo mùi đặc trưng của tinh dầu quế. Cho các nhận định sau đây về Cinnamaldehyde:"
    )
    stmts_c1_p2 = [
        ("a) Cinnamaldehyde có đồng phân hình học.", "ĐÚNG"),
        ("b) Cinnamaldehyde có khả năng tham gia phản ứng tráng bạc.", "ĐÚNG"),
        ("c) Ở điều kiện thường 1 mol Cinnamaldehyde có thể làm mất màu tối đa 1 mol bromine.", "SAI"),
        ("d) Cinnamaldehyde phản ứng với Na kim loại giải phóng khí hydrogen.", "SAI")
    ]
    for stmt, ans in stmts_c1_p2:
        p_st = add_p(doc, space_before=1, space_after=1)
        p_st.paragraph_format.left_indent = Cm(0.5)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- a) ĐÚNG: Cấu trúc -CH=CH- có 2 nhóm thế khác nhau ở mỗi carbon nối đôi nên có đồng phân cis-trans.\n"
            "- b) ĐÚNG: Phân tử chứa nhóm -CHO nên tráng bạc tốt với thuốc thử Tollens.\n"
            "- c) SAI: Có 1 liên kết đôi C=C (cộng Br~2~) và 1 nhóm -CHO (bị Br~2~ oxi hóa) nên 1 mol phản ứng tối đa với 2 mol Br~2~.\n"
            "- d) SAI: Cinnamaldehyde không có nguyên tử H linh động (không có nhóm -OH, -COOH) nên không phản ứng với Na."
        )

    # Câu 2 Phần II
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 2. Thực hiện thí nghiệm theo các bước như sau:\n"
        "Bước 1: Cho 2 mL dung dịch NaOH 10% và 2 mL dung dịch I~2~/KI vào ống nghiệm có chứa 0,5 mL CH~3~CHO 10%.\n"
        "Bước 2: Lắc nhẹ ống nghiệm, hơ nhẹ ống nghiệm trên ngọn lửa đèn cồn hoặc đặt vào cốc nước nóng, sau đó làm nguội."
    )
    stmts_c2_p2 = [
        ("a) Đây là thí nghiệm nghiên cứu phản ứng tạo iodoform.", "ĐÚNG"),
        ("b) Có thể thay CH~3~CHO bằng CH~3~COCH~3~ vẫn thu được hiện tượng tương tự.", "ĐÚNG"),
        ("c) Sau bước 2, thu được kết tủa màu trắng.", "SAI"),
        ("d) Trong phản ứng trên, I~2~ đóng vai trò là chất khử.", "SAI")
    ]
    for stmt, ans in stmts_c2_p2:
        p_st = add_p(doc, space_before=1, space_after=1)
        p_st.paragraph_format.left_indent = Cm(0.5)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- a) ĐÚNG: Đây là phản ứng tạo iodoform (CHI~3~) đặc trưng của hợp chất chứa nhóm methyl ketone / CH~3~CHO.\n"
            "- b) ĐÚNG: Acetone (CH~3~COCH~3~) cũng có nhóm CH~3~-C=O nên cho kết tủa vàng tương tự.\n"
            "- c) SAI: Iodoform (CHI~3~) là kết tủa màu vàng nhạt, không phải màu trắng.\n"
            "- d) SAI: I~2~ bị khử từ số oxi hóa 0 xuống -1 (I^-) nên I~2~ đóng vai trò là chất oxi hóa."
        )

    # Câu 3 Phần II
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 3. Malic acid là hợp chất hữu cơ có công thức cấu tạo HOOC-CH(OH)-CH~2~-COOH, là nguyên nhân chính gây nên vị chua của quả táo."
    )
    stmts_c3_p2 = [
        ("a) Cứ 1 mol malic acid phản ứng được với tối đa 2 mol NaHCO~3~.", "ĐÚNG"),
        ("b) Malic là hợp chất hữu cơ đa chức.", "SAI"),
        ("c) Công thức phân tử của Malic là C~5~H~8~O~4~.", "SAI"),
        ("d) Làm quỳ tím hóa đỏ.", "ĐÚNG")
    ]
    for stmt, ans in stmts_c3_p2:
        p_st = add_p(doc, space_before=1, space_after=1)
        p_st.paragraph_format.left_indent = Cm(0.5)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- a) ĐÚNG: Malic acid có 2 nhóm -COOH, mỗi nhóm tác dụng với 1 phân tử NaHCO~3~ giải phóng CO~2~.\n"
            "- b) SAI: Malic acid vừa có nhóm -OH vừa có nhóm -COOH nên là hợp chất hữu cơ tạp chức (không phải đa chức).\n"
            "- c) SAI: Công thức phân tử của malic acid là C~4~H~6~O~5~.\n"
            "- d) ĐÚNG: Có nhóm carboxyl -COOH có tính acid làm quỳ tím chuyển màu đỏ."
        )

    # Câu 4 Phần II
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 4. CH~3~(CH~2~)~3~COOH (chất A) là chất lỏng không màu, có mùi khó chịu. "
        "Đun nóng A với methanol (CH~3~OH) có mặt của chất xúc tác acid sau một vài giờ thu được hỗn hợp chứa A, "
        "methanol và một sản phẩm hữu cơ CH~3~(CH~2~)~3~COOCH~3~ (chất V) có mùi trái cây dễ chịu. "
        "Một hợp chất hữu cơ được xác định bằng máy quang phổ IR như sau:"
    )
    add_centered_picture(doc, os.path.join(IMG_DIR, "ir_spectrum_pentanoic.png"), width_cm=11.0)
    stmts_c4_p2 = [
        ("a) Dựa vào quang phổ IR xác định có nhóm chức -COOH của (A).", "ĐÚNG"),
        ("b) Phổ MS sẽ xác định được khối lượng của chất V thu được từ thí nghiệm trên.", "SAI"),
        ("c) Phản ứng của chất A với methanol tạo thành chất V là phản ứng thuận nghịch.", "ĐÚNG"),
        ("d) Dựa vào quang phổ IR có thể xác định được công thức phân tử của một hợp chất hữu cơ.", "SAI")
    ]
    for stmt, ans in stmts_c4_p2:
        p_st = add_p(doc, space_before=1, space_after=1)
        p_st.paragraph_format.left_indent = Cm(0.5)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- a) ĐÚNG: Phổ IR có dải hấp thụ rộng mạnh của liên kết O-H acid (2500 - 3300 cm^-1) và peak C=O (~1710 cm^-1) đặc trưng cho nhóm -COOH của acid A.\n"
            "- b) SAI: Phổ MS (khối phổ) xác định phân tử khối và mảnh ion của phân tử, không xác định khối lượng gam chất thu được trong phản ứng.\n"
            "- c) ĐÚNG: Phản ứng ester hóa là phản ứng thuận nghịch.\n"
            "- d) SAI: Phổ hồng ngoại IR chỉ dùng để xác định các nhóm chức trong phân tử, không xác định được công thức phân tử."
        )


    # =========================================================================
    # PHẦN III
    # =========================================================================
    add_section_header(
        doc,
        title="PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 6."
    )

    # Câu 1 Phần III
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 1. Cho các chất sau: CH~3~-CH~2~-CHO (1), CH~2~=CH-CHO (2), (CH~3~)~2~CH-CHO (3), CH~2~=CH-CH~2~-OH (4). "
        "Có bao nhiêu chất phản ứng hoàn toàn với lượng dư [H] cùng tạo ra CH~3~-CH~2~-CH~2~-OH?"
    )
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Các chất khi bị khử bởi lượng dư [H] tạo thành propan-1-ol (CH~3~CH~2~CH~2~OH) gồm:\n"
            "- (1) CH~3~CH~2~CHO + 2[H] -> CH~3~CH~2~CH~2~OH\n"
            "- (2) CH~2~=CH-CHO + 4[H] -> CH~3~CH~2~CH~2~OH\n"
            "- (4) CH~2~=CH-CH~2~OH + 2[H] -> CH~3~CH~2~CH~2~OH\n"
            "=> Có 3 chất: (1), (2), (4).\n"
            "=> Đáp số: 3."
        )

    # Câu 2 Phần III
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 2. Cho các tác nhân phản ứng:\n"
        "(a) H~2~/Pd, t°\n"
        "(b) LiAlH~4~\n"
        "(c) NaBH~4~\n"
        "(d) Dung dịch bromine\n"
        "Có bao nhiêu tác nhân ở trên có thể chuyển hóa RCHO thành RCH~2~OH?"
    )
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Chuyển hóa RCHO thành RCH~2~OH là phản ứng khử aldehyde thành alcohol bậc một:\n"
            "- (a) H~2~/Pd, t°: chất khử hydro hóa nhóm -CHO thành -CH~2~OH.\n"
            "- (b) LiAlH~4~: chất khử mạnh khử aldehyde thành alcohol bậc một.\n"
            "- (c) NaBH~4~: chất khử chọn lọc khử aldehyde thành alcohol bậc một.\n"
            "- (d) Dung dịch bromine: là chất oxi hóa chuyển aldehyde thành carboxylic acid RCOOH.\n"
            "=> Có 3 tác nhân: (a), (b), (c).\n"
            "=> Đáp số: 3."
        )

    # Câu 3 Phần III
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 3. Cho các phát biểu sau:\n"
        "(1) Aldehyde phản ứng được với nước bromine.\n"
        "(2) Ketone không phản ứng được với Cu(OH)~2~/OH^-.\n"
        "(3) Aldehyde tác dụng với dung dịch AgNO~3~/NH~3~ tạo ra bạc.\n"
        "(4) Trong các hợp chất carbonyl, chỉ aldehyde bị khử bởi NaBH~4~.\n"
        "Có bao nhiêu phát biểu đúng?"
    )
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- (1) Đúng: Aldehyde bị oxi hóa bởi nước bromine tạo carboxylic acid.\n"
            "- (2) Đúng: Ketone không có tính khử đối với Cu(OH)~2~/OH^- đun nóng.\n"
            "- (3) Đúng: Aldehyde tráng bạc tạo kim loại Ag.\n"
            "- (4) Sai: Cả aldehyde và ketone đều bị khử bởi NaBH~4~.\n"
            "=> Có 3 phát biểu đúng là (1), (2), (3).\n"
            "=> Đáp số: 3."
        )

    # Câu 4 Phần III
    p = add_p(doc)
    render_question_prompt(p, "Câu 4. Để điều chế CH~3~CHO từ Al~4~C~3~ cần ít nhất bao nhiêu phương trình phản ứng?")
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Chuỗi 3 phương trình phản ứng ngắn nhất:\n"
            "(1) Al~4~C~3~ + 12H~2~O -> 4Al(OH)~3~ + 3CH~4~\n"
            "(2) 2CH~4~ --(1500°C, lln)--> C~2~H~2~ + 3H~2~\n"
            "(3) C~2~H~2~ + H~2~O --(HgSO~4~, 80°C)--> CH~3~CHO\n"
            "=> Đáp số: 3."
        )

    # Câu 5 Phần III
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 5. Quế có vị cay, mùi thơm nồng, được sử dụng phổ biến làm gia vị, vị thuốc trong Đông y. "
        "Hợp chất hữu cơ cinnamaldehyde tạo mùi đặc trưng của quế. Trong phân tử cinnamaldehyde chứa vòng benzene có một nhóm thế. "
        "Cinnamaldehyde bị oxi hóa bởi thuốc thử Tollens và có đồng phân hình học dạng trans. "
        "Công thức cấu tạo của cinnamaldehyde là:"
    )
    add_centered_picture(doc, os.path.join(IMG_DIR, "cinnamaldehyde_trans.png"), width_cm=4.5)
    p_sub = add_p(doc)
    render_rich_text(p_sub, "Phân tử khối của cinnamaldehyde là bao nhiêu?")
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Công thức cấu tạo của cinnamaldehyde: C~6~H~5~-CH=CH-CHO.\n"
            "Công thức phân tử: C~9~H~8~O.\n"
            "Phân tử khối: M = 9 × 12 + 8 × 1 + 16 = 132 g/mol.\n"
            "=> Đáp số: 132."
        )

    # Câu 6 Phần III
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 6. Hợp chất hữu cơ X được dùng nhiều trong tổng hợp hữu cơ. Bằng phương pháp phân tích nguyên tố, "
        "người ta xác định được X chứa 62,07% C; 10,34% H; còn lại là O. "
        "Trên phổ MS của X, người ta thấy có peak ion phân tử [M^+] có giá trị m/z bằng 58. "
        "Trên phổ IR của X có một peak trong vùng 2380 – 2695 cm^-1 và một peak trong vùng 1740 – 1685 cm^-1. "
        "Chất X phản ứng với Cu(OH)~2~ trong môi trường kiềm để tạo ra kết tủa màu đỏ gạch. "
        "Xác định tổng số nhóm CH~2~ trong công thức cấu tạo của X."
    )
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- Đặt công thức phân tử của X là C~x~H~y~O~z~.\n"
            "- %O = 100% - 62,07% - 10,34% = 27,59%.\n"
            "- Phân tử khối M_X = 58.\n"
            "  x = (62,07 × 58) / (12 × 100) = 3\n"
            "  y = (10,34 × 58) / (1 × 100) = 6\n"
            "  z = (27,59 × 58) / (16 × 100) = 1\n"
            "  => Công thức phân tử là C~3~H~6~O.\n"
            "- Phổ IR có peak 2380 - 2695 cm^-1 (C-H aldehyde) và 1740 - 1685 cm^-1 (C=O aldehyde); X phản ứng với Cu(OH)~2~/OH^- tạo kết tủa đỏ gạch Cu~2~O => X là propanal: CH~3~-CH~2~-CHO.\n"
            "- Trong công thức CH~3~-CH~2~-CHO có đúng 1 nhóm -CH~2~-.\n"
            "=> Đáp số: 1."
        )


    # =========================================================================
    # PHẦN IV
    # =========================================================================
    add_section_header(
        doc,
        title="PHẦN IV. Câu hỏi tự luận.",
        subtitle="Thí sinh trình bày lời giải chi tiết vào giấy làm bài."
    )

    # Câu 1 Phần IV
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 1. (SBT - CTST) Trong các phản ứng sau đây, phản ứng nào không thể xảy ra? "
        "Hoàn thành phương trình hoá học các phản ứng xảy ra.\n"
        "a) CH~3~CH~2~-CO-CH~3~ + [O] ->\n"
        "b) CH~3~CH~2~-OH + [O] ->\n"
        "c) C~6~H~6~ + Cl~2~ --(FeBr~3~, t°)-->\n"
        "d) (CH~3~)~2~C(OH)-CH~2~CH~3~ + [O] ->"
    )
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- Phản ứng không thể xảy ra:\n"
            "  + a) CH~3~CH~2~-CO-CH~3~ là ketone no, không bị oxi hóa bởi các tác nhân oxi hóa thông thường.\n"
            "  + d) (CH~3~)~2~C(OH)-CH~2~CH~3~ là alcohol bậc ba, không bị oxi hóa.\n"
            "- Hoàn thành phương trình hoá học các phản ứng xảy ra:\n"
            "  b) CH~3~CH~2~-OH + O~2~ --(men giấm)--> CH~3~COOH + H~2~O\n"
            "     (hoặc: CH~3~CH~2~OH + CuO --(t°)--> CH~3~CHO + Cu + H~2~O)\n"
            "  c) C~6~H~6~ + Cl~2~ --(FeBr~3~, t°)--> C~6~H~5~Cl + HCl"
        )
        add_centered_picture(doc, os.path.join(IMG_DIR, "chlorobenzene_solution.png"), width_cm=8.5)
    else:
        add_student_writing_lines(doc, num_lines=8)

    # Câu 2 Phần IV
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 2. Cho thí nghiệm: Chuẩn độ acetic acid bằng dung dịch sodium hydroxide 0,15 M. "
        "Sau khi thực hiện thí nghiệm chuẩn độ này 3 lần, thu được bảng kết quả dưới đây:"
    )
    add_centered_picture(doc, os.path.join(IMG_DIR, "table_titration_de17.png"), width_cm=11.5)
    p_sub = add_p(doc)
    render_rich_text(p_sub, "Xác định nồng độ mol của dung dịch acetic acid.")
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "1. Thể tích trung bình của dung dịch NaOH qua 3 lần đo:\n"
            "   V_NaOH = (20,0 + 19,9 + 20,0) / 3 ≈ 20,0 mL.\n"
            "2. Thể tích mẫu acetic acid: V_acid = 6,0 mL.\n"
            "3. Phương trình phản ứng chuẩn độ:\n"
            "   CH~3~COOH + NaOH -> CH~3~COONa + H~2~O\n"
            "4. Áp dụng định luật đương lượng:\n"
            "   V_acid × C_acid = V_NaOH × C_NaOH\n"
            "   => C_acid = (V_NaOH × C_NaOH) / V_acid = (20,0 × 0,15) / 6,0 = 0,5 M.\n"
            "=> Đáp số: 0,5 M."
        )
    else:
        add_student_writing_lines(doc, num_lines=8)

    # Dòng kết thúc
    p_end = add_p(doc, space_before=6, space_after=6)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_end = p_end.add_run("---------- HẾT ----------")
    set_run_style(run_end, bold=True, color=COLOR_PRIMARY)

    # KIỂM TRA QUY TẮC: Tuyệt đối không dùng Table trong nội dung (chỉ có bảng Title Block ở đầu)
    assert len(doc.tables) <= 1, f"LỖI QUY CHUẨN: Có {len(doc.tables)} bảng trong tài liệu!"

    return doc

def main():
    os.makedirs(OUT_DIR, exist_ok=True)

    # 1. Sinh bản Học sinh
    doc_hs = build_exam_document(is_teacher=False)
    hs_path = os.path.join(OUT_DIR, "De_Kiem_Tra_Hoa_11_De17.docx")
    doc_hs.save(hs_path)
    print(f"[OK] Đã tạo bản Học sinh: {hs_path}")

    # 2. Sinh bản Giáo viên
    doc_gv = build_exam_document(is_teacher=True)
    gv_path = os.path.join(OUT_DIR, "De_Kiem_Tra_Hoa_11_De17_GiaoVien.docx")
    doc_gv.save(gv_path)
    print(f"[OK] Đã tạo bản Giáo viên: {gv_path}")

if __name__ == "__main__":
    main()
