import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

CROP_DIR = r"c:\Antigravity_Thanh\San_Pham\images_crop"
OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"

COLOR_HF_RED = RGBColor(211, 47, 47)      # Coral Red (#D32F2F)
COLOR_HF_GREEN_HEX = "8BC390"             # Sage Green for horizontal lines (#8BC390)

COLOR_PRIMARY = RGBColor(46, 125, 50)     # Dark Green / Teal for headers in body
COLOR_CORRECT = RGBColor(198, 40, 40)     # Red for correct answers
COLOR_SOLUTION = RGBColor(21, 101, 192)   # Blue for solutions
COLOR_GRAY = RGBColor(97, 97, 97)

def add_p_border_bottom(p, color="8BC390", sz="8"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_p_border_top(p, color="8BC390", sz="8"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def set_run_style(run, font_name="Times New Roman", font_size=12, bold=False, italic=False, subscript=False, superscript=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def render_rich_text(p, text, base_font="Times New Roman", base_size=12, base_bold=False, base_italic=False, base_color=None):
    tokens = re.split(r'(~[^~]+~|\^[^\^]+\^|\*[^\*]+\*)', text)
    for token in tokens:
        if not token:
            continue
        is_bold = base_bold
        is_italic = base_italic
        is_sub = False
        is_sup = False
        content = token

        if token.startswith('*') and token.endswith('*'):
            is_bold = True
            content = token[1:-1]
        elif token.startswith('~') and token.endswith('~'):
            is_sub = True
            content = token[1:-1]
        elif token.startswith('^') and token.endswith('^'):
            is_sup = True
            content = token[1:-1]

        run = p.add_run(content)
        set_run_style(run, font_name=base_font, font_size=base_size, bold=is_bold, italic=is_italic,
                      subscript=is_sub, superscript=is_sup, color=base_color)

def setup_header_footer(doc):
    normal_style = doc.styles['Normal']
    normal_style.paragraph_format.tab_stops.clear_all()
    for pos in [0.5, 1.5, 5.0, 9.5, 14.0]:
        normal_style.paragraph_format.tab_stops.add_tab_stop(Cm(pos), WD_TAB_ALIGNMENT.LEFT)

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    section.start_type = WD_SECTION_START.NEW_PAGE
    section.different_first_page_header_footer = False
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    sectPr = section._sectPr
    header = section.header
    for p in header.paragraphs:
        p.text = ""
    tbl_h = header.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl_h.cell(0, 0)
    c_right = tbl_h.cell(0, 1)
    c_left.width = Cm(8.0)
    c_right.width = Cm(10.0)

    p_hl = c_left.paragraphs[0]
    p_hl.paragraph_format.space_before = Pt(0)
    p_hl.paragraph_format.space_after = Pt(2)
    p_hl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_hl = p_hl.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)

    p_hr = c_right.paragraphs[0]
    p_hr.paragraph_format.space_before = Pt(0)
    p_hr.paragraph_format.space_after = Pt(2)
    p_hr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_hr1 = p_hr.add_run("Thầy ")
    set_run_style(r_hr1, font_size=10, italic=True, color=COLOR_HF_RED)
    r_hr2 = p_hr.add_run("TRẦN VĂN THẠNH")
    set_run_style(r_hr2, font_size=10, bold=True, italic=True, color=COLOR_HF_RED)
    r_hr3 = p_hr.add_run(" - 0777.470.803")
    set_run_style(r_hr3, font_size=10, italic=True, color=COLOR_HF_RED)

    add_p_border_bottom(p_hl, color=COLOR_HF_GREEN_HEX, sz="8")
    add_p_border_bottom(p_hr, color=COLOR_HF_GREEN_HEX, sz="8")

    footer = section.footer
    for p in footer.paragraphs:
        p.text = ""
    tbl_f = footer.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    cf_left = tbl_f.cell(0, 0)
    cf_right = tbl_f.cell(0, 1)
    cf_left.width = Cm(14.5)
    cf_right.width = Cm(3.5)

    p_fl = cf_left.paragraphs[0]
    p_fl.paragraph_format.space_before = Pt(2)
    p_fl.paragraph_format.space_after = Pt(0)
    p_fl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_fl = p_fl.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)

    p_fr = cf_right.paragraphs[0]
    p_fr.paragraph_format.space_before = Pt(2)
    p_fr.paragraph_format.space_after = Pt(0)
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_field(p_fr, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED)

    add_p_border_top(p_fl, color=COLOR_HF_GREEN_HEX, sz="8")
    add_p_border_top(p_fr, color=COLOR_HF_GREEN_HEX, sz="8")

def format_cell_borders(cell, top="none", bottom="none", left="none", right="none", color="auto", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    xml = f'<w:tcBorders {nsdecls("w")}>'
    for side, val in [("top", top), ("left", left), ("bottom", bottom), ("right", right)]:
        if val != "none":
            xml += f'<w:{side} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        else:
            xml += f'<w:{side} w:val="none"/>'
    xml += '</w:tcBorders>'
    tcPr.append(parse_xml(xml))

def build_title_block(doc):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    c0 = tbl.cell(0, 0)
    c1 = tbl.cell(0, 1)
    c0.width = Cm(7.5)
    c1.width = Cm(10.5)

    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(1)
    p0.paragraph_format.line_spacing = 1.15
    render_rich_text(p0, "*SỞ GD & ĐT THÀNH PHỐ HUẾ*\n*TRƯỜNG THPT NGUYỄN CHÍ THANH*\n*ĐỀ CHÍNH THỨC*\n*(Đề thi có 04 trang)*", base_size=10)

    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(1)
    p1.paragraph_format.line_spacing = 1.15
    render_rich_text(p1, "*ĐỀ KIỂM TRA GIỮA HỌC KÌ 1 - NĂM HỌC 2025–2026*\n*MÔN: HÓA HỌC – 12CTST*\n_Thời gian làm bài: 45 phút (không kể thời gian phát đề)_", base_size=10)

    for c in [c0, c1]:
        format_cell_borders(c, "none", "none", "none", "none")

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(4)
    p_info.paragraph_format.space_after = Pt(2)
    r_info = p_info.add_run("Họ và tên: ...................................................................... Số báo danh: .............   Mã đề 104")
    set_run_style(r_info, font_size=11, bold=True)

    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(1)
    p_note.paragraph_format.space_after = Pt(6)
    r_note = p_note.add_run("Cho biết nguyên tử khối của các nguyên tố (amu): H=1; O=16; C=12; Ag=108; Na=23; N=14.")
    set_run_style(r_note, font_size=10.5, italic=True)

    add_p_border_bottom(p_note, color="CCCCCC", sz="4")

def add_section_header(doc, title_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"■ {title_text}")
    set_run_style(r, font_size=11.5, bold=True, color=COLOR_PRIMARY)

def add_solution_box(doc, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    c = tbl.cell(0, 0)
    c.width = Cm(18.0)
    
    tcPr = c._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F4F9F4"/>')
    tcPr.append(shd)
    
    format_cell_borders(c, left="single", color=COLOR_HF_GREEN_HEX, sz="18")
    
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.2)
    p.paragraph_format.line_spacing = 1.15
    
    r_lbl = p.add_run("Lời giải chi tiết:\n")
    set_run_style(r_lbl, font_size=10.5, bold=True, italic=True, color=COLOR_SOLUTION)
    
    render_rich_text(p, text, base_size=10.5, base_italic=False, base_color=RGBColor(50, 50, 50))

def add_mcq_choices_grid(doc, choices, correct_idx=-1, is_teacher=False):
    max_len = max(len(c) for c in choices)
    if max_len <= 18:
        tbl = doc.add_table(rows=1, cols=4)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        widths = [Cm(4.5), Cm(4.5), Cm(4.5), Cm(4.5)]
        for i, (ch, w) in enumerate(zip(choices, widths)):
            cell = tbl.cell(0, i)
            cell.width = w
            format_cell_borders(cell, "none", "none", "none", "none")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            prefix = f"{chr(65+i)}. "
            is_corr = (i == correct_idx) and is_teacher
            r_pre = p.add_run(prefix)
            set_run_style(r_pre, font_size=11, bold=is_corr, color=COLOR_CORRECT if is_corr else None)
            render_rich_text(p, ch, base_size=11, base_bold=is_corr, base_color=COLOR_CORRECT if is_corr else None)
    elif max_len <= 45:
        tbl = doc.add_table(rows=2, cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        widths = [Cm(9.0), Cm(9.0)]
        for i in range(4):
            r_idx = i // 2
            c_idx = i % 2
            cell = tbl.cell(r_idx, c_idx)
            cell.width = widths[c_idx]
            format_cell_borders(cell, "none", "none", "none", "none")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            prefix = f"{chr(65+i)}. "
            is_corr = (i == correct_idx) and is_teacher
            r_pre = p.add_run(prefix)
            set_run_style(r_pre, font_size=11, bold=is_corr, color=COLOR_CORRECT if is_corr else None)
            render_rich_text(p, choices[i], base_size=11, base_bold=is_corr, base_color=COLOR_CORRECT if is_corr else None)
    else:
        for i, ch in enumerate(choices):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Cm(0.5)
            prefix = f"{chr(65+i)}. "
            is_corr = (i == correct_idx) and is_teacher
            r_pre = p.add_run(prefix)
            set_run_style(r_pre, font_size=11, bold=is_corr, color=COLOR_CORRECT if is_corr else None)
            render_rich_text(p, ch, base_size=11, base_bold=is_corr, base_color=COLOR_CORRECT if is_corr else None)

def add_image_centered(doc, img_name, width_cm=10.0):
    img_path = os.path.join(CROP_DIR, img_name)
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(img_path, width=Cm(width_cm))

def build_exam_document(is_teacher=False):
    doc = docx.Document()
    setup_header_footer(doc)
    build_title_block(doc)

    # ==================== PHẦN TRẮC NGHIỆM ====================
    add_section_header(doc, "A. PHẦN TRẮC NGHIỆM (7 điểm)")
    
    # Phần I
    p_p1 = doc.add_paragraph()
    p_p1.paragraph_format.space_before = Pt(4)
    p_p1.paragraph_format.space_after = Pt(4)
    render_rich_text(p_p1, "*Phần I. Câu trắc nghiệm nhiều phương án lựa chọn.* (Thí sinh trả lời từ câu 1 đến câu 16. Mỗi câu hỏi thí sinh chỉ chọn một phương án).", base_size=11, base_italic=True)

    # Q1
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 1.* Phản ứng nào dưới đây *không* thể hiện tính base của amine?")
    q1_ch = [
        "C~6~H~5~NH~2~ + HCl → C~6~H~5~NH~3~Cl.",
        "RNH~2~ + HNO~2~ → ROH + N~2~ ↑ + H~2~O.",
        "Fe^3+^ + 3RNH~2~ + 3H~2~O → Fe(OH)~3~ ↓ + 3RNH~3~^+^.",
        "RNH~2~ + H~2~O ⇌ RNH~3~^+^ + OH^-^."
    ]
    add_mcq_choices_grid(doc, q1_ch, correct_idx=1, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Phản ứng của amine bậc 1 với acid nitrous HNO~2~ sinh ra alcohol ROH và khí N~2~ là phản ứng thế nhóm amino bằng nhóm -OH, *không* thể hiện tính base của amine.\n*Chọn B.*")

    # Q2
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 2.* Nhận xét nào dưới đây là *không đúng* khi nói về glucose và fructose?")
    q2_ch = [
        "Đều tạo được kết tủa đỏ gạch Cu~2~O khi tác dụng với Cu(OH)~2~, đun nóng trong môi trường kiềm.",
        "Đều xảy ra phản ứng tráng bạc khi tác dụng với thuốc thử Tollens.",
        "Đều tạo được dung dịch màu xanh lam khi tác dụng với Cu(OH)~2~ trong môi trường kiềm.",
        "Đều làm mất màu nước bromine."
    ]
    add_mcq_choices_grid(doc, q2_ch, correct_idx=3, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Fructose không làm mất màu nước bromine vì fructose không có nhóm aldehyde (-CHO), và trong môi trường nước bromine có tính acid yếu nên không có sự chuyển hóa fructose thành glucose.\n*Chọn D.*")

    # Q3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 3.* Khi đun nóng chất X có công thức phân tử C~3~H~6~O~2~ với dung dịch NaOH thu được CH~3~COONa. Công thức cấu tạo của X là:")
    q3_ch = [
        "CH~3~COOC~2~H~5~",
        "CH~3~COOCH~3~",
        "HCOOC~2~H~5~",
        "C~2~H~5~COOH"
    ]
    add_mcq_choices_grid(doc, q3_ch, correct_idx=1, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Muối thu được là CH~3~COONa (chứa gốc CH~3~COO- 2C). Do X có 3 carbon nên gốc alcohol là -CH~3~. Vậy CTCT của X là CH~3~COOCH~3~ (methyl acetate).\n*Chọn B.*")

    # Q4
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 4.* Trong cây thuốc lá tự nhiên và khói thuốc lá chứa một amine rất độc, đó là nicotin với công thức cấu tạo như sau:")
    add_image_centered(doc, "fig_cau4_nicotin.png", width_cm=14.0)
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(1)
    p_sub.paragraph_format.space_after = Pt(2)
    render_rich_text(p_sub, "Nicotin làm tăng huyết áp và nhịp tim, có khả năng gây xơ vữa động mạch vành và suy giảm trí nhớ. Số nguyên tử carbon trong một phân tử nicotin là:")
    q4_ch = ["11.", "8.", "9.", "10."]
    add_mcq_choices_grid(doc, q4_ch, correct_idx=3, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Vòng pyridine có 5 nguyên tử C; vòng pyrrolidine có 4 nguyên tử C trong vòng và 1 nguyên tử C ở nhóm methyl (-CH~3~) liên kết với N.\nTổng số carbon trong phân tử nicotin = 5 + 4 + 1 = 10 (CTPT là C~10~H~14~N~2~).\n*Chọn D.*")

    # Q5
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 5.* Chất giặt rửa tổng hợp sodium laurylsulfate có công thức cấu tạo như sau:")
    add_image_centered(doc, "fig_cau5_sodium_laurylsulfate.png", width_cm=9.5)
    p_sub5 = doc.add_paragraph()
    p_sub5.paragraph_format.space_before = Pt(1)
    p_sub5.paragraph_format.space_after = Pt(2)
    render_rich_text(p_sub5, "Nhóm được khoanh tròn trong công thức trên là")
    q5_ch = ["đầu kị nước.", "đuôi ưa nước.", "đầu ưa nước.", "đuôi kị nước."]
    add_mcq_choices_grid(doc, q5_ch, correct_idx=2, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Phân tử sodium laurylsulfate gồm đuôi kị nước (mạch hydrocarbon dài CH~3~[CH~2~]~10~CH~2~-) và đầu ưa nước là nhóm ion sulfate mang điện tích âm -OSO~3~^-^Na^+^ (được khoanh tròn).\n*Chọn C.*")

    # Q6
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 6.* Chất nào sau đây là amine bậc hai?")
    q6_ch = [
        "(C~2~H~5~)~3~N.",
        "CH~3~CH(NH~2~)CH~3~.",
        "C~2~H~5~NH~2~.",
        "(C~2~H~5~)~2~NH."
    ]
    add_mcq_choices_grid(doc, q6_ch, correct_idx=3, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Amine bậc hai có 2 nguyên tử H trong phân tử NH~3~ được thay thế bằng 2 gốc hydrocarbon: (C~2~H~5~)~2~NH (diethylamine).\n*Chọn D.*")

    # Q7
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 7.* Bệnh nhân phải tiếp đường (truyền dung dịch đường vào tĩnh mạch), đó là loại đường nào?")
    q7_ch = ["Glucose.", "Cellulose.", "Saccharose.", "Fructose."]
    add_mcq_choices_grid(doc, q7_ch, correct_idx=0, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Dung dịch glucose 5% (huyết thanh ngọt) được dùng để truyền trực tiếp vào tĩnh mạch giúp cung cấp nhanh năng lượng cho người bệnh hoặc cơ thể bị suy nhược.\n*Chọn A.*")

    # Q8
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 8.* HCOOCH~3~ được sử dụng làm dung môi cho nitrocellulose và cellulose acetate, một chất trung gian trong sản xuất dược phẩm và thuốc xông hơi. Tên gọi của HCOOCH~3~ là:")
    q8_ch = ["methyl acetate.", "methyl formate.", "ethyl acetate.", "ethyl formate."]
    add_mcq_choices_grid(doc, q8_ch, correct_idx=1, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "HCOOCH~3~ có gốc alcohol là methyl (-CH~3~) và gốc acid là formate (HCOO-) → Tên gọi: methyl formate.\n*Chọn B.*")

    # Q9
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 9.* Xà phòng hoá hoàn toàn 17,24 gam chất béo cần vừa đủ 0,06 mol NaOH. Cô cạn dung dịch sau phản ứng thu được bao nhiêu gam xà phòng?")
    q9_ch = ["21,15 gam.", "14,12 gam.", "14,84 gam.", "17,8 gam."]
    add_mcq_choices_grid(doc, q9_ch, correct_idx=3, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Phương trình tổng quát: Chất béo + 3NaOH → Xà phòng + Glycerol\nn~glycerol~ = n~NaOH~ / 3 = 0,06 / 3 = 0,02 mol.\nÁp dụng định luật bảo toàn khối lượng:\nm~chất béo~ + m~NaOH~ = m~xà phòng~ + m~glycerol~\n→ m~xà phòng~ = 17,24 + 0,06 × 40 - 0,02 × 92 = 17,24 + 2,40 - 1,84 = 17,80 gam.\n*Chọn D.*")

    # Q10
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 10.* Chất nào sau đây là thành phần chính của xà phòng?")
    q10_ch = [
        "CH~3~[CH~2~]~16~COOCH~3~.",
        "CH~3~[CH~2~]~14~COONa.",
        "CH~3~COONa.",
        "CH~3~[CH~2~]~11~C~6~H~4~SO~3~Na."
    ]
    add_mcq_choices_grid(doc, q10_ch, correct_idx=1, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Thành phần chính của xà phòng là muối sodium hoặc potassium của các acid béo. CH~3~[CH~2~]~14~COONa là sodium palmitate.\n*Chọn B.*")

    # Q11
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 11.* Aniline tác dụng với (HNO~2~ + HCl) ở 0 – 5°C tạo muối diazonium để tổng hợp phẩm nhuộm azo và dược phẩm.\nC~6~H~5~NH~2~ + HONO + HCl —(0-5°C)→ X + 2H~2~O\nChất X có công thức cấu tạo là:")
    q11_ch = [
        "[C~6~H~5~N~2~H]^+^Cl^-^",
        "[C~6~H~5~NH~2~]^+^Cl^-^",
        "[C~6~H~5~N~2~]^+^Cl^-^",
        "[C~6~H~5~NH~3~]^+^Cl^-^"
    ]
    add_mcq_choices_grid(doc, q11_ch, correct_idx=2, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Phản ứng diazot hóa amine thơm bậc một ở nhiệt độ thấp (0 - 5°C) tạo muối benzenediazonium chloride: [C~6~H~5~N~2~]^+^Cl^-^.\n*Chọn C.*")

    # Q12
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 12.* Các gốc α-glucose trong phân tử tinh bột tạo dạng mạch amylopectin phân nhánh, xoắn. Phần phân nhánh liên kết với nhau bởi liên kết")
    q12_ch = [
        "β-1,2-glycoside.",
        "α-1,4-glycoside.",
        "α-1,3-glycoside.",
        "α-1,6-glycoside."
    ]
    add_mcq_choices_grid(doc, q12_ch, correct_idx=3, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Trong phân tử amylopectin, các gốc α-glucose liên kết với nhau bởi liên kết α-1,4-glycoside tạo chuỗi mạch; tại các điểm phân nhánh, chuỗi nhánh liên kết với mạch chính bằng liên kết α-1,6-glycoside.\n*Chọn D.*")

    # Q13
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 13.* Một số ester được dùng trong hương liệu, mĩ phẩm, bột giặt là nhờ các ester")
    q13_ch = [
        "có mùi thơm, an toàn với người sử dụng.",
        "đều có nguồn gốc từ thiên nhiên.",
        "có thể bay hơi nhanh sau khi sử dụng.",
        "là chất lỏng dễ bay hơi."
    ]
    add_mcq_choices_grid(doc, q13_ch, correct_idx=0, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Nhiều ester có mùi thơm quyến rũ, đặc trưng của hoa quả và an toàn với sức khỏe con người nên được ứng dụng rộng rãi làm chất tạo hương.\n*Chọn A.*")

    # Q14
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 14.* Cho dung dịch ethylamine lần lượt tác dụng với: dung dịch FeCl~3~; dung dịch HCl; Cu(OH)~2~; dung dịch NaOH. Số trường hợp xảy ra phản ứng là:")
    q14_ch = ["2", "4", "5", "3"]
    add_mcq_choices_grid(doc, q14_ch, correct_idx=3, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Ethylamine C~2~H~5~NH~2~ phản ứng được với 3 chất:\n1. Dung dịch FeCl~3~: tạo kết tủa nâu đỏ Fe(OH)~3~.\n2. Dung dịch HCl: tạo muối C~2~H~5~NH~3~Cl.\n3. Cu(OH)~2~: hòa tan tạo phức đồng có màu xanh lam thẫm [Cu(C~2~H~5~NH~2~)~4~](OH)~2~.\nKhông phản ứng với dung dịch NaOH.\n*Chọn D.*")

    # Q15
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 15.* Kết quả thí nghiệm của các dung dịch X, Y, Z với thuốc thử được ghi ở bảng sau:")
    
    # Bảng Câu 15
    tbl15 = doc.add_table(rows=4, cols=3)
    tbl15.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl15.autofit = False
    col_w = [Cm(3.0), Cm(7.0), Cm(7.5)]
    headers15 = ["Mẫu thử", "Thuốc thử", "Hiện tượng"]
    rows15 = [
        ["X", "Dung dịch I~2~", "Có màu xanh tím"],
        ["Y", "Cu(OH)~2~", "Có màu xanh lam"],
        ["Z", "Dung dịch AgNO~3~ trong NH~3~", "Tạo kết tủa Ag"]
    ]
    for c_i, h in enumerate(headers15):
        cell = tbl15.cell(0, c_i)
        cell.width = col_w[c_i]
        format_cell_borders(cell, "single", "single", "single", "single", color="B0BEC5")
        p_c = cell.paragraphs[0]
        p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_c.paragraph_format.space_before = Pt(1)
        p_c.paragraph_format.space_after = Pt(1)
        render_rich_text(p_c, f"*{h}*", base_size=10.5)
    for r_i, r_data in enumerate(rows15):
        for c_i, val in enumerate(r_data):
            cell = tbl15.cell(r_i+1, c_i)
            cell.width = col_w[c_i]
            format_cell_borders(cell, "single", "single", "single", "single", color="B0BEC5")
            p_c = cell.paragraphs[0]
            p_c.paragraph_format.space_before = Pt(1)
            p_c.paragraph_format.space_after = Pt(1)
            if c_i == 0:
                p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                render_rich_text(p_c, f"*{val}*", base_size=10.5)
            else:
                p_c.alignment = WD_ALIGN_PARAGRAPH.LEFT
                render_rich_text(p_c, val, base_size=10.5)

    p_sub15 = doc.add_paragraph()
    p_sub15.paragraph_format.space_before = Pt(2)
    p_sub15.paragraph_format.space_after = Pt(2)
    render_rich_text(p_sub15, "Các dung dịch X, Y, Z lần lượt là")
    q15_ch = [
        "Hồ tinh bột, saccharose, glucose.",
        "Cellulose, saccharose, glucose.",
        "Cellulose, glucose, saccharose.",
        "Hồ tinh bột, glucose, saccharose."
    ]
    add_mcq_choices_grid(doc, q15_ch, correct_idx=0, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "- Mẫu thử X cho màu xanh tím với dung dịch I~2~ → X là Hồ tinh bột.\n- Mẫu thử Y hòa tan Cu(OH)~2~ cho dung dịch màu xanh lam → Y là Saccharose.\n- Mẫu thử Z tráng bạc tạo Ag → Z là Glucose.\n*Chọn A.*")

    # Q16
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 16.* Công thức của triolein là:")
    q16_ch = [
        "(C~17~H~35~COO)~3~C~3~H~5~.",
        "(C~17~H~33~COO)~3~C~3~H~5~.",
        "(CH~3~COO)~3~C~3~H~5~.",
        "(HCOO)~3~C~3~H~5~."
    ]
    add_mcq_choices_grid(doc, q16_ch, correct_idx=1, is_teacher=is_teacher)
    if is_teacher:
        add_solution_box(doc, "Triolein là triester của glycerol với oleic acid (C~17~H~33~COOH), có công thức phân tử thu gọn là (C~17~H~33~COO)~3~C~3~H~5~.\n*Chọn B.*")

    # ==================== PHẦN II: ĐÚNG / SAI ====================
    add_section_header(doc, "Phần II. Câu trắc nghiệm đúng/sai. (2 điểm – 2 câu)")
    p_p2_intro = doc.add_paragraph()
    p_p2_intro.paragraph_format.space_before = Pt(1)
    p_p2_intro.paragraph_format.space_after = Pt(3)
    render_rich_text(p_p2_intro, "Thí sinh trả lời từ câu 1 đến câu 2. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.", base_size=11, base_italic=True)

    # Câu 1 Phần II
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 1.* Tinh bột là loại lương thực quan trọng và là nguyên liệu chủ yếu để sản xuất bánh, kẹo, rượu, bia, ... Cellulose là thành phần chính tạo nên màng tế bào thực vật, có nhiều trong gỗ, bông gòn dùng làm nguyên liệu sản xuất thuốc súng không khói.")
    
    tf1_items = [
        ("a", "Tinh bột và cellulose là đồng phân cấu tạo của nhau do cùng có công thức phân tử dạng (C~6~H~10~O~5~)~n~.", "Sai", "Do hệ số mắt xích n của tinh bột và cellulose khác nhau rất nhiều (n của cellulose lớn hơn rất nhiều so với tinh bột) nên công thức phân tử khác nhau, chúng không phải là đồng phân của nhau."),
        ("b", "Tinh bột và cellulose đều thuộc loại polysaccharide.", "Đúng", "Cả tinh bột và cellulose đều là các polysaccharide tạo thành từ nhiều mắt xích monosaccharide."),
        ("c", "Dung dịch hồ tinh bột tạo với iodine hợp chất màu xanh tím, cellulose không có tính chất này.", "Đúng", "Cấu trúc xoắn dạng lò xo của tinh bột giữ các phân tử iodine tạo hợp chất màu xanh tím; cellulose mạch thẳng không xoắn nên không có hiện tượng này."),
        ("d", "Thuỷ phân hoàn toàn 162 gam tinh bột hoặc cellulose đều thu được 180 gam sản phẩm là glucose.", "Đúng", "(C~6~H~10~O~5~)~n~ + nH~2~O → nC~6~H~12~O~6~. Cứ 162n gam polysaccharide thủy phân hoàn toàn tạo 180n gam glucose. Với 162 gam thì lượng glucose thu được đúng bằng 162 × (180/162) = 180 gam.")
    ]
    for tag, text_item, ans_tf, exp_tf in tf1_items:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.space_before = Pt(1)
        p_item.paragraph_format.space_after = Pt(1)
        p_item.paragraph_format.left_indent = Cm(0.5)
        prefix = f"{tag}. "
        if is_teacher:
            prefix += f"[{ans_tf}] "
            r_pre = p_item.add_run(prefix)
            set_run_style(r_pre, font_size=11, bold=True, color=COLOR_CORRECT if ans_tf == "Đúng" else COLOR_PRIMARY)
        else:
            r_pre = p_item.add_run(prefix)
            set_run_style(r_pre, font_size=11, bold=True)
        render_rich_text(p_item, text_item, base_size=11)
        if is_teacher:
            p_exp = doc.add_paragraph()
            p_exp.paragraph_format.space_before = Pt(0)
            p_exp.paragraph_format.space_after = Pt(2)
            p_exp.paragraph_format.left_indent = Cm(1.0)
            render_rich_text(p_exp, f"→ *{ans_tf}*: {exp_tf}", base_size=10, base_color=COLOR_GRAY, base_italic=True)

    # Câu 2 Phần II
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 2.* Aspirin được sử dụng làm thuốc giảm đau, hạ sốt. Sau khi uống, aspirin bị thuỷ phân trong cơ thể tạo thành salicylic acid. Salicylic acid ức chế quá trình sinh tổng hợp prostaglandin (chất gây đau, sốt và viêm khi nồng độ trong máu cao hơn mức bình thường).")
    add_image_centered(doc, "fig_cau2_p2_aspirin.png", width_cm=11.5)

    tf2_items = [
        ("a", "1 mol aspirin hay 1 mol salicylic acid đều phản ứng tối đa với 2 mol NaOH.", "Sai", "Aspirin chứa 1 nhóm -COOH và 1 nhóm ester của phenol nên 1 mol aspirin phản ứng tối đa với 3 mol NaOH (CH~3~COOC~6~H~4~COOH + 3NaOH → CH~3~COONa + NaOC~6~H~4~COONa + 2H~2~O). Salicylic acid chỉ phản ứng với 2 mol NaOH."),
        ("b", "Trong phân tử aspirin có số liên kết π và vòng là 5.", "Sai", "Vòng benzene gồm 1 vòng và 3 liên kết π; nhóm -COO- có 1 liên kết π; nhóm -COOH có 1 liên kết π. Tổng số liên kết π và vòng = 3 + 1 + 1 + 1 = 6."),
        ("c", "Aspirin có công thức phân tử C~9~H~8~O~4~.", "Đúng", "Aspirin có công thức cấu tạo CH~3~COOC~6~H~4~COOH, đếm số nguyên tử: 9 carbon, 8 hydrogen, 4 oxygen → C~9~H~8~O~4~."),
        ("d", "Aspirin được điều chế theo phản ứng sau:\n(CH~3~CO)~2~O + HOC~6~H~4~COOH —(H~2~SO~4~, t°)→ CH~3~COOC~6~H~4~COOH + CH~3~COOH\nĐể sản xuất 3 triệu viên thuốc aspirin cần tối thiểu 260 kg salicylic acid. Biết rằng mỗi viên thuốc có chứa 81 mg aspirin và hiệu suất phản ứng đạt 70%.", "Sai", "Khối lượng aspirin trong 3 triệu viên: m~aspirin~ = 3·10^6^ × 81 mg = 243 kg. Theo lý thuyết: m~acid(LT)~ = 243 × (138 / 180) = 186,3 kg. Với H = 70%, lượng salicylic acid thực tế cần tối thiểu = 186,3 / 0,7 = 266,14 kg (khác 260 kg).")
    ]
    for tag, text_item, ans_tf, exp_tf in tf2_items:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.space_before = Pt(1)
        p_item.paragraph_format.space_after = Pt(1)
        p_item.paragraph_format.left_indent = Cm(0.5)
        prefix = f"{tag}. "
        if is_teacher:
            prefix += f"[{ans_tf}] "
            r_pre = p_item.add_run(prefix)
            set_run_style(r_pre, font_size=11, bold=True, color=COLOR_CORRECT if ans_tf == "Đúng" else COLOR_PRIMARY)
        else:
            r_pre = p_item.add_run(prefix)
            set_run_style(r_pre, font_size=11, bold=True)
        render_rich_text(p_item, text_item, base_size=11)
        if is_teacher:
            p_exp = doc.add_paragraph()
            p_exp.paragraph_format.space_before = Pt(0)
            p_exp.paragraph_format.space_after = Pt(2)
            p_exp.paragraph_format.left_indent = Cm(1.0)
            render_rich_text(p_exp, f"→ *{ans_tf}*: {exp_tf}", base_size=10, base_color=COLOR_GRAY, base_italic=True)

    # ==================== PHẦN III: TRẢ LỜI NGẮN ====================
    add_section_header(doc, "Phần III. Câu trắc nghiệm yêu cầu trả lời ngắn. (4 câu – 1 điểm)")
    p_p3_intro = doc.add_paragraph()
    p_p3_intro.paragraph_format.space_before = Pt(1)
    p_p3_intro.paragraph_format.space_after = Pt(3)
    render_rich_text(p_p3_intro, "Thí sinh trả lời từ câu 1 đến câu 4.", base_size=11, base_italic=True)

    # Câu 1 Phần III
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 1.* Cho công thức cấu tạo thu gọn và tên gọi sau:")
    tbl_am = doc.add_table(rows=5, cols=3)
    tbl_am.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_am.autofit = False
    col_w_am = [Cm(2.5), Cm(7.5), Cm(7.5)]
    headers_am = ["Chất", "Công thức cấu tạo thu gọn", "Tên gọi"]
    rows_am = [
        ["(1)", "CH~3~N(CH~3~)~2~", "Aniline"],
        ["(2)", "CH~3~CH~2~NHCH~3~", "Trimethylamine"],
        ["(3)", "CH~3~CH~2~NH~2~", "Ethylamine"],
        ["(4)", "C~6~H~5~NH~2~", "N-methylethanamine"]
    ]
    for c_i, h in enumerate(headers_am):
        cell = tbl_am.cell(0, c_i)
        cell.width = col_w_am[c_i]
        format_cell_borders(cell, "single", "single", "single", "single", color="B0BEC5")
        p_c = cell.paragraphs[0]
        p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_c.paragraph_format.space_before = Pt(1)
        p_c.paragraph_format.space_after = Pt(1)
        render_rich_text(p_c, f"*{h}*", base_size=10.5)
    for r_i, r_data in enumerate(rows_am):
        for c_i, val in enumerate(r_data):
            cell = tbl_am.cell(r_i+1, c_i)
            cell.width = col_w_am[c_i]
            format_cell_borders(cell, "single", "single", "single", "single", color="B0BEC5")
            p_c = cell.paragraphs[0]
            p_c.paragraph_format.space_before = Pt(1)
            p_c.paragraph_format.space_after = Pt(1)
            if c_i == 0:
                p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                render_rich_text(p_c, f"*{val}*", base_size=10.5)
            else:
                p_c.alignment = WD_ALIGN_PARAGRAPH.LEFT
                render_rich_text(p_c, val, base_size=10.5)

    p_sub_am = doc.add_paragraph()
    p_sub_am.paragraph_format.space_before = Pt(2)
    p_sub_am.paragraph_format.space_after = Pt(2)
    render_rich_text(p_sub_am, "Sắp xếp các chất theo thứ tự tên gọi tương ứng?")
    if is_teacher:
        add_solution_box(doc, "Đối chiếu công thức cấu tạo thu gọn với tên gọi đúng:\n- (1) CH~3~N(CH~3~)~2~ là Trimethylamine (tên dòng 2).\n- (2) CH~3~CH~2~NHCH~3~ là N-methylethanamine (tên dòng 4).\n- (3) CH~3~CH~2~NH~2~ là Ethylamine (tên dòng 3).\n- (4) C~6~H~5~NH~2~ là Aniline (tên dòng 1).\nThứ tự các chất tương ứng với cột tên gọi từ trên xuống dưới (Aniline, Trimethylamine, Ethylamine, N-methylethanamine) là: *(4), (1), (3), (2)*.")

    # Câu 2 Phần III
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 2.* Tại một nhà máy rượu, cứ 10 tấn tinh bột (chứa 6,85% tạp chất trơ) sẽ sản xuất được 7,21 m^3^ ethanol 40° (cho khối lượng riêng của ethanol nguyên chất là 0,789 g/cm^3^). Hiệu suất của quá trình sản xuất là bao nhiêu phần trăm? *(Kết quả làm tròn đến hàng đơn vị)*")
    if is_teacher:
        add_solution_box(doc, "1. Khối lượng tinh bột nguyên chất: m = 10 × (100% - 6,85%) = 9,315 tấn = 9315 kg.\n2. Thể tích ethanol nguyên chất: V = 7,21 m^3^ × (40 / 100) = 2,884 m^3^ = 2884 lít.\n3. Khối lượng ethanol thực tế: m~TT~ = 2884 × 0,789 = 2275,476 kg.\n4. Theo sơ đồ phản ứng: (C~6~H~10~O~5~)~n~ → 2n C~2~H~5~OH\nCứ 162 kg tinh bột theo lý thuyết tạo ra 92 kg C~2~H~5~OH.\n→ m~ethanol (LT)~ = 9315 × (92 / 162) = 5290 kg.\n5. Hiệu suất cả quá trình: H = (2275,476 / 5290) × 100% ≈ *43%*.\n*Đáp số: 43*")

    # Câu 3 Phần III
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 3.* Một loại chất béo có chứa 80% triolein về khối lượng. Xà phòng hóa hoàn toàn 22,1 kg chất béo này trong dung dịch NaOH, đun nóng thu được x bánh xà phòng. Biết rằng trong mỗi bánh xà phòng có chứa 60 gam sodium oleate. Xác định giá trị của x?")
    add_image_centered(doc, "fig_cau3_p3_soap.png", width_cm=4.2)
    if is_teacher:
        add_solution_box(doc, "1. Khối lượng triolein nguyên chất: m~triolein~ = 22,1 × 80% = 17,68 kg = 17680 gam.\n2. n~triolein~ = 17680 / 884 = 20 mol.\n3. Phản ứng xà phòng hóa: (C~17~H~33~COO)~3~C~3~H~5~ + 3NaOH → 3C~17~H~33~COONa + C~3~H~5~(OH)~3~\n→ n~sodium oleate~ = 3 × 20 = 60 mol.\n4. Khối lượng sodium oleate: m = 60 × 304 = 18240 gam.\n5. Số bánh xà phòng thu được: x = 18240 / 60 = *304* bánh.\n*Đáp số: 304*")

    # Câu 4 Phần III
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 4.* Khi xà phòng hóa triglyceride X bằng dung dịch NaOH dư, đun nóng, thu được sản phẩm gồm glycerol, sodium stearate và sodium palmitate. Có bao nhiêu đồng phân cấu tạo thỏa mãn tính chất trên của X?")
    if is_teacher:
        add_solution_box(doc, "Triglyceride X tạo thành từ 2 loại acid béo là stearic acid (gốc S) và palmitic acid (gốc P). Do đó trong phân tử X phải có cả gốc S và gốc P:\n- Dạng 2 gốc S + 1 gốc P:\n   + Gốc P gắn ở carbon trung tâm (vị trí β): S - β(P) - S (1 đồng phân).\n   + Gốc P gắn ở carbon đầu mạch (vị trí α): P - β(S) - S (1 đồng phân).\n- Dạng 1 gốc S + 2 gốc P:\n   + Gốc S gắn ở carbon trung tâm (vị trí β): P - β(S) - P (1 đồng phân).\n   + Gốc S gắn ở carbon đầu mạch (vị trí α): S - β(P) - P (1 đồng phân).\nTổng cộng có 2 + 2 = *4* đồng phân cấu tạo thỏa mãn.\n*Đáp số: 4*")

    # ==================== PHẦN B: TỰ LUẬN ====================
    add_section_header(doc, "B. PHẦN TỰ LUẬN (3 câu – 3 điểm)")

    # Câu 1 Tự luận
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 1: (1 điểm)* Viết phương trình phản ứng hoàn thành sơ đồ chuyển hóa sau (ghi rõ điều kiện phản ứng nếu có):")
    p_sd = doc.add_paragraph()
    p_sd.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sd.paragraph_format.space_before = Pt(1)
    p_sd.paragraph_format.space_after = Pt(2)
    render_rich_text(p_sd, "*Tinh bột → Glucose → Ethanol → Acetic acid → Ethyl acetate.*", base_bold=True)
    if is_teacher:
        add_solution_box(doc, "Các phương trình phản ứng:\n(1) (C~6~H~10~O~5~)~n~ + nH~2~O —(H^+^, t° hoặc enzyme)—→ nC~6~H~12~O~6~\n(2) C~6~H~12~O~6~ —(men rượu, 30–35°C)—→ 2C~2~H~5~OH + 2CO~2~ ↑\n(3) C~2~H~5~OH + O~2~ —(men giấm)—→ CH~3~COOH + H~2~O\n(4) CH~3~COOH + C~2~H~5~OH ⇌(H~2~SO~4~ đặc, t°)⇌ CH~3~COOC~2~H~5~ + H~2~O")

    # Câu 2 Tự luận
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 2: (1 điểm)* Aniline có nhiều ứng dụng quan trọng và đa dạng trong các ngành công nghiệp, đặc biệt là ngành sản xuất phẩm nhuộm và dược phẩm. Aniline có thể được tổng hợp từ benzene theo sơ đồ chuyển hoá sau:")
    add_image_centered(doc, "fig_cau2_tuluan_aniline.png", width_cm=14.0)
    p_sub_ani = doc.add_paragraph()
    p_sub_ani.paragraph_format.space_before = Pt(1)
    p_sub_ani.paragraph_format.space_after = Pt(2)
    render_rich_text(p_sub_ani, "Theo sơ đồ trên, từ 1 tấn benzene sẽ điều chế được bao nhiêu kg aniline? Biết hiệu suất toàn bộ quá trình là 60%.")
    if is_teacher:
        add_solution_box(doc, "Sơ đồ chuyển hóa bảo toàn vòng benzene:\nC~6~H~6~ (M = 78 g/mol) → C~6~H~5~NH~2~ (M = 93 g/mol)\n1. Khối lượng benzene: m~benzene~ = 1 tấn = 1000 kg.\n2. Theo lý thuyết (hiệu suất 100%), khối lượng aniline thu được là:\nm~aniline (LT)~ = 1000 × (93 / 78) = 1192,308 kg.\n3. Do hiệu suất toàn bộ quá trình đạt 60%, khối lượng aniline thực tế thu được là:\nm~aniline (TT)~ = 1192,308 × 60% = *715,38 kg* (hoặc 715,4 kg).\n*Đáp số: 715,38 kg aniline*")

    # Câu 3 Tự luận
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, "*Câu 3: (1 điểm)*\na, Trong cây xanh, glucose được tổng hợp nhờ quá trình quang hợp. Viết phương trình phản ứng xảy ra?\nb, Một nhóm học sinh muốn thử nghiệm phản ứng tráng bạc lên kính bằng nguyên liệu đầu là glucose. Giả sử lớp bạc có diện tích là 100 cm^2^ và độ dày là 0,5 μm. Biết rằng khối lượng riêng của bạc là 10,49 g/cm^3^ và khối lượng mol của glucose là 180 g/mol. Khối lượng glucose cần dùng là bao nhiêu gam? (Giả thiết hiệu suất phản ứng tráng bạc là 92%)")
    if is_teacher:
        add_solution_box(doc, "a, Phương trình hóa học của quá trình quang hợp tạo glucose trong cây xanh:\n6CO~2~ + 6H~2~O —(ánh sáng, diệp lục)—→ C~6~H~12~O~6~ + 6O~2~ ↑\n\nb, Khối lượng glucose cần dùng:\n1. Thể tích lớp bạc cần mạ trên mặt kính:\nV = S × h = 100 cm^2^ × (0,5 × 10^-4^ cm) = 0,005 cm^3^.\n2. Khối lượng bạc cần tráng:\nm~Ag~ = V × D = 0,005 cm^3^ × 10,49 g/cm^3^ = 0,05245 gam.\n3. Số mol bạc cần tạo thành: n~Ag~ = 0,05245 / 108 mol.\n4. Phản ứng tráng bạc của glucose: C~6~H~12~O~6~ → 2Ag\n→ n~glucose (LT)~ = n~Ag~ / 2 = 0,05245 / (2 × 108) mol.\n5. Khối lượng glucose theo lý thuyết: m~glucose (LT)~ = [0,05245 / (2 × 108)] × 180 = 0,043708 gam.\n6. Do hiệu suất phản ứng tráng bạc là 92%, khối lượng glucose thực tế cần dùng là:\nm~glucose (cần dùng)~ = 0,043708 / 0,92 ≈ *0,0475 gam* (tương đương 47,5 mg).\n*Đáp số: a, Phương trình quang hợp; b, Khoảng 0,0475 gam glucose*")

    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_end.paragraph_format.space_before = Pt(10)
    p_end.paragraph_format.space_after = Pt(4)
    r_end = p_end.add_run("--------- HẾT ---------")
    set_run_style(r_end, font_size=11, bold=True)

    filename = "De_Kiem_Tra_Hoa_12_Ma104_GiaoVien.docx" if is_teacher else "De_Kiem_Tra_Hoa_12_Ma104.docx"
    save_path = os.path.join(OUT_DIR, filename)
    doc.save(save_path)
    print(f"Saved: {save_path}")

if __name__ == "__main__":
    build_exam_document(is_teacher=False)
    build_exam_document(is_teacher=True)
