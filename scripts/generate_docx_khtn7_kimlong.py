# -*- coding: utf-8 -*-
"""
Script tạo 3 file Word (.docx) chuẩn "FORM ĐỀ KIỂM TRA" cho:
ĐỀ CƯƠNG ÔN TẬP GIỮA KÌ I - KHTN 7
Trường THCS Kim Long - Năm học 2026 - 2027
Thầy TRẦN VĂN THẠNH - 0777.470.803

3 phiên bản:
1. Bản Học sinh: San_Pham/De_Cuong_KHTN_7_Giua_HK1_Kim_Long_HocSinh.docx
2. Bản Giáo viên: San_Pham/De_Cuong_KHTN_7_Giua_HK1_Kim_Long_GiaoVien.docx
3. Bản Dạy trên màn hình: San_Pham/De_Cuong_KHTN_7_Giua_HK1_Kim_Long_DayTrenManHinh.docx
"""

import os
import sys
import re
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.path.append(os.path.dirname(__file__))
sys.stdout.reconfigure(encoding='utf-8')

from khtn7_kimlong_data import (
    INFO,
    BAI1_MCQ,
    KHTN1_THEORY, KHTN1_MCQ, KHTN1_TF, KHTN1_SHORT, KHTN1_ESSAY,
    KHTN2_THEORY, KHTN2_MCQ, KHTN2_TF, KHTN2_SHORT, KHTN2_ESSAY,
    KHTN3_THEORY, KHTN3_QUESTIONS, KHTN3_NOTES
)

ROOT_DIR = r"c:\Antigravity_Thanh"
OUT_DIR = os.path.join(ROOT_DIR, "San_Pham")
IMAGE_DIR = os.path.join(OUT_DIR, "images_khtn7_kimlong")

ONEDRIVE_DIRS = [
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn7\KHTN7-IN-DEKIEMTRADINHKI\New folder (2)",
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn7\KHTN7-IN-DEKIEMTRADINHKI",
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn7"
]
os.makedirs(OUT_DIR, exist_ok=True)

# Bảng màu chuẩn quy định trong FORM_DE_KIEM_TRA
COLOR_HF_RED = RGBColor(211, 47, 47)        # Coral Red (#D32F2F) cho Header & Footer
COLOR_HF_GREEN_HEX = "8BC390"               # Sage Green (#8BC390) cho viền Header & Footer
COLOR_PRIMARY = RGBColor(46, 125, 50)       # Xanh lá đậm (#2E7D32) cho Tiêu đề Phân môn / Phần
COLOR_CORRECT = RGBColor(198, 40, 40)       # Đỏ (#C62828) cho đáp án đúng Bản Giáo Viên
COLOR_SOLUTION = RGBColor(21, 101, 192)     # Xanh lam (#1565C0) cho tiêu đề Lời giải
COLOR_SCREEN_TITLE = RGBColor(106, 27, 154) # Tím đậm (#6A1B9A) cho tiêu đề bản Dạy Màn Hình
COLOR_GRAY = RGBColor(117, 117, 117)
COLOR_BORDER_GREY_HEX = "CCCCCC"

def add_p_border_bottom(p, color="8BC390", sz="6"):
    """Thêm đường kẻ viền dưới cho đoạn văn bản."""
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_p_border_top(p, color="8BC390", sz="6"):
    """Thêm đường kẻ viền trên cho đoạn văn bản."""
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    """Cài đặt lề đệm cho ô bảng."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex="F1F8E9"):
    """Đặt màu nền cho ô bảng."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="CCCCCC", sz="4"):
    """Đặt viền thanh mảnh cho bảng."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def set_run_style(run, font_name="Times New Roman", font_size=10.5, bold=False, italic=False, subscript=False, superscript=False, color=None):
    """Cài đặt style chi tiết cho từng run text."""
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    """Chèn trường số trang tự động PAGE vào chân trang."""
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def render_rich_text(p, text, base_font="Times New Roman", base_size=10.5, base_bold=False, base_italic=False, base_color=None):
    """
    Sinh văn bản định dạng phong phú từ ký hiệu:
    - *bold*
    - _italic_
    - ~subscript~ (CO~2~, H~2~O, CaCO~3~)
    - ^superscript^ (2^3^, 2^4^, 100°C)
    """
    tokens = re.split(r'(~[^~]+~|\^[^\^]+\^|\*[^\*]+\*|_[^_]+_)', text)
    for token in tokens:
        if not token:
            continue

        is_bold = base_bold
        is_italic = base_italic
        is_sub = False
        is_sup = False
        content = token

        if token.startswith('*') and token.endswith('*'):
            is_bold = True
            content = token[1:-1]
        elif token.startswith('_') and token.endswith('_'):
            is_italic = True
            content = token[1:-1]
        elif token.startswith('~') and token.endswith('~'):
            is_sub = True
            content = token[1:-1]
        elif token.startswith('^') and token.endswith('^'):
            is_sup = True
            content = token[1:-1]

        run = p.add_run(content)
        set_run_style(run, font_name=base_font, font_size=base_size, bold=is_bold, italic=is_italic,
                      subscript=is_sub, superscript=is_sup, color=base_color)

def setup_header_footer(doc):
    """Cài đặt Header & Footer chuẩn theo FORM_DE_KIEM_TRA."""
    normal_style = doc.styles['Normal']
    normal_style.paragraph_format.tab_stops.clear_all()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    section.start_type = WD_SECTION_START.NEW_PAGE
    section.different_first_page_header_footer = False
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    # 1. Header (Bảng 1x2, rộng 18.0 cm)
    header = section.header
    for p in header.paragraphs:
        p.text = ""
    tbl_h = header.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl_h.cell(0, 0)
    c_right = tbl_h.cell(0, 1)
    c_left.width = Cm(8.0)
    c_right.width = Cm(10.0)

    p_hl = c_left.paragraphs[0]
    p_hl.paragraph_format.space_before = Pt(0)
    p_hl.paragraph_format.space_after = Pt(2)
    p_hl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_hl = p_hl.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)

    p_hr = c_right.paragraphs[0]
    p_hr.paragraph_format.space_before = Pt(0)
    p_hr.paragraph_format.space_after = Pt(2)
    p_hr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p_hr.add_run("Thầy ")
    set_run_style(r1, font_size=10, color=COLOR_HF_RED)
    r2 = p_hr.add_run("TRẦN VĂN THẠNH")
    set_run_style(r2, font_size=10, bold=True, color=COLOR_HF_RED)
    r3 = p_hr.add_run(" - 0777.470.803")
    set_run_style(r3, font_size=10, color=COLOR_HF_RED)

    add_p_border_bottom(p_hl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_bottom(p_hr, color=COLOR_HF_GREEN_HEX, sz="6")

    # 2. Footer (Bảng 1x2, rộng 18.0 cm)
    footer = section.footer
    for p in footer.paragraphs:
        p.text = ""
    tbl_f = footer.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_fl = tbl_f.cell(0, 0)
    c_fr = tbl_f.cell(0, 1)
    c_fl.width = Cm(14.5)
    c_fr.width = Cm(3.5)

    p_fl = c_fl.paragraphs[0]
    p_fl.paragraph_format.space_before = Pt(2)
    p_fl.paragraph_format.space_after = Pt(0)
    p_fl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_fl = p_fl.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)

    p_fr = c_fr.paragraphs[0]
    p_fr.paragraph_format.space_before = Pt(2)
    p_fr.paragraph_format.space_after = Pt(0)
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_field(p_fr, font_size=10, bold=True, color=COLOR_HF_RED)

    add_p_border_top(p_fl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_top(p_fr, color=COLOR_HF_GREEN_HEX, sz="6")

def build_title_block(doc, is_teacher=False, is_screen=False, base_size=10.5):
    """Khối tiêu đề đề cương theo FORM_DE_KIEM_TRA."""
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl.cell(0, 0)
    c_right = tbl.cell(0, 1)
    c_left.width = Cm(8.0)
    c_right.width = Cm(10.0)

    # Cột trái
    p_l1 = c_left.paragraphs[0]
    p_l1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l1.paragraph_format.space_before = Pt(0)
    p_l1.paragraph_format.space_after = Pt(1)
    render_rich_text(p_l1, "*TRƯỜNG THCS KIM LONG*", base_size=11, base_bold=True)

    p_l2 = c_left.add_paragraph()
    p_l2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l2.paragraph_format.space_before = Pt(0)
    p_l2.paragraph_format.space_after = Pt(1)
    render_rich_text(p_l2, "*TỔ KHOA HỌC TỰ NHIÊN*", base_size=10.5, base_bold=True)

    p_l3 = c_left.add_paragraph()
    p_l3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l3.paragraph_format.space_before = Pt(0)
    p_l3.paragraph_format.space_after = Pt(2)
    render_rich_text(p_l3, "---***---", base_size=9, base_color=COLOR_GRAY)

    # Cột phải
    p_r1 = c_right.paragraphs[0]
    p_r1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r1.paragraph_format.space_before = Pt(0)
    p_r1.paragraph_format.space_after = Pt(1)
    render_rich_text(p_r1, "*ĐỀ CƯƠNG ÔN TẬP GIỮA KÌ I*", base_size=12, base_bold=True, base_color=COLOR_HF_RED)

    p_r2 = c_right.add_paragraph()
    p_r2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r2.paragraph_format.space_before = Pt(0)
    p_r2.paragraph_format.space_after = Pt(1)
    render_rich_text(p_r2, "*MÔN: KHOA HỌC TỰ NHIÊN 7*", base_size=11.5, base_bold=True, base_color=RGBColor(21, 101, 192))

    p_r3 = c_right.add_paragraph()
    p_r3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r3.paragraph_format.space_before = Pt(0)
    p_r3.paragraph_format.space_after = Pt(2)
    render_rich_text(p_r3, "_NĂM HỌC 2026 - 2027_", base_size=10, base_italic=True)

    if is_teacher:
        p_tc = doc.add_paragraph()
        p_tc.paragraph_format.space_before = Pt(2)
        p_tc.paragraph_format.space_after = Pt(2)
        p_tc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rtc = p_tc.add_run("(BẢN GIÁO VIÊN -- CÓ ĐÁP ÁN VÀ LỜI GIẢI CHI TIẾT)")
        set_run_style(rtc, font_size=11, bold=True, color=COLOR_CORRECT)
    elif is_screen:
        p_sc = doc.add_paragraph()
        p_sc.paragraph_format.space_before = Pt(2)
        p_sc.paragraph_format.space_after = Pt(2)
        p_sc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rsc = p_sc.add_run("(BẢN TRÌNH CHIẾU / GIẢNG DẠY MÀN HÌNH -- CÓ KHUNG CHỪA CHỖ TRỐNG TƯƠNG TÁC)")
        set_run_style(rsc, font_size=11, bold=True, color=COLOR_SCREEN_TITLE)

    # Dòng thông tin học sinh
    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(4)
    p_info.paragraph_format.space_after = Pt(3)
    render_rich_text(p_info, "*Họ, tên học sinh:* .............................................................. *Lớp:* ............. *Phòng thi:* .............", base_size=base_size)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(4)
    add_p_border_bottom(p_div, color=COLOR_BORDER_GREY_HEX, sz="4")

def add_part_header(doc, title, base_size=12.5):
    """Tiêu đề lớn cho Phân môn hoặc Phần."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    render_rich_text(p, f"*{title}*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)

def add_chapter_title(doc, title, base_size=11.5):
    """Tiêu đề cho từng Chương / Bài lớn."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, f"*{title}*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)

def add_section_header(doc, title, base_size=11):
    """Tiêu đề mục con."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, f"*{title}*", base_size=base_size, base_bold=True, base_color=RGBColor(21, 101, 192))

def add_callout_box(doc, lines, base_size=10):
    """Khung tóm tắt lý thuyết trọng tâm."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    c.width = Cm(18.0)
    set_cell_margins(c, top=60, bottom=60, left=100, right=100)
    set_cell_shading(c, "F5F9F6")
    
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    render_rich_text(p, lines[0], base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)
    
    for line in lines[1:]:
        p_line = c.add_paragraph()
        p_line.paragraph_format.space_before = Pt(1)
        p_line.paragraph_format.space_after = Pt(1)
        render_rich_text(p_line, line, base_size=base_size)

def add_solution_block(doc, solution_text, base_size=10):
    """Trình bày Lời giải chi tiết theo chuẩn không dùng Table (thụt lề 0.6 cm)."""
    p_sol = doc.add_paragraph()
    p_sol.paragraph_format.left_indent = Cm(0.6)
    p_sol.paragraph_format.space_before = Pt(2)
    p_sol.paragraph_format.space_after = Pt(3)
    p_sol.paragraph_format.line_spacing = 1.15

    r_lbl = p_sol.add_run("Lời giải chi tiết:\n")
    set_run_style(r_lbl, font_size=base_size, bold=True, color=COLOR_SOLUTION)

    lines = solution_text.split('\n')
    for idx, line in enumerate(lines):
        line_str = line.strip()
        if not line_str:
            continue
        if idx > 0:
            p_sol = doc.add_paragraph()
            p_sol.paragraph_format.left_indent = Cm(0.6)
            p_sol.paragraph_format.space_before = Pt(1)
            p_sol.paragraph_format.space_after = Pt(1)
            p_sol.paragraph_format.line_spacing = 1.15
        render_rich_text(p_sol, line_str, base_size=base_size)

def add_screen_teaching_space(doc, num_lines=3):
    """Chừa khoảng trống rộng rãi cho bản Dạy Trên Màn Hình."""
    p_box = doc.add_paragraph()
    p_box.paragraph_format.left_indent = Cm(0.5)
    p_box.paragraph_format.space_before = Pt(3)
    p_box.paragraph_format.space_after = Pt(5)
    
    # Viền nét đứt màu xám nhẹ, viền trái nét liền tím
    pPr = p_box._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="dashed" w:sz="4" w:space="4" w:color="B0BEC5"/><w:bottom w:val="dashed" w:sz="4" w:space="4" w:color="B0BEC5"/><w:left w:val="single" w:sz="12" w:space="6" w:color="7E57C2"/></w:pBdr>'
    pPr.append(parse_xml(xml))
    
    r_hint = p_box.add_run("✎ Ghi chú & Lời giải trực tiếp trên màn hình:\n")
    set_run_style(r_hint, font_size=9.5, bold=True, color=COLOR_SCREEN_TITLE)
    for _ in range(num_lines):
        p_box.add_run("\n")

def add_writing_lines(doc, num_lines=3):
    """Thêm dòng chấm chấm cho học sinh làm bài tự luận."""
    for _ in range(num_lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("................................................................................................................................................................................................................")
        set_run_style(r, font_size=9, color=COLOR_GRAY)

def add_image_figure(doc, img_filename, width_cm=11.0):
    """Chèn hình ảnh căn giữa trang trọng."""
    img_path = os.path.join(IMAGE_DIR, img_filename)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(2)
        p_img.paragraph_format.space_after = Pt(3)
        r_img = p_img.add_run()
        r_img.add_picture(img_path, width=Cm(width_cm))

def render_mcq_options(doc, opts, correct_letter, is_teacher, base_size=10.5):
    """Bố trí 4 phương án A, B, C, D bằng Tab Stops - Tuyệt đối không dùng Table."""
    lens = [len(opt) for opt in opts]
    max_len = max(lens) if lens else 0

    if max_len <= 18:
        # 4 phương án trên 1 dòng
        p_opt = doc.add_paragraph()
        p_opt.paragraph_format.space_before = Pt(1)
        p_opt.paragraph_format.space_after = Pt(2)
        p_opt.paragraph_format.tab_stops.clear_all()
        p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(5.0), WD_TAB_ALIGNMENT.LEFT)
        p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(9.5), WD_TAB_ALIGNMENT.LEFT)
        p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(14.0), WD_TAB_ALIGNMENT.LEFT)

        for opt in opts:
            p_opt.add_run("\t")
            letter = opt.strip()[0]
            if is_teacher and letter == correct_letter:
                render_rich_text(p_opt, f"*{opt}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p_opt, opt, base_size=base_size)

    elif max_len <= 45:
        # 2 dòng x 2 phương án
        p1 = doc.add_paragraph()
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.tab_stops.clear_all()
        p1.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p1.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)

        for opt in opts[:2]:
            p1.add_run("\t")
            letter = opt.strip()[0]
            if is_teacher and letter == correct_letter:
                render_rich_text(p1, f"*{opt}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p1, opt, base_size=base_size)

        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(1)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.tab_stops.clear_all()
        p2.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p2.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)

        for opt in opts[2:]:
            p2.add_run("\t")
            letter = opt.strip()[0]
            if is_teacher and letter == correct_letter:
                render_rich_text(p2, f"*{opt}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p2, opt, base_size=base_size)

    else:
        # 4 dòng riêng biệt
        for opt in opts:
            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.left_indent = Cm(0.5)
            p_opt.paragraph_format.space_before = Pt(1)
            p_opt.paragraph_format.space_after = Pt(1)
            letter = opt.strip()[0]
            if is_teacher and letter == correct_letter:
                render_rich_text(p_opt, f"*{opt}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p_opt, opt, base_size=base_size)

def render_mcq_list(doc, questions, is_teacher, is_screen, base_size=10.5):
    """Sinh danh sách câu hỏi trắc nghiệm nhiều lựa chọn."""
    for q in questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(1)
        render_rich_text(p_q, f"*Câu {q['id']}:* {q['content']}", base_size=base_size)

        if "image" in q and q["image"]:
            add_image_figure(doc, q["image"], width_cm=8.5)

        if q.get("type") == "matching":
            # Bảng nối 2 cột
            tbl = doc.add_table(rows=1 + len(q["table_data"]), cols=2)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(tbl, color="B0BEC5", sz="4")
            
            # Header
            c0, c1 = tbl.cell(0, 0), tbl.cell(0, 1)
            c0.width, c1.width = Cm(8.5), Cm(9.5)
            set_cell_shading(c0, "E8F5E9"); set_cell_shading(c1, "E8F5E9")
            p0 = c0.paragraphs[0]; p1 = c1.paragraphs[0]
            render_rich_text(p0, "*Cột A*", base_size=10, base_bold=True, base_color=COLOR_PRIMARY)
            render_rich_text(p1, "*Cột B*", base_size=10, base_bold=True, base_color=COLOR_PRIMARY)
            
            for idx, (col_a, col_b) in enumerate(q["table_data"], start=1):
                ca, cb = tbl.cell(idx, 0), tbl.cell(idx, 1)
                ca.width, cb.width = Cm(8.5), Cm(9.5)
                pa, pb = ca.paragraphs[0], cb.paragraphs[0]
                render_rich_text(pa, col_a, base_size=10)
                render_rich_text(pb, col_b, base_size=10)
            
            if is_teacher:
                add_solution_block(doc, f"**Đáp án nối:** {q['correct']}\n{q['solution']}", base_size=10)
            elif is_screen:
                add_screen_teaching_space(doc, num_lines=2)
            else:
                p_ans = doc.add_paragraph()
                p_ans.paragraph_format.left_indent = Cm(0.5)
                p_ans.paragraph_format.space_before = Pt(2)
                p_ans.paragraph_format.space_after = Pt(2)
                render_rich_text(p_ans, "*Trả lời:* 1 - ..... ; 2 - ..... ; 3 - ..... ; 4 - .....", base_size=base_size)
            continue

        if "options" in q:
            render_mcq_options(doc, q["options"], q.get("correct", ""), is_teacher, base_size=base_size)

        if is_teacher and "solution" in q and q["solution"]:
            add_solution_block(doc, q["solution"], base_size=10)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=2)

def render_tf_list(doc, questions, is_teacher, is_screen, base_size=10.5):
    """Sinh danh sách câu hỏi đúng/sai."""
    for q in questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*Câu {q['id']}:* {q['content']}", base_size=base_size)

        sol_lines = []
        for item in q["sub_items"]:
            p_item = doc.add_paragraph()
            p_item.paragraph_format.left_indent = Cm(0.5)
            p_item.paragraph_format.space_before = Pt(1)
            p_item.paragraph_format.space_after = Pt(1)

            tag = f"{item['id']})"
            txt = f"{tag} {item['text']}"
            if is_teacher:
                ans_str = " (ĐÚNG)" if item["answer"] else " (SAI)"
                render_rich_text(p_item, txt, base_size=base_size)
                r_ans = p_item.add_run(ans_str)
                set_run_style(r_ans, font_size=base_size, bold=True, color=COLOR_CORRECT)
                sol_lines.append(f"• Ý {item['id']}: {'Đúng' if item['answer'] else 'Sai'}. {item.get('reason', '')}")
            else:
                render_rich_text(p_item, f"{txt} [   ]", base_size=base_size)

        if is_teacher and sol_lines:
            add_solution_block(doc, "\n".join(sol_lines), base_size=10)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=3)

def render_short_list(doc, questions, is_teacher, is_screen, base_size=10.5):
    """Sinh danh sách câu hỏi trả lời ngắn."""
    for q in questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*Câu {q['id']}:* {q['content']}", base_size=base_size)

        if "image" in q and q["image"]:
            add_image_figure(doc, q["image"], width_cm=11.0)

        if is_teacher:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Cm(0.5)
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(2)
            r_ans_lbl = p_ans.add_run("Đáp số / Trả lời: ")
            set_run_style(r_ans_lbl, font_size=base_size, bold=True, color=COLOR_CORRECT)
            render_rich_text(p_ans, f"*{q['answer']}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)

            if "solution" in q and q["solution"]:
                add_solution_block(doc, q["solution"], base_size=10)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=2)
        else:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Cm(0.5)
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(2)
            render_rich_text(p_ans, "*Đáp án:* .................................................................................................................................................", base_size=base_size)

def render_essay_list(doc, questions, is_teacher, is_screen, base_size=10.5):
    """Sinh danh sách câu hỏi tự luận."""
    for q in questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*Câu {q['id']}:* {q['content']}", base_size=base_size)

        if is_teacher:
            if "image" in q and q["image"]:
                add_image_figure(doc, q["image"], width_cm=7.5)
            add_solution_block(doc, q["solution"], base_size=10)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=5)
        else:
            add_writing_lines(doc, num_lines=4)

def render_bio_questions(doc, questions, is_teacher, is_screen, base_size=10.5):
    """Sinh danh sách câu hỏi phân môn Sinh học."""
    current_level = None
    for q in questions:
        lvl = q.get("level")
        if lvl and lvl != current_level:
            current_level = lvl
            add_section_header(doc, current_level, base_size=11)

        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*Câu {q['id']}:* {q['content']}", base_size=base_size)

        if is_teacher:
            add_solution_block(doc, q["solution"], base_size=10)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=4)
        else:
            add_writing_lines(doc, num_lines=3)

def render_summary_table(doc):
    """Bảng tổng hợp đáp án nhanh ở cuối Bản Giáo Viên."""
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(12)
    p_sec.paragraph_format.space_after = Pt(4)
    render_rich_text(p_sec, "*BẢNG TỔNG HỢP ĐÁP ÁN TRẮC NGHIỆM*", base_size=12, base_bold=True, base_color=COLOR_HF_RED)

    # 1. Bài 1 & KHTN 1 MCQ
    p1 = doc.add_paragraph()
    p1.paragraph_format.space_before = Pt(4)
    p1.paragraph_format.space_after = Pt(2)
    render_rich_text(p1, "*1. Bài 1: Phương pháp và kĩ năng học tập môn KHTN*", base_size=10.5, base_bold=True, base_color=COLOR_PRIMARY)
    
    tbl1 = doc.add_table(rows=2, cols=10)
    tbl1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl1, color="B0BEC5", sz="4")
    for i in range(10):
        tbl1.cell(0, i).width = Cm(1.8)
        tbl1.cell(1, i).width = Cm(1.8)
        set_cell_shading(tbl1.cell(0, i), "E8F5E9")
        p_h = tbl1.cell(0, i).paragraphs[0]
        p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        render_rich_text(p_h, f"*{i+1}*", base_size=10, base_bold=True)
        
        p_a = tbl1.cell(1, i).paragraphs[0]
        p_a.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ans = BAI1_MCQ[i].get("correct", "") if i < len(BAI1_MCQ) else ""
        render_rich_text(p_a, f"*{ans}*", base_size=10, base_bold=True, base_color=COLOR_CORRECT)

    # 2. KHTN 1 (Hóa học) MCQ
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(6)
    p2.paragraph_format.space_after = Pt(2)
    render_rich_text(p2, "*2. Phân môn KHTN 1 (Hóa học) - Trắc nghiệm nhiều lựa chọn*", base_size=10.5, base_bold=True, base_color=COLOR_PRIMARY)

    tbl2 = doc.add_table(rows=4, cols=8)
    tbl2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl2, color="B0BEC5", sz="4")
    # Row 0 & 1: Cau 1-8
    for i in range(8):
        c_h, c_a = tbl2.cell(0, i), tbl2.cell(1, i)
        c_h.width = Cm(2.25); c_a.width = Cm(2.25)
        set_cell_shading(c_h, "E8F5E9")
        ph = c_h.paragraphs[0]; pa = c_a.paragraphs[0]
        ph.alignment = WD_ALIGN_PARAGRAPH.CENTER; pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
        render_rich_text(ph, f"*{i+1}*", base_size=10, base_bold=True)
        ans = KHTN1_MCQ[i]["correct"]
        render_rich_text(pa, f"*{ans}*", base_size=10, base_bold=True, base_color=COLOR_CORRECT)
    # Row 2 & 3: Cau 9-16
    for i in range(8):
        idx = i + 8
        c_h, c_a = tbl2.cell(2, i), tbl2.cell(3, i)
        c_h.width = Cm(2.25); c_a.width = Cm(2.25)
        set_cell_shading(c_h, "E8F5E9")
        ph = c_h.paragraphs[0]; pa = c_a.paragraphs[0]
        ph.alignment = WD_ALIGN_PARAGRAPH.CENTER; pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
        render_rich_text(ph, f"*{idx+1}*", base_size=10, base_bold=True)
        ans = KHTN1_MCQ[idx]["correct"]
        render_rich_text(pa, f"*{ans}*", base_size=10, base_bold=True, base_color=COLOR_CORRECT)

    # 3. KHTN 2 (Vật lí) MCQ
    p3 = doc.add_paragraph()
    p3.paragraph_format.space_before = Pt(6)
    p3.paragraph_format.space_after = Pt(2)
    render_rich_text(p3, "*3. Phân môn KHTN 2 (Vật lí) - Trắc nghiệm nhiều lựa chọn*", base_size=10.5, base_bold=True, base_color=COLOR_PRIMARY)

    tbl3 = doc.add_table(rows=2, cols=10)
    tbl3.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl3, color="B0BEC5", sz="4")
    for i in range(10):
        tbl3.cell(0, i).width = Cm(1.8)
        tbl3.cell(1, i).width = Cm(1.8)
        set_cell_shading(tbl3.cell(0, i), "E8F5E9")
        ph = tbl3.cell(0, i).paragraphs[0]; pa = tbl3.cell(1, i).paragraphs[0]
        ph.alignment = WD_ALIGN_PARAGRAPH.CENTER; pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
        render_rich_text(ph, f"*{i+1}*", base_size=10, base_bold=True)
        ans = KHTN2_MCQ[i]["correct"]
        render_rich_text(pa, f"*{ans}*", base_size=10, base_bold=True, base_color=COLOR_CORRECT)

def generate_exam_document(is_teacher=False, is_screen=False):
    """Hàm cốt lõi sinh tài liệu Word theo cấu trúc đầy đủ."""
    doc = docx.Document()
    setup_header_footer(doc)
    build_title_block(doc, is_teacher=is_teacher, is_screen=is_screen, base_size=10.5)

    # ==========================================================================
    # BÀI 1: PHƯƠNG PHÁP VÀ KĨ NĂNG HỌC TẬP MÔN KHOA HỌC TỰ NHIÊN
    # ==========================================================================
    add_chapter_title(doc, "BÀI 1: PHƯƠNG PHÁP VÀ KĨ NĂNG HỌC TẬP MÔN KHOA HỌC TỰ NHIÊN", base_size=11.5)
    render_mcq_list(doc, BAI1_MCQ, is_teacher, is_screen)

    # ==========================================================================
    # PHÂN MÔN KHTN 1: HÓA HỌC
    # ==========================================================================
    add_part_header(doc, "PHÂN MÔN KHTN 1 (HÓA HỌC)", base_size=12.5)
    add_chapter_title(doc, "I. NỘI DUNG ÔN TẬP", base_size=11.5)
    for t in KHTN1_THEORY:
        lines = [f"*{t['title']}*"] + [f"• {it}" for it in t["items"]]
        add_callout_box(doc, lines, base_size=10)

    add_chapter_title(doc, "II. TRẮC NGHIỆM", base_size=11.5)
    add_section_header(doc, "Phần 1: Câu hỏi nhiều lựa chọn", base_size=11)
    render_mcq_list(doc, KHTN1_MCQ, is_teacher, is_screen)

    add_section_header(doc, "Phần 2: Trắc nghiệm đúng sai", base_size=11)
    render_tf_list(doc, KHTN1_TF, is_teacher, is_screen)

    add_section_header(doc, "Phần 3: Câu hỏi trả lời ngắn", base_size=11)
    render_short_list(doc, KHTN1_SHORT, is_teacher, is_screen)

    add_chapter_title(doc, "III. TỰ LUẬN", base_size=11.5)
    render_essay_list(doc, KHTN1_ESSAY, is_teacher, is_screen)

    # ==========================================================================
    # PHÂN MÔN KHTN 2: VẬT LÍ
    # ==========================================================================
    add_part_header(doc, "PHÂN MÔN KHTN 2 (VẬT LÍ)", base_size=12.5)
    add_chapter_title(doc, "A. HỆ THỐNG KIẾN THỨC TRỌNG TÂM", base_size=11.5)
    for t in KHTN2_THEORY:
        lines = [f"*{t['title']}*"] + [f"• {it}" for it in t["items"]]
        add_callout_box(doc, lines, base_size=10)

    add_chapter_title(doc, "B. CÂU HỎI ÔN TẬP", base_size=11.5)
    add_chapter_title(doc, "PHẦN I: TRẮC NGHIỆM", base_size=11.5)
    add_section_header(doc, "Phần 1: Trắc nghiệm nhiều lựa chọn", base_size=11)
    render_mcq_list(doc, KHTN2_MCQ, is_teacher, is_screen)

    add_section_header(doc, "Phần 2: Trắc nghiệm đúng - sai", base_size=11)
    render_tf_list(doc, KHTN2_TF, is_teacher, is_screen)

    add_section_header(doc, "Phần 3: Trắc nghiệm trả lời ngắn", base_size=11)
    render_short_list(doc, KHTN2_SHORT, is_teacher, is_screen)

    add_chapter_title(doc, "PHẦN II: TỰ LUẬN", base_size=11.5)
    render_essay_list(doc, KHTN2_ESSAY, is_teacher, is_screen)

    # ==========================================================================
    # PHÂN MÔN KHTN 3: SINH HỌC
    # ==========================================================================
    add_part_header(doc, "PHÂN MÔN KHTN 3 (SINH HỌC)", base_size=12.5)
    add_chapter_title(doc, "PHẦN A. HỆ THỐNG KIẾN THỨC TRỌNG TÂM", base_size=11.5)
    for t in KHTN3_THEORY:
        lines = [f"*{t['title']}*"] + [f"• {it}" for it in t["items"]]
        add_callout_box(doc, lines, base_size=10)

    add_chapter_title(doc, "PHẦN B. CÂU HỎI ÔN TẬP VÀ GỢI Ý TRẢ LỜI", base_size=11.5)
    render_bio_questions(doc, KHTN3_QUESTIONS, is_teacher, is_screen)

    # Lưu ý khi ôn tập
    add_section_header(doc, "LƯU Ý KHI ÔN TẬP", base_size=11)
    for note in KHTN3_NOTES:
        p_note = doc.add_paragraph()
        p_note.paragraph_format.left_indent = Cm(0.5)
        p_note.paragraph_format.space_before = Pt(2)
        p_note.paragraph_format.space_after = Pt(2)
        render_rich_text(p_note, f"_{note}_", base_size=10, base_italic=True)

    # Bản Giáo viên có thêm Bảng tổng hợp đáp án ở cuối
    if is_teacher:
        render_summary_table(doc)

    return doc

def main():
    print("=" * 70)
    print("BẮT ĐẦU TẠO 3 FILE WORD CHUẨN FORM ĐỀ KIỂM TRA CHO KHTN 7 KIM LONG")
    print("=" * 70)

    configs = [
        {
            "name": "Bản Học sinh",
            "file": "De_Cuong_KHTN_7_Giua_HK1_Kim_Long_HocSinh.docx",
            "is_teacher": False,
            "is_screen": False
        },
        {
            "name": "Bản Giáo viên (Có Lời Giải Chi Tiết)",
            "file": "De_Cuong_KHTN_7_Giua_HK1_Kim_Long_GiaoVien.docx",
            "is_teacher": True,
            "is_screen": False
        },
        {
            "name": "Bản Dạy Trên Màn Hình (Chừa Chỗ Trống Giảng Dạy)",
            "file": "De_Cuong_KHTN_7_Giua_HK1_Kim_Long_DayTrenManHinh.docx",
            "is_teacher": False,
            "is_screen": True
        }
    ]

    generated_files = []

    for cfg in configs:
        print(f"\n[+] Đang khởi tạo: {cfg['name']} -> {cfg['file']}...")
        doc = generate_exam_document(is_teacher=cfg["is_teacher"], is_screen=cfg["is_screen"])
        out_path = os.path.join(OUT_DIR, cfg["file"])
        doc.save(out_path)
        print(f"    -> Đã lưu thành công tại: {out_path} ({os.path.getsize(out_path):,} bytes)")
        generated_files.append((cfg["file"], out_path))

    # Sao chép vào các thư mục OneDrive
    print("\n[+] Đồng bộ các file Word sang OneDrive:")
    for od in ONEDRIVE_DIRS:
        if os.path.exists(od):
            for fname, src_path in generated_files:
                dest_path = os.path.join(od, fname)
                shutil.copy2(src_path, dest_path)
                print(f"    -> Đã sao chép vào: {dest_path}")
        else:
            print(f"    [!] Thư mục OneDrive không tồn tại: {od}")

    print("\n" + "=" * 70)
    print("HOÀN THÀNH TẠO 3 BẢN WORD KHTN 7 KIM LONG THEO ĐÚNG YÊU CẦU!")
    print("=" * 70)

if __name__ == "__main__":
    main()
