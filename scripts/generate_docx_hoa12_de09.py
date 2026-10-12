# -*- coding: utf-8 -*-
"""
Script tạo file Word (.docx) chuẩn "FORM ĐỀ KIỂM TRA" cho:
ĐỀ KIỂM TRA SỐ 9 - CHƯƠNG 3: HỢP CHẤT CHỨA NITROGEN (HÓA HỌC 12)
Gồm 2 phiên bản:
1. De_Kiem_Tra_Hoa_12_De09.docx (Bản Học Sinh)
2. De_Kiem_Tra_Hoa_12_De09_GiaoVien.docx (Bản Giáo Viên có lời giải chi tiết)
Tuân thủ 100% quy chuẩn: Header & Footer, Tab Stops (NO TABLE trong đề/lời giải), Sub/Superscript công thức hóa học.
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
IMG_DIR = r"c:\Antigravity_Thanh\San_Pham\images_crop_de09"

COLOR_HF_RED = RGBColor(211, 47, 47)      # Coral Red (#D32F2F)
COLOR_HF_GREEN_HEX = "8BC390"             # Sage Green (#8BC390)
COLOR_PRIMARY = RGBColor(46, 125, 50)     # Dark Green (#2E7D32)
COLOR_CORRECT = RGBColor(198, 40, 40)     # Red for correct answers (#C62828)
COLOR_SOLUTION = RGBColor(21, 101, 192)   # Blue for solutions (#1565C0)
COLOR_GRAY = RGBColor(97, 97, 97)

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def set_run_style(run, font_name="Times New Roman", font_size=12, bold=False, italic=False, subscript=False, superscript=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def add_chem_run(p, text, base_font="Times New Roman", base_size=12, bold=False, italic=False, color=None):
    tokens = re.split(r'(\^\{[^}]+\}|\^[0-9\+\-]+|~\{[^}]+\}|~[0-9]+)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('^'):
            content = tok[2:-1] if tok.startswith('^{') else tok[1:]
            r = p.add_run(content)
            set_run_style(r, font_name=base_font, font_size=base_size, bold=bold, italic=italic, superscript=True, color=color)
        elif tok.startswith('~'):
            content = tok[2:-1] if tok.startswith('~{') else tok[1:]
            r = p.add_run(content)
            set_run_style(r, font_name=base_font, font_size=base_size, bold=bold, italic=italic, subscript=True, color=color)
        else:
            r = p.add_run(tok)
            set_run_style(r, font_name=base_font, font_size=base_size, bold=bold, italic=italic, color=color)

def setup_page(doc):
    for sec in doc.sections:
        sec.top_margin = Cm(1.5)
        sec.bottom_margin = Cm(1.5)
        sec.left_margin = Cm(1.5)
        sec.right_margin = Cm(1.5)
        sec.header_distance = Cm(0.6)
        sec.footer_distance = Cm(0.6)
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)

def set_cell_border_bottom(cell, color="8BC390", sz="6"):
    tcPr = cell._tc.get_or_add_tcPr()
    xml = f'<w:tcBorders {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/></w:tcBorders>'
    tcPr.append(parse_xml(xml))

def set_cell_border_top(cell, color="8BC390", sz="6"):
    tcPr = cell._tc.get_or_add_tcPr()
    xml = f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/></w:tcBorders>'
    tcPr.append(parse_xml(xml))

def set_cell_margins_zero(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    xml = f'<w:tcMar {nsdecls("w")}><w:top w:w="0" w:type="dxa"/><w:bottom w:w="0" w:type="dxa"/><w:left w:w="0" w:type="dxa"/><w:right w:w="0" w:type="dxa"/></w:tcMar>'
    tcPr.append(parse_xml(xml))

def setup_header_footer(doc):
    sec = doc.sections[0]
    
    # HEADER
    hdr = sec.header
    htable = hdr.add_table(rows=1, cols=2, width=Cm(18.0))
    htable.autofit = False
    htable.columns[0].width = Cm(9.0)
    htable.columns[1].width = Cm(9.0)
    
    cell_l, cell_r = htable.rows[0].cells[0], htable.rows[0].cells[1]
    set_cell_margins_zero(cell_l)
    set_cell_margins_zero(cell_r)
    set_cell_border_bottom(cell_l, color=COLOR_HF_GREEN_HEX, sz="6")
    set_cell_border_bottom(cell_r, color=COLOR_HF_GREEN_HEX, sz="6")
    
    p_hl = cell_l.paragraphs[0]
    p_hl.paragraph_format.space_before = Pt(0)
    p_hl.paragraph_format.space_after = Pt(2)
    p_hl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_hl = p_hl.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)
    
    p_hr = cell_r.paragraphs[0]
    p_hr.paragraph_format.space_before = Pt(0)
    p_hr.paragraph_format.space_after = Pt(2)
    p_hr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_hr1 = p_hr.add_run("Thầy ")
    set_run_style(r_hr1, font_size=10, color=COLOR_HF_RED)
    r_hr2 = p_hr.add_run("TRẦN VĂN THẠNH")
    set_run_style(r_hr2, font_size=10, bold=True, color=COLOR_HF_RED)
    r_hr3 = p_hr.add_run(" - 0777.470.803")
    set_run_style(r_hr3, font_size=10, color=COLOR_HF_RED)
    
    # FOOTER
    ftr = sec.footer
    ftable = ftr.add_table(rows=1, cols=2, width=Cm(18.0))
    ftable.autofit = False
    ftable.columns[0].width = Cm(14.0)
    ftable.columns[1].width = Cm(4.0)
    
    cell_fl, cell_fr = ftable.rows[0].cells[0], ftable.rows[0].cells[1]
    set_cell_margins_zero(cell_fl)
    set_cell_margins_zero(cell_fr)
    set_cell_border_top(cell_fl, color=COLOR_HF_GREEN_HEX, sz="6")
    set_cell_border_top(cell_fr, color=COLOR_HF_GREEN_HEX, sz="6")
    
    p_fl = cell_fl.paragraphs[0]
    p_fl.paragraph_format.space_before = Pt(2)
    p_fl.paragraph_format.space_after = Pt(0)
    p_fl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_fl = p_fl.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)
    
    p_fr = cell_fr.paragraphs[0]
    p_fr.paragraph_format.space_before = Pt(2)
    p_fr.paragraph_format.space_after = Pt(0)
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_field(p_fr, font_size=10, bold=True, color=COLOR_HF_RED)

def add_header_block(doc, is_teacher=False):
    lines = [
        ("SỞ GIÁO DỤC VÀ ĐÀO TẠO THỪA THIÊN HUẾ", "ĐỀ KIỂM TRA ĐỊNH KỲ CHƯƠNG 3 (HỢP CHẤT CHỨA NITROGEN)", True, True, 10.5, 11),
        ("ĐỀ CHÍNH THỨC", "MÔN: HOÁ HỌC – LỚP 12", True, True, 11, 11),
        ("(Đề thi có 04 trang)", "Thời gian làm bài: 50 phút (không kể thời gian phát đề)", False, True, 10, 10.5),
        ("", "Mã đề thi: 109", False, True, 10, 11)
    ]
    
    for left_text, right_text, l_bold, r_bold, l_size, r_size in lines:
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.05
        
        if left_text:
            r_l = p.add_run(left_text)
            set_run_style(r_l, font_size=l_size, bold=l_bold, italic=("Đề thi có" in left_text))
        
        p.add_run("\t")
        
        if right_text:
            r_r = p.add_run(right_text)
            set_run_style(r_r, font_size=r_size, bold=r_bold, italic=("Thời gian" in right_text), color=COLOR_PRIMARY if "CHƯƠNG 3" in right_text else None)

    if is_teacher:
        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_t.paragraph_format.space_before = Pt(2)
        p_t.paragraph_format.space_after = Pt(2)
        r_t = p_t.add_run("(BẢN GIÁO VIÊN — CÓ LỜI GIẢI CHI TIẾT)")
        set_run_style(r_t, font_size=11, bold=True, color=COLOR_CORRECT)

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(4)
    p_info.paragraph_format.space_after = Pt(2)
    r_info = p_info.add_run("Họ, tên thí sinh: .......................................................................... Lớp: ................ Số báo danh: ...............")
    set_run_style(r_info, font_size=10.5, italic=True)
    
    p_mass = doc.add_paragraph()
    p_mass.paragraph_format.space_before = Pt(0)
    p_mass.paragraph_format.space_after = Pt(6)
    r_mass = p_mass.add_run("(Cho biết nguyên tử khối: H = 1; C = 12; N = 14; O = 16; Na = 23; S = 32; Cl = 35,5; Fe = 56; Cu = 64).")
    set_run_style(r_mass, font_size=10, italic=True)

def add_part_heading(doc, part_title, part_desc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.15
    r_title = p.add_run(part_title)
    set_run_style(r_title, font_size=11.5, bold=True, color=COLOR_PRIMARY)
    
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_before = Pt(0)
    p_desc.paragraph_format.space_after = Pt(4)
    r_desc = p_desc.add_run(part_desc)
    set_run_style(r_desc, font_size=10.5, italic=True)

def add_solution(doc, sol_text):
    p_sol = doc.add_paragraph()
    p_sol.paragraph_format.space_before = Pt(2)
    p_sol.paragraph_format.space_after = Pt(4)
    p_sol.paragraph_format.left_indent = Cm(0.6)
    p_sol.paragraph_format.line_spacing = 1.15
    
    r_hdr = p_sol.add_run("Lời giải chi tiết:\n")
    set_run_style(r_hdr, font_size=10.5, bold=True, italic=True, color=COLOR_SOLUTION)
    
    lines = sol_text.strip().split("\n")
    for i, line in enumerate(lines):
        line = line.strip()
        if not line:
            continue
        if i > 0:
            p_line = doc.add_paragraph()
            p_line.paragraph_format.space_before = Pt(0)
            p_line.paragraph_format.space_after = Pt(1.5)
            p_line.paragraph_format.left_indent = Cm(0.6)
            p_line.paragraph_format.line_spacing = 1.15
            add_chem_run(p_line, line, base_size=10.5, italic=False)
        else:
            add_chem_run(p_sol, line, base_size=10.5, italic=False)

def add_mcq_options(doc, opt_a, opt_b, opt_c, opt_d, correct_idx=None, is_teacher=False):
    opts = [('A', opt_a), ('B', opt_b), ('C', opt_c), ('D', opt_d)]
    max_len = max(len(text) for _, text in opts)
    
    if max_len < 22:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.tab_stops.add_tab_stop(Cm(4.5), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(13.5), WD_TAB_ALIGNMENT.LEFT)
        
        for idx, (label, text) in enumerate(opts):
            if idx > 0:
                p.add_run("\t")
            is_corr = (is_teacher and idx == correct_idx)
            r_lbl = p.add_run(f"{label}. ")
            set_run_style(r_lbl, font_size=11, bold=True, color=COLOR_CORRECT if is_corr else None)
            add_chem_run(p, text, base_size=11, bold=is_corr, color=COLOR_CORRECT if is_corr else None)
    elif max_len < 46:
        for row in [(0, 1), (2, 3)]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
            
            for col_idx, opt_num in enumerate(row):
                if col_idx > 0:
                    p.add_run("\t")
                label, text = opts[opt_num]
                is_corr = (is_teacher and opt_num == correct_idx)
                r_lbl = p.add_run(f"{label}. ")
                set_run_style(r_lbl, font_size=11, bold=True, color=COLOR_CORRECT if is_corr else None)
                add_chem_run(p, text, base_size=11, bold=is_corr, color=COLOR_CORRECT if is_corr else None)
    else:
        for idx, (label, text) in enumerate(opts):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Cm(0.6)
            is_corr = (is_teacher and idx == correct_idx)
            r_lbl = p.add_run(f"{label}. ")
            set_run_style(r_lbl, font_size=11, bold=True, color=COLOR_CORRECT if is_corr else None)
            add_chem_run(p, text, base_size=11, bold=is_corr, color=COLOR_CORRECT if is_corr else None)

def build_docx_de09(is_teacher=False):
    doc = docx.Document()
    setup_page(doc)
    setup_header_footer(doc)
    add_header_block(doc, is_teacher=is_teacher)
    
    # ==================== PHẦN I ====================
    add_part_heading(
        doc,
        "PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.",
        "Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án."
    )
    
    part1_data = [
        (
            1,
            "(SBT – KNTT). Trong phân tử amine, nguyên tử nitrogen có số cặp electron chưa liên kết là",
            None,
            "một cặp.", "hai cặp.", "ba cặp.", "không cặp.",
            0,
            "Nguyên tử nitrogen có cấu hình e lớp ngoài cùng là 2s^2 2p^3. Trong phân tử amine, nitrogen tạo 3 liên kết cộng hóa trị với các nguyên tử khác và còn lại đúng 1 cặp electron tự do chưa liên kết."
        ),
        (
            2,
            "Cho dung dịch methylamine phản ứng với dung dịch nào sau đây tạo được kết tủa?",
            None,
            "NaCl.", "Ba(NO~3~)~2~.", "FeCl~3~.", "K~2~SO~4~.",
            2,
            "Dung dịch methylamine có tính base do tạo ion OH^-^ trong nước. Khi tác dụng với dung dịch FeCl~3~, ion OH^-^ tạo kết tủa nâu đỏ Fe(OH)~3~:\n"
            "3CH~3~NH~2~ + FeCl~3~ + 3H~2~O -> Fe(OH)~3~↓ (nâu đỏ) + 3CH~3~NH~3~Cl."
        ),
        (
            3,
            "(SBT – CTST). Phát biểu nào sau đây không đúng về ứng dụng của amine?",
            None,
            "Các amine đều độc, chủ yếu dùng sản xuất thuốc diệt nấm mốc, kháng sinh.",
            "Amine được sử dụng nhiều trong bào chế dược phẩm, vitamin.",
            "Nhiều polymer như nylon-6,6, polyurethane có thể tổng hợp từ amine.",
            "Phẩm nhuộm azo và dược phẩm là ứng dụng quan trọng của aniline.",
            0,
            "Phát biểu A không đúng vì amine có rất nhiều ứng dụng quan trọng trong công nghiệp tổng hợp polymer, thuốc nhuộm và dược phẩm, chứ không phải chủ yếu dùng sản xuất thuốc diệt nấm mốc."
        ),
        (
            4,
            "Số nguyên tử oxygen có trong một phân tử glutamic acid là",
            None,
            "1.", "3.", "4.", "2.",
            2,
            "Glutamic acid có công thức HOOC-CH~2~-CH~2~-CH(NH~2~)-COOH (C~5~H~9~NO~4~). Phân tử có 2 nhóm carboxyl (-COOH) nên chứa 2 * 2 = 4 nguyên tử oxygen."
        ),
        (
            5,
            "(SBT – CĐ). Các amino acid tồn tại ở trạng thái ion lưỡng cực, do đó chúng",
            None,
            "có nhiệt độ nóng chảy cao và tan tốt trong nước.",
            "có nhiệt độ nóng chảy cao và ít tan trong nước.",
            "dễ nóng chảy và tan tốt trong nước.",
            "dễ nóng chảy và ít tan trong nước.",
            0,
            "Ở trạng thái ion lưỡng cực (mạng tinh thể ion với lực hút tĩnh điện mạnh), các amino acid là chất rắn có nhiệt độ nóng chảy cao và dễ tan trong nước do tương tác solvat hóa mạnh với các phân tử nước phân cực."
        ),
        (
            6,
            "Cho phản ứng sau:\n"
            "    H~2~NCH~2~COOH + C~2~H~5~OH  <-{HCl, t°}->  H~2~NCH~2~COOC~2~H~5~ + H~2~O\n"
            "Phản ứng trên thể hiện tính chất nào của amino acid cụ thể là glycine?",
            None,
            "Lưỡng tính.", "Acid mạnh.", "Base mạnh.", "Ester hóa.",
            3,
            "Phản ứng giữa nhóm carboxyl (-COOH) của amino acid với alcohol (C~2~H~5~OH) có xúc tác acid vô cơ tạo ester và nước là phản ứng ester hóa."
        ),
        (
            7,
            "Biểu diễn dạng kí hiệu của peptide: HOOC-CH~2~-NH-CO-CH(CH~3~)NH~2~ là:",
            None,
            "Val-Ala.", "Ala-Val.", "Ala-Gly.", "Gly-Ala.",
            2,
            "Theo quy ước, peptide được viết từ amino acid đầu N (nhóm amino tự do bên trái) sang đầu C (nhóm carboxyl tự do bên phải).\n"
            "- Đầu N có nhóm amino tự do là -CH(CH~3~)NH~2~ (gốc Alanine - Ala).\n"
            "- Đầu C có nhóm carboxyl tự do là HOOC-CH~2~-NH- (gốc Glycine - Gly).\n"
            "Do đó kí hiệu của peptide là Ala-Gly."
        ),
        (
            8,
            "Trong phân tử tetrapeptide Gly-Glu-Ala-Val, amino acid đầu N là",
            None,
            "Valine.", "Alanine.", "Acid glutamic.", "Glycine.",
            3,
            "Amino acid đầu N là amino acid còn nhóm amino tự do, theo quy ước được viết ở vị trí đầu tiên bên trái chuỗi peptide. Trong Gly-Glu-Ala-Val, amino acid đầu N là Glycine."
        ),
        (
            9,
            "Số gốc α-amino acid trong phân tử tripeptide mạch hở là:",
            None,
            "1.", "3.", "2.", "4.",
            1,
            "Tripeptide là hợp chất peptide được cấu tạo từ 3 gốc α-amino acid liên kết với nhau bằng 2 liên kết peptide."
        ),
        (
            10,
            "Các protein khác nhau đảm nhận nhiều vai trò thiết yếu khác nhau với sự sống như xây dựng tế bào, xúc tác sinh hoá, điều hoà chuyển hoá chất, vận chuyển chất, kháng thể... Protein là hợp chất cao phân tử được tạo thành từ một hay nhiều",
            None,
            "mắc xích.", "chuỗi polypeptide.", "monomer.", "phân tử.",
            1,
            "Protein là các đại phân tử sinh học (polymer sinh học) được cấu tạo từ một hay nhiều chuỗi polypeptide, mỗi chuỗi gồm hàng chục đến hàng nghìn gốc α-amino acid."
        ),
        (
            11,
            "(SBT – CTST). Chất có khả năng tạo phức với methylamine và ethylamine trong các chất sau đây là",
            None,
            "Ca(OH)~2~.", "Cu(OH)~2~.", "Al(OH)~3~.", "KOH.",
            1,
            "Các alkylamine mạch ngắn như methylamine, ethylamine dùng cặp electron tự do trên N hòa tan Cu(OH)~2~ tạo phức chất [Cu(RNH~2~)~4~](OH)~2~ có màu xanh lam đặc trưng."
        ),
        (
            12,
            "(SBT – CTST). Có bao nhiêu amino acid cần thiết phải cung cấp cho cơ thể thông qua thực phẩm, dinh dưỡng?",
            None,
            "9.", "20.", "10.", "18.",
            0,
            "Trong số 20 loại α-amino acid tham gia cấu tạo protein, cơ thể người có thể tự tổng hợp được 11 loại, còn 9 loại amino acid thiết yếu (essential amino acids) cơ thể không tự tổng hợp được mà phải thu nhận qua chế độ dinh dưỡng hàng ngày."
        ),
        (
            13,
            "(SBT – CĐ). Cho các chất có công thức cấu tạo sau:\n"
            "Trong các hợp chất trên, những hợp chất nào thuộc loại dipeptide?",
            ("cau13_dipeptides.png", 13.5),
            "Hợp chất (1) và (2).", "Hợp chất (1) và (3).", "Hợp chất (2) và (3).", "Hợp chất (2) và (4).",
            0,
            "- Chất (1) cấu tạo từ Ala và Gly (đều là α-amino acid) liên kết qua liên kết peptide -> dipeptide Ala-Gly.\n"
            "- Chất (2) cấu tạo từ 2 gốc Ala -> dipeptide Ala-Ala.\n"
            "- Chất (3) có liên kết ester -COO-, không có liên kết peptide.\n"
            "- Chất (4) có nhóm amino ở vị trí γ (không phải α-amino acid) nên không phải peptide chuẩn.\n"
            "=> Các hợp chất thuộc loại dipeptide là (1) và (2)."
        ),
        (
            14,
            "(SBT – CĐ). Trong cấu trúc phân tử của chất cho ở hình bên, liên kết peptide là",
            ("cau14_peptide_bond.png", 6.5),
            "liên kết (1).", "liên kết (3).", "liên kết (2).", "liên kết (4).",
            2,
            "Liên kết peptide là liên kết amide -CO-NH- giữa nhóm carboxyl của một α-amino acid với nhóm amino của một α-amino acid khác. Trên hình vẽ, mũi tên (2) chỉ trực tiếp vào liên kết cộng hóa trị giữa carbon của nhóm carbonyl (C=O) và nguyên tử nitrogen (N-H). Do đó (2) là liên kết peptide."
        ),
        (
            15,
            "(SBT – CTST). Loại dinh dưỡng nào sau đây không được cơ thể dự trữ để sử dụng?",
            None,
            "Tinh bột.", "Chất béo.", "Amino acid.", "Acid béo.",
            2,
            "Cơ thể người có thể dự trữ glycogen (tinh bột động vật ở gan và cơ) và lipid (mô mỡ). Cơ thể không có cơ chế dự trữ amino acid tự do; lượng amino acid dư thừa sẽ bị khử nhóm amino tạo urea thải ra ngoài, khung carbon chuyển thành năng lượng, glucose hoặc chất béo."
        ),
        (
            16,
            "(SBT – CĐ). Cho các hợp chất có công thức cấu tạo dưới đây:\n"
            "Những hợp chất nào trong số các chất trên thuộc loại α-amino acid?",
            ("cau16_amino_acids.png", 14.0),
            "Chất (2), chất (3) và chất (4).", "Chất (1) và chất (2).", "Chất (1) và chất (3).", "Chất (1), chất (2) và chất (4).",
            3,
            "α-amino acid là hợp chất có nhóm amino (-NH~2~) gắn trực tiếp vào carbon số 2 (carbon α, liền kề nhóm -COOH).\n"
            "- Chất (1): 2-aminobutanoic acid -> α-amino acid.\n"
            "- Chất (2): 2-aminopropanoic acid (alanine) -> α-amino acid.\n"
            "- Chất (3): 2-hydroxypropanamide -> amide, không phải amino acid.\n"
            "- Chất (4): ^+^H~3~N-CH~2~-COO^-^ (glycine dạng ion lưỡng cực) -> α-amino acid.\n"
            "=> Các chất thuộc loại α-amino acid là (1), (2) và (4)."
        ),
        (
            17,
            "(SBT – CĐ). Chất nào dưới đây không phải là amino acid?",
            None,
            "Lysine.", "Glycine.", "Aniline.", "Glutamic acid.",
            2,
            "Aniline (C~6~H~5~NH~2~) là một arylamine (amine thơm), không có nhóm carboxyl (-COOH) nên không phải là amino acid. Lysine, glycine, glutamic acid đều là các amino acid."
        ),
        (
            18,
            "Cho các phát biểu sau:\n"
            "(a) Sục khí CH~3~NH~2~ vào dung dịch FeCl~3~ có sinh ra kết tủa màu xanh.\n"
            "(b) Không thể dùng quỳ tím để phân biệt ba dung dịch: alanine, lysine, glutamic acid.\n"
            "(c) Dung dịch glutamic acid đổi màu quỳ tím thành xanh.\n"
            "(d) Nhỏ vài giọt nước bromine vào dung dịch aniline thu được kết tủa đỏ nâu.\n"
            "(e) Cho dung dịch HCl vào dung dịch glutamic acid thì không có phản ứng xảy ra.\n"
            "(g) Methylamine có lực base yếu hơn ammonia.\n"
            "Số phát biểu đúng là:",
            None,
            "5.", "4.", "0.", "3.",
            2,
            "(a) Sai, tạo kết tủa Fe(OH)~3~ màu nâu đỏ (không phải màu xanh).\n"
            "(b) Sai, dùng được quỳ tím: alanine không đổi màu, lysine làm quỳ hóa xanh, glutamic acid làm quỳ hóa đỏ.\n"
            "(c) Sai, glutamic acid làm quỳ tím đổi màu đỏ (hồng).\n"
            "(d) Sai, aniline tác dụng nước bromine tạo kết tủa trắng 2,4,6-tribromoaniline.\n"
            "(e) Sai, glutamic acid có nhóm -NH~2~ mang tính base nên tác dụng với HCl.\n"
            "(g) Sai, methylamine có nhóm methyl đẩy electron nên lực base mạnh hơn ammonia.\n"
            "=> Cả 6 phát biểu đều sai. Số phát biểu đúng là 0."
        )
    ]
    
    for item in part1_data:
        num, text, img_info, oa, ob, oc, od, corr_idx, sol = item
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        
        r_num = p.add_run(f"Câu {num}. ")
        set_run_style(r_num, font_size=11, bold=True)
        add_chem_run(p, text, base_size=11)
        
        if img_info:
            img_file, img_w = img_info
            img_path = os.path.join(IMG_DIR, img_file)
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(2)
                p_img.paragraph_format.space_after = Pt(3)
                p_img.add_run().add_picture(img_path, width=Cm(img_w))
                
        add_mcq_options(doc, oa, ob, oc, od, correct_idx=corr_idx, is_teacher=is_teacher)
            
        if is_teacher and sol:
            add_solution(doc, sol)

    # ==================== PHẦN II ====================
    add_part_heading(
        doc,
        "PHẦN II. Câu trắc nghiệm đúng sai.",
        "Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai."
    )
    
    part2_data = [
        (
            1,
            "(SBT – KNTT). Các phát biểu về tính chất hoá học của dung dịch methylamine:",
            None,
            [
                ("a", "Phản ứng với HCl tạo thành CH~3~NH~3~Cl.", True),
                ("b", "Hoà tan Cu(OH)~2~ tạo thành [Cu(CH~3~NH~2~)~4~](OH)~2~.", True),
                ("c", "Phản ứng với FeCl~3~ tạo thành kết tủa Fe(OH)~3~.", True),
                ("d", "Phản ứng với HNO~2~ tạo thành CH~3~N~2~^+^.", False)
            ],
            (
                "a) Đúng. CH~3~NH~2~ + HCl -> CH~3~NH~3~Cl (methylammonium chloride).\n"
                "b) Đúng. Methylamine hòa tan Cu(OH)~2~ tạo phức chất màu xanh lam thẫm: Cu(OH)~2~ + 4CH~3~NH~2~ -> [Cu(CH~3~NH~2~)~4~](OH)~2~.\n"
                "c) Đúng. 3CH~3~NH~2~ + FeCl~3~ + 3H~2~O -> Fe(OH)~3~↓ (nâu đỏ) + 3CH~3~NH~3~Cl.\n"
                "d) Sai. Alkylamine bậc 1 phản ứng với HNO~2~ tạo alcohol và giải phóng khí N~2~ ngay ở nhiệt độ phòng: CH~3~NH~2~ + HNO~2~ -> CH~3~OH + N~2~↑ + H~2~O. Muối diazonium CH~3~N~2~^+^ rất không bền, bị phân hủy ngay lập tức."
            )
        ),
        (
            2,
            "(SBT – CTST). Cho một số hợp chất chứa nguyên tố nitrogen như sau:",
            ("p2_cau2_compounds.png", 12.5),
            [
                ("a", "Dễ tan trong nước.", True),
                ("b", "Ở điều kiện thường, có 1 chất là chất khí, 1 chất lỏng và 1 chất rắn.", False),
                ("c", "Chất có nhiệt độ nóng chảy thấp nhất là CH~3~CH~2~NH~2~, cao nhất là HOOCCH~2~NHCOCH~2~NH~2~.", True),
                ("d", "Khả năng tan tốt trong nước của các chất là do sự phân cực mạnh trong phân tử.", True)
            ],
            (
                "Các chất trong hình gồm:\n"
                "- Chất (1): CH~3~-CH~2~-NH~2~ (ethylamine: chất khí ở điều kiện thường).\n"
                "- Chất (2): HOOC-CH~2~-CH~2~-NH~2~ (β-alanine: amino acid ở dạng ion lưỡng cực, chất rắn).\n"
                "- Chất (3): HOOC-CH~2~-NH-CO-CH~2~-NH~2~ (Gly-Gly: dipeptide ở dạng ion lưỡng cực, chất rắn).\n"
                "a) Đúng. Cả 3 chất đều phân cực mạnh và tạo liên kết hydrogen với nước nên dễ tan trong nước.\n"
                "b) Sai. Ở điều kiện thường có 1 chất khí (ethylamine) và 2 chất rắn (β-alanine và Gly-Gly), không có chất lỏng.\n"
                "c) Đúng. Ethylamine là chất khí nên nhiệt độ nóng chảy thấp nhất; dipeptide Gly-Gly có phân tử khối lớn nhất và mạng tinh thể ion/hydrogen dày đặc nhất nên nhiệt độ nóng chảy cao nhất.\n"
                "d) Đúng. Sự phân cực mạnh của các nhóm chức và ion lưỡng cực giúp các chất tan tốt trong dung môi phân cực như nước."
            )
        ),
        (
            3,
            "Một peptide có cấu trúc như sau:",
            ("p2_cau3_tripeptide.png", 12.0),
            [
                ("a", "Peptide trên chứa alanine, glutamic acid và glycine.", True),
                ("b", "Peptide trên có 3 liên kết peptide trong phân tử trên.", False),
                ("c", "Peptide trên có phản ứng màu biuret.", True),
                ("d", "Peptide trên là tetrapeptide.", False)
            ],
            (
                "Phân tích công thức cấu tạo của peptide:\n"
                "- Gốc 1 (đầu N): H~2~N-CH(CH~3~)-CO- là gốc của Alanine (Ala).\n"
                "- Gốc 2: -NH-CH(CH~2~CH~2~COOH)-CO- là gốc của Glutamic acid (Glu).\n"
                "- Gốc 3 (đầu C): -NH-CH~2~-COOH là gốc của Glycine (Gly).\n"
                "Do đó peptide trên là tripeptide Ala-Glu-Gly.\n"
                "a) Đúng. Peptide được cấu tạo từ 3 amino acid: Ala, Glu và Gly.\n"
                "b) Sai. Tripeptide chỉ có 3 - 1 = 2 liên kết peptide (-CO-NH-).\n"
                "c) Đúng. Từ tripeptide trở lên (có từ 2 liên kết peptide) đều cho phản ứng màu biuret tím với Cu(OH)~2~.\n"
                "d) Sai. Peptide được tạo từ 3 gốc amino acid nên là tripeptide, không phải tetrapeptide."
            )
        ),
        (
            4,
            "Bromelain và papain là những enzyme có tác dụng thuỷ phân protein. Bromelain có nhiều trong quả dứa (thơm). Bromelain được sử dụng để làm thuốc giảm sưng (viêm), đặc biệt là ở mũi và xoang sau phẫu thuật hoặc chấn thương, điều trị bệnh viêm loét đại tràng,...",
            ("p2_cau4_bromelain.png", 9.0),
            [
                ("a", "Thịt được ướp với nước ép dứa thì khi nấu sẽ nhanh mềm hơn.", True),
                ("b", "Do enzyme có nhiều trong quả dứa có tác dụng phân giải protein có trong thịt nên khiến cho thịt khi nấu cùng sẽ nhanh mềm hơn.", True),
                ("c", "Từ dứa có thể chiết xuất được Bromelain.", True),
                ("d", "Bromelain có chứa 7 nguyên tử oxygen.", False)
            ],
            (
                "a) Đúng. Nước ép dứa chứa enzyme bromelain giúp làm mềm thịt hiệu quả.\n"
                "b) Đúng. Enzyme bromelain xúc tác phản ứng thủy phân các liên kết peptide trong protein của mô cơ thịt, làm đứt gãy sợi protein giúp thịt mềm nhanh khi nấu.\n"
                "c) Đúng. Bromelain được chiết xuất chủ yếu từ thân và quả cây dứa (thơm).\n"
                "d) Sai. Bromelain là một enzyme phức tạp có bản chất là protein đại phân tử (gồm hàng trăm gốc amino acid), do đó phân tử chứa hàng trăm đến hàng nghìn nguyên tử oxygen, không thể chỉ có 7 nguyên tử oxygen."
            )
        )
    ]
    
    for item in part2_data:
        num, text, img_info, statements, sol = item
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        
        r_num = p.add_run(f"Câu {num}. ")
        set_run_style(r_num, font_size=11, bold=True)
        add_chem_run(p, text, base_size=11)
        
        if img_info:
            img_file, img_w = img_info
            img_path = os.path.join(IMG_DIR, img_file)
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(2)
                p_img.paragraph_format.space_after = Pt(3)
                p_img.add_run().add_picture(img_path, width=Cm(img_w))
                
        if num == 2:
            p_lead = doc.add_paragraph()
            p_lead.paragraph_format.space_before = Pt(1)
            p_lead.paragraph_format.space_after = Pt(2)
            add_chem_run(p_lead, "Về tính chất vật lí của các chất, xét tính đúng/sai của các phát biểu sau:", base_size=11, italic=True)
            
        for stmt_label, stmt_text, is_correct in statements:
            p_stmt = doc.add_paragraph()
            p_stmt.paragraph_format.space_before = Pt(1)
            p_stmt.paragraph_format.space_after = Pt(1.5)
            p_stmt.paragraph_format.line_spacing = 1.15
            p_stmt.paragraph_format.left_indent = Cm(0.6)
            
            r_slbl = p_stmt.add_run(f"{stmt_label}) ")
            set_run_style(r_slbl, font_size=11, bold=True)
            add_chem_run(p_stmt, stmt_text, base_size=11)
            
            if is_teacher:
                status_txt = " — (Đúng)" if is_correct else " — (Sai)"
                r_status = p_stmt.add_run(status_txt)
                set_run_style(r_status, font_size=11, bold=True, color=COLOR_CORRECT)
                
        if is_teacher and sol:
            add_solution(doc, sol)

    # ==================== PHẦN III ====================
    add_part_heading(
        doc,
        "PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn.",
        "Thí sinh trả lời từ câu 1 đến câu 6."
    )
    
    part3_data = [
        (
            1,
            "(SBT – CTST). Các amine có tính base nên dễ phản ứng với dung dịch acid. Khi cho 1 mol methylenediamine phản ứng với dung dịch acid HCl dư, số mol HCl tham gia phản ứng là bao nhiêu?",
            "2",
            "Methylenediamine có công thức là H~2~N-CH~2~-NH~2~, chứa 2 nhóm chức amine bậc 1.\n"
            "Phương trình phản ứng:\n"
            "    H~2~N-CH~2~-NH~2~ + 2HCl -> ClH~3~N-CH~2~-NH~3~Cl\n"
            "=> 1 mol methylenediamine phản ứng với tối đa 2 mol HCl.\n"
            "**Đáp số: 2.**"
        ),
        (
            2,
            "Trong các dung dịch: CH~3~-CH~2~-NH~2~, H~2~N-CH~2~-COOH, H~2~N-CH~2~-CH(NH~2~)-COOH, HOOC-CH~2~-CH~2~-CH(NH~2~)-COOH, số dung dịch làm xanh quỳ tím là",
            "2",
            "Xét sự đổi màu quỳ tím của các dung dịch:\n"
            "- CH~3~CH~2~NH~2~ (ethylamine): Amine có tính base -> làm xanh quỳ tím.\n"
            "- H~2~N-CH~2~-COOH (glycine: 1 nhóm NH~2~, 1 nhóm COOH) -> môi trường trung tính, không đổi màu quỳ tím.\n"
            "- H~2~N-CH~2~-CH(NH~2~)-COOH (diaminopropanoic acid: 2 nhóm NH~2~, 1 nhóm COOH) -> tính base chiếm ưu thế -> làm xanh quỳ tím.\n"
            "- HOOC-CH~2~-CH~2~-CH(NH~2~)-COOH (glutamic acid: 1 nhóm NH~2~, 2 nhóm COOH) -> tính acid chiếm ưu thế -> làm đỏ quỳ tím.\n"
            "=> Có 2 dung dịch làm xanh quỳ tím: CH~3~-CH~2~-NH~2~ và H~2~N-CH~2~-CH(NH~2~)-COOH.\n"
            "**Đáp số: 2.**"
        ),
        (
            3,
            "(SBT – CTST). Cho giá trị nhiệt độ nóng chảy (°C) của glycine, trilaurin, H~2~O, aniline, tristearin không theo thứ tự như sau: 0; -6,0; 46,5; 71,6; 262,0. Giá trị nào phù hợp với nhiệt độ nóng chảy của glycine?",
            "262",
            "Glycine tồn tại ở dạng ion lưỡng cực ^+^H~3~N-CH~2~-COO^-^ có mạng tinh thể ion với lực hút tĩnh điện rất mạnh nên có nhiệt độ nóng chảy cao nhất trong dãy, đạt 262,0 °C.\n"
            "(Nhiệt độ nóng chảy của các chất còn lại lần lượt là: aniline -6,0 °C; H~2~O 0 °C; trilaurin 46,5 °C; tristearin 71,6 °C).\n"
            "**Đáp số: 262.**"
        ),
        (
            4,
            "Cho một số hiện tượng thí nghiệm:\n"
            "(1) Đun nóng dung dịch lòng trắng trứng thấy hiện tượng đông tụ lại, tách ra khỏi dung dịch.\n"
            "(2) Trộn lẫn lòng trắng trứng, dung dịch NaOH và một ít CuSO~4~ thấy xuất hiện màu đỏ gạch đặc trưng.\n"
            "(3) Nhỏ vài giọt dung dịch HNO~3~ đặc vào dung dịch lòng trắng trứng thấy kết tủa màu vàng.\n"
            "(4) Đốt cháy một mẫu lòng trắng trứng thấy xuất hiện mùi khét như mùi tóc cháy.\n"
            "(5) Vắt chanh vào một cốc sữa, ta thấy hiện tượng đông tụ protein xuất hiện.\n"
            "Có bao nhiêu hiện tượng thí nghiệm mô tả đúng?",
            "4",
            "(1) Đúng, protein bị đông tụ bởi nhiệt độ.\n"
            "(2) Sai, phản ứng màu biuret tạo dung dịch phức chất màu tím đặc trưng, không phải màu đỏ gạch.\n"
            "(3) Đúng, phản ứng xanthoproteic giữa protein lòng trắng trứng và HNO~3~ đặc tạo kết tủa màu vàng.\n"
            "(4) Đúng, đốt cháy protein giải phóng các hợp chất chứa sulfur/nitrogen có mùi khét như tóc cháy.\n"
            "(5) Đúng, acid citric trong chanh làm giảm pH đến điểm đẳng điện làm đông tụ protein casein trong sữa.\n"
            "=> Có 4 hiện tượng mô tả đúng: (1), (3), (4), (5).\n"
            "**Đáp số: 4.**"
        ),
        (
            5,
            "Cho một loại protein chứa 3,2% sulfur về khối lượng. Giả sử trong phân tử chỉ chứa 2 nguyên tử sulfur, phân tử khối của loại protein đó là bao nhiêu?",
            "2000",
            "Phân tử chỉ chứa 2 nguyên tử S, do đó tổng khối lượng của sulfur là:\n"
            "    m(S) = 2 * 32 = 64 amu.\n"
            "Vì sulfur chiếm 3,2% khối lượng phân tử protein:\n"
            "    M = (64 * 100) / 3,2 = 2000 amu.\n"
            "**Đáp số: 2000.**"
        ),
        (
            6,
            "Thủy phân hoàn toàn 200 gam hỗn hợp gồm tơ tằm và lông cừu thu được 31,7 g glycine. Phần trăm khối lượng của glycine trong tơ tằm và lông cừu tương ứng là 43,6% và 6,6%. Thành phần phần trăm khối lượng tương ứng của tơ tằm trong hỗn hợp ban đầu là a%. Xác định giá trị của a.",
            "25",
            "Gọi khối lượng của tơ tằm và lông cừu trong hỗn hợp lần lượt là x và y (gam).\n"
            "Theo đề bài, ta có hệ phương trình:\n"
            "    x + y = 200\n"
            "    0,436x + 0,066y = 31,7\n"
            "Giải hệ phương trình ta được: x = 50 g và y = 150 g.\n"
            "Thành phần phần trăm khối lượng của tơ tằm trong hỗn hợp ban đầu là:\n"
            "    a = (50 / 200) * 100% = 25%.\n"
            "**Đáp số: 25.**"
        )
    ]
    
    for q in part3_data:
        num, text, ans, sol = q
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        
        r_num = p.add_run(f"Câu {num}. ")
        set_run_style(r_num, font_size=11, bold=True)
        add_chem_run(p, text, base_size=11)
        
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(1)
        p_ans.paragraph_format.space_after = Pt(3)
        p_ans.paragraph_format.left_indent = Cm(0.6)
        
        r_ans_lbl = p_ans.add_run("Đáp số: ")
        set_run_style(r_ans_lbl, font_size=11, italic=True)
        
        if is_teacher:
            r_ans_val = p_ans.add_run(ans)
            set_run_style(r_ans_val, font_size=11, bold=True, color=COLOR_CORRECT)
        else:
            r_dots = p_ans.add_run(". . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . .")
            set_run_style(r_dots, font_size=11, color=COLOR_GRAY)
            
        if is_teacher and sol:
            add_solution(doc, sol)

    # Summary table for teacher version
    if is_teacher:
        p_tbl_hdr = doc.add_paragraph()
        p_tbl_hdr.paragraph_format.space_before = Pt(12)
        p_tbl_hdr.paragraph_format.space_after = Pt(4)
        r_thdr = p_tbl_hdr.add_run("BẢNG TỔNG HỢP ĐÁP ÁN ĐỀ SỐ 09")
        set_run_style(r_thdr, font_size=11.5, bold=True, color=COLOR_PRIMARY)
        
        p_p1 = doc.add_paragraph()
        p_p1.paragraph_format.space_before = Pt(2)
        p_p1.paragraph_format.space_after = Pt(2)
        r_p1 = p_p1.add_run("Phần I: ")
        set_run_style(r_p1, font_size=10.5, bold=True)
        p1_ans_str = "1.A  2.C  3.A  4.C  5.A  6.D  7.C  8.D  9.B  10.B  11.B  12.A  13.A  14.C  15.C  16.D  17.C  18.C"
        r_p1_val = p_p1.add_run(p1_ans_str)
        set_run_style(r_p1_val, font_size=10.5, bold=True, color=COLOR_CORRECT)
        
        p_p2 = doc.add_paragraph()
        p_p2.paragraph_format.space_before = Pt(2)
        p_p2.paragraph_format.space_after = Pt(2)
        r_p2 = p_p2.add_run("Phần II: ")
        set_run_style(r_p2, font_size=10.5, bold=True)
        p2_ans_str = "Câu 1: a-Đ, b-Đ, c-Đ, d-S  |  Câu 2: a-Đ, b-S, c-Đ, d-Đ  |  Câu 3: a-Đ, b-S, c-Đ, d-S  |  Câu 4: a-Đ, b-Đ, c-Đ, d-S"
        r_p2_val = p_p2.add_run(p2_ans_str)
        set_run_style(r_p2_val, font_size=10.5, bold=True, color=COLOR_CORRECT)
        
        p_p3 = doc.add_paragraph()
        p_p3.paragraph_format.space_before = Pt(2)
        p_p3.paragraph_format.space_after = Pt(4)
        r_p3 = p_p3.add_run("Phần III: ")
        set_run_style(r_p3, font_size=10.5, bold=True)
        p3_ans_str = "Câu 1: 2  |  Câu 2: 2  |  Câu 3: 262  |  Câu 4: 4  |  Câu 5: 2000  |  Câu 6: 25"
        r_p3_val = p_p3.add_run(p3_ans_str)
        set_run_style(r_p3_val, font_size=10.5, bold=True, color=COLOR_CORRECT)

    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(10)
    p_end.paragraph_format.space_after = Pt(6)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("---------- HẾT ----------")
    set_run_style(r_end, font_size=11, bold=True)

    out_name = "De_Kiem_Tra_Hoa_12_De09_GiaoVien.docx" if is_teacher else "De_Kiem_Tra_Hoa_12_De09.docx"
    out_path = os.path.join(OUT_DIR, out_name)
    doc.save(out_path)
    print(f"Saved: {out_path}")

if __name__ == "__main__":
    build_docx_de09(is_teacher=False)
    build_docx_de09(is_teacher=True)
