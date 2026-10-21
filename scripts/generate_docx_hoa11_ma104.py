# -*- coding: utf-8 -*-
"""
Script tạo file Word (.docx) chuẩn "Form Đề Kiểm Tra" cho:
ĐỀ KIỂM TRA GIỮA HỌC KỲ I - MÔN HOÁ HỌC 11 - TRƯỜNG THPT NGUYỄN CHÍ THANH - MÃ ĐỀ 104
Áp dụng 100% quy chuẩn:
- Không dùng bảng trong nội dung
- Header & Footer chuẩn
- Tab stops chuẩn
- Subscript/superscript triệt để
- Xuất bản 2 file: Bản Học sinh và Bản Giáo viên
"""

import os
import sys
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Thêm đường dẫn scripts vào sys.path để import form_de_kiem_tra
sys.path.insert(0, os.path.dirname(__file__))
from form_de_kiem_tra import (
    setup_header_footer,
    build_title_block,
    add_section_header,
    render_question_prompt,
    render_options_no_table,
    render_statement,
    add_solution_no_table,
    add_student_writing_lines,
    render_rich_text,
    COLOR_CORRECT,
    COLOR_PRIMARY
)

def create_exam_doc_104(is_teacher=False):
    doc = docx.Document()
    setup_header_footer(doc)

    build_title_block(
        doc,
        school_line="SỞ GIÁO DỤC VÀ ĐÀO TẠO TP. HUẾ\nTRƯỜNG THPT NGUYỄN CHÍ THANH",
        exam_title="ĐỀ KIỂM TRA GIỮA HỌC KỲ I - NĂM HỌC 2025 - 2026",
        subject_line="MÔN: HOÁ HỌC. Lớp 11",
        num_pages="03",
        duration="45 phút",
        exam_code="104",
        atomic_masses="H = 1; C = 12; N = 14; O = 16; Na = 23; S = 32; Cl = 35,5; Ba = 137",
        is_teacher=is_teacher
    )

    # ==================== PHẦN I ====================
    add_section_header(
        doc,
        title="PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn: (16 câu - 4 điểm)",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 16. Mỗi câu thí sinh chỉ chọn một phương án."
    )

    p1_questions = [
        {
            "q": "Câu 1. Các dung dịch sau đây có cùng nồng độ 1M, dd nào dẫn điện tốt nhất?",
            "opts": [
                "A. Ba(OH)~2~",
                "*B. Al~2~(SO~4~)~3~*",
                "C. H~2~SO~4~",
                "D. NH~4~NO~3~"
            ],
            "sol": "► *Lời giải chi tiết:*\nKhả năng dẫn điện phụ thuộc vào tổng nồng độ các ion trong dung dịch. Cùng nồng độ 1M:\n- Ba(OH)~2~ → Ba^2+^ + 2OH^-^ (tổng nồng độ ion = 3M).\n- Al~2~(SO~4~)~3~ → 2Al^3+^ + 3SO~4~^2-^ (tổng nồng độ ion = 5M).\n- H~2~SO~4~ → 2H^+^ + SO~4~^2-^ (tổng nồng độ ion = 3M).\n- NH~4~NO~3~ → NH~4~^+^ + NO~3~^-^ (tổng nồng độ ion = 2M).\nDung dịch Al~2~(SO~4~)~3~ có tổng nồng độ ion lớn nhất nên dẫn điện tốt nhất.\n*Chọn B.*"
        },
        {
            "q": "Câu 2. Khí thải công nghiệp, khí thải động cơ (ô tô, xe máy ...) là nguyên nhân chủ yếu gây mưa acid. Dãy nào dưới đây gồm các chất có trong các loại khí thải này và là nguyên nhân trực tiếp gây mưa acid?",
            "opts": [
                "A. SO~2~, CO, NO.",
                "B. NO~2~, CO~2~, CO.",
                "*C. NO, NO~2~, SO~2~.*",
                "D. SO~2~, CO, NO~2~."
            ],
            "sol": "► *Lời giải chi tiết:*\nNguyên nhân trực tiếp gây mưa acid là các oxide của nitrogen (NO, NO~2~) và SO~2~ vì chúng phản ứng với oxygen và nước trong không khí tạo thành HNO~3~ và H~2~SO~4~. Khí CO không gây mưa acid.\n*Chọn C.*"
        },
        {
            "q": "Câu 3. Cho các phản ứng sau:\n(1) N~2~ + O~2~ → 2NO (t^o^, xt);\n(2) N~2~ + 3H~2~ ⇌ 2NH~3~ (t^o^, xt).\nTrong hai phản ứng trên thì nitrogen",
            "opts": [
                "*A. thể hiện tính khử và tính oxi hóa.*",
                "B. không thể hiện tính khử và tính oxi hóa.",
                "C. chỉ thể hiện tính oxi hóa.",
                "D. chỉ thể hiện tính khử."
            ],
            "sol": "► *Lời giải chi tiết:*\n- Trong (1): N~2~(0) → 2NO(+2), số oxi hóa tăng nên N~2~ là chất khử.\n- Trong (2): N~2~(0) → 2NH~3~(-3), số oxi hóa giảm nên N~2~ là chất oxi hóa.\nVậy nitrogen thể hiện tính khử và tính oxi hóa.\n*Chọn A.*"
        },
        {
            "q": "Câu 4. Nhận xét nào sau đây *không đúng*?",
            "opts": [
                "*A. Phản ứng một chiều là phản ứng luôn xảy ra không hoàn toàn.*",
                "B. Trong phản ứng thuận nghịch, các chất sản phẩm có thể phản ứng với nhau để tạo thành chất đầu.",
                "C. Trong phản ứng một chiều, các chất sản phẩm không phản ứng được với nhau tạo thành chất đầu.",
                "D. Phản ứng thuận nghịch là phản ứng xảy ra theo hai chiều trái ngược nhau trong cùng điều kiện."
            ],
            "sol": "► *Lời giải chi tiết:*\nPhản ứng một chiều là phản ứng xảy ra hoàn toàn (khi phản ứng kết thúc ít nhất một chất phản ứng đã hết), do đó phát biểu A 'luôn xảy ra không hoàn toàn' là sai.\n*Chọn A.*"
        },
        {
            "q": "Câu 5. Trong y học, vì sao nitrogen lỏng được dùng để bảo quản mẫu vật?",
            "opts": [
                "*A. Do nitrogen hóa lỏng ở nhiệt độ thấp.*",
                "B. Do nitrogen không duy trì sự hô hấp.",
                "C. Do nitrogen tan ít rất trong nước.",
                "D. Do nitrogen là chất khí không màu, không mùi, không vị."
            ],
            "sol": "► *Lời giải chi tiết:*\nNitrogen hóa lỏng ở nhiệt độ rất thấp (-196 °C), tạo môi trường siêu lạnh làm ngừng mọi quá trình sinh hóa và vi sinh vật, giúp bảo quản lâu dài tế bào, mô, mẫu sinh học.\n*Chọn A.*"
        },
        {
            "q": "Câu 6. Chất nào sau đây là chất điện li?",
            "opts": [
                "A. Glucose.",
                "B. Saccharose.",
                "*C. Sodium hydroxide.*",
                "D. Alcohol ethylic."
            ],
            "sol": "► *Lời giải chi tiết:*\nSodium hydroxide (NaOH) là base mạnh phân li hoàn toàn thành ion Na^+^ và OH^-^ khi tan trong nước nên là chất điện li. Glucose, saccharose, alcohol ethylic là các chất không điện li.\n*Chọn C.*"
        },
        {
            "q": "Câu 7. Cho m gam Na vào nước dư thu được 1,5 L dung dịch có pH=12. Giá trị của m là",
            "opts": [
                "A. 0,23 gam.",
                "B. 0,46 gam.",
                "C. 0,115 gam.",
                "*D. 0,345 gam.*"
            ],
            "sol": "► *Lời giải chi tiết:*\npH = 12 ⇒ pOH = 2 ⇒ [OH^-^] = 0,01 M.\nSố mol OH^-^: n~OH-~ = 0,01 × 1,5 = 0,015 mol.\nPhản ứng: Na + H~2~O → NaOH + 1/2 H~2~↑\n⇒ n~Na~ = n~OH-~ = 0,015 mol ⇒ m = 0,015 × 23 = 0,345 gam.\n*Chọn D.*"
        },
        {
            "q": "Câu 8. Chuẩn độ dung dịch NaOH chưa biết chính xác nồng độ (biết nồng độ trong khoảng gần với 0,1 M) bằng dung dịch HCl 0,1 M với chỉ thị phenolphtalein. Chất đóng vai trò là dung dịch chuẩn là?",
            "opts": [
                "A. Phenolphtalein.",
                "B. NaCl.",
                "C. NaOH.",
                "*D. HCl.*"
            ],
            "sol": "► *Lời giải chi tiết:*\nDung dịch chuẩn là dung dịch đã biết chính xác nồng độ dùng để chuẩn độ. Ở đây HCl 0,1 M là dung dịch chuẩn.\n*Chọn D.*"
        },
        {
            "q": "Câu 9. Để hạn chế hiện tượng phú dưỡng, ta có thể dùng bao nhiêu biện pháp sau đây?\n(1) Tạo điều kiện để nước trong kênh, rạch, ao, hồ được lưu thông.\n(2) Xử lí nước thải trước khi chảy vào kênh rạch ao hồ.\n(3) Sử dụng phân bón đúng lượng, đúng cách & đúng thời điểm.",
            "opts": [
                "A. 1.",
                "B. 2.",
                "C. 0.",
                "*D. 3.*"
            ],
            "sol": "► *Lời giải chi tiết:*\nCả 3 biện pháp đều có tác dụng hạn chế hiện tượng phú dưỡng nguồn nước.\n*Chọn D.*"
        },
        {
            "q": "Câu 10. Trong phản ứng tổng hợp NH~3~ theo quá trình Haber hỗn hợp N~2~ & H~2~ (tỉ lệ mol x : y) được nén ở áp suất cao và đưa vào tháp tổng hợp ammonia trong điều kiện nhiệt độ, áp suất và chất xúc tác thích hợp. Tỉ lệ mol x : y là ?",
            "opts": [
                "A. 1 : 2.",
                "B. 2 : 1.",
                "*C. 1 : 3.*",
                "D. 3 : 1."
            ],
            "sol": "► *Lời giải chi tiết:*\nPhương trình: N~2~ + 3H~2~ ⇌ 2NH~3~. Tỉ lệ mol chuẩn theo hệ số phản ứng được nạp vào tháp là n~N2~ : n~H2~ = 1 : 3.\n*Chọn C.*"
        },
        {
            "q": "Câu 11. Xét cân bằng:\n(1) H~2~(g) + I~2~(g) ⇌ 2HI(g)   (K~1~)\n(2) 2HI(g) ⇌ H~2~(g) + I~2~(g)   (K~2~)\nMối quan hệ giữa K~1~ và K~2~ là:",
            "opts": [
                "A. K~1~ = 2K~2~.",
                "B. K~1~ = K~2~.",
                "C. K~1~ = √K~2~.",
                "*D. K~1~ = K~2~^-1^.*"
            ],
            "sol": "► *Lời giải chi tiết:*\nK~1~ = [HI]^2^ / ([H~2~][I~2~]) và K~2~ = ([H~2~][I~2~]) / [HI]^2^ ⇒ K~1~ = 1 / K~2~ = K~2~^-1^.\n*Chọn D.*"
        },
        {
            "q": "Câu 12. Công thức Lewis của NH~3~ là:",
            "opts": [
                "A. Công thức A",
                "B. Công thức B",
                "*C. Công thức C*",
                "D. Công thức D"
            ],
            "image": "San_Pham/images_crop/de104_cau12_lewis.png",
            "sol": "► *Lời giải chi tiết:*\nTrong phân tử NH~3~, nguyên tử N có 5 electron hóa trị: 3 electron tham gia tạo 3 liên kết đơn N - H với 3 nguyên tử H, còn lại 1 cặp electron tự do trên nguyên tử N. Do đó công thức C là đúng.\n*Chọn C.*"
        },
        {
            "q": "Câu 13. Kim loại không bị hòa tan trong dung dịch acid HNO~3~ đặc nguội, nhưng tan được trong dung dịch NaOH là:",
            "opts": [
                "A. Mg",
                "*B. Al*",
                "C. Fe",
                "D. Pb"
            ],
            "sol": "► *Lời giải chi tiết:*\nNhôm (Al) bị thụ động hóa trong HNO~3~ đặc nguội nên không tan, nhưng tan được trong dung dịch kiềm mạnh NaOH: 2Al + 2NaOH + 6H~2~O → 2Na[Al(OH)~4~] + 3H~2~↑.\n*Chọn B.*"
        },
        {
            "q": "Câu 14. Cân bằng hoá học là một trạng thái của phản ứng thuận nghịch tại đó:",
            "opts": [
                "A. tốc độ phản ứng thuận bằng một nửa tốc độ phản ứng nghịch.",
                "B. tốc độ phản ứng không thay đổi.",
                "C. tốc độ phản ứng thuận bằng hai lần tốc độ phản ứng nghịch.",
                "*D. tốc độ phản ứng thuận bằng tốc độ phản ứng nghịch.*"
            ],
            "sol": "► *Lời giải chi tiết:*\nCân bằng hóa học là trạng thái tại đó tốc độ phản ứng thuận bằng tốc độ phản ứng nghịch (v~thuận~ = v~nghịch~).\n*Chọn D.*"
        },
        {
            "q": "Câu 15. Tổng hệ số (các số nguyên, tối giản) của tất cả các chất trong phương trình phản ứng giữa Cu với dung dịch HNO~3~ đặc là",
            "opts": [
                "A. 11.",
                "B. 9.",
                "*C. 10.*",
                "D. 8."
            ],
            "sol": "► *Lời giải chi tiết:*\nPhương trình: Cu + 4HNO~3~ → Cu(NO~3~)~2~ + 2NO~2~↑ + 2H~2~O.\nTổng hệ số: 1 + 4 + 1 + 2 + 2 = 10.\n*Chọn C.*"
        },
        {
            "q": "Câu 16. Có thể phân biệt muối ammonium với các muối khác bằng cách cho nó tác dụng với dung dịch kiềm đun nóng, vì",
            "opts": [
                "A. thoát ra một chất khí màu nâu đỏ, làm xanh giấy quỳ tím ẩm.",
                "B. thoát ra một chất khí màu lục nhạt.",
                "C. thoát ra chất khí không màu, không mùi.",
                "*D. thoát ra một chất khí không màu, mùi khai, làm xanh giấy quỳ tím ẩm.*"
            ],
            "sol": "► *Lời giải chi tiết:*\nPhản ứng: NH~4~^+^ + OH^-^ → NH~3~↑ + H~2~O (t^o^). Khí NH~3~ không màu, mùi khai đặc trưng, làm xanh quỳ tím ẩm.\n*Chọn D.*"
        }
    ]

    for item in p1_questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, item["q"])

        if "image" in item and os.path.exists(item["image"]):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(2)
            p_img.paragraph_format.space_after = Pt(2)
            p_img.add_run().add_picture(item["image"], width=Cm(12.0))

        render_options_no_table(doc, item["opts"], is_teacher=is_teacher)

        if is_teacher:
            add_solution_no_table(doc, item["sol"])

    # ==================== PHẦN II ====================
    add_section_header(
        doc,
        title="PHẦN II. Câu trắc nghiệm đúng sai: (2 câu - 2 điểm)",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 2. Trong mỗi ý a), b), c), d) ở mỗi câu thí sinh chọn đúng hoặc sai."
    )

    p2_questions = [
        {
            "lead": "Câu 1. Hiện tượng phú dưỡng là một biểu hiện của môi trường ao, hồ bị ô nhiễm do dư thừa các chất dinh dưỡng.",
            "stmts": [
                ("a) Chất thải công nghiệp, chất thải sinh hoạt, ... khi không được xử lí theo quy chuẩn, nếu thải vào sông, hồ cũng gây ra hiện tượng phú dưỡng.", "Đúng"),
                ("b) Thông thường, khi hàm lượng calcium trong nước đạt 300 µg/L và hàm lượng magnesium đạt 20 µg/L sẽ gây hiện tượng phú dưỡng.", "Sai"),
                ("c) Có thể quan sát được hiện tượng phú dưỡng thông qua sự xuất hiện dày đặc của tảo xanh trong nước.", "Đúng"),
                ("d) Hiện tượng phú dưỡng gây cản trở hấp thụ ánh sáng mặt trời vào nước, làm giảm sự quang hợp của thực vật thủy sinh.", "Đúng")
            ],
            "sol": "► *Lời giải chi tiết:*\na) Đúng vì nước thải chứa nhiều hợp chất N và P chưa qua xử lí là nguyên nhân chính gây phú dưỡng.\nb) Sai vì phú dưỡng xảy ra do dư thừa nitrogen (≥ 300 µg/L) và phosphorus (≥ 20 µg/L), không phải do calcium và magnesium.\nc) Đúng vì dư thừa dinh dưỡng kích thích tảo phát triển bùng nổ (nở hoa nước) tạo màu xanh dày đặc.\nd) Đúng vì lớp váng tảo che khuất ánh sáng mặt trời chiếu xuống các tầng nước sâu, làm suy giảm quang hợp."
        },
        {
            "lead": "Câu 2. Cho đồ thị biểu diễn tốc độ phản ứng thuận và nghịch theo thời gian:",
            "image": "San_Pham/images_crop/de104_do_thi_can_bang.png",
            "stmts": [
                ("a) Khi phản ứng đạt trạng thái cân bằng thì phản ứng thuận và phản ứng nghịch dừng lại.", "Sai"),
                ("b) Tại trạng thái cân bằng, tốc độ phản ứng thuận bằng tốc độ phản ứng nghịch.", "Đúng"),
                ("c) Đồ thị trên biểu diễn tốc độ của phản ứng một chiều phụ thuộc vào thời gian.", "Sai"),
                ("d) Khi phản ứng thuận nghịch xảy ra thì tốc độ phản ứng thuận tăng dần còn tốc độ phản ứng nghịch giảm dần.", "Sai")
            ],
            "sol": "► *Lời giải chi tiết:*\na) Sai vì cân bằng hóa học là cân bằng động, phản ứng thuận và nghịch vẫn tiếp tục xảy ra với tốc độ bằng nhau.\nb) Đúng vì tại cân bằng v~thuận~ = v~nghịch~ (thể hiện bằng đường ngang trên đồ thị).\nc) Sai vì đây là đồ thị của phản ứng thuận nghịch.\nd) Sai vì ban đầu tốc độ thuận lớn nhất sau đó giảm dần, tốc độ nghịch bằng 0 sau đó tăng dần."
        }
    ]

    for item in p2_questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, item["lead"])

        if "image" in item and os.path.exists(item["image"]):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(2)
            p_img.paragraph_format.space_after = Pt(2)
            p_img.add_run().add_picture(item["image"], width=Cm(8.0))

        for stmt, ans in item["stmts"]:
            p_s = doc.add_paragraph()
            p_s.paragraph_format.left_indent = Cm(0.5)
            p_s.paragraph_format.space_before = Pt(1)
            p_s.paragraph_format.space_after = Pt(1)
            render_statement(p_s, stmt, ans=ans, is_teacher=is_teacher)

        if is_teacher:
            add_solution_no_table(doc, item["sol"])

    # ==================== PHẦN III ====================
    add_section_header(
        doc,
        title="PHẦN III. Câu trắc nghiệm trả lời ngắn: (4 câu - 1 điểm)",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 4."
    )

    p3_questions = [
        {
            "q": "Câu 1. Cho dung dịch HNO~3~ tác dụng với các chất sau: NH~3~, CaCO~3~, Ag, NaOH. Có bao nhiêu phản ứng trong đó HNO~3~ đóng vai trò acid theo Bronsted?",
            "ans": "3",
            "sol": "► *Lời giải chi tiết:*\nTheo Brønsted, HNO~3~ là acid khi nhường proton H^+^:\n- NH~3~ + HNO~3~ → NH~4~NO~3~ (nhường H^+^ cho NH~3~: là acid).\n- CaCO~3~ + 2HNO~3~ → Ca(NO~3~)~2~ + CO~2~↑ + H~2~O (nhường H^+^ cho CO~3~^2-^: là acid).\n- Ag + 2HNO~3~ → AgNO~3~ + NO~2~↑ + H~2~O (phản ứng oxi hóa - khử, HNO~3~ đóng vai trò chất oxi hóa và môi trường).\n- NaOH + HNO~3~ → NaNO~3~ + H~2~O (nhường H^+^ cho OH^-^: là acid).\nVậy có 3 phản ứng: NH~3~, CaCO~3~, NaOH.\n► *Đáp án: 3*"
        },
        {
            "q": "Câu 2. Nhiệt phân các muối ammonium: NH~4~Cl, (NH~4~)~2~CO~3~, NH~4~NO~2~, NH~4~HCO~3~, NH~4~NO~3~. Có bao nhiêu trường hợp thu được khí NH~3~?",
            "ans": "3",
            "sol": "► *Lời giải chi tiết:*\n- NH~4~Cl → NH~3~↑ + HCl↑ (có tạo NH~3~).\n- (NH~4~)~2~CO~3~ → 2NH~3~↑ + CO~2~↑ + H~2~O (có tạo NH~3~).\n- NH~4~NO~2~ → N~2~↑ + 2H~2~O (không tạo NH~3~).\n- NH~4~HCO~3~ → NH~3~↑ + CO~2~↑ + H~2~O (có tạo NH~3~).\n- NH~4~NO~3~ → N~2~O↑ + 2H~2~O (không tạo NH~3~).\nCó 3 trường hợp thu được NH~3~: NH~4~Cl, (NH~4~)~2~CO~3~, NH~4~HCO~3~.\n► *Đáp án: 3*"
        },
        {
            "q": "Câu 3. Cho các chất (phân tử và ion): NaOH, H~2~S, NH~3~, CO~3~^2-^, S^2-^, NH~4~^+^, PO~4~^3-^. Theo quan điểm của Bronsted – Lowry có bao nhiêu chất có thể là base?",
            "ans": "5",
            "sol": "► *Lời giải chi tiết:*\nTheo Brønsted – Lowry, base là chất nhận proton H^+^:\n- NaOH (chứa OH^-^ nhận H^+^: base);\n- NH~3~ (nhận H^+^ tạo NH~4~^+^: base);\n- CO~3~^2-^ (nhận H^+^: base);\n- S^2-^ (nhận H^+^: base);\n- PO~4~^3-^ (nhận H^+^: base).\nH~2~S và NH~4~^+^ là các chất cho proton (acid).\nVậy có 5 chất/ion là base: NaOH, NH~3~, CO~3~^2-^, S^2-^, PO~4~^3-^.\n► *Đáp án: 5*"
        },
        {
            "q": "Câu 4. Cho cân bằng hóa học (trong bình kín) sau: N~2~(g) + 3H~2~(g) ⇌ 2NH~3~(g)   Δ~r~H~298~^o^ = -92 kJ. Trong các yếu tố: (1) Thêm một lượng N~2~ hoặc H~2~; (2) Thêm một lượng NH~3~; (3) Tăng nhiệt độ của phản ứng; (4) Tăng áp suất của phản ứng; (5) Dùng thêm chất xúc tác. Có bao nhiêu yếu tố làm cho tỉ khối của hỗn hợp khí trong bình so với H~2~ tăng lên?",
            "ans": "2",
            "sol": "► *Lời giải chi tiết:*\n- Khối lượng mol trung bình M~hh~ = m / n. Trong bình kín khối lượng m không đổi (khi không thêm bớt chất). Khi cân bằng chuyển dịch theo chiều thuận (chiều giảm số mol khí), số mol n giảm nên M~hh~ tăng, dẫn đến tỉ khối so với H~2~ tăng.\n- Yếu tố (4) Tăng áp suất làm cân bằng chuyển dịch chiều thuận làm giảm số mol khí ⇒ tỉ khối tăng.\n- Yếu tố (2) Thêm NH~3~ (M = 17) có khối lượng mol lớn hơn hỗn hợp nạp ban đầu (M~tb~ ≈ 8,5), dù phân hủy một phần thì tỉ khối của hệ vẫn tăng lên.\n- Yếu tố (3) tăng nhiệt độ làm chuyển dịch chiều nghịch (giảm tỉ khối); (5) xúc tác không làm chuyển dịch cân bằng.\nVậy có 2 yếu tố làm tăng tỉ khối: (2) và (4).\n► *Đáp án: 2*"
        }
    ]

    for item in p3_questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, item["q"])

        if is_teacher:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Cm(0.5)
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(1)
            render_rich_text(p_ans, f"► *Đáp án:* *{item['ans']}*", base_size=10.5, base_bold=True, base_color=COLOR_CORRECT)
            add_solution_no_table(doc, item["sol"])
        else:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Cm(0.5)
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(3)
            render_rich_text(p_ans, "*Trả lời:* ............................................................", base_size=10.5)

    # ==================== PHẦN IV ====================
    add_section_header(
        doc,
        title="PHẦN IV. TỰ LUẬN: (3 câu - 3 điểm)",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 3."
    )

    essay_questions = [
        {
            "num": "Câu 1: (1,0 điểm)",
            "prompt": "Viết biểu thức tính hằng số cân bằng (Kc) cho các phản ứng thuận nghịch sau:\n(1) H~2~(g) + F~2~(g) ⇌ 2HF(g)\n(2) Fe~2~O~3~(s) + 3CO(g) ⇌ 2Fe(s) + 3CO~2~(g)",
            "sol": "► *Lời giải chi tiết:*\n(1) Đối với phản ứng: H~2~(g) + F~2~(g) ⇌ 2HF(g)\nKc = [HF]^2^ / ([H~2~][F~2~])\n\n(2) Đối với phản ứng: Fe~2~O~3~(s) + 3CO(g) ⇌ 2Fe(s) + 3CO~2~(g)\nDo Fe~2~O~3~ và Fe là chất rắn nguyên chất nên nồng độ của chúng không xuất hiện trong biểu thức hằng số cân bằng Kc:\nKc = [CO~2~]^3^ / [CO]^3^",
            "lines": 6
        },
        {
            "num": "Câu 2: (1,0 điểm)",
            "prompt": "Trộn lẫn V mL dung dịch NaOH 0,03 M với V mL dung dịch HCl 0,01 M được 2V mL dung dịch Y. Dung dịch Y có pH là bao nhiêu?",
            "sol": "► *Lời giải chi tiết:*\n- Đổi thể tích: V mL = V.10^-3^ L.\n- Số mol OH^-^ = 0,03 × V.10^-3^ mol.\n- Số mol H^+^ = 0,01 × V.10^-3^ mol.\n- Phương trình ion rút gọn: H^+^ + OH^-^ → H~2~O\nVì n~OH-~ > n~H+~ nên OH^-^ dư sau phản ứng:\nn~OH- dư~ = (0,03 - 0,01) × V.10^-3^ = 0,02.10^-3^ V mol.\n- Tổng thể tích sau trộn: V~Y~ = 2V mL = 2V.10^-3^ L.\n- Nồng độ OH^-^ dư:\n[OH^-^] = (0,02.10^-3^ V) / (2V.10^-3^) = 0,01 M = 10^-2^ M.\n- Giá trị pOH và pH của dung dịch Y:\npOH = -log[OH^-^] = 2 ⇒ pH = 14 - pOH = 14 - 2 = 12.\n*Đáp số:* pH = 12.",
            "lines": 8
        },
        {
            "num": "Câu 3: (1,0 điểm)",
            "prompt": "Nước cường toan trong tiếng Anh là Aqua Regia (nghĩa là nước hoàng gia) – là hợp chất có tính ăn mòn mạnh, ở dạng lỏng, có màu vàng và dễ bay hơi, được sử dụng cho 1 số quy trình hóa học phân tích và để tinh chế vàng. Hãy nêu thành phần của nước cường toan và viết phản ứng của nước cường toan với vàng.",
            "sol": "► *Lời giải chi tiết:*\n- Thành phần của nước cường toan (Aqua Regia):\nLà hỗn hợp gồm dung dịch HNO~3~ đặc và HCl đặc trộn với nhau theo tỉ lệ thể tích:\nV~HNO3 đặc~ : V~HCl đặc~ = 1 : 3.\n\n- Khả năng hòa tan vàng và phương trình hóa học:\nVàng (Au) không tan trong HNO~3~ đặc hay HCl đặc riêng rẽ, nhưng tan được trong nước cường toan do tính oxi hóa mạnh của HNO~3~ kết hợp với khả năng tạo phức bền của ion Cl^-^:\nAu + HNO~3~ + 4HCl → HAuCl~4~ + NO↑ + 2H~2~O\n(hoặc viết dạng ion: Au + NO~3~^-^ + 4Cl^-^ + 4H^+^ → [AuCl~4~]^-^ + NO↑ + 2H~2~O).\nSản phẩm tạo thành là acid tetrachloroauric(III) HAuCl~4~ tan trong dung dịch và giải phóng khí NO.",
            "lines": 8
        }
    ]

    for item in essay_questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*{item['num']}*", base_size=10.5, base_bold=True)

        for line in item["prompt"].split("\n"):
            p_line = doc.add_paragraph()
            p_line.paragraph_format.left_indent = Cm(0.5)
            p_line.paragraph_format.space_before = Pt(1)
            p_line.paragraph_format.space_after = Pt(1)
            render_rich_text(p_line, line, base_size=10.5)

        if is_teacher:
            add_solution_no_table(doc, item["sol"])
        else:
            add_student_writing_lines(doc, num_lines=item["lines"])

    # Dòng kết thúc
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(10)
    p_end.paragraph_format.space_after = Pt(4)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_end, "---------- *HẾT* ----------", base_size=10.5, base_bold=True)

    p_rule = doc.add_paragraph()
    p_rule.paragraph_format.space_before = Pt(0)
    p_rule.paragraph_format.space_after = Pt(0)
    p_rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_rule, "_(Học sinh không sử dụng tài liệu và bảng tuần hoàn hoá học)_", base_size=9.5, base_italic=True)

    return doc

def main():
    out_dir = r"c:\Antigravity_Thanh\San_Pham"
    os.makedirs(out_dir, exist_ok=True)

    # 1. Tạo Bản Học sinh
    doc_student = create_exam_doc_104(is_teacher=False)
    path_student = os.path.join(out_dir, "De_Kiem_Tra_Hoa_11_Ma104.docx")
    doc_student.save(path_student)
    print(f"Đã tạo Bản Học sinh: {path_student}")

    # 2. Tạo Bản Giáo viên
    doc_teacher = create_exam_doc_104(is_teacher=True)
    path_teacher = os.path.join(out_dir, "De_Kiem_Tra_Hoa_11_Ma104_GiaoVien.docx")
    doc_teacher.save(path_teacher)
    print(f"Đã tạo Bản Giáo viên: {path_teacher}")

if __name__ == "__main__":
    main()
