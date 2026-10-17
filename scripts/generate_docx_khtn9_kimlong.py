# -*- coding: utf-8 -*-
"""
Script tạo file Word (.docx) chuẩn "FORM ĐỀ KIỂM TRA" cho Đề Cương Ôn Tập Giữa Kì I - KHTN 9
Trường THCS Kim Long - Năm học 2026 - 2027
Thầy TRẦN VĂN THẠNH - 0777.470.803

Tạo 3 phiên bản:
1. Bản Học sinh: San_Pham/De_Cuong_KHTN_9_Giua_HK1_Kim_Long_HocSinh.docx
2. Bản Giáo viên: San_Pham/De_Cuong_KHTN_9_Giua_HK1_Kim_Long_GiaoVien.docx
3. Bản Dạy trên màn hình: San_Pham/De_Cuong_KHTN_9_Giua_HK1_Kim_Long_DayTrenManHinh.docx
"""

import os
import sys
import re
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.path.append(os.path.dirname(__file__))
sys.stdout.reconfigure(encoding='utf-8')

from khtn9_kimlong_data import (
    INFO, BAI_1_QUESTIONS, HOA_PART1_QUESTIONS, HOA_PART2_QUESTIONS,
    HOA_PART3_QUESTIONS, HOA_ESSAY_QUESTIONS, VATLI_QUESTIONS
)

ROOT_DIR = r"c:\Antigravity_Thanh"
OUT_DIR = os.path.join(ROOT_DIR, "San_Pham")
ONEDRIVE_DIRS = [
    r"C:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn9\KHTN9-DEKIEMTRADINHKI",
    r"C:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn9"
]
os.makedirs(OUT_DIR, exist_ok=True)

# Bảng màu chuẩn quy định trong FORM_DE_KIEM_TRA
COLOR_HF_RED = RGBColor(211, 47, 47)        # Coral Red (#D32F2F) cho Header & Footer
COLOR_HF_GREEN_HEX = "8BC390"               # Sage Green (#8BC390) cho viền Header & Footer
COLOR_PRIMARY = RGBColor(46, 125, 50)       # Xanh lá đậm (#2E7D32) cho Tiêu đề Phân môn / Phần
COLOR_CORRECT = RGBColor(198, 40, 40)       # Đỏ (#C62828) cho đáp án đúng Bản Giáo Viên
COLOR_SOLUTION = RGBColor(21, 101, 192)     # Xanh lam (#1565C0) cho tiêu đề Lời giải
COLOR_SCREEN_TITLE = RGBColor(106, 27, 154) # Tím đậm (#6A1B9A) cho tiêu đề bản Dạy Màn Hình
COLOR_GRAY = RGBColor(97, 97, 97)
COLOR_BORDER_GREY_HEX = "CCCCCC"

def add_p_border_bottom(p, color="8BC390", sz="6"):
    """Thêm đường kẻ viền dưới cho đoạn văn bản."""
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_p_border_top(p, color="8BC390", sz="6"):
    """Thêm đường kẻ viền trên cho đoạn văn bản."""
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    """Cài đặt lề đệm cho ô bảng."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex="F1F8E9"):
    """Đặt màu nền cho ô bảng."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="CCCCCC", sz="4"):
    """Đặt viền thanh mảnh cho bảng."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def set_run_style(run, font_name="Times New Roman", font_size=10.5, bold=False, italic=False, subscript=False, superscript=False, color=None):
    """Cài đặt style chi tiết cho từng run text."""
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    """Chèn trường số trang tự động PAGE vào chân trang."""
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def render_rich_text(p, text, base_font="Times New Roman", base_size=10.5, base_bold=False, base_italic=False, base_color=None):
    """
    Sinh văn bản định dạng phong phú từ ký hiệu:
    - *bold*
    - _italic_
    - ~subscript~ (CO~2~, H~2~O, NaCl)
    - ^superscript^ (1s^2^, 4s^1^, H^+^)
    """
    tokens = re.split(r'(~[^~]+~|\^[^\^]+\^|\*[^\*]+\*|_[^_]+_)', text)
    for token in tokens:
        if not token:
            continue

        is_bold = base_bold
        is_italic = base_italic
        is_sub = False
        is_sup = False
        content = token

        if token.startswith('*') and token.endswith('*'):
            is_bold = True
            content = token[1:-1]
        elif token.startswith('_') and token.endswith('_'):
            is_italic = True
            content = token[1:-1]
        elif token.startswith('~') and token.endswith('~'):
            is_sub = True
            content = token[1:-1]
        elif token.startswith('^') and token.endswith('^'):
            is_sup = True
            content = token[1:-1]

        run = p.add_run(content)
        run_color = COLOR_CORRECT if (is_bold and not base_bold and not base_color) else base_color
        set_run_style(run, font_name=base_font, font_size=base_size, bold=is_bold, italic=is_italic,
                      subscript=is_sub, superscript=is_sup, color=run_color)

def setup_header_footer(doc):
    """Cài đặt Header & Footer chuẩn theo FORM_DE_KIEM_TRA."""
    normal_style = doc.styles['Normal']
    normal_style.paragraph_format.tab_stops.clear_all()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    section.start_type = WD_SECTION_START.NEW_PAGE
    section.different_first_page_header_footer = False
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    # 1. Header (Bảng 1x2, rộng 18.0 cm)
    header = section.header
    for p in header.paragraphs:
        p.text = ""
    tbl_h = header.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl_h.cell(0, 0)
    c_right = tbl_h.cell(0, 1)
    c_left.width = Cm(8.0)
    c_right.width = Cm(10.0)

    p_hl = c_left.paragraphs[0]
    p_hl.paragraph_format.space_before = Pt(0)
    p_hl.paragraph_format.space_after = Pt(2)
    p_hl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_hl = p_hl.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)

    p_hr = c_right.paragraphs[0]
    p_hr.paragraph_format.space_before = Pt(0)
    p_hr.paragraph_format.space_after = Pt(2)
    p_hr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p_hr.add_run("Thầy ")
    set_run_style(r1, font_size=10, color=COLOR_HF_RED)
    r2 = p_hr.add_run("TRẦN VĂN THẠNH")
    set_run_style(r2, font_size=10, bold=True, color=COLOR_HF_RED)
    r3 = p_hr.add_run(" - 0777.470.803")
    set_run_style(r3, font_size=10, color=COLOR_HF_RED)

    add_p_border_bottom(p_hl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_bottom(p_hr, color=COLOR_HF_GREEN_HEX, sz="6")

    # 2. Footer (Bảng 1x2, rộng 18.0 cm)
    footer = section.footer
    for p in footer.paragraphs:
        p.text = ""
    tbl_f = footer.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_fl = tbl_f.cell(0, 0)
    c_fr = tbl_f.cell(0, 1)
    c_fl.width = Cm(14.5)
    c_fr.width = Cm(3.5)

    p_fl = c_fl.paragraphs[0]
    p_fl.paragraph_format.space_before = Pt(2)
    p_fl.paragraph_format.space_after = Pt(0)
    p_fl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_fl = p_fl.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)

    p_fr = c_fr.paragraphs[0]
    p_fr.paragraph_format.space_before = Pt(2)
    p_fr.paragraph_format.space_after = Pt(0)
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_field(p_fr, font_size=10, bold=True, color=COLOR_HF_RED)

    add_p_border_top(p_fl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_top(p_fr, color=COLOR_HF_GREEN_HEX, sz="6")

def build_title_block(doc, is_teacher=False, is_screen=False, base_size=10.5):
    """Xây dựng Header Block không viền, chuẩn mực."""
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_l, cell_r = tbl.cell(0, 0), tbl.cell(0, 1)
    cell_l.width = Cm(8.0)
    cell_r.width = Cm(10.0)

    # Cột trái
    p_tl1 = cell_l.paragraphs[0]
    p_tl1.paragraph_format.space_before = Pt(0)
    p_tl1.paragraph_format.space_after = Pt(1)
    p_tl1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tl1, f"*{INFO['school']}*", base_size=10.5, base_bold=True)

    p_tl2 = cell_l.add_paragraph()
    p_tl2.paragraph_format.space_before = Pt(0)
    p_tl2.paragraph_format.space_after = Pt(1)
    p_tl2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tl2, f"*{INFO['year']}*", base_size=10, base_bold=True)

    p_tl3 = cell_l.add_paragraph()
    p_tl3.paragraph_format.space_before = Pt(0)
    p_tl3.paragraph_format.space_after = Pt(0)
    p_tl3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tl3, "_ĐỀ CƯƠNG ÔN TẬP CHÍNH THỨC_", base_size=9.5, base_italic=True)

    # Cột phải
    p_tr1 = cell_r.paragraphs[0]
    p_tr1.paragraph_format.space_before = Pt(0)
    p_tr1.paragraph_format.space_after = Pt(1)
    p_tr1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tr1, f"*{INFO['title']}*", base_size=11, base_bold=True)

    p_tr2 = cell_r.add_paragraph()
    p_tr2.paragraph_format.space_before = Pt(0)
    p_tr2.paragraph_format.space_after = Pt(1)
    p_tr2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tr2, f"*{INFO['subject']}*", base_size=12, base_bold=True, base_color=COLOR_PRIMARY)

    p_tr3 = cell_r.add_paragraph()
    p_tr3.paragraph_format.space_before = Pt(0)
    p_tr3.paragraph_format.space_after = Pt(0)
    p_tr3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tr3, f"_{INFO['time_str']}_", base_size=9.5, base_italic=True)

    # Dòng tiêu đề phụ cho Giáo viên / Màn hình
    if is_teacher:
        p_gv = doc.add_paragraph()
        p_gv.paragraph_format.space_before = Pt(4)
        p_gv.paragraph_format.space_after = Pt(2)
        p_gv.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rgv = p_gv.add_run("(BẢN GIÁO VIÊN -- CÓ ĐÁP ÁN & LỜI GIẢI CHI TIẾT)")
        set_run_style(rgv, font_size=11.5, bold=True, color=COLOR_CORRECT)
    elif is_screen:
        p_sc = doc.add_paragraph()
        p_sc.paragraph_format.space_before = Pt(4)
        p_sc.paragraph_format.space_after = Pt(2)
        p_sc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rsc = p_sc.add_run("(BẢN TRÌNH CHIẾU / GIẢNG DẠY MÀN HÌNH -- CÓ KHUNG GHI CHÚ BÀI DẠY)")
        set_run_style(rsc, font_size=11.5, bold=True, color=COLOR_SCREEN_TITLE)

    # Dòng thông tin học sinh
    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(4)
    p_info.paragraph_format.space_after = Pt(3)
    render_rich_text(p_info, "*Họ và tên thí sinh:* .................................................... *Lớp:* ......... *Trường:* ....................................................", base_size=base_size)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(4)
    add_p_border_bottom(p_div, color=COLOR_BORDER_GREY_HEX, sz="4")

def add_chapter_title(doc, title, subtitle=None, base_size=12):
    """Tiêu đề lớn cho từng Phân môn / Bài học."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    render_rich_text(p, f"*{title}*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)
    if subtitle:
        render_rich_text(p, f" {subtitle}", base_size=base_size * 0.9, base_italic=True)

def add_section_header(doc, title, base_size=11):
    """Tiêu đề mục: I. NỘI DUNG ÔN TẬP, II. TRẮC NGHIỆM, v.v."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, f"*{title}*", base_size=base_size, base_bold=True, base_color=RGBColor(21, 101, 192))

def add_solution_block(doc, solution_text, base_size=10):
    """Trình bày Lời giải chi tiết theo chuẩn không dùng Table."""
    p_sol = doc.add_paragraph()
    p_sol.paragraph_format.left_indent = Cm(0.6)
    p_sol.paragraph_format.space_before = Pt(2)
    p_sol.paragraph_format.space_after = Pt(3)
    p_sol.paragraph_format.line_spacing = 1.15

    r_lbl = p_sol.add_run("Lời giải chi tiết:\n")
    set_run_style(r_lbl, font_size=base_size, bold=True, color=COLOR_SOLUTION)

    lines = solution_text.split('\n')
    for idx, line in enumerate(lines):
        line_str = line.strip()
        if not line_str:
            continue
        if idx > 0:
            p_sol = doc.add_paragraph()
            p_sol.paragraph_format.left_indent = Cm(0.6)
            p_sol.paragraph_format.space_before = Pt(1)
            p_sol.paragraph_format.space_after = Pt(2)
            p_sol.paragraph_format.line_spacing = 1.15
        render_rich_text(p_sol, line_str, base_size=base_size)

def add_screen_teaching_space(doc, num_lines=3):
    """Chừa khoảng trống rộng rãi cho bản Dạy Trên Màn Hình."""
    p_box = doc.add_paragraph()
    p_box.paragraph_format.left_indent = Cm(0.5)
    p_box.paragraph_format.space_before = Pt(3)
    p_box.paragraph_format.space_after = Pt(5)
    
    # Viền nét đứt màu xám nhẹ, viền trái nét liền tím
    pPr = p_box._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="dashed" w:sz="4" w:space="4" w:color="B0BEC5"/><w:bottom w:val="dashed" w:sz="4" w:space="4" w:color="B0BEC5"/><w:left w:val="single" w:sz="12" w:space="6" w:color="7E57C2"/></w:pBdr>'
    pPr.append(parse_xml(xml))
    
    r_hint = p_box.add_run("✎ Ghi chú & Lời giải trực tiếp trên màn hình:\n")
    set_run_style(r_hint, font_size=9.5, bold=True, color=COLOR_SCREEN_TITLE)
    for _ in range(num_lines):
        p_box.add_run("\n")

def add_writing_lines(doc, num_lines=4):
    """Thêm dòng chấm chấm cho học sinh làm bài tự luận."""
    for _ in range(num_lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("................................................................................................................................................................................................................")
        set_run_style(r, font_size=9, color=COLOR_GRAY)

def add_callout_box(doc, title, content_lines, base_size=10.5):
    """Khung ghi chú công thức hoặc kiến thức trọng tâm."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    c.width = Cm(18.0)
    set_cell_margins(c, top=80, bottom=80, left=120, right=120)
    set_cell_shading(c, "F5F9F6")
    
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p, f"*{title}*", base_size=base_size + 0.5, base_bold=True, base_color=COLOR_PRIMARY)
    
    for line in content_lines:
        p_line = c.add_paragraph()
        p_line.paragraph_format.space_before = Pt(1)
        p_line.paragraph_format.space_after = Pt(1)
        p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        render_rich_text(p_line, line, base_size=base_size, base_bold=True)

def render_mcq_options(doc, opts, correct_letter, is_teacher, base_size=10.5):
    """Bố trí 4 phương án A, B, C, D bằng Tab Stops - Tuyệt đối không dùng Table."""
    lens = [len(opt) for opt in opts]
    max_len = max(lens)

    if max_len <= 18:
        # 4 phương án trên 1 dòng
        p_opt = doc.add_paragraph()
        p_opt.paragraph_format.space_before = Pt(1)
        p_opt.paragraph_format.space_after = Pt(2)
        p_opt.paragraph_format.tab_stops.clear_all()
        p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(5.0), WD_TAB_ALIGNMENT.LEFT)
        p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(9.5), WD_TAB_ALIGNMENT.LEFT)
        p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(14.0), WD_TAB_ALIGNMENT.LEFT)

        for opt in opts:
            p_opt.add_run("\t")
            letter = opt.strip()[0]
            if is_teacher and letter == correct_letter:
                render_rich_text(p_opt, f"*{opt}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p_opt, opt, base_size=base_size)

    elif max_len <= 45:
        # 2 dòng x 2 phương án
        p1 = doc.add_paragraph()
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.tab_stops.clear_all()
        p1.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p1.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)

        p1.add_run("\t")
        if is_teacher and opts[0].strip()[0] == correct_letter:
            render_rich_text(p1, f"*{opts[0]}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
        else:
            render_rich_text(p1, opts[0], base_size=base_size)

        p1.add_run("\t")
        if is_teacher and opts[1].strip()[0] == correct_letter:
            render_rich_text(p1, f"*{opts[1]}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
        else:
            render_rich_text(p1, opts[1], base_size=base_size)

        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(1)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.tab_stops.clear_all()
        p2.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p2.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)

        p2.add_run("\t")
        if is_teacher and opts[2].strip()[0] == correct_letter:
            render_rich_text(p2, f"*{opts[2]}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
        else:
            render_rich_text(p2, opts[2], base_size=base_size)

        p2.add_run("\t")
        if is_teacher and len(opts) > 3 and opts[3].strip()[0] == correct_letter:
            render_rich_text(p2, f"*{opts[3]}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
        elif len(opts) > 3:
            render_rich_text(p2, opts[3], base_size=base_size)

    else:
        # Mỗi phương án 1 dòng
        for opt in opts:
            p_line = doc.add_paragraph()
            p_line.paragraph_format.left_indent = Cm(0.5)
            p_line.paragraph_format.space_before = Pt(1)
            p_line.paragraph_format.space_after = Pt(1)
            letter = opt.strip()[0]
            if is_teacher and letter == correct_letter:
                render_rich_text(p_line, f"*{opt}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p_line, opt, base_size=base_size)

def build_exam_docx(output_path, mode="student"):
    """
    Hàm chính tạo file Word đề cương ôn tập
    mode: 'student' | 'teacher' | 'screen'
    """
    is_teacher = (mode == "teacher")
    is_screen = (mode == "screen")
    base_size = 11.5 if is_screen else 10.5

    doc = docx.Document()
    setup_header_footer(doc)
    build_title_block(doc, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    # =========================================================================
    # PHẦN MỞ ĐẦU: BÀI 1
    # =========================================================================
    add_chapter_title(
        doc,
        "Bài 1. Nhận biết một số dụng cụ, hoá chất. Thuyết trình một vấn đề khoa học",
        base_size=base_size + 1
    )

    for q in BAI_1_QUESTIONS:
        qid = q["id"]
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(1)
        p_q.paragraph_format.line_spacing = 1.15
        render_rich_text(p_q, f"*Câu {qid}.* {q['content']}", base_size=base_size)

        render_mcq_options(doc, q["options"], q["correct"], is_teacher, base_size=base_size)

        if is_teacher and "solution" in q:
            add_solution_block(doc, q["solution"], base_size=base_size * 0.95)

        if is_screen:
            add_screen_teaching_space(doc, num_lines=2)

    # =========================================================================
    # PHÂN MÔN KHTN 1 (HÓA HỌC)
    # =========================================================================
    add_chapter_title(
        doc,
        "PHÂN MÔN KHTN 1 (HÓA HỌC)",
        base_size=base_size + 1.5
    )

    # I. NỘI DUNG ÔN TẬP
    add_section_header(doc, "I. NỘI DUNG ÔN TẬP", base_size=base_size + 0.5)
    topics = [
        "- Bài 22. Giới thiệu về hợp chất hữu cơ",
        "- Bài 23. Alkane",
        "- Bài 24. Alkene",
        "- Bài 25. Nguồn nhiên liệu",
        "- Bài 26. Ethylic alcohol",
        "- Bài 27. Acetic acid (tiết 1)"
    ]
    for topic in topics:
        p_top = doc.add_paragraph()
        p_top.paragraph_format.left_indent = Cm(0.5)
        p_top.paragraph_format.space_before = Pt(0)
        p_top.paragraph_format.space_after = Pt(1)
        render_rich_text(p_top, topic, base_size=base_size)

    # II. TRẮC NGHIỆM
    add_section_header(doc, "II. TRẮC NGHIỆM", base_size=base_size + 0.5)

    # Phần 1: Câu hỏi nhiều lựa chọn (24 câu)
    p_p1_title = doc.add_paragraph()
    p_p1_title.paragraph_format.space_before = Pt(4)
    p_p1_title.paragraph_format.space_after = Pt(2)
    render_rich_text(p_p1_title, "*Phần 1: Câu hỏi nhiều lựa chọn (24 câu)*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)

    for q in HOA_PART1_QUESTIONS:
        qid = q["id"]
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(1)
        p_q.paragraph_format.line_spacing = 1.15
        render_rich_text(p_q, f"*Câu {qid}:* {q['content']}", base_size=base_size)

        render_mcq_options(doc, q["options"], q["correct"], is_teacher, base_size=base_size)

        if is_teacher and "solution" in q:
            add_solution_block(doc, q["solution"], base_size=base_size * 0.95)

        if is_screen:
            add_screen_teaching_space(doc, num_lines=2)

    # Phần 2: Trắc nghiệm đúng sai (4 câu)
    p_p2_title = doc.add_paragraph()
    p_p2_title.paragraph_format.space_before = Pt(6)
    p_p2_title.paragraph_format.space_after = Pt(2)
    render_rich_text(p_p2_title, "*Phần 2: Trắc nghiệm đúng sai (4 câu)*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)

    for q in HOA_PART2_QUESTIONS:
        qid = q["id"]
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(1)
        p_q.paragraph_format.line_spacing = 1.15
        render_rich_text(p_q, f"*Câu {qid}:* {q['title']}", base_size=base_size)

        for st in q["statements"]:
            p_st = doc.add_paragraph()
            p_st.paragraph_format.left_indent = Cm(0.5)
            p_st.paragraph_format.space_before = Pt(1)
            p_st.paragraph_format.space_after = Pt(1)
            p_st.paragraph_format.line_spacing = 1.15

            lbl = st["label"]
            txt = st["text"]
            ans_mark = ""
            if is_teacher:
                ans_str = " (Đúng)" if st["correct"] else " (Sai)"
                ans_color = COLOR_CORRECT if st["correct"] else RGBColor(0, 105, 92)
                render_rich_text(p_st, f"*{lbl})* {txt}", base_size=base_size)
                r_ans = p_st.add_run(ans_str)
                set_run_style(r_ans, font_size=base_size, bold=True, color=ans_color)
            else:
                render_rich_text(p_st, f"*{lbl})* {txt}", base_size=base_size)

        if is_teacher:
            exp_text = "\n".join([f"- Ý {st['label']}: {st['explain']}" for st in q["statements"]])
            add_solution_block(doc, exp_text, base_size=base_size * 0.95)

        if is_screen:
            add_screen_teaching_space(doc, num_lines=3)

    # Phần 3: Câu hỏi trả lời ngắn (4 câu)
    p_p3_title = doc.add_paragraph()
    p_p3_title.paragraph_format.space_before = Pt(6)
    p_p3_title.paragraph_format.space_after = Pt(2)
    render_rich_text(p_p3_title, "*Phần 3: Câu hỏi trả lời ngắn (4 câu)*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)

    for q in HOA_PART3_QUESTIONS:
        qid = q["id"]
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(1)
        p_q.paragraph_format.line_spacing = 1.15
        render_rich_text(p_q, f"*Câu {qid}:* {q['content']}", base_size=base_size)

        if is_teacher:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Cm(0.5)
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(1)
            r_ans_lbl = p_ans.add_run("Đáp số: ")
            set_run_style(r_ans_lbl, font_size=base_size, bold=True, color=COLOR_CORRECT)
            r_ans_val = p_ans.add_run(q["correct"])
            set_run_style(r_ans_val, font_size=base_size + 0.5, bold=True, color=COLOR_CORRECT)

            if "solution" in q:
                add_solution_block(doc, q["solution"], base_size=base_size * 0.95)
        else:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Cm(0.5)
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(3)
            r_ans_lbl = p_ans.add_run("Trả lời: ............................................................")
            set_run_style(r_ans_lbl, font_size=base_size, italic=True)

        if is_screen:
            add_screen_teaching_space(doc, num_lines=3)

    # III. TỰ LUẬN (4 câu)
    add_section_header(doc, "III. TỰ LUẬN", base_size=base_size + 0.5)

    for q in HOA_ESSAY_QUESTIONS:
        qid = q["id"]
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        render_rich_text(p_q, f"*Câu {qid}:* {q['content']}", base_size=base_size)

        if is_teacher:
            add_solution_block(doc, q["solution"], base_size=base_size * 0.95)
        else:
            add_writing_lines(doc, num_lines=4)

        if is_screen:
            add_screen_teaching_space(doc, num_lines=4)

    # =========================================================================
    # PHÂN MÔN KHTN 2 (VẬT LÍ)
    # =========================================================================
    add_chapter_title(
        doc,
        "PHÂN MÔN KHTN 2 (VẬT LÍ)",
        base_size=base_size + 1.5
    )

    # I. NỘI DUNG ÔN TẬP
    add_section_header(doc, "I. NỘI DUNG ÔN TẬP", base_size=base_size + 0.5)

    p_a = doc.add_paragraph()
    p_a.paragraph_format.space_before = Pt(3)
    p_a.paragraph_format.space_after = Pt(2)
    render_rich_text(p_a, "*A. LÍ THUYẾT:*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)

    # 1. Động năng
    p_dn = doc.add_paragraph()
    p_dn.paragraph_format.left_indent = Cm(0.3)
    p_dn.paragraph_format.space_before = Pt(3)
    p_dn.paragraph_format.space_after = Pt(1)
    render_rich_text(p_dn, "*I. ĐỘNG NĂNG*", base_size=base_size, base_bold=True)

    dn_bullets = [
        "- Động năng là dạng năng lượng của một vật có được do nó đang chuyển động.",
        "- Vật có khối lượng càng lớn chuyển động càng nhanh thì động năng càng lớn.",
        "- Động năng của một vật bằng một nửa tích của khối lượng và bình phương vận tốc của vật."
    ]
    for b in dn_bullets:
        p_b = doc.add_paragraph()
        p_b.paragraph_format.left_indent = Cm(0.6)
        p_b.paragraph_format.space_before = Pt(0)
        p_b.paragraph_format.space_after = Pt(1)
        render_rich_text(p_b, b, base_size=base_size)

    add_callout_box(
        doc,
        "Công thức tính động năng",
        [
            "W_đ = (1/2)·m·v^2",
            "Trong đó: W_đ là động năng [J]; m là khối lượng [kg]; v là tốc độ chuyển động [m/s]"
        ],
        base_size=base_size
    )

    # 2. Thế năng
    p_tn = doc.add_paragraph()
    p_tn.paragraph_format.left_indent = Cm(0.3)
    p_tn.paragraph_format.space_before = Pt(4)
    p_tn.paragraph_format.space_after = Pt(1)
    render_rich_text(p_tn, "*II. THẾ NĂNG*", base_size=base_size, base_bold=True)

    tn_bullets = [
        "- Thế năng trọng trường, hay gọi tắt là thế năng, là năng lượng vật có được khi ở trên cao so với mặt đất (hoặc so với một vị trí khác được chọn làm mốc để tính độ cao).",
        "- Vật có khối lượng càng lớn ở độ cao càng cao thì thế năng càng lớn.",
        "- Giá trị của thế năng phụ thuộc vào mốc chọn để tính độ cao, hay còn gọi là gốc thế năng.",
        "- Thông thường, gốc thế năng được chọn tại mặt đất."
    ]
    for b in tn_bullets:
        p_b = doc.add_paragraph()
        p_b.paragraph_format.left_indent = Cm(0.6)
        p_b.paragraph_format.space_before = Pt(0)
        p_b.paragraph_format.space_after = Pt(1)
        render_rich_text(p_b, b, base_size=base_size)

    add_callout_box(
        doc,
        "Công thức tính thế năng trọng trường",
        [
            "W_t = P·h = 10·m·h",
            "Trong đó: W_t là thế năng [J]; P là trọng lượng [N]; m là khối lượng [kg]; h là độ cao [m]"
        ],
        base_size=base_size
    )

    # 3. Cơ năng
    p_cn = doc.add_paragraph()
    p_cn.paragraph_format.left_indent = Cm(0.3)
    p_cn.paragraph_format.space_before = Pt(4)
    p_cn.paragraph_format.space_after = Pt(1)
    render_rich_text(p_cn, "*III. CƠ NĂNG*", base_size=base_size, base_bold=True)

    cn_bullets = [
        "- Trong thực tế, một vật có thể vừa có động năng, vừa có thế năng.",
        "- Cơ năng của vật chuyển động dưới tác dụng của trọng lực bằng tổng động năng và thế năng trọng trường của vật."
    ]
    for b in cn_bullets:
        p_b = doc.add_paragraph()
        p_b.paragraph_format.left_indent = Cm(0.6)
        p_b.paragraph_format.space_before = Pt(0)
        p_b.paragraph_format.space_after = Pt(1)
        render_rich_text(p_b, b, base_size=base_size)

    add_callout_box(
        doc,
        "Công thức tính cơ năng",
        [
            "W = W_t + W_đ = P·h + (1/2)·m·v^2 = 10·m·h + (1/2)·m·v^2  (J)"
        ],
        base_size=base_size
    )

    # 4. Công suất
    p_cs = doc.add_paragraph()
    p_cs.paragraph_format.left_indent = Cm(0.3)
    p_cs.paragraph_format.space_before = Pt(4)
    p_cs.paragraph_format.space_after = Pt(1)
    render_rich_text(p_cs, "*IV. CÔNG SUẤT*", base_size=base_size, base_bold=True)

    p_c1 = doc.add_paragraph()
    p_c1.paragraph_format.left_indent = Cm(0.6)
    p_c1.paragraph_format.space_before = Pt(1)
    p_c1.paragraph_format.space_after = Pt(1)
    render_rich_text(p_c1, "*1. Công cơ học:*", base_size=base_size, base_bold=True)

    add_callout_box(
        doc,
        "Công thức tính công",
        [
            "A = F·s",
            "Trong đó: F là lực tác dụng [N]; s là quãng đường dịch chuyển theo hướng của lực [m]; A là công cơ học [J]"
        ],
        base_size=base_size
    )

    p_c2 = doc.add_paragraph()
    p_c2.paragraph_format.left_indent = Cm(0.6)
    p_c2.paragraph_format.space_before = Pt(2)
    p_c2.paragraph_format.space_after = Pt(1)
    render_rich_text(p_c2, "*2. Công suất:*", base_size=base_size, base_bold=True)

    cs_bullets = [
        "- Định nghĩa công suất: Công suất là tốc độ thực hiện công, đặc trưng cho mức độ nhanh hay chậm của việc sinh công.",
        "- Đơn vị đo công suất: Oát (kí hiệu: W); kilôoát (kW); mêgaoát (MW); gigaoát (GW)."
    ]
    for b in cs_bullets:
        p_b = doc.add_paragraph()
        p_b.paragraph_format.left_indent = Cm(0.6)
        p_b.paragraph_format.space_before = Pt(0)
        p_b.paragraph_format.space_after = Pt(1)
        render_rich_text(p_b, b, base_size=base_size)

    add_callout_box(
        doc,
        "Công thức tính công suất & Đổi đơn vị",
        [
            "P = A / t",
            "Trong đó: A là công thực hiện [J]; t là thời gian thực hiện công [s]; P là công suất [W]",
            "1 kW = 10^3 W  |  1 MW = 10^6 W  |  1 GW = 10^9 W"
        ],
        base_size=base_size
    )

    # 5. Vòng năng lượng trên Trái Đất. Năng lượng hóa thạch
    p_vnl = doc.add_paragraph()
    p_vnl.paragraph_format.left_indent = Cm(0.3)
    p_vnl.paragraph_format.space_before = Pt(4)
    p_vnl.paragraph_format.space_after = Pt(1)
    render_rich_text(p_vnl, "*V. VÒNG NĂNG LƯỢNG TRÊN TRÁI ĐẤT. NĂNG LƯỢNG HÓA THẠCH*", base_size=base_size, base_bold=True)

    vnl_bullets = [
        "- *Mô tả vòng năng lượng giữa các vật sống trên Trái Đất:* Thực vật hấp thụ năng lượng mặt trời, chuyển hoá năng lượng mặt trời thành hoá năng lưu trữ trong lục lạp của chúng. Động vật lấy glucose từ quá trình phân giải thức ăn từ thực vật được phân giải thành khí CO~2~ và nước, đồng thời tạo ra năng lượng ATP.",
        "- *Mô tả sự chuyển hoá giữa các dạng năng lượng trong vòng tuần hoàn của nước:* Ánh sáng mặt trời chiếu xuống mặt đất và biển, làm nóng nước và khí quyển, tạo ra sự chênh lệch nhiệt độ giữa các vùng trên Trái Đất, gây ra sự chuyển động của khí quyển và các dòng hải lưu (dòng biển) trong đại dương, tạo ra gió và các hiện tượng thời tiết. Nước từ dạng lỏng bốc hơi dưới tác động của ánh sáng mặt trời, hơi nước ngưng tụ thành các đám mây và mưa. Thoát hơi nước ở thực vật trên cạn cũng chuyển vận một lượng hơi nước đáng kể vào bầu khí quyển. Nước ở bề mặt đất và nước ngầm có thể chảy xuống đại dương, khép kín vòng chuyển vận nước.",
        "- *Nguồn gốc năng lượng hoá thạch:* Năng lượng hoá thạch có nguồn gốc từ Mặt Trời.",
        "- *Ưu điểm:* Có thể khai thác với khối lượng lớn, dễ vận chuyển, công nghệ chuyển hoá thành các dạng năng lượng khác phổ biến với chi phí rẻ.",
        "- *Nhược điểm:* Khi sử dụng gây phát thải khí độc hại và ô nhiễm môi trường."
    ]
    for b in vnl_bullets:
        p_b = doc.add_paragraph()
        p_b.paragraph_format.left_indent = Cm(0.6)
        p_b.paragraph_format.space_before = Pt(1)
        p_b.paragraph_format.space_after = Pt(2)
        render_rich_text(p_b, b, base_size=base_size)

    # 6. Một số dạng năng lượng tái tạo (Bảng)
    p_tt = doc.add_paragraph()
    p_tt.paragraph_format.left_indent = Cm(0.3)
    p_tt.paragraph_format.space_before = Pt(4)
    p_tt.paragraph_format.space_after = Pt(2)
    render_rich_text(p_tt, "*VI. MỘT SỐ DẠNG NĂNG LƯỢNG TÁI TẠO*", base_size=base_size, base_bold=True)

    p_tt_def = doc.add_paragraph()
    p_tt_def.paragraph_format.left_indent = Cm(0.6)
    p_tt_def.paragraph_format.space_before = Pt(0)
    p_tt_def.paragraph_format.space_after = Pt(3)
    render_rich_text(p_tt_def, "- Năng lượng tái tạo là năng lượng đến từ các nguồn năng lượng có sẵn trong thiên nhiên, liên tục được bổ sung thông qua các quá trình tự nhiên.", base_size=base_size)

    # Bảng tổng hợp Ưu điểm - Nhược điểm
    tbl_tt = doc.add_table(rows=5, cols=3)
    tbl_tt.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_tt, color="B0BEC5", sz="4")

    headers_tt = ["Dạng năng lượng tái tạo", "Ưu điểm nổi bật", "Nhược điểm cần khắc phục"]
    widths_tt = [Cm(4.5), Cm(7.0), Cm(6.5)]

    for j, h in enumerate(headers_tt):
        cell = tbl_tt.cell(0, j)
        cell.width = widths_tt[j]
        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        set_cell_shading(cell, "E8F5E9")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_style(r, font_size=base_size * 0.95, bold=True, color=COLOR_PRIMARY)

    data_tt = [
        ("Năng lượng mặt trời", "- Nguồn năng lượng vô tận.\n- Không gây ô nhiễm môi trường.\n- Giảm chi phí điện lâu dài.", "- Chi phí ban đầu cao.\n- Phụ thuộc vào thời tiết.\n- Yêu cầu diện tích lớn để lắp đặt pin mặt trời."),
        ("Năng lượng từ gió", "- Nguồn năng lượng vô tận.\n- Không gây ô nhiễm môi trường.\n- Có thể kết hợp lắp đặt trên đất nông nghiệp.", "- Chi phí ban đầu cao.\n- Ảnh hưởng đến cảnh quan.\n- Phụ thuộc vào điều kiện gió."),
        ("Năng lượng từ sóng biển", "- Nguồn năng lượng vô tận.\n- Không gây ô nhiễm môi trường.\n- Tạo ra năng lượng ổn định.", "- Chi phí đầu tư cao.\n- Ảnh hưởng đến hệ sinh thái biển.\n- Yêu cầu công nghệ phức tạp."),
        ("Năng lượng từ dòng sông", "- Tận dụng dòng chảy tự nhiên.\n- Không gây ô nhiễm môi trường.\n- Có thể kết hợp với các nhà máy thủy điện.", "- Phụ thuộc vào lưu lượng nước theo mùa.\n- Có thể ảnh hưởng đến sinh thái sông.\n- Chi phí xây dựng công trình cao.")
    ]

    for idx, (col1, col2, col3) in enumerate(data_tt):
        row_idx = idx + 1
        c1, c2, c3 = tbl_tt.cell(row_idx, 0), tbl_tt.cell(row_idx, 1), tbl_tt.cell(row_idx, 2)
        c1.width = widths_tt[0]; c2.width = widths_tt[1]; c3.width = widths_tt[2]
        set_cell_margins(c1, top=50, bottom=50, left=80, right=80)
        set_cell_margins(c2, top=50, bottom=50, left=80, right=80)
        set_cell_margins(c3, top=50, bottom=50, left=80, right=80)

        p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        render_rich_text(p1, f"*{col1}*", base_size=base_size * 0.95, base_bold=True)

        p2 = c2.paragraphs[0]; p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        render_rich_text(p2, col2, base_size=base_size * 0.9)

        p3 = c3.paragraphs[0]; p3.alignment = WD_ALIGN_PARAGRAPH.LEFT
        render_rich_text(p3, col3, base_size=base_size * 0.9)

    # B. TRẮC NGHIỆM VẬT LÍ
    p_b_head = doc.add_paragraph()
    p_b_head.paragraph_format.space_before = Pt(8)
    p_b_head.paragraph_format.space_after = Pt(2)
    render_rich_text(p_b_head, "*B. TRẮC NGHIỆM VẬT LÍ:*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)

    p_p1_vl = doc.add_paragraph()
    p_p1_vl.paragraph_format.space_before = Pt(2)
    p_p1_vl.paragraph_format.space_after = Pt(2)
    render_rich_text(p_p1_vl, "*Phần 1: Câu hỏi nhiều lựa chọn*", base_size=base_size, base_bold=True, base_color=RGBColor(21, 101, 192))

    for q in VATLI_QUESTIONS:
        qid = q["id"]
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(1)
        p_q.paragraph_format.line_spacing = 1.15
        render_rich_text(p_q, f"*Câu {qid}.* {q['content']}", base_size=base_size)

        render_mcq_options(doc, q["options"], q["correct"], is_teacher, base_size=base_size)

        if is_teacher and "solution" in q:
            add_solution_block(doc, q["solution"], base_size=base_size * 0.95)

        if is_screen:
            add_screen_teaching_space(doc, num_lines=2)

    # =========================================================================
    # BẢNG TỔNG HỢP ĐÁP ÁN (DÀNH CHO BẢN GIÁO VIÊN)
    # =========================================================================
    if is_teacher:
        p_ans_head = doc.add_paragraph()
        p_ans_head.paragraph_format.space_before = Pt(14)
        p_ans_head.paragraph_format.space_after = Pt(4)
        p_ans_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_ah = p_ans_head.add_run("BẢNG TỔNG HỢP ĐÁP ÁN ĐỀ CƯƠNG ÔN TẬP")
        set_run_style(r_ah, font_size=12, bold=True, color=COLOR_CORRECT)

        # 1. Bảng Bài 1 (5 câu)
        p_tb1 = doc.add_paragraph()
        p_tb1.paragraph_format.space_before = Pt(4)
        p_tb1.paragraph_format.space_after = Pt(2)
        render_rich_text(p_tb1, "*1. Đáp án Bài 1 (Dụng cụ, hóa chất & Thuyết trình - 5 câu):*", base_size=10, base_bold=True, base_color=COLOR_PRIMARY)

        tbl_b1 = doc.add_table(rows=2, cols=5)
        tbl_b1.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl_b1, color="B0BEC5", sz="4")
        for i in range(5):
            c0 = tbl_b1.cell(0, i); c1 = tbl_b1.cell(1, i)
            c0.width = Cm(3.6); c1.width = Cm(3.6)
            set_cell_margins(c0, 30, 30, 40, 40); set_cell_margins(c1, 30, 30, 40, 40)
            set_cell_shading(c0, "E8F5E9")
            p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r0 = p0.add_run(f"Câu {i+1}")
            set_run_style(r0, font_size=9.5, bold=True)
            p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r1 = p1.add_run(BAI_1_QUESTIONS[i]["correct"])
            set_run_style(r1, font_size=11, bold=True, color=COLOR_CORRECT)

        # 2. Bảng Hóa học Phần 1 (24 câu)
        p_th1 = doc.add_paragraph()
        p_th1.paragraph_format.space_before = Pt(6)
        p_th1.paragraph_format.space_after = Pt(2)
        render_rich_text(p_th1, "*2. Đáp án Phân môn Hóa học - Phần 1 (Trắc nghiệm 24 câu):*", base_size=10, base_bold=True, base_color=COLOR_PRIMARY)

        tbl_h1 = doc.add_table(rows=4, cols=12)
        tbl_h1.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl_h1, color="B0BEC5", sz="4")
        # Hàng 1 & 2: Câu 1 đến 12
        for i in range(12):
            c0 = tbl_h1.cell(0, i); c1 = tbl_h1.cell(1, i)
            c0.width = Cm(1.5); c1.width = Cm(1.5)
            set_cell_margins(c0, 30, 30, 30, 30); set_cell_margins(c1, 30, 30, 30, 30)
            set_cell_shading(c0, "E8F5E9")
            p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r0 = p0.add_run(f"C{i+1}")
            set_run_style(r0, font_size=9, bold=True)
            p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r1 = p1.add_run(HOA_PART1_QUESTIONS[i]["correct"])
            set_run_style(r1, font_size=10, bold=True, color=COLOR_CORRECT)

        # Hàng 3 & 4: Câu 13 đến 24
        for i in range(12):
            c0 = tbl_h1.cell(2, i); c1 = tbl_h1.cell(3, i)
            c0.width = Cm(1.5); c1.width = Cm(1.5)
            set_cell_margins(c0, 30, 30, 30, 30); set_cell_margins(c1, 30, 30, 30, 30)
            set_cell_shading(c0, "E8F5E9")
            p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r0 = p0.add_run(f"C{i+13}")
            set_run_style(r0, font_size=9, bold=True)
            p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r1 = p1.add_run(HOA_PART1_QUESTIONS[i+12]["correct"])
            set_run_style(r1, font_size=10, bold=True, color=COLOR_CORRECT)

        # 3. Bảng Hóa học Phần 2 (Đúng / Sai)
        p_th2 = doc.add_paragraph()
        p_th2.paragraph_format.space_before = Pt(6)
        p_th2.paragraph_format.space_after = Pt(2)
        render_rich_text(p_th2, "*3. Đáp án Phân môn Hóa học - Phần 2 (Trắc nghiệm Đúng/Sai - 4 câu):*", base_size=10, base_bold=True, base_color=COLOR_PRIMARY)

        tbl_h2 = doc.add_table(rows=5, cols=5)
        tbl_h2.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl_h2, color="B0BEC5", sz="4")
        headers_p2 = ["Câu hỏi", "Lệnh hỏi a", "Lệnh hỏi b", "Lệnh hỏi c", "Lệnh hỏi d"]
        for j, h in enumerate(headers_p2):
            cell = tbl_h2.cell(0, j)
            cell.width = Cm(3.6)
            set_cell_margins(cell, 30, 30, 40, 40)
            set_cell_shading(cell, "E8F5E9")
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(h)
            set_run_style(r, font_size=9.5, bold=True)

        for q_idx, q in enumerate(HOA_PART2_QUESTIONS):
            row_idx = q_idx + 1
            c_qid = tbl_h2.cell(row_idx, 0)
            c_qid.width = Cm(3.6)
            set_cell_margins(c_qid, 30, 30, 40, 40)
            p = c_qid.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(f"Câu {q['id']}")
            set_run_style(r, font_size=9.5, bold=True)
            for st_idx, st in enumerate(q["statements"]):
                c_st = tbl_h2.cell(row_idx, st_idx + 1)
                c_st.width = Cm(3.6)
                set_cell_margins(c_st, 30, 30, 40, 40)
                p_st = c_st.paragraphs[0]; p_st.alignment = WD_ALIGN_PARAGRAPH.CENTER
                val_str = "Đ" if st["correct"] else "S"
                r_st = p_st.add_run(val_str)
                color_st = COLOR_CORRECT if st["correct"] else RGBColor(0, 105, 92)
                set_run_style(r_st, font_size=10, bold=True, color=color_st)

        # 4. Bảng Hóa học Phần 3 (Trả lời ngắn)
        p_th3 = doc.add_paragraph()
        p_th3.paragraph_format.space_before = Pt(6)
        p_th3.paragraph_format.space_after = Pt(2)
        render_rich_text(p_th3, "*4. Đáp án Phân môn Hóa học - Phần 3 (Trả lời ngắn - 4 câu):*", base_size=10, base_bold=True, base_color=COLOR_PRIMARY)

        tbl_h3 = doc.add_table(rows=2, cols=4)
        tbl_h3.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl_h3, color="B0BEC5", sz="4")
        for i in range(4):
            c0 = tbl_h3.cell(0, i); c1 = tbl_h3.cell(1, i)
            c0.width = Cm(4.5); c1.width = Cm(4.5)
            set_cell_margins(c0, 30, 30, 40, 40); set_cell_margins(c1, 30, 30, 40, 40)
            set_cell_shading(c0, "E8F5E9")
            p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r0 = p0.add_run(f"Câu {i+1}")
            set_run_style(r0, font_size=9.5, bold=True)
            p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r1 = p1.add_run(HOA_PART3_QUESTIONS[i]["correct"])
            set_run_style(r1, font_size=11, bold=True, color=COLOR_CORRECT)

        # 5. Bảng Trắc nghiệm Vật lí
        p_tvl = doc.add_paragraph()
        p_tvl.paragraph_format.space_before = Pt(6)
        p_tvl.paragraph_format.space_after = Pt(2)
        render_rich_text(p_tvl, "*5. Đáp án Trắc nghiệm Phân môn Vật lí (2 câu):*", base_size=10, base_bold=True, base_color=COLOR_PRIMARY)

        tbl_vl = doc.add_table(rows=2, cols=2)
        tbl_vl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl_vl, color="B0BEC5", sz="4")
        for i in range(2):
            c0 = tbl_vl.cell(0, i); c1 = tbl_vl.cell(1, i)
            c0.width = Cm(4.0); c1.width = Cm(4.0)
            set_cell_margins(c0, 30, 30, 40, 40); set_cell_margins(c1, 30, 30, 40, 40)
            set_cell_shading(c0, "E8F5E9")
            p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r0 = p0.add_run(f"Câu {i+1}")
            set_run_style(r0, font_size=9.5, bold=True)
            p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r1 = p1.add_run(VATLI_QUESTIONS[i]["correct"])
            set_run_style(r1, font_size=11, bold=True, color=COLOR_CORRECT)

    # Lưu tài liệu
    doc.save(output_path)
    print(f"Created: {output_path}")

def main():
    files = [
        ("De_Cuong_KHTN_9_Giua_HK1_Kim_Long_HocSinh.docx", "student"),
        ("De_Cuong_KHTN_9_Giua_HK1_Kim_Long_GiaoVien.docx", "teacher"),
        ("De_Cuong_KHTN_9_Giua_HK1_Kim_Long_DayTrenManHinh.docx", "screen"),
    ]
    for filename, mode in files:
        out_file = os.path.join(OUT_DIR, filename)
        build_exam_docx(out_file, mode=mode)
        
        # Sao chép sang thư mục OneDrive của Thầy
        for onedrive_folder in ONEDRIVE_DIRS:
            if os.path.exists(onedrive_folder):
                try:
                    dest = os.path.join(onedrive_folder, filename)
                    shutil.copy2(out_file, dest)
                    print(f"Copied to OneDrive successfully: {dest}")
                except Exception as e:
                    print(f"Copy to OneDrive error ({onedrive_folder}): {e}")

if __name__ == "__main__":
    main()
