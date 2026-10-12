# -*- coding: utf-8 -*-
"""
Script tạo file Word (.docx) chuẩn "FORM ĐỀ KIỂM TRA" cho:
ĐỀ KIỂM TRA SỐ 2 - CHƯƠNG 1: ESTER - LIPID (HÓA HỌC 12)
Gồm 2 phiên bản:
1. De_Kiem_Tra_Hoa_12_De02.docx (Bản Học Sinh)
2. De_Kiem_Tra_Hoa_12_De02_GiaoVien.docx (Bản Giáo Viên có lời giải chi tiết)
Tuân thủ 100% quy cách: Header & Footer, Tab Stops (NO TABLE), Sub/Superscript công thức hóa học.
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
IMG_DIR = r"c:\Antigravity_Thanh\San_Pham\images_crop_de02"

COLOR_HF_RED = RGBColor(211, 47, 47)      # Coral Red (#D32F2F)
COLOR_HF_GREEN_HEX = "8BC390"             # Sage Green (#8BC390)
COLOR_PRIMARY = RGBColor(46, 125, 50)     # Dark Green (#2E7D32)
COLOR_CORRECT = RGBColor(198, 40, 40)     # Red for correct answers (#C62828)
COLOR_SOLUTION = RGBColor(21, 101, 192)   # Blue for solutions (#1565C0)
COLOR_GRAY = RGBColor(97, 97, 97)

def add_p_border_bottom(p, color="8BC390", sz="6"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_p_border_top(p, color="8BC390", sz="6"):
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def set_run_style(run, font_name="Times New Roman", font_size=12, bold=False, italic=False, subscript=False, superscript=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def render_rich_text(p, text, base_font="Times New Roman", base_size=10.5, base_bold=False, base_italic=False, base_color=None):
    tokens = re.split(r'(~[^~]+~|\^[^\^]+\^|\*[^\*]+\*|_[^_]+_)', text)
    for token in tokens:
        if not token:
            continue
        is_bold = base_bold
        is_italic = base_italic
        is_sub = False
        is_sup = False
        content = token

        if token.startswith('*') and token.endswith('*'):
            is_bold = True
            content = token[1:-1]
        elif token.startswith('_') and token.endswith('_'):
            is_italic = True
            content = token[1:-1]
        elif token.startswith('~') and token.endswith('~'):
            is_sub = True
            content = token[1:-1]
        elif token.startswith('^') and token.endswith('^'):
            is_sup = True
            content = token[1:-1]

        run = p.add_run(content)
        run_color = COLOR_CORRECT if (is_bold and not base_bold and not base_color) else base_color
        set_run_style(run, font_name=base_font, font_size=base_size, bold=is_bold, italic=is_italic,
                      subscript=is_sub, superscript=is_sup, color=run_color)

def setup_header_footer(doc):
    normal_style = doc.styles['Normal']
    normal_style.paragraph_format.tab_stops.clear_all()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    section.start_type = WD_SECTION_START.NEW_PAGE
    section.different_first_page_header_footer = False
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    # 1. Header (bảng 1 hàng 2 cột, rộng 18.0 cm)
    header = section.header
    for p in header.paragraphs:
        p.text = ""
    tbl_h = header.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl_h.cell(0, 0)
    c_right = tbl_h.cell(0, 1)
    c_left.width = Cm(8.0)
    c_right.width = Cm(10.0)

    p_hl = c_left.paragraphs[0]
    p_hl.paragraph_format.space_before = Pt(0)
    p_hl.paragraph_format.space_after = Pt(2)
    p_hl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_hl = p_hl.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)

    p_hr = c_right.paragraphs[0]
    p_hr.paragraph_format.space_before = Pt(0)
    p_hr.paragraph_format.space_after = Pt(2)
    p_hr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p_hr.add_run("Thầy ")
    set_run_style(r1, font_size=10, color=COLOR_HF_RED)
    r2 = p_hr.add_run("TRẦN VĂN THẠNH")
    set_run_style(r2, font_size=10, bold=True, color=COLOR_HF_RED)
    r3 = p_hr.add_run(" - 0777.470.803")
    set_run_style(r3, font_size=10, color=COLOR_HF_RED)

    add_p_border_bottom(p_hl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_bottom(p_hr, color=COLOR_HF_GREEN_HEX, sz="6")

    # 2. Footer (bảng 1 hàng 2 cột, rộng 18.0 cm)
    footer = section.footer
    for p in footer.paragraphs:
        p.text = ""
    tbl_f = footer.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_fl = tbl_f.cell(0, 0)
    c_fr = tbl_f.cell(0, 1)
    c_fl.width = Cm(14.5)
    c_fr.width = Cm(3.5)

    p_fl = c_fl.paragraphs[0]
    p_fl.paragraph_format.space_before = Pt(2)
    p_fl.paragraph_format.space_after = Pt(0)
    p_fl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_fl = p_fl.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)

    p_fr = c_fr.paragraphs[0]
    p_fr.paragraph_format.space_before = Pt(2)
    p_fr.paragraph_format.space_after = Pt(0)
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_field(p_fr, font_size=10, bold=True, color=COLOR_HF_RED)

    add_p_border_top(p_fl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_top(p_fr, color=COLOR_HF_GREEN_HEX, sz="6")

def build_title_block_no_table(doc, is_teacher=False):
    headers_data = [
        ("SỞ GIÁO DỤC VÀ ĐÀO TẠO THỪA THIÊN HUẾ", "ĐỀ KIỂM TRA ĐỊNH KỲ CHƯƠNG 1 (ESTER - LIPID)", 10, True, 11, True),
        ("ĐỀ CHÍNH THỨC", "MÔN: HOÁ HỌC - LỚP 12", 10.5, True, 11, True),
        ("(Đề thi có 04 trang)", "Thời gian: 50 phút | MÃ ĐỀ THI: 102", 9.5, False, 10, True)
    ]
    for l_text, r_text, l_size, l_bold, r_size, r_bold in headers_data:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)

        rl = p.add_run(l_text)
        set_run_style(rl, font_size=l_size, bold=l_bold, italic=(not l_bold))

        p.add_run('\t')
        rr = p.add_run(r_text)
        set_run_style(rr, font_size=r_size, bold=r_bold, italic=(not r_bold))

    if is_teacher:
        p_gv = doc.add_paragraph()
        p_gv.paragraph_format.space_before = Pt(2)
        p_gv.paragraph_format.space_after = Pt(2)
        p_gv.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rgv = p_gv.add_run("(BẢN GIÁO VIÊN -- CÓ LỜI GIẢI CHI TIẾT)")
        set_run_style(rgv, font_size=11, bold=True, color=COLOR_CORRECT)

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(4)
    p_info.paragraph_format.space_after = Pt(2)
    render_rich_text(p_info, "*Họ, tên thí sinh:* ............................................................ *Lớp:* ............. *Số báo danh:* .............", base_size=10.5)

    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(0)
    p_note.paragraph_format.space_after = Pt(4)
    render_rich_text(p_note, "_(Cho biết nguyên tử khối: H = 1; C = 12; N = 14; O = 16; Na = 23; K = 39; Br = 80; Cl = 35,5)._", base_size=10, base_italic=True)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(4)
    add_p_border_bottom(p_div, color="CCCCCC", sz="4")

def add_section_header(doc, title, subtitle=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, f"*{title}*", base_size=11.5, base_bold=True, base_color=COLOR_PRIMARY)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(4)
        render_rich_text(p2, f"_{subtitle}_", base_size=10.5, base_italic=True)

def render_question_prompt(p, text, base_size=10.5):
    m = re.match(r'^(Câu\s+\d+\.)\s*(.*)$', text.strip(), re.DOTALL)
    if m:
        lbl = m.group(1)
        rest = m.group(2)
        r_lbl = p.add_run(lbl + " ")
        set_run_style(r_lbl, font_size=base_size, bold=True)
        render_rich_text(p, rest, base_size=base_size, base_bold=False)
    else:
        render_rich_text(p, text, base_size=base_size, base_bold=False)

def render_single_option(p, opt_str, is_teacher=False, base_size=10.5):
    clean_str = opt_str.strip()
    is_correct = False
    if clean_str.startswith('*') and clean_str.endswith('*'):
        is_correct = True
        clean_str = clean_str[1:-1].strip()

    m = re.match(r'^([A-D]\.\s*)(.*)$', clean_str)
    if m:
        lbl = m.group(1)
        body = m.group(2)
    else:
        lbl = ""
        body = clean_str

    if is_teacher and is_correct:
        if lbl:
            r_lbl = p.add_run(lbl)
            set_run_style(r_lbl, font_size=base_size, bold=True, color=COLOR_CORRECT)
        render_rich_text(p, body, base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
    else:
        if lbl:
            r_lbl = p.add_run(lbl)
            set_run_style(r_lbl, font_size=base_size, bold=True, color=None)
        render_rich_text(p, body, base_size=base_size, base_bold=False, base_color=None)

def render_options_no_table(doc, opts, is_teacher=False):
    clean_opts = [re.sub(r'[*_~^]', '', opt) for opt in opts]
    max_len = max(len(o) for o in clean_opts)

    if max_len <= 18:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(4.5), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(13.5), WD_TAB_ALIGNMENT.LEFT)

        for idx, opt in enumerate(opts):
            if idx > 0:
                p.add_run('\t')
            render_single_option(p, opt, is_teacher=is_teacher, base_size=10.5)
    elif max_len <= 38:
        p1 = doc.add_paragraph()
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.tab_stops.clear_all()
        p1.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
        render_single_option(p1, opts[0], is_teacher=is_teacher, base_size=10.5)
        p1.add_run('\t')
        render_single_option(p1, opts[1], is_teacher=is_teacher, base_size=10.5)

        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(1)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.tab_stops.clear_all()
        p2.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
        render_single_option(p2, opts[2], is_teacher=is_teacher, base_size=10.5)
        p2.add_run('\t')
        render_single_option(p2, opts[3], is_teacher=is_teacher, base_size=10.5)
    else:
        for opt in opts:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            render_single_option(p, opt, is_teacher=is_teacher, base_size=10.5)

def render_statement(p, stmt, ans=None, is_teacher=False, base_size=10.5):
    m = re.match(r'^([a-d]\))\s*(.*)$', stmt.strip())
    if m:
        lbl = m.group(1)
        body = m.group(2)
        r_lbl = p.add_run(lbl + " ")
        set_run_style(r_lbl, font_size=base_size, bold=True)
        render_rich_text(p, body, base_size=base_size, base_bold=False)
    else:
        render_rich_text(p, stmt, base_size=base_size, base_bold=False)

    if is_teacher and ans:
        p.add_run("  ➔ ")
        r_ans = p.add_run(f"({ans})")
        set_run_style(r_ans, font_size=base_size, bold=True, color=COLOR_CORRECT)

def add_solution_no_table(doc, solution_text):
    lines = solution_text.strip().split("\n")
    p_head = doc.add_paragraph()
    p_head.paragraph_format.left_indent = Cm(0.4)
    p_head.paragraph_format.space_before = Pt(3)
    p_head.paragraph_format.space_after = Pt(1)
    render_rich_text(p_head, "► *Lời giải chi tiết:*", base_size=10.5, base_bold=True, base_color=COLOR_SOLUTION)

    for line in lines:
        if line.startswith("*Lời giải chi tiết:*") or line.startswith("►"):
            continue
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        render_rich_text(p, line, base_size=10)

def build_exam_de02(is_teacher=False):
    doc = docx.Document()
    setup_header_footer(doc)
    build_title_block_no_table(doc, is_teacher=is_teacher)

    # =========================================================================
    # PHẦN I
    # =========================================================================
    add_section_header(doc, "PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.",
                       "Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án.")

    part1_questions = [
        {
            "num": 1,
            "q": "Câu 1. HCOOCH~3~ được sử dụng làm dung môi cho nitrocellulose và cellulose acetate, một chất trung gian trong sản xuất dược phẩm và thuốc xông hơi. Tên gọi của HCOOCH~3~ là",
            "opts": ["A. ethyl formate.", "B. ethyl acetate.", "*C. methyl formate.*", "D. methyl acetate."],
            "sol": "HCOOCH~3~ gồm gốc formate (HCOO-) và gốc methyl (-CH~3~), do đó tên gọi là methyl formate."
        },
        {
            "num": 2,
            "q": "Câu 2. Ester nào sau đây tác dụng với NaOH thu được ethyl alcohol?",
            "opts": ["*A. CH~3~COOC~2~H~5~.*", "B. CH~3~COOC~3~H~7~.", "C. C~2~H~5~COOCH~3~.", "D. HCOOCH~3~."],
            "sol": "Thuỷ phân ester RCOOR' trong NaOH: RCOOR' + NaOH ➔ RCOONa + R'OH.\nĐể thu được ethyl alcohol (C~2~H~5~OH) thì gốc R' phải là -C~2~H~5~.\nVậy ester là CH~3~COOC~2~H~5~ (ethyl acetate)."
        },
        {
            "num": 3,
            "q": "Câu 3. Công thức tổng quát của ester mạch hở được tạo thành từ carboxylic acid no, đơn chức và alcohol no, đơn chức là",
            "opts": ["A. C~n~H~2n-1~COOC~m~H~2m+1~.", "B. C~n~H~2n-1~COOC~m~H~2m-1~.", "C. C~n~H~2n+1~COOC~m~H~2m-1~.", "*D. C~n~H~2n+1~COOC~m~H~2m+1~.*"],
            "sol": "Acid no, đơn chức có gốc alkyl là C~n~H~2n+1~- (n ≥ 0), alcohol no đơn chức có gốc alkyl là -C~m~H~2m+1~ (m ≥ 1).\nCông thức ester no, đơn chức, mạch hở là C~n~H~2n+1~COOC~m~H~2m+1~ (hay C~x~H~2x~O~2~ với x = n + m + 1 ≥ 2)."
        },
        {
            "num": 4,
            "q": "Câu 4. Chất X có công thức cấu tạo CH~3~CH~2~COOCH~3~ là một chất lỏng trong suốt với mùi trái cây và mùi thơm của rum. Tên gọi của X là",
            "opts": ["A. propyl acetate.", "*B. methyl propionate.*", "C. methyl acetate.", "D. ethyl acetate."],
            "sol": "Gốc alcohol là methyl (-CH~3~), gốc acid là propionate (CH~3~CH~2~COO-).\nTên gọi của ester là methyl propionate (hoặc methyl propanoate)."
        },
        {
            "num": 5,
            "q": "Câu 5 (SBT - KNTT). Thuỷ phân methyl propanoate trong môi trường acid thu được sản phẩm gồm:",
            "opts": ["A. CH~3~COOH và C~3~H~7~OH.", "B. CH~3~COOH và C~2~H~5~OH.", "C. C~2~H~5~COOH và C~2~H~5~OH.", "*D. C~2~H~5~COOH và CH~3~OH.*"],
            "sol": "Methyl propanoate có công thức là C~2~H~5~COOCH~3~.\nPhản ứng thuỷ phân trong môi trường acid:\nC~2~H~5~COOCH~3~ + H~2~O ⇌ C~2~H~5~COOH + CH~3~OH."
        },
        {
            "num": 6,
            "q": "Câu 6. Công thức nào sau đây có thể là công thức của chất béo?",
            "opts": ["A. CH~3~COOCH~2~C~6~H~5~.", "B. C~15~H~31~COOCH~3~.", "C. (C~17~H~33~COO)~2~C~2~H~4~.", "*D. (C~17~H~35~COO)~3~C~3~H~5~.*"],
            "sol": "Chất béo là triester của glycerol với các acid béo (triglyceride), có dạng tổng quát là (RCOO)~3~C~3~H~5~.\nDo đó (C~17~H~35~COO)~3~C~3~H~5~ (tristearin) là chất béo."
        },
        {
            "num": 7,
            "q": "Câu 7. Acid nào sau đây là acid béo?",
            "opts": ["A. formic acid.", "*B. oleic acid.*", "C. acrylic acid.", "D. acetic acid."],
            "sol": "Acid béo là các monocarboxylic acid có mạch carbon dài, không phân nhánh, thường có số nguyên tử carbon chẵn (từ 12 đến 24 C).\nOleic acid có công thức C~17~H~33~COOH (18C, có 1 liên kết đôi cis) là acid béo tiêu biểu."
        },
        {
            "num": 8,
            "q": "Câu 8. Chất không phải acid béo là",
            "opts": ["*A. acetic acid.*", "B. palmitic acid.", "C. stearic acid.", "D. oleic acid."],
            "sol": "Acetic acid (CH~3~COOH) chỉ có 2 nguyên tử carbon trong phân tử nên không phải là acid béo.\nCác chất palmitic acid (C~15~H~31~COOH), stearic acid (C~17~H~35~COOH), oleic acid (C~17~H~33~COOH) đều là acid béo."
        },
        {
            "num": 9,
            "q": "Câu 9. Hỗn hợp các muối sodium của acid béo sinh ra ở trạng thái keo. Để tách muối này ra khỏi hỗn hợp, người ta thêm chất gì vào hỗn hợp?",
            "opts": ["A. tinh bột.", "*B. muối ăn.*", "C. nước.", "D. acid."],
            "sol": "Muối sodium của acid béo (xà phòng) tan một phần tạo hệ keo trong dung dịch.\nKhi thêm dung dịch NaCl bão hoà (muối ăn) nóng vào hỗn hợp sau phản ứng xà phòng hoá, do khối lượng riêng của xà phòng nhỏ hơn và độ tan giảm mạnh trong môi trường có nồng độ Na^+^ cao (hiệu ứng muối kết / ion chung), xà phòng sẽ đông tụ và nổi lên trên bề mặt, giúp tách ra dễ dàng."
        },
        {
            "num": 10,
            "q": "Câu 10 (SBT - CD). Ester có công thức phân tử là C~2~H~4~O~2~ được tạo thành từ methyl alcohol và carboxylic acid nào sau đây?",
            "opts": ["A. propionic acid.", "B. acetic acid.", "*C. formic acid.*", "D. oxalic acid."],
            "sol": "Ester tạo từ methyl alcohol (CH~3~OH) có công thức dạng RCOOCH~3~.\nVới công thức phân tử C~2~H~4~O~2~, ta suy ra R là H-.\nCông thức cấu tạo là HCOOCH~3~ (methyl formate), được tạo thành từ formic acid (HCOOH) và methyl alcohol."
        },
        {
            "num": 11,
            "q": "Câu 11. Hợp chất nào dưới đây được sử dụng làm xà phòng?",
            "opts": ["A. CH~3~COONa.", "*B. CH~3~(CH~2~)~14~COONa.*", "C. CH~3~(CH~2~)~12~COOCH~3~.", "D. CH~3~(CH~2~)~5~O(CH~2~)~5~CH~3~."],
            "sol": "Xà phòng là muối sodium hoặc potassium của acid béo (mạch hydrocarbon dài).\nCH~3~(CH~2~)~14~COONa chính là sodium palmitate (C~15~H~31~COONa), có tính chất hoạt động bề mặt và được sử dụng làm thành phần chính của xà phòng."
        },
        {
            "num": 12,
            "q": "Câu 12. Phản ứng xà phòng hóa là phản ứng xảy ra giữa",
            "opts": ["*A. chất béo với base.*", "B. chất béo với sulfuric acid đặc.", "C. chất béo với acid béo.", "D. chất béo với acid hữu cơ."],
            "sol": "Phản ứng xà phòng hóa là phản ứng thuỷ phân chất béo (hoặc ester) trong môi trường kiềm (dung dịch base như NaOH, KOH đun nóng), thu được xà phòng và glycerol."
        },
        {
            "num": 13,
            "q": "Câu 13. Chất nào cùng thành phần nguyên tố với chất béo?",
            "opts": ["A. dầu mỏ.", "*B. tinh dầu sả.*", "C. ethylene.", "D. dầu bôi trơn."],
            "sol": "Chất béo chứa 3 nguyên tố: C, H, O.\n- Dầu mỏ, dầu bôi trơn là hỗn hợp các hydrocarbon (chỉ gồm C, H).\n- Ethylene (C~2~H~4~) là hydrocarbon (chỉ gồm C, H).\n- Tinh dầu sả chứa các hợp chất terpene mang nhóm chức chứa oxygen (geraniol, citronellal, citral có thành phần C, H, O), do đó có cùng thành phần nguyên tố với chất béo."
        },
        {
            "num": 14,
            "q": "Câu 14. Nhận định nào sau đây đúng?",
            "opts": ["*A. Chất béo là triester của glycerol với acid béo.*", "B. Xà phòng là muối của acid vô cơ.", "C. Mỡ động vật, dầu thực vật thường tan tốt trong nước.", "D. Dung dịch xà phòng có môi trường acid."],
            "sol": "Theo định nghĩa chuẩn IUPAC: Chất béo là triester của glycerol với các acid béo (triglyceride).\n- Xà phòng là muối của acid béo (acid hữu cơ).\n- Mỡ động vật, dầu thực vật nhẹ hơn nước và không tan trong nước.\n- Dung dịch xà phòng có môi trường kiềm yếu do thủy phân anion acid béo (RCOO^- + H~2~O ⇌ RCOOH + OH^-)."
        },
        {
            "num": 15,
            "q": "Câu 15. Cho sơ đồ: X + Y ➔ (t°) Xà phòng + glycerol. Các chất X, Y có thể lần lượt là",
            "opts": ["A. alcohol và base.", "*B. ester và base.*", "C. base và acid.", "D. ester và alcohol."],
            "sol": "Chất béo là triester của glycerol với acid béo. Khi đun nóng chất béo (ester) với dung dịch kiềm (base) thì sinh ra xà phòng và glycerol.\nDo đó X là ester (chất béo) và Y là dung dịch base."
        },
        {
            "num": 16,
            "q": "Câu 16. Xét phân tử xà phòng có cấu tạo như dưới đây:",
            "img": os.path.join(IMG_DIR, "cau16_phan_tu_xa_phong.png"),
            "img_width": 14.5,
            "q_sub": "Nhóm được khoanh tròn trong công thức trên là",
            "opts": ["A. nhóm kị nước.", "B. nhóm dị thể.", "*C. nhóm ưa nước.*", "D. nhóm đồng thể."],
            "sol": "Phân tử xà phòng gồm hai phần:\n- Phần đuôi hydrocarbon mạch dài không phân cực: phần kị nước (ưa dầu mỡ).\n- Phần đầu carboxylate phân cực mang điện tích âm (-COO^-Na^+): phần ưa nước (dễ tương tác với các phân tử nước).\nNhóm được khoanh tròn đỏ ở đầu phân tử là nhóm -COO^-Na^+, tức là nhóm ưa nước."
        },
        {
            "num": 17,
            "q": "Câu 17. Thể tích H~2~ (đkc) cần để hydrogen hóa hoàn toàn 1 tấn triolein là",
            "opts": ["*A. 84128,96 L.*", "B. 760,18 lít.", "C. 7,6018 lít.", "D. 7601,8 lít."],
            "sol": "Khối lượng mol của triolein (C~17~H~33~COO)~3~C~3~H~5~ là:\nM = 17 × 3 + 33 × 3 + 44 × 3 + 41 = 884 g/mol.\n1 tấn = 10^6 gam ➔ n~triolein~ = 10^6 / 884 ≈ 1131,2217 mol.\nPhương trình phản ứng hydrogen hoá:\n(C~17~H~33~COO)~3~C~3~H~5~ + 3H~2~ ➔ (Ni, t°) (C~17~H~35~COO)~3~C~3~H~5~.\nn~H2~ = 3 × n~triolein~ = 3 × 1131,2217 ≈ 3393,665 mol.\nỞ điều kiện chuẩn (đkc, 25°C, 1 bar):\nV~H2~ = 3393,665 × 24,79 ≈ 84128,96 lít."
        },
        {
            "num": 18,
            "q": "Câu 18. Cho 1 mol triglyceride X tác dụng với dung dịch NaOH dư, thu được 1 mol glycerol, 1 mol sodium palmitate và 2 mol sodium oleate. Phát biểu nào sau đây sai?",
            "opts": ["A. Phân tử X có 5 liên kết π.", "B. Có 2 đồng phân cấu tạo thỏa mãn tính chất của X.", "*C. Công thức phân tử chất X là C~52~H~96~O~6~.*", "D. 1 mol X làm mất màu tối đa 2 mol Br~2~ trong dung dịch."],
            "sol": "Thủy phân X thu được 1 gốc palmitate (C~15~H~31~COO-) và 2 gốc oleate (C~17~H~33~COO-).\n- Công thức phân tử của X:\nSố C = 16 + 2 × 18 + 3 = 55.\nSố H = 31 + 2 × 33 + 5 = 102.\nSố O = 6.\n➔ Công thức phân tử của X là C~55~H~102~O~6~. Phát biểu C ghi C~52~H~96~O~6~ là SAI.\n- Liên kết π: 3 liên kết π trong 3 nhóm -COO- và 2 liên kết π C=C trong 2 gốc oleic, tổng cộng có 3 + 2 = 5 liên kết π ➔ A đúng.\n- Đồng phân: Gốc C~15~H~31~COO- có thể nằm ở vị trí α hoặc β của mạch glycerol ➔ có 2 đồng phân cấu tạo ➔ B đúng.\n- Có 2 nối đôi C=C nên 1 mol X phản ứng tối đa với 2 mol Br~2~ ➔ D đúng."
        }
    ]

    for q_data in part1_questions:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1)
        render_question_prompt(p, q_data["q"])

        if "img" in q_data:
            p_img = doc.add_paragraph()
            p_img.paragraph_format.space_before = Pt(2)
            p_img.paragraph_format.space_after = Pt(2)
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.add_run().add_picture(q_data["img"], width=Cm(q_data.get("img_width", 12)))
            if "q_sub" in q_data:
                p_sub = doc.add_paragraph()
                p_sub.paragraph_format.space_before = Pt(1)
                p_sub.paragraph_format.space_after = Pt(1)
                render_rich_text(p_sub, q_data["q_sub"])

        render_options_no_table(doc, q_data["opts"], is_teacher=is_teacher)

        if is_teacher and "sol" in q_data:
            add_solution_no_table(doc, q_data["sol"])

    # =========================================================================
    # PHẦN II
    # =========================================================================
    add_section_header(doc, "PHẦN II. Câu trắc nghiệm đúng sai.",
                       "Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.")

    part2_questions = [
        {
            "num": 1,
            "prompt": "Câu 1 (SBT - KNTT). Cho các phát biểu về acid béo và chất béo sau:",
            "stmts": [
                ("a) Acid béo thường có gốc hydrocarbon mạch dài, có số nguyên tử carbon lẻ.", "Sai",
                 "Acid béo trong tự nhiên hầu hết có mạch carbon dài không phân nhánh và có số nguyên tử carbon chẵn (từ 12 đến 24 carbon)."),
                ("b) Khi xà phòng hoá chất béo, sản phẩm thu được là glycerol và các acid đơn chức.", "Sai",
                 "Sản phẩm của phản ứng xà phòng hoá chất béo là glycerol và muối của acid béo (xà phòng), không phải acid tự do."),
                ("c) Chất béo có nhiều gốc acid no thường ở dạng rắn, còn chất béo có nhiều gốc acid không no thường ở dạng lỏng.", "Đúng",
                 "Chất béo chứa nhiều gốc acid béo no (như mỡ động vật) có nhiệt độ nóng chảy cao nên ở trạng thái rắn; chất béo chứa nhiều gốc acid béo không no (như dầu thực vật) có nhiệt độ nóng chảy thấp nên ở trạng thái lỏng."),
                ("d) Acid béo omega-3 là các acid béo không no có liên kết đôi đầu tiên ở vị trí số 3 nếu đánh số từ nhóm carboxyl (-COOH).", "Sai",
                 "Acid béo omega-3 được đánh số vị trí từ nguyên tử carbon của nhóm methyl (-CH~3~) ở cuối mạch (carbon ω), chứ không phải từ nhóm carboxyl (-COOH).")
            ]
        },
        {
            "num": 2,
            "prompt": "Câu 2. Người ta thực hiện phản ứng xà phòng hóa như hình vẽ dưới đây:",
            "img": os.path.join(IMG_DIR, "cau2_thi_nghiem_xa_phong_hoa.png"),
            "img_width": 10.5,
            "post_img": "Sau thí nghiệm, để nguội hỗn hợp trong cốc và quan sát.",
            "stmts": [
                ("a) Khi đã nguội, hỗn hợp trong cốc tách thành 2 lớp.", "Đúng",
                 "Sau khi thêm NaCl bão hòa và để nguội, xà phòng không tan trong dung dịch NaCl và có khối lượng riêng nhỏ hơn nên đông tụ nổi lên trên tạo thành lớp chất rắn màu trắng; lớp dưới là dung dịch chứa glycerol, NaCl và một ít kiềm dư."),
                ("b) Dung dịch NaCl bão hòa, nóng là xúc tác cho phản ứng xà phòng hóa.", "Sai",
                 "Dung dịch NaCl bão hòa không phải chất xúc tác; nó có tác dụng làm tăng khối lượng riêng của dung dịch và làm giảm độ tan của muối acid béo (hiệu ứng muối kết) giúp xà phòng tách lớp hoàn toàn."),
                ("c) Khi đã nguội, có một lớp glycerol nổi lên trên bề mặt hỗn hợp.", "Sai",
                 "Lớp chất rắn nổi lên trên bề mặt là xà phòng (muối của acid béo); còn glycerol tan rất tốt trong nước nên nằm ở lớp dung dịch phía dưới."),
                ("d) Phản ứng xảy ra trong thí nghiệm là phản ứng thủy phân ester trong môi trường kiềm.", "Đúng",
                 "Chất béo là triester của glycerol với các acid béo; phản ứng với dung dịch NaOH là phản ứng thủy phân ester trong môi trường kiềm (phản ứng xà phòng hóa).")
            ]
        },
        {
            "num": 3,
            "prompt": "Câu 3 (SBT - KNTT). Xét các phát biểu về xà phòng và chất giặt rửa:",
            "stmts": [
                ("a) Chất kị nước là những chất không tan trong dầu mỡ, hydrocarbon,...", "Sai",
                 "Chất kị nước là chất không tan trong nước nhưng tan tốt trong dầu mỡ và các dung môi hydrocarbon (tức chất ưa dầu)."),
                ("b) Muối sodium hoặc potassium của carboxylic acid được dùng làm xà phòng.", "Sai",
                 "Phải là muối sodium hoặc potassium của acid béo (mạch dài không phân nhánh từ 12-24 C) mới có khả năng tạo bọt và tẩy rửa làm xà phòng; các muối của carboxylic acid mạch ngắn như CH~3~COONa không làm xà phòng được."),
                ("c) Xà phòng bị giảm hoặc mất tác dụng tẩy rửa khi dùng nước cứng vì gốc acid béo dễ kết tủa với cation Ca^2+^, Mg^2+^.", "Đúng",
                 "Các anion acid béo kết hợp với ion Ca^2+^, Mg^2+^ tạo thành kết tủa không tan như (C~17~H~35~COO)~2~Ca↓, làm giảm lượng chất hoạt động bề mặt và đóng cặn trên vải."),
                ("d) Xà phòng, chất giặt rửa tan vào nước tạo dung dịch có sức căng bề mặt nhỏ làm cho vật cần giặt rửa dễ thấm ướt.", "Đúng",
                 "Chất hoạt động bề mặt làm giảm sức căng bề mặt tại mặt phân cách giữa nước và không khí/vải, làm dung dịch dễ len lỏi thấm ướt bề mặt sợi vải và đánh bật vết bẩn.")
            ]
        },
        {
            "num": 4,
            "prompt": "Câu 4 (SBT - KNTT). Có bốn ống nghiệm: ống (1) chứa 3 mL nước cất; ống (2) chứa 3 mL nước xà phòng; ống (3) chứa 3 mL nước xà phòng và 3 giọt dung dịch CaCl~2~ bão hoà; ống (4) chứa 3 mL nước giặt rửa tổng hợp và 3 giọt dung dịch CaCl~2~ bão hoà. Cho vào mỗi ống nghiệm 3 giọt dầu ăn, lắc đều.",
            "stmts": [
                ("a) Trong ống nghiệm (1), dầu ăn không tan và chìm xuống dưới.", "Sai",
                 "Dầu ăn không tan trong nước và có khối lượng riêng nhỏ hơn nước (d ≈ 0,9 g/mL) nên nổi lên trên mặt nước cất chứ không chìm."),
                ("b) Trong ống nghiệm (2), dầu ăn tan tạo hỗn hợp đồng nhất.", "Đúng",
                 "Nước xà phòng hoạt động bề mặt giúp nhũ hóa dầu ăn, tạo thành hệ nhũ tương bền, phân tán đồng nhất."),
                ("c) Trong ống nghiệm (3) có kết tủa xuất hiện.", "Đúng",
                 "Ống (3) có nước xà phòng phản ứng với Ca^2+^ trong CaCl~2~ sinh ra kết tủa calcium carboxylate không tan: 2RCOO^- + Ca^2+^ ➔ (RCOO)~2~Ca↓."),
                ("d) Trong ống nghiệm (4), dầu ăn tan tạo hỗn hợp đồng nhất.", "Đúng",
                 "Chất giặt rửa tổng hợp không tạo kết tủa với ion Ca^2+^ của nước cứng, do đó vẫn giữ nguyên khả năng nhũ hóa dầu ăn, tạo hỗn hợp phân tán đồng nhất.")
            ]
        }
    ]

    for p2_data in part2_questions:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1)
        render_question_prompt(p, p2_data["prompt"])

        if "img" in p2_data:
            p_img = doc.add_paragraph()
            p_img.paragraph_format.space_before = Pt(2)
            p_img.paragraph_format.space_after = Pt(2)
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.add_run().add_picture(p2_data["img"], width=Cm(p2_data.get("img_width", 10.5)))

        if "post_img" in p2_data:
            p_post = doc.add_paragraph()
            p_post.paragraph_format.space_before = Pt(1)
            p_post.paragraph_format.space_after = Pt(2)
            render_rich_text(p_post, p2_data["post_img"], base_italic=True)

        for stmt, ans, sol in p2_data["stmts"]:
            p_stmt = doc.add_paragraph()
            p_stmt.paragraph_format.left_indent = Cm(0.5)
            p_stmt.paragraph_format.space_before = Pt(1)
            p_stmt.paragraph_format.space_after = Pt(1)
            render_statement(p_stmt, stmt, ans=ans, is_teacher=is_teacher)

        if is_teacher:
            sol_full = "\n".join([f"• *{s[:2]}* là *{a}*: {sol}" for s, a, sol in p2_data["stmts"]])
            add_solution_no_table(doc, sol_full)

    # =========================================================================
    # PHẦN III
    # =========================================================================
    add_section_header(doc, "PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn.",
                       "Thí sinh trả lời từ câu 1 đến câu 6.")

    part3_questions = [
        {
            "num": 1,
            "q": "Câu 1. Cho các chất sau: C~15~H~31~COONa; C~15~H~31~COOK; CH~3~[CH~2~]~11~OSO~3~Na; CH~3~[CH~2~]~11~C~6~H~4~SO~3~Na; C~17~H~33~COOK. Có bao nhiêu chất là thành phần chính của chất giặt rửa tổng hợp?",
            "ans": "2",
            "sol": "• C~15~H~31~COONa, C~15~H~31~COOK, C~17~H~33~COOK là muối sodium/potassium của acid béo ➔ thành phần chính của xà phòng.\n• CH~3~[CH~2~]~11~OSO~3~Na (sodium lauryl sulfate - muối alkyl sulfate) và CH~3~[CH~2~]~11~C~6~H~4~SO~3~Na (sodium dodecylbenzenesulfonate - muối alkylbenzenesulfonate) là thành phần chính của chất giặt rửa tổng hợp (synthetic detergents).\n➔ Có 2 chất thỏa mãn."
        },
        {
            "num": 2,
            "q": "Câu 2 (SBT - CTST). Thuỷ phân hoàn toàn triglyceride X thu được glycerol và 3 acid béo là lauric acid, palmitic acid và oleic acid. X có thể có bao nhiêu công thức cấu tạo?",
            "ans": "3",
            "sol": "Khung phân tử glycerol có 3 nguyên tử carbon liên kết với nhau: C(α) - C(β) - C(α'), trong đó 2 vị trí α và α' đối xứng nhau qua vị trí β ở giữa.\nKhi gắn 3 gốc acid béo khác nhau (kí hiệu là A, B, C) vào khung glycerol:\n• Gốc A ở giữa (vị trí β): B-α, A-β, C-α'\n• Gốc B ở giữa (vị trí β): A-α, B-β, C-α'\n• Gốc C ở giữa (vị trí β): A-α, C-β, B-α'\n➔ Có đúng 3 đồng phân cấu tạo thỏa mãn."
        },
        {
            "num": 3,
            "q": "Câu 3 (SBT - CTST). Có bao nhiêu hợp chất hữu cơ đơn chức, phân tử chỉ chứa carbon, hydrogen, oxygen và có phân tử khối là 60?",
            "ans": "5",
            "sol": "Hợp chất đơn chức chứa C, H, O có phân tử khối M = 60:\n1. Phân tử có 1 nguyên tử O: C~x~H~y~O có 12x + y + 16 = 60 ➔ 12x + y = 44 ➔ x = 3, y = 8 (C~3~H~8~O). Các đồng phân đơn chức gồm alcohol và ether:\n   (1) CH~3~CH~2~CH~2~OH (propan-1-ol)\n   (2) CH~3~CH(OH)CH~3~ (propan-2-ol)\n   (3) CH~3~CH~2~OCH~3~ (ethyl methyl ether)\n2. Phân tử có 2 nguyên tử O: C~x~H~y~O~2~ có 12x + y + 32 = 60 ➔ 12x + y = 28 ➔ x = 2, y = 4 (C~2~H~4~O~2~). Các đồng phân đơn chức gồm carboxylic acid và ester:\n   (4) CH~3~COOH (acetic acid)\n   (5) HCOOCH~3~ (methyl formate)\n➔ Tổng cộng có 5 hợp chất hữu cơ đơn chức."
        },
        {
            "num": 4,
            "q": "Câu 4 (SBT - KNTT). Chỉ số xà phòng hoá là số miligam KOH dùng để xà phòng hoá hết triglyceride và trung hoà hết lượng acid tự do trong 1 g chất béo. Một chất béo chứa 3,55% stearic acid và 89% tristearin về khối lượng, còn lại là các chất không tham gia phản ứng với KOH. Tính chỉ số xà phòng hoá của chất béo trên.",
            "ans": "175",
            "sol": "Xét trong 1 gam chất béo:\n• Khối lượng stearic acid: m~stearic~ = 0,0355 g ➔ n~stearic~ = 0,0355 / 284 = 1,25 × 10^-4^ mol.\n• Khối lượng tristearin: m~tristearin~ = 0,89 g ➔ n~tristearin~ = 0,89 / 890 = 0,001 mol.\nLượng KOH cần dùng:\nn~KOH~ = n~stearic~ + 3 × n~tristearin~ = 1,25 × 10^-4^ + 3 × 0,001 = 0,003125 mol.\nKhối lượng KOH cần dùng:\nm~KOH~ = 0,003125 × 56 = 0,175 g = 175 mg.\n➔ Chỉ số xà phòng hoá của chất béo là 175."
        },
        {
            "num": 5,
            "q": "Câu 5 (SBT - CTST). Xà phòng hoá hoàn toàn 0,1 mol ester E đơn chức, mạch hở bằng dung dịch NaOH vừa đủ thu được muối F và alcohol G. Biết khối lượng muối F thu được lớn hơn khối lượng ester E đã tham gia xà phòng hoá. Tính khối lượng alcohol G thu được (theo gam).",
            "ans": "3,2",
            "sol": "Đặt công thức của ester đơn chức mạch hở E là RCOOR'.\nPhương trình phản ứng: RCOOR' + NaOH ➔ RCOONa + R'OH.\nTheo đề bài: m~muối~ > m~ester~ ➔ M~RCOONa~ > M~RCOOR'~ ➔ 23 > M~R'~.\nVì R' là gốc hydrocarbon của alcohol nên điều kiện M~R'~ < 23 chỉ có một trường hợp duy nhất là gốc methyl: -CH~3~ (M = 15).\nAlcohol G thu được là methanol (CH~3~OH).\nKhối lượng alcohol G: m~CH3OH~ = 0,1 × 32 = 3,2 gam."
        },
        {
            "num": 6,
            "q": "Câu 6. Chỉ số xà phòng hóa (Saponification value, viết tắt là SAP): Lượng KOH cần thiết (tính bằng miligam) để phản ứng hết với 1 gam dầu, mỡ theo phản ứng xà phòng hóa và trung hòa hết carboxylic acid tự do có trong dầu, mỡ.\nTrong kĩ thuật, để tính lượng NaOH dùng để xà phòng hóa người ta xác định hệ số a. Công thức như sau: a = SAP / 1402,5; m~NaOH~ = a × m~chất béo~ (gam).\nBiết loại dầu dừa có giá trị SAP là 255. Tính khối lượng của NaOH cần dùng để xà phòng hóa hết 200 gam dầu dừa. (Kết quả làm tròn đến hai chữ số thập phân).",
            "ans": "36,36",
            "sol": "Hệ số a được tính theo công thức:\na = SAP / 1402,5 = 255 / 1402,5 ≈ 0,181818.\nKhối lượng NaOH cần dùng để xà phòng hóa 200 gam dầu dừa:\nm~NaOH~ = a × m~dầu dừa~ = (255 / 1402,5) × 200 ≈ 36,3636 gam ≈ 36,36 gam."
        }
    ]

    for p3_data in part3_questions:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1)
        render_question_prompt(p, p3_data["q"])

        if is_teacher:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Cm(0.4)
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(1)
            render_rich_text(p_ans, f"► *Đáp án:* *{p3_data['ans']}*", base_size=10.5, base_bold=True, base_color=COLOR_CORRECT)
            add_solution_no_table(doc, p3_data["sol"])
        else:
            p_box = doc.add_paragraph()
            p_box.paragraph_format.left_indent = Cm(0.5)
            p_box.paragraph_format.space_before = Pt(1)
            p_box.paragraph_format.space_after = Pt(3)
            r_box = p_box.add_run("Đáp số: ............................................................ ")
            set_run_style(r_box, font_size=10, italic=True, color=COLOR_GRAY)

    # Dòng kết thúc đề
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(10)
    p_end.paragraph_format.space_after = Pt(10)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("---------- HẾT ----------")
    set_run_style(r_end, font_size=11, bold=True, italic=True)

    out_filename = "De_Kiem_Tra_Hoa_12_De02_GiaoVien.docx" if is_teacher else "De_Kiem_Tra_Hoa_12_De02.docx"
    out_path = os.path.join(OUT_DIR, out_filename)
    doc.save(out_path)
    print(f"Saved: {out_path}")

if __name__ == "__main__":
    os.makedirs(OUT_DIR, exist_ok=True)
    build_exam_de02(is_teacher=False)
    build_exam_de02(is_teacher=True)
