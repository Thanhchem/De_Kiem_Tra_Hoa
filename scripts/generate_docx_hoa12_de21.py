# -*- coding: utf-8 -*-
"""
Script tạo 2 file Word (.docx) chuẩn Form Đề Kiểm Tra cho:
ĐỀ SỐ 21 - CHƯƠNG 7: KIM LOẠI NHÓM IA & NHÓM IIA (HÓA HỌC 12)
Mã đề thi: 121
Bao gồm:
1. Bản Học sinh: De_Kiem_Tra_Hoa_12_De21.docx
2. Bản Giáo viên: De_Kiem_Tra_Hoa_12_De21_GiaoVien.docx (Có đáp án đỏ, lời giải chi tiết xanh, bảng đáp án)
"""

import os
import sys
sys.stdout.reconfigure(encoding='utf-8')
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"

COLOR_CORAL_RED = RGBColor(211, 47, 47)      # #D32F2F
COLOR_SAGE_GREEN = "8BC390"                 # #8BC390
COLOR_DEEP_BLUE = RGBColor(21, 101, 192)    # #1565C0
COLOR_DARK_GREEN = RGBColor(46, 125, 50)    # #2E7D32
COLOR_BLACK = RGBColor(0, 0, 0)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.59)     # 1.5 cm
        s.bottom_margin = Inches(0.59)  # 1.5 cm
        s.left_margin = Inches(0.59)    # 1.5 cm
        s.right_margin = Inches(0.59)   # 1.5 cm
        s.header_distance = Inches(0.24) # 0.6 cm
        s.footer_distance = Inches(0.24) # 0.6 cm

        # --- HEADER ---
        header = s.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.text = ""
        hp.paragraph_format.space_before = Pt(0)
        hp.paragraph_format.space_after = Pt(0)

        tbl_h = header.add_table(rows=1, cols=2, width=Inches(7.08)) # 18.0 cm
        tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
        c0, c1 = tbl_h.rows[0].cells
        c0.width = Inches(3.54)
        c1.width = Inches(3.54)
        set_cell_margins(c0, 60, 60, 0, 0)
        set_cell_margins(c1, 60, 60, 0, 0)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run("Tài liệu lưu hành nội bộ")
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(10)
        r0.font.italic = True
        r0.font.color.rgb = COLOR_CORAL_RED

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1_a = p1.add_run("Thầy ")
        r1_a.font.name = "Times New Roman"
        r1_a.font.size = Pt(10)
        r1_a.font.color.rgb = COLOR_CORAL_RED
        r1_b = p1.add_run("TRẦN VĂN THẠNH")
        r1_b.font.name = "Times New Roman"
        r1_b.font.size = Pt(10)
        r1_b.font.bold = True
        r1_b.font.color.rgb = COLOR_CORAL_RED
        r1_c = p1.add_run(" - 0777.470.803")
        r1_c.font.name = "Times New Roman"
        r1_c.font.size = Pt(10)
        r1_c.font.color.rgb = COLOR_CORAL_RED

        tblPr_h = tbl_h._tbl.tblPr
        tblBorders_h = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{COLOR_SAGE_GREEN}"/>'
            f'</w:tblBorders>'
        )
        tblPr_h.append(tblBorders_h)

        # --- FOOTER ---
        footer = s.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.text = ""
        fp.paragraph_format.space_before = Pt(0)
        fp.paragraph_format.space_after = Pt(0)

        tbl_f = footer.add_table(rows=1, cols=2, width=Inches(7.08))
        tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
        fc0, fc1 = tbl_f.rows[0].cells
        fc0.width = Inches(5.31)
        fc1.width = Inches(1.77)
        set_cell_margins(fc0, 60, 60, 0, 0)
        set_cell_margins(fc1, 60, 60, 0, 0)

        fp0 = fc0.paragraphs[0]
        fp0.paragraph_format.space_before = Pt(0)
        fp0.paragraph_format.space_after = Pt(0)
        fr0 = fp0.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
        fr0.font.name = "Times New Roman"
        fr0.font.size = Pt(9.5)
        fr0.font.italic = True
        fr0.font.color.rgb = COLOR_CORAL_RED

        fp1 = fc1.paragraphs[0]
        fp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp1.paragraph_format.space_before = Pt(0)
        fp1.paragraph_format.space_after = Pt(0)
        fr1_a = fp1.add_run("Trang ")
        fr1_a.font.name = "Times New Roman"
        fr1_a.font.size = Pt(10)
        fr1_a.font.bold = True
        fr1_a.font.color.rgb = COLOR_CORAL_RED

        fldSimple = OxmlElement('w:fldSimple')
        fldSimple.set(qn('w:instr'), 'PAGE')
        r_page = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        rFont = OxmlElement('w:rFonts')
        rFont.set(qn('w:ascii'), 'Times New Roman')
        rFont.set(qn('w:hAnsi'), 'Times New Roman')
        b = OxmlElement('w:b')
        c_el = OxmlElement('w:color')
        c_el.set(qn('w:val'), 'D32F2F')
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '20')
        rPr.append(rFont)
        rPr.append(b)
        rPr.append(c_el)
        rPr.append(sz)
        r_page.append(rPr)
        t_el = OxmlElement('w:t')
        t_el.text = "1"
        r_page.append(t_el)
        fldSimple.append(r_page)
        fp1._p.append(fldSimple)

        tblPr_f = tbl_f._tbl.tblPr
        tblBorders_f = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{COLOR_SAGE_GREEN}"/>'
            f'</w:tblBorders>'
        )
        tblPr_f.append(tblBorders_f)

def add_chem_run(paragraph, text, bold=False, italic=False, color=None, font_size=11):
    parts = re.split(r'(~.*?~|\^.*?\^)', text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run()
        run.font.name = "Times New Roman"
        run.font.size = Pt(font_size)
        run.bold = bold
        run.italic = italic
        if color:
            run.font.color.rgb = color

        if part.startswith('~') and part.endswith('~'):
            run.text = part[1:-1]
            run.font.subscript = True
        elif part.startswith('^') and part.endswith('^'):
            run.text = part[1:-1]
            run.font.superscript = True
        else:
            run.text = part

def create_title_block(doc, is_teacher=False):
    lines_info = [
        ("SỞ GIÁO DỤC VÀ ĐÀO TẠO THỪA THIÊN HUẾ", "ĐỀ KIỂM TRA ĐỊNH KỲ CHƯƠNG 7 (KIM LOẠI NHÓM IA & IIA)", True, True, COLOR_DARK_GREEN),
        ("TRƯỜNG THPT HAI BÀ TRƯNG", "MÔN: HOÁ HỌC -- LỚP 12", True, True, None),
        ("ĐỀ CHÍNH THỨC", "Thời gian làm bài: 50 phút (không kể thời gian phát đề)", True, False, None),
        ("(Đề thi có 04 trang)", "Mã đề thi: 121", False, True, None)
    ]

    for left_txt, right_txt, left_bold, right_bold, right_color in lines_info:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.08), WD_TAB_ALIGNMENT.RIGHT)

        r_l = p.add_run(left_txt)
        r_l.font.name = "Times New Roman"
        r_l.font.size = Pt(11)
        r_l.bold = left_bold
        if not left_bold:
            r_l.italic = True

        p.add_run("\t")

        r_r = p.add_run(right_txt)
        r_r.font.name = "Times New Roman"
        r_r.font.size = Pt(11)
        r_r.bold = right_bold
        if right_color:
            r_r.font.color.rgb = right_color
        if not right_bold:
            r_r.italic = True

    if is_teacher:
        p_note = doc.add_paragraph()
        p_note.paragraph_format.space_before = Pt(4)
        p_note.paragraph_format.space_after = Pt(4)
        p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_n = p_note.add_run("(BẢN GIÁO VIÊN -- CÓ LỜI GIẢI CHI TIẾT)")
        r_n.font.name = "Times New Roman"
        r_n.font.size = Pt(12)
        r_n.font.bold = True
        r_n.font.color.rgb = COLOR_CORAL_RED

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(3)
    p_info.paragraph_format.space_after = Pt(2)
    r_info = p_info.add_run("Họ, tên thí sinh: ..................................................................... Lớp: ........................ SBD: ............................")
    r_info.font.name = "Times New Roman"
    r_info.font.size = Pt(10.5)
    r_info.font.italic = True

    p_mass = doc.add_paragraph()
    p_mass.paragraph_format.space_before = Pt(0)
    p_mass.paragraph_format.space_after = Pt(5)
    r_mass = p_mass.add_run("(Cho biết nguyên tử khối: H = 1; C = 12; N = 14; O = 16; Na = 23; Mg = 24; Al = 27; P = 31; S = 32; Cl = 35,5; K = 39; Ca = 40; Fe = 56; Cu = 64; Zn = 65; Br = 80; Rb = 85,5; Ag = 108; Cs = 133; Ba = 137).")
    r_mass.font.name = "Times New Roman"
    r_mass.font.size = Pt(9.5)
    r_mass.font.italic = True

def add_part_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = COLOR_DARK_GREEN

def add_explanation(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.line_spacing = 1.15

    r_lbl = p.add_run("Lời giải chi tiết: ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(10)
    r_lbl.font.bold = True
    r_lbl.font.italic = True
    r_lbl.font.color.rgb = COLOR_DEEP_BLUE

    add_chem_run(p, text, font_size=10, italic=False, color=COLOR_DEEP_BLUE)

def render_mcq(doc, q_num, text, options, correct_key, explanation, is_teacher=False, img_path=None, img_width=Inches(3.5)):
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(3)
    p_q.paragraph_format.space_after = Pt(2)
    p_q.paragraph_format.line_spacing = 1.15

    r_lbl = p_q.add_run(f"Câu {q_num}: ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(11)
    r_lbl.bold = True

    add_chem_run(p_q, text, font_size=11)

    if img_path and os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(2)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture(img_path, width=img_width)

    max_len = max(len(opt[1]) for opt in options)
    if max_len <= 15:
        p_opts = doc.add_paragraph()
        p_opts.paragraph_format.space_before = Pt(0)
        p_opts.paragraph_format.space_after = Pt(2)
        p_opts.paragraph_format.line_spacing = 1.15
        p_opts.paragraph_format.tab_stops.add_tab_stop(Inches(1.77), WD_TAB_ALIGNMENT.LEFT)
        p_opts.paragraph_format.tab_stops.add_tab_stop(Inches(3.54), WD_TAB_ALIGNMENT.LEFT)
        p_opts.paragraph_format.tab_stops.add_tab_stop(Inches(5.31), WD_TAB_ALIGNMENT.LEFT)

        for i, (k, txt) in enumerate(options):
            is_c = (k == correct_key) and is_teacher
            r_k = p_opts.add_run(f"{k}. ")
            r_k.font.name = "Times New Roman"
            r_k.font.size = Pt(11)
            r_k.font.bold = is_c or True
            if is_c:
                r_k.font.color.rgb = COLOR_CORAL_RED
            add_chem_run(p_opts, txt, bold=is_c, color=COLOR_CORAL_RED if is_c else None, font_size=11)
            if i < 3:
                p_opts.add_run("\t")

    elif max_len <= 35:
        for row_opts in [options[:2], options[2:]]:
            p_opts = doc.add_paragraph()
            p_opts.paragraph_format.space_before = Pt(0)
            p_opts.paragraph_format.space_after = Pt(2)
            p_opts.paragraph_format.line_spacing = 1.15
            p_opts.paragraph_format.tab_stops.add_tab_stop(Inches(3.54), WD_TAB_ALIGNMENT.LEFT)

            for i, (k, txt) in enumerate(row_opts):
                is_c = (k == correct_key) and is_teacher
                r_k = p_opts.add_run(f"{k}. ")
                r_k.font.name = "Times New Roman"
                r_k.font.size = Pt(11)
                r_k.font.bold = is_c or True
                if is_c:
                    r_k.font.color.rgb = COLOR_CORAL_RED
                add_chem_run(p_opts, txt, bold=is_c, color=COLOR_CORAL_RED if is_c else None, font_size=11)
                if i == 0:
                    p_opts.add_run("\t")
    else:
        for k, txt in options:
            p_opts = doc.add_paragraph()
            p_opts.paragraph_format.space_before = Pt(0)
            p_opts.paragraph_format.space_after = Pt(2)
            p_opts.paragraph_format.line_spacing = 1.15
            is_c = (k == correct_key) and is_teacher
            r_k = p_opts.add_run(f"{k}. ")
            r_k.font.name = "Times New Roman"
            r_k.font.size = Pt(11)
            r_k.font.bold = is_c or True
            if is_c:
                r_k.font.color.rgb = COLOR_CORAL_RED
            add_chem_run(p_opts, txt, bold=is_c, color=COLOR_CORAL_RED if is_c else None, font_size=11)

    if is_teacher:
        add_explanation(doc, explanation)

def render_tf(doc, q_num, text, sub_items, explanation, is_teacher=False, img_path=None, img_width=Inches(3.5)):
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(3)
    p_q.paragraph_format.space_after = Pt(2)
    p_q.paragraph_format.line_spacing = 1.15

    r_lbl = p_q.add_run(f"Câu {q_num}: ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(11)
    r_lbl.bold = True

    add_chem_run(p_q, text, font_size=11)

    if img_path and os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(2)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture(img_path, width=img_width)

    for sub_k, sub_txt, is_true in sub_items:
        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(1.5)
        p_sub.paragraph_format.left_indent = Inches(0.2)
        p_sub.paragraph_format.line_spacing = 1.15

        r_sk = p_sub.add_run(f"{sub_k}) ")
        r_sk.font.name = "Times New Roman"
        r_sk.font.size = Pt(11)
        r_sk.bold = True

        add_chem_run(p_sub, sub_txt, font_size=11)

        if is_teacher:
            r_ans = p_sub.add_run(f" --> {'ĐÚNG' if is_true else 'SAI'}")
            r_ans.font.name = "Times New Roman"
            r_ans.font.size = Pt(10.5)
            r_ans.font.bold = True
            r_ans.font.color.rgb = COLOR_CORAL_RED

    if is_teacher:
        add_explanation(doc, explanation)

def render_short(doc, q_num, text, answer, explanation, is_teacher=False):
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(3)
    p_q.paragraph_format.space_after = Pt(2)
    p_q.paragraph_format.line_spacing = 1.15

    r_lbl = p_q.add_run(f"Câu {q_num}: ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(11)
    r_lbl.bold = True

    add_chem_run(p_q, text, font_size=11)

    if is_teacher:
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(0)
        p_ans.paragraph_format.space_after = Pt(1)
        p_ans.paragraph_format.left_indent = Inches(0.25)
        r_ak = p_ans.add_run("Đáp số: ")
        r_ak.font.name = "Times New Roman"
        r_ak.font.size = Pt(10.5)
        r_ak.font.bold = True
        r_ak.font.color.rgb = COLOR_CORAL_RED
        r_av = p_ans.add_run(str(answer))
        r_av.font.name = "Times New Roman"
        r_av.font.size = Pt(10.5)
        r_av.font.bold = True
        r_av.font.color.rgb = COLOR_CORAL_RED

        add_explanation(doc, explanation)
    else:
        p_dot = doc.add_paragraph()
        p_dot.paragraph_format.space_before = Pt(0)
        p_dot.paragraph_format.space_after = Pt(2)
        r_d = p_dot.add_run("Trả lời: .....................................................................................................................................................")
        r_d.font.name = "Times New Roman"
        r_d.font.size = Pt(10)
        r_d.font.italic = True

def add_answer_table(doc):
    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_before = Pt(8)
    p_t.paragraph_format.space_after = Pt(4)
    r_t = p_t.add_run("BẢNG ĐÁP ÁN THAM KHẢO")
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(11)
    r_t.font.bold = True
    r_t.font.color.rgb = COLOR_CORAL_RED

    p_p1 = doc.add_paragraph()
    p_p1.paragraph_format.space_before = Pt(0)
    p_p1.paragraph_format.space_after = Pt(2)
    r1 = p_p1.add_run("PHẦN I: ")
    r1.font.bold = True
    r1.font.color.rgb = COLOR_DARK_GREEN
    p_p1.add_run("1.C   2.D   3.B   4.B   5.D   6.C   7.C   8.A   9.D   10.D   11.C   12.C   13.B   14.A   15.C   16.D   17.C   18.A")

    p_p2 = doc.add_paragraph()
    p_p2.paragraph_format.space_before = Pt(2)
    p_p2.paragraph_format.space_after = Pt(2)
    r2 = p_p2.add_run("PHẦN II: ")
    r2.font.bold = True
    r2.font.color.rgb = COLOR_DARK_GREEN
    p_p2.add_run("Câu 1: a-Đ, b-S, c-Đ, d-Đ   |   Câu 2: a-S, b-Đ, c-Đ, d-Đ\nCâu 3: a-S, b-S, c-S, d-S   |   Câu 4: a-Đ, b-Đ, c-Đ, d-Đ")

    p_p3 = doc.add_paragraph()
    p_p3.paragraph_format.space_before = Pt(2)
    p_p3.paragraph_format.space_after = Pt(2)
    r3 = p_p3.add_run("PHẦN III: ")
    r3.font.bold = True
    r3.font.color.rgb = COLOR_DARK_GREEN
    p_p3.add_run("Câu 1: 3   |   Câu 2: 2   |   Câu 3: 2   |   Câu 4: 2   |   Câu 5: 26,4   |   Câu 6: 0,03")

# DỮ LIỆU ĐỀ 21
P1_DATA = [
    (1, "Chất nào sau đây tác dụng được với dung dịch KHCO~3~?",
     [("A", "K~2~SO~4~."), ("B", "KNO~3~."), ("C", "HCl."), ("D", "KCl.")],
     "C", "KHCO~3~ là muối của acid yếu (carbonic acid), tác dụng với acid mạnh HCl sinh ra khí CO~2~: KHCO~3~ + HCl -> KCl + CO~2~^ + H~2~O."),

    (2, "Dung dịch Na~2~CO~3~ tác dụng được với dung dịch nào sau đây?",
     [("A", "Na~2~SO~4~."), ("B", "KNO~3~."), ("C", "KOH."), ("D", "CaCl~2~.")],
     "D", "Dung dịch Na~2~CO~3~ tác dụng với CaCl~2~ tạo kết tủa trắng calcium carbonate: Na~2~CO~3~ + CaCl~2~ -> CaCO~3~v + 2NaCl."),

    (3, "Dung dịch nào sau đây phản ứng với dung dịch HCl dư tạo ra chất khí?",
     [("A", "Ba(OH)~2~."), ("B", "Na~2~CO~3~."), ("C", "K~2~SO~4~."), ("D", "Ca(NO~3~)~2~.")],
     "B", "Na~2~CO~3~ phản ứng với HCl sinh ra bọt khí carbon dioxide: Na~2~CO~3~ + 2HCl -> 2NaCl + CO~2~^ + H~2~O."),

    (4, "Trong cốc nước chứa 0,01 mol Na^+^; 0,02 mol Ca^2+^; 0,01 mol Mg^2+^; 0,05 mol HCO~3~^-^ và 0,02 mol Cl^-^. Đun sôi cốc nước hồi lâu, nước thu được là",
     [("A", "Nước cứng tạm thời."), ("B", "Nước cứng vĩnh cửu."), ("C", "Nước mềm."), ("D", "Nước cứng toàn phần.")],
     "B", "Khi đun sôi: 2HCO~3~^-^ -> CO~3~^2-^ + CO~2~^ + H~2~O. Với n(HCO~3~^-^) = 0,05 mol sinh ra 0,025 mol CO~3~^2-^.\n"
     "- Ion CO~3~^2-^ kết tủa với Ca^2+^: Ca^2+^ + CO~3~^2-^ -> CaCO~3~v (0,02 mol Ca^2+^ phản ứng hết với 0,02 mol CO~3~^2-^).\n"
     "- Lượng CO~3~^2-^ còn dư = 0,025 - 0,02 = 0,005 mol sẽ kết tủa một phần Mg^2+^: Mg^2+^ + CO~3~^2-^ -> MgCO~3~v.\n"
     "- Lượng Mg^2+^ còn lại trong dung dịch = 0,01 - 0,005 = 0,005 mol.\n"
     "Sau khi đun sôi, trong dung dịch vẫn còn chứa cation Mg^2+^ (0,005 mol) và anion Cl^-^ (0,02 mol) mà không còn ion HCO~3~^-^ => Nước thu được là nước cứng vĩnh cửu."),

    (5, "Phản ứng giải thích sự hình thành thạch nhũ trong các hang động núi đá vôi là:",
     [("A", "CaCO~3~ + CO~2~ + H~2~O -> Ca(HCO~3~)~2~."),
      ("B", "CaCO~3~ + 2HCl -> CaCl~2~ + CO~2~ + H~2~O."),
      ("C", "CaCO~3~ -(t°)-> CaO + CO~2~."),
      ("D", "Ca(HCO~3~)~2~ <=> CaCO~3~ + CO~2~ + H~2~O.")],
     "D", "Quá trình hình thành thạch nhũ là phản ứng phân huỷ muối calcium hydrogencarbonate nhỏ giọt từ vòm hang đá vôi: Ca(HCO~3~)~2~ -> CaCO~3~v + CO~2~^ + H~2~O. Chiều thuận là quá trình tạo thạch nhũ, chiều nghịch là quá trình xâm thực hoà tan đá vôi."),

    (6, "Nước thải công nghiệp thường chứa các ion kim loại nặng như Hg^2+^, Pb^2+^, Fe^3+^. Để xử lí triệt để và làm giảm nồng độ của các ion kim loại nặng với chi phí thấp nhất, người ta sử dụng dung dịch chất nào sau đây?",
     [("A", "HCl."), ("B", "NaCl."), ("C", "Ca(OH)~2~."), ("D", "K~2~SO~4~.")],
     "C", "Nước vôi tôi Ca(OH)~2~ là hoá chất rẻ tiền, dễ kiếm và cung cấp ion OH^-^ để kết tủa hầu hết các ion kim loại nặng dưới dạng hydroxide ít tan: Fe(OH)~3~v, Pb(OH)~2~v, Hg(OH)~2~ -> HgOv, giúp lọc tách dễ dàng khỏi nguồn nước thải."),

    (7, "Cho dãy các kim loại sau: Fe, Na, K, Cu và Li. Số kim loại trong dãy tác dụng được với nước ở nhiệt độ thường là",
     [("A", "1."), ("B", "2."), ("C", "3."), ("D", "4.")],
     "C", "Có 3 kim loại là Na, K, Li (đều thuộc nhóm IA) tác dụng mãnh liệt với nước ở nhiệt độ thường. Fe và Cu không phản ứng với nước ở nhiệt độ thường."),

    (8, "Nguyên tắc cơ bản để làm mềm nước cứng là làm giảm nồng độ của",
     [("A", "ion Ca^2+^, Mg^2+^."), ("B", "ion HCO~3~^-^."), ("C", "ion OH^-^, SO~4~^2-^."), ("D", "ion Ca^2+^.")],
     "A", "Tính cứng của nước do các cation Ca^2+^ và Mg^2+^ gây ra. Do đó nguyên tắc làm mềm nước cứng là làm giảm nồng độ của các cation Ca^2+^ và Mg^2+^ trong nước (bằng phương pháp kết tủa hoặc trao đổi ion)."),

    (9, "Kim loại nhóm IIA có nhiệt độ nóng chảy, nhiệt độ sôi, khối lượng riêng biến đổi không theo quy luật như kim loại kiềm nhóm IA, nguyên nhân chủ yếu là do các kim loại nhóm IIA có",
     [("A", "điện tích hạt nhân khác nhau."), ("B", "cấu hình electron khác nhau."), ("C", "bán kính nguyên tử khác nhau."), ("D", "kiểu mạng tinh thể khác nhau.")],
     "D", "Các kim loại nhóm IA đều kết tinh theo kiểu mạng lập phương tâm khối nên tính chất vật lí biến đổi tuần hoàn có quy luật. Trong khi đó, kim loại nhóm IIA có các kiểu mạng tinh thể khác nhau (Be, Mg: lục phương; Ca, Sr: lập phương tâm diện; Ba: lập phương tâm khối) nên các hằng số vật lí biến đổi không theo quy luật đơn điệu."),

    (10, "Khi để vôi sống trong không khí ẩm một thời gian sẽ có hiện tượng một phần bị chuyển hoá trở lại thành đá vôi. Khí nào sau đây có trong không khí gây ra hiện tượng trên?",
     [("A", "Oxygen."), ("B", "Methane."), ("C", "Nitrogen."), ("D", "Carbon dioxide.")],
     "D", "Vôi sống CaO hút ẩm trong không khí tạo Ca(OH)~2~, sau đó tác dụng với khí carbon dioxide (CO~2~) tạo thành calcium carbonate (đá vôi): CaO + H~2~O -> Ca(OH)~2~; Ca(OH)~2~ + CO~2~ -> CaCO~3~v + H~2~O."),

    (11, "Đun nóng tinh thể CaF~2~ với dung dịch H~2~SO~4~ đặc ở nhiệt độ khoảng 250 °C, thu được khí nào sau đây?",
     [("A", "SO~2~."), ("B", "F~2~."), ("C", "HF."), ("D", "H~2~S.")],
     "C", "Phản ứng điều chế hydrogen fluoride (HF) trong công nghiệp bằng phương pháp sunfat: CaF~2~(s) + H~2~SO~4~(đặc) -(~250 °C)-> CaSO~4~ + 2HF^(k)."),

    (12, "Cho dung dịch HCl vào dung dịch X thấy có hiện tượng sủi bọt khí; nếu cho dung dịch Ca(OH)~2~ vào dung dịch X thì sinh ra kết tủa trắng. Dung dịch X là",
     [("A", "Na~2~SO~4~."), ("B", "KNO~3~."), ("C", "Ca(HCO~3~)~2~."), ("D", "BaCl~2~.")],
     "C", "Ca(HCO~3~)~2~ tác dụng với acid HCl sinh khí CO~2~: Ca(HCO~3~)~2~ + 2HCl -> CaCl~2~ + 2CO~2~^ + 2H~2~O.\nTác dụng với dung dịch Ca(OH)~2~ sinh kết tủa trắng: Ca(HCO~3~)~2~ + Ca(OH)~2~ -> 2CaCO~3~v + 2H~2~O."),

    (13, "Cho một mẩu Na vào dung dịch MgSO~4~ dư, thu được kết tủa X và chất khí Y. Hai chất X, Y lần lượt là",
     [("A", "Mg và H~2~."), ("B", "Mg(OH)~2~ và H~2~."), ("C", "Mg và O~2~."), ("D", "Na~2~SO~4~ và H~2~.")],
     "B", "Na phản ứng với nước trước: 2Na + 2H~2~O -> 2NaOH + H~2~^ (khí Y là H~2~). Sau đó NaOH tác dụng với MgSO~4~: 2NaOH + MgSO~4~ -> Mg(OH)~2~v (kết tủa X) + Na~2~SO~4~."),

    (14, "Trong cốc nước chứa nhiều các ion sau: Ca^2+^, Mg^2+^, Cl^-^, SO~4~^2-^. Nước trong cốc trên thuộc loại",
     [("A", "có tính cứng vĩnh cửu."), ("B", "không có tính cứng."), ("C", "có tính cứng tạm thời."), ("D", "có tính cứng toàn phần.")],
     "A", "Nước cứng chứa các anion Cl^-^, SO~4~^2-^ cùng với cation Ca^2+^, Mg^2+^ (không có ion HCO~3~^-^) là nước cứng vĩnh cửu."),

    (15, "Nhận định nào không đúng về cấu tạo và tính chất của các kim loại nhóm IIA?",
     [("A", "Khối lượng riêng tương đối nhỏ, là những kim loại nhẹ hơn nhôm (trừ Ba)."),
      ("B", "Độ cứng cao hơn kim loại nhóm IA, nhưng mềm hơn kim loại nhôm."),
      ("C", "Mạng tinh thể của chúng đều có kiểu lập phương tâm khối."),
      ("D", "Nhiệt độ nóng chảy và nhiệt độ sôi tương đối thấp (trừ Be).")],
     "C", "Các kim loại nhóm IIA không cùng một kiểu mạng tinh thể: Be và Mg có mạng lục phương; Ca và Sr có mạng lập phương tâm diện; Ba có mạng lập phương tâm khối."),

    (16, "Cho các nhận định sau về tác hại của nước cứng:\n"
         "(1) Làm giảm bọt và giảm khả năng tẩy rửa khi giặt quần áo bằng xà phòng.\n"
         "(2) Làm đường ống dẫn nước bị đóng cặn, giảm lưu lượng dòng chảy.\n"
         "(3) Làm thức ăn lâu chín và giảm mùi vị đặc trưng khi nấu nướng.\n"
         "(4) Làm nồi hơi bị phủ cặn cách nhiệt, gây tốn nhiên liệu và có nguy cơ gây nổ lò hơi.\n"
         "Số nhận định đúng là",
     [("A", "1."), ("B", "2."), ("C", "3."), ("D", "4.")],
     "D", "Cả 4 nhận định trên đều là các tác hại tiêu biểu và nghiêm trọng của nước cứng trong đời sống sinh hoạt và sản xuất công nghiệp."),

    (17, "Phát biểu nào sau đây không đúng?",
     [("A", "Đun sôi nước cứng chỉ làm mềm nước cứng tạm thời."),
      ("B", "Dung dịch Na~2~CO~3~ làm mềm nước cứng tạm thời, nước cứng vĩnh cửu."),
      ("C", "Dung dịch NH~4~Cl làm mềm nước cứng."),
      ("D", "Cho dung dịch Ca(OH)~2~ vừa đủ làm mềm nước cứng tạm thời.")],
     "C", "Dung dịch NH~4~Cl không phản ứng tạo kết tủa với Ca^2+^ và Mg^2+^ nên không thể làm mềm nước cứng."),

    (18, "Cho sơ đồ chuyển hoá hoá học sau:",
     [("A", "BaCO~3~."), ("B", "CaCO~3~."), ("C", "BaSO~3~."), ("D", "MgCO~3~.")],
     "A", "Nhiệt phân X tạo ra 2 oxide: BaCO~3~ -(t°)-> BaO (B~1~) + CO~2~ (A~1~).\n"
     "- Khí A~1~ (CO~2~): CO~2~ + Ca(OH)~2~ -> CaCO~3~v (A~2~) + H~2~O; CaCO~3~ -(900 °C)-> CaO + CO~2~^.\n"
     "- Oxide B~1~ (BaO): BaO + 2HCl -> BaCl~2~ (B~2~) + H~2~O; BaCl~2~ + Na~2~SO~4~ -> BaSO~4~v (B~3~) + 2NaCl.\n"
     "Chất X là BaCO~3~.")
]

P2_DATA = [
    (1, "Trong tự nhiên, các nguyên tố nhóm IA chỉ tồn tại ở dạng hợp chất (chủ yếu là dạng muối). Sodium và potassium là hai nguyên tố phổ biến trong vỏ Trái Đất, có nhiều trong nước biển, muối mỏ, quặng halite (NaCl), quặng sylvinite (NaCl.KCl).",
     [("a", "Trong đời sống hàng ngày, sodium chloride được dùng làm gia vị, bảo quản và chế biến thực phẩm.", True),
      ("b", "Điện phân dung dịch NaCl bão hoà có màng ngăn điện cực được ứng dụng trong công nghiệp để sản xuất nước Javel.", False),
      ("c", "Trong công nghiệp, khí chlorine được sản xuất bằng cách điện phân dung dịch muối ăn bão hoà có màng ngăn xốp.", True),
      ("d", "Sodium carbonate là nguyên liệu quan trọng để sản xuất thuỷ tinh, xà phòng, bột giặt, giấy, sợi và chất tẩy rửa.", True)],
     None, None,
     "a) Đúng vì muối ăn NaCl là gia vị không thể thiếu và dùng bảo quản ướp thịt, cá.\n"
     "b) Sai vì sản xuất nước Javel bằng phương pháp điện phân dung dịch NaCl KHÔNG có màng ngăn để Cl~2~ tác dụng trực tiếp với dung dịch NaOH.\n"
     "c) Đúng vì điện phân dd NaCl bão hoà có màng ngăn xốp thu được Cl~2~ ở anode, H~2~ và NaOH ở cathode.\n"
     "d) Đúng vì Na~2~CO~3~ (soda) là hoá chất cơ bản của công nghiệp hoá chất vô cơ."),

    (2, "Khi nồng độ acid trong dịch vị dạ dày tăng cao có thể gây ra viêm loét dạ dày, tá tràng. Để giảm bớt nồng độ acid trong dịch vị dạ dày người ta thường dùng thuốc đau dạ dày có chứa thành phần chính là muối X.",
     [("a", "Muối X là sodium bicarbonate có công thức hoá học là Na~2~CO~3~.", False),
      ("b", "Muối X còn được dùng để tạo độ xốp cho các loại bánh nướng, làm mềm thực phẩm.", True),
      ("c", "Muối X bị phân huỷ khi đun nóng.", True),
      ("d", "Trong công nghiệp hoá chất, muối X được sản xuất chủ yếu bằng phương pháp Solvay.", True)],
     None, None,
     "a) Sai vì muối X là sodium hydrogencarbonate (baking soda) có công thức hoá học là NaHCO~3~ (không phải Na~2~CO~3~).\n"
     "b) Đúng vì NaHCO~3~ khi nướng sinh khí CO~2~ làm xốp ruột bánh mì, bánh quy.\n"
     "c) Đúng vì 2NaHCO~3~ -(t°)-> Na~2~CO~3~ + CO~2~^ + H~2~O.\n"
     "d) Đúng vì NaHCO~3~ là sản phẩm trung gian chính thu được từ tháp carbonate hoá trong phương pháp Solvay: NaCl + NH~3~ + CO~2~ + H~2~O -> NaHCO~3~v + NH~4~Cl."),

    (3, "Bảng sau đây cho biết thông số một số tính chất vật lí cơ bản của các kim loại nhóm IIA:",
     [("a", "Nhiệt độ sôi tăng dần từ Be đến Ba.", False),
      ("b", "Khối lượng riêng tăng dần đều đặn từ trên xuống dưới từ Mg đến Ba.", False),
      ("c", "Độ cứng của các kim loại tăng dần từ Be đến Ba.", False),
      ("d", "Nhiệt độ nóng chảy giảm dần đều đặn từ Be đến Ba.", False)],
     "San_Pham/images_crop_de21/p2_cau3_bang_tinh_chat.png", Inches(5.0),
     "a) Sai vì nhiệt độ sôi biến thiên không theo quy luật tăng dần: Be (2770 °C) -> Mg (1110 °C) -> Ca (1440 °C) -> Sr (1380 °C) -> Ba (1640 °C).\n"
     "b) Sai vì khối lượng riêng không tăng đều đặn: Mg (1,74 g/cm^3^) giảm xuống Ca (1,55 g/cm^3^) rồi mới tăng lên Sr (2,6 g/cm^3^) và Ba (3,5 g/cm^3^).\n"
     "c) Sai vì độ cứng có xu hướng giảm dần do bán kính ion tăng và mật độ electron giảm.\n"
     "d) Sai vì nhiệt độ nóng chảy không giảm đều đặn: Be (1280 °C) -> Mg (650 °C) -> Ca (838 °C) -> Sr (768 °C) -> Ba (714 °C) do các kim loại có cấu trúc mạng tinh thể khác nhau."),

    (4, "Cho các phản ứng hoá học xảy ra theo sơ đồ sau:\n"
        "(1) 2X~1~ + 2H~2~O -(đpdd, có màng ngăn)-> 2X~2~ + X~3~ + H~2~\n"
        "(2) 2X~2~ + X~4~ -> BaCO~3~v + K~2~CO~3~ + 2H~2~O\n"
        "(3) X~4~ + X~5~ -(1:1)-> BaSO~4~v + 2CO~2~^ + 2H~2~O",
     [("a", "Hợp chất X~2~ là potassium hydroxide (KOH).", True),
      ("b", "Đơn chất X~3~ là chất khí có màu vàng lục nhạt.", True),
      ("c", "Muối X~4~ là muối acid (muối hydrogencarbonate).", True),
      ("d", "Khi đun nóng dung dịch muối X~4~ thì thu được kết tủa trắng.", True)],
     None, None,
     "a) Đúng vì từ phản ứng (2) sinh ra K~2~CO~3~ nên X~2~ là KOH, tương ứng X~1~ là KCl: 2KCl + 2H~2~O -> 2KOH + Cl~2~^ + H~2~^.\n"
     "b) Đúng vì X~3~ là khí chlorine (Cl~2~), có màu vàng lục nhạt đặc trưng và mùi xốc.\n"
     "c) Đúng vì X~4~ là barium hydrogencarbonate Ba(HCO~3~)~2~, là muối acid của acid yếu H~2~CO~3~: 2KOH + Ba(HCO~3~)~2~ -> BaCO~3~v + K~2~CO~3~ + 2H~2~O.\n"
     "d) Đúng vì khi đun nóng, muối Ba(HCO~3~)~2~ bị nhiệt phân sinh kết tủa trắng BaCO~3~: Ba(HCO~3~)~2~ -(t°)-> BaCO~3~v + CO~2~^ + H~2~O.")
]

P3_DATA = [
    (1, "Trong công nghiệp, quá trình sản xuất soda bằng phương pháp Solvay đi từ các nguyên liệu ban đầu là muối ăn (NaCl), đá vôi (CaCO~3~) và ammonia (NH~3~). Có bao nhiêu phương trình hoá học chủ yếu đã xảy ra trong toàn bộ quy trình trên?",
     "3",
     "Có 3 phản ứng hoá học chủ yếu:\n"
     "(1) Nung đá vôi tạo CO~2~: CaCO~3~ -(t°)-> CaO + CO~2~^\n"
     "(2) Carbonate hoá tạo kết tủa: NaCl + NH~3~ + CO~2~ + H~2~O -> NaHCO~3~v + NH~4~Cl\n"
     "(3) Nhiệt phân tạo soda: 2NaHCO~3~ -(t°)-> Na~2~CO~3~ + CO~2~^ + H~2~O\n"
     "(Quá trình tái sinh NH~3~: 2NH~4~Cl + CaO -> CaCl~2~ + 2NH~3~^ + H~2~O). Đáp án chuẩn: 3 phương trình."),

    (2, "Cho các nguyên tố sau: Kr, Ca, Sr, Li, Ni, Rb. Có bao nhiêu nguyên tố là kim loại kiềm (nhóm IA) trong dãy trên?",
     "2",
     "Có 2 nguyên tố là kim loại kiềm thuộc nhóm IA: Li và Rb.\n(Ca, Sr thuộc nhóm IIA; Ni là kim loại chuyển tiếp nhóm VIIIB; Kr là khí hiếm nhóm VIIIA)."),

    (3, "Cho các chất sau: Ca(OH)~2~, Na~2~CO~3~, Na~3~PO~4~, NaOH, KCl. Có bao nhiêu chất có thể dùng để làm mềm nước có tính cứng vĩnh cửu?",
     "2",
     "Có 2 chất là Na~2~CO~3~ và Na~3~PO~4~:\n"
     "- Ca^2+^ + CO~3~^2-^ -> CaCO~3~v; Mg^2+^ + CO~3~^2-^ -> MgCO~3~v\n"
     "- 3Ca^2+^ + 2PO~4~^3-^ -> Ca~3~(PO~4~)~2~v; 3Mg^2+^ + 2PO~4~^3-^ -> Mg~3~(PO~4~)~2~v\n"
     "Ca(OH)~2~, NaOH và KCl không làm mềm được nước cứng vĩnh cửu."),

    (4, "Cho dung dịch của 4 chất sau đều trong suốt và không màu: NaCl, Na~2~SO~4~, KCl, LiNO~3~. Chỉ dùng ngọn lửa không màu của đèn khí thì có thể nhận biết được bao nhiêu dung dịch riêng biệt trong số 4 dung dịch trên?",
     "2",
     "Nhận biết được 2 dung dịch là KCl và LiNO~3~:\n"
     "- Hợp chất của potassium (KCl) nhuộm ngọn lửa thành màu tím nhạt.\n"
     "- Hợp chất của lithium (LiNO~3~) nhuộm ngọn lửa thành màu đỏ tía.\n"
     "- NaCl và Na~2~SO~4~ đều chứa ion Na^+^ nên đều nhuộm ngọn lửa thành màu vàng tươi, không thể phân biệt được hai dung dịch này với nhau."),

    (5, "Ở 20 °C, độ tan của NaCl trong nước là 35,9 gam trong 100 gam nước. Ở nhiệt độ này, dung dịch NaCl bão hoà có nồng độ phần trăm là a%. Xác định giá trị của a (làm tròn kết quả đến hàng phần mười).",
     "26,4",
     "Khối lượng chất tan m(ct) = 35,9 gam.\n"
     "Khối lượng dung môi m(dm) = 100 gam.\n"
     "Khối lượng dung dịch bão hoà: m(dd) = m(ct) + m(dm) = 35,9 + 100 = 135,9 gam.\n"
     "Nồng độ phần trăm của dung dịch bão hoà:\n"
     "C% = (m(ct) / m(dd)) * 100% = (35,9 / 135,9) * 100% ≈ 26,416% ≈ 26,4%.\n"
     "Vậy a = 26,4."),

    (6, "Một loại thuốc kháng acid dạ dày có tên thương mại là Rolaids chứa thành phần hoạt chất chính là dihydroxyaluminum sodium carbonate có công thức NaAl(OH)~2~CO~3~. Mỗi viên thuốc chứa 0,334 gam hoạt chất NaAl(OH)~2~CO~3~. Một cốc nước giải khát có thể tích 250 mL và có pH = 3,0. Giả thiết phản ứng trung hòa xảy ra hoàn toàn theo phương trình: NaAl(OH)~2~CO~3~ + 4H^+^ -> Na^+^ + Al^3+^ + CO~2~^ + 3H~2~O. Tính số viên thuốc Rolaids cần dùng để trung hòa hết lượng acid có trong cốc nước giải khát trên (làm tròn kết quả đến hai chữ số thập phân).",
     "0,03",
     "- Nồng độ ion H^+^ trong cốc nước giải khát:\n"
     "[H^+^] = 10^-pH^ = 10^-3^ M = 0,001 M.\n"
     "- Số mol ion H^+^ trong 250 mL (0,25 L) dung dịch:\n"
     "n(H^+^) = 0,25 * 10^-3^ = 2,5.10^-4^ mol.\n"
     "- Theo phương trình hoá học: n(NaAl(OH)~2~CO~3~) = n(H^+^) / 4 = 2,5.10^-4^ / 4 = 6,25.10^-5^ mol.\n"
     "- Khối lượng mol của NaAl(OH)~2~CO~3~:\n"
     "M = 23 + 27 + (16 + 1)*2 + 12 + 16*3 = 144 g/mol.\n"
     "- Khối lượng hoạt chất cần dùng:\n"
     "m = 6,25.10^-5^ * 144 = 0,009 gam.\n"
     "- Số viên thuốc Rolaids cần dùng:\n"
     "Số viên = 0,009 / 0,334 ≈ 0,0269 viên ≈ 0,03 viên.")
]

def build_exam_docx(is_teacher=False):
    doc = docx.Document()
    add_header_footer(doc)
    create_title_block(doc, is_teacher=is_teacher)

    # PHẦN I
    add_part_heading(doc, "PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn. Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án.")
    for q_num, text, options, c_key, expl in P1_DATA:
        img_p = "San_Pham/images_crop_de21/p1_cau18_so_do.png" if q_num == 18 else None
        img_w = Inches(3.8) if q_num == 18 else Inches(3.5)
        render_mcq(doc, q_num, text, options, c_key, expl, is_teacher=is_teacher, img_path=img_p, img_width=img_w)

    # PHẦN II
    add_part_heading(doc, "PHẦN II. Câu trắc nghiệm đúng sai. Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.")
    for q_num, text, sub_items, img_path, img_w, expl in P2_DATA:
        render_tf(doc, q_num, text, sub_items, expl, is_teacher=is_teacher, img_path=img_path, img_width=img_w)

    # PHẦN III
    add_part_heading(doc, "PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn. Thí sinh trả lời từ câu 1 đến câu 6.")
    for q_num, text, ans, expl in P3_DATA:
        render_short(doc, q_num, text, ans, expl, is_teacher=is_teacher)

    if is_teacher:
        add_answer_table(doc)

    out_name = "De_Kiem_Tra_Hoa_12_De21_GiaoVien.docx" if is_teacher else "De_Kiem_Tra_Hoa_12_De21.docx"
    out_path = os.path.join(OUT_DIR, out_name)
    doc.save(out_path)
    print(f"Đã lưu: {out_path}")

def main():
    print("--- Bắt đầu tạo 2 file Word cho ĐỀ SỐ 21 (Mã 121) ---")
    build_exam_docx(is_teacher=False)
    build_exam_docx(is_teacher=True)
    print("--- Hoàn tất sinh file Word cho ĐỀ SỐ 21! ---")

if __name__ == "__main__":
    main()
