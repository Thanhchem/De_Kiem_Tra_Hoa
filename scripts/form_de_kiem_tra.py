# -*- coding: utf-8 -*-
"""
FORM ĐỀ KIỂM TRA HÓA HỌC CHUẨN (Thầy TRẦN VĂN THẠNH)
Mô-đun hỗ trợ tạo file Word (.docx) chuẩn theo Form Đề Kiểm Tra.
Tuyệt đối KHÔNG dùng Table trong nội dung câu hỏi/đáp án/lời giải.
Header & Footer, Margins, Tab Stops theo đúng quy chuẩn 100%.
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# Bảng màu quy chuẩn
COLOR_HF_RED = RGBColor(211, 47, 47)        # Coral Red (#D32F2F) cho Header & Footer
COLOR_HF_GREEN_HEX = "8BC390"               # Sage Green (#8BC390) cho viền Header & Footer
COLOR_PRIMARY = RGBColor(46, 125, 50)       # Xanh lá đậm (#2E7D32) cho Tiêu đề Phần
COLOR_CORRECT = RGBColor(198, 40, 40)       # Đỏ (#C62828) cho đáp án đúng Bản Giáo Viên
COLOR_SOLUTION = RGBColor(21, 101, 192)     # Xanh lam (#1565C0) cho tiêu đề Lời giải
COLOR_GRAY = RGBColor(97, 97, 97)           # Xám nhạt
COLOR_BORDER_GREY_HEX = "CCCCCC"            # Xám viền phân cách tiêu đề

def add_p_border_bottom(p, color="8BC390", sz="6"):
    """Thêm đường kẻ dưới cho đoạn văn bản (dùng cho Header hoặc dòng phân cách)."""
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def add_p_border_top(p, color="8BC390", sz="6"):
    """Thêm đường kẻ trên cho đoạn văn bản (dùng cho Footer)."""
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))

def set_run_style(run, font_name="Times New Roman", font_size=12, bold=False, italic=False, subscript=False, superscript=False, color=None):
    """Cài đặt style chi tiết cho từng run text."""
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color

def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    """Chèn trường số trang tự động PAGE vào chân trang."""
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)
    
    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)

def render_rich_text(p, text, base_font="Times New Roman", base_size=10.5, base_bold=False, base_italic=False, base_color=None):
    """
    Hỗ trợ sinh văn bản phong phú từ cú pháp markdown:
    - *text*: in đậm
    - _text_: in nghiêng
    - ~text~: chỉ số dưới (Subscript, vd: H~2~O)
    - ^text^: chỉ số trên (Superscript, vd: Fe^3+^)
    """
    tokens = re.split(r'(~[^~]+~|\^[^\^]+\^|\*[^\*]+\*|_[^_]+_)', text)
    for token in tokens:
        if not token:
            continue
        is_bold = base_bold
        is_italic = base_italic
        is_sub = False
        is_sup = False
        content = token

        if token.startswith('*') and token.endswith('*'):
            is_bold = True
            content = token[1:-1]
        elif token.startswith('_') and token.endswith('_'):
            is_italic = True
            content = token[1:-1]
        elif token.startswith('~') and token.endswith('~'):
            is_sub = True
            content = token[1:-1]
        elif token.startswith('^') and token.endswith('^'):
            is_sup = True
            content = token[1:-1]

        run = p.add_run(content)
        run_color = COLOR_CORRECT if (is_bold and not base_bold and not base_color) else base_color
        set_run_style(run, font_name=base_font, font_size=base_size, bold=is_bold, italic=is_italic,
                      subscript=is_sub, superscript=is_sup, color=run_color)

def setup_header_footer(doc):
    """
    Cài đặt chuẩn trang A4, lề 1.5 cm, Header & Footer theo đúng quy chuẩn:
    Header: Bảng 1x2, rộng 18.0cm, viền dưới #8BC390.
    Footer: Bảng 1x2, rộng 18.0cm, viền trên #8BC390.
    """
    normal_style = doc.styles['Normal']
    normal_style.paragraph_format.tab_stops.clear_all()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    section.start_type = WD_SECTION_START.NEW_PAGE
    section.different_first_page_header_footer = False
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    # 1. Header (Bảng 1x2, rộng 18.0 cm)
    header = section.header
    for p in header.paragraphs:
        p.text = ""
    tbl_h = header.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl_h.cell(0, 0)
    c_right = tbl_h.cell(0, 1)
    c_left.width = Cm(8.0)
    c_right.width = Cm(10.0)

    p_hl = c_left.paragraphs[0]
    p_hl.paragraph_format.space_before = Pt(0)
    p_hl.paragraph_format.space_after = Pt(2)
    p_hl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_hl = p_hl.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)

    p_hr = c_right.paragraphs[0]
    p_hr.paragraph_format.space_before = Pt(0)
    p_hr.paragraph_format.space_after = Pt(2)
    p_hr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p_hr.add_run("Thầy ")
    set_run_style(r1, font_size=10, color=COLOR_HF_RED)
    r2 = p_hr.add_run("TRẦN VĂN THẠNH")
    set_run_style(r2, font_size=10, bold=True, color=COLOR_HF_RED)
    r3 = p_hr.add_run(" - 0777.470.803")
    set_run_style(r3, font_size=10, color=COLOR_HF_RED)

    add_p_border_bottom(p_hl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_bottom(p_hr, color=COLOR_HF_GREEN_HEX, sz="6")

    # 2. Footer (Bảng 1x2, rộng 18.0 cm)
    footer = section.footer
    for p in footer.paragraphs:
        p.text = ""
    tbl_f = footer.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_fl = tbl_f.cell(0, 0)
    c_fr = tbl_f.cell(0, 1)
    c_fl.width = Cm(14.5)
    c_fr.width = Cm(3.5)

    p_fl = c_fl.paragraphs[0]
    p_fl.paragraph_format.space_before = Pt(2)
    p_fl.paragraph_format.space_after = Pt(0)
    p_fl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_fl = p_fl.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)

    p_fr = c_fr.paragraphs[0]
    p_fr.paragraph_format.space_before = Pt(2)
    p_fr.paragraph_format.space_after = Pt(0)
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_field(p_fr, font_size=10, bold=True, color=COLOR_HF_RED)

    add_p_border_top(p_fl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_top(p_fr, color=COLOR_HF_GREEN_HEX, sz="6")

def build_title_block(doc, 
                      school_line="SỞ GIÁO DỤC VÀ ĐÀO TẠO THỪA THIÊN HUẾ",
                      exam_title="ĐỀ KIỂM TRA ĐỊNH KỲ",
                      subject_line="MÔN: HOÁ HỌC - LỚP 11",
                      num_pages="04",
                      duration="45 phút",
                      exam_code="101",
                      atomic_masses="H = 1; C = 12; N = 14; O = 16; Na = 23; Al = 27; S = 32; Cl = 35,5; K = 39; Ca = 40; Fe = 56; Cu = 64; Ba = 137",
                      is_teacher=False):
    """
    Xây dựng khối tiêu đề đầu đề thi (Title Block) dùng Tab Stop phải ở 18.0 cm.
    TUYỆT ĐỐI KHÔNG DÙNG BẢNG để tránh lỗi viền rườm rà.
    """
    headers_data = [
        (school_line, exam_title, 10, True, 11, True),
        ("ĐỀ CHÍNH THỨC", subject_line, 10.5, True, 11, True),
        (f"(Đề thi có {num_pages} trang)", f"Thời gian làm bài: {duration} | MÃ ĐỀ THI: {exam_code}", 9.5, False, 10, True)
    ]
    for l_text, r_text, l_size, l_bold, r_size, r_bold in headers_data:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)

        rl = p.add_run(l_text)
        set_run_style(rl, font_size=l_size, bold=l_bold, italic=(not l_bold))

        p.add_run('\t')
        rr = p.add_run(r_text)
        set_run_style(rr, font_size=r_size, bold=r_bold, italic=(not r_bold))

    if is_teacher:
        p_gv = doc.add_paragraph()
        p_gv.paragraph_format.space_before = Pt(2)
        p_gv.paragraph_format.space_after = Pt(2)
        p_gv.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rgv = p_gv.add_run("(BẢN GIÁO VIÊN -- CÓ LỜI GIẢI CHI TIẾT)")
        set_run_style(rgv, font_size=11, bold=True, color=COLOR_CORRECT)

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(4)
    p_info.paragraph_format.space_after = Pt(2)
    render_rich_text(p_info, "*Họ, tên thí sinh:* ............................................................ *Lớp:* ............. *Số báo danh:* .............", base_size=10.5)

    if atomic_masses:
        p_note = doc.add_paragraph()
        p_note.paragraph_format.space_before = Pt(0)
        p_note.paragraph_format.space_after = Pt(4)
        render_rich_text(p_note, f"_(Cho biết nguyên tử khối: {atomic_masses})._", base_size=10, base_italic=True)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(4)
    add_p_border_bottom(p_div, color=COLOR_BORDER_GREY_HEX, sz="4")

def add_section_header(doc, title, subtitle=None):
    """Tiêu đề phần thi (Phần I, Phần II, Phần III, Phần IV)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, f"*{title}*", base_size=11.5, base_bold=True, base_color=COLOR_PRIMARY)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(4)
        render_rich_text(p2, f"_{subtitle}_", base_size=10.5, base_italic=True)

def render_question_prompt(p, text, base_size=10.5):
    """In đậm tiền tố 'Câu X.' ở đầu câu hỏi, phần nội dung sau đó giữ kiểu bình thường."""
    m = re.match(r'^(Câu\s+\d+\.)\s*(.*)$', text.strip(), re.DOTALL)
    if m:
        lbl = m.group(1)
        rest = m.group(2)
        r_lbl = p.add_run(lbl + " ")
        set_run_style(r_lbl, font_size=base_size, bold=True)
        render_rich_text(p, rest, base_size=base_size, base_bold=False)
    else:
        render_rich_text(p, text, base_size=base_size, base_bold=False)

def render_single_option(p, opt_str, is_teacher=False, base_size=10.5):
    """
    In đậm nhãn A., B., C., D. cho mọi phương án.
    Nếu là bản giáo viên và là phương án đúng: in đậm và tô màu đỏ #C62828.
    """
    clean_str = opt_str.strip()
    is_correct = False
    if clean_str.startswith('*') and clean_str.endswith('*'):
        is_correct = True
        clean_str = clean_str[1:-1].strip()

    m = re.match(r'^([A-D]\.\s*)(.*)$', clean_str)
    if m:
        lbl = m.group(1)
        body = m.group(2)
    else:
        lbl = ""
        body = clean_str

    if is_teacher and is_correct:
        if lbl:
            r_lbl = p.add_run(lbl)
            set_run_style(r_lbl, font_size=base_size, bold=True, color=COLOR_CORRECT)
        render_rich_text(p, body, base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
    else:
        if lbl:
            r_lbl = p.add_run(lbl)
            set_run_style(r_lbl, font_size=base_size, bold=True, color=None)
        render_rich_text(p, body, base_size=base_size, base_bold=False, base_color=None)

def render_options_no_table(doc, opts, is_teacher=False):
    """
    Trình bày 4 phương án A, B, C, D hoàn toàn KHÔNG DÙNG BẢNG:
    - Chiều dài <= 18 ký tự: 1 dòng (Tab stops: 4.5 cm, 9.0 cm, 13.5 cm).
    - Chiều dài 19 - 38 ký tự: 2 dòng (Tab stop: 9.0 cm).
    - Chiều dài > 38 ký tự: 4 dòng riêng biệt (Left indent: 0.5 cm).
    """
    clean_opts = [re.sub(r'[*_~^]', '', opt) for opt in opts]
    max_len = max(len(o) for o in clean_opts)

    if max_len <= 18:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(4.5), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(13.5), WD_TAB_ALIGNMENT.LEFT)

        for idx, opt in enumerate(opts):
            if idx > 0:
                p.add_run('\t')
            render_single_option(p, opt, is_teacher=is_teacher, base_size=10.5)
    elif max_len <= 38:
        p1 = doc.add_paragraph()
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.tab_stops.clear_all()
        p1.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
        render_single_option(p1, opts[0], is_teacher=is_teacher, base_size=10.5)
        p1.add_run('\t')
        render_single_option(p1, opts[1], is_teacher=is_teacher, base_size=10.5)

        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(1)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.tab_stops.clear_all()
        p2.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.LEFT)
        render_single_option(p2, opts[2], is_teacher=is_teacher, base_size=10.5)
        p2.add_run('\t')
        render_single_option(p2, opts[3], is_teacher=is_teacher, base_size=10.5)
    else:
        for opt in opts:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            render_single_option(p, opt, is_teacher=is_teacher, base_size=10.5)

def render_statement(p, stmt, ans=None, is_teacher=False, base_size=10.5):
    """In đậm tiền tố 'a)', 'b)', 'c)', 'd)' ở Phần II; thêm kết luận đúng/sai cho Bản Giáo Viên."""
    m = re.match(r'^([a-d]\))\s*(.*)$', stmt.strip())
    if m:
        lbl = m.group(1)
        body = m.group(2)
        r_lbl = p.add_run(lbl + " ")
        set_run_style(r_lbl, font_size=base_size, bold=True)
        render_rich_text(p, body, base_size=base_size, base_bold=False)
    else:
        render_rich_text(p, stmt, base_size=base_size, base_bold=False)

    if is_teacher and ans:
        p.add_run("  ➔ ")
        r_ans = p.add_run(f"({ans})")
        set_run_style(r_ans, font_size=base_size, bold=True, color=COLOR_CORRECT)

def add_solution_no_table(doc, solution_text):
    """
    Trình bày lời giải chi tiết:
    - Tiêu đề '► Lời giải chi tiết:' màu xanh #1565C0, thụt lề 0.4 cm.
    - Thân bài thụt lề 0.8 cm, cỡ 10 pt.
    """
    lines = solution_text.strip().split("\n")
    p_head = doc.add_paragraph()
    p_head.paragraph_format.left_indent = Cm(0.4)
    p_head.paragraph_format.space_before = Pt(3)
    p_head.paragraph_format.space_after = Pt(1)
    render_rich_text(p_head, "► *Lời giải chi tiết:*", base_size=10.5, base_bold=True, base_color=COLOR_SOLUTION)

    for line in lines:
        if line.startswith("*Lời giải chi tiết:*") or line.startswith("►"):
            continue
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        render_rich_text(p, line, base_size=10)

def add_student_writing_lines(doc, num_lines=4):
    """Thêm các dòng chấm làm bài cho phần tự luận ở Bản Học Sinh."""
    for _ in range(num_lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("." * 105)
        set_run_style(r, font_size=10, color=RGBColor(180, 180, 180))
