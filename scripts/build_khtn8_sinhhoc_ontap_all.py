# -*- coding: utf-8 -*-
r"""
Script biên soạn trọn bộ 3 tài liệu chuẩn FORM ĐỀ KIỂM TRA cho:
ÔN TẬP GIỮA KÌ I – KHTN3 – LỚP 8 (PHÂN MÔN SINH HỌC)
1. De_Cuong_GK1_KHTN8_SinhHoc_HocSinh.docx
2. De_Cuong_GK1_KHTN8_SinhHoc_GiaoVien.docx
3. De_Cuong_GK1_KHTN8_SinhHoc_DayTrenManHinh.pdf (Slide 16:9 Widescreen cho Tivi 65 inch)
Và bản Word: De_Cuong_GK1_KHTN8_SinhHoc_DayTrenManHinh.docx
Thầy TRẦN VĂN THẠNH - 0777.470.803
"""

import os
import sys
import subprocess
import shutil

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# Import các hàm chuẩn từ form_de_kiem_tra
sys.path.insert(0, r"c:\Antigravity_Thanh\scripts")
from form_de_kiem_tra import (
    setup_header_footer,
    build_title_block,
    add_section_header,
    render_question_prompt,
    render_options_no_table,
    render_statement,
    add_solution_no_table,
    add_student_writing_lines,
    render_rich_text,
    set_run_style,
    COLOR_PRIMARY,
    COLOR_CORRECT,
    COLOR_SOLUTION,
    COLOR_HF_RED,
    COLOR_HF_GREEN_HEX
)

OUTPUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
SCRATCH_DIR = r"c:\Antigravity_Thanh\scratch"
os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(SCRATCH_DIR, exist_ok=True)

TARGET_ONEDRIVE_DIRS = [
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn8\KHTN8-IN-SINHHOC",
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn8\KHTN8-IN-DEKIEMTRADINHKI",
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn8"
]

# ==============================================================================
# DỮ LIỆU ĐỀ CƯƠNG KHTN3 - LỚP 8
# ==============================================================================

MCQ_DATA = [
    {
        "id": 1,
        "content": "Trong cơ thể người, hệ cơ quan nào có chức năng thực hiện quá trình trao đổi khí giữa cơ thể với môi trường ngoài?",
        "options": [
            "A. Hệ bài tiết.",
            "B. Hệ hô hấp.",
            "C. Hệ tiêu hóa.",
            "D. Hệ tuần hoàn."
        ],
        "answer": "B",
        "solution": "Hệ hô hấp gồm đường dẫn khí (mũi, họng, thanh quản, khí quản, phế quản) và hai lá phổi, có chức năng thực hiện quá trình trao đổi khí giữa cơ thể với môi trường ngoài (lấy khí O_2 cung cấp cho tế bào và thải khí CO_2 sinh ra từ quá trình trao đổi chất ra ngoài)."
    },
    {
        "id": 2,
        "content": "Khi tham gia các hoạt động như đi bộ, chạy, nhảy hoặc chơi thể thao, hệ vận động có chức năng chủ yếu nào sau đây?",
        "options": [
            "A. Điều hòa lượng đường trong máu.",
            "B. Vận chuyển các chất dinh dưỡng trong cơ thể.",
            "C. Giúp cơ thể vận động và duy trì tư thế.",
            "D. Thực hiện quá trình tiêu hóa thức ăn."
        ],
        "answer": "C",
        "solution": "Hệ vận động ở người gồm bộ xương và hệ cơ, có chức năng nâng đỡ, bảo vệ cơ thể, tạo hình dáng và giúp cơ thể di chuyển, vận động linh hoạt cũng như duy trì các tư thế đứng, ngồi, thăng bằng."
    },
    {
        "id": 3,
        "content": "Phát biểu nào sau đây đúng khi nói về chất dinh dưỡng?",
        "options": [
            "A. Chất dinh dưỡng là những chất cơ thể không cần sử dụng.",
            "B. Chất dinh dưỡng chỉ có trong các loại thực phẩm có nguồn gốc động vật.",
            "C. Chất dinh dưỡng chỉ có vai trò cung cấp năng lượng cho cơ thể.",
            "D. Chất dinh dưỡng là những chất được cơ thể sử dụng để duy trì các hoạt động sống, sinh trưởng và phát triển."
        ],
        "answer": "D",
        "solution": "Chất dinh dưỡng là các chất có trong thức ăn mà cơ thể người và động vật sử dụng để duy trì sự sống, sinh trưởng, phát triển và phục hồi các mô, tế bào tổn thương. Chất dinh dưỡng có trong cả thực phẩm nguồn gốc động vật và thực vật."
    },
    {
        "id": 4,
        "content": "Cơ quan nào sau đây thuộc ống tiêu hóa và có chức năng chứa, nhào trộn thức ăn với dịch tiêu hóa?",
        "options": [
            "A. Gan.",
            "B. Dạ dày.",
            "C. Tuyến tụy.",
            "D. Tuyến nước bọt."
        ],
        "answer": "B",
        "solution": "Dạ dày là đoạn phình to nhất của ống tiêu hóa, có lớp cơ dày thực hiện chức năng tiếp nhận, chứa thức ăn và co bóp nhào trộn đều thức ăn với dịch vị (chứa HCl và enzym pepsin). Gan, tuyến tụy và tuyến nước bọt là các tuyến tiêu hóa (tiết dịch tiêu hóa vào ống tiêu hóa)."
    }
]

SHORT_DATA = {
    "context": "Trong buổi tuyên truyền về sức khỏe học đường, giáo viên giới thiệu một số hình ảnh về các bệnh, tật liên quan đến hệ vận động như cong vẹo cột sống, loãng xương và các tổn thương ở khớp. Giáo viên cũng hướng dẫn học sinh những thói quen sinh hoạt và vận động phù hợp để bảo vệ hệ vận động.\nDựa vào kiến thức đã học, hãy trả lời các câu hỏi sau:",
    "sub_questions": [
        {
            "id": "a",
            "prompt": "Một người lớn tuổi được chẩn đoán mắc bệnh loãng xương. Bệnh này ảnh hưởng trực tiếp đến bộ phận nào của hệ vận động?\nQuy ước: 1. Xương.    2. Cơ.    3. Khớp.",
            "answer_short": "1 (hoặc Xương)",
            "solution": "Bệnh loãng xương là tình trạng suy giảm mật độ khoáng chất của xương (đặc biệt là calcium), khiến cấu trúc xương trở nên mỏng, xốp, giòn và dễ gãy. Bệnh này ảnh hưởng trực tiếp đến *Xương* (chọn đáp án 1)."
        },
        {
            "id": "b",
            "prompt": "Để bảo vệ hệ vận động và hạn chế nguy cơ mắc các bệnh, tật liên quan đến hệ vận động, có bao nhiêu biện pháp đúng trong các biện pháp sau?\n(1) Ăn uống đủ chất, bổ sung thực phẩm giàu calcium và vitamin D phù hợp với nhu cầu cơ thể.\n(2) Thường xuyên mang vác vật nặng vượt quá sức chịu đựng của cơ thể.\n(3) Duy trì tư thế đúng khi ngồi học, đứng và đi lại.\n(4) Tập luyện thể dục thể thao thường xuyên với cường độ phù hợp.",
            "answer_short": "3 (gồm các biện pháp 1, 3, 4)",
            "solution": "Có *3 biện pháp đúng* là (1), (3) và (4).\n• Biện pháp (2) là sai vì mang vác quá nặng sẽ gây tổn thương cột sống, cong vẹo cột sống và tổn thương các khớp."
        }
    ]
}

ESSAY_QUESTION = {
    "num": 6,
    "title": "Hãy kể tên 3 bệnh ở đường tiêu hóa mà em biết. Đề xuất các biện pháp giúp phòng tránh các bệnh về đường tiêu hóa.",
    "solution": (
        "1. KỂ TÊN 3 BỆNH Ở ĐƯỜNG TIÊU HÓA:\n"
        "   Học sinh kể tên 3 trong số các bệnh sau:\n"
        "   • Bệnh sâu răng: Do vi khuẩn lên men thức ăn thừa sinh axit phá hủy men răng, ngà răng.\n"
        "   • Bệnh viêm loét dạ dày - tá tràng: Do vi khuẩn Helicobacter pylori (Hp), ăn uống thất thường, căng thẳng.\n"
        "   • Bệnh tiêu chảy / ngộ độc thực phẩm: Do ăn phải thức ăn ôi thiu, nhiễm khuẩn hoặc độc tố.\n"
        "   • Bệnh táo bón: Do chế độ ăn ít chất xơ, uống ít nước, ít vận động.\n"
        "   • Bệnh viêm ruột thừa, kiết lị, trĩ...\n\n"
        "2. CÁC BIỆN PHÁP PHÒNG TRÁNH BỆNH ĐƯỜNG TIÊU HÓA:\n"
        "   • Thực hiện vệ sinh an toàn thực phẩm: Thực hiện \"ăn chín, uống sôi\", ăn thức ăn sạch, có nguồn gốc rõ ràng; không ăn thức ăn ôi thiu, nấm mốc.\n"
        "   • Vệ sinh cá nhân đúng cách: Rửa tay sạch bằng xà phòng trước khi ăn và sau khi đi vệ sinh; đánh răng ít nhất 2 lần/ngày (sau khi ăn và trước khi ngủ).\n"
        "   • Chế độ dinh dưỡng khoa học: Ăn uống điều độ, đúng giờ, ăn chậm nhai kĩ; bổ sung nhiều rau xanh, trái cây tươi (chất xơ) và uống đủ nước (1,5 - 2 lít/ngày); hạn chế đồ ăn quá cay nóng, nhiều dầu mỡ, đồ uống có cồn.\n"
        "   • Giữ tinh thần thoải mái, tránh lo âu căng thẳng kéo dài; rèn luyện thể dục thể thao đều đặn."
    )
}

# ==============================================================================
# 1. HÀM TẠO FILE WORD
# ==============================================================================

def generate_word_document(mode="student"):
    is_teacher = (mode == "teacher")
    is_screen = (mode == "screen")

    if mode == "teacher":
        filename = "De_Cuong_GK1_KHTN8_SinhHoc_GiaoVien.docx"
    elif mode == "screen":
        filename = "De_Cuong_GK1_KHTN8_SinhHoc_DayTrenManHinh.docx"
    else:
        filename = "De_Cuong_GK1_KHTN8_SinhHoc_HocSinh.docx"

    filepath = os.path.join(OUTPUT_DIR, filename)

    doc = docx.Document()
    setup_header_footer(doc)

    build_title_block(
        doc,
        school_line="TRƯỜNG TRUNG HỌC CƠ SỞ",
        exam_title="ÔN TẬP GIỮA KÌ I – KHTN 8 (PHÂN MÔN SINH HỌC)",
        subject_line="MÔN: KHOA HỌC TỰ NHIÊN 8 (KHTN3)",
        num_pages="02",
        duration="45 phút",
        exam_code="KHTN8-SH",
        atomic_masses=None,
        is_teacher=is_teacher
    )

    base_size = 11.5 if is_screen else 10.5

    # I. TRẮC NGHIỆM NHIỀU LỰA CHỌN
    add_section_header(doc, "I. TRẮC NGHIỆM NHIỀU LỰA CHỌN", "Chọn một phương án trả lời đúng nhất trong mỗi câu sau:")
    for q in MCQ_DATA:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_question_prompt(p_q, f"Câu {q['id']}. {q['content']}", base_size=base_size)

        render_options_no_table(doc, q["options"], is_teacher=is_teacher)

        if is_teacher and q.get("solution"):
            add_solution_no_table(doc, f"Đáp án đúng: *{q['answer']}*.\n{q['solution']}")
        elif is_screen:
            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.left_indent = Cm(0.6)
            p_sp.paragraph_format.space_before = Pt(2)
            p_sp.paragraph_format.space_after = Pt(2)
            render_rich_text(p_sp, "✎ Ghi chú / Giải thích: ........................................................................................................", base_size=base_size-1)

    # II. TRẮC NGHIỆM TRẢ LỜI NGẮN
    add_section_header(doc, "II. TRẮC NGHIỆM TRẢ LỜI NGẮN (0,5 ĐIỂM)", "Đọc đoạn thông tin và trả lời các câu hỏi bên dưới:")
    p_ctx_head = doc.add_paragraph()
    p_ctx_head.paragraph_format.space_before = Pt(4)
    p_ctx_head.paragraph_format.space_after = Pt(2)
    render_question_prompt(p_ctx_head, "Câu 5.", base_size=base_size)

    lines_ctx = SHORT_DATA["context"].split("\n")
    for l in lines_ctx:
        p_c = doc.add_paragraph()
        p_c.paragraph_format.left_indent = Cm(0.4)
        p_c.paragraph_format.space_before = Pt(1)
        p_c.paragraph_format.space_after = Pt(1)
        render_rich_text(p_c, l, base_size=base_size)

    for sub in SHORT_DATA["sub_questions"]:
        sub_lines = sub["prompt"].split("\n")
        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.left_indent = Cm(0.4)
        p_sub.paragraph_format.space_before = Pt(2)
        p_sub.paragraph_format.space_after = Pt(1)
        render_rich_text(p_sub, f"*{sub['id']})* {sub_lines[0]}", base_size=base_size)

        for extra_line in sub_lines[1:]:
            p_ex = doc.add_paragraph()
            p_ex.paragraph_format.left_indent = Cm(0.8)
            p_ex.paragraph_format.space_before = Pt(1)
            p_ex.paragraph_format.space_after = Pt(1)
            render_rich_text(p_ex, extra_line, base_size=base_size)

        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.left_indent = Cm(0.8)
        p_ans.paragraph_format.space_before = Pt(1)
        p_ans.paragraph_format.space_after = Pt(2)

        if is_teacher:
            render_rich_text(p_ans, f"➔ *Đáp án:* {sub['answer_short']}", base_size=base_size, base_color=COLOR_CORRECT)
            add_solution_no_table(doc, sub["solution"])
        elif is_screen:
            render_rich_text(p_ans, "➔ Trả lời: ...................................................................................................................", base_size=base_size)
        else:
            render_rich_text(p_ans, "Trả lời: .....................................................................................................................", base_size=base_size)

    # III. PHẦN TỰ LUẬN
    add_section_header(doc, "III. PHẦN TỰ LUẬN (1,0 ĐIỂM)", "Học sinh trình bày câu trả lời chi tiết vào bài làm:")
    p_tl = doc.add_paragraph()
    p_tl.paragraph_format.space_before = Pt(4)
    p_tl.paragraph_format.space_after = Pt(2)
    render_question_prompt(p_tl, f"Câu {ESSAY_QUESTION['num']}. {ESSAY_QUESTION['title']}", base_size=base_size)

    if is_teacher:
        add_solution_no_table(doc, ESSAY_QUESTION["solution"])
    elif is_screen:
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.left_indent = Cm(0.6)
        render_rich_text(p_sp, "✎ Bài làm / Phân tích chi tiết: ..............................................................................................", base_size=base_size)
    else:
        add_student_writing_lines(doc, num_lines=8)

    # Dòng kết thúc
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(12)
    p_end.paragraph_format.space_after = Pt(6)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("--- HẾT ---")
    set_run_style(r_end, font_size=11, bold=True, italic=True, color=COLOR_PRIMARY)

    doc.save(filepath)
    print(f"[OK] Đã tạo thành công file Word: {filepath} ({os.path.getsize(filepath):,} bytes)")
    return filepath


# ==============================================================================
# 2. HÀM TẠO FILE LATEX & PDF TRÌNH CHIẾU TIVI 16:9
# ==============================================================================

def generate_tivi_presentation_pdf():
    tex_filename = "De_Cuong_GK1_KHTN8_SinhHoc_DayTrenManHinh.tex"
    pdf_filename = "De_Cuong_GK1_KHTN8_SinhHoc_DayTrenManHinh.pdf"
    tex_path = os.path.join(SCRATCH_DIR, tex_filename)
    pdf_dest = os.path.join(OUTPUT_DIR, pdf_filename)

    tex_lines = []
    # Preamble chuẩn 16:9 widescreen
    tex_lines.append(r"""\documentclass{article}
\usepackage[utf8]{vietnam}
\usepackage{times}
\usepackage{amsmath,amssymb}
\usepackage{tabularx}
\usepackage{array}
\usepackage{booktabs}
\usepackage{xcolor}
\usepackage[
  paperwidth=29.7cm,
  paperheight=16.706cm,
  top=0.9cm,
  bottom=0.9cm,
  left=1.2cm,
  right=1.2cm,
  footskip=0pt,
  headheight=0pt,
  headsep=0pt
]{geometry}

\pagestyle{empty}
\setlength{\parindent}{0pt}
\setlength{\parskip}{0pt}

\definecolor{brandRed}{RGB}{211, 47, 47}
\definecolor{brandGreen}{RGB}{46, 125, 50}
\definecolor{brandBlue}{RGB}{21, 101, 192}
\definecolor{textColor}{RGB}{20, 20, 20}

\color{textColor}

\begin{document}
""")

    # --------------------------------------------------------------------------
    # SLIDE 0: TIÊU ĐỀ
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 0: TIÊU ĐỀ ĐỀ CƯƠNG
% ==============================================================================
\begin{center}
\vspace*{0.3cm}
{\fontsize{24}{30}\selectfont \textbf{\color{brandRed}TRƯỜNG TRUNG HỌC CƠ SỞ -- NĂM HỌC 2026 -- 2027}}\\[12pt]
{\fontsize{34}{42}\selectfont \textbf{\color{brandBlue}ÔN TẬP GIỮA KÌ I -- KHTN 8 (KHTN3)}}\\[8pt]
{\fontsize{28}{36}\selectfont \textbf{\color{brandGreen}PHÂN MÔN: SINH HỌC CƠ THỂ NGƯỜI}}\\[8pt]
{\fontsize{22}{28}\selectfont \textit{(Hệ hô hấp -- Hệ vận động -- Dinh dưỡng \& Hệ tiêu hóa)}}\\[18pt]
\rule{16cm}{1.5pt}\\[22pt]
{\fontsize{26}{34}\selectfont \textbf{Giáo viên hướng dẫn: Thầy TRẦN VĂN THẠNH -- 0777.470.803}}\\[8pt]
{\fontsize{22}{28}\selectfont \textit{Cơ sở 1: 6/15 Nguyễn Hoàng, Kim Long | Cơ sở 2: 24 Đặng Thái Thân, TP Huế}}\\[18pt]
{\fontsize{24}{30}\selectfont \textbf{CẤU TRÚC: TRẮC NGHIỆM ABCD -- TRẢ LỜI NGẮN -- TỰ LUẬN}}
\end{center}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 1: TRẮC NGHIỆM CÂU 1 & CÂU 2
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 1: TRẮC NGHIỆM CÂU 1 & CÂU 2
% ==============================================================================
\fontsize{26}{35}\selectfont
\textbf{Câu 1.} Trong cơ thể người, hệ cơ quan nào có chức năng thực hiện quá trình trao đổi khí giữa cơ thể với môi trường ngoài?
\vspace{4pt}

\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{A.} Hệ bài tiết. & \textbf{B.} Hệ hô hấp.\\[6pt]
\textbf{C.} Hệ tiêu hóa. & \textbf{D.} Hệ tuần hoàn.
\end{tabularx}

\vspace{16pt}

\textbf{Câu 2.} Khi tham gia các hoạt động như đi bộ, chạy, nhảy hoặc chơi thể thao, hệ vận động có chức năng chủ yếu nào sau đây?
\vspace{4pt}

\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{A.} Điều hòa lượng đường trong máu. & \textbf{B.} Vận chuyển các chất dinh dưỡng.\\[6pt]
\textbf{C.} Giúp cơ thể vận động và duy trì tư thế. & \textbf{D.} Thực hiện quá trình tiêu hóa thức ăn.
\end{tabularx}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 2: TRẮC NGHIỆM CÂU 3 & CÂU 4 (KHÔNG CÓ ĐÁP ÁN)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 2: TRẮC NGHIỆM CÂU 3 & CÂU 4 (KHÔNG CÓ ĐÁP ÁN)
% ==============================================================================
\fontsize{25}{34}\selectfont
\textbf{Câu 3.} Phát biểu nào sau đây đúng khi nói về chất dinh dưỡng?
\vspace{4pt}

\begin{itemize}
  \setlength{\itemsep}{3pt}
  \item[\textbf{A.}] Chất dinh dưỡng là những chất cơ thể không cần sử dụng.
  \item[\textbf{B.}] Chất dinh dưỡng chỉ có trong các loại thực phẩm có nguồn gốc động vật.
  \item[\textbf{C.}] Chất dinh dưỡng chỉ có vai trò cung cấp năng lượng cho cơ thể.
  \item[\textbf{D.}] Chất dinh dưỡng là những chất được cơ thể sử dụng để duy trì các hoạt động sống, sinh trưởng và phát triển.
\end{itemize}

\vspace{12pt}

\textbf{Câu 4.} Cơ quan nào sau đây thuộc ống tiêu hóa và có chức năng chứa, nhào trộn thức ăn với dịch tiêu hóa?
\vspace{4pt}

\begin{tabularx}{\linewidth}{@{}X X@{}}
\textbf{A.} Gan. & \textbf{B.} Dạ dày.\\[6pt]
\textbf{C.} Tuyến tụy. & \textbf{D.} Tuyến nước bọt.
\end{tabularx}
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 3: TRẢ LỜI NGẮN CÂU 5a (KHÔNG CÓ ĐÁP ÁN)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 3: TRẢ LỜI NGẮN CÂU 5a (KHÔNG CÓ ĐÁP ÁN)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}II. TRẮC NGHIỆM TRẢ LỜI NGẮN -- CÂU 5 (Ý a)}}
\vspace{8pt}

\fontsize{21}{28}\selectfont
\textbf{Ngữ cảnh:} Trong buổi tuyên truyền về sức khỏe học đường, giáo viên giới thiệu các bệnh, tật liên quan đến hệ vận động (cong vẹo cột sống, loãng xương, tổn thương khớp) và hướng dẫn những thói quen sinh hoạt, vận động phù hợp để bảo vệ hệ vận động.

\vspace{14pt}
\textbf{a)} Một người lớn tuổi được chẩn đoán mắc bệnh loãng xương. Bệnh này ảnh hưởng trực tiếp đến bộ phận nào của hệ vận động?\\
(\textit{Quy ước: 1. Xương; \quad 2. Cơ; \quad 3. Khớp}).

\vspace{16pt}
\textbf{Trả lời:} $\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots$
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 4: TRẢ LỜI NGẮN CÂU 5b (KHÔNG CÓ ĐÁP ÁN)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 4: TRẢ LỜI NGẮN CÂU 5b (KHÔNG CÓ ĐÁP ÁN)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}II. TRẮC NGHIỆM TRẢ LỜI NGẮN -- CÂU 5 (Ý b)}}
\vspace{8pt}

\fontsize{21}{28}\selectfont
\textbf{b)} Để bảo vệ hệ vận động và hạn chế nguy cơ mắc các bệnh, tật liên quan đến hệ vận động, có bao nhiêu biện pháp đúng trong các biện pháp sau?

\vspace{6pt}
\begin{itemize}
  \setlength{\itemsep}{4pt}
  \item (1) Ăn uống đủ chất, bổ sung thực phẩm giàu calcium và vitamin D phù hợp với nhu cầu cơ thể.
  \item (2) Thường xuyên mang vác vật nặng vượt quá sức chịu đựng của cơ thể.
  \item (3) Duy trì tư thế đúng khi ngồi học, đứng và đi lại.
  \item (4) Tập luyện thể dục thể thao thường xuyên với cường độ phù hợp.
\end{itemize}

\vspace{12pt}
\textbf{Trả lời:} $\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots\dots$
\newpage
""")

    # --------------------------------------------------------------------------
    # SLIDE 5: TỰ LUẬN CÂU 6 (KHÔNG CÓ ĐÁP ÁN - ĐỂ TRỐNG ĐỂ TƯƠNG TÁC GIẢNG BÀI)
    # --------------------------------------------------------------------------
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 5: TỰ LUẬN CÂU 6 (KHÔNG CÓ ĐÁP ÁN)
% ==============================================================================
{\fontsize{26}{32}\selectfont \textbf{\color{brandBlue}III. PHẦN TỰ LUẬN -- CÂU 6 (1,0 ĐIỂM)}}
\vspace{14pt}

\fontsize{23}{33}\selectfont
\textbf{Đề bài:}
\begin{enumerate}
  \setlength{\itemsep}{10pt}
  \item[1)] Hãy kể tên 3 bệnh ở đường tiêu hóa mà em biết.
  \item[2)] Đề xuất các biện pháp giúp phòng tránh các bệnh về đường tiêu hóa.
\end{enumerate}

\vspace*{2.5cm}
\begin{center}
{\fontsize{20}{26}\selectfont \textit{\color{gray}(Khoảng trống trình chiếu tương tác / Giáo viên và học sinh chữa bài trực tiếp trên bảng)}}
\end{center}

\end{document}
""")

    with open(tex_path, "w", encoding="utf-8") as f:
        f.write("".join(tex_lines))
    print(f"[OK] Đã ghi file LaTeX: {tex_path}")

    # Biên dịch pdflatex 2 lần để ổn định layout
    cmd = ["pdflatex", "-interaction=nonstopmode", "-output-directory", SCRATCH_DIR, tex_path]
    res1 = subprocess.run(cmd, capture_output=True, text=True)
    res2 = subprocess.run(cmd, capture_output=True, text=True)

    pdf_scratch = os.path.join(SCRATCH_DIR, pdf_filename)
    if os.path.exists(pdf_scratch):
        shutil.copy2(pdf_scratch, pdf_dest)
        print(f"[OK] Đã biên dịch và copy PDF thành công: {pdf_dest} ({os.path.getsize(pdf_dest):,} bytes)")
        return pdf_dest
    else:
        print("[ERROR] Không tìm thấy file PDF được tạo. Log:")
        print(res2.stdout[-1500:])
        return None

# ==============================================================================
# HÀM MAIN
# ==============================================================================

def main():
    print("=" * 75)
    print("BẮT ĐẦU TẠO 3 FILE CHO ĐỀ CƯƠNG ÔN TẬP GIỮA KÌ I KHTN 8 (SINH HỌC - KHTN3)")
    print("1. FILE HỌC SINH (.docx)")
    print("2. FILE GIÁO VIÊN (.docx có lời giải chi tiết)")
    print("3. FILE DẠY TRÊN MÀN HÌNH (Slide PDF 16:9 widescreen cho Tivi 65 inch)")
    print("=" * 75)

    f_hs = generate_word_document(mode="student")
    f_gv = generate_word_document(mode="teacher")
    f_screen_docx = generate_word_document(mode="screen")
    f_screen_pdf = generate_tivi_presentation_pdf()

    all_files = [f_hs, f_gv, f_screen_docx, f_screen_pdf]

    # Đồng bộ sang các thư mục OneDrive
    print("\n[+] Đồng bộ các file sang các thư mục OneDrive:")
    for od in TARGET_ONEDRIVE_DIRS:
        if os.path.exists(od):
            for src_path in all_files:
                if src_path and os.path.exists(src_path):
                    dest = os.path.join(od, os.path.basename(src_path))
                    try:
                        shutil.copy2(src_path, dest)
                        print(f"    -> Đã copy {os.path.basename(src_path)} vào: {od}")
                    except Exception as e:
                        print(f"    [!] Cảnh báo file đang mở hoặc bận tại {od}: {e}")

    print("\n" + "=" * 75)
    print("HOÀN THÀNH TẤT CẢ 3 FILE THEO ĐÚNG CHUẨN FORM ĐỀ KIỂM TRA!")
    print("=" * 75)

if __name__ == "__main__":
    main()
