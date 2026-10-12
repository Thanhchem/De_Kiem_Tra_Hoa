"""
Script tạo file Word (.docx) chuẩn "Form Đề Kiểm Tra" cho ĐỀ SỐ 23:
- Đề kiểm tra: Chương 8 - Sơ lược về dãy kim loại chuyển tiếp thứ nhất và Phức chất (Mã đề 123)
- Tạo 2 file:
  + San_Pham/De_Kiem_Tra_Hoa_12_De23.docx (Bản Học Sinh)
  + San_Pham/De_Kiem_Tra_Hoa_12_De23_GiaoVien.docx (Bản Giáo Viên)
"""

import os
import sys
import re
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# --- CÁC HÀM TIỆN ÍCH TẠO HEADER / FOOTER CHUẨN FORM ---

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.59)     # 1.5 cm
        s.bottom_margin = Inches(0.59)  # 1.5 cm
        s.left_margin = Inches(0.59)    # 1.5 cm
        s.right_margin = Inches(0.59)   # 1.5 cm
        s.header_distance = Inches(0.24) # 0.6 cm
        s.footer_distance = Inches(0.24) # 0.6 cm

        # --- HEADER ---
        header = s.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.text = ""
        hp.paragraph_format.space_after = Pt(0)
        hp.paragraph_format.space_before = Pt(0)

        tbl_h = header.add_table(rows=1, cols=2, width=Inches(7.08)) # 18.0 cm
        tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_h.autofit = False

        c0 = tbl_h.cell(0, 0)
        c1 = tbl_h.cell(0, 1)
        c0.width = Inches(3.54)
        c1.width = Inches(3.54)
        set_cell_margins(c0, 40, 40, 60, 60)
        set_cell_margins(c1, 40, 40, 60, 60)

        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run("Tài liệu lưu hành nội bộ")
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(10)
        r0.font.italic = True
        r0.font.color.rgb = RGBColor(211, 47, 47) # #D32F2F

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_after = Pt(0)
        r1_pre = p1.add_run("Thầy ")
        r1_pre.font.name = "Times New Roman"
        r1_pre.font.size = Pt(10)
        r1_pre.font.color.rgb = RGBColor(211, 47, 47)
        r1_bold = p1.add_run("TRẦN VĂN THẠNH")
        r1_bold.font.name = "Times New Roman"
        r1_bold.font.size = Pt(10)
        r1_bold.font.bold = True
        r1_bold.font.color.rgb = RGBColor(211, 47, 47)
        r1_post = p1.add_run(" - 0777.470.803")
        r1_post.font.name = "Times New Roman"
        r1_post.font.size = Pt(10)
        r1_post.font.color.rgb = RGBColor(211, 47, 47)

        # Border bottom cho Header table: #8BC390 (sz=6)
        tblPr_h = tbl_h._tbl.tblPr
        tblBorders_h = parse_xml(
            r'<w:tblBorders %s><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="8BC390"/><w:right w:val="none"/><w:insideH w:val="none"/><w:insideV w:val="none"/></w:tblBorders>' % nsdecls('w')
        )
        tblPr_h.append(tblBorders_h)

        # --- FOOTER ---
        footer = s.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.text = ""
        fp.paragraph_format.space_after = Pt(0)
        fp.paragraph_format.space_before = Pt(0)

        tbl_f = footer.add_table(rows=1, cols=2, width=Inches(7.08))
        tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_f.autofit = False

        fc0 = tbl_f.cell(0, 0)
        fc1 = tbl_f.cell(0, 1)
        fc0.width = Inches(5.31) # 13.5 cm
        fc1.width = Inches(1.77) # 4.5 cm
        set_cell_margins(fc0, 40, 40, 60, 60)
        set_cell_margins(fc1, 40, 40, 60, 60)

        fp0 = fc0.paragraphs[0]
        fp0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fp0.paragraph_format.space_after = Pt(0)
        fr0 = fp0.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
        fr0.font.name = "Times New Roman"
        fr0.font.size = Pt(9.5)
        fr0.font.italic = True
        fr0.font.color.rgb = RGBColor(211, 47, 47)

        fp1 = fc1.paragraphs[0]
        fp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp1.paragraph_format.space_after = Pt(0)
        fr1 = fp1.add_run("Trang ")
        fr1.font.name = "Times New Roman"
        fr1.font.size = Pt(10)
        fr1.font.bold = True
        fr1.font.color.rgb = RGBColor(211, 47, 47)
        # Dynamic Page Field
        fldSimple = OxmlElement('w:fldSimple')
        fldSimple.set(qn('w:instr'), 'PAGE')
        fp1._p.append(fldSimple)

        # Border top cho Footer table: #8BC390 (sz=6)
        tblPr_f = tbl_f._tbl.tblPr
        tblBorders_f = parse_xml(
            r'<w:tblBorders %s><w:top w:val="single" w:sz="6" w:space="0" w:color="8BC390"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/><w:insideH w:val="none"/><w:insideV w:val="none"/></w:tblBorders>' % nsdecls('w')
        )
        tblPr_f.append(tblBorders_f)

def add_title_block(doc, is_teacher=False):
    lines_info = [
        ("SỞ GD&ĐT THỪA THIÊN HUẾ", "ĐỀ KIỂM TRA ĐỊNH KỲ CHƯƠNG 8 -- HOÁ HỌC 12", True, True, RGBColor(21, 101, 192)),
        ("TRƯỜNG THPT CHUYÊN QUỐC HỌC", "MÔN: HOÁ HỌC -- LỚP 12 (NĂM HỌC 2026 - 2027)", True, True, None),
        ("ĐỀ CHÍNH THỨC", "Thời gian làm bài: 50 phút (không kể thời gian phát đề)", True, False, None),
        ("(Đề thi gồm 04 trang)", "Mã đề thi: 123", False, True, None)
    ]

    for left_txt, right_txt, left_bold, right_bold, right_color in lines_info:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.08), WD_TAB_ALIGNMENT.RIGHT)

        r_l = p.add_run(left_txt)
        r_l.font.name = "Times New Roman"
        r_l.font.size = Pt(10.5)
        r_l.bold = left_bold
        if not left_bold:
            r_l.italic = True

        p.add_run("\t")

        r_r = p.add_run(right_txt)
        r_r.font.name = "Times New Roman"
        r_r.font.size = Pt(10.5)
        r_r.bold = right_bold
        if right_color:
            r_r.font.color.rgb = right_color
        if not right_bold:
            r_r.italic = True

    # Student info line
    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(3)
    p_info.paragraph_format.space_after = Pt(3)
    p_info.paragraph_format.line_spacing = 1.15
    r_info = p_info.add_run("Họ, tên thí sinh: ..................................................................... Lớp: ................. SBD: .................")
    r_info.font.name = "Times New Roman"
    r_info.font.size = Pt(10)
    r_info.italic = True

    # Atomic masses
    p_mass = doc.add_paragraph()
    p_mass.paragraph_format.space_before = Pt(0)
    p_mass.paragraph_format.space_after = Pt(6)
    p_mass.paragraph_format.line_spacing = 1.15
    r_mass = p_mass.add_run("(Cho biết nguyên tử khối: H = 1; C = 12; N = 14; O = 16; Na = 23; Mg = 24; Al = 27; S = 32; Cl = 35,5; K = 39; Ca = 40; Cr = 52; Mn = 55; Fe = 56; Cu = 64; Zn = 65; Ag = 108; Ba = 137; Pt = 195).")
    r_mass.font.name = "Times New Roman"
    r_mass.font.size = Pt(9.5)
    r_mass.italic = True

    if is_teacher:
        p_tea = doc.add_paragraph()
        p_tea.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_tea.paragraph_format.space_before = Pt(2)
        p_tea.paragraph_format.space_after = Pt(6)
        r_tea = p_tea.add_run("(BẢN GIÁO VIÊN -- CÓ ĐÁP ÁN VÀ HƯỚNG DẪN GIẢI CHI TIẾT)")
        r_tea.font.name = "Times New Roman"
        r_tea.font.size = Pt(11)
        r_tea.bold = True
        r_tea.font.color.rgb = RGBColor(211, 47, 47)

def format_chemical_runs(paragraph, text, is_bold=False, is_italic=False, color_rgb=None):
    """
    Tách chuỗi chứa công thức hóa học có dấu _ (subscript) hoặc ^ (superscript).
    Ví dụ: H_2O, Fe^3+, SO_4^2-, v.v.
    """
    pattern = r'(\^\{[^\}]+\}|\^[0-9\+\-]+|\_\{[^\}]+\}|\_[0-9a-zA-Z\+\-]+)'
    tokens = re.split(pattern, text)
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
        run.font.bold = is_bold
        run.font.italic = is_italic
        if color_rgb:
            run.font.color.rgb = color_rgb

        if token.startswith('^'):
            val = token[2:-1] if token.startswith('^{') else token[1:]
            run.text = val
            run.font.superscript = True
        elif token.startswith('_'):
            val = token[2:-1] if token.startswith('_{') else token[1:]
            run.text = val
            run.font.subscript = True
        else:
            run.text = token

def add_part_heading(doc, part_title, part_desc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    r1 = p.add_run(part_title + " ")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(180, 20, 20)

    r2 = p.add_run(part_desc)
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(11)
    r2.font.italic = True

def add_mcq_options(doc, options, correct_opt=None, is_teacher=False):
    """
    Bố trí 4 phương án A, B, C, D bằng đoạn văn bản thuần túy (NO TABLE) với Tab Stops:
    4.5 cm (1.77 in), 9.0 cm (3.54 in), 13.5 cm (5.31 in).
    """
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.tab_stops.add_tab_stop(Inches(1.77), WD_TAB_ALIGNMENT.LEFT)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(3.54), WD_TAB_ALIGNMENT.LEFT)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(5.31), WD_TAB_ALIGNMENT.LEFT)

    keys = ['A', 'B', 'C', 'D']
    for idx, k in enumerate(keys):
        opt_text = options[idx]
        if idx > 0:
            p.add_run("\t")
        is_correct = (is_teacher and correct_opt == k)
        color = RGBColor(211, 47, 47) if is_correct else None

        r_lbl = p.add_run(f"{k}. ")
        r_lbl.font.name = "Times New Roman"
        r_lbl.font.size = Pt(11)
        r_lbl.font.bold = True
        if is_correct:
            r_lbl.font.color.rgb = color

        format_chemical_runs(p, opt_text, is_bold=is_correct, color_rgb=color)

def add_solution(doc, sol_text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.24) # 0.6 cm
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15

    r_hdr = p.add_run("Lời giải chi tiết: ")
    r_hdr.font.name = "Times New Roman"
    r_hdr.font.size = Pt(10.5)
    r_hdr.font.bold = True
    r_hdr.font.color.rgb = RGBColor(21, 101, 192) # #1565C0

    format_chemical_runs(p, sol_text, is_italic=True)

# --- DỮ LIỆU ĐỀ KIỂM TRA SỐ 23 ---

PART1_QUESTIONS = [
    {
        "num": 1,
        "text": "[KNTT - SBT] Sự hình thành các nguyên tố chuyển tiếp dãy thứ nhất là do có sự sắp xếp lần lượt các electron vào phân lớp nào sau đây?",
        "options": ["3d.", "4s.", "4p.", "3p."],
        "key": "A",
        "sol": "Các nguyên tố chuyển tiếp dãy thứ nhất (từ Sc đến Cu) có cấu hình electron với sự điền dần các electron vào phân lớp 3d."
    },
    {
        "num": 2,
        "text": "[KNTT - SBT] Ion nào sau đây không có electron trên phân lớp 3d và không có màu trong dung dịch nước?",
        "options": ["Fe^3+.", "Cr^3+.", "Ti^3+.", "Sc^3+."],
        "key": "D",
        "sol": "Sc (Z=21) có cấu hình [Ar]3d^1 4s^2. Khi tạo ion Sc^3+, nguyên tử nhường 3 electron tạo cấu hình khí hiếm [Ar] (không có electron 3d) nên ion Sc^3+ không màu trong nước."
    },
    {
        "num": 3,
        "text": "[KNTT - SBT] Oxide nào sau đây có màu trắng?",
        "options": ["Fe_2O_3.", "Cr_2O_3.", "Al_2O_3.", "CuO."],
        "key": "C",
        "sol": "Al_2O_3 là oxide màu trắng. Fe_2O_3 màu nâu đỏ, Cr_2O_3 màu lục thẫm, CuO màu đen."
    },
    {
        "num": 4,
        "text": "[CTST - SBT] Trong thí nghiệm xác định hàm lượng muối Fe(II) bằng dung dịch thuốc tím (KMnO_4) trong môi trường acid, nhận xét nào sau đây là không đúng?",
        "options": [
            "Ion Fe^2+ là chất bị oxi hóa.",
            "H_2SO_4 là chất tạo môi trường phản ứng.",
            "Ion MnO_4^- là chất bị khử.",
            "Dung dịch muối Fe(II) có màu vàng nhạt."
        ],
        "key": "D",
        "sol": "Dung dịch muối Fe(II) tinh khiết hầu như không màu hoặc có màu xanh rất nhạt. Dung dịch muối Fe(III) mới có màu vàng nâu/vàng nhạt."
    },
    {
        "num": 5,
        "text": "Để trở thành phối tử trong phức chất thì phân tử hoặc anion cần phải có",
        "options": [
            "các orbital trống.",
            "cặp electron hoá trị riêng.",
            "ít nhất 4 orbital trống.",
            "ít nhất hai cặp electron hoá trị riêng."
        ],
        "key": "B",
        "sol": "Phối tử là các phân tử hoặc anion có chứa cặp electron hoá trị tự do (chưa liên kết) để cho vào các orbital trống của nguyên tử/ion trung tâm tạo liên kết phối trí."
    },
    {
        "num": 6,
        "text": "Trong phức chất [Co(H_2O)_6]Cl_3, số phối tử liên kết với nguyên tử trung tâm là",
        "options": ["2.", "3.", "9.", "6."],
        "key": "D",
        "sol": "Cầu nội của phức chất là [Co(H_2O)_6]^3+, gồm nguyên tử trung tâm Co^3+ liên kết với 6 phân tử phối tử H_2O."
    },
    {
        "num": 7,
        "text": "Trong ion phức [Zn(NH_3)_4]^2+, ion phức có dạng hình học nào sau đây?",
        "options": ["Hình vuông phẳng.", "Hình lập phương.", "Hình bát diện.", "Hình tứ diện."],
        "key": "D",
        "sol": "Ion Zn^2+ có cấu hình d^10 bền vững, tham gia lai hoá sp^3 tạo cấu trúc dạng hình học tứ diện với 4 phối tử NH_3."
    },
    {
        "num": 8,
        "text": "Phức chất nào sau đây có cấu trúc hình học tứ diện?",
        "options": ["[Pt(NH_3)_4]^2+.", "[Zn(NH_3)_4]^2+.", "[Fe(H_2O)_6]^2+.", "[Co(H_2O)_6]^3+."],
        "key": "B",
        "sol": "[Zn(NH_3)_4]^2+ có số phối trí 4 và có cấu trúc tứ diện. [Pt(NH_3)_4]^2+ có cấu trúc vuông phẳng. Các phức chất [Fe(H_2O)_6]^2+ và [Co(H_2O)_6]^3+ có dạng bát diện."
    },
    {
        "num": 9,
        "text": "Chọn phát biểu sai trong các phát biểu sau:",
        "options": [
            "Phức chất có số phối trí 4, nguyên tử trung tâm lai hoá sp^3 thì có dạng hình học tứ diện.",
            "Phức chất có số phối trí 4, nguyên tử trung tâm lai hoá dsp^2 thì có dạng vuông phẳng.",
            "Phức chất có số phối tử là 6, nguyên tử trung tâm lai hoá d^2sp^3 hoặc sp^3d^2 thì có dạng bát diện.",
            "Phức chất có số phối trí là 4 thì thường có dạng hình học vuông phẳng."
        ],
        "key": "D",
        "sol": "Phát biểu D sai vì phức chất có số phối trí 4 có hai dạng hình học phổ biến là tứ diện (rất phổ biến với ion d^10 như Zn^2+) và vuông phẳng (phổ biến với Pt^2+, Pd^2+, Ni^2+), không thể khẳng định là thường có dạng vuông phẳng."
    },
    {
        "num": 10,
        "text": "Nhỏ vài giọt dung dịch NaCl vào ống nghiệm chứa dung dịch AgNO_3, thu được kết tủa trắng AgCl. Tiếp tục nhỏ dung dịch NH_3 đến dư thấy kết tủa tan, chứng tỏ phức chất [Ag(NH_3)_2]^+ đã được tạo thành. Dấu hiệu nhận biết của phản ứng tạo phức trên dựa vào",
        "options": ["sự ngưng tụ.", "sự đổi màu sắc.", "sự hoà tan kết tủa.", "sự thay đổi nhiệt độ sôi."],
        "key": "C",
        "sol": "Dấu hiệu rõ rệt của phản ứng tạo phức [Ag(NH_3)_2]^+ là sự hoà tan của kết tủa trắng AgCl trong dung dịch NH_3 dư."
    },
    {
        "num": 11,
        "text": "Nhỏ vài giọt dung dịch NaOH vào ống nghiệm chứa dung dịch AlCl_3, thấy xuất hiện kết tủa trắng keo, chứng tỏ phức chất [Al(OH)_3(H_2O)_3] đã tạo thành. Dấu hiệu nhận biết của phản ứng tạo phức trên dựa vào",
        "options": ["sự ngưng tụ.", "sự kết tủa.", "sự bay hơi.", "sự thay đổi nhiệt độ sôi."],
        "key": "B",
        "sol": "Dấu hiệu nhận biết là sự xuất hiện kết tủa trắng keo [Al(OH)_3(H_2O)_3] trong dung dịch."
    },
    {
        "num": 12,
        "text": "Hoàn thành phương trình phản ứng sau: NiCl_2(s) + ? -> [Ni(OH_2)_6]^2+(aq) + ?",
        "options": [
            "NiCl_2(s) + 6H_2O(l) -> [Ni(OH_2)_6]^2+(aq) + 2Cl^-(aq).",
            "NiCl_2(s) + 2H_2O(l) -> [Ni(OH_2)_6]^2+(aq) + 2Cl^-(aq).",
            "NiCl_2(s) + 6H_2O(l) -> [Ni(OH_2)_6]^2+(s) + 2Cl^-(aq).",
            "NiCl_2(s) + 3H_2O(l) -> [Ni(OH_2)_6]^2+(aq) + Cl^-(aq)."
        ],
        "key": "A",
        "sol": "Quá trình hoà tan và hydrate hoá muối NiCl_2: NiCl_2(s) + 6H_2O(l) -> [Ni(OH_2)_6]^2+(aq) + 2Cl^-(aq)."
    },
    {
        "num": 13,
        "text": "Phức chất của Fe^2+ cấu tạo nên nhân heme, thành phần của hemoglobin trong hồng cầu của máu. Trong hệ thống tuần hoàn, heme có vai trò chính trong việc vận chuyển và tích luỹ oxygen trong cơ thể. Nguyên tử/ion trung tâm trong phức chất này là",
        "options": ["Cl^-.", "Fe^2+.", "OH^-.", "NH_3."],
        "key": "B",
        "sol": "Trong nhân heme, ion Fe^2+ đóng vai trò là nguyên tử/ion trung tâm phối trí với vòng porphyrin."
    },
    {
        "num": 14,
        "text": "Phân tử phức chất cis-dichloro(ethylenediamine)platinum(II) có công thức cấu tạo như hình bên:\nCó bao nhiêu loại phối tử có trong phân tử phức chất đó?",
        "options": ["1.", "2.", "3.", "4."],
        "image": "San_Pham/images_crop_de23/p1_cau14_cis_pt_en.png",
        "key": "B",
        "sol": "Phức chất [Pt(en)Cl_2] có 2 loại phối tử: 2 phối tử chloride Cl^- (đơn càng) và 1 phối tử ethylenediamine H_2N-CH_2-CH_2-NH_2 (hai càng)."
    },
    {
        "num": 15,
        "text": "EDTA là một phối tử sáu càng có thể tạo phức bát diện bền với hầu hết các cation kim loại theo tỉ lệ mol phối tử : cation là 1 : 1. Số nguyên tử nitrogen có trong phân tử phối tử EDTA tham gia tạo liên kết phối trí với ion kim loại là",
        "options": ["2.", "4.", "6.", "8."],
        "key": "A",
        "sol": "Phối tử EDTA (ethylenediaminetetraacetate) có 6 vị trí phối trí gồm 2 nguyên tử nitrogen và 4 nguyên tử oxygen từ các nhóm carboxylate."
    },
    {
        "num": 16,
        "text": "Cho sơ đồ chuyển hoá sau:\nNhững phản ứng có sự tạo thành phức chất tan là",
        "options": [
            "(1) và (2).",
            "(1) và (3).",
            "(2) và (3).",
            "Cả (1), (2) và (3)."
        ],
        "image": "San_Pham/images_crop_de23/p1_cau16_so_do.png",
        "key": "B",
        "sol": "Phản ứng (1) tạo phức aqua tan [Cu(OH_2)_6]^2+ (màu xanh). Phản ứng (2) tạo kết tủa hydroxide [Cu(OH_2)_4(OH)_2] (màu xanh nhạt). Phản ứng (3) kết tủa tan trong NH_3 tạo dung dịch phức chất tan [Cu(NH_3)_4(OH_2)_2]^2+ (màu xanh lam). Do đó các phản ứng tạo phức tan là (1) và (3)."
    },
    {
        "num": 17,
        "text": "Cho sơ đồ chuyển hoá của bạc (Ag) sau:\nNhững phản ứng có sự tạo thành phức chất tan là",
        "options": [
            "(1), (2), (3).",
            "(1), (3), (4).",
            "(2), (3), (4).",
            "(1), (2), (4)."
        ],
        "image": "San_Pham/images_crop_de23/p1_cau17_so_do.png",
        "key": "B",
        "sol": "Phản ứng (1) tạo phức aqua tan [Ag(OH_2)_4]^+ (không màu). Phản ứng (2) tạo kết tủa đen Ag_2O. Phản ứng (3) kết tủa tan tạo phức diammine bạc [Ag(NH_3)_2]^+ (không màu). Phản ứng (4) chuyển thành phức cyano [Ag(CN)_2]^- bền hơn. Vậy các phản ứng tạo phức tan là (1), (3), (4)."
    },
    {
        "num": 18,
        "text": "Cho cấu hình electron nguyên tử của các nguyên tố: (X): [Ar]3d^2 4s^2; (Y): [Ar]4s^1; (Z): [Ar]3d^{10} 4s^2; (T): [Ar]3d^{10} 4s^2 4p^3. Trong các nhận định sau:\n(a) Nguyên tố (X) không phải là kim loại chuyển tiếp.\n(b) Nguyên tố (Y) là kim loại chuyển tiếp.\n(c) Nguyên tố (T) là kim loại chuyển tiếp.\n(d) Nguyên tố (Z) không phải là kim loại chuyển tiếp (theo định nghĩa IUPAC).\nSố nhận định đúng là",
        "options": ["2.", "1.", "3.", "4."],
        "key": "B",
        "sol": "• (X) có electron đang điền vào 3d (Ti) -> là kim loại chuyển tiếp -> (a) sai.\n• (Y) là K (kim loại nhóm IA) -> (b) sai.\n• (T) là As (phi kim nhóm VA) -> (c) sai.\n• (Z) là Zn có phân lớp d bão hoà (3d^{10}) ở cả nguyên tử và cation Zn^2+ -> theo IUPAC Zn không phải là kim loại chuyển tiếp -> (d) đúng.\nVậy có 1 nhận định đúng là (d)."
    }
]

PART2_QUESTIONS = [
    {
        "num": 1,
        "intro": "[CĐ – SBT] Cho các thông tin về thế điện cực chuẩn của các cặp oxi hoá – khử sau:",
        "table_data": [
            ("Cặp oxi hoá – khử", "Thế điện cực chuẩn (V)"),
            ("Fe^3+/Fe^2+", "+0,77"),
            ("Cr_2O_7^2- + 14H^+ / 2Cr^3+ + 7H_2O", "+1,33"),
            ("MnO_4^- + 8H^+ / Mn^2+ + 4H_2O", "+1,51")
        ],
        "note": "Biết Cr_2O_7^2-(aq) có màu da cam và Cr^3+(aq) có màu xanh lá cây.",
        "items": [
            ("a", "Trong môi trường acid, anion Cr_2O_7^2- (từ sự phân li của K_2Cr_2O_7) có tính oxi hoá mạnh hơn anion MnO_4^- (từ KMnO_4).", False,
             "Sai vì E°(MnO_4^-/Mn^2+) = +1,51 V > E°(Cr_2O_7^2-/2Cr^3+) = +1,33 V nên MnO_4^- có tính oxi hoá mạnh hơn Cr_2O_7^2-."),
            ("b", "Chuẩn độ được Fe^2+ trong dung dịch gồm Fe^2+, SO_4^2- và H^+ bằng dung dịch chất chuẩn KMnO_4.", True,
             "Đúng vì MnO_4^- oxi hoá hoàn toàn Fe^2+ lên Fe^3+ trong môi trường H_2SO_4 loãng và phản ứng tự chỉ thị điểm kết thúc bằng màu tím nhạt của MnO_4^- dư."),
            ("c", "Không chuẩn độ được Fe^2+ trong dung dịch gồm Fe^2+, SO_4^2- và H^+ bằng dung dịch chất chuẩn K_2Cr_2O_7.", False,
             "Sai vì E°(Cr_2O_7^2-/2Cr^3+) = 1,33 V > E°(Fe^3+/Fe^2+) = 0,77 V nên hoàn toàn chuẩn độ được Fe^2+ bằng dung dịch chuẩn K_2Cr_2O_7 (dùng chỉ thị diphenylamine)."),
            ("d", "Có diễn ra phản ứng oxi hoá – khử tự diễn biến theo phương trình: 6Fe^3+(aq) + 2Cr^3+(aq) + 7H_2O(l) -> 6Fe^2+(aq) + Cr_2O_7^2-(aq) + 14H^+(aq).", False,
             "Sai vì phản ứng tự phát phải xảy ra theo chiều chất oxi hoá mạnh hơn tác dụng chất khử mạnh hơn: Cr_2O_7^2- oxi hoá Fe^2+ thành Fe^3+ và Cr^3+.")
        ]
    },
    {
        "num": 2,
        "intro": "[CTST - SBT] Khi tìm hiểu về hợp chất phức chlorophyll (chất diệp lục), xét tính đúng/sai của các phát biểu sau:",
        "items": [
            ("a", "Chlorophyll chỉ được tìm thấy trong tế bào thực vật.", False,
             "Sai vì chlorophyll còn được tìm thấy trong các sinh vật quang dưỡng khác như vi khuẩn lam, tảo và vi khuẩn lưu huỳnh màu lục."),
            ("b", "Chlorophyll đóng vai trò hấp thụ năng lượng ánh sáng và xúc tác cho quá trình quang hợp.", True,
             "Đúng vì chlorophyll hấp thụ photon ánh sáng kích hoạt electron để biến đổi CO_2 và H_2O thành carbohydrate."),
            ("c", "Trong chất diệp lục chỉ chứa duy nhất một loại chlorophyll.", False,
             "Sai vì có nhiều loại chlorophyll khác nhau như chlorophyll a, chlorophyll b, chlorophyll c, chlorophyll d..."),
            ("d", "Chlorophyll là thành phần sắc tố quan trọng nhất trong lục lạp của lá cây.", True,
             "Đúng vì chlorophyll là sắc tố chính tạo nên màu xanh và thực hiện chức năng quang hợp chủ yếu của lá.")
        ]
    },
    {
        "num": 3,
        "intro": "[KNTT - SBT] Nhỏ muối thiocyanate (SCN^-) vào dung dịch muối Fe^3+ loãng, dung dịch từ màu vàng nhạt chuyển sang màu đỏ máu do một phối tử nước trong phức chất aqua bát diện của Fe^3+ bị thay thế bởi một phối tử SCN^-:",
        "items": [
            ("a", "Phức chất aqua của sắt(III) có công thức hoá học là [Fe(H_2O)_6]^3+.", True,
             "Đúng vì ion Fe^3+ trong nước tồn tại dưới dạng phức chất bát diện hexaaquairon(III) [Fe(H_2O)_6]^3+."),
            ("b", "Phức chất có màu đỏ máu là phức chất của Fe^3+ chứa 1 phối tử SCN^- và 6 phối tử nước.", False,
             "Sai vì khi 1 phối tử SCN^- thay thế 1 phối tử H_2O thì phức tạo thành chỉ còn 5 phối tử H_2O (tổng số phối trí là 6)."),
            ("c", "Phức chất màu đỏ máu có công thức hoá học là [Fe(H_2O)_5(SCN)]^2+.", True,
             "Đúng vì cation Fe^3+ (+3) kết hợp với 1 phối tử SCN^- (-1) và 5 phân tử H_2O (trung hoà) tạo cation phức có điện tích +2: [Fe(H_2O)_5(SCN)]^2+."),
            ("d", "Phức chất màu đỏ máu có điện tích là +3.", False,
             "Sai vì điện tích của ion phức là (+3) + (-1) = +2.")
        ]
    },
    {
        "num": 4,
        "intro": "Hemoglobin (Hb) là một loại metalloprotein chiếm khoảng 35% trọng lượng của hồng cầu, có nhiệm vụ vận chuyển oxygen từ phổi đến các mô. Cấu trúc gồm 4 tiểu đơn vị, mỗi tiểu đơn vị chứa một nhân heme (chiếm 5% khối lượng Hb) và một chuỗi polypeptide globin. Mỗi nhân heme chứa một ion Fe^2+ ở trung tâm vòng porphyrin và chỉ liên kết tối đa với 1 phân tử O_2. Đánh giá tính đúng/sai của các phát biểu sau:",
        "items": [
            ("a", "Dạng hình học của phức chất Fe(II) trong nhân heme khi liên kết với oxygen là bát diện.", True,
             "Đúng vì ion Fe^2+ liên kết với 4 nguyên tử N của porphyrin, 1 nguyên tử N của histidine và 1 phân tử O_2 tạo số phối trí 6 (dạng bát diện)."),
            ("b", "Mỗi phân tử hemoglobin có thể kết hợp tối đa với bốn phân tử oxygen.", True,
             "Đúng vì mỗi phân tử Hb có 4 nhân heme, mỗi nhân heme kết hợp được với 1 phân tử O_2."),
            ("c", "Trong các nhân heme, nguyên tử trung tâm Fe^2+ liên kết với phối tử qua các nguyên tử nitrogen.", True,
             "Đúng vì vòng porphyrin phối trí với Fe^2+ qua 4 nguyên tử N của các vòng pyrrole và liên kết với chuỗi globin qua nguyên tử N của gốc histidine."),
            ("d", "Trong phân tử hemoglobin, nguyên tố iron (Fe) chiếm 5% theo khối lượng.", False,
             "Sai vì bốn nhân heme chiếm 5% khối lượng Hb, nhưng trong nhân heme Fe chỉ chiếm một phần khối lượng, nên hàm lượng sắt trong hemoglobin chỉ khoảng 0,34% theo khối lượng.")
        ]
    }
]

PART3_QUESTIONS = [
    {
        "num": 1,
        "text": "Số phối tử không mang điện tích trong phức chất [Co(NH_3)_5Cl]Cl_2 là bao nhiêu?",
        "ans": "5",
        "sol": "Trong cầu nội phức chất [Co(NH_3)_5Cl]^2+, các phối tử gồm 5 phân tử NH_3 (phối tử trung hoà, không mang điện) và 1 ion Cl^- (phối tử mang điện tích âm -1). Vậy số phối tử không mang điện là 5."
    },
    {
        "num": 2,
        "text": "Số nguyên tử kim loại ở cầu ngoại trong một phân tử phức chất Na_2[Pt(CN)_4] là bao nhiêu?",
        "ans": "2",
        "sol": "Cầu ngoại của phức chất Na_2[Pt(CN)_4] gồm 2 cation Na^+, tức có 2 nguyên tử kim loại natri ở cầu ngoại."
    },
    {
        "num": 3,
        "text": "Trong dung dịch nước, ion Co^2+ tạo thành phức chất aqua có dạng hình học bát diện [Co(H_2O)_6]^2+. Số phối tử liên kết với nguyên tử trung tâm trong ion phức này là bao nhiêu?",
        "ans": "6",
        "sol": "Ion phức aqua dạng bát diện là [Co(H_2O)_6]^2+, có 6 phối tử phân tử nước H_2O liên kết với cation trung tâm Co^2+."
    },
    {
        "num": 4,
        "text": "Cho các phức chất và ion phức sau: [Co(NH_3)_6]^3+, [Cr(H_2O)_6]^3+, [Ni(CO)_4], [BF_4]^-, [Fe(CN)_6]^4-, [Ag(NH_3)_2]^+. Có bao nhiêu phức anion (phức chất mang điện tích âm) trong số các chất trên?",
        "ans": "2",
        "sol": "Các phức anion là các phức mang điện tích âm, gồm: [BF_4]^- (điện tích -1) và [Fe(CN)_6]^4- (điện tích -4). Tổng cộng có 2 phức anion."
    },
    {
        "num": 5,
        "text": "Cho 16,0 gam bột CuO tan hết trong dung dịch H_2SO_4 20% vừa đủ, sau đó làm nguội dung dịch thu được đến 10°C thấy tách ra m gam tinh thể CuSO_4.5H_2O. Biết độ tan của CuSO_4 ở 10°C là 17,4 gam trong 100 gam nước. Giá trị của m là bao nhiêu? (Làm tròn kết quả đến hàng phần mười).",
        "ans": "30,7",
        "sol": "• n(CuO) = 16,0 / 80 = 0,2 mol.\n• Phương trình phản ứng: CuO + H_2SO_4 -> CuSO_4 + H_2O\n• n(H_2SO_4) = 0,2 mol -> m(dd H_2SO_4 20%) = 0,2 * 98 / 0,20 = 98 gam.\n• Khối lượng dung dịch sau phản ứng = 16,0 + 98 = 114 gam.\n• n(CuSO_4 tạo thành) = 0,2 mol -> m(CuSO_4) = 0,2 * 160 = 32 gam.\n• Gọi số mol CuSO_4.5H_2O tách ra là x mol (m = 250x gam).\n  Khối lượng CuSO_4 còn lại trong dung dịch = 32 - 160x (gam).\n  Khối lượng H_2O còn lại trong dung dịch = (114 - 32) - 5*18x = 82 - 90x (gam).\n• Ở 10°C: Độ tan S = 17,4 gam / 100 gam H_2O -> (32 - 160x) / (82 - 90x) = 17,4 / 100\n  <=> 3200 - 16000x = 1426,8 - 1566x <=> 14434x = 1773,2 => x ≈ 0,12285 mol.\n• Khối lượng tinh thể m = 250 * 0,12285 ≈ 30,71 gam -> làm tròn 30,7 gam."
    },
    {
        "num": 6,
        "text": "Đem 9,80 gam muối Mohr FeSO_4.(NH_4)_2SO_4.6H_2O hoà tan trong nước cất rồi định mức thành 250 mL dung dịch trong bình định mức với sự có mặt của H_2SO_4 loãng. Tiến hành chuẩn độ 25,0 mL dung dịch này bằng dung dịch KMnO_4 chưa rõ nồng độ, thể tích dung dịch KMnO_4 tiêu tốn trung bình qua 3 lần chuẩn độ là 24,65 mL. Nồng độ mol/L của dung dịch KMnO_4 là bao nhiêu? (Làm tròn đến hai chữ số thập phân).",
        "ans": "0,02",
        "sol": "• Khối lượng mol muối Mohr: M = 56 + 96 + 2*18 + 96 + 6*18 = 392 g/mol.\n• Số mol Fe^2+ trong 250 mL: n = 9,80 / 392 = 0,025 mol = 25 mmol.\n• Trong 25,0 mL dung dịch lấy ra chuẩn độ có: n(Fe^2+) = 25 / 10 = 2,5 mmol.\n• Phương trình chuẩn độ:\n  5Fe^2+ + MnO_4^- + 8H^+ -> 5Fe^3+ + Mn^2+ + 4H_2O\n• n(KMnO_4) = n(Fe^2+) / 5 = 2,5 / 5 = 0,50 mmol.\n• C_M(KMnO_4) = 0,50 mmol / 24,65 mL ≈ 0,02028 mol/L -> làm tròn 0,02 M."
    }
]

# --- HÀM TẠO FILE DOCX CHÍNH ---

def generate_exam_docx(is_teacher=False, output_path=""):
    doc = Document()
    add_header_footer(doc)
    add_title_block(doc, is_teacher=is_teacher)

    # ==================== PHẦN I ====================
    add_part_heading(doc, "PHẦN I.", "Câu trắc nghiệm nhiều phương án lựa chọn. Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án.")

    for q in PART1_QUESTIONS:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        r_lbl = p_q.add_run(f"Câu {q['num']}: ")
        r_lbl.font.name = "Times New Roman"
        r_lbl.font.size = Pt(11)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = RGBColor(21, 101, 192)

        format_chemical_runs(p_q, q["text"])

        # Nếu có hình ảnh
        if "image" in q and os.path.exists(q["image"]):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(2)
            p_img.paragraph_format.space_after = Pt(3)
            # Tùy kích thước ảnh
            if q["num"] == 14:
                p_img.add_run().add_picture(q["image"], width=Inches(2.5))
            else:
                p_img.add_run().add_picture(q["image"], width=Inches(5.0))

        add_mcq_options(doc, q["options"], correct_opt=q["key"], is_teacher=is_teacher)

        if is_teacher and "sol" in q:
            add_solution(doc, q["sol"])

    # ==================== PHẦN II ====================
    add_part_heading(doc, "PHẦN II.", "Câu trắc nghiệm đúng sai. Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.")

    for q in PART2_QUESTIONS:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        r_lbl = p_q.add_run(f"Câu {q['num']}: ")
        r_lbl.font.name = "Times New Roman"
        r_lbl.font.size = Pt(11)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = RGBColor(21, 101, 192)

        format_chemical_runs(p_q, q["intro"])

        # Bảng dữ liệu thế điện cực nếu có ở Câu 1
        if "table_data" in q:
            tbl = doc.add_table(rows=len(q["table_data"]), cols=2)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit = False
            for r_idx, (c0_text, c1_text) in enumerate(q["table_data"]):
                row = tbl.rows[r_idx]
                cell0 = row.cells[0]
                cell1 = row.cells[1]
                cell0.width = Inches(4.5)
                cell1.width = Inches(2.2)
                set_cell_margins(cell0, 30, 30, 60, 60)
                set_cell_margins(cell1, 30, 30, 60, 60)

                p0 = cell0.paragraphs[0]
                p0.paragraph_format.space_before = Pt(1)
                p0.paragraph_format.space_after = Pt(1)
                p0.paragraph_format.line_spacing = 1.15
                format_chemical_runs(p0, c0_text, is_bold=(r_idx==0))

                p1 = cell1.paragraphs[0]
                p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p1.paragraph_format.space_before = Pt(1)
                p1.paragraph_format.space_after = Pt(1)
                p1.paragraph_format.line_spacing = 1.15
                format_chemical_runs(p1, c1_text, is_bold=(r_idx==0))

            # Border thanh lịch cho bảng
            tblPr = tbl._tbl.tblPr
            tblBorders = parse_xml(
                r'<w:tblBorders %s><w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/><w:left w:val="none"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/><w:right w:val="none"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/><w:insideV w:val="none"/></w:tblBorders>' % nsdecls('w')
            )
            tblPr.append(tblBorders)

        if "note" in q:
            p_n = doc.add_paragraph()
            p_n.paragraph_format.space_before = Pt(2)
            p_n.paragraph_format.space_after = Pt(2)
            p_n.paragraph_format.line_spacing = 1.15
            format_chemical_runs(p_n, q["note"], is_italic=True)

        for item_lbl, item_text, is_true, item_sol in q["items"]:
            p_it = doc.add_paragraph()
            p_it.paragraph_format.left_indent = Inches(0.24)
            p_it.paragraph_format.space_before = Pt(1)
            p_it.paragraph_format.space_after = Pt(1)
            p_it.paragraph_format.line_spacing = 1.15

            r_it_lbl = p_it.add_run(f"{item_lbl}) ")
            r_it_lbl.font.name = "Times New Roman"
            r_it_lbl.font.size = Pt(11)
            r_it_lbl.font.bold = True

            format_chemical_runs(p_it, item_text)

            if is_teacher:
                status_str = " (ĐÚNG)" if is_true else " (SAI)"
                r_st = p_it.add_run(status_str)
                r_st.font.name = "Times New Roman"
                r_st.font.size = Pt(11)
                r_st.font.bold = True
                r_st.font.color.rgb = RGBColor(46, 125, 50) if is_true else RGBColor(211, 47, 47)

                p_sol = doc.add_paragraph()
                p_sol.paragraph_format.left_indent = Inches(0.48)
                p_sol.paragraph_format.space_before = Pt(0)
                p_sol.paragraph_format.space_after = Pt(2)
                p_sol.paragraph_format.line_spacing = 1.15
                r_s_hdr = p_sol.add_run("• Giải thích: ")
                r_s_hdr.font.name = "Times New Roman"
                r_s_hdr.font.size = Pt(10)
                r_s_hdr.font.bold = True
                r_s_hdr.font.color.rgb = RGBColor(21, 101, 192)
                format_chemical_runs(p_sol, item_sol, is_italic=True)

    # ==================== PHẦN III ====================
    add_part_heading(doc, "PHẦN III.", "Câu trắc nghiệm yêu cầu trả lời ngắn. Thí sinh trả lời từ câu 1 đến câu 6.")

    for q in PART3_QUESTIONS:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        r_lbl = p_q.add_run(f"Câu {q['num']}: ")
        r_lbl.font.name = "Times New Roman"
        r_lbl.font.size = Pt(11)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = RGBColor(21, 101, 192)

        format_chemical_runs(p_q, q["text"])

        if is_teacher:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Inches(0.24)
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(1)
            r_ans_lbl = p_ans.add_run(f"Đáp số: {q['ans']}")
            r_ans_lbl.font.name = "Times New Roman"
            r_ans_lbl.font.size = Pt(11)
            r_ans_lbl.font.bold = True
            r_ans_lbl.font.color.rgb = RGBColor(211, 47, 47)

            add_solution(doc, q["sol"])
        else:
            # Dòng trả lời học sinh
            p_ans_line = doc.add_paragraph()
            p_ans_line.paragraph_format.left_indent = Inches(0.24)
            p_ans_line.paragraph_format.space_before = Pt(1)
            p_ans_line.paragraph_format.space_after = Pt(3)
            r_al = p_ans_line.add_run("Trả lời: ......................................................................................................................................................")
            r_al.font.name = "Times New Roman"
            r_al.font.size = Pt(10.5)
            r_al.font.italic = True
            r_al.font.color.rgb = RGBColor(120, 120, 120)

    # Bảng tổng hợp đáp án ở cuối bản giáo viên
    if is_teacher:
        doc.add_page_break()
        p_tbl_hdr = doc.add_paragraph()
        p_tbl_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_tbl_hdr.paragraph_format.space_before = Pt(6)
        p_tbl_hdr.paragraph_format.space_after = Pt(6)
        r_th = p_tbl_hdr.add_run("BẢNG TỔNG HỢP ĐÁP ÁN ĐỀ SỐ 23 (MÃ ĐỀ 123)")
        r_th.font.name = "Times New Roman"
        r_th.font.size = Pt(13)
        r_th.font.bold = True
        r_th.font.color.rgb = RGBColor(180, 20, 20)

        # Bảng Phần I: 18 câu (2 hàng: câu 1-9, câu 10-18)
        p_p1 = doc.add_paragraph()
        p_p1.paragraph_format.space_before = Pt(4)
        p_p1.paragraph_format.space_after = Pt(2)
        r_p1 = p_p1.add_run("1. PHẦN I: TRẮC NGHIỆM NHIỀU LỰA CHỌN (Mỗi câu đúng được 0,25 điểm)")
        r_p1.font.bold = True

        tbl_ans1 = doc.add_table(rows=4, cols=10)
        tbl_ans1.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_ans1.autofit = False

        headers1 = ["Câu", "1", "2", "3", "4", "5", "6", "7", "8", "9"]
        vals1 = ["Đ/A", "A", "D", "C", "D", "B", "D", "D", "B", "D"]
        headers2 = ["Câu", "10", "11", "12", "13", "14", "15", "16", "17", "18"]
        vals2 = ["Đ/A", "C", "B", "A", "B", "B", "A", "B", "B", "B"]

        for col_idx in range(10):
            # Row 0
            c = tbl_ans1.cell(0, col_idx)
            c.width = Inches(0.7)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(headers1[col_idx])
            r.font.bold = True

            # Row 1
            c = tbl_ans1.cell(1, col_idx)
            c.width = Inches(0.7)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(vals1[col_idx])
            r.font.bold = True
            if col_idx > 0:
                r.font.color.rgb = RGBColor(211, 47, 47)

            # Row 2
            c = tbl_ans1.cell(2, col_idx)
            c.width = Inches(0.7)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(headers2[col_idx])
            r.font.bold = True

            # Row 3
            c = tbl_ans1.cell(3, col_idx)
            c.width = Inches(0.7)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(vals2[col_idx])
            r.font.bold = True
            if col_idx > 0:
                r.font.color.rgb = RGBColor(211, 47, 47)

        tblPr = tbl_ans1._tbl.tblPr
        tblBorders = parse_xml(
            r'<w:tblBorders %s><w:top w:val="single" w:sz="6" w:space="0" w:color="8BC390"/><w:left w:val="single" w:sz="6" w:space="0" w:color="8BC390"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="8BC390"/><w:right w:val="single" w:sz="6" w:space="0" w:color="8BC390"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/><w:insideV w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/></w:tblBorders>' % nsdecls('w')
        )
        tblPr.append(tblBorders)

        # Bảng Phần II & III
        p_p2 = doc.add_paragraph()
        p_p2.paragraph_format.space_before = Pt(8)
        p_p2.paragraph_format.space_after = Pt(2)
        r_p2 = p_p2.add_run("2. PHẦN II: TRẮC NGHIỆM ĐÚNG SAI & PHẦN III: TRẢ LỜI NGẮN")
        r_p2.font.bold = True

        tbl_ans2 = doc.add_table(rows=5, cols=7)
        tbl_ans2.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_ans2.autofit = False

        h_p2 = ["Câu", "a", "b", "c", "d", "Phần III", "Đáp số"]
        rows_p2 = [
            ["Câu 1", "S", "Đ", "S", "S", "Câu 1", "5"],
            ["Câu 2", "S", "Đ", "S", "Đ", "Câu 2", "2"],
            ["Câu 3", "Đ", "S", "Đ", "S", "Câu 3", "6"],
            ["Câu 4", "Đ", "Đ", "Đ", "S", "Câu 4", "2"],
        ]

        for col_idx, h in enumerate(h_p2):
            c = tbl_ans2.cell(0, col_idx)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(h)
            r.font.bold = True

        for r_idx, row_vals in enumerate(rows_p2):
            for col_idx, val in enumerate(row_vals):
                c = tbl_ans2.cell(r_idx + 1, col_idx)
                p = c.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                r = p.add_run(val)
                r.font.bold = True
                if col_idx in [1, 2, 3, 4, 6]:
                    r.font.color.rgb = RGBColor(211, 47, 47)

        # Bổ sung dòng Câu 5 và 6 của Phần III
        p_extra = doc.add_paragraph()
        p_extra.paragraph_format.space_before = Pt(2)
        p_extra.paragraph_format.space_after = Pt(2)
        r_ex = p_extra.add_run("• Phần III - Câu 5: 30,7   |   • Phần III - Câu 6: 0,02")
        r_ex.font.bold = True
        r_ex.font.color.rgb = RGBColor(211, 47, 47)

        tblPr2 = tbl_ans2._tbl.tblPr
        tblBorders2 = parse_xml(
            r'<w:tblBorders %s><w:top w:val="single" w:sz="6" w:space="0" w:color="8BC390"/><w:left w:val="single" w:sz="6" w:space="0" w:color="8BC390"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="8BC390"/><w:right w:val="single" w:sz="6" w:space="0" w:color="8BC390"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/><w:insideV w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/></w:tblBorders>' % nsdecls('w')
        )
        tblPr2.append(tblBorders2)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Đã tạo thành công: {output_path}")

if __name__ == "__main__":
    generate_exam_docx(is_teacher=False, output_path="San_Pham/De_Kiem_Tra_Hoa_12_De23.docx")
    generate_exam_docx(is_teacher=True, output_path="San_Pham/De_Kiem_Tra_Hoa_12_De23_GiaoVien.docx")
