# -*- coding: utf-8 -*-
"""
Script tạo trọn bộ 3 file Word (.docx) chuẩn "FORM ĐỀ KIỂM TRA" cho Đề Cương Ôn Tập Giữa Kì I - KHTN 6
Năm học 2026 - 2027
Thầy TRẦN VĂN THẠNH - 0777.470.803

3 phiên bản:
1. Bản Học sinh: San_Pham/De_Cuong_KHTN_6_On_Tap_GK1_HocSinh.docx
2. Bản Giáo viên: San_Pham/De_Cuong_KHTN_6_On_Tap_GK1_GiaoVien.docx
3. Bản Dạy trên màn hình: San_Pham/De_Cuong_KHTN_6_On_Tap_GK1_DayTrenManHinh.docx
"""

import os
import sys
import re
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.path.append(os.path.dirname(__file__))
sys.stdout.reconfigure(encoding='utf-8')

from khtn6_ontap_gk1_data import (
    INFO,
    VATLI_THEORY, VATLI_MCQ, VATLI_SHORT, VATLI_TF, VATLI_ESSAY,
    HOA_THEORY,
    SINH_CHAPTER1_THEORY, SINH_CHAPTER1_EXERCISES,
    SINH_CHAPTER5_THEORY, SINH_CHAPTER5_EXERCISES,
    SINH_CHAPTER6_THEORY, SINH_CHAPTER6_EXERCISES
)

ROOT_DIR = r"c:\Antigravity_Thanh"
OUT_DIR = os.path.join(ROOT_DIR, "San_Pham")
os.makedirs(OUT_DIR, exist_ok=True)

ONEDRIVE_DIRS = [
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn6\KHTN6-DEKIEMTRADINHKI\New folder",
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn6\KHTN6-DEKIEMTRADINHKI",
    r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn6"
]

# Bảng màu chuẩn quy chuẩn trong GEMINI.md & FORM_DE_KIEM_TRA
COLOR_HF_RED = RGBColor(211, 47, 47)        # Coral Red (#D32F2F) cho Header & Footer
COLOR_HF_GREEN_HEX = "8BC390"               # Sage Green (#8BC390) cho viền Header & Footer
COLOR_PRIMARY = RGBColor(46, 125, 50)       # Xanh lá đậm (#2E7D32) cho Tiêu đề Phân môn / Phần
COLOR_CORRECT = RGBColor(198, 40, 40)       # Đỏ (#C62828) cho Đáp án đúng Bản Giáo Viên
COLOR_SOLUTION = RGBColor(21, 101, 192)     # Xanh lam (#1565C0) cho tiêu đề Lời giải
COLOR_SCREEN_TITLE = RGBColor(106, 27, 154) # Tím đậm (#6A1B9A) cho tiêu đề bản Dạy Màn Hình
COLOR_GRAY = RGBColor(110, 110, 110)
COLOR_BORDER_GREY_HEX = "CCCCCC"


def add_p_border_bottom(p, color="8BC390", sz="6"):
    """Thêm đường kẻ viền dưới cho đoạn văn bản."""
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))


def add_p_border_top(p, color="8BC390", sz="6"):
    """Thêm đường kẻ viền trên cho đoạn văn bản."""
    pPr = p._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="4" w:color="{color}"/></w:pBdr>'
    pPr.append(parse_xml(xml))


def set_run_style(run, font_name="Times New Roman", font_size=10.5, bold=False, italic=False, subscript=False, superscript=False, color=None):
    """Cài đặt style chi tiết cho từng run text."""
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True
    if color:
        run.font.color.rgb = color


def add_page_number_field(p, font_name="Times New Roman", font_size=10, bold=True, color=COLOR_HF_RED):
    """Chèn trường số trang tự động PAGE vào chân trang."""
    run_txt = p.add_run("Trang ")
    set_run_style(run_txt, font_name=font_name, font_size=font_size, bold=bold, color=color)

    r = p.add_run()
    set_run_style(r, font_name=font_name, font_size=font_size, bold=bold, color=color)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1); r._r.append(instr1); r._r.append(fld2); r._r.append(fld3)


def render_rich_text(p, text, base_font="Times New Roman", base_size=10.5, base_bold=False, base_italic=False, base_color=None):
    """
    Sinh văn bản định dạng phong phú từ ký hiệu:
    - *bold*
    - _italic_
    - ~subscript~ (CO~2~, H~2~O)
    - ^superscript^ (0^0C, 2^n)
    """
    tokens = re.split(r'(~[^~]+~|\^[^\^]+\^|\*[^\*]+\*|_[^_]+_)', text)
    for token in tokens:
        if not token:
            continue

        is_bold = base_bold
        is_italic = base_italic
        is_sub = False
        is_sup = False
        content = token

        if token.startswith('*') and token.endswith('*'):
            is_bold = True
            content = token[1:-1]
        elif token.startswith('_') and token.endswith('_'):
            is_italic = True
            content = token[1:-1]
        elif token.startswith('~') and token.endswith('~'):
            is_sub = True
            content = token[1:-1]
        elif token.startswith('^') and token.endswith('^'):
            is_sup = True
            content = token[1:-1]

        run = p.add_run(content)
        run_color = COLOR_CORRECT if (is_bold and not base_bold and not base_color) else base_color
        set_run_style(run, font_name=base_font, font_size=base_size, bold=is_bold, italic=is_italic,
                      subscript=is_sub, superscript=is_sup, color=run_color)


def setup_header_footer(doc):
    """Cài đặt Header & Footer chuẩn theo FORM_DE_KIEM_TRA."""
    normal_style = doc.styles['Normal']
    normal_style.paragraph_format.tab_stops.clear_all()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    section.start_type = WD_SECTION_START.NEW_PAGE
    section.different_first_page_header_footer = False
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    # 1. Header (Bảng 1x2, rộng 18.0 cm)
    header = section.header
    for p in header.paragraphs:
        p.text = ""
    tbl_h = header.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl_h.cell(0, 0)
    c_right = tbl_h.cell(0, 1)
    c_left.width = Cm(8.0)
    c_right.width = Cm(10.0)

    p_hl = c_left.paragraphs[0]
    p_hl.paragraph_format.space_before = Pt(0)
    p_hl.paragraph_format.space_after = Pt(2)
    p_hl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_hl = p_hl.add_run("Tài liệu lưu hành nội bộ")
    set_run_style(r_hl, font_size=10, italic=True, color=COLOR_HF_RED)

    p_hr = c_right.paragraphs[0]
    p_hr.paragraph_format.space_before = Pt(0)
    p_hr.paragraph_format.space_after = Pt(2)
    p_hr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p_hr.add_run("Thầy ")
    set_run_style(r1, font_size=10, color=COLOR_HF_RED)
    r2 = p_hr.add_run("TRẦN VĂN THẠNH")
    set_run_style(r2, font_size=10, bold=True, color=COLOR_HF_RED)
    r3 = p_hr.add_run(" - 0777.470.803")
    set_run_style(r3, font_size=10, color=COLOR_HF_RED)

    add_p_border_bottom(p_hl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_bottom(p_hr, color=COLOR_HF_GREEN_HEX, sz="6")

    # 2. Footer (Bảng 1x2, rộng 18.0 cm)
    footer = section.footer
    for p in footer.paragraphs:
        p.text = ""
    tbl_f = footer.add_table(rows=1, cols=2, width=Cm(18.0))
    tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_fl = tbl_f.cell(0, 0)
    c_fr = tbl_f.cell(0, 1)
    c_fl.width = Cm(14.5)
    c_fr.width = Cm(3.5)

    p_fl = c_fl.paragraphs[0]
    p_fl.paragraph_format.space_before = Pt(2)
    p_fl.paragraph_format.space_after = Pt(0)
    p_fl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_fl = p_fl.add_run("CS1: 6/15 Nguyễn Hoàng, P. Kim Long, TP Huế   |   CS2: 24 Đặng Thái Thân, TP Huế")
    set_run_style(r_fl, font_size=9.5, italic=True, color=COLOR_HF_RED)

    p_fr = c_fr.paragraphs[0]
    p_fr.paragraph_format.space_before = Pt(2)
    p_fr.paragraph_format.space_after = Pt(0)
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_field(p_fr, font_size=10, bold=True, color=COLOR_HF_RED)

    add_p_border_top(p_fl, color=COLOR_HF_GREEN_HEX, sz="6")
    add_p_border_top(p_fr, color=COLOR_HF_GREEN_HEX, sz="6")


def build_title_block(doc, is_teacher=False, is_screen=False, base_size=10.5):
    """Xây dựng Header Block không viền, chuẩn mực."""
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_l, cell_r = tbl.cell(0, 0), tbl.cell(0, 1)
    cell_l.width = Cm(8.0)
    cell_r.width = Cm(10.0)

    # Cột trái
    p_tl1 = cell_l.paragraphs[0]
    p_tl1.paragraph_format.space_before = Pt(0)
    p_tl1.paragraph_format.space_after = Pt(1)
    p_tl1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tl1, f"*{INFO['school']}*", base_size=10.5, base_bold=True)

    p_tl2 = cell_l.add_paragraph()
    p_tl2.paragraph_format.space_before = Pt(0)
    p_tl2.paragraph_format.space_after = Pt(1)
    p_tl2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tl2, f"*{INFO['year']}*", base_size=10, base_bold=True)

    p_tl3 = cell_l.add_paragraph()
    p_tl3.paragraph_format.space_before = Pt(0)
    p_tl3.paragraph_format.space_after = Pt(0)
    p_tl3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tl3, "_ĐỀ CƯƠNG ÔN TẬP CHÍNH THỨC_", base_size=9.5, base_italic=True)

    # Cột phải
    p_tr1 = cell_r.paragraphs[0]
    p_tr1.paragraph_format.space_before = Pt(0)
    p_tr1.paragraph_format.space_after = Pt(1)
    p_tr1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tr1, f"*{INFO['title']}*", base_size=11, base_bold=True)

    p_tr2 = cell_r.add_paragraph()
    p_tr2.paragraph_format.space_before = Pt(0)
    p_tr2.paragraph_format.space_after = Pt(1)
    p_tr2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tr2, f"*{INFO['subject']}*", base_size=12, base_bold=True, base_color=COLOR_PRIMARY)

    p_tr3 = cell_r.add_paragraph()
    p_tr3.paragraph_format.space_before = Pt(0)
    p_tr3.paragraph_format.space_after = Pt(0)
    p_tr3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_tr3, f"_{INFO['subtitle']}_", base_size=9.5, base_italic=True)

    # Dòng tiêu đề phụ cho Giáo viên / Màn hình
    if is_teacher:
        p_gv = doc.add_paragraph()
        p_gv.paragraph_format.space_before = Pt(4)
        p_gv.paragraph_format.space_after = Pt(2)
        p_gv.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rgv = p_gv.add_run("(BẢN GIÁO VIÊN -- CÓ ĐÁP ÁN & LỜI GIẢI CHI TIẾT)")
        set_run_style(rgv, font_size=11.5, bold=True, color=COLOR_CORRECT)
    elif is_screen:
        p_sc = doc.add_paragraph()
        p_sc.paragraph_format.space_before = Pt(4)
        p_sc.paragraph_format.space_after = Pt(2)
        p_sc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rsc = p_sc.add_run("(BẢN TRÌNH CHIẾU / GIẢNG DẠY MÀN HÌNH -- CÓ KHUNG GHI CHÚ BÀI DẠY)")
        set_run_style(rsc, font_size=11.5, bold=True, color=COLOR_SCREEN_TITLE)

    # Dòng thông tin học sinh
    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(4)
    p_info.paragraph_format.space_after = Pt(3)
    render_rich_text(p_info, "*Họ và tên thí sinh:* .................................................... *Lớp:* ......... *Trường:* ....................................................", base_size=base_size)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(4)
    add_p_border_bottom(p_div, color=COLOR_BORDER_GREY_HEX, sz="4")


def add_part_header(doc, title, base_size=12.5):
    """Tiêu đề lớn cho từng Phân môn."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    render_rich_text(p, f"*{title.upper()}*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)


def add_section_header(doc, title, base_size=11):
    """Tiêu đề mục: I. LÝ THUYẾT, II. TRẮC NGHIỆM..."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    render_rich_text(p, f"*{title}*", base_size=base_size, base_bold=True, base_color=RGBColor(21, 101, 192))


def add_callout_box(doc, lines, base_size=10):
    """Khung ghi chú nội dung ôn tập lý thuyết."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    c.width = Cm(18.0)
    
    tcPr = c._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', 60), ('w:bottom', 60), ('w:left', 120), ('w:right', 120)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F5F9F6"/>')
    tcPr.append(shd)

    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    render_rich_text(p, lines[0], base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)

    for line in lines[1:]:
        p_line = c.add_paragraph()
        p_line.paragraph_format.space_before = Pt(1)
        p_line.paragraph_format.space_after = Pt(1)
        render_rich_text(p_line, line, base_size=base_size)


def add_solution_block(doc, solution_text, base_size=10):
    """Trình bày Lời giải chi tiết theo chuẩn không dùng Table."""
    p_sol = doc.add_paragraph()
    p_sol.paragraph_format.left_indent = Cm(0.6)
    p_sol.paragraph_format.space_before = Pt(2)
    p_sol.paragraph_format.space_after = Pt(3)
    p_sol.paragraph_format.line_spacing = 1.15

    r_lbl = p_sol.add_run("Lời giải chi tiết:\n")
    set_run_style(r_lbl, font_size=base_size, bold=True, color=COLOR_SOLUTION)

    lines = solution_text.split('\n')
    for idx, line in enumerate(lines):
        line_str = line.strip()
        if not line_str:
            continue
        if idx > 0:
            p_sol = doc.add_paragraph()
            p_sol.paragraph_format.left_indent = Cm(0.6)
            p_sol.paragraph_format.space_before = Pt(1)
            p_sol.paragraph_format.space_after = Pt(2)
            p_sol.paragraph_format.line_spacing = 1.15
        render_rich_text(p_sol, line_str, base_size=base_size)


def add_screen_teaching_space(doc, num_lines=3):
    """Chừa khoảng trống rộng rãi cho bản Dạy Trên Màn Hình."""
    p_box = doc.add_paragraph()
    p_box.paragraph_format.left_indent = Cm(0.5)
    p_box.paragraph_format.space_before = Pt(3)
    p_box.paragraph_format.space_after = Pt(5)

    pPr = p_box._element.get_or_add_pPr()
    xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="dashed" w:sz="4" w:space="4" w:color="B0BEC5"/><w:bottom w:val="dashed" w:sz="4" w:space="4" w:color="B0BEC5"/><w:left w:val="single" w:sz="12" w:space="6" w:color="7E57C2"/></w:pBdr>'
    pPr.append(parse_xml(xml))

    r_hint = p_box.add_run("✎ Ghi chú & Lời giải trực tiếp trên màn hình:\n")
    set_run_style(r_hint, font_size=9.5, bold=True, color=COLOR_SCREEN_TITLE)
    for _ in range(num_lines):
        p_box.add_run("\n")


def add_writing_lines(doc, num_lines=3):
    """Thêm dòng chấm chấm cho học sinh làm bài tự luận."""
    for _ in range(num_lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("................................................................................................................................................................................................................")
        set_run_style(r, font_size=9, color=COLOR_GRAY)


def render_mcq_options(doc, opts, correct_letter, is_teacher, base_size=10.5):
    """Bố trí 4 phương án A, B, C, D bằng Tab Stops - Tuyệt đối không dùng Table."""
    lens = [len(opt) for opt in opts]
    max_len = max(lens) if lens else 0

    if max_len <= 18:
        # 4 phương án trên 1 dòng
        p_opt = doc.add_paragraph()
        p_opt.paragraph_format.space_before = Pt(1)
        p_opt.paragraph_format.space_after = Pt(2)
        p_opt.paragraph_format.tab_stops.clear_all()
        p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(5.0), WD_TAB_ALIGNMENT.LEFT)
        p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(9.5), WD_TAB_ALIGNMENT.LEFT)
        p_opt.paragraph_format.tab_stops.add_tab_stop(Cm(14.0), WD_TAB_ALIGNMENT.LEFT)

        for opt in opts:
            p_opt.add_run("\t")
            letter = opt.strip()[0]
            if is_teacher and letter == correct_letter:
                render_rich_text(p_opt, f"*{opt}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p_opt, opt, base_size=base_size)

    elif max_len <= 45:
        # 2 dòng x 2 phương án
        p1 = doc.add_paragraph()
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.tab_stops.clear_all()
        p1.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
        p1.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)

        p1.add_run("\t")
        if is_teacher and len(opts) > 0 and opts[0].strip()[0] == correct_letter:
            render_rich_text(p1, f"*{opts[0]}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
        elif len(opts) > 0:
            render_rich_text(p1, opts[0], base_size=base_size)

        if len(opts) > 1:
            p1.add_run("\t")
            if is_teacher and opts[1].strip()[0] == correct_letter:
                render_rich_text(p1, f"*{opts[1]}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p1, opts[1], base_size=base_size)

        if len(opts) > 2:
            p2 = doc.add_paragraph()
            p2.paragraph_format.space_before = Pt(1)
            p2.paragraph_format.space_after = Pt(2)
            p2.paragraph_format.tab_stops.clear_all()
            p2.paragraph_format.tab_stops.add_tab_stop(Cm(0.5), WD_TAB_ALIGNMENT.LEFT)
            p2.paragraph_format.tab_stops.add_tab_stop(Cm(9.2), WD_TAB_ALIGNMENT.LEFT)

            p2.add_run("\t")
            if is_teacher and opts[2].strip()[0] == correct_letter:
                render_rich_text(p2, f"*{opts[2]}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p2, opts[2], base_size=base_size)

            if len(opts) > 3:
                p2.add_run("\t")
                if is_teacher and opts[3].strip()[0] == correct_letter:
                    render_rich_text(p2, f"*{opts[3]}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
                else:
                    render_rich_text(p2, opts[3], base_size=base_size)
    else:
        # Mỗi phương án 1 dòng riêng
        for opt in opts:
            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.left_indent = Cm(0.5)
            p_opt.paragraph_format.space_before = Pt(1)
            p_opt.paragraph_format.space_after = Pt(1)
            letter = opt.strip()[0]
            if is_teacher and letter == correct_letter:
                render_rich_text(p_opt, f"*{opt}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
            else:
                render_rich_text(p_opt, opt, base_size=base_size)


def generate_document(is_teacher=False, is_screen=False):
    """Tạo tài liệu Word hoàn chỉnh theo cấu hình."""
    doc = docx.Document()
    setup_header_footer(doc)

    base_size = 11.5 if is_screen else 10.5
    build_title_block(doc, is_teacher=is_teacher, is_screen=is_screen, base_size=base_size)

    # ==========================================================================
    # PHẦN 1: PHÂN MÔN VẬT LÝ
    # ==========================================================================
    add_part_header(doc, "PHÂN MÔN VẬT LÝ", base_size=base_size + 2)

    # I. LÝ THUYẾT VẬT LÝ
    add_section_header(doc, "I. LÝ THUYẾT", base_size=base_size + 0.5)
    for q in VATLI_THEORY:
        p_qt = doc.add_paragraph()
        p_qt.paragraph_format.space_before = Pt(4)
        p_qt.paragraph_format.space_after = Pt(2)
        render_rich_text(p_qt, f"*{q['title']}*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)
        for item in q["items"]:
            p_item = doc.add_paragraph()
            p_item.paragraph_format.left_indent = Cm(0.4)
            p_item.paragraph_format.space_before = Pt(1)
            p_item.paragraph_format.space_after = Pt(1)
            render_rich_text(p_item, item, base_size=base_size)

    # II. TRẮC NGHIỆM NHIỀU LỰA CHỌN
    add_section_header(doc, "II. TRẮC NGHIỆM NHIỀU LỰA CHỌN", base_size=base_size + 0.5)
    for q in VATLI_MCQ:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*Câu {q['id']}:* {q['content']}", base_size=base_size)
        render_mcq_options(doc, q["options"], q["answer"], is_teacher, base_size=base_size)

        if is_teacher:
            add_solution_block(doc, f"Đáp án đúng: *{q['answer']}*.\n{q['solution']}", base_size=base_size - 0.5)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=2)

    # III. TRẮC NGHIỆM TRẢ LỜI NGẮN
    add_section_header(doc, "III. TRẮC NGHIỆM TRẢ LỜI NGẮN", base_size=base_size + 0.5)
    for q in VATLI_SHORT:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*Câu {q['id']}.* {q['content']}", base_size=base_size)

        for sub in q["sub_questions"]:
            p_sub = doc.add_paragraph()
            p_sub.paragraph_format.left_indent = Cm(0.4)
            p_sub.paragraph_format.space_before = Pt(2)
            p_sub.paragraph_format.space_after = Pt(2)
            render_rich_text(p_sub, f"*{sub['sub_id']})* {sub['text']}", base_size=base_size)

            if is_teacher:
                p_ans = doc.add_paragraph()
                p_ans.paragraph_format.left_indent = Cm(0.8)
                p_ans.paragraph_format.space_before = Pt(1)
                p_ans.paragraph_format.space_after = Pt(2)
                r_lbl = p_ans.add_run("Đáp số: ")
                set_run_style(r_lbl, font_size=base_size, bold=True, color=COLOR_CORRECT)
                render_rich_text(p_ans, f"*{sub['answer']}*", base_size=base_size, base_bold=True, base_color=COLOR_CORRECT)
                add_solution_block(doc, sub["solution"], base_size=base_size - 0.5)
            elif is_screen:
                add_screen_teaching_space(doc, num_lines=2)
            else:
                p_ans = doc.add_paragraph()
                p_ans.paragraph_format.left_indent = Cm(0.8)
                p_ans.paragraph_format.space_before = Pt(1)
                p_ans.paragraph_format.space_after = Pt(2)
                render_rich_text(p_ans, "*Trả lời:* .................................................................................................................................................", base_size=base_size)

    # IV. TRẮC NGHIỆM ĐÚNG SAI
    add_section_header(doc, "IV. TRẮC NGHIỆM ĐÚNG SAI", base_size=base_size + 0.5)
    for q in VATLI_TF:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*Câu {q['id']}:* {q['context']}", base_size=base_size)

        p_guide = doc.add_paragraph()
        p_guide.paragraph_format.left_indent = Cm(0.4)
        p_guide.paragraph_format.space_before = Pt(1)
        p_guide.paragraph_format.space_after = Pt(2)
        render_rich_text(p_guide, "_Mỗi phát biểu sau đây là Đúng hay Sai?_", base_size=base_size, base_italic=True)

        sol_lines = []
        for item in q["items"]:
            p_item = doc.add_paragraph()
            p_item.paragraph_format.left_indent = Cm(0.5)
            p_item.paragraph_format.space_before = Pt(1)
            p_item.paragraph_format.space_after = Pt(1)

            txt = f"*{item['id']})* {item['content']}"
            if is_teacher:
                render_rich_text(p_item, txt, base_size=base_size)
                ans_str = " ---> [ĐÚNG]" if item["answer"] else " ---> [SAI]"
                r_ans = p_item.add_run(ans_str)
                set_run_style(r_ans, font_size=base_size, bold=True, color=COLOR_CORRECT)
                sol_lines.append(f"• Ý {item['id']}: {'Đúng' if item['answer'] else 'Sai'}. {item.get('reason', '')}")
            else:
                render_rich_text(p_item, txt, base_size=base_size)

        if is_teacher and sol_lines:
            add_solution_block(doc, "\n".join(sol_lines), base_size=base_size - 0.5)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=3)

    # V. TỰ LUẬN VẬT LÝ
    add_section_header(doc, "V. TỰ LUẬN", base_size=base_size + 0.5)
    for q in VATLI_ESSAY:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        render_rich_text(p_q, f"*Câu {q['id']}:* {q['content']}", base_size=base_size)

        if is_teacher:
            add_solution_block(doc, q["solution"], base_size=base_size - 0.5)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=4)
        else:
            add_writing_lines(doc, num_lines=4)

    # ==========================================================================
    # PHẦN 2: PHÂN MÔN HOÁ HỌC
    # ==========================================================================
    add_part_header(doc, "PHÂN MÔN HOÁ HỌC", base_size=base_size + 2)
    for q in HOA_THEORY:
        p_qt = doc.add_paragraph()
        p_qt.paragraph_format.space_before = Pt(4)
        p_qt.paragraph_format.space_after = Pt(2)
        render_rich_text(p_qt, f"*{q['title']}*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)
        for item in q["items"]:
            p_item = doc.add_paragraph()
            p_item.paragraph_format.left_indent = Cm(0.4)
            p_item.paragraph_format.space_before = Pt(1)
            p_item.paragraph_format.space_after = Pt(1)
            render_rich_text(p_item, item, base_size=base_size)

        if is_screen:
            add_screen_teaching_space(doc, num_lines=2)

    # ==========================================================================
    # PHẦN 3: PHÂN MÔN SINH HỌC
    # ==========================================================================
    add_part_header(doc, "PHÂN MÔN SINH HỌC", base_size=base_size + 2)

    # Chương I: Mở đầu
    add_section_header(doc, "Chương I: Mở đầu", base_size=base_size + 0.5)
    for item in SINH_CHAPTER1_THEORY:
        p_it = doc.add_paragraph()
        p_it.paragraph_format.left_indent = Cm(0.4)
        p_it.paragraph_format.space_before = Pt(1)
        p_it.paragraph_format.space_after = Pt(2)
        render_rich_text(p_it, item, base_size=base_size)

    p_bt_hdr1 = doc.add_paragraph()
    p_bt_hdr1.paragraph_format.space_before = Pt(4)
    p_bt_hdr1.paragraph_format.space_after = Pt(2)
    render_rich_text(p_bt_hdr1, "*BÀI TẬP MINH HOẠ:*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)

    for ex in SINH_CHAPTER1_EXERCISES:
        p_ex = doc.add_paragraph()
        p_ex.paragraph_format.left_indent = Cm(0.3)
        p_ex.paragraph_format.space_before = Pt(3)
        p_ex.paragraph_format.space_after = Pt(2)
        render_rich_text(p_ex, f"*{ex['id']}.* {ex['question']}", base_size=base_size)

        if is_teacher:
            add_solution_block(doc, ex["solution"], base_size=base_size - 0.5)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=3)
        else:
            add_writing_lines(doc, num_lines=3)

    # Chương V. Tế bào – đơn vị cơ sở của sự sống
    add_section_header(doc, "Chương V. Tế bào – đơn vị cơ sở của sự sống", base_size=base_size + 0.5)
    for item in SINH_CHAPTER5_THEORY:
        p_it = doc.add_paragraph()
        p_it.paragraph_format.left_indent = Cm(0.4)
        p_it.paragraph_format.space_before = Pt(1)
        p_it.paragraph_format.space_after = Pt(1)
        render_rich_text(p_it, item, base_size=base_size)

    p_bt_hdr5 = doc.add_paragraph()
    p_bt_hdr5.paragraph_format.space_before = Pt(4)
    p_bt_hdr5.paragraph_format.space_after = Pt(2)
    render_rich_text(p_bt_hdr5, "*BÀI TẬP MINH HOẠ:*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)

    for ex in SINH_CHAPTER5_EXERCISES:
        p_ex = doc.add_paragraph()
        p_ex.paragraph_format.left_indent = Cm(0.3)
        p_ex.paragraph_format.space_before = Pt(3)
        p_ex.paragraph_format.space_after = Pt(2)
        render_rich_text(p_ex, f"*{ex['id']}.* {ex['question']}", base_size=base_size)

        if is_teacher:
            add_solution_block(doc, ex["solution"], base_size=base_size - 0.5)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=4)
        else:
            add_writing_lines(doc, num_lines=3)

    # Chương VI. Từ tế bào đến cơ thể
    add_section_header(doc, "Chương VI. Từ tế bào đến cơ thể", base_size=base_size + 0.5)
    for item in SINH_CHAPTER6_THEORY:
        p_it = doc.add_paragraph()
        p_it.paragraph_format.left_indent = Cm(0.4)
        p_it.paragraph_format.space_before = Pt(1)
        p_it.paragraph_format.space_after = Pt(1)
        render_rich_text(p_it, item, base_size=base_size)

    p_bt_hdr6 = doc.add_paragraph()
    p_bt_hdr6.paragraph_format.space_before = Pt(4)
    p_bt_hdr6.paragraph_format.space_after = Pt(2)
    render_rich_text(p_bt_hdr6, "*BÀI TẬP MINH HOẠ:*", base_size=base_size, base_bold=True, base_color=COLOR_PRIMARY)

    for ex in SINH_CHAPTER6_EXERCISES:
        p_ex = doc.add_paragraph()
        p_ex.paragraph_format.left_indent = Cm(0.3)
        p_ex.paragraph_format.space_before = Pt(3)
        p_ex.paragraph_format.space_after = Pt(2)
        render_rich_text(p_ex, f"*{ex['id']}.* {ex['question']}", base_size=base_size)

        if is_teacher:
            add_solution_block(doc, ex["solution"], base_size=base_size - 0.5)
        elif is_screen:
            add_screen_teaching_space(doc, num_lines=4)
        else:
            add_writing_lines(doc, num_lines=3)

    # Dòng kết thúc
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(12)
    p_end.paragraph_format.space_after = Pt(6)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    render_rich_text(p_end, "=== *HẾT* ===", base_size=11, base_bold=True, base_color=COLOR_PRIMARY)

    return doc


def main():
    print("=" * 70)
    print("BẮT ĐẦU TẠO 3 FILE WORD CHUẨN FORM ĐỀ KIỂM TRA CHO KHTN 6 (2026-2027)")
    print("=" * 70)

    configs = [
        {
            "name": "Bản Học sinh",
            "file": "De_Cuong_KHTN_6_On_Tap_GK1_HocSinh.docx",
            "is_teacher": False,
            "is_screen": False
        },
        {
            "name": "Bản Giáo viên (Có Lời Giải Chi Tiết)",
            "file": "De_Cuong_KHTN_6_On_Tap_GK1_GiaoVien.docx",
            "is_teacher": True,
            "is_screen": False
        },
        {
            "name": "Bản Dạy Trên Màn Hình (Chừa Khung Tương Tác)",
            "file": "De_Cuong_KHTN_6_On_Tap_GK1_DayTrenManHinh.docx",
            "is_teacher": False,
            "is_screen": True
        }
    ]

    generated_files = []

    for cfg in configs:
        print(f"\n[+] Đang khởi tạo: {cfg['name']} -> {cfg['file']}...")
        doc = generate_document(is_teacher=cfg["is_teacher"], is_screen=cfg["is_screen"])
        out_path = os.path.join(OUT_DIR, cfg["file"])
        doc.save(out_path)
        print(f"    -> Đã lưu thành công tại: {out_path} ({os.path.getsize(out_path):,} bytes)")
        generated_files.append((cfg["file"], out_path))

    # Đồng bộ sang OneDrive nếu có
    print("\n[+] Đồng bộ các file Word sang OneDrive:")
    for od in ONEDRIVE_DIRS:
        if os.path.exists(od):
            for fname, src_path in generated_files:
                dest_path = os.path.join(od, fname)
                shutil.copy2(src_path, dest_path)
                print(f"    -> Đã sao chép vào: {dest_path}")
        else:
            print(f"    [!] Thư mục OneDrive không tồn tại: {od}")

    print("\n" + "=" * 70)
    print("HOÀN TẤT TẠO TOÀN BỘ 3 FILE WORD THEO ĐÚNG FORM ĐỀ KIỂM TRA!")
    print("=" * 70)


if __name__ == "__main__":
    main()
