# -*- coding: utf-8 -*-
"""
Script sinh tự động 2 file Word (.docx) cho Đề số 12 (Chương 4: HYDROCARBON):
1. Bản Học sinh (De_Kiem_Tra_Hoa_11_De12.docx)
2. Bản Giáo viên (De_Kiem_Tra_Hoa_11_De12_GiaoVien.docx)

Tuân thủ nghiêm ngặt Quy chuẩn Form Đề Kiểm Tra (GEMINI.md):
- Không dùng bảng (Table) trong nội dung câu hỏi/đáp án/lời giải.
- Dùng Paragraph + Tab stops chuẩn cho các phương án trắc nghiệm.
- Công thức hóa học xử lý triệt để chỉ số trên/dưới.
- Header & Footer chuẩn với viền màu Sage Green (#8BC390) và chữ Coral Red (#D32F2F).
- Bản Giáo viên có ghi chú đỏ và lời giải thụt lề màu xanh #1565C0.
"""

import os
import sys

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH

# Đảm bảo import được module form_de_kiem_tra
sys.path.insert(0, os.path.dirname(__file__))
from form_de_kiem_tra import (
    setup_header_footer,
    build_title_block,
    add_section_header,
    render_question_prompt,
    render_options_no_table,
    render_statement,
    add_solution_no_table,
    add_student_writing_lines,
    render_rich_text,
    set_run_style,
    COLOR_CORRECT,
    COLOR_SOLUTION,
    COLOR_PRIMARY
)

OUT_DIR = r"c:\Antigravity_Thanh\San_Pham"
IMG_DIR = os.path.join(OUT_DIR, "images_crop_hoa11_de12")

def add_p(doc, space_before=3, space_after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    return p

def add_centered_picture(doc, img_path, width_cm=8.0):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(img_path, width=Cm(width_cm))

def build_exam_document(is_teacher=False):
    doc = docx.Document()
    setup_header_footer(doc)

    build_title_block(
        doc,
        school_line="SỞ GIÁO DỤC VÀ ĐÀO TẠO THỪA THIÊN HUẾ",
        exam_title="ĐỀ KIỂM TRA ĐỊNH KỲ CHƯƠNG 4 (HYDROCARBON)",
        subject_line="MÔN: HOÁ HỌC - LỚP 11",
        num_pages="04",
        duration="50 phút",
        exam_code="112",
        atomic_masses="H = 1; C = 12; N = 14; O = 16; Na = 23; Al = 27; S = 32; Cl = 35,5; K = 39; Ca = 40; Mn = 55; Fe = 56; Cu = 64; Zn = 65; Br = 80; Ag = 108. (Thể tích khí đo ở đkc 25 °C, 1 bar là 24,79 L/mol)",
        is_teacher=is_teacher
    )

    # =========================================================================
    # PHẦN I
    # =========================================================================
    add_section_header(
        doc,
        title="PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án."
    )

    # Câu 1
    p = add_p(doc)
    render_question_prompt(p, "Câu 1. (SBT – KNTT) Chất nào sau đây không phải là hydrocarbon?")
    opts_c1 = [
        "A. H~3~C - CH~3~",
        "B. H~2~C = CH~2~",
        "C. HC ≡ CH",
        "*D. CH~3~-CH~2~-OH*"
    ]
    render_options_no_table(doc, opts_c1, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "CH~3~-CH~2~-OH chứa nhóm chức alcohol (-OH) gồm 3 nguyên tố C, H, O nên là dẫn xuất chứa oxygen của hydrocarbon, không phải hydrocarbon (chỉ gồm C và H)."
        )

    # Câu 2
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 2. (SBT-CD) Những yếu tố nào sau đây không quyết định đến độ lớn của nhiệt độ sôi của các alkane?\n"
        "(a) Phân tử khối.\n"
        "(b) Tương tác van der Waals giữa các phân tử.\n"
        "(c) Độ tan trong nước.\n"
        "(d) Liên kết hydrogen giữa các phân tử."
    )
    opts_c2 = [
        "A. (a) và (b).",
        "B. (b) và (c).",
        "*C. (c) và (d).* ",
        "D. (a) và (d)."
    ]
    render_options_no_table(doc, opts_c2, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Nhiệt độ sôi của alkane phụ thuộc vào phân tử khối và độ lớn tương tác van der Waals giữa các phân tử. "
            "Alkane không tan trong nước và giữa các phân tử alkane không có liên kết hydrogen, do đó (c) và (d) không quyết định độ lớn nhiệt độ sôi."
        )

    # Câu 3
    p = add_p(doc)
    render_question_prompt(p, "Câu 3. [KNTT - SBT] Phát biểu nào sau đây về ứng dụng của alkane không đúng?")
    opts_c3 = [
        "A. Propane C~3~H~8~ và butane C~4~H~10~ được sử dụng làm khí đốt.",
        "B. Các alkane C~6~, C~7~, C~8~ là nguyên liệu để sản xuất một số hydrocarbon thơm.",
        "C. Các alkane lỏng được sử dụng làm nhiên liệu như xăng hay dầu diesel.",
        "*D. Các alkane từ C~11~ đến C~20~ được dùng làm nến và sáp.*"
    ]
    render_options_no_table(doc, opts_c3, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Các alkane rắn (thường có từ 20 nguyên tử carbon trở lên, sáp paraffin) mới được dùng làm nến và sáp. "
            "Các alkane từ C~11~ đến C~20~ ở trạng thái lỏng được dùng làm dầu hoả, nhiên liệu phản lực và dầu diesel."
        )

    # Câu 4
    p = add_p(doc)
    render_question_prompt(p, "Câu 4. (SBT-CD) Để tăng chất lượng của xăng, dầu, người ta thực hiện cách nào sau đây?")
    opts_c4 = [
        "*A. Thực hiện phản ứng reforming để thay đổi cấu trúc của các alkane mạch không nhánh thành hydrocarbon mạch nhánh hoặc mạch vòng có chỉ số octane cao.*",
        "B. Thực hiện phản ứng cracking để thay đổi cấu trúc các alkane mạch dài chuyển thành các alkene và alkane mạch ngắn hơn.",
        "C. Thực hiện phản ứng hydrogen hoá để chuyển các alkene thành alkane.",
        "D. Bổ sung thêm heptane vào xăng, dầu."
    ]
    render_options_no_table(doc, opts_c4, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Phản ứng reforming biến đổi alkane mạch thẳng thành alkane phân nhánh hoặc hydrocarbon thơm (chỉ số octane cao), giúp nhiên liệu cháy êm, chống kích nổ hiệu quả."
        )

    # Câu 5
    p = add_p(doc)
    render_question_prompt(p, "Câu 5. [KNTT – SBT] Hợp chất nào sau đây là một alkyne?")
    opts_c5 = [
        "A. CH~3~CH~2~CH~2~CH~3~.",
        "B. CH~3~CH=CH~2~.",
        "*C. CH~3~CH~2~C≡CH.*",
        "D. CH~2~=CH-CH=CH~2~."
    ]
    render_options_no_table(doc, opts_c5, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Alkyne là hydrocarbon không no, mạch hở, trong phân tử có chứa 1 liên kết ba C≡C. Do đó CH~3~CH~2~C≡CH (but-1-yne) là alkyne."
        )

    # Câu 6
    p = add_p(doc)
    render_question_prompt(p, "Câu 6. (SBT-CD) Các alkene không có tính chất vật lí đặc trưng nào sau đây?")
    opts_c6 = [
        "*A. Tan tốt trong nước và các dung môi hữu cơ.*",
        "B. Có khối lượng riêng nhỏ hơn khối lượng riêng của nước.",
        "C. Có nhiệt độ sôi thấp hơn alkane phân tử có cùng số nguyên tử carbon.",
        "D. Không dẫn điện."
    ]
    render_options_no_table(doc, opts_c6, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Phân tử các alkene không phân cực hoặc rất ít phân cực, nên không tan trong nước (dung môi phân cực mạnh), chỉ tan tốt trong các dung môi hữu cơ không phân cực."
        )

    # Câu 7
    p = add_p(doc)
    render_question_prompt(p, "Câu 7. [KNTT – SBT] Phản ứng nào sau đây không phải là phản ứng đặc trưng của hydrocarbon không no?")
    opts_c7 = [
        "A. Phản ứng cộng.",
        "B. Phản ứng trùng hợp.",
        "C. Phản ứng oxi hóa – khử.",
        "*D. Phản ứng thế.*"
    ]
    render_options_no_table(doc, opts_c7, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Phản ứng thế là phản ứng đặc trưng của hydrocarbon no (alkane). Hydrocarbon không no có đặc trưng là phản ứng cộng vào liên kết bội và phản ứng trùng hợp."
        )

    # Câu 8
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 8. (SBT-CD) Cho các chất có công thức cấu tạo sau:\n"
        "(1) ClCH~2~CH=CHCH~3~; (2) CH~3~CH=CHCH~3~; (3) BrCH~2~C(CH~3~)=C(CH~2~CH~3~)~2~;\n"
        "(4) ClCH~2~CH=CH~2~; (5) ClCH~2~CH=CHCH~2~CH~3~; (6) (CH~3~)~2~C=CH~2~.\n"
        "Trong số các chất trên, bao nhiêu chất có đồng phân hình học?"
    )
    opts_c8 = [
        "*A. 3.*",
        "B. 4.",
        "C. 5.",
        "D. 6."
    ]
    render_options_no_table(doc, opts_c8, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Điều kiện để chất có đồng phân hình học quanh liên kết C=C là mỗi carbon liên kết đôi liên kết với 2 nhóm nguyên tử khác nhau (R1 ≠ R2 và R3 ≠ R4).\n"
            "- (1) ClCH~2~CH=CHCH~3~: có đphh (ClCH~2~- ≠ H- và CH~3~- ≠ H-).\n"
            "- (2) CH~3~CH=CHCH~3~: có đphh (CH~3~- ≠ H- và CH~3~- ≠ H-).\n"
            "- (3) BrCH~2~C(CH~3~)=C(CH~2~CH~3~)~2~: C thứ hai liên kết 2 nhóm -CH~2~CH~3~ giống nhau -> không có đphh.\n"
            "- (4) ClCH~2~CH=CH~2~: C cuối liên kết 2 nguyên tử H -> không có đphh.\n"
            "- (5) ClCH~2~CH=CHCH~2~CH~3~: có đphh (ClCH~2~- ≠ H- và C~2~H~5~- ≠ H-).\n"
            "- (6) (CH~3~)~2~C=CH~2~: không có đphh.\n"
            "=> Có 3 chất (1), (2), (5)."
        )

    # Câu 9
    p = add_p(doc)
    render_question_prompt(p, "Câu 9. (SBT-CD) Phát biểu nào sau đây là không đúng?")
    opts_c9 = [
        "A. Hydrocarbon thơm là những hydrocarbon trong phân tử có chứa vòng benzene.",
        "*B. Các chất trong phân tử có vòng benzene được gọi là hydrocarbon thơm.*",
        "C. Những hydrocarbon trong phân tử có chứa vòng benzene được gọi là arene.",
        "D. Dãy đồng đẳng của benzene có công thức chung là C~n~H~2n-6~ (n ≥ 6)."
    ]
    render_options_no_table(doc, opts_c9, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Phát biểu B không đúng vì các chất chứa vòng benzene nhưng có thêm nhóm chức khác chứa O, N, halogen (ví dụ phenol C~6~H~5~OH, aniline C~6~H~5~NH~2~) là dẫn xuất thơm, không phải là hydrocarbon thơm (hydrocarbon chỉ chứa C và H)."
        )

    # Câu 10
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 10. (SBT – KNTT) Cho các chất sau: (X) _o_-bromotoluene; (Y) _m_-bromotoluene; (Z) _p_-bromotoluene. "
        "Sản phẩm chính của phản ứng giữa toluene với bromine ở nhiệt độ cao có mặt iron(III) bromide là"
    )
    opts_c10 = [
        "A. (X) và (Y).",
        "B. (Y) và (Z).",
        "*C. (X) và (Z).* ",
        "D. (Y)."
    ]
    render_options_no_table(doc, opts_c10, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Nhóm alkyl (-CH~3~) trên vòng thơm là nhóm đẩy electron, làm tăng mật độ điện tích tại vị trí _ortho_ và _para_, "
            "do đó phản ứng thế halogen ưu tiên định hướng vào vị trí _ortho_ (chất X) và _para_ (chất Z)."
        )

    # Câu 11
    p = add_p(doc)
    render_question_prompt(p, "Câu 11. (SBT – KNTT) Nhận xét nào sau đây về tính chất hoá học của benzene là không đúng?")
    opts_c11 = [
        "A. Benzene khó tham gia phản ứng cộng hơn ethylene.",
        "B. Benzene dễ tham gia phản ứng thế hơn so với phản ứng cộng.",
        "C. Benzene không bị oxi hoá bởi tác nhân oxi hoá thông thường.",
        "*D. Benzene làm mất màu dung dịch nước bromine ở điều kiện thường.*"
    ]
    render_options_no_table(doc, opts_c11, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Hệ liên hợp pi kín trong vòng thơm benzene rất bền vững, do đó benzene không phản ứng và không làm mất màu nước bromine ở điều kiện thường."
        )

    # Câu 12
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 12. (SBT-CD) Một số chất gây ô nhiễm môi trường như benzene, toluene có trong khí thải đốt cháy nhiên liệu xăng, dầu. "
        "Để giảm thiểu nguyên nhân gây ô nhiễm này cần"
    )
    opts_c12 = [
        "A. cấm sử dụng nhiên liệu xăng.",
        "*B. hạn chế sử dụng nhiên liệu hoá thạch.*",
        "C. thay xăng bằng khí gas.",
        "D. cấm sử dụng xe cá nhân."
    ]
    render_options_no_table(doc, opts_c12, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Hạn chế sử dụng nhiên liệu hoá thạch, chuyển đổi sang các nguồn năng lượng sạch (điện, hydro, năng lượng tái tạo) là biện pháp bền vững và thực tế nhất để giảm khí thải độc hại."
        )

    # Câu 13
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 13. (SBT-CD) Hydrocarbon Y có công thức cấu tạo: (CH~3~)~2~CHCH~2~CH~3~. "
        "Khi cho Y phản ứng với bromine theo tỉ lệ mol 1:1 có chiếu sáng có thể thu được bao nhiêu dẫn xuất monobromo là đồng phân cấu tạo của nhau?"
    )
    opts_c13 = [
        "A. 3.",
        "*B. 4.*",
        "C. 5.",
        "D. 6."
    ]
    render_options_no_table(doc, opts_c13, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Phân tử (CH~3~)~2~CH-CH~2~-CH~3~ có 4 vị trí nguyên tử C không tương đương (2 nhóm -CH~3~ ở C1, C2, C3, C4) nên tạo ra 4 dẫn xuất monobromo đồng phân cấu tạo."
        )

    # Câu 14
    p = add_p(doc)
    render_question_prompt(p, "Câu 14. (SBT-CD) Tên thay thế của hydrocarbon có công thức cấu tạo (CH~3~)~3~CCH~2~CH~2~CH~3~ là")
    opts_c14 = [
        "*A. 2,2-dimethylpentane.*",
        "B. 2,3-dimethylpentane.",
        "C. 2,2,3-trimethylbutane.",
        "D. 2,2-dimethylbutane."
    ]
    render_options_no_table(doc, opts_c14, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Mạch chính dài nhất gồm 5 nguyên tử carbon (pentane), đánh số từ đầu gần nhánh hơn: vị trí C số 2 mang hai nhóm methyl -> 2,2-dimethylpentane."
        )

    # Câu 15
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 15. (SBT-CD) Để phân biệt but-2-yne (CH~3~C≡CCH~3~) với but-1-yne (CH≡CCH~2~CH~3~) có thể dùng thuốc thử nào sau đây?"
    )
    opts_c15 = [
        "A. Dung dịch HCl.",
        "*B. Dung dịch AgNO~3~/NH~3~.*",
        "C. Nước bromine.",
        "D. Dung dịch KMnO~4~."
    ]
    render_options_no_table(doc, opts_c15, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "But-1-yne có liên kết ba đầu mạch (alk-1-yne) nên tạo kết tủa màu vàng nhạt AgC≡CCH~2~CH~3~ với dung dịch AgNO~3~/NH~3~. But-2-yne có liên kết ba giữa mạch nên không phản ứng."
        )

    # Câu 16
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 16. (SBT – KNTT) Đun nóng hydrocarbon thơm X có công thức phân tử C~8~H~10~ với dung dịch KMnO~4~ nóng thu được dung dịch có chứa C~6~H~5~COOK và K~2~CO~3~. Chất X là"
    )
    opts_c16 = [
        "A. _o_-xylene.",
        "B. _p_-xylene.",
        "*C. ethylbenzene.*",
        "D. styrene."
    ]
    render_options_no_table(doc, opts_c16, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Oxi hoá ethylbenzene C~6~H~5~CH~2~CH~3~ bằng KMnO~4~ đun nóng làm bẻ gãy mạch nhánh tại carbon benzylic: "
            "C~6~H~5~CH~2~CH~3~ + 4KMnO~4~ -> C~6~H~5~COOK + K~2~CO~3~ + 4MnO~2~ + KOH + 2H~2~O. Do đó X là ethylbenzene."
        )

    # Câu 17
    p = add_p(doc)
    render_question_prompt(p, "Câu 17. Có 4 tên gọi: _o_-xylene; _o_-dimethylbenzene; 1,2-dimethylbenzene; ethylbenzene. Đó là tên của mấy chất?")
    opts_c17 = [
        "A. 1 chất.",
        "*B. 2 chất.*",
        "C. 3 chất.",
        "D. 4 chất."
    ]
    render_options_no_table(doc, opts_c17, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- _o_-xylene, _o_-dimethylbenzene, 1,2-dimethylbenzene đều là tên gọi của cùng một chất: 1,2-(CH~3~)~2~C~6~H~4~.\n"
            "- ethylbenzene là chất thứ hai: C~6~H~5~CH~2~CH~3~.\n"
            "=> Có tất cả 2 chất."
        )

    # Câu 18
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 18. Khu vực trạm xăng dầu, khu vực có nhiều xe cơ giới qua lại, nơi có khói thuốc lá... lại được xem là nơi gây tổn hại đến sức khoẻ con người. Phát biểu nào sau đây đúng?"
    )
    opts_c18 = [
        "*A. Khu vực có trạm xăng dầu, khu vực có nhiều xe cơ giới qua lại, nơi có khói thuốc lá,... là nơi có nguồn hydrocarbon thơm, đây là các chất độc gây tổn hại đến sức khoẻ con người.*",
        "B. Trong xăng tồn tại ethanol.",
        "C. Trong thuốc lá tồn tại nicotine (là hợp chất chứa vòng benzene) không độc.",
        "D. Trong xăng không tồn tại benzene (C~6~H~6~)."
    ]
    render_options_no_table(doc, opts_c18, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Khói thuốc lá và khí thải xăng xe chứa hàm lượng đáng kể các hydrocarbon thơm độc hại (benzene, toluene, polycyclic aromatic hydrocarbons) gây ức chế tuỷ và nguy cơ ung thư cao."
        )

    # =========================================================================
    # PHẦN II
    # =========================================================================
    add_section_header(
        doc,
        title="PHẦN II. Câu trắc nghiệm đúng sai.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai."
    )

    # Câu 1
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 1. So sánh và giải thích nhiệt độ sôi của alkane mạch không phân nhánh với alkane mạch phân nhánh khi chúng có cùng số nguyên tử carbon trong phân tử."
    )
    stmts_c1_p2 = [
        ("a) Các phân tử alkane phân nhánh có nhiệt độ sôi thấp hơn các phân tử alkane không phân nhánh dù cho chúng có cùng số nguyên tử carbon.", "ĐÚNG"),
        ("b) Các phân tử alkane phân nhánh có nhiệt độ sôi cao hơn các phân tử alkane không phân nhánh dù cho chúng có cùng số nguyên tử carbon.", "SAI"),
        ("c) Các phân tử alkane không phân nhánh có diện tích bề mặt lớn hơn so với phân tử alkane phân nhánh nên tồn tại tương tác van der Waals giữa các phân tử lớn hơn, dẫn đến nhiệt độ sôi cao hơn.", "ĐÚNG"),
        ("d) Các phân tử alkane không phân nhánh có diện tích bề mặt nhỏ hơn so với phân tử alkane phân nhánh nên tồn tại tương tác van der Waals giữa các phân tử nhỏ hơn, dẫn đến nhiệt độ sôi thấp hơn.", "SAI")
    ]
    for stmt, ans in stmts_c1_p2:
        p_st = add_p(doc, space_before=1, space_after=1)
        p_st.paragraph_format.left_indent = Cm(0.5)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- Alkane mạch không phân nhánh có diện tích tiếp xúc phân tử lớn hơn alkane mạch nhánh (cấu trúc gọn hơn dạng hình cầu).\n"
            "- Do đó tương tác van der Waals giữa các phân tử mạch thẳng lớn hơn, cần nhiều năng lượng hơn để phân tán phân tử khi sôi, dẫn đến nhiệt độ sôi cao hơn.\n"
            "=> a: ĐÚNG, b: SAI, c: ĐÚNG, d: SAI."
        )

    # Câu 2
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 2. Không được dùng nước dập tắt đám cháy có mặt đất đèn (có thành phần chính là CaC~2~)."
    )
    stmts_c2_p2 = [
        ("a) CaC~2~ tác dụng được với nước có thể tạo ra lượng lớn acetylene (C~2~H~2~), là chất khí dễ cháy nổ.", "ĐÚNG"),
        ("b) Khi acetylene được đốt cháy trong không khí ở nhiệt độ cao, nó phát ra nhiều nhiệt và có thể làm đám cháy lớn hơn và toả ra nhiều nhiệt hơn nên gây nguy hiểm cho người dập tắt đám cháy và môi trường xung quanh.", "ĐÚNG"),
        ("c) Phản ứng của CaC~2~ với nước: CaC~2~ + 2H~2~O -> Ca(OH)~2~ + C~2~H~2~.", "ĐÚNG"),
        ("d) Phản ứng đốt cháy C~2~H~2~ với oxygen: 2C~2~H~2~ + 5O~2~ --t°--> 4CO~2~ + 2H~2~O, Δ~r~H°~298~ = -2602 kJ.", "ĐÚNG")
    ]
    for stmt, ans in stmts_c2_p2:
        p_st = add_p(doc, space_before=1, space_after=1)
        p_st.paragraph_format.left_indent = Cm(0.5)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- Đất đèn CaC~2~ phản ứng mãnh liệt với nước sinh ra khí acetylene (C~2~H~2~).\n"
            "- Acetylene cháy toả lượng nhiệt rất lớn (-2602 kJ cho 2 mol), làm đám cháy bùng phát dữ dội hơn và có nguy cơ nổ bình dập.\n"
            "=> Cả 4 ý a, b, c, d đều ĐÚNG."
        )

    # Câu 3
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 3. Nitrobenzene là sản phẩm hữu cơ được tạo thành khi nitro hoá benzene bằng hỗn hợp HNO~3~ đặc và H~2~SO~4~ đặc ở nhiệt độ ≤ 50 °C."
    )
    stmts_c3_p2 = [
        ("a) Nitrobenzene được ứng dụng trong nông nghiệp để tổng hợp chất kích thích tăng trưởng thực vật.", "SAI"),
        ("b) Nitrobenzene là nguyên liệu sản xuất aniline, tiền chất để tổng hợp polyaniline, cao su nhân tạo.", "ĐÚNG"),
        ("c) Nitrobenzene là nguyên liệu để sản xuất thuốc nổ TNB sử dụng trong quân sự, công nghiệp và khai thác mỏ.", "SAI"),
        ("d) Nitrobenzene là chất lỏng, màu trắng, không tan trong nước, sánh như dầu.", "SAI")
    ]
    for stmt, ans in stmts_c3_p2:
        p_st = add_p(doc, space_before=1, space_after=1)
        p_st.paragraph_format.left_indent = Cm(0.5)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- a) Sai: Nitrobenzene rất độc hại, không dùng làm chất kích thích tăng trưởng cây trồng.\n"
            "- b) Đúng: Khử nitrobenzene bằng H~2~/kim loại thu được aniline, nguyên liệu then chốt điều chế polyaniline và cao su nhân tạo.\n"
            "- c) Sai: TNB là trinitrobenzene, nitro hoá benzene chỉ dừng lại ở mono-nitrobenzene ở điều kiện thường; TNB khó điều chế và ít dùng hơn TNT.\n"
            "- d) Sai: Nitrobenzene là chất lỏng màu vàng nhạt (không phải màu trắng), mùi hạnh nhân đắng.\n"
            "=> a: SAI, b: ĐÚNG, c: SAI, d: SAI."
        )

    # Câu 4
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 4. Trước những năm 50 của thế kỷ XX, công nghiệp tổng hợp hữu cơ dựa trên nguyên liệu chính là acetylene. Ngày nay, người ta thường dùng ethylene."
    )
    stmts_c4_p2 = [
        ("a) Ethylene là nguyên liệu rẻ hơn, tiện lợi hơn nhiều so với acetylene (ethylene thu được từ quá trình khai thác và chế biến dầu mỏ).", "ĐÚNG"),
        ("b) Acetylene là nguyên liệu rẻ hơn, tiện lợi hơn nhiều so với ethylene (ethylene thu được từ quá trình khai thác và chế biến dầu mỏ).", "SAI"),
        ("c) Phương pháp điều chế các monomer để tổng hợp các polymer đi từ ethylene kinh tế hơn và ít ảnh hưởng đến môi trường.", "ĐÚNG"),
        ("d) Sơ đồ điều chế vinyl chloride từ ethylene dưới đây có chất thải ra môi trường chỉ là nước:", "ĐÚNG")
    ]
    for stmt, ans in stmts_c4_p2:
        p_st = add_p(doc, space_before=1, space_after=1)
        p_st.paragraph_format.left_indent = Cm(0.5)
        render_statement(p_st, stmt, ans=ans, is_teacher=is_teacher)
    add_centered_picture(doc, os.path.join(IMG_DIR, "so_do_vinyl_chloride.png"), width_cm=13.5)
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- Khai thác dầu khí cung cấp nguồn ethylene dồi dào, giá rẻ hơn hẳn acetylene.\n"
            "- Sơ đồ oxy-clo hoá khép kín tận dụng hoàn toàn HCl sinh ra: 2CH~2~=CH~2~ + Cl~2~ + 1/2 O~2~ -> 2CH~2~=CHCl + H~2~O. "
            "Chất thải duy nhất thải ra môi trường là nước H~2~O.\n"
            "=> a: ĐÚNG, b: SAI, c: ĐÚNG, d: ĐÚNG."
        )

    # =========================================================================
    # PHẦN III
    # =========================================================================
    add_section_header(
        doc,
        title="PHẦN III. Câu trắc nghiệm yêu cầu trả lời ngắn.",
        subtitle="Thí sinh trả lời từ câu 1 đến câu 6."
    )

    # Câu 1
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 1. Số dẫn xuất monobromine tối đa thu được khi cho isopentane tác dụng với bromine trong điều kiện có chiếu sáng là bao nhiêu?"
    )
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Isopentane có công thức (CH~3~)~2~CH-CH~2~-CH~3~ có 4 vị trí C không tương đương:\n"
            "1) CH~2~Br-CH(CH~3~)-CH~2~-CH~3~ (1-bromo-2-methylbutane)\n"
            "2) CH~3~-CBr(CH~3~)-CH~2~-CH~3~ (2-bromo-2-methylbutane, sản phẩm chính)\n"
            "3) CH~3~-CH(CH~3~)-CHBr-CH~3~ (2-bromo-3-methylbutane)\n"
            "4) CH~3~-CH(CH~3~)-CH~2~-CH~2~Br (1-bromo-3-methylbutane)\n"
            "=> Đáp số: 4."
        )

    # Câu 2
    p = add_p(doc)
    render_question_prompt(p, "Câu 2. Có bao nhiêu đồng phân alkyne ứng với công thức phân tử C~5~H~8~?")
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Các đồng phân alkyne mạch hở C~5~H~8~:\n"
            "(1) CH≡C-CH~2~-CH~2~-CH~3~ (pent-1-yne)\n"
            "(2) CH~3~-C≡C-CH~2~-CH~3~ (pent-2-yne)\n"
            "(3) CH≡C-CH(CH~3~)~2~ (3-methylbut-1-yne)\n"
            "=> Đáp số: 3."
        )

    # Câu 3
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 3. Terephthalic acid (_p_-HOOCC~6~H~4~COOH) là nguyên liệu để sản xuất poly(ethylene terephthalate) (PET, một loại polymer quan trọng được sử dụng làm sợi dệt và chai nhựa,...). Cho các phát biểu sau:\n"
        "a) Phương trình hoá học của phản ứng điều chế terephthalic acid từ _p_-xylene:"
    )
    add_centered_picture(doc, os.path.join(IMG_DIR, "so_do_terephthalic_acid.png"), width_cm=11.0)
    p_extra = add_p(doc, space_before=2, space_after=3)
    render_rich_text(p_extra, "b) _p_-xylene còn có tên gọi khác là 1,3-dimethylbenzene.\n"
                              "c) Terephthalic acid (_p_-HOOCC~6~H~4~COOH) là một chất hữu cơ tạp chức.\n"
                              "d) Trong phân tử _p_-xylene có chứa 8 nguyên tử carbon.\n"
                              "Số phát biểu đúng là bao nhiêu?")
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- a) Đúng: Oxi hoá _p_-xylene bằng xúc tác tạo terephthalic acid và nước.\n"
            "- b) Sai: _p_-xylene là 1,4-dimethylbenzene.\n"
            "- c) Sai: Terephthalic acid chỉ chứa một loại nhóm chức -COOH nên là chất đa chức, không phải tạp chức.\n"
            "- d) Đúng: _p_-xylene có CTPT là C~8~H~10~, chứa 8 nguyên tử carbon.\n"
            "=> Có 2 phát biểu đúng (a và d). Đáp số: 2."
        )

    # Câu 4
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 4. Acetylene dùng trong đèn xì oxygen-acetylene để hàn, cắt kim loại. Trong công nghiệp acetylene là nguyên liệu để sản xuất poly(vinyl chloride) dùng sản xuất nhựa PVC và nhiều hóa chất khác. Acetylene cháy theo phương trình hóa học sau:\n"
        "C~2~H~2~(g) + 2,5O~2~(g) --t°--> 2CO~2~(g) + H~2~O(g)\n"
        "Dựa vào nhiệt tạo thành chuẩn ở 298 K của các chất: Δ~f~H°~298~ của C~2~H~2~(g) = +227,0 kJ/mol; CO~2~(g) = -393,5 kJ/mol; H~2~O(g) = -241,82 kJ/mol. "
        "Hãy cho biết biến thiên enthalpy của phản ứng cháy trên là a kJ. Xác định a (làm tròn đến hàng đơn vị)."
    )
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Biến thiên enthalpy phản ứng cháy:\n"
            "Δ~r~H°~298~ = 2 × Δ~f~H°~298~(CO~2~, g) + 1 × Δ~f~H°~298~(H~2~O, g) - Δ~f~H°~298~(C~2~H~2~, g) - 2,5 × Δ~f~H°~298~(O~2~, g)\n"
            "= 2 × (-393,5) + (-241,82) - (+227,0) - 0 = -787,0 - 241,82 - 227,0 = -1255,82 kJ.\n"
            "Làm tròn đến hàng đơn vị: a = -1256.\n"
            "=> Đáp số: -1256 (hoặc -1255,8)."
        )

    # Câu 5
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 5. Khí gas chứa chủ yếu các thành phần chính: Propane (C~3~H~8~), butane (C~4~H~10~) và một số thành phần khác. Để tạo mùi cho gas nhà sản xuất đã pha trộn thêm chất tạo mùi đặc trưng như methanethiol (CH~3~SH), có mùi giống tỏi, hành tây. Thành phần khí gas, tỉ lệ hòa trộn phổ biến của propane : butane theo thứ tự 30 : 70 đến 50 : 50.\n"
        "Cho phương trình nhiệt hóa học sau:\n"
        "C~3~H~8~(g) + 5O~2~(g) --t°--> 3CO~2~(g) + 4H~2~O(g)    Δ~r~H°~298~ = -2220 kJ\n"
        "C~4~H~10~(g) + 13/2 O~2~(g) --t°--> 4CO~2~(g) + 5H~2~O(g)    Δ~r~H°~298~ = -2874 kJ\n"
        "Giả sử một hộ gia đình cần 6 000 kJ nhiệt mỗi ngày, sau bao nhiêu ngày sẽ sử dụng hết 1 bình gas 12 kg với tỉ lệ khối lượng propane : butane là 30 : 70 (thành phần khác không đáng kể)? (Biết hiệu suất hấp thụ nhiệt khoảng 60%, làm tròn đến hàng đơn vị)."
    )
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- Khối lượng từng khí trong bình 12 kg:\n"
            "  m(C~3~H~8~) = 12 × 0,30 = 3,6 kg = 3600 g => n(C~3~H~8~) = 3600 / 44 ≈ 81,818 mol.\n"
            "  m(C~4~H~10~) = 12 × 0,70 = 8,4 kg = 8400 g => n(C~4~H~10~) = 8400 / 58 ≈ 144,828 mol.\n"
            "- Tổng nhiệt lượng toả ra khi đốt cháy hết bình gas:\n"
            "  Q(toả) = 81,818 × 2220 + 144,828 × 2874 ≈ 181 636 + 416 236 = 597 872 kJ.\n"
            "- Nhiệt lượng có ích thực tế gia đình nhận được (H = 60%):\n"
            "  Q(có ích) = 597 872 × 60% ≈ 358 723 kJ.\n"
            "- Số ngày sử dụng:\n"
            "  N = 358 723 / 6000 ≈ 59,79 ≈ 60 ngày.\n"
            "=> Đáp số: 60."
        )

    # Câu 6
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 6. Một loại xăng chứa hỗn hợp hexane, heptane và 2,2,4-trimethylpentane (isooctane). Hóa hơi lượng xăng này được hơi xăng có tỉ khối so với H~2~ là 54,9. "
        "Tính tỉ lệ thể tích hơi xăng và không khí (20% thể tích O~2~, 80% thể tích N~2~) vừa đủ đốt cháy hết lượng xăng này là n~xăng~ / n~không khí~ = 1 / x. Xác định x (làm tròn đến hai chữ số thập phân)."
    )
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- Hỗn hợp xăng gồm các alkane có công thức chung C~n~H~2n+2~ (n trung bình).\n"
            "  M = 54,9 × 2 = 109,8 g/mol => 14n + 2 = 109,8 => n = 7,7.\n"
            "- Phương trình đốt cháy: C~n~H~2n+2~ + [(3n+1)/2] O~2~ -> n CO~2~ + (n+1) H~2~O.\n"
            "  n(O~2~) / n(xăng) = (3 × 7,7 + 1) / 2 = 12,05.\n"
            "- Không khí chứa 20% thể tích O~2~ => n(không khí) = 5 × n(O~2~) = 5 × 12,05 × n(xăng) = 60,25 × n(xăng).\n"
            "  => n(xăng) / n(không khí) = 1 / 60,25 => x = 60,25.\n"
            "=> Đáp số: 60,25."
        )

    # =========================================================================
    # PHẦN IV
    # =========================================================================
    add_section_header(
        doc,
        title="PHẦN IV. Câu hỏi tự luận.",
        subtitle="Thí sinh trình bày lời giải chi tiết vào giấy làm bài."
    )

    # Câu 1
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 1. (SBT - CTST) Nhu cầu sử dụng PVC trên toàn thế giới liên tục tăng trong các năm qua. Để thu được PVC, cần đi từ monomer là vinyl chloride. Có thể điều chế vinyl chloride từ acetylene hoặc ethylene. Một trong những cách điều chế vinyl chloride từ ethylene hiện nay là theo sơ đồ:\n"
        "C~2~H~4~ --(+Cl~2~)--> (X) --(-HCl, 500 °C)--> CH~2~=CH-Cl\n"
        "Viết các phương trình phản ứng minh hoạ cho sơ đồ trên."
    )
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "Các phương trình phản ứng minh hoạ:\n"
            "(1) Cộng chlorine vào ethylene tạo 1,2-dichloroethane (chất X):\n"
            "    CH~2~=CH~2~ + Cl~2~ --(FeCl~3~, 50-60 °C)--> CH~2~Cl-CH~2~Cl\n"
            "(2) Nhiệt phân tách HCl ở 500 °C tạo vinyl chloride:\n"
            "    CH~2~Cl-CH~2~Cl --(500 °C)--> CH~2~=CH-Cl + HCl\n"
            "(Lưu ý: trong phòng thí nghiệm có thể đun nóng CH~2~Cl-CH~2~Cl với KOH trong ethanol: CH~2~Cl-CH~2~Cl + KOH --(C~2~H~5~OH, t°)--> CH~2~=CH-Cl + KCl + H~2~O)."
        )
    else:
        add_student_writing_lines(doc, num_lines=8)

    # Câu 2
    p = add_p(doc)
    render_question_prompt(
        p,
        "Câu 2. (SBT-CD) Khí đốt hoá lỏng (Liquified Petroleum Gas, viết tắt là LPG) hay còn được gọi là gas, là hỗn hợp khí chủ yếu gồm propane (C~3~H~8~) và butane (C~4~H~10~) đã được hoá lỏng. Một loại gas dân dụng chứa khí hoá lỏng có tỉ lệ mol propane : butane là 40 : 60. Đốt cháy 1 lít khí gas này (ở 25 °C, 1 bar) thì toả ra một lượng nhiệt bằng bao nhiêu? Biết khi đốt cháy hoàn toàn 1 mol mỗi chất propane và butane toả ra lượng nhiệt tương ứng là 2 220 kJ và 2 875 kJ."
    )
    if is_teacher:
        add_solution_no_table(
            doc,
            "Lời giải chi tiết:\n"
            "- Ở điều kiện chuẩn (25 °C, 1 bar), 1 mol khí có thể tích là 24,79 lít.\n"
            "  Tổng số mol khí trong 1 lít gas: n(gas) = 1 / 24,79 ≈ 0,040339 mol.\n"
            "- Số mol từng khí theo tỉ lệ mol 40 : 60:\n"
            "  n(propane) = 0,040339 × 0,40 ≈ 0,016136 mol.\n"
            "  n(butane) = 0,040339 × 0,60 ≈ 0,024203 mol.\n"
            "- Lượng nhiệt toả ra khi đốt cháy 1 lít khí gas:\n"
            "  Q = 0,016136 × 2220 + 0,024203 × 2875 ≈ 35,822 + 69,584 = 105,41 kJ.\n"
            "  (Kết luận: Khi đốt cháy 1 lít khí gas toả ra khoảng 105,4 kJ nhiệt lượng)."
        )
    else:
        add_student_writing_lines(doc, num_lines=9)

    return doc

def check_and_save():
    os.makedirs(OUT_DIR, exist_ok=True)

    # 1. Bản Học sinh
    doc_hs = build_exam_document(is_teacher=False)
    tables_hs = len(doc_hs.tables)
    print(f"Bản Học sinh có {tables_hs} bảng (Chỉ được phép có 1 bảng Header block).")
    assert tables_hs <= 1, "LỖI: Bản Học sinh chứa bảng trong nội dung!"
    path_hs = os.path.join(OUT_DIR, "De_Kiem_Tra_Hoa_11_De12.docx")
    doc_hs.save(path_hs)
    print("Đã lưu Bản Học sinh:", path_hs)

    # 2. Bản Giáo viên
    doc_gv = build_exam_document(is_teacher=True)
    tables_gv = len(doc_gv.tables)
    print(f"Bản Giáo viên có {tables_gv} bảng (Chỉ được phép có 1 bảng Header block).")
    assert tables_gv <= 1, "LỖI: Bản Giáo viên chứa bảng trong nội dung!"
    path_gv = os.path.join(OUT_DIR, "De_Kiem_Tra_Hoa_11_De12_GiaoVien.docx")
    doc_gv.save(path_gv)
    print("Đã lưu Bản Giáo viên:", path_gv)

if __name__ == "__main__":
    check_and_save()
