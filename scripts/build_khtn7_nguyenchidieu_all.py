# -*- coding: utf-8 -*-
"""
Script biên soạn trọn bộ 3 file:
1. Đề cương Học sinh (.docx)
2. Đề cương Giáo viên có lời giải chi tiết (.docx)
3. Slide Dạy trên màn hình Tivi 65 inch tỉ lệ 16:9 (.pdf và .docx) - KHÔNG CÓ ĐÁP ÁN

ĐỀ CƯƠNG ÔN TẬP KIỂM TRA GIỮA HỌC KỲ I - MÔN KHOA HỌC TỰ NHIÊN 7
NĂM HỌC 2026 - 2027
TRƯỜNG THCS NGUYỄN CHÍ DIỂU
(Trọn vẹn 3 phân môn: KHTN 1 - Vật lí, KHTN 2 - Hóa học, KHTN 3 - Sinh học)
"""

import sys
if sys.stdout:
    sys.stdout.reconfigure(encoding='utf-8')
if sys.stderr:
    sys.stderr.reconfigure(encoding='utf-8')

import os
import shutil
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

# ==============================================================================
# HẰNG SỐ MÀU SẮC & ĐỊNH DẠNG THEO GEMINI.md (FORM ĐỀ KIỂM TRA)
# ==============================================================================
COLOR_PRIMARY_HEX = "D32F2F"       # Đỏ san hô (Coral Red)
COLOR_SECONDARY_HEX = "8BC390"     # Xanh xô thơm (Sage Green)
COLOR_TEXT_HEX = "1A1A1A"          # Đen than nhẹ (Off-black)
COLOR_SOLUTION_TITLE_HEX = "1565C0"# Xanh dương đậm cho tiêu đề lời giải
COLOR_GRAY_HEX = "757575"

COLOR_PRIMARY = RGBColor(0xD3, 0x2F, 0x2F)
COLOR_SECONDARY = RGBColor(0x8B, 0xC3, 0x90)
COLOR_TEXT = RGBColor(0x1A, 0x1A, 0x1A)
COLOR_SOLUTION_TITLE = RGBColor(0x15, 0x65, 0xC0)
COLOR_GRAY = RGBColor(0x75, 0x75, 0x75)

FONT_NAME = "Times New Roman"

BASE_DIR = r"c:\Antigravity_Thanh"
SAN_PHAM_DIR = os.path.join(BASE_DIR, "San_Pham")
SCRATCH_DIR = os.path.join(BASE_DIR, "scratch")
IMG_DIR = os.path.join(SAN_PHAM_DIR, "images_khtn7_nguyenchidieu")
IMG_CAU4 = os.path.join(IMG_DIR, "dothi_cau4_khtn7_nguyenchidieu.png")
IMG_CAU3 = os.path.join(IMG_DIR, "dothi_cau3_khtn7_nguyenchidieu.png")

ONEDRIVE_DIR1 = r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn7\KHTN7-IN-DEKIEMTRADINHKI"
ONEDRIVE_DIR2 = r"c:\Users\Admin\OneDrive\Tệp đính kèm\2026-2027\khtn7"

os.makedirs(SAN_PHAM_DIR, exist_ok=True)
os.makedirs(SCRATCH_DIR, exist_ok=True)
os.makedirs(IMG_DIR, exist_ok=True)

# ==============================================================================
# HÀM BỔ TRỢ ĐỊNH DẠNG WORD (.docx) CHUẨN GEMINI.md
# ==============================================================================

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders_horizontal_only(table, border_color_hex="8BC390", border_sz="6"):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color_hex}"/>\n'
        f'  <w:bottom w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color_hex}"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:insideH w:val="none"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.59)    # 1.5 cm
        s.bottom_margin = Inches(0.59) # 1.5 cm
        s.left_margin = Inches(0.59)   # 1.5 cm
        s.right_margin = Inches(0.59)  # 1.5 cm
        s.header_distance = Inches(0.24) # 0.6 cm
        s.footer_distance = Inches(0.24) # 0.6 cm

        # Header Table
        header = s.header
        header.is_linked_to_previous = False
        tbl_hdr = header.add_table(rows=1, cols=2, width=Inches(7.08))
        tbl_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders_horizontal_only(tbl_hdr, border_color_hex=COLOR_SECONDARY_HEX, border_sz="6")

        c0 = tbl_hdr.cell(0, 0)
        c0.width = Inches(3.54)
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r0 = p0.add_run("Tài liệu lưu hành nội bộ")
        r0.font.name = FONT_NAME
        r0.font.size = Pt(10)
        r0.font.italic = True
        r0.font.color.rgb = COLOR_PRIMARY

        c1 = tbl_hdr.cell(0, 1)
        c1.width = Inches(3.54)
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r1 = p1.add_run("Thầy ")
        r1.font.name = FONT_NAME
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_PRIMARY
        r1_b = p1.add_run("TRẦN VĂN THẠNH")
        r1_b.font.name = FONT_NAME
        r1_b.font.size = Pt(10)
        r1_b.font.bold = True
        r1_b.font.color.rgb = COLOR_PRIMARY
        r1_ph = p1.add_run(" - 0777.470.803")
        r1_ph.font.name = FONT_NAME
        r1_ph.font.size = Pt(10)
        r1_ph.font.color.rgb = COLOR_PRIMARY

        # Footer Table
        footer = s.footer
        footer.is_linked_to_previous = False
        tbl_ftr = footer.add_table(rows=1, cols=2, width=Inches(7.08))
        tbl_ftr.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders_horizontal_only(tbl_ftr, border_color_hex=COLOR_SECONDARY_HEX, border_sz="6")

        fc0 = tbl_ftr.cell(0, 0)
        fc0.width = Inches(5.3)
        fp0 = fc0.paragraphs[0]
        fp0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fr0 = fp0.add_run("CS1: 6/15 Nguyễn Hoàng, Kim Long   |   CS2: 24 Đặng Thái Thân, TP Huế")
        fr0.font.name = FONT_NAME
        fr0.font.size = Pt(9.5)
        fr0.font.italic = True
        fr0.font.color.rgb = COLOR_PRIMARY

        fc1 = tbl_ftr.cell(0, 1)
        fc1.width = Inches(1.78)
        fp1 = fc1.paragraphs[0]
        fp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fr1 = fp1.add_run("Trang ")
        fr1.font.name = FONT_NAME
        fr1.font.size = Pt(10)
        fr1.font.bold = True
        fr1.font.color.rgb = COLOR_PRIMARY

        fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        fp1._p.append(fldSimple)

def add_header_block(doc, is_teacher=False, is_screen=False):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c0 = tbl.cell(0, 0)
    c0.width = Inches(3.2)
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.line_spacing = 1.15
    p0.paragraph_format.space_after = Pt(2)

    r_ubnd = p0.add_run("TRƯỜNG THCS NGUYỄN CHÍ DIỂU\n")
    r_ubnd.font.name = FONT_NAME
    r_ubnd.font.size = Pt(10.5)
    r_ubnd.font.bold = True
    r_ubnd.font.color.rgb = COLOR_TEXT

    r_nh = p0.add_run("NĂM HỌC 2026 - 2027\n")
    r_nh.font.name = FONT_NAME
    r_nh.font.size = Pt(10)
    r_nh.font.italic = True

    r_de = p0.add_run("ĐỀ CƯƠNG CHÍNH THỨC")
    r_de.font.name = FONT_NAME
    r_de.font.size = Pt(10)
    r_de.font.bold = True
    r_de.font.color.rgb = COLOR_PRIMARY

    c1 = tbl.cell(0, 1)
    c1.width = Inches(3.88)
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.line_spacing = 1.15
    p1.paragraph_format.space_after = Pt(2)

    r_title = p1.add_run("ĐỀ CƯƠNG ÔN TẬP KIỂM TRA GIỮA HỌC KỲ I\n")
    r_title.font.name = FONT_NAME
    r_title.font.size = Pt(11)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY

    r_sub = p1.add_run("MÔN: KHOA HỌC TỰ NHIÊN - LỚP 7\n")
    r_sub.font.name = FONT_NAME
    r_sub.font.size = Pt(11)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_PRIMARY

    r_note = p1.add_run("(Trọn bộ 3 phân môn: Vật lí - Hóa học - Sinh học)")
    r_note.font.name = FONT_NAME
    r_note.font.size = Pt(10)
    r_note.font.bold = True
    r_note.font.color.rgb = COLOR_SOLUTION_TITLE

    # Notice
    if is_teacher:
        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_t.paragraph_format.space_before = Pt(4)
        p_t.paragraph_format.space_after = Pt(6)
        r_tn = p_t.add_run("(BẢN GIÁO VIÊN -- CÓ LỜI GIẢI CHI TIẾT VÀ ĐÁP ÁN ĐẦY ĐỦ)")
        r_tn.font.name = FONT_NAME
        r_tn.font.size = Pt(10.5)
        r_tn.font.bold = True
        r_tn.font.color.rgb = COLOR_PRIMARY
    elif is_screen:
        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_t.paragraph_format.space_before = Pt(4)
        p_t.paragraph_format.space_after = Pt(6)
        r_tn = p_t.add_run("(BẢN DẠY TRÊN MÀN HÌNH TIVI 65 INCH -- KHÔNG CÓ ĐÁP ÁN)")
        r_tn.font.name = FONT_NAME
        r_tn.font.size = Pt(10.5)
        r_tn.font.bold = True
        r_tn.font.color.rgb = COLOR_PRIMARY
    else:
        p_stu = doc.add_paragraph()
        p_stu.paragraph_format.space_before = Pt(4)
        p_stu.paragraph_format.space_after = Pt(6)
        r_s = p_stu.add_run("Họ và tên học sinh: ........................................................................ Lớp: .............. Số báo danh: .......................................")
        r_s.font.name = FONT_NAME
        r_s.font.size = Pt(10.5)
        r_s.font.italic = True

def add_part_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = COLOR_PRIMARY

def add_subpart_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = COLOR_SOLUTION_TITLE

def add_solution_paragraph(doc, text):
    lines = text.strip().split("\n")
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.space_before = Pt(2) if i == 0 else Pt(1)
        p.paragraph_format.space_after = Pt(2) if i == len(lines)-1 else Pt(1)
        p.paragraph_format.line_spacing = 1.15
        
        if i == 0:
            r_tag = p.add_run(">> Hướng dẫn giải chi tiết:\n")
            r_tag.font.name = FONT_NAME
            r_tag.font.size = Pt(10.5)
            r_tag.font.bold = True
            r_tag.font.color.rgb = COLOR_SOLUTION_TITLE
        
        parts = line.split("*")
        for idx, part in enumerate(parts):
            if not part:
                continue
            r = p.add_run(part)
            r.font.name = FONT_NAME
            r.font.size = Pt(10.5)
            if idx % 2 == 1:
                r.font.bold = True
                r.font.color.rgb = COLOR_PRIMARY
            else:
                r.font.color.rgb = COLOR_TEXT

def add_dotted_lines(doc, num_lines=4):
    for _ in range(num_lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("......................................................................................................................................................................................")
        r.font.name = FONT_NAME
        r.font.size = Pt(9.5)
        r.font.color.rgb = COLOR_GRAY

# ==============================================================================
# HÀM TẠO FILE WORD
# ==============================================================================
def create_word_document(mode="student"):
    doc = docx.Document()
    add_header_footer(doc)
    is_teacher = (mode == "teacher")
    is_screen = (mode == "screen")
    add_header_block(doc, is_teacher=is_teacher, is_screen=is_screen)

    # --------------------------------------------------------------------------
    # MỤC HƯỚNG DẪN DẠNG BÀI TẬP VÀ CÁCH LÀM
    # --------------------------------------------------------------------------
    p_guide = doc.add_paragraph()
    p_guide.paragraph_format.space_before = Pt(4)
    p_guide.paragraph_format.space_after = Pt(3)
    r_g = p_guide.add_run("HƯỚNG DẪN CẤU TRÚC ĐỀ VÀ PHƯƠNG PHÁP LÀM BÀI:")
    r_g.font.name = FONT_NAME
    r_g.font.size = Pt(11)
    r_g.font.bold = True
    r_g.font.color.rgb = COLOR_PRIMARY

    guides = [
        "1. Dạng 1: Trắc nghiệm nhiều lựa chọn (12 câu – 3,0 điểm): Đọc kĩ từ khóa trong câu hỏi (vật thể tự nhiên/nhân tạo, tính chất vật lí/hóa học, an toàn phòng thí nghiệm,...) để chọn phương án A, B, C hoặc D.",
        "2. Dạng 2: Trắc nghiệm Đúng/Sai (2 câu, mỗi câu 4 ý a, b, c, d – 2,0 điểm): Xác định từng nhận định là Đúng (ghi 'Đ') hoặc Sai (ghi 'S').",
        "3. Dạng 3: Trắc nghiệm trả lời ngắn (8 câu – 2,0 điểm): Trả lời dưới dạng con số (tối đa 4 kí tự kể cả dấu phẩy, ví dụ: 0,24).",
        "4. Dạng 4: Tự luận (4 – 5 câu – 3,0 điểm): Đọc kĩ câu hỏi, trả lời đúng trọng tâm, trình bày rõ ràng, sạch đẹp, có lời dẫn cụ thể."
    ]
    for g_text in guides:
        p_gd = doc.add_paragraph()
        p_gd.paragraph_format.left_indent = Inches(0.15)
        p_gd.paragraph_format.space_before = Pt(1)
        p_gd.paragraph_format.space_after = Pt(1)
        p_gd.paragraph_format.line_spacing = 1.15
        r_gt = p_gd.add_run(g_text)
        r_gt.font.name = FONT_NAME
        r_gt.font.size = Pt(10)

    # --------------------------------------------------------------------------
    # A. PHÂN MÔN KHTN 1: VẬT LÍ (TỐC ĐỘ & ĐỒ THỊ QUÃNG ĐƯỜNG - THỜI GIAN)
    # --------------------------------------------------------------------------
    add_part_heading(doc, "A. PHÂN MÔN KHTN 1 (VẬT LÍ)")
    add_subpart_heading(doc, "I. TÓM TẮT LÍ THUYẾT TRỌNG TÂM")
    phys_theories = [
        "1. Phương pháp và kĩ năng học tập môn KHTN: Nắm vững phương pháp tìm hiểu tự nhiên, thực hiện đúng các kĩ năng tiến trình: quan sát, phân loại, liên kết, đo, dự báo. Sử dụng thành thạo một số dụng cụ đo thông dụng.",
        "2. Tốc độ chuyển động: Tốc độ là đại lượng cho biết mức độ nhanh hay chậm của chuyển động, xác định bằng chiều dài quãng đường đi được trong một đơn vị thời gian. Công thức: v = s / t (s là quãng đường, t là thời gian, v là tốc độ). Đơn vị thường dùng: m/s, km/h (1 m/s = 3,6 km/h; 1 km/h = 1/3,6 m/s).",
        "3. Đồ thị quãng đường – thời gian (s - t): Mô tả mối liên hệ giữa quãng đường đi được và thời gian chuyển động. Đồ thị chuyển động thẳng đều là đường thẳng xiên góc đi qua gốc tọa độ. Đồ thị nằm ngang biểu diễn vật đứng yên.",
        "4. Phương pháp đo tốc độ: Mô tả sơ lược cách đo tốc độ bằng đồng hồ bấm giây và cổng quang điện trong phòng thí nghiệm; cách đo bằng thiết bị 'bắn tốc độ' (camera kết hợp radar/laser) của cảnh sát giao thông.",
        "5. Tốc độ và an toàn giao thông: Tốc độ ảnh hưởng trực tiếp đến khoảng cách an toàn và quãng đường phanh dừng xe. Người điều khiển phương tiện phải tuân thủ nghiêm quy định về tốc độ và giữ khoảng cách an toàn để tránh tai nạn."
    ]
    for pt in phys_theories:
        p_pt = doc.add_paragraph()
        p_pt.paragraph_format.left_indent = Inches(0.15)
        p_pt.paragraph_format.space_before = Pt(2)
        p_pt.paragraph_format.space_after = Pt(2)
        p_pt.paragraph_format.line_spacing = 1.15
        r_pt = p_pt.add_run(f"• {pt}")
        r_pt.font.name = FONT_NAME
        r_pt.font.size = Pt(10.5)

    add_subpart_heading(doc, "II. BÀI TẬP VẬN DỤNG")
    
    # Câu 1
    p_c1 = doc.add_paragraph()
    p_c1.paragraph_format.space_before = Pt(4)
    p_c1.paragraph_format.space_after = Pt(2)
    p_c1.paragraph_format.line_spacing = 1.15
    r_c1 = p_c1.add_run("Câu 1. ")
    r_c1.font.name = FONT_NAME
    r_c1.font.size = Pt(10.5)
    r_c1.font.bold = True
    p_c1.add_run("Thực hiện các yêu cầu sau:\n"
                 "a) Một người đi quãng đường dài 1,5 km với tốc độ 20 m/s. Thời gian để người đó đi hết quãng đường là bao nhiêu?\n"
                 "b) Khi nói tốc độ của vật A là 36 km/h, con số đó cho biết điều gì?\n"
                 "c) Đổi các đơn vị sau: 54 km/h = …………….. m/s;   10 m/s = …………….. km/h.")
    if is_teacher:
        add_solution_paragraph(doc, (
            "a) Đổi s = 1,5 km = 1500 m. Thời gian người đó đi hết quãng đường là:\n"
            "   t = s / v = 1500 / 20 = *75 (s)* = *1 phút 15 giây*.\n"
            "b) Con số 36 km/h cho biết: Trong mỗi giờ, vật A đi được quãng đường là *36 km* (hoặc trong mỗi giây vật A đi được 10 m).\n"
            "c) Đổi đơn vị:\n"
            "   • 54 km/h = 54 / 3,6 = *15 m/s*.\n"
            "   • 10 m/s = 10 × 3,6 = *36 km/h*."
        ))
    elif not is_screen:
        add_dotted_lines(doc, num_lines=4)

    # Câu 2
    p_c2 = doc.add_paragraph()
    p_c2.paragraph_format.space_before = Pt(4)
    p_c2.paragraph_format.space_after = Pt(2)
    p_c2.paragraph_format.line_spacing = 1.15
    r_c2 = p_c2.add_run("Câu 2. ")
    r_c2.font.name = FONT_NAME
    r_c2.font.size = Pt(10.5)
    r_c2.font.bold = True
    p_c2.add_run("Hai người đi xe đạp. Người thứ nhất đi quãng đường 300 m hết 1 phút. Người thứ hai đi quãng đường 7,5 km hết 0,5 h. Người nào đi nhanh hơn?")
    if is_teacher:
        add_solution_paragraph(doc, (
            "• Tốc độ của người thứ nhất:\n"
            "  Đổi t_1 = 1 phút = 60 s.  v_1 = s_1 / t_1 = 300 / 60 = *5 (m/s)* = 5 × 3,6 = *18 (km/h)*.\n"
            "• Tốc độ của người thứ hai:\n"
            "  v_2 = s_2 / t_2 = 7,5 / 0,5 = *15 (km/h)*.\n"
            "• So sánh: Vì v_1 (18 km/h) > v_2 (15 km/h) nên *Người thứ nhất đi nhanh hơn người thứ hai*."
        ))
    elif not is_screen:
        add_dotted_lines(doc, num_lines=4)

    # Câu 3
    p_c3 = doc.add_paragraph()
    p_c3.paragraph_format.space_before = Pt(4)
    p_c3.paragraph_format.space_after = Pt(2)
    p_c3.paragraph_format.line_spacing = 1.15
    r_c3 = p_c3.add_run("Câu 3. ")
    r_c3.font.name = FONT_NAME
    r_c3.font.size = Pt(10.5)
    r_c3.font.bold = True
    p_c3.add_run("Dựa vào số liệu trong bảng, hãy vẽ đồ thị quãng đường - thời gian của vật chuyển động như bảng sau:\n"
                 "[BẢNG SỐ LIỆU] Thời gian t (h): 1 | 2 | 3 | 4;   Quãng đường s (km): 60 | 120 | 180 | 240.")
    if os.path.exists(IMG_CAU3) and is_teacher:
        p_img3 = doc.add_paragraph()
        p_img3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img3.paragraph_format.space_before = Pt(3)
        p_img3.paragraph_format.space_after = Pt(3)
        run_img3 = p_img3.add_run()
        run_img3.add_picture(IMG_CAU3, width=Inches(3.2))

    if is_teacher:
        add_solution_paragraph(doc, (
            "• Nhận xét: Tốc độ của vật là v = s / t = 60 / 1 = 120 / 2 = 180 / 3 = 240 / 4 = *60 (km/h)* không đổi (chuyển động thẳng đều).\n"
            "• Đồ thị quãng đường - thời gian là đường thẳng đi qua gốc tọa độ O(0,0) và các điểm (1h; 60km), (2h; 120km), (3h; 180km), (4h; 240km)."
        ))
    elif not is_screen:
        add_dotted_lines(doc, num_lines=4)

    # Câu 4
    p_c4 = doc.add_paragraph()
    p_c4.paragraph_format.space_before = Pt(4)
    p_c4.paragraph_format.space_after = Pt(2)
    p_c4.paragraph_format.line_spacing = 1.15
    r_c4 = p_c4.add_run("Câu 4. ")
    r_c4.font.name = FONT_NAME
    r_c4.font.size = Pt(10.5)
    r_c4.font.bold = True
    p_c4.add_run("Quan sát đồ thị hình vẽ bên và thực hiện các yêu cầu sau:\n"
                 "a) Cho biết:\n"
                 "   - Quãng đường vật đi được trong thời gian 2 s đầu tiên.\n"
                 "   - Thời gian vật đi được quãng đường 8 m.\n"
                 "b) Tính tốc độ vật đi được trong mỗi giai đoạn (A, B, C, D).")
    
    if os.path.exists(IMG_CAU4):
        p_img4 = doc.add_paragraph()
        p_img4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img4.paragraph_format.space_before = Pt(3)
        p_img4.paragraph_format.space_after = Pt(3)
        run_img4 = p_img4.add_run()
        run_img4.add_picture(IMG_CAU4, width=Inches(3.5))

    if is_teacher:
        add_solution_paragraph(doc, (
            "a) Đọc số liệu từ đồ thị:\n"
            "   • Trong 2 s đầu tiên (t = 0 đến t = 2 s), dóng lên đồ thị được điểm có tung độ s = *4 m*.\n"
            "   • Quãng đường 8 m tương ứng với điểm trên đồ thị có tung độ s = 8 m, dóng xuống trục hoành được thời gian t = *6 s*.\n"
            "b) Tính tốc độ của vật trong mỗi giai đoạn:\n"
            "   • Giai đoạn A (từ 0 đến 2 s): v_A = (4 - 0) / (2 - 0) = *2 (m/s)*.\n"
            "   • Giai đoạn B (từ 2 s đến 4 s): Đồ thị nằm ngang, quãng đường không đổi (s = 4 m), vật đứng yên ==> v_B = *0 (m/s)*.\n"
            "   • Giai đoạn C (từ 4 s đến 6 s): v_C = (8 - 4) / (6 - 4) = 4 / 2 = *2 (m/s)*.\n"
            "   • Giai đoạn D (từ 6 s đến 8 s): v_D = (10 - 8) / (8 - 6) = 2 / 2 = *1 (m/s)*."
        ))
    elif not is_screen:
        add_dotted_lines(doc, num_lines=5)

    # Câu 5
    p_c5 = doc.add_paragraph()
    p_c5.paragraph_format.space_before = Pt(4)
    p_c5.paragraph_format.space_after = Pt(2)
    p_c5.paragraph_format.line_spacing = 1.15
    r_c5 = p_c5.add_run("Câu 5. ")
    r_c5.font.name = FONT_NAME
    r_c5.font.size = Pt(10.5)
    r_c5.font.bold = True
    p_c5.add_run("Một người đi xe máy đều trên quãng đường đầu dài 18 km hết 30 phút. Sau đó, ở quãng đường sau dài 1 200 m, người đó đi hết 5 phút 20 giây. Tính tốc độ trung bình của người đó trên cả 2 quãng đường.")
    if is_teacher:
        add_solution_paragraph(doc, (
            "• Đổi đơn vị:\n"
            "  Quãng đường đầu: s_1 = 18 km = 18 000 m; thời gian t_1 = 30 phút = 1800 s.\n"
            "  Quãng đường sau: s_2 = 1 200 m; thời gian t_2 = 5 phút 20 giây = 5 × 60 + 20 = 320 s.\n"
            "• Tổng quãng đường đi được: s = s_1 + s_2 = 18 000 + 1 200 = 19 200 (m) = 19,2 (km).\n"
            "• Tổng thời gian chuyển động: t = t_1 + t_2 = 1800 + 320 = 2120 (s) = 53/90 (h).\n"
            "• Tốc độ trung bình trên cả hai quãng đường:\n"
            "  v_tb = s / t = 19 200 / 2120 ≈ *9,06 (m/s)* ≈ *32,6 (km/h)*."
        ))
    elif not is_screen:
        add_dotted_lines(doc, num_lines=4)

    # --------------------------------------------------------------------------
    # B. PHÂN MÔN KHTN 2: HÓA HỌC (NGUYÊN TỬ & NGUYÊN TỐ HÓA HỌC)
    # --------------------------------------------------------------------------
    add_part_heading(doc, "B. PHÂN MÔN KHTN 2 (HÓA HỌC)")
    chem_questions = [
        {
            "id": 1,
            "title": "Câu 1",
            "content": (
                "a. Nguyên tử là gì?\n"
                "b. Khối lượng nguyên tử là gì? 1 amu có khối lượng bằng bao nhiêu gam?\n"
                "c. Nêu tóm tắt mô hình nguyên tử theo Rutherford - Bohr.\n"
                "d. Trình bày cấu tạo nguyên tử."
            ),
            "solution": (
                "a. *Nguyên tử* là hạt vô cùng nhỏ, trung hòa về điện, cấu tạo nên mọi chất.\n"
                "b. *Khối lượng nguyên tử* là khối lượng của một nguyên tử, tính theo đơn vị amu (atomic mass unit). 1 amu có khối lượng xấp xỉ bằng *1,6605 × 10^-24 gam*.\n"
                "c. *Mô hình Rutherford - Bohr*: Nguyên tử gồm hạt nhân ở tâm mang điện tích dương và các electron mang điện tích âm chuyển động xung quanh hạt nhân trên những quỹ đạo xác định tạo thành các lớp electron.\n"
                "d. *Cấu tạo nguyên tử*:\n"
                "   • Hạt nhân: ở tâm nguyên tử, gồm các hạt *proton* (p, mang điện tích +1) và *neutron* (n, không mang điện).\n"
                "   • Vỏ nguyên tử: gồm các *electron* (e, mang điện tích -1) chuyển động quanh hạt nhân.\n"
                "   • Trong nguyên tử: *Số proton = Số electron* (p = e). Khối lượng nguyên tử tập trung hầu hết ở hạt nhân."
            )
        },
        {
            "id": 2,
            "title": "Câu 2",
            "content": (
                "a. Nguyên tố hóa học là gì? Nêu cách biểu diễn kí hiệu hóa học.\n"
                "b. Học thuộc tên, kí hiệu hóa học của 20 nguyên tố ở bảng 3.1 SGK."
            ),
            "solution": (
                "a. • *Nguyên tố hóa học* là tập hợp những nguyên tử cùng loại, có cùng số proton trong hạt nhân.\n"
                "   • *Cách biểu diễn*: Mỗi nguyên tố được biểu diễn bằng 1 hoặc 2 chữ cái (chữ đầu viết in hoa, chữ thứ hai nếu có viết thường), gọi là kí hiệu hóa học (ví dụ: H, C, O, Na, Fe,...).\n"
                "b. *20 nguyên tố đầu tiên*: 1. H (Hydrogen), 2. He (Helium), 3. Li (Lithium), 4. Be (Beryllium), 5. B (Boron), 6. C (Carbon), 7. N (Nitrogen), 8. O (Oxygen), 9. F (Fluorine), 10. Ne (Neon), 11. Na (Sodium), 12. Mg (Magnesium), 13. Al (Aluminium), 14. Si (Silicon), 15. P (Phosphorus), 16. S (Sulfur), 17. Cl (Chlorine), 18. Ar (Argon), 19. K (Potassium), 20. Ca (Calcium)."
            )
        },
        {
            "id": 3,
            "title": "Câu 3",
            "content": (
                "a. Vẽ mô hình cấu tạo nguyên tử Calcium có điện tích hạt nhân là +20. Dựa vào mô hình em hãy cho biết số proton, số electron, số lớp electron của nguyên tử Calcium.\n"
                "b. Điền thông tin còn thiếu vào bảng sau:\n"
                "[BẢNG] Cột: Nguyên tử | Điện tích hạt nhân | Số p | Số e | Số n | Số lớp e | Số e lớp ngoài cùng | Khối lượng nguyên tử\n"
                "Hàng 1: Sodium | ? | 11 | ? | 12 | ? | ? | ?\n"
                "Hàng 2: Oxygen | +8 | ? | ? | 8 | ? | ? | ?"
            ),
            "solution": (
                "a. *Nguyên tử Calcium (Ca)*:\n"
                "   • Điện tích hạt nhân là +20 ==> Số proton = *20*, số electron = *20*.\n"
                "   • Sự phân bố electron vào 4 lớp: Lớp 1 có *2e*, Lớp 2 có *8e*, Lớp 3 có *8e*, Lớp 4 có *2e*.\n"
                "   • Số lớp electron là *4*; số electron ở lớp ngoài cùng là *2*.\n"
                "b. *Hoàn thành bảng thông tin*:\n"
                "   • *Sodium (Na)*: Điện tích hạt nhân: *+11*; Số p: *11*; Số e: *11*; Số n: *12*; Số lớp e: *3* (2, 8, 1); Số e lớp ngoài cùng: *1*; Khối lượng nguyên tử: *23 amu*.\n"
                "   • *Oxygen (O)*: Điện tích hạt nhân: *+8*; Số p: *8*; Số e: *8*; Số n: *8*; Số lớp e: *2* (2, 6); Số e lớp ngoài cùng: *6*; Khối lượng nguyên tử: *16 amu*."
            )
        },
        {
            "id": 4,
            "title": "Câu 4",
            "content": (
                "a. Hãy cho biết tên gọi của các nguyên tố sau: N, Si, Al, K, Mg, H, O, P.\n"
                "b. Viết kí hiệu hóa học của các nguyên tố sau: calcium, helium, carbon, sulfur, sodium, oxygen, potassium."
            ),
            "solution": (
                "a. *Tên gọi các nguyên tố*:\n"
                "   • N: *Nitrogen* \t\t• Si: *Silicon* \t\t• Al: *Aluminium* \t• K: *Potassium*\n"
                "   • Mg: *Magnesium* \t• H: *Hydrogen* \t• O: *Oxygen* \t\t• P: *Phosphorus*\n"
                "b. *Kí hiệu hóa học*:\n"
                "   • calcium: *Ca* \t\t• helium: *He* \t\t• carbon: *C* \t\t• sulfur: *S*\n"
                "   • sodium: *Na* \t\t• oxygen: *O* \t\t• potassium: *K*"
            )
        },
        {
            "id": 5,
            "title": "Câu 5",
            "content": "Biết rằng 4 nguyên tử Mg nặng bằng 3 nguyên tử nguyên tố X. Hãy cho biết tên, kí hiệu hóa học của X.",
            "solution": (
                "• Khối lượng nguyên tử của Magnesium là: m_Mg = 24 amu.\n"
                "• Theo đề bài: 4 × m_Mg = 3 × m_X  ==>  4 × 24 = 3 × m_X  ==>  96 = 3 × m_X.\n"
                "  ==> m_X = 96 / 3 = *32 (amu)*.\n"
                "• Nguyên tử có khối lượng nguyên tử bằng 32 amu là *Sulfur* (lưu huỳnh).\n"
                "• Tên nguyên tố: *Sulfur*; Kí hiệu hóa học: *S*."
            )
        },
        {
            "id": 6,
            "title": "Câu 6",
            "content": (
                "Một nguyên tử sodium có 11 proton và 12 neutron.\n"
                "a. Tính tổng số hạt của nguyên tử sodium.\n"
                "b. Tính khối lượng của nguyên tử sodium."
            ),
            "solution": (
                "a. Nguyên tử trung hòa về điện nên: số electron = số proton = 11 hạt.\n"
                "   Tổng số hạt cấu tạo nên nguyên tử sodium là: N = p + n + e = 11 + 12 + 11 = *34 (hạt)*.\n"
                "b. Khối lượng của nguyên tử sodium xấp xỉ bằng khối lượng của hạt nhân (bỏ qua khối lượng electron):\n"
                "   m_Na ≈ (p + n) = 11 + 12 = *23 (amu)* (hoặc 23 × 1,6605 × 10^-24 g ≈ *3,82 × 10^-23 g*)."
            )
        }
    ]

    for item in chem_questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        r_num = p_q.add_run(f"{item['title']}. ")
        r_num.font.name = FONT_NAME
        r_num.font.size = Pt(10.5)
        r_num.font.bold = True
        r_cnt = p_q.add_run(item["content"])
        r_cnt.font.name = FONT_NAME
        r_cnt.font.size = Pt(10.5)

        if is_teacher:
            add_solution_paragraph(doc, item["solution"])
        elif not is_screen:
            add_dotted_lines(doc, num_lines=4)

    # --------------------------------------------------------------------------
    # C. PHÂN MÔN KHTN 3: SINH HỌC (TRAO ĐỔI CHẤT & QUANG HỢP)
    # --------------------------------------------------------------------------
    add_part_heading(doc, "C. PHÂN MÔN KHTN 3 (SINH HỌC)")
    bio_questions = [
        {
            "id": 1,
            "title": "Câu 1",
            "content": "Khái niệm trao đổi chất và chuyển hoá năng lượng ở sinh vật. Cho ví dụ minh hoạ mỗi quá trình.",
            "solution": (
                "• *Trao đổi chất*: Là quá trình cơ thể sinh vật lấy các chất từ môi trường ngoài, biến đổi chúng thành các chất cần thiết cho tế bào và thải các chất bã, chất cặn ra môi trường. Ví dụ: Con người hít vào khí O2, uống nước, ăn thức ăn và bài tiết nước tiểu, phân, mồ hôi, thở ra CO2.\n"
                "• *Chuyển hoá năng lượng*: Là sự biến đổi năng lượng từ dạng này sang dạng khác trong tế bào và cơ thể. Ví dụ: Trong quang hợp, thực vật chuyển hóa quang năng (ánh sáng mặt trời) thành hóa năng trong các hợp chất hữu cơ (glucose, tinh bột); trong hô hấp, hóa năng biến đổi thành nhiệt năng và cơ năng."
            )
        },
        {
            "id": 2,
            "title": "Câu 2",
            "content": "Vai trò của trao đổi chất và chuyển hoá năng lượng ở sinh vật.",
            "solution": (
                "Vai trò cốt lõi của trao đổi chất và chuyển hóa năng lượng:\n"
                "1. *Cung cấp nguyên liệu*: Tạo ra các chất hữu cơ cấu tạo nên tế bào, mô, cơ quan và cơ thể sinh vật.\n"
                "2. *Cung cấp năng lượng*: Tạo ra năng lượng (ATP) để duy trì mọi hoạt động sống (sinh trưởng, phát triển, sinh sản, vận động, cảm ứng).\n"
                "3. *Đào thải chất cặn bã*: Loại bỏ chất độc hại sinh ra trong chuyển hóa ra môi trường ngoài, giúp duy trì cân bằng nội môi."
            )
        },
        {
            "id": 3,
            "title": "Câu 3",
            "content": "Quang hợp là gì? Xác định các chất tham gia/nguyên liệu, chất tạo thành/sản phẩm của quá trình quang hợp.",
            "solution": (
                "• *Khái niệm quang hợp*: Quang hợp là quá trình lá cây sử dụng năng lượng ánh sáng mặt trời đã được diệp lục hấp thu để tổng hợp chất hữu cơ (glucose, tinh bột) từ nước và khí carbon dioxide, đồng thời giải phóng khí oxygen.\n"
                "• *Chất tham gia (nguyên liệu)*: Nước (H2O) và Khí carbon dioxide (CO2).\n"
                "• *Chất tạo thành (sản phẩm)*: Chất hữu cơ (Glucose - C6H12O6) và Khí oxygen (O2).\n"
                "• *Phương trình chữ*: Nước + Carbon dioxide --(Ánh sáng / Diệp lục)--> Glucose + Oxygen."
            )
        },
        {
            "id": 4,
            "title": "Câu 4",
            "content": "Quá trình quang hợp diễn ra chủ yếu ở cơ quan nào? Vì sao?",
            "solution": (
                "• Cơ quan diễn ra quang hợp chủ yếu ở thực vật là *lá cây*.\n"
                "• *Giải thích lý do*:\n"
                "  - Phiến lá có dạng bản dẹt, rộng giúp hấp thu được nhiều ánh sáng mặt trời.\n"
                "  - Lớp thịt lá chứa mật độ rất lớn các bào quan *lục lạp* có chứa sắc tố *diệp lục* (chất có khả năng hấp thụ năng lượng quang năng).\n"
                "  - Hệ gân lá phân bố dày đặc giúp vận chuyển nước từ rễ lên và dẫn truyền sản phẩm quang hợp đến các bộ phận khác.\n"
                "  - Lớp biểu bì lá có nhiều *khí khổng* giúp trao đổi khí CO2, O2 và thoát hơi nước dễ dàng."
            )
        },
        {
            "id": 5,
            "title": "Câu 5",
            "content": "Trình bày một số yếu tố chủ yếu bên ngoài ảnh hưởng đến quang hợp.",
            "solution": (
                "Các yếu tố môi trường chủ yếu ảnh hưởng đến quang hợp:\n"
                "1. *Ánh sáng*: Cường độ ánh sáng và thành phần quang phổ ảnh hưởng trực tiếp đến hiệu suất quang hợp; cường độ ánh sáng tăng thì quang hợp tăng đến điểm bão hòa.\n"
                "2. *Nồng độ khí CO2*: Nồng độ CO2 ngoài môi trường từ 0,03% đến 0,2% thuận lợi cho quang hợp; nồng độ quá thấp hoặc quá cao đều làm giảm hoặc ngừng quang hợp.\n"
                "3. *Nước*: Vừa là nguyên liệu trực tiếp, vừa duy trì độ trương của tế bào khí khổng để trao đổi khí.\n"
                "4. *Nhiệt độ*: Nhiệt độ tối ưu cho hầu hết thực vật là từ 25 °C đến 35 °C. Nhiệt độ quá thấp hoặc quá cao (trên 40 °C) làm biến tính enzyme quang hợp.\n"
                "5. *Các nguyên tố khoáng*: Tham gia cấu trúc phân tử diệp lục (như N, Mg) và các enzyme xúc tác."
            )
        },
        {
            "id": 6,
            "title": "Câu 6",
            "content": (
                "Vận dụng kiến thức về quang hợp để giải thích một số hiện tượng liên quan. Ví dụ minh hoạ:\n"
                "a) Vì sao trong bể cá cảnh, người ta thường hay thả thêm các cây thuỷ sinh?\n"
                "b) Một số cây lá không có màu xanh lục (lẻ bạn, tía tô, …) thì có quang hợp được không?\n"
                "c) Cây xương rồng không có lá thì quang hợp bằng cơ quan nào?\n"
                "d) Vì sao phải trồng cây đúng mùa vụ?"
            ),
            "solution": (
                "a) *Thả cây thủy sinh vào bể cá*: Khi có ánh sáng, cây thủy sinh quang hợp hấp thu khí CO2 do cá thải ra và giải phóng khí O2 hòa tan vào nước, cung cấp dưỡng khí cho cá hô hấp; đồng thời hút bớt các chất hữu cơ cặn bã làm sạch nước.\n"
                "b) *Lá không có màu xanh lục (lẻ bạn, tía tô, huyết dụ...)*: Vẫn *quang hợp bình thường*, vì trong lá vẫn có sắc tố diệp lục, nhưng do hàm lượng các sắc tố khác (anthocyanin màu đỏ, tím; carotenoid màu vàng, cam) chiếm tỉ lệ quá cao nên đã che lấp màu xanh của diệp lục.\n"
                "c) *Cây xương rồng*: Lá biến đổi thành gai để hạn chế tối đa sự thoát hơi nước ở sa mạc; cơ quan thực hiện quang hợp thay cho lá là *thân cây* (thân cây mọng nước, có màu xanh chứa nhiều lục lạp).\n"
                "d) *Trồng cây đúng mùa vụ*: Giúp cây trồng đón nhận được các điều kiện ngoại cảnh (nhiệt độ, ánh sáng, lượng mưa) thích hợp nhất cho quá trình quang hợp, sinh trưởng và phát triển, hạn chế sâu bệnh, từ đó đạt năng suất và chất lượng thu hoạch cao nhất."
            )
        },
        {
            "id": 7,
            "title": "Câu 7",
            "content": "Em hãy đề xuất những biện pháp để bảo vệ cây xanh.",
            "solution": (
                "Các biện pháp thiết thực để bảo vệ cây xanh:\n"
                "1. Tích cực tham gia phong trào trồng cây gây rừng, phủ xanh đất trống đồi trọc.\n"
                "2. Chăm sóc, tưới nước, bón phân và bắt sâu bọ định kì cho cây xanh ở trường học, gia đình và khu dân cư.\n"
                "3. Không bẻ cành, hái hoa, khắc chữ lên thân cây hoặc chặt phá cây bừa bãi.\n"
                "4. Tuyên truyền cho bạn bè, người thân về vai trò to lớn của cây xanh đối với việc điều hòa khí hậu, lọc bụi và bảo vệ sự sống.\n"
                "5. Lên án và kịp thời thông báo cho cơ quan chức năng những hành vi chặt phá, đốt phá rừng trái phép."
            )
        }
    ]

    for item in bio_questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        r_num = p_q.add_run(f"{item['title']}. ")
        r_num.font.name = FONT_NAME
        r_num.font.size = Pt(10.5)
        r_num.font.bold = True
        r_cnt = p_q.add_run(item["content"])
        r_cnt.font.name = FONT_NAME
        r_cnt.font.size = Pt(10.5)

        if is_teacher:
            add_solution_paragraph(doc, item["solution"])
        elif not is_screen:
            add_dotted_lines(doc, num_lines=4)

    # Kết thúc
    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_end.paragraph_format.space_before = Pt(10)
    r_end = p_end.add_run("“HẾT -- Chúc các em ôn tập chu đáo và đạt điểm số cao nhất!”")
    r_end.font.name = FONT_NAME
    r_end.font.size = Pt(10)
    r_end.font.italic = True
    r_end.font.color.rgb = COLOR_PRIMARY

    # Đặt tên file xuất
    if mode == "student":
        filename = "De_Cuong_GK1_KHTN7_NguyenChiDieu_HocSinh.docx"
    elif mode == "teacher":
        filename = "De_Cuong_GK1_KHTN7_NguyenChiDieu_GiaoVien.docx"
    else:
        filename = "De_Cuong_GK1_KHTN7_NguyenChiDieu_DayTrenManHinh.docx"

    out_path = os.path.join(SAN_PHAM_DIR, filename)
    doc.save(out_path)
    print(f"[OK] Đã tạo thành công file Word: {out_path} ({os.path.getsize(out_path):,} bytes)")
    return out_path

# ==============================================================================
# HÀM TẠO FILE SLIDE PDF DẠY TRÊN MÀN HÌNH 16:9 CHO TIVI 65 INCH (KHÔNG CÓ ĐÁP ÁN)
# ==============================================================================
def create_screen_teaching_pdf():
    tex_path = os.path.join(SCRATCH_DIR, "De_Cuong_GK1_KHTN7_NguyenChiDieu_DayTrenManHinh.tex")
    img_cau4_unix = IMG_CAU4.replace("\\", "/")
    img_cau3_unix = IMG_CAU3.replace("\\", "/")

    tex_lines = [r"""\documentclass[12pt,landscape]{article}
\usepackage[paperwidth=29.7cm,paperheight=16.706cm,top=1.0cm,bottom=1.0cm,left=1.2cm,right=1.2cm,headheight=0pt,headsep=0pt,footskip=0pt]{geometry}
\usepackage[utf8]{vietnam}
\usepackage{times}
\usepackage{amsmath,amssymb}
\usepackage{graphicx}
\usepackage{xcolor}
\usepackage{tabularx}
\usepackage{enumitem}

\definecolor{brandRed}{RGB}{211, 47, 47}
\definecolor{brandBlue}{RGB}{21, 101, 192}
\definecolor{brandGreen}{RGB}{46, 125, 50}
\definecolor{offBlack}{RGB}{26, 26, 26}
\definecolor{grayText}{RGB}{117, 117, 117}

\pagestyle{empty}
\setlength{\parindent}{0pt}
\setlength{\parskip}{0pt}

\begin{document}
"""]

    # SLIDE 1: TIÊU ĐỀ
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 1: TIÊU ĐỀ
% ==============================================================================
\begin{center}
\vspace*{0.8cm}
{\fontsize{24}{30}\selectfont \textbf{\color{offBlack}TRƯỜNG THCS NGUYỄN CHÍ DIỂU -- NĂM HỌC 2026 -- 2027}}\\[12pt]
{\fontsize{34}{44}\selectfont \textbf{\color{brandRed}ĐỀ CƯƠNG ÔN TẬP KIỂM TRA GIỮA HỌC KỲ I}}\\[10pt]
{\fontsize{28}{36}\selectfont \textbf{\color{brandBlue}MÔN: KHOA HỌC TỰ NHIÊN 7}}\\[8pt]
{\fontsize{22}{28}\selectfont \textit{\color{grayText}(Trọn bộ 3 phân môn: KHTN 1 - Vật lí \quad|\quad KHTN 2 - Hóa học \quad|\quad KHTN 3 - Sinh học)}}\\[1.2cm]

{\fontsize{20}{26}\selectfont \textbf{Giáo viên hướng dẫn:} Thầy \textbf{\color{brandRed}TRẦN VĂN THẠNH} -- 0777.470.803}\\[6pt]
{\fontsize{18}{24}\selectfont \textit{Cơ sở 1: 6/15 Nguyễn Hoàng, Kim Long \quad|\quad Cơ sở 2: 24 Đặng Thái Thân, TP Huế}}\\[1.0cm]

{\fontsize{18}{24}\selectfont \textbf{\color{brandGreen}SLIDE TRÌNH CHIẾU TƯƠNG TÁC TIVI 65 INCH (KHÔNG HIỂN THỊ ĐÁP ÁN)}}
\end{center}
\newpage
""")

    # SLIDE 2: KHTN 1 - VẬT LÍ (CÂU 1 & CÂU 2)
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 2: VẬT LÍ (CÂU 1 & 2)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}A. PHÂN MÔN KHTN 1 (VẬT LÍ) -- BÀI TẬP TỐC ĐỘ}}
\vspace{10pt}

\fontsize{21}{28}\selectfont
\textbf{Câu 1.}
\begin{enumerate}[label=\alph*),itemsep=4pt]
  \item Một người đi quãng đường $1,5\text{ km}$ với tốc độ $20\text{ m/s}$. Thời gian để người đó đi hết quãng đường là bao nhiêu?
  \item Khi nói tốc độ của vật A là $36\text{ km/h}$, con số đó cho biết điều gì?
  \item Đổi các đơn vị sau: $54\text{ km/h} = \dots\dots\dots\dots\text{ m/s}$; \quad $10\text{ m/s} = \dots\dots\dots\dots\text{ km/h}$.
\end{enumerate}

\vspace{10pt}
\textbf{Câu 2.} Hai người đi xe đạp. Người thứ nhất đi quãng đường $300\text{ m}$ hết $1\text{ phút}$. Người thứ hai đi quãng đường $7,5\text{ km}$ hết $0,5\text{ h}$. Người nào đi nhanh hơn?

\vspace*{1.5cm}
\begin{center}
{\fontsize{18}{24}\selectfont \textit{\color{grayText}(Khoảng trống chữa bài: $v = \frac{s}{t}$, đổi đơn vị tốc độ)}}
\end{center}
\newpage
""")

    # SLIDE 3: KHTN 1 - CÂU 3 (VẼ ĐỒ THỊ TỪ BẢNG SỐ LIỆU)
    slide3_tex = r"""
% ==============================================================================
% SLIDE 3: VẬT LÍ (CÂU 3)
% ==============================================================================
{\fontsize{23}{28}\selectfont \textbf{\color{brandBlue}A. PHÂN MÔN KHTN 1 (VẬT LÍ) -- CÂU 3 (ĐỒ THỊ S - T)}}
\vspace{4pt}

\fontsize{20}{26}\selectfont
\textbf{Câu 3.} Dựa vào số liệu trong bảng, hãy vẽ đồ thị quãng đường -- thời gian của vật:
\vspace{4pt}

\begin{center}
\begin{tabular}{|l|c|c|c|c|}
\hline
\textbf{Thời gian $t$ (h)} & $1$ & $2$ & $3$ & $4$ \\ \hline
\textbf{Quãng đường $s$ (km)} & $60$ & $120$ & $180$ & $240$ \\ \hline
\end{tabular}
\end{center}

\vspace{2pt}
\begin{center}
\includegraphics[height=6.2cm,keepaspectratio]{__IMG_CAU3__}
\end{center}
\newpage
""".replace("__IMG_CAU3__", img_cau3_unix)
    tex_lines.append(slide3_tex)

    # SLIDE 4: KHTN 1 - CÂU 4 (ĐỌC ĐỒ THỊ 4 GIAI ĐOẠN)
    slide4_tex = r"""
% ==============================================================================
% SLIDE 4: VẬT LÍ (CÂU 4)
% ==============================================================================
{\fontsize{23}{28}\selectfont \textbf{\color{brandBlue}A. PHÂN MÔN KHTN 1 (VẬT LÍ) -- CÂU 4 (ĐỌC ĐỒ THỊ)}}
\vspace{4pt}

\begin{minipage}[c]{0.54\textwidth}
\fontsize{19}{25}\selectfont
\textbf{Câu 4.} Quan sát đồ thị hình bên và cho biết:
\begin{enumerate}[label=\alph*),itemsep=4pt]
  \item Quãng đường vật đi được trong $2\text{ s}$ đầu tiên?\\
        Thời gian vật đi được quãng đường $8\text{ m}$?
  \item Tính tốc độ của vật đi được trong mỗi giai đoạn (A, B, C, D).
\end{enumerate}

\vspace*{0.8cm}
\begin{center}
{\fontsize{17}{22}\selectfont \textit{\color{grayText}(Khoảng trống chữa bài trên Tivi 65 inch)}}
\end{center}
\end{minipage}
\hfill
\begin{minipage}[c]{0.44\textwidth}
\centering
\includegraphics[height=10.0cm,keepaspectratio]{__IMG_CAU4__}
\end{minipage}
\newpage
""".replace("__IMG_CAU4__", img_cau4_unix)
    tex_lines.append(slide4_tex)

    # SLIDE 5: KHTN 1 - CÂU 5 (TỐC ĐỘ TRUNG BÌNH 2 GIAI ĐOẠN)
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 5: VẬT LÍ (CÂU 5)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}A. PHÂN MÔN KHTN 1 (VẬT LÍ) -- CÂU 5}}
\vspace{12pt}

\fontsize{22}{30}\selectfont
\textbf{Câu 5.} Một người đi xe máy đều trên quãng đường đầu dài $18\text{ km}$ hết $30\text{ phút}$. Sau đó, ở quãng đường sau dài $1\,200\text{ m}$, người đó đi hết $5\text{ phút } 20\text{ giây}$.\\
Tính tốc độ trung bình của người đó trên cả hai quãng đường (theo đơn vị $\text{m/s}$ và $\text{km/h}$).

\vspace*{2.5cm}
\begin{center}
{\fontsize{19}{25}\selectfont \textit{\color{grayText}(Khoảng trống tính: $s = s_1 + s_2$, $t = t_1 + t_2 \implies v_{tb} = \frac{s}{t}$)}}
\end{center}
\newpage
""")

    # SLIDE 6: KHTN 2 - HÓA HỌC (CÂU 1 & CÂU 2)
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 6: HÓA HỌC (CÂU 1 & 2)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}B. PHÂN MÔN KHTN 2 (HÓA HỌC) -- NGUYÊN TỬ \& NGUYÊN TỐ}}
\vspace{10pt}

\fontsize{21}{28}\selectfont
\textbf{Câu 1.}
\begin{enumerate}[label=\alph*),itemsep=2pt]
  \item Nguyên tử là gì?
  \item Khối lượng nguyên tử là gì? $1\text{ amu}$ có khối lượng bằng bao nhiêu gam?
  \item Nêu tóm tắt mô hình nguyên tử theo Rutherford -- Bohr.
  \item Trình bày cấu tạo nguyên tử ($p, n, e$).
\end{enumerate}

\vspace{10pt}
\textbf{Câu 2.}
\begin{enumerate}[label=\alph*),itemsep=2pt]
  \item Nguyên tố hóa học là gì? Nêu cách biểu diễn kí hiệu hóa học.
  \item Đọc tên và kí hiệu hóa học của 20 nguyên tố đầu tiên trong Bảng tuần hoàn.
\end{enumerate}
\newpage
""")

    # SLIDE 7: KHTN 2 - HÓA HỌC (CÂU 3: MÔ HÌNH CALCIUM & ĐIỀN BẢNG)
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 7: HÓA HỌC (CÂU 3)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}B. PHÂN MÔN KHTN 2 (HÓA HỌC) -- CÂU 3 (MÔ HÌNH NGUYÊN TỬ)}}
\vspace{8pt}

\fontsize{20}{26}\selectfont
\textbf{a)} Vẽ mô hình cấu tạo nguyên tử Calcium ($+20$). Dựa vào mô hình hãy cho biết số proton, số electron, số lớp electron và số electron lớp ngoài cùng của Calcium.\\[8pt]
\textbf{b)} Điền thông tin còn thiếu vào bảng sau:

\vspace{6pt}
\begin{center}
\begin{tabular}{|l|c|c|c|c|c|c|c|}
\hline
\textbf{Nguyên tử} & \textbf{Điện tích HN} & \textbf{Số p} & \textbf{Số e} & \textbf{Số n} & \textbf{Số lớp e} & \textbf{Số e ngoài} & \textbf{NTK (amu)} \\ \hline
Sodium (Na) & $\dots$ & $11$ & $\dots$ & $12$ & $\dots$ & $\dots$ & $\dots$ \\ \hline
Oxygen (O) & $+8$ & $\dots$ & $\dots$ & $8$ & $\dots$ & $\dots$ & $\dots$ \\ \hline
\end{tabular}
\end{center}

\vspace*{1.2cm}
\begin{center}
{\fontsize{18}{24}\selectfont \textit{\color{grayText}(Học sinh lên bảng hoàn thành bảng và vẽ các lớp electron)}}
\end{center}
\newpage
""")

    # SLIDE 8: KHTN 2 - HÓA HỌC (CÂU 4, 5, 6)
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 8: HÓA HỌC (CÂU 4, 5, 6)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}B. PHÂN MÔN KHTN 2 (HÓA HỌC) -- BÀI TẬP VẬN DỤNG}}
\vspace{8pt}

\fontsize{20}{26}\selectfont
\textbf{Câu 4.}
\begin{enumerate}[label=\alph*),itemsep=2pt]
  \item Cho biết tên gọi các nguyên tố sau: $\text{N, Si, Al, K, Mg, H, O, P}$.
  \item Viết kí hiệu hóa học của: calcium, helium, carbon, sulfur, sodium, oxygen, potassium.
\end{enumerate}

\vspace{8pt}
\textbf{Câu 5.} Biết rằng 4 nguyên tử $\text{Mg}$ nặng bằng 3 nguyên tử nguyên tố $\text{X}$. Hãy cho biết tên và kí hiệu hóa học của nguyên tố $\text{X}$.\\[8pt]
\textbf{Câu 6.} Một nguyên tử sodium có $11\text{ proton}$ và $12\text{ neutron}$.
\begin{enumerate}[label=\alph*),itemsep=2pt]
  \item Tính tổng số hạt cấu tạo nên nguyên tử sodium.
  \item Tính khối lượng nguyên tử sodium.
\end{enumerate}
\newpage
""")

    # SLIDE 9: KHTN 3 - SINH HỌC (CÂU 1, 2, 3)
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 9: SINH HỌC (CÂU 1, 2, 3)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}C. PHÂN MÔN KHTN 3 (SINH HỌC) -- TRAO ĐỔI CHẤT \& QUANG HỢP}}
\vspace{10pt}

\fontsize{21}{28}\selectfont
\textbf{Câu 1.} Khái niệm trao đổi chất và chuyển hoá năng lượng ở sinh vật? Cho ví dụ minh hoạ cho mỗi quá trình.\\[12pt]
\textbf{Câu 2.} Trình bày vai trò của trao đổi chất và chuyển hoá năng lượng đối với sự sống của sinh vật.\\[12pt]
\textbf{Câu 3.} Quang hợp là gì? Xác định các chất tham gia (nguyên liệu) và chất tạo thành (sản phẩm) của quá trình quang hợp. Viết phương trình quang hợp.

\vspace*{1.5cm}
\begin{center}
{\fontsize{18}{24}\selectfont \textit{\color{grayText}(Khoảng trống trình bày phương trình và phân tích vai trò của quang hợp)}}
\end{center}
\newpage
""")

    # SLIDE 10: KHTN 3 - SINH HỌC (CÂU 4 & CÂU 5)
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 10: SINH HỌC (CÂU 4 & 5)
% ==============================================================================
{\fontsize{24}{30}\selectfont \textbf{\color{brandBlue}C. PHÂN MÔN KHTN 3 (SINH HỌC) -- ĐẶC ĐIỂM QUANG HỢP}}
\vspace{12pt}

\fontsize{21}{29}\selectfont
\textbf{Câu 4.} Quá trình quang hợp ở thực vật diễn ra chủ yếu ở cơ quan nào? Vì sao cơ quan đó lại thích nghi tốt nhất cho chức năng quang hợp?\\[16pt]
\textbf{Câu 5.} Trình bày các yếu tố chủ yếu từ môi trường bên ngoài ảnh hưởng đến quá trình quang hợp ở thực vật (ánh sáng, nồng độ khí $\text{CO}_2$, nước, nhiệt độ, chất khoáng).

\vspace*{2.0cm}
\begin{center}
{\fontsize{18}{24}\selectfont \textit{\color{grayText}(Khoảng trống thảo luận và liên hệ kiến thức thực nghiệm)}}
\end{center}
\newpage
""")

    # SLIDE 11: KHTN 3 - SINH HỌC (CÂU 6 & CÂU 7 - VẬN DỤNG THỰC TIỄN)
    tex_lines.append(r"""
% ==============================================================================
% SLIDE 11: SINH HỌC (CÂU 6 & 7)
% ==============================================================================
{\fontsize{23}{28}\selectfont \textbf{\color{brandBlue}C. PHÂN MÔN KHTN 3 (SINH HỌC) -- VẬN DỤNG THỰC TIỄN}}
\vspace{6pt}

\fontsize{19}{25}\selectfont
\textbf{Câu 6.} Vận dụng kiến thức về quang hợp để giải thích các hiện tượng thực tế:
\begin{enumerate}[label=\alph*),itemsep=3pt]
  \item Vì sao trong bể nuôi cá cảnh, người ta thường thả thêm các cây thuỷ sinh?
  \item Một số cây có lá không có màu xanh lục (lẻ bạn, tía tô,...) thì có quang hợp được không?
  \item Cây xương rồng không có lá thì quang hợp bằng cơ quan nào?
  \item Vì sao người nông dân phải gieo trồng cây đúng thời vụ?
\end{enumerate}

\vspace{6pt}
\textbf{Câu 7.} Em hãy đề xuất các biện pháp cụ thể, thiết thực để bảo vệ và phát triển cây xanh trong trường học và địa phương.
\newpage
""")

    tex_lines.append(r"\end{document}")

    with open(tex_path, "w", encoding="utf-8") as f:
        f.write("\n".join(tex_lines))
    print(f"[OK] Đã ghi file LaTeX slide: {tex_path}")

    cmd = ["pdflatex", "-interaction=nonstopmode", "-output-directory", SCRATCH_DIR, tex_path]
    res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    pdf_temp = os.path.join(SCRATCH_DIR, "De_Cuong_GK1_KHTN7_NguyenChiDieu_DayTrenManHinh.pdf")
    pdf_out = os.path.join(SAN_PHAM_DIR, "De_Cuong_GK1_KHTN7_NguyenChiDieu_DayTrenManHinh.pdf")

    if os.path.exists(pdf_temp):
        shutil.copy2(pdf_temp, pdf_out)
        print(f"[OK] Đã biên dịch và copy PDF thành công: {pdf_out} ({os.path.getsize(pdf_out):,} bytes)")
    else:
        print(f"[ERR] Không tạo được PDF. Mã lỗi: {res.returncode}")
    return pdf_out

# ==============================================================================
# HÀM ĐỒNG BỘ ONEDRIVE
# ==============================================================================
def sync_to_onedrive():
    files_to_sync = [
        "De_Cuong_GK1_KHTN7_NguyenChiDieu_HocSinh.docx",
        "De_Cuong_GK1_KHTN7_NguyenChiDieu_GiaoVien.docx",
        "De_Cuong_GK1_KHTN7_NguyenChiDieu_DayTrenManHinh.docx",
        "De_Cuong_GK1_KHTN7_NguyenChiDieu_DayTrenManHinh.pdf",
    ]
    print("\n[+] Đồng bộ các file sang các thư mục OneDrive:")
    for dest_dir in [ONEDRIVE_DIR1, ONEDRIVE_DIR2]:
        if os.path.exists(dest_dir):
            for fname in files_to_sync:
                src_f = os.path.join(SAN_PHAM_DIR, fname)
                dst_f = os.path.join(dest_dir, fname)
                if os.path.exists(src_f):
                    try:
                        shutil.copy2(src_f, dst_f)
                        print(f"    -> Đã copy {fname} vào: {dest_dir}")
                    except Exception as e:
                        print(f"    [!] Cảnh báo file đang mở/khóa tại {dest_dir}: {e}")
        else:
            print(f"    [-] Không tìm thấy thư mục: {dest_dir}")

# ==============================================================================
# MAIN
# ==============================================================================
if __name__ == "__main__":
    print("=" * 75)
    print("BẮT ĐẦU TẠO 3 FILE CHO ĐỀ CƯƠNG ÔN TẬP GIỮA KỲ I KHTN 7 (THCS NGUYỄN CHÍ DIỂU)")
    print("TRỌN BỘ 3 PHÂN MÔN: KHTN 1 (VẬT LÍ) - KHTN 2 (HÓA HỌC) - KHTN 3 (SINH HỌC)")
    print("1. FILE HỌC SINH (.docx)")
    print("2. FILE GIÁO VIÊN (.docx có lời giải chi tiết)")
    print("3. FILE DẠY TRÊN MÀN HÌNH (Slide PDF 16:9 widescreen cho Tivi 65 inch)")
    print("=" * 75)

    create_word_document(mode="student")
    create_word_document(mode="teacher")
    create_word_document(mode="screen")
    create_screen_teaching_pdf()
    sync_to_onedrive()

    print("\n" + "=" * 75)
    print("HOÀN THÀNH TẤT CẢ 3 FILE THEO ĐÚNG CHUẨN FORM ĐỀ KIỂM TRA!")
    print("=" * 75)
