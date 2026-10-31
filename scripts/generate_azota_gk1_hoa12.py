# -*- coding: utf-8 -*-
"""
Chuyển đổi Đề Kiểm Tra Giữa Học Kì I - Hóa Học 12 sang định dạng chuẩn AZOTA:
- Không có bảng tiêu đề rườm rà hay nhãn [Mã NH: ...]
- Câu hỏi bắt đầu bằng "Câu X." (in đậm, màu xanh dương #0000FF)
- Phần I:
  + Phương án A., B., C., D. (nhãn in đậm, màu xanh dương #0000FF)
  + Đáp án đúng: gạch chân nhãn (ví dụ: A. được gạch chân)
  + Lời giải bắt đầu bằng dòng "Lời giải:" rồi xuống dòng ghi nội dung giải thích.
- Phần II:
  + Các ý a), b), c), d) ở đầu dòng.
  + Ý Đúng: gạch chân nhãn (ví dụ: a) được gạch chân); Ý Sai: không gạch chân nhãn.
  + Không ghi "→ [Đúng]" hay "[HH1.1 - Biết]" trong câu.
  + Lời giải bắt đầu bằng dòng "Lời giải" rồi xuống dòng ghi giải thích từng ý.
- Phần III:
  + Câu hỏi bắt đầu bằng "Câu X."
  + Ngay dưới câu hỏi là dòng "Lời giải", dòng tiếp theo "Đáp án: <số>", dòng tiếp theo "Giải thích:" và nội dung giải.
"""

import os
import re
import copy
import json
import docx
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from generate_5_de_gk1_hoa12 import (
    BANK_PATH,
    OUT_DIR,
    BANK_DIR,
    TESTS_SPEC,
    STEM_PATCHES,
    P2_ITEM_PATCHES,
    ANSWER_PATCHES,
    SOLUTION_PATCHES,
    add_chem_runs,
    parse_abcd_from_opt_lines,
    add_data_table_to_doc,
)
from extract_q_assets import extract_chapter_assets

COLOR_AZOTA_BLUE = RGBColor(0, 0, 255)  # #0000FF chuẩn Azota trong ảnh mẫu

# Bổ sung bản vá lỗi mất công thức Equation cho Đề 02 và các câu liên quan
EXTRA_OPT_PATCHES = {
    ("C2", 220): ["A. Rau răm.\tB. Hạt gạo.\tC. Hạt óc chó.\tD. Quả nho chín."],
    ("C2", 39): [
        "A. dung dịch [Ag(NH3)2]OH, to.\tB. Cu(OH)2.\tC. H2 (xúc tác Ni, to).\tD. Nước bromine."
    ]
}

EXTRA_ANSWER_PATCHES = {
    ("C1", 164): "65,6",
}

EXTRA_SOLUTION_PATCHES = {
    ("C1", 164): (
        "n(isoamyl alcohol) = (30,0 × 0,810) / 88 = 24,3 / 88 mol ≈ 0,2761 mol.\n"
        "n(CH3COOH) = (20,0 × 1,049) / 60 = 20,98 / 60 mol ≈ 0,3497 mol.\n"
        "Vì n(alcohol) < n(acid) nên hiệu suất tính theo isoamyl alcohol:\n"
        "n(ester) = (26,9 × 0,876) / 130 = 23,5644 / 130 mol ≈ 0,1813 mol.\n"
        "Hiệu suất phản ứng: h% = [(23,5644 / 130) / (24,3 / 88)] × 100% ≈ 65,6%."
    ),
    ("C1", 157): (
        "Đánh số nguyên tử carbon từ đầu nhóm methyl (–CH3) là C1 thì liên kết đôi C=C đầu tiên xuất hiện ở vị trí C6 nên chất X thuộc loại acid béo omega-6 (n = 6)."
    ),
    ("C1", 178): (
        "n(KOH) = 1,6 × 2 = 3,2 mol; n(glycerol) = 96,6 / 92 = 1,05 mol.\n"
        "n(KOH phản ứng với triglyceride) = 3 × n(glycerol) = 3,15 mol.\n"
        "n(KOH phản ứng với acid béo tự do) = 3,2 - 3,15 = 0,05 mol → n(H2O) = 0,05 mol.\n"
        "Bảo toàn khối lượng: m(X) + m(KOH) = m(muối) + m(glycerol) + m(H2O)\n"
        "⇒ 948,7 + 3,2 × 56 = m(muối) + 96,6 + 0,05 × 18 ⇒ m(muối) = 1030,4 gam.\n"
        "Khối lượng xà phòng thu được: m(xà phòng) = 1030,4 / 74,66% ≈ 1380 gam."
    ),
    ("C2", 110): (
        "Các chất tham gia phản ứng với thuốc thử Tollens gồm 5 chất: methanal, formic acid, methyl formate, glucose, fructose."
    ),
    ("C1", 185): (
        "Sơ đồ: (C6H10O5)n → nC6H12O6 → 2nC2H5OH.\n"
        "Thể tích C2H5OH nguyên chất trong 400 L xăng E10: V(C2H5OH) = 400 × 10% = 40 L.\n"
        "Khối lượng C2H5OH: m(C2H5OH) = 40 × 0,79 = 31,6 kg.\n"
        "Khối lượng sắn khô cần dùng: m = (31,6 × 162) / (2 × 46 × 0,80 × 0,70 × 0,72) ≈ 138 kg."
    ),
    ("C3", 153): (
        "Có 3 đồng phân amine bậc ba ứng với công thức phân tử C5H13N:\n"
        "(1) CH3–N(CH3)–CH2–CH2–CH3\n"
        "(2) CH3–N(CH3)–CH(CH3)2\n"
        "(3) CH3–CH2–N(CH3)–CH2–CH3"
    ),
    ("C2", 93): (
        "a) Đúng: Phương trình chuẩn độ: CH3CH(OH)COOH + NaOH → CH3CH(OH)COONa + H2O. Lượng lactic acid nhiều nhất ở mẫu 3 (200 mL): m = 0,1 × (10,5 / 1000) × (200 / 5) × 90 = 3,78 gam.\n"
        "b) Đúng: Nếu thay bằng dung dịch NaOH 0,2 M thì V(dd NaOH 0,2 M) = 10,5 / 2 = 5,25 mL < 6 mL.\n"
        "c) Sai: Ở nhiệt độ quá cao (45 °C) lượng lactic acid thu được giảm xuống so với ở 37 °C nên giả thuyết ban đầu là sai.\n"
        "d) Đúng: Trong các mức nhiệt độ khảo sát, nhiệt độ tối ưu cho quá trình lên men tạo lactic acid là 37 °C."
    ),
    ("C2", 98): (
        "a) Sai: Phân tử cellulose cấu tạo mạch không phân nhánh từ các đơn vị β-glucose qua liên kết β-1,4-glycoside.\n"
        "b) Sai: [C6H7O2(OH)3]n + 3nHNO3 → [C6H7O2(ONO2)3]n + 3nH2O. Ta có V(dd HNO3) = (825 × 3 × 63 × 100) / (297 × 63 × 1,4 × 0,9) ≈ 661 L (khác 590 L).\n"
        "c) Đúng: Mỗi mắt xích C6H10O5 có 3 nhóm –OH tự do nên viết được dưới dạng [C6H7O2(OH)3]n.\n"
        "d) Sai: Cellulose không có phản ứng tạo màu xanh tím với iodine."
    ),
}


def setup_azota_doc(doc):
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)


def add_azota_options_p1(doc, opts, correct_ans):
    """
    Trình bày 4 phương án A, B, C, D đúng chuẩn ảnh mẫu Azota:
    - Nhãn A., B., C., D. in đậm, màu xanh dương (#0000FF).
    - Phương án đúng: gạch chân nhãn (ví dụ A. có gạch chân).
    - Nội dung phương án: chữ đen bình thường, không gạch chân.
    """
    a_txt = opts.get("A", "")
    b_txt = opts.get("B", "")
    c_txt = opts.get("C", "")
    d_txt = opts.get("D", "")
    max_len = max(len(a_txt), len(b_txt), len(c_txt), len(d_txt))

    def render_opt(p, letter, txt, prefix_tab=False):
        if prefix_tab:
            p.add_run("\t")
        is_corr = (letter == correct_ans)
        r_lbl = p.add_run(f"{letter}.")
        r_lbl.bold = True
        r_lbl.underline = is_corr
        r_lbl.font.name = "Times New Roman"
        r_lbl.font.size = Pt(12)
        r_lbl.font.color.rgb = COLOR_AZOTA_BLUE

        r_sp = p.add_run(" ")
        r_sp.font.name = "Times New Roman"
        r_sp.font.size = Pt(12)

        add_chem_runs(p, txt, bold=False, underline=False, color=None, size=12)

    if max_len <= 38:
        # 2 hàng x 2 cột giống hệt ảnh mẫu Azota
        for pair in [("A", "B"), ("C", "D")]:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(8.5), WD_TAB_ALIGNMENT.LEFT)
            render_opt(p, pair[0], opts.get(pair[0], ""), False)
            render_opt(p, pair[1], opts.get(pair[1], ""), True)
    else:
        # 4 hàng riêng biệt nếu phương án dài
        for letter in ["A", "B", "C", "D"]:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            render_opt(p, letter, opts.get(letter, ""), False)


def clean_solution_lines(sol_text):
    lines = []
    for raw in sol_text.split("\n"):
        ln = raw.strip()
        if not ln:
            continue
        # Bỏ dấu bullet đầu dòng nếu có
        ln = re.sub(r"^[•\-\+]\s*", "", ln).strip()
        # Bỏ dòng lặp "Đáp số ..." hoặc "Đáp án: ..." ở đầu nếu trùng
        if re.match(r"^(Đáp số|Đáp án)\s*[:.]?\s*[\d,\.a-dA-D]+", ln, flags=re.IGNORECASE):
            continue
        lines.append(ln)
    return lines


def generate_azota_docx(test_cfg, bank_map, assets_map):
    doc = docx.Document()
    setup_azota_doc(doc)
    de_num = test_cfg["de_num"]

    # ================= PHẦN I =================
    p_sec1 = doc.add_paragraph()
    p_sec1.paragraph_format.space_before = Pt(4)
    p_sec1.paragraph_format.space_after = Pt(4)
    r_s1 = p_sec1.add_run("PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn. ")
    r_s1.bold = True
    r_s1.font.name = "Times New Roman"
    r_s1.font.size = Pt(12)
    r_s1_sub = p_sec1.add_run(
        f"Thí sinh trả lời từ câu 1 đến câu {len(test_cfg['P1'])}. Mỗi câu hỏi thí sinh chỉ chọn một phương án."
    )
    r_s1_sub.font.name = "Times New Roman"
    r_s1_sub.font.size = Pt(12)

    for idx, (ch_code, orig_qnum) in enumerate(test_cfg["P1"], 1):
        q = copy.deepcopy(bank_map[(ch_code, orig_qnum)])
        if (ch_code, orig_qnum) in STEM_PATCHES:
            q["stem_lines"] = STEM_PATCHES[(ch_code, orig_qnum)]
        if (ch_code, orig_qnum) in EXTRA_OPT_PATCHES:
            q["opt_lines"] = EXTRA_OPT_PATCHES[(ch_code, orig_qnum)]
        if (ch_code, orig_qnum) in ANSWER_PATCHES:
            q["answer"] = ANSWER_PATCHES[(ch_code, orig_qnum)]
        if (ch_code, orig_qnum) in EXTRA_ANSWER_PATCHES:
            q["answer"] = EXTRA_ANSWER_PATCHES[(ch_code, orig_qnum)]
        if (ch_code, orig_qnum) in SOLUTION_PATCHES:
            q["solution"] = SOLUTION_PATCHES[(ch_code, orig_qnum)]
        if (ch_code, orig_qnum) in EXTRA_SOLUTION_PATCHES:
            q["solution"] = EXTRA_SOLUTION_PATCHES[(ch_code, orig_qnum)]

        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(3)
        r_qnum = p_q.add_run(f"Câu {idx}. ")
        r_qnum.bold = True
        r_qnum.font.name = "Times New Roman"
        r_qnum.font.size = Pt(12)
        r_qnum.font.color.rgb = COLOR_AZOTA_BLUE

        stem_lines = q["stem_lines"]
        c_int = int(ch_code[1])
        q_assets = assets_map.get(c_int, {}).get(orig_qnum, {"images": [], "tables": []})
        img_placed_early = False
        img_triggers = ("như sau:", "dưới đây:", "hình vẽ sau:", "hình sau:", "sơ đồ sau:")

        if stem_lines:
            add_chem_runs(p_q, stem_lines[0], size=12)
            if any(trig in stem_lines[0].lower() for trig in img_triggers) and q_assets["images"]:
                for img_path in q_assets["images"]:
                    if os.path.exists(img_path):
                        p_img = doc.add_paragraph()
                        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_img.paragraph_format.space_before = Pt(2)
                        p_img.paragraph_format.space_after = Pt(2)
                        p_img.add_run().add_picture(img_path, width=Cm(6.5))
                img_placed_early = True
            for extra_ln in stem_lines[1:]:
                p_ex = doc.add_paragraph()
                p_ex.paragraph_format.space_before = Pt(1)
                p_ex.paragraph_format.space_after = Pt(2)
                add_chem_runs(p_ex, extra_ln, size=12)

        if not img_placed_early:
            for img_path in q_assets["images"]:
                if os.path.exists(img_path):
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_before = Pt(2)
                    p_img.paragraph_format.space_after = Pt(2)
                    p_img.add_run().add_picture(img_path, width=Cm(6.5))
        for tbl_data in q_assets["tables"]:
            add_data_table_to_doc(doc, tbl_data)

        opts = parse_abcd_from_opt_lines(q["opt_lines"])
        add_azota_options_p1(doc, opts, q["answer"])

        # Lời giải Phần I theo mẫu Azota:
        # Lời giải:
        # Nội dung lời giải...
        p_lg = doc.add_paragraph()
        p_lg.paragraph_format.space_before = Pt(3)
        p_lg.paragraph_format.space_after = Pt(2)
        r_lg = p_lg.add_run("Lời giải:")
        r_lg.font.name = "Times New Roman"
        r_lg.font.size = Pt(12)

        sol_text = (q.get("solution") or "").strip()
        sol_lines = clean_solution_lines(sol_text)
        if not sol_lines:
            ans = q.get("answer", "")
            chosen_txt = opts.get(ans, "")
            sol_lines = [f"Chọn đáp án {ans} ({chosen_txt})."]

        for ln in sol_lines:
            p_ln = doc.add_paragraph()
            p_ln.paragraph_format.space_before = Pt(1)
            p_ln.paragraph_format.space_after = Pt(2)
            add_chem_runs(p_ln, ln, size=12)

    # ================= PHẦN II =================
    p_sec2 = doc.add_paragraph()
    p_sec2.paragraph_format.space_before = Pt(8)
    p_sec2.paragraph_format.space_after = Pt(4)
    r_s2 = p_sec2.add_run("PHẦN II. Câu trắc nghiệm đúng sai. ")
    r_s2.bold = True
    r_s2.font.name = "Times New Roman"
    r_s2.font.size = Pt(12)
    r_s2_sub = p_sec2.add_run(
        f"Thí sinh trả lời từ câu 1 đến câu {len(test_cfg['P2'])}. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai."
    )
    r_s2_sub.font.name = "Times New Roman"
    r_s2_sub.font.size = Pt(12)

    for idx, (ch_code, orig_qnum) in enumerate(test_cfg["P2"], 1):
        q = copy.deepcopy(bank_map[(ch_code, orig_qnum)])
        if (ch_code, orig_qnum) in STEM_PATCHES:
            q["stem_lines"] = STEM_PATCHES[(ch_code, orig_qnum)]
        if (ch_code, orig_qnum) in P2_ITEM_PATCHES:
            q["p2_items"] = P2_ITEM_PATCHES[(ch_code, orig_qnum)]
        if (ch_code, orig_qnum) in ANSWER_PATCHES:
            q["answer"] = ANSWER_PATCHES[(ch_code, orig_qnum)]
        if (ch_code, orig_qnum) in EXTRA_ANSWER_PATCHES:
            q["answer"] = EXTRA_ANSWER_PATCHES[(ch_code, orig_qnum)]
        if (ch_code, orig_qnum) in SOLUTION_PATCHES:
            q["solution"] = SOLUTION_PATCHES[(ch_code, orig_qnum)]
        if (ch_code, orig_qnum) in EXTRA_SOLUTION_PATCHES:
            q["solution"] = EXTRA_SOLUTION_PATCHES[(ch_code, orig_qnum)]

        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(3)
        r_qnum = p_q.add_run(f"Câu {idx}. ")
        r_qnum.bold = True
        r_qnum.font.name = "Times New Roman"
        r_qnum.font.size = Pt(12)
        r_qnum.font.color.rgb = COLOR_AZOTA_BLUE

        stem_lines = q["stem_lines"]
        if stem_lines:
            add_chem_runs(p_q, stem_lines[0], size=12)
            for extra_ln in stem_lines[1:]:
                p_ex = doc.add_paragraph()
                p_ex.paragraph_format.space_before = Pt(1)
                p_ex.paragraph_format.space_after = Pt(2)
                add_chem_runs(p_ex, extra_ln, size=12)

        c_int = int(ch_code[1])
        q_assets = assets_map.get(c_int, {}).get(orig_qnum, {"images": [], "tables": []})
        for img_path in q_assets["images"]:
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(2)
                p_img.paragraph_format.space_after = Pt(2)
                p_img.add_run().add_picture(img_path, width=Cm(8.5))
        for tbl_data in q_assets["tables"]:
            add_data_table_to_doc(doc, tbl_data)

        # Parse a:Đ - b:S - c:Đ - d:S
        ans_map = {}
        for part_ans in q["answer"].split("-"):
            part_ans = part_ans.strip()
            if ":" in part_ans:
                k_a, v_a = part_ans.split(":", 1)
                ans_map[k_a.strip()] = v_a.strip()

        for key in ["a", "b", "c", "d"]:
            stmt_txt = q.get("p2_items", {}).get(key, "")
            p_it = doc.add_paragraph()
            p_it.paragraph_format.space_before = Pt(2)
            p_it.paragraph_format.space_after = Pt(2)
            is_true = (ans_map.get(key) == "Đ")
            r_k = p_it.add_run(f"{key})")
            r_k.underline = is_true
            r_k.font.name = "Times New Roman"
            r_k.font.size = Pt(12)

            r_sp = p_it.add_run(" ")
            r_sp.font.name = "Times New Roman"
            r_sp.font.size = Pt(12)

            add_chem_runs(p_it, stmt_txt, size=12)

        # Lời giải Phần II theo mẫu Azota:
        # Lời giải
        # Nội dung giải thích...
        p_lg = doc.add_paragraph()
        p_lg.paragraph_format.space_before = Pt(3)
        p_lg.paragraph_format.space_after = Pt(2)
        r_lg = p_lg.add_run("Lời giải")
        r_lg.font.name = "Times New Roman"
        r_lg.font.size = Pt(12)

        sol_text = (q.get("solution") or "").strip()
        sol_lines = clean_solution_lines(sol_text)
        for ln in sol_lines:
            p_ln = doc.add_paragraph()
            p_ln.paragraph_format.space_before = Pt(1)
            p_ln.paragraph_format.space_after = Pt(2)
            add_chem_runs(p_ln, ln, size=12)

    # ================= PHẦN III =================
    p_sec3 = doc.add_paragraph()
    p_sec3.paragraph_format.space_before = Pt(8)
    p_sec3.paragraph_format.space_after = Pt(4)
    r_s3 = p_sec3.add_run("PHẦN III. Câu trắc nghiệm trả lời ngắn. ")
    r_s3.bold = True
    r_s3.font.name = "Times New Roman"
    r_s3.font.size = Pt(12)
    r_s3_sub = p_sec3.add_run(f"Thí sinh trả lời từ câu 1 đến câu {len(test_cfg['P3'])}.")
    r_s3_sub.font.name = "Times New Roman"
    r_s3_sub.font.size = Pt(12)

    for idx, (ch_code, orig_qnum) in enumerate(test_cfg["P3"], 1):
        q = copy.deepcopy(bank_map[(ch_code, orig_qnum)])
        if (ch_code, orig_qnum) in STEM_PATCHES:
            q["stem_lines"] = STEM_PATCHES[(ch_code, orig_qnum)]
        if (ch_code, orig_qnum) in ANSWER_PATCHES:
            q["answer"] = ANSWER_PATCHES[(ch_code, orig_qnum)]
        if (ch_code, orig_qnum) in EXTRA_ANSWER_PATCHES:
            q["answer"] = EXTRA_ANSWER_PATCHES[(ch_code, orig_qnum)]
        if (ch_code, orig_qnum) in SOLUTION_PATCHES:
            q["solution"] = SOLUTION_PATCHES[(ch_code, orig_qnum)]
        if (ch_code, orig_qnum) in EXTRA_SOLUTION_PATCHES:
            q["solution"] = EXTRA_SOLUTION_PATCHES[(ch_code, orig_qnum)]

        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(3)
        r_qnum = p_q.add_run(f"Câu {idx}. ")
        r_qnum.bold = True
        r_qnum.font.name = "Times New Roman"
        r_qnum.font.size = Pt(12)
        r_qnum.font.color.rgb = COLOR_AZOTA_BLUE

        stem_lines = q["stem_lines"]
        if stem_lines:
            add_chem_runs(p_q, stem_lines[0], size=12)
            for extra_ln in stem_lines[1:]:
                p_ex = doc.add_paragraph()
                p_ex.paragraph_format.space_before = Pt(1)
                p_ex.paragraph_format.space_after = Pt(2)
                add_chem_runs(p_ex, extra_ln, size=12)

        c_int = int(ch_code[1])
        q_assets = assets_map.get(c_int, {}).get(orig_qnum, {"images": [], "tables": []})
        for img_path in q_assets["images"]:
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(2)
                p_img.paragraph_format.space_after = Pt(2)
                p_img.add_run().add_picture(img_path, width=Cm(7.5))
        for tbl_data in q_assets["tables"]:
            add_data_table_to_doc(doc, tbl_data)

        # Lời giải Phần III theo mẫu Azota:
        # Lời giải
        # Đáp án: <số>
        # Giải thích:
        # <nội dung giải thích>
        p_lg = doc.add_paragraph()
        p_lg.paragraph_format.space_before = Pt(3)
        p_lg.paragraph_format.space_after = Pt(2)
        r_lg = p_lg.add_run("Lời giải")
        r_lg.font.name = "Times New Roman"
        r_lg.font.size = Pt(12)

        p_da = doc.add_paragraph()
        p_da.paragraph_format.space_before = Pt(1)
        p_da.paragraph_format.space_after = Pt(2)
        r_da = p_da.add_run(f"Đáp án: {q.get('answer', '')}")
        r_da.font.name = "Times New Roman"
        r_da.font.size = Pt(12)

        p_gt = doc.add_paragraph()
        p_gt.paragraph_format.space_before = Pt(1)
        p_gt.paragraph_format.space_after = Pt(2)
        r_gt = p_gt.add_run("Giải thích:")
        r_gt.font.name = "Times New Roman"
        r_gt.font.size = Pt(12)

        sol_text = (q.get("solution") or "").strip()
        sol_lines = clean_solution_lines(sol_text)
        for ln in sol_lines:
            p_ln = doc.add_paragraph()
            p_ln.paragraph_format.space_before = Pt(1)
            p_ln.paragraph_format.space_after = Pt(2)
            add_chem_runs(p_ln, ln, size=12)

    fname = f"De_Kiem_Tra_GK1_Hoa_12_De{de_num:02d}_Azota.docx"
    out_p1 = os.path.join(OUT_DIR, fname)
    out_p2 = os.path.join(BANK_DIR, fname)
    doc.save(out_p1)
    doc.save(out_p2)
    print(f"  + Saved Azota file: {out_p1}")


if __name__ == "__main__":
    with open(BANK_PATH, "r", encoding="utf-8") as f:
        bank_data = json.load(f)

    bank_map = {}
    for ch_code in ["C1", "C2", "C3"]:
        for q in bank_data[ch_code]["questions"]:
            bank_map[(ch_code, q["q_num"])] = q

    assets_map = {c: extract_chapter_assets(c) for c in [1, 2, 3]}

    for t_cfg in TESTS_SPEC:
        generate_azota_docx(t_cfg, bank_map, assets_map)
    print("\nHOÀN TẤT CHUYỂN ĐỔI ĐỊNH DẠNG AZOTA!")
